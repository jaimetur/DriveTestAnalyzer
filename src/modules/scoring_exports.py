"""Reference-style editable scoring matrices using the existing CDR master."""
from __future__ import annotations

from copy import deepcopy
import csv
import json
from io import BytesIO, StringIO
from math import ceil, floor, isclose, isfinite, log10
from pathlib import Path
from typing import Any, Callable
from textwrap import wrap

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

from src.modules.cdr_reporting import (
    _named_slide_layout, _remove_all_slides, _set_slide_header,
    _set_structural_slide_text,
)
from src.modules.column_names import campaign_sort_key, compact_campaign_value
from src.modules import report_structure
from src.modules.scoring import most_reliable_result
from src.modules.operator_roles import focus_operators
from src.modules.scoring_points_loss import map_background, map_boundaries
from src.modules.scoring_insight_slides import (
    TITLE_BLUE, add_area_summary_slides, add_campaign_comparison_slides, add_kpi_gap_profile_slides,
    add_location_card_slides, add_points_loss_slides, add_points_per_kpi_slides, add_trend_slides, scoring_scope_labels,
)
from src.modules.scoring_summary_slides import (
    add_campaign_scope_slides, add_gap_parts_slides, add_levers_slides, add_location_ranking_slides,
    add_location_versus_slides, add_points_at_stake_slides, add_position_matrix_slides, add_serving_loss_slides,
    add_sites_slides, add_standings_slides, add_stretch_slides, latest_campaign_items,
)
from src.modules.scoring_area_summary import build_area_summary
from src.modules.scoring_config import BEST_NETWORK_SCORING, MOST_RELIABLE_SCORING, SCORING_LABELS
from src.modules.scoring_insights import operator_styles
from src.modules.scoring_pptx_scorecard import apply_scorecard, style_legend_band
from src.modules.scoring_themes import DEFAULT_THEME, active_theme, themed, use_theme
from src.modules.theme_titles import style_title
from src.modules.scoring_act_slides import (
    add_act_plan_slides, add_owner_card_slides, add_place_ranking_slides, add_route_focus_slides, owner_names,
    owner_palette,
)
from src.modules.scoring_vendors import normalize_scoring_vendor_result
from src.modules.scoring_views import (
    THRESHOLD_COLORS, _hierarchy_display_value, gap_priority_order_key, latest_campaign_columns, build_scoring_views, refresh_baseline_warning, scoring_coverage_notes,
)
from src.modules.scoring_pptx_allocation import (
    _environment_display_label, add_maximum_allocation_donut, category_maximum_allocations,
    maximum_allocations_from_configuration, scaled_environment_allocations,
)

_FONT = 'Ericsson Hilda'
_SCALING_NOTICE_COLOR = '#FF8C1A'
_WHITE = '#FFFFFF'
_NEUTRAL = '#ECEFF1'
_GAP_SUMMARY_HEADER = '#455B65'
_GAP_SUMMARY_COMPARISON_HEADER = '#DDEBE6'
_GAP_SUMMARY_CATEGORY = '#E6F0F7'
_SCORING_EXPANDED_CATEGORY = '#DCE5E9'
_KPI_TYPE_COLORS = {'Reliable': '#D8EFCA', 'Diff': '#FFF2CC'}
_SUBTITLE_COLOR = '#245A96'


def _theme_rgb(color: str) -> RGBColor:
    """A colour written in the NetCheck palette, in the active theme (see scoring_themes.themed)."""
    return RGBColor.from_string(themed(color).lstrip('#'))
_LEGACY_CAMPAIGN_WARNING = 'Campaigns are scored separately; the supplied Tableau Prep flow pools campaigns.'
_SCOPE_FILTER_FIELDS = ('Operator', 'Vendor', 'Region', 'Cluster', 'City', 'Campaign')
_SCORING_SECTION_TITLES = {BEST_NETWORK_SCORING: 'Best Network Scoring',
                           MOST_RELIABLE_SCORING: 'Most Reliable Network Scoring'}


def _table_for_mode(matrix: dict, table_mode: str) -> dict:
    """Project saved KPI rows without changing chart inputs or weighted totals."""
    if table_mode not in {'expanded', 'summary'}:
        raise ValueError('Table mode must be expanded or summary.')
    projected = deepcopy(matrix)
    row_key = 'category_rows' if table_mode == 'summary' else 'expanded_rows'
    total_key = 'category_total' if table_mode == 'summary' else 'expanded_total'
    projected['rows'] = deepcopy(matrix.get(row_key, matrix.get('rows', [])))
    projected['total'] = deepcopy(matrix.get(total_key, matrix.get('total', {})))
    projected['table_mode'] = table_mode
    return projected


def export_scoring_csv(job: dict[str, Any], result: dict[str, Any], kind: str,
                       table_mode: str, operator_mapping_groups: list[dict[str, Any]] | None = None,
                       *, environment: str = 'all', vendor_mapping_groups: list[dict[str, Any]] | None = None) -> str:
    """Export the displayed table mode in a flat, explicit comparison format.

    Legacy CSV URLs without table_mode keep their original raw-result schema.
    Category scores are sums; category and final GAPs average available KPI gaps.
    """
    if kind not in {'scoring', 'gap'}:
        raise ValueError('CSV table kind must be scoring or gap.')
    views = build_scoring_views(job, result, operator_mapping_groups, vendor_mapping_groups=vendor_mapping_groups)
    matrix_key = 'hierarchy_score_tables' if kind == 'scoring' else 'hierarchy_gap_tables'
    matrices = views.get(matrix_key) or views.get('score_tables' if kind == 'scoring' else 'gap_summary_tables', [])
    _validate_environment(views, environment)
    if environment != 'all':
        matrices = [matrix for matrix in matrices if matrix.get('context', {}).get('environment') == environment]
    fields = ['table_type', 'row_type', 'environment', 'operator', 'vendor', 'region', 'cluster', 'city', 'campaign',
              'reference_operator', 'category', 'kpi_code', 'kpi', 'kpi_type', 'kpi_value',
              'score_points', 'max_points', 'score_weight_percent', 'gap_points', 'complete',
              'gap_partial', 'gap_environments']
    output = StringIO(newline='')
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writeheader()
    for original in matrices:
        matrix = _table_for_mode(original, table_mode)
        columns = {column['id']: column for column in matrix.get('hierarchy_columns', [])}
        operators = matrix['operators']
        if kind == 'gap':
            operators = [operator for operator in operators
                         if not columns.get(operator, {}).get('is_reference')
                         and operator != matrix['baseline_operator']]
        rows = matrix['rows'] + [{
            'row_type': 'total', 'category': 'Total', 'kpi': 'Weighted score / Average KPI GAP',
            'kpi_code': '', 'kpi_type': '', **matrix.get('total', {}),
        }]
        for row in rows:
            for operator in operators:
                column = columns.get(operator, {})
                context = {**matrix.get('context', {}), **column.get('context', {})}
                value = row.get('values', {}).get(operator, {})
                writer.writerow({
                    'table_type': kind, 'row_type': row.get('row_type', 'kpi'),
                    **{field: context.get(field) for field in ('environment', 'vendor', 'region', 'cluster', 'city', 'campaign')},
                    'operator': column.get('operator', operator), 'reference_operator': matrix['baseline_operator'],
                    **{field: row.get(field, '') for field in ('category', 'kpi_code', 'kpi', 'kpi_type')},
                    'kpi_value': value.get('value'), 'score_points': value.get('points'),
                    'max_points': row.get('max_points'), 'score_weight_percent': row.get('weight_percent'),
                    'gap_points': row.get('gaps', {}).get(operator), 'complete': value.get('complete', ''),
                    'gap_partial': row.get('gap_partial', {}).get(operator, False),
                    'gap_environments': json.dumps(row.get('gap_environments', {}).get(operator, [])),
                })
    return output.getvalue()


def _validate_environment(views: dict, environment: str) -> None:
    """Validate against the saved result, including custom methodology environments."""
    available = {matrix.get('context', {}).get('environment')
                 for key in ('score_tables', 'hierarchy_score_tables')
                 for matrix in views.get(key, [])}
    if environment != 'all' and environment not in available:
        raise ValueError('The selected environment is not available in this scoring job.')


def _export_environment_views(views: dict, environment: str) -> dict:
    """Put the existing aggregate first; never synthesize missing environment points."""
    _validate_environment(views, environment)
    projected = dict(views)
    for key in ('score_tables', 'gap_summary_tables', 'gap_tables',
                'hierarchy_score_tables', 'hierarchy_gap_tables'):
        matrices = views.get(key, [])
        if environment != 'all':
            matrices = [matrix for matrix in matrices
                        if matrix.get('context', {}).get('environment') == environment]
        matrices = sorted(matrices, key=lambda matrix: matrix.get('context', {}).get('environment') != 'Combined')
        projected[key] = []
        for matrix in matrices:
            if matrix.get('context', {}).get('environment') == 'Combined':
                matrix = {**matrix, 'context': {**matrix['context'], 'environment': 'All Environments'}}
            projected[key].append(matrix)
    return projected


def _number(value: Any) -> str:
    if value is None:
        return 'N/A'
    formatted = f'{float(value):.2f}'
    return '0.00' if formatted == '-0.00' else formatted


def _gap_number(row: dict, operator: str | None = None) -> str:
    """Mark comparisons based on only the common available environments."""
    value = row.get('gaps', {}).get(operator) if operator is not None else row.get('gap_points')
    partial = row.get('gap_partial', {})
    partial = partial.get(operator, False) if isinstance(partial, dict) else partial
    return _number(value) + ('*' if value is not None and partial else '')


def _score_total_label(rows: list[dict], score_ids: list[str], gap_ids: list[str], *, show_gap_values: bool) -> str:
    """Include the partial-coverage legend only when a displayed value is marked."""
    label = 'Weighted score / Average KPI GAP' if show_gap_values else 'Weighted score'
    partial = any(
        (value.get('points') is not None and not value.get('complete'))
        for row in rows for key in score_ids
        for value in [row.get('values', {}).get(key, {})]
    ) or any('*' in _gap_number(row, key) for row in rows for key in gap_ids)
    return label + ('; * incomplete' if partial else '')


def _scope(context: dict[str, Any], *, environment_label: bool = False) -> str:
    return ' · '.join(
        str(value) if key == 'environment' and not environment_label
        else f'{key.replace("_", " ").title()}: {value}'
        for key, value in context.items() if value is not None and str(value).strip()
    )


def _chart_subtitle(matrix: dict[str, Any]) -> str:
    environment = str(matrix.get('context', {}).get('environment') or '').strip()
    return matrix.get('environment_subtitle', _environment_display_label(environment))


def _scoring_chart_title(matrix: dict[str, Any], suffix: str) -> str:
    return f"{matrix.get('scoring_label') or SCORING_LABELS[BEST_NETWORK_SCORING]} Scoring {suffix}"


# Titles stay on one line: the template title is 40pt in Ericsson Hilda, where a character is about
# 0.41em wide (bold focus included), and its placeholder is 27.5 cm wide. Only a title wider than
# the placeholder at 40pt is made smaller, measured with a small margin over the font's width.
TITLE_FONT_PT = 40
TITLE_MIN_FONT_PT = 20
# The subtitle (the scoring, environment, area, operators and vendors) is 14pt bold, about 0.6em a character: a longer
# one is made smaller so it does not wrap over the content of the slide.
SUBTITLE_FONT_PT = 14
SUBTITLE_CHARACTER_EM = .6
SUBTITLE_MIN_FONT_PT = 8


def _fit_slide_title_on_one_line(title_shape, paragraph) -> None:
    """Make a title that would wrap smaller, so it fits its placeholder on one line (measured with the metrics of a
    real font, its bold focus included)."""
    from src.modules.slide_titles import fit_paragraph_on_one_line

    if fit_paragraph_on_one_line(title_shape, paragraph, TITLE_FONT_PT, TITLE_MIN_FONT_PT) is not None:
        title_shape.text_frame.word_wrap = False


def _subtitle_font_size(title_shape, paragraph) -> float:
    """The size of a slide subtitle that fits its placeholder on one line, at most 14pt."""
    text = paragraph.text.strip()
    if not text or title_shape.width is None:
        return SUBTITLE_FONT_PT
    text_frame = title_shape.text_frame
    insets = (text_frame.margin_left if text_frame.margin_left is not None else Emu(72000)) + (text_frame.margin_right or 0)
    width_pt = (int(title_shape.width) - int(insets)) / 12700
    size = width_pt / (len(text) * SUBTITLE_CHARACTER_EM)
    return SUBTITLE_FONT_PT if size >= SUBTITLE_FONT_PT else max(SUBTITLE_MIN_FONT_PT, floor(size * 2) / 2)


def _slide(presentation, title: str, subtitle: str):
    layout = _named_slide_layout(presentation, 'Title Only')
    if layout is None:
        raise ValueError("The PowerPoint template needs a 'Title Only' layout.")
    slide = presentation.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    # On the white background, the dark logos of the master instead of the white ones of the layout.
    from src.modules.template_logos import use_dark_logos
    use_dark_logos(slide)
    _set_slide_header(slide, title, subtitle)
    title_shape = slide.shapes.title
    if title_shape is not None:
        paragraphs = title_shape.text_frame.paragraphs
        if paragraphs:
            _emphasize_title(paragraphs[0])
        # The title and subtitle in the style of the theme (see theme_titles), the subtitle on one line.
        sizes = [_subtitle_font_size(title_shape, paragraph) for paragraph in paragraphs[1:]]
        style_title(slide, subtitle_size=min(sizes) if sizes else None)
        if paragraphs:
            _fit_slide_title_on_one_line(title_shape, paragraphs[0])
    return slide


def _emphasize_title(paragraph) -> None:
    """The subject of a title in the title colour of the theme and, in bold and in its accent, what tells it from the
    slides next to it: the part after " — " ("GAP Analysis — 3 vs EE", "Scoring Tables — Summary") or the focus
    ("Best Network Scoring per City")."""
    text = paragraph.text
    if ' — ' in text:
        heading, separator, focus = text.partition(' — ')
        parts = ((heading + separator, False), (focus, True))
    else:
        heading, separator, focus = text.partition(' per ')
        if not separator:
            return
        parts = ((heading + ' ', False), ('per ' + focus, True))
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    paragraph.clear()
    for text, bold in parts:
        run = paragraph.add_run()
        if properties is not None:
            run._r.insert(0, deepcopy(properties))
        run.text = text
        run.font.bold = bold
        run.font.color.rgb = (RGBColor.from_string(active_theme()['title_accent'].lstrip('#')) if bold
                              else _theme_rgb(TITLE_BLUE))


def _text(slide, text: str, top: float, *, left: float = .55, width: float = 12.2,
          height: float = .4, size: float = 10, color: str = '#263746', align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    box.text = text
    for paragraph in box.text_frame.paragraphs:
        paragraph.font.name = _FONT
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = _theme_rgb(color)
        paragraph.alignment = align
    return box


def _add_individual_gap_notes(slide, matrix: dict, rows: list[dict]) -> None:
    """Show the KPI mean and the mean of valid column totals beside the table."""
    if matrix.get('hierarchy_columns'):
        comparisons = [
            (row.get('gaps', {}).get(column['id']),
             row.get('gap_partial', {}).get(column['id'], False))
            for row in rows for column in matrix['hierarchy_columns']
        ]
    else:
        comparisons = [(row.get('gap_points'), row.get('gap_partial', False)) for row in rows]
    valid = [(float(value), partial) for value, partial in comparisons
             if value is not None and isfinite(float(value))]
    summed_gap = sum(value for value, _ in valid) if valid else None
    average = summed_gap / len(valid) if valid else None
    if matrix.get('hierarchy_columns'):
        valid_column_count = sum(
            any(value is not None and isfinite(float(value))
                for row in rows
                for value in [row.get('gaps', {}).get(column['id'])])
            for column in matrix['hierarchy_columns']
        )
    else:
        valid_column_count = 1 if valid else 0
    total = summed_gap / valid_column_count if valid_column_count else None
    marker = '*' if any(partial for _, partial in valid) else ''
    value_text = _number(average) + marker
    total_text = _number(total) + marker
    positive_name, negative_name = active_theme()['gap_names']
    columns = matrix.get('hierarchy_columns') or []
    order = ('KPIs ordered by GAP in the latest campaign' if len(latest_campaign_columns(columns)) < len(columns)
             else 'KPIs ordered by GAP')
    box = _text(
        slide, f'{order}\n\nAverage KPI GAP: {value_text} points\n\n'
        f'Average total KPI GAP: {total_text} points\n\n'
        f'Operator − reference\n{positive_name}: positive\n{negative_name}: negative\n\n{matrix["note"]}',
        1.8, left=10.2, width=2.5, height=4.65, size=13,
    )
    box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, label, value, text in (
        (2, 'Average KPI GAP: ', average, value_text),
        (4, 'Average total KPI GAP: ', total, total_text),
    ):
        paragraph = box.text_frame.paragraphs[index]
        paragraph.clear()
        paragraph.add_run().text = label
        value_run = paragraph.add_run()
        value_run.text = text
        value_run.font.bold = True
        value_run.font.color.rgb = _theme_rgb(
            '228B22' if value is not None and value > 0 else
            'CC2424' if value is not None and value < 0 else '263746'
        )
        paragraph.add_run().text = ' points'


def _add_gap_priority_arrow(slide, table, *, header_rows: int = 1) -> None:
    """Span the table from just below the Priority label to its bottom edge."""
    frame = table._graphic_frame
    label = _text(slide, 'Priority', frame.top.inches, left=frame.left.inches - .59,
                  width=.52, height=.28, size=7, color='#245A96', align=PP_ALIGN.CENTER)
    label.name = 'GAP KPI Priority Label'
    label.fill.solid()
    label.fill.fore_color.rgb = RGBColor(255, 255, 255)
    label.line.color.rgb = _theme_rgb(_SUBTITLE_COLOR)
    label.line.width = Pt(1)
    label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    label.text_frame.paragraphs[0].font.name = 'Arial'
    label.text_frame.paragraphs[0].font.bold = True
    body_top = label.top + label.height + Inches(.06)
    body_bottom = frame.top + frame.height
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, frame.left - Inches(.38), body_top,
        Inches(.12), max(Inches(.2), body_bottom - body_top),
    )
    arrow.name = 'GAP KPI Priority Arrow'
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = _theme_rgb(_SUBTITLE_COLOR)
    arrow.line.fill.background()
    arrow._element.spPr.append(OxmlElement('a:effectLst'))
    for reference in arrow._element.xpath('./p:style/a:effectRef'):
        reference.set('idx', '0')


def _add_gap_color_scale(slide, matrix: dict, *, left: float, width: float, top: float = 6.8) -> None:
    """Add an editable red-white-green bar with this comparison's actual GAP range, on one line under the table: its
    heading, the range behind the reference, the bar (zero in its middle) and the range ahead."""
    colors = matrix.get('gap_scale_colors') or {
        'negative': '#E57373', 'zero': '#FFFFFF', 'positive': '#70AD47',
    }
    maximum = float(matrix.get('gap_scale_max') or 0.0)

    def blend(start: str, end: str, amount: float) -> str:
        channels = [
            round(int(start[index:index + 2], 16) * (1 - amount)
                  + int(end[index:index + 2], 16) * amount)
            for index in (1, 3, 5)
        ]
        return '#' + ''.join(f'{channel:02X}' for channel in channels)

    heading = _text(slide, 'GAP color scale', top, left=left, width=width * .15, size=7.5, height=.14)
    heading.text_frame.paragraphs[0].font.bold = True
    bar_left, bar_width, bar_top, bar_height = left + width * .32, width * .36, top + .01, .12
    segments = 24
    for index in range(segments):
        midpoint = (index + .5) / segments
        if midpoint <= .5:
            color = blend(colors['negative'], colors['zero'], midpoint * 2)
        else:
            color = blend(colors['zero'], colors['positive'], (midpoint - .5) * 2)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(bar_left + bar_width * index / segments), Inches(bar_top),
            Inches(bar_width / segments + .005), Inches(bar_height),
        )
        shape.name = f'GAP Color Scale Segment {index + 1}'
        shape.fill.solid()
        shape.fill.fore_color.rgb = _theme_rgb(color)
        shape.line.fill.background()
    _text(slide, f'−{maximum:.2f} (operator trails)', top + .01, left=left + width * .15,
          width=width * .16, size=6.5, height=.13, align=PP_ALIGN.RIGHT)
    _text(slide, '0', top + .01, left=bar_left + bar_width * .45,
          width=bar_width * .1, size=6.5, height=.13, align=PP_ALIGN.CENTER)
    _text(slide, f'+{maximum:.2f} (operator leads)', top + .01, left=left + width * .69,
          width=width * .16, size=6.5, height=.13)


def _cell(cell, text: str, *, color: str = _WHITE, foreground: str = '#17232D',
          size: float = 9, bold: bool = False, left: bool = False) -> None:
    if str(text) == 'N/A':
        text = 'N/A*'
    if str(text) == 'N/A*' or (str(text).endswith('*') and str(text)[:-1].replace('.', '').replace(',', '').lstrip('+-').isdigit()):
        foreground = '#C62828'
        bold = True
    cell.text = text
    cell.margin_left = cell.margin_right = Inches(.025)
    cell.margin_top = cell.margin_bottom = 0
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = _theme_rgb(color)
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.name = _FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = _theme_rgb(foreground)
        paragraph.alignment = PP_ALIGN.LEFT if left else PP_ALIGN.CENTER
        paragraph.line_spacing = 1.0
        paragraph.space_before = paragraph.space_after = 0
    properties = cell._tc.get_or_add_tcPr()
    border_namespace = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    edges = ('lnL', 'lnR', 'lnT', 'lnB')
    # The cells are set apart by the lines of the theme (white gaps in most themes), not drawn on a grid.
    theme = active_theme()
    border_width = str(int(float(theme.get('cell_border_pt', 1.5)) * 12700))
    border_color = str(theme.get('cell_border', '#FFFFFF')).lstrip('#').upper()
    for edge_index, edge in enumerate(edges):
        existing = properties.find(f'{border_namespace}{edge}')
        if existing is not None:
            properties.remove(existing)
        line = OxmlElement(f'a:{edge}')
        line.set('w', border_width)
        fill = OxmlElement('a:solidFill')
        rgb = OxmlElement('a:srgbClr')
        rgb.set('val', border_color)
        fill.append(rgb)
        line.append(fill)
        properties.insert(edge_index, line)


def _operator_color(matrix: dict, operator: str) -> str:
    color = matrix.get('operator_styles', {}).get(operator, {}).get('color')
    if color:
        return color
    for column in matrix.get('hierarchy_columns', []):
        if column.get('operator') == operator and column.get('color'):
            return column['color']
    return '#365F91'


def _operator_label(matrix: dict, operator: str) -> str:
    return matrix.get('operator_styles', {}).get(operator, {}).get('label', operator)


def _campaigns_for_export(job: dict[str, Any], result: dict[str, Any]) -> list[str]:
    """Prefer campaign values in scored contexts over the selected-source union."""
    campaigns: list[str] = []

    def collect(value: Any) -> None:
        if value is None:
            return
        # Shown as everywhere else (2026-Q2-SA); spellings of one campaign are listed once.
        label = compact_campaign_value(str(value).strip())
        if label and label.casefold() not in {item.casefold() for item in campaigns}:
            campaigns.append(label)

    for key in ('scoring', 'score_rows', 'totals', 'charts'):
        rows = result.get(key)
        if not isinstance(rows, list):
            continue
        for row in rows:
            if not isinstance(row, dict):
                continue
            campaign = next((value for field, value in row.items()
                             if str(field).casefold() in {'campaign', 'period', 'quarter'} and value is not None), None)
            collect(campaign)
    if campaigns:
        return sorted(campaigns, key=campaign_sort_key)

    source_metadata = job.get('source_metadata', [])
    if isinstance(source_metadata, dict):
        source_metadata = source_metadata.get('sources', [])
    if isinstance(source_metadata, list):
        for source in source_metadata:
            if not isinstance(source, dict):
                continue
            source_campaigns = source.get('campaigns', [])
            for campaign in source_campaigns if isinstance(source_campaigns, list) else [source_campaigns]:
                collect(campaign)
    if not campaigns:
        fallback = job.get('campaigns', [])
        for campaign in fallback if isinstance(fallback, list) else [fallback]:
            collect(campaign)
    return sorted(campaigns, key=campaign_sort_key)


def _filter_value(value: Any) -> str:
    if isinstance(value, (list, tuple, set)):
        return ', '.join(str(item).strip() for item in value if item is not None and str(item).strip())
    return str(value).strip() if value is not None else ''


def _scope_filter_fields(job: dict[str, Any]) -> tuple[str, ...]:
    configured = job.get('aggregation_hierarchy')
    if not isinstance(configured, list):
        configuration = job.get('configuration')
        configured = configuration.get('aggregation_hierarchy') if isinstance(configuration, dict) else None
    if not isinstance(configured, list):
        return _SCOPE_FILTER_FIELDS
    canonical = {field.casefold(): field for field in _SCOPE_FILTER_FIELDS}
    fields = []
    for value in configured:
        field = canonical.get(str(value).strip().casefold())
        if field and field not in fields:
            fields.append(field)
    fields.extend(field for field in _SCOPE_FILTER_FIELDS if field not in fields)
    return tuple(fields)


def _canonical_scope_filter_labels(payload: Any, fields: tuple[str, ...] = _SCOPE_FILTER_FIELDS) -> list[str] | None:
    if not isinstance(payload, dict) or any(not isinstance(value, list) for value in payload.values()):
        return None
    normalized = {str(key).strip().casefold(): value for key, value in payload.items()}
    if payload and not any(field.casefold() in normalized for field in fields):
        return None
    return [
        f'{field}: {_filter_value(normalized.get(field.casefold(), [])) or "All"}'
        for field in fields
    ]


def _context_filter_labels(job: dict[str, Any]) -> list[str]:
    labels: list[str] = []
    seen: set[str] = set()
    reserved = {'period', 'quarter', 'nr mode', 'aggregation', 'aggregation level', 'split by', 'baseline'}

    def append(label: Any, value: Any, operator: Any = None) -> None:
        name = str(label or '').strip().replace('_', ' ')
        rendered_value = _filter_value(value)
        if not name or not rendered_value or name.casefold() in reserved:
            return
        formatted_name = ' '.join(
            part.upper() if part.casefold() == 'ran' else part.capitalize() for part in name.split()
        )
        description = f'{formatted_name} {operator} {rendered_value}' if operator else f'{formatted_name}: {rendered_value}'
        identity = description.casefold()
        if identity not in seen:
            seen.add(identity)
            labels.append(description)

    for payload_name in ('context_filters', 'filters'):
        payload = job.get(payload_name)
        if isinstance(payload, dict):
            canonical_labels = (
                _canonical_scope_filter_labels(payload, _scope_filter_fields(job))
                if payload_name == 'context_filters' else None
            )
            if canonical_labels is not None:
                for label in canonical_labels:
                    append(*label.split(': ', 1))
                continue
            if any(key in payload for key in ('column', 'field', 'name', 'key')):
                field = payload.get('column') or payload.get('field') or payload.get('name') or payload.get('key')
                append(field, payload.get('value', payload.get('values')), payload.get('operator'))
            else:
                for name, value in payload.items():
                    if isinstance(value, dict):
                        append(name, value.get('value', value.get('values')), value.get('operator'))
                    else:
                        append(name, value)
        elif isinstance(payload, list):
            for item in payload:
                if isinstance(item, dict):
                    field = item.get('column') or item.get('field') or item.get('name') or item.get('key')
                    append(field, item.get('value', item.get('values')), item.get('operator'))
                elif isinstance(item, str) and item.strip() and item.casefold() not in seen:
                    seen.add(item.casefold())
                    labels.append(item.strip())
        elif isinstance(payload, str) and payload.strip() and payload.casefold() not in seen:
            seen.add(payload.casefold())
            labels.append(payload.strip())
    return labels


def _scoring_filter_subtitle(job: dict[str, Any], environment: str | None = None) -> str:
    nr_mode = str(job.get('nr_mode') or 'NSA').strip()
    mode_label = {'nsa': 'Non-Standalone', 'sa': 'Standalone'}.get(nr_mode.casefold(), nr_mode)
    levels = job.get('aggregation_levels') or job.get('levels') or []
    if isinstance(levels, str):
        levels = [levels]
    level_labels = [str(level).strip() for level in levels if str(level).strip()] if isinstance(levels, list) else []
    plan = job.get('_scoring_display_selections')
    if plan is None:
        from src.modules.cdr_report_filenames import scoring_display_selections
        plan = scoring_display_selections(job.get('context_filters'), aggregation_values=level_labels,
                                          nr_mode_label=nr_mode)
    def display(field):
        selected = ', '.join(plan['values'][field])
        return selected + (', ...' if selected else '...') if plan['omitted'][field] else selected
    return '\n'.join([mode_label, *[
        f"{'Split by' if field == 'aggregation' else field.title()}: {display(field)}"
        for field in ('aggregation', 'operator', 'vendor', 'region', 'city')
    ]])


def prepare_scoring_display_selections(job: dict[str, Any], result: dict[str, Any],
                                       template_path: Path) -> dict:
    """Compute shared presentation and download-name selections without changing filters."""
    from src.modules.cdr_report_filenames import scoring_display_selections
    return scoring_display_selections(
        job.get('context_filters'),
        campaign_values=_campaigns_for_export(job, result),
        aggregation_values=job.get('aggregation_levels') or job.get('levels'),
        nr_mode_label=job.get('nr_mode') or 'NSA',
    )


def _style_scoring_intro_field(paragraph, text: str, *, bold_label: bool = False) -> None:
    """Show field labels in white and their values in bold yellow."""
    label, separator, value = text.partition(':')
    paragraph.clear()
    label_run = paragraph.add_run()
    label_run.text = label + separator + (' ' if value.strip() else '')
    label_run.font.color.rgb = RGBColor.from_string('FFFFFF')
    label_run.font.bold = bold_label
    if value.strip():
        value_run = paragraph.add_run()
        value_run.text = value.strip()
        value_run.font.color.rgb = RGBColor.from_string('FFC700')
        value_run.font.bold = True


def _fit_scoring_intro_subtitle(presentation, slide, layout_name: str, subtitle: str, campaign_text: str) -> None:
    """The title (its first line), the NR Mode as its subtitle and, under the line of the layout, the Campaigns and the
    Split by and filters of a scoring cover or transition; report_structure gives the title and subtitle their places
    and sizes, as on every cover and transition."""
    title_shape = next(
        (shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3}),
        None,
    )
    if title_shape is None:
        return
    subtitle_shape = next(
        (shape for shape in slide.placeholders if shape.placeholder_format.type == 4),
        None,
    )
    if subtitle_shape is None:
        # Transition layouts may omit the subtitle; clone the template's real placeholder.
        title_layout = _named_slide_layout(presentation, 'Title Page')
        source = next((shape for shape in title_layout.placeholders
                       if shape.placeholder_format.type == 4), None) if title_layout else None
        if source is None:
            raise ValueError("The PowerPoint template needs a subtitle placeholder for scoring exports.")
        slide.shapes.clone_placeholder(source)
        subtitle_shape = next(shape for shape in slide.placeholders
                              if shape.placeholder_format.type == 4)

    title_text = title_shape.text_frame.text.split('\x0b', 1)[0].split('\n', 1)[0]
    title_shape.text_frame.text = title_text
    lines = subtitle.splitlines()
    subtitle_shape.text_frame.text = lines[0] if lines else ''
    mode = subtitle_shape.text_frame.paragraphs[0]
    mode.alignment = PP_ALIGN.LEFT
    properties = mode._p.get_or_add_pPr()
    properties.set('marL', '0')
    properties.set('indent', '0')
    for bullet in properties.xpath('./a:buNone'):
        properties.remove(bullet)
    properties.insert_element_before(OxmlElement('a:buNone'), 'a:tabLst', 'a:defRPr', 'a:extLst')

    # The Campaigns, then the Split by and each filter, under the line of the layout as the blocks of every cover.
    details = [campaign_text, *lines[1:]]
    filters_shape = _text(slide, '', 5.92, left=.52, width=12.29, height=.21 * len(details) + .1, size=12, color=_WHITE)
    filters_shape.name = 'Scoring Aggregations and Filters'
    frame = filters_shape.text_frame
    frame.clear()
    # Each line fits one line: its values are cut to the room of the template (prepare_scoring_display_selections).
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, line in enumerate(details):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(12 if index == 0 else 11)
        properties = paragraph._p.get_or_add_pPr()
        properties.set('marL', '0')
        properties.set('indent', '0')
        properties.insert(0, OxmlElement('a:buNone'))
        _style_scoring_intro_field(paragraph, line, bold_label=True)
        paragraph.line_spacing = 1.0
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(1)


def _set_shape_geometry(shape, height: int) -> None:
    """Materialize inherited placeholder geometry before resizing its height."""
    left, top, width = shape.left, shape.top, shape.width
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height


def _scoring_environment_title(environment: str, configuration: dict) -> str:
    scope = configuration.get('scope', {}).get('environments', {}).get(environment, {})
    if scope.get('display_name'):
        return str(scope['display_name'])
    selectors = scope.get('source_filters') or {}
    filters = [str(selectors.get(field) or scope.get(key) or '').strip()
               for field, key in (('G_Level_1', 'g_level_1'), ('G_Level_2', 'g_level_2'))]
    filters = [value for value in filters if value]
    if filters:
        filters = ['Connecting Roads' if value.casefold() in {'connectionroad', 'connectingroads'}
                   else value for value in filters]
        return ' - '.join(filters)
    return _environment_display_label(environment)


SECTION_TAG_PREFIX = 'section:'


def _tag_slides(presentation, start: int, section: str) -> None:
    """Name the slides from ``start`` that have no section yet after ``section`` (the Word version keeps the
    slides of a section together on its pages)."""
    for slide in list(presentation.slides)[start:]:
        if not str(slide._element.cSld.get('name') or '').startswith(SECTION_TAG_PREFIX):
            slide._element.cSld.set('name', f'{SECTION_TAG_PREFIX}{section}')


def _add_scoring_intro_slides(
    presentation, job: dict[str, Any], result: dict[str, Any], *, environment: str | None = None,
    scoring: str | None = None, title: str | None = None, subtitle: str | None = None, level: int | None = None,
) -> None:
    """The cover of the export (without ``level``) or a transition of ``level``: a scoring or one of its environments."""
    configuration = job.get('configuration') or result.get('configuration') or {}
    if title is not None:
        pass
    elif environment:
        title = _scoring_block_title(scoring, _scoring_environment_title(environment, configuration))
    else:
        title = _SCORING_SECTION_TITLES[scoring] if scoring else 'Scoring & GAP Analysis'
    subtitle = _scoring_filter_subtitle(job, environment) if subtitle is None else subtitle
    campaigns = _campaigns_for_export(job, result)
    plan = job.get('_scoring_display_selections')
    campaign_value = ', '.join(plan['values']['campaign']) if plan else ', '.join(campaigns) or 'All Campaigns'
    if plan and plan['omitted']['campaign']:
        campaign_value += ', ...' if campaign_value else '...'
    campaign_text = f'Campaigns: {campaign_value}'
    layout_name = 'Title Page'
    layout = _named_slide_layout(presentation, layout_name)
    if layout is None:
        raise ValueError(f"The PowerPoint template needs a '{layout_name}' layout for scoring exports.")
    slide = presentation.slides.add_slide(layout)
    _tag_slides(presentation, len(presentation.slides) - 1, 'cover')
    _set_structural_slide_text(slide, title, subtitle)
    _fit_scoring_intro_subtitle(presentation, slide, layout_name, subtitle, campaign_text)
    if level is not None:
        report_structure.mark(slide, level)
    # Combined scores scaled because environments have no results are announced in orange, above the line.
    notices = [str(notice) for notice in result.get('notices') or []]
    if notices and (environment is None or str(environment).casefold() in {'combined', 'all', 'all environments'}):
        notice_shape = _text(slide, ' '.join(notices), 4.75, left=.52, width=12.29, height=.6,
                             size=12, color=_SCALING_NOTICE_COLOR)
        notice_shape.name = 'Scoring Scaling Notice'
        for paragraph in notice_shape.text_frame.paragraphs:
            paragraph.font.bold = True


def _add_scenario_transition(presentation, name: str, job: dict[str, Any]) -> None:
    """The transition of a scenario of a report: its name, with its NR Mode and the filters that choose part of the
    country or of the operators (none: the NR Mode alone)."""
    layout = _named_slide_layout(presentation, 'Transition') or _named_slide_layout(presentation, 'Title Page')
    slide = presentation.slides.add_slide(layout)
    _tag_slides(presentation, len(presentation.slides) - 1, 'cover')
    lines = _scoring_filter_subtitle(job).splitlines()
    chosen = [line for line in lines[1:] if not line.startswith('Split by') and ': ' in line
              and line.split(': ', 1)[1].strip() and not line.split(': ', 1)[1].startswith('All')]
    _set_structural_slide_text(slide, name, ' · '.join([lines[0], *chosen]) if lines else '')
    report_structure.mark(slide, 1)


def _header_foreground(color: str) -> str:
    red, green, blue = (int(themed(color).lstrip('#')[index:index + 2], 16) for index in (0, 2, 4))
    return '#17232D' if .299 * red + .587 * green + .114 * blue > 155 else _WHITE


def _operator_pages(table: dict) -> list[list[str]]:
    operators = table['operators']
    baseline = table['baseline_operator']
    if len(operators) <= 5:
        return [operators]
    if baseline in operators:
        compared = [operator for operator in operators if operator != baseline]
        return [[operator for operator in operators if operator in compared[start:start + 4] or operator == baseline]
                for start in range(0, len(compared), 4)]
    return [operators[start:start + 5] for start in range(0, len(operators), 5)]


def _scalar_matrix_as_hierarchy(matrix: dict) -> dict:
    converted = deepcopy(matrix)
    converted['hierarchy_levels'] = ['Operator']
    converted['hierarchy_columns'] = [
        {'id': operator, 'operator': operator, 'path': [{'level': 'Operator', 'value': _operator_label(matrix, operator)}],
         'color': _operator_color(matrix, operator), 'is_reference': operator == matrix['baseline_operator']}
        for operator in matrix['operators']
    ]
    return converted


def _kpi_type_color(metric: dict) -> str:
    """Keep the same KPI classification colors across Score and GAP tables."""
    return _KPI_TYPE_COLORS.get(str(metric.get('kpi_type', '')), _NEUTRAL)


_SUMMARY_BEST_COLOR = '#C6EFCE'
_SUMMARY_WORST_COLOR = '#FFC7CE'


def _add_summary_operator_legend(slide, table_right: float) -> None:
    """Explain the best and worst score highlights above Summary tables on the right."""
    for index, (label, color) in enumerate((('Best operator', _SUMMARY_BEST_COLOR),
                                            ('Worst operator', _SUMMARY_WORST_COLOR))):
        table = slide.shapes.add_table(
            1, 1, Inches(table_right - 2.65 + index * 1.35), Inches(1.25), Inches(1.3), Inches(.18),
        ).table
        _cell(table.cell(0, 0), label, color=color, size=8)


def _summary_extreme_operators(values: dict[str, dict]) -> tuple[set[str], set[str]]:
    """Return tied best and worst operators when the values establish a ranking."""
    available = {}
    for operator, value in values.items():
        points = value.get('points') if isinstance(value, dict) else value
        try:
            points = float(points)
        except (TypeError, ValueError):
            continue
        if isfinite(points):
            available[operator] = points
    if len(available) < 2:
        return set(), set()
    best = max(available.values())
    worst = min(available.values())
    if isclose(best, worst, rel_tol=1e-9, abs_tol=1e-9):
        return set(), set()
    return (
        {operator for operator, points in available.items()
         if isclose(points, best, rel_tol=1e-9, abs_tol=1e-9)},
        {operator for operator, points in available.items()
         if isclose(points, worst, rel_tol=1e-9, abs_tol=1e-9)},
    )


def _summary_extreme_columns(values: dict[str, dict], columns: list[dict]) -> tuple[set[str], set[str]]:
    """Rank all score leaves across the entire Summary row, preserving ties."""
    return _summary_extreme_operators({column['id']: values.get(column['id'], {}) for column in columns})


def _score_tables(presentation, matrices: list[dict], legend: list[dict], *, gap_layout: str = 'end',
                  show_gap_values: bool = True, title: str = 'Scoring Tables — Breakdown',
                  show_kpi_values: bool = False) -> None:
    if gap_layout == 'adjacent' or show_kpi_values:
        matrices = [_scalar_matrix_as_hierarchy(matrix) for matrix in matrices]
        _hierarchy_score_tables(presentation, matrices, legend, gap_layout=gap_layout,
                                show_gap_values=show_gap_values, title=title, show_kpi_values=show_kpi_values)
        return

    for matrix in matrices:
        metric_pages = [matrix['rows']]
        is_summary = matrix.get('table_mode') == 'summary'
        fixed_count = 4 if is_summary else 5
        weight_index = 2 if is_summary else 3
        maximum_index = weight_index + 1
        for operators in [matrix['operators']]:
            baseline = matrix['baseline_operator']
            comparisons = ([name for name in operators if name != baseline]
                           if show_gap_values and baseline in matrix['operators'] else [])
            for page_index, metrics in enumerate(metric_pages):
                total_label = _score_total_label(
                    [*metrics, matrix['total']], operators, comparisons, show_gap_values=show_gap_values,
                )
                page_label = f' · KPIs {page_index + 1}/{len(metric_pages)}' if len(metric_pages) > 1 else ''
                slide = _slide(presentation, title, _chart_subtitle(matrix) + page_label)
                include_total = page_index == len(metric_pages) - 1
                headers = ['CATEGORY', 'CATEGORY' if is_summary else 'KPI',
                           *([] if is_summary else ['Type of KPI']), 'Score weight\n(%)', 'Max score',
                           *[_operator_label(matrix, name) for name in operators],
                           *[f'GAP {_operator_label(matrix, name)}\n− {_operator_label(matrix, baseline)}' for name in comparisons]]
                count = len(metrics) + 1 + int(include_total)
                table = slide.shapes.add_table(count, len(headers), Inches(.55), Inches(1.55),
                                               presentation.slide_width - Inches(1.1), Inches(5.2)).table
                widths = [1.22, 2.60, *([] if is_summary else [.75]), .68, .79]
                remaining = (presentation.slide_width / 914400 - 1.1 - sum(widths)) / max(1, len(headers) - fixed_count)
                for index, column in enumerate(table.columns):
                    column.width = Inches(widths[index] if index < fixed_count else remaining)
                table.rows[0].height = Inches(.46)
                metric_font, body_heights = _hierarchy_metric_layout(
                    metrics, total_label,
                    available_height=5.2 - .46, kpi_width=widths[1],
                )
                for row, height in zip(list(table.rows)[1:], body_heights):
                    row.height = Inches(height)
                numeric_texts = [
                    _number(value.get('points')) + ('*' if value.get('points') is not None and not value.get('complete') else '')
                    for metric in metrics + [matrix['total']] for value in metric.get('values', {}).values()
                ] + [_gap_number(metric, operator) for metric in metrics + [matrix['total']] for operator in comparisons]
                numeric_font = _hierarchy_content_font(
                    numeric_texts, remaining, maximum=min(11, min(body_heights) * 72 * .82),
                )
                colors = [_GAP_SUMMARY_HEADER, _GAP_SUMMARY_HEADER,
                          *([] if is_summary else [_GAP_SUMMARY_HEADER]), '#5B9A44', '#C9483F',
                          *[_operator_color(matrix, name) for name in operators], *['#F6D55C'] * len(comparisons)]
                for index, header in enumerate(headers):
                    _cell(table.cell(0, index), header, color=colors[index],
                          foreground=_header_foreground(colors[index]), size=8.5, bold=True)
                category_color = _GAP_SUMMARY_CATEGORY if is_summary else _SCORING_EXPANDED_CATEGORY
                for row_index, metric in enumerate(metrics, 1):
                    subtotal = metric.get('row_type') == 'category'
                    row_color = ('#E3E6E7' if is_summary else category_color) if subtotal else None
                    best, worst = (_summary_extreme_operators(metric.get('values', {}))
                                   if matrix.get('table_mode') == 'summary' and subtotal else (set(), set()))
                    _cell(table.cell(row_index, 0), metric['category'], color=category_color,
                          size=metric_font, left=True, bold=True)
                    _cell(table.cell(row_index, 1), metric['kpi'], color=row_color or '#E7E8E9',
                          bold=True, left=True, size=metric_font)
                    if not is_summary:
                        _cell(table.cell(row_index, 2), metric.get('kpi_type', ''),
                              color=row_color or _kpi_type_color(metric), size=metric_font, bold=subtotal)
                    _cell(table.cell(row_index, weight_index), _number(metric['weight_percent']) + '%', color=row_color or _WHITE,
                          size=metric_font, bold=subtotal)
                    _cell(table.cell(row_index, maximum_index), _number(metric['max_points']), color=row_color or _WHITE,
                          size=metric_font, bold=subtotal)
                    for offset, operator in enumerate(operators, fixed_count):
                        value = metric['values'][operator]
                        partial = value['points'] is not None and not value['complete']
                        score_color = (_SUMMARY_BEST_COLOR if operator in best else
                                       _SUMMARY_WORST_COLOR if operator in worst else
                                       row_color or value.get('color', _WHITE if value['complete'] else _NEUTRAL))
                        _cell(table.cell(row_index, offset), _number(value['points']) + ('*' if partial else ''),
                              color=score_color,
                              size=numeric_font, bold=subtotal)
                    for offset, operator in enumerate(comparisons, fixed_count + len(operators)):
                        _cell(table.cell(row_index, offset), _gap_number(metric, operator),
                              color=row_color or metric.get('gap_colors', {}).get(operator, '#FFF0D8'),
                              size=numeric_font, bold=True)
                first = 0
                while first < len(metrics):
                    last = first
                    while last + 1 < len(metrics) and metrics[last + 1]['category'] == metrics[first]['category']:
                        last += 1
                    if last > first:
                        cell = table.cell(first + 1, 0)
                        cell.merge(table.cell(last + 1, 0))
                        _cell(cell, metrics[first]['category'], color=category_color,
                              size=metric_font, left=True, bold=True)
                    first = last + 1
                if include_total:
                    total = matrix['total']
                    index = count - 1
                    total_color = '#D8DFE4' if is_summary else category_color
                    best, worst = (_summary_extreme_operators(total.get('values', {}))
                                   if is_summary else (set(), set()))
                    _cell(table.cell(index, 0), 'TOTAL', color=total_color, bold=True, left=True, size=metric_font)
                    _cell(table.cell(index, 1), total_label, color=total_color, size=metric_font, left=True)
                    if not is_summary:
                        _cell(table.cell(index, 2), '', color=total_color, size=metric_font)
                    for column, value in ((weight_index, _number(total['weight_percent']) + '%'),
                                          (maximum_index, _number(total['max_points']))):
                        _cell(table.cell(index, column), value, color=total_color, bold=True, size=numeric_font)
                    for column, operator in enumerate(operators, fixed_count):
                        value = total['values'][operator]
                        score_color = (_SUMMARY_BEST_COLOR if operator in best else
                                       _SUMMARY_WORST_COLOR if operator in worst else total_color)
                        _cell(table.cell(index, column), _number(value['points']) + ('*' if value['points'] is not None and not value['complete'] else ''),
                              color=score_color, bold=True, size=numeric_font)
                    for column, operator in enumerate(comparisons, fixed_count + len(operators)):
                        _cell(table.cell(index, column), _gap_number(total, operator), color=total_color, bold=True, size=numeric_font)
                apply_scorecard(
                    table, header_rows=1, metrics=metrics, total=matrix['total'] if include_total else None,
                    total_label=total_label, font=_FONT, label_width=2.6 if is_summary else 3.2,
                    value_columns=[*[(column, 'Score', operator) for column, operator in enumerate(operators, fixed_count)],
                                   *[(column, 'GAP', operator) for column, operator
                                     in enumerate(comparisons, fixed_count + len(operators))]],
                )
                if is_summary:
                    _add_summary_operator_legend(slide, presentation.slide_width.inches - .55)
                else:
                    for index, item in enumerate(legend):
                        legend_table = slide.shapes.add_table(1, 1, Inches(.55 + index * 1.15), Inches(7.02), Inches(1.1), Inches(.18)).table
                        _cell(legend_table.cell(0, 0), item.get('band', ''), color=item['color'], size=8)
                        style_legend_band(legend_table.cell(0, 0), item.get('band', ''), 8, _FONT)
                _text(slide, f'GAP = operator − reference; ±{matrix.get("gap_scale_max", 0):.2f} points; '
                      f'{active_theme()["gap_names"][0].lower()} + / {active_theme()["gap_names"][1].lower()} −',
                      7.03, left=6.5, width=6.2, size=8.5, height=.18)
                _text(slide, matrix['coverage_note'], 7.27, size=8, height=.18)


def _format_chart(chart, *, maximum: float, labels=XL_DATA_LABEL_POSITION.OUTSIDE_END) -> None:
    # Axis identifiers are unsigned integers in DrawingML, including cross references.
    axis_references = chart._chartSpace.xpath('.//c:axId | .//c:crossAx')
    axis_ids = {value: str(index + 1) for index, value in enumerate(
        dict.fromkeys(reference.get('val') for reference in axis_references),
    )}
    for axis_reference in axis_references:
        axis_reference.set('val', axis_ids[axis_reference.get('val')])
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.name = 'Arial'
    chart.legend.font.size = Pt(11)
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name = 'Arial'
        axis.tick_labels.font.size = Pt(10)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = max(1, maximum) * 1.12
    chart.plots[0].has_data_labels = True
    data_labels = chart.plots[0].data_labels
    data_labels.position = labels
    data_labels.number_format = '0.0'
    data_labels.font.name = 'Arial'
    data_labels.font.size = Pt(9)
    # Single-level categories need no native hierarchy label dividers.
    for flag in chart.category_axis._element.xpath('./c:noMultiLvlLbl'):
        flag.set('val', '1' if chart.plots[0].categories.depth <= 1 else '0')
    for gridline in chart.category_axis._element.xpath('./c:majorGridlines | ./c:minorGridlines'):
        gridline.getparent().remove(gridline)
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.minor_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.format.line.fill.background()
    if labels == XL_DATA_LABEL_POSITION.CENTER:
        data_labels.font.color.rgb = RGBColor.from_string('FFFFFF')


def _hide_chart_legend_entries(chart, indexes: set[int]) -> None:
    """Hide selected native series legend entries without changing the chart data."""
    if not indexes:
        return
    legend = chart.legend._element
    position = 1 if len(legend) and legend[0].tag.endswith('}legendPos') else 0
    for index in sorted(indexes):
        entry = OxmlElement('c:legendEntry')
        entry_index = OxmlElement('c:idx')
        entry_index.set('val', str(index))
        delete = OxmlElement('c:delete')
        delete.set('val', '1')
        entry.extend((entry_index, delete))
        legend.insert(position, entry)
        position += 1


def _fill_series(series, color: str) -> None:
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = _theme_rgb(color)
    series.format.line.color.rgb = RGBColor.from_string('FFFFFF')
    series.format.line.width = Pt(.5)


# A Breakdown chart keeps every category on one slide up to this number of bars; above it, the
# categories are spread over as few slides as possible, each with a similar number of bars.
BREAKDOWN_BARS_PER_SLIDE = 50


def _breakdown_pages(categories: list, bars_per_category: int, split: bool = True) -> list[list]:
    """The categories of each Breakdown slide: all on one slide unless they need more than
    BREAKDOWN_BARS_PER_SLIDE bars (and splitting is allowed)."""
    bars_per_category = max(1, bars_per_category)
    if not split or len(categories) * bars_per_category <= BREAKDOWN_BARS_PER_SLIDE:
        return [list(categories)] if categories else []
    per_slide = max(1, BREAKDOWN_BARS_PER_SLIDE // bars_per_category)
    slides = ceil(len(categories) / per_slide)
    size = ceil(len(categories) / slides)
    return [list(categories[start:start + size]) for start in range(0, len(categories), size)]


def _charts(presentation, matrices: list[dict]) -> None:
    for matrix in matrices:
        categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
        chunks = _breakdown_pages(categories, len(matrix['operators']), matrix.get('_split_charts') is not False)
        if len(chunks) > 1:
            for index, chunk in enumerate(chunks, 1):
                page = dict(matrix, rows=[row for row in matrix['rows'] if row['category'] in chunk],
                            _stacked_chart_page=f'Page {index} of {len(chunks)}')
                _charts(presentation, [page])
            continue
        slide = _slide(presentation, _scoring_chart_title(matrix, 'per Category (Breakdown)'), _chart_subtitle(matrix))
        _stacked_chart_page_note(slide, matrix)
        data = CategoryChartData()
        data.categories = categories
        maximum = 0.0
        for operator in matrix['operators']:
            values = []
            for category in categories:
                cells = [row['values'][operator] for row in matrix['rows'] if row['category'] == category]
                available = [cell['points'] for cell in cells if cell['points'] is not None]
                value = sum(available) if available else None
                values.append(value)
                maximum = max(maximum, value or 0)
            data.add_series(_operator_label(matrix, operator), values)
        chart_width = presentation.slide_width - Inches(1.4)
        chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(1.65),
                                             chart_width, Inches(4.65), data)
        chart = chart_shape.chart
        _format_chart(chart, maximum=maximum)
        chart.plots[0].gap_width = 140
        chart.plots[0].overlap = -20
        bar_count = len(categories) * len(matrix['operators'])
        chart.plots[0].data_labels.font.size = Pt(min(7, max(4, 140 / max(1, bar_count))))
        chart.plots[0].data_labels.number_format = '0' if bar_count > 40 else '0.0'
        chart.plots[0].has_data_labels = bar_count <= 70
        chart.legend.font.size = Pt(10)
        for series, operator in zip(chart.series, matrix['operators']):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(_operator_color(matrix, operator).lstrip('#'))
        slide.shapes.add_group_shape([chart_shape]).name = 'Scoring Category Comparison Chart'
        _text(slide, matrix['coverage_note'], 6.95, size=9)


def _add_family_allocation_donut(slide, matrix: dict) -> None:
    voice_categories = {'CLASSIC CALLS', 'WHATSAPP CALLS', 'MULTI RAB'}
    maximums = [
        sum(row.get('max_points') or 0 for row in matrix['rows']
            if (row.get('source_kind') in {'voice', 'speech'}
                if row.get('source_kind') else row['category'] in voice_categories) == is_voice)
        for is_voice in (True, False)
    ]
    allocations = matrix.get('environment_allocations') or {
        str(matrix['context'].get('environment') or 'All Environments'): {
            'voice': maximums[0], 'data': maximums[1],
        },
    }
    return add_maximum_allocation_donut(
        slide, matrix, allocations, left=9.7, top=1.75, width=3.0, height=5.0,
    )


def _add_category_allocation_donut(slide, matrix: dict) -> None:
    voice_categories = {'CLASSIC CALLS', 'WHATSAPP CALLS', 'MULTI RAB'}
    maximums = [
        sum(row.get('max_points') or 0 for row in matrix['rows']
            if (row.get('source_kind') in {'voice', 'speech'}
                if row.get('source_kind') else row['category'] in voice_categories) == is_voice)
        for is_voice in (True, False)
    ]
    allocations = matrix.get('environment_allocations') or {
        str(matrix['context'].get('environment') or 'All Environments'): {
            'voice': maximums[0], 'data': maximums[1],
        },
    }
    return add_maximum_allocation_donut(
        slide, matrix, allocations, left=9.7, top=1.75, width=3.0, height=5.0,
        category_allocations=category_maximum_allocations(matrix),
    )


def _stacked_category_chart(presentation, matrix: dict) -> None:
    """Compare KPI category contributions as one stacked bar per operator."""
    pages = _stacked_chart_pages(matrix)
    if len(pages) > 1:
        for page in pages:
            _stacked_category_chart(presentation, page)
        return
    slide = _slide(presentation, _scoring_chart_title(matrix, 'per Category'), _chart_subtitle(matrix))
    categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
    operators = list(matrix['operators'])
    _stacked_chart_page_note(slide, matrix)
    data = CategoryChartData()
    data.categories = [_operator_label(matrix, operator) for operator in operators]
    series_specs = []
    maximum = 0.0
    for operator_index, operator in enumerate(operators):
        for category_index, category in enumerate(categories):
            cells = [row['values'][operator] for row in matrix['rows'] if row['category'] == category]
            available = [cell['points'] for cell in cells if cell['points'] is not None]
            value = sum(available) if available else None
            values = [value if index == operator_index else None for index in range(len(operators))]
            label = _operator_label(matrix, operator) if category_index == 0 else f'{operator} · {category}'
            data.add_series(label, values)
            series_specs.append((operator, category_index, category_index > 0))
    totals = [matrix['total']['values'][operator]['points'] for operator in operators]
    data.add_series('TOTAL SCORE', totals)
    chart_width = Inches(8.9)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Inches(.55), Inches(1.65), chart_width, Inches(4.65), data,
    )
    chart = chart_shape.chart
    maximum = max(
        (sum(row['values'][operator]['points'] or 0 for row in matrix['rows']
             if row['values'][operator]['points'] is not None) for operator in operators),
        default=0.0,
    )
    _format_chart(chart, maximum=matrix.get('_stacked_chart_maximum', maximum), labels=XL_DATA_LABEL_POSITION.CENTER)
    chart.legend.font.size = Pt(9)
    hidden_legend_entries = set()
    for series_index, (operator, category_index, hide_from_legend) in enumerate(series_specs):
        series = chart.series[series_index]
        _fill_series(
            series, _hierarchy_chart_color(_operator_color(matrix, operator), category_index, len(categories)),
        )
        if hide_from_legend:
            hidden_legend_entries.add(series_index)
    _hide_chart_legend_entries(chart, hidden_legend_entries | {len(series_specs)})
    _format_stacked_segment_labels(chart)
    _add_total_labels(chart)
    category_key = _add_category_shade_bar(
        slide, categories, left=Inches(.55), top=Inches(6.38), width=chart_width, height=Inches(.4),
    )
    hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, [
        {'id': operator, 'operator': _operator_label(matrix, operator)} for operator in operators
    ])
    chart_group = [chart_shape] + ([category_key] if category_key is not None else []) + ([hierarchy_grid] if hierarchy_grid is not None else [])
    slide.shapes.add_group_shape(chart_group + chart_shape._scoring_axis_labels).name = 'Scoring Stacked Operator Chart'
    _add_category_allocation_donut(slide, matrix)
    _text(slide, matrix['coverage_note'], 6.95, size=9)


def _stacked_chart_pages(matrix: dict, *, limit: int = 20) -> list[dict]:
    """Keep preferred operators together and preserve hierarchy order within each page."""
    columns = matrix.get('hierarchy_columns') or [
        {'id': operator, 'operator': operator} for operator in matrix['operators']
    ]
    if matrix.get('_split_charts') is False or len(columns) <= limit:
        return [matrix]
    grouped = {}
    for column in columns:
        grouped.setdefault(column['operator'], []).append(column)
    preferred = {'3', 'three', 'threeuk', 'ee', 'vfuk', 'vodafoneuk'}
    groups = sorted(grouped.items(), key=lambda group: (
        ''.join(character for character in _operator_label(matrix, group[0]).lower() if character.isalnum()) not in preferred
        and ''.join(character for character in str(group[0]).lower() if character.isalnum()) not in preferred,
        list(grouped).index(group[0]),
    ))
    pages = []
    current = []
    for _, group in groups:
        for start in range(0, len(group), limit):
            chunk = group[start:start + limit]
            if current and len(current) + len(chunk) > limit:
                pages.append(current)
                current = []
            current.extend(chunk)
    if current:
        pages.append(current)
    projected = []
    positions = {column['id']: index for index, column in enumerate(columns)}
    for index, page in enumerate(pages, 1):
        page = sorted(page, key=lambda column: positions[column['id']])
        item = dict(matrix)
        item['operators'] = list(dict.fromkeys(column['operator'] for column in page))
        if matrix.get('hierarchy_columns'):
            item['hierarchy_columns'] = page
        item['_stacked_chart_page'] = f'Page {index} of {len(pages)}'
        item['_stacked_chart_maximum'] = max(
            (value['points'] for value in matrix['total']['values'].values() if value['points'] is not None),
            default=matrix['total']['max_points'],
        )
        projected.append(item)
    return projected


def _stacked_chart_page_note(slide, matrix: dict) -> None:
    if matrix.get('_stacked_chart_page'):
        _text(slide, matrix['_stacked_chart_page'], 1.42, left=.65, width=8.65, height=.2, size=9)


def _add_chart_hierarchy_grid(slide, chart_shape, columns: list[dict]) -> Any:
    """Align an editable aggregation grid with the chart's explicit plot area."""
    if not columns:
        return None
    if not columns[0].get('path'):
        columns = [dict(column, path=[{'level': 'Operator', 'value': column['operator']}])
                   for column in columns]
    depth = len(columns[0]['path'])
    chart = chart_shape.chart
    chart_shape._scoring_axis_labels = []
    if depth == 1:
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NEXT_TO_AXIS
        return None
    chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    chart.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    plot = chart._chartSpace.chart.plotArea
    layout = plot.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}layout')
    if layout is None:
        layout = OxmlElement('c:layout')
        plot.insert(0, layout)
    for child in list(layout):
        layout.remove(child)
    narrow_leaves = chart_shape.width.inches * .89 / len(columns) < .65
    row_heights = [.30 if narrow_leaves and index == depth - 1 else .22 for index in range(depth)]
    grid_cells = {}
    for depth_index in range(depth):
        start = 0
        while start < len(columns):
            prefix = columns[start]['path'][:depth_index + 1]
            end = start
            while end + 1 < len(columns) and columns[end + 1]['path'][:depth_index + 1] == prefix:
                end += 1
            value = str(_hierarchy_display_value(columns[start]['path'][depth_index]))
            cell_width = (chart_shape.width.inches - .75) * (end - start + 1) / len(columns)
            longest_line = max(len(line) for line in value.split('\n'))
            font_size = min(8, max(4, (cell_width - .02) * 72 / max(1, longest_line * .52)))
            grid_cells[depth_index, start] = (end, value, font_size)
            start = end + 1
        level_font = min(font for (level, _), (_, _, font) in grid_cells.items() if level == depth_index)
        for (level, start), (end, value, _) in list(grid_cells.items()):
            if level != depth_index:
                continue
            cell_width = (chart_shape.width.inches - .75) * (end - start + 1) / len(columns)
            characters = max(1, floor((cell_width - .02) * 72 / (level_font * .56)))
            lines = [part for line in value.split('\n') for part in (wrap(line, characters) or [''])]
            row_heights[depth_index] = max(row_heights[depth_index], len(lines) * level_font * 1.2 / 72 + .04)
            grid_cells[depth_index, start] = (end, '\n'.join(lines), level_font)
    grid_height = sum(row_heights)
    chart_shape.height -= Inches(grid_height + .1)
    plot_left = .55 / chart_shape.width.inches
    plot_width = 1 - .75 / chart_shape.width.inches
    manual = OxmlElement('c:manualLayout')
    for name, value in (('layoutTarget', 'inner'), ('xMode', 'edge'), ('yMode', 'edge'),
                        ('wMode', 'factor'), ('hMode', 'factor'),
                        ('x', str(plot_left)), ('y', '.13'), ('w', str(plot_width)), ('h', '.87')):
        element = OxmlElement('c:' + name)
        element.set('val', value)
        manual.append(element)
    layout.append(manual)
    left = chart_shape.left + Inches(.55)
    width = chart_shape.width - Inches(.75)
    top = chart_shape.top + chart_shape.height
    maximum = chart.value_axis.maximum_scale
    tick_base = 10 ** floor(log10(maximum / 6))
    tick_step = min((1, 2, 5, 10), key=lambda step: abs(step * tick_base - maximum / 6)) * tick_base
    chart.value_axis.major_unit = tick_step
    axis_labels = []
    plot_top = chart_shape.top.inches + chart_shape.height.inches * .13
    plot_height = chart_shape.height.inches * .87
    for index in range(floor(maximum / tick_step) + 1):
        value = index * tick_step
        label = _text(slide, f'{value:g}', plot_top + plot_height * (1 - value / maximum) - .08,
                      left=left / 914400 - .48, width=.40, height=.18, size=10,
                      color='#17232D', align=PP_ALIGN.RIGHT)
        label.text_frame.margin_left = label.text_frame.margin_right = 0
        label.name = 'Scoring Chart Value Axis Tick'
        axis_labels.append(label)
    chart_shape._scoring_axis_labels = axis_labels
    grid = slide.shapes.add_table(depth, len(columns), left, top, width, Inches(grid_height))
    grid.name = 'Scoring Chart Aggregation Grid'
    row_top = top / 914400
    for row_index, depth_index in enumerate(reversed(range(depth))):
        row = grid.table.rows[row_index]
        row.height = Inches(row_heights[depth_index])
        start = 0
        while start < len(columns):
            end, value, grid_font = grid_cells[depth_index, start]
            cell = grid.table.cell(row_index, start)
            if end > start:
                cell.merge(grid.table.cell(row_index, end))
            _cell(cell, value,
                  color='#EDF3F7' if row_index % 2 == 0 else '#E1EBF1',
                  foreground='#344858', size=grid_font, bold=depth_index == 0)
            cell.margin_left = cell.margin_right = Inches(.01)
            cell.text_frame.word_wrap = False
            start = end + 1
        level = columns[0]['path'][depth_index]['level']
        label = _text(slide, 'Category (0)' if level == 'Category' else level, row_top,
                      left=chart_shape.left.inches, width=.51,
                      height=row_heights[depth_index], size=7, color='#344858', align=PP_ALIGN.RIGHT)
        label.name = 'Scoring Chart Aggregation Level ' + level
        label.text_frame.margin_left = label.text_frame.margin_right = 0
        label.text_frame.word_wrap = False
        label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        row_top += row_heights[depth_index]
    return grid


def _format_stacked_segment_labels(chart) -> None:
    """Use contrasting labels and omit numbers that cannot fit thin segments."""
    chart.plots[0].gap_width = 45
    leaf_count = len(chart.series[-1].values)
    chart.plots[0].data_labels.font.size = Pt(min(9, max(4, 108 / max(1, leaf_count))))
    plot = chart._chartSpace.xpath('.//c:barChart')[0]
    defaults = plot.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}dLbls')
    for series in list(chart.series)[:-1]:
        labels = deepcopy(defaults)
        color = _header_foreground(str(series.format.fill.fore_color.rgb)).lstrip('#')
        for properties in labels.xpath('.//a:defRPr | .//a:rPr'):
            for fill in properties.xpath('./a:solidFill'):
                properties.remove(fill)
            fill = OxmlElement('a:solidFill')
            rgb = OxmlElement('a:srgbClr')
            rgb.set('val', color)
            fill.append(rgb)
            properties.insert(0, fill)
        for index, value in enumerate(series.values):
            if value is not None and value < chart.value_axis.maximum_scale * .04:
                point = OxmlElement('c:dLbl')
                point_index = OxmlElement('c:idx')
                point_index.set('val', str(index))
                delete = OxmlElement('c:delete')
                delete.set('val', '1')
                point.extend((point_index, delete))
                labels.insert(0, point)
        category = series._element.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}cat')
        series._element.insert(series._element.index(category), labels)


def _best_network(presentation, matrices: list[dict]) -> None:
    voice_categories = {'CLASSIC CALLS', 'WHATSAPP CALLS', 'MULTI RAB'}
    for matrix in [page for source in matrices for page in _stacked_chart_pages(source)]:
        slide = _slide(presentation, _scoring_chart_title(matrix, 'per Service'), _chart_subtitle(matrix))
        _stacked_chart_page_note(slide, matrix)
        hierarchy_columns = matrix.get('hierarchy_columns', [])
        if hierarchy_columns:
            chart_columns = hierarchy_columns
            operators = list(dict.fromkeys(column['operator'] for column in chart_columns))
        else:
            operators = list(matrix['operators'])
            chart_columns = [
                {'id': operator, 'operator': operator}
                for operator in operators
            ]
        data = CategoryChartData()
        if hierarchy_columns:
            _add_hierarchy_chart_categories(data, hierarchy_columns)
        else:
            data.categories = [_operator_label(matrix, operator) for operator in operators]
        maximums = []
        family_rows = {}
        for family in ('Voice', 'Data'):
            metrics = [row for row in matrix['rows']
                       if (row.get('source_kind') in {'voice', 'speech'}
                           if row.get('source_kind') else row['category'] in voice_categories) == (family == 'Voice')]
            family_rows[family] = metrics
            maximums.append(sum(row['max_points'] or 0 for row in metrics))

        voice_series_count = len(operators)
        for family in ('Data', 'Voice'):
            for operator in operators:
                values = []
                for column in chart_columns:
                    if column['operator'] != operator:
                        values.append(None)
                        continue
                    available = [row['values'][column['id']]['points'] for row in family_rows[family]
                                 if row['values'][column['id']]['points'] is not None]
                    values.append(sum(available) if available else None)
                label = _operator_label(matrix, operator)
                name = f'Voice · {label}' if family == 'Voice' else label
                data.add_series(name, values)
        totals = [matrix['total']['values'][column['id']]['points'] for column in chart_columns]
        data.add_series('Total', totals)
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_STACKED, Inches(.65), Inches(1.75),
            Inches(8.65), Inches(4.25), data,
        )
        chart = chart_shape.chart
        _format_chart(chart, maximum=matrix.get('_stacked_chart_maximum')
                      or max((value for value in totals if value is not None), default=0)
                      or matrix['total']['max_points'], labels=XL_DATA_LABEL_POSITION.CENTER)
        chart.legend.font.size = Pt(8.5)
        if hierarchy_columns:
            chart.category_axis.tick_labels.font.size = Pt(9)
        legend_hidden = set(range(voice_series_count, 2 * voice_series_count)) | {2 * voice_series_count}
        _hide_chart_legend_entries(chart, legend_hidden)
        for operator_index, operator in enumerate(operators):
            color = _operator_color(matrix, operator).lstrip('#')
            channels = [int(color[index:index + 2], 16) for index in (0, 2, 4)]
            voice_color = '#{:02X}{:02X}{:02X}'.format(
                *(round(channel + (255 - channel) * .45) for channel in channels)
            )
            _fill_series(chart.series[operator_index], color)
            _fill_series(chart.series[voice_series_count + operator_index], voice_color)
        _format_stacked_segment_labels(chart)
        _add_total_labels(chart)
        service_heading = _text(slide, 'Service types', 6.22, left=.65, width=8.65,
                                height=.22, size=9, color='#4A5B65')
        service_heading.name = 'Scoring Chart Service Key Heading'
        service_heading.text_frame.paragraphs[0].font.bold = True
        service_key = _add_category_shade_bar(
            slide, ['Data', 'Voice'], left=Inches(.65), top=Inches(6.48),
            width=Inches(8.65), height=Inches(.35),
        )
        service_key.name = 'Scoring Chart Service Shade Key'
        for index, color in enumerate(('555555', 'C4C4C4')):
            cell = service_key.table.cell(0, index)
            cell.fill.fore_color.rgb = RGBColor.from_string(color)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)
                paragraph.font.color.rgb = RGBColor.from_string(
                    _header_foreground(color).lstrip('#'))
        grid_columns = chart_columns if hierarchy_columns else [
            dict(column, operator=_operator_label(matrix, column['operator'])) for column in chart_columns
        ]
        hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, grid_columns)
        slide.shapes.add_group_shape(
            [chart_shape, service_heading, service_key] + ([hierarchy_grid] if hierarchy_grid is not None else [])
            + chart_shape._scoring_axis_labels,
        ).name = 'Scoring Best Network Service Chart'
        _add_family_allocation_donut(slide, matrix)
        _text(slide, matrix['coverage_note'], 7.03, size=9, height=.35)


def _add_total_labels(chart) -> None:
    """Overlay an invisible native line series with labels at each stacked total."""
    plot_area = chart._chartSpace.chart.plotArea
    bar_plot = plot_area.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}barChart')
    series = bar_plot.findall('{http://schemas.openxmlformats.org/drawingml/2006/chart}ser')[-1]
    bar_plot.remove(series)
    for element in list(series):
        if element.tag.endswith('}invertIfNegative'):
            series.remove(element)
    shape_properties = series.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr')
    if shape_properties is not None:
        series.remove(shape_properties)
    shape_properties = OxmlElement('c:spPr')
    line = OxmlElement('a:ln')
    line.append(OxmlElement('a:noFill'))
    shape_properties.append(line)
    series.insert(3, shape_properties)
    marker = series.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}marker')
    if marker is not None:
        series.remove(marker)
    marker = OxmlElement('c:marker')
    symbol = OxmlElement('c:symbol')
    symbol.set('val', 'none')
    marker.append(symbol)
    series.insert(4, marker)
    line_plot = OxmlElement('c:lineChart')
    grouping = OxmlElement('c:grouping')
    grouping.set('val', 'standard')
    line_plot.extend([grouping, series])
    labels = deepcopy(bar_plot.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}dLbls'))
    labels.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}dLblPos').set('val', 't')
    for run_properties in labels.xpath('.//a:defRPr | .//a:rPr'):
        run_properties.set('b', '1')
        label_size = min(12, max(4, 144 / max(1, len(chart.series[-1].values))))
        run_properties.set('sz', str(round(label_size * 100)))
        for fill in run_properties.xpath('./a:solidFill'):
            run_properties.remove(fill)
        fill = OxmlElement('a:solidFill')
        color = OxmlElement('a:srgbClr')
        color.set('val', '000000')
        fill.append(color)
        run_properties.insert(0, fill)
    line_plot.append(labels)
    for axis_id in bar_plot.findall('{http://schemas.openxmlformats.org/drawingml/2006/chart}axId'):
        line_plot.append(deepcopy(axis_id))
    plot_area.insert(plot_area.index(bar_plot) + 1, line_plot)


def _hierarchy_chart(presentation, matrix: dict) -> None:
    pages = _stacked_chart_pages(matrix)
    if len(pages) > 1:
        for page in pages:
            _hierarchy_chart(presentation, page)
        return
    columns = matrix.get('hierarchy_columns', [])
    slide = _slide(presentation, _scoring_chart_title(matrix, 'per Category'), _chart_subtitle(matrix))
    _stacked_chart_page_note(slide, matrix)
    data = CategoryChartData()
    _add_hierarchy_chart_categories(data, columns)
    categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
    operators = list(dict.fromkeys(column['operator'] for column in columns))
    category_rows = {
        category: [row for row in matrix['rows'] if row['category'] == category]
        for category in categories
    }
    series_specs = []
    for operator in operators:
        for category_index, category in enumerate(categories):
            values = []
            for column in columns:
                if column['operator'] != operator:
                    values.append(None)
                    continue
                cells = [row['values'][column['id']] for row in category_rows[category]]
                available = [cell['points'] for cell in cells if cell['points'] is not None]
                values.append(sum(available) if available else None)
            operator_label = _operator_label(matrix, operator)
            name = operator_label if category_index == 0 else f'{operator_label} · {category}'
            data.add_series(name, values)
            series_specs.append((operator, category_index, category_index > 0))

    data.add_series('TOTAL SCORE', [matrix['total']['values'][column['id']]['points'] for column in columns])

    maximum = max((sum(row['values'][column['id']]['points'] or 0
                       for row in matrix['rows']
                       if row['values'][column['id']]['points'] is not None)
                   for column in columns), default=0.0)
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Inches(.55), Inches(1.92), Inches(8.9), Inches(4.48), data,
    )
    chart = chart_shape.chart
    _format_chart(chart, maximum=matrix.get('_stacked_chart_maximum', maximum), labels=XL_DATA_LABEL_POSITION.CENTER)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(8)
    chart.category_axis.tick_labels.font.size = Pt(9)
    hidden_legend_entries = set()
    for series_index, (operator, category_index, hide_from_legend) in enumerate(series_specs):
        series = chart.series[series_index]
        color = _operator_color(matrix, operator)
        _fill_series(series, _hierarchy_chart_color(color, category_index, len(categories)))
        if hide_from_legend:
            hidden_legend_entries.add(series_index)
    _hide_chart_legend_entries(chart, hidden_legend_entries | {len(series_specs)})
    _format_stacked_segment_labels(chart)
    _add_total_labels(chart)

    category_key = _add_category_shade_bar(
        slide, categories, left=Inches(.55), top=Inches(6.46),
        width=Inches(8.9), height=Inches(.4),
    )
    hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, columns)
    chart_group = [chart_shape] + ([category_key] if category_key is not None else []) + ([hierarchy_grid] if hierarchy_grid is not None else [])
    slide.shapes.add_group_shape(chart_group + chart_shape._scoring_axis_labels).name = 'Scoring Chart With Category Key'
    _add_category_allocation_donut(slide, matrix)
    _text(slide, matrix['coverage_note'], 7.06, size=8, height=.2)


def _hierarchy_category_comparison_chart(presentation, matrix: dict) -> None:
    """Show each category with readable hierarchy labels and one series per operator."""
    categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
    category_groups = _breakdown_pages(categories, len(matrix['hierarchy_columns']), matrix.get('_split_charts') is not False)
    for page_index, category_group in enumerate(category_groups, 1):
        metrics = [row for row in matrix['rows'] if row['category'] in category_group]
        columns = [dict(column, comparison_category=category,
                        path=[{'level': 'Category', 'value': category}] + column['path'])
                   for category in category_group for column in matrix['hierarchy_columns']]
        maximum = max((sum(row['values'][column['id']]['points'] or 0 for row in metrics
                           if row['category'] == column['comparison_category'])
                       for column in columns), default=0.0)
        operators = list(dict.fromkeys(column['operator'] for column in columns))
        slide = _slide(presentation, _scoring_chart_title(matrix, 'per Category (Breakdown)'), _chart_subtitle(matrix))
        heading = ('All Categories' if len(category_group) == len(categories)
                   else category_group[0] if len(category_group) == 1 else ', '.join(category_group))
        if len(category_groups) > 1:
            heading = f'{heading} · Page {page_index} of {len(category_groups)}'
        _text(slide, heading, 1.42, left=.7, width=11.9, height=.25, size=12, color='#4A5B65')
        data = CategoryChartData()
        _add_hierarchy_chart_categories(data, columns)
        for operator in operators:
            values = []
            for column in columns:
                points = [row['values'][column['id']]['points'] for row in metrics
                          if row['category'] == column['comparison_category']
                          and row['values'][column['id']]['points'] is not None]
                values.append(sum(points) if points and column['operator'] == operator else None)
            data.add_series(_operator_label(matrix, operator), values)
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.7), Inches(1.8),
            presentation.slide_width - Inches(1.4), Inches(4.85), data,
        )
        chart = chart_shape.chart
        _format_chart(chart, maximum=maximum)
        chart.plots[0].gap_width = 120
        chart.plots[0].overlap = 100
        chart.plots[0].data_labels.font.size = Pt(min(10, max(4, chart_shape.width.inches * .89 * 72 / max(1, len(columns)) / 3.2)))
        chart.plots[0].data_labels.number_format = '0' if len(columns) > 40 else '0.0'
        chart.plots[0].has_data_labels = len(columns) <= 70
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.legend.font.size = Pt(11)
        for series, operator in zip(chart.series, operators):
            _fill_series(series, _operator_color(matrix, operator))
        hierarchy_grid = _add_chart_hierarchy_grid(slide, chart_shape, columns)
        slide.shapes.add_group_shape([chart_shape, hierarchy_grid] + chart_shape._scoring_axis_labels).name = 'Scoring Hierarchy Category Comparison Chart'
        _text(slide, matrix['coverage_note'], 6.95, size=9)


def _add_category_shade_bar(slide, categories: list[str], *, left: int, top: int, width: int, height: int):
    """Create a native continuous grayscale key bar grouped with its editable chart."""
    if not categories:
        return None
    key_shape = slide.shapes.add_table(1, len(categories), left, top, width, height)
    key_shape.name = 'Scoring Chart Category Shade Key'
    table = key_shape.table
    for index, category in enumerate(categories):
        color = _hierarchy_chart_color('#606060', index, len(categories)).lstrip('#')
        cell = table.cell(0, index)
        cell.text = str(category)
        cell.margin_left = cell.margin_right = Inches(.035)
        cell.margin_top = cell.margin_bottom = 0
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(color)
        cell.text_frame.word_wrap = True
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = _FONT
            paragraph.font.size = Pt(8)
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor.from_string(_header_foreground(color).lstrip('#'))
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.line_spacing = 1.0
            paragraph.space_before = paragraph.space_after = 0
    return key_shape


def _hierarchy_chart_color(color: str, category_index: int, category_count: int = 7) -> str:
    """Use lighter shades of each operator color to distinguish stacked KPI groups."""
    channels = [int(color.lstrip('#')[index:index + 2], 16) for index in (0, 2, 4)]
    tint = .72 * category_index / max(1, category_count - 1)
    return '#{:02X}{:02X}{:02X}'.format(*(round(channel + (255 - channel) * tint) for channel in channels))


def _wrapped_line_count(text: Any, width_inches: float, font_size: float) -> int:
    capacity = max(1, int((width_inches * 72 - 3.6) / max(.1, font_size * .65)))
    lines = 1
    used = 0
    for word in str(text or '').split():
        word_length = len(word)
        if used and used + 1 + word_length > capacity:
            lines += 1
            used = 0
        if word_length > capacity:
            extra_lines = ceil(word_length / capacity) - 1
            lines += extra_lines
            used = word_length % capacity
        else:
            used += (1 if used else 0) + word_length
    return lines


def _hierarchy_metric_layout(metrics: list[dict], total_label: str | None, *, available_height: float,
                             kpi_width: float) -> tuple[float, list[float]]:
    texts = [str(row.get('kpi', '')) for row in metrics]
    if total_label is not None:
        texts.append(total_label)
    font_size = min(7.5, max(2.5, available_height * 72 / max(1, len(texts)) * .82))
    line_counts = [_wrapped_line_count(text, kpi_width, font_size) for text in texts]
    for _ in range(3):
        base_height = available_height / max(1, sum(line_counts))
        font_size = min(7.5, max(2.5, base_height * 72 * .82))
        updated = [_wrapped_line_count(text, kpi_width, font_size) for text in texts]
        if updated == line_counts:
            break
        line_counts = updated
    base_height = available_height / max(1, sum(line_counts))
    return font_size, [base_height * lines for lines in line_counts]


def _hierarchy_content_font(texts: list[str], width_inches: float, *, maximum: float = 7.5) -> float:
    glyph_units = {
        '.': .28, '-': .4, '+': .48, '*': .36,
    }
    widest = max(
        (sum(glyph_units.get(char, .56 if char.isdigit() else .52) for char in str(text)) for text in texts),
        default=1.0,
    )
    usable_width = max(1.0, width_inches * 72 - 3.6)
    return min(maximum, max(2.5, usable_width / max(1.0, widest)))


def _add_hierarchy_chart_categories(data: CategoryChartData, columns: list[dict]) -> None:
    """Build native multi-level chart categories from hierarchy values."""
    nodes = {}
    for column in columns:
        parent = None
        prefix = ()
        for item in column.get('path', []):
            value = _hierarchy_display_value(item)
            prefix += (item.get('value'),)
            if prefix in nodes:
                parent = nodes[prefix]
                continue
            node = data.add_category(value) if parent is None else parent.add_sub_category(value)
            nodes[prefix] = node
            parent = node


def _hierarchy_header_groups(table, columns: list[dict], levels: list[str], *, start_col: int,
                             leaf_width: int, header_rows: int, physical_column_width: float,
                             leaf_label: str = 'Score', header_color: str | None = None,
                             gap_reference: str | None = None) -> None:
    for level_index, level in enumerate(levels):
        start = 0
        while start < len(columns):
            representative = columns[start]
            prefix = tuple((entry['level'], entry.get('value')) for entry in representative['path'][:level_index + 1])
            end = start
            while end + 1 < len(columns):
                next_prefix = tuple(
                    (entry['level'], entry.get('value')) for entry in columns[end + 1]['path'][:level_index + 1]
                )
                if next_prefix != prefix:
                    break
                end += 1
            left_column = start_col + start * leaf_width
            right_column = start_col + (end + 1) * leaf_width - 1
            cell = table.cell(level_index, left_column)
            if right_column > left_column:
                cell.merge(table.cell(level_index, right_column))
            label = _hierarchy_display_value(representative['path'][level_index])
            if level == 'Operator' and gap_reference:
                label = f'{label} − {gap_reference}'
            color = representative['color'] if level == 'Operator' else (header_color or '#E6ECFA')
            header_size = min(7.5, max(5.0, physical_column_width * leaf_width
                                      * (end - start + 1) * 12))
            _cell(cell, str(label), color=color, foreground=_header_foreground(color), size=header_size, bold=True)
            start = end + 1
    subheader_row = len(levels)
    for column_index, column in enumerate(columns):
        color = column['color']
        leaf_font_size = min(7.5, max(5.0, physical_column_width * 12))
        _cell(table.cell(subheader_row, start_col + column_index * leaf_width), leaf_label,
              color=color, foreground=_header_foreground(color), size=leaf_font_size, bold=True)
        if leaf_width == 2:
            _cell(table.cell(subheader_row, start_col + column_index * leaf_width + 1), 'GAP',
                  color='#F6D55C', foreground='#17232D', size=leaf_font_size, bold=True)


def _score_column_plan(columns: list[dict], gap_layout: str, *, show_gap_values: bool = True,
                       show_kpi_values: bool = False) -> list[tuple[dict, str]]:
    """Leaf columns of a scoring table; KPI values, when shown, precede each Score column."""
    scores = [item for column in columns
              for item in ([(column, 'Value')] if show_kpi_values else []) + [(column, 'Score')]]
    if not show_gap_values:
        return scores
    gaps = [(column, 'GAP') for column in columns if not column.get('is_reference')]
    if gap_layout == 'end':
        return scores + gaps
    return [item for column in columns for item in (
        ([(column, 'Value')] if show_kpi_values else [])
        + ([(column, 'Score')] if column.get('is_reference') else [(column, 'Score'), (column, 'GAP')])
    )]


def _score_column_headers(table, plan: list[tuple[dict, str]], levels: list[str],
                          *, gap_layout: str, leaf_width: float, fixed_count: int = 5) -> None:
    score_count = sum(kind in {'Score', 'Value'} for _, kind in plan)
    score_label = 'Value / Score' if any(kind == 'Value' for _, kind in plan) else 'Score'
    groups = [(score_label, 0, score_count), ('GAP', score_count, len(plan))]
    if gap_layout == 'adjacent' and any(kind == 'GAP' for _, kind in plan):
        groups = [('Score / GAP', 0, len(plan))]
    for label, start, end in groups:
        if start == end:
            continue
        cell = table.cell(0, fixed_count + start)
        if end - start > 1:
            cell.merge(table.cell(0, fixed_count + end - 1))
        _cell(cell, label, color='#F6D55C' if label == 'GAP' else '#455B65',
              foreground='#17232D' if label == 'GAP' else _WHITE, size=7, bold=True)
    for depth, level in enumerate(levels):
        start = 0
        while start < len(plan):
            column, kind = plan[start]
            prefix = tuple((entry['level'], entry.get('value')) for entry in column['path'][:depth + 1])
            end = start
            while end + 1 < len(plan):
                next_column, next_kind = plan[end + 1]
                next_prefix = tuple((entry['level'], entry.get('value'))
                                    for entry in next_column['path'][:depth + 1])
                if next_prefix != prefix or (gap_layout == 'end' and next_kind != kind
                                             and {kind, next_kind} != {'Value', 'Score'}):
                    break
                end += 1
            cell = table.cell(depth + 1, fixed_count + start)
            if end > start:
                cell.merge(table.cell(depth + 1, fixed_count + end))
            value = _hierarchy_display_value(column['path'][depth])
            color = column['color'] if level == 'Operator' else '#E6ECFA'
            size = min(7, max(2.5, leaf_width * (end - start + 1) * 12))
            _cell(cell, str(value) if value is not None else 'Not specified', color=color,
                  foreground=_header_foreground(color), size=size, bold=True)
            start = end + 1
    for index, (column, kind) in enumerate(plan, fixed_count):
        color = '#F6D55C' if kind == 'GAP' else '#E6ECFA' if kind == 'Value' else column['color']
        cell = table.cell(len(levels) + 1, index)
        _cell(cell, kind, color=color, foreground=_header_foreground(color),
              size=min(7, max(2.5, leaf_width * 12)), bold=True)


def _hierarchy_score_tables(presentation, matrices: list[dict], legend: list[dict], *, gap_layout: str = 'end',
                            show_gap_values: bool = True, title: str = 'Scoring Tables — Breakdown',
                            show_kpi_values: bool = False) -> None:
    for matrix in matrices:
        is_summary = matrix.get('table_mode') == 'summary'
        fixed_count = 4 if is_summary else 5
        weight_index = 2 if is_summary else 3
        maximum_index = weight_index + 1
        columns = matrix['hierarchy_columns']
        levels = matrix['hierarchy_levels']
        metrics = matrix['rows']
        plan = _score_column_plan(columns, gap_layout, show_gap_values=show_gap_values,
                                  show_kpi_values=show_kpi_values and not is_summary)
        total_label = _score_total_label(
            [*metrics, matrix['total']],
            [column['id'] for column, kind in plan if kind == 'Score'],
            [column['id'] for column, kind in plan if kind == 'GAP'],
            show_gap_values=show_gap_values,
        )
        slide = _slide(presentation, title, _chart_subtitle(matrix))
        header_rows = len(levels) + 2
        row_count = header_rows + len(metrics) + 1
        table_height = 5.0
        table_width = presentation.slide_width / 914400 - 1.1
        fixed_widths = [0.9, 2.15, *([] if is_summary else [0.65]), 0.47, 0.52]
        table = slide.shapes.add_table(row_count, fixed_count + len(plan), Inches(.55), Inches(1.5),
                                       Inches(table_width), Inches(table_height)).table
        leaf_width = (table_width - sum(fixed_widths)) / max(1, len(plan))
        for index, column in enumerate(table.columns):
            column.width = Inches(fixed_widths[index] if index < fixed_count else leaf_width)
        header_level_height = min(.23, 1.05 / max(1, len(levels)))
        table.rows[0].height = Inches(.22)
        for row_index in range(1, len(levels) + 1):
            table.rows[row_index].height = Inches(header_level_height)
        table.rows[header_rows - 1].height = Inches(.22)
        available_body_height = table_height - header_level_height * len(levels) - .44
        metric_font, body_heights = _hierarchy_metric_layout(
            metrics, total_label,
            available_height=available_body_height, kpi_width=fixed_widths[1],
        )
        for row_index, row_height in enumerate(body_heights, header_rows):
            table.rows[row_index].height = Inches(row_height)
        _score_column_headers(table, plan, levels, gap_layout=gap_layout, leaf_width=leaf_width,
                              fixed_count=fixed_count)
        static_headers = [('CATEGORY', _GAP_SUMMARY_HEADER),
                          ('CATEGORY' if is_summary else 'KPI', _GAP_SUMMARY_HEADER),
                          *([] if is_summary else [('Type of KPI', _GAP_SUMMARY_HEADER)]),
                          ('Score weight\n(%)', '#5B9A44'), ('Max score', '#C9483F')]
        for index, (label, color) in enumerate(static_headers):
            cell = table.cell(0, index)
            cell.merge(table.cell(header_rows - 1, index))
            _cell(cell, label, color=color, foreground=_header_foreground(color), size=7, bold=True)
        data_font = min(7, metric_font)
        total = matrix['total']
        numeric_texts = []
        for metric in metrics + [total]:
            for column, kind in plan:
                value = metric['values'][column['id']]
                numeric_texts.append(_gap_number(metric, column['id']) if kind == 'GAP'
                                     else _kpi_value_text(metric, value) if kind == 'Value'
                                     else _number(value['points']) + ('*' if value['points'] is not None and not value['complete'] else ''))
        numeric_font = _hierarchy_content_font(
            numeric_texts, leaf_width, maximum=min(11, min(body_heights) * 72 * .82),
        )
        category_color = _GAP_SUMMARY_CATEGORY if is_summary else _SCORING_EXPANDED_CATEGORY
        for row_offset, metric in enumerate(metrics, header_rows):
            subtotal = metric.get('row_type') == 'category'
            row_color = ('#E3E6E7' if is_summary else category_color) if subtotal else None
            best, worst = (_summary_extreme_columns(
                metric.get('values', {}), [column for column, kind in plan if kind == 'Score'],
            ) if matrix.get('table_mode') == 'summary' and subtotal else (set(), set()))
            _cell(table.cell(row_offset, 0), metric['category'], color=category_color,
                  size=data_font, left=True, bold=True)
            _cell(table.cell(row_offset, 1), metric['kpi'], color=row_color or '#E7E8E9',
                  bold=subtotal, left=True, size=metric_font)
            if not is_summary:
                _cell(table.cell(row_offset, 2), metric.get('kpi_type', ''),
                      color=row_color or _kpi_type_color(metric), size=data_font, bold=subtotal)
            _cell(table.cell(row_offset, weight_index), _number(metric['weight_percent']) + '%', color=row_color or _WHITE,
                  size=data_font, bold=subtotal)
            _cell(table.cell(row_offset, maximum_index), _number(metric['max_points']), color=row_color or _WHITE,
                  size=data_font, bold=subtotal)
            for index, (column, kind) in enumerate(plan, fixed_count):
                leaf_id = column['id']
                value = metric['values'][leaf_id]
                if kind == 'GAP':
                    _cell(table.cell(row_offset, index), _gap_number(metric, leaf_id),
                          color=row_color or metric.get('gap_colors', {}).get(leaf_id, THRESHOLD_COLORS['Unavailable']),
                          size=numeric_font, bold=True)
                elif kind == 'Value':
                    _cell(table.cell(row_offset, index), _kpi_value_text(metric, value),
                          color=row_color or _WHITE, size=numeric_font)
                else:
                    partial = value['points'] is not None and not value['complete']
                    score_color = (_SUMMARY_BEST_COLOR if leaf_id in best else
                                   _SUMMARY_WORST_COLOR if leaf_id in worst else
                                   row_color or value.get('color', _NEUTRAL))
                    _cell(table.cell(row_offset, index), _number(value['points']) + ('*' if partial else ''),
                          color=score_color, size=numeric_font, bold=subtotal)
        total_index = row_count - 1
        total_color = '#D8DFE4' if is_summary else category_color
        best, worst = (_summary_extreme_columns(
            total.get('values', {}), [column for column, kind in plan if kind == 'Score'],
        ) if is_summary else (set(), set()))
        for index, label in enumerate(('TOTAL', total_label, *([] if is_summary else ['']),
                                      _number(total['weight_percent']) + '%', _number(total['max_points']))):
            _cell(table.cell(total_index, index), label, color=total_color, bold=True, left=index < 2, size=data_font)
        for index, (column, kind) in enumerate(plan, fixed_count):
            value = total['values'][column['id']]
            text = (_gap_number(total, column['id']) if kind == 'GAP' else '' if kind == 'Value' else
                    _number(value['points']) + ('*' if value['points'] is not None and not value['complete'] else ''))
            color = (_SUMMARY_BEST_COLOR if kind == 'Score' and column['id'] in best else
                     _SUMMARY_WORST_COLOR if kind == 'Score' and column['id'] in worst else total_color)
            _cell(table.cell(total_index, index), text, color=color, size=numeric_font, bold=True)
        _merge_category_cells(table, metrics, header_rows,
                              color=category_color, include_subtotals=True, bold_categories=True)
        apply_scorecard(
            table, header_rows=header_rows, metrics=metrics, total=total, total_label=total_label, font=_FONT,
            value_columns=[(index, kind, column['id']) for index, (column, kind) in enumerate(plan, fixed_count)],
            drop_leaf_header=gap_layout == 'end' and not any(kind == 'Value' for _column, kind in plan),
            label_width=2.6 if is_summary else 3.0,
        )
        if is_summary:
            _add_summary_operator_legend(slide, presentation.slide_width.inches - .55)
        else:
            for index, item in enumerate(legend):
                legend_table = slide.shapes.add_table(
                    1, 1, Inches(presentation.slide_width.inches - .55 - len(legend) * 1.15 + .05 + index * 1.15),
                    Inches(1.25), Inches(1.1), Inches(.18),
                ).table
                _cell(legend_table.cell(0, 0), item.get('band', ''), color=item['color'], size=8)
                # As the cells of the band look in the theme.
                style_legend_band(legend_table.cell(0, 0), item.get('band', ''), 8, _FONT)
        _text(slide, matrix['coverage_note'], 7.27, size=8, height=.18)


def _kpi_value_text(metric: dict, value: dict) -> str:
    """Measured KPI value of a Breakdown row; category subtotals have none."""
    if metric.get('row_type') == 'category':
        return ''
    measured = value.get('value')
    return _number(measured) if isinstance(measured, (int, float)) and not isinstance(measured, bool) else 'N/A'


def _merge_category_cells(table, metrics: list[dict], start_row: int,
                          *, color: str = '#E6F0F7', include_subtotals: bool = False,
                          bold_categories: bool = False,
                          alternate_colors: tuple[str, str] | None = None) -> None:
    first = 0
    run_index = 0
    while first < len(metrics):
        last = first
        while (last + 1 < len(metrics) and metrics[last + 1]['category'] == metrics[first]['category']
               and (include_subtotals or metrics[last + 1].get('row_type') != 'category')):
            last += 1
        cell = table.cell(start_row + first, 0)
        run_color = alternate_colors[run_index % 2] if alternate_colors else color
        if alternate_colors:
            cell.fill.solid()
            cell.fill.fore_color.rgb = _theme_rgb(run_color)
        if last > first:
            cell.merge(table.cell(start_row + last, 0))
            subtotal = metrics[first].get('row_type') == 'category'
            _cell(cell, metrics[first]['category'], color=run_color,
                  size=7, left=True, bold=bold_categories or subtotal)
        run_index += 1
        first = last + 1


def _hierarchy_gap_projection(matrix: dict, columns: list[dict]) -> dict:
    """Keep a selected set of leaves in a hierarchy GAP view."""
    projected = deepcopy(matrix)
    leaf_ids = [column['id'] for column in columns]
    projected['hierarchy_columns'] = deepcopy(columns)
    projected['operators'] = leaf_ids
    projected['operator_styles'] = {
        leaf_id: projected['operator_styles'][leaf_id]
        for leaf_id in leaf_ids if leaf_id in projected.get('operator_styles', {})
    }
    projected['rows'] = [
        row | {
            'gaps': {leaf_id: row.get('gaps', {}).get(leaf_id) for leaf_id in leaf_ids},
            'gap_colors': {leaf_id: row.get('gap_colors', {}).get(leaf_id, THRESHOLD_COLORS['Unavailable'])
                           for leaf_id in leaf_ids},
        }
        for row in projected.get('rows', [])
    ]
    compared = [column for column in columns if not column.get('is_reference')]
    if len({column['operator'] for column in compared}) == 1:
        # One operator: its KPIs ordered by their GAP in the latest campaign.
        if projected.get('table_mode') == 'summary':
            projected['rows'].sort(key=lambda row: gap_priority_order_key(row, compared))
        else:
            kpis = [row for row in projected['rows'] if row.get('row_type') != 'category']
            projected['rows'] = sorted(kpis, key=lambda row: gap_priority_order_key(row, compared))
    projected['total'] = {
        'gaps': {leaf_id: matrix.get('total', {}).get('gaps', {}).get(leaf_id) for leaf_id in leaf_ids},
        'gap_partial': {leaf_id: matrix.get('total', {}).get('gap_partial', {}).get(leaf_id, False)
                        for leaf_id in leaf_ids},
        'gap_environments': {leaf_id: matrix.get('total', {}).get('gap_environments', {}).get(leaf_id, [])
                             for leaf_id in leaf_ids},
    }
    return projected


# The GAP tables grow down to their colour scale (at 6.8 in), their text up to this size.
GAP_TABLE_BOTTOM = 6.7
GAP_TABLE_MAXIMUM_FONT = 9.0
# A line of a row, as a multiple of its text size, and what a row adds around its lines, in points.
_ROW_LINE = 1.2
_ROW_PADDING = .5


def _text_points(text: str, size: float, bold: bool | None) -> float:
    """About how wide ``text`` is at ``size`` points (Arial-like widths, a little wider when bold)."""
    width = 0.0
    for char in text:
        width += (.28 if char in ' .,:;|!il' else .34 if char in '[]()/-' else .89 if char in '%@MW' else
                  .56 if char.isdigit() or char.islower() else .68 if char.isupper() else .6)
    return width * size * (1.06 if bold else 1.0)


def _fill_gap_table(table, *, header_rows: int, bottom: float = GAP_TABLE_BOTTOM,
                    maximum: float = GAP_TABLE_MAXIMUM_FONT) -> None:
    """Make the text of the rows under the header of a scorecard GAP table as large as they fit down to ``bottom``
    (inches), never smaller than it is: every run keeps its size relative to the others, the KPI labels (first column)
    wrap when they must and the values never do; the rows then share the room left, up to half their height more."""
    frame = table._graphic_frame
    rows = list(table.rows)
    columns = [column.width / 12700 for column in table.columns]
    available = (Inches(bottom) - frame.top - sum(row.height for row in rows[:header_rows])) / 12700
    body = []
    for row in rows[header_rows:]:
        cells = []
        for index, cell in enumerate(row.cells):
            if cell.is_spanned:
                continue
            span = cell.span_width if cell.is_merge_origin else 1
            runs = [run for paragraph in cell.text_frame.paragraphs for run in paragraph.runs
                    if run.text and run.font.size is not None]
            # The cell margins (.04 in each side) and a little room for the renderer.
            cells.append((index == 0, sum(columns[index:index + span]) * .96 - 6, runs))
        body.append(cells)
    sizes = [run.font.size.pt for cells in body for _wraps, _width, runs in cells for run in runs]
    if not sizes or available <= 0:
        return
    largest = max(sizes)

    def heights(scale: float) -> list[float] | None:
        result = []
        for cells in body:
            lines, tallest = 1, 0.0
            for wraps, width, runs in cells:
                if not runs:
                    continue
                needed = sum(_text_points(run.text, run.font.size.pt * scale, run.font.bold) for run in runs)
                tallest = max(tallest, *(run.font.size.pt * scale for run in runs))
                if needed > width:
                    if not wraps:
                        return None
                    lines = max(lines, ceil(needed * 1.08 / max(1.0, width)))
            result.append(lines * (tallest or largest * scale) * _ROW_LINE + _ROW_PADDING)
        return result

    scale = maximum / largest
    while scale > 1:
        layout = heights(scale)
        if layout is not None and sum(layout) <= available:
            break
        scale *= .97
    else:
        return
    grow = min(1.5, available / sum(layout))
    for row, height, cells in zip(rows[header_rows:], layout, body):
        row.height = Pt(height * grow)
        for _wraps, _width, runs in cells:
            for run in runs:
                run.font.size = Pt(round(run.font.size.pt * scale, 1))
        for cell in row.cells:
            cell.margin_top = cell.margin_bottom = 0
    frame.height = sum(row.height for row in rows)


def _add_gap_total_row(table, rows: list[dict], column_ids: list[str | None], *, size: float) -> None:
    """Sum valid KPI GAPs independently for each displayed column."""
    index = len(table.rows) - 1
    label = table.cell(index, 0)
    label.merge(table.cell(index, 2))
    _cell(label, 'Total KPI GAP', color='#D7E0E5', size=size, bold=True, left=True)
    for offset, column_id in enumerate(column_ids, 3):
        values = [row.get('gaps', {}).get(column_id) if column_id is not None
                  else row.get('gap_points') for row in rows]
        valid = [float(value) for value in values if value is not None and isfinite(float(value))]
        partial = len(valid) < len(values) or any(
            row.get('gap_partial', {}).get(column_id, False) if column_id is not None
            else row.get('gap_partial', False) for row in rows
        )
        text = _number(sum(valid) if valid else None) + ('*' if valid and partial else '')
        _cell(table.cell(index, offset), text, color='#D7E0E5', size=size, bold=True)


def _hierarchy_gap_tables(presentation, matrices: list[dict], *,
                          title: str = 'GAP Analysis — All vs reference',
                          show_priority: bool = False) -> None:
    for matrix in matrices:
        columns = matrix['hierarchy_columns']
        levels = matrix['hierarchy_levels']
        metrics = [row for row in matrix['rows'] if row.get('row_type') != 'category']
        slide = _slide(presentation, title, _chart_subtitle(matrix))
        header_rows = len(levels) + 1
        row_count = header_rows + len(metrics) + 1
        table_height = 4.65
        table_width = 9.3 if show_priority else 12.03
        fixed_widths = [1.0, 2.55, .7]
        table = slide.shapes.add_table(row_count, 3 + len(columns), Inches(.65), Inches(1.55),
                                       Inches(table_width), Inches(table_height)).table
        leaf_width = (table_width - sum(fixed_widths)) / max(1, len(columns))
        for index, column in enumerate(table.columns):
            column.width = Inches(fixed_widths[index] if index < 3 else leaf_width)
        header_level_height = min(.25, 1.05 / max(1, len(levels)))
        subheader_height = .25
        for row_index in range(len(levels)):
            table.rows[row_index].height = Inches(header_level_height)
        table.rows[len(levels)].height = Inches(subheader_height)
        available_body_height = table_height - header_level_height * len(levels) - subheader_height
        metric_font, body_heights = _hierarchy_metric_layout(
            metrics, None, available_height=available_body_height - .28,
            kpi_width=fixed_widths[1],
        )
        for row_index, row_height in enumerate(body_heights, header_rows):
            table.rows[row_index].height = Inches(row_height)
        _hierarchy_header_groups(table, columns, levels, start_col=3, leaf_width=1,
                                 header_rows=header_rows, physical_column_width=leaf_width,
                                 leaf_label='GAP', gap_reference=_operator_label(matrix, matrix['baseline_operator']),
                                 header_color=_GAP_SUMMARY_COMPARISON_HEADER)
        data_font = min(7, metric_font)
        for index, label in enumerate(('Category', 'KPI', 'Type of KPI')):
            cell = table.cell(0, index)
            if header_rows > 1:
                cell.merge(table.cell(header_rows - 1, index))
            _cell(cell, label, color=_GAP_SUMMARY_HEADER,
                  foreground=_WHITE, size=7, bold=True)
        gap_texts = [
            _gap_number(row, column['id'])
            for row in metrics for column in columns
        ]
        leaf_font = _hierarchy_content_font(gap_texts, leaf_width, maximum=data_font)
        for row_offset, row in enumerate(metrics, header_rows):
            subtotal = row.get('row_type') == 'category'
            row_color = _GAP_SUMMARY_CATEGORY if subtotal else None
            _cell(table.cell(row_offset, 0), row['category'], color=row_color or _GAP_SUMMARY_CATEGORY,
                  foreground='#17232D', size=data_font, bold=True, left=True)
            _cell(table.cell(row_offset, 1), row['kpi'], color=row_color or '#E6ECFA', size=metric_font, left=True,
                  bold=subtotal)
            _cell(table.cell(row_offset, 2), row['kpi_type'],
                  color=row_color or _kpi_type_color(row),
                  size=data_font, bold=subtotal)
            for column_index, column in enumerate(columns, 3):
                leaf_id = column['id']
                _cell(table.cell(row_offset, column_index), _gap_number(row, leaf_id),
                      color=row_color or row['gap_colors'].get(leaf_id, THRESHOLD_COLORS['Unavailable']),
                      size=leaf_font, bold=True)
        _merge_category_cells(table, metrics, header_rows, include_subtotals=True,
                              bold_categories=True, alternate_colors=('#E6F0F7', '#D7E5EE'))
        table.rows[len(table.rows) - 1].height = Inches(.28)
        _add_gap_total_row(table, metrics, [column['id'] for column in columns], size=leaf_font)
        header_rows = apply_scorecard(
            table, header_rows=header_rows, metrics=metrics, total={}, total_label='Total KPI GAP', font=_FONT,
            value_columns=[(index, 'GAP', column['id']) for index, column in enumerate(columns, 3)],
            # The KPIs take the width the GAP columns do not need (each keeps about an inch).
            drop_leaf_header=True, label_width=max(3.0, min(4.6, table_width - 1.0 * len(columns))),
        )
        _fill_gap_table(table, header_rows=header_rows)
        _text(slide, matrix['note'], 7.05, height=.25, size=9)
        if show_priority:
            _add_individual_gap_notes(slide, matrix, metrics)
            _add_gap_priority_arrow(slide, table, header_rows=header_rows)
        _add_gap_color_scale(slide, matrix, left=.65, width=table_width)


def _gap_summary_tables(presentation, matrices: list[dict]) -> None:
    """Export the combined comparison before individual operator GAP slides."""
    for matrix in matrices:
        operators = matrix['operators']
        pages = [[row for row in matrix['rows'] if row.get('row_type') != 'category']]
        for page_index, rows in enumerate(pages):
            page_label = f' · Page {page_index + 1}/{len(pages)}' if len(pages) > 1 else ''
            reference = _operator_label(matrix, matrix['baseline_operator'])
            slide = _slide(presentation, 'GAP Analysis — All vs ' + reference,
                           _chart_subtitle(matrix) + page_label)
            table = slide.shapes.add_table(len(rows) + 2, len(operators) + 3, Inches(.65), Inches(1.65),
                                           Inches(12.03), Inches(4.55)).table
            widths = [1.45, 4.05, 1.1] + [5.43 / max(1, len(operators))] * len(operators)
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            headers = ['Category', 'NETCHECK KPIs', 'Type of KPI'] + [
                f'{_operator_label(matrix, operator)} − {reference}' for operator in operators]
            table.rows[0].height = Inches(.43)
            metric_font, body_heights = _hierarchy_metric_layout(
                rows, None, available_height=4.55 - .43 - .28, kpi_width=4.05,
            )
            for row, height in zip(list(table.rows)[1:], body_heights):
                row.height = Inches(height)
            for index, header in enumerate(headers):
                color = _operator_color(matrix, operators[index - 3]) if index >= 3 else _GAP_SUMMARY_HEADER
                _cell(table.cell(0, index), header, color=color,
                      foreground=_header_foreground(color), size=10, bold=True)
            for index, row in enumerate(rows, 1):
                subtotal = row.get('row_type') == 'category'
                row_color = _GAP_SUMMARY_CATEGORY if subtotal else None
                _cell(table.cell(index, 0), row['category'], color=row_color or _GAP_SUMMARY_CATEGORY,
                      foreground='#17232D', size=metric_font, bold=True, left=True)
                _cell(table.cell(index, 1), row['kpi'], color=row_color or '#E6ECFA', size=metric_font,
                      left=True, bold=subtotal)
                _cell(table.cell(index, 2), row['kpi_type'],
                      color=row_color or _kpi_type_color(row),
                      size=metric_font, bold=subtotal)
                for column, operator in enumerate(operators, 3):
                    _cell(table.cell(index, column), _gap_number(row, operator),
                          color=row_color or row['gap_colors'][operator], size=metric_font, bold=True)
            _text(slide, matrix['note'], 7.12, height=.25, size=9)
            _merge_category_cells(table, rows, 1, include_subtotals=True, bold_categories=True,
                                  alternate_colors=('#E6F0F7', '#D7E5EE'))
            table.rows[len(table.rows) - 1].height = Inches(.28)
            _add_gap_total_row(table, rows, operators, size=metric_font)
            apply_scorecard(table, header_rows=1, metrics=rows, total={}, total_label='Total KPI GAP', font=_FONT,
                            value_columns=[(index, 'GAP', operator) for index, operator in enumerate(operators, 3)],
                            label_width=4.2)
            _fill_gap_table(table, header_rows=1)
            _add_gap_color_scale(slide, matrix, left=.65, width=12.03)


def _gap_tables(presentation, matrices: list[dict]) -> None:
    for matrix in matrices:
        pages = [[row for row in matrix['rows'] if row.get('row_type') != 'category']]
        for page_index, rows in enumerate(pages):
            page_label = f' · Page {page_index + 1}/{len(pages)}' if len(pages) > 1 else ''
            comparison = f'{_operator_label(matrix, matrix["operator"])} vs {_operator_label(matrix, matrix["baseline_operator"])}'
            subtitle = _chart_subtitle(matrix) + page_label
            slide = _slide(presentation, 'GAP Analysis — ' + comparison, subtitle)
            if not rows:
                _text(slide, 'No comparable KPI gaps are available. See the scoring matrix for missing values.', 1.8, height=1, size=16)
                continue
            table = slide.shapes.add_table(len(rows) + 2, 4, Inches(.65), Inches(1.65), Inches(9.3),
                                           Inches(4.55)).table
            for column, width in zip(table.columns, (1.65, 4.45, 1.5, 1.7)):
                column.width = Inches(width)
            table.rows[0].height = Inches(.43)
            metric_font, body_heights = _hierarchy_metric_layout(
                rows, None, available_height=4.55 - .43 - .28, kpi_width=4.45,
            )
            for row, height in zip(list(table.rows)[1:], body_heights):
                row.height = Inches(height)
            headers = ['Category', 'NETCHECK KPIs', 'Type of KPI',
                       f'{_operator_label(matrix, matrix["operator"])} − {_operator_label(matrix, matrix["baseline_operator"])}']
            for index, header in enumerate(headers):
                color = _operator_color(matrix, matrix['operator']) if index == 3 else _GAP_SUMMARY_HEADER
                _cell(table.cell(0, index), header, color=color,
                      foreground=_header_foreground(color), size=11, bold=True)
            for index, row in enumerate(rows, 1):
                subtotal = row.get('row_type') == 'category'
                row_color = _GAP_SUMMARY_CATEGORY if subtotal else None
                _cell(table.cell(index, 0), row['category'], color=row_color or _GAP_SUMMARY_CATEGORY,
                      foreground='#17232D', size=metric_font, bold=True, left=True)
                _cell(table.cell(index, 1), row['kpi'], color=row_color or '#E6ECFA', size=metric_font,
                      left=True, bold=subtotal)
                _cell(table.cell(index, 2), row['kpi_type'],
                      color=row_color or _kpi_type_color(row),
                      size=metric_font, bold=subtotal)
                _cell(table.cell(index, 3), _gap_number(row), color=row_color or row.get('gap_color', '#FFF6C2'),
                      size=metric_font, bold=True)
            _add_individual_gap_notes(slide, matrix, rows)
            _merge_category_cells(table, rows, 1, include_subtotals=True, bold_categories=True,
                                  alternate_colors=('#E6F0F7', '#D7E5EE'))
            table.rows[len(table.rows) - 1].height = Inches(.28)
            _add_gap_total_row(table, rows, [None], size=metric_font)
            apply_scorecard(table, header_rows=1, metrics=rows, total={}, total_label='Total KPI GAP', font=_FONT,
                            value_columns=[(3, 'GAP', None)], label_width=5.6)
            _fill_gap_table(table, header_rows=1)
            _add_gap_priority_arrow(slide, table)
            _add_gap_color_scale(slide, matrix, left=.65, width=9.3)
            _text(slide, 'KPIs are ordered by GAP, from lowest to highest. GAP Priority does not affect this order.', 7.12, height=.25, size=9)


def _scoring_block_title(scoring: str | None, title: str) -> str:
    """Name the scoring in block titles when the document carries both scorings."""
    return f'{SCORING_LABELS[scoring]} — {title}' if scoring else title


def _scoring_block_subtitle(scoring: str | None, environment: str, configuration: dict) -> str:
    """Slide subtitle of a block: the scoring, and the environment unless it is All Environments."""
    label = SCORING_LABELS[scoring or BEST_NETWORK_SCORING]
    if configuration.get('_scenario_name'):
        label = f"{configuration['_scenario_name']} · {label}"
    if str(environment).casefold() in {'combined', 'all', 'all environments'}:
        return label
    return f'{label} — {_scoring_environment_title(environment, configuration)}'


def _scoring_slide_notes(result: dict[str, Any]) -> list[str]:
    coverage_notes = scoring_coverage_notes(result)
    coverage_details = [f"{'All Environments' if environment == 'Combined' else environment} — {note}"
                        for environment, notes in coverage_notes.items() for note in notes]
    slide_warnings = [str(notice) for notice in result.get('notices', [])]
    slide_warnings.extend(str(warning) for warning in result.get('warnings', [])
                          if str(warning) != _LEGACY_CAMPAIGN_WARNING)
    slide_warnings.extend(coverage_details)
    return slide_warnings


def _option(options: dict | None, group: str, key: str, default: bool = True) -> bool:
    """A report scenario option; without a scenario every slide is included. Options that moved to another group are
    still read where configurations saved before kept them (the location cards in charts, the Points Lost Map in gap),
    and the Geographical Areas where they were called Locations."""
    if options is None:
        return default
    moved = {('geographical_areas', 'cards'): ('charts', 'location_cards'), ('points', 'map'): ('gap', 'points_loss_map'),
             ('points', 'kpis'): ('gap', 'profile')}
    supplied = options.get(group) or (options.get('locations') if group == 'geographical_areas' else None) or {}
    if key not in supplied and (group, key) in moved:
        old_group, old_key = moved[(group, key)]
        return bool((options.get(old_group) or {}).get(old_key, default))
    return bool(supplied.get(key, default))


def _everyone(_operator: str, _label: str | None = None) -> bool:
    return True


def focus_matcher(focus: Any, operators: Any = ()) -> Callable[..., bool]:
    """Whether an operator (by its name or label) is one of the Focus operators: the labels with the Focus role in the
    Operator Maps and the names of the CDRs they stand for (see operator_roles.focus_operators), or patterns (see
    value_patterns), whatever the case. Without Focus operators, or when none of ``operators`` (those of the results) is
    one of them, every operator is."""
    from src.modules.value_patterns import ValueSet

    chosen = ValueSet(focus or [])

    def matches(operator: str, label: str | None = None) -> bool:
        return chosen.accepts([operator, label])
    names = [str(name) for name in operators or [] if str(name).strip()]
    if not chosen or (names and not any(matches(name) for name in names)):
        return _everyone
    return matches


def _first_ranking_group(rankings: list[dict]) -> list[dict]:
    """The location rankings of the most specific level (City, then Cluster, then Region) and its first group."""
    for level in ('City', 'Cluster', 'Region'):
        items = [item for item in rankings if item['level'] == level]
        if items:
            title = items[0].get('title') or ''
            return [item for item in items if (item.get('title') or '') == title]
    return []


def _add_insight_slides(presentation, views: dict, environment: str, scoring: str | None, configuration: dict,
                        *, stage: str, options: dict | None = None, focus: Callable[..., bool] = _everyone) -> None:
    """Add the insight slides of one stage of an environment block, in the order of the results tabs.

    ``summary`` (after the cover): where each operator stands, its position, the biggest levers and how the measurements
    ran; ``geographical_areas`` (after the scoring tables): the location cards, the locations ranked and against the
    reference; ``trends`` (after the Geographical Areas): the scoring trend; ``gaps`` (after the GAP tables): the GAP by part and the KPI
    GAP profiles; ``points``: where the points go, the points-lost maps, Vendors, kinds of site, sites and route
    stretches; ``comparison``: the campaign comparison.
    """
    insights = views.get('insights') or {}
    key = 'Combined' if str(environment).casefold() in {'all environments', 'combined', 'all'} else environment
    # The insights summarise All Environments; a single exported environment gets its own.
    if views.get('_export_environment') == 'all' and key != 'Combined':
        return
    label = SCORING_LABELS[scoring or BEST_NETWORK_SCORING]
    subtitle = _scoring_block_subtitle(scoring, key, configuration)
    scope = views.get('_scope')
    items = {name: [item for item in values if isinstance(item, dict) and item.get('environment') == key]
             for name, values in insights.items() if isinstance(values, list)}
    done = views.setdefault('_insight_slides', set())
    if (key, stage) in done:
        return
    done.add((key, stage))
    environment_labels = {name: _scoring_environment_title(name, configuration)
                          for name in configuration.get('scope', {}).get('environments', {})}

    # The Focus operators (see focus_matcher): those of Where to act and its performance owners, places and route, the
    # locations ranked and against the reference, the GAP by part, the KPI GAP Profile and the Points Lost views; the
    # other slides show every operator.
    compared = focus
    # The operators each section of the page has enabled (sent with the export of a Saved Filter); without them, the
    # Focus operators (every operator in the sections they do not choose), with the operators enabled or disabled in
    # any section of the page (_enabled and _disabled), as the page shows them.
    choices = views.get('_operator_choices') or {}
    enabled = {str(name).casefold() for name in choices.get('_enabled') or []}
    disabled = {str(name).casefold() for name in choices.get('_disabled') or []}

    def toggled(operator: Any, operator_label: Any = None) -> bool | None:
        keys = {str(operator).casefold(), str(operator_label or '').casefold()} - {''}
        return True if keys & enabled else False if keys & disabled else None

    def chosen(section: str, default: Callable[..., bool] | None = None) -> Callable[..., bool]:
        names = {str(name).casefold() for name in choices.get(section) or []}
        if not names:
            base = default or compared
            if not enabled and not disabled:
                return base

            def with_toggles(operator, operator_label=None):
                state = toggled(operator, operator_label)
                return base(operator, operator_label) if state is None else state
            return with_toggles
        return lambda operator, operator_label=None: (str(operator).casefold() in names
                                                      or str(operator_label or '').casefold() in names)

    def compared_items(section: str, entries: list[dict]) -> list[dict]:
        """The items of a section whose page leaves the reference out by default (unless it was enabled)."""
        return entries if choices.get(section) else [
            entry for entry in entries
            if not entry.get('is_reference') or toggled(entry.get('operator'), entry.get('label'))]

    if stage == 'summary':
        # As on the page: how the measurements ran, Where to act and its performance owners first, then the standings, positions
        # and levers, and the places ranked and the route that loses most.
        if _option(options, 'summary', 'scope'):
            add_campaign_scope_slides(presentation, items.get('campaign_scope', []), subtitle=subtitle, new_slide=_slide,
                                      environment_labels=environment_labels)
        # Where to act: for the operator chosen on the page, else for every Focus operator (every operator but the
        # reference without them, or when they are only the reference).
        act_items = {name: items.get(name, []) for name in ('act_plans', 'owner_cards', 'place_rankings', 'route_focus')}

        def act_operators_of(accepted: Callable[..., bool]) -> list[str]:
            return list(dict.fromkeys(item['operator'] for entries in act_items.values() for item in entries
                                      if accepted(item['operator'], item.get('label'))
                                      and (choices.get('act-operator') or not item.get('is_reference'))))
        act_operators = act_operators_of(chosen('act-operator'))
        if not act_operators and not choices.get('act-operator'):
            act_operators = act_operators_of(_everyone)
        act_items = {name: sorted((item for item in entries if item['operator'] in act_operators),
                                  key=lambda item: act_operators.index(item['operator'])) for name, entries in act_items.items()}
        # The colours of the page: every owner of the results, in the order of their colours.
        owner_color = owner_palette(insights.get('owner_order') or owner_names(act_items))
        if _option(options, 'summary', 'act'):
            add_act_plan_slides(presentation, act_items['act_plans'], subtitle=subtitle, new_slide=_slide, scope=scope)
        if _option(options, 'summary', 'owners'):
            add_owner_card_slides(presentation, act_items['owner_cards'], subtitle=subtitle, new_slide=_slide,
                                  owner_color=owner_color, scope=scope)
        if _option(options, 'summary', 'standings'):
            add_standings_slides(presentation, items.get('standings', []), scoring_label=label, subtitle=subtitle,
                                 new_slide=_slide, scope=scope, compared=chosen('summary-operators', _everyone))
        if _option(options, 'summary', 'position'):
            add_position_matrix_slides(presentation, items.get('position_matrices', []), subtitle=subtitle, new_slide=_slide,
                                       environment_labels=environment_labels, scope=scope,
                                       compared=chosen('summary-operators', _everyone))
        if _option(options, 'summary', 'levers'):
            add_levers_slides(presentation, items.get('levers', []), subtitle=subtitle, new_slide=_slide, scope=scope,
                              compared=chosen('summary-operators', _everyone))
        if _option(options, 'summary', 'places'):
            add_place_ranking_slides(presentation, act_items['place_rankings'], subtitle=subtitle, new_slide=_slide,
                                     owner_color=owner_color, scope=scope)
        if _option(options, 'summary', 'route'):
            boundaries = views.get('_points_loss_boundaries')
            # The routes of the operators chosen on the page (in their own section, or in any section), else those of the
            # operators of Where to act.
            route_items = act_items['route_focus']
            if choices.get('route-operators') or enabled or disabled:
                wanted = chosen('route-operators')
                route_items = [item for item in items.get('route_focus', []) if wanted(item['operator'], item.get('label'))]
            add_route_focus_slides(presentation, route_items, subtitle=subtitle, new_slide=_slide,
                                   owner_color=owner_color, boundaries=boundaries('Area') if boundaries else None, scope=scope)
        return
    if stage == 'trends':
        if _option(options, 'trends', 'trend'):
            add_trend_slides(presentation, items.get('campaign_trends', []), scoring_label=label,
                             subtitle=subtitle, new_slide=_slide, scope=scope)
        return
    if stage == 'geographical_areas':
        if _option(options, 'geographical_areas', 'cards'):
            add_location_card_slides(presentation, items.get('location_cards', []), scoring_label=label,
                                     subtitle=subtitle, new_slide=_slide, scope=scope)
        rankings = _first_ranking_group(items.get('location_rankings', []))
        if _option(options, 'geographical_areas', 'ranking'):
            add_location_ranking_slides(presentation, compared_items('ranking-operators', rankings), subtitle=subtitle,
                                        new_slide=_slide, scope=scope, compared=chosen('ranking-operators'))
        if _option(options, 'geographical_areas', 'versus'):
            add_location_versus_slides(presentation, rankings, subtitle=subtitle, new_slide=_slide, scope=scope,
                                       compared=chosen('ranking-operators'))
        return
    styles = operator_styles(views.get('score_tables', []))
    if stage == 'gaps':
        if _option(options, 'gap', 'parts'):
            add_gap_parts_slides(presentation, compared_items('gap-parts-operators', items.get('gap_parts', [])),
                                 subtitle=subtitle, new_slide=_slide, scope=scope, compared=chosen('gap-parts-operators'))
        if _option(options, 'gap', 'profile'):
            # The slides of one campaign show its latest one (of each operator and place).
            profiles = [item for item in latest_campaign_items(items.get('kpi_gap_profiles', []))
                        if chosen('gap-insight-operators')(item['operator'], item.get('label'))]
            add_kpi_gap_profile_slides(
                presentation, profiles, scoring_label=label, subtitle=subtitle,
                environment_labels=environment_labels,
                new_slide=_slide, operator_colors={name: style.get('color') for name, style in styles.items()},
                scope=scope,
            )
        return
    if stage == 'points':
        # The slides of one campaign show its latest one (of each operator and place).
        if _option(options, 'points', 'stake'):
            add_points_at_stake_slides(presentation, latest_campaign_items(items.get('points_at_stake', [])), subtitle=subtitle,
                                       new_slide=_slide, scope=scope)
        if _option(options, 'points', 'kpis'):
            profiles = [item for item in latest_campaign_items(items.get('kpi_gap_profiles', []))
                        if chosen('gap-insight-operators')(item['operator'], item.get('label'))]
            add_points_per_kpi_slides(
                presentation, profiles, scoring_label=label, subtitle=subtitle, environment_labels=environment_labels,
                new_slide=_slide, operator_colors={name: style.get('color') for name, style in styles.items()}, scope=scope,
            )
        if _option(options, 'points', 'map'):
            # The reference operator has its map too when it is one of the operators to compare.
            maps = [item for item in items.get('points_loss_maps', []) if compared(item['operator'], item.get('label'))]
            add_points_loss_slides(presentation, latest_campaign_items(maps), insights.get('points_loss_background') or [],
                                   scoring_label=label, subtitle=subtitle, new_slide=_slide,
                                   boundaries=views.get('_points_loss_boundaries'),
                                   outlines=views.get('_points_loss_outlines'),
                                   environment_labels=environment_labels, scope=scope)
        if _option(options, 'points', 'vendors'):
            add_serving_loss_slides(presentation, compared_items('vendor_losses-operators', items.get('vendor_losses', [])),
                                    title='Points Lost — per Vendor', column='Vendor', subtitle=subtitle, new_slide=_slide,
                                    scope=scope, compared=chosen('vendor_losses-operators'))
        if _option(options, 'points', 'site_types'):
            add_serving_loss_slides(presentation, compared_items('site_type_losses-operators', items.get('site_type_losses', [])),
                                    title='Points Lost — per kind of site', column='Kind of site', subtitle=subtitle,
                                    new_slide=_slide, scope=scope, compared=chosen('site_type_losses-operators'))
        if _option(options, 'points', 'sites'):
            add_sites_slides(presentation, compared_items('sites-operators', latest_campaign_items(items.get('sites_to_review', []))),
                             subtitle=subtitle, new_slide=_slide, scope=scope, compared=chosen('sites-operators'))
        if _option(options, 'points', 'stretches'):
            # The map areas (ITL3 in the UK) as the land under each route, as in the route of Where to act.
            boundaries = views.get('_points_loss_boundaries')
            add_stretch_slides(presentation, latest_campaign_items(items.get('route_stretches', [])),
                               insights.get('points_loss_background') or [],
                               subtitle=subtitle, new_slide=_slide, scope=scope, compared=chosen('stretches-operators'),
                               boundaries=boundaries('Area') if boundaries else None)
        return
    if stage == 'comparison' and _option(options, 'tables', 'campaign_comparison'):
        add_campaign_comparison_slides(presentation, items.get('campaign_comparisons', []), scoring_label=label,
                                       subtitle=subtitle, new_slide=_slide, scope=scope)


def _add_scoring_block(presentation, job: dict[str, Any], result: dict[str, Any],
                       operator_mapping_groups: list[dict[str, Any]] | None, *, scoring: str | None,
                       table_mode: str, gap_layout: str, environment: str, show_gap_values: bool,
                       split_charts: bool, vendor_mapping_groups: list[dict[str, Any]] | None,
                       options: dict | None = None, section_cover: bool = True,
                       environments: list[str] | None = None, level: int = 1) -> None:
    """Add the charts, tables and GAP slides of one scoring; scoring is None when the job has only Best Network.
    Its cover is a transition of ``level``, and each environment one of the level below it.

    ``options`` are the content options of a report scenario: All Environments only or
    a block per environment, and the charts, tables and GAP slides to include.
    ``environments`` keeps only those environments (Combined for All Environments).
    """
    selected_environments = ({'combined' if str(name).casefold() in {'all', 'all environments', 'combined'}
                              else str(name).casefold() for name in environments} if environments else None)
    # The Focus operators of the Operator Maps (see focus_matcher), among the operators of the results.
    focus = focus_matcher(focus_operators(operator_mapping_groups),
                          {row.get('operator') or row.get('Operator') for row in result.get('scoring') or []
                           if isinstance(row, dict)})
    combined_only = (options is not None and options.get('environments') != 'split') or selected_environments == {'combined'}
    if options is not None:
        show_gap_values = _option(options, 'tables', 'gap_values', False)
    show_kpi_values = _option(options, 'tables', 'kpi_values', False)
    views = _export_environment_views(
        build_scoring_views(job, result, operator_mapping_groups, vendor_mapping_groups=vendor_mapping_groups), environment,
    )
    refresh_baseline_warning(result, views, str(job.get('baseline_operator') or 'EE'))
    views['_export_environment'] = 'all' if combined_only else environment
    # The operators each section of the page has enabled, when the export comes from the page (see _add_insight_slides).
    views['_operator_choices'] = job.get('_operator_choices') or {}
    boundary_cache: dict[str, dict] = {}

    def points_loss_boundaries(field: str) -> dict:
        if field not in boundary_cache:
            boundary_cache[field] = map_boundaries(result, field)
        return boundary_cache[field]
    views['_points_loss_boundaries'] = points_loss_boundaries
    views['_points_loss_outlines'] = lambda field: map_background(result, field)
    configuration = job.get('configuration') or result.get('configuration') or {}
    # The slides name the area, operators and vendors of their table, chart or location, as the web titles do.
    scope_filters = job.get('context_filters') or {}
    scope_operators = list(dict.fromkeys(str(operator) for matrix in views.get('score_tables', [])
                                         for operator in matrix.get('operators') or []))
    views['_scope'] = lambda context=None: scoring_scope_labels(context, scope_filters, scope_operators,
                                                                omit_context=True)
    for matrix_key in ('score_tables', 'gap_summary_tables', 'gap_tables',
                       'hierarchy_score_tables', 'hierarchy_gap_tables'):
        for matrix in views.get(matrix_key, []):
            matrix['environment_subtitle'] = ' · '.join((
                _scoring_block_subtitle(scoring, matrix['context']['environment'], configuration),
                *scoring_scope_labels(matrix['context'], scope_filters, scope_operators),
            ))
            matrix['scoring_label'] = SCORING_LABELS[scoring or BEST_NETWORK_SCORING]
    environment_allocations = maximum_allocations_from_configuration(configuration)
    combined_allocations = scaled_environment_allocations(environment_allocations, result.get('environment_scaling'))
    for matrix_key in ('score_tables', 'hierarchy_score_tables'):
        for matrix in views.get(matrix_key, []):
            matrix['_split_charts'] = split_charts
            combined = str(matrix['context'].get('environment') or '').casefold() in {'combined', 'all', 'all environments'}
            matrix['environment_allocations'] = combined_allocations if combined else environment_allocations
            matrix['environment_labels'] = {
                name: str(scope.get('display_name') or _environment_display_label(name))
                for name, scope in configuration.get('scope', {}).get('environments', {}).items()
            }
    environment_level = level
    if scoring is not None and section_cover:
        _add_scoring_intro_slides(presentation, job, result, scoring=scoring, level=level)
        environment_level = level + 1
    # Each slide is named after its section (summary, charts, tables, geographical_areas, gap, points, comparison).
    untagged = [len(presentation.slides)]

    def section(name: str) -> None:
        _tag_slides(presentation, untagged[0], name)
        untagged[0] = len(presentation.slides)
    subtitle = 'All Environments' if environment == 'all' else _scoring_environment_title(environment, configuration)
    matrices = views['score_tables']
    hierarchy_matrices = views.get('hierarchy_score_tables', [])

    def keep(environment_name: str) -> bool:
        combined = environment_name.casefold() in {'all environments', 'combined'}
        if selected_environments is not None:
            return 'combined' in selected_environments if combined else environment_name.casefold() in selected_environments
        return not combined_only or combined

    if selected_environments is not None and not any(
            keep(str(matrix['context'].get('environment') or 'Unspecified'))
            for matrix in (hierarchy_matrices or matrices)):
        raise ValueError('None of the selected environments is available for this Saved Filter.')

    area_summaries = build_area_summary(
        job, result, operator_mapping_groups, vendor_mapping_groups=vendor_mapping_groups,
    )

    def environment_cover(environment_name: str) -> None:
        # With All Environments only, the scoring cover already introduces the block; the Summary follows it.
        if not combined_only:
            _add_scoring_intro_slides(presentation, job, result, environment=environment_name, scoring=scoring,
                                      level=environment_level)
        _add_insight_slides(presentation, views, environment_name, scoring, configuration, stage='summary', options=options,
                            focus=focus)
        section('summary')

    trended: set[str] = set()

    def geographical_areas(environment_name: str) -> None:
        _add_insight_slides(presentation, views, environment_name, scoring, configuration, stage='geographical_areas',
                            options=options, focus=focus)
        section('geographical_areas')
        # The Scoring Trends follow the Geographical Areas: the scoring trend, then the National & Geographical Areas Trends.
        _add_insight_slides(presentation, views, environment_name, scoring, configuration, stage='trends', options=options,
                            focus=focus)
        if environment_name not in trended and _option(options, 'trends', 'areas'):
            trended.add(environment_name)
            add_area_summary_slides(
                presentation, [item for item in area_summaries if item['environment'] == environment_name
                               or (item['environment'] == 'Combined' and environment_name == 'All Environments')],
                scoring_label=SCORING_LABELS[scoring or BEST_NETWORK_SCORING],
                subtitle=_scoring_block_subtitle(scoring, environment_name, configuration), new_slide=_slide,
            )
        section('trends')

    def closing_stages(environment_name: str) -> None:
        for stage in ('gaps', 'points'):
            _add_insight_slides(presentation, views, environment_name, scoring, configuration, stage=stage, options=options,
                                focus=focus)
            section('gap' if stage == 'gaps' else stage)

    def score_tables(matrix: dict, render, environment_name: str) -> None:
        for mode, title, key in (('summary', 'Scoring Tables — Summary', 'summary'),
                                 ('expanded', 'Scoring Tables — Breakdown', 'breakdown')):
            if _option(options, 'tables', key):
                render(presentation, [_table_for_mode(matrix, mode)], views.get('threshold_legend', []),
                       gap_layout=gap_layout, show_gap_values=show_gap_values, title=title,
                       show_kpi_values=show_kpi_values and mode == 'expanded')
        section('tables')
        # The comparison of the campaigns closes the Scoring Tables, as on the page.
        _add_insight_slides(presentation, views, environment_name, scoring, configuration, stage='comparison', options=options,
                            focus=focus)
        section('comparison')

    use_hierarchy = bool(hierarchy_matrices) and len(hierarchy_matrices[0].get('hierarchy_levels', [])) > 1
    if use_hierarchy:
        gap_by_environment = {
            matrix['context'].get('environment'): matrix
            for matrix in views.get('hierarchy_gap_tables', [])
        }
        grouped_hierarchy: dict[str, list[dict[str, Any]]] = {}
        for matrix in hierarchy_matrices:
            grouped_hierarchy.setdefault(str(matrix['context'].get('environment') or 'Unspecified'), []).append(matrix)
        for current_environment, environment_matrices in grouped_hierarchy.items():
            if not keep(current_environment):
                continue
            environment_cover(current_environment)
            for matrix in environment_matrices:
                if _option(options, 'charts', 'service'):
                    _best_network(presentation, [matrix])
                if _option(options, 'charts', 'category'):
                    _hierarchy_chart(presentation, matrix)
                if _option(options, 'charts', 'breakdown'):
                    _hierarchy_category_comparison_chart(presentation, matrix)
                section('charts')
                score_tables(matrix, _hierarchy_score_tables, current_environment)
                geographical_areas(current_environment)
                gap_matrix = gap_by_environment.get(matrix['context'].get('environment'))
                if gap_matrix is not None:
                    # Finish the full comparison block before advancing to the next environment.
                    reference = _operator_label(gap_matrix, gap_matrix['baseline_operator'])
                    # Every operator against the reference; one by one, the Focus operators.
                    compared_columns = [column for column in gap_matrix['hierarchy_columns'] if not column.get('is_reference')]
                    if compared_columns and _option(options, 'gap', 'all_operators'):
                        all_matrix = _hierarchy_gap_projection(_table_for_mode(gap_matrix, 'expanded'), compared_columns)
                        _hierarchy_gap_tables(presentation, [all_matrix], title=f'GAP Analysis — All vs {reference}')
                    if _option(options, 'gap', 'individual'):
                        for operator in dict.fromkeys(column['operator'] for column in compared_columns
                                                      if focus(column['operator'], _operator_label(gap_matrix, column['operator']))):
                            operator_columns = [column for column in gap_matrix['hierarchy_columns']
                                                if column['operator'] == operator]
                            operator_matrix = _hierarchy_gap_projection(_table_for_mode(gap_matrix, 'expanded'),
                                                                        operator_columns)
                            _hierarchy_gap_tables(
                                presentation, [operator_matrix],
                                title=f'GAP Analysis — {operator} vs {reference}', show_priority=True,
                            )
                section('gap')
                closing_stages(current_environment)
    elif matrices:
        grouped_scores: dict[str, list[dict[str, Any]]] = {}
        for matrix in matrices:
            grouped_scores.setdefault(str(matrix['context'].get('environment') or 'Unspecified'), []).append(matrix)
        for current_environment, environment_matrices in grouped_scores.items():
            if not keep(current_environment):
                continue
            environment_cover(current_environment)
            for matrix in environment_matrices:
                if _option(options, 'charts', 'service'):
                    _best_network(presentation, [matrix])
                if _option(options, 'charts', 'category'):
                    _stacked_category_chart(presentation, matrix)
                if _option(options, 'charts', 'breakdown'):
                    _charts(presentation, [matrix])
                section('charts')
                score_tables(matrix, _score_tables, current_environment)
                geographical_areas(current_environment)
                if _option(options, 'gap', 'all_operators'):
                    _gap_summary_tables(presentation, [_table_for_mode(table, 'expanded') for table in views['gap_summary_tables']
                                                       if table['context'] == matrix['context']])
                if _option(options, 'gap', 'individual'):
                    _gap_tables(presentation, [_table_for_mode(table, 'expanded') for table in views['gap_tables']
                                               if table['context'] == matrix['context']
                                               and focus(table.get('operator', ''), _operator_label(table, table.get('operator', '')))])
                section('gap')
            closing_stages(current_environment)
    else:
        slide = _slide(presentation, 'Scoring Tables', subtitle)
        _text(slide, 'No scoring measurements are available for this saved job.', 1.8, height=1, size=16)


def export_scoring_powerpoint(job: dict[str, Any], result: dict[str, Any], template_path: Path,
                             operator_mapping_groups: list[dict[str, Any]] | None = None,
                             *, table_mode: str = 'expanded', gap_layout: str = 'end',
                             environment: str = 'all', show_gap_values: bool = True,
                             split_charts: bool = True,
                             vendor_mapping_groups: list[dict[str, Any]] | None = None,
                             theme: str = DEFAULT_THEME, scorings: list[str] | None = None,
                             environments: list[str] | None = None) -> bytes:
    """Export saved points as comparison matrices, charts and prioritized gaps, in the colours of ``theme``.

    ``scorings`` (best_network, most_reliable) and ``environments`` (Combined for All Environments, or the names of
    the environments) keep only those blocks; without them, every scoring and environment is exported.
    """
    with use_theme(theme):
        return _export_scoring_powerpoint(
            job, result, template_path, operator_mapping_groups, table_mode=table_mode, gap_layout=gap_layout,
            environment=environment, show_gap_values=show_gap_values, split_charts=split_charts,
            vendor_mapping_groups=vendor_mapping_groups, scorings=scorings, environments=environments)


def _export_scoring_powerpoint(job: dict[str, Any], result: dict[str, Any], template_path: Path,
                               operator_mapping_groups: list[dict[str, Any]] | None, *, table_mode: str,
                               gap_layout: str, environment: str, show_gap_values: bool, split_charts: bool,
                               vendor_mapping_groups: list[dict[str, Any]] | None, scorings: list[str] | None = None,
                               environments: list[str] | None = None) -> bytes:
    if not template_path.is_file():
        raise ValueError('The configured CDR PowerPoint template is missing.')
    if table_mode not in {'expanded', 'summary'}:
        raise ValueError('Table mode must be expanded or summary.')
    if gap_layout not in {'end', 'adjacent'}:
        raise ValueError('GAP layout must be end or adjacent.')
    job = dict(job)
    result = normalize_scoring_vendor_result(result, operator_mapping_groups)
    if '_scoring_display_selections' not in job:
        job['_scoring_display_selections'] = prepare_scoring_display_selections(job, result, template_path)
    presentation = Presentation(template_path)
    _remove_all_slides(presentation)
    _add_scoring_intro_slides(presentation, job, result)
    blocks: list[tuple[str | None, dict[str, Any], dict[str, Any]]] = [(None, job, result)]
    reliable = most_reliable_result(result, str(job.get('baseline_operator') or 'EE'), job.get('configuration'))
    if reliable is not None:
        # Both scorings share the cover, filters and aggregation; each one has its own block.
        blocks = [(BEST_NETWORK_SCORING, job, result),
                  (MOST_RELIABLE_SCORING, {**job, 'configuration': reliable['configuration']}, reliable)]
    if scorings:
        if reliable is None and MOST_RELIABLE_SCORING in scorings and BEST_NETWORK_SCORING not in scorings:
            raise ValueError('This Saved Filter has no Most Reliable scoring.')
        if reliable is not None:
            blocks = [block for block in blocks if block[0] in scorings]
    for scoring, block_job, block_result in blocks:
        # The cover keeps the notes of the first scoring.
        first_slide = len(presentation.slides) if block_result is not result else 0
        _add_scoring_block(
            presentation, block_job, block_result, operator_mapping_groups, scoring=scoring,
            table_mode=table_mode, gap_layout=gap_layout, environment=environment,
            show_gap_values=show_gap_values, split_charts=split_charts,
            vendor_mapping_groups=vendor_mapping_groups, environments=environments,
        )
        notes = '\n'.join(_scoring_slide_notes(block_result))
        for index, slide in enumerate(presentation.slides):
            if index >= first_slide:
                slide.notes_slide.notes_text_frame.text = notes
    report_structure.finish(presentation)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def export_scoring_report(entries: list[dict[str, Any]], template_path: Path,
                          operator_mapping_groups: list[dict[str, Any]] | None = None, *,
                          gap_layout: str = 'end', split_charts: bool = True,
                          vendor_mapping_groups: list[dict[str, Any]] | None = None,
                          theme: str = DEFAULT_THEME) -> bytes:
    """One document with the scenarios of a scoring report, in the colours of ``theme``.

    Each entry has a normalized ``scenario`` and the completed ``job`` and ``result`` of
    its filters and aggregation. Every scenario adds, for each scoring it includes, a
    cover with its filters followed by the slides its options choose.
    """
    with use_theme(theme):
        return _export_scoring_report(entries, template_path, operator_mapping_groups, gap_layout=gap_layout,
                                      split_charts=split_charts, vendor_mapping_groups=vendor_mapping_groups)


def _export_scoring_report(entries: list[dict[str, Any]], template_path: Path,
                           operator_mapping_groups: list[dict[str, Any]] | None, *, gap_layout: str,
                           split_charts: bool, vendor_mapping_groups: list[dict[str, Any]] | None) -> bytes:
    if not template_path.is_file():
        raise ValueError('The configured CDR PowerPoint template is missing.')
    if not entries:
        raise ValueError('A scoring report needs at least one scenario.')
    presentation = Presentation(template_path)
    _remove_all_slides(presentation)
    several = len(entries) > 1
    prepared = []
    for entry in entries:
        job = dict(entry['job'])
        result = normalize_scoring_vendor_result(entry['result'], operator_mapping_groups)
        if '_scoring_display_selections' not in job:
            job['_scoring_display_selections'] = prepare_scoring_display_selections(job, result, template_path)
        prepared.append((entry['scenario'], job, result))
    first_scenario, first_job, first_result = prepared[0]
    if several:
        mode = _scoring_filter_subtitle(first_job).split('\n', 1)[0]
        names = ', '.join(scenario['name'] for scenario, _job, _result in prepared)
        _add_scoring_intro_slides(presentation, first_job, first_result,
                                  subtitle=f'{mode}\nScenarios: {names}')
    else:
        _add_scoring_intro_slides(presentation, first_job, first_result)
    for scenario, job, result in prepared:
        if several:
            # Each scenario opens with its transition; its scorings and environments are the levels below it.
            _add_scenario_transition(presentation, scenario['name'], job)
        baseline = str(job.get('baseline_operator') or 'EE')
        blocks = [(BEST_NETWORK_SCORING, job, result)]
        reliable = most_reliable_result(result, baseline, job.get('configuration'))
        if reliable is not None:
            blocks.append((MOST_RELIABLE_SCORING, {**job, 'configuration': reliable['configuration']}, reliable))
        for scoring, block_job, block_result in blocks:
            options = scenario['scorings'][scoring]
            if not options.get('enabled'):
                continue
            if several:
                block_job = {**block_job, '_scenario_name': scenario['name']}
                configuration = dict(block_job.get('configuration') or block_result.get('configuration') or {})
                configuration['_scenario_name'] = scenario['name']
                block_job['configuration'] = configuration
            first_slide = len(presentation.slides)
            _add_scoring_block(
                presentation, block_job, block_result, operator_mapping_groups, scoring=scoring,
                table_mode='expanded', gap_layout=gap_layout, environment='all',
                show_gap_values=False, split_charts=split_charts,
                vendor_mapping_groups=vendor_mapping_groups, options=options,
                environments=options.get('environment_names') or None, level=2 if several else 1,
            )
            notes = '\n'.join(_scoring_slide_notes(block_result))
            for index, slide in enumerate(presentation.slides):
                if index >= first_slide or index == 0:
                    if index == 0 and slide.notes_slide.notes_text_frame.text:
                        continue
                    slide.notes_slide.notes_text_frame.text = notes
    report_structure.finish(presentation)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
