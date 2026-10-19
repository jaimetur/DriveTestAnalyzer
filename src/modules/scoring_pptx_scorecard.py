"""The scorecard design of the Scoring and GAP tables of the PowerPoint and Word exports, in the style of each theme.

As the page draws them (applyScorecardLayout in scoring.js), the tables have no Category, Type of KPI, Score weight or
Max score columns: each category opens a section (its subtotal, when the table has one), the KPI carries the badge of
its type (R or D) with its weight and maximum, the operators are underlined (or filled) in their colour, every score
fills its cell as far as its share of the maximum and every GAP is a bar either side of zero as long as the GAP.
"""

from __future__ import annotations

from copy import deepcopy
from typing import Any

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from src.modules.scoring_themes import active_theme, blend, gap_share, themed

# The cell treatments of each theme: 'bar' fills the share of the maximum, 'solid' tints the cell with its band,
# 'heat' tints it by its share and 'text' writes the score in the colour of its band.
STYLES: dict[str, dict[str, Any]] = {
    # NetCheck: as the NetCheck reports draw their tables (Excel-like data bars fading to white, operator strips).
    'netcheck': {'score': 'excel', 'bar_mix': 1, 'card': '#FFFFFF', 'label': '#FFFFFF', 'ink': '#17232D', 'muted': '#8A96A0',
                 'header': '#FFFFFF', 'header_ink': '#17232D', 'section': '#FFFFFF', 'section_ink': '#17232D',
                 'section_rule': '#C9D2DB', 'total': '#FFFFFF', 'total_ink': '#17232D', 'total_rule': '#17232D',
                 'gap_card': '#FFFFFF', 'gain_bar': '#4CA65A', 'loss_bar': '#E8414F', 'fade': True, 'gain': '#2E8540',
                 'loss': '#C62835', 'reliable': '#1F6FB2', 'diff': '#C25E00', 'best': '#DFF3E4', 'worst': '#FBE1DF',
                 'operators': 'fill'},
    # Umlaut: white tables with thin rules, each score a bar in the colour of its band, operators on strips.
    'umlaut': {'score': 'bar', 'bar_mix': .55, 'card': '#FFFFFF', 'label': '#FFFFFF', 'ink': '#1C1C28', 'muted': '#8C8C99',
               'header': '#FFFFFF', 'header_ink': '#3F3F4C', 'section': '#F4F2F4', 'section_ink': '#1C1C28',
               'section_rule': '#E4E2E6', 'total': '#FFFFFF', 'total_ink': '#1C1C28', 'total_rule': '#1C1C28',
               'gap_card': '#FFFFFF', 'gain_bar': '#93C47D', 'loss_bar': '#E39B94', 'gain': '#2F7D4A', 'loss': '#A6192E',
               'reliable': '#5B7A99', 'diff': '#B86E00', 'best': '#DCEFD6', 'worst': '#F6D5D2', 'operators': 'fill'},
    'ericsson': {'score': 'bar', 'bar_mix': .38, 'card': '#F2F2F2', 'label': '#FFFFFF', 'ink': '#242424', 'muted': '#767676',
                 'header': '#FFFFFF', 'header_ink': '#767676', 'header_label': '#242424', 'section': '#FFFFFF',
                 'section_ink': '#242424', 'section_rule': '#242424', 'total': '#FFFFFF', 'total_ink': '#242424',
                 'total_rule': '#242424', 'gap_card': '#F2F2F2', 'gain_bar': '#B7EDD6', 'loss_bar': '#FFC1C1',
                 'gain': '#0C9B5B', 'loss': '#CC2828', 'reliable': '#0068BF', 'diff': '#CC7007', 'best': '#B7EDD6',
                 'worst': '#FFC1C1', 'operators': 'underline',
                 'gap_marker': True, 'legend_fill': True},
    'vodafone': {'score': 'solid', 'bar_mix': .26, 'card': '#F4F4F4', 'label': '#FFFFFF', 'ink': '#333333', 'muted': '#7E7E7E',
                 'header': '#FFFFFF', 'header_ink': '#4A4D4E', 'section': '#F2F2F2', 'section_ink': '#333333',
                 'total': '#4A4D4E', 'total_ink': '#FFFFFF', 'gap_card': '#F4F4F4', 'gain_bar': '#A3E4ED',
                 'loss_bar': '#F8B8B8', 'gain': '#00697C', 'loss': '#B80000', 'reliable': '#00697C', 'diff': '#B86E00',
                 'best': '#D6F3F7', 'worst': '#FDD9D9', 'operators': 'fill', 'legend_fill': True},
    'midnight': {'score': 'bar', 'bar_mix': .7, 'card': '#131F33', 'label': '#0F1726', 'ink': '#DCE6F2', 'muted': '#7F93B0',
                 'header': '#0F1726', 'header_ink': '#8FB3D9', 'section': '#16233A', 'section_ink': '#8FD3FF',
                 'total': '#0B1220', 'total_ink': '#E6EDF7', 'total_rule': '#3DB9FF', 'gap_card': '#131F33',
                 'gain_bar': '#1E9C78', 'loss_bar': '#B33A55', 'gain': '#2EE6A6', 'loss': '#FF4D6D', 'reliable': '#2EE6A6',
                 'diff': '#FFC145', 'best': '#123A35', 'worst': '#3A1A2A', 'operators': 'underline', 'border': '#0F1726'},
    'paper': {'score': 'text', 'bar_mix': 0, 'card': '#FFFEFB', 'label': '#FFFEFB', 'ink': '#1A1A1A', 'muted': '#8A8376',
              'header': '#FFFEFB', 'header_ink': '#1A1A1A', 'section': '#FFFEFB', 'section_ink': '#1A1A1A',
              'section_rule': '#1A1A1A', 'total': '#FFFEFB', 'total_ink': '#1A1A1A', 'total_rule': '#1A1A1A',
              'gap_card': '#F7F4EE', 'gain_bar': '#B4C7DC', 'loss_bar': '#E8BCB0', 'gain': '#1F4E79', 'loss': '#B23A1E',
              'reliable': '#2F6B3A', 'diff': '#B5561D', 'best': '#FFFEFB', 'worst': '#FFFEFB', 'operators': 'underline',
              'italic': True},
    'heatmap': {'score': 'heat', 'bar_mix': 0, 'card': '#F6F8FA', 'label': '#FFFFFF', 'ink': '#1F2A37', 'muted': '#8A96A0',
                'header': '#FFFFFF', 'header_ink': '#5B6B78', 'section': '#F3F5F7', 'section_ink': '#1F2A37',
                'total': '#EEF1F4', 'total_ink': '#1F2A37', 'gap_card': '#F6F8FA', 'gain_bar': '#A9DCD5',
                'loss_bar': '#F4B9B3', 'gain': '#1F6F66', 'loss': '#A63A33', 'reliable': '#1F6F66', 'diff': '#C06A0A',
                'best': '#D3EDE4', 'worst': '#F9D5D3', 'operators': 'underline'},
}
FIXED_ROLES = {'CATEGORY': 'category', 'KPI': 'kpi', 'NETCHECK KPIS': 'kpi', 'TYPE OF KPI': 'type',
               'SCORE WEIGHT\n(%)': 'weight', 'SCORE WEIGHT (%)': 'weight', 'MAX SCORE': 'maximum'}
# The NetCheck colours of the header cells that are not an operator: groups, levels, GAP and the fixed headers.
_NEUTRAL_HEADERS = ('#455B65', '#DDEBE6', '#F6D55C', '#E6ECFA', '#5B9A44', '#C9483F', '#E6F0F7', '#D7E5EE')


def style() -> dict[str, Any]:
    return STYLES.get(active_theme()['id'], STYLES['netcheck'])


def _hex(color: Any) -> str:
    return '#' + str(color).lstrip('#').upper()


def _rgb(color: str) -> RGBColor:
    return RGBColor.from_string(_hex(color).lstrip('#'))


def _cell_fill(cell) -> str | None:
    try:
        if cell.fill.type is not None:
            return _hex(cell.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _font_size(cell, default: float) -> float:
    for paragraph in cell.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                return run.font.size.pt
        if paragraph.font.size is not None:
            return paragraph.font.size.pt
    return default


def _replace_fill(cell, element) -> None:
    properties = cell._tc.get_or_add_tcPr()
    existing = [child for child in properties if child.tag in {qn('a:solidFill'), qn('a:gradFill'), qn('a:noFill')}]
    if existing:
        existing[0].addprevious(element)
        for child in existing:
            properties.remove(child)
    else:
        properties.append(element)


def solid(cell, color: str) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(color)


# The width of the edge between two colours of a bar. PowerPoint draws the fills of the slides on screen with their stops
# at whole percents (its PDF export does not), so two stops closer than that merge and the colour runs on to the next
# stop: a bar to the left of the middle of a cell would fade to the right of it.
HARD_EDGE = .01


def gradient(cell, stops: list[tuple[float, str]]) -> None:
    """A horizontal fill with hard stops: each colour from its position to the next one, the stops at whole percents
    and at least one percent apart (see HARD_EDGE)."""
    fill = etree.Element(qn('a:gradFill'))
    fill.set('rotWithShape', '1')
    stop_list = etree.SubElement(fill, qn('a:gsLst'))
    previous = -1
    for position, color in stops:
        percent = max(previous + 1, int(round(max(0.0, min(1.0, position)) * 100)))
        if percent > 100:
            continue
        previous = percent
        stop = etree.SubElement(stop_list, qn('a:gs'))
        stop.set('pos', str(percent * 1000))
        etree.SubElement(stop, qn('a:srgbClr')).set('val', _hex(color).lstrip('#'))
    linear = etree.SubElement(fill, qn('a:lin'))
    linear.set('ang', '0')
    linear.set('scaled', '0')
    _replace_fill(cell, fill)


def share_bar(cell, color: str, share: float, rest: str) -> None:
    share = max(0.0, min(1.0, share))
    if share <= 0:
        solid(cell, rest)
    elif share >= 1:
        solid(cell, color)
    else:
        gradient(cell, [(0, color), (share, color), (share + HARD_EDGE, rest), (1, rest)])


def fading_bar(cell, color: str, share: float, rest: str) -> None:
    """An Excel-like data bar: the colour fading towards white along the share of the cell."""
    share = max(0.0, min(1.0, share))
    if share <= 0:
        solid(cell, rest)
        return
    end = blend('#FFFFFF', color, .14)
    stops = [(0, color), (share, end)] + ([(share + HARD_EDGE, rest), (1, rest)] if share < 1 else [])
    gradient(cell, stops)


def diverging_bar(cell, color: str, share: float, rest: str, gain: bool, *, fade: bool = False) -> None:
    """A bar from the middle of the cell: to the right for a gain, to the left for a loss (fading outwards when
    ``fade``)."""
    half = max(0.0, min(1.0, share)) / 2
    if half <= 0:
        solid(cell, rest)
        return
    low, high = (.5, .5 + half) if gain else (.5 - half, .5)
    light = blend('#FFFFFF', color, .14) if fade else color
    inner, outer = (color, light) if gain else (light, color)
    gradient(cell, [(0, rest), (low, rest), (low + HARD_EDGE, inner), (high, outer), (high + HARD_EDGE, rest), (1, rest)])


def border(cell, edge: str, color: str, width_pt: float) -> None:
    """Set one edge of a cell (lnL, lnR, lnT or lnB), keeping the schema order of the cell properties."""
    properties = cell._tc.get_or_add_tcPr()
    order = ('lnL', 'lnR', 'lnT', 'lnB')
    existing = properties.find(qn(f'a:{edge}'))
    if existing is not None:
        properties.remove(existing)
    line = etree.Element(qn(f'a:{edge}'))
    line.set('w', str(int(width_pt * 12700)))
    fill = etree.SubElement(line, qn('a:solidFill'))
    etree.SubElement(fill, qn('a:srgbClr')).set('val', _hex(color).lstrip('#'))
    position = 0
    for name in order[:order.index(edge)]:
        if properties.find(qn(f'a:{name}')) is not None:
            position += 1
    properties.insert(position, line)


def write(cell, runs: list[tuple[str, dict[str, Any]]], *, align=PP_ALIGN.CENTER, font: str | None = None) -> None:
    """Write the runs of a cell: each with its size, colour, bold, italic and underline."""
    frame = cell.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = 1.0
    paragraph.space_before = paragraph.space_after = 0
    for text, options in runs:
        if not text:
            continue
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(options.get('size', 7))
        run.font.bold = bool(options.get('bold'))
        run.font.italic = bool(options.get('italic'))
        if options.get('underline'):
            run.font.underline = True
        if font:
            run.font.name = font
        run.font.color.rgb = _rgb(options.get('color', '#1F2A37'))
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(.04 if align == PP_ALIGN.LEFT else .025)


def drop_columns(table, columns: set[int]) -> None:
    """Remove grid columns from a table, keeping the cells merged across them."""
    for index in sorted(columns, reverse=True):
        grid = table._tbl.tblGrid
        grid.remove(grid.gridCol_lst[index])
        for row in table._tbl.tr_lst:
            cells = row.tc_lst
            cell = cells[index]
            span = int(cell.get('gridSpan', 1))
            if cell.get('hMerge') == '1':
                for left in range(index - 1, -1, -1):
                    origin = cells[left]
                    if origin.get('hMerge') != '1':
                        remaining = int(origin.get('gridSpan', 1)) - 1
                        if remaining > 1:
                            origin.set('gridSpan', str(remaining))
                        else:
                            origin.attrib.pop('gridSpan', None)
                        break
            elif span > 1 and index + 1 < len(cells):
                heir = deepcopy(cell)
                if span - 1 > 1:
                    heir.set('gridSpan', str(span - 1))
                else:
                    heir.attrib.pop('gridSpan', None)
                row.replace(cells[index + 1], heir)
            row.remove(cell)


def drop_row(table, index: int) -> None:
    """Remove a row whose cells are not the origin of a vertical merge, shortening the merges across it."""
    rows = table._tbl.tr_lst
    for column, cell in enumerate(rows[index].tc_lst):
        if cell.get('vMerge') != '1':
            continue
        for above in range(index - 1, -1, -1):
            origin = rows[above].tc_lst[column]
            if origin.get('vMerge') != '1':
                remaining = int(origin.get('rowSpan', 1)) - 1
                if remaining > 1:
                    origin.set('rowSpan', str(remaining))
                else:
                    origin.attrib.pop('rowSpan', None)
                break
    table._tbl.remove(rows[index])


def _band_color(value: dict) -> str | None:
    band = str(value.get('threshold_band') or '')
    strong = active_theme().get('band_strong', {})
    return strong.get(band)


def _band_fill(value: dict) -> str | None:
    """The colour of the band of a score in the legend of the table (the light colour of the theme)."""
    return active_theme().get('bands', {}).get(str(value.get('threshold_band') or ''))


def style_legend_band(cell, band: str, size: float, font: str | None = None) -> None:
    """A band of the legend of a Scoring table drawn as the cells of the band in the active theme: the data bar of
    NetCheck, the bar of Umlaut and Midnight, the heat tint of Heatmap or the coloured, underlined text of Paper (the
    themes whose cells take the colour of the legend keep it; Unavailable keeps its colour)."""
    current = style()
    strong = active_theme().get('band_strong', {}).get(band)
    if strong is None or current.get('legend_fill'):
        return
    mode = current['score']
    run = {'size': size, 'color': current['ink'], 'bold': False}
    if mode == 'excel':
        fading_bar(cell, strong, 1.0, current['card'])
    elif mode == 'bar':
        solid(cell, blend(current['card'], strong, current['bar_mix']))
    elif mode == 'solid':
        solid(cell, blend('#FFFFFF', strong, current['bar_mix']))
    elif mode == 'heat':
        # The tint of a score at half its maximum.
        solid(cell, blend('#FFFFFF', strong, .29))
    else:
        solid(cell, current['card'])
        run.update(color=strong, bold=True, underline=True)
    write(cell, [(cell.text_frame.text, run)], font=font)


def _score_share(value: dict, maximum: Any) -> float | None:
    try:
        points, top = float(value.get('points')), float(maximum)
    except (TypeError, ValueError):
        return None
    return max(0.0, min(1.0, points / top)) if top > 0 else None


def _style_score(cell, value: dict, maximum: Any, size: float, font: str | None, rank: str | None = None) -> None:
    current = style()
    text = cell.text_frame.text
    color = _band_color(value)
    share = _score_share(value, maximum)
    run = {'size': size, 'color': current['ink'], 'bold': False}
    if text.endswith('*'):
        run.update(color='#C62828', bold=True)
    mode = current['score']
    if rank:
        solid(cell, current[rank])
        run['bold'] = True
    elif color is None or share is None:
        solid(cell, current['card'])
    elif mode == 'excel':
        fading_bar(cell, color, share, current['card'])
    elif mode == 'bar':
        # Themes with legend_fill fill the cell with the colour of its band in the legend (Ericsson).
        fill = _band_fill(value) if current.get('legend_fill') else None
        share_bar(cell, fill or blend(current['card'], color, current['bar_mix']), share, current['card'])
    elif mode == 'solid':
        fill = _band_fill(value) if current.get('legend_fill') else None
        solid(cell, fill or blend('#FFFFFF', color, current['bar_mix']))
    elif mode == 'heat':
        solid(cell, blend('#FFFFFF', color, .14 + share * .3))
    else:
        solid(cell, current['card'])
        run.update(color=color, bold=True, underline=True)
    write(cell, [(text, run)], font=font)


def _style_gap(cell, gap: Any, color: Any, size: float, font: str | None) -> None:
    current = style()
    text = cell.text_frame.text
    try:
        number = float(gap)
    except (TypeError, ValueError):
        number = None
    run = {'size': size, 'color': current['muted'], 'bold': False}
    runs = []
    if number is None or number == 0:
        solid(cell, current['gap_card'])
    else:
        gain = number > 0
        diverging_bar(cell, current['gain_bar' if gain else 'loss_bar'], max(gap_share(color), .08),
                      current['gap_card'], gain, fade=bool(current.get('fade')))
        if current.get('gap_marker'):
            # The figure in black after a small ▲ or ▼ in the colour of the GAP.
            runs.append(('▲ ' if gain else '▼ ', {'size': size * .68, 'color': current['gain' if gain else 'loss']}))
            run.update(color=current['ink'], bold=True)
        else:
            run.update(color=current['gain' if gain else 'loss'], bold=True)
    if text.endswith('*'):
        run.update(color='#C62828', bold=True)
    write(cell, [*runs, (text, run)], font=font)


def _label_runs(name: str, *, size: float, kpi_type: str = '', meta: str = '', category: str = '',
                section: bool = False, total: bool = False) -> list[tuple[str, dict[str, Any]]]:
    current = style()
    ink = current['total_ink'] if total else current['section_ink'] if section else current['ink']
    italic = bool(current.get('italic')) and (section or total)
    runs: list[tuple[str, dict[str, Any]]] = []
    if category:
        runs.append((f'{category.upper()} · ', {'size': size * .8, 'color': current['section_ink'], 'bold': True}))
    badge = {'reliable': 'R', 'diff': 'D'}.get(kpi_type.strip().lower())
    if badge and not section and not total:
        runs.append((f'{badge}  ', {'size': size, 'color': current[kpi_type.strip().lower()], 'bold': True}))
    runs.append((name, {'size': size, 'color': ink, 'bold': section or total, 'italic': italic}))
    if meta:
        runs.append((f'   {meta}', {'size': size * .8, 'color': current['total_ink'] if total else current['muted']}))
    return runs


def _meta(metric: dict) -> str:
    from src.modules.scoring_exports import _number
    parts = []
    if metric.get('weight_percent') is not None:
        parts.append(f"{_number(metric['weight_percent'])}%")
    if metric.get('max_points') is not None:
        parts.append(f"{_number(metric['max_points'])} pts")
    return ' · '.join(parts)


def _origin(table, row: int, column: int):
    """The cell that holds the content of a horizontally merged cell."""
    while column > 0 and table.cell(row, column)._tc.get('hMerge') == '1':
        column -= 1
    return table.cell(row, column)


def _style_headers(table, header_rows: int, label_column: int, font: str | None) -> None:
    current = style()
    neutral = {_hex(themed(color)) for color in _NEUTRAL_HEADERS} | {_hex(color) for color in _NEUTRAL_HEADERS}
    for row_index in range(header_rows):
        for column_index, cell in enumerate(table.rows[row_index].cells):
            if cell.is_spanned:
                continue
            text = cell.text_frame.text
            size = _font_size(cell, 7)
            fill = _cell_fill(cell)
            if row_index == 0 and column_index == label_column:
                solid(cell, current['header'])
                write(cell, [({'CATEGORY': 'Category', 'NETCHECK KPIS': 'NetCheck KPIs'}.get(text.strip().upper(), text),
                              {'size': size, 'color': current.get('header_label', current['muted']), 'bold': True})],
                      align=PP_ALIGN.LEFT, font=font)
                continue
            if fill is None or fill in neutral:
                solid(cell, current['header'])
                write(cell, [(text, {'size': size, 'color': current['header_ink'], 'bold': True,
                                     'italic': bool(current.get('italic'))})], font=font)
                continue
            # An operator: underlined in its colour, or filled with it.
            if current['operators'] == 'fill':
                solid(cell, fill)
                red, green, blue = (int(fill[index:index + 2], 16) for index in (1, 3, 5))
                ink = '#17232D' if .299 * red + .587 * green + .114 * blue > 155 else '#FFFFFF'
                write(cell, [(text, {'size': size, 'color': ink, 'bold': True})], font=font)
            else:
                solid(cell, current['header'])
                ink = current.get('header_label') or (
                    current['ink'] if current['header'] != '#FFFFFF' and current['header'] != '#FFFEFB' else fill)
                write(cell, [(text, {'size': size, 'color': ink, 'bold': True,
                                     'italic': bool(current.get('italic'))})], font=font)
                border(cell, 'lnB', fill, 2.25)
                # Under every column of the operator, so the rule shows how far they go: PowerPoint draws the bottom
                # line of merged cells under their first column only, so the cells under them have it on top too.
                spanned = column_index + 1
                while spanned < len(table.columns) and table.cell(row_index, spanned)._tc.get('hMerge') == '1':
                    border(table.cell(row_index, spanned), 'lnB', fill, 2.25)
                    spanned += 1
                if row_index + 1 < len(table.rows):
                    for column in range(column_index, spanned):
                        border(table.cell(row_index + 1, column), 'lnT', fill, 2.25)


def _column_groups(table, header_rows: int, label_column: int) -> tuple[dict[int, tuple[int, bool]], int]:
    """The first column of each group of columns of the header, with the header row of the group and whether it is an
    operator (or the Score and GAP groups above them) or a group below one, such as a route, city or campaign. Read
    before the header is styled, while the operators still have their colours; also returns the row of the operators."""
    neutral = {_hex(themed(color)) for color in _NEUTRAL_HEADERS} | {_hex(color) for color in _NEUTRAL_HEADERS}
    origins = [(row, column, cell.span_width, (_cell_fill(cell) or '') not in neutral | {''})
               for row in range(header_rows) for column, cell in enumerate(table.rows[row].cells)
               if column != label_column and not cell.is_spanned]
    # The first row in the colours of the operators (the leaf header row under them may carry them too).
    operator_row = min((row for row, _column, _span, operator in origins if operator), default=-1)
    starts: dict[int, tuple[int, bool]] = {}
    for row, column, span, _operator in origins:
        if span > 1 or row == operator_row:
            first_row, major = starts.get(column, (row, False))
            starts[column] = (min(first_row, row), major or row <= operator_row)
    return starts, operator_row


def _group_lines(table, starts: dict[int, tuple[int, bool]], operator_row: int, header_rows: int) -> None:
    """Where each group of columns starts, as on the page: a line down the header rows below its label and the body,
    stronger before an operator than before a route, city or campaign, and the label of each group below an operator
    ruled across its columns, so each group shows where it ends."""
    color = style()['muted']
    first = min(starts, default=None)
    rows = len(table.rows)
    for column, (row, major) in starts.items():
        if column == first:
            continue
        width = 1.5 if major else .75
        for below in range(row + 1, rows):
            border(table.cell(below, column), 'lnL', color, width)
            border(table.cell(below, column - 1), 'lnR', color, width)
    for row in range(operator_row + 1, header_rows):
        for column, cell in enumerate(table.rows[row].cells):
            if cell.is_spanned or cell.span_width < 2 or column not in starts:
                continue
            # PowerPoint draws the bottom line of merged cells under their first column only.
            for spanned in range(column, column + cell.span_width):
                border(table.cell(row, spanned), 'lnB', color, 1)
                if row + 1 < rows:
                    border(table.cell(row + 1, spanned), 'lnT', color, 1)


def apply_scorecard(table, *, header_rows: int, metrics: list[dict], value_columns: list[tuple[int, str, str | None]],
                    total: dict | None = None, total_label: str = '', gap_column: str = 'gaps',
                    drop_leaf_header: bool = False, font: str | None = None, label_width: float | None = None) -> int:
    """Give a Scoring or GAP table of the exports the scorecard design of the active theme.

    ``metrics`` are the rows under the header, in order (the total, when ``total`` is given, is the last row);
    ``value_columns`` are (column, 'Score' | 'GAP' | 'Value', leaf id or None for a single GAP) in the columns of
    the table as built. Returns the number of header rows left.
    """
    current = style()
    headers = [cell.text_frame.text.strip().upper() for cell in table.rows[0].cells]
    roles = {index: FIXED_ROLES[text] for index, text in enumerate(headers) if text in FIXED_ROLES}
    label_column = next((index for index, role in roles.items() if role == 'kpi'), None)
    if label_column is None:
        label_column = next((index for index, role in roles.items() if role == 'category'), 0)
    dropped = {index for index, role in roles.items() if index != label_column}
    label_sizes = [_font_size(table.cell(header_rows + offset, label_column), 7) for offset in range(len(metrics))]
    value_sizes = [_font_size(table.cell(header_rows + offset, value_columns[0][0]), 7) if value_columns else 7
                   for offset in range(len(metrics) + (1 if total is not None else 0))]
    total_index = len(table.rows) - 1
    best, worst = _hex(themed('#C6EFCE')), _hex(themed('#FFC7CE'))
    ranks = {(row, column): ('best' if _cell_fill(table.cell(row, column)) == best else
                             'worst' if _cell_fill(table.cell(row, column)) == worst else None)
             for row in range(header_rows, len(table.rows)) for column, kind, _leaf in value_columns if kind == 'Score'}
    width = sum(column.width for column in table.columns)

    # The sections: contiguous categories open a section; KPIs ordered by GAP carry their category instead.
    categories = [str(metric.get('category') or '') for metric in metrics]
    runs: list[tuple[str, int, int]] = []
    for offset, category in enumerate(categories):
        if runs and runs[-1][0] == category:
            runs[-1] = (category, runs[-1][1], offset)
        else:
            runs.append((category, offset, offset))
    summary = all(metric.get('row_type') == 'category' for metric in metrics) if metrics else False
    grouped = not summary and len({run[0] for run in runs}) == len(runs)
    has_subtotals = any(metric.get('row_type') == 'category' for metric in metrics)
    added = 0 if not grouped or has_subtotals else len(runs)
    body_rows = len(metrics) + (1 if total is not None else 0)
    factor = body_rows / (body_rows + added) if added else 1.0

    # Rows, before the columns go: each label and value in the style of the theme.
    value_index = {column: (kind, leaf) for column, kind, leaf in value_columns}
    for offset, metric in enumerate(metrics):
        row_index = header_rows + offset
        subtotal = metric.get('row_type') == 'category' and not summary
        label = table.cell(row_index, label_column)
        name = metric.get('category') if subtotal else str(metric.get('kpi') or label.text_frame.text)
        size = max(4.5, label_sizes[offset] * factor)
        solid(label, current['section'] if subtotal else current['label'])
        write(label, _label_runs(str(name), size=size, kpi_type='' if summary else str(metric.get('kpi_type') or ''),
                                 meta=_meta(metric), section=subtotal,
                                 category='' if grouped or summary else categories[offset]),
              align=PP_ALIGN.LEFT, font=font)
        for column, (kind, leaf) in value_index.items():
            cell = table.cell(row_index, column)
            number_size = max(4.5, value_sizes[offset] * factor)
            if subtotal:
                solid(cell, current['section'])
                write(cell, [(cell.text_frame.text, {'size': number_size, 'color': current['section_ink'], 'bold': True})],
                      font=font)
            elif kind == 'Score':
                _style_score(cell, (metric.get('values') or {}).get(leaf) or {}, metric.get('max_points'),
                             number_size, font, rank=ranks.get((row_index, column)))
            elif kind == 'GAP':
                gap = (metric.get(gap_column) or {}).get(leaf) if leaf is not None else metric.get('gap_points')
                color = ((metric.get('gap_colors') or {}).get(leaf) if leaf is not None else metric.get('gap_color'))
                _style_gap(cell, gap, color, number_size, font)
            else:
                solid(cell, current['label'])
                write(cell, [(cell.text_frame.text, {'size': number_size, 'color': current['muted']})], font=font)
        if subtotal and current.get('section_rule'):
            for cell in table.rows[row_index].cells:
                border(cell, 'lnT', current['section_rule'], 1.5)
    if total is not None:
        size = max(4.5, (label_sizes[-1] if label_sizes else 7) * factor)
        label = _origin(table, total_index, label_column)
        solid(label, current['total'])
        write(label, _label_runs(total_label or 'Total', size=size, meta=_meta(total), total=True), align=PP_ALIGN.LEFT,
              font=font)
        for column, (kind, _leaf) in value_index.items():
            cell = table.cell(total_index, column)
            rank = ranks.get((total_index, column))
            solid(cell, current[rank] if rank else current['total'])
            text = cell.text_frame.text
            write(cell, [(text, {'size': max(4.5, value_sizes[-1] * factor), 'bold': True,
                                 'color': '#C62828' if text.endswith('*') else current['ink'] if rank else current['total_ink']})],
                  font=font)
        if current.get('total_rule'):
            for cell in table.rows[total_index].cells:
                border(cell, 'lnT', current['total_rule'], 2)
    groups, operator_row = _column_groups(table, header_rows, label_column)
    _style_headers(table, header_rows, label_column, font)
    _group_lines(table, groups, operator_row, header_rows)

    # The columns: Category, Type of KPI, Score weight and Max score go; the KPI keeps room for its badge and meta.
    drop_columns(table, dropped)
    if drop_leaf_header and header_rows > 1:
        drop_row(table, header_rows - 1)
        header_rows -= 1
    label_after = label_column - sum(1 for index in dropped if index < label_column)
    columns = list(table.columns)
    wanted = Inches(label_width) if label_width else int(width * .26)
    others = max(1, len(columns) - 1)
    for index, column in enumerate(columns):
        column.width = int(wanted if index == label_after else (width - wanted) / others)
    # The table keeps its width exactly, so what lines up with its right edge still does.
    columns[-1].width = columns[-1].width + width - sum(column.width for column in columns)

    # The sections: subtotals move above their KPIs; tables without subtotals get a row with the category.
    if grouped:
        rows = table._tbl.tr_lst
        body = rows[header_rows:header_rows + len(metrics)]
        heights = [int(row.get('h', 0)) for row in body]
        for category, first, last in runs:
            run_rows = body[first:last + 1]
            subtotal = next((row for row, metric in zip(run_rows, metrics[first:last + 1])
                             if metric.get('row_type') == 'category'), None)
            if subtotal is not None:
                run_rows[0].addprevious(subtotal)
                continue
            section = deepcopy(run_rows[0])
            section.set('h', str(heights[first]))
            run_rows[0].addprevious(section)
        for row in table._tbl.tr_lst[header_rows:]:
            if row.get('h'):
                row.set('h', str(int(int(row.get('h')) * factor)))
        if added:
            index = header_rows
            for category, first, last in runs:
                section_cells = table.rows[index].cells
                kpis = last - first + 1
                for column, cell in enumerate(section_cells):
                    solid(cell, current['section'])
                    if column == label_after:
                        write(cell, _label_runs(category or 'Unclassified', size=max(4.5, label_sizes[first] * factor),
                                                meta=f"{kpis} KPI{'s' if kpis != 1 else ''}", section=True),
                              align=PP_ALIGN.LEFT, font=font)
                    else:
                        write(cell, [], font=font)
                    if current.get('section_rule'):
                        border(cell, 'lnT', current['section_rule'], 1.5)
                index += kpis + 1
    if current.get('border'):
        for row in table.rows:
            for cell in row.cells:
                for edge in ('lnL', 'lnR', 'lnT', 'lnB'):
                    properties = cell._tc.get_or_add_tcPr()
                    line = properties.find(qn(f'a:{edge}'))
                    if line is not None and line.find(qn('a:solidFill')) is not None and \
                            line.find(qn('a:solidFill'))[0].get('val') == 'FFFFFF':
                        line.find(qn('a:solidFill'))[0].set('val', current['border'].lstrip('#'))
    return header_rows
