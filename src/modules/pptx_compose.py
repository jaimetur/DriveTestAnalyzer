"""One PowerPoint from the slides of several decks built on the same template, such as a Complete Report.

Slides are copied with everything they use (pictures, charts with their workbooks, media) and keep their layout:
the layout of the same name in the destination template, or a copy of it when the template has none.

A component is a chart or table of a module: the slides with the same title whose titles end in a page number
(`· 1/3`, `· Page 2/4`, `· KPIs 1/2`) are one component, and `merge_onto_slide` puts their content on one slide,
side by side or one under the other and as small as needed, so a component never spans several slides. The pages of a
table become one table (`join_table_pages`), which `fit_table` shows in up to four tables side by side when it is long.
Small consecutive components share a slide in the same way, each one under its own caption.
"""
from __future__ import annotations

import copy
import math
import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import PartFactory, _Relationship
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Inches, Pt

from src.modules.slide_titles import text_width_pt

SECTION_PREFIX = 'section:'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P_NS = 'http://schemas.openxmlformats.org/presentationml/2006/main'
C_NS = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
# Layouts of covers (a report, section or block cover) and of the closing slide of the template.
COVER_LAYOUTS = {'title page', 'black title page', 'transition'}
CLOSING_LAYOUTS = {'black logo end slide'}
# The page number at the end of the title of a slide that continues a chart or table on the next slides.
PAGE_SUFFIX = re.compile(r'(?:\s*·\s*(?:Part\s+|Page\s+|KPIs\s+)?\d+\s*/\s*\d+)+\s*$', re.IGNORECASE)
# Font sizes (hundredths of a point) never go below this when content is made smaller.
MINIMUM_FONT = 500
_TEXT_PROPERTIES = (f'{{{A_NS}}}rPr', f'{{{A_NS}}}defRPr', f'{{{A_NS}}}endParaRPr')
_SHAPE_TAGS = {f'{{{P_NS}}}{name}' for name in ('sp', 'pic', 'graphicFrame', 'grpSp', 'cxnSp')}
_SKIPPED_RELATIONSHIPS = {RT.NOTES_SLIDE, RT.SLIDE, RT.SLIDE_MASTER, RT.NOTES_MASTER, RT.HANDOUT_MASTER}


def layout_name(slide: Any) -> str:
    return str(slide.slide_layout.name or '').strip().casefold()


def is_cover(slide: Any) -> bool:
    """A cover or a transition: a slide of a cover layout, or a title alone on Title Only (the transitions of the PPT
    Templates, such as Coverage (RSRP) in a Dashboard)."""
    name = layout_name(slide)
    return name in COVER_LAYOUTS or (name == 'title only' and not content_shapes(slide))


def is_closing(slide: Any) -> bool:
    return layout_name(slide) in CLOSING_LAYOUTS


def slide_tag(slide: Any) -> str:
    """The section a module named the slide after (``section:<name>``), or ''."""
    name = str(slide._element.cSld.get('name') or '')
    return name[len(SECTION_PREFIX):] if name.startswith(SECTION_PREFIX) else ''


def title_shape(slide: Any) -> Any:
    for shape in slide.placeholders:
        if shape.placeholder_format.type in {1, 3} and shape.has_text_frame:
            return shape
    return None


def title_parts(slide: Any) -> tuple[str, str]:
    """The title of a slide and the subtitle written under it (the other paragraphs of the title placeholder)."""
    shape = title_shape(slide)
    if shape is None:
        return '', ''
    paragraphs = [' '.join(''.join(run.text for run in paragraph.runs).split()) for paragraph in shape.text_frame.paragraphs]
    paragraphs = [text for text in paragraphs if text]
    return (paragraphs[0] if paragraphs else ''), ' · '.join(paragraphs[1:])


def base_title(text: str) -> str:
    return PAGE_SUFFIX.sub('', text).strip()


def is_continued(slide: Any) -> bool:
    """Whether the slide is one page of a chart or table spread over several slides."""
    title, subtitle = title_parts(slide)
    return bool(PAGE_SUFFIX.search(title) or PAGE_SUFFIX.search(subtitle))


def content_shapes(slide: Any) -> list[Any]:
    """The shapes of a slide besides its title: tables, charts, pictures, text and drawings (empty placeholders left
    out)."""
    title = title_shape(slide)
    shapes = []
    for shape in slide.shapes:
        if title is not None and shape.shape_id == title.shape_id:
            continue
        if shape.is_placeholder and shape.has_text_frame and not shape.text_frame.text.strip():
            continue
        if shape.width is None or shape.height is None or shape.left is None or shape.top is None:
            continue
        shapes.append(shape)
    return shapes


def bounds(shapes: list[Any]) -> tuple[int, int, int, int] | None:
    """(left, top, right, bottom) of the shapes in EMU."""
    if not shapes:
        return None
    return (min(shape.left for shape in shapes), min(shape.top for shape in shapes),
            max(shape.left + shape.width for shape in shapes), max(shape.top + shape.height for shape in shapes))


def content_frame(slide: Any, slide_width: int, slide_height: int, items_bounds: list[tuple[int, int, int, int]]) -> tuple[int, int, int, int]:
    """Where merged content goes: from the top of the content (under the title) to the bottom of the content area of
    the layout, across the width of the slide less its margins."""
    margin = Inches(0.45)
    placeholders = [shape for shape in slide.slide_layout.placeholders if shape.placeholder_format.type not in {1, 3, 4}
                    and shape.top is not None and shape.top > Inches(1)]
    title = title_shape(slide)
    title_bottom = (title.top + title.height) if title is not None and title.top is not None else Inches(1.6)
    # The content starts where the modules start it under their title.
    top = min([item[1] for item in items_bounds] or [title_bottom])
    bottoms = [item[3] for item in items_bounds] + [shape.top + shape.height for shape in placeholders]
    bottom = min(slide_height - Inches(0.3), max(bottoms + [slide_height - Inches(0.6)]))
    return margin, top, slide_width - 2 * margin, max(Inches(1), bottom - top)


# -- copying slides between decks ----------------------------------------------
class DeckComposer:
    """A deck from a template (without its slides) that slides of other decks are copied into."""

    def __init__(self, template: Any, *, slide_height: int | None = None):
        self.presentation = Presentation(template)
        slide_ids = self.presentation.slides._sldIdLst
        for slide_id in list(slide_ids):
            self.presentation.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
        if slide_height:
            self.presentation.slide_height = slide_height
        self.package = self.presentation.part.package
        self._names = {str(part.partname) for part in self.package.iter_parts()}
        self._layouts: dict[int, Any] = {}

    @property
    def width(self) -> int:
        return int(self.presentation.slide_width)

    @property
    def height(self) -> int:
        return int(self.presentation.slide_height)

    def _partname(self, source: str) -> PackURI:
        template = re.sub(r'\d*(\.\w+)$', r'%d\1', source)
        number = 1
        while template % number in self._names:
            number += 1
        self._names.add(template % number)
        return PackURI(template % number)

    @staticmethod
    def _relate(part: Any, rId: str, reltype: str, target: Any, external: bool = False) -> None:
        part.rels._rels[rId] = _Relationship(part.rels._base_uri, rId, reltype,
                                             RTM.EXTERNAL if external else RTM.INTERNAL, target)

    def _clone_part(self, part: Any, cache: dict[int, Any]) -> Any:
        """A copy of ``part`` and of the parts it uses in this deck; pictures are shared when identical."""
        if id(part) in cache:
            return cache[id(part)]
        if str(part.content_type).startswith('image/'):
            try:
                clone = self.package.get_or_add_image_part(BytesIO(part.blob))
            except Exception:  # noqa: BLE001 - pictures Pillow cannot read (such as the WMF logos) are copied as they are.
                clone = None
            if clone is not None:
                self._names.add(str(clone.partname))
                cache[id(part)] = clone
                return clone
        clone = PartFactory(self._partname(str(part.partname)), part.content_type, self.package, part.blob)
        cache[id(part)] = clone
        for rId, relationship in part.rels.items():
            if relationship.is_external:
                self._relate(clone, rId, relationship.reltype, relationship.target_ref, external=True)
            elif relationship.reltype not in _SKIPPED_RELATIONSHIPS:
                self._relate(clone, rId, relationship.reltype, self._clone_part(relationship.target_part, cache))
        return clone

    def _layout(self, source_layout: Any) -> Any:
        """The layout of this deck for a layout of another deck: the same one of the template, else one with its name,
        else a copy of it."""
        key = id(source_layout.part)
        if key in self._layouts:
            return self._layouts[key]
        name = str(source_layout.name or '').strip().casefold()
        partname = str(source_layout.part.partname)
        candidates = [layout for master in self.presentation.slide_masters for layout in master.slide_layouts
                      if str(layout.name or '').strip().casefold() == name]
        layout = next((item for item in candidates if str(item.part.partname) == partname), None) or (candidates[0] if candidates else None)
        if layout is None:
            layout = self._copy_layout(source_layout)
        self._layouts[key] = layout
        return layout

    def _copy_layout(self, source_layout: Any) -> Any:
        master = self.presentation.slide_masters[0]
        clone = PartFactory(self._partname(str(source_layout.part.partname)), source_layout.part.content_type, self.package,
                            source_layout.part.blob)
        cache: dict[int, Any] = {}
        for rId, relationship in source_layout.part.rels.items():
            if relationship.reltype == RT.SLIDE_MASTER:
                self._relate(clone, rId, relationship.reltype, master.part)
            elif relationship.is_external:
                self._relate(clone, rId, relationship.reltype, relationship.target_ref, external=True)
            elif relationship.reltype not in _SKIPPED_RELATIONSHIPS:
                self._relate(clone, rId, relationship.reltype, self._clone_part(relationship.target_part, cache))
        rId = master.part.relate_to(clone, RT.SLIDE_LAYOUT)
        used = [int(item.get('id')) for item in self.presentation.part._element.iter(f'{{{P_NS}}}sldMasterId')]
        used += [int(item.get('id')) for other in self.presentation.slide_masters
                 for item in other._element.iter(f'{{{P_NS}}}sldLayoutId')]
        layout_list = master._element.find(f'{{{P_NS}}}sldLayoutIdLst')
        entry = layout_list.makeelement(f'{{{P_NS}}}sldLayoutId', {'id': str(max(used + [2147483648]) + 1), f'{{{R_NS}}}id': rId})
        layout_list.append(entry)
        return clone.slide_layout

    def add_slide(self, source: Any, source_size: tuple[int, int] | None = None) -> Any:
        """A copy of ``source`` (a slide of another deck) at the end of this deck, with its speaker notes."""
        cache: dict[int, Any] = {}
        clone = PartFactory(self._partname('/ppt/slides/slide1.xml'), source.part.content_type, self.package, source.part.blob)
        # The slide joins the deck first, so the pictures it shares with other slides are found and named after them.
        rId = self.presentation.part.relate_to(clone, RT.SLIDE)
        self.presentation.slides._sldIdLst.add_sldId(rId)
        dropped = set()
        for rId, relationship in source.part.rels.items():
            if relationship.reltype == RT.SLIDE_LAYOUT:
                self._relate(clone, rId, relationship.reltype, self._layout(source.slide_layout).part)
            elif relationship.is_external:
                self._relate(clone, rId, relationship.reltype, relationship.target_ref, external=True)
            elif relationship.reltype in _SKIPPED_RELATIONSHIPS:
                dropped.add(rId)
            else:
                self._relate(clone, rId, relationship.reltype, self._clone_part(relationship.target_part, cache))
        slide = clone.slide
        # Links to other slides of the source deck have no target here.
        for element in list(slide._element.iter(f'{{{A_NS}}}hlinkClick', f'{{{A_NS}}}hlinkHover')):
            if element.get(f'{{{R_NS}}}id') in dropped:
                element.getparent().remove(element)
        if source.has_notes_slide and source.notes_slide.notes_text_frame is not None:
            text = source.notes_slide.notes_text_frame.text
            if text.strip():
                slide.notes_slide.notes_text_frame.text = text
        if source_size and source_size != (self.width, int(self.presentation.slide_height)):
            factor = min(self.width / source_size[0], self.presentation.slide_height / source_size[1])
            if abs(factor - 1) > 0.01:
                for shape in list(slide.shapes):
                    scale_shape(slide, shape, factor, shape.left * factor if shape.left is not None else None,
                                shape.top * factor if shape.top is not None else None)
        return slide

    def add_layout_slide(self, name: str) -> Any:
        layout = next((layout for master in self.presentation.slide_masters for layout in master.slide_layouts
                       if str(layout.name or '').strip().casefold() == name.casefold()), None)
        if layout is None:
            layout = self.presentation.slide_layouts[0]
        slide = self.presentation.slides.add_slide(layout)
        # python-pptx names a new slide after the number of slides, which may be the name of a copied slide.
        slide.part.partname = self._partname('/ppt/slides/slide1.xml')
        return slide

    def remove_slide(self, slide: Any) -> None:
        slide_ids = self.presentation.slides._sldIdLst
        for slide_id in list(slide_ids):
            if self.presentation.part.related_part(slide_id.rId) is slide.part:
                self.presentation.part.drop_rel(slide_id.rId)
                slide_ids.remove(slide_id)
                return

    def save(self) -> bytes:
        output = BytesIO()
        self.presentation.save(output)
        return output.getvalue()


# -- scaling and moving content -------------------------------------------------
def _scale_fonts(element: Any, factor: float) -> None:
    if abs(factor - 1) < 0.005:
        return
    for node in element.iter(*_TEXT_PROPERTIES):
        size = node.get('sz')
        if size and size.isdigit():
            node.set('sz', str(max(min(int(size), MINIMUM_FONT), int(round(int(size) * factor)))))


def _scale_table(element: Any, factor: float) -> None:
    for column in element.iter(f'{{{A_NS}}}gridCol'):
        if column.get('w'):
            column.set('w', str(int(int(column.get('w')) * factor)))
    for row in element.iter(f'{{{A_NS}}}tr'):
        if row.get('h'):
            row.set('h', str(int(int(row.get('h')) * factor)))
    for cell in element.iter(f'{{{A_NS}}}tcPr'):
        for side in ('marL', 'marR', 'marT', 'marB'):
            if cell.get(side):
                cell.set(side, str(int(int(cell.get(side)) * factor)))


def _scale_chart_fonts(slide: Any, element: Any, factor: float) -> None:
    for chart in element.iter(f'{{{C_NS}}}chart'):
        rId = chart.get(f'{{{R_NS}}}id')
        if rId and rId in slide.part.rels:
            _scale_fonts(slide.part.related_part(rId)._element, factor)


def scale_shape(slide: Any, shape: Any, factor: float, left: float | None, top: float | None) -> None:
    """Move ``shape`` to (left, top) at ``factor`` of its size, with its text, table and chart fonts."""
    width, height = shape.width, shape.height
    if left is not None:
        shape.left = int(left)
    if top is not None:
        shape.top = int(top)
    if width is not None and height is not None:
        shape.width, shape.height = int(width * factor), int(height * factor)
    element = shape._element
    _scale_fonts(element, factor)
    if element.find(f'.//{{{A_NS}}}tbl') is not None:
        _scale_table(element, factor)
    _scale_chart_fonts(slide, element, factor)


def _copy_shape(source_slide: Any, shape: Any, target_slide: Any) -> Any:
    """A copy of ``shape`` on ``target_slide`` with its geometry written out and the pictures and charts it uses."""
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    element = copy.deepcopy(shape._element)
    # Copied placeholders become plain shapes: the target layout may not have them.
    for placeholder in list(element.iter(f'{{{P_NS}}}ph')):
        placeholder.getparent().remove(placeholder)
    mapping = {}
    for node in element.iter():
        for attribute, value in list(node.attrib.items()):
            if not attribute.startswith(f'{{{R_NS}}}') or value not in source_slide.part.rels:
                continue
            if value not in mapping:
                relationship = source_slide.part.rels[value]
                mapping[value] = (target_slide.part.rels.get_or_add_ext_rel(relationship.reltype, relationship.target_ref)
                                  if relationship.is_external else
                                  target_slide.part.rels.get_or_add(relationship.reltype, relationship.target_part))
            node.set(attribute, mapping[value])
    # Every shape of a slide (those inside a group too) has its own ID.
    shape_id = target_slide.shapes._next_shape_id
    for properties in element.iter(f'{{{P_NS}}}cNvPr'):
        properties.set('id', str(shape_id))
        shape_id += 1
    target_slide.shapes._spTree.append(element)
    copied = target_slide.shapes[-1]
    copied.left, copied.top, copied.width, copied.height = left, top, width, height
    return copied


@dataclass
class Block:
    """Content to place: shapes of a slide and the caption written above them (a title and a subtitle)."""
    slide: Any
    shapes: list[Any]
    caption: tuple[str, str] = ('', '')
    box: tuple[int, int, int, int] = field(init=False)

    def __post_init__(self) -> None:
        self.box = bounds(self.shapes) or (0, 0, 1, 1)

    @property
    def width(self) -> int:
        return max(1, self.box[2] - self.box[0])

    @property
    def height(self) -> int:
        return max(1, self.box[3] - self.box[1])

    @property
    def caption_height(self) -> int:
        title, subtitle = self.caption
        return (Inches(0.34) if title else 0) + (Inches(0.24) if subtitle else 0)


@dataclass
class Arrangement:
    scale: float
    columns: int
    # (left, top) of the content of each block and of its caption.
    places: list[tuple[int, int, int, int]]


def arrange(blocks: list[Block], frame: tuple[int, int, int, int], gap: int = Inches(0.25)) -> Arrangement:
    """The grid (one to four columns, blocks in reading order) that shows the blocks largest in ``frame``."""
    left, top, width, height = frame
    best: Arrangement | None = None
    count = len(blocks)
    for columns in range(1, min(count, 4) + 1):
        rows = math.ceil(count / columns)
        cell = (width - gap * (columns - 1)) / columns
        grid = [blocks[index:index + columns] for index in range(0, count, columns)]
        captions = sum(max(block.caption_height for block in row) for row in grid)
        content = sum(max(block.height for block in row) for row in grid)
        available = height - gap * (rows - 1) - captions
        if available <= 0:
            continue
        scale = min(1.0, available / content, *(cell / block.width for block in blocks))
        if best is not None and scale <= best.scale * 1.02:
            continue
        places = []
        cursor = top
        for row in grid:
            caption = max(block.caption_height for block in row)
            for column, block in enumerate(row):
                x = left + column * (cell + gap)
                places.append((int(x), int(cursor + caption), int(x), int(cursor)))
            cursor += caption + max(block.height for block in row) * scale + gap
        best = Arrangement(scale, columns, places)
    return best or Arrangement(min(1.0, height / max(block.height for block in blocks)), 1,
                               [(left, top, left, top) for _block in blocks])


def _caption(slide: Any, left: int, top: int, width: int, title: str, subtitle: str) -> None:
    if not title and not subtitle:
        return
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(Inches(0.3)))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    first = True
    for text, size, bold, colour in ((title, 13, True, RGBColor(0x1F, 0x24, 0x33)), (subtitle, 9, False, RGBColor(0x5B, 0x63, 0x77))):
        if not text:
            continue
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        run = paragraph.add_run()
        run.text = text
        run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, colour
    box.height = Emu(Inches(0.3 if not subtitle else 0.52))


def set_title(slide: Any, title: str, subtitle: str | None = None) -> None:
    """Write ``title`` (and ``subtitle``, when given) in the title placeholder keeping its formatting."""
    shape = title_shape(slide)
    if shape is None:
        return
    paragraphs = shape.text_frame.paragraphs

    def write(paragraph: Any, text: str) -> None:
        runs = paragraph.runs
        if runs:
            runs[0].text = text
            for run in runs[1:]:
                run._r.getparent().remove(run._r)
        else:
            paragraph.text = text

    write(paragraphs[0], title)
    if subtitle is None:
        return
    if len(paragraphs) > 1:
        write(paragraphs[1], subtitle)
        for paragraph in paragraphs[2:]:
            paragraph._p.getparent().remove(paragraph._p)
        if not subtitle:
            paragraphs[1]._p.getparent().remove(paragraphs[1]._p)


def merge_onto_slide(composer: DeckComposer, target: Any, blocks: list[Block], *, captions: bool,
                     frame: tuple[int, int, int, int] | None = None) -> float:
    """Put the content of ``blocks`` on ``target`` as large as it fits, replacing the content it had, and return the
    scale it needed. ``captions`` writes the title of each block above its content."""
    if not captions:
        blocks = [Block(block.slide, block.shapes) for block in blocks]
    frame = frame or content_frame(target, composer.width, composer.height, [block.box for block in blocks])
    arrangement = arrange(blocks, frame)
    originals = content_shapes(target)
    cell_width = int((frame[2] - Inches(0.25) * (arrangement.columns - 1)) / arrangement.columns)
    for block, (x, y, caption_x, caption_y) in zip(blocks, arrangement.places):
        if captions:
            _caption(target, caption_x, caption_y, cell_width, *block.caption)
        for shape in block.shapes:
            copied = _copy_shape(block.slide, shape, target)
            scale_shape(target, copied, arrangement.scale,
                        x + (shape.left - block.box[0]) * arrangement.scale, y + (shape.top - block.box[1]) * arrangement.scale)
    for shape in originals:
        shape._element.getparent().remove(shape._element)
    return arrangement.scale


def _row_texts(row: Any) -> list[str]:
    return [' '.join(''.join(node.text or '' for node in cell.iter(f'{{{A_NS}}}t')).split()) for cell in row.iter(f'{{{A_NS}}}tc')]


def _table_rows(table: Any) -> list[Any]:
    return table._element.find(f'.//{{{A_NS}}}tbl').findall(f'{{{A_NS}}}tr')


def join_table_pages(composer: DeckComposer, pages: list[Any]) -> bool:
    """Join the pages of a table into the table of the first page: when every page has one table and the same other
    shapes (such as a legend), the rows of the next pages with the same columns (without their header) are added to
    it, and the pages of a table too wide for one slide (parts with other columns, each starting with the same first
    column) add their columns to it; the shapes under it move down. Returns whether the pages were joined; the other
    pages are then removed."""
    tables, extras = [], []
    for page in pages:
        shapes = content_shapes(page)
        found = [shape for shape in shapes if shape.has_table]
        if len(found) != 1:
            return False
        tables.append(found[0])
        extras.append(sorted(' '.join(shape.text_frame.text.split()) for shape in shapes
                             if shape is not found[0] and shape.has_text_frame))
    if any(item != extras[0] for item in extras):
        return False
    # Consecutive pages with the same header are one part: the rows of a part continue over its pages.
    parts: list[list[Any]] = []
    for table in tables:
        if parts and _row_texts(_table_rows(table)[0]) == _row_texts(_table_rows(parts[-1][0])[0]):
            parts[-1].append(table)
        else:
            parts.append([table])
    joined: list[list[Any]] = []
    for part in parts:
        first_rows = _table_rows(part[0])
        rows = list(first_rows)
        for table in part[1:]:
            page_rows = _table_rows(table)
            header = 0
            while header < min(len(page_rows), len(first_rows)) - 1 and _row_texts(page_rows[header]) == _row_texts(first_rows[header]):
                header += 1
            rows.extend(page_rows[header:])
        joined.append(rows)
    # The parts of a wide table show the same rows: the same count and the same first column.
    if len(parts) > 1 and any(len(rows) != len(joined[0]) or [_row_texts(row)[:1] for row in rows[1:]] != [_row_texts(row)[:1] for row in joined[0][1:]]
                              for rows in joined[1:]):
        return False
    shape = tables[0]
    first = shape._element.find(f'.//{{{A_NS}}}tbl')
    bottom, right = shape.top + shape.height, shape.left + shape.width
    for row in joined[0][len(_table_rows(shape)):]:
        first.append(copy.deepcopy(row))
    grid = first.find(f'{{{A_NS}}}tblGrid')
    for part, rows in zip(parts[1:], joined[1:]):
        for column in part[0]._element.find(f'.//{{{A_NS}}}tblGrid').findall(f'{{{A_NS}}}gridCol')[1:]:
            grid.append(copy.deepcopy(column))
        for target, row in zip(first.findall(f'{{{A_NS}}}tr'), rows):
            for cell in row.findall(f'{{{A_NS}}}tc')[1:]:
                target.append(copy.deepcopy(cell))
    if len(parts) > 1:
        # The parts may have their own font size: the joined table takes the smallest.
        sizes = [node for node in first.iter(*_TEXT_PROPERTIES) if (node.get('sz') or '').isdigit()]
        smallest = min((int(node.get('sz')) for node in sizes), default=0)
        for node in sizes:
            node.set('sz', str(smallest))
    height = sum(int(row.get('h') or 0) for row in first.findall(f'{{{A_NS}}}tr'))
    width = sum(int(column.get('w') or 0) for column in grid.findall(f'{{{A_NS}}}gridCol'))
    delta = max(0, height - shape.height)
    shape.height = max(shape.height, height)
    shape.width = max(shape.width, width)
    for other in content_shapes(pages[0]):
        if other.shape_id != shape.shape_id and other.top >= bottom - Inches(0.05):
            other.top = other.top + delta
        elif other.shape_id != shape.shape_id and other.left >= right - Inches(0.05):
            other.left = other.left + max(0, width - (right - shape.left))
    for page in pages[1:]:
        composer.remove_slide(page)
    return True


# Text is measured with Arial (as wide as the fonts LibreOffice uses for the template's Ericsson Hilda when it is not
# installed, see slide_titles) with room for a substituted font; a block or square drawn as a bar is this wide in ems.
TEXT_ROOM, BLOCK_CHARACTER_EM = 1.15, 0.8
# Height of a line of text in ems.
LINE_EM = 1.25
# A table continues in another column beside it only when each column keeps at least these rows.
MINIMUM_FLOW_ROWS = 6


def _scaled_size(size: int, factor: float) -> int:
    """A font size made ``factor`` times smaller as _scale_fonts makes it."""
    return max(min(size, MINIMUM_FONT), int(round(size * factor)))


def _is_block(character: str) -> bool:
    return '\u2580' <= character <= '\u25ff'


def _text_ems(text: str, header: bool) -> float:
    """The width of ``text`` in ems; a header (in bold) wraps onto up to two lines, so its longest word or half of it
    counts."""
    def ems(value: str) -> float:
        rest = ''.join(character for character in value if not _is_block(character))
        return (len(value) - len(rest)) * BLOCK_CHARACTER_EM + text_width_pt(rest, 1, bold=header) * TEXT_ROOM
    if not header:
        return ems(text)
    return max(max((ems(word) for word in text.split()), default=0), ems(text) / 2)


@dataclass
class _TableCells:
    """What sizes a table: per cell its text width in ems, its font size and its margins, and the height of each row."""
    cells: list[list[tuple[float, int, int, int]]]
    heights: list[int]
    count: int

    def widths(self, factor: float) -> list[int]:
        widths = [0] * self.count
        for row in self.cells:
            for column, (ems, size, sides, _ends) in enumerate(row[:self.count]):
                widths[column] = max(widths[column], int(ems * _scaled_size(size, factor) / 100 * 12700 + sides * factor))
        return widths

    def row_heights(self, factor: float) -> list[int]:
        return [max(int(height * factor), max((int((2 if index == 0 else 1) * LINE_EM * _scaled_size(size, factor) / 100 * 12700 + ends * factor)
                                                for _ems, size, _sides, ends in row), default=0))
                for index, (row, height) in enumerate(zip(self.cells, self.heights))]


def _table_cells(table: Any, fallback_height: int) -> _TableCells:
    rows = table.findall(f'{{{A_NS}}}tr')
    cells = []
    for index, row in enumerate(rows):
        values = []
        for cell, text in zip(row.findall(f'{{{A_NS}}}tc'), _row_texts(row)):
            properties = cell.find(f'{{{A_NS}}}tcPr')
            # DrawingML cell margins default to 0.1" at the sides and 0.05" at the top and bottom.
            margins = {side: int(properties.get(side) or default) if properties is not None else default
                       for side, default in (('marL', 91440), ('marR', 91440), ('marT', 45720), ('marB', 45720))}
            values.append((_text_ems(text, index == 0), _cell_size(cell), margins['marL'] + margins['marR'], margins['marT'] + margins['marB']))
        cells.append(values)
    return _TableCells(cells, [int(row.get('h') or 0) or fallback_height for row in rows], len(table.findall(f'.//{{{A_NS}}}gridCol')))


def _cell_size(cell: Any) -> int:
    """The largest font size of a table cell in hundredths of a point (18pt, the DrawingML default, when it has none)."""
    sizes = [int(node.get('sz')) for node in cell.iter(*_TEXT_PROPERTIES) if (node.get('sz') or '').isdigit()]
    return max(sizes or [1800])


def fit_table(slide: Any, frame: tuple[int, int, int, int], gap: int = Inches(0.25)) -> float | None:
    """Show the only content of ``slide``, a table, as large as it fits in ``frame``: its rows in one to four tables
    side by side (each under the header) and its columns as wide as their text needs, sharing the width of the
    frame. Returns the scale of its text, or None when the slide has other content or the table merges cells."""
    shapes = content_shapes(slide)
    if len(shapes) != 1 or not shapes[0].has_table:
        return None
    shape = shapes[0]
    table = shape._element.find(f'.//{{{A_NS}}}tbl')
    rows = table.findall(f'{{{A_NS}}}tr')
    merged = ('gridSpan', 'rowSpan', 'hMerge', 'vMerge')
    if len(rows) < 2 or any(cell.get(name) for cell in table.iter(f'{{{A_NS}}}tc') for name in merged):
        return None
    sizes = _table_cells(table, shape.height // len(rows))
    left, top, width, height = frame
    body = len(rows) - 1

    def fits(factor: float, per: int, cell: float) -> bool:
        heights = sizes.row_heights(factor)
        tallest = heights[0] + max(sum(heights[1 + start:1 + start + per]) for start in range(0, body, per))
        return tallest <= height and sum(sizes.widths(factor)) <= cell

    best: tuple[float, int, int] | None = None
    for columns in range(1, 5):
        per = math.ceil(body / columns)
        if columns > 1 and (per < MINIMUM_FLOW_ROWS or math.ceil(body / per) < columns):
            break
        cell = (width - gap * (columns - 1)) / columns
        # The largest scale that fits (both the widths and the heights grow with it).
        low, high = 0.2, 1.0
        if fits(high, per, cell):
            low = high
        else:
            for _step in range(16):
                middle = (low + high) / 2
                low, high = (middle, high) if fits(middle, per, cell) else (low, middle)
        # Another column beside it only for clearly larger text.
        if best is None or low > best[0] * 1.15:
            best = (low, columns, per)
    scale, columns, per = best
    cell = int((width - gap * (columns - 1)) / columns)
    needed = sizes.widths(scale)
    table_width = cell if columns > 1 else min(cell, max(shape.width, sum(needed)))
    total = max(1, sum(needed))
    for column, value in zip(table.findall(f'.//{{{A_NS}}}gridCol'), needed):
        column.set('w', str(int(value * table_width / total)))
    _scale_fonts(shape._element, scale)
    for row, row_height in zip(rows, sizes.row_heights(scale)):
        row.set('h', str(row_height))
    for properties in table.iter(f'{{{A_NS}}}tcPr'):
        for side in ('marL', 'marR', 'marT', 'marB'):
            if properties.get(side):
                properties.set(side, str(int(int(properties.get(side)) * scale)))
    starts = list(range(0, body, per))
    copies = [shape._element]
    for _start in starts[1:]:
        element = copy.deepcopy(shape._element)
        copies[-1].addnext(element)
        copies.append(element)
    shape_id = slide.shapes._next_shape_id
    for index, (element, start) in enumerate(zip(copies, starts)):
        if index:
            for properties in element.iter(f'{{{P_NS}}}cNvPr'):
                properties.set('id', str(shape_id))
                shape_id += 1
        element_table = element.find(f'.//{{{A_NS}}}tbl')
        # Each table keeps the header and its own rows.
        for position, row in enumerate(element_table.findall(f'{{{A_NS}}}tr')[1:]):
            if not start <= position < start + per:
                element_table.remove(row)
    placed = [item for item in slide.shapes if item._element in copies]
    for index, item in enumerate(sorted(placed, key=lambda item: copies.index(item._element))):
        item.left, item.top = int(left + index * (cell + gap)), int(top)
        item.width = table_width
        item.height = sum(int(row.get('h') or 0) for row in item._element.find(f'.//{{{A_NS}}}tbl').findall(f'{{{A_NS}}}tr'))
    return scale


def fits_together(composer: DeckComposer, blocks: list[Block], minimum_scale: float = 0.9) -> bool:
    """Whether the blocks fit on one slide, each under its caption, at least at ``minimum_scale`` of their size."""
    if any(not block.shapes for block in blocks):
        return False
    target = blocks[0].slide
    frame = content_frame(target, composer.width, composer.height, [block.box for block in blocks])
    return arrange(blocks, frame).scale >= minimum_scale


def tag(slide: Any, section: str) -> None:
    slide._element.cSld.set('name', f'{SECTION_PREFIX}{section}')


def deck(content: bytes) -> Any:
    return Presentation(BytesIO(content))
