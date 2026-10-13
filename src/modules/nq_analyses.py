"""The analyses of the previous tool: reading its workbooks and translating them onto the Analysis Center.

The analyses were kept in workbooks (one sheet per operator and NR mode, plus the data tests) whose rows repeat the
CDR export of each call and add the analysis after it. Every visible sheet with a ``JOIN_ID`` column is read; its
columns fill the fields of the Analysis Center by their names (or the names the workbooks give them), each comment
column becomes a comment of its own, and whatever has no field travels in an import note of the call, so nothing of
the analyses is lost. Values are checked as if they were chosen in the call: a reviewed table of equivalences turns
the values of the workbooks into those of the catalog (or into new ones).
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field as dataclass_field
from datetime import date, datetime, timedelta
from typing import Any

MAX_ROWS = 50000
# A sheet ends after this many empty rows in a row (the workbooks are formatted far below their data).
EMPTY_ROWS_END = 200
# Cells that say nothing.
EMPTY_VALUES = {'', '.', '-', '--', '#n/a', 'n/a', 'na', 'none', 'null', 'nan', '<na>'}
# Text and list cells that say nothing either (the first analyses wrote 0 for "no initial category").
EMPTY_TEXTS = {'0'}

# What the columns of the workbooks of the previous tool fill when their name is not the name of a field, by identity
# (letters and digits): a follow-up value, a field of the Analysis Center, a comment, or a text added to a field.
COLUMN_ALIASES: dict[str, tuple[str, str]] = {
    'nqcallstatus': ('tracking', 'status'),
    'teamresponsible': ('tracking', 'team'), 'responsiblestream': ('tracking', 'team'),
    'assignee': ('tracking', 'assignee'), 'asignee': ('tracking', 'assignee'),
    'rootcategory': ('tracking', 'root_category'), 'selectedrootcategory': ('tracking', 'root_category'),
    'finalcategoryclassification': ('tracking', 'root_category'),
    'rootcause': ('tracking', 'root_cause'), 'selectedrootcause': ('tracking', 'root_cause'),
    'finalcausesubcategory': ('tracking', 'root_cause'), 'subcategory': ('tracking', 'root_cause'),
    # The people who analysed and reviewed the calls are not users of the tool: their names are kept as they are.
    'team': ('field', 'analyst'), 'analyst': ('field', 'analyst'), 'reviewer': ('field', 'reviewer'),
    'priority': ('field', 'priority'),
    'analysisfinished': ('field', 'analysis_status'),
    'status': ('field', 'planned_status'),
    'evaluation': ('field', 'implementation_status'),
    'streamresponsible': ('field', 'implementation_stream'),
    'actionproposal': ('field', 'implementation_proposal'),
    'site': ('field', 'implementation_site'),
    'servingcell': ('field', 'serving_cell_when_failure'),
    'band': ('field', 'band_when_failure'),
    'problemsolutionlocation': ('field', 'problem_location'),
    'findingsbucketization': ('field', 'findings'),
    'initialcategoryclassification': ('field', 'initial_root_category'),
    'initialcausesubcategory': ('field', 'initial_root_cause'),
    'description': ('field', 'description'),
    'planneddate': ('field', 'planned_date'),
    'plannedfy': ('field', 'planned_fy'),
    'findingscomments': ('append', 'findings'),
    'proposedmeasurecomments': ('append', 'proposed_measure'),
    'comments': ('comments', ''), 'comment': ('comments', ''), 'commentsonlytofilloutbybnpnqstream': ('comments', ''),
}
# Where the analysis starts in a sheet that repeats the CDR export before it: its first column of these (or the first
# column that fills a value). The columns before it are the CDR, which the tool already has.
ANALYSIS_MARKERS = {'uniqueid', 'priority', 'itemserialnumber'}
# When the analysts last changed the analysis of a row (the date of the analysis of a data area): the date the
# import gives the changes of the call.
WHEN_COLUMNS = ('datelastmodification', 'date')
# Built-in translations of list values and teams, by field: the old names of the values of the catalog.
BUILTIN_VALUES: dict[str, dict[str, str]] = {
    'analysis_status': {'yes': 'Finished', 'finished': 'Finished', 'no': 'Pending'},
    'implementation_status': {'solved': 'Implemented'},
    'priority': {str(number): f'P{number}' for number in range(1, 6)},
    'implementation_stream': {'e2e stream': 'E2E Team', 'successful call-excluded': 'No Team'},
    'team': {'e2e stream': 'E2E Team', 'successful call-excluded': 'No Team'},
    'domain': {'core 2g/4g/5g': 'Core'},
}
FISCAL_PERIOD = re.compile(r'^(?:FY\s*\d{2,4}|\d{4}|Q[1-4][\s\-]*\d{2,4})$', re.IGNORECASE)
JOIN_TOKEN = re.compile(r'0x[0-9a-fA-F]+|[0-9a-fA-F]{24,}')


def identity(value: Any) -> str:
    """A name reduced to its lowercase letters and digits, without accents."""
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'[^a-z0-9]', '', text.casefold())


def cell_text(value: Any) -> str:
    """A cell as text, keeping its line breaks (long texts and comments)."""
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d') if not (value.hour or value.minute or value.second) else value.isoformat(sep=' ')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace('\r\n', '\n').strip()
    return '' if text.casefold() in {'nan', 'none', 'null', '<na>'} else text


def is_empty(text: str) -> bool:
    return text.strip().casefold() in EMPTY_VALUES


@dataclass
class Table:
    """One sheet (or the CSV file): its header, its rows with their number in the sheet, and whether it is the Excel
    export of Non-Qualified Calls (its first row names the sections)."""
    sheet: str
    header: list[str]
    rows: list[tuple[int, list[str]]] = dataclass_field(default_factory=list)
    export: bool = False


def _header_row(rows: list[list[str]]) -> int | None:
    return next((index for index, row in enumerate(rows[:3]) if any(identity(value) in {'joinid', 'callkey'} for value in row)), None)


def _table_from_rows(sheet: str, rows: Any) -> Table | None:
    first: list[list[str]] = []
    table: Table | None = None
    empty = 0
    for number, row in enumerate(rows, start=1):
        cells = [cell_text(value) for value in row]
        if table is None:
            first.append(cells)
            header = _header_row(first)
            if header is not None:
                table = Table(sheet, first[header], export=header > 0)
            elif len(first) >= 3:
                return None
            continue
        if not any(cells):
            empty += 1
            if empty >= EMPTY_ROWS_END:
                break
            continue
        empty = 0
        table.rows.append((number, cells))
        if len(table.rows) >= MAX_ROWS:
            break
    return table


def read_tables(content: bytes, file_name: str) -> tuple[list[Table], list[str]]:
    """Every visible sheet with a JOIN_ID (or Call Key) column in its first three rows, and the sheets left out."""
    if file_name.casefold().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        tables, skipped = [], []
        for sheet in workbook.worksheets:
            if sheet.sheet_state != 'visible':
                continue
            table = _table_from_rows(sheet.title, sheet.iter_rows(values_only=True))
            if table is None:
                skipped.append(sheet.title)
            else:
                tables.append(table)
        return tables, skipped
    text = content.decode('utf-8-sig', errors='replace')
    header = text.split('\n', 1)[0]
    delimiter = max(',;\t|', key=header.count)
    table = _table_from_rows('CSV', csv.reader(io.StringIO(text), delimiter=delimiter if header.count(delimiter) else ','))
    return ([table] if table else []), ([] if table else ['CSV'])


@dataclass
class Layout:
    """What each column of a table fills: index -> (kind, ref); the comment and note columns; the identity columns."""
    targets: dict[int, tuple[str, str]]
    comments: list[int]
    appends: dict[int, str]
    notes: list[int]
    cdr: list[str]
    read_only: list[str]
    join: int | None
    key: int | None
    direction: int | None
    when: int | None = None

    def recognized(self, header: list[str]) -> list[str]:
        return [header[index] for index in sorted({*self.targets, *self.comments, *self.appends})]


def resolve_layout(table: Table, definitions: list[dict[str, Any]], tracking_refs: tuple[str, ...]) -> Layout:
    """Which value of the Analysis Center each column fills: the name of a field (its label or key) first, then the names
    the workbooks of the previous tool use. A second column for the same value, and every other column of the analysis,
    goes to the import note; the CDR repeated before the analysis is left out (the tool has it)."""
    names: dict[str, tuple[str, str]] = {}
    read_only: set[str] = set()
    fields = {field['key']: field for field in definitions}
    for field in definitions:
        if field['source'] == 'user':
            target: tuple[str, str] | None = ('field', field['key'])
        elif field['source'] == 'tracking' and field['source_ref'] == 'comments':
            target = ('comments', '')
        elif field['source'] == 'tracking' and field['source_ref'] in tracking_refs:
            target = ('tracking', field['source_ref'])
        else:
            target = None
        for name in (field['label'], field['key']):
            if target is None:
                read_only.add(identity(name))
            else:
                names.setdefault(identity(name), target)
    for alias, target in COLUMN_ALIASES.items():
        # An alias only fills a field the workspace still has.
        if target[0] in {'field', 'append'} and target[1] not in fields:
            continue
        names.setdefault(alias, target)
    header = [identity(name) for name in table.header]
    join = next((index for index, name in enumerate(header) if name == 'joinid'), None)
    key = next((index for index, name in enumerate(header) if name == 'callkey'), None)
    direction = next((index for index, name in enumerate(header) if name == 'calldirection'), None)

    def target_of(index: int) -> tuple[str, str] | None:
        name = header[index]
        # Vendor fills Vendor 1 when Failure only next to Band (in the analysis); elsewhere it repeats the CDR.
        if name == 'vendor':
            return ('field', 'vendor_when_failure') if index and header[index - 1] == 'band' and 'vendor_when_failure' in fields else None
        return names.get(name)

    start = 0
    if not table.export:
        marks = [index for index, name in enumerate(header) if name in ANALYSIS_MARKERS or (
            index not in {join, key} and target_of(index) is not None)]
        start = min(marks) if marks else 0
    layout = Layout({}, [], {}, [], [], [], join, key, direction)
    if not table.export:
        layout.when = next((index for name in WHEN_COLUMNS for index, column in enumerate(header)
                            if column == name and index >= start), None)
    taken: set[tuple[str, str]] = set()
    for index, name in enumerate(header):
        if not name or index in {join, key, layout.when}:
            continue
        target = target_of(index)
        if index < start:
            layout.cdr.append(table.header[index])
        elif target and target[0] == 'comments':
            layout.comments.append(index)
        elif target and target[0] == 'append':
            layout.appends[index] = target[1]
        elif target and target not in taken:
            layout.targets[index] = target
            taken.add(target)
        elif table.export:
            # The export of Non-Qualified Calls is imported back: its other columns come from the CDR or the RCA.
            if name in read_only or target:
                layout.read_only.append(table.header[index])
        else:
            layout.notes.append(index)
    return layout


def join_ids(text: str) -> list[str]:
    """The JOIN_IDs of a cell: one, or several (an analysis of an area lists the calls it covers)."""
    tokens = JOIN_TOKEN.findall(text)
    if len(tokens) > 1:
        return list(dict.fromkeys(tokens))
    return [text.strip()] if text.strip() else []


def failed_party(value: str, direction: str) -> str | None:
    """MOC and MTC (the originating and the terminating side) as the UE that failed, from the direction of the call."""
    side = identity(value)
    if side not in {'moc', 'mtc', 'mo', 'mt'}:
        return None
    caller, callee = ('B', 'A') if identity(direction) == 'ba' else ('A', 'B')
    return f'UE_{caller} (MO)' if side.startswith('mo') else f'UE_{callee} (MT)'


def parse_date(text: str) -> str | None:
    for pattern in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y/%m/%d', '%d/%m/%y'):
        try:
            return datetime.strptime(text.strip()[:10], pattern).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def parse_moment(text: str) -> str | None:
    """A date (or date and time) of a cell as an ISO timestamp in the local time zone, at noon without a time; None when
    the cell is not a date or is in the future."""
    value = text.strip()
    day = parse_date(value)
    if day is None:
        return None
    moment = None
    if len(value) > 10 and value[:10] == day:
        try:
            moment = datetime.fromisoformat(value.replace(' ', 'T'))
        except ValueError:
            moment = None
        if moment is not None and not (moment.hour or moment.minute or moment.second):
            moment = None
    moment = (moment or datetime.fromisoformat(f'{day}T12:00:00')).replace(tzinfo=None).astimezone()
    return moment.isoformat(timespec='seconds') if moment <= datetime.now().astimezone() else None


def moment_after(text: str, days: int) -> str | None:
    """The day of a date (or date and time) some days later, at noon in the local time zone, as an ISO timestamp; None
    when the cell is not a date or the moment is in the future."""
    day = parse_date(text)
    if day is None:
        return None
    moment = (datetime.fromisoformat(f'{day}T12:00:00') + timedelta(days=days)).astimezone()
    return moment.isoformat(timespec='seconds') if moment <= datetime.now().astimezone() else None


def canonical_names(values: list[str]) -> dict[str, str]:
    """The spelling of each name the most rows use, for every way the same name is written (Raj Kumar, Rajkumar)."""
    groups: dict[str, Counter] = {}
    for value in values:
        if value.strip():
            groups.setdefault(identity(value), Counter())[value.strip()] += 1
    return {key: counts.most_common(1)[0][0] for key, counts in groups.items()}


def list_choice(value: str, options: list[str]) -> str | None:
    """The option of a list a value names: the same name, else the same letters and digits (RF OPT (Hot Spots)_WP2a)."""
    exact = {option.casefold(): option for option in options}
    if value.casefold() in exact:
        return exact[value.casefold()]
    loose = {identity(option): option for option in options}
    return loose.get(identity(value))


def band_name(value: str) -> str:
    """Bands as the catalog writes them: L800 is LTE800, N3500 is NR3500."""
    match = re.fullmatch(r'\s*([LN])\s*(\d{3,4})\s*', value, re.IGNORECASE)
    if not match:
        return value
    return f"{'LTE' if match.group(1).upper() == 'L' else 'NR'}{match.group(2)}"
