"""Root Cause Analysis of the NQ Analysis Center: the Ericsson RCA results and the recommendation.

Three sources describe the root cause of a Non-Qualified Call:

- the **NetCheck RCA** of the CDR (Failure Classification, Category and Subcategory);
- the **Ericsson RCA** (a Python script of the team) writes, for each JOIN_ID, the categories and
  subcategories of both sides and a suggested category and cause; its CSV or Excel files are
  imported and linked to the calls by JOIN_ID;
- the **analysts**, who choose the Selected Root Category and Root Cause.

Both RCAs keep their own category → cause hierarchy. The recommendation maps each one onto the
root cause catalog of the workspace (names first, then the keywords of the domains, categories
and causes), adds the keyword reading of the comments and what the analysts chose before for the
same RCA values, and proposes the category and cause most sources agree on.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter
from threading import Lock
from typing import Any

from src.modules.nq_catalog import RCA_REFS, keyword_found, normalized_text

RCA_FILE_FORMAT_HELP = (
    'A CSV (comma, semicolon or tab separated) or Excel file with a JOIN_ID column and the RCA columns '
    'Auto_RCA_Category_A, Auto_RCA_Subcategory_A, Auto_RCA_Category_B, Auto_RCA_Subcategory_B, '
    'RCA_Category, RCA_Subcategory, RCA_Suggested_Category, RCA_Suggested_Cause and RCA_Evidence, or the RCA_Summary.csv '
    'of the Auto RCA output of the Ericsson RCA (JoinID, UE, RCA_Category, RCA_Subcategory, Final_RCA_Category, Final_RCA_Subcategory and '
    'RCA_Evidence: the rows of the A and B sides fill the fields of each side, and the rows of one JOIN_ID are combined); '
    'any other column is kept and shown with the call.'
)
RCA_TEMPLATE_COLUMNS = ('JOIN_ID', 'Auto_RCA_Category_A', 'Auto_RCA_Subcategory_A', 'Auto_RCA_Category_B',
                        'Auto_RCA_Subcategory_B', 'RCA_Suggested_Category', 'RCA_Suggested_Cause', 'RCA_Evidence')
MAX_RCA_ROWS = 500_000
MAX_EXTRA_COLUMNS = 80
MAX_COMBINED_TEXT = 2000
# Column identities (letters and digits only) of the known RCA columns: the value they hold and their tier. A row
# takes its suggested category and cause from the highest tier it has one of them in (Final_RCA_* before
# RCA_Suggested_*), or else from its initial RCA_Category and RCA_Subcategory; the columns of a lower tier are also
# kept with the call.
_RCA_COLUMNS = {
    'joinid': ('join_id', 0), 'ue': ('side', 0), 'side': ('side', 0),
    'autorcacategorya': ('auto_rca_category_a', 0), 'autorcasubcategorya': ('auto_rca_subcategory_a', 0),
    'autorcacategoryb': ('auto_rca_category_b', 0), 'autorcasubcategoryb': ('auto_rca_subcategory_b', 0),
    'rcacategory': ('rca_initial_category', 0), 'rcasubcategory': ('rca_initial_cause', 0),
    'rcasuggestedcategory': ('rca_category', 1), 'suggestedcategory': ('rca_category', 1),
    'rcasuggestedcause': ('rca_cause', 1), 'rcacause': ('rca_cause', 1), 'suggestedcause': ('rca_cause', 1),
    'finalrcacategory': ('rca_category', 2), 'finalrcasubcategory': ('rca_cause', 2), 'finalrcacause': ('rca_cause', 2),
    'rcaevidence': ('rca_evidence', 0), 'evidence': ('rca_evidence', 0),
}
_SIDES = {'a': 'a', 'uea': 'a', 'b': 'b', 'ueb': 'b'}
# Categories that name no cause (the script found none): another row of the same JOIN_ID suggests better.
_GENERIC = re.compile(r'^(other\b|unknown$|unclear$|n/?a$|none$|-+$)', re.IGNORECASE)
# A row without a result (no category, or UNKNOWN when the script had no measurements of the session).
_NO_RESULT = re.compile(r'^(unknown)?$', re.IGNORECASE)
# The domain proposed for a category or cause missing from the catalog, by the words of its name (when the domain
# exists): "Other" results are radio results when they name a technology.
_DOMAIN_HINTS = (('unknown', ''), ('unclear', ''), ('mobility', 'RAN'), ('handover', 'RAN'), ('coverage', 'RF'),
                 ('interference', 'RF'))
_TECHNOLOGY = re.compile(r'\b(lte|nr|gsm|umts|wcdma|2g|3g|4g|5g)\b', re.IGNORECASE)
SOURCE_LABELS = {'netcheck': 'NetCheck RCA', 'script': 'Ericsson RCA', 'comments': 'Comments', 'learned': 'Previous decisions'}
SOURCE_WEIGHTS = {'netcheck': 3, 'learned': 3, 'script': 2, 'comments': 1}


def _identity(value: Any) -> str:
    return re.sub(r'[^a-z0-9]', '', str(value or '').casefold())


def _cell(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r'\s+', ' ', str(value)).strip()
    return '' if text.casefold() in {'nan', 'none', 'null', '<na>'} else text[:300]


def _table(content: bytes, file_name: str) -> list[list[Any]]:
    if file_name.casefold().endswith(('.xlsx', '.xlsm')):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(row) for row in workbook.worksheets[0].iter_rows(values_only=True)]
    text = content.decode('utf-8-sig', errors='replace')
    # The separator is the one the header row uses most: the values (evidence texts) may hold any of them.
    header = text.split('\n', 1)[0]
    delimiter = max(',;\t|', key=header.count)
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter if header.count(delimiter) else ',')]


def read_rca_file(content: bytes, file_name: str) -> dict[str, Any]:
    """The RCA results of a file: one entry per JOIN_ID (its rows combined), the recognized columns and the rest."""
    rows = _table(content, file_name)
    if not rows:
        raise ValueError('The RCA file is empty.')
    headers = [_cell(value) for value in rows[0]]
    columns: dict[int, tuple[str, int]] = {}
    extra: dict[int, str] = {}
    for index, header in enumerate(headers):
        known = _RCA_COLUMNS.get(_identity(header))
        if known and known not in columns.values():
            columns[index] = known
        elif header and len(extra) < MAX_EXTRA_COLUMNS:
            extra[index] = header
    keys = {key for key, _tier in columns.values()}
    if 'join_id' not in keys:
        raise ValueError('The RCA file needs a JOIN_ID column.')
    if not keys & set(RCA_REFS):
        raise ValueError('The RCA file has none of the RCA columns (Auto_RCA_Category_A… RCA_Suggested_Cause, Final_RCA_Category).')
    # A category or cause column of a lower tier stays with the call beside the one it gives way to.
    top = {key: max(tier for other, tier in columns.values() if other == key) for key in keys}
    for index, (key, tier) in columns.items():
        if tier < top[key] and len(extra) < MAX_EXTRA_COLUMNS:
            extra[index] = headers[index]
    join_index = next(index for index, (key, _tier) in columns.items() if key == 'join_id')
    side_index = next((index for index, (key, _tier) in columns.items() if key == 'side'), None)
    groups: dict[str, dict[str, Any]] = {}
    skipped = 0
    for row in rows[1:MAX_RCA_ROWS + 1]:
        cells = {index: _cell(value) for index, value in enumerate(row)}
        join_id = cells.get(join_index, '')
        if not join_id:
            if any(cells.values()):
                skipped += 1
            continue
        found = [tier for index, (key, tier) in columns.items() if key in {'rca_category', 'rca_cause'} and cells.get(index)]
        tier = max(found, default=None)
        values: dict[str, str] = {}
        for index, (key, column_tier) in columns.items():
            if key in {'join_id', 'side'} or not cells.get(index):
                continue
            if key in {'rca_category', 'rca_cause'} and column_tier != tier:
                continue
            values.setdefault(key, cells[index])
        if tier is None:
            # Without a suggestion the initial category and subcategory suggest.
            values.update({key: values[initial] for key, initial in (('rca_category', 'rca_initial_category'),
                                                                     ('rca_cause', 'rca_initial_cause')) if values.get(initial)})
        side = _SIDES.get(_identity(cells.get(side_index, ''))) if side_index is not None else None
        extras = {name: cells[index] for index, name in extra.items() if cells.get(index)}
        groups.setdefault(join_id.casefold(), {'join_id': join_id, 'rows': []})['rows'].append(
            {'side': side, 'values': values, 'extras': extras})
    results = [{'join_id': group['join_id'], 'values': _combine_rows(group['rows'])} for group in groups.values()]
    present = {key for result in results for key in result['values']}
    return {'results': results, 'columns': [label for key, label in RCA_REFS.items() if key in present],
            'extra': list(extra.values()), 'skipped': skipped,
            'duplicates': sum(1 for group in groups.values() if len(group['rows']) > 1)}


def _combine_rows(rows: list[dict[str, Any]]) -> dict[str, str]:
    """The values of a JOIN_ID from its rows: the A and B sides fill their own fields with their own (initial)
    diagnosis, the suggestion (and the initial category and subcategory) is the first one naming a cause (the A side
    first), and the evidence and other columns keep the value of every row."""
    # A side repeated without a result (the session without its measurements) adds nothing to a row with one.
    def result(row: dict[str, Any]) -> str:
        return row['values'].get('rca_category') or row['values'].get('rca_initial_category') or ''
    found = {row['side'] for row in rows if not _NO_RESULT.match(result(row))}
    rows = [row for row in rows if row['side'] not in found or not _NO_RESULT.match(result(row))]
    ordered = sorted(enumerate(rows), key=lambda item: ({'a': 0, 'b': 1}.get(item[1]['side'], 2), item[0]))
    labels = {}
    for position, (index, row) in enumerate(ordered, start=1):
        labels[index] = row['side'].upper() if row['side'] else str(position)
    values: dict[str, str] = {}
    for _index, row in ordered:
        for key, value in row['values'].items():
            if key.startswith('auto_rca_'):
                values.setdefault(key, value)
        if row['side']:
            # Each side keeps its own diagnosis (the initial one): the final one is that of the whole call.
            initial = bool(row['values'].get('rca_initial_category'))
            category = row['values'].get('rca_initial_category' if initial else 'rca_category')
            cause = row['values'].get('rca_initial_cause' if initial else 'rca_cause')
            if category:
                values.setdefault(f"auto_rca_category_{row['side']}", category)
            if cause:
                values.setdefault(f"auto_rca_subcategory_{row['side']}", cause)
    for category_key, cause_key in (('rca_initial_category', 'rca_initial_cause'), ('rca_category', 'rca_cause')):
        pairs = [(row['values'].get(category_key, ''), row['values'].get(cause_key, '')) for _index, row in ordered
                 if row['values'].get(category_key) or row['values'].get(cause_key)]
        if pairs:
            category, cause = next((pair for pair in pairs if not _GENERIC.match(pair[0] or pair[1])), pairs[0])
            values.update({key: value for key, value in ((category_key, category), (cause_key, cause)) if value})

    def combined(found: list[tuple[int, str]], name: str = '') -> str:
        distinct = list(dict.fromkeys(value for _index, value in found))
        if len(distinct) <= 1:
            return distinct[0] if distinct else ''
        # The rows of different categories say which one they explain.
        categories = {row['values'].get('rca_category', '') for row in rows}
        def label(index: int) -> str:
            text = labels[index]
            category = rows[index]['values'].get('rca_category', '')
            return f'{text} ({category})' if name == 'evidence' and len(categories) > 1 and category else text
        return ' | '.join(f'{label(index)}: {value}' for index, value in found)[:MAX_COMBINED_TEXT]

    evidence = combined([(index, row['values']['rca_evidence']) for index, row in ordered if row['values'].get('rca_evidence')], 'evidence')
    if evidence:
        values['rca_evidence'] = evidence
    names = list(dict.fromkeys(name for _index, row in ordered for name in row['extras']))
    for name in names:
        value = combined([(index, row['extras'][name]) for index, row in ordered if row['extras'].get(name)])
        if value:
            values[f'extra:{name}'] = value
    return values


def proposed_domain(name: str, domains: list[str]) -> str:
    """The domain of the catalog a missing category or cause belongs to, by the words of its name ('' when unsure)."""
    words = normalized_text(name)
    known = {domain.casefold(): domain for domain in domains}
    for word, domain in _DOMAIN_HINTS:
        if word in words:
            return known.get(domain.casefold(), '')
    if words.startswith('other') and _TECHNOLOGY.search(words):
        return known.get('ran', '')
    return ''


def unmapped_values(results: list[dict[str, Any]], root: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    """The script categories and causes the catalog has no entry for (by name or keywords), with their number of
    results and the domain proposed for them, the categories first."""
    pairs: Counter = Counter()
    for result in results:
        values = result['values']
        found = {(values.get('rca_category', ''), values.get('rca_cause', '')),
                 (values.get('auto_rca_category_a', ''), values.get('auto_rca_subcategory_a', '')),
                 (values.get('auto_rca_category_b', ''), values.get('auto_rca_subcategory_b', ''))}
        pairs.update(pair for pair in found if any(pair))
    missing: dict[tuple[str, str], dict[str, Any]] = {}
    domains = [domain['name'] for domain in root.get('domains') or []]
    for (category, cause), count in pairs.items():
        mapped = map_to_catalog(category, cause, root)
        for kind, name, hit in (('category', category, mapped['category']), ('cause', cause, mapped['cause'])):
            if not name or hit:
                continue
            entry = missing.setdefault((kind, name.casefold()), {'kind': kind, 'name': name, 'domain': proposed_domain(name, domains), 'results': 0})
            entry['results'] += count
    return sorted(missing.values(), key=lambda item: (item['kind'] != 'category', -item['results'], item['name'].casefold()))


def rca_template() -> bytes:
    """An empty RCA results file with its columns and one example row."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(RCA_TEMPLATE_COLUMNS)
    writer.writerow(['0x0123ABCD', 'RF Problems', 'Coverage problem', 'RF Problems', 'Coverage problem',
                     'Poor Coverage LTE', 'NLOS due to clutter or terrain profile', 'LTE PCell RSRP=-118 dBm < RSRP_Low=-115 dBm'])
    return output.getvalue().encode('utf-8')


# -- mapping onto the root cause catalog --------------------------------------------------------
def _by_name(items: list[dict[str, Any]], value: str) -> dict[str, Any] | None:
    key = normalized_text(value)
    return next((item for item in items if normalized_text(item['name']) == key), None) if key else None


def _first_match(items: list[dict[str, Any]], text: str, match: str, domain: str = '') -> dict[str, Any] | None:
    for item in items:
        if domain and item.get('domain') and item['domain'] != domain:
            continue
        if keyword_found(item.get('keywords') or [], text, match):
            return item
    return None


# The results of the catalog and the rule of one request, by their text: many calls share the same NetCheck and
# script values, so each combination is matched once. The catalog and rule objects themselves are the key.
_MEMO_SIZE = 8
_memos: list[tuple[Any, Any, dict[Any, Any]]] = []
_memos_guard = Lock()


def _memo(root: Any, rule: Any) -> dict[Any, Any]:
    with _memos_guard:
        for entry in _memos:
            if entry[0] is root and entry[1] is rule:
                return entry[2]
        memo: dict[Any, Any] = {}
        _memos.insert(0, (root, rule, memo))
        del _memos[_MEMO_SIZE:]
        return memo


def map_to_catalog(category: str, cause: str, root: dict[str, list[dict[str, Any]]], match: str = 'words',
                   extra_text: str = '') -> dict[str, str]:
    """The catalog category and cause a pair of RCA values stands for: same names first, then keywords."""
    memo = _memo(root, None)
    key = ('map', category, cause, match, extra_text)
    if key not in memo:
        memo[key] = _map_to_catalog(category, cause, root, match, extra_text)
    return dict(memo[key])


def _map_to_catalog(category: str, cause: str, root: dict[str, list[dict[str, Any]]], match: str, extra_text: str) -> dict[str, str]:
    found_category = _by_name(root['categories'], category) or _first_match(
        root['categories'], ' | '.join(normalized_text(value) for value in (category, cause, extra_text) if value), match)
    found_cause = _by_name(root['causes'], cause) or _first_match(root['causes'], normalized_text(cause), match) \
        or _first_match(root['causes'], normalized_text(category), match)
    domain = (found_category or {}).get('domain') or (found_cause or {}).get('domain') or ''
    return {'domain': domain, 'category': (found_category or {}).get('name', ''), 'cause': (found_cause or {}).get('name', '')}


def suggest_from_text(call: Any, root: dict[str, list[dict[str, Any]]], rule: dict[str, Any],
                      comments: str = '') -> dict[str, str] | None:
    """The domain, category and cause the CDR fields (or, else, the comments) suggest with the workspace rule.

    1. A domain keyword in the rule's domain fields (Failure Classification) chooses the domain.
    2. Category keywords in those fields, the cause fields and the technology choose the category,
       and cause keywords in the cause fields (Failure Category, Subcategory, Comment) the cause:
       in the domain found, when the rule keeps them in it.
    3. Without a category and cause, the comments are read the same way.
    """
    def text(fields: list[str]) -> str:
        values = []
        for field in fields:
            try:
                values.append(normalized_text(call[field]))
            except (KeyError, IndexError, TypeError):
                continue
        return ' | '.join(value for value in values if value)

    domain_text = text(rule.get('domain_fields') or [])
    cause_text = text(rule.get('cause_fields') or [])
    technology = text(['technology', 'failure_technology'])
    memo = _memo(root, rule)
    key = ('suggest', domain_text, cause_text, technology, comments)
    if key not in memo:
        memo[key] = _suggest(domain_text, cause_text, technology, comments, root, rule)
    return dict(memo[key]) if memo[key] else None


def _suggest(domain_text: str, cause_text: str, technology: str, comments: str, root: dict[str, list[dict[str, Any]]],
             rule: dict[str, Any]) -> dict[str, str] | None:
    match = rule.get('match', 'words')
    domain = next((item for item in root['domains'] if keyword_found(item.get('keywords') or [], domain_text, match)), None)
    scope = domain['name'] if domain and rule.get('causes_in_domain', True) else ''
    category_text = ' | '.join(value for value in (domain_text, cause_text, technology) if value)
    category = _first_match(root['categories'], category_text, match, scope)
    cause = _first_match(root['causes'], cause_text, match, scope)
    if category or cause:
        found_domain = (domain or {}).get('name') or (category or {}).get('domain') or (cause or {}).get('domain') or ''
        return {'domain': found_domain, 'category': (category or {}).get('name', ''), 'cause': (cause or {}).get('name', ''),
                'source': 'netcheck'}
    notes = normalized_text(comments) if rule.get('comments', True) else ''
    if notes:
        category = _first_match(root['categories'], notes, match, scope)
        cause = _first_match(root['causes'], notes, match, scope)
        if category or cause:
            found_domain = (domain or {}).get('name') or (category or {}).get('domain') or (cause or {}).get('domain') or ''
            return {'domain': found_domain, 'category': (category or {}).get('name', ''), 'cause': (cause or {}).get('name', ''),
                    'source': 'comments'}
    if domain is not None:
        return {'domain': domain['name'], 'category': '', 'cause': '', 'source': 'netcheck'}
    if notes and rule.get('comment_domain', True):
        commented = next((item for item in root['domains'] if keyword_found(item.get('keywords') or [], notes, match)), None)
        if commented is not None:
            return {'domain': commented['name'], 'category': '', 'cause': '', 'source': 'comments'}
    return None


def learned_key(source: str, category: str, cause: str) -> tuple[str, str, str]:
    return source, normalized_text(category), normalized_text(cause)


def recommend(call: Any, rca: dict[str, str], root: dict[str, list[dict[str, Any]]], rule: dict[str, Any],
              comments: str = '', learned: dict[tuple[str, str, str], dict[tuple[str, str], set[str]]] | None = None) -> dict[str, Any]:
    """The root cause each source proposes for a call and the recommended Selected Root Category and Cause.

    ``rca`` holds the imported script values of the call; ``learned`` holds, for the RCA values of
    the calls already decided, the calls that chose each category and cause. Each other call with
    the same script or NetCheck values counts once.
    """
    match = rule.get('match', 'words')
    learned = learned or {}
    # NetCheck first: it weighs more and its card comes first.
    sources: list[dict[str, Any]] = []
    netcheck_category = _value(call, 'failure_category') or _value(call, 'failure_classification')
    netcheck_cause = _value(call, 'failure_subcategory')
    suggestion = suggest_from_text(call, root, rule, comments)
    if netcheck_category or netcheck_cause or (suggestion and suggestion['source'] == 'netcheck'):
        mapped = suggestion if suggestion and suggestion['source'] == 'netcheck' else {'domain': '', 'category': '', 'cause': ''}
        sources.append({'source': 'netcheck', 'raw_category': netcheck_category, 'raw_cause': netcheck_cause,
                        'domain': mapped['domain'], 'category': mapped['category'], 'cause': mapped['cause']})
    script_category = rca.get('rca_category') or rca.get('auto_rca_category_a') or ''
    script_cause = rca.get('rca_cause') or rca.get('auto_rca_subcategory_a') or ''
    if script_category or script_cause:
        mapped = map_to_catalog(script_category, script_cause, root, match)
        sources.append({'source': 'script', 'raw_category': script_category, 'raw_cause': script_cause, **mapped})
    if suggestion and suggestion['source'] == 'comments':
        sources.append({'source': 'comments', 'raw_category': '', 'raw_cause': '', 'domain': suggestion['domain'],
                        'category': suggestion['category'], 'cause': suggestion['cause']})
    voters: dict[tuple[str, str], set[str]] = {}
    own = _value(call, 'call_key')
    for source, category, cause in (('netcheck', netcheck_category, netcheck_cause), ('script', script_category, script_cause)):
        if category or cause:
            for choice, keys in learned.get(learned_key(source, category, cause), {}).items():
                voters.setdefault(choice, set()).update(keys - {own})
    voters = {choice: keys for choice, keys in voters.items() if keys}
    if voters:
        (category, cause), keys = max(sorted(voters.items()), key=lambda item: len(item[1]))
        count = len(keys)
        if count >= 2:
            item = _by_name(root['categories'], category)
            reason = _by_name(root['causes'], cause)
            sources.append({'source': 'learned', 'raw_category': '', 'raw_cause': '', 'count': count,
                            'domain': (item or reason or {}).get('domain', ''), 'category': (item or {}).get('name', ''),
                            'cause': (reason or {}).get('name', '')})
    for source in sources:
        source['label'] = SOURCE_LABELS[source['source']]
    return {'sources': sources, 'recommendation': _consensus(sources, root)}


def _value(call: Any, key: str) -> str:
    try:
        return str(call[key] or '').strip()
    except (KeyError, IndexError, TypeError):
        return ''


def _consensus(sources: list[dict[str, Any]], root: dict[str, list[dict[str, Any]]]) -> dict[str, Any] | None:
    """The category most sources (weighted) propose and, among the sources that agree on it, the best cause."""
    category_votes: Counter = Counter()
    for source in sources:
        if source['category']:
            category_votes[source['category']] += SOURCE_WEIGHTS[source['source']]
    cause_votes: Counter = Counter()
    best_category = category_votes.most_common(1)[0][0] if category_votes else ''
    for source in sources:
        if source['cause'] and (not best_category or source['category'] in {best_category, ''}):
            cause_votes[source['cause']] += SOURCE_WEIGHTS[source['source']]
    best_cause = cause_votes.most_common(1)[0][0] if cause_votes else ''
    if not best_category and not best_cause:
        return None
    agreeing = [source['source'] for source in sources
                if (best_category and source['category'] == best_category) or (not best_category and source['cause'] == best_cause)]
    if len(agreeing) >= 2:
        confidence = 'high'
    elif agreeing and agreeing[0] in {'netcheck', 'learned'}:
        confidence = 'medium'
    else:
        confidence = 'low'
    category = _by_name(root['categories'], best_category)
    cause = _by_name(root['causes'], best_cause)
    return {'category': best_category, 'cause': best_cause,
            'domain': (category or {}).get('domain') or (cause or {}).get('domain') or '',
            'sources': agreeing, 'confidence': confidence}
