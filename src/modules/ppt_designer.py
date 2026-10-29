"""Workspace design definitions and shared-table, template-driven previews."""
from __future__ import annotations

import copy
import json
import math
import html
import os
import re
import shutil
import sqlite3
import tempfile
import zipfile
from collections import OrderedDict, defaultdict
from contextlib import contextmanager, nullcontext
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta, timezone
from hashlib import sha256
from io import BytesIO
from time import monotonic
from pathlib import Path
from threading import Condition, Event, Lock, RLock, Thread
from types import SimpleNamespace
from typing import Any, Literal
from uuid import uuid4

import pandas as pd
from src.modules.column_names import campaign_sort_key, column_identity, compact_campaign_value, sort_vendor_values, vendor_filter_column, vendor_filter_value, vendor_filter_values, vendor_match_values
from src.modules import ppt_design_cache, report_structure
from src.modules.deletion_rights import (
    PPT_DESIGNER, can_delete, creator_of, creators, deletion_refusal, forget_creator, record_creator,
)
from src.modules.scoring_themes import THEMES, canvas_chart_style, load_user_theme, normalize_theme, save_user_theme, themes_payload, use_theme
from src.modules.theme_titles import style_title
from src.modules.value_maps import ValueMapper, field_kind
from src.modules.cdr_auto_selection import (
    LATEST_CAMPAIGN, LATEST_TWO_CAMPAIGNS, campaign_identity, campaign_presets, cdr_campaigns, candidate_cdrs, describe_cdrs,
    resolve_campaign_filter, resolve_cdrs, sort_campaigns,
)
from src.modules.cdr_facets import FILTER_FIELDS as CASCADING_FILTER_FIELDS, cascading_options
from src.modules.nr_mode import NR_MODE_FIELD, NR_MODES, NSA_VS_SA, dataset_nr_mode, nr_modes_of
from src.modules.report_template_modes import (
    DESIGN_NR_MODE_LABELS, NR_MODE_TOKEN_VALUES, chart_filter_nr_modes, design_nr_mode_value, entry_for_nr_mode, migrate_design_definition,
    missing_nr_mode_fields, missing_nr_mode_note,
)
from src.modules.output_layout import DESIGNS_FOLDER, module_output_dir
from fastapi import BackgroundTasks, Body, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from pydantic import BaseModel, Field, model_validator

from src.modules.ppt_naming import renamed_design
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from starlette.background import BackgroundTask

from src.modules.cdr_reporting import (
    BOTH_NR_MODES_ATTR, _catalog_spec, _clear_commentary, _clear_end_slide, _layout_chart_frames, _legend_dimensions, _named_slide_layout,
    _remove_all_slides, _remove_template_chart_placeholders, _render_canvas_payload,
    _set_commentary, _set_slide_header, _set_structural_slide_text, catalog_chart_payload,
    catalog_kpi_fields, expand_dynamic_layouts, entry_dynamic_fields, dynamic_chart_title, _normalise_operator_label, _normalise_vendor, prepare_multivendor_catalog_entry,
    ensure_vendor_group, normalise_operator_aliases, parse_catalog_filters,
    parse_catalog_grouping,
    prepare_catalog_chart_preview_frame, render_catalog_chart_preview, render_unavailable_source_chart,
    split_calculated_dimension_aliases, chart_title_parts, _select_dynamic_chart_frame, STRUCTURAL_LAYOUT_TYPES,
    CATALOG_SOURCE_KINDS, CatalogEntry, ELEMENT_CHART_TYPES,
)
from src.modules.report_layouts import DYNAMIC_LAYOUTS, canonical_layout_name, compact_grid_layout, grid_layout_name

from src.modules import cdr_stage, cdr_store
from src.modules.design_column_cache import (
    WEIGHT_COLUMN, CacheSource, ColumnCache, Segment, SelectionTerms, StaleSegment, build_from_copies, count_rows,
    default_memory_bytes, expand_weighted_rows, read_frame, read_grouped_frame,
)
from src.modules.rf_catalog_source import RF_CATALOG_FIELDS, rf_source_columns, rf_union_source
from src.modules.repository import Repository, workspace_read_connection
from src.modules.cdr_report_filenames import build_cdr_report_filename
from src.modules.runtime_config import ignore_event_time_filtering
from src.modules.ingestion import apply_operator_mappings
from src.modules.value_patterns import ValueSet, is_pattern
from src.modules.operator_roles import (
    drop_excluded_operators, excluded_operators, focus_operators, out_of_scope_operators, out_of_scope_pairs,
)
from src.modules import comment_writers, design_comments
from src.modules import design_images, design_slides

KINDS = ('data', 'voice', 'speech')


def existing_dataset_selection(definition: Any, dataset_kinds: dict[int, str]) -> Any:
    """A Design without the CDRs that are not in its workspace with their type (deleted, or from another workspace).

    When none of its chosen CDRs is left, the Design uses the latest CDRs again, like a Design without CDRs.
    """
    if not isinstance(definition, dict) or not isinstance(definition.get('datasets'), dict):
        return definition
    stored = definition['datasets']
    kept: dict[str, list] = {}
    for kind, ids in stored.items():
        values = ids if isinstance(ids, list) else []
        kept[kind] = [value for value in values if _dataset_id(value) is not None and dataset_kinds.get(_dataset_id(value)) == kind]
    if all(kept.get(kind) == (ids if isinstance(ids, list) else []) for kind, ids in stored.items()):
        return definition
    if any(kept.values()):
        return {**definition, 'datasets': kept}
    return {key: value for key, value in definition.items() if key != 'datasets'}


def _dataset_id(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None
STATE_KEY = 'ppt_designs_v2'
# A PPT Template opened in the viewer and the Designer (PPT Templates → Edit): a Design of each user kept apart from
# the Designs (never listed, exported, transferred or used by another module), whose changes go to its template.
TEMPLATE_PREVIEWS_STATE_KEY = 'ppt_template_previews_v1'
TEMPLATE_PREVIEW_PREFIX = 'template-preview-'
UNAVAILABLE_COMBINED_COLUMNS_STATE_KEY = 'design_unavailable_combined_columns_v1'
LEGACY_STATE_KEY = 'ppt_design_sets_v1'
DEFAULT_FILTERS_MIGRATION_KEY = 'ppt_design_default_filters_v6'
# The Main Cities choice of a City filter: the Main Cities of the workspace whenever the Design is used, so a change
# of the Main Cities in Workspace Config reaches the Designs saved with it.
MAIN_CITIES = '@main-cities'
# Designs saved before kept the Main Cities of that moment: once per workspace, a City filter of exactly the current
# Main Cities becomes the Main Cities choice.
MAIN_CITIES_MIGRATION_KEY = 'ppt_design_main_cities_reference_v1'
DESIGN_PPT_JOBS_TABLE = 'design_ppt_jobs'
# The slide set a user is changing and has not saved, per Design (Designer): it survives a reload and is not
# part of exports, transfers or backups (it is the user's work in progress, not the Design).
DESIGN_DRAFTS_TABLE = 'design_drafts'
# The pictures the Image elements of Designs show (logos, legends, diagrams), by id; they travel with the Designs.
DESIGN_IMAGES_TABLE = 'design_images'
DESIGN_IMAGE_TYPES = {'image/png': 'PNG', 'image/jpeg': 'JPEG', 'image/gif': 'GIF'}
DESIGN_IMAGE_MAX_BYTES = 8 * 1024 * 1024
# The colour theme of PPT Designer: per user in each workspace, as in the other themed modules.
DESIGN_THEME_MODULE = 'ppt_designer'


def design_viewer_tokens(theme_name: str) -> dict[str, str]:
    """The CSS variables of the Design viewer in a theme, set on the viewer and changed at once with the theme: the
    slide title and subtitle in the theme's title style (as on the content slides of the Design PPTs, the part that
    tells the slide apart in its focus style) and the chart cards in its rules."""
    from src.modules.scoring_themes import get_theme
    from src.modules.theme_titles import title_style
    theme = get_theme(theme_name)
    style = title_style(theme)
    family = lambda name: f'"{name}", {theme["font"]}' if name else theme['font']
    weight = lambda bold: '700' if bold else '400'
    return {'--ds-title-font': family(style.get('font')), '--ds-title-color': style['color'],
            '--ds-title-weight': weight(style.get('bold')), '--ds-title-transform': 'uppercase' if style.get('caps') else 'none',
            '--ds-title-style': 'italic' if style.get('italic') else 'normal',
            '--ds-title-focus-font': family(style.get('focus_font') or style.get('font')),
            '--ds-title-focus-color': style['focus_color'], '--ds-title-focus-weight': weight(style.get('focus_bold', style.get('bold'))),
            '--ds-title-focus-style': 'italic' if style.get('focus_italic', style.get('italic')) else 'normal',
            '--ds-subtitle-font': family(style.get('subtitle_font') or style.get('font')),
            '--ds-subtitle-color': style['subtitle_color'], '--ds-subtitle-weight': weight(style.get('subtitle_bold')),
            '--ds-subtitle-style': 'italic' if style.get('subtitle_italic') else 'normal',
            '--ds-card-line': theme['line'], '--ds-accent': theme['title_accent']}


class DesignDefinition(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    template: str = Field(min_length=1)
    # The NR Mode of the CDRs the Design reads, part of its selection: its template serves any NR Mode. NSA+SA
    # (NSA & SA) reads the CDRs of both, whose rows keep the NR Mode of their CDR.
    nr_mode: Literal['NSA', 'SA', 'NSA+SA'] = 'NSA'
    # With NSA & SA, the Campaign filter of the SA CDRs (the Campaign filter chooses the NSA ones); empty, the
    # Campaign filter chooses both. Each has its Latest presets, resolved in its own NR Mode.
    sa_campaigns: list[str] = Field(default_factory=list)
    scope: Literal['single', 'multivendor'] = 'single'
    vendor_comparison: Literal['operator_vendor', 'vendor_only'] = 'operator_vendor'
    datasets: dict[str, list[int]] = Field(default_factory=dict)
    filters: dict[str, list[str]] = Field(default_factory=dict)
    custom_fields: list[str] = Field(default_factory=list)
    hidden_filters: list[str] = Field(default_factory=list)
    slide_comments: dict[str, list[str]] = Field(default_factory=dict)
    # Per slide, the comments added from Generate comments: {'generated': [their texts], 'engine', 'model',
    # 'generated_at', 'generated_by', 'selection': the selection (campaigns, filters and CDRs) they were written for}.
    slide_comment_meta: dict[str, dict[str, Any]] = Field(default_factory=dict)
    date_from: date | Literal['Oldest'] | None = 'Oldest'
    date_to: date | Literal['Newest'] | None = 'Newest'
    # The CDRs follow the NR Mode and the Campaign filter (see ``cdr_auto_selection``); Designs saved
    # before (without it) chose their CDRs by hand and are migrated once (``migrate_auto_cdrs``).
    auto_cdrs: bool = False
    # Its slide set (design_slides): what it changes of its template, by id; its own slide set once detached from its
    # template (``template`` then names the template it came from); the values of the template parameters by id.
    slide_overrides: dict[str, Any] = Field(default_factory=dict)
    own_slides: dict[str, Any] | None = None
    parameter_values: dict[str, Any] = Field(default_factory=dict)

    @model_validator(mode='before')
    @classmethod
    def without_template_nr_mode(cls, value: Any) -> Any:
        """A Design saved, exported or sent with the NR Mode of before (technology, template_technology), or as a Dashboard
        of PPT Dashboards (the fields of its slides were design and design_overrides)."""
        if isinstance(value, dict):
            value = migrate_design_definition(renamed_design(value))
            if isinstance(value, dict) and value.get('nr_mode') is not None:
                value = {**value, 'nr_mode': design_nr_mode_value(value['nr_mode'])}
        return value


class DesignThemeRequest(BaseModel):
    theme: str


AUTO_COMMENT_MODES = ('off', 'empty', 'generated', 'all')


class DesignPptExportRequest(BaseModel):
    definition: DesignDefinition | None = None
    # Auto comments: write generated comments into the PPT (not into the Design) on the slides without comments,
    # in place of the generated ones or in place of every comment.
    auto_comments: Literal['off', 'empty', 'generated', 'all'] = 'off'
    # Their writer (empty: the default one of Super Admin).
    auto_comments_engine: str = ''
    preparation_token: str | None = None
    # Legend entries hidden in the viewer: {"design": {label: hidden},
    # "charts": {chart key: {label: hidden}}}; the PPT shows them hidden too.
    legend_visibility: dict[str, object] | None = None
    selected_operators: list[str] | None = None
    selected_vendors: list[str] | None = None
    selected_regions: list[str] | None = None
    selected_cities: list[str] | None = None


class DesignFilterOptionsRequest(BaseModel):
    definition: DesignDefinition
    field: str = Field(min_length=1, max_length=255)


class DesignFilterOptionsBatchRequest(BaseModel):
    definition: DesignDefinition
    fields: list[str] = Field(min_length=1, max_length=64)


class DesignChartFilterPreviewRequest(BaseModel):
    filters: str = ''
    chart_title: str | None = None
    cdr_source: str | None = None
    dataset_ids: list[int] | str | None = None
    kpi: str | None = None
    chart_type: str | None = None
    rows_split: str | None = None
    columns_split: str | None = None
    legend: str | None = None
    legend_position: str | None = None
    legend_format: str | None = None
    axis_x_range: str | None = None
    axis_y_range: str | None = None
    label_position: str | None = None
    label_format: str | None = None
    exclude_null_empty: str | bool | None = None
    exclude_zero: str | bool | None = None
    dynamic_rows_field: str | None = None
    dynamic_columns_field: str | None = None
    dynamic_field: str | None = None


class DesignComments(BaseModel):
    slide_comments: dict[str, list[str]] = Field(default_factory=dict)
    # Comments added from the generated suggestions of one slide: their slide, texts (as added, after any edit),
    # engine and model, and the preparation they were generated from (its selection marks them out of date later).
    generated_slide: str = ''
    generated_added: int = Field(default=0, ge=0)
    generated_texts: list[str] = Field(default_factory=list)
    generated_engine: str = ''
    generated_model: str = ''
    token: str = ''


class GeneratedCommentChart(BaseModel):
    index: int
    hidden: list[str] = Field(default_factory=list)
    shown: list[str] = Field(default_factory=list)


class GenerateComments(BaseModel):
    token: str
    slide: int
    # The writer: rules, kiro_cli or claude (empty: the default one of Super Admin).
    engine: str = ''
    # The charts the viewer shows on the slide, with the legend items hidden and shown on each one.
    charts: list[GeneratedCommentChart] = Field(default_factory=list)


class GeneratedCommentSlide(BaseModel):
    number: int
    charts: list[GeneratedCommentChart] = Field(default_factory=list)


class GenerateAllComments(BaseModel):
    token: str
    # empty: only the slides without comments; generated: in place of the generated comments, keeping those written by
    # hand; all: in place of every comment.
    mode: Literal['empty', 'generated', 'all'] = 'empty'
    engine: str = ''
    # The slides the viewer shows (a dynamic grid can hide some cards); every content slide when none is given.
    slides: list[GeneratedCommentSlide] = Field(default_factory=list)


class DesignName(BaseModel):
    name: str = Field(min_length=1, max_length=120)


class DesignSlidesPreview(BaseModel):
    """A slide set the Designer shows before it is saved, laid out on the selection of an open Design (``token``):
    the template slide set as edited (Template scope), the Design slide set as edited in This Design scope (its
    overrides are worked out from it, starting from ``overrides``) and the values of the parameters."""
    token: str = Field(min_length=1)
    # The Design being edited: its slide set as saved now (detached or not, its overrides) is the start of the preview,
    # whatever it was when the selection was prepared.
    design_id: str = ''
    template_slide_set: dict[str, Any] | None = None
    slide_set: dict[str, Any] | None = None
    overrides: dict[str, Any] | None = None
    parameter_values: dict[str, Any] | None = None


class DesignSlidesSave(BaseModel):
    """What the Designer saves: the template slide set (Template scope, checked against the template it was read from),
    and the Design slide set edited in This Design scope (or its own slide set once detached)."""
    template_slide_set: dict[str, Any] | None = None
    template_updated_at: str = ''
    slide_set: dict[str, Any] | None = None
    overrides: dict[str, Any] | None = None
    # Keep mine: the overrides flagged Template changed keep their value and stop being flagged.
    keep_template_changes: bool = False


class DesignSlidesTemplate(BaseModel):
    name: str = Field(min_length=1, max_length=120)


class DesignSlidesAddChart(BaseModel):
    """A chart built elsewhere (Chart Builder) added to a Design: on a free place of one of its slides, or on a new
    slide at its end (``slide_id`` empty)."""
    element: dict[str, Any]
    slide_id: str = ''
    slide_title: str = ''


class DesignSlidesDraft(BaseModel):
    """The Designer's unsaved work on a Design: the template slide set, the Design slide set and overrides as edited,
    what of them changed, and the template version they started from."""
    state: dict[str, Any]


class DesignTemplateSelection(BaseModel):
    nr_mode: Literal['NSA', 'SA', 'NSA+SA'] = 'NSA'
    template: str = Field(min_length=1)

    @model_validator(mode='before')
    @classmethod
    def without_template_nr_mode(cls, value: Any) -> Any:
        if isinstance(value, dict):
            mode = value.get('nr_mode', value.get('technology'))
            value = {**{key: item for key, item in value.items() if key != 'technology'},
                     'nr_mode': design_nr_mode_value(mode)}
        return value


def identity(value):
    return column_identity(value)


def template_uids(task_repository) -> dict[str, str]:
    """The identity of each PPT Template of a workspace by name."""
    return task_repository.report_template_uids()


def follow_template_identities(designs: dict, templates: dict[str, str]) -> dict:
    """The Designs with the identity (``template_uid``) of their PPT Template: one whose template was renamed since,
    here or by an import, takes its current name, and one without it (saved before, or imported) takes that of the
    template it names."""
    by_uid = {uid: name for name, uid in templates.items() if uid}
    followed = {}
    for design_id, definition in designs.items():
        if isinstance(definition, dict) and definition.get('template'):
            current = by_uid.get(str(definition.get('template_uid') or ''))
            if current:
                if current != definition['template']:
                    definition = {**definition, 'template': current}
            elif templates.get(str(definition['template'])):
                definition = {**definition, 'template_uid': templates[str(definition['template'])]}
        followed[design_id] = definition
    return followed


def rename_template_in_designs(designs: dict, previous_template: str, template: str) -> list[str]:
    """The Designs that use a renamed template take its new name, and their names follow it; returns their IDs."""
    changed: list[str] = []
    for design_id, raw_definition in designs.items():
        if not isinstance(raw_definition, dict) or str(raw_definition.get('template') or '') != previous_template:
            continue
        updated = dict(raw_definition)
        updated['template'] = template
        name = str(updated.get('name') or '')
        updated['name'] = re.sub(re.escape(previous_template), template, name, flags=re.IGNORECASE)
        designs[design_id] = updated
        changed.append(design_id)
    return changed


def design_nr_mode(definition) -> str:
    """Return the NR Mode (``NSA``, ``SA`` or ``NSA+SA``) of the CDRs a PPT definition reads."""
    if isinstance(definition, dict):
        value = definition.get('nr_mode') or definition.get('technology') or definition.get('template_technology')
    else:
        value = getattr(definition, 'nr_mode', None)
    return design_nr_mode_value(value)


# The options of the Campaign filter of the SA CDRs of an NSA & SA Design (its ``sa_campaigns``).
SA_CAMPAIGN_FIELD = 'SA Campaign'


def entry_uses_nr_mode(entry) -> bool:
    """Whether a chart row filters, splits or labels its series by the NR Mode field."""
    key = identity(NR_MODE_FIELD)
    fields = [*re.split(r'\s*×\s*|\s+\bx\b\s+|,|/', f'{entry.rows_split} × {entry.columns_split} × {entry.legend}'),
              *entry_dynamic_fields(entry)]
    return any(identity(field) == key for field in fields if str(field).strip()) or 'nrmode' in identity(entry.filters or '')


def chart_nr_mode_unread(entry, definition) -> str:
    """Why a chart of one NR Mode shows nothing in a Design that does not read it (empty when it reads it)."""
    modes = chart_filter_nr_modes(entry)
    read = nr_modes_of(design_nr_mode(definition))
    if not entry.source_kind or entry.structural_type or any(mode in read for mode in modes):
        return ''
    return (f"This chart shows the {' & '.join(modes) or 'CDRs of no NR Mode'} CDRs and the Design reads the "
            f"{DESIGN_NR_MODE_LABELS[design_nr_mode(definition)]} ones: choose NR Mode NSA & SA to show it.")


def dataset_matches_nr_mode(dataset, nr_mode) -> bool:
    """Whether a serialized CDR belongs to the requested NR Mode (either one for NSA & SA)."""
    mode = dataset_nr_mode(dataset.get('dataset_kind'), dataset.get('nr_mode'), dataset.get('file_name'))
    return mode is None or mode in nr_modes_of(design_nr_mode_value(nr_mode))


def chart_nr_mode_rows(entry, definition, rows_by_kind):
    """The CDRs a chart reads: in a Design of NSA & SA, those of the one NR Mode its filters keep (a CDR of no known NR
    Mode stays), so a chart of SA reads no NSA CDR; a chart of both NR Modes, or split by NR Mode, reads them all."""
    if not DESIGN_NR_MODE_READS or design_nr_mode(definition) != NSA_VS_SA:
        return rows_by_kind
    modes = [mode for mode in chart_filter_nr_modes(entry) if mode in NR_MODES]
    if len(modes) != 1:
        return rows_by_kind
    narrowed = {kind: [row for row in rows if dataset_matches_nr_mode(row, modes[0])] for kind, rows in rows_by_kind.items()}
    # (Without a CDR of its NR Mode it reads them all and its filter leaves it empty, as before.)
    return narrowed if any(narrowed.values()) else rows_by_kind


FILTER_COLUMNS = {
    'Market': ('market',), 'Operator': ('operator',), 'Operator_Vendor': ('Operator_Vendor',), 'Vendor': ('Vendor',),
    # Vendor_Operator is Operator_Vendor the other way round: it reads that column (see mapped_filter_options).
    'Vendor_Operator': ('Operator_Vendor',),
    'Region': ('Region', 'G_Level_2', 'G Level 2'), 'Cluster': ('Cluster',),
    'City': ('City', 'G_Level_4', 'G Level 4'), 'Campaign': ('Campaign', 'campaign'),
    'Session Type': ('session_type',),
    'RAT': ('RAT_A', 'RAT', 'Sample_RAT_A'),
    'Call Status': ('Call_Status', 'call_status', 'status'),
}
# The CDR catalogue (repository.cdr_catalogue_values) whose spellings a pattern of a filter is matched against.
PATTERN_CATALOGUE_FIELDS = {
    identity('Operator'): 'operators', identity('Vendor'): 'vendors_only', identity('Operator_Vendor'): 'vendors',
    identity('Vendor_Operator'): 'vendor_operators', identity('Region'): 'regions', identity('Cluster'): 'clusters',
    identity('City'): 'cities', identity('Campaign'): 'campaigns',
}
ADAPTATIVE_FILTER_FIELDS = (
    'Operator', 'Operator_Vendor', 'Vendor_Operator', 'Vendor', 'Market', 'Region', 'Cluster', 'City', 'Campaign', 'RAT',
    'Session Type', 'Call Status',
)
DESIGN_RENDER_CACHE_VERSION = 22
DESIGN_SELECTION_CACHE_VERSION = 18
# Only the most recently used Design snapshots keep their chart frames in
# memory; older snapshots rebuild them on demand from the combined tables.
# Also read under its name of PPT Dashboards (DRIVETEST_ANALYZER_DASHBOARD_FRAME_CACHE_SNAPSHOTS).
DESIGN_FRAME_CACHE_SNAPSHOTS = max(1, int(os.environ.get('DRIVETEST_ANALYZER_DESIGN_FRAME_CACHE_SNAPSHOTS')
                                          or os.environ.get('DRIVETEST_ANALYZER_DASHBOARD_FRAME_CACHE_SNAPSHOTS') or 3))
# Grouped compact reads for eligible charts; 0 restores complete reads.
DESIGN_COMPACT_READS = os.environ.get('DRIVETEST_ANALYZER_COMPACT_READS', '1') != '0'
# Per-CDR column cache: any universe or comparison mode is assembled from it in
# memory; 0 restores SQLite reads for every chart.
DESIGN_COLUMN_CACHE = os.environ.get('DRIVETEST_ANALYZER_COLUMN_CACHE', '1') != '0'
# Order-independent charts filter each distinct row of their columns once, with its number of rows (see read_frame).
DESIGN_WEIGHTED_READS = os.environ.get('DRIVETEST_ANALYZER_WEIGHTED_READS', '1') != '0'
# While nobody uses the application for this long, every PPT Design and PPT Template of the open workspace is prepared
# and its chart models worked out and kept (the warm-up, see install_design_routes); DRIVETEST_ANALYZER_DESIGN_WARMUP=0
# turns it off.
DESIGN_WARMUP = os.environ.get('DRIVETEST_ANALYZER_DESIGN_WARMUP', '1') != '0'
DESIGN_WARMUP_IDLE_SECONDS = float(os.environ.get('DRIVETEST_ANALYZER_DESIGN_WARMUP_IDLE_SECONDS') or 5 * 60)
# In a Design of NSA & SA, a chart of one NR Mode (NR Mode IN (SA)) reads only the CDRs of that NR Mode; 0 reads both
# and leaves its filter to drop the rows of the other one (the same chart).
DESIGN_NR_MODE_READS = os.environ.get('DRIVETEST_ANALYZER_NR_MODE_READS', '1') != '0'
# Design source reads: larger page cache, memory-mapped I/O and a sequential
# scan when the selected CDRs hold at least this share of a combined table.
DESIGN_READ_CACHE_KIB = 262_144
DESIGN_READ_MMAP_BYTES = 2 * 1024 ** 3
DESIGN_SEQUENTIAL_SCAN_SHARE = 0.2
DESIGN_SELECTION_CACHE_LIMIT = 128
DESIGN_PROFILE_SELECTION_THRESHOLD = 100_000
DESIGN_CHART_MODEL_CACHE_VERSION = 21
DESIGN_CHART_MODELS_KEPT = 2000
DESIGN_PREVIEW_MANIFESTS_KEPT = 256
DESIGN_CHART_RENDER_WORKERS = max(1, min(2, (os.cpu_count() or 2) - 1))
DESIGN_PREVIEW_MANIFEST_VERSION = 9
DESIGN_DATE_BOUNDS_CACHE_VERSION = 1


def design_cache_dir(workspace: str | Path) -> Path:
    return Path(workspace).parent / '.design-data-cache'


def column_cache_dir(workspace: str | Path) -> Path:
    return design_cache_dir(workspace) / 'columns'


# Template "filters" that configure a chart instead of selecting rows.
CHART_SETTING_FILTERS = frozenset({'threshold', 'buckets', 'binsize', 'classcolours', 'histogramoperator'})
# Chart types whose models only depend on the multiset of rows, so SQLite can
# group identical rows and the loader expands them back before charting.
COMPACT_READ_CHART_TYPES = frozenset({
    'cdf line', 'histogram line', 'histogram bars', 'average vertical bars', 'median vertical bars',
    'distribution stacked vertical bars', 'threshold stacked vertical bars',
    '100% stacked vertical bars', 'count stacked horizontal bars',
})
# Fields the chart pipeline always consults and fields derived from the measure.
COMPACT_READ_IMPLICIT_FIELDS = ('Operator', 'operator', 'Campaign', 'period', 'Period', 'Quarter')
COMPACT_READ_DERIVED_FIELDS = frozenset({'ratebucket', 'valuebucket'})
COMPACT_READ_VENDOR_FIELDS = frozenset({'vendor', 'vendoronly', 'vendorv3', 'operatorvendor', 'opvendor'})


def resolve_filter_column(frame, field):
    field = vendor_filter_column(field)
    columns = {identity(column): column for column in frame.columns}
    aliases = next((values for label, values in FILTER_COLUMNS.items() if identity(label) == identity(field)), (field,))
    candidates = [columns[identity(alias)] for alias in aliases if identity(alias) in columns]
    return next((
        column for column in candidates
        if frame[column].fillna('').astype(str).str.strip().ne('').any()
    ), candidates[0] if candidates else None)


def filter_value_series(frame, field):
    """Resolve geographic aliases per row, skipping empty higher-priority values."""
    field = vendor_filter_column(field)
    columns = {identity(column): column for column in frame.columns}
    aliases = next((values for label, values in FILTER_COLUMNS.items() if identity(label) == identity(field)), (field,))
    candidates = [columns[identity(alias)] for alias in aliases if identity(alias) in columns]
    if not candidates:
        return None
    values = frame[candidates[0]].fillna('').astype(str)
    if identity(field) in {identity('City'), identity('Region')}:
        for column in candidates[1:]:
            empty = values.str.strip().eq('')
            if not empty.any():
                break
            values = values.where(~empty, frame[column].fillna('').astype(str))
    return values


def filter_mask(frame, definition, exclude=None):
    """Filter row indices so adaptive facets never copy complete, wide CDR tables."""
    columns = {identity(column): column for column in frame.columns}
    mask = pd.Series(True, index=frame.index)
    for field, values in definition.filters.items():
        if identity(field) == identity(exclude):
            continue
        field_values = filter_value_series(frame, field)
        if field_values is None:
            return pd.Series(False, index=frame.index)
        if identity(vendor_filter_column(field)) in {'vendor', 'operatorvendor'}:
            values = vendor_match_values(field, values, frame.attrs.get('operator_mappings', {}))
        # Patterns (VF*, Not Lebara) take every value they match (see value_patterns); an explicit empty selection none.
        chosen = ValueSet(values)
        mask &= chosen.mask(field_values.astype(str)) if chosen else pd.Series(False, index=frame.index)
    concrete_from = definition.date_from if isinstance(definition.date_from, date) and not ignore_event_time_filtering() else None
    concrete_to = definition.date_to if isinstance(definition.date_to, date) and not ignore_event_time_filtering() else None
    if concrete_from or concrete_to:
        time_column = next((columns[key] for key in ('eventstarttime', 'teststarttime', 'timestamp', 'datetime', 'date') if key in columns), None)
        if time_column is None:
            return pd.Series(False, index=frame.index)
        times = pd.to_datetime(frame[time_column], errors='coerce', utc=True, format='mixed')
        if concrete_from:
            mask &= times >= pd.Timestamp(concrete_from, tz='UTC')
        if concrete_to:
            mask &= times < pd.Timestamp(concrete_to + timedelta(days=1), tz='UTC')
    return mask


def filter_frame(frame, definition, exclude=None):
    """Apply selections, including explicit empty selections and missing fields."""
    if not definition.filters and not isinstance(definition.date_from, date) and not isinstance(definition.date_to, date):
        return frame
    return frame.loc[filter_mask(frame, definition, exclude)]



def dynamic_chart_values(entry) -> list[list[str]]:
    """Dynamic values of a grid card, each with its displayed form.

    Vendor grids display operator-only values as "<Operator> - All", so
    both forms identify the card when a value is hidden in a legend or axis.
    """
    values = [str(value) for value in (entry.dynamic_row_value, entry.dynamic_column_value) if value is not None]
    if not values and entry.dynamic_value is not None:
        values = [str(entry.dynamic_value)]
    label, _title = chart_title_parts(entry)
    displayed = label.split(' / ') if len(values) > 1 else [label]
    if len(displayed) != len(values):
        displayed = values
    return [list(dict.fromkeys(item for item in pair if item)) for pair in zip(values, displayed)]


def legend_value_hidden(hidden, shown, value: str) -> bool:
    """Whether a label is hidden by itself or by any of its " · " levels.

    Mirrors the Canvas renderer so the viewer, the PPT and its viewer agree.
    """
    hidden, shown = set(hidden), set(shown)
    return value not in shown and any(
        part in hidden and part not in shown for part in (value, *value.split(' · '))
    )


def dynamic_card_hidden(values: list[list[str]], hidden, shown) -> bool:
    """A dynamic grid card disappears when any of its dynamic values is hidden."""
    return any(legend_value_hidden(hidden, shown, value) for candidates in values for value in candidates)


GRID_LAYOUT_PATTERN = re.compile(r'Title \+ (\d+) rows \+ (\d+) columns(?: \+ comments (down|right))?')


def reflowed_dynamic_grid(layout: str, hidden_cards: list[bool]) -> tuple[str, list[int | None]] | None:
    """Close the rows and columns of hidden dynamic grid cards.

    Returns the smaller grid layout ('' when every card is hidden) and, per
    card in row-major order, its new cell or None when hidden. Returns None
    when nothing is hidden or the slide is not a complete grid.
    """
    match = GRID_LAYOUT_PATTERN.fullmatch(canonical_layout_name(layout or ''))
    if not match or not any(hidden_cards):
        return None
    rows, columns, comments = int(match.group(1)), int(match.group(2)), match.group(3) or ''
    if rows * columns != len(hidden_cards):
        return None
    visible_rows = [row for row in range(rows) if not all(hidden_cards[row * columns:(row + 1) * columns])]
    visible_columns = [column for column in range(columns) if not all(hidden_cards[column::columns])]
    cells = [
        None if hidden else visible_rows.index(card // columns) * len(visible_columns) + visible_columns.index(card % columns)
        for card, hidden in enumerate(hidden_cards)
    ]
    return (grid_layout_name(len(visible_rows), len(visible_columns), comments) if visible_rows else ''), cells


def slide_geometry(entries) -> list[list[float]] | None:
    """The places a slide set gives a slide of its own (a free layout: rectangles in percent of the chart area of its
    layout), or None when it uses those of its layout."""
    for entry in entries:
        options = design_slides.row_format(entry).get('slide_options')
        geometry = options.get('geometry') if isinstance(options, dict) else None
        if isinstance(geometry, list) and geometry and all(isinstance(rect, list) and len(rect) == 4 for rect in geometry):
            return [[float(value) for value in rect] for rect in geometry]
    return None


def explicit_places(entries, capacity: int) -> list[int] | None:
    """The place of each chart row of a slide when the slide set gives them (a place can stay empty): None for rows that
    fill the places in order (a template written by hand, or the cards of a dynamic grid, which repeat one element)."""
    rows = [design_slides.row_format(entry) for entry in entries]
    ids = [row.get('id') for row in rows]
    if not rows or not all(ids) or len(set(ids)) != len(ids):
        return None
    places = [row.get('place', position) for position, row in enumerate(rows)]
    if not all(isinstance(place, int) and 0 <= place < capacity for place in places) or len(set(places)) != len(places):
        return None
    return places


def repeat_field(entry) -> str:
    """The field a chart is repeated by in its own place (Repeat by of the Designer: small multiples), or ''."""
    repeat = design_slides.row_format(entry).get('repeat')
    return str(repeat.get('field') or '').strip() if isinstance(repeat, dict) else ''


def subdivided(rect, count: int) -> list[list[float]]:
    """A place divided into a grid of ``count`` smaller places (as square as it can), in reading order."""
    columns = max(1, math.ceil(math.sqrt(count * max(rect[2], 1) / max(rect[3], 1) / 1.6)))
    columns = min(columns, count)
    rows = max(1, math.ceil(count / columns))
    gap_x, gap_y = min(1.2, rect[2] / columns / 10), min(1.6, rect[3] / rows / 10)
    width, height = (rect[2] - gap_x * (columns - 1)) / columns, (rect[3] - gap_y * (rows - 1)) / rows
    return [[rect[0] + (index % columns) * (width + gap_x), rect[1] + (index // columns) * (height + gap_y), width, height]
            for index in range(count)]


def frame_from_rect(rect, frames) -> tuple[int, int, int, int]:
    """A rectangle in percent of the chart area of a layout (its chart frames' bounds) as a frame of the slide (EMU)."""
    left, top = min(frame[0] for frame in frames), min(frame[1] for frame in frames)
    width = max(frame[0] + frame[2] for frame in frames) - left
    height = max(frame[1] + frame[3] for frame in frames) - top
    x, y, w, h = (float(value) for value in rect)
    return int(left + x / 100 * width), int(top + y / 100 * height), int(w / 100 * width), int(h / 100 * height)


def normalized_frame_positions(frames) -> list[list[float]]:
    """Chart frames as percentages of their combined bounds, as the viewer lays them out."""
    if not frames:
        return []
    left, top = min(frame[0] for frame in frames), min(frame[1] for frame in frames)
    width = max(frame[0] + frame[2] for frame in frames) - left
    height = max(frame[1] + frame[3] for frame in frames) - top
    return [[(x - left) / width * 100, (y - top) / height * 100, w / width * 100, h / height * 100] for x, y, w, h in frames]


@dataclass
class Snapshot:
    workspace: str
    owner: str
    entries: list
    frames: dict
    multivendor: bool
    definition: DesignDefinition
    dimensions: tuple
    selection_id: int
    selection_key: str
    payload: dict[str, object]
    chart_frames: dict[int, pd.DataFrame] = field(default_factory=dict)
    filtered_frames: dict[str, pd.DataFrame] = field(default_factory=dict)
    chart_payloads: dict[int, dict[str, object]] = field(default_factory=dict)
    frame_locks: dict[str, RLock] = field(default_factory=dict)
    cancelled: object = None
    # Operator/Vendor mapping fingerprint the cached frames were built with.
    mapping_key: str = ''
    # Last time a user opened one of its charts; the preparations of exports leave it at 0.
    viewed_at: float = 0.0
    # Selected CDR rows and combined source per chart source kind; they do not
    # change while the prepared universe is open.
    sources: dict[str, tuple] = field(default_factory=dict)
    # The values of each filter field in the selection (as read, before the value maps): the Designer lays out the
    # dynamic grids of a slide set it previews on this same selection.
    options: dict[str, list] = field(default_factory=dict)


def install_design_routes(core):
    app = core.app
    lock = RLock()
    snapshots = OrderedDict()
    images = OrderedDict()
    direct_preparation_tasks: dict[str, dict] = {}
    direct_preparation_results: OrderedDict[str, dict] = OrderedDict()
    initialized_ppt_job_databases: set[str] = set()
    design_ppt_runs: dict[tuple[str, int], str] = {}
    design_ppt_data_tokens: dict[tuple[str, int, str], str] = {}
    # Whether each Design has CDRs to show, by the signature of its inputs (see design_status_payload).
    design_status_cache: dict[tuple[str, str], tuple[str, bool]] = {}
    interactive_chart_preview_jobs: OrderedDict[str, dict] = OrderedDict()
    # Jobs writing the comments of every slide of a Design (Generate comments → All slides).
    comment_jobs: OrderedDict[str, dict] = OrderedDict()
    # Kiro Briefs prepared for download (the facts of every slide of a Design).
    comment_briefs: OrderedDict[str, dict] = OrderedDict()
    design_work_gate = Lock()
    # The prepared universes of the Designs (by workspace, Design and definition), so that reopening one (or exporting
    # it) reuses its preparation: the newest DESIGN_PREVIEW_MANIFESTS_KEPT in memory, and all in the cache folder of the
    # workspace (see stored_preview_manifest).
    preview_manifests: OrderedDict[str, dict] = OrderedDict()
    # The Canvas models of the charts: the newest DESIGN_CHART_MODELS_KEPT in memory, and all in the cache folder.
    canvas_models: OrderedDict[str, dict] = OrderedDict()

    def workspace_key():
        if not core.active_workspace:
            raise HTTPException(400, 'Open a workspace before using PPT Designer.')
        return str(Path(core.repository.db_path).resolve())

    def design_user(user=Depends(core.current_user)):
        if core.active_workspace and user.role != 'super-admin' and not core.repository.user_has_workspace_access(user.username, core.active_workspace.id):
            raise HTTPException(403, 'You do not have access to the active workspace.')
        return user

    def design_editor_user(user=Depends(design_user)):
        if user.role not in {'user-editor', 'admin', 'super-admin'}:
            raise HTTPException(403, 'Editor access required.')
        return user

    def bound_repository():
        return Repository(Path(workspace_key()), core.repository.global_db_path)

    def ensure_design_ppt_jobs(task_repository):
        database_key = str(Path(task_repository.db_path).resolve())
        with lock:
            first_check = database_key not in initialized_ppt_job_databases
            initialized_ppt_job_databases.add(database_key)
        with task_repository.connection() as connection:
            connection.execute(f'''CREATE TABLE IF NOT EXISTS {DESIGN_PPT_JOBS_TABLE} (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                design_id TEXT NOT NULL,
                design_name TEXT NOT NULL,
                template_name TEXT NOT NULL,
                nr_mode TEXT NOT NULL DEFAULT 'nsa',
                scope TEXT NOT NULL DEFAULT 'single',
                regions_json TEXT NOT NULL DEFAULT '[]',
                selections_json TEXT NOT NULL DEFAULT '{{}}',
                filters_json TEXT NOT NULL DEFAULT '[]',
                legend_visibility_json TEXT NOT NULL DEFAULT '{{}}',
                output_file TEXT NOT NULL,
                output_path TEXT NOT NULL,
                created_by TEXT NOT NULL,
                created_at TEXT NOT NULL,
                started_at TEXT,
                status TEXT NOT NULL DEFAULT 'queued',
                progress INTEGER NOT NULL DEFAULT 0,
                slide_count INTEGER NOT NULL DEFAULT 0,
                chart_count INTEGER NOT NULL DEFAULT 0,
                last_error TEXT NOT NULL DEFAULT '',
                finished_at TEXT
            )''')
            columns = {str(row['name']) for row in connection.execute(
                f'PRAGMA table_info({DESIGN_PPT_JOBS_TABLE})'
            ).fetchall()}
            if 'vendor_comparison' not in columns:
                connection.execute(f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN vendor_comparison TEXT NOT NULL DEFAULT 'operator_vendor'")
            if 'scope' not in columns:
                connection.execute(
                    f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN scope TEXT NOT NULL DEFAULT 'single'"
                )
            if 'nr_mode' not in columns:
                connection.execute(
                    f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN nr_mode TEXT NOT NULL DEFAULT 'nsa'"
                )
            if 'filters_json' not in columns:
                connection.execute(
                    f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN filters_json TEXT NOT NULL DEFAULT '[]'"
                )
            if 'regions_json' not in columns:
                connection.execute(
                    f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN regions_json TEXT NOT NULL DEFAULT '[]'"
                )
            if 'selections_json' not in columns:
                connection.execute(
                    f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN selections_json TEXT NOT NULL DEFAULT '{{}}'"
                )
            if 'started_at' not in columns:
                connection.execute(
                    f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN started_at TEXT"
                )
            if 'legend_visibility_json' not in columns:
                connection.execute(
                    f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN legend_visibility_json TEXT NOT NULL DEFAULT '{{}}'"
                )
            if 'auto_comments' not in columns:
                connection.execute(f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN auto_comments TEXT NOT NULL DEFAULT 'off'")
            if 'auto_comments_engine' not in columns:
                connection.execute(f"ALTER TABLE {DESIGN_PPT_JOBS_TABLE} ADD COLUMN auto_comments_engine TEXT NOT NULL DEFAULT ''")
            if first_check:
                connection.execute(
                    f'''UPDATE {DESIGN_PPT_JOBS_TABLE}
                        SET status = 'failed', progress = 100,
                            last_error = 'PPT export was interrupted by an application restart.',
                            finished_at = ?
                        WHERE status IN ('queued', 'processing')''',
                    (datetime.now(timezone.utc).isoformat(),),
                )

    def design_ppt_job(task_repository, job_id):
        ensure_design_ppt_jobs(task_repository)
        with task_repository.connection() as connection:
            return connection.execute(
                f'SELECT * FROM {DESIGN_PPT_JOBS_TABLE} WHERE id = ?', (job_id,),
            ).fetchone()

    def update_design_ppt_job(task_repository, job_id, **changes):
        if not changes:
            return
        assignments = ', '.join(f'{key} = ?' for key in changes)
        with task_repository.connection() as connection:
            connection.execute(
                f'UPDATE {DESIGN_PPT_JOBS_TABLE} SET {assignments} WHERE id = ?',
                (*changes.values(), job_id),
            )

    def normalized_legend_visibility(value) -> dict[str, dict]:
        """Keep a bounded {"design": {label: hidden}, "charts": {key: {label: hidden}}} state."""
        def states(raw) -> dict[str, bool]:
            if not isinstance(raw, dict):
                return {}
            return {str(label)[:300]: bool(hidden) for label, hidden in list(raw.items())[:2000]}

        value = value if isinstance(value, dict) else {}
        charts = value.get('charts') if isinstance(value.get('charts'), dict) else {}
        return {
            'design': states(value.get('design')),
            'charts': {str(key)[:600]: states(raw) for key, raw in list(charts.items())[:2000] if states(raw)},
        }

    def chart_legend_keys(entries) -> dict[int, str]:
        """Stable chart keys (slide, title, occurrence) shared by the viewer and the PPT."""
        keys = {}
        occurrences: dict[tuple[int, str], int] = defaultdict(int)
        for index, entry in sorted(enumerate(entries), key=lambda item: (item[1].slide, item[0])):
            if entry.structural_type:
                continue
            title = str(entry.chart_title or '')
            occurrence = occurrences[(entry.slide, title)]
            occurrences[(entry.slide, title)] += 1
            keys[index] = f'{entry.slide}|{title}|{occurrence}'
        return keys

    def chart_legend_lists(visibility, chart_key) -> tuple[list[str], list[str]]:
        """Hidden and explicitly shown legend labels or values of one chart.

        The chart's own choices come first, then the Design-wide ones. The
        renderer hides every series, row or column with a hidden value at any
        aggregation level, exactly as the viewer does.
        """
        chart_states = visibility.get('charts', {}).get(chart_key, {}) if chart_key else {}
        hidden = [label for label, value in chart_states.items() if value]
        shown = [label for label, value in chart_states.items() if not value]
        for label, value in visibility.get('design', {}).items():
            if label not in chart_states:
                (hidden if value else shown).append(label)
        return hidden, shown

    def design_filter_lines(definition: dict, task_repository, selection_labels: dict[str, str] | None = None) -> list[str]:
        lines = []
        for kind in KINDS:
            datasets = design_datasets(definition, kind, task_repository)
            names = [str(dataset.get('file_name') or dataset.get('display_name') or dataset.get('id')) for dataset in datasets]
            if names:
                lines.append(f'CDR {kind.title()}: {", ".join(names)}')
        date_from = '' if ignore_event_time_filtering() else str(definition.get('date_from') or '').strip()
        date_to = '' if ignore_event_time_filtering() else str(definition.get('date_to') or '').strip()
        if date_from and date_to:
            lines.append(f'Date: {date_from} to {date_to}')
        elif date_from:
            lines.append(f'Date from: {date_from}')
        elif date_to:
            lines.append(f'Date to: {date_to}')
        filters = definition.get('filters')
        if not isinstance(filters, dict):
            filters = {}
        if selection_labels is not None:
            # The Campaigns line already reads "Campaign: …" or "Campaigns: …".
            if selection_labels.get('Campaign'):
                lines.append(selection_labels['Campaign'])
            for field in ('Operator', 'Vendor', 'Region', 'City'):
                if selection_labels.get(field):
                    lines.append(f'{field}: {selection_labels[field]}')
        for field, values in filters.items():
            if selection_labels is not None and identity(field) in {'campaign', 'operator', 'vendor', 'vendoronly', 'region', 'city'}:
                continue
            if isinstance(values, (list, tuple, set)):
                selected = [str(value).strip() for value in values if str(value).strip()]
            elif str(values).strip():
                selected = [str(values).strip()]
            else:
                selected = []
            if selected:
                label = vendor_filter_column(field) if identity(vendor_filter_column(field)) in {'vendor', 'operatorvendor'} else re.sub(r'[_-]+', ' ', str(field)).strip().title() or 'Filter'
                lines.append(f'{label}: {", ".join(selected)}')
        return lines

    def design_ppt_job_filters_with_all_selections(filters: list[str]) -> list[str]:
        """Show All Operators and All Vendors for jobs that stored no such line.

        Jobs that recorded their Region or City selection omitted Operator and
        Vendor when nothing specific was selected; that means every value.
        """
        fields = [str(line).partition(':')[0].strip() for line in filters]
        geography = next((index for index, field in enumerate(fields) if field in {'Region', 'City'}), None)
        if geography is None:
            return filters
        missing = [
            f'{field}: {label}'
            for field, label in (('Operator', 'All Operators'), ('Vendor', 'All Vendors'))
            if field not in fields
        ]
        return [*filters[:geography], *missing, *filters[geography:]]

    def serialize_design_ppt_job(row):
        output_path = Path(str(row['output_path'] or ''))
        output_file = str(row['output_file'] or '')
        timestamp_match = re.match(r'^(\d{8}_\d{6})', output_file)
        charts_dir = output_path.parent / 'design-charts'
        ready = str(row['status']) == 'ready' and output_path.is_file()
        charts_ready = ready and (charts_dir / 'manifest.json').is_file()
        job_id = int(row['id'])
        duration = None
        if row['finished_at'] and row['started_at']:
            try:
                duration = max(0, (
                    datetime.fromisoformat(str(row['finished_at']))
                    - datetime.fromisoformat(str(row['started_at']))
                ).total_seconds())
            except ValueError:
                pass
        try:
            filters = json.loads(str(row['filters_json'] or '[]'))
            filters = [str(item) for item in filters if str(item).strip()] if isinstance(filters, list) else []
        except (TypeError, json.JSONDecodeError):
            filters = []
        filters = design_ppt_job_filters_with_all_selections(filters)
        if str(row['scope']) == 'multivendor':
            filters.append('Vendor comparison: ' + ('Vendor' if row['vendor_comparison'] == 'vendor_only' else 'Operator_Vendor'))
        cover = design_ppt_job_cover(row, {})
        if charts_ready and (not cover['regions'] or not cover['cities']):
            try:
                manifest = json.loads((charts_dir / 'manifest.json').read_text(encoding='utf-8'))
            except (OSError, ValueError, json.JSONDecodeError):
                pass
            else:
                cover = design_ppt_job_cover(row, manifest if isinstance(manifest, dict) else {})
        return {
            'id': job_id, 'design_id': str(row['design_id']),
            'design_name': str(row['design_name']),
            'timestamp': timestamp_match.group(1) if timestamp_match else '',
            'template': str(row['template_name'] or ''),
            'nr_mode': str(row['nr_mode'] or 'nsa').upper(),
            'scope': 'Multivendor Comparison' if str(row['scope'] or '').casefold() == 'multivendor' else 'Operator Comparison',
            'cover': cover,
            'filters': filters,
            'generated_by': str(row['created_by']),
            'date': core._local_report_date(row['created_at']).replace(' ', '\n', 1),
            'status': str(row['status']), 'progress': int(row['progress'] or 0),
            'duration_seconds': duration,
            'slides': int(row['slide_count'] or 0), 'charts': int(row['chart_count'] or 0),
            'error': str(row['last_error'] or ''),
            'download_url': f'/api/ppt-designer/ppt-jobs/{job_id}/download' if ready else None,
            'charts_url': f'/api/ppt-designer/ppt-jobs/{job_id}/charts' if charts_ready else None,
            'charts_api_url': f'/api/ppt-designer/ppt-jobs/{job_id}/charts.json' if charts_ready else None,
            'charts_download_url': f'/api/ppt-designer/ppt-jobs/{job_id}/charts.zip' if charts_ready else None,
            'retry_url': f'/api/ppt-designer/ppt-jobs/{job_id}/retry' if str(row['status']) in {'ready', 'failed', 'stopped'} else None,
            'stop_url': f'/api/ppt-designer/ppt-jobs/{job_id}/stop' if str(row['status']) in {'queued', 'processing'} else None,
            'delete_url': f'/api/ppt-designer/ppt-jobs/{job_id}/delete',
        }

    def design_preview_fingerprint(raw_definition, task_repository):
        """Include the current Report Template content in the prepared-preview identity."""
        definition = DesignDefinition.model_validate(
            runtime_design_definition(raw_definition, task_repository)
        )
        # A Design detached from its template reads its own slide set.
        template_content = (json.dumps(definition.own_slides, sort_keys=True).encode() if definition.own_slides
                            else task_repository.report_template_content(definition.template))
        material = {
            'definition': definition.model_dump(mode='json'),
            'template_content': sha256(template_content).hexdigest(),
            'chart_mappings': task_repository.chart_mapping_settings(),
        }
        return sha256(json.dumps(material, sort_keys=True, default=str).encode()).hexdigest()

    def design_export_snapshot(design_id, raw_definition, task_repository):
        workspace = str(Path(task_repository.db_path).resolve())
        definition = DesignDefinition.model_validate(
            runtime_design_definition(raw_definition, task_repository)
        )
        preview = restore_matching_preview_manifest(workspace, design_id, definition)
        snapshot = snapshots.get(str(preview.get('token') or '')) if preview else None
        if snapshot is None:
            raise ValueError('The PPT dataset is not prepared.')
        return validate_design_export_snapshot(snapshot, task_repository)

    def validate_design_export_snapshot(snapshot, task_repository):
        selected_sources(snapshot.definition, task_repository)
        return snapshot

    def design_chart_render_size(placement) -> tuple[int, int]:
        """Return a placeholder-ratio Canvas size with one unchanged logical axis."""
        _left, _top, width, height = placement
        frame_ratio = float(width) / float(height)
        if frame_ratio >= 1600 / 900:
            return min(4096, max(320, round(900 * frame_ratio))), 900
        return 1600, min(4096, max(240, round(1600 / frame_ratio)))

    def design_ppt_scope_label(scope: str) -> str:
        """Return the human-readable scope used on PPT covers."""
        return 'Multivendor Comparison' if str(scope).casefold() == 'multivendor' else 'Operator Comparison'

    def values_present_in_cdrs(values: list[str], available: list[str]) -> list[str]:
        """Keep only values that exist in the selected CDRs.

        Saved filters can hold values of other CDRs (for example an NSA-only
        Operator in an SA Design). When none of the values exist, the
        selection is kept so it never reads as an unrestricted "All".
        """
        present = {str(value).strip().casefold() for value in available if str(value).strip()}
        if not present:
            return values
        kept = [value for value in values if str(value).strip().casefold() in present]
        return kept or values

    def design_ppt_selection_values(
        filters: dict, field: str, requested: list[str], available: list[str],
    ) -> list[str]:
        configured = next((values for key, values in filters.items() if identity(vendor_filter_column(key)) == identity(vendor_filter_column(field))), None)
        values = configured if configured is not None else requested or available
        if isinstance(values, (list, tuple, set)):
            values = list(dict.fromkeys(str(value).strip() for value in values if str(value).strip()))
        else:
            values = [str(values).strip()] if str(values).strip() else []
        return values_present_in_cdrs(values, available)

    def design_ppt_selection_label(
        field: str, values: list[str], available: list[str], alternatives: tuple[list[str], ...] = (),
    ) -> str:
        # "All" means the selection covers every value of the CDR catalogue or
        # of the Design's own filter options (which a saved "select every
        # value" filter comes from), ignoring case and surrounding spaces.
        selected = {str(value).strip().casefold() for value in values if str(value).strip()}
        field = vendor_filter_column(field) if identity(vendor_filter_column(field)) == 'vendor' else field
        all_label = {'Operator': 'All Operators', 'Vendor': 'All Vendors', 'Region': 'All Regions', 'City': 'All Cities'}[field]
        # No specific Operator or Vendor means the export is not restricted
        # by them. Region and City stay blank so covers and file names only
        # name geography the CDRs actually provide.
        if not selected and field in {'Operator', 'Vendor'}:
            return all_label
        candidates = [available, *alternatives]
        if selected and any(
            candidate and selected >= {str(value).strip().casefold() for value in candidate if str(value).strip()}
            for candidate in candidates
        ):
            return all_label
        return ', '.join(values)

    def design_ppt_cover_label(field: str, label: str, count: int) -> str:
        if not label or label.startswith('All '):
            return label
        singular, plural = {'Region': ('Region', 'Regions'), 'City': ('City', 'Cities')}[field]
        return f'{singular if count == 1 else plural}: {label}'

    def design_ppt_job_cover(row, manifest: dict) -> dict[str, str]:
        """Return a PPT job's cover lines, rebuilding them for older jobs."""
        cover = manifest.get('cover')
        if isinstance(cover, dict):
            return {key: str(cover.get(key) or '') for key in ('campaigns', 'scope', 'regions', 'cities')}
        try:
            lines = json.loads(str(row['filters_json'] or '[]'))
        except (TypeError, ValueError, json.JSONDecodeError):
            lines = []
        labels = {}
        for line in lines if isinstance(lines, list) else []:
            field, _separator, value = str(line).partition(': ')
            if field in {'Region', 'City'} and value:
                labels[field] = value
        return {
            'campaigns': '',
            'scope': design_ppt_scope_label(str(row['scope'] or 'single')),
            'regions': design_ppt_cover_label('Region', labels.get('Region', ''), len(labels.get('Region', '').split(', '))),
            'cities': design_ppt_cover_label('City', labels.get('City', ''), len(labels.get('City', '').split(', '))),
        }

    def add_design_ppt_cover_geography(
        slide, campaigns: str, scope: str, regions: str, cities: str, slide_height: int,
    ) -> None:
        """Place the cover Campaigns above the decorative line and the Scope and geography below it."""
        title = next((shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3}), None)
        subtitle = next((shape for shape in slide.placeholders if shape.placeholder_format.type == 4), None)
        # At the left of the title once report_structure.finish places the cover.
        left = report_structure.TITLE_LEFT
        width = title.width if title is not None else Inches(11)
        margin_left = title.text_frame.margin_left if title is not None else 0
        anchor = subtitle or title
        anchor_bottom = anchor.top + anchor.height if anchor is not None else None

        def add_line(name: str, value: str, color: RGBColor, size: int, top: int, height: int) -> None:
            box = slide.shapes.add_textbox(left, top, width, height)
            box.name = f'design-ppt-{name}'
            box.text_frame.margin_left = margin_left
            box.text_frame.margin_top = 0
            box.text_frame.margin_bottom = 0
            box.text_frame.word_wrap = True
            box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            paragraph = box.text_frame.paragraphs[0]
            # White on the red of the cover: the label bold, its values regular.
            label, separator, rest = value.partition(': ')
            for text, bold in ((label + separator, True), (rest, False)) if separator else ((value, True),):
                if not text:
                    continue
                run = paragraph.add_run()
                run.text = text
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = color

        # Like the Design viewer, Scope, Regions and Cities follow the
        # template's decorative line under the subtitle.
        white = RGBColor(255, 255, 255)
        details = [
            (name, value, white)
            for name, value in (('scope', scope), ('region', regions), ('city', cities))
            if value
        ]
        first_top = (
            anchor_bottom + Inches(0.29) if anchor_bottom is not None
            else slide_height - Inches(0.35 + 0.38 * len(details))
        )
        first_top = min(first_top, slide_height - Inches(0.4 + 0.38 * len(details)))
        if campaigns:
            # The larger Campaigns line sits right above the decorative
            # line. Title and subtitle move up just enough to make room for it.
            campaign_height = Inches(0.42)
            separator = next((
                shape for shape in slide.slide_layout.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.LINE and anchor_bottom is not None
                and anchor_bottom - Inches(0.1) <= shape.top <= first_top
            ), None)
            if separator is not None:
                campaign_top = separator.top - Inches(0.06) - campaign_height
                shift = max(0, anchor_bottom + Inches(0.04) - campaign_top)
                for placeholder in (title, subtitle):
                    if placeholder is not None and shift:
                        placeholder.left, placeholder.top, placeholder.width, placeholder.height = (
                            placeholder.left, placeholder.top - shift, placeholder.width, placeholder.height,
                        )
            else:
                campaign_top = anchor_bottom + Inches(0.06) if anchor_bottom is not None else first_top
                first_top = max(first_top, campaign_top + campaign_height + Inches(0.08))
            add_line('campaigns', campaigns, white, 20, campaign_top, campaign_height)
        for index, (name, value, color) in enumerate(details):
            add_line(name, value, color, 16, first_top + Inches(0.38 * index), Inches(0.34))

    def add_design_cell_title(slide, entry, placement):
        """Write a small-cell chart title as slide text and return the remaining chart frame."""
        from pptx.dml.color import RGBColor
        from pptx.util import Emu, Pt

        left, top, width, height = placement
        size = 10 if width >= Emu(2_400_000) else 9
        title_height = min(int(height * .3), int(Pt(size) * 2.6))
        box = slide.shapes.add_textbox(left, top, width, title_height)
        box.name = 'Chart Title'
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Pt(2)
        frame.margin_top = frame.margin_bottom = 0
        label, base = chart_title_parts(entry)
        paragraph = frame.paragraphs[0]
        for text, bold in ((label, True), ((' · ' if label else '') + base, False)):
            if not text:
                continue
            run = paragraph.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(0x1D, 0x33, 0x45)
        return left, top + title_height, width, height - title_height

    def add_design_element(slide, entry, placement, task_repository) -> None:
        """A text box (its Markdown as paragraphs, **bold** kept, lines starting with - or * as bullets) or a picture
        (fitted in its place, keeping its ratio) of a Design slide set."""
        if entry.chart_type.casefold() == design_slides.DRAFT_CHART_TYPE.casefold():
            # A chart still being built: its place stays empty.
            return
        left, top, width, height = placement
        content = design_slides.row_format(entry).get('content')
        content = content if isinstance(content, dict) else {}
        if entry.chart_type.casefold() == 'image':
            data = design_image(task_repository, str(content.get('image') or ''))
            if data is None:
                return
            from PIL import Image as PilImage
            with PilImage.open(BytesIO(data[0])) as picture:
                ratio = picture.width / max(picture.height, 1)
            fitted_width, fitted_height = (width, int(width / ratio)) if width / max(height, 1) < ratio else (int(height * ratio), height)
            slide.shapes.add_picture(BytesIO(data[0]), left + (width - fitted_width) // 2, top + (height - fitted_height) // 2,
                                     fitted_width, fitted_height)
            return
        from pptx.enum.text import PP_ALIGN
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        size = {'small': 12, 'medium': 16, 'large': 22}.get(str(content.get('size') or 'medium').casefold(), 16)
        colour = str(content.get('colour') or '#263B4A').lstrip('#')
        align = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(str(content.get('align') or ''), PP_ALIGN.LEFT)
        lines = str(content.get('text') or '').splitlines() or ['']
        for number, raw in enumerate(lines):
            paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
            paragraph.alignment = align
            text = raw.rstrip()
            heading = re.match(r'^(#{1,3})\s+(.*)$', text)
            bullet = re.match(r'^\s*[-*]\s+(.*)$', text)
            if heading:
                text = heading.group(2)
            elif bullet:
                text = '• ' + bullet.group(1)
            for part_number, part in enumerate(re.split(r'\*\*', text)):
                if not part:
                    continue
                run = paragraph.add_run()
                run.text = part
                run.font.size = Pt(size + (6 if heading else 0))
                run.font.bold = bool(heading) or part_number % 2 == 1
                try:
                    run.font.color.rgb = RGBColor.from_string(colour.upper()[:6])
                except ValueError:
                    pass

    def design_image(task_repository, image_id: str) -> tuple[bytes, str] | None:
        if not re.fullmatch(r'[0-9a-f]{32}', image_id or ''):
            return None
        ensure_design_images(task_repository)
        with task_repository.connection() as connection:
            row = connection.execute(
                f'SELECT data, content_type FROM {DESIGN_IMAGES_TABLE} WHERE id = ?', (image_id,)).fetchone()
        return (bytes(row['data']), str(row['content_type'])) if row else None

    def ensure_design_images(task_repository) -> None:
        with task_repository.connection() as connection:
            design_images.ensure_table(connection)

    @app.post('/api/ppt-designer/images')
    async def upload_design_slides_image(file: UploadFile = File(...), user=Depends(design_editor_user)):
        """Keep a picture for an Image element of a slide set (PNG, JPEG or GIF, up to 8 MB)."""
        task_repository = bound_repository()
        data = await file.read(DESIGN_IMAGE_MAX_BYTES + 1)
        if len(data) > DESIGN_IMAGE_MAX_BYTES:
            raise HTTPException(413, 'The picture is larger than 8 MB.')
        from PIL import Image as PilImage, UnidentifiedImageError
        try:
            with PilImage.open(BytesIO(data)) as picture:
                kind, width, height = picture.format, picture.width, picture.height
        except (UnidentifiedImageError, OSError) as exc:
            raise HTTPException(400, 'Choose a PNG, JPEG or GIF picture (an SVG can be saved as PNG first).') from exc
        content_type = next((mime for mime, name in DESIGN_IMAGE_TYPES.items() if name == kind), None)
        if content_type is None:
            raise HTTPException(400, 'Choose a PNG, JPEG or GIF picture (an SVG can be saved as PNG first).')
        ensure_design_images(task_repository)
        image_id = uuid4().hex
        name = Path(file.filename or 'picture').name[:200]
        with task_repository.connection() as connection:
            connection.execute(
                f'INSERT INTO {DESIGN_IMAGES_TABLE} (id, name, content_type, data, width, height, created_by, created_at) '
                'VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
                (image_id, name, content_type, sqlite3.Binary(data), width, height, user.username,
                 datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')),
            )
        task_repository.add_log(user.username, 'design_image_uploaded', json.dumps({
            'id': image_id, 'name': name, 'type': content_type, 'width': width, 'height': height, 'bytes': len(data),
        }))
        return {'id': image_id, 'name': name, 'width': width, 'height': height, 'url': f'/api/ppt-designer/images/{image_id}'}

    @app.get('/api/ppt-designer/images/{image_id}')
    def design_slides_image(image_id: str, user=Depends(design_user)):
        found = design_image(bound_repository(), image_id)
        if found is None:
            raise HTTPException(404, 'The picture is no longer available.')
        return Response(found[0], media_type=found[1], headers={'Cache-Control': 'private, max-age=86400'})

    def add_design_chart_picture(slide, png: bytes, placement) -> None:
        """Fill a same-ratio chart placeholder without PowerPoint cropping or distortion."""
        left, top, width, height = placement
        slide.shapes.add_picture(BytesIO(png), left, top, width, height)

    def ensure_export_universe_cached(task_repository, workspace, design_id, definition, user, *, cancelled, progress):
        """Prepare (in memory) the universe of the saved Design behind an export when it is not prepared yet."""
        with lock:
            saved = read_designs(task_repository).get(design_id)
        if not isinstance(saved, dict):
            return
        universe = DesignDefinition.model_validate(runtime_design_definition(saved, task_repository))
        universe.datasets = {kind: list(ids) for kind, ids in definition.datasets.items()}
        universe.scope = definition.scope
        universe.date_from, universe.date_to = 'Oldest', 'Newest'
        if restore_matching_preview_manifest(workspace, design_id, universe, materialize=False) is not None:
            return
        try:
            with design_work_gate:
                preview = build_preview(universe, user, workspace=workspace, cancelled=cancelled, progress=progress)
        except HTTPException:
            # The export's own preparation reports an invalid selection.
            return
        with lock:
            snapshot = snapshots.get(preview['token'])
        if snapshot is not None:
            persist_preview_manifest(
                workspace, design_id,
                design_preview_fingerprint(snapshot.definition.model_dump(mode='json'), task_repository),
                preview['token'],
            )

    def auto_slide_comments(snapshot, chart_entries, models, legend_visibility, legend_keys, layout, task_repository,
                            engine: str, username: str) -> tuple[list[str], str]:
        """The comments the writer of a PPT job writes for one of its slides from the chart models of the job, and the
        writer that wrote them (the rules when the LLM writer fails)."""
        charts = []
        for index, entry in chart_entries:
            model = models.get(index)
            if not model or model.get('type') not in design_comments.SUPPORTED_MODEL_TYPES:
                continue
            hidden, shown = chart_legend_lists(legend_visibility, legend_keys.get(index))
            charts.append({
                'index': index, 'title': entry.chart_title or entry.slide_title, 'kpi': entry.kpi, 'model': model,
                'hidden': hidden, 'shown': shown,
                'subject': ' · '.join(value for value in (entry.dynamic_row_value, entry.dynamic_column_value) if value),
            })
        if not charts:
            return [], design_comments.ENGINE_RULES
        suggestions, written = write_slide_comments(charts, layout or '', task_repository, engine, username)
        return [item['text'] for item in suggestions], written['engine']

    def merge_auto_comments(comments: list[str], automatic: list[str], mode: str, generated: list[str]) -> list[str]:
        """The comments of a PPT slide with its Auto comments: only on a slide without comments (``empty``), in place of
        the comments generated before, keeping those written by hand (``generated``), or in place of every comment
        (``all``). A slide the rules have nothing to say about keeps its comments."""
        if not automatic or (mode == 'empty' and comments):
            return comments
        manual = [text for text in comments if text not in set(generated)] if mode == 'generated' else []
        return [*manual, *[text for text in automatic if text not in manual]][:50]

    def render_design_ppt_job(
        job_id, run_token, task_repository, snapshot, definition, user,
        destination, design_id, preview_fingerprint, cover_regions, cover_cities, cover_campaigns='', theme=None,
    ):
        run_key = (str(Path(task_repository.db_path).resolve()), job_id)
        # The colour theme of the titles of the content slides and of the charts: that of the report generating the PPT
        # (a Bundle Report Template or a Reporting Job), else the one the user chose for PPT Designer.
        theme = normalize_theme(theme) if theme else load_user_theme(task_repository, user.username, DESIGN_THEME_MODULE)
        chart_style = canvas_chart_style(theme)

        def run_is_active():
            with lock:
                if design_ppt_runs.get(run_key) != run_token:
                    return False
            current = design_ppt_job(task_repository, job_id)
            return current is not None and str(current['status']) in {'queued', 'processing'}

        try:
            if not run_is_active():
                return
            # The preparation step already marked the job as started.
            update_design_ppt_job(task_repository, job_id, status='processing')
            if snapshot is None:
                update_design_ppt_job(
                    task_repository, job_id, progress=1, last_error='',
                )

                def preparation_progress(percent, _detail):
                    if run_is_active():
                        update_design_ppt_job(
                            task_repository, job_id,
                            # Preparation is visible without dominating the
                            # progress bar; the following Canvas rendering
                            # phase has a known model count and owns most of it.
                            progress=max(12, min(24, 12 + round(percent * 0.12))),
                        )

                workspace = str(Path(task_repository.db_path).resolve())
                # Prepare the chosen CDR universe (the saved Design filters
                # over all dates) before applying the export's own selections.
                ensure_export_universe_cached(
                    task_repository, workspace, design_id, definition, user,
                    cancelled=lambda: not run_is_active(),
                    progress=lambda percent, _detail: run_is_active() and update_design_ppt_job(
                        task_repository, job_id, progress=max(1, min(12, 1 + round(percent * 0.11))),
                    ),
                )
                preview = restore_matching_preview_manifest(workspace, design_id, definition)
                if preview is None:
                    with design_work_gate:
                        preview = build_preview(
                            definition, user, workspace=workspace,
                            cancelled=lambda: not run_is_active(), progress=preparation_progress,
                        )
                token = str(preview['token'])
                with lock:
                    snapshot = snapshots.get(token)
                    design_ppt_data_tokens[(
                        str(Path(task_repository.db_path).resolve()), job_id, user.username,
                    )] = token
                if snapshot is None:
                    raise RuntimeError('The refreshed PPT snapshot could not be created.')
                persist_preview_manifest(
                    str(Path(task_repository.db_path).resolve()), design_id,
                    design_preview_fingerprint(
                        snapshot.definition.model_dump(mode='json'), task_repository,
                    ), token,
                )
                snapshot = validate_design_export_snapshot(snapshot, task_repository)
            update_design_ppt_job(task_repository, job_id, progress=24, last_error='')
            with lock:
                snapshot_token = next((
                    token for token, candidate in snapshots.items() if candidate is snapshot
                ), '')
            if not snapshot_token:
                raise RuntimeError('The prepared PPT snapshot has expired.')
            indexes = [
                index for index, entry in enumerate(snapshot.entries)
                if design_datasets(snapshot.definition, entry.source_kind, task_repository)
            ]
            chart_model_errors: dict[int, str] = {}
            job_row = design_ppt_job(task_repository, job_id)
            auto_mode = str(job_row['auto_comments'] or 'off') if job_row is not None else 'off'
            auto_engine = comment_writers.resolve_engine(
                core.repository, str(job_row['auto_comments_engine'] or '') if job_row is not None else '',
            )
            auto_engines: set[str] = set()
            auto_models: dict[int, dict] = {}
            auto_written: dict[str, list[str]] = {}
            for position, index in enumerate(indexes):
                if not run_is_active():
                    return
                model_progress_start = 25 + round(position * 65 / max(len(indexes), 1))
                model_progress_end = 25 + round((position + 1) * 65 / max(len(indexes), 1))
                update_design_ppt_job(
                    task_repository, job_id,
                    progress=model_progress_start, last_error='',
                )
                model_finished = Event()

                def pulse_canvas_progress():
                    # A single Canvas model can take much longer than the
                    # others.  It has no internal callback, so advance only
                    # within its known model slice while it is genuinely
                    # running; completion still owns the final increment.
                    elapsed_seconds = 0
                    while not model_finished.wait(1):
                        elapsed_seconds += 1
                        span = max(0, model_progress_end - model_progress_start - 1)
                        fraction = min(0.9, elapsed_seconds / 18)
                        progress = model_progress_start + round(span * fraction)
                        if run_is_active():
                            try:
                                update_design_ppt_job(task_repository, job_id, progress=progress, last_error='')
                            except sqlite3.Error:
                                # Progress is advisory.  A transient SQLite
                                # writer contention must never affect the
                                # Canvas model or cause the PPT job to fail.
                                pass

                pulse = Thread(target=pulse_canvas_progress, name=f'design-ppt-model-progress-{job_id}', daemon=True)
                pulse.start()
                try:
                    try:
                        model = chart_model(
                            snapshot_token, index, user,
                            expected_workspace=str(Path(task_repository.db_path).resolve()),
                        )
                        if auto_mode != 'off' and design_comments.needs_rebuild(model):
                            # Auto comments read the facts a model cached before them lacks.
                            model = chart_model(
                                snapshot_token, index, user,
                                expected_workspace=str(Path(task_repository.db_path).resolve()), force=True,
                            )
                        auto_models[index] = model
                    except (HTTPException, ValueError) as exc:
                        # A template can legitimately contain a chart for a
                        # field absent from one selected CDR type.  Retain a
                        # visible placeholder for that chart instead of
                        # discarding the complete PPT job.
                        chart_model_errors[index] = str(getattr(exc, 'detail', exc))
                finally:
                    model_finished.set()
                    pulse.join(timeout=0.1)
                update_design_ppt_job(
                    task_repository, job_id,
                    progress=model_progress_end, last_error='',
                )
            update_design_ppt_job(task_repository, job_id, progress=90, last_error='')
            presentation = Presentation(core.settings.powerpoint_template)
            _remove_all_slides(presentation)
            # Legend entries hidden in the viewer stay hidden in the PPT charts.
            job_row = design_ppt_job(task_repository, job_id)
            try:
                legend_visibility = normalized_legend_visibility(json.loads(str(job_row['legend_visibility_json'] or '{}')))
            except (KeyError, IndexError, TypeError, ValueError):
                legend_visibility = normalized_legend_visibility({})
            legend_keys = chart_legend_keys(snapshot.entries)
            grouped = defaultdict(list)
            for index, entry in enumerate(snapshot.entries):
                grouped[entry.slide].append((index, entry))
            # Dynamic grid cards whose value is hidden leave the slide and the
            # remaining cards fill a smaller grid, exactly as in the viewer; a
            # slide whose cards are all hidden is left out.
            hidden_entries: set[int] = set()
            reflowed_layouts: dict[int, str] = {}
            reflowed_cells: dict[int, dict[int, int]] = {}
            for slide_number, slide_entries in grouped.items():
                cards = [(index, entry) for index, entry in slide_entries if not entry.structural_type]
                hidden_cards = []
                for index, entry in cards:
                    values = dynamic_chart_values(entry)
                    hidden_cards.append(bool(values) and dynamic_card_hidden(
                        values, *chart_legend_lists(legend_visibility, legend_keys.get(index)),
                    ))
                reflow = reflowed_dynamic_grid(slide_entries[0][1].layout, hidden_cards)
                if reflow is None:
                    continue
                reflowed_layouts[slide_number], cells = reflow
                hidden_entries.update(index for (index, _entry), cell in zip(cards, cells) if cell is None)
                reflowed_cells[slide_number] = {
                    index: cell for (index, _entry), cell in zip(cards, cells) if cell is not None
                }
            charts_dir = destination.parent / 'design-charts'
            charts_dir.mkdir(parents=True, exist_ok=True)
            manifest = []
            focus_rows = {
                entry_index: (_entry.template_index if _entry.template_index is not None else editor_index)
                for editor_index, (entry_index, _entry) in enumerate(
                    sorted(enumerate(snapshot.entries), key=lambda item: (item[1].slide, item[0]))
                )
            }
            chart_total = sum(
                1 for entries in grouped.values() for index, entry in entries
                if index not in hidden_entries
                and entry.source_kind and design_datasets(snapshot.definition, entry.source_kind, task_repository)
            )
            rendered = 0
            for slide_number in sorted(grouped):
                if not run_is_active():
                    return
                slide_entries = grouped[slide_number]
                header = slide_entries[0][1]
                _comment_key, comments = design_slide_comments(
                    snapshot.definition, [entry for _index, entry in slide_entries],
                )
                if header.structural_type:
                    layout = _named_slide_layout(presentation, header.layout or 'Title Page')
                    if layout is None:
                        raise ValueError(f"Slide {slide_number}: layout '{header.layout}' is unavailable.")
                    slide = presentation.slides.add_slide(layout)
                    if header.structural_type == 'end slide':
                        _clear_end_slide(slide)
                        continue
                    is_cover = header.structural_type == 'title slide' and slide_number == min(grouped)
                    _set_structural_slide_text(slide, header.slide_title, header.slide_subtitle)
                    if is_cover:
                        add_design_ppt_cover_geography(
                            slide, cover_campaigns, design_ppt_scope_label(snapshot.definition.scope),
                            cover_regions, cover_cities, presentation.slide_height,
                        )
                    else:
                        # The other title and transition rows open the parts of the Design: 01, 02…
                        report_structure.mark(slide, 1)
                    _set_commentary(slide, comments)
                    continue
                chart_entries = [
                    (index, entry) for index, entry in slide_entries
                    if index not in hidden_entries
                    and (entry.chart_type.casefold() in ELEMENT_CHART_TYPES
                         or entry.source_kind and design_datasets(snapshot.definition, entry.source_kind, task_repository))
                ]
                if not chart_entries:
                    continue
                layout_name = reflowed_layouts.get(slide_number) or header.layout
                layout = _named_slide_layout(presentation, layout_name)
                placements = _layout_chart_frames(layout)
                if slide_number in reflowed_cells:
                    if layout is None or len(placements) < len(chart_entries):
                        raise ValueError(f'Slide {slide_number}: the PowerPoint layout has insufficient chart placeholders.')
                    placements = [placements[reflowed_cells[slide_number][index]] for index, _entry in chart_entries]
                elif layout is not None and placements and all(
                        isinstance(design_slides.row_format(entry).get('rect'), list) for _index, entry in chart_entries):
                    # Each chart in its own place (Repeat by), as the viewer lays them out.
                    placements = [frame_from_rect(design_slides.row_format(entry)['rect'], placements) for _index, entry in chart_entries]
                elif layout is not None and placements:
                    # As the viewer lays them out: the places of a free layout, and each element on its own place.
                    slide_rows = [entry for _index, entry in slide_entries if not entry.structural_type]
                    geometry = slide_geometry(slide_rows)
                    places_frames = [frame_from_rect(rect, placements) for rect in geometry] if geometry else placements
                    places = explicit_places(slide_rows, len(places_frames))
                    if places is not None:
                        place_of = {id(entry): place for entry, place in zip(slide_rows, places)}
                        placements = [places_frames[place_of[id(entry)]] for _index, entry in chart_entries]
                    else:
                        placements = places_frames
                if layout is None or len(placements) < len(chart_entries):
                    raise ValueError(f'Slide {slide_number}: the PowerPoint layout has insufficient chart placeholders.')
                slide = presentation.slides.add_slide(layout)
                _set_slide_header(slide, header.slide_title, header.slide_subtitle)
                with use_theme(theme):
                    style_title(slide)
                if auto_mode != 'off':
                    comment_key = design_slide_comment_key([entry for _index, entry in slide_entries])
                    automatic, writer = auto_slide_comments(
                        snapshot, [item for item in chart_entries if item[1].chart_type.casefold() not in ELEMENT_CHART_TYPES],
                        auto_models, legend_visibility, legend_keys, layout_name, task_repository,
                        auto_engine, user.username,
                    )
                    auto_engines.add(writer)
                    merged = merge_auto_comments(comments, automatic, auto_mode,
                                                 (snapshot.definition.slide_comment_meta.get(comment_key) or {}).get('generated') or [])
                    if merged != comments:
                        comments = merged
                        auto_written[comment_key] = merged
                _clear_commentary(slide)
                _set_commentary(slide, comments)
                _remove_template_chart_placeholders(slide)
                compact_titles = compact_grid_layout(layout_name)
                for chart_number, ((index, entry), placement) in enumerate(zip(chart_entries, placements, strict=False), 1):
                    if not run_is_active():
                        return
                    if entry.chart_type.casefold() in ELEMENT_CHART_TYPES:
                        # A text box or picture of the slide set: native, editable PowerPoint shapes.
                        add_design_element(slide, entry, placement, task_repository)
                        continue
                    if compact_titles:
                        # Small grid cells: a readable text title above a title-less chart image.
                        placement = add_design_cell_title(slide, entry, placement)
                    update_design_ppt_job(
                        task_repository, job_id,
                        progress=91 + round(rendered * 4 / max(chart_total, 1)), chart_count=rendered,
                    )
                    render_width, render_height = design_chart_render_size(placement)
                    chart_error = chart_model_errors.get(index)
                    payload = None
                    if chart_error:
                        png = render_unavailable_source_chart(
                            entry, f'Chart unavailable: {chart_error}',
                            width=render_width, height=render_height,
                        )
                        hover_targets = []
                    else:
                        try:
                            # Its own copy, as the renderer may change it (the model stays as the browser has it).
                            payload = {**copy.deepcopy(auto_models[index]), 'style': chart_style}
                            if compact_titles:
                                payload = {**payload, 'hide_title': True}
                            hidden_items, shown_items = chart_legend_lists(legend_visibility, legend_keys.get(index))
                            if hidden_items:
                                payload = {**payload, 'hidden_legend_items': hidden_items, 'shown_legend_items': shown_items}
                            png, hover_targets = _render_canvas_payload(
                                payload, width=render_width, height=render_height,
                            )
                        except Exception as exc:
                            # A renderer defect in one chart must not discard
                            # every other slide in an otherwise valid export.
                            chart_error = f'Chart rendering failed: {exc}'
                            chart_model_errors[index] = chart_error
                            png = render_unavailable_source_chart(
                                entry, chart_error, width=render_width, height=render_height,
                            )
                            hover_targets = []
                    if not run_is_active():
                        return
                    file_name = f'slide-{slide_number:03d}-chart-{chart_number:02d}.png'
                    hover_file = f'slide-{slide_number:03d}-chart-{chart_number:02d}.hover.json'
                    model_file = f'slide-{slide_number:03d}-chart-{chart_number:02d}.model.json'
                    (charts_dir / file_name).write_bytes(png)
                    (charts_dir / hover_file).write_text(
                        json.dumps(hover_targets, ensure_ascii=False), encoding='utf-8',
                    )
                    if payload is not None:
                        (charts_dir / model_file).write_text(
                            json.dumps(payload, ensure_ascii=False, separators=(',', ':')), encoding='utf-8',
                        )
                    add_design_chart_picture(slide, png, placement)
                    manifest.append({
                        'slide': slide_number, 'title': entry.chart_title or header.slide_title,
                        'source': entry.cdr_source, 'chart_type': entry.chart_type,
                        'entry_index': index, 'focus_row': focus_rows[index],
                        'file': file_name, 'hover_file': hover_file,
                        'model_file': model_file if payload is not None else '', 'error': chart_error,
                    })
                    rendered += 1
                    update_design_ppt_job(
                        task_repository, job_id,
                        progress=91 + round(rendered * 4 / max(chart_total, 1)), chart_count=rendered,
                    )
            if not run_is_active():
                return
            viewer_slides = []
            for slide in snapshot.payload.get('slides', []):
                number = slide.get('number')
                if number not in reflowed_layouts:
                    viewer_slides.append(slide)
                    continue
                layout_name = reflowed_layouts[number]
                if not layout_name:
                    continue
                positions = normalized_frame_positions(_layout_chart_frames(_named_slide_layout(presentation, layout_name)))
                charts = []
                for chart in slide.get('charts', []):
                    cell = reflowed_cells[number].get(chart.get('index'))
                    if cell is None:
                        continue
                    charts.append({**chart, 'position': positions[cell]} if cell < len(positions) else dict(chart))
                viewer_slides.append({
                    **slide, 'layout': layout_name, 'compact_titles': compact_grid_layout(layout_name), 'charts': charts,
                })
            (charts_dir / 'manifest.json').write_text(
                json.dumps({
                    'generate_tooltips': True,
                    'design_id': design_id,
                    'preview_fingerprint': preview_fingerprint,
                    # The job's viewer shows the comments of the PPT, Auto comments included.
                    'definition': {**snapshot.definition.model_dump(mode='json'), 'slide_comments': {
                        **snapshot.definition.slide_comments, **auto_written,
                    }},
                    # Keep the exact viewer slide structure alongside the
                    # immutable PNG/Canvas assets so a completed PPT job can
                    # later be opened as its own Design snapshot.
                    'slides': viewer_slides,
                    'charts': manifest,
                    # The exact cover lines, so the job's viewer snapshot shows
                    # the same Campaigns, Scope, Regions and Cities as the PPT cover.
                    'cover': {
                        'campaigns': cover_campaigns,
                        'scope': design_ppt_scope_label(snapshot.definition.scope),
                        'regions': cover_regions, 'cities': cover_cities,
                    },
                }, ensure_ascii=False),
                encoding='utf-8',
            )
            destination.parent.mkdir(parents=True, exist_ok=True)
            update_design_ppt_job(task_repository, job_id, progress=96, chart_count=rendered)
            report_structure.finish(presentation)
            presentation.save(destination)
            if auto_mode != 'off':
                task_repository.add_log(user.username, 'auto_design_comments', json.dumps({
                    'design_id': design_id, 'job_id': job_id, 'mode': auto_mode, 'requested_engine': auto_engine,
                    'engines': sorted(auto_engines), 'slides_written': len(auto_written),
                }))
            if not run_is_active():
                shutil.rmtree(destination.parent, ignore_errors=True)
                return
            update_design_ppt_job(
                task_repository, job_id, status='ready', progress=100,
                slide_count=len(presentation.slides), chart_count=rendered,
                finished_at=datetime.now(timezone.utc).isoformat(),
            )
            if chart_model_errors:
                task_repository.add_log(user.username, 'export_design_ppt_warning', json.dumps({
                    'design_id': design_id, 'job_id': job_id,
                    'placeholder_charts': [
                        {'slide': snapshot.entries[index].slide, 'error': error}
                        for index, error in chart_model_errors.items()
                    ],
                }, ensure_ascii=False))
        except Exception as exc:
            if run_is_active():
                update_design_ppt_job(
                    task_repository, job_id, status='failed', progress=100, last_error=str(exc),
                    finished_at=datetime.now(timezone.utc).isoformat(),
                )
                task_repository.add_log(user.username, 'export_design_ppt_failed', json.dumps({
                    'design_id': design_id, 'job_id': job_id, 'error': str(exc),
                }, ensure_ascii=False))

    def city_filter_key(filters) -> str | None:
        return next((field for field in filters if identity(field) == identity('City')), None)

    def with_main_cities(filters, task_repository):
        """Filters whose City holds the Main Cities choice, with the Main Cities of the workspace now (every city when
        the workspace has none)."""
        key = city_filter_key(filters)
        values = list(filters.get(key) or []) if key else []
        if MAIN_CITIES not in values:
            return filters
        cities = list(dict.fromkeys([*(value for value in values if value != MAIN_CITIES), *task_repository.list_main_cities()]))
        if cities:
            filters[key] = cities
        else:
            filters.pop(key)
        return filters

    def main_cities_reference(definition, main_cities):
        """A Design whose City filter is exactly the Main Cities (saved before as their names) with the Main Cities choice."""
        if not isinstance(definition, dict) or not main_cities:
            return definition
        filters = dict(definition.get('filters') or {})
        key = city_filter_key(filters)
        values = [str(value) for value in filters.get(key) or []] if key else []
        if not values or {value.casefold() for value in values} != {city.casefold() for city in main_cities}:
            return definition
        filters[key] = [MAIN_CITIES]
        return {**definition, 'filters': filters}

    def read_designs(task_repository):
        stored = task_repository.get_workspace_state(STATE_KEY)
        if stored is None:
            stored = task_repository.get_workspace_state(LEGACY_STATE_KEY)
        if stored is None:
            return {}
        designs = json.loads(stored or '{}')
        if not isinstance(designs, dict):
            return designs
        reset_defaults = task_repository.get_workspace_state(DEFAULT_FILTERS_MIGRATION_KEY) != '1'
        try:
            dataset_kinds = {int(row['id']): str(row['dataset_kind'] or '') for row in task_repository.list_datasets()}
        except (AttributeError, sqlite3.Error):
            dataset_kinds = None
        try:
            template_renames, template_uids_now = task_repository.report_template_merge_state()
        except (AttributeError, sqlite3.Error):
            template_renames, template_uids_now = {}, {}
        # Designs of before kept an NR Mode of their own and of their template: it is now part of their selection.
        migrated = {
            design_id: migrate_auto_cdrs(normalize_design_filters(
                migrate_design_definition(
                    definition if dataset_kinds is None else existing_dataset_selection(definition, dataset_kinds),
                    template_renames, template_uids_now),
                reset_defaults=reset_defaults,
            ), task_repository)
            for design_id, definition in designs.items()
        }
        reference_main_cities = task_repository.get_workspace_state(MAIN_CITIES_MIGRATION_KEY) != '1'
        if reference_main_cities:
            main_cities = task_repository.list_main_cities()
            migrated = {design_id: main_cities_reference(definition, main_cities) for design_id, definition in migrated.items()}
        # Each Design refers to its PPT Template by its identity, so a renamed template is still its template.
        try:
            migrated = follow_template_identities(migrated, template_uids(task_repository))
        except (AttributeError, sqlite3.Error):
            pass
        if migrated != designs or reset_defaults or task_repository.get_workspace_state(STATE_KEY) is None:
            task_repository.set_workspace_state(STATE_KEY, json.dumps(migrated))
        if reset_defaults:
            task_repository.set_workspace_state(DEFAULT_FILTERS_MIGRATION_KEY, '1')
        if reference_main_cities:
            task_repository.set_workspace_state(MAIN_CITIES_MIGRATION_KEY, '1')
        return migrated

    def is_template_preview(design_id) -> bool:
        return str(design_id or '').startswith(TEMPLATE_PREVIEW_PREFIX)

    def template_preview_id(username: str) -> str:
        return f'{TEMPLATE_PREVIEW_PREFIX}{sha256(str(username).casefold().encode()).hexdigest()[:16]}'

    def read_template_previews(task_repository) -> dict:
        try:
            previews = json.loads(task_repository.get_workspace_state(TEMPLATE_PREVIEWS_STATE_KEY) or '{}')
        except (TypeError, ValueError):
            previews = {}
        return previews if isinstance(previews, dict) else {}

    def design_exists(task_repository, design_id: str) -> bool:
        """Whether a Design (or a PPT Template open as one) is stored, read as saved: no migration of every Design."""
        if is_template_preview(design_id):
            return design_id in read_template_previews(task_repository)
        for key in (STATE_KEY, LEGACY_STATE_KEY):
            stored = task_repository.get_workspace_state(key)
            if stored is None:
                continue
            try:
                designs = json.loads(stored or '{}')
            except (TypeError, ValueError):
                return design_id in read_designs(task_repository)
            return isinstance(designs, dict) and design_id in designs
        return False

    def read_designs_and_previews(task_repository) -> dict:
        """The Designs and the PPT Templates open as one (for what reads or changes an open Design)."""
        return {**read_designs(task_repository), **read_template_previews(task_repository)}

    def store_design_entry(task_repository, design_id, definition, designs=None) -> None:
        """Store one Design, or the PPT Template open as one, in its own store (``designs``: the Designs just read)."""
        if is_template_preview(design_id):
            previews = read_template_previews(task_repository)
            previews[design_id] = definition
            task_repository.set_workspace_state(TEMPLATE_PREVIEWS_STATE_KEY, json.dumps(previews))
            return
        designs = read_designs(task_repository) if designs is None else designs
        designs[design_id] = definition
        task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))

    def design_slide_comment_key(entries) -> str:
        """Return a position-independent identity for one template slide."""
        signature = [
            {
                'title': entry.slide_title,
                'subtitle': entry.slide_subtitle,
                'layout': entry.layout,
                'structural_type': entry.structural_type or '',
                'chart_title': entry.chart_title,
            }
            for entry in entries
        ]
        encoded = json.dumps(signature, ensure_ascii=False, sort_keys=True, separators=(',', ':'))
        return f"slide:{sha256(encoded.encode('utf-8')).hexdigest()[:24]}"

    def design_slide_comments(definition: DesignDefinition, entries) -> tuple[str, list[str]]:
        """Read stable comments while retaining numeric keys from older Designs."""
        key = design_slide_comment_key(entries)
        legacy_key = str(entries[0].slide) if entries else ''
        comments = definition.slide_comments.get(key, definition.slide_comments.get(legacy_key, []))
        return key, list(comments or [])

    def rendered_slides(entries, nr_mode) -> dict[str, list]:
        """The chart rows of each slide of a template as a Design of an NR Mode shows them (its comments are kept by
        the titles it shows)."""
        slides: dict[str, list] = defaultdict(list)
        for entry in entries:
            slides[str(entry.slide)].append(entry_for_nr_mode(entry, nr_mode))
        return slides

    def reconcile_template_slide_comments(
        template_name: str, old_entries, new_entries,
        task_repository=None, design_ids=None,
    ) -> int:
        """Keep the comments of each slide when a template (or, with ``design_ids``, the slide set of those Designs)
        changes the titles they are kept by: legacy numeric keys move to the stable ones, and a slide whose titles
        changed takes the comments of the slide with the same id (design_slides). Comments of slides the rows do not
        name (the pages of a dynamic grid) stay as they are."""
        if not old_entries:
            return 0
        # Titles written alike before and after (tokens included) read alike in every NR Mode: the Designs need not
        # be read.
        old_raw: dict[str, list] = defaultdict(list)
        new_raw: dict[str, list] = defaultdict(list)
        for entries, slides in ((old_entries, old_raw), (new_entries, new_raw)):
            for entry in entries:
                slides[str(entry.slide)].append(entry)
        if design_ids is None and {slide: design_slide_comment_key(entries) for slide, entries in old_raw.items()} == {
                slide: design_slide_comment_key(entries) for slide, entries in new_raw.items()}:
            return 0
        old_slide_set = new_slide_set = None

        def slide_id(entries) -> str:
            return str(design_slides.row_format(entries[0]).get('slide') or '') if entries else ''

        changed = 0
        with lock:
            source_repository = task_repository or bound_repository()
            designs = read_designs(source_repository)
            for design_id, raw_definition in designs.items():
                if design_ids is not None:
                    if design_id not in design_ids:
                        continue
                    old_rows, new_rows = old_entries, new_entries
                elif str(raw_definition.get('template') or '') != template_name or raw_definition.get('own_slides'):
                    continue
                elif design_slides.has_overrides(raw_definition.get('slide_overrides')):
                    # Its slides are those of the template with its differences.
                    if old_slide_set is None:
                        old_slide_set = design_slides.slides_from_entries(old_entries)
                        new_slide_set = design_slides.slides_from_entries(new_entries)
                    overrides = raw_definition.get('slide_overrides')
                    old_rows = design_slides.entries_from_slides(design_slides.resolve_slides(old_slide_set, overrides)[0])
                    new_rows = design_slides.entries_from_slides(design_slides.resolve_slides(new_slide_set, overrides)[0])
                else:
                    old_rows, new_rows = old_entries, new_entries
                nr_mode = design_nr_mode(raw_definition)
                old_slides, new_slides = rendered_slides(old_rows, nr_mode), rendered_slides(new_rows, nr_mode)
                old_keys = {number: design_slide_comment_key(entries) for number, entries in old_slides.items()}
                new_keys = {number: design_slide_comment_key(entries) for number, entries in new_slides.items()}
                if old_keys == new_keys:
                    continue
                available_keys, known_keys = set(new_keys.values()), set(old_keys.values())
                new_by_id = {slide_id(entries): new_keys[number] for number, entries in new_slides.items() if slide_id(entries)}
                original = raw_definition.get('slide_comments') or {}
                migrated = {
                    key: list(values)
                    for key, values in original.items()
                    if values and key.startswith('slide:') and (key in available_keys or key not in known_keys)
                }
                moved: dict[str, str] = {}
                for legacy_number, entries in old_slides.items():
                    key = old_keys[legacy_number]
                    values = original.get(key, original.get(legacy_number, []))
                    target = key if key in available_keys else new_by_id.get(slide_id(entries))
                    if target and values and (target == key or target not in migrated):
                        migrated[target] = list(values)
                        moved[target] = key
                original_meta = raw_definition.get('slide_comment_meta') or {}
                migrated_meta = {key: value for key, value in original_meta.items()
                                 if key in migrated and (key in available_keys or key not in known_keys)}
                for target, source in moved.items():
                    if target not in migrated_meta and source in original_meta:
                        migrated_meta[target] = original_meta[source]
                if migrated != original or migrated_meta != original_meta:
                    raw_definition['slide_comments'] = migrated
                    raw_definition['slide_comment_meta'] = migrated_meta
                    changed += 1
            if changed:
                source_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
        return changed

    def normalize_design_filters(definition, *, reset_defaults=False):
        """Migrate retired filters and apply the current defaults once per workspace."""
        if not isinstance(definition, dict):
            return definition
        normalized = dict(definition)
        # Scope, source CDRs and dates are stored independently from filters so
        # a Design can restore an explicitly saved Dataset Universe.
        filters = normalized.get('filters')
        retired_filter_keys = {identity('Technology'), identity('Zone')}
        migrated_filters = {}
        region_values = []
        for field, values in (filters.items() if isinstance(filters, dict) else []):
            field_key = identity(field)
            if field_key == identity('Zone'):
                region_values.extend(values or [])
            elif field_key == identity('Region'):
                region_values = list(values or []) + region_values
            elif field_key not in retired_filter_keys:
                migrated_filters[vendor_filter_column(field)] = values
        if region_values:
            migrated_filters['Region'] = list(dict.fromkeys(region_values))
        normalized['filters'] = migrated_filters
        custom_fields = normalized.get('custom_fields')
        normalized['custom_fields'] = [
            vendor_filter_column(field) for field in custom_fields or []
            if identity(field) not in retired_filter_keys and identity(field) != identity('Region')
        ]
        default_keys = {identity(field) for field in ADAPTATIVE_FILTER_FIELDS}
        hidden_filters = normalized.get('hidden_filters')
        normalized['hidden_filters'] = [
            vendor_filter_column(field) for field in hidden_filters or []
            if identity(field) not in retired_filter_keys
            and (not reset_defaults or identity(vendor_filter_column(field)) not in default_keys)
        ]
        return normalized

    def default_campaign_filter(definition) -> list[str]:
        """The Campaign of a new Design: its newest two campaigns, or the newest one for Operator_Vendor."""
        multivendor = definition.get('scope') == 'multivendor' and definition.get('vendor_comparison', 'operator_vendor') == 'operator_vendor'
        return [LATEST_CAMPAIGN if multivendor else LATEST_TWO_CAMPAIGNS]

    def campaign_filter_key(filters) -> str | None:
        return next((field for field in filters if identity(field) == identity('Campaign')), None)

    def follow_nr_mode_campaigns(updated: dict, previous_nr_mode: str, nr_mode: str) -> None:
        """The Campaign filters of a Design whose NR Mode changes: the campaigns of an NR Mode it keeps reading stay
        (the Campaign filter chooses the NSA CDRs and SA Campaign the SA ones with NSA & SA), the others take the newest
        campaigns of the new NR Mode."""
        filters = dict(updated.get('filters') or {})
        key = campaign_filter_key(filters)
        current = list(filters.pop(key) or []) if key else []
        by_mode = {'NSA': current if previous_nr_mode in ('NSA', NSA_VS_SA) else [],
                   'SA': (list(updated.get('sa_campaigns') or []) or current) if previous_nr_mode == NSA_VS_SA
                   else current if previous_nr_mode == 'SA' else []}
        default = default_campaign_filter(updated)
        if nr_mode == NSA_VS_SA:
            filters['Campaign'] = by_mode['NSA'] or default
            updated['sa_campaigns'] = by_mode['SA'] or default
        else:
            filters['Campaign'] = by_mode.get(nr_mode) or default
            updated['sa_campaigns'] = []
        updated['filters'] = filters

    def migrate_auto_cdrs(definition, task_repository):
        """A Design saved when its CDRs were chosen by hand chooses them by Campaign.

        Its saved CDRs become the campaigns they hold, unless it already restricted the Campaign; a
        Design without saved CDRs (which used the newest ones) takes Latest 2 campaigns, or Latest
        campaign for Multivendor Operator_Vendor, as its saved CDRs did.
        """
        if not isinstance(definition, dict) or definition.get('auto_cdrs'):
            return definition
        migrated = dict(definition)
        filters = dict(migrated.get('filters') or {})
        if campaign_filter_key(filters) is None:
            datasets = migrated.get('datasets') if isinstance(migrated.get('datasets'), dict) else {}
            ids = [int(value) for values in datasets.values() for value in (values or []) if str(value).lstrip('-').isdigit()]
            campaigns = sort_campaigns(value for values in cdr_campaigns(task_repository, ids).values() for value in values) if ids else []
            filters['Campaign'] = campaigns or default_campaign_filter(migrated)
        migrated['filters'] = filters
        migrated.pop('datasets', None)
        migrated['auto_cdrs'] = True
        return migrated

    def runtime_design_definition(definition, task_repository):
        """A Design definition with its CDRs (see _runtime_design_definition), the CDRs included worked out once."""
        with cdr_stage.remembered_inclusion():
            return _runtime_design_definition(definition, task_repository)

    def _runtime_design_definition(definition, task_repository):
        """A Design definition with its CDRs: those of its NR Mode and Campaign filter.

        The CDRs are those the combined CDR tables include (the Final CDR of each campaign, else its
        newest Weekly or Daily CDRs, or the Active? choice of Workspace), of the Design's NR
        Mode and of the campaigns of its Campaign filter (every campaign without one). Latest campaign
        and Latest 2 campaigns are replaced by the newest campaigns now.
        """
        effective = migrate_auto_cdrs(normalize_design_filters(definition), task_repository)
        effective.setdefault('scope', 'single')
        nr_mode = design_nr_mode(effective)
        filters = with_main_cities(dict(effective.get('filters') or {}), task_repository)
        key = campaign_filter_key(filters)
        campaigns = list(filters.get(key) or []) if key else []
        if nr_mode == NSA_VS_SA:
            # NSA & SA: the CDRs of each NR Mode, chosen by its own Campaign filter (the SA one, or the Campaign one).
            sa_campaigns = [str(value) for value in effective.get('sa_campaigns') or [] if str(value).strip()] or campaigns
            by_mode = {'NSA': resolve_campaign_filter(task_repository, 'NSA', campaigns),
                       'SA': resolve_campaign_filter(task_repository, 'SA', sa_campaigns)}
            datasets = {kind: [] for kind in KINDS}
            for mode, mode_campaigns in by_mode.items():
                for kind, ids in resolve_cdrs(task_repository, mode, mode_campaigns, KINDS).items():
                    datasets.setdefault(kind, []).extend(ids)
            effective['datasets'] = datasets
            # Their campaigns, unless an NR Mode takes every campaign (its CDRs are already chosen).
            campaigns = sort_campaigns({*by_mode['NSA'], *by_mode['SA']}) if all(by_mode.values()) else []
            if campaigns:
                filters[key or 'Campaign'] = campaigns
            elif key:
                filters.pop(key)
            key = campaign_filter_key(filters)
        else:
            if key and campaign_presets(campaigns):
                campaigns = resolve_campaign_filter(task_repository, nr_mode, campaigns)
                if campaigns:
                    filters[key] = campaigns
                else:
                    # CDRs without campaigns: every CDR of the NR Mode.
                    filters.pop(key)
            effective['datasets'] = resolve_cdrs(task_repository, nr_mode, campaigns, KINDS)
        if key and campaigns:
            # CDRs that hold only chosen campaigns need no Campaign predicate on their rows: the selection
            # stays a complete, unfiltered universe whose counts and options are known.
            ids = [dataset_id for values in effective['datasets'].values() for dataset_id in values]
            held = {campaign_identity(value) for values in cdr_campaigns(task_repository, ids).values() for value in values}
            if held and held <= {campaign_identity(value) for value in campaigns}:
                filters.pop(key)
        effective['filters'] = filters
        effective.setdefault('date_from', 'Oldest')
        effective.setdefault('date_to', 'Newest')
        return effective

    def auto_cdr_definition(definition: DesignDefinition, task_repository) -> DesignDefinition:
        """A Design sent by the browser with the CDRs of its NR Mode and Campaign filter."""
        return DesignDefinition.model_validate(runtime_design_definition(definition.model_dump(mode='json'), task_repository))

    def automatic_scope_universe(definition: DesignDefinition, task_repository) -> DesignDefinition:
        """The Design of a library PPT export, with automatic dates."""
        raw_definition = definition.model_dump(mode='json')
        for field in ('date_from', 'date_to'):
            raw_definition.pop(field, None)
        return DesignDefinition.model_validate(runtime_design_definition(raw_definition, task_repository))

    def cascading_design_options(options, definition, task_repository):
        """The Operator, Vendor, Region, Cluster, City and Campaign lists of every CDR of the NR Mode.

        Each lists only the values the other filters leave, read from the CDR Filter Facets (never the
        CDR rows); the other filters keep the values of the Design's CDRs.
        """
        raw = definition.model_dump(mode='json') if hasattr(definition, 'model_dump') else dict(definition)
        nr_mode = design_nr_mode(raw)
        candidates = [row['id'] for row in candidate_cdrs(task_repository, nr_mode)]
        # The Campaign lists every campaign of the NR Mode; the other filters, the values of the CDRs it chooses.
        resolved = [int(value) for values in (raw.get('datasets') or {}).values() for value in values]
        filters = raw.get('filters') or {}
        fields = [field for field in options if field in CASCADING_FILTER_FIELDS] or list(CASCADING_FILTER_FIELDS)
        cascading = {}
        if 'Campaign' in fields and nr_mode == NSA_VS_SA:
            # NSA & SA: the Campaign filter lists the campaigns of the NSA CDRs and SA Campaign those of the SA ones.
            for mode, field in (('NSA', 'Campaign'), ('SA', SA_CAMPAIGN_FIELD)):
                ids = [row['id'] for row in candidate_cdrs(task_repository, mode)]
                cascading[field] = cascading_options(task_repository, ids, filters, ['Campaign']).get('Campaign', [])
        elif 'Campaign' in fields:
            cascading.update(cascading_options(task_repository, candidates, filters, ['Campaign']))
        if others := [field for field in fields if field != 'Campaign']:
            cascading.update(cascading_options(task_repository, resolved or candidates, filters, others))
        return {**options, **cascading}

    def template_slide_set(template_name, task_repository):
        """The slide set of a template (design_slides) and its rows as written, or (None, None) when it is gone."""
        row = next((row for row in task_repository.list_report_templates() if row['name'] == template_name), None)
        if row is None:
            return None, None
        entries = core.load_template_catalogue(bytes(row['content']), task_repository=task_repository)
        return design_slides.slides_from_entries(entries), entries

    def design_slides_state(definition, task_repository, template=None):
        """The slide set a Design shows: (template slide set or None when detached, resolved slide set, report). ``template``:
        the template slide set to read instead of the stored one (the Designer previews one being edited)."""
        if definition.own_slides:
            slide_set = copy.deepcopy(definition.own_slides)
            return None, slide_set, {'dropped': [], 'template_changed': []}
        if template is None:
            try:
                template, _entries = template_slide_set(definition.template, task_repository)
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        if template is None:
            raise HTTPException(400, 'The selected template is no longer available.')
        resolved, report = design_slides.resolve_slides(template, definition.slide_overrides)
        return template, resolved, report

    def catalogue(definition, task_repository=None, template=None):
        """The chart rows of a Design: its template with its overrides (or its own slide set once detached) and the
        values of its parameters, their titles naming the NR Mode it reads. ``template``: see design_slides_state."""
        task_repository = task_repository or core.repository
        nr_mode = design_nr_mode(definition)
        if template is None and not definition.own_slides and not design_slides.has_overrides(definition.slide_overrides):
            rows = task_repository.list_report_templates()
            row = next((row for row in rows if row['name'] == definition.template), None)
            if row is None:
                raise HTTPException(400, 'The selected template is no longer available.')
            try:
                entries = core.load_template_catalogue(bytes(row['content']), task_repository=task_repository)
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
            parameters = (design_slides.slides_from_entries(entries)['parameters']
                          if any('"parameters"' in entry.format for entry in entries) else [])
        else:
            _template, resolved, _report = design_slides_state(definition, task_repository, template)
            parameters = resolved.get('parameters') or []
            try:
                # Checked as a template written by hand would be (a slide of the Design included).
                entries = core.load_template_catalogue(
                    core.catalogue_csv(design_slides.entries_from_slides(resolved)), task_repository=task_repository)
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        entries = design_slides.apply_parameters(entries, parameters, definition.parameter_values)
        entries = design_slides.apply_tokens(entries, selection_token_texts(definition, entries, task_repository))
        # The titles name the NR Mode as each chart shows it (entry_for_nr_mode); a text box names that of the Design.
        mode_texts = {token.strip('{}').casefold(): text for token, text in NR_MODE_TOKEN_VALUES[design_nr_mode_value(nr_mode)].items()}
        entries = design_slides.apply_tokens(entries, mode_texts, titles=False)
        return [entry_for_nr_mode(entry, nr_mode) for entry in entries]

    # The field behind each selection token of the titles and text boxes, and what it reads when not filtered.
    SELECTION_TOKEN_FIELDS = {
        'operators': ('Operator', 'All Operators'), 'vendors': ('Vendor', 'All Vendors'), 'regions': ('Region', 'All Regions'),
        'clusters': ('Cluster', 'All Clusters'), 'cities': ('City', 'All Cities'),
    }

    def selection_token_texts(definition, entries, task_repository) -> dict[str, str]:
        """What the selection tokens the rows name read for a Design (design_slides.SELECTION_TOKENS): its name,
        its Scope, its Campaigns and the Operators, Vendors, Regions, Clusters and Cities it is filtered by, or All… when
        it is not filtered by them or keeps every value of its CDRs. Only the tokens named are read."""
        wanted = design_slides.used_tokens(entries) & ({token.casefold() for token in design_slides.SELECTION_TOKENS}
                                                       | set(design_slides.LEGACY_SELECTION_TOKENS))
        if not wanted:
            return {}
        texts: dict[str, str] = {}
        if wanted & {'design', 'dashboard'}:
            texts['design'] = texts['dashboard'] = definition.name
        if 'scope' in wanted:
            texts['scope'] = design_ppt_scope_label(definition.scope)
        try:
            selected_by_kind = selected_sources(definition, task_repository) if wanted - {'design', 'dashboard', 'scope'} else {}
        except (HTTPException, ValueError):
            # No CDR to read yet (a slide set checked before its CDRs are ready): the fields read as not filtered.
            selected_by_kind = {}
        if 'campaigns' in wanted:
            campaigns = design_campaign_values(definition, selected_by_kind, task_repository, wait=False)[0]
            texts['campaigns'] = ', '.join(campaigns) or 'All Campaigns'
        fields = [token for token in SELECTION_TOKEN_FIELDS if token in wanted]
        if not fields:
            return texts
        filters = definition.filters if isinstance(definition.filters, dict) else {}
        selected_ids = [int(row['id']) for rows in selected_by_kind.values() for row in rows]
        if selected_ids:
            core.backfill_cdr_catalogues(selected_ids, task_repository)
        catalogue_values = dict(task_repository.cdr_catalogue_values(selected_ids)) if selected_ids else {}
        for token in fields:
            field, all_label = SELECTION_TOKEN_FIELDS[token]
            configured = next((values for key, values in filters.items()
                               if identity(vendor_filter_column(key)) == identity(vendor_filter_column(field))), None)
            values = configured if isinstance(configured, list) else ([configured] if configured else [])
            values = list(dict.fromkeys(str(value).strip() for value in values if str(value).strip()))
            if not values:
                texts[token] = all_label
                continue
            if token == 'operators' and selected_ids:
                # As the Operators of its filters name them (Operator Maps).
                available = profile_filter_options(unfiltered_definition(definition), [], selected_by_kind, ['Operator'], task_repository)['Operator']
            else:
                available = catalogue_values.get('vendors_only' if token == 'vendors' else token) or []
            values = values_present_in_cdrs(values, available)
            every = {str(value).strip().casefold() for value in available if str(value).strip()}
            texts[token] = all_label if every and {value.casefold() for value in values} >= every else ', '.join(values)
        return texts

    def validate(definition, task_repository=None):
        if not definition.name.strip():
            raise HTTPException(400, 'Enter a PPT name.')
        if (
            not ignore_event_time_filtering()
            and isinstance(definition.date_from, date)
            and isinstance(definition.date_to, date)
            and definition.date_from > definition.date_to
        ):
            raise HTTPException(400, 'The start date must not follow the end date.')
        if set(definition.datasets) - set(KINDS):
            raise HTTPException(400, 'Unsupported dataset type.')
        entries = catalogue(definition, task_repository)
        return entries

    def add_dynamic_grid_metadata(payload, entries, deck=None) -> None:
        """Add each grid card's dynamic values and the smaller grids it can reflow to.

        Hiding a dynamic value removes its grid row or column; the viewer lays
        the remaining cards out on the same smaller grid the PPT export uses.
        """
        grids = set()
        for slide in payload.get('slides', []):
            for chart in slide.get('charts', []):
                index = chart.get('index')
                values = dynamic_chart_values(entries[index]) if isinstance(index, int) and 0 <= index < len(entries) else []
                if values:
                    chart['dynamic_values'] = values
            grid = GRID_LAYOUT_PATTERN.fullmatch(canonical_layout_name(slide.get('layout') or ''))
            if grid and any(chart.get('dynamic_values') for chart in slide.get('charts', [])):
                grids.add((int(grid.group(1)), int(grid.group(2)), grid.group(3) or ''))
        names = {
            grid_layout_name(rows, columns, comments)
            for row_limit, column_limit, comments in grids
            for rows in range(1, row_limit + 1) for columns in range(1, column_limit + 1)
        }
        if names and deck is None:
            deck = Presentation(core.settings.powerpoint_template)
        payload['grid_layouts'] = {
            name: normalized_frame_positions(_layout_chart_frames(_named_slide_layout(deck, name)))
            for name in sorted(names)
        }

    def expand_design_layouts(entries, definition, options, task_repository):
        effective_entries = [prepare_multivendor_catalog_entry(entry, definition.vendor_comparison) if definition.scope == "multivendor" else entry for entry in entries]
        dynamic_fields = {field for entry in effective_entries for field in entry_dynamic_fields(entry) if field}
        mappings = task_repository.chart_mapping_settings()['operator_mappings']
        values_by_field = {}
        for field in dynamic_fields:
            vendor_mappings = task_repository.chart_mapping_settings()['vendor_mappings']
            normalize = ((lambda value: _normalise_operator_label(value, mappings)) if identity(field) == 'operator' else
                         (lambda value: _normalise_vendor(value, mappings, vendor_mappings)) if identity(field) in {'vendor', 'vendoronly', 'operatorvendor'} else str)
            raw_values = next((values for key, values in options.items() if identity(key) == identity(field)), [])
            values = {normalize(value) for value in raw_values if str(value).strip()}
            selected_values = next((items for key, items in definition.filters.items() if identity(key) == identity(field)), [])
            if selected_values:
                values &= {normalize(value) for value in selected_values}
            if definition.scope == 'multivendor' and definition.vendor_comparison != 'vendor_only' and identity(field) in {'operatorvendor', 'vendor'}:
                values = {value for value in values if not any(term in value.casefold() for term in ('mixed', 'other'))}
            values_by_field[field] = sorted(values, key=str.casefold)
        vendor_families = {}
        if definition.scope == 'multivendor' and definition.vendor_comparison == 'operator_vendor':
            settings = task_repository.chart_mapping_settings()
            with task_repository.connection() as connection:
                for kind, ids in definition.datasets.items():
                    if not ids:
                        continue
                    columns = task_repository.list_reporting_row_columns(kind)
                    vendor_column = resolve_sql_column(columns, 'Operator_Vendor')
                    family_column = resolve_sql_column(columns, 'Vendor')
                    if not vendor_column or not family_column:
                        continue
                    quote = task_repository._quote_identifier
                    with cdr_store.combined_scope(ids):
                        rows = connection.execute(
                            f'SELECT DISTINCT {quote(vendor_column)}, {quote(family_column)} FROM {quote(task_repository.reporting_rows_table_name(kind))} '
                            f'WHERE dataset_id IN ({", ".join("?" for _ in ids)})', ids,
                        ).fetchall()
                    for vendor, family in rows:
                        if vendor and family:
                            vendor_families[_normalise_vendor(vendor, mappings, settings['vendor_mappings'])] = str(family)
        chart_mapping_settings = task_repository.chart_mapping_settings()
        return expand_dynamic_layouts(entries, values_by_field, multivendor=definition.scope == 'multivendor',
                                      operator_mappings=mappings, vendor_mappings=chart_mapping_settings['vendor_mappings'], vendor_comparison=definition.vendor_comparison, vendor_families=vendor_families,
                                      mapping_groups=chart_mapping_settings)

    @app.get('/ppt-designer', response_class=HTMLResponse)
    def page(request: Request, user=Depends(design_user)):
        workspace_key()
        ready = [core.serialize_dataset_row(row) for row in core.repository.list_datasets() if row['status'] == 'ready']
        core.backfill_cdr_catalogues(row['id'] for row in ready if row.get('dataset_kind') in KINDS)
        catalogue = core.repository.cdr_catalogue_values(
            row['id'] for row in ready if row.get('dataset_kind') in KINDS
        )
        return core.render_template(request, 'ppt_designer.html', {
            'user': user, 'design_datasets': {kind: [row for row in ready if row.get('dataset_kind') == kind] for kind in KINDS},
            # Any template serves the CDRs of any NR Mode.
            'design_templates': [{'name': row['name'], 'identifier': row['identifier'], 'active': row['active']}
                                    for row in core.report_catalogue_options()],
            'design_workspace_id': core.active_workspace.id,
            'calculated_dimensions': core.calculated_dimensions_json(core.load_workspace_calculated_dimensions()),
            'design_filter_fields': ADAPTATIVE_FILTER_FIELDS,
            'design_filter_aliases': {
                field: list(aliases) for field, aliases in FILTER_COLUMNS.items() if len(aliases) > 1
            },
            'design_ignore_event_time_filtering': ignore_event_time_filtering(),
            'design_multivendor_available': len(catalogue['vendors']) >= 2,
            'dashboard_main_cities': core.repository.list_main_cities(),
            # The colour theme of the charts and slide titles of the viewer and of the Design PPTs.
            'design_themes': themes_payload()['themes'],
            'design_theme': load_user_theme(bound_repository(), user.username, DESIGN_THEME_MODULE),
            'canvas_chart_styles': {key: canvas_chart_style(key) for key in THEMES},
            'design_viewer_tokens': {key: design_viewer_tokens(key) for key in THEMES},
            'design_blank_template': design_slides.BLANK_PPT_TEMPLATE,
            **templates_library_context(),
            'catalogue_notice': request.query_params.get('catalogue_notice') or None,
            'catalogue_error': request.query_params.get('catalogue_error') or None,
        })

    def templates_library_context() -> dict:
        """PPT Templates (the button of the library): the templates of the workspace, by name."""
        catalogues = core.report_catalogue_options()
        return {
            'workspace_catalogues': sorted((
                {**catalogue, 'created_at': core.format_local_timestamp(catalogue.get('created_at')),
                 'updated_at': core.format_local_timestamp(catalogue.get('updated_at'))}
                for catalogue in catalogues
            ), key=lambda item: (core.nr_mode_free_sort_name(item['name']), str(item['name']).casefold())),
            'template_names': [str(catalogue['name']) for catalogue in catalogues],
        }

    @app.get('/api/ppt-designer/templates-library')
    def templates_library(user=Depends(design_editor_user)):
        """PPT Templates as they are now (a template saved from the Designer meanwhile included): the list of the dialog
        and the templates a new PPT Design can choose."""
        workspace_key()
        context = {'user': user, **templates_library_context(), 'catalogue_notice': None, 'catalogue_error': None}
        return {
            'html': core.templates.get_template('ppt_templates_panel.html').render(context),
            'templates': [{'name': row['name'], 'identifier': row['identifier'], 'active': row['active']}
                          for row in core.report_catalogue_options()],
        }

    @app.get('/api/ppt-designer/theme')
    def design_theme(user=Depends(design_user)):
        """The colour theme the user chose for the PPT Designer charts and slide titles in this workspace."""
        theme = load_user_theme(bound_repository(), user.username, DESIGN_THEME_MODULE)
        return {'theme': theme, 'style': canvas_chart_style(theme)}

    @app.put('/api/ppt-designer/theme')
    def design_theme_save(payload: DesignThemeRequest, user=Depends(design_user)):
        if payload.theme.strip().casefold() not in THEMES:
            raise HTTPException(400, f'Unknown theme: {payload.theme}.')
        task_repository = bound_repository()
        previous = load_user_theme(task_repository, user.username, DESIGN_THEME_MODULE)
        theme = save_user_theme(task_repository, user.username, payload.theme, DESIGN_THEME_MODULE)
        if theme != previous:
            task_repository.try_add_log(user.username, 'ppt_designer_theme_changed', json.dumps({'theme': theme, 'previous': previous}))
        return {'theme': theme, 'style': canvas_chart_style(theme)}

    @app.get('/api/ppt-designer')
    def list_designs(user=Depends(design_user)):
        workspace_key()
        with lock:
            task_repository = bound_repository()
            designs = read_designs(task_repository)
            known = creators(task_repository, PPT_DESIGNER)
            # The PPT Template the user has open as a Design (PPT Templates → Edit), which the page never lists.
            preview_id = template_preview_id(user.username)
            preview = read_template_previews(task_repository).get(preview_id)
        if isinstance(preview, dict):
            designs = {**designs, preview_id: {**preview, 'template_preview': True}}
            known = {**known, preview_id: user.username}
        # Each Design with who created it: only they, admins and super-admins delete it.
        return JSONResponse(
                {design_id: {**definition, 'created_by': known.get(str(design_id).casefold(), '')}
                 if isinstance(definition, dict) else definition for design_id, definition in designs.items()},
            headers={'Cache-Control': 'no-store, max-age=0, must-revalidate'},
        )

    @app.get('/api/ppt-designer/{design_id}/export')
    def export_design_definition(design_id: str, user=Depends(design_user)):
        """Create the same Design archive accepted by Admin Import."""
        task_repository = bound_repository()
        with lock:
            definition = read_designs(task_repository).get(design_id)
        if not isinstance(definition, dict):
            raise HTTPException(404, 'Design not found.')
        workspace = core.active_workspace
        workspace_name = str(workspace.name if workspace else 'Workspace')
        archive_path = f'workspaces/{workspace_name}/designs/designs.json'
        # (With the pictures of its own slides, which another workspace does not have.)
        payload = json.dumps({
            'format': 'drivetest-analyzer-designs', 'version': 1,
            'designs': {design_id: definition},
            'images': design_images.portable_images(task_repository, design_images.referenced_images(definition)),
        }, ensure_ascii=False, indent=2).encode('utf-8')
        temporary = tempfile.NamedTemporaryFile(prefix='design-export-', suffix='.zip', delete=False)
        temporary.close()
        archive = Path(temporary.name)
        try:
            with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED, compresslevel=1) as bundle:
                bundle.writestr('manifest.json', json.dumps({
                    'format': 'drivetest-analyzer-export', 'version': 1, 'kind': 'designs',
                    'components': ['workspace_components'], 'workspace_components': ['designs'],
                    'source_workspace': {'id': workspace.id, 'name': workspace_name} if workspace else {},
                    'archive_path': archive_path,
                }, indent=2, sort_keys=True))
                bundle.writestr(archive_path, payload)
        except Exception:
            archive.unlink(missing_ok=True)
            raise
        safe_name = re.sub(r'[^A-Za-z0-9._-]+', '_', str(definition.get('name') or design_id)).strip('._') or 'Design'
        task_repository.try_add_log(user.username, 'export_design_definition', json.dumps({
            'id': design_id, 'name': str(definition.get('name') or ''),
        }))
        return FileResponse(
            archive, filename=f'{safe_name}_design.zip', media_type='application/zip',
            background=BackgroundTask(archive.unlink, missing_ok=True),
        )

    def design_ppt_output_dir(task_repository) -> Path:
        return module_output_dir(Path(task_repository.db_path).parent / 'output', DESIGNS_FOLDER)

    def remove_design_ppt_job_folder(task_repository, output_path) -> None:
        """Delete one job folder, never anything outside the Design output root."""
        text = str(output_path or '').strip()
        if not text:
            return
        folder = Path(text).parent.resolve()
        root = design_ppt_output_dir(task_repository).resolve()
        if folder == root or root not in folder.parents:
            return
        shutil.rmtree(folder, ignore_errors=True)

    def design_ppt_campaign_suffix(campaigns: list[str]) -> str:
        """Return ``2026-Q1_vs_2026-Q2`` for exports that compare several Campaigns."""
        return '_vs_'.join(campaigns) if len(campaigns) > 1 else ''

    def design_ppt_unique_output(
        task_repository, export_time, nr_mode_label, design_name, scope_label, zone_label, job_id=None,
        campaign_suffix='',
    ):
        """Return a unique ``timestamp - NR Mode - Design - Scope - Regions - Campaigns`` file and folder."""
        for offset in range(60):
            output_file = build_cdr_report_filename(
                export_time + timedelta(seconds=offset), nr_mode_label, design_name,
                scope_label, zone_label, campaign_suffix,
            )
            job_dir = design_ppt_output_dir(task_repository) / Path(output_file).stem
            with task_repository.connection() as connection:
                already_queued = connection.execute(
                    f'SELECT 1 FROM {DESIGN_PPT_JOBS_TABLE} WHERE output_file = ? AND id IS NOT ? LIMIT 1',
                    (output_file, job_id),
                ).fetchone()
            if not already_queued and not job_dir.exists():
                return output_file, job_dir / output_file
        raise HTTPException(503, 'No unique PowerPoint filename is available for this export.')

    def queue_design_ppt_export(
        design_id, user, *, reuse_job_id=None, export_definition: DesignDefinition | None = None,
        preparation_token: str | None = None, selected_operators: list[str] | None = None,
        selected_vendors: list[str] | None = None, selected_regions: list[str] | None = None,
        selected_cities: list[str] | None = None, legend_visibility: dict | None = None, theme: str | None = None,
        auto_comments: str | None = None, auto_comments_engine: str | None = None,
    ):
        """Create the export job immediately; all preparation runs in the background job (``theme``: the colour theme
        of a report generating it, else the user's PPT Designer theme; ``auto_comments``: the Auto comments mode, kept
        by a relaunched job when None)."""
        task_repository = bound_repository()
        workspace = workspace_key()
        # Preparation never blocks an export: the queued job waits for the
        # shared Design work slot and prepares its universe.
        with lock:
            stored_definition = read_designs(task_repository).get(design_id)
        if not isinstance(stored_definition, dict):
            raise HTTPException(404, 'Design not found.')
        raw_definition = export_definition.model_dump(mode='json') if export_definition is not None else dict(stored_definition)
        # Its slide set is the one saved (the Designer saves it), whatever the viewer sends.
        raw_definition.update({key: copy.deepcopy(stored_definition.get(key, default))
                               for key, default in (('slide_overrides', {}), ('own_slides', None))})
        if design_nr_mode(raw_definition) != design_nr_mode(stored_definition):
            # The comments of a Design were written for the CDRs of its own NR Mode: generated in another one (by a
            # Reporting Job or a Bundle Report Template), its slides go without them, unless auto comments write them.
            raw_definition = {**raw_definition, 'slide_comments': {}, 'slide_comment_meta': {}}
        explicit_selections = {
            field: list(dict.fromkeys(str(value).strip() for value in selected if str(value).strip()))
            for field, selected in (
                ('Operator', selected_operators), ('Vendor', selected_vendors),
                ('Region', selected_regions), ('City', selected_cities),
            ) if selected is not None
        }
        # An active snapshot supplied by the viewer is only looked up in memory.
        snapshot = None
        if preparation_token and not explicit_selections:
            with lock:
                candidate = snapshots.get(preparation_token)
            if (
                candidate is not None
                and candidate.workspace == workspace
                and candidate.owner in {user.username, '*'}
                and (export_definition is None or candidate.definition.scope == export_definition.scope and candidate.definition.vendor_comparison == export_definition.vendor_comparison)
            ):
                snapshot = candidate
        design_name = str(raw_definition.get('name') or design_id)
        nr_mode = design_nr_mode(raw_definition)
        # The file names the NR Mode the Design reads (NSA & SA for both).
        nr_mode_label = DESIGN_NR_MODE_LABELS[nr_mode]
        scope = 'multivendor' if str(raw_definition.get('scope') or 'single') == 'multivendor' else 'single'
        export_time = datetime.now()
        ensure_design_ppt_jobs(task_repository)
        # A provisional, always valid output path; the background job adds the
        # selected Regions once it has resolved them.
        output_file, destination = design_ppt_unique_output(
            task_repository, export_time, nr_mode_label, design_name, design_ppt_scope_label(scope), '', reuse_job_id,
        )
        common = {
            'design_name': design_name, 'template_name': str(raw_definition.get('template') or ''),
            'nr_mode': nr_mode.casefold(), 'scope': scope, 'filters_json': '[]',
            'output_file': output_file, 'output_path': str(destination),
            'vendor_comparison': raw_definition.get('vendor_comparison') or 'operator_vendor',
        }
        if reuse_job_id is None:
            with task_repository.connection() as connection:
                cursor = connection.execute(
                    f'''INSERT INTO {DESIGN_PPT_JOBS_TABLE} (
                        design_id, design_name, template_name, nr_mode, scope, regions_json, selections_json, filters_json, output_file, output_path,
                        created_by, created_at, status, progress
                    ) VALUES (?, ?, ?, ?, ?, '[]', '{{}}', ?, ?, ?, ?, ?, 'queued', 0)''',
                    (
                        design_id, common['design_name'], common['template_name'], common['nr_mode'], common['scope'],
                        common['filters_json'], output_file, str(destination), user.username,
                        datetime.now(timezone.utc).isoformat(),
                    ),
                )
                job_id = int(cursor.lastrowid)
            update_design_ppt_job(task_repository, job_id, vendor_comparison=common['vendor_comparison'])
            previous_output = ''
        else:
            job_id = int(reuse_job_id)
            previous_output = str(design_ppt_job(task_repository, job_id)['output_path'] or '')
            update_design_ppt_job(
                task_repository, job_id, **common, created_at=datetime.now(timezone.utc).isoformat(),
                status='queued', progress=0, slide_count=0, chart_count=0,
                last_error='', started_at=None, finished_at=None,
            )
        if auto_comments is not None:
            update_design_ppt_job(task_repository, job_id,
                                     auto_comments=auto_comments if auto_comments in AUTO_COMMENT_MODES else 'off',
                                     auto_comments_engine=str(auto_comments_engine or ''))
        if legend_visibility is not None:
            # A relaunched job keeps the legend state it was created with.
            update_design_ppt_job(
                task_repository, job_id,
                legend_visibility_json=json.dumps(normalized_legend_visibility(legend_visibility), ensure_ascii=False),
            )
        run_token = uuid4().hex
        with lock:
            design_ppt_runs[(str(Path(task_repository.db_path).resolve()), job_id)] = run_token
            snapshot_token = next((token for token, value in snapshots.items() if value is snapshot), '')
            if snapshot_token:
                design_ppt_data_tokens[(str(Path(task_repository.db_path).resolve()), job_id, user.username)] = snapshot_token
        core.submit_background_task(
            prepare_and_render_design_ppt_job,
            job_id, run_token, task_repository, user, design_id, raw_definition,
            explicit_selections, snapshot, export_time, previous_output, theme,
        )
        task_repository.add_log(user.username, 'export_design_ppt', json.dumps({
            'design_id': design_id, 'job_id': job_id, 'output': str(destination),
            'auto_comments': str(design_ppt_job(task_repository, job_id)['auto_comments'] or 'off'),
            'theme': normalize_theme(theme) if theme else load_user_theme(task_repository, user.username, DESIGN_THEME_MODULE),
        }))
        return job_id

    def stop_reporting_design_job(task_repository, job_id) -> None:
        """Stop the PPT Designer job a Reporting Job run waits for, when the run is stopped."""
        row = design_ppt_job(task_repository, job_id)
        if row is not None and str(row['status']) in {'queued', 'processing'}:
            update_design_ppt_job(task_repository, job_id, status='stopped', finished_at=datetime.now(timezone.utc).isoformat())

    def design_parameters(raw_definition, task_repository) -> list[dict[str, Any]]:
        """The parameters of the slides of a Design (those of its template, or its own when detached)."""
        try:
            _template, resolved, _report = design_slides_state(DesignDefinition.model_validate(raw_definition), task_repository)
        except (HTTPException, ValueError):
            return []
        return [dict(item) for item in resolved.get('parameters') or []]

    # Reporting Jobs and Bundle Report Templates generate Design PPTs through the PPT jobs.
    core.design_reporting = SimpleNamespace(
        definition_model=DesignDefinition,
        list_designs=lambda task_repository: read_designs(task_repository),
        queue_export=lambda design_id, user, definition=None, theme=None, auto_comments=None, auto_comments_engine=None: queue_design_ppt_export(
            design_id, user, export_definition=DesignDefinition.model_validate(definition) if definition is not None else None,
            theme=theme, auto_comments=auto_comments, auto_comments_engine=auto_comments_engine,
        ),
        job=design_ppt_job,
        serialize_job=lambda row: serialize_design_ppt_job(row),
        stop_job=stop_reporting_design_job,
        parameters=lambda task_repository, raw_definition: design_parameters(raw_definition, task_repository),
    )

    def prepare_and_render_design_ppt_job(
        job_id, run_token, task_repository, user, design_id, raw_definition,
        explicit_selections, snapshot, export_time, previous_output, theme=None,
    ):
        """Resolve the export universe, labels and file name, then render the PPT."""
        run_key = (str(Path(task_repository.db_path).resolve()), job_id)
        with lock:
            if design_ppt_runs.get(run_key) != run_token:
                return
        try:
            update_design_ppt_job(
                task_repository, job_id, status='processing', progress=1,
                started_at=datetime.now(timezone.utc).isoformat(),
            )
            if previous_output:
                remove_design_ppt_job_folder(task_repository, previous_output)
            # The CDRs of its NR Mode and Campaign filter, chosen now.
            raw_definition = runtime_design_definition(dict(raw_definition), task_repository)
            if explicit_selections:
                raw_filters = dict(raw_definition.get('filters') or {})
                for field, selected in explicit_selections.items():
                    for existing in list(raw_filters):
                        if identity(vendor_filter_column(existing)) == identity(vendor_filter_column(field)):
                            del raw_filters[existing]
                    if selected:
                        raw_filters[field] = selected
                raw_definition['filters'] = raw_filters
            snapshot_definition = DesignDefinition.model_validate(raw_definition)
            # A reused snapshot carries resolved calendar dates; the job Filters
            # column keeps the requested (possibly symbolic) range either way.
            requested_dates = {key: raw_definition.get(key) for key in ('date_from', 'date_to')}
            if snapshot is not None:
                try:
                    snapshot = validate_design_export_snapshot(snapshot, task_repository)
                    raw_definition = snapshot.definition.model_dump(mode='json')
                except (ValueError, HTTPException):
                    snapshot = None
            design_name = str(raw_definition.get('name') or design_id)
            preview_fingerprint = design_preview_fingerprint(raw_definition, task_repository)
            selected_by_kind = selected_sources(snapshot_definition, task_repository)
            catalogue = design_geography_catalogue(snapshot_definition, selected_by_kind, task_repository)
            # Available values come only from the selected CDRs, never from
            # the Design's saved filters.
            profile_options = profile_filter_options(
                unfiltered_definition(snapshot_definition), [], selected_by_kind,
                ['Operator', 'Region', 'City'], task_repository,
            )
            available = {
                'Operator': profile_options['Operator'], 'Vendor': catalogue['vendors'],
                'Region': catalogue['regions'], 'City': catalogue['cities'],
            }
            requested = {field: explicit_selections.get(field, []) for field in ('Operator', 'Vendor', 'Region', 'City')}
            filters = raw_definition.get('filters') or {}
            selections = {
                field: design_ppt_selection_values(filters, field, requested[field], values)
                for field, values in available.items()
            }
            selection_labels = {
                field: design_ppt_selection_label(
                    field, selections[field], values,
                    # Dialog selections are chosen from the catalogue; only saved
                    # Design filters come from the filter options.
                    (profile_options[field],) if field in {'Region', 'City'} and field not in explicit_selections else (),
                )
                for field, values in available.items()
            }
            campaigns = design_campaign_values(snapshot_definition, selected_by_kind, task_repository)[0]
            cover_campaigns = design_campaigns_label(campaigns)
            filters_json = json.dumps(design_filter_lines(
                {**raw_definition, **requested_dates}, task_repository,
                {**selection_labels, 'Campaign': cover_campaigns},
            ), ensure_ascii=False)
            selected_regions, selected_cities = selections['Region'], selections['City']
            region_label, city_label = selection_labels['Region'], selection_labels['City']
            cover_regions = design_ppt_cover_label('Region', region_label, len(selected_regions))
            cover_cities = design_ppt_cover_label('City', city_label, len(selected_cities))
            zone_label = region_label if region_label == 'All Regions' else ' + '.join(selected_regions)
            output_file, destination = design_ppt_unique_output(
                task_repository, export_time, DESIGN_NR_MODE_LABELS[design_nr_mode(raw_definition)], design_name,
                design_ppt_scope_label(raw_definition.get('scope') or 'single'), zone_label, job_id,
                design_ppt_campaign_suffix(campaigns),
            )
            update_design_ppt_job(
                task_repository, job_id, regions_json=json.dumps(selected_regions, ensure_ascii=False),
                selections_json=json.dumps(selections, ensure_ascii=False), filters_json=filters_json,
                output_file=output_file, output_path=str(destination),
            )
        except Exception as exc:
            with lock:
                active = design_ppt_runs.get(run_key) == run_token
            if active:
                detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
                update_design_ppt_job(
                    task_repository, job_id, status='failed', last_error=str(detail),
                    finished_at=datetime.now(timezone.utc).isoformat(),
                )
            return
        render_design_ppt_job(
            job_id, run_token, task_repository, snapshot, snapshot_definition, user, destination,
            design_id, preview_fingerprint, cover_regions, cover_cities, cover_campaigns, theme,
        )

    @app.post('/api/ppt-designer/{design_id}/export-ppt')
    def export_design_ppt(
        design_id: str, request: DesignPptExportRequest | None = None,
        user=Depends(design_user),
    ):
        job_id = queue_design_ppt_export(
            design_id, user,
            export_definition=request.definition if request else None,
            preparation_token=request.preparation_token if request else None,
            selected_operators=request.selected_operators if request else None,
            selected_vendors=request.selected_vendors if request else None,
            selected_regions=request.selected_regions if request else None,
            selected_cities=request.selected_cities if request else None,
            legend_visibility=request.legend_visibility if request else None,
            auto_comments=request.auto_comments if request else None,
            auto_comments_engine=request.auto_comments_engine if request else None,
        )
        return JSONResponse({'job_id': job_id, 'status': 'queued'}, status_code=202)

    @app.get('/api/ppt-designer/ppt-jobs')
    def design_ppt_jobs(user=Depends(design_user)):
        task_repository = bound_repository()
        ensure_design_ppt_jobs(task_repository)
        with task_repository.connection() as connection:
            rows = connection.execute(
                # Newest first: a relaunched job takes the date of its relaunch and moves to the top.
                f'SELECT * FROM {DESIGN_PPT_JOBS_TABLE} ORDER BY created_at DESC, id DESC LIMIT 100'
            ).fetchall()
        return {'jobs': [serialize_design_ppt_job(row) for row in rows]}

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/download')
    def download_design_ppt(job_id: int, user=Depends(design_user)):
        task_repository = bound_repository()
        row = design_ppt_job(task_repository, job_id)
        path = Path(str(row['output_path'])) if row else None
        if row is None or str(row['status']) != 'ready' or not path.is_file():
            raise HTTPException(404, 'PowerPoint is not available.')
        task_repository.try_add_log(user.username, 'download_design_ppt', json.dumps({'job_id': job_id, 'file': path.name}))
        return FileResponse(
            path, filename=path.name,
            media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation',
        )

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/charts')
    def open_design_ppt_charts(job_id: int, user=Depends(design_user)):
        row = design_ppt_job(bound_repository(), job_id)
        if row is None:
            raise HTTPException(404, 'PPT charts are not available.')
        charts_dir = Path(str(row['output_path'])).parent / 'design-charts'
        try:
            manifest = json.loads((charts_dir / 'manifest.json').read_text(encoding='utf-8'))
        except (AttributeError, OSError, json.JSONDecodeError):
            raise HTTPException(404, 'PPT charts are not available.')
        cards = ''.join(
            f'<article><h2>{html.escape(str(item.get("title") or "Chart"))}</h2>'
            f'<img src="/api/ppt-designer/ppt-jobs/{job_id}/charts/{html.escape(str(item.get("file") or ""))}" alt=""></article>'
            for item in manifest.get('charts', []) if isinstance(item, dict)
        )
        return HTMLResponse(
            '<!doctype html><html><head><title>PPT charts</title><style>'
            'body{font-family:Arial;margin:24px;background:#f5f2f8}article{margin:0 0 24px;padding:16px;background:white;border-radius:12px}'
            'img{display:block;width:100%;height:auto}h2{font-size:18px}</style></head><body>' + cards + '</body></html>'
        )

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/charts.json')
    def design_ppt_charts_manifest(job_id: int, user=Depends(design_user)):
        row = design_ppt_job(bound_repository(), job_id)
        if row is None:
            raise HTTPException(404, 'PPT charts are not available.')
        charts_dir = Path(str(row['output_path'])).parent / 'design-charts'
        try:
            manifest = json.loads((charts_dir / 'manifest.json').read_text(encoding='utf-8'))
        except (AttributeError, OSError, json.JSONDecodeError):
            raise HTTPException(404, 'PPT charts are not available.')
        charts = []
        definition_available = isinstance(manifest.get('definition'), dict)
        fallback_focus_rows = []
        if any(not isinstance(item.get('focus_row'), int) for item in manifest.get('charts', []) if isinstance(item, dict)):
            try:
                definition = DesignDefinition(
                    name=str(row['design_name']), template=str(row['template_name']),
                    nr_mode=design_nr_mode_value(row['nr_mode']),
                    scope='multivendor' if str(row['scope']).casefold() == 'multivendor' else 'single',
                )
                entries = validate(definition, bound_repository())
                focus_by_index = {
                    entry_index: editor_index
                    for editor_index, (entry_index, _entry) in enumerate(
                        sorted(enumerate(entries), key=lambda value: (value[1].slide, value[0]))
                    )
                }
                unused = set(range(len(entries)))
                for item in manifest.get('charts', []):
                    matched = next((
                        entry_index for entry_index, entry in enumerate(entries)
                        if entry_index in unused
                        and entry.slide == int(item.get('slide') or 0)
                        and str(entry.chart_title or entry.slide_title) == str(item.get('title') or '')
                        and str(entry.cdr_source) == str(item.get('source') or '')
                        and str(entry.chart_type) == str(item.get('chart_type') or '')
                    ), None)
                    fallback_focus_rows.append(focus_by_index.get(matched) if matched is not None else None)
                    if matched is not None:
                        unused.discard(matched)
            except (TypeError, ValueError):
                fallback_focus_rows = []
        for chart_index, item in enumerate(manifest.get('charts', [])):
            if not isinstance(item, dict):
                continue
            chart_file = str(item.get('file') or '')
            if not re.fullmatch(r'slide-\d+-chart-\d+\.png', chart_file):
                continue
            model_file = str(item.get('model_file') or '')
            model_available = bool(
                re.fullmatch(r'slide-\d+-chart-\d+\.model\.json', model_file)
                and (charts_dir / model_file).is_file()
            )
            charts.append({
                'index': len(charts),
                'entry_index': int(item['entry_index']) if isinstance(item.get('entry_index'), int) else None,
                'slide': int(item.get('slide') or 0),
                'title': str(item.get('title') or 'Chart'),
                'source': str(item.get('source') or ''),
                'chart_type': str(item.get('chart_type') or ''),
                'focus_row': int(item['focus_row']) if isinstance(item.get('focus_row'), int) else (
                    fallback_focus_rows[chart_index] if chart_index < len(fallback_focus_rows) else None
                ),
                'image_url': f'/api/ppt-designer/ppt-jobs/{job_id}/charts/{chart_file}',
                'payload_url': f'/api/ppt-designer/ppt-jobs/{job_id}/chart-models/{model_file}' if model_available else None,
                'data_url': f'/ppt-jobs/{job_id}/data/{chart_index}' if definition_available and isinstance(item.get('entry_index'), int) else None,
            })
        chart_by_entry = {
            int(chart['entry_index']): chart for chart in charts
            if isinstance(chart.get('entry_index'), int)
        }
        slides = []
        for source_slide in manifest.get('slides', []):
            if not isinstance(source_slide, dict):
                continue
            slide = {
                key: source_slide.get(key) for key in ('number', 'title', 'subtitle', 'layout', 'structural_type')
            }
            slide['charts'] = []
            for source_chart in source_slide.get('charts', []):
                if not isinstance(source_chart, dict):
                    continue
                entry_index = source_chart.get('index')
                chart = chart_by_entry.get(entry_index)
                if chart is None:
                    continue
                slide['charts'].append({**source_chart, **chart, 'available': bool(chart.get('payload_url'))})
            slides.append(slide)
        # Older PPT jobs did not retain their viewer slide structure. They can
        # still be opened as simple slide groups when Canvas models are present.
        if not slides:
            grouped = defaultdict(list)
            for chart in charts:
                grouped[int(chart.get('slide') or 0)].append({**chart, 'available': bool(chart.get('payload_url'))})
            slides = [
                {'number': number, 'title': f'Slide {number}', 'subtitle': '', 'layout': '', 'structural_type': '', 'charts': items}
                for number, items in sorted(grouped.items())
            ]
        return {
            'job': serialize_design_ppt_job(row),
            'generate_tooltips': bool(manifest.get('generate_tooltips')),
            'charts': charts,
            'definition': manifest.get('definition') if definition_available else None,
            'cover': design_ppt_job_cover(row, manifest),
            'slides': slides,
            'viewer_available': bool(slides) and any(
                chart.get('available') for slide in slides for chart in slide.get('charts', [])
            ),
        }

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/charts/{chart_index}/filter-context')
    def design_ppt_chart_filter_context(
        job_id: int, chart_index: int, prepare: bool = False, user=Depends(design_user),
    ):
        """Load PPT chart metadata immediately and restore its data only on demand."""
        task_repository = bound_repository()
        row = design_ppt_job(task_repository, job_id)
        charts_dir = Path(str(row['output_path'])).parent / 'design-charts' if row else None
        try:
            manifest = json.loads((charts_dir / 'manifest.json').read_text(encoding='utf-8'))
            definition = DesignDefinition.model_validate(manifest['definition'])
            entry_index = int(manifest['charts'][chart_index]['entry_index'])
            entry = validate(definition, task_repository)[entry_index]
        except (AttributeError, IndexError, KeyError, OSError, TypeError, ValueError, json.JSONDecodeError):
            raise HTTPException(404, 'Chart filters are not available for this PowerPoint Job.')
        workspace = str(Path(task_repository.db_path).resolve())
        token = ''
        if prepare:
            preview = restore_preview_manifest(
                workspace, str(manifest.get('design_id') or row['design_id']),
                str(manifest.get('preview_fingerprint') or ''),
            ) or build_preview(definition, user, workspace=workspace)
            token = str(preview['token'])
        hidden = {identity('dataset_id'), identity('source_row_id')}
        datasets_by_source = {f'cdr-{kind}': [] for kind in KINDS}
        columns_by_source = {f'cdr-{kind}': [] for kind in KINDS}
        for dataset in task_repository.list_datasets():
            kind = str(dataset['dataset_kind'] or '').casefold()
            if kind in KINDS and dataset['status'] == 'ready':
                datasets_by_source[f'cdr-{kind}'].append({'value': str(dataset['id']), 'label': str(dataset['file_name'])})
                columns_by_source[f'cdr-{kind}'].extend(
                    str(column) for column in task_repository.list_dataset_row_columns(int(dataset['id']))
                )
        for source, source_columns in columns_by_source.items():
            columns_by_source[source] = sorted({
                column for column in source_columns if identity(column) not in hidden
            }, key=str.casefold)
        columns = columns_by_source.get(f'cdr-{entry.source_kind}', [])
        datasets_by_source['cdr-all'] = [row for source in KINDS for row in datasets_by_source[f'cdr-{source}']]
        columns_by_source['cdr-all'] = sorted({*rf_union_source(task_repository, [kind for kind in KINDS if design_datasets(definition, kind, task_repository)])[1]}, key=str.casefold)
        if entry.source_kind == 'all':
            columns = columns_by_source['cdr-all']
        template_available = user.role in {'user-editor', 'admin', 'super-admin'} and any(
            str(item['name']) == definition.template
            for item in task_repository.list_report_templates()
        )
        return JSONResponse({
            'token': token, 'chart_index': entry_index, 'cdr_source': entry.cdr_source,
            'chart_type': entry.chart_type, 'chart_title': entry.chart_title, 'kpi': entry.kpi,
            'dataset_ids': [str(row['id']) for row in design_datasets(definition, entry.source_kind, task_repository)],
            'filters': entry.filters, 'rows_split': entry.rows_split,
            'columns_split': entry.columns_split, 'legend': entry.legend,
            'legend_position': entry.legend_position,
            'legend_format': entry.legend_format,
            'axis_x_range': entry.axis_x_range, 'axis_y_range': entry.axis_y_range,
            'label_position': entry.label_position,
            'label_format': entry.label_format,
            'exclude_null_empty': entry.exclude_null_empty, 'exclude_zero': entry.exclude_zero, 'dynamic_rows_field': entry_dynamic_fields(entry)[0], 'dynamic_columns_field': entry_dynamic_fields(entry)[1],
            'template_available': template_available, 'datasets_by_source': datasets_by_source,
            'columns_by_source': columns_by_source, 'columns': columns,
        })

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/chart-models/{model_file}')
    def design_ppt_chart_model(job_id: int, model_file: str, user=Depends(design_user)):
        if not re.fullmatch(r'slide-\d+-chart-\d+\.model\.json', model_file):
            raise HTTPException(404, 'Chart model not found.')
        row = design_ppt_job(bound_repository(), job_id)
        path = Path(str(row['output_path'])).parent / 'design-charts' / model_file if row else None
        try:
            payload = json.loads(path.read_text(encoding='utf-8'))
        except (AttributeError, OSError, json.JSONDecodeError):
            raise HTTPException(404, 'Chart model not found.')
        return JSONResponse(payload)

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/data/{chart_index}')
    def design_ppt_chart_data(
        job_id: int, chart_index: int, page: int = 0, download: bool = False,
        column_filters: str = '', include_filter_values: bool = False, filter_column: str = '',
        user=Depends(design_user),
    ):
        task_repository = bound_repository()
        row = design_ppt_job(task_repository, job_id)
        charts_dir = Path(str(row['output_path'])).parent / 'design-charts' if row else None
        try:
            manifest = json.loads((charts_dir / 'manifest.json').read_text(encoding='utf-8'))
            chart = manifest['charts'][chart_index]
            definition = DesignDefinition.model_validate(manifest['definition'])
            entry_index = int(chart['entry_index'])
        except (AttributeError, IndexError, KeyError, OSError, TypeError, ValueError, json.JSONDecodeError):
            raise HTTPException(404, 'The filtered dataset is not available for this PowerPoint Job.')
        workspace = str(Path(task_repository.db_path).resolve())
        token_key = (workspace, job_id, user.username)
        with lock:
            token = design_ppt_data_tokens.get(token_key)
            cached_snapshot = snapshots.get(token) if token else None
        if cached_snapshot is None:
            preview = restore_preview_manifest(
                workspace,
                str(manifest.get('design_id') or row['design_id']),
                str(manifest.get('preview_fingerprint') or ''),
            )
            if preview is None:
                preview = build_preview(definition, user, workspace=workspace)
            token = str(preview['token'])
            with lock:
                design_ppt_data_tokens[token_key] = token
        snapshot, entry, _ = snapshot_chart(
            token, entry_index, user, include_frame=False, expected_workspace=workspace,
        )
        filter_column = vendor_filter_column(filter_column) if filter_column else ''
        selected_column_filters = {vendor_filter_column(key): values for key, values in parse_chart_column_filters(column_filters).items()}
        if not download:
            projected_page = reporting_chart_data_page(
                snapshot, entry, page, selected_column_filters, include_filter_values, filter_column,
            )
            if projected_page is not None:
                visible, total, chart_total, filter_values = projected_page
                return {
                    'columns': list(visible.columns),
                    'rows': visible.fillna('').astype(str).values.tolist(),
                    'total': total, 'chart_total': chart_total,
                    'filter_values': filter_values, 'page': max(page, 0),
                    'column_classes': chart_dataset_column_classes(snapshot, entry, visible.columns),
                    'column_metadata': chart_dataset_column_metadata(snapshot, entry, visible.columns),
                    'page_size': 100, 'unfiltered_total': chart_total,
                }
        _, _, frame = snapshot_chart(token, entry_index, user, expected_workspace=workspace, complete=True)
        frame = unique_chart_dataset_rows(snapshot, entry, frame)
        chart_total = len(frame)
        values_frame = apply_chart_column_filters(
            frame, {column: values for column, values in selected_column_filters.items() if identity(column) != identity(filter_column)},
        )
        available_filter_values = chart_dataset_filter_values(values_frame) if filter_column else {}
        filter_values = next(
            (values for column, values in available_filter_values.items() if identity(column) == identity(filter_column)), []
        ) if filter_column else (chart_dataset_filter_values(frame) if include_filter_values else {})
        frame = apply_chart_column_filters(frame, selected_column_filters)
        if download:
            task_repository.try_add_log(user.username, 'export_design_chart_data_csv', json.dumps({
                'job_id': job_id, 'chart_index': chart_index, 'rows': len(frame),
            }))
            return Response(
                frame.to_csv(index=False), media_type='text/csv',
                headers={'Content-Disposition': 'attachment; filename="design-chart-data.csv"'},
            )
        page = max(page, 0)
        visible = frame.iloc[page * 100:(page + 1) * 100].fillna('').astype(str)
        return {
            'columns': list(visible.columns), 'rows': visible.values.tolist(),
            'total': len(frame), 'chart_total': chart_total,
            'filter_values': filter_values, 'page': page,
            'column_classes': chart_dataset_column_classes(snapshot, entry, visible.columns),
            'column_metadata': chart_dataset_column_metadata(snapshot, entry, visible.columns),
            'page_size': 100, 'unfiltered_total': chart_total,
        }

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/charts/{chart_file}')
    def design_ppt_chart(job_id: int, chart_file: str, user=Depends(design_user)):
        if not re.fullmatch(r'slide-\d+-chart-\d+\.png', chart_file):
            raise HTTPException(404, 'Chart not found.')
        row = design_ppt_job(bound_repository(), job_id)
        path = Path(str(row['output_path'])).parent / 'design-charts' / chart_file if row else None
        if path is None or not path.is_file():
            raise HTTPException(404, 'Chart not found.')
        return FileResponse(path, media_type='image/png')

    @app.get('/api/ppt-designer/ppt-jobs/{job_id}/charts.zip')
    def download_design_ppt_charts(job_id: int, user=Depends(design_user)):
        task_repository = bound_repository()
        row = design_ppt_job(task_repository, job_id)
        charts_dir = Path(str(row['output_path'])).parent / 'design-charts' if row else None
        if charts_dir is None or not (charts_dir / 'manifest.json').is_file():
            raise HTTPException(404, 'PPT charts are not available.')
        task_repository.try_add_log(user.username, 'download_design_ppt_charts', json.dumps({'job_id': job_id}))
        temporary = tempfile.NamedTemporaryFile(prefix='design-ppt-charts-', suffix='.zip', delete=False)
        temporary.close()
        archive = Path(temporary.name)
        with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as bundle:
            for path in charts_dir.glob('*.png'):
                bundle.write(path, path.name)
        return FileResponse(
            archive, filename=f'{Path(str(row["output_file"])).stem}_charts.zip',
            media_type='application/zip', background=BackgroundTask(archive.unlink, missing_ok=True),
        )

    @app.post('/api/ppt-designer/ppt-jobs/{job_id}/stop')
    def stop_design_ppt(job_id: int, user=Depends(design_user)):
        task_repository = bound_repository()
        row = design_ppt_job(task_repository, job_id)
        if row is None or str(row['status']) not in {'queued', 'processing'}:
            raise HTTPException(409, 'Only queued or processing PPT exports can be stopped.')
        update_design_ppt_job(
            task_repository, job_id, status='stopped',
            finished_at=datetime.now(timezone.utc).isoformat(),
        )
        task_repository.try_add_log(user.username, 'stop_design_ppt_job', json.dumps({'job_id': job_id}))
        return {'stopped': job_id}

    @app.post('/api/ppt-designer/ppt-jobs/{job_id}/retry')
    def retry_design_ppt(job_id: int, user=Depends(design_user)):
        task_repository = bound_repository()
        row = design_ppt_job(task_repository, job_id)
        if row is None or str(row['status']) not in {'ready', 'failed', 'stopped'}:
            raise HTTPException(409, 'This PPT export cannot be relaunched.')
        try:
            saved_selections = json.loads(str(row['selections_json'] or '{}'))
        except (TypeError, json.JSONDecodeError):
            saved_selections = {}
        if not isinstance(saved_selections, dict):
            saved_selections = {}
        try:
            manifest_path = Path(str(row['output_path'])).parent / 'design-charts' / 'manifest.json'
            manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
            retry_definition = DesignDefinition.model_validate(manifest['definition'])
        except (KeyError, OSError, TypeError, ValueError, json.JSONDecodeError):
            # Failed or stopped Jobs may not have reached manifest creation.
            # Preserve at least their recorded Scope while rebuilding the
            # current saved Design universe.
            raw_definition = read_designs(task_repository).get(str(row['design_id']))
            if not isinstance(raw_definition, dict):
                raise HTTPException(404, 'Design not found.')
            raw_definition = dict(raw_definition)
            raw_definition['scope'] = (
                'multivendor' if str(row['scope'] or '').casefold() == 'multivendor' else 'single'
            )
            if saved_selections:
                raw_definition['filters'] = {
                    **(raw_definition.get('filters') or {}),
                    **{field: values for field, values in saved_selections.items() if isinstance(values, list)},
                }
            retry_definition = DesignDefinition.model_validate(
                runtime_design_definition(raw_definition, task_repository)
            )
        try:
            selected_regions = json.loads(str(row['regions_json'] or '[]'))
        except (TypeError, json.JSONDecodeError):
            selected_regions = []
        retry_definition.vendor_comparison = str(row['vendor_comparison'] or 'operator_vendor')
        queue_design_ppt_export(
            str(row['design_id']), user, reuse_job_id=job_id,
            export_definition=retry_definition,
            selected_operators=saved_selections.get('Operator') if isinstance(saved_selections.get('Operator'), list) else None,
            selected_vendors=saved_selections.get('Vendor', saved_selections.get('Vendor_Only')),
            selected_regions=saved_selections.get('Region') if isinstance(saved_selections.get('Region'), list) else selected_regions,
            selected_cities=saved_selections.get('City') if isinstance(saved_selections.get('City'), list) else None,
        )
        return JSONResponse({'job_id': job_id, 'status': 'queued'}, status_code=202)

    @app.post('/api/ppt-designer/ppt-jobs/delete-all')
    def delete_all_design_ppts(user=Depends(design_editor_user)):
        task_repository = bound_repository()
        ensure_design_ppt_jobs(task_repository)
        with task_repository.connection() as connection:
            rows = connection.execute(
                f'SELECT id, output_path FROM {DESIGN_PPT_JOBS_TABLE}'
            ).fetchall()
            connection.execute(f'DELETE FROM {DESIGN_PPT_JOBS_TABLE}')
        database_key = str(Path(task_repository.db_path).resolve())
        with lock:
            for row in rows:
                design_ppt_runs.pop((database_key, int(row['id'])), None)
        for row in rows:
            remove_design_ppt_job_folder(task_repository, row['output_path'])
        task_repository.add_log(user.username, 'delete_all_design_ppts', json.dumps({'count': len(rows)}))
        return {'deleted': len(rows)}

    @app.post('/api/ppt-designer/ppt-jobs/{job_id}/delete')
    def delete_design_ppt(job_id: int, user=Depends(design_editor_user)):
        task_repository = bound_repository()
        row = design_ppt_job(task_repository, job_id)
        if row is None:
            raise HTTPException(404, 'PPT export job not found.')
        if str(row['status']) in {'queued', 'processing'}:
            raise HTTPException(409, 'A running PPT export cannot be deleted.')
        remove_design_ppt_job_folder(task_repository, row['output_path'])
        with task_repository.connection() as connection:
            connection.execute(f'DELETE FROM {DESIGN_PPT_JOBS_TABLE} WHERE id = ?', (job_id,))
        task_repository.try_add_log(user.username, 'delete_design_ppt', json.dumps({
            'job_id': job_id, 'file': str(row['output_file'] or ''),
        }))
        return {'deleted': job_id}

    def ensure_unique_design_name(designs, design_id, name) -> None:
        """PPT names are unique in a workspace, whatever the NR Mode each one reads."""
        if any(
            key != design_id and isinstance(item, dict)
            and str(item.get('name') or '').strip().casefold() == str(name).strip().casefold()
            for key, item in designs.items()
        ):
            raise HTTPException(409, 'A Design with this name already exists.')

    @app.put('/api/ppt-designer/template-preview')
    def open_template_preview(definition: DesignDefinition, user=Depends(design_editor_user)):
        """Open a PPT Template in the viewer and the Designer (PPT Templates → Edit) as a Design of the user kept apart
        from the Designs: the selection sent (NR Mode, Campaign, filters) shows its slides, and what the Designer changes
        goes to the template."""
        workspace_key()
        task_repository = bound_repository()
        if template_row_of(definition.template, task_repository) is None:
            raise HTTPException(404, 'PPT Template not found.')
        design_id = template_preview_id(user.username)
        stored = {**definition.model_dump(mode='json'), 'name': definition.template, 'slide_overrides': {},
                  'own_slides': None, 'slide_comments': {}, 'slide_comment_meta': {}}
        stored.pop('datasets', None)
        with lock:
            store_design_entry(task_repository, design_id, stored)
        task_repository.add_log(user.username, 'open_template_preview', json.dumps({
            'id': design_id, 'template': definition.template, 'nr_mode': definition.nr_mode,
        }))
        return {'id': design_id, 'definition': {**stored, 'template_preview': True, 'created_by': user.username}}

    @app.put('/api/ppt-designer/{design_id}')
    def save_design(design_id: str, definition: DesignDefinition, user=Depends(design_editor_user)):
        workspace_key()
        preview = is_template_preview(design_id)
        if preview and design_id != template_preview_id(user.username):
            raise HTTPException(403, 'This PPT Template is open by another user.')
        if definition.template == design_slides.BLANK_PPT_TEMPLATE and not definition.own_slides:
            # A PPT Design created blank: its cover, as slides of its own from the start (it follows no template).
            definition = definition.model_copy(update={
                'own_slides': design_slides.slides_from_entries(core.cover_catalog_entries(definition.name.strip()))})
        with lock:
            task_repository = bound_repository()
            validate(definition, task_repository)
            designs = read_template_previews(task_repository) if preview else read_designs(task_repository)
            if not preview:
                ensure_unique_design_name(designs, design_id, definition.name)
            created = design_id not in designs
            definition.name = definition.name.strip()
            saved_definition = migrate_auto_cdrs(normalize_design_filters(definition.model_dump(mode='json')), task_repository)
            # The CDRs are chosen from the NR Mode and Campaign each time: they are not saved.
            saved_definition.pop('datasets', None)
            saved_definition['slide_comment_meta'] = comment_meta(
                saved_definition.get('slide_comment_meta'), saved_definition.get('slide_comments') or {},
            )
            saved_definition = follow_template_identities({design_id: saved_definition}, template_uids(task_repository))[design_id]
            if not created:
                # Its slide set changes only from the Designer (save_design_slides): the filters saved from a viewer
                # opened before never take it back.
                for key, default in (('slide_overrides', {}), ('own_slides', None)):
                    saved_definition[key] = copy.deepcopy(designs[design_id].get(key, default))
            store_design_entry(task_repository, design_id, saved_definition, designs)
            if created and not preview:
                record_creator(task_repository, PPT_DESIGNER, design_id, user.username)
            creator = user.username if preview else creator_of(task_repository, PPT_DESIGNER, design_id)
            task_repository.add_log(user.username, 'save_template_preview' if preview else 'save_design', json.dumps({
                'id': design_id, 'name': definition.name, 'nr_mode': definition.nr_mode, 'template': definition.template,
            }))
        # As the list gives it, with who created it: the page compares it with the Design it has open.
        return {'id': design_id, 'definition': {**saved_definition, 'created_by': creator}}

    @app.patch('/api/ppt-designer/{design_id}/name')
    def rename_design(design_id: str, payload: DesignName, user=Depends(design_editor_user)):
        name = payload.name.strip()
        if not name:
            raise HTTPException(400, 'Enter a PPT name.')
        with lock:
            task_repository = bound_repository()
            designs = read_designs(task_repository)
            if design_id not in designs:
                raise HTTPException(404, 'Design not found.')
            ensure_unique_design_name(designs, design_id, name)
            designs[design_id]['name'] = name
            task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
            task_repository.add_log(user.username, 'rename_design', json.dumps({'id': design_id, 'name': name}))
        return {'id': design_id, 'name': name}

    @app.patch('/api/ppt-designer/{design_id}/template')
    def change_design_template(
        design_id: str, payload: DesignTemplateSelection, user=Depends(design_editor_user),
    ):
        """Change the NR Mode of the CDRs and the template while retaining the saved filters.

        The saved CDR universe is kept when the NR Mode is unchanged. A new NR
        Mode selects the latest CDRs of that mode instead: the campaigns of the
        other NR Mode do not apply. Any template serves any NR Mode.
        """
        workspace_key()
        with lock:
            task_repository = bound_repository()
            designs = read_designs(task_repository)
            current = designs.get(design_id)
            if current is None:
                raise HTTPException(404, 'Design not found.')
            updated = dict(current)
            previous_nr_mode, previous_template = design_nr_mode(current), str(current.get('template') or '')
            nr_mode_changed = previous_nr_mode != payload.nr_mode
            updated['nr_mode'] = payload.nr_mode
            updated['template'] = payload.template
            dropped_overrides = {}
            if payload.template != previous_template:
                # What it changed of its template (by the ids of that template), or its own slide set once detached,
                # does not apply to another template: it follows the chosen one as it is.
                dropped_overrides = design_slides.overrides_summary(current.get('slide_overrides'))
                updated['slide_overrides'], updated['own_slides'] = {}, None
            if nr_mode_changed:
                # The campaigns of the other NR Mode do not apply: the newest ones of the new NR Mode do.
                updated.pop('datasets', None)
                follow_nr_mode_campaigns(updated, previous_nr_mode, payload.nr_mode)
                updated['auto_cdrs'] = True
            definition = DesignDefinition.model_validate(updated)
            validate(definition, task_repository)
            saved_definition = migrate_auto_cdrs(normalize_design_filters(definition.model_dump(mode='json')), task_repository)
            saved_definition.pop('datasets', None)
            saved_definition = follow_template_identities({design_id: saved_definition}, template_uids(task_repository))[design_id]
            designs[design_id] = saved_definition
            task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
            task_repository.add_log(user.username, 'change_design_template', json.dumps({
                'id': design_id, 'name': definition.name,
                'nr_mode': payload.nr_mode, 'previous_nr_mode': previous_nr_mode,
                'template': payload.template, 'previous_template': previous_template,
                **({'dropped_slide_overrides': dropped_overrides, 'detached_before': bool(current.get('own_slides'))}
                   if payload.template != previous_template and (any(dropped_overrides.values()) or current.get('own_slides')) else {}),
            }))
        return {'id': design_id, 'definition': saved_definition, 'invalidated': True}

    @app.delete('/api/ppt-designer/{design_id}')
    def delete_design(design_id: str, user=Depends(design_editor_user)):
        with lock:
            task_repository = bound_repository()
            designs = read_designs(task_repository)
            if design_id not in designs:
                raise HTTPException(404, 'Design not found.')
            if not can_delete(user, creator_of(task_repository, PPT_DESIGNER, design_id)):
                raise HTTPException(403, deletion_refusal('Design'))
            del designs[design_id]
            task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
            forget_creator(task_repository, PPT_DESIGNER, design_id)
            task_repository.add_log(user.username, 'delete_design', json.dumps({'id': design_id}))
        return {'deleted': True}

    # ------------------------------------------------------------------
    # Designer (ticket #26): the slide set of a Design and of its template
    # ------------------------------------------------------------------

    def template_row_of(template_name, task_repository):
        return next((row for row in task_repository.list_report_templates() if str(row['name']) == str(template_name)), None)

    def can_edit_design_slides(user, design_id, task_repository) -> bool:
        """Admins and super-admins edit the slides of any Design, a user-editor only those they created, a user-viewer none."""
        return user.role in {'user-editor', 'admin', 'super-admin'} and can_delete(
            user, creator_of(task_repository, PPT_DESIGNER, design_id))

    def can_edit_template_slides(user, template_row) -> bool:
        """The same rule for a PPT Template, by its creator."""
        return template_row is not None and user.role in {'user-editor', 'admin', 'super-admin'} and can_delete(
            user, template_row['created_by'])

    def designer_refusal(label: str) -> str:
        return f'Only admins and super-admins, or the user who created it, can change the slides of this {label}.'

    def checked_slide_set(slide_set, label='template') -> dict:
        """A slide set sent by the Designer, checked as its template CSV would be (the same messages), with its
        parameters completed."""
        if not isinstance(slide_set, dict) or not isinstance(slide_set.get('slides'), list):
            raise HTTPException(400, f'The slides of the {label} are not valid.')
        slide_set = design_slides.template_slides_of(slide_set)
        try:
            slide_set['parameters'] = design_slides.validate_parameters(slide_set.get('parameters') or [])
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        if not slide_set['slides']:
            raise HTTPException(400, f'The {label} needs at least one slide.')
        if empty := design_slides.slides_without_elements(slide_set):
            raise HTTPException(400, f'Add an element to the slide "{empty[0]}" or delete it.')
        slide_ids = [slide.get('id') for slide in slide_set['slides']]
        element_ids = [element.get('id') for slide in slide_set['slides'] for element in slide.get('elements') or []]
        if not all(slide_ids) or len(set(slide_ids)) != len(slide_ids) or not all(element_ids) or len(set(element_ids)) != len(element_ids):
            raise HTTPException(400, f'Each slide and element of the {label} needs an id of its own.')
        try:
            core.parse_catalog_csv(design_slides.slides_csv(slide_set))
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        return slide_set

    def template_with_ids(template_name, task_repository, username='system'):
        """The slide set of a template, with ids of its own stored first when it has none yet (the positional ones it was
        read with, so nothing overridden changes id)."""
        row = template_row_of(template_name, task_repository)
        if row is None:
            return None, None
        entries = core.load_template_catalogue(bytes(row['content']), task_repository=task_repository)
        if entries and not design_slides.has_slide_ids(entries):
            content = core.catalogue_csv(design_slides.with_slide_ids(entries))
            task_repository.set_report_template_content(str(row['name']), content)
            task_repository.add_log(username, 'report_template_slide_ids_assigned', json.dumps({
                'template': str(row['name']), 'rows': len(entries),
            }))
            row = template_row_of(template_name, task_repository)
            entries = core.load_template_catalogue(bytes(row['content']), task_repository=task_repository)
        return design_slides.slides_from_entries(entries), row

    def slide_set_payload(design_id, definition, task_repository, user, *, template=None, template_row=None,
                       report=None, resolved=None):
        """What the Designer reads: the slide set shown, its template's slide set, overrides, report, rights."""
        if resolved is None:
            template, resolved, report = design_slides_state(definition, task_repository, template)
        template_row = template_row if template_row is not None else (None if definition.own_slides else template_row_of(definition.template, task_repository))
        designs = read_designs(task_repository)
        using = [
            {'id': key, 'name': str(item.get('name') or ''), 'nr_mode': design_nr_mode(item)}
            for key, item in designs.items()
            if isinstance(item, dict) and not item.get('own_slides') and str(item.get('template') or '') == definition.template
        ] if not definition.own_slides else []
        return {
            'design_id': design_id,
            'name': definition.name,
            # A PPT Template open as a Design (PPT Templates → Edit): the Designer changes only its template.
            'template_preview': is_template_preview(design_id),
            'detached': bool(definition.own_slides),
            'template': {
                'name': definition.template,
                'slide_set': template,
                'updated_at': str(template_row['updated_at'] or '') if template_row is not None else '',
                'created_by': str(template_row['created_by'] or '') if template_row is not None else '',
                'designs': sorted(using, key=lambda item: item['name'].casefold()),
            },
            'overrides': design_slides.normalized_overrides(definition.slide_overrides),
            'summary': design_slides.overrides_summary(definition.slide_overrides),
            'slide_set': resolved,
            'report': report or {'dropped': [], 'template_changed': []},
            'parameter_values': definition.parameter_values,
            'rights': {
                # A PPT Template open as a Design changes only its template.
                'design': not is_template_preview(design_id) and can_edit_design_slides(user, design_id, task_repository),
                'template': not definition.own_slides and can_edit_template_slides(user, template_row),
            },
        }

    def stored_design(design_id, task_repository):
        designs = read_designs_and_previews(task_repository)
        if design_id not in designs:
            raise HTTPException(404, 'Design not found.')
        return DesignDefinition.model_validate(designs[design_id])

    def write_design_slides(design_id, task_repository, **changes) -> dict:
        """Store slide set fields of one Design (re-read under the lock: its filters and comments saved meanwhile stay)."""
        with lock:
            designs = read_designs_and_previews(task_repository)
            if design_id not in designs:
                raise HTTPException(404, 'Design not found.')
            designs[design_id] = {**designs[design_id], **copy.deepcopy(changes)}
            store_design_entry(task_repository, design_id, designs[design_id])
            return designs[design_id]

    layouts_read: dict[tuple[str, float], list[dict[str, Any]]] = {}

    def slide_layouts() -> list[dict[str, Any]]:
        """The layouts of the slide master the Designer offers: the places of each one (fractions of the slide), read once
        for each version of the master."""
        master = Path(core.settings.powerpoint_template)
        key = (str(master), master.stat().st_mtime if master.exists() else 0.0)
        if key in layouts_read:
            return layouts_read[key]
        deck = Presentation(core.settings.powerpoint_template)
        layouts = []
        for name in core.catalogue_layout_names():
            try:
                places = normalized_frame_positions(_layout_chart_frames(_named_slide_layout(deck, name)))
            except (KeyError, ValueError):
                continue
            grid = GRID_LAYOUT_PATTERN.fullmatch(canonical_layout_name(name))
            layouts.append({
                'name': name, 'places': places,
                'grid': [int(grid.group(1)), int(grid.group(2))] if grid else None,
                'comments': (grid.group(3) or '') if grid else ('right' if name.endswith('right') else 'down' if 'comments' in name else ''),
                'dynamic': name in DYNAMIC_LAYOUTS,
                'structural': STRUCTURAL_LAYOUT_TYPES.get(name.casefold(), ''),
            })
        layouts_read.clear()
        layouts_read[key] = layouts
        return layouts

    @app.get('/api/ppt-designer/slides/layouts')
    def design_slides_layouts(user=Depends(design_user)):
        """The layouts of the slide master the Designer offers: the places of each one (fractions of the slide)."""
        return {'layouts': slide_layouts()}

    GALLERY_SLIDES = 6

    @app.get('/api/ppt-designer/templates-gallery')
    def templates_gallery(user=Depends(design_user)):
        """The PPT Templates a new PPT Design can follow, for the gallery of Create PPT Design: each with its first slides in
        small (the places of each one and the elements in them) and how many slides and charts it has."""
        task_repository = bound_repository()
        places = {layout['name']: layout['places'] for layout in slide_layouts()}

        def small(slide):
            geometry = (slide.get('options') or {}).get('geometry')
            return {
                'kind': slide['kind'], 'title': slide.get('title') or '',
                'places': geometry if isinstance(geometry, list) and geometry else places.get(slide.get('layout') or '', []),
                'elements': [{'kind': element.get('kind') or 'chart', 'chart_type': element.get('chart_type') or '',
                              'place': element.get('place', index), 'title': element.get('chart_title') or ''}
                             for index, element in enumerate(slide.get('elements') or []) if not element.get('hidden')],
            }

        gallery = []
        for option in core.report_catalogue_options():
            slide_set, _entries = template_slide_set(option['name'], task_repository)
            if slide_set is None:
                continue
            slides = slide_set['slides']
            gallery.append({
                'name': option['name'], 'identifier': option['identifier'], 'active': bool(option['active']),
                'slides': len(slides), 'charts': sum((element.get('kind') or 'chart') in {'chart', 'kpi'}
                                                     for slide in slides for element in slide.get('elements') or []),
                'preview': [small(slide) for slide in slides[:GALLERY_SLIDES]],
            })
        return {'templates': gallery}

    def designer_snapshot(token, user):
        workspace = workspace_key()
        with lock:
            snapshot = snapshots.get(token)
        if snapshot is None or snapshot.workspace != workspace or snapshot.owner not in {user.username, '*'}:
            raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
        return snapshot, Repository(Path(workspace), core.repository.global_db_path)

    def field_categories(task_repository) -> dict[str, str]:
        """The section of the CDR Field Dictionary each CDR column belongs to (its first one), by column identity."""
        from src.modules import cdr_field_dictionary
        from src.modules.cdr_field_profiles import workspace_field_profile
        categories: dict[str, str] = {}
        try:
            rows = cdr_field_dictionary.entries(workspace_field_profile(task_repository))
        except (AttributeError, KeyError, TypeError, ValueError, sqlite3.Error):
            return categories
        for row in rows:
            for column in row.get('defaults') or []:
                categories.setdefault(identity(column), str(row['section']))
        return categories

    def stored_column_types(snapshot, kind, task_repository) -> dict[str, str]:
        """The type of each column of the selected CDRs of a kind, as their columnar copies keep them: number or text."""
        types: dict[str, str] = {}
        root = cdr_store.store_root(task_repository.db_path)
        for row in design_datasets(snapshot.definition, kind, task_repository):
            manifest = cdr_store.read_manifest(cdr_store.manifest_file(root, kind, int(row['id']))) or {}
            for column, column_type in manifest.get('columns') or []:
                kind_of = 'number' if str(column_type).split(':')[0] in {'BIGINT', 'DOUBLE'} else 'text'
                # A column a CDR keeps as text and another as numbers is text.
                types[identity(column)] = 'text' if types.get(identity(column)) == 'text' else kind_of
        return types

    @app.get('/api/ppt-designer/slides/fields')
    def design_slides_fields(token: str, user=Depends(design_user)):
        """The fields of each source of the CDRs (Data, Voice, Speech and RF All) the Designer offers for an element:
        their names (``columns_by_source``) and, for the chart builder, their type (number or text) and the section of
        the CDR Field Dictionary they belong to (``fields``)."""
        snapshot, task_repository = designer_snapshot(token, user)
        try:
            _database_path, _table_name, columns = reporting_source(snapshot, 'all', task_repository)
        except (HTTPException, ValueError):
            columns = []
        _datasets, columns_by_source = source_columns(task_repository, columns)
        dimensions = core.load_repository_calculated_dimensions(task_repository)
        categories = field_categories(task_repository)
        # The Auto-calculated Fields of the workspace are listed together, also those the CDRs already hold as columns.
        categories.update({identity(dimension.name): 'Auto-calculated Fields' for dimension in dimensions})
        internal = {identity(name) for name in ('dataset_kind', 'source_file', '__rowid', '__dataset_id')}
        fields: dict[str, list[dict[str, str]]] = {}
        all_types: dict[str, str] = {}
        for kind in KINDS:
            types = stored_column_types(snapshot, kind, task_repository)
            all_types.update({key: value for key, value in types.items() if all_types.get(key) != 'text'})
            fields[f'cdr-{kind}'] = [
                {'name': column, 'type': types.get(identity(column), 'text'), 'category': categories.get(identity(column), 'Other fields')}
                for column in columns_by_source.get(f'cdr-{kind}', []) if identity(column) not in internal
            ]
        fields['cdr-all'] = [
            {'name': column, 'type': all_types.get(identity(column), 'number' if identity(column) in {identity(field) for field in RF_CATALOG_FIELDS} else 'text'),
             'category': categories.get(identity(column), 'Other fields')}
            for column in columns_by_source.get('cdr-all', []) if identity(column) not in internal
        ]
        return {'columns_by_source': columns_by_source, 'fields': fields,
                'calculated_fields': [dimension.name for dimension in dimensions]}

    @app.get('/api/ppt-designer/slides/values')
    def design_slides_values(token: str, source: str, field: str, user=Depends(design_user)):
        """The values of a field in the selection of an open Design, with their number of samples (most first): read
        as the rows of a chart are (the same filters, maps and excluded operators)."""
        snapshot, _task_repository = designer_snapshot(token, user)
        if str(source).casefold() not in CATALOG_SOURCE_KINDS:
            raise HTTPException(400, 'Unsupported source.')
        entry = CatalogEntry(slide=1, slide_title='Values', slide_subtitle='', layout='Title + 1 rows + 1 columns', chart_title='',
                             cdr_source=source, kpi=field, chart_type='Table', legend='', filters='', rows_split=field,
                             columns_split='', calculated_dimensions=snapshot.dimensions)
        temporary = uuid4().hex
        with lock:
            snapshots[temporary] = replace(snapshot, owner=user.username, entries=[entry], chart_frames={}, filtered_frames={},
                                           chart_payloads={}, cancelled=None)
        try:
            _snapshot, _entry, frame = snapshot_chart(temporary, 0, user)
        finally:
            with lock:
                snapshots.pop(temporary, None)
        column = resolve_filter_column(frame, field) if frame is not None and not frame.empty else None
        if column is None:
            return {'field': field, 'values': []}
        counts = frame[column].dropna().astype(str).str.strip()
        counts = counts[counts.ne('')].value_counts()
        return {'field': field, 'values': [{'value': value, 'count': int(count)} for value, count in counts.head(2000).items()]}

    @app.get('/api/ppt-designer/{design_id}/slides')
    def design_slides_read(design_id: str, edit: bool = False, user=Depends(design_user)):
        """The slide set of a Design; ``edit`` (entering the Designer) first stores ids of its own in its template when it has
        none yet, so what a Design changes is kept by ids that do not move."""
        workspace_key()
        task_repository = bound_repository()
        definition = stored_design(design_id, task_repository)
        template, template_row = (None, None)
        if not definition.own_slides:
            template_row = template_row_of(definition.template, task_repository)
            if edit and (can_edit_design_slides(user, design_id, task_repository) or can_edit_template_slides(user, template_row)):
                template, template_row = template_with_ids(definition.template, task_repository, user.username)
                ensure_design_drafts(task_repository)
            else:
                template, _entries = template_slide_set(definition.template, task_repository)
            if template is None:
                raise HTTPException(400, 'The selected template is no longer available.')
        return slide_set_payload(design_id, definition, task_repository, user, template=template, template_row=template_row)

    @app.post('/api/ppt-designer/slides/preview')
    def design_slides_preview(request: DesignSlidesPreview, user=Depends(design_editor_user)):
        """Lay out a slide set being edited on the selection of an open Design: a new preview (token) whose charts
        are read from the same prepared CDR rows, at once."""
        workspace = workspace_key()
        with lock:
            snapshot = snapshots.get(request.token)
        if snapshot is None or snapshot.workspace != workspace or snapshot.owner not in {user.username, '*'}:
            raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
        task_repository = Repository(Path(workspace), core.repository.global_db_path)
        definition = snapshot.definition.model_copy(deep=True)
        if request.design_id:
            stored = read_designs(task_repository).get(request.design_id)
            if isinstance(stored, dict):
                definition.own_slides = copy.deepcopy(stored.get('own_slides'))
                definition.slide_overrides = copy.deepcopy(stored.get('slide_overrides') or {})
        template = None
        if request.template_slide_set is not None and not definition.own_slides:
            template = checked_slide_set(request.template_slide_set)
        elif not definition.own_slides:
            template, _entries = template_slide_set(definition.template, task_repository)
            if template is None:
                raise HTTPException(400, 'The selected template is no longer available.')
        if request.parameter_values is not None:
            definition.parameter_values = dict(request.parameter_values)
        if request.slide_set is not None:
            if definition.own_slides:
                definition.own_slides = checked_slide_set(request.slide_set, 'Design')
            else:
                definition.slide_overrides = design_slides.overrides_between(
                    template, request.slide_set, request.overrides if request.overrides is not None else definition.slide_overrides)
        elif request.overrides is not None:
            definition.slide_overrides = design_slides.normalized_overrides(request.overrides)
        entries = catalogue(definition, task_repository, template=template)
        deck = Presentation(core.settings.powerpoint_template)
        options = dict(snapshot.options or snapshot.payload.get('options') or {})
        for field_name in {repeat_field(entry) for entry in entries} - {''}:
            # A field chosen for Repeat by in the Designer and not read with the selection: its values are read now.
            if not any(identity(key) == identity(field_name) for key in options):
                try:
                    options[field_name] = load_filter_options(definition, field_name, task_repository)
                except HTTPException:
                    options[field_name] = []
        expanded, slides = laid_out_slides(
            entries, definition, options, task_repository, selected_sources(definition, task_repository), deck)
        payload = {**snapshot.payload, 'slides': slides}
        payload.pop('grid_layouts', None)
        add_dynamic_grid_metadata(payload, expanded, deck)
        token = uuid4().hex
        with lock:
            snapshots[token] = replace(
                snapshot, owner=user.username, entries=expanded, frames={}, definition=definition, payload=payload, options=options,
                chart_frames={}, filtered_frames={}, chart_payloads={}, frame_locks={}, sources=dict(snapshot.sources),
                cancelled=None, viewed_at=0.0,
            )
            while len(snapshots) > 128:
                snapshots.popitem(last=False)
        _template, resolved, report = design_slides_state(definition, task_repository, template)
        return {
            **payload, 'token': token, 'slide_set': resolved, 'report': report,
            'overrides': design_slides.normalized_overrides(definition.slide_overrides),
            'summary': design_slides.overrides_summary(definition.slide_overrides),
        }

    @app.put('/api/ppt-designer/{design_id}/slides')
    def save_design_slides(design_id: str, request: DesignSlidesSave, user=Depends(design_editor_user)):
        """Save the slide set of a Design: its template (Template scope) and its own differences (This Design)."""
        workspace_key()
        task_repository = bound_repository()
        definition = stored_design(design_id, task_repository)
        template_row = None if definition.own_slides else template_row_of(definition.template, task_repository)
        details: dict[str, Any] = {'id': design_id, 'name': definition.name}
        template = None
        if request.template_slide_set is not None:
            if definition.own_slides:
                raise HTTPException(400, 'This Design is detached from its template.')
            if not can_edit_template_slides(user, template_row):
                raise HTTPException(403, designer_refusal('PPT Template'))
            if request.template_updated_at and str(template_row['updated_at'] or '') != request.template_updated_at:
                raise HTTPException(409, f'The template "{definition.template}" changed since the Designer opened it: '
                                         'reload the Designer and apply your changes again.')
            template = checked_slide_set(request.template_slide_set)
            previous_content = bytes(template_row['content'] or b'')
            previous_slide_set = design_slides.slides_from_entries(
                core.load_template_catalogue(previous_content, task_repository=task_repository))
            content = design_slides.slides_csv(template)
            try:
                core.persist_report_template_for_request(definition.template, content, updated_by=user.username)
            except TimeoutError as exc:
                raise HTTPException(409, str(exc)) from exc
            new_entries = core.parse_catalog_csv(content)
            reconcile_template_slide_comments(definition.template, core.parse_catalog_csv(previous_content, validate_filters=False),
                                              new_entries, task_repository)
            invalidate_workspace_chart_models(task_repository.db_path)
            details['template'] = {
                'name': definition.template, 'before': design_slides.slides_summary(previous_slide_set),
                'after': design_slides.slides_summary(template),
                'designs': sum(1 for item in read_designs(task_repository).values()
                                  if isinstance(item, dict) and not item.get('own_slides') and item.get('template') == definition.template),
            }
            task_repository.add_log(user.username, 'report_template_saved_from_designer', json.dumps(details['template']))
            template_row = template_row_of(definition.template, task_repository)
        elif not definition.own_slides:
            template, template_row = template_with_ids(definition.template, task_repository, user.username)
            if template is None:
                raise HTTPException(400, 'The selected template is no longer available.')
        if request.slide_set is not None or request.overrides is not None or request.keep_template_changes:
            if not can_edit_design_slides(user, design_id, task_repository):
                raise HTTPException(403, designer_refusal('Design'))
            _base, before, _report = design_slides_state(definition, task_repository, template)
            if definition.own_slides:
                own = checked_slide_set(request.slide_set, 'Design') if request.slide_set is not None else definition.own_slides
                definition = DesignDefinition.model_validate(write_design_slides(design_id, task_repository, own_slides=own))
                details['slides'] = design_slides.slides_summary(own)
            else:
                previous = request.overrides if request.overrides is not None else definition.slide_overrides
                overrides = (design_slides.overrides_between(template, request.slide_set, previous) if request.slide_set is not None
                             else design_slides.normalized_overrides(previous))
                if request.keep_template_changes:
                    overrides = design_slides.keep_overrides(overrides, template)
                candidate = definition.model_copy(update={'slide_overrides': overrides}, deep=True)
                # Refused as a template would be (a slide of the Design with a chart that cannot be drawn).
                catalogue(candidate, task_repository, template=template)
                stored = overrides if design_slides.has_overrides(overrides) else {}
                definition = DesignDefinition.model_validate(
                    write_design_slides(design_id, task_repository, slide_overrides=stored))
                details['overrides'] = design_slides.overrides_summary(stored)
            _base, after, _report = design_slides_state(definition, task_repository, template)
            # Its slides retitled here keep their comments.
            reconcile_template_slide_comments(definition.template, design_slides.entries_from_slides(before),
                                              design_slides.entries_from_slides(after), task_repository,
                                              design_ids={design_id})
        if len(details) > 2:
            task_repository.add_log(user.username, 'design_slides_saved', json.dumps(details))
        # Saved: the draft of this user is the slide set now.
        ensure_design_drafts(task_repository)
        with task_repository.connection() as connection:
            connection.execute(f'DELETE FROM {DESIGN_DRAFTS_TABLE} WHERE design_id = ? AND username = ?',
                               (design_id, user.username))
        return slide_set_payload(design_id, definition, task_repository, user, template=template, template_row=template_row)

    @app.post('/api/ppt-designer/{design_id}/detach')
    def detach_design(design_id: str, user=Depends(design_editor_user)):
        """Detach: the Design keeps the slides it shows as its own and stops following its template."""
        workspace_key()
        task_repository = bound_repository()
        definition = stored_design(design_id, task_repository)
        if definition.own_slides:
            raise HTTPException(400, 'This Design is already detached from its template.')
        if not can_edit_design_slides(user, design_id, task_repository):
            raise HTTPException(403, designer_refusal('Design'))
        _template, resolved, _report = design_slides_state(definition, task_repository)
        own = design_slides.template_slides_of(resolved)
        definition = DesignDefinition.model_validate(
            write_design_slides(design_id, task_repository, own_slides=own, slide_overrides={}))
        task_repository.add_log(user.username, 'design_detached', json.dumps({
            'id': design_id, 'name': definition.name, 'template': definition.template,
            'slides': design_slides.slides_summary(own),
        }))
        return slide_set_payload(design_id, definition, task_repository, user)

    @app.post('/api/ppt-designer/{design_id}/save-as-template')
    def save_design_slides_as_template(design_id: str, payload: DesignSlidesTemplate,
                                          user=Depends(design_editor_user)):
        """Save as template: the slide set the Design shows (its template with its differences) as a new template."""
        workspace_key()
        task_repository = bound_repository()
        definition = stored_design(design_id, task_repository)
        name = payload.name.strip()
        if any(str(row['name']).casefold() == name.casefold() for row in task_repository.list_report_templates()):
            raise HTTPException(409, 'A PPT Template with this name already exists.')
        _template, resolved, _report = design_slides_state(definition, task_repository)
        slide_set = checked_slide_set(design_slides.template_slides_of(resolved))
        task_repository.add_report_template(name, design_slides.slides_csv(slide_set), updated_by=user.username)
        task_repository.add_log(user.username, 'report_template_saved_from_design', json.dumps({
            'design': design_id, 'name': definition.name, 'template': name,
            'from_template': '' if definition.own_slides else definition.template, 'slides': design_slides.slides_summary(slide_set),
        }))
        return {'template': name}

    def detach_template_designs(template_name: str, username: str) -> int:
        """Before a template is deleted, the Designs that follow it keep the slides they show as their own."""
        workspace_key()
        task_repository = bound_repository()
        detached = []
        with lock:
            designs = read_designs(task_repository)
            template = None
            for design_id, raw in designs.items():
                if not isinstance(raw, dict) or raw.get('own_slides') or str(raw.get('template') or '') != template_name:
                    continue
                if template is None:
                    template, _entries = template_slide_set(template_name, task_repository)
                    if template is None:
                        return 0
                resolved, _report = design_slides.resolve_slides(template, raw.get('slide_overrides'))
                designs[design_id] = {**raw, 'own_slides': design_slides.template_slides_of(resolved), 'slide_overrides': {}}
                detached.append(design_id)
            if detached:
                task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
                task_repository.add_log(username, 'design_detached', json.dumps({
                    'template': template_name, 'designs': detached, 'reason': 'template deleted',
                }))
        return len(detached)

    core.ppt_design_detach_template_designs = detach_template_designs

    layout_capacity_cache: dict[str, int] = {}

    def layout_capacity(layout: str) -> int:
        """How many charts a layout of the slide master places."""
        if layout not in layout_capacity_cache:
            deck = Presentation(core.settings.powerpoint_template)
            try:
                layout_capacity_cache[layout] = len(_layout_chart_frames(_named_slide_layout(deck, layout)))
            except (KeyError, ValueError):
                layout_capacity_cache[layout] = 0
        return layout_capacity_cache[layout]

    def free_places(slide) -> list[int]:
        visible = [element for element in slide.get('elements') or [] if not element.get('hidden')]
        used = {element.get('place', index) for index, element in enumerate(visible)}
        geometry = (slide.get('options') or {}).get('geometry')
        capacity = len(geometry) if isinstance(geometry, list) and geometry else layout_capacity(slide.get('layout') or '')
        return [place for place in range(capacity) if place not in used]

    @app.get('/api/ppt-designer/targets')
    def design_slides_targets(user=Depends(design_editor_user)):
        """The Designs whose slides the user can edit, with their content slides and how many free places each one has (Add to
        Design of Chart Builder)."""
        workspace_key()
        task_repository = bound_repository()
        result = []
        for design_id, raw in read_designs(task_repository).items():
            if not isinstance(raw, dict) or not can_edit_design_slides(user, design_id, task_repository):
                continue
            try:
                _template, resolved, _report = design_slides_state(DesignDefinition.model_validate(raw), task_repository)
            except HTTPException:
                continue
            slides = [{'id': slide['id'], 'title': slide.get('title') or 'Slide', 'free': len(free_places(slide))}
                      for slide in resolved.get('slides') or []
                      if (slide.get('kind') or 'content') == 'content' and not slide.get('hidden') and not slide.get('optional_off')]
            result.append({'id': design_id, 'name': str(raw.get('name') or ''), 'nr_mode': design_nr_mode(raw), 'slides': slides})
        # The PPT Templates whose slides the user can edit (a chart added to one reaches every Design that follows it).
        templates = []
        for row in task_repository.list_report_templates():
            if not can_edit_template_slides(user, row):
                continue
            try:
                slide_set = design_slides.slides_from_entries(core.load_template_catalogue(bytes(row['content'] or b''), task_repository=task_repository))
            except (ValueError, HTTPException):
                continue
            templates.append({'name': str(row['name']), 'slides': [
                {'id': slide['id'], 'title': slide.get('title') or 'Slide', 'free': len(free_places(slide))}
                for slide in slide_set.get('slides') or [] if (slide.get('kind') or 'content') == 'content']})
        return {'designs': sorted(result, key=lambda item: item['name'].casefold()),
                'templates': sorted(templates, key=lambda item: item['name'].casefold())}

    def added_element(payload: DesignSlidesAddChart, **extra) -> dict[str, Any]:
        """The chart of Chart Builder as an element of slides: its columns, its format and its One chart per, and a KPI Card as
        a KPI card with its settings."""
        element = {field: payload.element.get(field, False if field in {'exclude_null_empty', 'exclude_zero'} else '')
                   for field in design_slides.ELEMENT_FIELDS}
        element.update(id=design_slides.new_id('e'), kind='chart', **extra)
        if str(element['chart_type']).casefold() == design_slides.KIND_CHART_TYPES['kpi'].casefold():
            element['kind'] = 'kpi'
            if isinstance(payload.element.get('kpi_card'), dict):
                element['kpi_card'] = payload.element['kpi_card']
        for name in ('style', 'repeat'):
            if isinstance(payload.element.get(name), dict) and payload.element[name]:
                element[name] = payload.element[name]
        element['legend_position'] = str(element['legend_position'] or '').casefold()
        if not str(element['kpi']).strip() or not str(element['cdr_source']).strip():
            raise HTTPException(400, 'The chart needs a source and a KPI.')
        return element

    def place_added_element(slide_set: dict[str, Any], element: dict[str, Any], payload: DesignSlidesAddChart, label: str) -> str:
        """The element on the first free place of the slide chosen, or on a new slide at the end; the title of that slide."""
        if payload.slide_id:
            slide = next((item for item in slide_set.get('slides') or [] if item['id'] == payload.slide_id), None)
            if slide is None or (slide.get('kind') or 'content') != 'content':
                raise HTTPException(404, f'The slide is no longer in the {label}.')
            places = free_places(slide)
            if not places:
                raise HTTPException(409, f'The slide "{slide.get("title") or "Slide"}" has no free place: choose another slide or a new one.')
            element['place'] = places[0]
            slide.setdefault('elements', []).append(element)
            return slide.get('title') or 'Slide'
        title = payload.slide_title.strip() or element['chart_title'] or 'New slide'
        element['place'] = 0
        slide_set.setdefault('slides', []).append({
            'id': design_slides.new_id('s'), 'kind': 'content', 'title': title, 'subtitle': '',
            'layout': 'Title + 1 rows + 1 columns + comments down', 'optional': False, 'elements': [element],
        })
        return title

    @app.post('/api/ppt-designer/{design_id}/add-chart')
    def add_chart_to_design(design_id: str, payload: DesignSlidesAddChart, user=Depends(design_editor_user)):
        """Add to Design: the chart becomes part of this Design only (a difference of it from its template)."""
        workspace_key()
        task_repository = bound_repository()
        if not can_edit_design_slides(user, design_id, task_repository):
            raise HTTPException(403, designer_refusal('Design'))
        definition = stored_design(design_id, task_repository)
        template = None if definition.own_slides else template_with_ids(definition.template, task_repository, user.username)[0]
        _base, resolved, _report = design_slides_state(definition, task_repository, template)
        element = added_element(payload, origin='design')
        title = place_added_element(resolved, element, payload, 'Design')
        if definition.own_slides:
            write_design_slides(design_id, task_repository, own_slides=checked_slide_set(resolved, 'Design'))
        else:
            overrides = design_slides.overrides_between(template, resolved, definition.slide_overrides)
            catalogue(definition.model_copy(update={'slide_overrides': overrides}, deep=True), task_repository, template=template)
            write_design_slides(design_id, task_repository, slide_overrides=overrides)
        task_repository.add_log(user.username, 'design_chart_added', json.dumps({
            'id': design_id, 'name': definition.name, 'slide': title, 'new_slide': not payload.slide_id,
            'chart': element['chart_title'] or element['kpi'], 'chart_type': element['chart_type'], 'source': element['cdr_source'],
        }))
        return {'design': definition.name, 'slide': title}

    @app.post('/api/ppt-designer/templates/{template_name}/add-chart')
    def add_chart_to_template(template_name: str, payload: DesignSlidesAddChart, user=Depends(design_editor_user)):
        """Add to PPT Template: the chart becomes part of the template, and so of every PPT Design that follows it."""
        workspace_key()
        task_repository = bound_repository()
        template, row = template_with_ids(template_name, task_repository, user.username)
        if template is None:
            raise HTTPException(404, f'The PPT Template "{template_name}" is no longer in the workspace.')
        if not can_edit_template_slides(user, row):
            raise HTTPException(403, designer_refusal('PPT Template'))
        slide_set = copy.deepcopy(template)
        element = added_element(payload)
        title = place_added_element(slide_set, element, payload, 'PPT Template')
        content = design_slides.slides_csv(checked_slide_set(slide_set))
        previous_content = bytes(row['content'] or b'')
        try:
            core.persist_report_template_for_request(template_name, content, updated_by=user.username)
        except TimeoutError as exc:
            raise HTTPException(409, str(exc)) from exc
        reconcile_template_slide_comments(template_name, core.parse_catalog_csv(previous_content, validate_filters=False),
                                          core.parse_catalog_csv(content), task_repository)
        invalidate_workspace_chart_models(task_repository.db_path)
        task_repository.add_log(user.username, 'template_chart_added', json.dumps({
            'template': template_name, 'slide': title, 'new_slide': not payload.slide_id,
            'designs': sum(1 for item in read_designs(task_repository).values()
                           if isinstance(item, dict) and not item.get('own_slides') and item.get('template') == template_name),
            'chart': element['chart_title'] or element['kpi'], 'chart_type': element['chart_type'], 'source': element['cdr_source'],
        }))
        return {'template': template_name, 'slide': title}

    def ensure_design_drafts(task_repository) -> None:
        with task_repository.connection() as connection:
            connection.execute(
                f'CREATE TABLE IF NOT EXISTS {DESIGN_DRAFTS_TABLE} ('
                'design_id TEXT NOT NULL, username TEXT NOT NULL, state TEXT NOT NULL, updated_at TEXT NOT NULL, '
                'PRIMARY KEY (design_id, username))'
            )

    @app.get('/api/ppt-designer/{design_id}/slides/draft')
    def design_slides_draft(design_id: str, user=Depends(design_editor_user)):
        task_repository = bound_repository()
        ensure_design_drafts(task_repository)
        with task_repository.connection() as connection:
            row = connection.execute(
                f'SELECT state, updated_at FROM {DESIGN_DRAFTS_TABLE} WHERE design_id = ? AND username = ?',
                (design_id, user.username),
            ).fetchone()
        if row is None:
            return {'draft': None}
        try:
            state = json.loads(row['state'])
        except ValueError:
            state = None
        return {'draft': state, 'updated_at': str(row['updated_at'])}

    @app.put('/api/ppt-designer/{design_id}/slides/draft')
    def save_design_slides_draft(design_id: str, payload: DesignSlidesDraft, user=Depends(design_editor_user)):
        task_repository = bound_repository()
        stored_design(design_id, task_repository)
        encoded = json.dumps(payload.state)
        if len(encoded) > 8_000_000:
            raise HTTPException(413, 'The slides are too large to keep as a draft.')
        ensure_design_drafts(task_repository)
        now = datetime.now(timezone.utc).astimezone().isoformat(timespec='seconds')
        with task_repository.connection() as connection:
            connection.execute(
                f'INSERT INTO {DESIGN_DRAFTS_TABLE} (design_id, username, state, updated_at) VALUES (?, ?, ?, ?) '
                'ON CONFLICT(design_id, username) DO UPDATE SET state = excluded.state, updated_at = excluded.updated_at',
                (design_id, user.username, encoded, now),
            )
        return {'updated_at': now}

    @app.delete('/api/ppt-designer/{design_id}/slides/draft')
    def delete_design_slides_draft(design_id: str, discarded: bool = False, user=Depends(design_editor_user)):
        """Forget the draft of a Design (saved, or ``discarded`` by the user: logged in App Events)."""
        task_repository = bound_repository()
        ensure_design_drafts(task_repository)
        with task_repository.connection() as connection:
            deleted = connection.execute(
                f'DELETE FROM {DESIGN_DRAFTS_TABLE} WHERE design_id = ? AND username = ?',
                (design_id, user.username),
            ).rowcount
        if discarded:
            task_repository.add_log(user.username, 'design_draft_discarded', json.dumps({
                'id': design_id, 'draft': bool(deleted),
            }))
        return {'deleted': bool(deleted)}

    @app.patch('/api/ppt-designer/{design_id}/comments')
    def save_design_comments(design_id: str, payload: DesignComments, user=Depends(design_editor_user)):
        with lock:
            task_repository = bound_repository()
            designs = read_designs_and_previews(task_repository)
            if design_id not in designs:
                raise HTTPException(404, 'Design not found.')
            comments = {
                str(slide): [str(comment).strip() for comment in values if str(comment).strip()][:50]
                for slide, values in payload.slide_comments.items()
            }
            saved = {slide: values for slide, values in comments.items() if values}
            designs[design_id]['slide_comments'] = saved
            meta = comment_meta(designs[design_id].get('slide_comment_meta'), saved)
            if payload.generated_added and payload.generated_slide:
                snapshot = snapshots.get(payload.token) if payload.token else None
                present = set(saved.get(payload.generated_slide, []))
                added = [text for text in (str(item).strip() for item in payload.generated_texts) if text in present]
                if not payload.generated_texts:
                    # Without their texts, the comments just added are the last ones of the slide.
                    added = saved.get(payload.generated_slide, [])[-payload.generated_added:]
                previous = (meta.get(payload.generated_slide) or {}).get('generated') or []
                meta[payload.generated_slide] = {
                    'generated': list(dict.fromkeys([*previous, *added])),
                    'engine': payload.generated_engine or design_comments.ENGINE_RULES,
                    'model': payload.generated_model, 'generated_at': datetime.now(timezone.utc).isoformat(),
                    'generated_by': user.username,
                    'selection': snapshot.selection_key if snapshot is not None and snapshot.workspace == workspace_key() else '',
                }
            designs[design_id]['slide_comment_meta'] = meta
            store_design_entry(task_repository, design_id, designs[design_id])
            task_repository.add_log(user.username, 'save_design_comments', json.dumps({'id': design_id}))
            if payload.generated_added and payload.generated_slide:
                task_repository.add_log(user.username, 'add_generated_design_comments', json.dumps({
                    'id': design_id, 'name': designs[design_id].get('name', ''), 'slide': payload.generated_slide,
                    'added': payload.generated_added, 'engine': payload.generated_engine or design_comments.ENGINE_RULES,
                    'comments': len(saved.get(payload.generated_slide, [])),
                }))
        return {'slide_comments': saved, 'slide_comment_meta': meta}

    def comment_meta(meta, comments) -> dict[str, dict]:
        """The generated-comment metadata that still applies: a generated comment edited or removed by hand is no longer
        generated, and a slide without generated comments has none."""
        kept = {}
        for slide, info in (meta or {}).items():
            if not isinstance(info, dict):
                continue
            present = set(comments.get(slide, []))
            generated = [text for text in info.get('generated') or [] if text in present]
            if generated:
                kept[slide] = {**info, 'generated': generated}
        return kept

    def comment_snapshot(token: str, user):
        with lock:
            snapshot = snapshots.get(token)
        if snapshot is None or snapshot.workspace != workspace_key() or snapshot.owner not in {user.username, '*'}:
            raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
        return snapshot

    def slide_comment_charts(token: str, snapshot, slide: dict, requested, user) -> tuple[list[dict], list[str]]:
        """The charts of a slide the comments read, with the legend items the viewer hides on each one, and the
        titles of those they cannot read."""
        available = {int(chart['index']): chart for chart in slide.get('charts', []) if chart.get('available')}
        requested = requested or [GeneratedCommentChart(index=index) for index in available]
        charts, skipped = [], []
        for chart in requested:
            if chart.index not in available:
                continue
            model = chart_model(token, chart.index, user)
            if design_comments.needs_rebuild(model):
                # A model cached before the comments read their facts is calculated again once.
                model = chart_model(token, chart.index, user, force=True)
            entry = snapshot.entries[chart.index]
            if model.get('type') == 'empty':
                continue
            if model.get('type') not in design_comments.SUPPORTED_MODEL_TYPES:
                skipped.append(entry.chart_title or entry.chart_type)
                continue
            charts.append({
                'index': chart.index, 'title': entry.chart_title or entry.slide_title, 'kpi': entry.kpi,
                'model': model, 'hidden': chart.hidden, 'shown': chart.shown,
                # The value of a dynamic grid card (its Operator, Vendor…) names what its map or table shows.
                'subject': ' · '.join(value for value in (entry.dynamic_row_value, entry.dynamic_column_value) if value),
            })
        return charts, skipped

    def write_slide_comments(charts: list[dict], layout: str, task_repository, engine: str, username: str) -> tuple[list[dict], dict]:
        """The comments of one slide by its writer: the rules, or an LLM writer that rewrites their facts (checked
        against them; the rules comments when it fails). Returns them and what wrote them."""
        brief = design_comments.slide_brief(charts, focus_operators=focus_operators(task_repository), layout=layout or '')
        if engine not in comment_writers.SERVER_ENGINES:
            return brief['rules'], {'engine': design_comments.ENGINE_RULES, 'model': '', 'warning': '', 'problems': []}
        written = comment_writers.write_comments(engine, brief, repository=core.repository,
                                                 config_dir=core.application_config_dir, username=username)
        if written['engine'] == design_comments.ENGINE_RULES:
            return brief['rules'], written
        return [{'text': item, 'kind': written['engine'], 'chart_index': None, 'chart_title': ''} for item in written['texts']], written

    @app.post('/api/ppt-designer/{design_id}/comments/generate')
    def generate_design_comments(design_id: str, payload: GenerateComments, user=Depends(design_editor_user)):
        """Suggest the comments of one slide from the charts the viewer shows; nothing is saved until they are added."""
        started = monotonic()
        task_repository = bound_repository()
        name = read_designs(task_repository).get(design_id, {}).get('name')
        if name is None:
            raise HTTPException(404, 'Design not found.')
        snapshot = comment_snapshot(payload.token, user)
        slide = next((item for item in snapshot.payload.get('slides', []) if item.get('number') == payload.slide), None)
        if slide is None:
            raise HTTPException(404, 'Slide not found.')
        engine = comment_writers.resolve_engine(core.repository, payload.engine)
        try:
            charts, skipped = slide_comment_charts(payload.token, snapshot, slide, payload.charts, user)
            suggestions, written = write_slide_comments(charts, slide.get('layout') or '', task_repository, engine, user.username)
        except HTTPException:
            raise
        except Exception as exc:
            task_repository.add_log(user.username, 'generate_design_comments_failed', json.dumps({
                'id': design_id, 'name': name, 'slide': payload.slide, 'error': str(exc),
            }))
            raise HTTPException(500, f'The comments of this slide could not be generated: {exc}') from exc
        task_repository.add_log(user.username, 'generate_design_comments', json.dumps({
            'id': design_id, 'name': name, 'slide': payload.slide, 'comment_key': slide.get('comment_key', ''),
            'requested_engine': engine, 'engine': written['engine'], 'model': written.get('model', ''),
            'charts': len(charts), 'skipped_charts': skipped, 'suggestions': len(suggestions),
            **{key: written[key] for key in ('input_tokens', 'output_tokens') if key in written},
            'left_out': written.get('problems', []), 'warning': written.get('warning', ''),
            'duration_ms': round((monotonic() - started) * 1000),
        }))
        return {
            'comment_key': slide.get('comment_key', ''), 'engine': written['engine'], 'model': written.get('model', ''),
            'suggestions': suggestions, 'skipped': skipped, 'warning': written.get('warning', ''),
        }

    @app.get('/api/ppt-designer/comment-engines')
    def design_comment_engines(user=Depends(design_user)):
        """The writers a user can choose (Rules and those Super Admin enabled), the default one, and whether the
        Kiro Brief is enabled."""
        return comment_writers.available_engines(core.repository)

    def apply_generated_comments(definition: dict, written: dict[str, list[str]], mode: str, legacy: dict[str, str],
                                 info: dict, writers: dict[str, dict] | None = None) -> dict[str, int]:
        """Write the generated comments of several slides into a Design definition: on the slides without comments
        (``empty``), in place of the generated ones keeping those written by hand (``generated``), or in place of every
        comment (``all``). Returns how many slides were written, kept and left without comments."""
        comments = dict(definition.get('slide_comments') or {})
        meta = comment_meta(definition.get('slide_comment_meta'), comments)
        counts = {'written': 0, 'kept': 0, 'empty': 0, 'comments': 0}
        for key, texts in written.items():
            current = list(comments.get(key, comments.get(legacy.get(key, ''), [])))
            if not texts:
                counts['empty'] += 1
                continue
            if mode == 'empty' and current:
                counts['kept'] += 1
                continue
            generated = set((meta.get(key) or {}).get('generated') or [])
            manual = [text for text in current if text not in generated] if mode == 'generated' else []
            comments[key] = [*manual, *[text for text in texts if text not in manual]][:50]
            comments.pop(legacy.get(key, ''), None)
            meta[key] = {**info, **((writers or {}).get(key) or {}), 'generated': [text for text in texts if text in comments[key]]}
            counts['written'] += 1
            counts['comments'] += len(meta[key]['generated'])
        definition['slide_comments'] = {key: values for key, values in comments.items() if values}
        definition['slide_comment_meta'] = comment_meta(meta, definition['slide_comments'])
        return counts

    def run_design_comment_job(job_id: str, design_id: str, payload: GenerateAllComments, user, workspace: str) -> None:
        """Generate the comments of every slide of a Design and write them as its mode says."""
        started = monotonic()
        task_repository = Repository(Path(workspace), core.repository.global_db_path)

        def update(**changes):
            with lock:
                job = comment_jobs.get(job_id)
                if job is not None:
                    job.update(changes)

        try:
            snapshot = comment_snapshot(payload.token, user) if workspace == workspace_key() else None
            if snapshot is None:
                raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
            requested = {item.number: item for item in payload.slides}
            slides = [slide for slide in snapshot.payload.get('slides', [])
                      if not slide.get('structural_type') and (not requested or slide.get('number') in requested)]
            written: dict[str, list[str]] = {}
            writers: dict[str, dict] = {}
            legacy: dict[str, str] = {}
            skipped: list[str] = []
            warnings: list[str] = []
            engine = comment_writers.resolve_engine(core.repository, payload.engine)
            for position, slide in enumerate(slides, 1):
                spec = requested.get(slide.get('number'))
                charts, not_read = slide_comment_charts(payload.token, snapshot, slide, spec.charts if spec else [], user)
                skipped.extend(not_read)
                key = str(slide.get('comment_key') or slide.get('number'))
                suggestions, info_written = (write_slide_comments(charts, slide.get('layout') or '', task_repository, engine, user.username)
                                             if charts else ([], {'engine': design_comments.ENGINE_RULES, 'model': ''}))
                written[key] = [item['text'] for item in suggestions]
                writers[key] = {'engine': info_written['engine'], 'model': info_written.get('model', '')}
                if info_written.get('warning'):
                    warnings.append(info_written['warning'])
                legacy[key] = str(slide.get('number'))
                update(done=position)
            with lock:
                designs = read_designs(task_repository)
                if design_id not in designs:
                    raise HTTPException(404, 'Design not found.')
                info = {
                    'engine': design_comments.ENGINE_RULES, 'model': '',
                    'generated_at': datetime.now(timezone.utc).isoformat(), 'generated_by': user.username,
                    'selection': snapshot.selection_key,
                }
                counts = apply_generated_comments(designs[design_id], written, payload.mode, legacy, info, writers)
                task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
                definition = designs[design_id]
            task_repository.add_log(user.username, 'generate_all_design_comments', json.dumps({
                'id': design_id, 'name': definition.get('name', ''), 'mode': payload.mode, 'requested_engine': engine,
                'engines': sorted({item['engine'] for item in writers.values()}), 'warnings': sorted(set(warnings))[:5],
                'slides': len(slides), **counts,
                'skipped_charts': sorted(set(skipped)), 'duration_ms': round((monotonic() - started) * 1000),
            }))
            update(state='ready', result={
                **counts, 'slides': len(slides), 'warnings': sorted(set(warnings))[:5],
                'slide_comments': definition.get('slide_comments', {}),
                'slide_comment_meta': definition.get('slide_comment_meta', {}),
            })
        except HTTPException as exc:
            update(state='failed', status=exc.status_code, detail=str(exc.detail))
        except Exception as exc:
            task_repository.add_log(user.username, 'generate_all_design_comments_failed', json.dumps({
                'id': design_id, 'mode': payload.mode, 'error': str(exc),
            }))
            update(state='failed', status=500, detail=f'The comments of this Design could not be generated: {exc}')

    @app.post('/api/ppt-designer/{design_id}/comments/generate-all')
    def generate_all_design_comments(
        design_id: str, payload: GenerateAllComments, background_tasks: BackgroundTasks, user=Depends(design_editor_user),
    ):
        """Start writing the comments of every slide of a Design; the page follows the job until it is done."""
        task_repository = bound_repository()
        if design_id not in read_designs(task_repository):
            raise HTTPException(404, 'Design not found.')
        snapshot = comment_snapshot(payload.token, user)
        total = sum(1 for slide in snapshot.payload.get('slides', []) if not slide.get('structural_type')
                    and (not payload.slides or slide.get('number') in {item.number for item in payload.slides}))
        job_id = uuid4().hex
        with lock:
            comment_jobs[job_id] = {
                'state': 'processing', 'owner': user.username, 'workspace': workspace_key(), 'done': 0, 'total': total,
            }
            while len(comment_jobs) > 16:
                removable = next((key for key, value in comment_jobs.items() if value.get('state') != 'processing'), None)
                if removable is None:
                    break
                comment_jobs.pop(removable, None)
        background_tasks.add_task(run_design_comment_job, job_id, design_id, payload, user, workspace_key())
        return JSONResponse({'state': 'processing', 'job_id': job_id, 'done': 0, 'total': total}, status_code=202)

    def run_kiro_brief_job(job_id: str, design_id: str, payload: GenerateAllComments, user, workspace: str) -> None:
        """Read the facts of every slide for the Kiro Brief: Kiro IDE writes the comments from them."""
        task_repository = Repository(Path(workspace), core.repository.global_db_path)

        def update(**changes):
            with lock:
                job = comment_jobs.get(job_id)
                if job is not None:
                    job.update(changes)

        try:
            snapshot = comment_snapshot(payload.token, user) if workspace == workspace_key() else None
            if snapshot is None:
                raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
            requested = {item.number: item for item in payload.slides}
            slides = [slide for slide in snapshot.payload.get('slides', [])
                      if not slide.get('structural_type') and (not requested or slide.get('number') in requested)]
            focus = focus_operators(task_repository)
            briefed = []
            for position, slide in enumerate(slides, 1):
                spec = requested.get(slide.get('number'))
                charts, _skipped = slide_comment_charts(payload.token, snapshot, slide, spec.charts if spec else [], user)
                if charts:
                    brief = design_comments.slide_brief(charts, focus_operators=focus, layout=slide.get('layout') or '')
                    brief.pop('rules', None)
                    briefed.append({'comment_key': str(slide.get('comment_key') or slide.get('number')), 'number': slide.get('number'),
                                    'title': slide.get('title', ''), 'subtitle': slide.get('subtitle', ''), 'layout': slide.get('layout', ''),
                                    **brief})
                update(done=position)
            name = read_designs(task_repository).get(design_id, {}).get('name', design_id)
            brief_id = uuid4().hex
            with lock:
                comment_briefs[brief_id] = {
                    'owner': user.username, 'workspace': workspace, 'design_id': design_id, 'name': name,
                    'selection': snapshot.selection_key, 'slides': briefed,
                    'generated_at': datetime.now(timezone.utc).isoformat(timespec='seconds'),
                }
                while len(comment_briefs) > 16:
                    comment_briefs.popitem(last=False)
            task_repository.add_log(user.username, 'export_kiro_comment_brief', json.dumps({
                'id': design_id, 'name': name, 'slides': len(briefed),
            }))
            update(state='ready', result={'brief_id': brief_id, 'slides': len(briefed)})
        except HTTPException as exc:
            update(state='failed', status=exc.status_code, detail=str(exc.detail))
        except Exception as exc:
            update(state='failed', status=500, detail=f'The Kiro Brief could not be prepared: {exc}')

    @app.post('/api/ppt-designer/{design_id}/comments/kiro-brief')
    def start_kiro_brief(design_id: str, payload: GenerateAllComments, background_tasks: BackgroundTasks,
                         user=Depends(design_editor_user)):
        """Start reading the facts of every slide for the Kiro Brief; the page downloads it when the job is done."""
        if not comment_writers.available_engines(core.repository)['kiro_brief']:
            raise HTTPException(403, 'Super Admin has not enabled the Kiro Brief.')
        if design_id not in read_designs(bound_repository()):
            raise HTTPException(404, 'Design not found.')
        snapshot = comment_snapshot(payload.token, user)
        total = sum(1 for slide in snapshot.payload.get('slides', []) if not slide.get('structural_type'))
        job_id = uuid4().hex
        with lock:
            comment_jobs[job_id] = {'state': 'processing', 'owner': user.username, 'workspace': workspace_key(), 'done': 0, 'total': total}
        background_tasks.add_task(run_kiro_brief_job, job_id, design_id, payload, user, workspace_key())
        return JSONResponse({'state': 'processing', 'job_id': job_id, 'done': 0, 'total': total}, status_code=202)

    @app.get('/api/ppt-designer/comment-briefs/{brief_id}.zip')
    def download_kiro_brief(brief_id: str, user=Depends(design_user)):
        with lock:
            brief = comment_briefs.get(brief_id)
        if brief is None or brief['owner'] != user.username or brief['workspace'] != workspace_key():
            raise HTTPException(404, 'Kiro Brief not found: prepare it again.')
        content = comment_writers.kiro_brief_archive(brief)
        name = re.sub(r'[^\w.-]+', '_', f"Kiro brief - {brief['name']}").strip('_')
        return Response(content, media_type='application/zip', headers={'Content-Disposition': f'attachment; filename="{name}.zip"'})

    @app.post('/api/ppt-designer/{design_id}/comments/kiro-import')
    async def import_kiro_comments(
        design_id: str, files: list[UploadFile] = File(...), mode: str = Form('empty'), user=Depends(design_editor_user),
    ):
        """Check the comments Kiro IDE wrote from a Kiro Brief against its facts and write those that keep to them."""
        if mode not in {'empty', 'generated', 'all'}:
            raise HTTPException(400, 'Choose which comments the imported ones replace.')
        uploads = [(upload.filename or '', await upload.read()) for upload in files]
        task_repository = bound_repository()
        try:
            comments_document, facts_document = comment_writers.read_kiro_import(uploads)
            written, problems = comment_writers.check_kiro_comments(comments_document, facts_document)
        except ValueError as exc:
            task_repository.add_log(user.username, 'import_kiro_comments_failed', json.dumps({'id': design_id, 'error': str(exc)}))
            raise HTTPException(400, str(exc)) from exc
        with lock:
            designs = read_designs(task_repository)
            if design_id not in designs:
                raise HTTPException(404, 'Design not found.')
            legacy = {str(slide.get('comment_key')): str(slide.get('number')) for slide in facts_document.get('slides') or []}
            info = {'engine': comment_writers.ENGINE_KIRO_BRIEF, 'model': '', 'generated_by': user.username,
                    'generated_at': datetime.now(timezone.utc).isoformat(), 'selection': str(facts_document.get('selection') or '')}
            counts = apply_generated_comments(designs[design_id], written, mode, legacy, info)
            task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
            definition = designs[design_id]
        task_repository.add_log(user.username, 'import_kiro_comments', json.dumps({
            'id': design_id, 'name': definition.get('name', ''), 'mode': mode, 'slides': len(written), **counts,
            'left_out': len(problems), 'problems': problems[:5],
        }))
        return {**counts, 'slides': len(written), 'left_out': len(problems), 'problems': problems[:5],
                'slide_comments': definition.get('slide_comments', {}), 'slide_comment_meta': definition.get('slide_comment_meta', {})}

    @app.get('/api/ppt-designer/comment-jobs/{job_id}')
    def design_comment_job(job_id: str, user=Depends(design_user)):
        with lock:
            job = comment_jobs.get(job_id)
            if job is None or job.get('owner') != user.username or job.get('workspace') != workspace_key():
                raise HTTPException(404, 'Comment job not found.')
            result = dict(job)
        if result.get('state') == 'failed':
            raise HTTPException(int(result.get('status') or 500), str(result.get('detail') or 'The comments could not be generated.'))
        return JSONResponse({key: result.get(key) for key in ('state', 'done', 'total', 'result')}, headers={'Cache-Control': 'no-store'})

    def design_datasets(definition, kind, task_repository):
        """Return the selected CDRs of one kind that match the Design's NR Mode."""
        if kind == 'all':
            return [row for source in KINDS for row in design_datasets(definition, source, task_repository)]
        datasets = definition.get('datasets', {}) if isinstance(definition, dict) else definition.datasets
        nr_mode = design_nr_mode(definition)
        return [
            row for row in core._optional_reporting_datasets(list((datasets or {}).get(kind, [])), kind, task_repository)
            # A Design only analyses CDRs of its own NR Mode.
            if dataset_matches_nr_mode(row, nr_mode)
        ]

    def selected_sources(definition, task_repository):
        selected_by_kind = {}
        for kind in KINDS:
            selected = design_datasets(definition, kind, task_repository)
            if definition.scope == 'multivendor' and any(not row.get('vendor_mapping_applied') for row in selected):
                raise HTTPException(400, 'Map Vendors for every selected CDR before using Multivendor Comparison.')
            if selected:
                selected_by_kind[kind] = selected
        if not selected_by_kind:
            raise HTTPException(400, f'No ready {definition.nr_mode} CDR is in the combined CDR tables for the chosen Campaign.')
        if definition.scope == 'multivendor':
            selected_ids = [int(row['id']) for rows in selected_by_kind.values() for row in rows]
            if len(task_repository.cdr_catalogue_values(selected_ids)['vendors']) < 2:
                raise HTTPException(400, 'Vendor Comparison requires at least two distinct Vendors in the selected CDRs.')
        return selected_by_kind

    def is_complete_unfiltered_universe(definition, selected_by_kind, task_repository, date_bounds):
        """Whether the selection exactly matches every row in each combined CDR table."""
        if definition.filters:
            return False
        if not ignore_event_time_filtering() and date_bounds and (
            str(definition.date_from) != date_bounds['min']
            or str(definition.date_to) != date_bounds['max']
        ):
            return False
        included = core.combined_dataset_ids(task_repository)
        for kind, selected in selected_by_kind.items():
            selected_ids = {int(row['id']) for row in selected}
            ready_ids = {
                int(row['id']) for row in task_repository.list_datasets()
                if row['status'] == 'ready' and str(row['dataset_kind'] or '').casefold() == kind
                and int(row['id']) in included
            }
            if selected_ids != ready_ids:
                return False
        return bool(selected_by_kind)

    def resolve_sql_column(columns, field):
        lookup = {identity(column): column for column in columns}
        if identity(field) == 'vendor' and 'vendor' in lookup:
            return lookup['vendor']
        aliases = next((values for label, values in FILTER_COLUMNS.items() if identity(label) == identity(field)), (field,))
        return next((lookup[identity(alias)] for alias in aliases if identity(alias) in lookup), None)

    def filter_sql_value_expression(task_repository, columns, field):
        """Return a text value expression with row-level fallback aliases."""
        field = vendor_filter_column(field)
        lookup = {identity(column): column for column in columns}
        aliases = next((values for label, values in FILTER_COLUMNS.items() if identity(label) == identity(field)), (field,))
        resolved = list(dict.fromkeys(
            lookup[identity(alias)] for alias in aliases if identity(alias) in lookup
        ))
        if not resolved:
            return None
        quote = task_repository._quote_identifier
        if identity(field) in {identity('City'), identity('Region')}:
            candidates = ', '.join(
                f"NULLIF(TRIM(CAST({quote(column)} AS TEXT)), '')" for column in resolved
            )
            return f"COALESCE({candidates}, '')"
        return f"COALESCE(CAST({quote(resolved[0])} AS TEXT), '')"

    def filter_sql_normalized_expression(task_repository, columns, field):
        """Return an indexed normalized expression when one physical column owns the filter."""
        field = vendor_filter_column(field)
        lookup = {identity(column): column for column in columns}
        aliases = next((values for label, values in FILTER_COLUMNS.items() if identity(label) == identity(field)), (field,))
        resolved = list(dict.fromkeys(
            lookup[identity(alias)] for alias in aliases if identity(alias) in lookup
        ))
        if not resolved:
            return None
        quote = task_repository._quote_identifier
        if len(resolved) == 1:
            return f"LOWER(TRIM(CAST({quote(resolved[0])} AS TEXT)))"
        value_expression = filter_sql_value_expression(task_repository, columns, field)
        return f"LOWER(TRIM({value_expression}))" if value_expression is not None else None

    def pattern_filter_values(task_repository, mapper, field_name, values, catalogue):
        """The CDR spellings a filter with patterns (VF*, -Three) takes: those of the CDR catalogues whose spelling or
        Operator / Vendor Maps label it accepts (see value_patterns), so -Three also leaves out the 3 of the CDRs. A
        filter whose patterns take no known spelling keeps them: it then takes no value instead of every value."""
        field = vendor_filter_column(field_name)
        source = PATTERN_CATALOGUE_FIELDS.get(identity(field))
        kind = field_kind(field)
        if source and catalogue.get(source) is None:
            catalogue[source] = task_repository.cdr_catalogue_values().get(source, [])
        chosen = ValueSet(values)
        taken = [value for value in catalogue.get(source) or [] if chosen.accepts([value, mapper.map(kind, value)])]
        return taken or [value for value in values if is_pattern(value)]

    def selection_filter_terms(task_repository, columns, definition, exclude=None):
        """Return each user-selected Design filter as (normalized SQL or None, values).

        A ``None`` expression selects no row. SQLite queries and the per-CDR
        column cache both evaluate these terms, so they always agree.
        """
        terms = []
        excluded = identity(exclude) if exclude else ''
        mapper = ValueMapper.from_repository(task_repository)
        catalogue = {}
        for field_name, values in definition.filters.items():
            # Patterns (VF*, -Three) become the spellings they take: these terms compare values as they are.
            if any(is_pattern(value) for value in values or []):
                values = pattern_filter_values(task_repository, mapper, field_name, values, catalogue)
            if identity(vendor_filter_column(field_name)) in {'vendor', 'operatorvendor'}:
                values = vendor_match_values(field_name, values, task_repository.list_operator_mappings())
            # Operators and Vendors are chosen with their mapped labels: match every source spelling.
            kind = field_kind(vendor_filter_column(field_name))
            if kind == 'vendor_operator' and values:
                values = mapper.sources('operator_vendor', mapper.operator_vendors(values))
            elif kind and values:
                values = mapper.sources(kind, values)
            if identity(field_name) == excluded:
                continue
            normalized_expression = filter_sql_normalized_expression(
                task_repository, columns, field_name,
            )
            if normalized_expression is None or not values:
                terms.append((None, ()))
                continue
            terms.append((normalized_expression, tuple(str(value).strip().lower() for value in values)))
        return terms

    def selection_filter_where(task_repository, columns, definition, exclude=None):
        """Build only the user-selected PPT filter predicates."""
        clauses, params = [], []
        for normalized_expression, values in selection_filter_terms(task_repository, columns, definition, exclude):
            if normalized_expression is None:
                clauses.append('0')
                continue
            value_placeholders = ', '.join('?' for _ in values)
            clauses.append(f"{normalized_expression} IN ({value_placeholders})")
            params.extend(values)
        return ' AND '.join(f'({clause})' for clause in clauses) or '1', params

    def selection_date_expression(task_repository, columns):
        date_column = next((resolve_sql_column(columns, candidate) for candidate in ('event_start_time', 'Test_Start_Time', 'Timestamp', 'Date') if resolve_sql_column(columns, candidate)), None)
        return f"date(CAST({task_repository._quote_identifier(date_column)} AS TEXT))" if date_column else None

    def selection_date_term(task_repository, columns, definition):
        """Return the universe's (event-date SQL or None, low, high) bounds, if any."""
        concrete_from = definition.date_from if isinstance(definition.date_from, date) and not ignore_event_time_filtering() else None
        concrete_to = definition.date_to if isinstance(definition.date_to, date) and not ignore_event_time_filtering() else None
        if not concrete_from and not concrete_to:
            return None
        return (
            selection_date_expression(task_repository, columns),
            concrete_from.isoformat() if concrete_from else None,
            concrete_to.isoformat() if concrete_to else None,
        )

    def selection_sheet_term(task_repository, columns):
        """Return the (source sheet column, ignored sheet keys) exclusion, if any."""
        source_sheet = resolve_sql_column(columns, 'source_sheet')
        if source_sheet and core.CDR_IGNORED_SHEET_KEYS:
            return task_repository._quote_identifier(source_sheet), sorted(core.CDR_IGNORED_SHEET_KEYS)
        return None

    def selection_where(
        task_repository, kind, dataset_ids, definition, exclude=None, *, columns=None, include_dataset_scope=True,
    ):
        columns = columns or task_repository.list_reporting_row_columns(kind)
        clauses = []
        params = []
        if include_dataset_scope:
            placeholders = ', '.join('?' for _ in dataset_ids)
            clauses.append(f"dataset_id IN ({placeholders})")
            params.extend(int(dataset_id) for dataset_id in dataset_ids)
        filter_where, filter_params = selection_filter_where(
            task_repository, columns, definition, exclude=exclude,
        )
        if filter_where != '1':
            clauses.append(filter_where)
            params.extend(filter_params)
        for operator_expression, excluded_names in selection_excluded_operators(task_repository, columns):
            # The operators the Operator Maps exclude are left out of every Design.
            clauses.append(f"NOT COALESCE({operator_expression} IN ({', '.join('?' for _ in excluded_names)}), 0)")
            params.extend(sorted(excluded_names))
        for mode_datasets, operator_expression, excluded_names in selection_out_of_scope_operators(task_repository, columns):
            # And those out of scope in the NR Mode of each CDR (the NR Modes of the Operator Maps).
            clauses.append(f"NOT COALESCE(dataset_id IN ({', '.join('?' for _ in mode_datasets)}) "
                           f"AND {operator_expression} IN ({', '.join('?' for _ in excluded_names)}), 0)")
            params.extend([*sorted(mode_datasets), *sorted(excluded_names)])
        date_term = selection_date_term(task_repository, columns, definition)
        if date_term is not None:
            date_expression, low, high = date_term
            if date_expression is None:
                clauses.append('0')
            else:
                if low:
                    clauses.append(f"{date_expression} >= date(?)")
                    params.append(low)
                if high:
                    clauses.append(f"{date_expression} <= date(?)")
                    params.append(high)
        sheet_term = selection_sheet_term(task_repository, columns)
        if sheet_term is not None:
            source_sheet, ignored = sheet_term
            clauses.append(f"({source_sheet} IS NULL OR LOWER(TRIM(CAST({source_sheet} AS TEXT))) NOT IN ({', '.join('?' for _ in ignored)}))")
            params.extend(ignored)
        return ' AND '.join(f'({clause})' for clause in clauses) or '1', params

    def selection_cache_terms(task_repository, columns, definition):
        """The universe's row predicates for the per-CDR column cache."""
        date_term = selection_date_term(task_repository, columns, definition)
        sheet_term = selection_sheet_term(task_repository, columns)
        return SelectionTerms(
            values=tuple(
                (expression, frozenset(values))
                for expression, values in selection_filter_terms(task_repository, columns, definition)
            ),
            date=date_term,
            excluded_sheets=(
                (f'LOWER(TRIM(CAST({sheet_term[0]} AS TEXT)))', frozenset(sheet_term[1])) if sheet_term else None
            ),
            excluded_values=selection_excluded_operators(task_repository, columns),
            dataset_excluded_values=selection_out_of_scope_operators(task_repository, columns),
        )

    def selection_excluded_operators(task_repository, columns):
        """The (normalized Operator SQL, excluded names) term of the operators the Operator Maps exclude, if any."""
        names = excluded_operators(task_repository)
        expression = filter_sql_normalized_expression(task_repository, columns, 'Operator') if names else None
        return ((expression, frozenset(name.lower() for name in names)),) if expression is not None else ()

    def selection_out_of_scope_operators(task_repository, columns):
        """The (CDRs of an NR Mode, normalized Operator SQL, names) terms of the operators out of scope in each NR Mode
        (the NR Modes of the Operator Maps), if any."""
        groups = task_repository.list_operator_mapping_groups()
        if not out_of_scope_pairs(groups):
            return ()
        expression = filter_sql_normalized_expression(task_repository, columns, 'Operator')
        if expression is None:
            return ()
        modes = task_repository.ready_cdr_nr_modes()
        terms = []
        for mode in NR_MODES:
            names = frozenset(name.lower() for name in out_of_scope_operators(groups, mode))
            datasets = frozenset(dataset_id for dataset_id, value in modes.items() if value == mode)
            if names and datasets:
                terms.append((datasets, expression, names))
        return tuple(terms)

    def design_combined_requested_columns(definition, dimensions, kind):
        default_filter_keys = {identity(field) for field in ADAPTATIVE_FILTER_FIELDS}
        # The combined-table contract owns fixed preview columns, Default
        # Filter aliases and saved-template KPI fields. Some of those aliases
        # are alternatives, so comparing the whole contract with one physical
        # CDR schema would report false missing columns forever. Design
        # preparation only owns genuinely Design-specific extensions.
        candidates = {
            field for field in definition.custom_fields
            if identity(field) not in default_filter_keys
        }
        candidates.update(
            field for field in definition.filters
            if identity(field) not in default_filter_keys
        )
        dimensions_by_name = {identity(dimension.name): dimension for dimension in dimensions}
        return {
            field for field in candidates
            if (
                identity(field) not in dimensions_by_name
                or f'cdr-{kind}' in dimensions_by_name[identity(field)].sources
            )
        }

    def design_combined_columns_missing(definition, task_repository, dimensions, selected_by_kind):
        """Whether opening this Design must add a genuinely new source column."""
        for kind in selected_by_kind:
            requested = design_combined_requested_columns(definition, dimensions, kind)
            current = {identity(column) for column in task_repository.list_reporting_row_columns(kind)}
            if not {identity(column) for column in requested}.issubset(current):
                return True
        return False

    # Per CDR type and revision of its combined table, kept in the workspace (a restart keeps it): the columns the charts
    # of Designs asked for that none of its CDRs has.
    def unavailable_combined_columns(task_repository) -> dict:
        try:
            stored = json.loads(task_repository.get_workspace_state(UNAVAILABLE_COMBINED_COLUMNS_STATE_KEY) or '{}')
        except (TypeError, ValueError):
            stored = {}
        return stored if isinstance(stored, dict) else {}

    def remember_unavailable_combined_columns(task_repository, key: str, columns: set[str]) -> None:
        stored = unavailable_combined_columns(task_repository)
        if set(stored.get(key) or []) == columns:
            return
        stored[key] = sorted(columns)
        # Only the latest revisions matter.
        stored = dict(list(stored.items())[-24:])
        task_repository.set_workspace_state(UNAVAILABLE_COMBINED_COLUMNS_STATE_KEY, json.dumps(stored))

    def ensure_combined_filter_columns(
        definition, task_repository, dimensions, selected_by_kind, chart_entries=None,
    ):
        selected_dimension_keys = {identity(name) for name in definition.custom_fields}
        active_dimensions = tuple(
            dimension for dimension in dimensions if identity(dimension.name) in selected_dimension_keys
        )
        for kind in selected_by_kind:
            requested = design_combined_requested_columns(definition, dimensions, kind)
            if chart_entries:
                requested.update(core.reporting_query_columns(
                    kind, chart_entries, definition.scope == 'multivendor',
                ))
            current = {identity(column) for column in task_repository.list_reporting_row_columns(kind)}
            wanted = {identity(column) for column in requested}
            # Columns asked for that no CDR of the type has stay missing: until the combined table changes they are not
            # looked for again in every CDR at each preparation.
            unavailable_key = f"{kind}:{task_repository.get_workspace_state(f'combined_reporting_updated_{kind}') or ''}"
            if wanted.issubset(current | set(unavailable_combined_columns(task_repository).get(unavailable_key) or [])):
                continue
            rows_changed = columns_changed = False
            included = core.combined_dataset_ids(task_repository, kind)
            all_ready_sources = [
                row for row in task_repository.list_datasets()
                if row['status'] == 'ready' and str(row['dataset_kind'] or '').casefold() == kind
                and int(row['id']) in included
            ]
            for row in all_ready_sources:
                # A newly requested Design-specific column belongs to the
                # shared source, not merely the currently selected universe.
                # Future dataset selections can therefore stay SQL-only.
                change = task_repository.copy_dataset_rows_to_reporting(
                    row['id'], kind, sorted(requested, key=str.casefold),
                )
                rows_changed = rows_changed or change == 'rows'
                columns_changed = columns_changed or change == 'columns'
            current = {identity(column) for column in task_repository.list_reporting_row_columns(kind)}
            remember_unavailable_combined_columns(task_repository, unavailable_key, wanted - current)
            missing_dimensions = tuple(
                dimension for dimension in active_dimensions
                if f'cdr-{kind}' in dimension.sources and identity(dimension.name) not in current
            )
            if missing_dimensions:
                core._incremental_auto_field_table_update(
                    task_repository, task_repository.reporting_rows_table_name(kind),
                    f'cdr-{kind}', (), missing_dimensions, {},
                )
                columns_changed = True
            if rows_changed or columns_changed:
                # Filling new columns keeps prepared Design universes valid;
                # only reinserted rows advance their cache revision.
                core.mark_combined_reporting_updated(task_repository, kind, rows_changed=rows_changed)

    def persistent_selection_key(definition, task_repository, dimensions, selected_by_kind):
        versions = {
            kind: sorted([
                (row['id'], row.get('updated_at'), row.get('processed_at'), row.get('normalization_version'), row.get('row_count'))
                for row in selected
            ], key=lambda row: row[0])
            for kind, selected in selected_by_kind.items()
        }
        revisions = {kind: task_repository.get_workspace_state(f'combined_reporting_updated_{kind}') for kind in selected_by_kind}
        selection_definition = {
            # Named as the selections of before were cached, which stay valid.
            'technology': definition.nr_mode.casefold(),
            'datasets': {
                kind: sorted(set(dataset_ids))
                for kind, dataset_ids in definition.datasets.items()
            },
            'filters': {
                field: sorted(set(values))
                for field, values in definition.filters.items()
            },
            'custom_fields': sorted(set(definition.custom_fields), key=str.casefold),
            'hidden_filters': sorted(set(definition.hidden_filters), key=str.casefold),
            'date_from': definition.date_from,
            'date_to': definition.date_to,
        }
        payload = {
            # Scope changes only how charts group the already selected rows.
            # Keep that presentation setting out of the selection cache key.
            'schema': DESIGN_SELECTION_CACHE_VERSION,
            'definition': selection_definition,
            'versions': versions,
            'combined_revisions': revisions,
            'dimensions': core.calculated_dimensions_json(dimensions),
        }
        return sha256(json.dumps(payload, sort_keys=True, default=str).encode()).hexdigest()

    def selected_date_bounds(task_repository, selected_by_kind, progress=None):
        """Return the inclusive calendar bounds across the selected source datasets."""
        if callable(progress):
            progress(45, 'Checking cached Dataset Universe date limits')
        sources = {}
        for kind, selected in selected_by_kind.items():
            columns = task_repository.list_reporting_row_columns(kind)
            date_column = next((
                resolve_sql_column(columns, candidate)
                for candidate in ('event_start_time', 'Test_Start_Time', 'Timestamp', 'Date')
                if resolve_sql_column(columns, candidate)
            ), None)
            sources[kind] = {
                'datasets': sorted((
                    row['id'], row.get('updated_at'), row.get('processed_at'),
                    row.get('normalization_version'), row.get('row_count'),
                ) for row in selected),
                'revision': task_repository.get_workspace_state(f'combined_reporting_updated_{kind}'),
                'date_column': date_column,
            }
        cache_signature = sha256(json.dumps({
            'schema': DESIGN_DATE_BOUNDS_CACHE_VERSION, 'sources': sources,
        }, sort_keys=True, default=str).encode()).hexdigest()
        cache_state_key = f'design_date_bounds_v{DESIGN_DATE_BOUNDS_CACHE_VERSION}'
        try:
            cached_bounds = json.loads(task_repository.get_workspace_state(cache_state_key) or '{}')
            if not isinstance(cached_bounds, dict):
                cached_bounds = {}
            if cache_signature in cached_bounds:
                if callable(progress):
                    progress(52, 'Restoring cached Dataset Universe date limits')
                return cached_bounds[cache_signature]
        except (TypeError, ValueError, json.JSONDecodeError):
            cached_bounds = {}
        lower = upper = None
        if callable(progress):
            progress(48, 'Reading Dataset Universe date limits from the combined CDR tables')
        with task_repository.connection() as connection:
            for kind, selected in selected_by_kind.items():
                date_column = sources[kind]['date_column']
                if not date_column:
                    continue
                dataset_ids = [int(row['id']) for row in selected]
                placeholders = ', '.join('?' for _ in dataset_ids)
                table = task_repository._quote_identifier(task_repository.reporting_rows_table_name(kind))
                column = task_repository._quote_identifier(date_column)
                with cdr_store.combined_scope(dataset_ids):
                    row = connection.execute(
                        f'SELECT MIN(date({column})) AS first_date, MAX(date({column})) AS last_date '
                        f'FROM {table} WHERE dataset_id IN ({placeholders})', dataset_ids,
                    ).fetchone()
                for value, is_lower in ((row['first_date'], True), (row['last_date'], False)):
                    try:
                        parsed = date.fromisoformat(str(value)) if value else None
                    except ValueError:
                        parsed = None
                    if parsed is not None and (lower is None or parsed < lower) and is_lower:
                        lower = parsed
                    if parsed is not None and (upper is None or parsed > upper) and not is_lower:
                        upper = parsed
        bounds = None if lower is None or upper is None else {'min': lower.isoformat(), 'max': upper.isoformat()}
        # This result is immutable while the selected datasets and their
        # combined-table revisions are unchanged.  Persist a small bounded
        # map so restarting the server does not repeat a full date scan just
        # to validate a Design preview manifest.
        cached_bounds[cache_signature] = bounds
        if len(cached_bounds) > 128:
            cached_bounds = dict(list(cached_bounds.items())[-128:])
        task_repository.set_workspace_state(cache_state_key, json.dumps(cached_bounds, separators=(',', ':')))
        if callable(progress):
            progress(52, 'Caching Dataset Universe date limits')
        return bounds

    def apply_selected_date_bounds(definition, bounds):
        if not bounds:
            return
        lower, upper = date.fromisoformat(bounds['min']), date.fromisoformat(bounds['max'])
        definition.date_from = lower if definition.date_from == 'Oldest' or definition.date_from is None or not lower <= definition.date_from <= upper else definition.date_from
        definition.date_to = upper if definition.date_to == 'Newest' or definition.date_to is None or not lower <= definition.date_to <= upper else definition.date_to

    def combined_filter_options(
        definition, selected_by_kind, fields, task_repository, connection,
    ):
        """Calculate every adaptive facet with one combined-table scan per CDR kind."""
        options = {field_name: set() for field_name in fields}
        active_filter_keys = {identity(field_name) for field_name in definition.filters}
        facet_groups: dict[str | None, list[str]] = {None: []}
        for field_name in fields:
            excluded = field_name if identity(field_name) in active_filter_keys else None
            facet_groups.setdefault(excluded, []).append(field_name)
        for kind, selected in selected_by_kind.items():
            columns = task_repository.list_reporting_row_columns(kind)
            available = {
                field_name: filter_sql_value_expression(task_repository, columns, field_name)
                for field_name in fields
            }
            available = {
                field_name: expression for field_name, expression in available.items()
                if expression is not None
            }
            if not available:
                continue
            table = task_repository._quote_identifier(task_repository.reporting_rows_table_name(kind))
            base_definition = definition.model_copy(deep=True)
            base_definition.filters = {}
            base_where, base_params = selection_where(
                task_repository, kind, [int(row['id']) for row in selected],
                base_definition, columns=columns,
            )
            group_predicates = {
                excluded: selection_filter_where(task_repository, columns, definition, exclude=excluded)
                for excluded in facet_groups
            }
            aggregates, aggregate_parameters = [], []
            field_order = []
            for field_name in fields:
                expression = available.get(field_name)
                if expression is None:
                    continue
                excluded = field_name if identity(field_name) in active_filter_keys else None
                predicate, predicate_params = group_predicates[excluded]
                aggregates.append(
                    f'json_group_array(DISTINCT {expression}) '
                    f'FILTER (WHERE ({predicate})) AS facet_{len(field_order)}'
                )
                aggregate_parameters.extend(predicate_params)
                field_order.append(field_name)
            # One scan reading only the columns of the facets and filters (the rows of an emptied combined table are
            # given from the columnar copies, column by column).
            query = f'SELECT {", ".join(aggregates)} FROM {table} WHERE {base_where}'
            with cdr_store.combined_scope([int(row['id']) for row in selected]):
                row = connection.execute(query, [*aggregate_parameters, *base_params]).fetchone()
            for index, field_name in enumerate(field_order):
                encoded_values = row[f'facet_{index}'] if row else '[]'
                options[field_name].update(
                    str(value) for value in json.loads(encoded_values or '[]') if value is not None
                )
        return {field_name: sorted(values, key=str.casefold) for field_name, values in options.items()}

    def profile_filter_options(definition, dimensions, selected_by_kind, fields, task_repository, checkpoint=None):
        """Load large-design facet values from selected CDR profiles.

        ``checkpoint`` runs between fields, so a background preparation stops
        promptly when a user-requested preparation is waiting.
        """
        options = {field_name: set() for field_name in fields}
        incomplete_fields = set()
        dataset_ids = [int(dataset['id']) for selected in selected_by_kind.values() for dataset in selected]
        vendor_catalogues = task_repository.cdr_catalogues_by_dataset(dataset_ids)
        missing_vendor_catalogues = set(task_repository.missing_cdr_vendor_only_ids(dataset_ids))
        missing_cluster_catalogues = set(task_repository.missing_cdr_cluster_ids(dataset_ids))
        for selected in selected_by_kind.values():
            for dataset in selected:
                try:
                    stored = dataset.get('filter_options') or json.loads(dataset.get('filter_options_json') or '{}')
                except (TypeError, json.JSONDecodeError):
                    stored = {}
                lookup = {identity(column): values for column, values in stored.items()} if isinstance(stored, dict) else {}
                # One schema read per CDR instead of one per field and alias.
                dataset_columns = set(task_repository.list_dataset_row_columns(int(dataset['id'])))
                for field_name in fields:
                    if callable(checkpoint):
                        checkpoint()
                    if identity(field_name) in {identity('Vendor'), identity('Vendor_Only')} and int(dataset['id']) not in missing_vendor_catalogues:
                        options[field_name].update(vendor_catalogues.get(int(dataset['id']), {}).get('vendors_only', []))
                        continue
                    if identity(field_name) in {identity('Operator_Vendor'), identity('OP_Vendor')} and int(dataset['id']) in vendor_catalogues:
                        options[field_name].update(vendor_catalogues.get(int(dataset['id']), {}).get('vendors', []))
                        continue
                    if identity(field_name) == identity('Cluster') and int(dataset['id']) not in missing_cluster_catalogues:
                        options[field_name].update(vendor_catalogues.get(int(dataset['id']), {}).get('clusters', []))
                        continue
                    aliases = next((values for label, values in FILTER_COLUMNS.items() if identity(label) == identity(field_name)), (field_name,))
                    resolved = next((
                        column for alias in aliases
                        if dataset_columns and (column := task_repository._resolve_dataset_row_column_name(dataset_columns, alias))
                    ), None)
                    if resolved is None:
                        continue
                    lookup_keys = list(dict.fromkeys([
                        identity(resolved), identity(vendor_filter_column(field_name)),
                        *(identity(alias) for alias in aliases),
                    ]))
                    values = next((lookup[key] for key in lookup_keys if isinstance(lookup.get(key), list)), None)
                    if not isinstance(values, list):
                        incomplete_fields.add(field_name)
                        continue
                    if len(values) >= 50:
                        values = dataset_distinct_values(task_repository, dataset, resolved)
                    options[field_name].update(str(value) for value in values if value is not None)
        dimensions_by_name = {identity(dimension.name): dimension for dimension in dimensions}
        for field_name in fields:
            dimension = dimensions_by_name.get(identity(field_name))
            if dimension:
                incomplete_fields.discard(field_name)
                options[field_name].update(str(rule.value) for rule in dimension.rules if str(rule.value).strip())
                if str(dimension.default).strip():
                    options[field_name].add(str(dimension.default))
            # Saved filter values are not options by themselves: only values the
            # selected CDRs (or calculated dimensions) provide are offered.
        missing_fields = [
            field_name for field_name in fields
            if not options[field_name] or field_name in incomplete_fields
        ]
        if missing_fields:
            # Older upload profiles do not contain every Design facet. Read
            # only the missing catalogues from the narrow combined tables. An
            # adaptive facet must ignore its own current restriction so users
            # can see and select values outside the saved selection.
            with task_repository.connection() as connection:
                queried = combined_filter_options(
                    definition, selected_by_kind, missing_fields, task_repository, connection,
                )
            for field_name, values in queried.items():
                options[field_name].update(values)
        return {field_name: sorted(values, key=str.casefold) for field_name, values in options.items()}

    def materialize_selection(
        definition, task_repository, dimensions, selected_by_kind, fields, *,
        use_profile_options=False, known_full_row_counts=None, progress=None, checkpoint=None,
    ):
        cache_key = persistent_selection_key(definition, task_repository, dimensions, selected_by_kind)
        estimated_rows = sum(sum(int(row.get('row_count') or 0) for row in selected) for selected in selected_by_kind.values())
        use_profile_options = use_profile_options or estimated_rows > DESIGN_PROFILE_SELECTION_THRESHOLD
        if callable(progress):
            progress(55, 'Checking the filtered PPT selection cache')
        # Design preparation is already serialized by design_work_gate.
        # Do not retain the shared task-registry lock while SQLite counts or
        # facet queries run: progress polling needs that lock to report the
        # exact live phase during a potentially long combined-table scan.
        with task_repository.connection() as connection:
            cached = connection.execute(
                'SELECT id, options_json, row_counts_json, universe_row_counts_json FROM design_filter_selections WHERE cache_key = ?',
                (cache_key,),
            ).fetchone()
            if cached:
                options = json.loads(cached['options_json'])
                missing_fields = [field for field in fields if field not in options]
                if missing_fields:
                    options.update(combined_filter_options(
                        definition, selected_by_kind, missing_fields, task_repository, connection,
                    ))
                    connection.execute('UPDATE design_filter_selections SET options_json = ? WHERE id = ?',
                                       (json.dumps(options), cached['id']))
                if callable(progress):
                    progress(78, 'Restoring cached row counts and filter options')
                connection.execute('UPDATE design_filter_selections SET last_accessed_at = CURRENT_TIMESTAMP WHERE id = ?', (cached['id'],))
                return (
                    int(cached['id']), cache_key,
                    options, json.loads(cached['row_counts_json']),
                    json.loads(cached['universe_row_counts_json']), True,
                )
            known_counts = known_full_row_counts if isinstance(known_full_row_counts, dict) else None
            if callable(progress):
                progress(
                    62,
                    'Using materialized combined CDR row counts for the complete Dataset Universe'
                    if known_counts is not None else 'Counting reduced and filtered Universe rows in the combined CDR tables',
                )
            row_counts = dict(known_counts or {})
            universe_row_counts = dict(known_counts or {})
            selected_groups = list(selected_by_kind.items())
            for group_index, (kind, selected) in enumerate(selected_groups):
                if known_counts is not None:
                    continue
                if callable(progress):
                    progress(
                        62 + round(group_index * 6 / max(len(selected_groups), 1)),
                        f'Counting Filtered Universe CDR-{kind.title()} rows in the combined table',
                    )
                dataset_ids = [int(row['id']) for row in selected]
                columns = task_repository.list_reporting_row_columns(kind)
                cached_counts = cached_selection_counts(task_repository, kind, selected, definition, columns)
                if cached_counts is not None:
                    universe_row_counts[kind], row_counts[kind] = cached_counts
                    continue
                universe_definition = definition.model_copy(deep=True)
                universe_definition.filters = {}
                universe_where, universe_params = selection_where(
                    task_repository, kind, dataset_ids, universe_definition, columns=columns,
                )
                filtered_where, filtered_params = selection_filter_where(
                    task_repository, columns, definition,
                )
                table = task_repository._quote_identifier(task_repository.reporting_rows_table_name(kind))
                with cdr_store.combined_scope(dataset_ids):
                    counts = connection.execute(
                        f'''SELECT COUNT(*) AS universe_count,
                                   SUM(CASE WHEN {filtered_where} THEN 1 ELSE 0 END) AS filtered_count
                            FROM {table} WHERE {universe_where}''',
                        [*filtered_params, *universe_params],
                    ).fetchone()
                universe_row_counts[kind] = int(counts['universe_count'] or 0)
                row_counts[kind] = int(counts['filtered_count'] or 0)
            # A full, unfiltered CDR selection can use its upload-time profile
            # catalogue. This avoids wide DISTINCT scans while row counts and
            # chart previews remain exact from the combined reporting tables.
            if use_profile_options:
                if callable(progress):
                    progress(70, 'Loading filter options from Dataset profiles')
                options = profile_filter_options(
                    definition, dimensions, selected_by_kind, fields, task_repository, checkpoint=checkpoint,
                )
            else:
                if callable(progress):
                    progress(70, 'Calculating adaptive filter options in one pass per CDR table')
                options = combined_filter_options(
                    definition, selected_by_kind, fields, task_repository, connection,
                )
            # Expensive counts and facet scans above are read-only. Start the
            # coordinated Workspace write transaction only once their final
            # cache payload is ready.
            if callable(progress):
                progress(78, 'Caching filtered row counts and filter options')
            # Another process (for example a Reporting worker) may have cached
            # the same selection meanwhile; keep a single row per cache key.
            connection.execute(
                'INSERT INTO design_filter_selections '
                '(cache_key, options_json, row_counts_json, universe_row_counts_json) VALUES (?, ?, ?, ?) '
                'ON CONFLICT(cache_key) DO UPDATE SET options_json = excluded.options_json, '
                'row_counts_json = excluded.row_counts_json, universe_row_counts_json = excluded.universe_row_counts_json',
                (cache_key, json.dumps(options), json.dumps(row_counts), json.dumps(universe_row_counts)),
            )
            selection_id = int(connection.execute(
                'SELECT id FROM design_filter_selections WHERE cache_key = ?', (cache_key,)).fetchone()['id'])
            stale = connection.execute(
                'SELECT id FROM design_filter_selections ORDER BY last_accessed_at DESC, id DESC LIMIT -1 OFFSET ?',
                (DESIGN_SELECTION_CACHE_LIMIT,),
            ).fetchall()
            for row in stale:
                connection.execute('DELETE FROM design_filter_selections WHERE id = ?', (row['id'],))
            return selection_id, cache_key, options, row_counts, universe_row_counts, True

    def mapped_values(field, values, mapper):
        """The values of one filter with their Operator and Vendor Maps labels (Vendor_Operator from Operator_Vendor)."""
        kind = field_kind(vendor_filter_column(field))
        if kind == 'vendor_operator':
            return mapper.values('vendor_operator', mapper.vendor_operators(mapper.values('operator_vendor', values)))
        return mapper.values(kind, values) if kind else values

    def mapped_filter_options(options, task_repository):
        """Operator, Vendor, Operator_Vendor and Vendor_Operator values with their Operator and Vendor Maps labels."""
        mapper = ValueMapper.from_repository(task_repository)
        return {field: mapped_values(field, values, mapper) for field, values in options.items()}

    def load_filter_options_batch(definition, field_names, task_repository):
        """Load several unrestricted filter catalogues of the selected CDRs at once.

        Values come from the CDR upload profiles; only fields missing from
        them are read from the combined tables, in one pass per CDR table.
        """
        fields = list(dict.fromkeys(str(field).strip() for field in field_names if str(field).strip()))
        validate(definition, task_repository)
        dimensions = core.load_repository_calculated_dimensions(task_repository)
        selected_by_kind = selected_sources(definition, task_repository)
        requested_definition = unfiltered_definition(definition).model_copy(deep=True)
        for field in fields:
            known_default = any(identity(field) == identity(item) for item in ADAPTATIVE_FILTER_FIELDS)
            if not known_default and not any(identity(field) == identity(item) for item in requested_definition.custom_fields):
                requested_definition.custom_fields.append(field)
        ensure_combined_filter_columns(requested_definition, task_repository, dimensions, selected_by_kind)
        options = mapped_filter_options(
            profile_filter_options(requested_definition, dimensions, selected_by_kind, fields, task_repository), task_repository,
        )
        return {field: [str(value) for value in options.get(field, [])] for field in fields}

    def load_filter_options(definition, field_name, task_repository):
        """Load one adaptive filter catalogue without preparing PPT charts."""
        field_name = field_name.strip()
        validate(definition, task_repository)
        dimensions = core.load_repository_calculated_dimensions(task_repository)
        selected_by_kind = selected_sources(definition, task_repository)
        requested_definition = definition.model_copy(deep=True)
        known_default = any(identity(field_name) == identity(field) for field in ADAPTATIVE_FILTER_FIELDS)
        if not known_default and not any(identity(field_name) == identity(field) for field in requested_definition.custom_fields):
            requested_definition.custom_fields.append(field_name)
        ensure_combined_filter_columns(requested_definition, task_repository, dimensions, selected_by_kind)
        apply_selected_date_bounds(requested_definition, selected_date_bounds(task_repository, selected_by_kind))
        values = set()
        resolved = False
        with task_repository.connection() as connection:
            for kind, selected in selected_by_kind.items():
                columns = task_repository.list_reporting_row_columns(kind)
                value_expression = filter_sql_value_expression(task_repository, columns, field_name)
                if value_expression is None:
                    continue
                resolved = True
                where, params = selection_where(
                    task_repository, kind, [int(row['id']) for row in selected],
                    requested_definition, exclude=field_name, columns=columns,
                )
                table = task_repository._quote_identifier(task_repository.reporting_rows_table_name(kind))
                with cdr_store.combined_scope([int(row['id']) for row in selected]):
                    rows = connection.execute(
                        f'SELECT DISTINCT {value_expression} AS value '
                        f'FROM {table} WHERE {where}', params,
                    ).fetchall()
                values.update(str(row['value']) for row in rows)
        if not resolved:
            raise HTTPException(400, f'The selected CDRs do not contain the {field_name} field.')
        if field_kind(vendor_filter_column(field_name)):
            return mapped_values(field_name, values, ValueMapper.from_repository(task_repository))
        return sorted(values, key=str.casefold)

    def unfiltered_definition(definition):
        """Return the Design without its filters, to read the values its CDRs provide."""
        return definition.model_copy(update={'filters': {}})

    def design_geography_catalogue(definition, selected_by_kind, task_repository) -> dict[str, list[str]]:
        """Return Vendor, Region and City values of the selected CDRs.

        Values come from the per-CDR catalogue. A Region or City catalogue can
        be empty when the CDR stores that field under another alias, so the
        Design's own filter values (which resolve those aliases) complete it.
        """
        selected_ids = [int(row['id']) for rows in selected_by_kind.values() for row in rows]
        if not selected_ids:
            return {'vendors': [], 'regions': [], 'cities': []}
        core.backfill_cdr_catalogues(selected_ids, task_repository)
        catalogue = dict(task_repository.cdr_catalogue_values(selected_ids))
        cached = task_repository.cdr_catalogues_by_dataset(selected_ids)
        catalogue['vendors'] = sort_vendor_values({value for item in cached.values() for value in (item.get('vendors_only') or [])})
        missing = [(field, key) for field, key in (('Region', 'regions'), ('City', 'cities')) if not catalogue.get(key)]
        if missing:
            fallback = profile_filter_options(
                unfiltered_definition(definition), [], selected_by_kind, [field for field, _key in missing], task_repository,
            )
            for field, key in missing:
                catalogue[key] = sorted(
                    {str(value).strip() for value in fallback.get(field) or [] if str(value).strip()},
                    key=str.casefold,
                )
        return catalogue

    campaign_backfills: set[tuple[str, int]] = set()

    def schedule_campaign_backfill(dataset_ids: list[int], task_repository) -> None:
        """Read missing CDR Campaigns once in the background; the viewer polls for them."""
        database = str(Path(task_repository.db_path).resolve())
        with lock:
            pending = [dataset_id for dataset_id in dataset_ids if (database, dataset_id) not in campaign_backfills]
            campaign_backfills.update((database, dataset_id) for dataset_id in pending)
        if not pending:
            return

        def run() -> None:
            try:
                core.backfill_cdr_campaigns(pending, task_repository)
            finally:
                with lock:
                    campaign_backfills.difference_update((database, dataset_id) for dataset_id in pending)

        Thread(target=run, name='design-campaign-backfill', daemon=True).start()

    def design_campaign_values(
        definition, selected_by_kind, task_repository, *, wait: bool = True,
    ) -> tuple[list[str], bool]:
        """Return the compact Campaigns (for example ``2026-Q2`` or ``2026-Q2-SA``) behind a Design.

        A Campaign filter limits them to its values; otherwise every Campaign of
        the selected CDRs is read from their cached catalogue. CDRs catalogued
        before Campaigns were cached are read once: immediately when ``wait``
        is set, otherwise in the background, reporting that some are pending.
        """
        filters = definition.filters if isinstance(definition.filters, dict) else {}
        configured = next((values for key, values in filters.items() if identity(key) == identity('Campaign')), None)
        if isinstance(configured, str):
            configured = [configured]
        configured = [str(value).strip() for value in configured or [] if str(value).strip()]
        dataset_ids = [int(row['id']) for rows in selected_by_kind.values() for row in rows]
        pending = False
        if dataset_ids:
            core.backfill_cdr_catalogues(dataset_ids, task_repository)
            missing = task_repository.missing_cdr_campaign_ids(dataset_ids)
            if missing and wait:
                core.backfill_cdr_campaigns(missing, task_repository)
            elif missing:
                schedule_campaign_backfill(missing, task_repository)
                pending = True
        available = task_repository.cdr_catalogue_values(dataset_ids)['campaigns'] if dataset_ids else []
        # A Campaign filter only names the Campaigns its CDRs actually contain.
        values = values_present_in_cdrs(configured, available) if configured else available
        labels = {compact_campaign_value(value) for value in values}
        return sorted((label for label in labels if label), key=campaign_sort_key), pending

    def design_campaigns_label(campaigns: list[str]) -> str:
        if not campaigns:
            return ''
        return f"{'Campaign' if len(campaigns) == 1 else 'Campaigns'}: {', '.join(campaigns)}"

    def load_design_geography_options(definition, task_repository) -> dict[str, list[str]]:
        """Load Vendor and geography catalogues from the per-CDR persisted cache."""
        validate(definition, task_repository)
        selected_by_kind = selected_sources(definition, task_repository)
        options = design_geography_catalogue(definition, selected_by_kind, task_repository)
        operators = profile_filter_options(
            unfiltered_definition(definition), [], selected_by_kind, ['Operator'], task_repository,
        )['Operator']
        campaigns, campaigns_pending = design_campaign_values(
            definition, selected_by_kind, task_repository, wait=False,
        )
        mapper = ValueMapper.from_repository(task_repository)
        return {
            'campaigns': campaigns,
            'campaigns_pending': campaigns_pending,
            'operators': mapper.values('operator', operators),
            'vendors': mapper.values('vendor', options['vendors']),
            'regions': options['regions'],
            'cities': options['cities'],
        }

    def expand_repeated_elements(entries, definition, options, deck):
        """Repeat by (Designer): each chart repeated by a field becomes one chart per value of the field in the selection
        (its filter values first), each in its share of the chart's place; the other charts of its slide keep theirs. The
        places are written on the rows (``rect``), so the viewer and the PPT export lay them out alike."""
        if not any(repeat_field(entry) for entry in entries):
            return entries
        deck = deck or Presentation(core.settings.powerpoint_template)
        by_slide: dict[int, list[int]] = defaultdict(list)
        for index, entry in enumerate(entries):
            by_slide[entry.slide].append(index)
        result: list = []
        for _number, indexes in sorted(by_slide.items()):
            rows = [entries[index] for index in indexes]
            charts = [row for row in rows if not row.structural_type]
            if not any(repeat_field(row) for row in charts) or rows[0].layout in DYNAMIC_LAYOUTS:
                result.extend(rows)
                continue
            positions = slide_geometry(charts) or normalized_frame_positions(_layout_chart_frames(_named_slide_layout(deck, rows[0].layout)))
            places = explicit_places(charts, len(positions)) or list(range(len(charts)))
            for row, place in zip(charts, places):
                rect = positions[place] if place < len(positions) else None
                data = design_slides.row_format(row)
                field_name = repeat_field(row)
                if rect is None:
                    result.append(row)
                    continue
                if not field_name:
                    result.append(replace(row, format=design_slides.format_cell({**data, 'rect': rect})))
                    continue
                values = next((items for key, items in options.items() if identity(key) == identity(field_name)), [])
                chosen = next((items for key, items in definition.filters.items() if identity(key) == identity(field_name)), [])
                values = [str(value) for value in values if str(value).strip() and (not chosen or str(value) in {str(item) for item in chosen})]
                values = sorted(dict.fromkeys(values), key=str.casefold)[:12]
                if not values:
                    result.append(replace(row, format=design_slides.format_cell({**data, 'rect': rect})))
                    continue
                for value, part in zip(values, subdivided(rect, len(values))):
                    filters = design_slides.join_conditions(design_slides.filter_conditions(row.filters) + [f'{field_name} IN ({value})'])
                    result.append(replace(row, filters=filters,
                                          chart_title=f'{row.chart_title} – {value}' if row.chart_title else value,
                                          format=design_slides.format_cell({**data, 'rect': part, 'repeat_value': value})))
        return result

    def laid_out_slides(entries, definition, options, task_repository, selected_by_kind, deck):
        """The rows of a Design with its dynamic layouts expanded, and its slides as the viewer shows them: their
        charts, positions on the layout and whether each chart has CDRs to read."""
        slides = OrderedDict()
        entries = expand_design_layouts(entries, definition, options, task_repository)
        entries = expand_repeated_elements(entries, definition, options, deck)
        slide_entries_by_number: dict[int, list] = defaultdict(list)
        for entry in entries:
            slide_entries_by_number[entry.slide].append(entry)
        fields_by_mode: dict[str, dict[str, set[str]]] = {}
        for editor_index, (index, entry) in enumerate(sorted(enumerate(entries), key=lambda item: (item[1].slide, item[0]))):
            editor_index = entry.template_index if entry.template_index is not None else editor_index
            row = design_slides.row_format(entry)
            slide = slides.setdefault(entry.slide, {
                'number': entry.slide,
                'comment_key': design_slide_comment_key(slide_entries_by_number[entry.slide]),
                'title': entry.slide_title,
                'subtitle': entry.slide_subtitle,
                'layout': entry.layout,
                'structural_type': entry.structural_type or '',
                'compact_titles': compact_grid_layout(entry.layout),
                'charts': [],
                'focus_row': editor_index,
                # The slide of the slide set it shows (a dynamic grid shows one slide of the slide set in several pages).
                'source_slide': str(row.get('slide') or ''),
            })
            if not entry.structural_type:
                title_label, title_base = chart_title_parts(entry)
                chart = {
                    'index': index, 'title': entry.chart_title, 'title_label': title_label, 'title_base': title_base,
                    'source': entry.source_kind,
                    'cdr_source': entry.cdr_source, 'chart_type': entry.chart_type,
                    'available': bool(selected_by_kind) if entry.source_kind == 'all' else entry.source_kind in selected_by_kind, 'focus_row': editor_index,
                    'element': str(row.get('id') or ''), 'place': row.get('place'),
                }
                kind = next((key for key, chart_type in design_slides.KIND_CHART_TYPES.items()
                             if chart_type.casefold() == entry.chart_type.casefold()), 'chart')
                if kind != 'chart':
                    chart['kind'] = kind
                if entry.chart_type.casefold() == design_slides.DRAFT_CHART_TYPE.casefold():
                    # A chart still being built reads no CDR: the viewer shows its empty place.
                    chart.update(kind='draft', available=False, content={})
                elif kind in {'text', 'image'}:
                    # Read no CDR: the viewer shows the text or the picture (their content), never a chart model.
                    content = dict(row.get('content') or {}) if isinstance(row.get('content'), dict) else {}
                    if kind == 'image' and content.get('image'):
                        content['url'] = f"/api/ppt-designer/images/{content['image']}"
                    chart.update(available=False, content=content)
                elif unread := chart_nr_mode_unread(entry, definition):
                    # A chart of one NR Mode (NR Mode IN (SA)) in a Design that reads the other one.
                    chart.update(available=False, unavailable=unread)
                elif design_nr_mode(definition) == NSA_VS_SA and entry.source_kind in selected_by_kind:
                    # A field only the CDRs of one NR Mode have: the chart says which NR Mode has no samples.
                    if entry.source_kind not in fields_by_mode:
                        fields_by_mode[entry.source_kind] = core.cdr_fields_by_nr_mode(
                            task_repository, entry.source_kind, [int(row['id']) for row in selected_by_kind[entry.source_kind]])
                    if missing := missing_nr_mode_fields(entry, fields_by_mode[entry.source_kind], NSA_VS_SA):
                        chart['nr_mode_note'] = missing_nr_mode_note(missing)
                slide['charts'].append(chart)
        for slide in slides.values():
            charts = slide['charts']
            for chart in charts:
                chart.pop('place', None)
            chart_rows = [entries[chart['index']] for chart in charts]
            # The places of a free layout (Designer), else those of the layout; each element on its place (a place can
            # stay empty), else the charts fill the places in order (the cards of a dynamic grid too).
            rects = [design_slides.row_format(row).get('rect') for row in chart_rows]
            if chart_rows and all(isinstance(rect, list) and len(rect) == 4 for rect in rects):
                # Each chart has its own place (a chart repeated in its place: Repeat by).
                for chart, rect in zip(charts, rects):
                    chart['position'] = rect
                continue
            positions = slide_geometry(chart_rows) or normalized_frame_positions(_layout_chart_frames(_named_slide_layout(deck, slide['layout'])))
            places = explicit_places(chart_rows, len(positions))
            if places is not None:
                for chart, place in zip(charts, places):
                    chart['position'] = positions[place]
            elif positions and len(positions) >= len(charts):
                for chart, position in zip(charts, positions):
                    chart['position'] = position
        return entries, list(slides.values())

    def build_preview(definition, user, *, workspace: str | None = None, cancelled=None, progress=None):
        def ensure_not_cancelled():
            if callable(cancelled) and cancelled():
                raise RuntimeError('PPT preparation cancelled.')

        def report(percent, detail):
            # Every progress step is also a cancellation point.
            ensure_not_cancelled()
            if callable(progress):
                progress(percent, detail)

        with cdr_stage.remembered_inclusion():
            return _build_preview(definition, user, workspace=workspace, cancelled=cancelled, report=report,
                                  ensure_not_cancelled=ensure_not_cancelled)

    def _build_preview(definition, user, *, workspace, cancelled, report, ensure_not_cancelled):
        workspace = workspace or workspace_key()
        report(5, 'Validating the Design template and selected source datasets')
        ensure_not_cancelled()
        task_repository = Repository(Path(workspace), core.repository.global_db_path)
        entries = validate(definition, task_repository)
        report(6, 'Reading the Auto-calculated Fields and the selected CDRs')
        dimensions = core.load_repository_calculated_dimensions(task_repository)
        selected_by_kind = selected_sources(definition, task_repository)
        if any(entry.source_kind == 'all' for entry in entries):
            report(8, 'Adding the RF fields of the selected CDRs to the combined CDR tables')
            # A virtual source needs the physical RF fields of every selected
            # CDR type, including templates added after CDR ingestion.
            for kind, datasets in selected_by_kind.items():
                for dataset in datasets:
                    available = task_repository.list_dataset_row_columns(int(dataset['id']))
                    requested = [column for column in available if identity(column) in {identity(field) for field in rf_source_columns(kind)}]
                    task_repository.copy_dataset_rows_to_reporting(int(dataset['id']), kind, requested)
        if any(any(entry_dynamic_fields(entry)) or entry_uses_nr_mode(entry) for entry in entries):
            report(10, 'Checking the fields of the charts in the combined CDR tables')
            ensure_combined_filter_columns(definition, task_repository, dimensions, selected_by_kind, chart_entries=entries)
        report(12, 'Checking the Design fields in the combined CDR tables')
        selected_source_rows = {
            kind: sum(int(dataset.get('row_count') or 0) for dataset in selected)
            for kind, selected in selected_by_kind.items()
        }
        if design_combined_columns_missing(definition, task_repository, dimensions, selected_by_kind):
            report(20, 'Adding new Design-specific fields to the combined CDR tables')
            ensure_combined_filter_columns(definition, task_repository, dimensions, selected_by_kind)
        ensure_not_cancelled()
        date_bounds = selected_date_bounds(task_repository, selected_by_kind, progress=report)
        apply_selected_date_bounds(definition, date_bounds)
        selected_dataset_ids = {dataset_id for ids in definition.datasets.values() for dataset_id in ids}
        available_fields = {column for dataset_id in selected_dataset_ids for column in task_repository.list_dataset_row_columns(dataset_id)}
        available_fields.update(dimension.name for dimension in dimensions)
        hidden_filter_keys = {identity(field) for field in definition.hidden_filters}
        fields = {field for field in ADAPTATIVE_FILTER_FIELDS if identity(field) not in hidden_filter_keys}
        fields.update(definition.custom_fields)
        fields.update(field for entry in entries for field in entry_dynamic_fields(prepare_multivendor_catalog_entry(entry, definition.vendor_comparison) if definition.scope == "multivendor" else entry) if field)
        fields.update(field for entry in entries if (field := repeat_field(entry)))
        fields = sorted(fields, key=str.casefold)
        use_profile_options = bool(
            date_bounds
            and not definition.filters
            and str(definition.date_from) == date_bounds['min']
            and str(definition.date_to) == date_bounds['max']
        )
        complete_unfiltered_universe = is_complete_unfiltered_universe(
            definition, selected_by_kind, task_repository, date_bounds,
        )
        selection_id, selection_key, options, row_counts, universe_rows, rows_exact = materialize_selection(
            definition, task_repository, dimensions, selected_by_kind, fields,
            use_profile_options=use_profile_options,
            known_full_row_counts=selected_source_rows if complete_unfiltered_universe else None,
            progress=report, checkpoint=ensure_not_cancelled,
        )
        report(82, 'Preparing the Design slide structure and chart positions')
        deck = Presentation(core.settings.powerpoint_template)
        entries, slides = laid_out_slides(entries, definition, options, task_repository, selected_by_kind, deck)
        token = uuid4().hex
        payload = {
            'slides': slides,
            'options': cascading_design_options(mapped_filter_options(options, task_repository), definition, task_repository),
            # The CDRs the Design uses, chosen from its NR Mode and Campaign.
            'cdrs': describe_cdrs(task_repository, {kind: [int(row['id']) for row in rows] for kind, rows in selected_by_kind.items()}),
            'filter_fields': list(ADAPTATIVE_FILTER_FIELDS),
            'available_fields': sorted(available_fields, key=str.casefold),
            'rows': row_counts,
            'universe_rows': universe_rows,
            'rows_exact': rows_exact,
            'date_bounds': date_bounds,
            # The selection behind the slides: generated comments written for another one are out of date.
            'selection_key': selection_key,
        }
        add_dynamic_grid_metadata(payload, entries, deck)
        with lock:
            snapshots[token] = Snapshot(
                workspace, user.username, entries, {}, definition.scope == 'multivendor',
                definition.model_copy(deep=True), tuple(dimensions), selection_id, selection_key, payload,
                cancelled=cancelled, options=options,
            )
            # Keep recent explicitly prepared snapshots lightweight enough for
            # users to switch between them without rebuilding their filters.
            while len(snapshots) > 128:
                snapshots.popitem(last=False)
        ensure_not_cancelled()
        report(90, 'PPT selection and slide structure are ready')
        return {**payload, 'token': token}

    def preview_manifest_key(workspace: str, design_id: str, fingerprint: str) -> str:
        return f'{workspace}:{ppt_design_cache.preparation_key(DESIGN_PREVIEW_MANIFEST_VERSION, design_id, fingerprint)}'

    # The preparations of the Designs (preview manifests) and the models of their charts are kept in memory and in the
    # cache folder of the workspace (ppt_design_cache), so a restart restores them instead of working them out again.
    # Their keys change with what they were made from (the CDRs and their revisions, the selection, the template and
    # format, the mappings, the renderer).
    def stored_preview_manifest(workspace: str, key: str, design_id: str) -> dict | None:
        """The preparation kept under ``key``, in memory or else in the cache folder (then kept in memory too)."""
        with lock:
            manifest = preview_manifests.get(key)
        if manifest is not None:
            return manifest
        manifest = ppt_design_cache.read(ppt_design_cache.preparation_path(workspace, design_id, key.rsplit(':', 1)[-1]))
        if not isinstance(manifest, dict):
            return None
        with lock:
            preview_manifests[key] = manifest
            while len(preview_manifests) > DESIGN_PREVIEW_MANIFESTS_KEPT:
                preview_manifests.popitem(last=False)
        return manifest

    def stored_chart_model(workspace: str, model_key: str) -> dict | None:
        payload = ppt_design_cache.read(ppt_design_cache.chart_path(workspace, model_key.rsplit(':', 1)[-1]))
        if not isinstance(payload, dict):
            return None
        with lock:
            canvas_models[model_key] = payload
            while len(canvas_models) > DESIGN_CHART_MODELS_KEPT:
                canvas_models.popitem(last=False)
        return payload

    def store_chart_model(workspace: str, model_key: str, payload: dict) -> None:
        ppt_design_cache.write(ppt_design_cache.chart_path(workspace, model_key.rsplit(':', 1)[-1]), payload)

    def forget_stored_chart_model(workspace: str, model_key: str) -> None:
        ppt_design_cache.forget(ppt_design_cache.chart_path(workspace, model_key.rsplit(':', 1)[-1]))

    def clear_workspace_design_caches(workspace: str | Path) -> Path | None:
        """Clear workspace cache: the preparations and chart models of its Designs, in memory and in its cache folder
        (set aside at once; the folder returned is deleted by the caller)."""
        database_path = str(Path(workspace).resolve())
        with lock:
            for key in [key for key in canvas_models if key.startswith(f'{database_path}:')]:
                canvas_models.pop(key, None)
            preview_manifests.clear()
        return ppt_design_cache.clear(database_path)

    def move_stored_caches(workspace: str | Path) -> dict[str, int]:
        """Move the preparations and chart models earlier versions kept in the workspace database to its cache folder
        (the preparations of an older format are left out)."""
        def preparation_name(manifest: dict) -> str | None:
            if manifest.get('version') != DESIGN_PREVIEW_MANIFEST_VERSION or not manifest.get('design_id'):
                return None
            return ppt_design_cache.preparation_key(
                DESIGN_PREVIEW_MANIFEST_VERSION, str(manifest['design_id']), str(manifest.get('fingerprint') or ''))
        return ppt_design_cache.move_database_rows(workspace, preparation_name)

    def persist_preview_manifest(
        workspace: str,
        design_id: str,
        fingerprint: str,
        token: str,
    ) -> None:
        with lock:
            snapshot = snapshots.get(token)
            if snapshot is None:
                return
            manifest = {
                'version': DESIGN_PREVIEW_MANIFEST_VERSION,
                'design_id': design_id,
                'fingerprint': fingerprint,
                'definition': snapshot.definition.model_dump(mode='json'),
                'selection_id': snapshot.selection_id,
                'selection_key': snapshot.selection_key,
                'payload': snapshot.payload,
            }
        key = preview_manifest_key(workspace, design_id, fingerprint)
        with lock:
            preview_manifests[key] = manifest
            while len(preview_manifests) > DESIGN_PREVIEW_MANIFESTS_KEPT:
                preview_manifests.popitem(last=False)
        ppt_design_cache.write(ppt_design_cache.preparation_path(workspace, design_id, key.rsplit(':', 1)[-1]), manifest)

    def restore_preview_manifest(
        workspace: str,
        design_id: str,
        fingerprint: str,
        *,
        materialize: bool = True,
        manifest: dict | None = None,
    ) -> dict | None:
        try:
            if manifest is None:
                manifest = stored_preview_manifest(workspace, preview_manifest_key(workspace, design_id, fingerprint), design_id)
                if manifest is None:
                    return None
            if (
                manifest.get('version') != DESIGN_PREVIEW_MANIFEST_VERSION
                or manifest.get('design_id') != design_id
                or manifest.get('fingerprint') != fingerprint
            ):
                return None
            definition = DesignDefinition.model_validate(manifest['definition'])
            task_repository = Repository(Path(workspace), core.repository.global_db_path)
            entries = validate(definition, task_repository)
            dimensions = core.load_repository_calculated_dimensions(task_repository)
            selected_by_kind = selected_sources(definition, task_repository)
            selection_key = persistent_selection_key(definition, task_repository, dimensions, selected_by_kind)
            if selection_key != manifest.get('selection_key'):
                return None
            selection_id = int(manifest['selection_id'])
            with task_repository.connection() as connection:
                selection = connection.execute(
                    'SELECT id FROM design_filter_selections WHERE id = ? AND cache_key = ?',
                    (selection_id, selection_key),
                ).fetchone()
            if selection is None:
                return None
            payload = manifest['payload']
            entries = expand_design_layouts(entries, definition, payload.get('options', {}), task_repository)
            entries = expand_repeated_elements(entries, definition, payload.get('options', {}), None)
            if not isinstance(payload, dict) or not isinstance(payload.get('slides'), list):
                return None
            if 'grid_layouts' not in payload:
                # Universes prepared by earlier versions gain the dynamic grid
                # metadata without being prepared again.
                add_dynamic_grid_metadata(payload, entries)
        except (KeyError, OSError, TypeError, ValueError, json.JSONDecodeError, sqlite3.Error):
            return None
        if not materialize:
            return dict(payload)
        definition_dump = definition.model_dump(mode='json')
        with lock:
            # Reopening the same cached universe reuses its live snapshot (and
            # any chart frames it already holds) instead of adding a copy.
            for existing_token, existing in reversed(snapshots.items()):
                if (
                    existing.workspace == workspace and existing.owner == '*'
                    and existing.selection_key == selection_key
                    and existing.selection_id == selection_id
                    and existing.definition.model_dump(mode='json') == definition_dump
                ):
                    snapshots.move_to_end(existing_token)
                    return {**existing.payload, 'token': existing_token}
        token = uuid4().hex
        with lock:
            snapshots[token] = Snapshot(
                workspace, '*', entries, {}, definition.scope == 'multivendor',
                definition, tuple(dimensions), selection_id, selection_key,
                payload,
            )
            while len(snapshots) > 128:
                snapshots.popitem(last=False)
        return {**payload, 'token': token}

    def preview_cache_identity(definition):
        """Return the parts of a PPT definition that affect its prepared data."""
        identity_definition = definition.model_dump(mode='json')
        identity_definition.pop('name', None)
        identity_definition.pop('slide_comments', None)
        identity_definition.pop('slide_comment_meta', None)
        identity_definition['datasets'] = {
            kind: sorted(set(dataset_ids))
            for kind, dataset_ids in identity_definition.get('datasets', {}).items()
            if dataset_ids
        }
        return identity_definition

    def restore_matching_preview_manifest(
        workspace: str, design_id: str, definition: DesignDefinition, *, materialize: bool = True,
    ):
        """Restore a persistent preview by definition, including non-default session universes."""
        task_repository = Repository(Path(workspace), core.repository.global_db_path)
        requested = definition.model_copy(deep=True)
        try:
            apply_selected_date_bounds(requested, selected_date_bounds(
                task_repository, selected_sources(requested, task_repository),
            ))
            requested_fingerprint = design_preview_fingerprint(
                requested.model_dump(mode='json'), task_repository,
            )
            requested_identity = preview_cache_identity(requested)
            # Manifests are kept under a key derived from their fingerprint, so the matching one is found directly.
            manifest = stored_preview_manifest(workspace, preview_manifest_key(workspace, design_id, requested_fingerprint), design_id)
            if manifest is None:
                return None
            stored_definition = DesignDefinition.model_validate(manifest['definition'])
            if preview_cache_identity(stored_definition) != requested_identity:
                return None
            return restore_preview_manifest(
                workspace, design_id, requested_fingerprint, materialize=materialize, manifest=manifest,
            )
        except (KeyError, OSError, TypeError, ValueError, json.JSONDecodeError, sqlite3.Error):
            return None

    def run_direct_preparation(preparation_id, definition, design_id, workspace, user, cancellation):
        """Prepare one requested Design universe and return its preview payload."""
        task_repository = Repository(Path(workspace), core.repository.global_db_path)

        def preparation_cancelled():
            return bool(cancellation['requested'])

        # How long each phase of the preparation takes, recorded in App Events with the whole duration.
        phase_times = [('Waiting for an available PPT preparation slot', monotonic())]

        def update_preparation_progress(percent, detail):
            if detail and detail != phase_times[-1][0]:
                phase_times.append((detail, monotonic()))
            with lock:
                task = direct_preparation_tasks.get(preparation_id)
                if task is not None:
                    task.update(progress=percent, detail=detail)

        gate_acquired = False
        while not gate_acquired:
            if cancellation['requested']:
                raise RuntimeError('PPT preparation cancelled.')
            gate_acquired = design_work_gate.acquire(timeout=0.2)
        try:
            with lock:
                task = direct_preparation_tasks.get(preparation_id)
                if task is not None:
                    if cancellation['requested']:
                        raise RuntimeError('PPT preparation cancelled.')
                    task.update(
                        status='processing',
                        started_at=datetime.now(timezone.utc).timestamp(),
                        progress=1,
                        detail='Starting the PPT preparation worker',
                    )
            preview = build_preview(
                definition, user, workspace=workspace, cancelled=preparation_cancelled,
                progress=update_preparation_progress,
            )
        finally:
            design_work_gate.release()
        if design_id:
            update_preparation_progress(94, 'Saving the reusable PPT preparation cache')
            with lock:
                prepared_snapshot = snapshots.get(preview['token'])
            if prepared_snapshot is None:
                raise RuntimeError('The prepared PPT snapshot has expired.')
            persist_preview_manifest(
                workspace,
                design_id,
                design_preview_fingerprint(
                    prepared_snapshot.definition.model_dump(mode='json'), task_repository,
                ),
                preview['token'],
            )
        update_preparation_progress(100, 'PPT dataset and filters are ready')
        record_preparation_times(task_repository, user, design_id, definition, phase_times)
        return preview

    def record_preparation_times(task_repository, user, design_id, definition, phase_times) -> None:
        """App Event of a finished preparation: how long it took and its phases of 0.1 s or more, slowest first."""
        try:
            ended = monotonic()
            spans = [(detail, (phase_times[index + 1][1] if index + 1 < len(phase_times) else ended) - started)
                     for index, (detail, started) in enumerate(phase_times)]
            slow = sorted((span for span in spans if span[1] >= 0.1), key=lambda span: -span[1])
            task_repository.add_log(getattr(user, 'username', 'system'), 'design_prepared', json.dumps({
                'id': design_id, 'name': definition.name, 'seconds': round(ended - phase_times[0][1], 2),
                'phases': '; '.join(f'{detail} {seconds:.1f} s' for detail, seconds in slow),
            }))
        except (sqlite3.Error, OSError):
            pass  # The preparation stands; only its App Event is missing.

    def remember_preparation_result(preparation_id, workspace, username, result) -> None:
        with lock:
            direct_preparation_results[preparation_id] = {
                **result, 'workspace': workspace, 'owner': username,
            }
            while len(direct_preparation_results) > 256:
                direct_preparation_results.popitem(last=False)

    @app.post('/api/ppt-designer/prepare')
    def prepare(
        definition: DesignDefinition,
        design_id: str | None = None,
        rendering_only: bool = False,
        preparation_id: str | None = None,
        use_scope_universe: bool = False,
        background: bool = False,
        user=Depends(design_user),
    ):
        """Prepare a Design universe.

        With ``background=1`` the request returns immediately (202) and the
        browser polls ``/preparation-progress/{id}``. Long preparations on slow
        servers therefore never exceed reverse-proxy response timeouts.
        """
        definition = (automatic_scope_universe if use_scope_universe else auto_cdr_definition)(definition, bound_repository())
        if design_id:
            # The slide set of a saved Design is the one stored (the Designer saves it): a viewer opened before does
            # not send it back older.
            stored = read_designs_and_previews(bound_repository()).get(design_id)
            if isinstance(stored, dict):
                definition = definition.model_copy(update={
                    'slide_overrides': copy.deepcopy(stored.get('slide_overrides') or {}),
                    'own_slides': copy.deepcopy(stored.get('own_slides')),
                })
        preparation_id = preparation_id or f'design-preparation:{uuid4().hex}'
        workspace = workspace_key()
        cancellation = {'requested': False}
        with lock:
            # A newer request from the same user supersedes that user's
            # earlier preparation; other users' preparations are unaffected.
            for task_id, task in direct_preparation_tasks.items():
                if task_id == preparation_id or task.get('workspace') != workspace or task.get('owner') != user.username:
                    continue
                existing_cancellation = task.get('cancellation')
                if isinstance(existing_cancellation, dict):
                    existing_cancellation['requested'] = True
            direct_preparation_results.pop(preparation_id, None)
            direct_preparation_tasks[preparation_id] = {
                'id': preparation_id, 'workspace': workspace, 'name': definition.name,
                'design_id': design_id, 'owner': user.username,
                'rendering_only': rendering_only, 'cancellation': cancellation,
                'status': 'queued', 'queued_at': datetime.now(timezone.utc).timestamp(),
                'started_at': None,
                'progress': 0, 'detail': 'Waiting for an available PPT preparation slot',
            }

        def execute():
            try:
                return run_direct_preparation(preparation_id, definition, design_id, workspace, user, cancellation)
            except HTTPException:
                raise
            except (ValueError, KeyError) as exc:
                raise HTTPException(400, str(exc)) from exc
            except RuntimeError as exc:
                if cancellation['requested']:
                    raise HTTPException(409, 'PPT preparation was interrupted.') from exc
                raise
            finally:
                with lock:
                    direct_preparation_tasks.pop(preparation_id, None)

        if not background:
            return execute()

        def execute_in_background():
            try:
                preview = execute()
                result = {'status': 'completed', 'token': str(preview['token'])}
            except HTTPException as exc:
                result = {'status': 'failed', 'code': exc.status_code, 'detail': str(exc.detail)}
            except Exception as exc:
                result = {'status': 'failed', 'code': 500, 'detail': f'PPT preparation failed: {exc}'}
            remember_preparation_result(preparation_id, workspace, user.username, result)

        Thread(target=execute_in_background, name='design-preparation', daemon=True).start()
        return JSONResponse({'preparation_id': preparation_id, 'status': 'queued'}, status_code=202)

    @app.post('/api/ppt-designer/filter-options')
    def filter_options(request: DesignFilterOptionsRequest, user=Depends(design_user)):
        try:
            task_repository = bound_repository()
            definition = auto_cdr_definition(request.definition, task_repository)
            values = load_filter_options(definition, request.field, task_repository)
            if request.field.strip() in CASCADING_FILTER_FIELDS:
                values = cascading_design_options({request.field.strip(): values}, definition, task_repository)[request.field.strip()]
            return {'field': request.field.strip(), 'values': values}
        except (ValueError, KeyError) as exc:
            raise HTTPException(400, str(exc)) from exc

    @app.post('/api/ppt-designer/filter-options/batch')
    def filter_options_batch(request: DesignFilterOptionsBatchRequest, user=Depends(design_user)):
        """Return the values the selected CDRs offer for several filters in one request."""
        try:
            task_repository = bound_repository()
            definition = auto_cdr_definition(request.definition, task_repository)
            options = load_filter_options_batch(definition, request.fields, task_repository)
            return {'options': cascading_design_options(options, definition, task_repository)}
        except (ValueError, KeyError) as exc:
            raise HTTPException(400, str(exc)) from exc

    @app.post('/api/ppt-designer/auto-cdrs')
    def design_auto_cdrs(definition: DesignDefinition, user=Depends(design_user)):
        """The CDRs a Design selection uses now and the lists of its Campaign, Operator, Vendor, Region, Cluster and City.

        Chosen from its NR Mode and Campaign and read from the CDR Filter Facets, without preparing the Design
        or reading CDR rows, so the CDRs and these filters follow every change before it is applied.
        """
        try:
            task_repository = bound_repository()
            resolved = auto_cdr_definition(definition, task_repository)
            options = cascading_design_options({field: [] for field in CASCADING_FILTER_FIELDS}, resolved, task_repository)
            return {'options': options, 'cdrs': describe_cdrs(task_repository, resolved.datasets)}
        except (ValueError, KeyError) as exc:
            raise HTTPException(400, str(exc)) from exc

    @app.post('/api/ppt-designer/geography-options')
    def geography_options(definition: DesignDefinition, user=Depends(design_user)):
        try:
            task_repository = bound_repository()
            # The values of the CDRs of the chosen campaigns (the campaigns and those CDRs come from /auto-cdrs).
            return load_design_geography_options(auto_cdr_definition(definition, task_repository), task_repository)
        except (ValueError, KeyError) as exc:
            raise HTTPException(400, str(exc)) from exc

    @app.get('/api/ppt-designer/prepared/{token}')
    def prepared_preview(token: str, user=Depends(design_user)):
        with lock:
            snapshot = snapshots.get(token)
            if snapshot is None or snapshot.workspace != workspace_key() or snapshot.owner not in {user.username, '*'}:
                raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
            return {**snapshot.payload, 'token': token}

    @app.get('/api/ppt-designer/preparation-progress/{preparation_id}')
    def preparation_progress(preparation_id: str, user=Depends(design_user)):
        workspace = workspace_key()
        with lock:
            task = direct_preparation_tasks.get(preparation_id)
            if task is None:
                result = direct_preparation_results.get(preparation_id)
                if result is not None and result.get('workspace') == workspace and result.get('owner') == user.username:
                    # A finished background preparation reports its outcome;
                    # the browser loads a completed preview with /prepared.
                    return {
                        key: value for key, value in result.items() if key not in {'workspace', 'owner'}
                    } | {'progress': 100 if result.get('status') == 'completed' else 0}
            if task is None or task.get('workspace') != workspace:
                raise HTTPException(404, 'PPT preparation is no longer active.')
            return {
                'progress': task.get('progress', 0),
                'detail': task.get('detail') or 'Preparing PPT dataset',
                'status': task.get('status', 'processing'),
            }

    def has_preview_manifest(design_id: str) -> bool:
        """Whether a preparation of the Design is kept, in memory or in the cache folder of the workspace."""
        with lock:
            if any(manifest.get('design_id') == design_id for manifest in preview_manifests.values()):
                return True
        return ppt_design_cache.has_preparation(workspace_key(), design_id)

    def prefetched_design_response(design_id: str, definition: DesignDefinition, user, started=None, definition_seconds=0.0,
                                   found_seconds=0.0, queued_seconds=None):
        workspace = workspace_key()
        started = monotonic() if started is None else started
        restoring = monotonic()
        restored = restore_matching_preview_manifest(workspace, design_id, definition)
        # How long opening a Design waits for this lookup (Checking persistent PPT cache), in App Events: how long the
        # request waited for this handler, finding the Design and its kept preparation, working out its CDRs, then
        # restoring the preparation.
        try:
            bound_repository().add_log(user.username, 'design_cache_checked', json.dumps({
                'id': design_id, 'name': definition.name, 'restored': restored is not None,
                **({} if queued_seconds is None else {'queued_seconds': round(queued_seconds, 2)}),
                'seconds': round(monotonic() - started, 2), 'found_seconds': round(found_seconds, 2),
                'definition_seconds': round(definition_seconds, 2), 'restore_seconds': round(monotonic() - restoring, 2),
            }))
        except (sqlite3.Error, OSError):
            pass  # The lookup stands; only its App Event is missing.
        if restored is not None:
            return restored
        raise HTTPException(409, 'PPT preparation is required.')

    @app.get('/api/ppt-designer/prefetched/{design_id}')
    def prefetched_design(design_id: str, user=Depends(design_user)):
        task_repository = bound_repository()
        with lock:
            raw_definition = read_designs_and_previews(task_repository).get(design_id)
            if not isinstance(raw_definition, dict):
                raise HTTPException(404, 'Design not found.')
        if not has_preview_manifest(design_id):
            raise HTTPException(409, 'PPT preparation is required.')
        definition = DesignDefinition.model_validate(runtime_design_definition(raw_definition, task_repository))
        return prefetched_design_response(design_id, definition, user)

    @app.post('/api/ppt-designer/prefetched/{design_id}')
    def prefetched_design_for_definition(
        design_id: str,
        definition: DesignDefinition,
        request: Request,
        use_scope_universe: bool = False,
        user=Depends(design_user),
    ):
        started = monotonic()
        received = request.scope.get('received_at')
        task_repository = bound_repository()
        if not design_exists(task_repository, design_id):
            raise HTTPException(404, 'Design not found.')
        # With no preparation of the Design kept there is nothing to look up: it is prepared at once.
        if not has_preview_manifest(design_id):
            raise HTTPException(409, 'PPT preparation is required.')
        checked = monotonic()
        with cdr_stage.remembered_inclusion():
            definition = (automatic_scope_universe if use_scope_universe else auto_cdr_definition)(definition, task_repository)
            return prefetched_design_response(design_id, definition, user, started, monotonic() - checked,
                                              found_seconds=checked - started,
                                              queued_seconds=None if received is None else started - received)

    @app.post('/api/ppt-designer/opened/{design_id}')
    def design_opened(design_id: str, timeline: dict[str, Any] = Body(default_factory=dict), user=Depends(design_user)):
        """App Event of a Design opened in the viewer: how long the browser waited for the cache lookup, for the
        prepared Design and for the charts of the slide it opened on."""
        task_repository = bound_repository()
        with lock:
            raw_definition = read_designs_and_previews(task_repository).get(design_id)
        if not isinstance(raw_definition, dict):
            raise HTTPException(404, 'Design not found.')
        numbers = {key: round(float(timeline[key]), 2) for key in ('before_lookup_seconds', 'lookup_seconds', 'lookup_browser_wait_seconds',
                                                                   'lookup_server_seconds', 'prepared_seconds', 'slide_seconds')
                   if isinstance(timeline.get(key), (int, float))}
        task_repository.add_log(user.username, 'design_opened', json.dumps({
            'id': design_id, 'name': raw_definition.get('name', ''), **numbers,
            'restored': bool(timeline.get('restored')), 'slide': int(timeline.get('slide') or 0),
            'charts': int(timeline.get('charts') or 0),
        }))
        return {'recorded': True}

    def reporting_source(snapshot, kind, task_repository):
        """Return the shared combined CDR table used directly by Design queries."""
        selected = design_datasets(snapshot.definition, kind, task_repository)
        if not selected:
            raise HTTPException(400, 'Unavailable source type: select a matching CDR dataset.')
        if kind == 'all':
            table, columns = rf_union_source(task_repository, [source for source in KINDS if design_datasets(snapshot.definition, source, task_repository)])
            return Path(snapshot.workspace), table, columns
        columns = task_repository.list_reporting_row_columns(kind)
        if not columns:
            raise ValueError(f'The selected {kind.title()} CDRs have no materialised reporting columns.')
        return Path(snapshot.workspace), task_repository.reporting_rows_table_name(kind), columns

    def chart_query_columns(entry, multivendor):
        reported = core.reporting_query_columns(entry.source_kind, [entry], multivendor)
        explicit = set(catalog_kpi_fields(entry.kpi))
        explicit.update(field for field in entry_dynamic_fields(entry) if field)
        explicit.update(_legend_dimensions(entry.legend))
        explicit.update(parse_catalog_grouping(entry.rows_split).dimensions)
        explicit.update(parse_catalog_grouping(entry.columns_split).dimensions)
        explicit.update(condition.column for condition in parse_catalog_filters(entry.filters))
        for dimension in entry.calculated_dimensions:
            explicit.update(dimension.default_from)
            for rule in dimension.rules:
                for condition in rule.conditions:
                    explicit.update(split_calculated_dimension_aliases(condition.column))
        selected = list(reported)
        selected.extend(('dataset_id', 'source_row_id', 'Vendor'))
        # RF companions use the same source selection. A stable projection of
        # all radio fields lets RSRP/SINR and LTE/NR share the raw-frame cache,
        # rather than rereading millions of observations for each measure.
        radio_columns = (set(RF_CATALOG_FIELDS) if entry.source_kind == 'all'
                         else rf_source_columns(entry.source_kind))
        radio_identities = {identity(column) for column in radio_columns}
        if any(identity(column) in radio_identities for column in catalog_kpi_fields(entry.kpi)):
            selected.extend(radio_columns)
            return sorted(set(selected), key=lambda column: (identity(column), column))
        return list(dict.fromkeys(selected))

    def chart_filter_sql(entry, columns, multivendor, task_repository):
        """Push safe physical template filters into SQLite before DataFrame creation."""
        clauses = []
        parameters = []
        complete = True
        quote = lambda value: '"' + str(value).replace('"', '""') + '"'
        for condition in parse_catalog_filters(entry.filters):
            normalized = identity(condition.column)
            if normalized in {'threshold', 'buckets', 'binsize', 'classcolours', 'histogramoperator'}:
                continue
            if multivendor and normalized in {'operator', 'operatorvendor', 'vendorv3'}:
                complete = False
                continue
            column = next((
                resolve_sql_column(columns, candidate.strip())
                for candidate in split_calculated_dimension_aliases(condition.column)
                if resolve_sql_column(columns, candidate.strip())
            ), None)
            if not column:
                complete = False
                continue
            values = [str(value) for value in condition.values]
            if normalized == 'vendor':
                mappings = task_repository.list_operator_mappings()
                values = vendor_filter_values(values, mappings)
            if normalized == 'operator':
                workspace_mappings = task_repository.list_operator_mappings()
                canonical_values = {
                    workspace_mappings.get(value.strip().casefold(), value.strip()).casefold()
                    for value in values
                }
                values = list(dict.fromkeys([
                    *values,
                    *(
                        source for source, canonical in workspace_mappings.items()
                        if canonical.casefold() in canonical_values
                    ),
                    *(
                        canonical for canonical in workspace_mappings.values()
                        if canonical.casefold() in canonical_values
                    ),
                ]))
            selected = (
                f"COALESCE(CAST({quote(column)} AS TEXT), '') COLLATE NOCASE"
                if '' in values
                else f'{quote(column)} COLLATE NOCASE'
            )
            if normalized in {'operatorvendor', 'vendorv3'} and condition.operator not in {'CONTAINS', 'NOT CONTAINS'}:
                complete = False
                continue
            if condition.operator in {'=', '!='} and len(values) > 1:
                placeholders = ', '.join('?' for _value in values)
                expression = f'{selected} IN ({placeholders})'
                clauses.append(f'NOT ({expression})' if condition.operator == '!=' else expression)
                parameters.extend(values)
            elif condition.operator in {'=', '!='}:
                clauses.append(f'{selected} {condition.operator} ?')
                parameters.append(values[0])
            elif condition.operator in {'IN', 'NOT IN'}:
                placeholders = ', '.join('?' for _value in values)
                expression = f'{selected} IN ({placeholders})'
                clauses.append(f'NOT ({expression})' if condition.operator == 'NOT IN' else expression)
                parameters.extend(values)
            elif condition.operator in {'CONTAINS', 'NOT CONTAINS'}:
                checks = [f'instr(LOWER(COALESCE(CAST({quote(column)} AS TEXT), \'\')), LOWER(?)) > 0' for _value in values]
                expression = f'({" OR ".join(checks)})'
                clauses.append(f'NOT {expression}' if condition.operator == 'NOT CONTAINS' else expression)
                parameters.extend(values)
            else:
                complete = False
        return ' AND '.join(clauses), parameters, complete

    def chart_aggregation_columns(entry, columns, multivendor, filters_applied):
        """Return physical dimensions that SQLite can aggregate without changing chart semantics."""
        if (
            multivendor
            or not filters_applied
            or entry.calculated_dimensions
            or entry.chart_type.casefold() != '100% stacked vertical bars'
        ):
            return None
        requested = [
            *parse_catalog_grouping(entry.rows_split).dimensions,
            *parse_catalog_grouping(entry.columns_split).dimensions,
            *catalog_kpi_fields(entry.kpi),
        ]
        resolved = [resolve_sql_column(columns, name) for name in requested]
        if any(column is None for column in resolved):
            return None
        return list(dict.fromkeys(resolved))

    def reporting_read_connection(database_path):
        """Read connection with a larger page cache and memory-mapped I/O."""
        connection = workspace_read_connection(database_path)
        connection.execute(f'PRAGMA cache_size=-{DESIGN_READ_CACHE_KIB}')
        connection.execute(f'PRAGMA mmap_size={DESIGN_READ_MMAP_BYTES}')
        return connection

    def selection_share(database_path, table, where, parameters) -> float:
        """Share of a table's rows matched by the dataset part of a selection."""
        connection = workspace_read_connection(database_path)
        try:
            dataset_clause = re.match(r'\s*\(?(dataset_id IN \([?, ]+\))', where)
            if not dataset_clause:
                return 0.0
            count = dataset_clause.group(1).count('?')
            with cdr_store.combined_scope(list(parameters)[:count]):
                selected = connection.execute(
                    f'SELECT count(*) FROM "{table}" WHERE {dataset_clause.group(1)}', list(parameters)[:count],
                ).fetchone()[0]
            total = connection.execute(f'SELECT count(*) FROM "{table}"').fetchone()[0]
            return selected / total if total else 0.0
        except sqlite3.Error:
            return 0.0
        finally:
            connection.close()

    def sequential_scan_source(database_path, branch, where, parameters):
        match = re.search(r'FROM "(reporting_rows_[a-z]+)"$', branch)
        if match and selection_share(database_path, match.group(1), where, parameters) >= DESIGN_SEQUENTIAL_SCAN_SHARE:
            return branch + ' NOT INDEXED'
        return branch

    def sequential_scan_table(database_path, table_name, where, parameters):
        if selection_share(database_path, table_name, where, parameters) >= DESIGN_SEQUENTIAL_SCAN_SHARE:
            return '"' + str(table_name).replace('"', '""') + '" NOT INDEXED'
        return table_name

    def calculated_dimension_fields(dimension) -> set[str]:
        """The fields an Auto-calculated Field reads, row by row: its sources, the fields of its default and of the
        conditions of its rules and expression (each alias)."""
        fields = {dimension.name}
        for value in (*dimension.sources, *dimension.default_from):
            fields.update(split_calculated_dimension_aliases(value))
        conditions = [condition for rule in dimension.rules for condition in rule.conditions]
        expressions = [dimension.expression] if dimension.expression is not None else []
        while expressions:
            expression = expressions.pop()
            for branch in expression.branches:
                conditions.extend(condition for group in branch.conditions for condition in group)
                if not isinstance(branch.result, str) and branch.result is not None:
                    expressions.append(branch.result)
            if not isinstance(expression.otherwise, str) and expression.otherwise is not None:
                expressions.append(expression.otherwise)
        for condition in conditions:
            fields.update(split_calculated_dimension_aliases(condition.column))
        return {field.strip() for field in fields if field and field.strip()}

    def compact_chart_columns(entry, source_columns, multivendor):
        """Return the only physical columns a chart reads, or None when unknown.

        An Auto-calculated Field reads the columns of its conditions (row by
        row): those of the source that it names are read instead. Charts
        referencing other semantic fields keep the complete read.
        Order-independent chart types read these columns grouped; the others
        read them ungrouped, in the same row order as the complete read.
        Vendor Comparisons also read the Operator and Vendor identities their
        comparison groups are built from.
        """
        by_identity = defaultdict(list)
        for column in source_columns:
            by_identity[identity(column)].append(column)
        calculated = {identity(dimension.name): dimension for dimension in entry.calculated_dimensions}
        referenced = [
            *catalog_kpi_fields(entry.kpi),
            *(field for field in entry_dynamic_fields(entry) if field),
            *_legend_dimensions(entry.legend),
            *parse_catalog_grouping(entry.rows_split).dimensions,
            *parse_catalog_grouping(entry.columns_split).dimensions,
            *(condition.column for condition in parse_catalog_filters(entry.filters)
              if identity(condition.column) not in CHART_SETTING_FILTERS),
        ]
        selected = []
        expanded: set[str] = set()
        seen: list[str] = []
        # Whether the chart names an Auto-calculated Field or a chart setting, which it once read with every column.
        named_whole = False
        while referenced:
            name = referenced.pop(0)
            seen.append(name)
            key = identity(name)
            if key in CHART_SETTING_FILTERS:
                # A chart setting named as a legend (Buckets, Threshold) is no column of any read.
                named_whole = True
                continue
            if key in calculated:
                named_whole = True
                if key not in expanded:
                    expanded.add(key)
                    # Its fields that the source has (an alias it lacks is not read either way), and the other
                    # Auto-calculated Fields it reads.
                    fields = calculated_dimension_fields(calculated[key])
                    selected.extend(column for field in fields for column in by_identity.get(identity(field), []))
                    referenced.extend(field for field in fields if identity(field) in calculated)
                continue
            if key in by_identity:
                selected.extend(by_identity[key])
            elif key not in COMPACT_READ_DERIVED_FIELDS:
                return None
        if multivendor or any(identity(name) in COMPACT_READ_VENDOR_FIELDS for name in seen):
            for key in COMPACT_READ_VENDOR_FIELDS:
                selected.extend(by_identity.get(key, []))
        if multivendor:
            for key in ('operator', 'subscriber'):
                selected.extend(by_identity.get(key, []))
        for name in COMPACT_READ_IMPLICIT_FIELDS:
            selected.extend(by_identity.get(identity(name), []))
        if named_whole and entry.chart_type.casefold() == 'map':
            # Such a map still shows where its series lie by City or Region (as when it read every column).
            for name in ('City', 'Region'):
                selected.extend(by_identity.get(identity(name), []))
        return sorted(set(selected), key=lambda column: (identity(column), column))

    def where_scope(where, parameters):
        """The CDRs a selection WHERE keeps (its leading dataset_id IN clause, see selection_where), as the scope of an
        emptied combined table (see cdr_store.combined_scope); every CDR without one."""
        match = re.match(r'\s*\(?\s*dataset_id IN \(([?, ]+)\)', str(where))
        if not match:
            return nullcontext()
        return cdr_store.combined_scope(list(parameters)[:match.group(1).count('?')])

    def load_reporting_frame(
        database_path, table_name, source_columns, requested_columns, where, parameters,
        aggregation_columns=None, expand_groups=False,
    ):
        lookup = {identity(column): column for column in source_columns}
        selected_columns = []
        for requested in aggregation_columns or requested_columns:
            actual = lookup.get(identity(requested))
            if actual and actual not in selected_columns:
                selected_columns.append(actual)
        if not selected_columns:
            return pd.DataFrame()
        quote = lambda column: '"' + str(column).replace('"', '""') + '"'
        select_clause = ", ".join(quote(column) for column in selected_columns)
        if aggregation_columns:
            select_clause += ', COUNT(*) AS "__catalog_weight"'
        elif expand_groups:
            # Identical rows travel once with their count and are expanded
            # below, so the chart receives exactly the same rows.
            select_clause += ', COUNT(*) AS "__catalog_rows"'
        group_by = f' GROUP BY {", ".join(quote(column) for column in selected_columns)}' if aggregation_columns or expand_groups else ''
        branches = getattr(table_name, 'branches', ())
        if expand_groups and not aggregation_columns:
            # Grouped reads do not depend on row order: when the selection
            # covers much of a table, a sequential scan avoids random row lookups.
            branches = tuple(sequential_scan_source(database_path, branch, where, parameters) for branch in branches) if branches else ()
            if not branches and not table_name.startswith('('):
                table_name = sequential_scan_table(database_path, table_name, where, parameters)
        if branches:
            # Filter each CDR type separately: SQLite flattens each simple
            # subquery, uses the dataset index and computes only the needed columns.
            query = ' UNION ALL '.join(f'SELECT {select_clause} FROM ({branch}) WHERE {where}{group_by}' for branch in branches)
            parameters = list(parameters) * len(branches)
            if aggregation_columns:
                columns_text = ", ".join(quote(column) for column in selected_columns)
                query = (f'SELECT {columns_text}, SUM("__catalog_weight") AS "__catalog_weight" '
                         f'FROM ({query}) GROUP BY {columns_text}')
        else:
            source = table_name if table_name.startswith('(') or table_name.endswith(' NOT INDEXED') else quote(table_name)
            query = f'SELECT {select_clause} FROM {source} WHERE {where}{group_by}'
        connection = reporting_read_connection(database_path)
        try:
            with where_scope(where, parameters):
                frame = pd.read_sql_query(query, connection, params=parameters)
        finally:
            connection.close()
        return expand_grouped_rows(frame) if expand_groups and not aggregation_columns else frame

    def expand_grouped_rows(frame):
        """Repeat each grouped row by its count, restoring every source row."""
        frame = frame.copy()
        counts = frame.pop('__catalog_rows').to_numpy()
        return frame.loc[frame.index.repeat(counts)].reset_index(drop=True)

    column_caches: dict[str, ColumnCache] = {}
    column_cache_prefetch: OrderedDict = OrderedDict()

    def workspace_column_cache(workspace) -> ColumnCache:
        """The active workspace's per-CDR column cache; other workspaces release theirs."""
        key = str(Path(workspace).resolve())
        with lock:
            cache = column_caches.get(key)
            if cache is None:
                for other in list(column_caches):
                    column_caches.pop(other).forget()
                # In memory only: the columns of the CDRs come from their columnar copies (see build_from_copies).
                cache = column_caches[key] = ColumnCache(
                    column_cache_dir(workspace), memory_bytes=default_memory_bytes(), disk_bytes=0, persist=False,
                )
            return cache

    def forget_column_caches() -> None:
        with lock:
            for cache in column_caches.values():
                cache.forget()

    def plain_reporting_query(table_name, selected_columns, where, parameters):
        """The ungrouped query ``load_reporting_frame`` runs for these columns."""
        quote = lambda column: '"' + str(column).replace('"', '""') + '"'
        select_clause = ", ".join(quote(column) for column in selected_columns)
        branches = getattr(table_name, 'branches', ())
        if branches:
            query = ' UNION ALL '.join(f'SELECT {select_clause} FROM ({branch}) WHERE {where}' for branch in branches)
            return query, list(parameters) * len(branches)
        source = table_name if table_name.startswith('(') else quote(table_name)
        return f'SELECT {select_clause} FROM {source} WHERE {where}', list(parameters)

    planned_row_orders: OrderedDict = OrderedDict()

    def planned_row_order(database_path, query, parameters):
        """The row order of SQLite's plan for a query, when the column cache reproduces it.

        Searching a per-dataset index visits the CDRs in ascending order and,
        within each, the event date (date index) and table order; a full scan
        returns table order. Charts whose rows follow any other plan keep
        reading from SQLite, so their row order never changes.
        """
        key = (str(database_path), query, tuple(parameters))
        with lock:
            if key in planned_row_orders:
                planned_row_orders.move_to_end(key)
                return planned_row_orders[key]
        order = sqlite_plan_row_order(database_path, query, parameters)
        with lock:
            planned_row_orders[key] = order
            while len(planned_row_orders) > 512:
                planned_row_orders.popitem(last=False)
        return order

    def sqlite_plan_row_order(database_path, query, parameters):
        connection = reporting_read_connection(database_path)
        try:
            plan = [str(row[3]) for row in connection.execute(f'EXPLAIN QUERY PLAN {query}', parameters).fetchall()]
        except sqlite3.Error:
            return None
        finally:
            connection.close()
        modes = set()
        for detail in plan:
            if detail in {'COMPOUND QUERY', 'LEFT-MOST SUBQUERY', 'UNION ALL'}:
                continue
            match = re.fullmatch(r'SEARCH (reporting_rows_[a-z]+) USING INDEX (\S+) \(dataset_id=\?(?: AND .*)?\)', detail)
            if match and match.group(2) == f'idx_{match.group(1)}_dataset_event_time_date':
                modes.add('dataset_date_rowid')
            elif match and match.group(2) == f'idx_{match.group(1)}_dataset_id_rows':
                modes.add('dataset_rowid')
            elif re.fullmatch(r'SCAN reporting_rows_[a-z]+', detail):
                modes.add('rowid')
            else:
                return None
        return modes.pop() if len(modes) == 1 else None

    def cache_prefetch_expressions(task_repository, columns, definition):
        """Filter expressions cached with every CDR, so later filter changes skip SQLite."""
        fields = tuple(dict.fromkeys((*ADAPTATIVE_FILTER_FIELDS, 'Vendor', *definition.custom_fields)))
        key = (tuple(columns), fields)
        with lock:
            cached = column_cache_prefetch.get(key)
            if cached is not None:
                column_cache_prefetch.move_to_end(key)
                return cached
        expressions = [filter_sql_normalized_expression(task_repository, columns, field) for field in fields]
        expressions.append(selection_date_expression(task_repository, columns))
        sheet_term = selection_sheet_term(task_repository, columns)
        if sheet_term:
            expressions.append(f'LOWER(TRIM(CAST({sheet_term[0]} AS TEXT)))')
        result = tuple(dict.fromkeys(expression for expression in expressions if expression))
        with lock:
            column_cache_prefetch[key] = result
            while len(column_cache_prefetch) > 64:
                column_cache_prefetch.popitem(last=False)
        return result

    def column_cache_sources(task_repository, kind, table_name, definition, rows_by_kind=None):
        """Return (source, selected CDR rows) for each combined table behind a chart source."""
        projections = getattr(table_name, 'branch_projections', ())
        sources = projections or ((kind, str(table_name), ()),)
        result = []
        for source_kind, table, projection in sources:
            physical = task_repository.list_reporting_row_columns(source_kind)
            source = CacheSource(
                source_kind, table, tuple(projection),
                has_event_time=any(str(column).casefold() == 'event_start_time' for column in physical),
            )
            rows = rows_by_kind.get(source_kind, []) if rows_by_kind is not None else design_datasets(definition, source_kind, task_repository)
            result.append((source, rows))
        return result

    def column_cache_segments(task_repository, kind, table_name, definition, rows_by_kind=None):
        """Cache segments of the selected CDRs, keyed by their data revision."""
        dimensions = sha256(json.dumps(
            task_repository.list_calculated_dimensions(), sort_keys=True, default=str,
        ).encode()).hexdigest()
        segments = []
        for source, rows in column_cache_sources(task_repository, kind, table_name, definition, rows_by_kind):
            # Same revision contract as every Design cache: the CDR version,
            # its combined table's row and column revisions and the calculated
            # dimensions, so a cached column is never stale.
            revision = [
                task_repository.get_workspace_state(f'combined_reporting_updated_{source.kind}'),
                task_repository.get_workspace_state(f'combined_reporting_columns_updated_{source.kind}'),
                dimensions,
            ]
            for row in sorted(rows, key=lambda item: int(item['id'])):
                version = json.dumps([
                    row.get('updated_at'), row.get('processed_at'), row.get('normalization_version'),
                    row.get('row_count'), *revision,
                ], default=str)
                segments.append(Segment(source, int(row['id']), version))
        return segments

    def cached_selection_counts(task_repository, kind, selected, definition, columns):
        """Universe and filtered row counts of one CDR type, from the column cache.

        They equal the SQLite counts: the universe is the selected CDRs within
        the dates and sheets, and the filtered rows also match every filter.
        """
        if not DESIGN_COLUMN_CACHE or not selected:
            return None
        universe_definition = definition.model_copy(deep=True)
        universe_definition.filters = {}
        universe = selection_cache_terms(task_repository, columns, universe_definition)
        selection = selection_cache_terms(task_repository, columns, definition)
        database_path = Path(task_repository.db_path)
        segments = column_cache_segments(
            task_repository, kind, task_repository.reporting_rows_table_name(kind), definition, {kind: selected},
        )
        prefetch = cache_prefetch_expressions(task_repository, columns, definition)
        cache = workspace_column_cache(database_path)
        connect = lambda: reporting_read_connection(database_path)
        builder = copies_builder(task_repository, segments)
        for _attempt in range(2):
            try:
                return count_rows(cache, connect, segments, universe, selection, prefetch=lambda _source: prefetch,
                                  builder=builder)
            except StaleSegment:
                for segment in segments:
                    cache.invalidate(segment)
            except (sqlite3.Error, OSError, ValueError, KeyError):
                return None
        return None

    distinct_value_memo: OrderedDict = OrderedDict()

    def dataset_distinct_values(task_repository, dataset, column):
        """Every distinct value of one CDR column, kept in memory per CDR version.

        The scan of a CDR's own rows only depends on that CDR, so a changed
        Design filter reuses it instead of reading the CDR again.
        """
        version = [dataset.get('id'), dataset.get('updated_at'), dataset.get('processed_at'),
                   dataset.get('normalization_version'), dataset.get('row_count'), column]
        key = (str(task_repository.db_path), sha256(json.dumps(version, default=str).encode()).hexdigest())
        with lock:
            if key in distinct_value_memo:
                distinct_value_memo.move_to_end(key)
                return list(distinct_value_memo[key])
        values = task_repository.list_distinct_dataset_row_values(int(dataset['id']), column, limit=None)
        with lock:
            distinct_value_memo[key] = list(values)
            while len(distinct_value_memo) > 512:
                distinct_value_memo.popitem(last=False)
        return list(values)

    def copies_builder(task_repository, segments):
        """The builder of the columns of ``segments`` from the columnar copies of their CDRs, in memory only, when every
        copy holds its CDR as it is (see build_from_copies); None otherwise."""
        dataset_ids = {segment.dataset_id for segment in segments}
        if not dataset_ids or not dataset_ids <= cdr_store.fresh_copies(task_repository.db_path, dataset_ids):
            return None
        return lambda segment, expressions: build_from_copies(task_repository.db_path, segment, expressions)

    def column_cache_frame(snapshot, kind, task_repository, database_path, table_name, source_columns,
                           requested_columns, terms, order, rows_by_kind=None, weighted=False):
        """Assemble a chart's source rows from the per-CDR column cache (``weighted``: each distinct row once, see
        read_frame)."""
        lookup = {identity(column): column for column in source_columns}
        selected_columns = []
        for requested in requested_columns:
            actual = lookup.get(identity(requested))
            if actual and actual not in selected_columns:
                selected_columns.append(actual)
        if not selected_columns:
            return pd.DataFrame()
        quote = task_repository._quote_identifier
        segments = column_cache_segments(task_repository, kind, table_name, snapshot.definition, rows_by_kind)
        prefetch = cache_prefetch_expressions(task_repository, source_columns, snapshot.definition)
        cache = workspace_column_cache(snapshot.workspace)
        connect = lambda: reporting_read_connection(database_path)
        builder = copies_builder(task_repository, segments)
        for _attempt in range(2):
            try:
                if weighted:
                    # DuckDB groups the selected rows of each copy (see read_grouped_frame); the complete reads of the
                    # other charts keep their columns cached across charts.
                    grouped = read_grouped_frame(
                        cache, database_path, segments, [(column, quote(column)) for column in selected_columns],
                        terms, order, builder,
                    )
                    if grouped is not None:
                        return grouped
                return read_frame(
                    cache, connect, segments, [(column, quote(column)) for column in selected_columns],
                    terms, order, prefetch=lambda _source: prefetch, builder=builder, weighted=weighted,
                )
            except StaleSegment:
                # Rows changed under an unchanged revision: rebuild the CDRs.
                for segment in segments:
                    cache.invalidate(segment)
        raise StaleSegment('The selected CDR rows keep changing.')

    def parse_chart_column_filters(encoded):
        try:
            payload = json.loads(encoded or '{}')
        except (TypeError, json.JSONDecodeError):
            return {}
        if not isinstance(payload, dict):
            return {}
        return {
            str(column): [str(value) for value in values]
            for column, values in payload.items()
            if isinstance(values, list)
        }

    def chart_dataset_filter_values(frame):
        return {
            str(column): sorted(
                {'' if pd.isna(value) else str(value) for value in frame[column]},
                key=lambda value: value.casefold(),
            )
            for column in frame.columns
        }

    def chart_dataset_column_classes(snapshot, entry, columns):
        task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
        datasets = []
        for dataset_id in [row['id'] for row in design_datasets(snapshot.definition, entry.source_kind, task_repository)]:
            dataset = task_repository.get_dataset(int(dataset_id))
            if dataset:
                datasets.append(core.serialize_dataset_row(dataset))
        return core._chart_preview_column_classes(columns, datasets, task_repository)

    def chart_dataset_column_metadata(snapshot, entry, columns):
        task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
        datasets = []
        for dataset_id in [row['id'] for row in design_datasets(snapshot.definition, entry.source_kind, task_repository)]:
            dataset = task_repository.get_dataset(int(dataset_id))
            if dataset:
                datasets.append(core.serialize_dataset_row(dataset))
        return core._chart_preview_column_metadata(
            columns, datasets, entry.source_kind, task_repository,
        )

    def apply_chart_column_filters(frame, column_filters):
        result = frame
        lookup = {identity(column): column for column in frame.columns}
        for requested, values in column_filters.items():
            column = lookup.get(identity(vendor_filter_column(requested)))
            if column is None or not values:
                return result.iloc[0:0]
            if identity(vendor_filter_column(requested)) in {'vendor', 'operatorvendor'}:
                values = vendor_match_values(requested, values, frame.attrs.get('operator_mappings', {}))
            accepted = {str(value).strip().casefold() for value in values}
            result = result.loc[
                result[column].map(
                    lambda value: '' if pd.isna(value) else str(value).strip().casefold()
                ).isin(accepted)
            ]
        return result

    def chart_row_universe(snapshot, entry, selected=None):
        """Return the strongest known upper bound for one chart's source rows."""
        if selected is None:
            task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
            selected = design_datasets(snapshot.definition, entry.source_kind, task_repository)
        limits = []
        source_total = sum(int(row.get('row_count') or 0) for row in selected)
        if source_total:
            limits.append(source_total)
        filtered_rows = (snapshot.payload.get('rows') or {}).get(entry.source_kind)
        if filtered_rows is not None:
            limits.append(max(0, int(filtered_rows)))
        return min(limits) if limits else 0

    def reporting_chart_data_page(
        snapshot, entry, page, column_filters=None, include_filter_values=False, filter_column='',
    ):
        """Read one exact chart-data page without rebuilding its full DataFrame."""
        spec = _catalog_spec(entry)
        if entry.source_kind == 'all' or spec.get('operators'):
            return None
        # This SQL-only pagination path reads physical CDR values. Chart frames
        # intentionally replace Operator aliases (and their derived Subscriber
        # and Vendor presentation labels) in memory, so using the SQL values
        # for the dropdown would make a visible canonical value impossible to
        # select. Use the shared mapped frame whenever such a column is shown.
        if any(
            identity(column) in {'operator', 'subscriber', 'vendor', 'vendoronly', 'operatorvendor'}
            for column in chart_query_columns(entry, snapshot.multivendor)
        ):
            return None
        task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
        selected = design_datasets(snapshot.definition, entry.source_kind, task_repository)
        database_path, table_name, source_columns = reporting_source(
            snapshot, entry.source_kind, task_repository,
        )
        where, parameters = selection_where(
            task_repository, entry.source_kind, [int(row['id']) for row in selected], snapshot.definition,
            columns=source_columns,
        )
        source_filters = (
            (spec.get('sessions'), ('Session_Type', 'session_type', 'Test_Name', 'Test_Type')),
            (spec.get('tests'), ('Test_Name', 'test_name', 'Type_of_Test', 'Test_Type')),
            (spec.get('directions'), ('Direction', 'direction', 'Call_Direction')),
            ((spec.get('city_scope'),) if spec.get('city_scope') else None, ('city', 'City', 'G_Level_1', 'G_Level_2')),
        )
        for values, candidates in source_filters:
            if not values:
                continue
            column = next((resolve_sql_column(source_columns, name) for name in candidates if resolve_sql_column(source_columns, name)), None)
            if not column:
                continue
            expression = ' OR '.join(
                f'LOWER(COALESCE(CAST({task_repository._quote_identifier(column)} AS TEXT), \'\')) LIKE ?'
                for _value in values
            )
            where = f'({where}) AND ({expression})'
            parameters.extend(f'%{str(value).casefold()}%' for value in values)
        template_where, template_parameters, filters_applied = chart_filter_sql(
            entry, source_columns, snapshot.multivendor, task_repository,
        )
        if not filters_applied:
            return None
        if template_where:
            where = f'({where}) AND ({template_where})'
            parameters.extend(template_parameters)
        lookup = {identity(column): column for column in source_columns}
        selected_columns = []
        for requested in chart_query_columns(entry, snapshot.multivendor):
            actual = lookup.get(identity(requested))
            if actual and actual not in selected_columns:
                selected_columns.append(actual)
        dataset_column = lookup.get(identity('dataset_id'))
        source_row_column = lookup.get(identity('source_row_id'))
        if not selected_columns or not dataset_column or not source_row_column:
            return None
        quote = lambda column: '"' + str(column).replace('"', '""') + '"'
        identity_columns = f'{quote(dataset_column)}, {quote(source_row_column)}'
        page = max(0, int(page))
        connection = workspace_read_connection(database_path)
        # Every query below reads the rows of the selected CDRs only (an emptied combined table is given them alone).
        scope = cdr_store.combined_scope([int(row['id']) for row in selected])
        scope.__enter__()
        try:
            chart_total = int(connection.execute(
                f'SELECT COUNT(*) FROM ('
                f'SELECT 1 FROM {quote(table_name)} WHERE {where} GROUP BY {identity_columns}'
                f')',
                parameters,
            ).fetchone()[0])
            chart_total = min(chart_total, chart_row_universe(snapshot, entry, selected))
            filter_values = {}
            if include_filter_values:
                values_query = ', '.join(
                    f'json_group_array(DISTINCT COALESCE(CAST({quote(column)} AS TEXT), \'\')) AS filter_{index}'
                    for index, column in enumerate(selected_columns)
                )
                values_row = connection.execute(
                    f'SELECT {values_query} FROM {quote(table_name)} WHERE {where}', parameters,
                ).fetchone()
                filter_values = {
                    column: sorted(
                        (str(value) for value in json.loads(values_row[index] or '[]')),
                        key=lambda value: value.casefold(),
                    )
                    for index, column in enumerate(selected_columns)
                }
            if filter_column:
                requested_filter_column = lookup.get(identity(filter_column))
                if requested_filter_column:
                    values_where = where
                    values_parameters = list(parameters)
                    for requested, values in (column_filters or {}).items():
                        if identity(requested) == identity(filter_column):
                            continue
                        column = lookup.get(identity(requested))
                        if column is None or not values:
                            values_where = f'({values_where}) AND 0'
                            continue
                        placeholders = ', '.join('?' for _value in values)
                        values_where = (
                            f'({values_where}) AND '
                            f'LOWER(TRIM(COALESCE(CAST({quote(column)} AS TEXT), \'\'))) IN ({placeholders})'
                        )
                        values_parameters.extend(str(value).strip().lower() for value in values)
                    filter_values = sorted(
                        (
                            str(row[0]) for row in connection.execute(
                                f'SELECT DISTINCT COALESCE(CAST({quote(requested_filter_column)} AS TEXT), \'\') '
                                f'FROM {quote(table_name)} WHERE {values_where}',
                                values_parameters,
                            )
                        ),
                        key=str.casefold,
                    )
            filtered_where = where
            filtered_parameters = list(parameters)
            for requested, values in (column_filters or {}).items():
                column = lookup.get(identity(requested))
                if column is None or not values:
                    filtered_where = f'({filtered_where}) AND 0'
                    continue
                placeholders = ', '.join('?' for _value in values)
                filtered_where = (
                    f'({filtered_where}) AND '
                    f'LOWER(TRIM(COALESCE(CAST({quote(column)} AS TEXT), \'\'))) IN ({placeholders})'
                )
                filtered_parameters.extend(str(value).strip().lower() for value in values)
            total = chart_total
            if column_filters:
                total = int(connection.execute(
                    f'SELECT COUNT(*) FROM ('
                    f'SELECT 1 FROM {quote(table_name)} WHERE {filtered_where} GROUP BY {identity_columns}'
                    f')',
                    filtered_parameters,
                ).fetchone()[0])
                total = min(total, chart_total)
            page_size = min(100, max(0, total - page * 100))
            query = (
                f'SELECT {", ".join(quote(column) for column in selected_columns)} '
                f'FROM {quote(table_name)} WHERE {filtered_where} GROUP BY {identity_columns} '
                f'ORDER BY {identity_columns} LIMIT ? OFFSET ?'
            )
            visible = pd.read_sql_query(
                query, connection, params=(*filtered_parameters, page_size, page * 100),
            )
        finally:
            scope.__exit__(None, None, None)
            connection.close()
        if snapshot.multivendor:
            visible = ensure_vendor_group(visible)
        mapping_settings = task_repository.chart_mapping_settings()
        operator_mappings = mapping_settings['operator_mappings']
        visible = apply_operator_mappings(visible, operator_mappings)
        visible = drop_excluded_operators(visible, excluded_operators(mapping_settings))
        visible.attrs.update(mapping_settings)
        visible = normalise_operator_aliases(visible)
        return visible, total, chart_total, filter_values

    def unique_chart_dataset_rows(snapshot, entry, frame):
        """Keep one preview row per source row and never exceed its filtered CDR universe."""
        lookup = {identity(column): column for column in frame.columns}
        identity_columns = [
            lookup[key] for key in (identity('dataset_id'), identity('source_row_id')) if key in lookup
        ]
        result = frame.drop_duplicates(subset=identity_columns, keep='first') if identity_columns else frame
        universe = chart_row_universe(snapshot, entry)
        return result.iloc[:universe] if len(result) > universe else result

    def release_idle_snapshot_frames() -> None:
        """Bound chart-frame memory to the most recently used snapshots.

        Every snapshot keeps its lightweight payload, but only the latest few
        keep their pandas frames; older ones rebuild them on demand. Without
        this bound, many opened universes could exhaust server memory.
        """
        with lock:
            # Snapshots a user is viewing keep their frames first, so the
            # preparations of other Designs (exports) cannot evict the open
            # Design and force every remaining chart to reread its CDR rows.
            ordered = list(snapshots.values())
            viewed = sorted((item for item in ordered if item.viewed_at), key=lambda item: item.viewed_at)
            unviewed = [item for item in ordered if not item.viewed_at]
            recent = [*viewed[-DESIGN_FRAME_CACHE_SNAPSHOTS:], *unviewed[-1:]]
            for snapshot in ordered:
                if any(snapshot is item for item in recent):
                    continue
                if snapshot.frames or snapshot.filtered_frames or snapshot.chart_frames:
                    snapshot.frames.clear()
                    snapshot.filtered_frames.clear()
                    snapshot.chart_frames.clear()
                    snapshot.frame_locks.clear()

    def snapshot_chart(token, index, user, *, include_frame=True, expected_workspace: str | None = None, complete=False):
        """Return a chart's snapshot, entry and prepared rows.

        ``complete`` returns every source column, as the chart dataset views
        need, instead of the compact grouped read used to build chart models.
        """
        frame_key = ('complete', index) if complete else index
        with lock:
            snapshot = snapshots.get(token)
            if snapshot is not None:
                snapshots.move_to_end(token)
        if snapshot is None or snapshot.workspace != (expected_workspace or workspace_key()) or snapshot.owner not in {user.username, '*'}:
            raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
        if index < 0 or index >= len(snapshot.entries):
            raise HTTPException(404, 'Chart not found.')
        entry = snapshot.entries[index]
        if not include_frame:
            return snapshot, entry, None
        with lock:
            frame = snapshot.chart_frames.get(frame_key)
        if frame is None:
            task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
            with lock:
                source = snapshot.sources.get(entry.source_kind)
            if source is None:
                rows_by_kind = {
                    kind: design_datasets(snapshot.definition, kind, task_repository)
                    for kind in (KINDS if entry.source_kind == 'all' else (entry.source_kind,))
                }
                source = (rows_by_kind, *reporting_source(snapshot, entry.source_kind, task_repository))
                with lock:
                    snapshot.sources[entry.source_kind] = source
            rows_by_kind, database_path, table_name, source_columns = source
            rows_by_kind = chart_nr_mode_rows(entry, snapshot.definition, rows_by_kind)
            selected = [row for rows in rows_by_kind.values() for row in rows]
            where, parameters = selection_where(
                task_repository, entry.source_kind, [int(row['id']) for row in selected], snapshot.definition,
                columns=source_columns,
            )
            requested_columns = chart_query_columns(entry, snapshot.multivendor)
            template_where, template_parameters, template_filters_applied = chart_filter_sql(
                entry, source_columns, snapshot.multivendor, task_repository,
            )
            if entry.source_kind == 'all':
                # Every RF chart over the combined CDR types then shares one
                # loaded frame; its own filters are applied in memory below.
                template_where, template_parameters, template_filters_applied = '', [], False
            if template_where:
                where = f'({where}) AND ({template_where})'
                parameters.extend(template_parameters)
            aggregation_columns = chart_aggregation_columns(
                entry, source_columns, snapshot.multivendor, template_filters_applied,
            )
            compact_columns = None if complete or aggregation_columns or not DESIGN_COMPACT_READS else compact_chart_columns(
                entry, source_columns, snapshot.multivendor,
            )
            grouped_read = bool(compact_columns) and entry.chart_type.casefold() in COMPACT_READ_CHART_TYPES
            if compact_columns and not grouped_read:
                # Row order matters (for example map point order): keep the
                # complete read's filters and order, only with fewer columns.
                requested_columns = compact_columns
            elif compact_columns:
                # Template filters run in memory on the compact rows, so charts
                # reading the same fields share one grouped read.
                where, parameters = selection_where(
                    task_repository, entry.source_kind, [int(row['id']) for row in selected], snapshot.definition,
                    columns=source_columns,
                )
                template_filters_applied = False
                requested_columns = compact_columns
            # Key on the physical columns actually read: requested names that do
            # not exist in the source (such as Bin Size) do not change the frame.
            source_lookup = {identity(column): column for column in source_columns}
            read_columns = sorted({
                source_lookup[identity(column)] for column in (aggregation_columns or requested_columns)
                if identity(column) in source_lookup
            })
            # The per-CDR column cache returns the rows SQLite would return for
            # any universe or comparison mode. Order-independent charts take them
            # in table order; the others only when the column cache reproduces
            # SQLite's plan order, so no chart's row order changes.
            cache_order = None
            if DESIGN_COLUMN_CACHE and not aggregation_columns:
                ordered_columns = list(dict.fromkeys(
                    source_lookup[identity(column)] for column in requested_columns if identity(column) in source_lookup
                ))
                if ordered_columns:
                    # A chart and its data table view then share one frame.
                    cache_order = planned_row_order(
                        database_path, *plain_reporting_query(table_name, ordered_columns, where, parameters),
                    )
                if cache_order is None and not complete and entry.chart_type.casefold() in COMPACT_READ_CHART_TYPES:
                    cache_order = 'dataset_rowid'
            cache_terms = None
            if cache_order:
                # Template filters then run in memory on the cached rows.
                template_filters_applied = False
                cache_terms = selection_cache_terms(task_repository, source_columns, snapshot.definition)
                selection_identity = {
                    'datasets': sorted(int(row['id']) for row in selected),
                    'values': [[expression, sorted(values)] for expression, values in cache_terms.values],
                    'date': cache_terms.date,
                    'sheets': [cache_terms.excluded_sheets[0], sorted(cache_terms.excluded_sheets[1])] if cache_terms.excluded_sheets else None,
                    'excluded': [[expression, sorted(values)] for expression, values in cache_terms.excluded_values],
                    **({'out_of_scope': [[sorted(datasets), expression, sorted(values)]
                                         for datasets, expression, values in cache_terms.dataset_excluded_values]}
                       if cache_terms.dataset_excluded_values else {}),
                    'order': cache_order,
                }
            else:
                selection_identity = {'where': where, 'parameters': parameters}
            # An order-independent chart read from the column cache filters each distinct row once, with the number
            # of its rows, and has its rows back before its model is calculated (see read_frame).
            weighted_read = bool(cache_order) and grouped_read and DESIGN_WEIGHTED_READS
            raw_key = sha256(json.dumps({
                'kind': entry.source_kind, 'columns': read_columns, **selection_identity,
                'aggregation_columns': aggregation_columns, 'compact': grouped_read and not cache_order,
                'weighted': weighted_read,
                'chart_mappings': task_repository.chart_mapping_settings(),
            }, sort_keys=True, default=str).encode()).hexdigest()
            with lock:
                raw_frame = snapshot.frames.get(raw_key)
                frame_lock = snapshot.frame_locks.setdefault(raw_key, RLock())
            if raw_frame is None:
                with frame_lock:
                    with lock:
                        raw_frame = snapshot.frames.get(raw_key)
                    if raw_frame is None:
                        if cache_order:
                            try:
                                raw_frame = column_cache_frame(
                                    snapshot, entry.source_kind, task_repository, database_path, table_name,
                                    source_columns, requested_columns, cache_terms, cache_order, rows_by_kind,
                                    weighted=weighted_read,
                                )
                            except (sqlite3.Error, StaleSegment, OSError, ValueError, KeyError):
                                # SQLite remains the source of truth; its rows are
                                # filtered again in memory, which changes nothing.
                                raw_frame = load_reporting_frame(
                                    database_path, table_name, source_columns, requested_columns, where, parameters,
                                    aggregation_columns, expand_groups=grouped_read,
                                )
                        else:
                            raw_frame = load_reporting_frame(
                                database_path, table_name, source_columns, requested_columns, where, parameters,
                                aggregation_columns, expand_groups=grouped_read,
                            )
                        if snapshot.multivendor:
                            raw_frame = ensure_vendor_group(raw_frame)
                        mapping_settings = task_repository.chart_mapping_settings()
                        operator_mappings = mapping_settings['operator_mappings']
                        raw_frame = apply_operator_mappings(raw_frame, operator_mappings)
                        # The operators the Operator Maps exclude are left out of every Design.
                        raw_frame = drop_excluded_operators(raw_frame, excluded_operators(mapping_settings))
                        raw_frame.attrs.update(mapping_settings)
                        # NSA & SA: its campaigns that name SA are drawn as SA series (see _sa_series_keys).
                        raw_frame.attrs[BOTH_NR_MODES_ATTR] = design_nr_mode(snapshot.definition) == NSA_VS_SA
                        raw_frame = normalise_operator_aliases(raw_frame)
                        raw_frame.attrs['operator_aliases_normalized'] = True
                        with lock:
                            snapshot.frames[raw_key] = raw_frame
            # A CDF and its companion Average/Median chart normally use the
            # same source, KPI and template filters. Share that expensive
            # filtering pass while keeping each chart's aggregation and visual
            # model independent.
            # Cells of a dynamic grid share one filtered frame; each cell then
            # keeps only its own row/column values, which is a pure row filter.
            dynamic_cell = any(value is not None for value in (entry.dynamic_row_value, entry.dynamic_column_value, entry.dynamic_value))
            base_entry = replace(entry, dynamic_row_value=None, dynamic_column_value=None, dynamic_value=None) if dynamic_cell else entry
            # The filtered rows depend only on what the filtering pipeline reads:
            # source, effective filters (chart-setting pseudo filters such as Bin
            # Size are ignored there), calculated fields, the measured column and
            # the exclusion flags. Chart types sharing them (a CDF and its
            # histogram, classes or thresholds) reuse one filtering pass.
            filter_spec = _catalog_spec(base_entry)
            effective_filters = tuple(
                condition for condition in parse_catalog_filters(base_entry.filters)
                if identity(condition.column) not in CHART_SETTING_FILTERS
            )
            filtered_key = sha256(repr((
                raw_key, entry.cdr_source, effective_filters, filter_spec.get('metric'),
                entry.exclude_null_empty, entry.exclude_zero,
                (filter_spec.get('kind'), filter_spec.get('x_metric'), filter_spec.get('colour_metric'))
                if entry.exclude_null_empty or entry.exclude_zero else None,
                entry.calculated_dimensions, snapshot.multivendor, entry.vendor_comparison, entry_dynamic_fields(entry),
            )).encode()).hexdigest()
            with lock:
                prepared_frame = snapshot.filtered_frames.get(filtered_key)
                filter_lock = snapshot.frame_locks.setdefault(f'filtered:{filtered_key}', RLock())
            if prepared_frame is None:
                with filter_lock:
                    with lock:
                        prepared_frame = snapshot.filtered_frames.get(filtered_key)
                    if prepared_frame is None:
                        try:
                            prepared_frame = prepare_catalog_chart_preview_frame(
                                raw_frame, base_entry, multivendor=snapshot.multivendor,
                                template_filters_applied=template_filters_applied,
                            )[0]
                        except ValueError as exc:
                            raise HTTPException(400, str(exc)) from exc
                        with lock:
                            prepared_frame = snapshot.filtered_frames.setdefault(filtered_key, prepared_frame)
            if dynamic_cell:
                render_entry = prepare_multivendor_catalog_entry(entry) if snapshot.multivendor else entry
                prepared_frame = _select_dynamic_chart_frame(prepared_frame, render_entry)
            if weighted_read and WEIGHT_COLUMN in prepared_frame.columns:
                prepared_frame = expand_weighted_rows(prepared_frame)
            with lock:
                frame = snapshot.chart_frames.setdefault(frame_key, prepared_frame)
            release_idle_snapshot_frames()
        return snapshot, entry, frame

    def chart_model(
        token: str,
        index: int,
        user,
        *,
        expected_workspace: str | None = None,
        force: bool = False,
    ):
        snapshot, entry, _ = snapshot_chart(token, index, user, include_frame=False, expected_workspace=expected_workspace)
        model_key = canvas_model_key(snapshot, entry)
        mapping_key = sha256(json.dumps(
            Repository(Path(snapshot.workspace), core.repository.global_db_path).chart_mapping_settings(),
            sort_keys=True,
        ).encode()).hexdigest()
        if force:
            with lock:
                snapshot.chart_payloads.pop(index, None)
                # The prepared frame carries the mapping aliases, ordering and
                # theme colours that existed when it was first built. Dropping
                # only the Canvas payload would immediately rebuild it from
                # that stale frame after an Admin mapping change.
                snapshot.chart_frames.pop(index, None)
                snapshot.chart_frames.pop(('complete', index), None)
                # Filtered/raw frames are mapping-dependent too, but reloading
                # them re-reads every CDR row of the chart's type. Drop them
                # only when the workspace mappings changed since they were
                # built (or when that is unknown).
                if snapshot.mapping_key != mapping_key:
                    snapshot.frames.clear()
                    snapshot.filtered_frames.clear()
                    snapshot.frame_locks.clear()
                canvas_models.pop(model_key, None)
            forget_stored_chart_model(snapshot.workspace, model_key)
        with lock:
            snapshot.mapping_key = snapshot.mapping_key or mapping_key
            if force:
                snapshot.mapping_key = mapping_key
        with lock:
            payload = snapshot.chart_payloads.get(index)
            if payload is None and model_key in canvas_models:
                canvas_models.move_to_end(model_key)
                payload = canvas_models[model_key]
        if payload is None and not force:
            payload = stored_chart_model(snapshot.workspace, model_key)
        if payload is None:
            snapshot, entry, frame = snapshot_chart(token, index, user, expected_workspace=expected_workspace)
            payload = catalog_chart_payload(
                frame, entry, multivendor=snapshot.multivendor, prefiltered=True,
            )
            # The format the Designer gave the chart (series colours, legend, labels, axes): the renderer applies it, in
            # the viewer and in the PPT alike.
            style = design_slides.row_format(entry).get('style')
            if isinstance(style, dict) and style:
                payload = {**payload, 'format_overrides': style}
            keep = not (callable(snapshot.cancelled) and snapshot.cancelled())
            with lock:
                payload = snapshot.chart_payloads.setdefault(index, payload)
                if keep:
                    canvas_models[model_key] = payload
                    while len(canvas_models) > DESIGN_CHART_MODELS_KEPT:
                        canvas_models.popitem(last=False)
            if keep:
                store_chart_model(snapshot.workspace, model_key, payload)
        with lock:
            snapshot.chart_payloads.setdefault(index, payload)
        return payload

    def canvas_model_key(snapshot, entry, operator_mapping_key: str | None = None) -> str:
        """The key of the Canvas model of one chart entry (see canvas_models)."""
        entry_key = sha256(repr(entry).encode()).hexdigest()
        if operator_mapping_key is None:
            task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
            operator_mapping_key = sha256(json.dumps(
                task_repository.chart_mapping_settings(), sort_keys=True,
            ).encode()).hexdigest()
        filename = sha256(
            f'{DESIGN_CHART_MODEL_CACHE_VERSION}:{snapshot.selection_key}:{snapshot.definition.scope}:'
            f'{entry_key}:{operator_mapping_key}'.encode()
        ).hexdigest()
        return f'{snapshot.workspace}:{filename}'

    visible_chart_condition = Condition(Lock())
    visible_chart_requests: dict[str, int] = {}

    def chart_model_ready(snapshot, index) -> bool:
        """Whether a chart model can be served without calculating it."""
        with lock:
            if index in snapshot.chart_payloads:
                return True
        try:
            if not 0 <= index < len(snapshot.entries):
                return False
            key = canvas_model_key(snapshot, snapshot.entries[index])
            with lock:
                if key in canvas_models:
                    return True
            return ppt_design_cache.exists(ppt_design_cache.chart_path(snapshot.workspace, key.rsplit(':', 1)[-1]))
        except (OSError, sqlite3.Error):
            return False

    @contextmanager
    def visible_chart_work(workspace: str):
        with visible_chart_condition:
            visible_chart_requests[workspace] = visible_chart_requests.get(workspace, 0) + 1
        try:
            yield
        finally:
            with visible_chart_condition:
                visible_chart_requests[workspace] = max(0, visible_chart_requests.get(workspace, 1) - 1)
                visible_chart_condition.notify_all()

    def wait_for_visible_charts(workspace: str) -> None:
        """Hold a preload while the charts of the open slide are being calculated."""
        with visible_chart_condition:
            while visible_chart_requests.get(workspace, 0) and not core.APP_SHUTTING_DOWN.is_set():
                visible_chart_condition.wait(timeout=0.5)

    @app.get('/api/ppt-designer/chart/{token}/{index}')
    def interactive_chart(token: str, index: int, priority: str = 'high', user=Depends(design_user)):
        with lock:
            snapshot = snapshots.get(token)
            if snapshot is not None:
                snapshot.viewed_at = monotonic()
        workspace = workspace_key()
        preload = priority == 'low'
        if preload and snapshot is not None and not chart_model_ready(snapshot, index):
            # Preloaded neighbours wait for the slide the user is looking at.
            wait_for_visible_charts(workspace)
        with nullcontext() if preload else visible_chart_work(workspace):
            payload = chart_model(token, index, user)
        # The server keeps the chart models in memory. A browser cache could
        # outlive a renderer change or an explicit Design refresh and
        # reintroduce an obsolete ordering/model.
        return JSONResponse(payload, headers={'Cache-Control': 'no-store'})

    @app.post('/api/ppt-designer/chart/{token}/{index}/refresh')
    def refresh_interactive_chart(token: str, index: int, user=Depends(design_user)):
        """Invalidate and rebuild only the requested Canvas chart model."""
        payload = chart_model(token, index, user, force=True)
        return JSONResponse(payload, headers={'Cache-Control': 'no-store'})

    @app.post('/api/ppt-designer/charts/{token}/refresh')
    def refresh_interactive_charts(token: str, user=Depends(design_user)):
        """Invalidate and rebuild every available Canvas model in a PPT snapshot."""
        with lock:
            snapshot = snapshots.get(token)
        workspace = workspace_key()
        if snapshot is None or snapshot.workspace != workspace or snapshot.owner not in {user.username, '*'}:
            raise HTTPException(410, 'PPT preview expired. Refresh the Design.')
        indexes = list(dict.fromkeys(
            int(chart['index'])
            for slide in snapshot.payload.get('slides', [])
            for chart in slide.get('charts', [])
            if chart.get('available')
        ))
        with ThreadPoolExecutor(
            max_workers=DESIGN_CHART_RENDER_WORKERS,
            thread_name_prefix='ppt-designer-refresh',
        ) as chart_executor:
            futures = [
                chart_executor.submit(
                    chart_model, token, index, user, expected_workspace=workspace, force=True,
                )
                for index in indexes
            ]
            for future in as_completed(futures):
                future.result()
        core.invalidate_workspace_size_cache(Path(workspace).parent)
        return {'refreshed': len(indexes)}

    @app.get('/api/ppt-designer/chart/{token}/{index}/filter-context')
    def interactive_chart_filter_context(token: str, index: int, user=Depends(design_user)):
        """Return filter fields available to an expanded PPT chart."""
        snapshot, entry, _ = snapshot_chart(token, index, user, include_frame=False)
        task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
        _database_path, _table_name, columns = reporting_source(snapshot, entry.source_kind, task_repository)
        hidden = {identity('dataset_id'), identity('source_row_id')}
        datasets_by_source, columns_by_source = source_columns(task_repository, columns)
        template_available = user.role in {'user-editor', 'admin', 'super-admin'} and any(
            str(row['name']) == snapshot.definition.template
            for row in task_repository.list_report_templates()
        )
        return JSONResponse({
            'cdr_source': entry.cdr_source,
            'chart_type': entry.chart_type,
            'chart_title': entry.chart_title,
            'kpi': entry.kpi,
            'dataset_ids': [str(row['id']) for row in design_datasets(snapshot.definition, entry.source_kind, task_repository)],
            'filters': entry.filters,
            'rows_split': entry.rows_split,
            'columns_split': entry.columns_split,
            'legend': entry.legend,
            'legend_position': entry.legend_position,
            'legend_format': entry.legend_format,
            'axis_x_range': entry.axis_x_range, 'axis_y_range': entry.axis_y_range,
            'label_position': entry.label_position,
            'label_format': entry.label_format,
            'exclude_null_empty': entry.exclude_null_empty, 'exclude_zero': entry.exclude_zero, 'dynamic_rows_field': entry_dynamic_fields(entry)[0], 'dynamic_columns_field': entry_dynamic_fields(entry)[1],
            'template_available': template_available,
            'datasets_by_source': datasets_by_source, 'columns_by_source': columns_by_source,
            'columns': columns_by_source.get(f'cdr-{entry.source_kind}', [str(column) for column in columns if identity(column) not in hidden]),
        })

    def source_columns(task_repository, all_columns):
        """The ready CDRs and the fields of each source (cdr-data, cdr-voice, cdr-speech and cdr-all, whose fields are
        ``all_columns``), without the internal row identities."""
        hidden = {identity('dataset_id'), identity('source_row_id')}
        datasets_by_source = {f'cdr-{kind}': [] for kind in KINDS}
        columns_by_source = {f'cdr-{kind}': [] for kind in KINDS}
        for row in task_repository.list_datasets():
            kind = str(row['dataset_kind'] or '').casefold()
            if kind in KINDS and row['status'] == 'ready':
                datasets_by_source[f'cdr-{kind}'].append({
                    'value': str(row['id']), 'label': str(row['file_name']),
                })
                columns_by_source[f'cdr-{kind}'].extend(
                    str(column) for column in task_repository.list_dataset_row_columns(int(row['id']))
                )
        for source, source_columns_list in columns_by_source.items():
            columns_by_source[source] = sorted({
                column for column in source_columns_list if identity(column) not in hidden
            }, key=str.casefold)
        datasets_by_source['cdr-all'] = [row for kind in KINDS for row in datasets_by_source[f'cdr-{kind}']]
        columns_by_source['cdr-all'] = [str(column) for column in all_columns if identity(column) not in hidden]
        return datasets_by_source, columns_by_source

    def build_interactive_chart_filter_preview(
        token: str,
        index: int,
        request: DesignChartFilterPreviewRequest,
        user,
    ):
        """Build one expanded chart with temporary, unsaved template filters."""
        snapshot, entry, _ = snapshot_chart(token, index, user, include_frame=False)
        changes = {
            key: value for key, value in request.model_dump().items()
            if value is not None and key in {
                'filters', 'chart_title', 'cdr_source', 'kpi', 'chart_type', 'rows_split',
                'columns_split', 'legend', 'legend_position', 'legend_format', 'axis_x_range', 'axis_y_range', 'label_position', 'label_format', 'exclude_null_empty', 'exclude_zero', 'dynamic_rows_field', 'dynamic_columns_field', 'dynamic_field',
            }
        }
        # The template owns these required chart attributes. Custom dropdowns
        # can briefly report an empty value while their available fields are
        # being rebuilt, which must not erase the template KPI/source/type.
        for key in ('cdr_source', 'kpi', 'chart_type'):
            if not str(changes.get(key, '')).strip():
                changes.pop(key, None)
        if 'label_position' in changes:
            try:
                changes['label_position'] = core.parse_label_position(str(changes['label_position']))
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        if 'label_format' in changes:
            try:
                changes['label_format'] = core.parse_label_format(str(changes['label_format']))
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        if 'legend_format' in changes:
            try:
                changes['legend_format'] = core.parse_label_format(str(changes['legend_format']))
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        for key, label in (('exclude_null_empty', 'Exclude Null/Empty'), ('exclude_zero', 'Exclude Zero')):
            if key in changes:
                try:
                    changes[key] = core.parse_template_boolean(changes[key], label)
                except ValueError as exc:
                    raise HTTPException(400, str(exc)) from exc
        preview_entry = replace(entry, **changes)
        try:
            core.parse_axis_range(preview_entry.axis_x_range, 'x')
            core.parse_axis_range(preview_entry.axis_y_range, 'y')
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        if not preview_entry.source_kind:
            raise HTTPException(400, 'Select a valid CDR type.')
        task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
        preview_definition = snapshot.definition.model_copy(deep=True)
        if request.dataset_ids is not None:
            values = request.dataset_ids if isinstance(request.dataset_ids, list) else str(request.dataset_ids).split(',')
            try:
                ids = list(dict.fromkeys(int(str(value).strip()) for value in values if str(value).strip()))
                if preview_entry.source_kind == 'all':
                    for kind in KINDS:
                        preview_definition.datasets[kind] = [value for value in ids if (task_repository.get_dataset(value) and str(task_repository.get_dataset(value)['dataset_kind']).casefold() == kind)]
                else:
                    preview_definition.datasets[preview_entry.source_kind] = ids
            except ValueError as exc:
                raise HTTPException(400, 'Selected datasets must have valid identifiers.') from exc
        # An expanded chart may introduce any available field. Materialize it
        # in the shared combined table before querying that table directly.
        preview_snapshot = replace(
            snapshot, entries=[preview_entry], definition=preview_definition,
            frames={}, chart_frames={}, filtered_frames={}, chart_payloads={}, frame_locks={},
        )
        selected = design_datasets(preview_definition, preview_entry.source_kind, task_repository)
        ensure_combined_filter_columns(
            preview_definition, task_repository, snapshot.dimensions,
            selected_sources(preview_definition, task_repository), [preview_entry],
        )
        database_path, table_name, source_columns = reporting_source(
            preview_snapshot, preview_entry.source_kind, task_repository,
        )
        where, parameters = selection_where(
            task_repository, preview_entry.source_kind, [int(row['id']) for row in selected],
            preview_definition, columns=source_columns,
        )
        requested_columns = chart_query_columns(preview_entry, snapshot.multivendor)
        template_where, template_parameters, template_filters_applied = chart_filter_sql(
            preview_entry, source_columns, snapshot.multivendor, task_repository,
        )
        if template_where:
            where = f'({where}) AND ({template_where})'
            parameters.extend(template_parameters)
        aggregation_columns = chart_aggregation_columns(
            preview_entry, source_columns, snapshot.multivendor, template_filters_applied,
        )
        raw_frame = load_reporting_frame(
            database_path, table_name, source_columns, requested_columns, where, parameters,
            aggregation_columns,
        )
        if snapshot.multivendor:
            raw_frame = ensure_vendor_group(raw_frame)
        mapping_settings = task_repository.chart_mapping_settings()
        operator_mappings = mapping_settings['operator_mappings']
        raw_frame = apply_operator_mappings(raw_frame, operator_mappings)
        raw_frame = drop_excluded_operators(raw_frame, excluded_operators(mapping_settings))
        raw_frame.attrs.update(mapping_settings)
        raw_frame = normalise_operator_aliases(raw_frame)
        raw_frame.attrs['operator_aliases_normalized'] = True
        try:
            frame, _ = prepare_catalog_chart_preview_frame(
                raw_frame, preview_entry, multivendor=snapshot.multivendor,
                template_filters_applied=template_filters_applied,
            )
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        return catalog_chart_payload(
            frame, preview_entry, multivendor=snapshot.multivendor, prefiltered=True,
        )

    def run_interactive_chart_filter_preview(
        job_id: str,
        token: str,
        index: int,
        request: DesignChartFilterPreviewRequest,
        user,
    ) -> None:
        """Finish a potentially slow preview without holding a proxy request open."""
        try:
            payload = build_interactive_chart_filter_preview(token, index, request, user)
            result = {'state': 'ready', 'payload': payload}
        except HTTPException as exc:
            result = {'state': 'failed', 'status': exc.status_code, 'detail': str(exc.detail)}
        except Exception as exc:
            result = {'state': 'failed', 'status': 500, 'detail': str(exc)}
        with lock:
            job = interactive_chart_preview_jobs.get(job_id)
            if job is not None:
                job.update(result)

    @app.post('/api/ppt-designer/chart/{token}/{index}/filter-preview')
    def interactive_chart_filter_preview(
        token: str,
        index: int,
        request: DesignChartFilterPreviewRequest,
        background_tasks: BackgroundTasks,
        background: bool = False,
        user=Depends(design_user),
    ):
        """Render one temporary chart, optionally outside the proxy request."""
        if not background:
            return JSONResponse(
                build_interactive_chart_filter_preview(token, index, request, user),
                headers={'Cache-Control': 'no-store'},
            )
        # Validate access before returning a job identifier. The worker repeats
        # the check so an expired snapshot still becomes an actionable error.
        snapshot_chart(token, index, user, include_frame=False)
        job_id = uuid4().hex
        with lock:
            interactive_chart_preview_jobs[job_id] = {
                'state': 'processing', 'owner': user.username, 'workspace': workspace_key(),
            }
            while len(interactive_chart_preview_jobs) > 32:
                removable = next((
                    key for key, value in interactive_chart_preview_jobs.items()
                    if value.get('state') != 'processing'
                ), None)
                if removable is None:
                    break
                interactive_chart_preview_jobs.pop(removable, None)
        background_tasks.add_task(
            run_interactive_chart_filter_preview, job_id, token, index, request, user,
        )
        return JSONResponse(
            {'state': 'processing', 'job_id': job_id}, status_code=202,
            headers={'Cache-Control': 'no-store'},
        )

    @app.get('/api/ppt-designer/chart-preview-jobs/{job_id}')
    def interactive_chart_filter_preview_status(job_id: str, user=Depends(design_user)):
        """Return a temporary Chart Definition preview once its query completes."""
        with lock:
            job = interactive_chart_preview_jobs.get(job_id)
            if (
                job is None
                or job.get('owner') != user.username
                or job.get('workspace') != workspace_key()
            ):
                raise HTTPException(404, 'Chart preview job not found.')
            result = dict(job)
        if result.get('state') == 'processing':
            return JSONResponse(
                {'state': 'processing', 'job_id': job_id}, status_code=202,
                headers={'Cache-Control': 'no-store'},
            )
        if result.get('state') == 'failed':
            raise HTTPException(int(result.get('status') or 500), str(result.get('detail') or 'Unable to render chart preview.'))
        return JSONResponse(result['payload'], headers={'Cache-Control': 'no-store'})

    @app.post('/api/ppt-designer/chart/{token}/{index}/update-template')
    def update_interactive_chart_template(
        token: str,
        index: int,
        request: DesignChartFilterPreviewRequest,
        user=Depends(design_editor_user),
    ):
        """Persist the expanded Chart Definition into its source template row."""
        snapshot, _entry, _ = snapshot_chart(token, index, user, include_frame=False)
        task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
        template_name, nr_mode = snapshot.definition.template, design_nr_mode(snapshot.definition)
        template = next((
            row for row in task_repository.list_report_templates()
            if str(row['name']) == template_name
        ), None)
        if template is None:
            raise HTTPException(404, 'The Report Template used by this chart is no longer available.')
        if not can_edit_template_slides(user, template):
            raise HTTPException(403, designer_refusal('PPT Template'))
        # The template as written (its titles may name the NR Mode with tokens, which the Design shows resolved).
        try:
            entries = core.load_template_catalogue(bytes(template['content']), task_repository=task_repository)
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        source_index = _entry.template_index if _entry.template_index is not None else index
        element_id = design_slides.row_format(_entry).get('id')
        if element_id and (snapshot.definition.own_slides or design_slides.has_overrides(snapshot.definition.slide_overrides)):
            # A Design with a slide set of its own: its rows are not those of the template, the ids are.
            if snapshot.definition.own_slides:
                raise HTTPException(409, 'This Design is detached from its template: change the chart from the Designer.')
            source_index = next((position for position, entry in enumerate(entries)
                                 if design_slides.row_format(entry).get('id') == element_id), None)
            if source_index is None:
                raise HTTPException(409, 'This chart is only in this Design: change it from the Designer.')
        if source_index >= len(entries):
            raise HTTPException(404, 'The chart row is no longer available in the Report Template.')
        changes = {
            key: value for key, value in request.model_dump().items()
            if value is not None and key in {
                'filters', 'chart_title', 'cdr_source', 'kpi', 'chart_type', 'rows_split',
                'columns_split', 'legend', 'legend_position', 'legend_format', 'axis_x_range', 'axis_y_range', 'label_position', 'label_format', 'exclude_null_empty', 'exclude_zero', 'dynamic_rows_field', 'dynamic_columns_field', 'dynamic_field',
            }
        }
        for key in ('cdr_source', 'kpi', 'chart_type'):
            if not str(changes.get(key, '')).strip():
                changes.pop(key, None)
        if 'label_position' in changes:
            try:
                changes['label_position'] = core.parse_label_position(str(changes['label_position']))
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        if 'label_format' in changes:
            try:
                changes['label_format'] = core.parse_label_format(str(changes['label_format']))
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        if 'legend_format' in changes:
            try:
                changes['legend_format'] = core.parse_label_format(str(changes['legend_format']))
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
        for key, label in (('exclude_null_empty', 'Exclude Null/Empty'), ('exclude_zero', 'Exclude Zero')):
            if key in changes:
                try:
                    changes[key] = core.parse_template_boolean(changes[key], label)
                except ValueError as exc:
                    raise HTTPException(400, str(exc)) from exc
        if changes.get("chart_title") in {_entry.chart_title, entry_for_nr_mode(entries[source_index], nr_mode).chart_title}:
            # The title as the chart shows it (dynamic values, the NR Mode resolved): the template keeps its own.
            changes["chart_title"] = entries[source_index].chart_title
        for field in ('dynamic_rows_field', 'dynamic_columns_field', 'dynamic_field', 'rows_split', 'columns_split', 'legend'):
            if field in changes and changes[field] == getattr(_entry, field):
                changes[field] = getattr(entries[source_index], field)
        updated_entry = replace(entries[source_index], **changes)
        try:
            core.parse_axis_range(updated_entry.axis_x_range, 'x')
            core.parse_axis_range(updated_entry.axis_y_range, 'y')
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        entries[source_index] = updated_entry
        try:
            task_repository.set_report_template_content(
                template_name, core.catalogue_csv(entries), updated_by=user.username,
            )
        except (OSError, ValueError) as exc:
            raise HTTPException(503, f'Unable to update the Report Template: {exc}') from exc
        # Keep every live view of this Report Template consistent with the row
        # just persisted. The Design viewer, an expanded chart and a
        # restored/PPT chart can legitimately use different snapshot tokens;
        # updating only the request's snapshot makes another one reopen the
        # previous Chart Definition even though the CSV was saved correctly.
        with lock:
            matching_snapshots = [
                candidate for candidate in snapshots.values()
                if (
                    candidate.workspace == snapshot.workspace
                    and candidate.definition.template == template_name
                )
            ]
            for candidate in matching_snapshots:
                for chart_index, current in enumerate(candidate.entries):
                    current_source = current.template_index if current.template_index is not None else chart_index
                    if current_source != source_index:
                        continue
                    candidate_mode = design_nr_mode(candidate.definition)
                    shown_update = entry_for_nr_mode(updated_entry, candidate_mode)
                    effective_update = (prepare_multivendor_catalog_entry(shown_update, candidate.definition.vendor_comparison)
                                        if candidate.multivendor else shown_update)
                    replacement = replace(
                        effective_update, slide=current.slide, layout=current.layout, template_index=current.template_index,
                        dynamic_value=current.dynamic_value,
                        dynamic_row_value=current.dynamic_row_value, dynamic_column_value=current.dynamic_column_value,
                    )
                    replacement = replace(replacement, chart_title=dynamic_chart_title(replacement))
                    candidate.entries[chart_index] = replacement
                    candidate.chart_frames.pop(chart_index, None)
                    candidate.chart_frames.pop(('complete', chart_index), None)
                    candidate.chart_payloads.pop(chart_index, None)
                    candidate.filtered_frames.clear()
                    for slide in candidate.payload.get('slides', []):
                        for chart in slide.get('charts', []):
                            if chart.get('index') == chart_index:
                                chart.update({'title': replacement.chart_title, 'source': replacement.source_kind,
                                              'cdr_source': replacement.cdr_source, 'chart_type': replacement.chart_type})
        task_repository.add_log(user.username, 'update_design_chart_template', json.dumps({
            'template': template_name, 'nr_mode': nr_mode, 'chart_index': index,
        }))
        return {
            'template': template_name, 'chart_index': index,
            'synced_snapshots': len(matching_snapshots),
            'updated_at': str(next((
                row['updated_at'] for row in task_repository.list_report_templates()
                if str(row['name']) == template_name
            ), '')),
        }

    def design_status_signature(workspace, raw_definition, task_repository, shared) -> str:
        """Fingerprint every input that can change a Design's cached universes."""
        definition = raw_definition if isinstance(raw_definition, dict) else {}
        template = str(definition.get('template') or '')
        template_key = template
        if template_key not in shared['templates']:
            try:
                content = task_repository.report_template_content(template)
                shared['templates'][template_key] = sha256(content).hexdigest()
            except Exception:
                shared['templates'][template_key] = ''
        return sha256(json.dumps({
            # The identity of the template is not an input: its content is.
            'definition': {key: value for key, value in definition.items() if key != 'template_uid'},
            'template': shared['templates'][template_key],
            'shared': shared['signature'],
        }, sort_keys=True, default=str).encode()).hexdigest()

    def design_status_shared_state(workspace, task_repository) -> dict:
        datasets = [
            (row['id'], row['status'], row['dataset_kind'], row['nr_mode'], row['updated_at'], row['processed_at'])
            for row in task_repository.list_datasets()
        ]
        data_signature = sha256(json.dumps({
            'datasets': datasets,
            'revisions': {kind: task_repository.get_workspace_state(f'combined_reporting_updated_{kind}') for kind in KINDS},
            'dimension_definitions': core.calculated_dimensions_json(core.load_repository_calculated_dimensions(task_repository)),
            'mappings': task_repository.chart_mapping_settings(),
        }, sort_keys=True, default=str).encode()).hexdigest()
        return {'signature': data_signature, 'data_signature': data_signature, 'templates': {}}

    def design_status_payload(requested_definitions=None, username='system', active_design_id=None):
        """The state of each Design: being prepared or rendered for a user, Ready when it has CDRs to show (it is
        prepared when opened, from the columnar copies of its CDRs), or Data needed."""
        workspace = workspace_key()
        task_repository = Repository(Path(workspace), core.repository.global_db_path)
        designs = read_designs(task_repository)
        result = {}
        with lock:
            direct_jobs = {}
            for design_id in designs:
                direct_jobs[design_id] = next((
                    dict(task) for task in direct_preparation_tasks.values()
                    if task.get('workspace') == workspace and task.get('design_id') == design_id
                ), None)
        try:
            shared = design_status_shared_state(workspace, task_repository)
        except (OSError, sqlite3.Error, TypeError, ValueError):
            shared = {'signature': uuid4().hex, 'data_signature': '', 'templates': {}}
        for design_id, raw_definition in designs.items():
            direct_job = direct_jobs.get(design_id)
            if direct_job is not None:
                rendering_only = bool(direct_job.get('rendering_only'))
                result[design_id] = {
                    'state': 'rendering' if rendering_only else 'loading-data',
                    'label': 'Rendering' if rendering_only else 'Loading data',
                }
                continue
            key = (workspace, design_id)
            signature = design_status_signature(workspace, raw_definition, task_repository, shared)
            with lock:
                cached = design_status_cache.get(key)
            if cached is not None and cached[0] == signature:
                has_data = cached[1]
            else:
                try:
                    definition = DesignDefinition.model_validate(runtime_design_definition(raw_definition, task_repository))
                    has_data = any(definition.datasets.values())
                except (HTTPException, KeyError, OSError, sqlite3.Error, TypeError, ValueError):
                    has_data = False
                with lock:
                    design_status_cache[key] = (signature, has_data)
            result[design_id] = {'state': 'ready', 'label': 'Ready'} if has_data else {'state': 'data-needed', 'label': 'Data needed'}
        return JSONResponse(result, headers={'Cache-Control': 'no-store, max-age=0, must-revalidate'})

    @app.get('/api/ppt-designer/statuses')
    def design_statuses(active_design_id: str | None = None, user=Depends(design_user)):
        return design_status_payload(username=user.username, active_design_id=active_design_id)

    @app.post('/api/ppt-designer/statuses')
    def design_statuses_for_session(
        definitions: dict[str, DesignDefinition], active_design_id: str | None = None,
        user=Depends(design_user),
    ):
        """Report the states of the Designs."""
        return design_status_payload(definitions, username=user.username, active_design_id=active_design_id)

    def cancel_workspace_design_tasks(workspace: str | Path) -> list:
        """Cancel explicitly requested PPT preparation for a workspace."""
        database_path = str(Path(workspace).resolve())
        with lock:
            for task in direct_preparation_tasks.values():
                if task.get('workspace') == database_path:
                    cancellation = task.get('cancellation')
                    if isinstance(cancellation, dict):
                        cancellation['requested'] = True
            stale_tokens = [token for token, snapshot in snapshots.items() if snapshot.workspace == database_path]
            for token in stale_tokens:
                snapshots.pop(token, None)
            column_cache = column_caches.pop(database_path, None)
        if column_cache is not None:
            column_cache.forget()
        return []

    def invalidate_workspace_chart_models(workspace: str | Path) -> None:
        """Rebuild open PPT charts after a mapping colour/order change."""
        database_path = str(Path(workspace).resolve())
        with lock:
            for snapshot in snapshots.values():
                if snapshot.workspace != database_path:
                    continue
                # Raw/filtered frames retain the mapping aliases and their
                # grouping order. Keep the prepared selection and slide token,
                # but rebuild all mapping-dependent chart data and Canvas
                # payloads on the next request.
                snapshot.frames.clear()
                snapshot.filtered_frames.clear()
                snapshot.chart_frames.clear()
                snapshot.chart_payloads.clear()
                snapshot.frame_locks.clear()

    def design_task_payloads(workspace):
        database_path = str(workspace.database_path.resolve())
        with lock:
            direct_tasks = [
                task for task in direct_preparation_tasks.values()
                if task['workspace'] == database_path
                and not bool((task.get('cancellation') or {}).get('requested'))
            ]
            return [{
                'id': task['id'],
                'design_name': task['name'],
                'label': (
                    'Waiting to render PPT Charts' if task.get('rendering_only') else 'Waiting to prepare PPT dataset'
                ) if task.get('status') == 'queued' else (
                    'Rendering PPT Charts' if task.get('rendering_only') else 'Preparing PPT dataset'
                ),
                'detail': (
                    'Rendering charts with the current PPT scope'
                    if task.get('rendering_only')
                    else str(task.get('detail') or 'Building filtered PPT selection')
                ),
                'status': task.get('status', 'processing'),
                'queued_at': task.get('queued_at'),
                'started_at': task.get('started_at'),
                'progress': task.get('progress', 0),
                'stop_task_id': f'design-prepare:{task["id"]}',
                'stop_url': f'/api/background-tasks/{workspace.id}/stop',
            } for task in direct_tasks]

    def stop_design_task(workspace_path, task_id):
        database_path = str(Path(workspace_path).resolve())
        stopped = False
        with lock:
            direct = direct_preparation_tasks.get(task_id)
            if direct and direct.get('workspace') == database_path:
                cancellation = direct.get('cancellation')
                if isinstance(cancellation, dict):
                    cancellation['requested'] = True
                    stopped = True
        return stopped

    def rename_template_designs(previous_template: str, template: str, username: str) -> int:
        """Keep Design template references and display names aligned after a template rename."""
        workspace_key()
        with lock:
            task_repository = bound_repository()
            designs = read_designs(task_repository)
            changed = [(design_id, designs[design_id])
                       for design_id in rename_template_in_designs(designs, previous_template, template)]
            if changed:
                task_repository.set_workspace_state(STATE_KEY, json.dumps(designs))
                task_repository.add_log(username, 'rename_design_template_references', json.dumps({
                    'previous_template': previous_template,
                    'template': template, 'designs': len(changed),
                }))
        return len(changed)

    # ------------------------------------------------------------------
    # Warm-up: while nobody uses the application, every PPT Design and PPT Template of the open workspace is prepared
    # and the models of its charts are worked out and kept (in its cache folder, ppt_design_cache), so after a
    # restart they open at once. A chart keeps its model while it, its CDR selection and the mappings stay the same (its
    # key), so a changed Design or template only works out its changed charts, and a changed selection all of them. It
    # pauses as soon as someone uses the application or other background work starts, stops from its card, and goes on
    # from where it was the next time the server is idle (what it kept is not worked out again).
    # ------------------------------------------------------------------
    warmups: dict[str, dict] = {}
    warmup_user = SimpleNamespace(username='system', role='super-admin')

    def warmup_signature(task_repository) -> str:
        """What the warm-up reads: the Designs, the templates and the template previews, the CDRs, the mappings and the
        chart models version. Unchanged since a finished warm-up, there is nothing to work out."""
        with task_repository.connection() as connection:
            datasets = [tuple(row) for row in connection.execute(
                'SELECT dataset_id, status, updated_at, processed_at, row_count, in_combined FROM dataset_profiles '
                'ORDER BY dataset_id').fetchall()]
        templates = [(str(row['name']), sha256(bytes(row['content'] or b'')).hexdigest())
                     for row in task_repository.list_report_templates()]
        return sha256(json.dumps([
            DESIGN_CHART_MODEL_CACHE_VERSION, DESIGN_PREVIEW_MANIFEST_VERSION,
            task_repository.get_workspace_state(STATE_KEY), task_repository.get_workspace_state(TEMPLATE_PREVIEWS_STATE_KEY),
            datasets, templates, task_repository.chart_mapping_settings(),
            [task_repository.get_workspace_state(f'combined_reporting_updated_{kind}') for kind in KINDS],
        ], sort_keys=True, default=str).encode()).hexdigest()

    def warmup_targets(task_repository) -> list[tuple[str, str | None, dict]]:
        """(name, design id, definition) of every PPT Design, then of every PPT Template as each user who opened one
        opens it (their NR Mode and Campaign), or with the defaults of PPT Templates → Edit when nobody has."""
        targets = [(f"PPT Design {raw.get('name') or design_id}", design_id, raw)
                   for design_id, raw in sorted(read_designs(task_repository).items(), key=lambda item: str(item[1].get('name') or '').casefold())
                   if isinstance(raw, dict)]
        previews = [(design_id, raw) for design_id, raw in read_template_previews(task_repository).items() if isinstance(raw, dict)]
        default = {'nr_mode': 'NSA', 'scope': 'single', 'auto_cdrs': True, 'filters': {'Campaign': ['@latest-2-campaigns']},
                   'sa_campaigns': [], 'custom_fields': [], 'hidden_filters': [], 'date_from': 'Oldest', 'date_to': 'Newest'}
        for row in sorted(task_repository.list_report_templates(), key=lambda row: str(row['name']).casefold()):
            name = str(row['name'])
            for design_id, preview in previews or [(None, default)]:
                definition = {**preview, 'name': name, 'template': name, 'slide_overrides': {}, 'own_slides': None,
                              'slide_comments': {}, 'slide_comment_meta': {}}
                targets.append((f'PPT Template {name}', design_id, definition))
        return targets

    def warmup_task(workspace_path: str) -> dict | None:
        with lock:
            job = warmups.get(workspace_path)
            return dict(job) if job else None

    def warm_up_target(job, workspace, task_repository, design_id, raw, should_pause) -> None:
        """Prepare one Design or template (or restore its kept preparation) and work out the charts it lacks."""
        definition = auto_cdr_definition(DesignDefinition.model_validate(raw), task_repository)
        with lock:
            known = set(snapshots)
        preview = restore_matching_preview_manifest(workspace, design_id, definition) if design_id else None
        if preview is None:
            while not design_work_gate.acquire(timeout=0.5):
                if should_pause():
                    return
            try:
                preview = build_preview(definition, warmup_user, workspace=workspace, cancelled=should_pause)
            finally:
                design_work_gate.release()
            if design_id:
                with lock:
                    built = snapshots.get(preview['token'])
                if built is not None:
                    persist_preview_manifest(workspace, design_id, design_preview_fingerprint(
                        built.definition.model_dump(mode='json'), task_repository), preview['token'])
                    with lock:
                        snapshots.pop(preview['token'], None)
                    # The charts are worked out on the kept preparation, as a user opening the Design reads it.
                    preview = restore_matching_preview_manifest(workspace, design_id, definition) or preview
        token = str(preview['token'])
        try:
            with lock:
                snapshot = snapshots.get(token)
            if snapshot is None:
                return
            indexes = [index for index, entry in enumerate(snapshot.entries)
                       if entry.source_kind and design_datasets(snapshot.definition, entry.source_kind, task_repository)]
            with lock:
                job.update(charts=len(indexes), chart=0)
            for position, index in enumerate(indexes):
                if should_pause():
                    return
                with lock:
                    job['chart'] = position + 1
                if chart_model_ready(snapshot, index):
                    with lock:
                        job['reused'] += 1
                    continue
                try:
                    chart_model(token, index, warmup_user, expected_workspace=workspace)
                    with lock:
                        job['computed'] += 1
                except (HTTPException, ValueError, KeyError, OSError, sqlite3.Error):
                    # A chart its selection cannot draw is left to the viewer, which says why.
                    with lock:
                        job['failed'] += 1
        finally:
            # The warm-up keeps the models, not the frames: its own preparations leave the memory.
            if token not in known:
                with lock:
                    snapshots.pop(token, None)

    def run_warmup(workspace: str, activity_at: float) -> None:
        task_repository = Repository(Path(workspace), core.repository.global_db_path)
        with lock:
            job = warmups[workspace]

        def should_pause() -> bool:
            with core.APPLICATION_ACTIVITY_LOCK:
                used = core.LAST_INTERACTIVE_APPLICATION_ACTIVITY != activity_at
            active = core.active_workspace and str(Path(core.active_workspace.database_path).resolve()) == workspace
            reason = ('stopped' if job.get('stop_requested') else 'in use' if used
                      else 'background work' if not (core.BACKGROUND_TASK_SCHEDULER.is_idle and core.EXPORT_TASK_SCHEDULER.is_idle)
                      else 'shutdown' if core.APP_SHUTTING_DOWN.is_set() or not active else '')
            if reason:
                with lock:
                    job['pause_reason'] = job.get('pause_reason') or reason
            return bool(reason)

        started = monotonic()
        signature = ''
        try:
            signature = warmup_signature(task_repository)
            targets = warmup_targets(task_repository)
            with lock:
                job.update(status='processing', started_at=datetime.now(timezone.utc).timestamp(), total=len(targets))
            for number, (name, design_id, raw) in enumerate(targets, 1):
                if should_pause():
                    break
                with lock:
                    job.update(done=number - 1, current=name, charts=0, chart=0)
                try:
                    warm_up_target(job, workspace, task_repository, design_id, raw, should_pause)
                except RuntimeError:
                    break  # Its preparation was cancelled: someone uses the application.
                except (HTTPException, ValueError, KeyError, TypeError) as exc:
                    # A Design without CDRs of its selection, or with a template that is gone: the next one.
                    with lock:
                        job['skipped'] += 1
                        job['last_error'] = f'{name}: {getattr(exc, "detail", exc)}'
            else:
                with lock:
                    job.update(done=len(targets), signature=signature)
            reason = job.get('pause_reason') or ''
            with lock:
                job.update(status='stopped' if reason == 'stopped' else 'paused' if reason else 'ready',
                           finished_at=datetime.now(timezone.utc).timestamp(), current='' if not reason else job['current'])
        except Exception as exc:  # Warming up is optional: an idle period tries again (see warm_up_when_idle).
            with lock:
                job.update(status='failed', last_error=str(exc), finished_at=datetime.now(timezone.utc).timestamp(),
                           failed_signature=signature)
        cache_files, cache_bytes = ppt_design_cache.usage(workspace)
        task_repository.try_add_log('system', 'design_warmup_finished', json.dumps({
            'status': job['status'], 'reason': job.get('pause_reason') or '', 'targets': job.get('total', 0),
            'done': job.get('done', 0), 'charts_computed': job['computed'], 'charts_kept': job['reused'],
            'charts_failed': job['failed'], 'skipped': job['skipped'], 'seconds': round(monotonic() - started, 1),
            'cache_files': cache_files, 'cache_mb': round(cache_bytes / 1024 ** 2, 1),
            **({'last_error': job['last_error']} if job.get('last_error') else {}),
        }))

    def warm_up_when_idle() -> None:
        """Start the warm-up of the open workspace while the server is idle, unless it runs, or nothing changed since
        the last one finished."""
        workspace_entry = core.active_workspace
        if not workspace_entry or core.APP_SHUTTING_DOWN.is_set() or not DESIGN_WARMUP:
            return
        workspace = str(Path(workspace_entry.database_path).resolve())
        task_repository = Repository(Path(workspace), core.repository.global_db_path)
        with lock:
            previous = warmups.get(workspace) or {}
        if previous.get('status') in {'queued', 'processing'}:
            return
        try:
            signature = warmup_signature(task_repository)
        except (sqlite3.Error, OSError):
            return
        if previous.get('status') == 'ready' and previous.get('signature') == signature:
            return
        # After a failure it waits half an hour before trying again, unless something it reads changed.
        if (previous.get('status') == 'failed' and previous.get('failed_signature') == signature
                and datetime.now(timezone.utc).timestamp() - float(previous.get('finished_at') or 0) < 30 * 60):
            return
        with core.APPLICATION_ACTIVITY_LOCK:
            activity_at = core.LAST_INTERACTIVE_APPLICATION_ACTIVITY
        with lock:
            warmups[workspace] = {
                'id': uuid4().hex, 'status': 'queued', 'queued_at': datetime.now(timezone.utc).timestamp(),
                'started_at': None, 'finished_at': None, 'total': 0, 'done': 0, 'current': '', 'charts': 0, 'chart': 0,
                'computed': 0, 'reused': 0, 'failed': 0, 'skipped': 0, 'stop_requested': False, 'pause_reason': '',
                'last_error': '', 'signature': previous.get('signature') if previous.get('status') == 'ready' else '',
                'resumed': previous.get('status') in {'paused', 'stopped', 'failed'},
            }
        task_repository.try_add_log('system', 'design_warmup_started', json.dumps({
            'resumed': warmups[workspace]['resumed'], 'idle_seconds': round(core.application_idle_seconds()),
        }))
        Thread(target=run_warmup, args=(workspace, activity_at), name='design-warmup', daemon=True).start()

    def warmup_task_payloads(workspace) -> list[dict]:
        """The warm-up in the background tasks card while it runs (with Stop), and a moment after it pauses or ends."""
        job = warmup_task(str(Path(workspace.database_path).resolve()))
        if not job:
            return []
        running = job['status'] in {'queued', 'processing'}
        if not running and datetime.now(timezone.utc).timestamp() - float(job.get('finished_at') or 0) > 8:
            return []
        if job['status'] == 'queued':
            detail = 'Waiting to prepare the PPT Designs and PPT Templates'
        elif running:
            charts = f": chart {job['chart']} of {job['charts']}" if job.get('charts') else ': preparing its data'
            detail = f"{job['current']}{charts} ({job['done'] + 1} of {job['total']})"
        elif job['status'] == 'ready':
            detail = f"Ready: {job['computed']} charts worked out, {job['reused']} already kept"
        elif job['status'] == 'failed':
            detail = f"Failed: {job.get('last_error') or 'unknown error'}; it tries again when the server is idle"
        else:
            detail = {'stopped': 'Stopped', 'in use': 'Paused while the application is in use',
                      'background work': 'Paused while other background work runs',
                      'shutdown': 'Paused'}.get(job.get('pause_reason') or '', 'Paused') + \
                ': it goes on from where it was the next time the server is idle'
        fraction = (job['chart'] / job['charts']) if job.get('charts') else 0
        # (Finished, it stays a moment with what it did: ready shows as completed once it goes, paused, stopped and
        # failed do not.)
        return [{
            'id': f"design-warmup:{job['id']}", 'label': 'Warming up PPT Designs', 'detail': detail,
            'status': job['status'] if running or job['status'] in {'paused', 'stopped', 'failed'} else 'processing',
            'progress': 100 if job['status'] == 'ready' else (
                round(100 * (job['done'] + fraction) / job['total']) if job.get('total') else None),
            'queued_at': job.get('queued_at'), 'started_at': job.get('started_at'), 'completed_at': job.get('finished_at'),
            **({'stop_task_id': f"design-warmup:{job['id']}", 'stop_url': f'/api/background-tasks/{workspace.id}/stop'}
               if running else {}),
        }]

    def stop_warmup(workspace_path, job_id: str) -> bool:
        with lock:
            job = warmups.get(str(Path(workspace_path).resolve()))
            if not job or job.get('id') != job_id or job.get('status') not in {'queued', 'processing'}:
                return False
            job['stop_requested'] = True
        return True

    core.register_idle_task(warm_up_when_idle, DESIGN_WARMUP_IDLE_SECONDS)
    core.BACKGROUND_TASK_PROVIDERS.append(warmup_task_payloads)
    core.ppt_design_warmup_stop = stop_warmup
    core.ppt_design_warm_up_now = warm_up_when_idle
    core.ppt_design_warmup_task = warmup_task

    core.ppt_design_tasks = design_task_payloads
    core.ppt_design_cancel_workspace_tasks = cancel_workspace_design_tasks
    core.ppt_design_clear_workspace_caches = clear_workspace_design_caches
    core.ppt_design_move_stored_caches = move_stored_caches
    core.ppt_design_invalidate_chart_models = invalidate_workspace_chart_models
    core.ppt_design_stop_task = stop_design_task
    core.ppt_design_rename_template_references = rename_template_designs
    core.ppt_design_reconcile_template_slide_comments = reconcile_template_slide_comments

    @app.get('/api/ppt-designer/preview/{token}/{index}.png')
    def chart(token: str, index: int, user=Depends(design_user)):
        snapshot, entry, _ = snapshot_chart(token, index, user, include_frame=False)
        entry_key = sha256(repr(entry).encode()).hexdigest()
        task_repository = Repository(Path(snapshot.workspace), core.repository.global_db_path)
        operator_mapping_key = sha256(json.dumps(
            task_repository.chart_mapping_settings(), sort_keys=True,
        ).encode()).hexdigest()
        key = (snapshot.selection_key, snapshot.definition.scope, entry_key, operator_mapping_key)
        # PIL chart renderers have no shared global canvas. Let browser image
        # requests for the same slide render concurrently instead of placing
        # every chart behind one process-wide lock; only cache mutation needs
        # synchronization.
        with lock:
            png = images.get(key)
        if png is None:
            _, _, frame = snapshot_chart(token, index, user)
            try:
                rendered = render_catalog_chart_preview(frame, entry, multivendor=snapshot.multivendor, prefiltered=True)
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
            with lock:
                png = images.setdefault(key, rendered)
                images[key] = png
                while len(images) > 100:
                    images.popitem(last=False)
        # The same Design token can be explicitly refreshed. Never let a
        # browser keep a previously rendered chart image for an hour after the
        # server has rebuilt its model or mapping-dependent frame.
        return Response(png, media_type='image/png', headers={'Cache-Control': 'no-store, max-age=0, must-revalidate'})

    @app.get('/api/ppt-designer/data/{token}/{index}')
    def chart_data(
        token: str, index: int, page: int = 0, download: bool = False,
        column_filters: str = '', include_filter_values: bool = False, filter_column: str = '',
        user=Depends(design_user),
    ):
        snapshot, entry, _ = snapshot_chart(token, index, user, include_frame=False)
        filter_column = vendor_filter_column(filter_column) if filter_column else ''
        selected_column_filters = {vendor_filter_column(key): values for key, values in parse_chart_column_filters(column_filters).items()}
        if not download:
            projected_page = reporting_chart_data_page(
                snapshot, entry, page, selected_column_filters, include_filter_values, filter_column,
            )
            if projected_page is not None:
                visible, total, chart_total, filter_values = projected_page
                return {
                    'columns': list(visible.columns),
                    'rows': visible.fillna('').astype(str).values.tolist(),
                    'total': total, 'chart_total': chart_total,
                    'filter_values': filter_values, 'page': max(page, 0),
                    'column_classes': chart_dataset_column_classes(snapshot, entry, visible.columns),
                    'column_metadata': chart_dataset_column_metadata(snapshot, entry, visible.columns),
                    'page_size': 100, 'unfiltered_total': chart_total,
                }
        _, _, frame = snapshot_chart(token, index, user, complete=True)
        frame = unique_chart_dataset_rows(snapshot, entry, frame)
        chart_total = len(frame)
        values_frame = apply_chart_column_filters(
            frame, {column: values for column, values in selected_column_filters.items() if identity(column) != identity(filter_column)},
        )
        available_filter_values = chart_dataset_filter_values(values_frame) if filter_column else {}
        filter_values = next(
            (values for column, values in available_filter_values.items() if identity(column) == identity(filter_column)), []
        ) if filter_column else (chart_dataset_filter_values(frame) if include_filter_values else {})
        frame = apply_chart_column_filters(frame, selected_column_filters)
        if download:
            bound_repository().try_add_log(user.username, 'export_design_chart_data_csv', json.dumps({
                'design': snapshot.definition.name, 'chart_index': index, 'rows': len(frame),
            }))
            return Response(frame.to_csv(index=False), media_type='text/csv', headers={'Content-Disposition': 'attachment; filename="design-chart-data.csv"'})
        page = max(page, 0)
        visible = frame.iloc[page*100:(page+1)*100].fillna('').astype(str)
        return {
            'columns': list(visible.columns), 'rows': visible.values.tolist(),
            'total': len(frame), 'chart_total': chart_total,
            'filter_values': filter_values, 'page': page,
            'column_classes': chart_dataset_column_classes(snapshot, entry, visible.columns),
            'column_metadata': chart_dataset_column_metadata(snapshot, entry, visible.columns),
            'page_size': 100, 'unfiltered_total': chart_total,
        }
