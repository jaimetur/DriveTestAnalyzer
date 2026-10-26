"""Non-Qualified Calls report: Executive Summary, NQ Prioritization, Progress Status and Root Cause Analysis in PowerPoint;
its Word version and its section of a Bundle Report Template are composed from it (see
bundle_report_templates.module_document)."""

from datetime import datetime
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from src.modules.network_insights_export import _content_frame, _slide, _table_slides

FILTER_LABELS = (
    ('nr_mode', 'NR Mode'), ('datasets', 'CDRs'), ('service', 'Service'), ('campaign', 'Campaign'), ('operator', 'Operator'),
    ('operator_vendor', 'Operator_Vendor'), ('vendor_operator', 'Vendor_Operator'), ('vendor', 'Vendor'),
    ('region', 'Region'), ('cluster', 'Cluster'),
    ('city', 'City'), ('technology', 'Technology'), ('test_name', 'Test Name'), ('result', 'Result'),
    ('failure_classification', 'Failure Classification'), ('failure_category', 'Failure Category'),
    ('status', 'Status'), ('status_mode', 'Status Set'), ('team', 'Team'), ('assignee', 'Assignee'), ('root_domain', 'Root Domain'),
    ('root_category', 'Root Category'), ('root_cause', 'Root Cause'), ('call_type', 'Call Type'), ('period', 'Period'),
    ('age', 'Age of Open Calls'), ('root_pair', 'Root Domain · Cause'),
    ('effective_domain', 'Root Domain (with suggestions)'), ('effective_cause', 'Root Domain · Cause (with suggestions)'),
    ('node', 'eNB / gNB'), ('place', 'Repetition Place'), ('hotspot', 'Hotspot'), ('state', 'Follow-up'), ('rca_state', 'Root Cause'),
    ('version', 'CDR Version'),
)
STATE_LABELS = {
    'closed': 'Closed', 'attended': 'Attended', 'not_attended': 'Not attended', 'with_team': 'With team',
    'assigned': 'Assigned', 'commented': 'Commented', 'labelled': 'With a root cause', 'suggested': 'Suggested root cause',
}
FLAG_LABELS = {'open_only': 'Open calls only', 'without_comments': 'Without comments'}
GRANULARITY_LABELS = {'week': 'Week', 'month': 'Month', 'quarter': 'Quarter', 'year': 'Year'}
RASPBERRY = RGBColor(0xB0, 0x23, 0x4F)


VERSION_LABELS = {
    'current': 'Latest version', 'changed': 'Changed between CDRs', 'newer': 'Newer version in another CDR',
    'not_in_final': 'Not in the Final CDR', 'qualified': 'Completed in a newer CDR',
}


def selection_lines(filters: dict[str, Any], dataset_names: dict[str, str], granularity: str,
                    field_labels: dict[str, str] | None = None) -> list[str]:
    """Readable filters of the report; empty filters include every value."""
    lines = []
    for key, label in FILTER_LABELS:
        values = [str(value) for value in filters.get(key) or [] if str(value).strip()]
        if key == 'datasets':
            values = [dataset_names.get(value, value) for value in values]
        empty = 'No cause' if key == 'root_cause' else 'Not classified' if key in {'root_domain', 'root_category', 'effective_domain', 'root_pair'} \
            else 'Unassigned'
        values = [empty if value == '__unassigned__' else value for value in values]
        # "RF||Coverage" and "month:2026-07" read "RF · Coverage" and "2026-07".
        values = [value.partition(':')[2] if key == 'period' else value.strip('|').replace('||', ' · ') for value in values]
        values = [STATE_LABELS.get(value, value) if key in {'state', 'rca_state'} else value for value in values]
        values = [VERSION_LABELS.get(value, value) if key == 'version' else value for value in values]
        values = [{'auto': 'Automatic', 'manual': 'By hand'}.get(value, value) if key == 'status_mode' else value for value in values]
        # A hotspot is chosen by its key, which tells nothing to read.
        if key == 'hotspot' and values:
            values = [f"{len(values)} hotspot{'s' if len(values) > 1 else ''} of the NQ Prioritization"]
        if values:
            lines.append(f'{label}: {", ".join(values)}')
    for key, chosen in (filters.get('fields') or {}).items():
        values = ['Empty' if value == '__unassigned__' else str(value) for value in chosen or [] if str(value).strip()]
        if values:
            lines.append(f'{(field_labels or {}).get(key, key)}: {", ".join(values)}')
    # The optional and CDR columns of the table ("cdr:<name>" is the CDR column <name>).
    column_labels = {'join_id': 'JOIN_ID', 'nr_mode': 'NR Mode', 'cell_id': 'Cell ID', 'direction': 'Direction', 'end_time': 'End Time'}
    for key, chosen in (filters.get('columns') or {}).items():
        values = ['Empty' if value == '__unassigned__' else str(value) for value in chosen or [] if str(value).strip()]
        if values:
            lines.append(f'{column_labels.get(key, str(key).removeprefix("cdr:"))}: {", ".join(values)}')
    lines.extend(text for key, text in FLAG_LABELS.items() if filters.get(key))
    if filters.get('search'):
        lines.append(f"Search: {filters['search']}")
    if not lines:
        lines.append('Filters: every Non-Qualified Call of the workspace')
    lines.append(f'Progress periods: {GRANULARITY_LABELS.get(granularity, "Month")}')
    lines.append(f'Generated: {datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %Z")}')
    return lines


def _share(count: int, total: int) -> str:
    return f'{count * 100 / total:.1f}%' if total else '0.0%'


def _days(value: Any) -> str:
    return '—' if value is None else f'{value:.1f}'


def summary_table(summary: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    total = summary['total']
    return ['Indicator', 'Calls', 'Share'], [
        ['Non-Qualified Calls', f'{total:,}', '100%' if total else '0%'],
        ['Open', f"{summary['open']:,}", _share(summary['open'], total)],
        ['Closed', f"{summary['closed']:,}", _share(summary['closed'], total)],
        ['Attended', f"{summary.get('attended', 0):,}", _share(summary.get('attended', 0), total)],
        ['With Team', f"{summary['with_team']:,}", _share(summary['with_team'], total)],
        ['Assigned', f"{summary['assigned']:,}", _share(summary['assigned'], total)],
        ['Commented', f"{summary['commented']:,}", _share(summary['commented'], total)],
    ]


def breakdown_table(breakdown: dict[str, Any], total: int) -> tuple[list[str], list[list[str]]]:
    empty = 'Unassigned' if breakdown['field'] == 'team' else 'Not classified'
    return [breakdown['label'].removeprefix('By ').strip(), 'Calls', 'Share'], [
        [item.get('label') or item['value'] or empty, f"{item['count']:,}", _share(item['count'], total)]
        for item in breakdown['items']
    ]


def progress_summary_table(summary: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    return ['Indicator', 'Value'], [
        ['Attended (followed up or commented)', f"{summary['attended']:,}"],
        ['Not attended yet', f"{summary['not_attended']:,}"],
        ['Open', f"{summary['open']:,}"],
        ['Closed', f"{summary['closed']:,}"],
        ['Closure rate', f"{summary['closure_rate']:.1f}%"],
        ['Comments', f"{summary['comments']:,}"],
        ['Follow-up changes', f"{summary['changes']:,}"],
        ['Contributors', f"{summary['contributors']:,}"],
        ['Average days from the call to its first follow-up', _days(summary['avg_days_to_first_follow_up'])],
        ['Average days from the call to its closing', _days(summary['avg_days_to_close'])],
    ]


def timeline_table(progress: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    statuses = progress['statuses']
    columns = ['Period', 'Detected', 'Attended', 'Comments', *[f'→ {status}' for status in statuses],
               'Closed', 'Reopened', 'Open backlog', 'Avg days to close']
    rows = [[period['period'], str(period['detected']), str(period['attended']), str(period['comments']),
             *[str(period['statuses'].get(status, 0)) for status in statuses],
             str(period['closed']), str(period['reopened']), str(period['open_backlog']), _days(period['avg_days_to_close'])]
            for period in progress['periods']]
    return columns, rows


def workload_table(title: str, rows: list[dict[str, Any]]) -> tuple[list[str], list[list[str]]]:
    return [title, 'Open', 'Closed', 'Total'], [[row['name'], str(row['open']), str(row['closed']), str(row['total'])] for row in rows]


def activity_table(rows: list[dict[str, Any]]) -> tuple[list[str], list[list[str]]]:
    return ['User', 'Changes', 'Comments', 'Calls closed'], [
        [row['name'], str(row['changes']), str(row['comments']), str(row['closed'])] for row in rows
    ]


def report_tables(summary: dict[str, Any], breakdowns: list[dict[str, Any]], progress: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    """Every table of the report, in order: Executive Summary, then Progress Status."""
    tables = [('Executive Summary', *summary_table(summary))]
    tables += [(breakdown['label'], *breakdown_table(breakdown, summary['total'])) for breakdown in breakdowns]
    tables.append(('Progress Status', *progress_summary_table(progress['summary'])))
    tables.append((f"Progress per {GRANULARITY_LABELS.get(progress['granularity'], 'Week')}", *timeline_table(progress)))
    tables.append(('Age of Open Calls', ['Age since the call', 'Open calls'], [[item['value'], str(item['count'])] for item in progress['aging']]))
    tables.append(('Team Workload', *workload_table('Team', progress['teams'])))
    tables.append(('Assignee Workload', *workload_table('Assignee', progress['assignees'])))
    tables.append(('User Activity', *activity_table(progress['activity'])))
    return tables


def _text_slide(presentation: Presentation, title: str, lines: list[str]) -> None:
    slide = _slide(presentation, title)
    left, top, width, height = _content_frame(slide)
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    shape.text_frame.word_wrap = True
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for index, line in enumerate(lines):
        paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(15)


def _pie_slide(presentation: Presentation, title: str, distributions: list[dict[str, Any]]) -> None:
    """Up to three doughnut charts side by side, each with its own legend."""
    slide = _slide(presentation, title)
    left, top, width, height = _content_frame(slide)
    shown = [item for item in distributions if item['items']][:3]
    if not shown:
        return
    column = width / len(shown)
    for index, distribution in enumerate(shown):
        data = CategoryChartData()
        items = distribution['items'][:10]
        data.categories = [item['value'] for item in items]
        data.add_series('Calls', [item['count'] for item in items])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(left + column * index), Inches(top),
                                       Inches(column - 0.15), Inches(height), data)
        chart = frame.chart
        chart.has_title = True
        chart.chart_title.text_frame.text = distribution['label']
        chart.chart_title.text_frame.paragraphs[0].font.size = Pt(13)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(9)
        plot.data_labels.number_format = '0'
        plot.data_labels.number_format_is_linked = False
        for point_index, item in enumerate(items):
            color = str(item.get('color') or '').lstrip('#')
            if len(color) == 6:
                fill = plot.series[0].points[point_index].format.fill
                fill.solid()
                fill.fore_color.rgb = RGBColor.from_string(color.upper())


def _timeline_chart_slide(presentation: Presentation, progress: dict[str, Any]) -> None:
    periods = progress['periods'][-24:]
    if not periods:
        return
    slide = _slide(presentation, f"Attended and Closed per {GRANULARITY_LABELS.get(progress['granularity'], 'Week')}",
                   'The latest 24 periods')
    left, top, width, height = _content_frame(slide)
    data = CategoryChartData()
    data.categories = [period['period'] for period in periods]
    for key, label in (('detected', 'Detected'), ('attended', 'Attended'), ('closed', 'Closed')):
        data.add_series(label, [period[key] for period in periods])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(left), Inches(top), Inches(width), Inches(height), data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(10)
    for series, color in zip(chart.series, ('8A6B76', 'E08A1E', '2E8B57'), strict=False):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color)
    chart.category_axis.tick_labels.font.size = Pt(9)
    chart.value_axis.tick_labels.font.size = Pt(9)


def with_colors(distributions: list[dict[str, Any]], options: dict[str, Any]) -> list[dict[str, Any]]:
    colors = {('status', item['name']): item['color'] for item in options['statuses']}
    colors.update({('team', item['name']): item['color'] for item in options['teams']})
    return [{**distribution, 'items': [{**item, 'color': colors.get((distribution['field'], item['value']), '')}
                                       for item in distribution['items']]} for distribution in distributions]


def _cover_selection(cover, lines: list[str]) -> None:
    """The report selection (filters, progress periods and generation time) below the cover's line."""
    from src.modules.exports import cover_details, remove_empty_placeholders

    cover_details(cover, lines)
    remove_empty_placeholders(cover)


# The root domain of the calls a source gives no root cause (as non_qualified_calls.NOT_CLASSIFIED).
NOT_CLASSIFIED = 'Not classified'


def classified_root_stats(stats: dict[str, Any]) -> dict[str, Any]:
    """The Root Cause Analysis of the calls the source gives a root cause: Not classified leaves its domains, categories,
    causes and charts, and every group of calls counts only those (as the page without Show Not classified)."""
    def classified(item: dict[str, Any]) -> bool:
        return item.get('value') != NOT_CLASSIFIED and item.get('domain') != NOT_CLASSIFIED

    columns = [column for column in stats.get('columns') or [] if column != NOT_CLASSIFIED]

    def groups(rows: list[dict[str, Any]], ordered: bool = False) -> list[dict[str, Any]]:
        kept = [{**row, 'total': sum(int((row.get('counts') or {}).get(column) or 0) for column in columns)} for row in rows or []]
        kept = [row for row in kept if row['total']]
        return sorted(kept, key=lambda row: -row['total']) if ordered else kept

    return {**stats, 'columns': columns,
            **{key: [item for item in stats.get(key) or [] if classified(item)] for key in ('domains', 'categories', 'causes')},
            'call_types': groups(stats.get('call_types')), 'nr_modes': groups(stats.get('nr_modes')),
            'technologies': groups(stats.get('technologies'), ordered=True)}


# The panels of the page and the charts and tables of each one, in the order of the page, as a Bundle Report Template
# chooses them (see bundle_report_templates): (panel key, panel title, ((component key, label), ...)).
REPORT_PANELS: tuple[tuple[str, str, tuple[tuple[str, str], ...]], ...] = (
    ('summary', 'Summary', (('executive', 'Executive Summary'),)),
    ('prioritization', 'NQ Prioritization & Queues', (('queue', 'Next to analyse'), ('rules', 'Prioritization Rules'),
                                                      ('kpis', 'Scoring points lost by KPI'), ('hotspots', 'Hotspots'),
                                                      ('issues', 'Repeated Issues Within Same Cluster'))),
    ('rates', 'NQ Rate by Campaign and Operator', (('rates', 'NQ rate of each campaign and operator'),)),
    ('failures', 'Failures by Type, Technology & Serving Cell', (
        ('types', 'Failure types of each service'), ('technology', 'Technology and classification'),
        ('serving', 'Serving Vendor and host network'), ('sites', 'Sites with the most Non-Qualified Calls'))),
    ('sources', 'Root Cause Sources', (('sources', 'The root cause sources compared'),)),
    ('root', 'Root Cause Analysis', (('dta', 'DT Analyzer RCA'), ('netcheck', 'NetCheck RCA'), ('script', 'Ericsson RCA'),
                                     ('analysts', 'Analysts'))),
    ('lifecycle', 'Lifecycle & Phases', (('pipeline', 'NQ Call Status Pipeline'), ('phases', 'Phases'),
                                         ('times', 'Time in each status and most frequent moves'), ('stalled', 'Stalled open calls'))),
    ('progress', 'Progress View', (('overview', 'Follow-up Overview'), ('timeline', 'Progress per period'),
                                   ('workload', 'Age and Workload'), ('details', 'Details per period'), ('activity', 'User Activity'))),
)
REPORT_SECTION_PREFIX = 'section:'


def _lifecycle_component(title: Any) -> str:
    subtitle = str(getattr(title, 'subtitle', '') or '')
    if subtitle.startswith('NQ Call Status Pipeline'):
        return 'pipeline'
    if subtitle.startswith('Time in Each Status'):
        return 'times'
    return 'stalled' if subtitle.startswith('Stalled Open Calls') else 'phases'


def _progress_component(title: Any) -> str:
    subtitle = str(getattr(title, 'subtitle', '') or '')
    for start, key in (('Follow-up Overview', 'overview'), ('Progress per', 'timeline'), ('Age and Workload', 'workload'),
                       ('Details per', 'details'), ('User Activity', 'activity')):
        if subtitle.startswith(start):
            return key
    return 'overview'


def _failure_component(subtitle: str) -> str:
    for end, key in (('Failure type', 'types'), ('Technology and classification', 'technology'),
                     ('Serving Vendor and host network', 'serving')):
        if subtitle.endswith(end):
            return key
    return 'sites'


def keyed_report_panels(summary: dict[str, Any], breakdowns: list[dict[str, Any]], progress: dict[str, Any],
                        options: dict[str, Any]) -> list[tuple[str, str, Any, Any]]:
    """The panels of ``report_panels`` with the panel and component of REPORT_PANELS each one belongs to:
    (panel key, component key, title, image)."""
    from src.modules.non_qualified_calls_visuals import (
        SlideTitle, executive_summary_panels, lifecycle_panels, progress_title, progress_view_panels, rate_panels,
        rca_insight_panels, root_cause_panels, table_panels,
    )

    label = GRANULARITY_LABELS.get(progress['granularity'], 'Week')
    # The detail tables of the Progress View, each page with its number in the subtitle.
    tables = []
    for subtitle, columns, rows in detail_tables(progress):
        pages = table_panels(subtitle, columns, rows)
        tables.extend(('progress', 'activity' if subtitle == 'User Activity' else 'details',
                       progress_title(label, subtitle + (f' · {number}/{len(pages)}' if len(pages) > 1 else '')), image)
                      for number, (_title, image) in enumerate(pages, start=1))
    # The Root Cause Sources that compare the sources, then the Root Cause Analysis of each one with its tables. As on
    # the page by default, the charts leave out the Not classified calls and the phases the calls without a value,
    # unless the options ask for them (Show Not classified, Show Empty Status).
    analysis = [('root', str(stats.get('source') or 'analysts'), *panel) for stats in options.get('root_causes') or []
                for shown in [stats if options.get('show_not_classified') else classified_root_stats(stats)]
                for panel in (*root_cause_panels(shown), *root_cause_table_panels(shown))]
    sources = [('sources', 'sources', *panel) for panel in rca_insight_panels(options['rca_insights'])] if options.get('rca_insights') else []
    rates = [('rates', 'rates', *panel) for panel in rate_panels(options['rates'])] if options.get('rates') else []
    # Failures by Type, Technology & Serving Cell, each table on its own panels with what it shows in the subtitle.
    failures = []
    for subtitle, columns, rows in failure_tables(options.get('failures') or {}):
        pages = table_panels(subtitle, columns, rows)
        failures.extend(('failures', _failure_component(subtitle),
                         SlideTitle(FAILURES_TITLE, '', subtitle + (f' · {number}/{len(pages)}' if len(pages) > 1 else '')), image)
                        for number, (_title, image) in enumerate(pages, start=1))
    # The NQ Prioritization, each table on its own panels with what it shows in the subtitle.
    prioritization = []
    for component, subtitle, columns, rows, numeric in prioritization_tables(options.get('prioritization') or {}):
        pages = table_panels(subtitle, columns, rows, numeric=numeric)
        prioritization.extend(('prioritization', component,
                               SlideTitle(PRIORITIZATION_TITLE, '', subtitle + (f' · {number}/{len(pages)}' if len(pages) > 1 else '')), image)
                              for number, (_title, image) in enumerate(pages, start=1))
    life = options.get('lifecycle')
    if life and not options.get('show_empty'):
        life = {**life, 'phases': [{**phase, 'empty': 0} for phase in life.get('phases') or []]}
    lifecycle = [('lifecycle', _lifecycle_component(title), title, image)
                 for title, image in (lifecycle_panels(life, options.get('sections')) if life else [])]
    return [*[('summary', 'executive', *panel) for panel in executive_summary_panels(summary, breakdowns, options)],
            *prioritization, *rates, *failures, *sources, *analysis, *lifecycle,
            *[('progress', _progress_component(title), title, image) for title, image in progress_view_panels(progress, options, label)],
            *tables]


def report_panels(summary: dict[str, Any], breakdowns: list[dict[str, Any]], progress: dict[str, Any],
                  options: dict[str, Any]) -> list[tuple[str, Any]]:
    """The Executive Summary, NQ Prioritization, NQ rates, Root Cause Sources and Analysis, Lifecycle and Progress View
    drawn like the page, as panels."""
    return [(title, image) for _panel, _component, title, image in keyed_report_panels(summary, breakdowns, progress, options)]


FAILURES_TITLE = 'Failures by Type, Technology & Serving Cell'
PRIORITIZATION_TITLE = 'NQ Prioritization & Queues'
# The slides keep the hotspots and the sites with a repeated issue with the most calls (the page lists every one).
MAX_EXPORT_ISSUES = 25


def _points(value: Any) -> str:
    return '—' if value is None else f'{float(value):.2f}'


def prioritization_tables(data: dict[str, Any]) -> list[tuple[str, str, list[str], list[list[str]], set[int]]]:
    """The NQ Prioritization as on the page: the calls to analyse next in the queue of each team, the calls of each step of
    its strategy, the KPIs whose points the calls lose and the repeated issues ((component, subtitle, columns, rows, numeric
    columns))."""
    tables = []
    strategy = data.get('strategy') or {}
    queue = [[row['join_id'] or '—', row['operator'], row['service'], row['nr_mode'] or '—', row['campaign'] or '—',
              row['result'] or '—', row['rule'] or '—', row['priority'] or '—', _points(row['points']),
              '—' if row['repetitions'] is None else f"{row['repetitions']:,}", row['place'] or '—'] for row in data.get('queue') or []]
    if queue:
        tables.append(('queue', 'Next to analyse', ['JOIN_ID', 'Operator', 'Service', 'NR Mode', 'Campaign', 'Result', 'Rule',
                                                    'Priority', 'Points lost', 'Repetitions', 'Place'], queue, {8, 9}))
    for item in data.get('queues') or []:
        rows = [[f"{row['position']:,}", row['join_id'] or '—', row['priority'] or '—', row['operator'], row['service'],
                 row['nr_mode'] or '—', row['campaign'] or '—', row['rule'] or '—', _points(row.get('gain')),
                 _points(row['points']), row['level'] or '—'] for row in item.get('rows') or []]
        if rows:
            tables.append(('queue', f"Next to analyse · {item['name']} queue",
                           ['#', 'JOIN_ID', 'Priority', 'Operator', 'Service', 'NR Mode', 'Campaign', 'Step', 'Expected Gain',
                            'Points lost', 'Scoring Impact'], rows, {0, 8, 9}))
    # The steps of each list in a table of its own and the calls out of scope and those no step takes in a third, as the
    # tabs of the page.
    shown = [row for row in data.get('rules') or [] if row['calls'] or row['name']]
    final_list = any(row.get('list') == 'final' for row in shown)
    groups = (('campaign', 'List 1 · Weekly CDRs' if final_list else 'List 1 · Weekly & Final CDRs'), ('final', 'List 2 · Final CDRs'),
              ('outside', 'Outside the lists'))
    for key, title in groups:
        rows = [row for row in shown if (row.get('list') if row.get('kind') == 'step' else 'outside') == key]
        if rows:
            tables.append(('rules', f"{strategy.get('label') or 'Strategy'} · {title}",
                           ['Priority', 'Step', 'NR Mode', 'Operators', 'Call Type', 'Area', 'Condition', 'Calls', 'Open', 'Waiting',
                            'Points lost', 'Queue'], [[*_step_fields(row)[:7], f"{row['calls']:,}", f"{row['open']:,}", f"{row.get('waiting', 0):,}",
                                                      _points(row['points']), _step_fields(row)[7]] for row in rows], {7, 8, 9, 10}))
    kpis = [[row['kpi'], row.get('category') or '—', row['environment'], f"{row['calls']:,}", _points(row['points'])]
            for row in data.get('kpis') or []]
    if kpis:
        tables.append(('kpis', 'Scoring points lost by KPI', ['KPI', 'Category', 'Environment', 'Calls', 'Points lost'], kpis, {3, 4}))
    spots = data.get('hotspots') or []
    spot_rows = [[row['operator'], row['nr_mode'] or '—', row['campaign'] or '—',
                  (row['place'] or '—') + (f" +{row['places'] - 1}" if row.get('places', 0) > 1 else ''), hotspot_time(row),
                  f"{row['calls']:,}", *(f"{row['kinds'].get(kind, 0):,}" for kind in ('Voice', 'WhatsApp', 'Data')),
                  _points(row['points']), f"{row['open']:,}", row['category'] or '—'] for row in spots[:MAX_EXPORT_ISSUES]]
    if spot_rows:
        # The title tells how the failures of a hotspot are joined, as the page does.
        join = 'AND' if (data.get('settings') or {}).get('hotspot_join') == 'and' else 'OR'
        tables.append(('hotspots', f'Hotspots (Consecutives {join} Clustered Issues)'
                       + (f' · the {len(spot_rows)} with the most calls of {len(spots):,}' if len(spots) > len(spot_rows) else ''),
                       ['Operator', 'NR Mode', 'Campaign', 'Site', 'Time', 'Calls', 'Voice', 'WhatsApp', 'Data', 'Points lost', 'Open',
                        'Main Root Category'], spot_rows, {5, 6, 7, 8, 9, 10}))
    # The sites with a repeated issue in the columns of the hotspots: their NR Modes and campaigns (with their calls when
    # there are several) and the days of their first and last call.
    def shares(counts: dict[str, int]) -> str:
        return ' · '.join(f'{name} {count:,}' for name, count in counts.items()) if len(counts) > 1 else next(iter(counts), '—')
    issues = [[row['operator'], shares(row['nr_modes']), shares(row.get('campaign_names') or {}), row['place'],
               hotspot_time({'first': str(row.get('first') or '')[:10], 'last': str(row.get('last') or '')[:10]}), f"{row['calls']:,}",
               *(f"{(row.get('kinds') or {}).get(kind, 0):,}" for kind in ('Voice', 'WhatsApp', 'Data')), _points(row['points']),
               f"{row['open']:,}", row['category'] or '—'] for row in (data.get('issues') or [])[:MAX_EXPORT_ISSUES]]
    if issues:
        count = len(data.get('issues') or [])
        tables.append(('issues', 'Repeated Issues Within Same Cluster' + (f' · the {len(issues)} sites with the most calls of {count:,}' if count > len(issues) else ''),
                       ['Operator', 'NR Mode', 'Campaign', 'Site', 'Time', 'Calls', 'Voice', 'WhatsApp', 'Data', 'Points lost', 'Open',
                        'Main Root Category'], issues, {5, 6, 7, 8, 9, 10}))
    return tables


def hotspot_time(row: dict[str, Any]) -> str:
    """When the calls of a hotspot started: 2026-07-14 10:05–10:31 (the date once when it is the same day); 2026-07-14–2026-08-02
    for days alone."""
    first, last = str(row.get('first') or '')[:16], str(row.get('last') or '')[:16]
    if not first:
        return '—'
    if last == first or not last:
        return first
    return f'{first}–{last[11:]}' if len(first) > 10 and last[:10] == first[:10] else f'{first}–{last}'


def _step_fields(row: dict[str, Any]) -> list[str]:
    """The fields of a step as the Prioritization Rules show them (Priority, Step, NR Mode, Operators, Call Type, Area,
    Condition and Queue, the calls and the points going before the Queue); the calls out of scope and those no step takes say
    so in the Condition."""
    if row.get('kind') != 'step':
        return ['—', _step_name(row), '—', '—', '—', '—', row.get('label') or row['conditions'], row['criterion'] or '—']
    return [row.get('priority') or '—', _step_name(row), 'Any' if row.get('nr_mode') == 'any' else row.get('nr_mode') or '—',
            row.get('operators') or '—', row.get('call_type') or '—', row.get('area') or '—', row.get('condition') or '—',
            row.get('queue') or '—']


def _step_name(row: dict[str, Any]) -> str:
    """A row of the steps of the strategy as the page names it: its number (in the table of its list), Out of scope or No step."""
    if row.get('kind') == 'step' and row.get('number'):
        return str(row['number'])
    if row.get('kind') == 'out_of_scope':
        return 'Out of scope'
    return row.get('name') or 'No step'


def failure_tables(failures: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    """Failures by Type, Technology & Serving Cell as on the page: the failure types of each service, the technology at
    the failure against the classification and the Vendor and host network of each operator, and the sites with the
    most Non-Qualified Calls."""
    tables = []
    for item in failures.get('types') or []:
        tables.append((f"{item['label']} · Failure type", ['Operator', 'Environment', 'Calls', 'Non-Qualified', *item['results']],
                       [[row['operator'], row['environment'], f"{row['calls']:,}", f"{row['nq']:,}",
                         *[(f"{row['rates'][result]:.2f}% · {row['counts'][result]:,}" if row['rates'].get(result) is not None
                            else f"{row['counts'][result]:,}") if row['counts'].get(result) else '—' for result in item['results']]]
                        for row in item['rows']]))
    for item in failures.get('technology') or []:
        classes = item['classes'][:6]
        rows = [[operator, name, *[f"{counts.get(label, 0):,}" if counts.get(label) else '—' for label in classes],
                 f"{sum(counts.values()):,}"]
                for operator in item['operators'] for name, counts in sorted(
                    item['cells'].get(operator, {}).items(), key=lambda pair: -sum(pair[1].values()))]
        if rows:
            tables.append((f"{item['label']} · Technology and classification", ['Operator', 'Technology', *classes, 'Total'], rows))
    for item in failures.get('serving') or []:
        rows = [[operator['operator'], row['vendor'], row['host_network'], f"{row['calls']:,}", f"{row['calls_share']}%",
                 f"{row['nq']:,}", f"{row['nq_share']}%", '—' if row['rate'] is None else f"{row['rate']:.2f}%"]
                for operator in item['operators'] for row in operator['rows'][:8]]
        if rows:
            tables.append((f"{item['label']} · Serving Vendor and host network",
                           ['Operator', 'Vendor', 'Host network', 'Calls', 'Share of calls', 'Non-Qualified', 'Share of NQ', 'NQ rate'], rows))
    rows = [[item['operator'], row['site_id'], row['host_network'] or '—', row['vendor'] or '—', row['city'] or '—', f"{row['calls']:,}",
             f"{row['nq']:,}", f"{row['rate']:.2f}%"] for item in failures.get('sites') or [] for row in item['rows'][:10]]
    if rows:
        tables.append(('Sites with the most Non-Qualified Calls',
                       ['Operator', 'Site', 'Host network', 'Vendor', 'City or route', 'Calls', 'Non-Qualified', 'NQ rate'], rows))
    return tables


def node_table(stats: dict[str, Any]) -> tuple[list[str], list[list[str]]]:
    return ['eNB / gNB', 'Operator', 'Site', 'Host', 'Vendor', 'Calls', 'Main Domain'], [
        [row['node'], row['operator'], row['site'] or '—', row['host'] or '—', row['vendor'] or '—',
         f"{row['calls']:,}", row['top_domain'] or '—'] for row in stats['nodes']
    ]


def root_cause_table_panels(stats: dict[str, Any]) -> list[tuple[str, Any]]:
    """The Technology bars on one slide, and the eNB / gNB with most calls on another (the first ones that fit)."""
    from src.modules.non_qualified_calls_visuals import root_cause_title, table_panels, technology_panel

    nodes = 'eNB / gNB with Most Non-Qualified Calls'
    return [(root_cause_title(stats, 'Root Domains by Technology'), technology_panel(stats)),
            (root_cause_title(stats, nodes), table_panels(nodes, *node_table(stats), single=True, numeric={5})[0][1])]


def detail_tables(progress: dict[str, Any]) -> list[tuple[str, list[str], list[list[str]]]]:
    """The figures the panels summarise: every period of the timeline and the activity of each user."""
    label = GRANULARITY_LABELS.get(progress['granularity'], 'Week')
    return [(f'Details per {label}', *timeline_table(progress)), ('User Activity', *activity_table(progress['activity']))]


def export_powerpoint(destination: Path, lines: list[str], summary: dict[str, Any], breakdowns: list[dict[str, Any]],
                      progress: dict[str, Any], options: dict[str, Any], template: Path | None = None,
                      components: Any = None) -> Path:
    """The Executive Summary and Progress Status in PowerPoint. ``components`` (``panel.component`` keys of
    REPORT_PANELS) keeps only those charts and tables, as a Bundle Report Template chooses them; every slide is named after
    its panel."""
    from PIL import Image

    from src.config import settings
    from src.modules.cdr_reporting import _named_slide_layout, _remove_all_slides, _set_structural_slide_text
    from src.modules.non_qualified_calls_visuals import SlideTitle

    presentation = Presentation(template or settings.powerpoint_template)
    _remove_all_slides(presentation)
    cover = presentation.slides.add_slide(_named_slide_layout(presentation, 'Title Page'))
    _set_structural_slide_text(cover, 'Non-Qualified Calls', 'Executive Summary and Progress Status')
    _cover_selection(cover, lines)
    cover._element.cSld.set('name', f'{REPORT_SECTION_PREFIX}cover')
    wanted = None if components is None else set(components)
    for panel, component, title, image in keyed_report_panels(summary, breakdowns, progress, options):
        if wanted is not None and f'{panel}.{component}' not in wanted:
            continue
        heading = title if isinstance(title, SlideTitle) else SlideTitle(title)
        slide = _slide(presentation, heading.heading, heading.subtitle)
        slide._element.cSld.set('name', f'{REPORT_SECTION_PREFIX}{panel}')
        if heading.bold:
            # The title ends in its part in bold, such as the source of the root causes.
            paragraph = slide.shapes.title.text_frame.paragraphs[0]
            paragraph.text = f'{heading.text} · '
            run = paragraph.add_run()
            run.text = heading.bold
            run.font.bold = True
            from src.modules.slide_titles import fit_slide_title
            fit_slide_title(slide)
        left, top, width, height = _content_frame(slide)
        with Image.open(image) as raster:
            ratio = raster.width / raster.height
        image.seek(0)
        picture_width = min(width, height * ratio)
        slide.shapes.add_picture(image, Inches(left + (width - picture_width) / 2), Inches(top),
                                 width=Inches(picture_width), height=Inches(picture_width / ratio))
    closing = _named_slide_layout(presentation, 'Black logo end slide')
    if closing is not None:
        _set_structural_slide_text(presentation.slides.add_slide(closing), '', '')
    from src.modules import report_structure
    report_structure.finish(presentation)
    presentation.save(destination)
    return destination
