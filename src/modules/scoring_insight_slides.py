"""Insight slides of the Scoring & GAP Analysis exports.

Location cards, KPI GAP profiles, scoring trends, campaign comparisons, the National & Areas and the points-lost maps
use the insights of the scoring views (see ``scoring_insights``), in the palette of the web page: Data in teal and Voice
in amber, losses and GAPs behind the reference in red and gains in green.
"""
from __future__ import annotations

from math import ceil, floor
from typing import Any, Callable

from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.modules.scoring_themes import active_theme, themed
from src.modules.scoring_vendors import scoring_vendor_name

VOICE_COLOR = '#F2A900'
DATA_COLOR = '#0E6B66'
TITLE_BLUE = '#1450A8'
TEXT_COLOR = '#17232D'
CARD_COLOR = '#F2F3F5'
LOSS_COLOR = '#E8414F'
WIN_COLOR = '#4CA65A'
# KPI GAP profiles for every location the location cards show (one per City, Cluster or Region).
MAX_PROFILE_SCOPES = 12
PROFILE_TABLE_BOTTOM = 6.68
MAX_GROUP_SLIDES = 6
MAX_COMPARISONS = 8
_FONT = 'Arial'
FALLBACK_COLORS = ('#E60000', '#0B6E8F', '#7A3DB8', '#00A3AD', '#F2A900', '#4CA65A', '#8C564B', '#5B6770')


def _rgb(color: str) -> RGBColor:
    """The colour in the active theme (see scoring_themes.themed)."""
    return RGBColor.from_string(themed(color).lstrip('#'))


def _points(value: float | None, digits: int = 0) -> str:
    return 'N/A' if value is None else f'{value:,.{digits}f}'


def _box(slide, left, top, width, height, color, *, shape=MSO_SHAPE.RECTANGLE):
    box = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb(color)
    box.line.fill.background()
    box.shadow.inherit = False
    return box


def _unsigned_axis_ids(chart) -> None:
    """Axis identifiers are unsigned integers in DrawingML; the bar chart templates of python-pptx use negative ones."""
    references = chart._chartSpace.xpath('.//c:axId | .//c:crossAx')
    numbers = {value: str(index + 1) for index, value in enumerate(dict.fromkeys(reference.get('val') for reference in references))}
    for reference in references:
        reference.set('val', numbers[reference.get('val')])


def _lighten(color: str, amount: float) -> RGBColor:
    red, green, blue = (int(themed(color).lstrip('#')[index:index + 2], 16) for index in (0, 2, 4))
    return RGBColor(*(round(value + (255 - value) * amount) for value in (red, green, blue)))


def _gradient_bar(slide, left, top, width, height, color: str, *, reverse: bool = False):
    """A data bar like the Excel ones: the colour fading to white (towards the left when ``reverse``), with a border of
    the same colour."""
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    bar.fill.gradient()
    bar.fill.gradient_angle = 180 if reverse else 0
    stops = bar.fill.gradient_stops
    stops[0].position, stops[1].position = 0, 1
    stops[0].color.rgb = _rgb(color)
    stops[1].color.rgb = _lighten(color, .88)
    bar.line.color.rgb = _rgb(color)
    bar.line.width = Pt(.5)
    bar.shadow.inherit = False
    return bar


def _write(frame, paragraphs: list[tuple[str, float, bool, str]], *, align=PP_ALIGN.LEFT) -> None:
    """Fill a text frame with (text, size, bold, color) paragraphs."""
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.08)
    frame.margin_top = frame.margin_bottom = Inches(.04)
    frame.clear()
    for index, (text, size, bold, color) in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.name = _FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)


def _write_runs(frame, runs: list[tuple[str, bool, str | None]], *, size: float, color: str) -> None:
    """One paragraph of (text, bold, color) runs; a run without a colour uses ``color``."""
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.clear()
    paragraph = frame.paragraphs[0]
    for text, bold, run_color in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = _FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(run_color or color)


def _legend_key(slide, left: float, top: float, items: list[tuple[str, str]]) -> None:
    for label, color in items:
        _box(slide, left, top + .03, .16, .16, color)
        box = slide.shapes.add_textbox(Inches(left + .22), Inches(top), Inches(1.1), Inches(.24))
        _write(box.text_frame, [(label, 9, False, '#4A5B65')])
        box.text_frame.margin_left = 0
        left += 1.35


def _block_profiles(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Keep the GAP profiles of at most a few contexts: the latest campaign first."""
    contexts = list(dict.fromkeys(tuple(sorted(item['context'].items())) for item in items))
    if len(contexts) > MAX_PROFILE_SCOPES:
        from src.modules.scoring_insights import campaign_sort_key
        campaigns = sorted({item['context'].get('campaign') for item in items if item['context'].get('campaign')},
                           key=campaign_sort_key)
        if campaigns:
            items = [item for item in items if item['context'].get('campaign') == campaigns[-1]]
            contexts = list(dict.fromkeys(tuple(sorted(item['context'].items())) for item in items))
    if len(contexts) > MAX_PROFILE_SCOPES:
        return []
    return items


# The levels that name the area of a chart or table, the most specific first.
AREA_LEVELS = (('City', 'Cities'), ('Cluster', 'Clusters'), ('Region', 'Regions'))
# What a chart or table without area, operator or vendor filters shows.
SCOPE_FALLBACKS = ('National', 'All Operators', 'All Vendors')
# The area, operators and vendors to name on a slide, for the context of the slide (see scoring_scope_labels).
ScopeLabels = Callable[[dict[str, Any] | None], list[str]]


def _level_values(source: Any, level: str) -> list[str]:
    """The values of a level (City, Region, Operator…) in a context or in filters (a mapping, or column filters)."""
    if isinstance(source, dict):
        value = next((item for key, item in source.items() if str(key).casefold() == level.casefold()), None)
        if isinstance(value, set):
            values = sorted(value, key=str)
        else:
            values = list(value) if isinstance(value, (list, tuple)) else [value]
    elif isinstance(source, list):
        values = []
        for item in source:
            if isinstance(item, dict) and str(item.get('column') or item.get('field') or '').casefold() == level.casefold():
                value = item.get('value', item.get('values'))
                values.extend(value if isinstance(value, (list, tuple)) else [value])
    else:
        values = []
    return list(dict.fromkeys(str(value).strip() for value in values if value is not None and str(value).strip()))


def _scope_values(context: dict[str, Any] | None, filters: Any, level: str, operators: Any = ()) -> list[str]:
    """The values of a level in the context of a chart or table, else in the filters (vendors without operator)."""
    for source in (context, filters):
        values = _level_values(source, level)
        if values and level == 'Vendor':
            return list(dict.fromkeys(scoring_vendor_name(value, operators) for value in values))
        if values:
            return values
    return []


def _scope_label(values: list[str], level: str, plural: str) -> str:
    shown = [*values[:2], f'{len(values) - 2} more'] if len(values) > 3 else values
    names = f"{', '.join(shown[:-1])} & {shown[-1]}" if len(shown) > 1 else shown[0]
    return f'{plural if len(values) > 1 else level}: {names}'


def scoring_area_label(context: dict[str, Any] | None, filters: Any) -> str:
    """The area a chart or table shows, as the web titles name it (for example City: London): the City, Cluster or
    Region of its context, else the City, Cluster or Region filters of the scoring, else National. A long list keeps
    its first two names and counts the others."""
    for level, plural in AREA_LEVELS:
        values = _scope_values(context, filters, level)
        if values:
            return _scope_label(values, level, plural)
    return 'National'


def scoring_scope_labels(context: dict[str, Any] | None, filters: Any, operators: Any = (), *,
                         omit_context: bool = False) -> list[str]:
    """The area, operators and vendors a chart or table shows, as the pills of the web titles: each from its context,
    else from the filters, else National, All Operators or All Vendors. ``operators`` take the operator out of
    Operator_Vendor values. With ``omit_context`` the values the slide already names from its context are left out."""
    operator_values = _scope_values(context, filters, 'Operator')
    vendor_values = _scope_values(context, filters, 'Vendor', operators)
    labels = [
        scoring_area_label(context, filters),
        _scope_label(operator_values, 'Operator', 'Operators') if operator_values else 'All Operators',
        _scope_label(vendor_values, 'Vendor', 'Vendors') if vendor_values else 'All Vendors',
    ]
    if omit_context and context:
        told = set(scoring_scope_labels(context, None, operators))
        labels = [label for label in labels if label in SCOPE_FALLBACKS or label not in told]
    return labels


def _with_scope(subtitle: str, scope: ScopeLabels | None, context: dict[str, Any] | None = None) -> str:
    """A slide subtitle followed by the area, operators and vendors that ``scope`` gives for the slide's context."""
    return ' · '.join(part for part in (subtitle, *(scope(context) if scope else ())) if part)


def _context_text(context: dict[str, Any]) -> str:
    from src.modules.scoring_insights import _display
    names = {'vendor': 'Vendor', 'region': 'Region', 'cluster': 'Cluster', 'city': 'City', 'campaign': 'Campaign'}
    return ' · '.join(f'{label}: {_display(label, context[field])}' for field, label in names.items() if field in context)


def _latest_campaign_groups(groups: list[dict[str, Any]]) -> list[dict[str, Any]]:
    from src.modules.scoring_insights import campaign_sort_key
    campaigns = [group.get('campaign') for group in groups if group.get('campaign')]
    if len(set(campaigns)) <= 1:
        return groups
    latest = max(campaigns, key=campaign_sort_key)
    return [group for group in groups if group.get('campaign') in (None, latest)]


def add_location_card_slides(presentation, groups: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                             new_slide: Callable, scope: ScopeLabels | None = None) -> None:
    """One slide per group: a card per City, Cluster or Region with each operator's Data and Voice points as a bar,
    from the highest score.

    With several campaigns only the latest one gets cards; the scoring trend shows the others.
    """
    groups = _latest_campaign_groups(groups)
    for group in groups[:MAX_GROUP_SLIDES]:
        level = group['level']
        plural = {'City': 'cities', 'Cluster': 'clusters', 'Region': 'regions'}[level]
        slide = new_slide(presentation, f'{scoring_label} Scoring per {level}',
                          ' · '.join(part for part in (_with_scope(subtitle, scope), group['title']) if part))
        cards = group['cards']
        maximum = group['maximum']
        voice_max, data_max = group.get('max_voice') or 0, group.get('max_data') or 0
        note = slide.shapes.add_textbox(Inches(.45), Inches(1.45), Inches(12.4), Inches(.3))
        note.name = 'Scoring Location Cards Note'
        _write_runs(note.text_frame, [
            (f'The {scoring_label} scoring of the selected {plural}, each from its highest score: ', False, None),
            (f'{_points(maximum)} points', True, None),
            (f' at most (Voice {_points(voice_max)} · Data {_points(data_max)})', False, None),
            (', the environments measured in each one scaled to them.' if group['scaled'] else '.', False, None)],
            size=9.5, color='#4A5B65')
        count = len(cards)
        columns = count if count <= 3 else 4 if count <= 8 else ceil(count / 3)
        rows = ceil(count / columns)
        gap = .14
        area_left, area_top, area_width, area_height = .45, 1.85, 12.45, 4.95
        width = (area_width - gap * (columns - 1)) / columns
        height = (area_height - gap * (rows - 1)) / rows
        label_size = 8 if rows <= 2 else 7
        for index, card in enumerate(cards):
            left = area_left + (index % columns) * (width + gap)
            top = area_top + (index // columns) * (height + gap)
            background = _box(slide, left, top, width, height, CARD_COLOR)
            background.name = f'Scoring Location Card {card["label"]}'
            heading = slide.shapes.add_textbox(Inches(left + .08), Inches(top + .04), Inches(width - .16), Inches(.3))
            _write(heading.text_frame, [(str(card['label']), 11 if rows <= 2 else 9.5, True, TITLE_BLUE)])
            operators = sorted(card['operators'], key=lambda item: item['total'])
            data = CategoryChartData()
            data.categories = [str(item['label']) for item in operators]
            data.add_series('Data', [item['data'] for item in operators])
            data.add_series('Voice', [item['voice'] for item in operators])
            # An empty segment after each bar carries its total.
            data.add_series('Total', [max(1, maximum) * .2 for _item in operators])
            chart_shape = slide.shapes.add_chart(
                XL_CHART_TYPE.BAR_STACKED, Inches(left + .03), Inches(top + .34),
                Inches(width - .06), Inches(height - .38), data,
            )
            chart_shape.name = f'Scoring Location Chart {card["label"]}'
            _format_card_chart(chart_shape.chart, maximum, label_size, operators)
        _legend_key(slide, 5.6, 6.98, [('Data', DATA_COLOR), ('Voice', VOICE_COLOR)])


def _format_card_chart(chart, maximum: float, label_size: float, operators: list[dict[str, Any]]) -> None:
    from src.modules.scoring_exports import _fill_series
    _unsigned_axis_ids(chart)
    chart.has_legend = False
    chart.font.name = _FONT
    chart.font.size = Pt(label_size)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = max(1, maximum) * 1.2
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.tick_labels.font.size = Pt(label_size + .5)
    chart.category_axis.tick_labels.font.bold = True
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.format.line.color.rgb = _rgb('#D5DBE2')
    plot = chart.plots[0]
    plot.gap_width = 55
    _fill_series(chart.series[0], DATA_COLOR)
    _fill_series(chart.series[1], VOICE_COLOR)
    # The total at the start of the empty segment after each bar.
    total = chart.series[2]
    total.format.fill.background()
    total.format.line.fill.background()
    for index, item in enumerate(operators):
        label = total.points[index].data_label
        label.position = XL_DATA_LABEL_POSITION.INSIDE_BASE
        label.text_frame.text = f"{item['total']:,.0f}{'' if item.get('complete', True) else '*'}"
        for paragraph in label.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(label_size + .5)
                run.font.bold = True
                run.font.color.rgb = _rgb(TEXT_COLOR)


def _table(slide, rows: int, columns: int, left: float, top: float, width: float, row_height: float, widths: list[float]):
    shape = slide.shapes.add_table(rows, columns, Inches(left), Inches(top), Inches(width), Inches(row_height * rows))
    table = shape.table
    for index, column_width in enumerate(widths):
        table.columns[index].width = Inches(column_width)
    for row in table.rows:
        row.height = Inches(row_height)
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    tbl_pr.set('firstRow', '0')
    tbl_pr.set('bandRow', '0')
    return shape, table


def _cell(cell, text: str, *, size: float, bold: bool = False, color: str = '#263746', fill: str | None = '#FFFFFF',
          align=PP_ALIGN.LEFT, underline: bool = False) -> None:
    cell.margin_left = cell.margin_right = Inches(.04)
    cell.margin_top = cell.margin_bottom = Inches(0)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(fill)
    frame = cell.text_frame
    frame.word_wrap = False
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = _FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    run.font.color.rgb = _rgb(color)


def _environment_label(name: str, labels: dict[str, str]) -> str:
    return labels.get(name) or name


def _profile_rows(profile: dict[str, Any], kind: str) -> list[tuple[str, dict[str, Any]]]:
    """The rows of a profile table: each category (the ones that lose most first) and then its KPIs."""
    rows = profile['to_maximum'] if kind == 'maximum' else profile['to_reference']
    total_key = 'total_gap_to_maximum' if kind == 'maximum' else 'total_gap_to_reference'
    groups: dict[tuple, list[dict[str, Any]]] = {}
    for row in rows:
        groups.setdefault((row['service'], row.get('category') or ''), []).append(row)
    ordered = sorted(groups.items(), key=lambda item: sum(row[total_key] or 0 for row in item[1]) * (-1 if kind == 'maximum' else 1))
    result: list[tuple[str, dict[str, Any]]] = []
    for (service, category), members in ordered:
        result.append(('group', {'service': service, 'category': category, 'rows': members}))
        result.extend(('kpi', row) for row in members)
    return result


def _profile_table(slide, profile: dict[str, Any], *, kind: str, left: float, width: float,
                   environment_labels: dict[str, str], header_color: str) -> None:
    lines = _profile_rows(profile, kind)
    columns = profile['columns']
    title = (f"{profile['label']} · points lost per KPI" if kind == 'maximum'
             else f"{profile['label']} · GAP to {profile['reference_label']} per KPI")
    count = max(1, len(lines))
    top = 1.62
    # The banner and the column headers have explicit heights, so the bars drawn under the values line up with their
    # rows whatever the KPI count.
    banner_height, header_height = .3, .3
    row_height = min(.22, (PROFILE_TABLE_BOTTOM - top - banner_height - header_height) / count)
    size = 7 if row_height >= .17 else 6
    environment_width = .7
    total_width = 1.0
    kpi_width = width - environment_width * len(columns) - total_width
    widths = [kpi_width, *([environment_width] * len(columns)), total_width]
    gap_key = 'gap_to_maximum' if kind == 'maximum' else 'gap_to_reference'
    total_key = 'total_gap_to_maximum' if kind == 'maximum' else 'total_gap_to_reference'
    # Data bars in every value column of the KPI rows, each column with its own scale; the values are written over
    # them, right-aligned.
    kpis = [row for line, row in lines if line == 'kpi']
    bar_columns = [[row[gap_key].get(name) for row in kpis] for name in columns]
    bar_columns.append([row[total_key] for row in kpis])
    for offset, values in enumerate(bar_columns):
        column = 1 + offset
        scale = max((abs(value) for value in values if value is not None), default=0) or 1
        bar_left = left + sum(widths[:column]) + .03
        bar_width = widths[column] - .06
        line_index = 0
        for index, (line, row) in enumerate(lines):
            if line != 'kpi':
                continue
            value = values[line_index]
            line_index += 1
            if value is None or value == 0:
                continue
            color = LOSS_COLOR if kind == 'maximum' or value < 0 else WIN_COLOR
            bar_top = top + banner_height + header_height + index * row_height + .02
            if kind == 'maximum':
                bar = _gradient_bar(slide, bar_left, bar_top, max(.02, bar_width * abs(value) / scale), row_height - .04, color)
            else:
                # A GAP grows from the middle of its cell: to the left when behind the reference, to the right when ahead.
                length = max(.02, bar_width / 2 * abs(value) / scale)
                middle = bar_left + bar_width / 2
                bar = _gradient_bar(slide, middle - length if value < 0 else middle, bar_top, length, row_height - .04,
                                    color, reverse=value < 0)
            bar.name = 'Scoring GAP Bar'
    shape, table = _table(slide, len(lines) + 2, len(widths), left, top, width, row_height, widths)
    table.rows[0].height = Inches(banner_height)
    table.rows[1].height = Inches(header_height)
    shape.name = f"Scoring KPI GAP {'Maximum' if kind == 'maximum' else 'Reference'} Table"
    table.cell(0, 0).merge(table.cell(0, len(widths) - 1))
    _cell(table.cell(0, 0), title, size=size + 3, bold=True, color='#FFFFFF', fill=header_color, align=PP_ALIGN.RIGHT)
    headers = ['KPI', *[_environment_label(name, environment_labels) for name in columns], 'Total']
    for column, text in enumerate(headers):
        _cell(table.cell(1, column), text, size=size, bold=True, fill='#E9ECEF',
              align=PP_ALIGN.LEFT if column < 1 else PP_ALIGN.CENTER)
        table.cell(1, column).text_frame.word_wrap = True
    for row_index, (line, row) in enumerate(lines, start=2):
        if line == 'group':
            fill = '#F1F4F7'
            _cell(table.cell(row_index, 0), f"{row['category']} · {row['service']}", size=size, bold=True, fill=fill)
            for offset, name in enumerate(columns):
                _cell(table.cell(row_index, 1 + offset), _points(sum(item[gap_key].get(name) or 0 for item in row['rows']), 2),
                      size=size, bold=True, fill=fill, align=PP_ALIGN.RIGHT)
            _cell(table.cell(row_index, len(widths) - 1), _points(sum(item[total_key] or 0 for item in row['rows']), 2), size=size,
                  bold=True, fill=fill, align=PP_ALIGN.RIGHT)
            continue
        underline = bool(profile.get('underline_most_reliable') and row.get('most_reliable'))
        _cell(table.cell(row_index, 0), f"   {row.get('kpi') or ''}", size=size, fill='#FFFFFF', underline=underline)
        for offset, name in enumerate(columns):
            value = row[gap_key].get(name)
            _cell(table.cell(row_index, 1 + offset), _points(value, 2), size=size, fill=None, align=PP_ALIGN.RIGHT)
        _cell(table.cell(row_index, len(widths) - 1), _points(row[total_key], 2), size=size, bold=True, fill=None,
              align=PP_ALIGN.RIGHT)


def _profile_pairs(profiles: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    """The profiles two by two, never mixing two locations or campaigns on a slide."""
    by_context: dict[tuple, list[dict[str, Any]]] = {}
    for profile in profiles:
        by_context.setdefault(tuple(sorted(profile['context'].items())), []).append(profile)
    return [members[start:start + 2] for members in by_context.values() for start in range(0, len(members), 2)]


def _profile_note(slide, runs: list[tuple[str, bool, str | None]]) -> None:
    note = slide.shapes.add_textbox(Inches(.35), Inches(PROFILE_TABLE_BOTTOM + .14), Inches(12.6), Inches(.36))
    note.name = 'Scoring KPI GAP Profile Note'
    _write_runs(note.text_frame, runs, size=8.5, color='#4A5B65')


def add_kpi_gap_profile_slides(presentation, profiles: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                               environment_labels: dict[str, str], new_slide: Callable,
                               operator_colors: dict[str, str], scope: ScopeLabels | None = None) -> None:
    """The GAP of each compared operator to the reference per KPI, grouped by category, two operators per slide (the
    points they lose against the maximum are in the Points Lost slides, see add_points_per_kpi_slides)."""
    compared = [profile for profile in _block_profiles(profiles) if profile['operator'] != profile['reference']]
    for pair in _profile_pairs(compared):
        first = pair[0]
        context = _context_text(first['context'])
        names = ' & '.join(profile['label'] for profile in pair)
        slide = new_slide(presentation, f"KPI GAP Profile — {names} vs {first['reference_label']}", ' · '.join(
            part for part in (_with_scope(subtitle, scope, first['context']), context) if part))
        runs: list[tuple[str, bool, str | None]] = []
        for index, profile in enumerate(pair):
            header = profile.get('color') or operator_colors.get(profile['operator']) or '#C8102E'
            _profile_table(slide, profile, kind='reference', left=.35 + index * 6.43, width=6.2 if len(pair) == 2 else 12.6,
                           environment_labels=environment_labels, header_color=header)
            gap = sum(row['total_gap_to_reference'] for row in profile['to_reference'])
            runs += [('; ' if index else '', False, None), (profile['label'], True, header), (' is ', False, None),
                     (f'{gap:+,.2f} points', True, None), (f" against {profile['reference_label']}", False, None)]
        runs.append(('. The GAP of a KPI is the points of the operator minus those of the reference; each category sums its '
                     'KPIs, the categories furthest behind first.', False, None))
        if first.get('underline_most_reliable'):
            runs.append((' Underlined KPIs are used for the Most Reliable Network scoring.', False, None))
        _profile_note(slide, runs)


def add_points_per_kpi_slides(presentation, profiles: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                              environment_labels: dict[str, str], new_slide: Callable,
                              operator_colors: dict[str, str], scope: ScopeLabels | None = None) -> None:
    """The points each operator (the reference included) loses against the maximum of every KPI, grouped by category,
    two operators per slide."""
    for pair in _profile_pairs(_block_profiles(profiles)):
        first = pair[0]
        context = _context_text(first['context'])
        slide = new_slide(presentation, f"Points lost per KPI — {' & '.join(profile['label'] for profile in pair)}",
                          ' · '.join(part for part in (_with_scope(subtitle, scope, first['context']), context) if part))
        runs: list[tuple[str, bool, str | None]] = []
        for index, profile in enumerate(pair):
            header = profile.get('color') or operator_colors.get(profile['operator']) or '#C8102E'
            _profile_table(slide, profile, kind='maximum', left=.35 + index * 6.43, width=6.2 if len(pair) == 2 else 12.6,
                           environment_labels=environment_labels, header_color=header)
            lost = sum(row['total_gap_to_maximum'] for row in profile['to_maximum'])
            runs += [('; ' if index else '', False, None), (profile['label'], True, header), (' loses ', False, None),
                     (f'{lost:,.2f} {scoring_label} points', True, None)]
        runs.append(('. The points lost of a KPI are its maximum minus the points scored; each category sums its KPIs, the '
                     'categories that lose most first.', False, None))
        if first.get('underline_most_reliable'):
            runs.append((' Underlined KPIs are used for the Most Reliable Network scoring.', False, None))
        _profile_note(slide, runs)


def add_trend_slides(presentation, trends: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                     new_slide: Callable, scope: ScopeLabels | None = None) -> None:
    """A line chart of each operator's scoring across the campaigns of every group, each line ending in the name and
    latest score of its operator."""
    for trend in trends[:MAX_GROUP_SLIDES]:
        slide = new_slide(presentation, f'{scoring_label} Scoring Trend',
                          ' · '.join(part for part in (_with_scope(subtitle, scope), trend['title']) if part))
        data = CategoryChartData()
        data.categories = trend['campaigns']
        for series in trend['series']:
            data.add_series(series['label'], series['points'])
        chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(.55), Inches(1.6), Inches(12.2),
                                             Inches(5.3), data)
        chart_shape.name = 'Scoring Trend Chart'
        chart = chart_shape.chart
        chart.font.name = _FONT
        chart.font.size = Pt(10)
        chart.has_legend = False
        values = [value for series in trend['series'] for value in series['points'] if value is not None]
        low = min(values, default=0)
        chart.value_axis.minimum_scale = max(0, floor((low - 25) / 50) * 50)
        chart.value_axis.maximum_scale = trend['maximum']
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb('#E6EAEE')
        chart.value_axis.major_gridlines.format.line.dash_style = 4
        chart.value_axis.format.line.fill.background()
        chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
        chart.category_axis.format.line.color.rgb = _rgb('#C3CCD5')
        for index, (plot_series, series) in enumerate(zip(chart.series, trend['series'])):
            color = series.get('color') or FALLBACK_COLORS[index % len(FALLBACK_COLORS)]
            plot_series.format.line.color.rgb = _rgb(color)
            plot_series.format.line.width = Pt(3 if series.get('is_reference') else 2.25)
            plot_series.smooth = False
            plot_series.marker.size = 6
            plot_series.marker.format.fill.solid()
            plot_series.marker.format.fill.fore_color.rgb = _rgb('#FFFFFF')
            plot_series.marker.format.line.color.rgb = _rgb(color)
            last = max((position for position, value in enumerate(series['points']) if value is not None), default=None)
            if last is None:
                continue
            point = plot_series.points[last]
            point.marker.size = 9
            point.marker.format.fill.solid()
            point.marker.format.fill.fore_color.rgb = _rgb(color)
            label = point.data_label
            label.position = XL_DATA_LABEL_POSITION.RIGHT
            label.text_frame.text = f"{series['label']} {series['points'][last]:,.0f}"
            for paragraph in label.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.bold = True
                    run.font.color.rgb = _rgb(color)


def _delta_color(value: float | None, scale: float) -> str:
    if value is None or scale <= 0 or value == 0:
        return '#FFFFFF'
    ratio = min(1.0, abs(value) / scale)
    start = (255, 255, 255)
    color = active_theme()['loss' if value < 0 else 'gain']
    end = tuple(int(color[index:index + 2], 16) for index in (1, 3, 5))
    return '#{:02X}{:02X}{:02X}'.format(*(round(a + (b - a) * ratio) for a, b in zip(start, end)))


def _comparison_table(slide, comparison: dict[str, Any], *, left: float, width: float) -> None:
    rows = comparison['rows']
    top = 2.2
    row_height = min(.2, (7.05 - top) / (len(rows) + 2.6))
    size = 7 if row_height >= .17 else 6
    widths = [.5, width - .5 - 4 * .62 - .8, .62, .62, .62, .62, .8]
    heading = slide.shapes.add_textbox(Inches(left), Inches(1.86), Inches(width), Inches(.32))
    title = f"{comparison['label']}" + (f" · {comparison['title']}" if comparison['title'] else '')
    _write(heading.text_frame, [(title, 12, True, '#17232D')], align=PP_ALIGN.CENTER)
    shape, table = _table(slide, len(rows) + 2, len(widths), left, top, width, row_height, widths)
    shape.name = 'Scoring Campaign Comparison Table'
    previous, latest = comparison['previous_campaign'], comparison['latest_campaign']
    headers = ['Service', 'KPI', f'{previous}\nKPI value', f'{previous}\nPoints', f'{latest}\nKPI value',
               f'{latest}\nPoints', 'Δ Points']
    table.rows[0].height = Inches(row_height * 1.6)
    for column, text in enumerate(headers):
        _cell(table.cell(0, column), text, size=size, bold=True, fill='#E9ECEF',
              align=PP_ALIGN.LEFT if column < 2 else PP_ALIGN.CENTER)
        table.cell(0, column).text_frame.word_wrap = True
    scale = max((abs(row['delta']) for row in rows if row['delta'] is not None), default=0)
    for index, row in enumerate(rows, start=1):
        fill = '#FFFFFF'
        _cell(table.cell(index, 0), row['service'].upper(), size=size, fill=fill)
        _cell(table.cell(index, 1), str(row.get('kpi') or ''), size=size, fill=fill)
        for column, key, digits in ((2, 'previous_value', 2), (3, 'previous_points', 2),
                                    (4, 'latest_value', 2), (5, 'latest_points', 2)):
            _cell(table.cell(index, column), _points(row[key], digits), size=size, fill=fill, align=PP_ALIGN.RIGHT)
        _cell(table.cell(index, 6), _points(row['delta'], 2), size=size, bold=True,
              fill=_delta_color(row['delta'], scale), align=PP_ALIGN.RIGHT)
    last = len(rows) + 1
    table.cell(last, 0).merge(table.cell(last, 5))
    _cell(table.cell(last, 0), 'Scoring Points Gap:', size=size + 2, bold=True, fill='#FFFFFF', align=PP_ALIGN.RIGHT)
    total = comparison['total_delta']
    _cell(table.cell(last, 6), f'{total:+,.2f}', size=size + 2, bold=True,
          color=LOSS_COLOR if total < 0 else WIN_COLOR, fill='#FFFFFF', align=PP_ALIGN.RIGHT)


def add_campaign_comparison_slides(presentation, comparisons: list[dict[str, Any]], *, scoring_label: str,
                                   subtitle: str, new_slide: Callable, scope: ScopeLabels | None = None) -> None:
    """Each operator's KPI points in the two latest campaigns, two comparisons per slide."""
    selected = comparisons[:MAX_COMPARISONS]
    for start in range(0, len(selected), 2):
        pair = selected[start:start + 2]
        first = pair[0]
        slide = new_slide(presentation, f'{scoring_label} Campaign Comparison',
                          f"{_with_scope(subtitle, scope)} · {first['previous_campaign']} → {first['latest_campaign']}")
        intro = slide.shapes.add_textbox(Inches(.45), Inches(1.52), Inches(12.4), Inches(.3))
        _write(intro.text_frame, [(
            'KPI values and scoring points of the two latest campaigns; Δ Points is the latest minus the previous '
            'campaign, from the largest loss to the largest gain.', 9, False, '#4A5B65')])
        intro.text_frame.margin_left = 0
        width = 6.15 if len(pair) == 2 else 8.0
        lefts = [.45, 6.78] if len(pair) == 2 else [2.65]
        for comparison, left in zip(pair, lefts):
            _comparison_table(slide, comparison, left=left, width=width)


MAX_SUMMARY_AREAS = 6
MAX_SUMMARY_SLIDES = 4
NATIONAL_FILL = '#FFF4D6'
AREA_FILLS = ('#FFFFFF', '#F2F3F5')


def _summary_number(value: float | None, *, signed: bool = False) -> str:
    if value is None:
        return '—'
    return f'{value:+,.1f}' if signed else f'{value:,.1f}'


def _signed_color(value: float | None) -> str:
    if value is None or round(value, 1) == 0:
        return '#4A5B65'
    return WIN_COLOR if value > 0 else LOSS_COLOR


def _area_summary_table(slide, summary: dict[str, Any], areas: list[dict[str, Any]], *, left: float,
                        width: float) -> bool:
    """The score of each operator per area, with its GAP to the reference and its change from the previous campaign."""
    from src.modules.scoring_insights import _display
    latest, previous = summary['latest_campaign'], summary['previous_campaign']
    reference = summary.get('reference_label') or 'reference'
    headers = ['Area', 'Operator', f"Score\n{_display('Campaign', latest)}" if latest else 'Score']
    if previous:
        headers.append(f"Δ vs\n{_display('Campaign', previous)}")
    headers.append(f'Δ vs\n{reference}')
    fixed = [1.0] * (len(headers) - 2)
    names = (width - sum(fixed)) / 2
    widths = [names, names, *fixed]
    count = sum(len(area['rows']) for area in areas)
    top, header_height = 1.95, .46
    row_height = min(.3, (6.8 - top - header_height) / max(1, count))
    size = 10 if row_height >= .3 else 9 if row_height >= .26 else 8
    shape, table = _table(slide, count + 1, len(headers), left, top, width, row_height, widths)
    shape.name = 'Scoring Area Summary Table'
    table.rows[0].height = Inches(header_height)
    for column, text in enumerate(headers):
        _cell(table.cell(0, column), text, size=size, bold=True, fill='#E9ECEF',
              align=PP_ALIGN.LEFT if column < 2 else PP_ALIGN.CENTER)
        table.cell(0, column).text_frame.word_wrap = True
    incomplete = False
    row = 1
    for index, area in enumerate(areas):
        fill = NATIONAL_FILL if area['kind'] == 'National' else AREA_FILLS[index % 2]
        first = row
        for item in area['rows']:
            color = item.get('color') or '#263746'
            marker = '' if item['complete'] else '*'
            incomplete = incomplete or not item['complete']
            _cell(table.cell(row, 1), item['label'], size=size, bold=item['is_reference'], color=color, fill=fill)
            _cell(table.cell(row, 2), _summary_number(item['points']) + marker, size=size, bold=True, color=color,
                  fill=fill, align=PP_ALIGN.CENTER)
            values = ([item['delta']] if previous else []) + [None if item['is_reference'] else item['gap']]
            for offset, value in enumerate(values, start=3):
                _cell(table.cell(row, offset), '' if item['is_reference'] and offset == len(headers) - 1
                      else _summary_number(value, signed=True), size=size, color=_signed_color(value), fill=fill,
                      align=PP_ALIGN.CENTER)
            row += 1
        if row - 1 > first:
            table.cell(first, 0).merge(table.cell(row - 1, 0))
        _cell(table.cell(first, 0), area['label'], size=size + 1, bold=True, fill=fill)
        table.cell(first, 0).vertical_anchor = MSO_ANCHOR.TOP
        table.cell(first, 0).margin_top = Inches(.04)
    return incomplete


def _area_summary_chart(slide, area: dict[str, Any], campaigns: list[str], *, left: float, top: float,
                        width: float, height: float, low: float, high: float, unit: float) -> None:
    """The points of each operator of one area in every campaign."""
    from src.modules.scoring_insights import _display
    heading = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(.26))
    _write(heading.text_frame, [(area['label'], 10, True, TITLE_BLUE)])
    data = CategoryChartData()
    data.categories = [_display('Campaign', campaign) for campaign in campaigns]
    for item in area['rows']:
        data.add_series(item['label'], item['trend'])
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(left), Inches(top + .22), Inches(width),
                                         Inches(height - .22), data)
    chart_shape.name = f'Scoring Area Summary Chart {area["label"]}'
    chart = chart_shape.chart
    chart.has_legend = False
    chart.font.name = _FONT
    chart.font.size = Pt(7)
    chart.value_axis.minimum_scale = low
    chart.value_axis.maximum_scale = high
    chart.value_axis.major_unit = unit
    chart.value_axis.tick_labels.number_format = '0'
    chart.value_axis.tick_labels.number_format_is_linked = False
    chart.value_axis.major_gridlines.format.line.color.rgb = _rgb('#E3E6E9')
    chart.value_axis.format.line.fill.background()
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    chart.category_axis.format.line.color.rgb = _rgb('#BFC6CC')
    for index, (plot_series, item) in enumerate(zip(chart.series, area['rows'])):
        color = item.get('color') or FALLBACK_COLORS[index % len(FALLBACK_COLORS)]
        plot_series.smooth = False
        plot_series.format.line.color.rgb = _rgb(color)
        plot_series.format.line.width = Pt(1.75)
        plot_series.marker.size = 4
        plot_series.marker.format.fill.solid()
        plot_series.marker.format.fill.fore_color.rgb = _rgb(color)
        plot_series.marker.format.line.color.rgb = _rgb(color)
        # The latest points labelled at the end of each line, as the Crowd charts do.
        last = max((position for position, value in enumerate(item['trend']) if value is not None), default=None)
        if last is not None:
            label = plot_series.points[last].data_label
            label.position = XL_DATA_LABEL_POSITION.RIGHT
            label.text_frame.text = f"{item['trend'][last]:,.0f}"
            for paragraph in label.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(7)
                    run.font.bold = True
                    run.font.color.rgb = _rgb(color)


def add_area_summary_slides(presentation, summaries: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                            new_slide: Callable) -> None:
    """The National scoring next to the chosen cities and each area: a table with each operator's score,
    its change from the previous campaign and its GAP to the reference, and with several campaigns a
    trend chart per area, up to six areas per slide.
    """
    for summary in summaries:
        areas = summary['areas']
        pages = [areas[start:start + MAX_SUMMARY_AREAS]
                 for start in range(0, len(areas), MAX_SUMMARY_AREAS)][:MAX_SUMMARY_SLIDES]
        campaigns = summary['campaigns']
        trend = len(campaigns) >= 2
        values = [value for area in areas for item in area['rows'] for value in item['trend'] if value is not None]
        # The same scale in every chart, in round steps.
        lowest, highest = min(values, default=0), max(values, default=1)
        padding = max(1.0, (highest - lowest) * .15)
        unit = next((step for step in (1, 2, 5, 10, 20, 25, 50, 100, 200, 250, 500)
                     if step * 5 >= highest - lowest + 2 * padding), 1000)
        low = floor((lowest - padding) / unit) * unit
        high = ceil((highest + padding) / unit) * unit
        for number, page in enumerate(pages, start=1):
            suffix = f' · {number}/{len(pages)}' if len(pages) > 1 else ''
            slide = new_slide(presentation, f'{scoring_label} — National & Areas', subtitle + suffix)
            width = 6.3 if trend else 8.4
            incomplete = _area_summary_table(slide, summary, page, left=.45 if trend else 2.45, width=width)
            split = summary.get('time_split') or 'Campaign'
            period = ('every selected CDR aggregated' if split == 'All'
                      else 'per campaign' if split == 'Campaign' else f'{split.lower()} periods of the test start time')
            notes = ['Each area is scored with its own measurements, as the National scoring is: the area scores are '
                     'not shares of the National points',
                     f'time split: {period}',
                     f"maximum scoring: {_points(summary['maximum'])} points"
                     + (' (environments without measurements in an area are scaled)' if summary['scaled'] else '')]
            if incomplete:
                notes.append('* incomplete KPI coverage')
            intro = slide.shapes.add_textbox(Inches(.45), Inches(1.5), Inches(12.4), Inches(.36))
            details = ' · '.join(notes[1:])
            _write(intro.text_frame, [(f'{notes[0]}. {details[:1].upper()}{details[1:]}.', 9, False, '#4A5B65')])
            intro.text_frame.margin_left = 0
            if trend:
                columns = 2
                rows = ceil(len(page) / columns)
                gap = .12
                chart_left, chart_top, chart_width, chart_height = 6.95, 1.95, 5.95, 4.9
                cell_width = (chart_width - gap) / columns
                cell_height = (chart_height - gap * (rows - 1)) / rows
                for index, area in enumerate(page):
                    _area_summary_chart(slide, area, campaigns,
                                        left=chart_left + (index % columns) * (cell_width + gap),
                                        top=chart_top + (index // columns) * (cell_height + gap),
                                        width=cell_width, height=cell_height, low=low, high=high, unit=unit)
                legend = list(summary['operators'].values())
                step = min(1.35, chart_width / max(1, len(legend)))
                for index, style in enumerate(legend):
                    key_left = chart_left + index * step
                    _box(slide, key_left, 6.98, .14, .14, style.get('color') or FALLBACK_COLORS[index % len(FALLBACK_COLORS)])
                    key = slide.shapes.add_textbox(Inches(key_left + .18), Inches(6.94), Inches(step - .2), Inches(.24))
                    _write(key.text_frame, [(style['label'], 8, False, '#4A5B65')])
                    key.text_frame.margin_left = 0


MAX_LOSS_BARS = 30
MAX_LOSS_MAPS = 8


def _loss_ramp() -> tuple[tuple[int, int, int], tuple[int, int, int]]:
    """The scale of the points lost of the active theme (yellow to red in NetCheck)."""
    low, high = active_theme().get('loss_ramp') or ('#F9D65C', '#C8102E')
    return tuple(int(low[index:index + 2], 16) for index in (1, 3, 5)), tuple(int(high[index:index + 2], 16) for index in (1, 3, 5))


def _loss_color(ratio: float) -> tuple[int, int, int]:
    ratio = max(0.0, min(1.0, ratio))
    low_color, high_color = _loss_ramp()
    return tuple(round(low + (high - low) * ratio) for low, high in zip(low_color, high_color))


def render_points_loss_map(areas: list[dict[str, Any]], background: list[list[float]], *,
                           width: int = 1500, height: int = 1300, labels: int = 12) -> bytes:
    """PNG map: the test locations in grey, each place as a bubble and each route as dots, coloured by points lost."""
    from io import BytesIO
    from math import cos, radians, sqrt
    from PIL import Image, ImageDraw, ImageFont

    coordinates = [(lat, lon) for lat, lon in background]
    for area in areas:
        if area.get('latitude') is not None:
            coordinates.append((area['latitude'], area['longitude']))
        coordinates.extend((lat, lon) for lat, lon in area.get('route') or [])
    image = Image.new('RGB', (width, height), (255, 255, 255))
    if not coordinates:
        output = BytesIO()
        image.save(output, format='PNG')
        return output.getvalue()
    latitudes = [lat for lat, _lon in coordinates]
    longitudes = [lon for _lat, lon in coordinates]
    scale_x = cos(radians(sum(latitudes) / len(latitudes)))
    min_x, max_x = min(longitudes) * scale_x, max(longitudes) * scale_x
    min_y, max_y = min(latitudes), max(latitudes)
    margin = 60
    span = max(max_x - min_x, 1e-6) / (width - 2 * margin), max(max_y - min_y, 1e-6) / (height - 2 * margin)
    unit = max(span)
    offset_x = (width - (max_x - min_x) / unit) / 2
    offset_y = (height - (max_y - min_y) / unit) / 2

    def project(lat: float, lon: float) -> tuple[float, float]:
        return offset_x + (lon * scale_x - min_x) / unit, height - offset_y - (lat - min_y) / unit

    draw = ImageDraw.Draw(image)
    for lat, lon in background:
        x, y = project(lat, lon)
        draw.ellipse((x - 2, y - 2, x + 2, y + 2), fill=(214, 220, 226))
    peak = max((area['points'] for area in areas), default=0) or 1
    for area in sorted(areas, key=lambda item: item['points']):
        color = _loss_color(area['points'] / peak)
        if area.get('kind') == 'route' and area.get('route'):
            for lat, lon in area['route']:
                x, y = project(lat, lon)
                draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=color)
        elif area.get('latitude') is not None:
            x, y = project(area['latitude'], area['longitude'])
            radius = 8 + 34 * sqrt(area['points'] / peak)
            draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=color, outline=(255, 255, 255), width=3)
    try:
        font = ImageFont.truetype('Arial.ttf', 26)
    except OSError:
        font = ImageFont.load_default()
    places = [area for area in areas if area.get('kind') != 'route' and area.get('latitude') is not None]
    for area in places[:labels]:
        x, y = project(area['latitude'], area['longitude'])
        radius = 8 + 34 * sqrt(area['points'] / peak)
        draw.text((x + radius + 6, y - 14), area['name'], fill=(23, 35, 45), font=font,
                  stroke_width=3, stroke_fill=(255, 255, 255))
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def _ring_centroid(ring: list[list[float]]) -> tuple[float, float, float]:
    """Area and centroid (longitude, latitude) of one ring (shoelace formula)."""
    area = cx = cy = 0.0
    for (x1, y1), (x2, y2) in zip(ring, ring[1:] + ring[:1]):
        cross = x1 * y2 - x2 * y1
        area += cross
        cx += (x1 + x2) * cross
        cy += (y1 + y2) * cross
    if abs(area) < 1e-12:
        return 0.0, ring[0][0], ring[0][1]
    return abs(area) / 2, cx / (3 * area), cy / (3 * area)


def render_points_loss_choropleth(boundaries: dict[str, list[list[list[float]]]], values: dict[str, float], *,
                                  height: int = 1300, labels: int = 8, label_room: int = 420,
                                  places: list[tuple[str, str]] | None = None,
                                  outlines: list[list[list[float]]] | None = None) -> bytes:
    """PNG map: every polygon of the layer, those with losses coloured by the points lost.

    The areas that lose most are labelled, or ``places`` (a label and the area it is drawn on).
    """
    from io import BytesIO
    from math import cos, radians
    from PIL import Image, ImageDraw, ImageFont

    rings = [(name, ring) for name, polygons in boundaries.items() for ring in polygons if len(ring) > 2]
    longitudes = [x for _name, ring in rings for x, _y in ring]
    latitudes = [y for _name, ring in rings for _x, y in ring]
    if not rings:
        image = Image.new('RGB', (height, height), (255, 255, 255))
        output = BytesIO()
        image.save(output, format='PNG')
        return output.getvalue()
    west, east, south, north = min(longitudes), max(longitudes), min(latitudes), max(latitudes)
    scale_x = cos(radians((south + north) / 2))
    margin = 30
    unit = (north - south) / (height - 2 * margin) or 1e-6
    # Room on the right for the labels of the areas near the east coast.
    width = max(int((east - west) * scale_x / unit) + 2 * margin + label_room, height // 3)
    image = Image.new('RGB', (width, height), (255, 255, 255))
    draw = ImageDraw.Draw(image)

    def project(x: float, y: float) -> tuple[float, float]:
        return margin + (x - west) * scale_x / unit, height - margin - (y - south) / unit

    # The outlines of the countries, in grey under their areas.
    for ring in outlines or []:
        if len(ring) > 2:
            draw.polygon([project(x, y) for x, y in ring], fill=(241, 243, 245), outline=(225, 228, 232))
    peak = max(values.values(), default=0) or 1
    floor_value = min((value for value in values.values() if value > 0), default=0)
    span = (peak - floor_value) or 1
    for name, ring in rings:
        value = values.get(name, 0)
        # When every area loses the same points, they all lose most: the strongest colour.
        fill = _loss_color((value - floor_value) / span if peak > floor_value else 1) if value > 0 else (234, 236, 239)
        draw.polygon([project(x, y) for x, y in ring], fill=fill, outline=(255, 255, 255))
    try:
        font = ImageFont.truetype('Arial.ttf', 30)
    except OSError:
        font = ImageFont.load_default()
    named = places if places is not None else [
        (name, name) for name, value in sorted(values.items(), key=lambda item: -item[1]) if value > 0]
    for label, name in named[:labels]:
        polygons = boundaries.get(name) or []
        if not polygons:
            continue
        _area, x, y = max((_ring_centroid(ring) for ring in polygons), key=lambda item: item[0])
        px, py = project(x, y)
        draw.ellipse((px - 5, py - 5, px + 5, py + 5), fill=(23, 35, 45))
        draw.text((px + 10, py - 16), label, fill=(23, 35, 45), font=font, stroke_width=3, stroke_fill=(255, 255, 255))
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def _loss_scale(slide, left: float, top: float, width: float, low: float, high: float) -> None:
    """Colour scale under the map, from the smallest to the largest loss of an area."""
    scale = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + .55), Inches(top), Inches(width - 1.1), Inches(.13))
    scale.name = 'Scoring Points Lost Scale'
    scale.fill.gradient()
    scale.fill.gradient_angle = 0
    stops = scale.fill.gradient_stops
    stops[0].position, stops[1].position = 0, 1
    stops[0].color.rgb, stops[1].color.rgb = (RGBColor(*color) for color in _loss_ramp())
    scale.line.fill.background()
    scale.shadow.inherit = False
    for text, position, align in ((_points(low, 1), left, PP_ALIGN.RIGHT), (_points(high, 1), left + width - .5, PP_ALIGN.LEFT)):
        label = slide.shapes.add_textbox(Inches(position), Inches(top - .06), Inches(.5), Inches(.25))
        label.text_frame.margin_left = label.text_frame.margin_right = 0
        _write(label.text_frame, [(text, 8, False, '#6B7A85')], align=align)


def _loss_bars(slide, bars: list[dict[str, Any]], *, left: float, top: float, width: float, height: float) -> None:
    from pptx.chart.data import CategoryChartData

    data = CategoryChartData()
    data.categories = [area['name'] for area in reversed(bars)]
    data.add_series('Points lost', [round(area['points'], 2) for area in reversed(bars)])
    chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(left), Inches(top),
                                         Inches(width), Inches(height), data)
    chart_shape.name = 'Scoring Points Lost Bars'
    chart = chart_shape.chart
    chart.has_title = False
    chart.has_legend = False
    chart.font.name = _FONT
    chart.font.size = Pt(7 if len(bars) > 20 else 8)
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = _rgb('#E3E6E9')
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
    plot = chart.plots[0]
    plot.gap_width = 35
    plot.has_data_labels = True
    plot.data_labels.number_format = '0.0'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(7)
    plot.data_labels.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
    peak = max((area['points'] for area in bars), default=0) or 1
    low = min((area['points'] for area in bars), default=0)
    series = plot.series[0]
    for index, area in enumerate(reversed(bars)):
        point = series.points[index]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor(*_loss_color((area['points'] - low) / ((peak - low) or 1)))


LOSS_FIELD_TITLES = {'Area': 'Map Area', 'City': 'City', 'Cluster': 'Cluster', 'Region': 'Region'}


def _short_environment(label: str) -> str:
    """The short name of an environment: the last word after its last " - " (Drive - Connecting Roads reads Roads)."""
    words = str(label).split(' - ')[-1].split()
    return words[-1] if words else str(label)


def add_points_loss_slides(presentation, maps: list[dict[str, Any]], background: list[list[float]], *,
                           scoring_label: str, subtitle: str, new_slide: Callable,
                           boundaries: Callable[[str], dict] | None = None,
                           environment_labels: dict[str, str] | None = None,
                           outlines: Callable[[str], list] | None = None, scope: ScopeLabels | None = None) -> None:
    """Points lost per area, like the NetCheck points-loss slides: a map and the areas that lose most.

    The main slide colours the map areas (the ITL3 areas in the UK, the workspace's Map Areas
    elsewhere) by the points of the city or Connecting Roads route that loses most among those
    measured in each one, as the web map does by default, and ranks the cities and routes; without
    map areas the cities are bubbles on the test locations. Clusters and Regions with polygons get
    their own slide.
    """
    from io import BytesIO

    boundaries = boundaries or (lambda _field: {})
    outlines = outlines or (lambda _field: [])
    by_series: dict[tuple, dict[str, dict[str, Any]]] = {}
    for item in maps:
        series = (item['operator'], tuple(sorted((key, str(value)) for key, value in item['context'].items())))
        by_series.setdefault(series, {})[item['field']] = item
    slides = []
    for fields in by_series.values():
        city = fields.get('City')
        if city is not None:
            layer = fields.get('Area')
            slides.append((city, 'City', layer if layer is not None and boundaries('Area') else None))
        for field in ('Cluster', 'Region'):
            if fields.get(field) is not None and boundaries(field):
                slides.append((fields[field], field, fields[field]))
    for ranking, field, layer in slides[:MAX_LOSS_MAPS]:
        context = _context_text(ranking['context'])
        level = LOSS_FIELD_TITLES.get(field, field)
        slide = new_slide(presentation, f"Points Lost per {level} — {ranking['label']}",
                          ' · '.join(part for part in (_with_scope(subtitle, scope, ranking['context']), context) if part))
        bars = [area for area in ranking['areas'] if area['name'] != 'Not specified'][:MAX_LOSS_BARS]
        # All environments: each city or route names the environments where it loses points (not per Region or Cluster).
        labels = environment_labels or {}
        bars = [{**area, 'name': f"{area['name']} ({', '.join(_short_environment(labels.get(name, name)) for name in area['environments'])})"}
                if area.get('environments') and field == 'City' else area for area in bars]
        share = sum(area['points'] for area in bars) / ranking['total'] if ranking['total'] else 0
        plural = {'City': 'cities and routes', 'Cluster': 'clusters', 'Region': 'regions'}.get(field, 'areas')
        by_rows = layer is not None and layer['field'] == 'Area' and field == 'City'
        area_label = (layer or {}).get('area_label') or LOSS_FIELD_TITLES.get((layer or {}).get('field'), 'area')
        map_text = f"The map colours each {area_label} by the points " if layer is not None else 'The map shows where '
        # The operator, large and in its colour, above the map and the ranking.
        heading = slide.shapes.add_textbox(Inches(.45), Inches(1.35), Inches(12.4), Inches(.5))
        heading.name = 'Scoring Points Lost Operator'
        _write(heading.text_frame, [(str(ranking['label']), 24, True, ranking.get('color') or '#17232D')])
        intro = slide.shapes.add_textbox(Inches(.45), Inches(1.85), Inches(12.4), Inches(.5))
        intro.text_frame.word_wrap = True
        _write_runs(intro.text_frame, [
            (map_text, False, None), (ranking['label'], True, ranking.get('color') or None),
            # The points lost and the share of the listed rows are bold, as on the web page.
            *([(' loses in the city or route that loses most among those measured in it (', False, None),
               (f"{_points(ranking['total'], 1)} {scoring_label} points", True, None), (' lost in total). ', False, None)]
              if by_rows else
              [(' loses in it (', False, None), (f"{_points(ranking['total'], 1)} {scoring_label} points", True, None),
               (' lost in total). ', False, None)] if layer is not None
              else [(f' loses its {scoring_label} points (', False, None), (_points(ranking['total'], 1), True, None),
                    (' in total). ', False, None)]),
            (f'The {len(bars)} {plural} listed', True, None),
            (' account for ', False, None), (f'{share * 100:.0f}%', True, None), (' of the points lost.', False, None),
        ], size=11, color='#17232D')
        map_top, map_height, map_width = 2.35, 4.5, 6.0
        if layer is not None:
            # As the web map does by default: the points of the city or route that loses most among those
            # measured in each area.
            from src.modules.scoring_points_loss import row_area_links, row_area_values
            values = (row_area_values(ranking['areas']) if by_rows else
                      {area['name']: area['points'] for area in layer['areas'] if area['name'] != 'Not specified'})
            # The cities and routes that lose most are labelled once, on their main area.
            if by_rows:
                links = row_area_links(ranking['areas'])
                places = [(area['name'], links[area['name']][0]) for area in ranking['areas'] if links.get(area['name'])]
            else:
                places = None
            image = render_points_loss_choropleth(boundaries(layer['field']), values, places=places,
                                                  outlines=outlines(layer['field']))
        else:
            image = render_points_loss_map(ranking['areas'], background)
        from PIL import Image
        with Image.open(BytesIO(image)) as opened:
            ratio = opened.width / opened.height
        width = min(map_width, map_height * ratio)
        height = width / ratio
        picture = slide.shapes.add_picture(BytesIO(image), Inches(.45 + (map_width - width) / 2), Inches(map_top),
                                           Inches(width), Inches(height))
        picture.name = 'Scoring Points Lost Map'
        _loss_bars(slide, bars, left=6.55, top=2.4, width=6.35, height=4.55)
        if layer is not None and values:
            positive = [value for value in values.values() if value > 0]
            _loss_scale(slide, .45, 6.95, map_width, min(positive), max(positive))
        else:
            legend = slide.shapes.add_textbox(Inches(.45), Inches(6.9), Inches(6.0), Inches(.3))
            legend.text_frame.margin_left = 0
            _write(legend.text_frame, [('Bubbles are cities and dots the Connecting Roads routes; colour and size grow '
                                        'with the points lost. Grey: test locations.', 8, False, '#6B7A85')])
