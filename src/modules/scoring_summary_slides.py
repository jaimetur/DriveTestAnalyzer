"""Summary, Geographical Areas, GAP by part and Points Lost slides of the Scoring & GAP Analysis exports.

They draw, as the web page does, where each operator stands, its position in every environment and service, the KPIs
that would close its GAP, how the measurements ran, every location ranked and against the reference, the GAP of each
part of the scoring over the campaigns, where the points go and who serves the tests where they are lost (Vendors,
kinds of site, sites and route stretches). Losses and GAPs behind the reference are red and gains green; Data points
are teal and Voice points amber.
"""
from __future__ import annotations

from io import BytesIO
from math import ceil, cos, radians, sqrt
from typing import Any, Callable

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.modules.scoring_insight_slides import (
    DATA_COLOR, FALLBACK_COLORS, VOICE_COLOR, ScopeLabels, _box, _cell, _context_text, _gradient_bar, _points, _rgb, _table,
    _unsigned_axis_ids, _with_scope, _write, _write_runs,
)
from src.modules.scoring_themes import themed

INK = '#17232D'
MUTED = '#5B6B78'
CARD = '#F6F8FA'
LOSS_COLOR = '#E8414F'
GAIN_COLOR = '#4CA65A'
# The position of an operator: green for the first, light green for the second, light red in between, red for the last.
RANK_COLORS = ('#2F8A3E', '#7CBF73', '#F0868E', '#E8414F')
STAKE_COLORS = {'failed': '#A51C30', 'below_par': '#E8414F', 'quality': '#F4A3AA'}
MAX_SCOPE_SLIDES = 6
MAX_OPERATOR_SLIDES = 4
MAX_TABLE_ROWS = 22
CONTENT_TOP = 1.55
CONTENT_BOTTOM = 7.05


# Compared operators share a slide: up to three side by side, four in two rows of two, five or six in two rows of three
# rather than leaving one or two on a slide of their own; more spread evenly over slides of four.
OPERATORS_PER_SLIDE, MAX_OPERATORS_PER_SLIDE = 4, 6
PANEL_GAP_X, PANEL_GAP_Y = .25, .2
# The panels end above the note at the bottom of the slide.
PANELS_BOTTOM = CONTENT_BOTTOM - .32


def operator_groups(items: list) -> list[list]:
    """The operators (or their items) of each slide: up to six together, else slides of four at most sharing them
    evenly, so a slide never has one alone."""
    count = len(items)
    if count <= MAX_OPERATORS_PER_SLIDE:
        return [list(items)] if items else []
    slides = ceil(count / OPERATORS_PER_SLIDE)
    per, extra = divmod(count, slides)
    groups, start = [], 0
    for index in range(slides):
        size = per + (1 if index < extra else 0)
        groups.append(list(items[start:start + size]))
        start += size
    return groups


def operator_frames(count: int) -> list[tuple[float, float, float, float]]:
    """(left, top, width, height) of the panel of each operator of a slide: one centred, two or three side by side,
    four in two rows of two, five or six in two rows of three."""
    if count <= 1:
        return [(2.45, CONTENT_TOP, 8.4, PANELS_BOTTOM - CONTENT_TOP)]
    columns = 2 if count in (2, 4) else 3
    rows = 1 if count <= 3 else 2
    width = (12.45 - PANEL_GAP_X * (columns - 1)) / columns
    height = (PANELS_BOTTOM - CONTENT_TOP - PANEL_GAP_Y * (rows - 1)) / rows
    return [(.45 + (index % columns) * (width + PANEL_GAP_X), CONTENT_TOP + (index // columns) * (height + PANEL_GAP_Y), width, height)
            for index in range(count)]


def _fitted_rows(height: float, rows: int, *, largest: float, smallest: float = .2) -> float:
    """The height of each of ``rows`` table rows in ``height`` (between ``smallest`` and ``largest``)."""
    return max(smallest, min(largest, height / max(1, rows)))


def _font(row_height: float, size: float, largest: float) -> float:
    """``size`` made smaller in proportion with rows lower than ``largest``, down to 6 pt."""
    return max(6.0, round(size * min(1.0, row_height / largest) * 2) / 2)


def _signed(value: float | None, digits: int = 1) -> str:
    return '—' if value is None else f'{value:+,.{digits}f}'.replace('-', '−')


def _ordinal(number: int) -> str:
    suffix = 'th' if 10 <= number % 100 <= 20 else {1: 'st', 2: 'nd', 3: 'rd'}.get(number % 10, 'th')
    return f'{number}{suffix}'


def _color(value: Any, index: int = 0) -> str:
    text = str(value or '')
    return text if text.startswith('#') and len(text) == 7 else FALLBACK_COLORS[index % len(FALLBACK_COLORS)]


def _tone(value: float | None) -> str:
    if value is None or abs(value) < .05:
        return MUTED
    return GAIN_COLOR if value > 0 else LOSS_COLOR


def _blend(color: str, alpha: float) -> str:
    """The colour (in the active theme) over white at ``alpha`` opacity."""
    red, green, blue = (int(themed(color).lstrip('#')[index:index + 2], 16) for index in (0, 2, 4))
    return '#{:02X}{:02X}{:02X}'.format(*(round(255 + (value - 255) * alpha) for value in (red, green, blue)))


def _text(slide, text: str, left: float, top: float, width: float, height: float, *, size: float = 10,
          bold: bool = False, color: str = INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name: str | None = None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    _write(box.text_frame, [(text, size, bold, color)], align=align)
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    box.text_frame.vertical_anchor = anchor
    if name:
        box.name = name
    return box


def _subtitle(subtitle: str, scope: ScopeLabels | None, context: dict[str, Any] | None, *extra: str) -> str:
    return ' · '.join(part for part in (_with_scope(subtitle, scope, context), *extra) if part)


def _diverging(slide, left: float, top: float, width: float, height: float, value: float | None, scale: float, *,
               label: bool = True, size: float = 8) -> None:
    """A bar from the middle of a track: red to the left (behind), green to the right (ahead), and its value."""
    track = width - (.55 if label else 0)
    middle = left + track / 2
    line = _box(slide, middle - .005, top, .01, height, '#C3CCD5')
    line.name = 'Scoring Diverging Axis'
    if value:
        share = min(1.0, abs(value) / scale) * track / 2 if scale else 0
        bar = _box(slide, middle - share if value < 0 else middle, top + height * .18, max(.01, share), height * .64,
                   LOSS_COLOR if value < 0 else GAIN_COLOR)
        bar.name = 'Scoring Diverging Bar'
    if label:
        _text(slide, _signed(value), left + track + .05, top - .03, .5, height + .06, size=size, bold=True,
              color=_tone(value), anchor=MSO_ANCHOR.MIDDLE)


def _banner(slide, text_runs: list[tuple[str, bool, str | None]], left: float, top: float, width: float, color: str) -> None:
    """The heading of a table: a light band with a stripe in the colour of its operator."""
    _box(slide, left, top, width, .3, _blend(color, .1)).name = 'Scoring Table Banner'
    _box(slide, left, top, .06, .3, color).name = 'Scoring Table Banner Stripe'
    box = slide.shapes.add_textbox(Inches(left + .12), Inches(top + .03), Inches(width - .16), Inches(.26))
    _write_runs(box.text_frame, text_runs, size=10, color=INK)


def _row_bars(slide, values: list[float | None], *, left: float, top: float, width: float, row_height: float,
              color: str | None = None, scale: float | None = None) -> None:
    """Data bars under the values of a table column (gradient, like Excel), one per row. Bars of one colour (such as
    points lost) grow from the left; signed GAPs grow from the middle: to the left when behind, to the right when ahead."""
    peak = scale if scale is not None else max((abs(value) for value in values if value), default=0)
    if not peak:
        return
    for index, value in enumerate(values):
        if not value:
            continue
        bar_top, bar_height = top + index * row_height + .025, row_height - .05
        if color:
            bar = _gradient_bar(slide, left + .03, bar_top, max(.02, (width - .06) * min(1, abs(value) / peak)), bar_height, color)
        else:
            length = max(.02, (width - .06) / 2 * min(1, abs(value) / peak))
            middle = left + width / 2
            bar = _gradient_bar(slide, middle - length if value < 0 else middle, bar_top, length, bar_height,
                                LOSS_COLOR if value < 0 else GAIN_COLOR, reverse=value < 0)
        bar.name = 'Scoring Row Bar'


def _to_front(slide, shape) -> None:
    """Draw a table over the bars added after it, so its values stay readable over them."""
    tree = slide.shapes._spTree
    tree.remove(shape._element)
    tree.append(shape._element)


def _header(table, headers: list[str], *, size: float, numeric_from: int = 1) -> None:
    for column, text in enumerate(headers):
        _cell(table.cell(0, column), text, size=size, bold=True, fill='#E9ECEF',
              align=PP_ALIGN.LEFT if column < numeric_from else PP_ALIGN.CENTER)
        table.cell(0, column).text_frame.word_wrap = True


def _pages(rows: list[Any], size: int = MAX_TABLE_ROWS) -> list[list[Any]]:
    return [rows[start:start + size] for start in range(0, len(rows), size)] or [[]]


# ---- Summary ----------------------------------------------------------------------------------------------------------
def add_standings_slides(presentation, standings: list[dict[str, Any]], *, scoring_label: str, subtitle: str,
                         new_slide: Callable, scope: ScopeLabels | None = None,
                         compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """Where each operator stands: every operator from the highest score (its Data and Voice points), then a panel
    per compared operator with its position, its GAP to the reference by part and a sentence that says it."""
    for item in standings[:MAX_SCOPE_SLIDES]:
        slide = new_slide(presentation, f'{scoring_label} — Standings',
                          _subtitle(subtitle, scope, item['context'], item.get('title') or '', item.get('campaign') or ''))
        rows = item['operators']
        maximum = item['maximum'] or 1
        row_height = .34 if len(rows) <= 5 else .27
        for index, row in enumerate(rows):
            color = _color(row.get('color'), index)
            top = CONTENT_TOP + index * row_height
            pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.45), Inches(top + .04), Inches(.48), Inches(row_height - .08))
            pill.fill.background()
            pill.line.color.rgb = _rgb(color)
            pill.line.width = Pt(1.25)
            pill.shadow.inherit = False
            _write(pill.text_frame, [(_ordinal(row['rank']), 8, True, color)], align=PP_ALIGN.CENTER)
            pill.text_frame.margin_top = pill.text_frame.margin_bottom = 0
            pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            _text(slide, f"{row['label']}{' (reference)' if row.get('is_reference') else ''}", 1.05, top, 2.0, row_height,
                  size=10, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            track_left, track_width = 3.1, 6.6
            _box(slide, track_left, top + .08, track_width, row_height - .16, '#EEF1F4').name = 'Scoring Leaderboard Track'
            data_width = track_width * max(0, row['data']) / maximum
            voice_width = track_width * max(0, row['voice']) / maximum
            _box(slide, track_left, top + .08, max(.01, data_width), row_height - .16, DATA_COLOR).name = 'Scoring Leaderboard Data'
            _box(slide, track_left + data_width, top + .08, max(.01, voice_width), row_height - .16,
                 VOICE_COLOR).name = 'Scoring Leaderboard Voice'
            _text(slide, f"{row['total']:,.1f}{'' if row.get('complete', True) else '*'}", 9.8, top, .9, row_height, size=11,
                  bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            if row.get('delta') is not None and item.get('previous_campaign'):
                _text(slide, f"{_signed(row['delta'])} vs {item['previous_campaign']}", 10.85, top, 2.0, row_height, size=8.5,
                      bold=True, color=_tone(row['delta']), anchor=MSO_ANCHOR.MIDDLE)
        legend_top = CONTENT_TOP + len(rows) * row_height + .02
        legend = slide.shapes.add_textbox(Inches(.45), Inches(legend_top), Inches(12.4), Inches(.22))
        legend.name = 'Scoring Leaderboard Legend'
        _write_runs(legend.text_frame, [('■ ', False, DATA_COLOR), ('Data   ', False, None), ('■ ', False, VOICE_COLOR),
                                        ('Voice', False, None), (f'   ·   Points of {maximum:,.0f}', False, None)],
                    size=8, color=MUTED)
        # The tiles of the chosen operators, the reference included (as the page shows by default).
        panels = [row for row in rows if compared(row['operator'], row['label'])][:4]
        if not panels:
            continue
        top = legend_top + .32
        height = CONTENT_BOTTOM - top
        gap = .18
        width = (12.45 - gap * (len(panels) - 1)) / len(panels)
        for index, row in enumerate(panels):
            left = .45 + index * (width + gap)
            color = _color(row.get('color'), index)
            _box(slide, left, top, width, height, CARD).name = 'Scoring Standing Panel'
            _box(slide, left, top, .07, height, color).name = 'Scoring Standing Stripe'
            inner = left + .2
            _text(slide, row['label'], inner, top + .08, width - .3, .28, size=12, bold=True, color=color)
            _text(slide, f"{row['total']:,.1f}", inner, top + .36, 1.6, .45, size=24, bold=True)
            _text(slide, f"{_ordinal(row['rank'])} of {row['of']} · of {maximum:,.0f} points", inner + 1.55, top + .5,
                  width - 1.8, .25, size=9, color=MUTED)
            line_top = top + .9
            if row.get('gap_to_reference') is not None:
                box = slide.shapes.add_textbox(Inches(inner), Inches(line_top), Inches(width - .3), Inches(.25))
                _write_runs(box.text_frame, [(f"GAP to {item['reference_label']} ", False, MUTED),
                                             (_signed(row['gap_to_reference']), True, _tone(row['gap_to_reference'])),
                                             (f"  (Data {_signed(row.get('data_gap'))} · Voice {_signed(row.get('voice_gap'))})", False, MUTED)],
                            size=9.5, color=INK)
                line_top += .27
            if row['rank'] > 1 and row.get('above_label') and row['above_label'] != item['reference_label']:
                _text(slide, f"Behind {row['above_label']}: {_signed(row['gap_to_above'])} · {_signed(row['gap_to_leader'])} to the leader",
                      inner, line_top, width - .3, .25, size=9.5, color=MUTED)
                line_top += .27
            if row.get('below_label') and row['below_label'] != item['reference_label']:
                _text(slide, f"Ahead of {row['below_label']}: {_signed(row['lead_over_below'])}", inner, line_top, width - .3, .25,
                      size=9.5, color=MUTED)
                line_top += .27
            parts = list((row.get('parts') or {}).items())
            if parts and row.get('gap_to_reference') is not None:
                scale = max(max(abs(value) for _name, value in parts), 1)
                _text(slide, f"GAP to {item['reference_label']} by part".upper(), inner, line_top + .05, width - .3, .2, size=7.5,
                      bold=True, color=MUTED)
                line_top += .28
                for name, value in parts:
                    _text(slide, name, inner, line_top, 1.45, .22, size=8.5, anchor=MSO_ANCHOR.MIDDLE)
                    _diverging(slide, inner + 1.45, line_top + .02, min(2.6, width - 1.75), .18, value, scale)
                    line_top += .24
            takeaway = slide.shapes.add_textbox(Inches(inner), Inches(max(line_top + .08, top + height - 1.05)),
                                                Inches(width - .35), Inches(.95))
            takeaway.name = 'Scoring Standing Takeaway'
            _write(takeaway.text_frame, [(row['takeaway'], 9, True, INK)])
            takeaway.text_frame.margin_left = 0


def add_position_matrix_slides(presentation, matrices: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                               environment_labels: dict[str, str], scope: ScopeLabels | None = None,
                               compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """Each compared operator's position and share of the points available in Voice, Data and overall, in every
    environment and in all of them, with the points it missed (up to six operators per slide, see operator_groups)."""
    for item in matrices[:MAX_SCOPE_SLIDES]:
        operators = [style for style in item['operators'] if compared(style['operator'], style['label'])] or item['operators'][:1]
        services = [service for service in ('Voice', 'Data', 'Overall') if any(row['cells'].get(service) for row in item['rows'])]
        for group in operator_groups(operators):
            slide = new_slide(presentation, 'Position — per environment',
                              _subtitle(subtitle, scope, item['context'], item.get('title') or '', item.get('campaign') or ''))
            for index, (style, (left, top, width, height)) in enumerate(zip(group, operator_frames(len(group)))):
                color = _color(style.get('color'), index)
                _text(slide, style['label'], left, top, width, .3, size=13 if len(group) <= 2 else 11, bold=True, color=color)
                rows = item['rows']
                row_height = _fitted_rows(height - .38, len(rows) + 1, largest=.7, smallest=.3)
                label_width = min(1.9, width * .32)
                shape, table = _table(slide, len(rows) + 1, len(services) + 1, left, top + .36, width, row_height,
                                      [label_width, *([(width - label_width) / len(services)] * len(services))])
                shape.name = 'Scoring Position Table'
                size = _font(row_height, 9, .55)
                _header(table, ['', *services], size=size)
                for row_index, row in enumerate(rows, start=1):
                    name = 'All Environments' if row['environment'] == 'Combined' else environment_labels.get(row['environment'], row['environment'])
                    _cell(table.cell(row_index, 0), name, size=size, bold=row['environment'] == 'Combined')
                    for column, service in enumerate(services, start=1):
                        value = (row['cells'].get(service) or {}).get(style['operator'])
                        cell = table.cell(row_index, column)
                        if not value:
                            _cell(cell, '—', size=size, align=PP_ALIGN.CENTER)
                            continue
                        rank_index = 0 if value['rank'] == 1 else 3 if value['rank'] == value['of'] else 1 if value['rank'] == 2 else 2
                        fill = RANK_COLORS[rank_index]
                        _cell(cell, f"{_ordinal(value['rank'])} · {value['percent']:.0f}%", size=_font(row_height, 10, .55), bold=True,
                              color='#FFFFFF', fill=fill, align=PP_ALIGN.CENTER)
                        paragraph = cell.text_frame.add_paragraph()
                        paragraph.alignment = PP_ALIGN.CENTER
                        run = paragraph.add_run()
                        run.text = f"{value['missed']:,.1f} of {value['maximum']:,.0f} missed"
                        run.font.size = Pt(_font(row_height, 7.5, .55))
                        run.font.color.rgb = _rgb('#FFFFFF')
            legend_top = CONTENT_BOTTOM - .3
            for index, (label, color) in enumerate((('1st', RANK_COLORS[0]), ('2nd', RANK_COLORS[1]), ('3rd', RANK_COLORS[2]),
                                                     ('Last', RANK_COLORS[3]))):
                _box(slide, .45 + index * 1.0, legend_top + .04, .16, .16, color)
                _text(slide, label, .66 + index * 1.0, legend_top, .7, .24, size=8.5, color=MUTED)
            _text(slide, 'Each tile: the position among the operators measured and the share of the points available; under it, '
                         'the points missed. All Environments scales the environments measured to the full maximum.',
                  4.6, legend_top - .05, 8.3, .4, size=8, color=MUTED)


def add_levers_slides(presentation, levers: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                      scope: ScopeLabels | None = None,
                      compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """The KPIs where matching the reference wins most points, up to six compared operators per slide."""
    by_scope: dict[str, list[dict[str, Any]]] = {}
    for item in levers:
        if compared(item['operator'], item['label']):
            by_scope.setdefault(repr(sorted(item['context'].items())), []).append(item)
    for items in list(by_scope.values())[:MAX_SCOPE_SLIDES]:
        for group in operator_groups(items):
            first = group[0]
            slide = new_slide(presentation, f"Biggest levers — vs {first['reference_label']}",
                              _subtitle(subtitle, scope, first['context'], first.get('title') or '', first.get('campaign') or ''))
            for index, (item, (left, top, width, height)) in enumerate(zip(group, operator_frames(len(group)))):
                color = _color(item.get('color'), index)
                _banner(slide, [(item['label'], True, color),
                                (f" · {-item['total_behind']:,.1f} points behind {item['reference_label']} in {item['count']} KPIs", False, None)],
                        left, top, width, color)
                rows = item['levers']
                row_height = _fitted_rows(height - .38, len(rows) + 1, largest=.36)
                size = _font(row_height, 8, .36)
                # Narrow panels give the KPI most of the width and leave its category out.
                widths = ([width - 3.2, 1.0, 1.0, 1.2] if width >= 6 else
                          [width * share for share in (.46, .16, .16, .22)])
                shape, table = _table(slide, len(rows) + 1, 4, left, top + .38, width, row_height, widths)
                shape.name = 'Scoring Levers Table'
                _header(table, ['KPI', item['label'], item['reference_label'], 'Points to win'], size=size + .5)
                for row_index, row in enumerate(rows, start=1):
                    _cell(table.cell(row_index, 0), f"{row['kpi']} · {row['category']}" if width >= 5 else str(row['kpi']), size=size)
                    table.cell(row_index, 0).text_frame.word_wrap = True
                    _cell(table.cell(row_index, 1), _points(row['value'], 2), size=size, align=PP_ALIGN.RIGHT)
                    _cell(table.cell(row_index, 2), _points(row['reference_value'], 2), size=size, align=PP_ALIGN.RIGHT)
                    _cell(table.cell(row_index, 3), f"{-row['gap']:,.2f}", size=size, bold=True, fill=None, align=PP_ALIGN.RIGHT)
                _row_bars(slide, [-row['gap'] for row in rows], left=left + sum(widths[:3]), top=top + .38 + row_height,
                          width=widths[3], row_height=row_height, color=LOSS_COLOR)
                _to_front(slide, shape)
            _text(slide, 'Points to win: the GAP of the KPI to the reference (the points the operator would add with its value).',
                  .45, CONTENT_BOTTOM - .25, 12.4, .25, size=8, color=MUTED)


def _fitted_size(text: str, width: float, maximum: float, minimum: float = 10) -> float:
    """The largest size up to ``maximum`` at which ``text`` fits on one line ``width`` inches wide (bold figures)."""
    estimate = width * 72 / max(1, .62 * len(str(text)))
    return max(minimum, min(maximum, round(estimate * 2) / 2))


def add_campaign_scope_slides(presentation, scopes: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                              environment_labels: dict[str, str]) -> None:
    """How the measurements ran: campaigns and dates, places, environments and the tests of every operator."""
    for item in scopes[:1]:
        slide = new_slide(presentation, 'How the measurements ran', subtitle)
        tiles = [('Campaigns', item.get('latest_campaign') or '—',
                  f"{len(item['campaigns'])} campaigns: {', '.join(item['campaigns'])}" if len(item['campaigns']) > 1 else 'One campaign')]
        window = item.get('window')
        if window:
            tiles.append(('Measurements', f"{window['first']} → {window['last']}", f"{window['days']} days with tests"))
        if item.get('places'):
            tiles.append(('Places', f"{item['places']['cities']} cities", f"{item['places']['routes']} Connecting Roads routes"))
        environments = [environment for environment in item['environments'] if environment['max_points'] > 0]
        # Each environment on a line of its own.
        tiles.append(('Environments', str(len(environments)),
                      [f"{environment['label']} {environment['max_points']:,.0f} pts" for environment in environments]))
        width = (12.45 - .18 * (len(tiles) - 1)) / len(tiles)
        lines = max(len(detail) if isinstance(detail, list) else 1 for _label, _value, detail in tiles)
        tile_height = max(1.25, .88 + .16 * lines)
        for index, (label, value, detail) in enumerate(tiles):
            left = .45 + index * (width + .18)
            _box(slide, left, CONTENT_TOP, width, tile_height, CARD).name = 'Scoring Scope Tile'
            _text(slide, label.upper(), left + .15, CONTENT_TOP + .1, width - .3, .22, size=8, bold=True, color=MUTED)
            # A long value (such as the dates of the measurements) gets a smaller size, so it keeps one line.
            _text(slide, value, left + .15, CONTENT_TOP + .34, width - .3, .4, size=_fitted_size(value, width - .46, 17), bold=True)
            box = _text(slide, '', left + .15, CONTENT_TOP + .78, width - .3, tile_height - .86, size=8.5, color=MUTED)
            _write(box.text_frame, [(line, 8.5, False, MUTED) for line in (detail if isinstance(detail, list) else [detail])])
            box.text_frame.margin_left = box.text_frame.margin_right = 0
        tests = item.get('tests') or {}
        names = list(dict.fromkeys(name for kinds in tests.values() for counts in kinds.values() for name in counts))
        kinds = [kind for kind in ('data', 'voice', 'speech') if any(kind in value for value in tests.values())]
        operators = [style for style in item['operators'] if style['operator'] in tests]
        if not operators or not names:
            continue
        labels = {'data': 'Data tests', 'voice': 'Calls', 'speech': 'Speech samples'}
        headers = ['Operator', *[f"{labels[kind]} · {environment_labels.get(name, name).split(' - ')[-1]}" for kind in kinds for name in names]]
        top = CONTENT_TOP + tile_height + .4
        _text(slide, f"Tests of each operator in {item.get('latest_campaign') or 'the measurements'}, as scored in each environment.",
              .45, top - .35, 12.4, .3, size=9, color=MUTED)
        widths = [1.8, *([(12.45 - 1.8) / (len(headers) - 1)] * (len(headers) - 1))]
        shape, table = _table(slide, len(operators) + 1, len(headers), .45, top, 12.45, .36, widths)
        shape.name = 'Scoring Scope Tests Table'
        _header(table, headers, size=8.5)
        for row_index, style in enumerate(operators, start=1):
            _cell(table.cell(row_index, 0), style['label'], size=9, bold=True, color=_color(style.get('color'), row_index))
            for column, (kind, name) in enumerate(((kind, name) for kind in kinds for name in names), start=1):
                _cell(table.cell(row_index, column), f"{tests[style['operator']].get(kind, {}).get(name, 0):,}", size=9,
                      align=PP_ALIGN.RIGHT)


# ---- Geographical Areas ----------------------------------------------------------------------------------------------
def add_location_ranking_slides(presentation, rankings: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                                scope: ScopeLabels | None = None,
                                compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """Every location ranked by its share of the points, per compared operator."""
    chosen = [item for item in rankings if compared(item['operator'], item['label'])]
    for item in chosen[:MAX_OPERATOR_SLIDES]:
        plural = {'City': 'Cities and routes', 'Cluster': 'Clusters', 'Region': 'Regions'}.get(item['level'], 'Locations')
        pages = _pages(item['rows'])
        has_vendors = any(row.get('vendors') for row in item['rows'])
        peak = max((row['lost'] for row in item['rows']), default=0)
        for number, rows in enumerate(pages, start=1):
            suffix = f' · {number}/{len(pages)}' if len(pages) > 1 else ''
            slide = new_slide(presentation, f"{plural} ranked — {item['label']}",
                              _subtitle(subtitle, scope, item['context'], item.get('title') or '', item.get('campaign') or '') + suffix)
            color = _color(item.get('color'))
            headers = ['#', 'Location', *(['Vendors (share of tests)'] if has_vendors else []), 'Points %',
                       f"Δ vs {item['previous_campaign']}" if item.get('previous_campaign') else 'Δ', 'Voice %', 'Data %', 'Points lost']
            widths = [.4, 2.6, *([2.4] if has_vendors else []), 1.9, 1.1, 1.0, 1.0, 1.9]
            scale = 12.45 / sum(widths)
            widths = [width * scale for width in widths]
            row_height = min(.3, (CONTENT_BOTTOM - CONTENT_TOP - .45) / (len(rows) + 1))
            size = 9 if row_height >= .26 else 8
            shape, table = _table(slide, len(rows) + 1, len(headers), .45, CONTENT_TOP, 12.45, row_height, widths)
            shape.name = 'Scoring Location Ranking Table'
            _header(table, headers, size=size, numeric_from=3 if has_vendors else 2)
            start = (number - 1) * MAX_TABLE_ROWS
            for row_index, row in enumerate(rows, start=1):
                values = [str(start + row_index), f"{row['location']}{' (route)' if row['kind'] == 'route' else ''}",
                          *([' · '.join(f"{vendor['name']} {vendor['share']:.0f}%" for vendor in row.get('vendors') or [])] if has_vendors else []),
                          f"{row['percent']:.1f}%{'' if row.get('complete', True) else '*'}", _signed(row['delta']),
                          '—' if row['voice_percent'] is None else f"{row['voice_percent']:.1f}%",
                          '—' if row['data_percent'] is None else f"{row['data_percent']:.1f}%", f"{row['lost']:,.1f}"]
                for column, text in enumerate(values):
                    numeric = column >= (3 if has_vendors else 2)
                    _cell(table.cell(row_index, column), text, size=size - (1 if has_vendors and column == 2 else 0),
                          bold=column in (len(values) - 5,), color=_tone(row['delta']) if column == len(values) - 4 else INK,
                          fill=None if column in (len(values) - 5, len(values) - 1) else '#FFFFFF',
                          align=PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT)
            percent_column = len(headers) - 5
            body_top = CONTENT_TOP + row_height
            _row_bars(slide, [row['percent'] for row in rows], left=.45 + sum(widths[:percent_column]), top=body_top,
                      width=widths[percent_column], row_height=row_height, color=color, scale=100)
            _row_bars(slide, [row['lost'] for row in rows], left=.45 + sum(widths[:-1]), top=body_top, width=widths[-1],
                      row_height=row_height, color=LOSS_COLOR, scale=peak)
            _to_front(slide, shape)
            _text(slide, 'Points %: the share of the points available in each location (its environments scaled to the full maximum); '
                         'Δ: the change since the previous campaign in points of percentage.', .45, CONTENT_BOTTOM - .2, 12.4, .25,
                  size=8, color=MUTED)


def add_location_versus_slides(presentation, rankings: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                               scope: ScopeLabels | None = None,
                               compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """Each compared operator against the reference in every location: the GAP of its share of the points, overall,
    in Voice and in Data."""
    chosen = [item for item in rankings if not item.get('is_reference') and compared(item['operator'], item['label'])
              and any(row.get('gap') is not None for row in item['rows'])]
    for item in chosen[:MAX_OPERATOR_SLIDES]:
        rows = sorted((row for row in item['rows'] if row.get('gap') is not None), key=lambda row: row['gap'])
        for number, page in enumerate(_pages(rows), start=1):
            slide = new_slide(presentation, f"Per location — {item['label']} vs {item['reference_label']}",
                              _subtitle(subtitle, scope, item['context'], item.get('title') or '', item.get('campaign') or '')
                              + (f' · {number}' if len(rows) > MAX_TABLE_ROWS else ''))
            scale = max((abs(value) for row in rows for value in (row['gap'], row.get('voice_gap'), row.get('data_gap'))
                         if value is not None), default=1) or 1
            headers = ['Location', 'Points %', 'Voice %', 'Data %', item['label'], item['reference_label']]
            widths = [2.6, 2.75, 2.75, 2.75, .8, .8]
            row_height = min(.3, (CONTENT_BOTTOM - CONTENT_TOP - .45) / (len(page) + 1))
            shape, table = _table(slide, len(page) + 1, len(headers), .45, CONTENT_TOP, 12.45, row_height, widths)
            shape.name = 'Scoring Location Versus Table'
            _header(table, headers, size=9)
            for row_index, row in enumerate(page, start=1):
                _cell(table.cell(row_index, 0), row['location'], size=9)
                for column in (1, 2, 3):
                    _cell(table.cell(row_index, column), '', size=9, fill=None)
                _cell(table.cell(row_index, 4), f"{row['percent']:.1f}%", size=9, align=PP_ALIGN.RIGHT)
                _cell(table.cell(row_index, 5), f"{row['reference_percent']:.1f}%", size=9, align=PP_ALIGN.RIGHT)
                top = CONTENT_TOP + row_index * row_height + row_height * .2
                for column, key in ((1, 'gap'), (2, 'voice_gap'), (3, 'data_gap')):
                    left = .45 + sum(widths[:column]) + .1
                    _diverging(slide, left, top, widths[column] - .2, row_height * .6, row.get(key), scale)
            _to_front(slide, shape)
            _text(slide, f"The GAP in points of percentage of the points available: to the left {item['label']} trails "
                         f"{item['reference_label']}, to the right it is ahead.", .45, CONTENT_BOTTOM - .2, 12.4, .25, size=8, color=MUTED)


# ---- GAP Analysis -----------------------------------------------------------------------------------------------------
def add_gap_parts_slides(presentation, items: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                         scope: ScopeLabels | None = None,
                         compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """The GAP of each part of the scoring over the campaigns, up to six compared operators per slide."""
    by_scope: dict[str, list[dict[str, Any]]] = {}
    for item in items:
        if compared(item['operator'], item['label']):
            by_scope.setdefault(repr(sorted(item['context'].items())), []).append(item)
    for scope_items in list(by_scope.values())[:MAX_SCOPE_SLIDES]:
        for group in operator_groups(scope_items):
            first = group[0]
            slide = new_slide(presentation, f"GAP by part — vs {first['reference_label']}",
                              _subtitle(subtitle, scope, first['context'], first.get('title') or ''))
            for index, (item, (left, top, width, height)) in enumerate(zip(group, operator_frames(len(group)))):
                color = _color(item.get('color'), index)
                _banner(slide, [(item['label'], True, color), (f" minus {item['reference_label']}", False, None)], left,
                        top, width, color)
                campaigns = [campaign or 'GAP' for campaign in item['campaigns']]
                with_latest = len(campaigns) > 1
                rows = [(part['label'], part['values']) for part in item['parts']] + [('Total', item['totals'])]
                values = [value for _label, line in rows for value in line if value is not None]
                scale = max((abs(value) for value in values), default=1) or 1
                part_width = min(1.55, width * .27)
                latest_width = min(1.7, width * .27) if with_latest else 0
                value_width = (width - part_width - latest_width) / len(campaigns)
                widths = [part_width, *([value_width] * len(campaigns)), *([latest_width] if with_latest else [])]
                row_height = _fitted_rows(height - .38, len(rows) + 1, largest=.38)
                size = _font(row_height, 8.5, .38)
                shape, table = _table(slide, len(rows) + 1, len(widths), left, top + .38, width, row_height, widths)
                shape.name = 'Scoring GAP Parts Table'
                _header(table, ['Part', *campaigns, *(['Latest'] if with_latest else [])], size=size)
                for row_index, (label, line) in enumerate(rows, start=1):
                    total = label == 'Total'
                    _cell(table.cell(row_index, 0), label, size=size, bold=total)
                    for column, value in enumerate(line, start=1):
                        fill = '#FFFFFF' if value is None else _blend(LOSS_COLOR if value < 0 else GAIN_COLOR, min(1, abs(value) / scale) * .55)
                        _cell(table.cell(row_index, column), _signed(value), size=size, bold=total, fill=fill, align=PP_ALIGN.RIGHT)
                    if with_latest:
                        _cell(table.cell(row_index, len(widths) - 1), '', size=size, fill=None)
                        _diverging(slide, left + sum(widths[:-1]) + .08, top + .38 + row_index * row_height + row_height * .26,
                                   latest_width - .12, row_height * .48, line[-1], scale, size=max(6.0, size - 1))
            _text(slide, ('Reliability: the KPIs typed Reliable (calls and tests that complete, POLQA below 1.6, the lowest throughputs); '
                          'quality: the others. ' if first.get('by_reliability') else 'Each part is a category of KPIs. ')
                  + 'Red: behind the reference; green: ahead.', .45, CONTENT_BOTTOM - .25, 12.4, .25, size=8, color=MUTED)


# ---- Points Lost ------------------------------------------------------------------------------------------------------
def add_points_at_stake_slides(presentation, items: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                               scope: ScopeLabels | None = None) -> None:
    """Where the points go: the points each operator lost on failed tests, on tests below the KPI thresholds and as a
    quality shortfall."""
    for item in items[:3]:
        slide = new_slide(presentation, 'Points Lost — where the points go',
                          _subtitle(subtitle, scope, item['context'], item.get('title') or '', item.get('campaign') or ''))
        operators = item['operators']
        data = CategoryChartData()
        data.categories = [row['label'] for row in reversed(operators)]
        for kind in item['kinds']:
            data.add_series(kind['label'], [round(row['parts'][kind['key']], 1) for row in reversed(operators)])
        height = min(3.4, .55 + .5 * len(operators))
        chart_shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(.45), Inches(CONTENT_TOP), Inches(12.45),
                                             Inches(height), data)
        chart_shape.name = 'Scoring Points At Stake Chart'
        chart = chart_shape.chart
        _unsigned_axis_ids(chart)
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.font.size = Pt(9)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb('#E3E6E9')
        chart.category_axis.major_tick_mark = XL_TICK_MARK.NONE
        plot = chart.plots[0]
        plot.gap_width = 45
        plot.has_data_labels = True
        plot.data_labels.number_format = '0'
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.font.size = Pt(8)
        plot.data_labels.font.color.rgb = _rgb('#FFFFFF')
        plot.data_labels.position = XL_DATA_LABEL_POSITION.CENTER
        for series, kind in zip(chart.series, item['kinds']):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = _rgb(STAKE_COLORS[kind['key']])
        top = CONTENT_TOP + height + .15
        headers = ['Operator', *[kind['label'] for kind in item['kinds']], 'Points lost', 'Points', 'KPIs that lose most']
        widths = [1.7, 1.45, 1.45, 1.45, 1.2, 1.0, 4.2]
        row_height = min(.36, (CONTENT_BOTTOM - top - .3) / (len(operators) + 1))
        shape, table = _table(slide, len(operators) + 1, len(headers), .45, top, 12.45, row_height, widths)
        shape.name = 'Scoring Points At Stake Table'
        _header(table, headers, size=8.5)
        for row_index, row in enumerate(operators, start=1):
            _cell(table.cell(row_index, 0), row['label'], size=8.5, bold=True, color=_color(row.get('color'), row_index))
            for column, kind in enumerate(item['kinds'], start=1):
                _cell(table.cell(row_index, column), f"{row['parts'][kind['key']]:,.1f}", size=8.5, align=PP_ALIGN.RIGHT)
            _cell(table.cell(row_index, 4), f"{row['lost']:,.1f}", size=8.5, bold=True, align=PP_ALIGN.RIGHT)
            _cell(table.cell(row_index, 5), f"{row['points']:,.1f}", size=8.5, align=PP_ALIGN.RIGHT)
            worst = sorted((kpi for kpis in (row.get('kpis') or {}).values() for kpi in kpis), key=lambda kpi: -kpi['lost'])[:2]
            _cell(table.cell(row_index, 6), ' · '.join(f"{kpi['kpi']} {kpi['lost']:,.1f}" for kpi in worst), size=7.5)
        _text(slide, 'Failed tests: the Reliable KPIs; below the threshold: tests that passed under the threshold of a KPI; '
                     'quality shortfall: what the averages, medians and 90th percentiles miss of their best score.',
              .45, CONTENT_BOTTOM - .22, 12.4, .25, size=8, color=MUTED)


def latest_campaign_items(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Of the items of every campaign (their context names it), those of the latest campaign (and those without a
    campaign): the slides of a single campaign show the latest one, without the operators it did not measure."""
    from src.modules.scoring_insights import campaign_sort_key
    campaigns = {item['context'].get('campaign') for item in items if item['context'].get('campaign')}
    if len(campaigns) <= 1:
        return list(items)
    latest = max(campaigns, key=campaign_sort_key)
    return [item for item in items if item['context'].get('campaign') in (None, latest)]


def add_serving_loss_slides(presentation, items: list[dict[str, Any]], *, title: str, column: str, subtitle: str,
                            new_slide: Callable, scope: ScopeLabels | None = None,
                            compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """The points each compared operator lost per Vendor or kind of serving site, with its share of the tests (up to
    six operators per slide)."""
    chosen = [item for item in latest_campaign_items(items) if compared(item['operator'], item['label'])]
    by_scope: dict[str, list[dict[str, Any]]] = {}
    for item in chosen:
        by_scope.setdefault(repr(sorted(item['context'].items())), []).append(item)
    for scope_items in list(by_scope.values())[:MAX_SCOPE_SLIDES]:
        for group in operator_groups(scope_items):
            first = group[0]
            slide = new_slide(presentation, title, _subtitle(subtitle, scope, first['context'], _context_text(first['context'])))
            for index, (item, (left, top, width, height)) in enumerate(zip(group, operator_frames(len(group)))):
                color = _color(item.get('color'), index)
                _banner(slide, [(item['label'], True, color), (f" · {item['total']:,.1f} points lost", False, None)], left,
                        top, width, color)
                # As many rows as keep a legible size (the Vendors or kinds of site that lose most first).
                rows = item['rows'][:min(16, max(3, int((height - .38) / .2) - 1))]
                with_tests = item.get('has_tests')
                # Narrow panels leave the cities and routes out.
                narrow = width < 5
                headers = [column, 'Points lost', 'Share of points', *(['Share of tests'] if with_tests else []),
                           *([] if narrow else ['Cities and routes that lose most'])]
                if narrow:
                    shares = (.36, .24, .2, .2) if with_tests else (.44, .3, .26)
                    widths = [width * share for share in shares]
                else:
                    widths = [1.45, 1.1, .8, *([.8] if with_tests else []), width - (4.15 if with_tests else 3.35)]
                row_height = _fitted_rows(height - .38, len(rows) + 1, largest=.34)
                size = _font(row_height, 8, .34)
                shape, table = _table(slide, len(rows) + 1, len(headers), left, top + .38, width, row_height, widths)
                shape.name = 'Scoring Serving Loss Table'
                _header(table, headers, size=size)
                for row_index, row in enumerate(rows, start=1):
                    _cell(table.cell(row_index, 0), row['name'], size=size, color=INK if row['identified'] else MUTED)
                    _cell(table.cell(row_index, 1), f"{row['points']:,.2f}", size=size, bold=True, fill=None, align=PP_ALIGN.RIGHT)
                    _cell(table.cell(row_index, 2), f"{row['share']:.1f}%", size=size, align=PP_ALIGN.RIGHT)
                    offset = 3
                    if with_tests:
                        above = row['tests_share'] is not None and row['share'] > row['tests_share'] * 1.25 and row['share'] - row['tests_share'] > 3
                        _cell(table.cell(row_index, 3), '—' if row['tests_share'] is None else f"{row['tests_share']:.1f}%", size=size,
                              color=LOSS_COLOR if above else INK, bold=above, align=PP_ALIGN.RIGHT)
                        offset = 4
                    if not narrow:
                        _cell(table.cell(row_index, offset), ' · '.join(f"{city['name']} {city['points']:,.1f}"
                                                                         for city in (row.get('cities') or [])[:2]), size=max(6.0, size - 1))
                _row_bars(slide, [row['points'] for row in rows], left=left + widths[0], top=top + .38 + row_height,
                          width=widths[1], row_height=row_height, color=LOSS_COLOR)
                _to_front(slide, shape)
            _text(slide, 'Each test belongs to the serving cell of its start (Not identified: no serving cell in the inventories). '
                         'In red, a share of the points lost well above the share of the tests.', .45, CONTENT_BOTTOM - .22,
                  12.4, .25, size=8, color=MUTED)


def add_sites_slides(presentation, items: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                     scope: ScopeLabels | None = None,
                     compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """The sites to review of each compared operator, High priority first."""
    # Every location first with each of its operators, so the slides cover the operators of the first locations.
    chosen = sorted((item for item in items if compared(item['operator'], item['label'])),
                    key=lambda item: repr(sorted(item['context'].items())))
    priority_colors = {'High': '#A51C30', 'Medium': LOSS_COLOR, 'Low': MUTED}
    for item in chosen[:MAX_OPERATOR_SLIDES * 2]:
        rows = sorted(item['rows'], key=lambda row: (['High', 'Medium', 'Low'].index(row['priority']), -row['points']))[:18]
        high = sum(1 for row in item['rows'] if row['priority'] == 'High')
        slide = new_slide(presentation, f"Sites to review — {item['label']}",
                          _subtitle(subtitle, scope, item['context'], item.get('campaign') or ''))
        color = _color(item.get('color'))
        _banner(slide, [(item['label'], True, color),
                        (f" · {len(item['rows'])} sites hold {item['listed_points']:,.1f} of {item['total']:,.1f} points lost · {high} High",
                         False, None)], .45, CONTENT_TOP, 12.45, color)
        headers = ['Priority', 'Site', 'Network', 'Vendor', 'City or route', 'Points lost', 'Failed tests',
                   *[f'{campaign} points' for campaign in item['earlier_campaigns']]]
        widths = [.9, 1.1, 1.4, 1.2, 2.4, 1.6, 1.0, *([1.0] * len(item['earlier_campaigns']))]
        scale = 12.45 / sum(widths)
        widths = [width * scale for width in widths]
        row_height = min(.3, (CONTENT_BOTTOM - CONTENT_TOP - .75) / (len(rows) + 1))
        shape, table = _table(slide, len(rows) + 1, len(headers), .45, CONTENT_TOP + .38, 12.45, row_height, widths)
        shape.name = 'Scoring Sites Table'
        _header(table, headers, size=8.5, numeric_from=5)
        for row_index, row in enumerate(rows, start=1):
            _cell(table.cell(row_index, 0), row['priority'], size=8.5, bold=True, color=priority_colors[row['priority']])
            for column, text in enumerate((row['site'], row.get('network') or '—', row.get('vendor') or '—', row.get('city') or '—'), start=1):
                _cell(table.cell(row_index, column), str(text), size=8.5)
            _cell(table.cell(row_index, 5), f"{row['points']:,.2f}", size=8.5, bold=True, fill=None, align=PP_ALIGN.RIGHT)
            _cell(table.cell(row_index, 6), f"{row['failures']:,.0f}" if row.get('failures') else '—', size=8.5, align=PP_ALIGN.RIGHT)
            for offset, entry in enumerate(row['history'], start=7):
                _cell(table.cell(row_index, offset), f"{entry['points']:,.1f}" if entry.get('points') else '—', size=8.5,
                      align=PP_ALIGN.RIGHT)
        _row_bars(slide, [row['points'] for row in rows], left=.45 + sum(widths[:5]), top=CONTENT_TOP + .38 + row_height,
                  width=widths[5], row_height=row_height, color=LOSS_COLOR)
        _to_front(slide, shape)
        _text(slide, f"High: {item['priority_points']:,.0f}+ points lost and losses on the same site in an earlier campaign; Medium: one of "
                     'the two; Low: neither. Please review their cells: configuration, alarms and changes on the days of the measurements.',
              .45, CONTENT_BOTTOM - .22, 12.4, .25, size=8, color=MUTED)


def render_stretch_map(route: dict[str, Any], color: str, background: list[list[float]], *, width: int = 1100,
                       height: int = 900) -> bytes:
    """PNG map of a route: the test locations in grey, the route dashed and a circle per stretch (area: points lost)."""
    from PIL import Image, ImageDraw, ImageFont

    points = [*(route.get('path') or []), *([entry[0], entry[1]] for entry in route['all_stretches'])]
    image = Image.new('RGB', (width, height), (248, 250, 252))
    if not points:
        output = BytesIO()
        image.save(output, format='PNG')
        return output.getvalue()
    south, north = min(point[0] for point in points) - .15, max(point[0] for point in points) + .15
    west, east = min(point[1] for point in points) - .25, max(point[1] for point in points) + .25
    factor = cos(radians((south + north) / 2))
    margin = 60
    unit = max((east - west) * factor / (width - 2 * margin), (north - south) / (height - 2 * margin)) or 1
    offset_x = (width - (east - west) * factor / unit) / 2
    offset_y = (height - (north - south) / unit) / 2

    def project(lat: float, lon: float) -> tuple[float, float]:
        return offset_x + (lon - west) * factor / unit, height - offset_y - (lat - south) / unit

    draw = ImageDraw.Draw(image)
    for lat, lon in background:
        if south <= lat <= north and west <= lon <= east:
            x, y = project(lat, lon)
            draw.ellipse((x - 2, y - 2, x + 2, y + 2), fill=(213, 219, 226))
    # The tests of the route (a sample of their positions), darker than the other tests.
    for lat, lon in route.get('path') or []:
        x, y = project(lat, lon)
        draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill=(148, 163, 184))
    red, green, blue = (int(color.lstrip('#')[index:index + 2], 16) for index in (0, 2, 4))
    fill = (round(255 + (red - 255) * .6), round(255 + (green - 255) * .6), round(255 + (blue - 255) * .6))
    peak = max((item[2] for item in route['all_stretches']), default=0) or 1
    top = {(row['latitude'], row['longitude']): index for index, row in enumerate(route['stretches'], start=1)}
    try:
        font = ImageFont.truetype('Arial Bold.ttf', 26)
    except OSError:
        font = ImageFont.load_default()
    for lat, lon, lost, *_owner in sorted(route['all_stretches'], key=lambda item: item[2]):
        x, y = project(lat, lon)
        radius = 10 + 52 * sqrt(lost / peak)
        draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill=fill, outline=(255, 255, 255), width=4)
        rank = top.get((lat, lon))
        if rank:
            draw.text((x, y), str(rank), fill=(255, 255, 255), font=font, anchor='mm')
    output = BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def add_stretch_slides(presentation, items: list[dict[str, Any]], background: list[list[float]], *, subtitle: str,
                       new_slide: Callable, scope: ScopeLabels | None = None,
                       compared: Callable[[str, str], bool] = lambda _operator, _label: True) -> None:
    """Where along its routes each compared operator loses its points: the two routes that lose most."""
    chosen = [item for item in latest_campaign_items(items) if not item.get('is_reference') and compared(item['operator'], item['label'])]
    for item in chosen[:MAX_OPERATOR_SLIDES]:
        color = _color(item.get('color'))
        for route in item['routes'][:2]:
            slide = new_slide(presentation, f"Route stretches — {item['label']}",
                              _subtitle(subtitle, scope, item['context'], route['route'], _context_text(item['context'])))
            image = render_stretch_map(route, color, background)
            picture = slide.shapes.add_picture(BytesIO(image), Inches(.45), Inches(CONTENT_TOP), Inches(5.6), Inches(4.58))
            picture.name = 'Scoring Route Stretch Map'
            box = slide.shapes.add_textbox(Inches(6.35), Inches(CONTENT_TOP), Inches(6.55), Inches(.5))
            _write_runs(box.text_frame, [(item['label'], True, color), (f" loses {route['points']:,.1f} points on {route['route']}; "
                                                                        'the stretches with most of them:', False, None)],
                        size=11, color=INK)
            rows = route['stretches']
            widths = [.4, 3.4, 1.5, 1.25]
            shape, table = _table(slide, len(rows) + 1, 4, 6.35, CONTENT_TOP + .6, 6.55, .34, widths)
            shape.name = 'Scoring Route Stretch Table'
            _header(table, ['#', 'Stretch', 'Points lost', 'Main Vendor'], size=9, numeric_from=2)
            for row_index, row in enumerate(rows, start=1):
                _cell(table.cell(row_index, 0), str(row_index), size=9, align=PP_ALIGN.CENTER)
                _cell(table.cell(row_index, 1), row['label'], size=9)
                _cell(table.cell(row_index, 2), f"{row['points']:,.2f}", size=9, bold=True, fill=None, align=PP_ALIGN.RIGHT)
                _cell(table.cell(row_index, 3), row.get('vendor') or '—', size=9)
            _row_bars(slide, [row['points'] for row in rows], left=6.35 + widths[0] + widths[1], top=CONTENT_TOP + .6 + .34,
                      width=widths[2], row_height=.34, color=LOSS_COLOR)
            _to_front(slide, shape)
            _text(slide, 'A stretch gathers the tests of the route within about 20 km, named after the nearest city of the measurements; '
                         'circle area: points lost.', .45, CONTENT_BOTTOM - .22, 12.4, .25, size=8, color=MUTED)


__all__ = [name for name in dir() if name.startswith('add_') or name in ('render_stretch_map', 'latest_campaign_items')]
