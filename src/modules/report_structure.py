"""The covers and transitions of the PowerPoint reports, alike in every module.

A report starts with its cover; its transitions open each part of it (a module of a Bundle Report Template, a scenario, a
scoring, an environment, a Design or a transition of its PPT Template). Every cover has one size of title and
subtitle and every transition another, smaller one, whatever the module; the subtitle starts right below the room of
two title lines, so a long title wraps into it. Each transition shows its counter above the title: 01, 02… for the first
level and 2.1, 2.1.1… below it.

A module marks each of its transitions with its level (``mark``), kept in the slide (the name of the box of its counter),
so a Bundle Report Template that copies the slides of a module can move them one or two levels down (``shift``) before
``finish`` writes every counter and size.
"""
from __future__ import annotations

from typing import Any

from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

COUNTER_PREFIX = 'Transition number · level '
SUBTITLE_NAME = 'Transition subtitle'
COVER_TITLE_PT, COVER_SUBTITLE_PT = 54, 24
TRANSITION_TITLE_PT, TRANSITION_SUBTITLE_PT, COUNTER_PT = 40, 20, 48
TITLE_LEFT, TITLE_TOP, RIGHT_MARGIN = Inches(0.52), Inches(2.32), Inches(0.52)
# The room of two title lines: the subtitle starts below it.
TWO_TITLE_LINES = {'cover': Inches(1.95), 'transition': Inches(1.45)}
COUNTER_HEIGHT, COUNTER_GAP = Inches(0.95), Inches(1.02)
SUBTITLE_HEIGHT = Inches(0.8)


def _counter(slide: Any) -> Any:
    return next((shape for shape in slide.shapes if str(shape.name).startswith(COUNTER_PREFIX)), None)


def _title(slide: Any) -> Any:
    return next((shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3} and shape.has_text_frame),
                None)


def mark(slide: Any, level: int) -> None:
    """Make the slide a transition of ``level`` (1 for the first level), its counter written by ``finish``."""
    counter = _counter(slide)
    if counter is None:
        counter = slide.shapes.add_textbox(TITLE_LEFT, TITLE_TOP - COUNTER_GAP, Inches(4), COUNTER_HEIGHT)
    counter.name = f'{COUNTER_PREFIX}{max(1, int(level))}'


def level(slide: Any) -> int | None:
    """The level of a transition, or None for any other slide."""
    counter = _counter(slide)
    if counter is None:
        return None
    try:
        return max(1, int(str(counter.name)[len(COUNTER_PREFIX):]))
    except ValueError:
        return 1


def shift(slide: Any, by: int) -> None:
    """A transition ``by`` levels down (a module's transitions inside a Bundle Report Template)."""
    current = level(slide)
    if current is not None:
        mark(slide, current + by)


def _place(slide: Any, kind: str, width: int) -> None:
    """The title at its place and sizes, and the subtitle right below the room of two title lines: the subtitle
    placeholder moved there, or the subtitle written under the title taken out into its own box."""
    title = _title(slide)
    if title is None:
        return
    title_pt = COVER_TITLE_PT if kind == 'cover' else TRANSITION_TITLE_PT
    subtitle_pt = COVER_SUBTITLE_PT if kind == 'cover' else TRANSITION_SUBTITLE_PT
    room = TWO_TITLE_LINES[kind]
    frame = title.text_frame
    paragraphs = list(frame.paragraphs)
    extra = '\n'.join(''.join(run.text for run in paragraph.runs) for paragraph in paragraphs[1:]).strip()
    for paragraph in paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(title_pt)
        for run in paragraph.runs:
            run.font.size = Pt(title_pt)
    # A placeholder placed by its layout gets its whole geometry written: one value alone would leave it no width.
    title.left, title.top, title.width, title.height = TITLE_LEFT, TITLE_TOP, width - TITLE_LEFT - RIGHT_MARGIN, room
    subtitle = next((shape for shape in slide.placeholders if shape.placeholder_format.type == 4 and shape.has_text_frame),
                    None) or next((shape for shape in slide.shapes if shape.name == SUBTITLE_NAME), None)
    if subtitle is None and extra:
        subtitle = slide.shapes.add_textbox(TITLE_LEFT, TITLE_TOP + room, title.width, SUBTITLE_HEIGHT)
        subtitle.name = SUBTITLE_NAME
        sub_frame = subtitle.text_frame
        sub_frame.margin_left = frame.margin_left
        sub_frame.margin_top = sub_frame.margin_bottom = 0
        run = sub_frame.paragraphs[0].add_run()
        run.text = extra
        run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_4
    if subtitle is None:
        return
    height = subtitle.height or SUBTITLE_HEIGHT
    subtitle.left, subtitle.top, subtitle.width, subtitle.height = TITLE_LEFT, TITLE_TOP + room, title.width, height
    subtitle.text_frame.word_wrap = True
    subtitle.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    for paragraph in subtitle.text_frame.paragraphs:
        paragraph.font.size = Pt(subtitle_pt)
        for run in paragraph.runs:
            run.font.size = Pt(subtitle_pt)


def _write_counter(slide: Any, label: str) -> None:
    counter = _counter(slide)
    counter.left, counter.top, counter.width, counter.height = TITLE_LEFT, TITLE_TOP - COUNTER_GAP, Inches(4), COUNTER_HEIGHT
    frame = counter.text_frame
    frame.clear()
    title = _title(slide)
    frame.margin_left = title.text_frame.margin_left if title is not None else 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    run = frame.paragraphs[0].add_run()
    run.text = label
    run.font.size, run.font.bold = Pt(COUNTER_PT), True
    run.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_1


def label(counters: list[int]) -> str:
    """01, 02… for the first level; 2.1, 2.1.1… below it."""
    return f'{counters[0]:02d}' if len(counters) == 1 else '.'.join(str(value) for value in counters)


class Counter:
    """The counters of the transitions of a report, in order: ``next(level)`` gives the label of the next one."""

    def __init__(self) -> None:
        self.counters: list[int] = []

    def next(self, current: int) -> str:
        self.counters = self.counters[:current]
        while len(self.counters) < current:
            self.counters.append(0)
        self.counters[current - 1] += 1
        # A level below one the report does not have (a transition of a Design without its cover) counts from 1.
        return label([value or 1 for value in self.counters])


def finish(presentation: Any) -> None:
    """Write the counter of every transition, in order, and give the cover and the transitions their sizes and places.
    The cover is the first slide when it is a cover layout and no transition. The slides of a dark theme turn their
    content dark (theme_dark)."""
    from src.modules.pptx_compose import is_cover

    width = presentation.slide_width
    counter = Counter()
    for index, slide in enumerate(presentation.slides):
        current = level(slide)
        if current is None:
            if index == 0 and is_cover(slide):
                _place(slide, 'cover', width)
            continue
        _write_counter(slide, counter.next(current))
        _place(slide, 'transition', width)
    # The content of the slides of a dark theme (Midnight) on their dark background.
    from src.modules.theme_dark import darken_presentation
    darken_presentation(presentation)
