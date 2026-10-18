"""Word versions of the PowerPoint reports: each slide becomes a landscape page with the same content.

Slide titles become headings, tables become Word tables with their colours, native charts are drawn as
images from their data, and pictures and text boxes are copied in reading order.

The Bundle Report Template and the exports of Scoring & GAP Analysis, Network Insights and Non-Qualified Calls pack their
components on portrait pages instead (see bundle_report_templates.compose_word), with the helpers here.
"""

from __future__ import annotations

import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from concurrent.futures import ThreadPoolExecutor
from functools import lru_cache
from threading import BoundedSemaphore
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn as pptx_qn

from src.utils.fonts import load_image_font

EMU_PER_INCH = 914400
# The logos of the slide master and of the dark layouts of the template.
TEMPLATE_LOGO_NAMES = frozenset({'Logo', 'Graphic 6', 'Logo (white)', 'Operators logo'})
CHART_PALETTE = ['#245a96', '#e08a1e', '#2e8b57', '#b0234f', '#6a63c9', '#0f6f7d', '#b85b20', '#7b8790']
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}


def _colour(fill: Any) -> str | None:
    """The solid fill colour of a cell or series as #RRGGBB, when it has one."""
    try:
        if fill.type is not None and fill.fore_color and fill.fore_color.type is not None:
            return f'#{fill.fore_color.rgb}'
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _foreground(hex_colour: str | None) -> str:
    if not hex_colour:
        return '#17232D'
    red, green, blue = (int(hex_colour[index:index + 2], 16) for index in (1, 3, 5))
    return '#FFFFFF' if red * 0.299 + green * 0.587 + blue * 0.114 < 150 else '#17232D'


def _shade(cell: Any, hex_colour: str) -> None:
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), hex_colour.lstrip('#'))
    cell._tc.get_or_add_tcPr().append(shading)


def _copy_table(document: Document, source: Any, scale: float) -> None:
    rows, columns = len(source.rows), len(source.columns)
    table = document.add_table(rows=rows, cols=columns)
    table.style = 'Table Grid'
    table.autofit = False
    for index, column in enumerate(source.columns):
        table.columns[index].width = int(column.width * scale)
    for row_index, row in enumerate(source.rows):
        for column_index in range(columns):
            origin = source.cell(row_index, column_index)
            target = table.cell(row_index, column_index)
            target.width = int(source.columns[column_index].width * scale)
            text = ' '.join(origin.text.splitlines())
            fill = _colour(origin.fill)
            if fill:
                _shade(target, fill)
            paragraph = target.paragraphs[0]
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            run.font.size = Pt(8)
            run.font.name = 'Arial'
            first = next((item for item in origin.text_frame.paragraphs[0].runs), None) if origin.text_frame.paragraphs else None
            run.bold = bool(first and first.font.bold) or row_index == 0
            try:
                colour = first.font.color.rgb if first is not None and first.font.color and first.font.color.type is not None else None
            except AttributeError:
                colour = None
            run.font.color.rgb = RGBColor.from_string(str(colour)) if colour else RGBColor.from_string(_foreground(fill).lstrip('#'))


def _chart_image(chart: Any, width_in: float, height_in: float) -> BytesIO | None:
    """A clustered bar chart drawn from the chart's own categories, series values and colours."""
    try:
        plot = chart.plots[0]
        categories = [str(category) for category in plot.categories]
        series = [(item.name, [float(value or 0) for value in item.values], _colour(item.format.fill)) for item in plot.series]
    except (IndexError, AttributeError, TypeError, ValueError):
        return None
    if not categories or not series:
        return None
    dpi = 150
    width, height = int(width_in * dpi), int(height_in * dpi)
    image = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(image)
    font, small = load_image_font(int(9 * dpi / 72)), load_image_font(int(7.5 * dpi / 72))
    left, right, top, bottom = int(0.6 * dpi), width - int(0.2 * dpi), int(0.45 * dpi), height - int(0.55 * dpi)
    values = [value for _name, numbers, _colour_value in series for value in numbers]
    low, high = min(0.0, *values), max(0.0, *values)
    if high == low:
        high = low + 1
    zero = bottom - (0 - low) / (high - low) * (bottom - top)
    for share in (0, 0.25, 0.5, 0.75, 1):
        y = bottom - share * (bottom - top)
        draw.line([left, y, right, y], fill='#e6eaef', width=1)
        draw.text((left - 6, y), f'{low + (high - low) * share:,.1f}', font=small, fill='#5b6377', anchor='rm')
    slot = (right - left) / len(categories)
    bar = max(2.0, min(slot * 0.8 / len(series), 0.35 * dpi))
    for index, category in enumerate(categories):
        start = left + index * slot + (slot - bar * len(series)) / 2
        for position, (_name, numbers, colour) in enumerate(series):
            value = numbers[index] if index < len(numbers) else 0
            y = bottom - (value - low) / (high - low) * (bottom - top)
            x = start + position * bar
            draw.rectangle([x, min(y, zero), x + bar - 1, max(y, zero)], fill=colour or CHART_PALETTE[position % len(CHART_PALETTE)])
        label = category if draw.textlength(category, font=small) <= slot else category[:max(1, int(len(category) * slot / max(1, draw.textlength(category, font=small))) - 1)] + '…'
        draw.text((left + index * slot + slot / 2, bottom + 6), label, font=small, fill='#3d4656', anchor='ma')
    x = left
    for position, (name, _numbers, colour) in enumerate(series):
        draw.rectangle([x, int(0.12 * dpi), x + 12, int(0.12 * dpi) + 12], fill=colour or CHART_PALETTE[position % len(CHART_PALETTE)])
        draw.text((x + 18, int(0.12 * dpi) - 2), str(name), font=font, fill='#17232D')
        x += 30 + draw.textlength(str(name), font=font)
    output = BytesIO()
    image.save(output, 'PNG')
    output.seek(0)
    return output


def _slide_title(slide: Any) -> tuple[str, Any]:
    for shape in slide.placeholders:
        if shape.placeholder_format.type in TITLE_TYPES and shape.has_text_frame and shape.text_frame.text.strip():
            return ' '.join(shape.text_frame.text.split()), shape
    return '', None


def _soffice() -> str | None:
    """LibreOffice, from the environment, the PATH or its usual macOS location."""
    candidates = [os.environ.get('DRIVETEST_ANALYZER_SOFFICE', ''), shutil.which('soffice') or '', shutil.which('libreoffice') or '',
                  '/Applications/LibreOffice.app/Contents/MacOS/soffice']
    return next((path for path in candidates if path and Path(path).is_file() and os.access(path, os.X_OK)), None)


def slide_images(content: bytes, dpi: int = 160) -> list[bytes] | None:
    """Every slide as a PNG, rendered by LibreOffice through PDF, or None when LibreOffice is unavailable."""
    return document_page_images(content, 'pptx', dpi)


def document_page_images(content: bytes, suffix: str, dpi: int = 160) -> list[bytes] | None:
    """Every page of a PowerPoint (``pptx``) or Word (``docx``) document as a PNG, rendered by LibreOffice through
    PDF, or None when LibreOffice is unavailable."""
    pdf = document_pdf(content, suffix)
    return None if pdf is None else pdf_page_images(pdf, dpi)


def without_unused_layouts(content: bytes) -> bytes:
    """A PowerPoint without the slide layouts none of its slides uses. LibreOffice converts every layout of a deck,
    and the template has well over a hundred, so a deck converts in seconds instead of about a minute."""
    try:
        presentation = Presentation(BytesIO(content))
        used = {id(slide.slide_layout.part) for slide in presentation.slides}
        for master in presentation.slide_masters:
            layouts = master._element.find(pptx_qn('p:sldLayoutIdLst'))
            for entry in list(layouts if layouts is not None else []):
                rId = entry.get(pptx_qn('r:id'))
                if id(master.part.related_part(rId)) not in used:
                    master.part.drop_rel(rId)
                    layouts.remove(entry)
        output = BytesIO()
        presentation.save(output)
        return output.getvalue()
    except Exception:  # noqa: BLE001 - a deck that cannot be read this way is converted as it is.
        return content


# On Linux, the fonts that stand in for those of a document that are not installed, in this order.
FALLBACK_FONTS = ('Liberation Sans', 'DejaVu Sans', 'Noto Sans')
_FONT_ATTRIBUTE = re.compile(rb'((?:typeface|w:ascii|w:hAnsi|w:cs|w:eastAsia)=")([^"+][^"]*)(")')


@lru_cache(maxsize=1)
def installed_fonts() -> frozenset[str] | None:
    """The font families fontconfig knows (lower case), on Linux; None elsewhere, where LibreOffice finds a font that
    stands in for a missing one by itself."""
    if not sys.platform.startswith('linux') or not shutil.which('fc-list'):
        return None
    try:
        listing = subprocess.run(['fc-list', ':', 'family'], check=True, capture_output=True, text=True, timeout=30).stdout
    except (OSError, subprocess.SubprocessError):
        return None
    return frozenset(name.strip().casefold() for line in listing.splitlines() for name in line.split(',') if name.strip())


def with_installed_fonts(content: bytes) -> bytes:
    """The document (PowerPoint or Word) with every font that is not installed replaced by one that is. On Linux
    LibreOffice lays out the text of a missing font with metrics that do not match the font it draws, so the letters
    of the preview come out spaced and overlapping; the documents themselves keep their fonts."""
    installed = installed_fonts()
    fallback = next((font for font in FALLBACK_FONTS if installed and font.casefold() in installed), None)
    if fallback is None:
        return content

    def replace(match: re.Match[bytes]) -> bytes:
        name = match.group(2).decode('utf-8', 'ignore').strip()
        return match.group(0) if name.casefold() in installed else match.group(1) + fallback.encode() + match.group(3)

    output = BytesIO()
    try:
        with zipfile.ZipFile(BytesIO(content)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename.endswith('.xml'):
                    data = _FONT_ATTRIBUTE.sub(replace, data)
                target.writestr(info, data)
    except zipfile.BadZipFile:
        return content
    return output.getvalue()


def libreoffice_workers() -> int:
    """How many LibreOffice conversions run at once: the processors of the machine, at most 4, or
    DRIVETEST_ANALYZER_LIBREOFFICE_WORKERS (each one takes a few hundred MB of memory)."""
    configured = os.environ.get('DRIVETEST_ANALYZER_LIBREOFFICE_WORKERS', '').strip()
    if configured.isdigit() and int(configured) > 0:
        return min(8, int(configured))
    try:
        processors = len(os.sched_getaffinity(0))
    except AttributeError:
        processors = os.cpu_count() or 1
    return max(1, min(4, processors))


# Every LibreOffice conversion of the application waits for one of these (several reports or previews at once share
# the processors instead of overloading the machine).
_CONVERSIONS = BoundedSemaphore(libreoffice_workers())
# A deck is shared among several conversions at once when each gets at least this many slides.
SLIDES_PER_CONVERSION = 6
# The time a conversion may take: a fixed part (starting LibreOffice) and a part per slide (or page).
CONVERSION_TIMEOUT_SECONDS = 180
CONVERSION_SECONDS_PER_SLIDE = 12


def _slide_count(content: bytes) -> int:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            return sum(1 for name in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', name))
    except zipfile.BadZipFile:
        return 0


# What a slide costs LibreOffice besides its content, in bytes of slide XML (about a tenth of a second).
_SLIDE_BASE_WEIGHT = 30000


def _deck_parts(content: bytes, parts: int) -> list[bytes]:
    """The deck split into ``parts`` decks of consecutive slides that take LibreOffice about as long (weighed by the
    size of their XML and of their charts: a chart of many bars or a large table cost much more than a few shapes),
    each with only its layouts."""
    presentation = Presentation(BytesIO(content))
    weights = [_SLIDE_BASE_WEIGHT + len(slide.part.blob)
               + sum(len(rel.target_part.blob) for rel in slide.part.rels.values()
                     if not rel.is_external and rel.reltype.endswith('/chart'))
               for slide in presentation.slides]
    total = sum(weights)
    cumulative = [sum(weights[:end]) for end in range(len(weights) + 1)]
    bounds = [0]
    for part in range(1, parts):
        # Each part ends where the slides before it come closest to its share; every part keeps at least one slide.
        candidates = range(bounds[-1] + 1, len(weights) - (parts - part) + 1)
        bounds.append(min(candidates, key=lambda end: abs(cumulative[end] - total * part / parts)))
    bounds.append(len(weights))
    decks = []
    for first, last in zip(bounds, bounds[1:]):
        part = Presentation(BytesIO(content))
        slide_ids = part.slides._sldIdLst
        for position, slide_id in enumerate(list(slide_ids)):
            if not first <= position < last:
                part.part.drop_rel(slide_id.rId)
                slide_ids.remove(slide_id)
        output = BytesIO()
        part.save(output)
        decks.append(without_unused_layouts(output.getvalue()))
    return decks


def _convert_to_pdf(soffice: str, content: bytes, suffix: str, slides: int) -> bytes | None:
    """One LibreOffice conversion to PDF, or None when it fails or takes too long."""
    with _CONVERSIONS, tempfile.TemporaryDirectory() as folder:
        source = Path(folder) / f'report.{suffix}'
        source.write_bytes(content)
        # A private profile lets the conversion run while the user (or another conversion) has LibreOffice open.
        profile = (Path(folder) / 'profile').as_uri()
        try:
            subprocess.run([soffice, f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'pdf',
                            '--outdir', folder, str(source)], check=True, capture_output=True,
                           timeout=CONVERSION_TIMEOUT_SECONDS + CONVERSION_SECONDS_PER_SLIDE * max(1, slides))
        except (OSError, subprocess.SubprocessError):
            return None
        pdf = Path(folder) / 'report.pdf'
        return pdf.read_bytes() if pdf.is_file() else None


def document_pdf(content: bytes, suffix: str) -> bytes | None:
    """A PowerPoint (``pptx``) or Word (``docx``) document as PDF, converted by LibreOffice, or None when LibreOffice
    is unavailable. A long deck is shared among several conversions that run at once (see libreoffice_workers), whose
    PDFs are joined in order."""
    import importlib.util

    soffice = _soffice()
    # The pages are rendered with PyMuPDF.
    if soffice is None or importlib.util.find_spec('pymupdf') is None:
        return None
    if suffix == 'pptx':
        content = without_unused_layouts(content)
    content = with_installed_fonts(content)
    slides = _slide_count(content) if suffix == 'pptx' else 0
    parts = min(libreoffice_workers(), slides // SLIDES_PER_CONVERSION) if suffix == 'pptx' else 1
    if parts <= 1:
        return _convert_to_pdf(soffice, content, suffix, slides or 20)
    decks = _deck_parts(content, parts)
    with ThreadPoolExecutor(max_workers=len(decks)) as pool:
        pdfs = list(pool.map(lambda deck: _convert_to_pdf(soffice, deck, suffix, _slide_count(deck)), decks))
    if any(pdf is None for pdf in pdfs):
        return None
    import pymupdf

    with pymupdf.open() as joined:
        for pdf in pdfs:
            with pymupdf.open(stream=pdf, filetype='pdf') as part:
                joined.insert_pdf(part)
        return joined.tobytes()


def pdf_page_images(pdf: bytes, dpi: int, pages: list[int] | None = None) -> list[bytes]:
    """The pages of a PDF (every one, or those at ``pages``) as PNGs at ``dpi``."""
    import pymupdf

    with pymupdf.open(stream=pdf, filetype='pdf') as document:
        chosen = range(document.page_count) if pages is None else [page for page in pages if 0 <= page < document.page_count]
        return [document[page].get_pixmap(dpi=dpi).tobytes('png') for page in chosen]


# Empty bands taller than this share of the slide height shrink to the band kept between two blocks.
_EMPTY_BAND = 0.025
_KEPT_BAND = 0.012


def compact_slide_image(image: bytes) -> bytes:
    """The slide image without its empty bands: the blank rows at the top and bottom go and every taller blank band
    between two blocks (such as the space between a table and its footnote) shrinks to a narrow gap."""
    with Image.open(BytesIO(image)) as raster:
        rgb = raster.convert('RGB')
    pixels = np.asarray(rgb).astype(np.int16)
    background = pixels[1, 1]
    filled = (np.abs(pixels - background).max(axis=2) > 16).any(axis=1)
    if not filled.any():
        return image
    height = len(filled)
    empty_band, kept_band = max(4, int(height * _EMPTY_BAND)), max(2, int(height * _KEPT_BAND))
    rows: list[int] = []
    first, last = int(np.argmax(filled)), height - int(np.argmax(filled[::-1]))
    start = max(0, first - kept_band)
    row = first
    rows.extend(range(start, first))
    while row < last:
        if filled[row]:
            rows.append(row)
            row += 1
            continue
        end = row
        while end < last and not filled[end]:
            end += 1
        rows.extend(range(row, end) if end - row <= empty_band else range(row, row + kept_band))
        row = end
    rows.extend(range(last, min(height, last + kept_band)))
    output = BytesIO()
    Image.fromarray(np.asarray(rgb)[rows]).save(output, format='PNG')
    return output.getvalue()


_PAGE_WIDTH, _PAGE_HEIGHT = 8.27, 11.69


# The Ericsson mark drawn behind the covers and transitions of the template: three bars, smaller and at the bottom
# right of a portrait page, below the text.
WATERMARK_PREFIX = 'Ericsson watermark'
PORTRAIT_WATERMARK_HEIGHT = Inches(3.6)


def _place_portrait_watermark(layout, width: int, height: int) -> None:
    bars = sorted((shape for shape in layout.shapes if str(shape.name).startswith(WATERMARK_PREFIX)),
                  key=lambda shape: shape.top)
    if len(bars) != 3 or None in (bars[0].width, bars[0].height):
        return
    angle = math.radians((360 - bars[0].rotation) % 360)
    length, thickness = bars[0].width, bars[0].height
    spacing = (bars[2].top - bars[0].top) / 2
    mark_height = 2 * spacing + length * math.sin(angle) + thickness * math.cos(angle)
    scale = PORTRAIT_WATERMARK_HEIGHT / mark_height if mark_height else 1
    length, thickness, spacing = length * scale, thickness * scale, spacing * scale
    mark_width = length * math.cos(angle) + thickness * math.sin(angle)
    centre_x = width - Inches(0.55) - mark_width / 2
    middle = height - Inches(0.6) - PORTRAIT_WATERMARK_HEIGHT / 2
    for index, bar in enumerate(bars):
        centre_y = middle + (index - 1) * spacing
        bar.left, bar.top = int(centre_x - length / 2), int(centre_y - thickness / 2)
        bar.width, bar.height = int(length), int(thickness)


def portrait_covers(content: bytes, positions: list[int]) -> list[bytes] | None:
    """The slides at ``positions`` (covers) as portrait A4 pages, in that order, rendered at once; None when
    LibreOffice cannot render them."""
    presentation = Presentation(BytesIO(content))
    slide_ids = presentation.slides._sldIdLst
    every = list(slide_ids)
    kept = [every[position] for position in dict.fromkeys(positions) if 0 <= position < len(every)]
    if not kept:
        return []
    for slide_id in every:
        if slide_id not in kept:
            presentation.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
    for slide_id in kept:  # In the order asked for.
        slide_ids.remove(slide_id)
        slide_ids.append(slide_id)
    old_width, old_height = presentation.slide_width, presentation.slide_height
    width, height = Inches(_PAGE_WIDTH), Inches(_PAGE_HEIGHT)
    presentation.slide_width, presentation.slide_height = width, height
    lower = int((height - old_height) * 0.3)
    # Placeholders that take their place from the layout get it written out first, on every cover before any layout
    # moves: moving only their top would give them no width or height, and a moved layout would move them twice.
    for slide in presentation.slides:
        for shape in slide.placeholders:
            geometry = (shape.left, shape.top, shape.width, shape.height)
            if None not in geometry:
                shape.left, shape.top, shape.width, shape.height = geometry
    done: set[int] = set()
    for slide in presentation.slides:
        layout = slide.slide_layout
        for container, move_down in ((slide, True), (layout, True), (layout.slide_master, False)):
            # Each layout and master moves once, whatever the number of covers that use it.
            if container is not slide and id(container.part) in done:
                continue
            done.add(id(container.part))
            if container is layout:
                _place_portrait_watermark(layout, width, height)
            for shape in container.shapes:
                if shape.left is None or shape.width is None or str(shape.name).startswith(WATERMARK_PREFIX):
                    continue
                if shape.left > old_width / 2:
                    # A logo at the right edge stays at the right edge.
                    shape.left = max(0, shape.left - (old_width - width))
                elif shape.left + shape.width > width:
                    shape.width = max(Inches(1), width - shape.left - Inches(0.4))
                if move_down and shape.top is not None and shape.top > Inches(1) and shape.left <= old_width / 2:
                    shape.top = shape.top + lower
    output = BytesIO()
    presentation.save(output)
    return slide_images(output.getvalue(), dpi=150)


def crop_to_content(image: bytes, box: tuple[float, float, float, float] | None = None) -> tuple[bytes, int, int]:
    """The part of a slide image inside ``box`` (left, top, right, bottom as shares of the image) without the blank
    margins around what it shows; returns the PNG and its width and height in pixels."""
    with Image.open(BytesIO(image)) as raster:
        rgb = raster.convert('RGB')
    if box is not None:
        width, height = rgb.size
        left, top, right, bottom = (max(0, int(box[0] * width)), max(0, int(box[1] * height)),
                                    min(width, int(math.ceil(box[2] * width))), min(height, int(math.ceil(box[3] * height))))
        if right - left > 4 and bottom - top > 4:
            rgb = rgb.crop((left, top, right, bottom))
    pixels = np.asarray(rgb).astype(np.int16)
    background = pixels[0, 0]
    filled = np.abs(pixels - background).max(axis=2) > 16
    rows, columns = np.where(filled.any(axis=1))[0], np.where(filled.any(axis=0))[0]
    if rows.size and columns.size:
        pad = max(2, int(min(rgb.size) * 0.01))
        rgb = rgb.crop((max(0, int(columns[0]) - pad), max(0, int(rows[0]) - pad),
                        min(rgb.width, int(columns[-1]) + pad + 1), min(rgb.height, int(rows[-1]) + pad + 1)))
    output = BytesIO()
    rgb.save(output, format='PNG')
    return output.getvalue(), rgb.width, rgb.height


@dataclass
class PackItem:
    """What a packed page holds: a ``cover`` (an image that fills a page of its own, without margins), a heading
    (``heading1`` opens a section and starts a new page, ``heading2`` opens a part of it, both with optional
    ``subtitle`` lines) or an ``image`` of ``width`` × ``height`` inches."""
    kind: str
    text: str = ''
    subtitle: list[str] = field(default_factory=list)
    image: bytes = b''
    width: float = 0.0
    height: float = 0.0
    # An image may come without its title (``untitled``, of ``untitled_width`` × ``untitled_height`` inches), shown
    # when it follows on the same page an image of the same ``group`` (the same title and subtitle).
    group: str = ''
    untitled: bytes = b''
    untitled_width: float = 0.0
    untitled_height: float = 0.0

    def without_title(self) -> 'PackItem':
        return PackItem('image', self.text, image=self.untitled, width=self.untitled_width, height=self.untitled_height,
                        group=self.group)

    @property
    def fixed_height(self) -> float:
        heading = {'heading1': 0.46, 'heading2': 0.38}.get(self.kind, 0.0)
        return heading + 0.21 * len(self.subtitle) if self.kind != 'image' else 0.0


_PACK_GAP = 0.15
_ROW_EXTRA = 0.06


def _shelves(images: list[PackItem], width: float, reorder: bool) -> list[list[PackItem]]:
    """Rows of images side by side: in their order (a new row when the next one does not fit beside the others), or
    the tallest first, each one in the first row with room for it."""
    rows: list[list[PackItem]] = []
    if not reorder:
        for item in images:
            if rows and sum(other.width for other in rows[-1]) + _PACK_GAP * len(rows[-1]) + item.width <= width + 1e-6:
                rows[-1].append(item)
            else:
                rows.append([item])
        return rows
    for item in sorted(images, key=lambda image: -image.height):
        row = next((row for row in rows if sum(other.width for other in row) + _PACK_GAP * len(row) + item.width <= width + 1e-6), None)
        if row is None:
            rows.append([item])
        else:
            row.append(item)
    return rows


def page_layout(items: list[PackItem], width: float, height: float) -> tuple[list[list[PackItem]], float]:
    """The rows of a page and the scale its images need to fit: in their order, or with the images between two
    headings reordered when that fits them larger (the page holds the same items either way)."""
    best: tuple[list[list[PackItem]], float] | None = None
    for reorder in (False, True):
        rows: list[list[PackItem]] = []
        run: list[PackItem] = []
        for item in items:
            if item.kind == 'image':
                run.append(item)
                continue
            rows.extend(_shelves(run, width, reorder))
            run = []
            rows.append([item])
        rows.extend(_shelves(run, width, reorder))
        fixed = sum(row[0].fixed_height for row in rows if row[0].kind != 'image')
        images = sum(max(item.height for item in row) + _ROW_EXTRA for row in rows if row[0].kind == 'image')
        gaps = _PACK_GAP * max(0, len(rows) - 1)
        scale = 1.0 if not images else min(1.0, max(0.05, (height - fixed - gaps) / images))
        if best is None or scale > best[1] + 1e-6:
            best = (rows, scale)
    return best or ([], 1.0)


def pack_pages(items: list[PackItem], width: float, height: float, minimum_scale: float = 0.88) -> list[tuple[list[list[PackItem]], float]]:
    """Pages of consecutive items: each page takes the next items while they fit (a little smaller, down to
    ``minimum_scale``), so an item never goes after one that follows it; a section heading starts a new page, a
    heading never ends a page, and an item larger than a page fills one page on its own, made smaller. An image that
    follows on the same page one of its group goes without its title, which that one already shows."""
    pages: list[list[PackItem]] = []
    current: list[PackItem] = []

    def close() -> list[PackItem]:
        # Headings at the end of a page go to the next one, with the item they introduce.
        trailing: list[PackItem] = []
        while current and current[-1].kind != 'image':
            trailing.insert(0, current.pop())
        if current:
            pages.append(list(current))
        return trailing

    for item in items:
        if item.kind == 'cover':
            # A cover fills a page of its own; what follows it starts the next page.
            trailing = close()
            if trailing:
                pages.append(trailing)
            pages.append([item])
            current = []
            continue
        if item.kind == 'heading1' and current:
            current = close()
            if current:
                pages.append(current)
                current = []
        shown = item
        if item.kind == 'image' and item.untitled and item.group and current and current[-1].kind == 'image' \
                and current[-1].group == item.group:
            shown = item.without_title()
        candidate = [*current, shown]
        if not any(entry.kind == 'image' for entry in current) or page_layout(candidate, width, height)[1] >= minimum_scale:
            current = candidate
            continue
        current = [*close(), item]
    trailing = close()
    if trailing:
        pages.append(trailing)
    return [page_layout(page, width, height) for page in pages]


def _cover_page(document: Any, cover: bytes) -> Any:
    """The first page of ``document`` as ``cover`` without margins, so its background reaches every edge; returns the
    section of the pages that follow it."""
    section = document.sections[0]
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = 0
    section.header_distance = section.footer_distance = 0
    paragraph = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.add_run().add_picture(BytesIO(cover), width=Inches(_PAGE_WIDTH - 0.01), height=Inches(_PAGE_HEIGHT - 0.05))
    following = document.add_section(WD_SECTION.NEW_PAGE)
    # The break of the cover section goes into the paragraph of the cover, so no empty line follows it.
    break_paragraph = document.paragraphs[-1]._p
    paragraph._p.get_or_add_pPr().append(break_paragraph.pPr.sectPr)
    break_paragraph.getparent().remove(break_paragraph)
    return following


def _margins(section: Any, inches: float) -> None:
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(inches)
    if not inches:
        section.header_distance = section.footer_distance = 0


def _next_section(document: Any, inches: float) -> Any:
    """A new section from a new page with ``inches`` margins. python-docx ends the previous section with an empty
    paragraph that holds its properties: they go into the last paragraph of that section instead, so it adds no line
    (after a full page, that line would take a blank page)."""
    section = document.add_section(WD_SECTION.NEW_PAGE)
    breaker = document.paragraphs[-1]._p
    previous = breaker.getprevious()
    if previous is not None and previous.tag == qn('w:p'):
        previous.get_or_add_pPr().append(breaker.pPr.sectPr)
        breaker.getparent().remove(breaker)
    _margins(section, inches)
    return section


_CONTENT_MARGIN = 0.45


def packed_component_document(items: list[PackItem], cover: bytes | None = None,
                              page_items: list[list[PackItem]] | None = None) -> bytes:
    """Portrait pages of covers, headings and component images packed by ``pack_pages``, after a full-page ``cover``.
    Every cover fills a page without margins (the report cover and the cover of each section). ``page_items`` receives
    the items of each page, in order (to name the pages of the document)."""
    if cover is not None:
        items = [PackItem('cover', image=cover), *items]
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(_PAGE_WIDTH), Inches(_PAGE_HEIGHT)
    width = _PAGE_WIDTH - 2 * _CONTENT_MARGIN
    height = _PAGE_HEIGHT - 2 * _CONTENT_MARGIN - 0.3
    previous_cover = None
    pages = pack_pages(items, width, height)
    if page_items is not None:
        page_items.extend([item for row in rows for item in row] for rows, _scale in pages)
    for page_number, (rows, scale) in enumerate(pages):
        is_cover = rows[0][0].kind == 'cover'
        if page_number == 0:
            _margins(section, 0 if is_cover else _CONTENT_MARGIN)
        elif is_cover or previous_cover:
            # A cover has a section of its own without margins; the pages after it, one with margins.
            section = _next_section(document, 0 if is_cover else _CONTENT_MARGIN)
        new_section = page_number == 0 or is_cover or previous_cover
        previous_cover = is_cover
        if is_cover:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.add_run().add_picture(BytesIO(rows[0][0].image), width=Inches(_PAGE_WIDTH - 0.01),
                                            height=Inches(_PAGE_HEIGHT - 0.05))
            continue
        for row_number, row in enumerate(rows):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.page_break_before = row_number == 0 and not new_section
            paragraph.paragraph_format.space_before = Pt(0)
            item = row[0]
            if item.kind != 'image':
                paragraph.style = document.styles['Heading 1' if item.kind == 'heading1' else 'Heading 2']
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.keep_with_next = True
                paragraph.add_run(item.text)
                for line in item.subtitle:
                    note = document.add_paragraph()
                    note.paragraph_format.space_before = note.paragraph_format.space_after = Pt(0)
                    note.paragraph_format.keep_with_next = True
                    run = note.add_run(line)
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x5B, 0x63, 0x77)
                continue
            paragraph.alignment = 1
            paragraph.paragraph_format.space_after = Pt(_PACK_GAP * 72)
            for index, image in enumerate(row):
                if index:
                    paragraph.add_run('\u2002')
                paragraph.add_run().add_picture(BytesIO(image.image), width=Inches(image.width * scale),
                                                height=Inches(image.height * scale))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pptx_to_docx(content: bytes, title: str = '') -> bytes:
    """Convert a generated PowerPoint into a landscape Word document with one page per slide.

    With LibreOffice, every page shows the slide exactly as in PowerPoint (layout, colours, tables and charts);
    without it, the slide content is rebuilt: headings, tables with their colours and charts drawn from their data.
    """
    images = slide_images(content, dpi=160)
    if images:
        document = Document()
        section = document.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Inches(11.69), Inches(8.27)
        section.left_margin = section.right_margin = Inches(0.5)
        section.top_margin = section.bottom_margin = Inches(0.5)
        usable_width = (section.page_width - section.left_margin - section.right_margin) / EMU_PER_INCH
        usable_height = (section.page_height - section.top_margin - section.bottom_margin) / EMU_PER_INCH - 0.3
        for index, image in enumerate(images):
            with Image.open(BytesIO(image)) as raster:
                ratio = raster.width / raster.height
            width = min(usable_width, usable_height * ratio)
            paragraph = document.paragraphs[0] if index == 0 and document.paragraphs else document.add_paragraph()
            paragraph.alignment = 1
            paragraph.paragraph_format.page_break_before = index > 0
            paragraph.add_run().add_picture(BytesIO(image), width=Inches(width))
        output = BytesIO()
        document.save(output)
        return output.getvalue()
    presentation = Presentation(BytesIO(content))
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.left_margin = section.right_margin = Inches(0.6)
    section.top_margin = section.bottom_margin = Inches(0.6)
    usable = (section.page_width - section.left_margin - section.right_margin) / EMU_PER_INCH
    scale = usable / (presentation.slide_width / EMU_PER_INCH)
    if title:
        document.add_heading(title, level=0)
    first_content = True
    for slide in presentation.slides:
        heading_text, title_shape = _slide_title(slide)
        shapes = sorted((shape for shape in slide.shapes if shape is not title_shape),
                        key=lambda shape: (round((shape.top or 0) / (EMU_PER_INCH * 0.25)), shape.left or 0))
        # The logos of the template a slide draws itself (see template_logos) are not content.
        body = [shape for shape in shapes if shape.has_table or getattr(shape, 'has_chart', False)
                or (shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.name not in TEMPLATE_LOGO_NAMES)
                or (shape.has_text_frame and shape.text_frame.text.strip())]
        if not heading_text and not body:
            continue
        if heading_text:
            heading = document.add_heading(heading_text, level=1)
            heading.paragraph_format.page_break_before = not first_content
        first_content = False
        for shape in body:
            if shape.has_table:
                _copy_table(document, shape.table, scale)
                document.add_paragraph()
            elif getattr(shape, 'has_chart', False):
                image = _chart_image(shape.chart, min(usable, shape.width / EMU_PER_INCH * scale), shape.height / EMU_PER_INCH * scale)
                if image is not None:
                    document.add_picture(image, width=Inches(min(usable, shape.width / EMU_PER_INCH * scale)))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                document.add_picture(BytesIO(shape.image.blob), width=Inches(min(usable, shape.width / EMU_PER_INCH * scale)))
            else:
                for paragraph in shape.text_frame.paragraphs:
                    text = ''.join(run.text for run in paragraph.runs).strip()
                    if text:
                        # Keep the colour of coloured notes, such as the orange scaling notice.
                        run = document.add_paragraph().add_run(text)
                        font = paragraph.runs[0].font if paragraph.runs else paragraph.font
                        # Only an explicit colour is kept (theme colours have no RGB value of their own).
                        colour = font.color.rgb if font.color and font.color.type == MSO_COLOR_TYPE.RGB else (
                            paragraph.font.color.rgb if paragraph.font.color and paragraph.font.color.type == MSO_COLOR_TYPE.RGB else None)
                        if colour is not None and str(colour).upper() not in {'FFFFFF'}:
                            run.font.color.rgb = RGBColor.from_string(str(colour))
                        run.font.bold = bool(paragraph.font.bold or (paragraph.runs and paragraph.runs[0].font.bold))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
