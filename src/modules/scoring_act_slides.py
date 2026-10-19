"""The slides of where to act in the Summary of the Scoring & GAP Analysis exports (see scoring_where_to_act): the plan
of an operator, its performance owners, its Cities, Towns, Roads and Railways ranked and the route that loses most, as
benchmark reports present them. Every colour comes from the active theme; each performance owner keeps one colour of
its categories on every slide.
"""
from __future__ import annotations

from io import BytesIO
from typing import Any, Callable

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.modules.scoring_insight_slides import ScopeLabels, _box as _styled_box, _cell, _rgb, _table, _write_runs
from src.modules.scoring_summary_slides import (
    CONTENT_BOTTOM, CONTENT_TOP, INK, LOSS_COLOR, MUTED, _header, _row_bars, _subtitle, _text, _to_front,
)
from src.modules.scoring_themes import active_theme, themed

LEFT, WIDTH = .45, 12.45
PLACE_ROWS_PER_SLIDE = 18


def _box(slide, left: float, top: float, width: float, height: float, color: str, **options):
    """A shape without the style of the theme, which draws a shadow under it in LibreOffice (the Word version)."""
    shape = _styled_box(slide, left, top, width, height, color, **options)
    style = shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        shape._element.remove(style)
    return shape


def owner_palette(names: list[str]) -> Callable[[str], str]:
    """One colour of the categories of the theme per performance owner (by name, the same on every slide and on the page:
    in the order of scoring_where_to_act.owner_order)."""
    from src.modules.scoring_where_to_act import owner_order
    owners = owner_order(names)
    categories = active_theme().get('categories') or ['#4472C4', '#7030A0', '#C55A11', '#5B9BD5', '#A64D79', '#548235']
    return lambda name: categories[owners.index(name) % len(categories)] if name in owners else '#9AA3AD'


def owner_names(views: dict[str, list[dict[str, Any]]]) -> list[str]:
    names = [card['owner'] for item in views.get('owner_cards') or [] for card in item['cards']]
    names += [owner['name'] for item in views.get('place_rankings') or [] for group in item['groups']
              for row in group['rows'] for owner in row['owners']]
    names += [owner['name'] for item in views.get('route_focus') or [] for route in item['routes'] for owner in route['owners']
              if not owner.get('unidentified')]
    return names


def _headline(slide, text: str, top: float = CONTENT_TOP) -> float:
    """The headline under the title in the accent of the theme; returns where the content starts."""
    _text(slide, text, LEFT, top, WIDTH, .34, size=14, color=active_theme()['title_accent'], name='Scoring Act Headline')
    return top + .46


def _footnote(slide, text: str) -> None:
    _text(slide, text, LEFT, CONTENT_BOTTOM - .22, WIDTH, .25, size=8, color=MUTED)


def _pill(slide, text: str, left: float, top: float, color: str, height: float = .22, size: float = 8) -> float:
    """An owner in its colour; returns its width."""
    width = .2 + len(text) * .065
    pill = _box(slide, left, top, width, height, color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    pill.adjustments[0] = .2
    pill.name = 'Scoring Owner Pill'
    frame = pill.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _write_runs(frame, [(text, True, '#FFFFFF')], size=size, color='#FFFFFF')
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    return width


def add_act_plan_slides(presentation, items: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                        scope: ScopeLabels | None = None) -> None:
    """Where to act: the themes that hold the points lost, what to do and the view to open."""
    stripes = active_theme().get('categories') or ['#5B7A99']
    loss = active_theme().get('loss_text') or LOSS_COLOR
    for item in items:
        slide = new_slide(presentation, f"Where to act — {item['label']}", _subtitle(subtitle, scope, item.get('context'),
                                                                                   item.get('campaign') or ''))
        top = _headline(slide, item['headline'])
        headers = (('Theme', LEFT + .25), ('What to do', LEFT + 3.35), ('Deep-dive', LEFT + 10.0))
        for text, left in headers:
            _text(slide, text, left, top, 2.5, .25, size=9, color=MUTED)
        top += .3
        height = min(.78, (CONTENT_BOTTOM - .4 - top) / max(1, len(item['rows'])) - .1)
        for index, row in enumerate(item['rows']):
            y = top + index * (height + .1)
            card = _box(slide, LEFT, y, WIDTH, height, '#FFFFFF')
            card.line.color.rgb = _rgb('#E3E8EC')
            card.name = 'Scoring Act Row'
            _box(slide, LEFT, y, .07, height, stripes[index % len(stripes)]).name = 'Scoring Act Stripe'
            # The points of the scoring stand out in the loss colour of the theme, as on the page.
            points = [('Within the points above', False)] if row.get('points') is None else [
                (f"{row['points']:,.1f} points", True), (' lost', False)]
            box = slide.shapes.add_textbox(Inches(LEFT + .25), Inches(y + .08), Inches(2.9), Inches(height - .12))
            _write_runs(box.text_frame, [(row['theme'], True, INK)], size=12, color=INK)
            paragraph = box.text_frame.add_paragraph()
            for text, highlighted in points:
                run = paragraph.add_run()
                run.text = text
                run.font.size = Pt(9)
                run.font.bold = highlighted
                run.font.color.rgb = _rgb(loss if highlighted else MUTED)
            parts = row.get('action_parts') or [{'text': row['action'], 'points': False}]
            action = slide.shapes.add_textbox(Inches(LEFT + 3.35), Inches(y + .06), Inches(6.4), Inches(height - .1))
            _write_runs(action.text_frame, [(part['text'], bool(part['points']), loss if part['points'] else None) for part in parts],
                        size=10, color=INK)
            action.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            _text(slide, row['deep_dive'], LEFT + 10.0, y + .06, 2.35, height - .1, size=9, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
        _footnote(slide, 'Points lost: the points of the scoring lost against the maximum of each KPI (not the GAP to the reference), '
                         'each place counting by its share of the tests; the sites to review and the route stretches are within them.')


def add_owner_card_slides(presentation, items: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                          owner_color: Callable[[str], str], scope: ScopeLabels | None = None) -> None:
    """Where to act, by performance owner: a card per owner with its points lost, what drives them and where."""
    for item in items:
        slide = new_slide(presentation, f"Where to act, by performance owner — {item['label']}",
                          _subtitle(subtitle, scope, item.get('context'), item.get('campaign') or ''))
        top = _headline(slide, "Each performance owner's points lost, what drives them and where they are")
        cards = item['cards'][:6]
        gap = .25
        width = (WIDTH - gap * (len(cards) - 1)) / max(1, len(cards))
        height = CONTENT_BOTTOM - .45 - top
        peak = max([value for card in cards for value in (card['settings'] or {}).values()] + [.01])
        for index, card in enumerate(cards):
            left = LEFT + index * (width + gap)
            color = owner_color(card['owner'])
            frame = _box(slide, left, top, width, height, '#FFFFFF')
            frame.line.color.rgb = _rgb('#D9DEE3')
            frame.name = 'Scoring Owner Card'
            _box(slide, left, top, width, 1.05, color).name = 'Scoring Owner Header'
            _text(slide, card['owner'], left + .15, top + .08, width - .3, .3, size=12, bold=True, color='#FFFFFF')
            box = slide.shapes.add_textbox(Inches(left + .15), Inches(top + .38), Inches(width - .3), Inches(.6))
            _write_runs(box.text_frame, [(f"{card['points']:,.1f}", True, '#FFFFFF'), ('  National points lost', False, '#FFFFFF')],
                        size=11, color='#FFFFFF')
            box.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
            facts = []
            if card.get('main'):
                facts.append(('Main issue: ', card['main']))
            if card.get('areas'):
                facts.append(('Biggest areas: ', ', '.join(f"{area['name']} ({area['setting'].lower()}) {area['points']:,.1f}"
                                                           for area in card['areas'])))
            cities = sorted({site.get('city') for site in card.get('high_sites') or [] if site.get('city')})
            facts.append(('High sites: ', f"{len(card['high_sites'])} ({'; '.join(cities)})" if card.get('high_sites') else 'none'))
            text = slide.shapes.add_textbox(Inches(left + .15), Inches(top + 1.15), Inches(width - .3), Inches(height - 2.4))
            text.text_frame.word_wrap = True
            for position, (label, value) in enumerate(facts):
                if position:
                    text.text_frame.add_paragraph()
                paragraph = text.text_frame.paragraphs[-1]
                paragraph.space_after = Pt(4)
                for content, bold in ((label, True), (value, False)):
                    run = paragraph.add_run()
                    run.text = content
                    run.font.size = Pt(9)
                    run.font.bold = bold
                    run.font.color.rgb = _rgb(INK)
            where_top = top + height - 1.15
            _box(slide, left + .15, where_top - .08, width - .3, .01, '#D9DEE3')
            _text(slide, 'Where its points are', left + .15, where_top, width - .3, .22, size=8.5, bold=True, color=MUTED)
            if card['settings'] is None:
                _text(slide, 'Recalculate the Saved Filter (Scoring Calculation) to see them.', left + .15, where_top + .27, width - .3, .22,
                      size=8.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
            for row, (setting, value) in enumerate((card['settings'] or {}).items()):
                y = where_top + .27 + row * .26
                _text(slide, setting, left + .15, y, 1.35, .22, size=8.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)
                span = (width - 2.15) * value / peak
                if span > .005:
                    _box(slide, left + 1.5, y + .04, max(.03, span), .14, color).name = 'Scoring Owner Bar'
                _text(slide, f'{value:,.1f}', left + width - .6, y, .45, .22, size=8.5, bold=True, color=INK,
                      align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        note = ('Performance owner: the Region of the tests. Points lost: the National points of the scoring lost against the '
                'maximum in its tests, not the own score of its Region (see the National & Geographical Areas Trends). '
                'High site: 1+ point lost and losses on the same site in an earlier campaign.')
        if item.get('unidentified', 0) > .05:
            note += f" {item['unidentified']:,.1f} points have no Region."
        _footnote(slide, note)


def _band(value: float | None) -> str:
    strong = active_theme()['band_strong']
    if value is None:
        return '#C3CCD5'
    return strong['High'] if value >= 85 else strong['Medium'] if value >= 75 else strong['Low']


def add_place_ranking_slides(presentation, items: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                             owner_color: Callable[[str], str], scope: ScopeLabels | None = None) -> None:
    """The Cities, Towns, Roads and Railways of an operator ranked: their owners, points, change, Voice and Data shares
    and points lost (one bar scale on every ranking)."""
    for item in items:
        peak = max([row['lost'] for group in item['groups'] for row in group['rows']] + [.01])
        previous = item.get('previous_campaign') or 'previous'
        for group in item['groups']:
            pages = [group['rows'][start:start + PLACE_ROWS_PER_SLIDE] for start in range(0, len(group['rows']), PLACE_ROWS_PER_SLIDE)]
            for number, rows in enumerate(pages, start=1):
                title = f"{group['title']} — {item['label']}" + (f' ({number} of {len(pages)})' if len(pages) > 1 else '')
                slide = new_slide(presentation, title, _subtitle(subtitle, scope, item.get('context'), item.get('campaign') or ''))
                top = _headline(slide, group['headline'])
                legend = slide.shapes.add_textbox(Inches(LEFT), Inches(top - .08), Inches(WIDTH), Inches(.25))
                _write_runs(legend.text_frame, [('Colour: ', True, INK), ('85% or more High, 75–85% Medium, under 75% Low. ', False, MUTED),
                                                (f'vs {previous}: ', True, INK), ('change in Points % (–: not measured). ', False, MUTED),
                                                ('Points lost: ', True, INK), ('one bar scale on every ranking.', False, MUTED)],
                            size=8.5, color=MUTED)
                columns = (('Place', LEFT, 2.75), ('Owner', LEFT + 2.8, 2.9), ('Points %', LEFT + 5.75, 1.6), (f'vs {previous}', LEFT + 7.4, .8),
                           ('Voice %', LEFT + 8.25, 1.45), ('Data %', LEFT + 9.75, 1.45), ('Points lost', LEFT + 11.25, 1.2))
                header_top = top + .25
                for text, left, width in columns:
                    _text(slide, text, left, header_top, width, .22, size=8.5, bold=True, color=MUTED)
                _box(slide, LEFT, header_top + .26, WIDTH, .01, '#C9D2DB')
                row_top = header_top + .32
                row_height = min(.3, (CONTENT_BOTTOM - .35 - row_top) / max(1, len(rows)))
                size = 9 if row_height >= .26 else 8
                for index, row in enumerate(rows):
                    y = row_top + index * row_height
                    if index % 2:
                        _box(slide, LEFT, y, WIDTH, row_height, '#F6F7F9').name = 'Scoring Place Band'
                    _text(slide, f"{row['location']}{' ' + row['note'] if row['note'] else ''}", LEFT + .05, y, 2.7, row_height,
                          size=size, color=INK, anchor=MSO_ANCHOR.MIDDLE)
                    x = LEFT + 2.8
                    for owner in row['owners']:
                        x += _pill(slide, f"{owner['name']} {owner['share']:.0f}%", x, y + (row_height - .2) / 2,
                                   owner_color(owner['name']), height=.2, size=7.5) + .06
                    for value, left in ((row['percent'], LEFT + 5.75), (row.get('voice_percent'), LEFT + 8.25),
                                        (row.get('data_percent'), LEFT + 9.75)):
                        if value is None:
                            _text(slide, '—', left, y, 1.0, row_height, size=size, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
                            continue
                        _box(slide, left, y + row_height * .18, max(.02, 1.0 * min(1.0, value / 100)), row_height * .64,
                             _band(value)).name = 'Scoring Place Bar'
                        _text(slide, f'{value:.0f}%', left + 1.03, y, .45, row_height, size=size, bold=left == LEFT + 5.75,
                              color=INK, anchor=MSO_ANCHOR.MIDDLE)
                    delta = row.get('delta')
                    _text(slide, '–' if delta is None else f'{delta:+.0f}', LEFT + 7.4, y, .7, row_height, size=size, bold=True,
                          color=MUTED if delta is None or abs(delta) < .5 else (active_theme()['gain_text'] if delta > 0
                                                                                else active_theme()['loss_text']),
                          anchor=MSO_ANCHOR.MIDDLE)
                    span = .8 * row['lost'] / peak
                    if span > .005:
                        _box(slide, LEFT + 11.25, y + row_height * .22, max(.02, span), row_height * .56, themed(INK)).name = 'Scoring Stake Bar'
                    _text(slide, f"{row['lost']:,.1f}", LEFT + 11.25 + max(.02, span) + .05, y, .5, row_height, size=size,
                          color=INK, anchor=MSO_ANCHOR.MIDDLE)
                _footnote(slide, 'Points %: the share of the points of the place; Points lost: the points it loses, by its '
                                 'share of the tests (national estimates).')


def render_owner_route_map(route: dict[str, Any], owner_color: Callable[[str], str], boundaries: dict[str, list] | None, *,
                           width: int = 1100, height: int = 1100) -> bytes:
    """PNG map of a route of Where to act, drawn as the map of the Route stretches (render_stretch_map): the stretches
    with the most points lost numbered and named after their nearest city, and the owner of each stretch as the
    ring of its circle."""
    from src.modules.scoring_summary_slides import render_stretch_map
    return render_stretch_map(route, '#64748B', [], boundaries=boundaries, width=width, height=height,
                              ranked=route['top'], ring_color=owner_color, measure='Points lost')


def add_route_focus_slides(presentation, items: list[dict[str, Any]], *, subtitle: str, new_slide: Callable,
                           owner_color: Callable[[str], str], boundaries: dict[str, list] | None = None,
                           scope: ScopeLabels | None = None) -> None:
    """The route that loses most for each operator, on one slide: the map, the points of each owner, the five stretches
    with the most points and what to review."""
    for item in items:
        if not item['routes']:
            continue
        route = item['routes'][0]
        slide = new_slide(presentation, f"{route['route']}: where the points are — {item['label']}",
                          _subtitle(subtitle, scope, item.get('context')))
        top = _headline(slide, route['headline'])
        map_size = CONTENT_BOTTOM - 1.15 - top
        picture = slide.shapes.add_picture(BytesIO(render_owner_route_map(route, owner_color, boundaries)),
                                           Inches(LEFT), Inches(top), Inches(map_size), Inches(map_size))
        picture.name = 'Scoring Owner Route Map'
        legend_top = top + map_size + .06
        x = LEFT
        for owner in [owner for owner in route['owners'] if not owner.get('unidentified')]:
            _box(slide, x, legend_top + .05, .14, .14, owner_color(owner['name']))
            _text(slide, owner['name'], x + .2, legend_top, 1.4, .24, size=8.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)
            x += .35 + len(owner['name']) * .075
        right = LEFT + map_size + .35
        side = LEFT + WIDTH - right
        _text(slide, 'Points lost by owner', right, top, side, .3, size=11, bold=True, color=INK)
        widths = [side * .4, side * .3, side * .3]
        rows = route['owners']
        shape, table = _table(slide, len(rows) + 1, 3, right, top + .32, side, .3, widths)
        shape.name = 'Scoring Route Owners Table'
        _header(table, ['Owner', 'Share of tests', 'Points'], size=9, numeric_from=1)
        for index, owner in enumerate(rows, start=1):
            _cell(table.cell(index, 0), owner['name'], size=9, bold=True,
                  color='#FFFFFF' if not owner.get('unidentified') else INK,
                  fill=owner_color(owner['name']) if not owner.get('unidentified') else '#E9ECEF')
            _cell(table.cell(index, 1), '–' if owner.get('tests_share') is None else f"{owner['tests_share']:.0f}%", size=9,
                  align=PP_ALIGN.RIGHT)
            _cell(table.cell(index, 2), f"{owner['points']:,.1f}", size=9, bold=True, fill=None, align=PP_ALIGN.RIGHT)
        _row_bars(slide, [owner['points'] for owner in rows], left=right + widths[0] + widths[1], top=top + .32 + .3,
                  width=widths[2], row_height=.3, color=LOSS_COLOR)
        _to_front(slide, shape)
        stretch_top = top + .32 + .3 * (len(rows) + 1) + .35
        _text(slide, f"The {len(route['top'])} stretches with the most points", right, stretch_top, side, .3, size=11, bold=True, color=INK)
        widths = [.35, side - .35 - 1.0 - 1.45, 1.0, 1.45]
        shape, table = _table(slide, len(route['top']) + 1, 4, right, stretch_top + .32, side, .3, widths)
        shape.name = 'Scoring Route Stretches Table'
        _header(table, ['#', 'Stretch', 'Points', 'Main owner'], size=9, numeric_from=2)
        for index, row in enumerate(route['top'], start=1):
            _cell(table.cell(index, 0), str(index), size=9, bold=True, color='#FFFFFF', fill='#1C1C28', align=PP_ALIGN.CENTER)
            _cell(table.cell(index, 1), row['label'], size=9)
            _cell(table.cell(index, 2), f"{row['points']:,.2f}", size=9, bold=True, fill=None, align=PP_ALIGN.RIGHT)
            owner = row.get('region') or ''
            _cell(table.cell(index, 3), owner or '—', size=9, bold=True, color='#FFFFFF' if owner else INK,
                  fill=owner_color(owner) if owner else '#FFFFFF', align=PP_ALIGN.CENTER)
        _row_bars(slide, [row['points'] for row in route['top']], left=right + widths[0] + widths[1],
                  top=stretch_top + .32 + .3, width=widths[2], row_height=.3, color=LOSS_COLOR)
        _to_front(slide, shape)
        if route.get('review'):
            review_top = CONTENT_BOTTOM - .82
            _box(slide, LEFT, review_top, WIDTH, .42, '#F2F4F6').name = 'Scoring Act Review'
            _box(slide, LEFT, review_top, .07, .42, active_theme()['loss']).name = 'Scoring Act Review Stripe'
            box = slide.shapes.add_textbox(Inches(LEFT + .2), Inches(review_top + .08), Inches(WIDTH - .3), Inches(.3))
            _write_runs(box.text_frame, [('Please review: ', True, active_theme()['loss']), (route['review'], False, INK)],
                        size=11, color=INK)
        _footnote(slide, 'A stretch gathers the tests of the route within about 20 km, named after the nearest city of the '
                         'measurements; circle area: points lost, in the colour of its owner.')
