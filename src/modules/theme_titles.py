"""The titles and subtitles of the content slides in the style of the active theme (see scoring_themes).

Each theme gives the titles its font, size and weight, the part of a title that tells it from the slides next to it
(its focus: what follows " — ", " · " or ": ", or the start of " per ") in bold and in its accent, NetCheck the first part in
capitals as the NetCheck reports, and the subtitles (the paragraphs under the title) their font, size, weight and
colour. Scoring & GAP Analysis and Network Insights style every content slide so; their covers and transitions keep the
style of the template (see report_structure).
"""
from __future__ import annotations

from copy import deepcopy
from typing import Any

from pptx.dml.color import RGBColor
from pptx.util import Pt

from src.modules.scoring_themes import active_theme

FOCUS_SEPARATORS = (' — ', ' · ', ': ')


def _rgb(colour: Any) -> RGBColor:
    return RGBColor.from_string(str(colour).lstrip('#').upper())


def title_style(theme: dict[str, Any] | None = None) -> dict[str, Any]:
    """The style of the titles of a theme with its colours: title, focus and subtitle."""
    theme = theme or active_theme()
    titles = theme.get('titles') or {}
    return {**titles, 'color': theme['title'], 'focus_color': titles.get('focus_color') or theme['title_accent'],
            'subtitle_color': titles.get('subtitle_color') or theme['subtitle']}


def _title_shape(slide: Any) -> Any:
    return next((shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3} and shape.has_text_frame), None)


def split_focus(text: str) -> tuple[str, str]:
    """The heading of a title and its focus: what follows " — ", " · " or ": ", or the start of " per "."""
    for separator in FOCUS_SEPARATORS:
        if separator in text:
            heading, found, focus = text.partition(separator)
            return heading + found, focus
    heading, found, focus = text.partition(' per ')
    return (heading + ' ', 'per ' + focus) if found else (text, '')


def _runs(paragraph: Any, parts: list[tuple[str, bool]]) -> list[tuple[Any, bool]]:
    """Write ``parts`` (text, is focus) as the runs of the paragraph, keeping the properties of its first run."""
    properties = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    written = []
    for text, focus in parts:
        if not text:
            continue
        run = paragraph.add_run()
        if properties is not None:
            run._r.insert(0, deepcopy(properties))
        run.text = text
        written.append((run, focus))
    return written


def style_title(slide: Any, *, subtitle_size: float | None = None, theme: dict[str, Any] | None = None) -> None:
    """Give the title (its first paragraph) and the subtitle (the other paragraphs) of a content slide the style of the
    theme. A title written in runs keeps them: its bold runs are its focus; otherwise its focus is found in its text.
    ``subtitle_size`` replaces the size of the subtitle of the theme (a subtitle made smaller to fit its line)."""
    from src.modules.slide_titles import fit_slide_title

    shape = _title_shape(slide)
    if shape is None or not shape.text_frame.paragraphs:
        return
    style = title_style(theme)
    paragraphs = shape.text_frame.paragraphs
    title = paragraphs[0]
    runs = [run for run in title.runs if run.text]
    if len(runs) > 1 and any(run.font.bold for run in runs):
        parts = [(run, bool(run.font.bold)) for run in runs]
    else:
        heading, focus = split_focus(title.text)
        parts = _runs(title, [(heading, False), (focus, True)])
    for run, focus in parts:
        font = run.font
        font.name = style.get('font') or None
        font.size = Pt(style.get('size', 28))
        font.bold = bool(style.get('focus_bold') if focus else style.get('bold'))
        font.italic = bool(style.get('focus_italic') if focus else style.get('italic'))
        font.color.rgb = _rgb(style['focus_color'] if focus else style['color'])
        # The heading in capitals without changing its text (NetCheck).
        if style.get('caps') and not focus:
            font._rPr.set('cap', 'all')
        elif font._rPr.get('cap') is not None:
            del font._rPr.attrib['cap']
    title.font.size = Pt(style.get('size', 28))
    fit_slide_title(slide, style.get('size', 28))
    for paragraph in paragraphs[1:]:
        text = ''.join(run.text for run in paragraph.runs)
        lead, separator, rest = text.partition(' · ')
        written = (_runs(paragraph, [(lead + separator, True), (rest, False)]) if style.get('lead_bold') and separator
                   else [(run, False) for run in paragraph.runs])
        size = Pt(subtitle_size or style.get('subtitle_size', 16))
        paragraph.font.size = size
        for run, lead_part in written:
            font = run.font
            font.name = style.get('subtitle_font') or style.get('font') or None
            font.size = size
            font.bold = bool(style.get('subtitle_bold') or lead_part)
            font.italic = bool(style.get('subtitle_italic'))
            font.color.rgb = _rgb(style['subtitle_color'])
