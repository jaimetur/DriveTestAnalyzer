from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from time import sleep
from typing import Callable, Iterable

import pandas as pd
from openpyxl import load_workbook
try:
    import pycountry
except ImportError:  # pragma: no cover - packaged installations include it.
    pycountry = None

from src.modules.cdr_field_profiles import NETCHECK_PROFILE, CdrFieldProfile, field_profile
from src.modules.column_names import (
    OPERATOR_VENDOR_FIELD, VENDOR_FIELD, campaign_parts, clean_column_name, column_identity, mapped_vendor_only_value,
    operator_vendor_value, resolve_column_name,
)


# The worksheets of NetCheck workbooks that never hold CDR rows (each CDR type lists its own, see cdr_field_profiles).
CDR_IGNORED_SHEETS = NETCHECK_PROFILE.ignored_sheets
CDR_IGNORED_SHEET_KEYS = frozenset(name.strip().casefold() for name in CDR_IGNORED_SHEETS)

# XLSX parsing is implemented largely in Python by openpyxl.  A short pause at
# predictable intervals releases the GIL to keep the web server responsive
# while a large CDR is being processed in the background.
# Prioritize an interactive server over finishing a background CDR import a
# few seconds sooner.  Large workbooks otherwise keep the interpreter busy
# enough that every browser request feels stalled.
EXCEL_READ_YIELD_EVERY_ROWS = 100
EXCEL_READ_YIELD_SECONDS = 0.01


@dataclass(slots=True)
class DatasetSummary:
    rows: int
    columns: list[str]
    numeric_columns: list[str]
    categorical_columns: list[str]


def _first_available_series(df: pd.DataFrame, candidates: Iterable[str], default: object = pd.NA) -> pd.Series:
    for column in candidates:
        resolved = resolve_column_name(df.columns, column)
        if resolved is not None:
            return df[resolved]
    return pd.Series([default] * len(df), index=df.index, dtype='object')


def _numeric_column(df: pd.DataFrame, candidates: Iterable[str]) -> pd.Series:
    """Return the first matching column as numeric values, preserving row alignment."""
    return pd.to_numeric(_first_available_series(df, candidates), errors='coerce')


def add_vfuk_gcid_column(df: pd.DataFrame) -> pd.DataFrame:
    """Materialise the VFUK mapping GCID for the 4G and 5G source worksheets.

    4G cells take their ECI, ``eNodeB ID × 256 + Cell ID``, and 5G cells their NCI,
    ``gNodeB ID × 2^(36 − gNodeB ID length) + Cell ID`` (22-bit gNodeB IDs). The Local Cell ID
    is used only without a Cell ID: Huawei and Nokia cells carry a different Local Cell ID, which
    is not part of the identity the cell broadcasts and the CDRs report.
    """
    from src.modules.cell_identity import lte_eci_series, nr_nci_series

    dataset = df.copy()
    source_sheet = _first_available_series(dataset, ['source_sheet']).astype(str).str.strip().str.casefold()
    cell_id = _first_available_series(dataset, ['Cell ID', 'Cell_ID', 'cellId'])
    local_cell_id = _first_available_series(dataset, ['Local Cell ID', 'Local_Cell_ID'])
    gcid = pd.Series(pd.NA, index=dataset.index, dtype='Int64')

    lte = lte_eci_series(_first_available_series(dataset, ['eNodeB ID', 'eNodeB_ID', 'eNB ID']), cell_id, local_cell_id)
    gcid = gcid.mask(source_sheet.eq('4g'), lte)
    nr = nr_nci_series(
        _first_available_series(dataset, ['gNodeB ID', 'gNodeB_ID', 'gNB ID']), cell_id, local_cell_id,
        _first_available_series(dataset, ['gNodeB ID Length', 'gNodeB_ID_Length']),
    )
    gcid = gcid.mask(source_sheet.eq('5g'), nr)

    dataset['GCID'] = gcid.astype('Int64')
    return dataset


def add_three_gcid_column(df: pd.DataFrame) -> pd.DataFrame:
    """Materialise the 3UK mapping GCID from its source ECI column."""
    dataset = df.copy()
    dataset['GCID'] = _first_available_series(dataset, ['CId___ECI', 'Cid__ECI'])
    return dataset


def _parse_campaign_dimension(value: object, part: str) -> object:
    year, quarter, _mode = campaign_parts(value)
    if part == 'year':
        return int(year) if year else pd.NA
    if part == 'quarter':
        return quarter or pd.NA
    if part == 'period':
        return f'{year}-{quarter}' if year and quarter else pd.NA
    return pd.NA


ISO_COUNTRY_CODES = frozenset((
    'AD AE AF AG AI AL AM AO AQ AR AS AT AU AW AX AZ BA BB BD BE BF BG BH BI BJ BL BM BN BO BQ BR BS BT BV BW BY BZ '
    'CA CC CD CF CG CH CI CK CL CM CN CO CR CU CV CW CX CY CZ DE DJ DK DM DO DZ EC EE EG EH ER ES ET FI FJ FK FM FO FR '
    'GA GB GD GE GF GG GH GI GL GM GN GP GQ GR GS GT GU GW GY HK HM HN HR HT HU ID IE IL IM IN IO IQ IR IS IT JE JM JO '
    'JP KE KG KH KI KM KN KP KR KW KY KZ LA LB LC LI LK LR LS LT LU LV LY MA MC MD ME MF MG MH MK ML MM MN MO MP MQ MR '
    'MS MT MU MV MW MX MY MZ NA NC NE NF NG NI NL NO NP NR NU NZ OM PA PE PF PG PH PK PL PM PN PR PS PT PW PY QA RE RO '
    'RS RU RW SA SB SC SD SE SG SH SI SJ SK SL SM SN SO SR SS ST SV SX SY SZ TC TD TF TG TH TJ TK TL TM TN TO TR TT TV TW '
    'TZ UA UG UM US UY UZ VA VC VE VG VI VN VU WF WS YE YT ZA ZM ZW UK'
).split())
COUNTRY_NAME_TO_CODE = {
    'united kingdom': 'UK', 'great britain': 'UK', 'england': 'UK', 'germany': 'DE',
    'spain': 'ES', 'france': 'FR', 'italy': 'IT', 'portugal': 'PT', 'ireland': 'IE',
    'netherlands': 'NL', 'belgium': 'BE', 'switzerland': 'CH', 'austria': 'AT',
    'poland': 'PL', 'sweden': 'SE', 'norway': 'NO', 'denmark': 'DK', 'finland': 'FI',
    'united states': 'US', 'united states of america': 'US', 'canada': 'CA', 'mexico': 'MX',
    'brazil': 'BR', 'argentina': 'AR', 'chile': 'CL', 'colombia': 'CO',
    'australia': 'AU', 'new zealand': 'NZ', 'india': 'IN', 'china': 'CN', 'japan': 'JP',
    'south korea': 'KR', 'singapore': 'SG', 'south africa': 'ZA', 'turkey': 'TR',
    'united arab emirates': 'AE', 'saudi arabia': 'SA',
}
if pycountry is not None:
    for country in pycountry.countries:
        for attribute in ('name', 'official_name', 'common_name'):
            name = getattr(country, attribute, '')
            if name:
                COUNTRY_NAME_TO_CODE.setdefault(str(name).casefold(), str(country.alpha_2))
# NetCheck uses UK rather than ISO's GB in campaign labels.
COUNTRY_NAME_TO_CODE.update({'united kingdom': 'UK', 'great britain': 'UK', 'england': 'UK'})


# Longest country names first, so "united kingdom" wins over shorter names.
_COUNTRY_NAME_PATTERNS = tuple(
    (re.compile(rf'(?:^|[_\- ]){re.escape(country_name)}(?:$|[_\- ])'), code)
    for country_name, code in sorted(COUNTRY_NAME_TO_CODE.items(), key=lambda item: -len(item[0]))
)
_COUNTRY_CODE_PATTERN = re.compile(r'(?:^|[_\- ])([A-Za-z]{2})(?=$|[_\- ])')


def _campaign_market(value: object) -> object:
    text = '' if value is None else str(value).strip()
    return _campaign_market_text(text)


@lru_cache(maxsize=4096)
def _campaign_market_text(text: str) -> object:
    if text.casefold() in {'<na>', 'nan', 'nat', 'none'}:
        return pd.NA
    lowered = text.casefold()
    for pattern, code in _COUNTRY_NAME_PATTERNS:
        if pattern.search(lowered):
            return code
    code_matches = list(_COUNTRY_CODE_PATTERN.finditer(text))
    country_codes = [match.group(1).upper() for match in code_matches if match.group(1).upper() in ISO_COUNTRY_CODES]
    non_mode_codes = [code for code in country_codes if code != 'SA']
    if non_mode_codes:
        return non_mode_codes[0]
    if country_codes == ['SA'] and text.upper().startswith('SA'):
        return 'SA'
    return pd.NA


def _has_values(series: pd.Series) -> bool:
    return series.notna().any() and series.fillna('').astype(str).str.strip().ne('').any()


def _ensure_dataset_field(dataset: pd.DataFrame, target: str, candidates: Iterable[str], default: object = pd.NA) -> str:
    actual = resolve_column_name(dataset.columns, target)
    if actual is not None and _has_values(dataset[actual]):
        return actual
    source = next(
        (resolved for candidate in candidates if (resolved := resolve_column_name(dataset.columns, candidate)) is not None and _has_values(dataset[resolved])),
        None,
    )
    values = dataset[source] if source is not None else pd.Series([default] * len(dataset), index=dataset.index, dtype='object')
    if actual is None:
        dataset[target] = values
        return target
    dataset[actual] = values
    return actual


def ensure_vendor_fields(dataset: pd.DataFrame, operator_column: str) -> pd.DataFrame:
    """Store Operator_Vendor (<Operator>_<Vendor>) and Vendor (the vendor alone).

    The source CDR's Vendor, or Operator_Vendor when the CDR already has it, is
    the operator-specific vendor. Files with the former Vendor_Only field keep
    their Vendor as Operator_Vendor. Operators without a vendor store
    "<Operator> - All" in both fields.
    """
    columns = list(dataset.columns)
    composite = resolve_column_name(columns, OPERATOR_VENDOR_FIELD)
    vendor_columns = [column for column in columns if column_identity(column) in {'vendor', 'vendoronly', 'vendor2'}]
    source = composite or next(
        (column for column in vendor_columns if column_identity(column) == 'vendor'), None,
    )
    position = min((columns.index(column) for column in [*vendor_columns, *([composite] if composite else [])]), default=len(columns))
    operators = dataset[operator_column].fillna('').astype(str).str.strip()
    raw_values = dataset[source] if source else pd.Series([''] * len(dataset), index=dataset.index)
    raw = raw_values.fillna('').astype(str).str.strip()
    # Each distinct (vendor, operator) pair is resolved once.
    resolved = {}
    for pair in set(zip(raw, operators, strict=False)):
        operator_vendor = operator_vendor_value(*pair)
        resolved[pair] = (operator_vendor, mapped_vendor_only_value(operator_vendor, pair[1]))
    rows = [resolved[pair] for pair in zip(raw, operators, strict=False)]
    composite_values = [row[0] for row in rows]
    vendor_values = [row[1] for row in rows]
    dataset = dataset.drop(columns=[*vendor_columns, *([composite] if composite else [])])
    position = min(position, len(dataset.columns))
    dataset.insert(position, OPERATOR_VENDOR_FIELD, composite_values)
    dataset.insert(position + 1, VENDOR_FIELD, vendor_values)
    return dataset


def ensure_fixed_cdr_fields(dataset: pd.DataFrame, profile: CdrFieldProfile = NETCHECK_PROFILE) -> pd.DataFrame:
    """Add the CDR-Main fields to every row, each from the columns of the CDR type's field profile."""
    sources = profile.fixed_sources
    legacy_report_columns = [column for column in dataset.columns if column_identity(column) == 'reportvendor']
    if legacy_report_columns:
        dataset = dataset.drop(columns=legacy_report_columns)
    campaign = _ensure_dataset_field(dataset, 'Campaign', sources['Campaign'])
    benchmark = _ensure_dataset_field(dataset, 'Benchmark', (*sources['Benchmark'], campaign))
    campaign_values = dataset[campaign]
    benchmark_values = dataset[benchmark]

    def extracted(part: str) -> pd.Series:
        primary = campaign_values.map({value: _parse_campaign_dimension(value, part) for value in pd.unique(campaign_values)})
        fallback = benchmark_values.map({value: _parse_campaign_dimension(value, part) for value in pd.unique(benchmark_values)})
        return primary.where(primary.notna(), fallback)

    year = _ensure_dataset_field(dataset, 'Campaign_Year', sources['Campaign_Year'])
    quarter = _ensure_dataset_field(dataset, 'Campaign_Quarter', sources['Campaign_Quarter'])
    existing_year = dataset[year].copy()
    existing_quarter = dataset[quarter].copy()
    parsed_year = extracted('year')
    parsed_quarter = extracted('quarter')
    dataset[year] = parsed_year.where(parsed_year.notna(), existing_year)
    dataset[quarter] = parsed_quarter.where(parsed_quarter.notna(), existing_quarter)
    period = _ensure_dataset_field(dataset, 'Period', sources['Period'])
    existing_period = dataset[period].copy()
    dataset[period] = [
        f'{year_value}-{quarter_value}' if pd.notna(year_value) and pd.notna(quarter_value) else previous
        for year_value, quarter_value, previous in zip(dataset[year], dataset[quarter], existing_period, strict=False)
    ]
    market = _ensure_dataset_field(dataset, 'Market', sources['Market'])
    existing_market = dataset[market].copy()
    # Campaign labels repeat on every row: derive each distinct label once.
    primary_market = campaign_values.map({value: _campaign_market(value) for value in pd.unique(campaign_values)})
    fallback_market = benchmark_values.map({value: _campaign_market(value) for value in pd.unique(benchmark_values)})
    parsed_market = primary_market.where(primary_market.notna(), fallback_market)
    dataset[market] = parsed_market.where(parsed_market.notna(), existing_market)
    _ensure_dataset_field(dataset, 'Region', sources['Region'])
    _ensure_dataset_field(dataset, 'Zone', sources['Zone'])
    _ensure_dataset_field(dataset, 'City', sources['City'])
    operator = _ensure_dataset_field(dataset, 'Operator', sources['Operator'])
    _ensure_dataset_field(dataset, 'Subscriber', (*sources['Subscriber'], operator))
    dataset = ensure_vendor_fields(dataset, operator)
    technology_fields, session_fields, status_fields = profile.fixed_groups
    for fields, candidates in (technology_fields, session_fields):
        for field in fields:
            _ensure_dataset_field(dataset, field, candidates)
    _ensure_dataset_field(dataset, 'Test_Name', sources['Test_Name'])
    fields, candidates = status_fields
    for field in fields:
        _ensure_dataset_field(dataset, field, candidates)
    return dataset


def canonical_operator_value(value: object, mappings: dict[str, str]) -> object:
    """Return the current workspace label for one non-blank Operator value."""
    if pd.isna(value) or not str(value).strip():
        return value
    text = str(value).strip()
    return mappings.get(text.casefold(), text)


def apply_operator_mappings(dataset: pd.DataFrame, mappings: dict[str, str]) -> pd.DataFrame:
    """Apply workspace operator aliases without changing blank values."""
    if not mappings:
        return dataset
    result = dataset.copy()
    operator_column = resolve_column_name(result.columns, 'Operator')
    subscriber_column = resolve_column_name(result.columns, 'Subscriber')
    subscriber_is_derived = bool(
        operator_column and subscriber_column
        and result[operator_column].fillna('').astype(str).equals(result[subscriber_column].fillna('').astype(str))
    )
    for column in result.columns:
        if column_identity(column) != 'operator':
            continue
        # Map each distinct label once; missing values keep their original form.
        values = result[column]
        present = values.notna()
        canonical = {value: canonical_operator_value(value, mappings) for value in pd.unique(values[present])}
        result[column] = values.map(canonical).where(present, values)
    if subscriber_is_derived and operator_column and subscriber_column:
        result[subscriber_column] = result[operator_column]
    return result


def _count_items(value: object) -> float:
    if pd.isna(value):
        return 0.0
    text = str(value).strip()
    if not text:
        return 0.0
    if '->' in text:
        return float(text.count('->') + 1)
    if '],' in text:
        return float(text.count('],') + 1)
    if ',' in text and '[' in text:
        return float(text.count(',') + 1)
    return 1.0


def _average_columns(df: pd.DataFrame, candidates: Iterable[str]) -> pd.Series:
    numeric_columns = [column for column in candidates if column in df.columns]
    if not numeric_columns:
        return pd.Series([pd.NA] * len(df), index=df.index, dtype='Float64')
    numeric_frame = df[numeric_columns].apply(pd.to_numeric, errors='coerce')
    return numeric_frame.mean(axis=1).astype('Float64')


def _sheet_has_only_empty_values(values: tuple[object, ...]) -> bool:
    return all(value is None or str(value).strip() == '' for value in values)


def _make_unique_headers(headers: Iterable[object], side_pattern: str = '') -> list[str]:
    """Return stable, unique column labels while preserving the first occurrence.

    With a side pattern (see CdrFieldProfile.side_header_pattern), a repeated header takes the side of the nearest
    header before it that matches the pattern: the second Sequence of a call is "Side2 Sequence".
    """
    side_header = re.compile(side_pattern, re.IGNORECASE) if side_pattern else None
    counts: dict[str, int] = {}
    unique_headers: list[str] = []
    used: set[str] = set()
    side = ''
    for index, value in enumerate(headers, start=1):
        base_name = clean_column_name(value) if value is not None and str(value).strip() else ''
        if not base_name:
            base_name = f'Unnamed_{index}'
        identity = base_name.casefold()
        counts[identity] = counts.get(identity, 0) + 1
        occurrence = counts[identity]
        name = base_name
        if occurrence > 1:
            sided = f'{side} {base_name}' if side else ''
            name = sided if sided and sided.casefold() not in used else f'{base_name}_Duplicate_{occurrence}'
        if side_header and (match := side_header.match(base_name)):
            side = re.sub(r'\s+', '', match.group(0)).capitalize()
        used.add(name.casefold())
        unique_headers.append(name)
    return unique_headers


def _operator_sheet_names(workbook, profile: CdrFieldProfile = NETCHECK_PROFILE) -> list[str]:
    """Return the contiguous operator-sheet block following KPI Definition.

    NetCheck CDR workbooks place their operator extracts immediately after the
    KPI Definition worksheet (the sheet of the CDR type's field profile).
    Later worksheets are helper material and must never become CDR rows.
    Every operator sheet in that block is imported, including
    technology-specific ones such as ``Vodafone VoNR``.
    """
    ignored = frozenset(name.strip().casefold() for name in profile.ignored_sheets)
    anchor = profile.operator_sheets_after.strip().casefold()
    worksheets = list(workbook.worksheets)
    definition_index = next(
        (
            index for index, sheet in enumerate(worksheets)
            if anchor and sheet.title.strip().casefold() == anchor
        ),
        None,
    )
    if definition_index is None:
        return [
            sheet.title for sheet in worksheets
            if sheet.sheet_state == 'visible' and sheet.title.strip().casefold() not in ignored
        ]

    candidates: list[str] = []
    for sheet in worksheets[definition_index + 1:]:
        key = sheet.title.strip().casefold()
        if sheet.sheet_state != 'visible' or key in ignored:
            # Ranking/helper sheets commonly sit between KPI Definition and
            # the first operator export. Skip that leading metadata, then stop
            # at the first helper sheet after the operator block has started.
            if candidates:
                break
            continue
        candidates.append(sheet.title)
    return candidates


def get_excel_sheet_columns(file_path: Path, sheet_name: str) -> list[str]:
    """Return a worksheet's source headers using the ingestion naming rules."""
    if file_path.suffix.lower() not in {'.xlsx', '.xlsm'}:
        return []
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        return []
    for row in workbook[sheet_name].iter_rows(values_only=True):
        values = tuple(row or ())
        if not _sheet_has_only_empty_values(values):
            return _make_unique_headers(values)
    return []


def get_dataset_source_columns(file_path: Path, profile: CdrFieldProfile = NETCHECK_PROFILE) -> list[str]:
    """Return every physical source header using the same names as ingestion."""
    file_path = file_path.resolve()
    stat = file_path.stat()
    return list(_cached_dataset_source_columns(str(file_path), stat.st_mtime_ns, stat.st_size, profile.cdr_type))


@lru_cache(maxsize=128)
def _cached_dataset_source_columns(
    file_path_value: str, modified_ns: int, file_size: int, cdr_type: str = NETCHECK_PROFILE.cdr_type,
) -> tuple[str, ...]:
    """Cache source headers without retaining source data in memory."""
    del modified_ns, file_size
    profile = field_profile(cdr_type)
    file_path = Path(file_path_value)
    suffix = file_path.suffix.lower()
    if suffix == '.csv':
        for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
            try:
                return tuple(clean_column_name(column) for column in pd.read_csv(file_path, nrows=0, encoding=encoding).columns)
            except UnicodeDecodeError:
                continue
        return ()
    if suffix in {'.xlsx', '.xlsm'}:
        workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
        try:
            columns: list[str] = []
            for sheet_name in _operator_sheet_names(workbook, profile):
                worksheet = workbook[sheet_name]
                for row in worksheet.iter_rows(values_only=True):
                    values = tuple(row or ())
                    if _sheet_has_only_empty_values(values):
                        continue
                    for column in _make_unique_headers(values, profile.side_header_pattern):
                        if column not in columns:
                            columns.append(column)
                    break
            return tuple(columns)
        finally:
            workbook.close()
    if suffix == '.xls':
        return tuple(clean_column_name(column) for column in pd.read_excel(file_path, nrows=0).columns)
    return ()


def _read_openxml_sheet(
    worksheet, progress_callback: Callable[[int], None] | None, progress_state: dict[str, int], total_rows: int,
    side_pattern: str = '',
) -> pd.DataFrame:
    rows_iter = worksheet.iter_rows(values_only=True)
    header: list[str] | None = None
    records: list[dict[str, object]] = []

    def advance_progress() -> None:
        progress_state['processed_rows'] += 1
        if progress_state['processed_rows'] % EXCEL_READ_YIELD_EVERY_ROWS == 0:
            sleep(EXCEL_READ_YIELD_SECONDS)
        if not progress_callback or total_rows <= 0:
            return
        progress = min(55, 14 + int((progress_state['processed_rows'] / total_rows) * 41))
        if progress > progress_state['last_progress']:
            progress_state['last_progress'] = progress
            progress_callback(progress)

    for row in rows_iter:
        advance_progress()
        values = tuple(row or ())
        if _sheet_has_only_empty_values(values):
            continue
        header = _make_unique_headers(values, side_pattern)
        break

    if not header:
        return pd.DataFrame()

    for row in rows_iter:
        advance_progress()
        values = tuple(row or ())
        if _sheet_has_only_empty_values(values):
            continue
        padded = list(values[:len(header)])
        if len(padded) < len(header):
            padded.extend([None] * (len(header) - len(padded)))
        records.append(dict(zip(header, padded, strict=False)))

    if not records:
        return pd.DataFrame(columns=header)
    return pd.DataFrame.from_records(records, columns=header)


def infer_dataset_kind(df: pd.DataFrame, file_name: str = '', profile: CdrFieldProfile = NETCHECK_PROFILE) -> str:
    """The CDR type of a file: a word of its name, else the columns of the CDR type's field profile it has most."""
    lower_name = file_name.lower()
    normalized_name = lower_name.replace('_', ' ').replace('-', ' ')
    if 'mapping' in normalized_name:
        return 'mapping'
    for word, kind in profile.kind_name_words:
        if word in normalized_name:
            return kind
    columns = set(df.columns)
    scores = {kind: len(columns & markers) for kind, markers in profile.kind_markers.items()}
    best_kind = max(scores, key=scores.get)
    if scores[best_kind] > 0:
        return best_kind
    return 'generic'


def _normalise_dataset(df: pd.DataFrame, file_path: Path, profile: CdrFieldProfile = NETCHECK_PROFILE) -> pd.DataFrame:
    sources = profile.normalised_sources
    dataset = df.copy()
    dataset.columns = [clean_column_name(column) for column in dataset.columns]
    if profile.derive_columns is not None:
        dataset = profile.derive_columns(dataset, file_path.name)
    source_vendor = resolve_column_name(dataset.columns, OPERATOR_VENDOR_FIELD) or resolve_column_name(dataset.columns, VENDOR_FIELD)
    # Whether the source CDR brings a Vendor for every row; mapped Vendors replace it.
    vendor_values_complete = bool(
        source_vendor and not dataset.empty
        and dataset[source_vendor].fillna('').astype(str).str.strip().ne('').all()
    )
    dataset_kind = infer_dataset_kind(dataset, file_path.name, profile)
    dataset['dataset_kind'] = dataset_kind
    dataset['source_file'] = file_path.name

    dataset['campaign'] = _first_available_series(dataset, sources['campaign'])
    derived_market = dataset['campaign'].map(lambda value: _parse_campaign_dimension(value, 'market'))
    derived_period = dataset['campaign'].map(lambda value: _parse_campaign_dimension(value, 'period'))
    existing_market = _first_available_series(dataset, sources['market'])
    existing_period = _first_available_series(dataset, sources['period'])
    dataset['market'] = existing_market.where(existing_market.notna(), derived_market)
    dataset['period'] = existing_period.where(existing_period.notna(), derived_period)
    dataset['campaign_year'] = dataset['campaign'].map(lambda value: _parse_campaign_dimension(value, 'year'))
    dataset['campaign_quarter'] = dataset['campaign'].map(lambda value: _parse_campaign_dimension(value, 'quarter'))

    for field in ('operator', 'session_type', 'test_name', 'direction', 'region', 'city', 'status'):
        dataset[field] = _first_available_series(dataset, sources[field])
    if profile.status_values:
        # The provider's results in the tool's words (SUCCESS is Completed); any other result is kept.
        status_text = dataset['status'].astype(str).str.strip().str.casefold()
        converted = status_text.map(profile.status_values)
        dataset['status'] = converted.where(converted.notna(), dataset['status'])

    dataset['disturbed'] = _first_available_series(dataset, sources['disturbed']).astype(str).str.lower().eq('yes')
    dataset['impaired'] = _first_available_series(dataset, sources['impaired']).astype(str).str.lower().eq('yes')
    dataset['dropped'] = dataset['status'].astype(str).str.contains('drop', case=False, na=False)
    if profile.dropped_flag in dataset.columns:
        dataset['dropped'] = dataset['dropped'] | dataset[profile.dropped_flag].astype(str).str.lower().eq('yes')
    if profile.unsustainable_flag in dataset.columns:
        dataset['unsustainable_call'] = dataset[profile.unsustainable_flag].astype(str).str.lower().eq('yes')
    else:
        dataset['unsustainable_call'] = False

    status_normalized = dataset['status'].astype(str).str.strip().str.lower()
    dataset['success'] = status_normalized.isin({'completed', 'success', 'ok', 'passed'})
    dataset['failure'] = dataset['status'].notna() & ~dataset['success']

    dataset['event_start_time'] = pd.to_datetime(
        _first_available_series(dataset, sources['event_start_time']), errors='coerce', format='mixed',
    )
    dataset['event_end_time'] = pd.to_datetime(
        _first_available_series(dataset, sources['event_end_time']), errors='coerce', format='mixed',
    )
    dataset['hour_bucket'] = dataset['event_start_time'].dt.hour
    dataset['day_bucket'] = dataset['event_start_time'].dt.day

    for field in ('setup_time_seconds', 'duration_seconds', 'quality_score', 'throughput_mbps', 'latency_ms'):
        dataset[field] = pd.to_numeric(_first_available_series(dataset, sources[field]), errors='coerce')
        if field in profile.normalised_scales:
            dataset[field] = dataset[field] * profile.normalised_scales[field]
    for field, columns in profile.averaged_sources.items():
        dataset[field] = _average_columns(dataset, columns)
    dataset['handovers'] = _first_available_series(dataset, sources['handovers']).map(_count_items)
    dataset['technology_primary'] = _first_available_series(dataset, sources['technology_primary'])
    dataset['technology_secondary'] = _first_available_series(dataset, sources['technology_secondary'])

    dataset = ensure_fixed_cdr_fields(dataset, profile)
    dataset.attrs['vendor_values_complete'] = vendor_values_complete
    start_column = resolve_column_name(dataset.columns, 'Event_Start_Time') or 'event_start_time'
    end_column = resolve_column_name(dataset.columns, 'Event_End_Time') or 'event_end_time'
    start_times = pd.to_datetime(dataset[start_column], errors='coerce', format='mixed')
    end_times = pd.to_datetime(dataset[end_column], errors='coerce', format='mixed')
    dataset[start_column] = start_times.dt.strftime('%Y-%m-%d %H:%M:%S.%f').where(start_times.notna(), pd.NA)
    dataset[end_column] = end_times.dt.strftime('%Y-%m-%d %H:%M:%S.%f').where(end_times.notna(), pd.NA)
    hour_column = resolve_column_name(dataset.columns, 'Hour_Bucket') or 'hour_bucket'
    day_column = resolve_column_name(dataset.columns, 'Day_Bucket') or 'day_bucket'
    dataset[hour_column] = start_times.dt.hour.astype('Int64')
    dataset[day_column] = start_times.dt.day.astype('Int64')

    return dataset


def _capture_source_columns(target: list[str] | None, columns: Iterable[object]) -> None:
    if target is None:
        return
    for column in columns:
        cleaned = clean_column_name(column)
        if cleaned not in target:
            target.append(cleaned)


def _load_excel_dataset(
    file_path: Path,
    progress_callback: Callable[[int], None] | None = None,
    source_columns: list[str] | None = None,
    profile: CdrFieldProfile = NETCHECK_PROFILE,
) -> pd.DataFrame:
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    candidate_sheets = _operator_sheet_names(workbook, profile)
    total_rows = sum(max(workbook[sheet_name].max_row or 0, 1) for sheet_name in candidate_sheets) or 1

    if progress_callback:
        progress_callback(14)

    data_frames: list[pd.DataFrame] = []
    progress_state = {'processed_rows': 0, 'last_progress': 14}
    for sheet_name in candidate_sheets:
        sheet = _read_openxml_sheet(workbook[sheet_name], progress_callback, progress_state, total_rows, profile.side_header_pattern)
        _capture_source_columns(source_columns, sheet.columns)
        sheet = sheet.dropna(axis=0, how='all').dropna(axis=1, how='all')
        if sheet.empty:
            continue
        if all(str(column).startswith('Unnamed:') for column in sheet.columns):
            continue
        sheet['source_sheet'] = sheet_name
        data_frames.append(sheet)

    if not data_frames:
        df = pd.read_excel(file_path)
        _capture_source_columns(source_columns, df.columns)
        if progress_callback:
            progress_callback(55)
        return df

    if len(data_frames) == 1 and data_frames[0]['source_sheet'].nunique() == 1:
        dataset = _normalise_dataset(data_frames[0], file_path, profile)
        if progress_callback:
            progress_callback(55)
        return dataset

    combined = pd.concat(data_frames, ignore_index=True, sort=False)
    dataset = _normalise_dataset(combined, file_path, profile)
    if progress_callback:
        progress_callback(55)
    return dataset


def _read_csv_dataset(file_path: Path) -> pd.DataFrame:
    """Read CSV sources using UTF-8 first, then common Windows legacy encodings."""
    decode_error: UnicodeDecodeError | None = None
    for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            return pd.read_csv(file_path, encoding=encoding)
        except UnicodeDecodeError as exc:
            decode_error = exc
    if decode_error:
        raise decode_error
    raise ValueError(f'Unable to read CSV file: {file_path.name}')


def load_dataset(
    file_path: Path,
    progress_callback: Callable[[int], None] | None = None,
    source_columns: list[str] | None = None,
    profile: CdrFieldProfile = NETCHECK_PROFILE,
) -> pd.DataFrame:
    suffix = file_path.suffix.lower()
    if suffix == '.csv':
        source = _read_csv_dataset(file_path)
        _capture_source_columns(source_columns, source.columns)
        df = _normalise_dataset(source, file_path, profile)
        if progress_callback:
            progress_callback(55)
        return df
    if suffix in {'.xlsx', '.xlsm'}:
        return _load_excel_dataset(
            file_path, progress_callback=progress_callback, source_columns=source_columns, profile=profile,
        )
    if suffix == '.xls':
        source = pd.read_excel(file_path)
        _capture_source_columns(source_columns, source.columns)
        df = _normalise_dataset(source, file_path, profile)
        if progress_callback:
            progress_callback(55)
        return df
    from src.modules.cell_configuration import is_cell_configuration_package, read_cell_configuration

    if is_cell_configuration_package(file_path):
        # Cell Configuration KMZ/KML files, alone or in a ZIP package: one row per cell Placemark.
        source = read_cell_configuration(file_path, progress_callback=progress_callback)
        _capture_source_columns(source_columns, source.columns)
        df = _normalise_dataset(source, file_path, profile)
        if progress_callback:
            progress_callback(55)
        return df
    raise ValueError(f'Unsupported file type: {suffix}')


def summarise_dataset(df: pd.DataFrame) -> DatasetSummary:
    numeric_columns = [
        column for column in df.columns
        if not pd.api.types.is_bool_dtype(df[column])
        and pd.to_numeric(df[column], errors='coerce').notna().any()
    ]
    categorical_columns = [column for column in df.columns.tolist() if column not in numeric_columns]
    return DatasetSummary(
        rows=len(df.index),
        columns=df.columns.tolist(),
        numeric_columns=numeric_columns,
        categorical_columns=categorical_columns,
    )
