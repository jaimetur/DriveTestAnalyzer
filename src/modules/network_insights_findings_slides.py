"""The Key Findings of Network Insights as benchmark slides: a bold title, the headline under it in the accent of the
theme, a pill with the tests it covers, the charts drawn as native shapes (stacked class bars, spread marks, ranked
bars and bandwidth ladders), what it means in a box with an accent bar, and the basis and footnote.

Every colour comes from the active theme (see scoring_themes): its five-step scale for the classes, its series for the
bands and its accent for the headline. The Word version renders these same slides (see bundle_report_templates.compose_word).
"""
from __future__ import annotations

from collections.abc import Callable
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.modules.scoring_themes import active_theme, blend

LEFT, RIGHT = 0.52, 12.81
WIDTH = RIGHT - LEFT
CONTENT_TOP = 1.86
# The headline wraps before the pill under the logos.
TITLE_WIDTH = 9.95
PILL_TOP, PILL_HEIGHT = 1.13, 0.3
FOOTNOTE_BOTTOM = 6.98
# Groups per slide: the signal views list every group in each environment.
SIGNAL_GROUPS_PER_SLIDE = 5
GROUPS_PER_SLIDE = 8


def _hex(value: Any, fallback: str = '#888888') -> str:
    text = str(value or '').strip().lstrip('#')
    return f'#{text.upper()}' if len(text) == 6 else fallback


def _rgb(value: Any) -> RGBColor:
    return RGBColor.from_string(_hex(value).lstrip('#'))


def _luminance(value: Any) -> float:
    text = _hex(value).lstrip('#')
    red, green, blue = (int(text[index:index + 2], 16) / 255 for index in (0, 2, 4))
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


def palette(theme: dict[str, Any] | None = None) -> dict[str, Any]:
    """The colours of the slides in a theme: dark text on the white slide (the title colour of a dark theme), its
    accent, muted text, soft panels, its scale and its series."""
    theme = theme or active_theme()
    light = _luminance(theme.get('surface', '#FFFFFF')) > .5
    ink = _hex(theme.get('ink') if light else theme.get('title'), '#1F2A37')
    soft = _hex(theme.get('soft_header'), '#EDF2F6')
    if _luminance(soft) < .8:
        soft = blend(soft, '#FFFFFF', .85)
    # A white panel stands apart from the slide in the shade of the groups of the theme.
    if _luminance(soft) > .99:
        soft = _hex(theme.get('group'), '#F5F5F5')
    return {
        'ink': ink, 'muted': _hex(theme.get('muted'), '#8A96A0'), 'accent': _hex(theme.get('title_accent'), '#1450A8'),
        'soft': soft, 'line': _hex(theme.get('line'), '#E3E8EC'),
        'scale': [_hex(colour) for colour in theme.get('scale') or ('#C9483F', '#E8A060', '#F2D06B', '#93C47D', '#3E8E4F')],
        'series': [_hex(colour) for colour in theme.get('series') or ('#5B7A99', '#A9C4DB', '#8BA88E', '#C9B27C', '#B48C7C', '#8E87A0', '#DADDE0')],
        'track': blend(_hex(theme.get('muted'), '#8A96A0'), '#FFFFFF', .45),
    }


def ladder_colours(colours: dict[str, Any], count: int) -> list[str]:
    """``count`` colours from the best end of the scale to its worst (the widest bandwidth class first)."""
    scale = list(reversed(colours['scale']))
    result = []
    for index in range(count):
        position = index / max(1, count - 1) * (len(scale) - 1)
        low = int(position)
        high = min(len(scale) - 1, low + 1)
        result.append(blend(scale[low], scale[high], position - low))
    return result


def rank_colour(colours: dict[str, Any], rank: float | None) -> str:
    """The colour of a rank from 0 (the best of the groups) to 1 (the worst)."""
    scale = colours['scale']
    if rank is None:
        return colours['track']
    return scale[round((1 - rank) * (len(scale) - 1))]


def _on(fill: str, colours: dict[str, Any]) -> str:
    return '#FFFFFF' if _luminance(fill) < .6 else colours['ink']


def _box(slide, left: float, top: float, width: float, height: float, fill: str, shape=MSO_SHAPE.RECTANGLE,
         line: str | None = None, radius: float | None = None):
    item = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(max(width, 0.001)), Inches(max(height, 0.001)))
    item.fill.solid()
    item.fill.fore_color.rgb = _rgb(fill)
    if line:
        item.line.color.rgb = _rgb(line)
        item.line.width = Pt(1)
    else:
        item.line.fill.background()
    if radius is not None and item.adjustments:
        item.adjustments[0] = radius
    # Without the style of the theme, which draws a shadow under every shape in LibreOffice (the Word version).
    style = item._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        item._element.remove(style)
    item.text_frame.margin_left = item.text_frame.margin_right = 0
    item.text_frame.margin_top = item.text_frame.margin_bottom = 0
    return item


def _write(frame, parts: list[tuple[str, dict[str, Any]]], size: float, colour: str, align=PP_ALIGN.LEFT,
           anchor=MSO_ANCHOR.TOP, bold: bool = False) -> None:
    """Runs of text, each ``(text, {'bold', 'colour', 'size'})``; a text starting a new paragraph with ``\\n``."""
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for text, style in parts:
        if text.startswith('\n'):
            paragraph = frame.add_paragraph()
            paragraph.alignment = align
            paragraph.space_before = Pt(style.get('space', 4))
            text = text[1:]
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(style.get('size', size))
        run.font.bold = style.get('bold', bold)
        run.font.color.rgb = _rgb(style.get('colour', colour))


def _text(slide, left: float, top: float, width: float, height: float, text: str | list[tuple[str, dict[str, Any]]],
          size: float, colour: str, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bold: bool = False, wrap: bool = True):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(max(width, 0.01)), Inches(max(height, 0.01)))
    frame = box.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    _write(frame, [(text, {})] if isinstance(text, str) else text, size, colour, align, anchor, bold)
    frame.word_wrap = wrap
    return box


def _header(slide, view: dict[str, Any], colours: dict[str, Any]) -> None:
    """The headline in the accent of the theme under the title, and the pill of the tests at the right."""
    from src.modules.slide_titles import fit_paragraph_on_one_line
    from src.modules.theme_titles import style_title, title_style

    title = next((shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3} and shape.has_text_frame), None)
    if title is not None:
        title.left, title.top, title.width, title.height = title.left, title.top, Inches(TITLE_WIDTH), title.height
        # The title and the headline under it in the style of the theme, the headline on one line when it fits.
        style_title(slide)
        paragraphs = title.text_frame.paragraphs
        if len(paragraphs) > 1:
            size = title_style().get('subtitle_size', 16)
            fit_paragraph_on_one_line(title, paragraphs[1], default_size=size, minimum=min(size, 13))
    tag = str(view.get('tag') or '')
    if tag:
        width = 0.5 + len(tag) * 0.075
        pill = _box(slide, RIGHT - width, PILL_TOP, width, PILL_HEIGHT, '#FFFFFF', MSO_SHAPE.ROUNDED_RECTANGLE,
                    line=colours['ink'], radius=0.5)
        pill.name = 'Key Finding tag'
        _write(pill.text_frame, [(tag, {})], 9, colours['ink'], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, bold=True)


def _footer(slide, view: dict[str, Any], scope: str, colours: dict[str, Any]) -> float:
    """Basis, what it means and the footnote at the bottom; returns where the content above them ends."""
    footnote = ' '.join(part for part in (str(view.get('footnote') or ''), scope) if part)
    bottom = FOOTNOTE_BOTTOM
    if footnote:
        _text(slide, LEFT, bottom - 0.24, WIDTH, 0.22, footnote, 8, colours['muted'])
        bottom -= 0.3
    meaning = str(view.get('meaning') or '')
    if meaning:
        # About 150 characters fit one line at 11pt across the box.
        height = 0.4 if len(meaning) <= 140 else 0.6
        _box(slide, LEFT, bottom - height, WIDTH, height, colours['soft'])
        _box(slide, LEFT, bottom - height, 0.06, height, colours['accent'])
        _text(slide, LEFT + 0.2, bottom - height, WIDTH - 0.3, height,
              [('What it means: ', {'bold': True, 'colour': colours['accent']}), (meaning, {})], 11, colours['ink'],
              anchor=MSO_ANCHOR.MIDDLE)
        bottom -= height + 0.12
    basis = str(view.get('basis') or '')
    if basis:
        _text(slide, LEFT, bottom - 0.24, WIDTH, 0.24, [('Basis: ', {'bold': True, 'colour': colours['ink']}), (basis, {})],
              9, colours['muted'])
        bottom -= 0.3
    return bottom


def _legend(slide, labels: list[str], fills: list[str], left: float, top: float, width: float, colours: dict[str, Any],
            size: float = 8.5, prefix: str = '') -> None:
    x = left
    if prefix:
        _text(slide, x, top, 1.0, 0.2, prefix, size, colours['ink'], anchor=MSO_ANCHOR.MIDDLE)
        x += 0.95
    step = (left + width - x) / max(1, len(labels))
    for label, fill in zip(labels, fills):
        _box(slide, x, top + 0.04, 0.13, 0.13, fill)
        _text(slide, x + 0.19, top, step - 0.22, 0.2, label, size, colours['muted'], anchor=MSO_ANCHOR.MIDDLE, wrap=False)
        x += step


def _stacked_bar(slide, left: float, top: float, width: float, height: float, shares: list[float], fills: list[str],
                 colours: dict[str, Any], size: float) -> None:
    x = left
    total = sum(shares) or 100
    for share, fill in zip(shares, fills):
        span = width * share / total
        if span <= 0.004:
            continue
        segment = _box(slide, x, top, span, height, fill)
        if share >= 4.5 and span >= 0.2:
            _write(segment.text_frame, [(f'{share:.0f}', {})], size, _on(fill, colours), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        x += span


def _row_label(slide, text: str, left: float, top: float, width: float, height: float, colours: dict[str, Any],
               size: float, colour: str | None = None, bold: bool = False) -> None:
    _text(slide, left, top, width, height, text, size, colour or colours['ink'], PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, bold, wrap=False)


def _signal_panel(slide, panel: dict[str, Any], groups: list[str], left: float, top: float, width: float, bottom: float,
                  colours: dict[str, Any]) -> None:
    _text(slide, left, top, width, 0.26, panel['title'], 11, colours['ink'], bold=True)
    legend_top = bottom - 0.24
    _legend(slide, panel['classes'], colours['scale'], left, legend_top, width, colours)
    blocks = [{**block, 'rows': [row for row in block['rows'] if row['group'] in groups]} for block in panel['environments']]
    blocks = [block for block in blocks if block['rows']]
    units = sum(1 + len(block['rows']) for block in blocks) + 0.4 * max(0, len(blocks) - 1)
    area_top, area_bottom = top + 0.34, legend_top - 0.08
    unit = min(0.27, (area_bottom - area_top) / max(1, units))
    bar_left, label_width = left + 1.2, 1.1
    bar_width = width - 1.25
    size = 8 if unit >= 0.2 else 7
    y = area_top
    for block in blocks:
        _text(slide, bar_left, y, bar_width, unit, block['environment'], 9.5, colours['ink'], anchor=MSO_ANCHOR.BOTTOM, bold=True)
        y += unit
        for row in block['rows']:
            _row_label(slide, row['group'], left, y, label_width, unit, colours, size + 1)
            _stacked_bar(slide, bar_left, y + unit * .1, bar_width, unit * .8, row['shares'], colours['scale'], colours, size)
            y += unit
        y += unit * 0.4


def _card(slide, left: float, top: float, width: float, height: float, colours: dict[str, Any]):
    return _box(slide, left, top, width, height, colours['soft'], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)


def _spread_panels(slide, view: dict[str, Any], groups: list[str], colours: dict[str, Any], group_colours: dict[str, str],
                   bottom: float) -> None:
    spread, narrow = view['panels']
    rows = [row for row in spread['rows'] if row['group'] in groups]
    shares = {row['group']: row for row in narrow['rows']}
    gap = 0.3
    left_width = WIDTH * 0.6 - gap / 2
    right_left = LEFT + left_width + gap
    right_width = RIGHT - right_left
    top = CONTENT_TOP
    height = bottom - top - 0.12
    for x, w, panel in ((LEFT, left_width, spread), (right_left, right_width, narrow)):
        _card(slide, x, top, w, height, colours)
        _text(slide, x + 0.22, top + 0.16, w - 0.4, 0.26, panel['title'], 11, colours['ink'], bold=True)
        if panel.get('note'):
            _text(slide, x + 0.22, top + height - 0.36, w - 0.4, 0.24, panel['note'], 8, colours['muted'])
    area_top, area_bottom = top + 0.62, top + height - 0.45
    step = min(0.72, (area_bottom - area_top) / max(1, len(rows)))
    maximum = max((row['p90'] or 0 for row in rows), default=1) * 1.06 or 1
    track_left, track_right = LEFT + 2.0, LEFT + left_width - 0.75
    narrow_max = max((shares.get(row['group'], {}).get('share') or 0 for row in rows), default=1) or 1
    bar_left, bar_span = right_left + 1.55, right_width - 2.45
    for index, row in enumerate(rows):
        centre = area_top + step * (index + 0.62)
        colour = group_colours.get(row['group']) or colours['ink']
        for x in (LEFT + 0.22, right_left + 0.22):
            _text(slide, x, centre - 0.15, 1.5, 0.3, row['group'], 10.5, colour, anchor=MSO_ANCHOR.MIDDLE, bold=True, wrap=False)
        position = lambda value: track_left + (track_right - track_left) * (value or 0) / maximum
        if row['p10'] is not None and row['p90'] is not None:
            _box(slide, position(row['p10']), centre - 0.04, max(0.02, position(row['p90']) - position(row['p10'])), 0.08, colours['track'])
            _text(slide, position(row['p10']) - 0.62, centre - 0.12, 0.55, 0.24, f"{row['p10']:g}", 8.5, colours['muted'],
                  PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
            _text(slide, position(row['p90']) + 0.07, centre - 0.12, 0.6, 0.24, f"{row['p90']:g}", 8.5, colours['muted'],
                  anchor=MSO_ANCHOR.MIDDLE)
        if row['median'] is not None:
            _box(slide, position(row['median']) - 0.04, centre - 0.17, 0.08, 0.34, colours['ink'])
            _text(slide, position(row['median']) - 0.4, centre - 0.42, 0.8, 0.22, f"{row['median']:g}", 9.5, colours['ink'],
                  PP_ALIGN.CENTER, MSO_ANCHOR.BOTTOM, bold=True)
        figures = shares.get(row['group'])
        if figures and figures.get('share') is not None:
            fill = rank_colour(colours, figures.get('rank'))
            span = max(0.03, bar_span * figures['share'] / narrow_max)
            _box(slide, bar_left, centre - 0.13, span, 0.26, fill)
            _text(slide, bar_left + span + 0.08, centre - 0.13, 0.8, 0.26, f"{figures['share']:.1f}%", 9.5, colours['ink'],
                  anchor=MSO_ANCHOR.MIDDLE, bold=True)


def _side_card(slide, card: dict[str, Any] | None, left: float, top: float, width: float, colours: dict[str, Any]) -> None:
    if not card or not card.get('lines'):
        return
    parts: list[tuple[str, dict[str, Any]]] = [(card['title'], {'bold': True, 'size': 12})]
    for line in card['lines']:
        parts.append((f"\n{line['lead']} ", {'bold': True, 'space': 9}))
        parts.append((line['text'], {}))
    height = 0.75 + 0.48 * len(card['lines'])
    _card(slide, left, top, width, height, colours)
    _text(slide, left + 0.2, top + 0.18, width - 0.4, height - 0.3, parts, 10, colours['ink'])


def _bar_panels(slide, view: dict[str, Any], groups: list[str], fills: list[str], colours: dict[str, Any], bottom: float,
                side: dict[str, Any] | None) -> None:
    """The 100% bars of the technology and bands views: one panel per test mode and, at the right, their card."""
    panels = view['panels']
    side_width = 3.35 if side and side.get('lines') else 0
    _legend(slide, view['classes'], fills, LEFT, CONTENT_TOP - 0.02, (WIDTH - side_width - 0.3) if side_width else WIDTH * .75,
            colours)
    top = CONTENT_TOP + 0.32
    available = WIDTH - (side_width + 0.35 if side_width else 0)
    gap = 0.35
    width = (available - gap * (len(panels) - 1)) / max(1, len(panels))
    for index, panel in enumerate(panels):
        left = LEFT + index * (width + gap)
        _text(slide, left, top, width, 0.26, panel['title'], 11, colours['ink'], bold=True)
        rows = [row for row in panel['rows'] if row['group'] in groups]
        area_top, area_bottom = top + 0.4, bottom - 0.15
        step = min(0.82, (area_bottom - area_top) / max(1, len(rows)))
        axis = left + 1.2
        _box(slide, axis - 0.01, area_top, 0.01, step * len(rows), colours['line'])
        for position, row in enumerate(rows):
            y = area_top + step * position
            _row_label(slide, row['group'], left, y, 1.08, step, colours, 9.5)
            _stacked_bar(slide, axis, y + step * .18, width - 1.25, step * .64, row['shares'], fills, colours, 9)
    if side_width:
        _side_card(slide, side, RIGHT - side_width, CONTENT_TOP + 0.05, side_width, colours)


def _ladder(slide, view: dict[str, Any], groups: list[str], colours: dict[str, Any], group_colours: dict[str, str],
            bottom: float) -> None:
    columns = [column for column in view['columns'] if column['group'] in groups]
    classes = view['classes']
    fills = ladder_colours(colours, len(classes))
    legend_top = bottom - 0.3
    chart_top, chart_bottom = CONTENT_TOP + 0.1, legend_top - 0.62
    chart_left = LEFT + 0.55
    height = chart_bottom - chart_top
    for share in (0, 50, 100):
        y = chart_bottom - height * share / 100
        _box(slide, chart_left, y, RIGHT - chart_left, 0.008, colours['line'])
        _text(slide, LEFT, y - 0.1, 0.48, 0.2, f'{share}%', 8, colours['muted'], PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    slot = (RIGHT - chart_left) / max(1, len(columns))
    widest = max((item['width'] for item in classes), default=1) or 1
    leader = max(columns, key=lambda column: column.get('wide_share') or 0) if columns else None
    for index, column in enumerate(columns):
        centre = chart_left + slot * (index + 0.5)
        y = chart_bottom
        for share, item, fill in zip(column['shares'], classes, fills):
            span = height * share / 100
            if span <= 0.003:
                continue
            width = slot * 0.82 * item['width'] / widest
            _box(slide, centre - width / 2, y - span, width, span, fill)
            y -= span
        colour = group_colours.get(column['group']) or colours['ink']
        _text(slide, centre - slot / 2, chart_bottom + 0.12, slot, 0.3,
              [(f"{column['group']}  ", {'bold': True, 'colour': colour}),
               (f"{column['wide_share']:g}% at 100 MHz+" if column.get('wide_share') is not None else '',
                {'bold': column is leader})], 10.5, colours['ink'], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    _legend(slide, [item['label'] for item in classes], fills, LEFT, legend_top, WIDTH * 0.72, colours, prefix='Primary bandwidth, MHz:')
    if view.get('note'):
        _text(slide, LEFT + WIDTH * 0.74, legend_top, WIDTH * 0.26, 0.2, view['note'], 8.5, colours['muted'], anchor=MSO_ANCHOR.MIDDLE)


def _pages(groups: list[str], per_slide: int) -> list[list[str]]:
    """The groups split evenly over as few slides as hold ``per_slide`` each."""
    if len(groups) <= per_slide:
        return [groups]
    count = -(-len(groups) // per_slide)
    size = -(-len(groups) // count)
    return [groups[index:index + size] for index in range(0, len(groups), size)]


def add_finding_slides(presentation, view: dict[str, Any], findings: dict[str, Any], scope: str,
                       new_slide: Callable[[Any, str, str], Any]) -> None:
    """The slides of one Key Findings view: ``new_slide(presentation, title, subtitle)`` adds a slide of the module."""
    colours = palette()
    group_colours = {group: _hex(colour, '') for group, colour in (findings.get('colours') or {}).items() if colour}
    pages = _pages(list(findings.get('groups') or []), SIGNAL_GROUPS_PER_SLIDE if view['kind'] == 'signal' else GROUPS_PER_SLIDE)
    for number, groups in enumerate(pages, start=1):
        title = view['title'] + (f' · {number}/{len(pages)}' if len(pages) > 1 else '')
        slide = new_slide(presentation, title, view.get('headline') or '')
        _header(slide, view, colours)
        bottom = _footer(slide, view, scope, colours)
        if view['kind'] == 'signal':
            gap = 0.4
            width = (WIDTH - gap) / 2
            for index, panel in enumerate(view['panels']):
                _signal_panel(slide, panel, groups, LEFT + index * (width + gap), CONTENT_TOP, width, bottom - 0.05, colours)
        elif view['kind'] == 'spread':
            _spread_panels(slide, view, groups, colours, group_colours, bottom)
        elif view['kind'] == 'technology':
            _bar_panels(slide, view, groups, [colours['series'][5], colours['scale'][0]], colours, bottom, view.get('reading'))
        elif view['kind'] == 'bands':
            _bar_panels(slide, view, groups, colours['series'][:len(view['classes'])], colours, bottom, view.get('side'))
        elif view['kind'] == 'ladder':
            _ladder(slide, view, groups, colours, group_colours, bottom)


BAND_ROWS_PER_SLIDE = 12


def add_band_share_slides(presentation, observed: list[dict[str, Any]], group_colours: dict[str, str], scope: str,
                          new_slide: Callable[[Any, str, str], Any], title: str) -> None:
    """The spectrum observed as bars: per group and technology, the share of its samples on each band, each band in a
    colour of the series of the theme, LTE and NR side by side."""
    colours = palette()
    technologies = [technology for technology in ('LTE', 'NR') if any(row.get('technology') == technology for row in observed)]
    if not technologies:
        return
    groups = list(dict.fromkeys(str(row.get('operator') or '') for row in observed))
    pages = _pages(groups, BAND_ROWS_PER_SLIDE)
    for number, page in enumerate(pages, start=1):
        slide = new_slide(presentation, title + (f' · {number}/{len(pages)}' if len(pages) > 1 else ''), scope)
        gap = 0.4
        width = (WIDTH - gap * (len(technologies) - 1)) / len(technologies)
        for index, technology in enumerate(technologies):
            left = LEFT + index * (width + gap)
            rows = [row for row in observed if row.get('technology') == technology]
            bands = list(dict.fromkeys(str(row.get('band') or '') for row in rows))
            series = colours['series']
            fills = {band: series[position] if position < len(series) else blend(series[position % len(series)], '#FFFFFF', .45)
                     for position, band in enumerate(bands)}
            _text(slide, left, CONTENT_TOP, width, 0.28, f'{technology}: serving band, % of the samples', 11, colours['ink'], bold=True)
            legend_top = 6.62
            labels = [f"{band} · {row['frequency_mhz']:g} MHz" if (row := next((item for item in rows if item.get('band') == band), {})).get('frequency_mhz')
                      else band for band in bands]
            _legend(slide, labels, [fills[band] for band in bands], left, legend_top, width, colours, size=8)
            area_top, area_bottom = CONTENT_TOP + 0.4, legend_top - 0.12
            step = min(0.42, (area_bottom - area_top) / max(1, len(page)))
            for position, group in enumerate(page):
                y = area_top + step * position
                _row_label(slide, group, left, y, 1.6, step, colours, 9, group_colours.get(group) or colours['ink'], bold=True)
                x, bar_width = left + 1.7, width - 1.75
                for row in [item for item in rows if item.get('operator') == group]:
                    span = bar_width * float(row.get('share') or 0) / 100
                    if span <= 0.004:
                        continue
                    segment = _box(slide, x, y + step * .16, span, step * .68, fills[str(row.get('band') or '')])
                    if span >= 0.42:
                        _write(segment.text_frame, [(str(row.get('band') or ''), {})], 8, _on(fills[str(row.get('band') or '')], colours),
                               PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, bold=True)
                    x += span
