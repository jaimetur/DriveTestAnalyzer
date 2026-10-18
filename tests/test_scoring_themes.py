"""Colour themes of Scoring & GAP Analysis: the NetCheck colours translated into each theme, the exports and the
theme each user chooses."""
from __future__ import annotations

from io import BytesIO
import json

from pptx import Presentation

import src.DriveTestAnalyzer as core
from src.modules.scoring_exports import export_scoring_powerpoint
from src.modules.scoring_pptx_scorecard import STYLES
from src.modules.scoring_reports import normalize_report_configuration
from src.modules.scoring_themes import DEFAULT_THEME, THEMES, blend, normalize_theme, themed, themes_payload, use_theme
from src.modules.scoring_views import THRESHOLD_COLORS, gap_color
from tests.scoring_synthetic import synthetic_result
from tests.test_scoring_exports import TEMPLATE


def _channels(color: str) -> tuple[int, ...]:
    return tuple(int(color.lstrip('#')[index:index + 2], 16) for index in (0, 2, 4))


def test_every_theme_has_every_token():
    tokens = set(THEMES[DEFAULT_THEME])
    assert DEFAULT_THEME == 'netcheck' and list(THEMES) == ['netcheck', 'umlaut', 'ericsson', 'vodafone', 'midnight', 'paper', 'heatmap']
    for name, theme in THEMES.items():
        assert set(theme) == tokens, name
        assert set(theme['bands']) == set(THRESHOLD_COLORS), name
        for token in ('ranks', 'stake', 'loss_ramp', 'environments', 'categories'):
            assert len(theme[token]) == len(THEMES[DEFAULT_THEME][token]), (name, token)
    payload = themes_payload()
    assert payload['default'] == 'netcheck' and [theme['id'] for theme in payload['themes']] == list(THEMES)


def test_the_rings_of_the_maximum_score_donuts_have_colours_of_their_own():
    # The outer ring (environments) never repeats a colour of the inner one (Voice and Data, or the KPI categories).
    for name, theme in THEMES.items():
        assert set(theme['environments'][:2]).isdisjoint(theme['categories']), name
        if name != DEFAULT_THEME:
            assert set(theme['environments']).isdisjoint(theme['categories']), name


def test_netcheck_colours_translate_into_the_active_theme():
    # NetCheck is the default: its colours stay as they are.
    assert themed('#1450A8') == '#1450A8' and themed(THRESHOLD_COLORS['Low']) == THRESHOLD_COLORS['Low']
    with use_theme('Ericsson') as theme:
        assert theme['id'] == 'ericsson'
        assert themed('#1450A8') == theme['title']
        # The exports write some colours without the #; the translation keeps their form.
        assert themed('5B9A44') == theme['weight'].lstrip('#')
        assert themed(THRESHOLD_COLORS['Low']) == theme['bands']['Low']
        assert themed('#E8414F') == theme['loss'] and themed('#0E6B66') == theme['data']
        # Every shade of the NetCheck GAP scale takes the same place in the scale of the theme.
        shade = themed(gap_color(-1, 2))
        expected = blend('#FFFFFF', theme['gap_negative'], (.5 ** .85) * .85)
        assert max(abs(a - b) for a, b in zip(_channels(shade), _channels(expected))) <= 3
        # Operator colours and the other colours are left as they are.
        assert themed('#E60000') == '#E60000' and themed('#17232D') == '#17232D'
    assert themed('#1450A8') == '#1450A8'
    assert normalize_theme('Colour blind') == 'netcheck' and normalize_theme(' Vodafone ') == 'vodafone'


def test_small_gaps_stay_almost_white():
    assert gap_color(0, 2) == '#FFFFFF'
    assert min(_channels(gap_color(.02, 2))) >= 240
    assert gap_color(2, 2) == blend('#FFFFFF', THEMES['netcheck']['gap_positive'], .85)
    assert gap_color(-2, 2) == blend('#FFFFFF', THEMES['netcheck']['gap_negative'], .85)


def test_exports_take_the_colours_of_their_theme():
    job, result = synthetic_result(levels=('Operator',))
    for name in ('netcheck', 'vodafone'):
        theme = THEMES[name]
        deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, theme=name)))
        summary = next(slide for slide in deck.slides if slide.shapes.title is not None
                       and slide.shapes.title.text_frame.text.startswith('Scoring Tables — Summary'))
        runs = summary.shapes.title.text_frame.paragraphs[0].runs
        assert str(runs[0].font.color.rgb) == theme['title'].lstrip('#')
        assert str(runs[-1].font.color.rgb) == theme['title_accent'].lstrip('#') and runs[-1].font.bold
        # The Summary table in the scorecard design of the theme: light header, operators underlined, its total.
        table = next(shape.table for shape in summary.shapes if shape.has_table and shape.table.cell(0, 0).text == 'Category')
        design = STYLES[name]
        assert str(table.cell(0, 0).fill.fore_color.rgb) == design['header'].lstrip('#')
        total = table.cell(len(table.rows) - 1, 0)
        assert total.text.startswith('Weighted score') and str(total.fill.fore_color.rgb) == design['total'].lstrip('#')
    # After the export the default theme is active again.
    assert themed('#1450A8') == '#1450A8'


def test_report_configurations_keep_their_theme():
    assert normalize_report_configuration({'scenarios': [{}], 'theme': 'Paper'})['theme'] == 'paper'
    # Themes of earlier versions that are gone fall back to NetCheck.
    assert normalize_report_configuration({'scenarios': [{}], 'theme': 'ocean'})['theme'] == 'netcheck'
    assert normalize_report_configuration({'scenarios': [{}]})['theme'] == 'netcheck'
    assert normalize_report_configuration({'scenarios': [{}], 'theme': 'unknown'})['theme'] == 'netcheck'


def test_each_user_keeps_a_theme_and_the_page_offers_every_theme(client):
    client.cookies.clear()
    assert client.post('/login', data={'username': 'super', 'password': 'super123'}, follow_redirects=False).status_code == 303
    page = client.get('/scoring').text
    assert 'data-scoring-theme-select' in page and 'id="scoring-themes-data"' in page
    assert 'data-scoring-theme="netcheck"' in page
    for theme in THEMES.values():
        assert f'>{theme["label"]}</option>' in page
    assert client.get('/api/scoring/theme').json()['theme'] == 'netcheck'
    assert client.put('/api/scoring/theme', json={'theme': 'vodafone'}).json() == {'theme': 'vodafone'}
    assert client.get('/api/scoring/theme').json()['theme'] == 'vodafone'
    assert 'data-scoring-theme="vodafone"' in client.get('/scoring').text
    assert client.put('/api/scoring/theme', json={'theme': 'sunset'}).status_code == 400
    changes = [json.loads(row['details']) for row in core.repository.list_logs() if row['action'] == 'scoring_theme_changed']
    assert changes[0] == {'theme': 'vodafone', 'previous': 'netcheck'}


def test_each_theme_styles_the_titles_and_subtitles_of_the_slides():
    from src.modules.cdr_reporting import _remove_all_slides
    from src.modules.network_insights_export import _slide
    from src.modules.theme_titles import split_focus, title_style

    assert split_focus('Scoring Tables — Summary') == ('Scoring Tables — ', 'Summary')
    assert split_focus('Best Network Scoring per Category') == ('Best Network Scoring ', 'per Category')
    assert split_focus('Overview') == ('Overview', '')
    for name in THEMES:
        assert set(THEMES[name]['titles']) == set(THEMES['netcheck']['titles']), name
    deck = Presentation(TEMPLATE)
    _remove_all_slides(deck)
    with use_theme('netcheck'):
        slide = _slide(deck, 'RF Quality · NR quality per carrier', 'National · All Operators')
        heading, focus = slide.shapes.title.text_frame.paragraphs[0].runs
        # NetCheck: the heading in capitals (without changing its text), the focus in bold, both blue, in Montserrat.
        assert heading.text == 'RF Quality · ' and heading.font._rPr.get('cap') == 'all' and not heading.font.bold
        assert focus.font.bold and focus.font._rPr.get('cap') is None and focus.font.name == 'Montserrat'
        assert str(focus.font.color.rgb) == THEMES['netcheck']['title_accent'].lstrip('#')
        lead = slide.shapes.title.text_frame.paragraphs[1].runs[0]
        assert lead.text == 'National · ' and lead.font.bold
    with use_theme('umlaut'):
        slide = _slide(deck, 'Overview', 'National')
        title, subtitle = slide.shapes.title.text_frame.paragraphs[:2]
        assert title.runs[0].font.bold and title.runs[0].font.name == 'Arial'
        assert str(title.runs[0].font.color.rgb) == '1C1C28' and str(subtitle.runs[0].font.color.rgb) == 'A6192E'
        assert title_style()['subtitle_color'] == '#A6192E'
    with use_theme('paper'):
        slide = _slide(deck, 'Spectrum · Deployed', 'National')
        assert slide.shapes.title.text_frame.paragraphs[0].runs[-1].font.italic
        assert slide.shapes.title.text_frame.paragraphs[1].runs[0].font.name == 'Georgia'


def test_light_scoring_slides_leave_out_the_mark_of_the_red_layouts():
    from src.modules.cdr_reporting import _remove_all_slides
    from src.modules.scoring_exports import _slide

    deck = Presentation(TEMPLATE)
    _remove_all_slides(deck)
    slide = _slide(deck, 'Scoring Tables — Summary', 'Best Network')
    names = [shape.name for shape in slide.shapes]
    # The dark logos of the master, without the Ericsson mark drawn on the red of the covers and transitions.
    assert 'Logo' in names and not any(name.startswith('Ericsson watermark') for name in names)


def test_ericsson_titles_are_light_with_the_part_that_changes_in_bold():
    from src.modules.cdr_reporting import _remove_all_slides
    from src.modules.network_insights_export import _slide
    from src.modules.pptx_compose import set_title

    deck = Presentation(TEMPLATE)
    _remove_all_slides(deck)
    with use_theme('ericsson'):
        slide = _slide(deck, 'RF Quality · NR quality per carrier', 'National · All Operators')
        heading, focus = slide.shapes.title.text_frame.paragraphs[0].runs
        # Ericsson Hilda, black: only the part that changes from one slide to the next in bold.
        assert heading.font.name == 'Ericsson Hilda' and heading.font.bold is False
        assert focus.font.name == 'Ericsson Hilda' and focus.font.bold
        assert str(heading.font.color.rgb) == str(focus.font.color.rgb) == '0C0C0C'
        only = _slide(deck, 'Overview', 'National').shapes.title.text_frame.paragraphs[0].runs
        assert [(run.text, run.font.bold) for run in only] == [('Overview', False)]
    # A component on several slides of a composed report keeps the light heading and its focus in bold.
    set_title(slide, 'RF Quality · NR quality per carrier · 1/2', 'National · All Operators')
    heading, focus = slide.shapes.title.text_frame.paragraphs[0].runs
    assert (heading.text, heading.font.bold, heading.font.name) == ('RF Quality · ', False, 'Ericsson Hilda')
    assert (focus.text, focus.font.bold, focus.font.name) == ('NR quality per carrier · 1/2', True, 'Ericsson Hilda')


def test_a_title_styled_twice_keeps_its_heading_and_focus():
    from src.modules.cdr_reporting import _remove_all_slides
    from src.modules.network_insights_export import _slide
    from src.modules.pptx_compose import set_title
    from src.modules.theme_titles import style_title

    deck = Presentation(TEMPLATE)
    _remove_all_slides(deck)
    for name in ('vodafone', 'heatmap'):
        with use_theme(name):
            slide = _slide(deck, 'LTE signal at the test: RSRP and SINR', 'National')
            style_title(slide)
            heading, focus = slide.shapes.title.text_frame.paragraphs[0].runs
            # The heading in the title colour and only the focus in the accent, also styled again.
            assert str(heading.font.color.rgb) == THEMES[name]['title'].lstrip('#')
            assert str(focus.font.color.rgb) == THEMES[name]['title_accent'].lstrip('#')
            set_title(slide, 'LTE signal at the test: RSRP and SINR · 1/2', 'National')
            heading, focus = slide.shapes.title.text_frame.paragraphs[0].runs
            assert str(heading.font.color.rgb) == THEMES[name]['title'].lstrip('#') and focus.text == 'RSRP and SINR · 1/2'
