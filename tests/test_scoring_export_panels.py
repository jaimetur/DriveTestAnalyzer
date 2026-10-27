"""The Scoring Export Panels of Workspace Config: the panels the Scoring PowerPoint and Word documents include and the
operators each one exports (every operator or the Focus operators), and Export All Operators in Focus panels."""
from __future__ import annotations

import json
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import src.DriveTestAnalyzer as app_module
from src.modules.repository import Repository
from src.modules.scoring_export_panels import (
    EXPORT_PANELS, default_export_panels, export_panel_catalogue, export_panel_changes, exportable_options,
    load_export_panels, normalize_export_panels,
)
from src.modules.scoring_exports import _ranking_groups, export_scoring_powerpoint
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_exports import TEMPLATE
from tests.test_scoring_where_to_act import owned  # noqa: F401 (fixture)

FOCUS = [{'canonical': 'Three UK', 'aliases': [], 'position': 0, 'color': '#000000', 'role': 'focus'}]


def _deck(job, result, **keys):
    # All Environments carries every panel (each environment repeats them).
    return Presentation(BytesIO(export_scoring_powerpoint({**job, **keys}, result, TEMPLATE, FOCUS, environments=['Combined'])))


def _titles(deck) -> list[str]:
    return [slide.shapes.title.text_frame.paragraphs[0].text for slide in deck.slides if slide.shapes.title is not None]


def test_the_panels_of_each_tab_and_the_operators_they_export_by_default():
    panels = default_export_panels()
    assert panels['summary']['act'] == {'exportable': True, 'operators': 'focus'}
    assert panels['charts']['service'] == {'exportable': True, 'operators': 'all'}
    assert panels['tables']['campaign_comparison']['operators'] == panels['points']['map']['operators'] == 'focus'
    # Unknown values keep the default; settings saved for other panels are ignored.
    normalized = normalize_export_panels({'panels': {'summary': {'act': {'exportable': False, 'operators': 'everyone'}},
                                                     'unknown': {'panel': {'exportable': False}}}})
    assert normalized['summary']['act'] == {'exportable': False, 'operators': 'focus'} and 'unknown' not in normalized
    assert export_panel_changes(panels, normalized) == ['Summary > Where to act: not exportable']
    # The tabs and panels of the Scoring section of the Bundle Report Templates.
    from src.modules.bundle_report_templates import SCORING_SUBMODULES
    assert [(group['key'], group['title'], [panel['key'] for panel in group['panels']]) for group in export_panel_catalogue()] == [
        (key, title, [name for name, _label in components]) for key, title, components in SCORING_SUBMODULES]


def test_the_options_of_a_document_leave_out_the_panels_that_are_not_exportable():
    assert exportable_options(None, default_export_panels()) is None
    options = exportable_options(None, normalize_export_panels({'points': {'map': {'exportable': False}}}))
    assert options['points']['map'] is False and options['points']['sites'] is True
    # The other options of a report are kept, and the panels saved before in another group are read from it.
    saved = {'enabled': True, 'environments': 'split', 'tables': {'kpi_values': True}, 'gap': {'points_loss_map': False}}
    merged = exportable_options(saved, default_export_panels())
    assert merged['environments'] == 'split' and merged['tables']['kpi_values'] is True and merged['points']['map'] is False


def test_the_export_leaves_out_the_panels_that_are_not_exportable(owned):  # noqa: F811
    job, result, _insights = owned
    titles = _titles(_deck(job, result))
    assert any(': where the points are — ' in title for title in titles) and any(title.endswith('per Service') for title in titles)
    panels = normalize_export_panels({'summary': {'route': {'exportable': False}}, 'charts': {'service': {'exportable': False}}})
    titles = _titles(_deck(job, result, _export_panels=panels))
    assert not any(': where the points are — ' in title for title in titles)
    assert not any(title.endswith('per Service') for title in titles) and any(title.endswith('per Category') for title in titles)


def test_the_focus_operators_panels_show_every_operator_when_asked(owned):  # noqa: F811
    job, result, _insights = owned

    def act(deck):
        return sorted(title.split(' — ', 1)[1] for title in _titles(deck) if title.startswith('Where to act — '))
    assert act(_deck(job, result)) == ['Three UK']
    # Every operator (the reference too) with Export All Operators in Focus panels, or when the panel exports them all.
    every = act(_deck(job, result, _all_operators=True))
    assert every == ['EE', 'O2 UK', 'Three UK', 'Vodafone UK']
    assert act(_deck(job, result, _export_panels=normalize_export_panels({'summary': {'act': {'operators': 'all'}}}))) == every


def test_a_chart_of_the_focus_operators(owned):  # noqa: F811
    job, result, _insights = owned

    def service(deck):
        """The pages of the Scoring per Service chart and the texts of their slides (the operators of the chart are in the
        grid of its levels, in the group of the chart)."""
        def texts(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from texts(shape.shapes)
                elif getattr(shape, 'has_table', False) and shape.has_table:
                    yield from (cell.text for row in shape.table.rows for cell in row.cells)
                elif shape.has_text_frame:
                    yield shape.text_frame.text
        slides = [slide for slide in deck.slides if slide.shapes.title is not None
                  and slide.shapes.title.text_frame.paragraphs[0].text == 'Best Network Scoring per Service']
        return len(slides), ' '.join(text for slide in slides for text in texts(slide.shapes))
    everyone, everyone_text = service(_deck(job, result))
    focus, focus_text = service(_deck(job, result, _export_panels=normalize_export_panels({'charts': {'service': {'operators': 'focus'}}})))
    assert 0 < focus < everyone
    assert all(name in everyone_text for name in ('EE', 'O2 UK', 'Three UK', 'Vodafone UK'))
    assert 'Three UK' in focus_text and not any(name in focus_text for name in ('O2 UK', 'Vodafone UK'))


def test_the_points_lost_map_follows_the_operators_of_the_page(owned):  # noqa: F811
    job, result, _insights = owned

    def maps(deck):
        return {title.rsplit(' — ', 1)[1] for title in _titles(deck) if title.startswith('Points Lost per City — ')}
    assert maps(_deck(job, result)) == {'Three UK'}
    # The operators chosen on the page with those of the KPI GAP Profile.
    assert maps(_deck(job, result, _operator_choices={'gap-insight-operators': ['Vodafone UK']})) == {'Vodafone UK'}


def test_every_level_and_group_of_the_geographical_areas_is_exported():
    items = [{'level': 'Region', 'title': '', 'operator': 'A'}, {'level': 'City', 'title': 'Vendor: X', 'operator': 'A'},
             {'level': 'City', 'title': 'Vendor: Y', 'operator': 'A'}, {'level': 'City', 'title': 'Vendor: X', 'operator': 'B'},
             {'level': 'Cluster', 'title': '', 'operator': 'A'}]
    # From the largest level to the smallest (Region, Cluster, City), as the page lists them.
    assert [[(item['level'], item['title'], item['operator']) for item in group] for group in _ranking_groups(items)] == [
        [('Region', '', 'A')], [('Cluster', '', 'A')], [('City', 'Vendor: X', 'A'), ('City', 'Vendor: X', 'B')],
        [('City', 'Vendor: Y', 'A')]]
    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    assert "const GEOGRAPHICAL_LEVEL_ORDER = {National: 0, Region: 1, Cluster: 2, City: 3};" in script
    assert '.sort((left, right) => geographicalLevelRank(left) - geographicalLevelRank(right));' in script


def test_reports_and_bundle_report_templates_carry_export_all_operators():
    from src.modules.bundle_report_templates import _normalize_options, catalogue_payload
    from src.modules.scoring_reports import default_scenario, normalize_report_configuration

    assert normalize_report_configuration({'scenarios': [default_scenario()], 'all_operators': True})['all_operators'] is True
    # (Configurations saved before: no.)
    assert normalize_report_configuration({'scenarios': [default_scenario()]})['all_operators'] is False
    assert _normalize_options('scoring', {})['all_operators'] is False
    assert _normalize_options('scoring', {'all_operators': True})['all_operators'] is True
    # The Bundle editor lists the exportable panels with the operators they export.
    panels = normalize_export_panels({'summary': {'route': {'exportable': False}},
                                      'charts': {key: {'exportable': False} for key in ('service', 'category', 'breakdown')}})
    scoring = next(section for section in catalogue_payload([], panels)['sections'] if section['key'] == 'scoring')
    items = {item['key']: item for item in scoring['items']}
    summary = {component['key']: component for component in items['summary']['components']}
    assert 'charts' not in items and 'route' not in summary
    assert summary['act']['operators'] == 'focus' and summary['scope']['operators'] == 'all'


def test_workspace_config_saves_the_scoring_export_panels(client, tmp_path):
    response = client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=False)
    assert response.status_code == 303
    workspace = app_module.active_workspace
    repository = Repository(workspace.database_path, app_module.repository.global_db_path,
                            app_module.workspace_registry.registry_path)
    repository.initialize()
    repository.replace_scoring_configuration(scoring_configuration())
    page = client.get('/workspace-config').text
    # A card per results tab with an Exportable switch and the operators to export, the default marked.
    assert 'id="scoring-export-panels"' in page and '<span>What to Export</span>' in page
    assert 'id="scoring-export-group-points"' in page and 'role="switch" name="exportable:summary:act"' in page
    assert 'name="operators:summary:act" value="focus" checked' in page
    assert '<label class="scoring-export-mode is-focus" title="Focus Operators (default)">' in page
    form = {f'exportable:{group}:{key}': 'on' for group, _title, items in EXPORT_PANELS for key, _label, _mode in items
            if (group, key) != ('summary', 'route')}
    form.update({f'operators:{group}:{key}': mode for group, _title, items in EXPORT_PANELS for key, _label, mode in items})
    form.update({'operators:charts:service': 'focus', 'action': 'save'})
    response = client.post('/workspace-config/scoring-export-panels/save', data=form, follow_redirects=False)
    assert response.status_code == 303 and 'scoring_export_panels_notice' in response.headers['location']
    saved = load_export_panels(repository)
    assert saved['summary']['route']['exportable'] is False and saved['charts']['service']['operators'] == 'focus'
    assert client.get('/api/scoring/export-panels').json()['panels'] == saved
    # The change is in App Events.
    event = next(json.loads(row['details']) for row in app_module.repository.list_logs(limit=20)
                 if row['action'] == 'save_scoring_export_panels')
    assert sorted(event['changes']) == ['Scoring Charts > Scoring per Service: Focus Operators',
                                        'Summary > The route that loses most: where the points are: not exportable']
    # They travel as an element of their own: exported, imported and transferred, and in backups.
    assert 'scoring-export-panels' in app_module.WORKSPACE_ELEMENT_EXPORT_TARGETS
    assert app_module.archive_workspace_components_for_target('scoring-export-panels') == ['scoring_export_panels']
    assert 'scoring_export_panels' in app_module.full_workspace_archive_components()
    assert 'Scoring Export Panels (from active workspace)' in client.get('/admin').text
    exported = client.get('/admin/import-export/export?export_target=scoring-export-panels')
    assert exported.status_code == 200, exported.text
    with zipfile.ZipFile(BytesIO(exported.content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        document = json.loads(archive.read(manifest['archive_path']))
    assert manifest['kind'] == 'scoring-export-panels'
    assert document == {'format': 'drivetest-analyzer-scoring-export-panels', 'version': 1, 'panels': saved}

    def reset() -> None:
        response = client.post('/workspace-config/scoring-export-panels/save', data={'action': 'reset'}, follow_redirects=False)
        assert response.status_code == 303 and load_export_panels(repository) == default_export_panels()
    reset()
    package = tmp_path / 'scoring-export-panels.zip'
    package.write_bytes(exported.content)
    app_module._apply_import_archive(package, app_module.read_import_manifest(package), destination_workspace_ids=[workspace.id])
    assert load_export_panels(repository) == saved
    backup_path = app_module.create_recurring_database_backup({
        'components': ['scoring_export_panels'], 'workspace_ids': [workspace.id], 'backup_path': str(tmp_path), 'max_backups': 5})
    assert app_module._backup_archive_components(backup_path) == ['scoring_export_panels']
    reset()
    app_module.restore_database_backup(backup_path, ['scoring_export_panels'])
    assert load_export_panels(repository) == saved
    # The Scoring & GAP Analysis Methodologies do not carry them.
    assert 'export_panels' not in json.loads(app_module._scoring_configuration_archive_payload(workspace))
