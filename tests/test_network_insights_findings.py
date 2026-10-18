"""Key Findings of Network Insights: the six benchmark views of the Data tests, their slides in the colours of the
theme, and the theme each user chooses."""
from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
from PIL import Image
from pptx import Presentation

import src.DriveTestAnalyzer as core
from src.modules.cdr_reporting import _remove_all_slides
from src.modules.network_insights_export import (
    REPORT_PANELS, export_network_insights_powerpoint, report_component, table_colours,
)
from src.modules.network_insights_findings import findings_payload, lte_band_group, mode_of_test
from src.modules.network_insights_findings_slides import add_finding_slides, palette
from src.modules.network_insights_export import _slide
from src.modules.scoring_themes import THEMES, use_theme

TEMPLATE = Path(__file__).resolve().parents[1] / 'assets/powerpoint-templates/Template_VodafoneThree.pptx'
P_STYLE = '{http://schemas.openxmlformats.org/presentationml/2006/main}style'


def _samples(count: int = 240) -> pd.DataFrame:
    """Data tests of two operators: A strong, B weak (lower RSRP, less 5G, narrower bandwidth), in Cities and Towns,
    driven and walked."""
    generator = np.random.default_rng(7)
    rows = []
    for index in range(count):
        operator = 'A' if index % 2 == 0 else 'B'
        strong = operator == 'A'
        nr = index % 4 != 1  # B uses 5G in half of its tests, A always.
        rows.append({
            'kind': 'Data', 'operator': operator, 'environment': 'Cities' if index % 3 else 'Towns',
            'mode': 'Walk' if index % 5 == 0 else 'Drive', 'download': index % 3 != 2,
            'lte_rsrp': generator.normal(-85 if strong else -103, 4), 'lte_sinr': generator.normal(14 if strong else 4, 2),
            'nr_rsrp': generator.normal(-90 if strong else -99, 4) if nr else np.nan,
            'nr_sinr': generator.normal(12 if strong else 6, 2) if nr else np.nan,
            'lte_band': 'B7' if strong else ('B3' if index % 3 else 'B38'), 'nr_band': ('n78' if strong or index % 8 else 'n28') if nr else '',
            'nr_time': 6.0 if nr else 0.0, 'test_duration': 7.5,
            'lte_aggregated_bandwidth': 60.0 if strong else 20.0, 'lte_bandwidth': 20.0, 'nr_bandwidth': (100.0 if strong else 10.0) if nr else np.nan,
        })
    return pd.DataFrame(rows)


def _findings() -> dict:
    samples = _samples()
    return findings_payload(samples, samples['operator'], ['A', 'B'], {'A': '#E60000', 'B': '#1D7A8C'})


def test_modes_of_tests_and_main_4g_bands() -> None:
    assert [mode_of_test(value) for value in ('Drive', 'walk test', 'Train', '', None)] == ['Drive', 'Walk', 'Train', '', '']
    assert [lte_band_group(band) for band in ('B7', 'B38', 'B1', 'B3', 'B40', 'B20', 'B32', '')] == [
        '2600 MHz', '2600 MHz TDD', '2100 MHz', '1800 MHz', '2300 MHz', '700–900 MHz', 'Other', '']


def test_findings_compare_the_groups_with_headlines_from_their_figures() -> None:
    findings = _findings()
    views = {view['key']: view for view in findings['views']}
    assert list(views) == ['lte_signal', 'nr_signal', 'bandwidth_spread', 'technology', 'bands', 'total_bandwidth']
    assert all(view['tag'] == 'Data · drive and walk' for view in views.values())
    lte = views['lte_signal']
    # B is weaker and noisier in every environment.
    assert lte['headline'] == "B's LTE signal is the weakest of the two in every environment; its LTE SINR is the lowest in every environment"
    assert lte['meaning'].startswith('in cities ') and "of B's tests ran below −100 dBm on LTE." in lte['meaning']
    rsrp = lte['panels'][0]
    assert rsrp['classes'] == ['< −110', '−110 to −100', '−100 to −90', '−90 to −80', '≥ −80']
    assert [block['environment'] for block in rsrp['environments']] == ['Cities', 'Towns']
    assert all(abs(sum(row['shares']) - 100) < 0.5 for block in rsrp['environments'] for row in block['rows'])
    assert views['nr_signal']['headline'].startswith("A's 5G quality is the best of the two")
    spread = views['bandwidth_spread']
    assert spread['headline'] == "A's median download had 160 MHz in total (B 25); 0.0% under 40 MHz, the fewest"
    narrow = {row['group']: row for row in spread['panels'][1]['rows']}
    assert narrow['A']['share'] == 0 and narrow['A']['rank'] == 0 and narrow['B']['rank'] == 1
    technology = views['technology']
    assert [panel['mode'] for panel in technology['panels']] == ['Drive', 'Walk']
    drive = {row['group']: row for row in technology['panels'][0]['rows']}
    assert drive['A']['shares'][0] == 100 and 40 < drive['B']['shares'][0] < 60
    assert drive['A']['time_share'] == 80
    assert technology['reading']['lines'][0]['lead'] == 'A:'
    bands = views['bands']
    assert bands['classes'][0] == '2600 MHz' and bands['classes'][-1] == 'Other'
    assert bands['headline'].startswith("4G mostly on 2600 MHz, the most on A's downloads (100%)")
    assert bands['side']['title'] == '5G on 3.5 GHz'
    ladder = views['total_bandwidth']
    assert ladder['headline'] == "100% of A's downloads had 100 MHz or more in 4G and 5G together, the most of the two"
    assert [item['label'] for item in ladder['classes']][0] == '200+ MHz'


def test_findings_without_data_tests_have_no_views() -> None:
    samples = _samples(20).assign(kind='Voice')
    assert findings_payload(samples, samples['operator'], ['A', 'B'], {})['views'] == []


def _png(*_args) -> bytes:
    output = BytesIO()
    Image.new('RGB', (160, 90), (230, 236, 245)).save(output, 'PNG')
    return output.getvalue()


def test_finding_slides_are_drawn_with_shapes_in_the_theme() -> None:
    findings = _findings()
    with use_theme('vodafone'):
        presentation = Presentation(TEMPLATE)
        _remove_all_slides(presentation)
        for view in findings['views']:
            add_finding_slides(presentation, view, findings, 'Scope: National.', _slide)
        colours = palette()
    assert colours['accent'] == THEMES['vodafone']['title_accent']
    assert len(presentation.slides) == 6
    for slide, view in zip(presentation.slides, findings['views']):
        title = next(shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3})
        heading, headline = title.text_frame.paragraphs[:2]
        assert heading.text == view['title'] and heading.runs[0].font.bold
        assert headline.text == view['headline']
        assert str(headline.runs[0].font.color.rgb) == THEMES['vodafone']['title_accent'].lstrip('#')
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        assert 'Data · drive and walk' in texts
        assert any(text.startswith('What it means: ') for text in texts)
        assert any(text.endswith('Scope: National.') for text in texts)
        # No shape keeps the style of the theme, which draws shadows in LibreOffice.
        assert not [shape for shape in slide.shapes if not shape.is_placeholder and shape._element.find(P_STYLE) is not None]
    signal = presentation.slides[0]
    scale = {colour.lstrip('#').upper() for colour in THEMES['vodafone']['scale']}
    fills = {str(shape.fill.fore_color.rgb) for shape in signal.shapes if shape.shape_type == 1 and shape.fill.type == 1}
    assert scale <= fills


def test_key_findings_are_a_panel_of_the_export_after_the_overview(tmp_path) -> None:
    assert [key for key, _title, _items in REPORT_PANELS][:2] == ['overview', 'findings']
    assert report_component('LTE signal at the test: RSRP and SINR') == 'findings.lte_signal'
    assert report_component('Total bandwidth · 1/2') == 'findings.total_bandwidth'
    assert report_component('Bandwidth: spread and narrow tests') == 'findings.bandwidth_spread'
    assert report_component('Bandwidth of each test: LTE + NR') == 'spectrum.bandwidth'
    analysis = {'rf_rows': [], 'maps': {}, 'spectrum': {}, 'charts': {}, 'findings': _findings()}
    output = tmp_path / 'findings.pptx'
    export_network_insights_powerpoint(output, analysis, {}, {'campaigns': ['2026-Q3']}, _png, template=TEMPLATE)
    titles = [next((shape.text_frame.paragraphs[0].text for shape in slide.placeholders if shape.placeholder_format.type in {1, 3}), '')
              for slide in Presentation(output).slides]
    # Right after the Overview.
    start = titles.index('Overview') + 1
    assert titles[start:start + 6] == [view['title'] for view in analysis['findings']['views']]
    chosen = tmp_path / 'chosen.pptx'
    export_network_insights_powerpoint(chosen, analysis, {}, {}, _png, template=TEMPLATE, components=['findings.technology'])
    texts = [shape.text_frame.text for slide in Presentation(chosen).slides for shape in slide.shapes if shape.has_text_frame]
    assert any(text.startswith('Technology in use') for text in texts)
    assert not any(text.startswith('Bands in use') for text in texts)


def test_export_tables_take_the_colours_of_the_theme() -> None:
    with use_theme('vodafone'):
        colours = table_colours()
    assert colours['header'] == THEMES['vodafone']['header'].lstrip('#')
    assert colours['odd'] == THEMES['vodafone']['row_alt'].lstrip('#')
    with use_theme('midnight'):
        # The text stays dark on the white slides of a dark theme.
        assert table_colours()['text'] == THEMES['midnight']['title'].lstrip('#')


def test_each_user_keeps_a_network_insights_theme(client) -> None:
    client.cookies.clear()
    assert client.post('/login', data={'username': 'super', 'password': 'super123'}, follow_redirects=False).status_code == 303
    page = client.get('/network-insights').text
    assert 'id="ni-theme"' in page and 'id="ni-findings"' in page and 'js/network_insights_findings.js' in page
    config = json.loads(page.split('<script id="ni-config" type="application/json">', 1)[1].split('</script>', 1)[0])
    assert config['theme'] == 'netcheck' and [theme['id'] for theme in config['themes']] == list(THEMES)
    assert client.get('/api/network-insights/theme').json() == {'theme': 'netcheck'}
    assert client.put('/api/network-insights/theme', json={'theme': 'paper'}).json() == {'theme': 'paper'}
    assert client.get('/api/network-insights/theme').json() == {'theme': 'paper'}
    # The theme of Scoring stays apart.
    assert client.get('/api/scoring/theme').json()['theme'] == 'netcheck'
    assert client.put('/api/network-insights/theme', json={'theme': 'sunset'}).status_code == 400
    changes = [json.loads(row['details']) for row in core.repository.list_logs() if row['action'] == 'network_insights_theme_changed']
    assert changes[0] == {'theme': 'paper', 'previous': 'netcheck'}


def test_complete_reports_and_reporting_jobs_choose_a_network_insights_theme() -> None:
    from src.modules import complete_reports
    from src.modules.report_tasks import _network_entry

    assert complete_reports.default_options('network_insights')['theme'] == 'netcheck'
    definition = complete_reports.normalize_definition({'sections': [
        {'key': 'scoring', 'options': {'theme': 'Vodafone'}},
        {'key': 'network_insights', 'items': [{'key': 'overview'}, {'key': 'spectrum'}], 'options': {'theme': 'paper'}}]})
    network = next(section for section in definition['sections'] if section['key'] == 'network_insights')
    # One theme for the whole report: that of Scoring & GAP Analysis.
    assert network['options']['theme'] == 'vodafone'
    # Reports saved before the Key Findings have them right after the Overview.
    assert [item['key'] for item in network['items']][:3] == ['overview', 'findings', 'spectrum']
    assert _network_entry({'theme': 'midnight', 'selection': {}})['theme'] == 'midnight'
    assert _network_entry({'selection': {}})['theme'] == 'netcheck'


def test_a_headline_names_only_the_values_the_tests_have() -> None:
    # The CDRs log the 5G RSRP but not the 5G SINR: the 5G headline names the weakest signal only.
    samples = _samples().assign(nr_sinr=np.nan)
    views = {view['key']: view for view in findings_payload(samples, samples['operator'], ['A', 'B'], {})['views']}
    assert views['nr_signal']['headline'] == "B's 5G signal is the weakest in every environment"
    assert views['nr_signal']['meaning'].startswith("B's 5G is the weakest (median RSRP −")
