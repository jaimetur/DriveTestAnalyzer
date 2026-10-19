"""The scorecard design of the Scoring and GAP tables (page, PowerPoint and Word), the export choices of a Saved
Filter and the compact Word reports."""
from __future__ import annotations

import json
import re
import shutil
import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.oxml.ns import qn

import src.modules.pptx_to_docx as pptx_to_docx_module
from src.modules.scoring_exports import SECTION_TAG_PREFIX, export_scoring_powerpoint
from src.modules.scoring_pptx_scorecard import STYLES, drop_columns
from src.modules.scoring_themes import THEMES
from tests.scoring_synthetic import synthetic_result
from tests.test_scoring_exports import TEMPLATE

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCORING_SCRIPT = PROJECT_ROOT / 'src/web_interface/static/js/scoring.js'
SCORING_TEMPLATE = PROJECT_ROOT / 'src/web_interface/templates/scoring.html'

# A small DOM with what the scorecard layout of the page uses.
FAKE_DOM = r"""
class TextNode { constructor(text) { this.tagName = '#text'; this.textContent = String(text); this.parentElement = null; this.children = []; } }
class Element {
  constructor(tag) {
    this.tagName = tag.toUpperCase(); this.children = []; this.parentElement = null; this.dataset = {};
    this.className = ''; this.title = ''; this.colSpan = 1; this.rowSpan = 1; this.scope = '';
    const styles = {};
    this.style = {backgroundColor: '', setProperty: (key, value) => { styles[key] = String(value); },
                  removeProperty: key => { delete styles[key]; if (key === 'background-color') this.style.backgroundColor = ''; },
                  getPropertyValue: key => styles[key] || ''};
    this.classList = {
      add: (...names) => { for (const name of names) if (!this.classList.contains(name)) this.className = `${this.className} ${name}`.trim(); },
      contains: name => this.className.split(/\s+/).includes(name),
    };
  }
  get childNodes() { return [...this.children]; }
  get textContent() { return this.children.map(child => child.textContent).join(''); }
  set textContent(value) { this.replaceChildren(new TextNode(value)); }
  detach(node) { if (node.parentElement) node.parentElement.children.splice(node.parentElement.children.indexOf(node), 1); node.parentElement = this; }
  append(...nodes) { for (const node of nodes) { this.detach(node); this.children.push(node); } }
  prepend(...nodes) { for (const node of nodes.reverse()) { this.detach(node); this.children.unshift(node); } }
  insertBefore(node, reference) { this.detach(node); const index = reference ? this.children.indexOf(reference) : -1;
    this.children.splice(index < 0 ? this.children.length : index, 0, node); }
  replaceChildren(...nodes) { for (const child of this.children) child.parentElement = null; this.children = []; this.append(...nodes); }
  remove() { if (this.parentElement) { this.parentElement.children.splice(this.parentElement.children.indexOf(this), 1); this.parentElement = null; } }
  querySelector(selector) { return all(this).slice(1).find(node => node.classList?.contains(selector.slice(1))) || null; }
  get rows() { return this.children.filter(node => node.tagName === 'TR'); }
  get cells() { return this.children.filter(node => node.tagName === 'TD' || node.tagName === 'TH'); }
  get tHead() { return this.children.find(node => node.tagName === 'THEAD') || null; }
  get tBodies() { return this.children.filter(node => node.tagName === 'TBODY'); }
  get tFoot() { return this.children.find(node => node.tagName === 'TFOOT') || null; }
}
function all(node) { return [node, ...(node.children || []).flatMap(all)]; }
const document = {createElement: tag => new Element(tag)};
function make(tag, text = '', options = {}) {
  const node = new Element(tag);
  if (text) node.textContent = text;
  for (const [key, value] of Object.entries(options)) {
    if (key === 'dataset') Object.assign(node.dataset, value);
    else if (key === 'background') node.style.backgroundColor = value;
    else node[key] = value;
  }
  return node;
}
function row(...cells) { const tr = make('tr'); tr.append(...cells); return tr; }
function describe(table) {
  const head = table.tHead.rows.map(tr => tr.cells.map(cell => [cell.textContent, cell.className, cell.colSpan, cell.rowSpan]));
  const body = table.tBodies[0].rows.map(tr => ({
    section: tr.classList.contains('scoring-section-row'),
    cells: tr.cells.map(cell => ({column: cell.dataset.column || '', text: cell.textContent, colSpan: cell.colSpan,
      badge: cell.querySelector('.scoring-kpi-badge')?.textContent || '', meta: cell.querySelector('.scoring-kpi-meta')?.textContent || '',
      tag: cell.querySelector('.scoring-kpi-category')?.textContent || '', name: cell.querySelector('.scoring-kpi-name')?.textContent || '',
      share: cell.style.getPropertyValue('--share'), tint: cell.style.getPropertyValue('--cell-tint'), background: cell.style.backgroundColor})),
  }));
  const foot = table.tFoot ? table.tFoot.rows.map(tr => tr.cells.map(cell => [cell.dataset.column || '', cell.textContent, cell.colSpan])) : [];
  return {className: table.className, head, body, foot};
}
"""


def _function_source(script: str, name: str) -> str:
    start = script.index(f'  function {name}(')
    end = script.index('\n  }', start) + len('\n  }')
    return script[start:end]


def _run_node(program: str) -> dict:
    node = shutil.which('node')
    if not node:
        pytest.skip('Node.js is required for the scorecard checks.')
    script = SCORING_SCRIPT.read_text(encoding='utf-8')
    helpers = '\n'.join(_function_source(script, name) for name in (
        'scorecardNumber', 'decorateScorecardLabel', 'decorateScorecardRow', 'applyScorecardLayout', 'markScorecardGroupEdges'))
    completed = subprocess.run([node, '-e', FAKE_DOM + helpers + program], capture_output=True, text=True, check=False)
    assert completed.returncode == 0, completed.stderr
    return json.loads(completed.stdout)


def test_scoring_tables_read_as_a_scorecard():
    result = _run_node(r"""
const table = make('table', '', {className: 'scoring-comparison-table'});
const thead = make('thead');
thead.append(
  row(make('th', 'CATEGORY', {rowSpan: 2, dataset: {column: 'category'}}), make('th', 'KPI', {rowSpan: 2, dataset: {column: 'kpi'}}),
      make('th', 'Type of KPI', {rowSpan: 2, dataset: {column: 'type'}}), make('th', 'Score weight (%)', {rowSpan: 2, dataset: {column: 'weight'}}),
      make('th', 'Max score', {rowSpan: 2, dataset: {column: 'maximum'}}),
      make('th', 'Score', {colSpan: 2, className: 'scoring-column-group scoring-score-group'})),
  row(make('th', 'EE', {className: 'scoring-operator-header', background: 'rgb(23, 160, 159)', dataset: {column: 'score'}}),
      make('th', 'O2', {className: 'scoring-operator-header', dataset: {column: 'score'}})));
const tbody = make('tbody');
const kpi = (category, name, type, weight, maximum, points, band, subtotal = false) => {
  const cells = [];
  if (category) cells.push(make('td', category.name, {className: 'scoring-category-cell', rowSpan: category.span, dataset: {column: 'category'}}));
  cells.push(make('td', name, {dataset: {column: 'kpi'}}));
  cells.push(make('td', type, {className: 'scoring-kpi-type-cell', dataset: {column: 'type', kpiType: type.toLowerCase() || 'neutral'}}));
  cells.push(make('td', weight, {dataset: {column: 'weight', numeric: 'true'}}), make('td', maximum, {dataset: {column: 'maximum', numeric: 'true'}}));
  for (const value of points) cells.push(make('td', value, {dataset: {column: 'score', numeric: 'true', band}, background: band ? '#DDEFD2' : ''}));
  const tr = row(...cells);
  if (subtotal) tr.classList.add('scoring-category-subtotal');
  return tr;
};
tbody.append(
  kpi({name: 'Voice', span: 3}, 'Call Setup', 'Reliable', '50.00', '500.00', ['400.00', '350.00'], 'High'),
  kpi(null, 'Drop Rate', 'Diff', '30.00', '300.00', ['150.00', '300.00'], 'Low'),
  kpi(null, 'Voice total', '—', '80.00', '800.00', ['550.00', '650.00'], '', true),
  kpi({name: 'Data', span: 2}, 'Throughput', 'Reliable', '20.00', '200.00', ['200.00', 'N/A*'], 'UltraHigh'),
  kpi(null, 'Data total', '—', '20.00', '200.00', ['200.00', 'N/A*'], '', true));
const tfoot = make('tfoot');
tfoot.append(row(make('th', 'Total', {dataset: {column: 'category'}}), make('td', 'Weighted score', {dataset: {column: 'kpi'}}),
  make('td', '—', {dataset: {column: 'type'}}), make('td', '100.00', {dataset: {column: 'weight'}}), make('td', '1000.00', {dataset: {column: 'maximum'}}),
  make('td', '750.00', {dataset: {column: 'score'}}), make('td', '650.00', {dataset: {column: 'score'}})));
table.append(thead, tbody, tfoot);
applyScorecardLayout(table);
process.stdout.write(JSON.stringify(describe(table)));
""")
    assert 'scoring-scorecard' in result['className']
    # The header starts with the operators: no Category, Type, Weight or Max columns, no lone Score group.
    assert [[cell[0] for cell in row] for row in result['head']] == [['KPI', 'EE', 'O2']]
    assert 'scoring-scorecard-corner' in result['head'][0][0][1] and result['head'][0][0][3] == 1
    # Each category opens a section with its subtotal; each KPI has its badge, weight and maximum.
    rows = result['body']
    assert [row['section'] for row in rows] == [True, False, False, True, False]
    labels = [row['cells'][0] for row in rows]
    assert [label['name'] for label in labels] == ['Voice', 'Call Setup', 'Drop Rate', 'Data', 'Throughput']
    assert [label['badge'] for label in labels] == ['', 'R', 'D', '', 'R']
    assert labels[1]['meta'] == '50.00% · 500.00 pts' and labels[0]['meta'] == '80.00% · 800.00 pts'
    assert all(len(row['cells']) == 3 for row in rows)
    # Every score keeps the share of its maximum and its band colour as a tint for the theme.
    call_setup = rows[1]['cells']
    assert (call_setup[1]['share'], call_setup[2]['share']) == ('0.800', '0.700')
    assert call_setup[1]['tint'] == '#DDEFD2' and call_setup[1]['background'] == ''
    assert rows[4]['cells'][2]['share'] == ''
    assert [cell[:2] for cell in result['foot'][0]] == [['kpi', 'Weighted score100.00% · 1000.00 pts'], ['score', '750.00'], ['score', '650.00']]


def test_gap_tables_get_sections_or_carry_the_category_of_each_kpi():
    result = _run_node(r"""
function gapTable(order) {
  const table = make('table', '', {className: 'scoring-comparison-table scoring-priority-table'});
  const thead = make('thead');
  thead.append(row(make('th', 'Category', {dataset: {column: 'category'}}), make('th', 'KPI', {dataset: {column: 'kpi'}}),
    make('th', 'Type of KPI', {dataset: {column: 'type'}}), make('th', 'VF − EE', {className: 'scoring-gap-header scoring-gap-operator-header', dataset: {column: 'gap'}})));
  const tbody = make('tbody');
  let previous = null;
  for (const [category, name, gap, span] of order) {
    const cells = [];
    if (category !== previous) cells.push(make('td', category, {className: 'scoring-category-cell', rowSpan: span, dataset: {column: 'category'}}));
    previous = category;
    cells.push(make('td', name, {dataset: {column: 'kpi'}}), make('td', 'Diff', {dataset: {column: 'type', kpiType: 'diff'}}),
      make('td', gap, {dataset: {column: 'gap', gap: gap.startsWith('−') ? 'loss' : 'gain'}, background: '#F4C7C7'}));
    tbody.append(row(...cells));
  }
  const tfoot = make('tfoot');
  tfoot.append(row(make('td', 'Total KPI GAP', {colSpan: 3}), make('td', '−1.00', {dataset: {column: 'gap'}})));
  table.append(thead, tbody, tfoot);
  applyScorecardLayout(table);
  return describe(table);
}
process.stdout.write(JSON.stringify({
  grouped: gapTable([['Voice', 'Call Setup', '−2.00', 2], ['Voice', 'Drop Rate', '+1.00', 2], ['Data', 'Throughput', '−0.50', 1]]),
  ordered: gapTable([['Voice', 'Call Setup', '−2.00', 1], ['Data', 'Throughput', '−0.50', 1], ['Voice', 'Drop Rate', '+1.00', 1]]),
}));
""")
    grouped = result['grouped']['body']
    assert [row['section'] for row in grouped] == [True, False, False, True, False]
    assert grouped[0]['cells'][0]['name'] == 'Voice' and grouped[0]['cells'][0]['meta'] == '2 KPIs'
    assert grouped[0]['cells'][0]['colSpan'] == 2 and grouped[3]['cells'][0]['meta'] == '1 KPI'
    assert grouped[1]['cells'][1]['tint'] == '#F4C7C7'
    # Ordered by GAP, the categories are apart: no sections, each KPI with its category.
    ordered = result['ordered']['body']
    assert not any(row['section'] for row in ordered)
    assert [row['cells'][0]['tag'] for row in ordered] == ['Voice', 'Data', 'Voice']
    assert result['ordered']['foot'][0][0] == ['kpi', 'Total KPI GAP', 1]


def test_every_theme_styles_the_scorecard_and_gaps_share_the_card_with_a_bar():
    template = SCORING_TEMPLATE.read_text(encoding='utf-8')
    for theme in THEMES:
        assert f'[data-scoring-theme="{theme}"] .scoring-comparison-table.scoring-scorecard' in template, theme
    # Every GAP is a card with a bar either side of zero, in the colours of the theme.
    assert re.search(r'td\[data-gap="gain"\]::after \{ left: 50%; background: var\(--sc-gain-bar\);', template)
    assert re.search(r'td\[data-gap="loss"\]::after \{ right: 50%; background: var\(--sc-loss-bar\);', template)
    # Ericsson fills each score as far as its share of the maximum; Midnight is the only dark theme.
    assert 'calc(var(--share, 0) * 100%), #f2f2f2 0)' in template
    assert THEMES['midnight']['surface'] == '#0F1726'
    for name, theme in THEMES.items():
        if name == 'midnight':
            continue
        red, green, blue = (int(theme['header'][index:index + 2], 16) for index in (1, 3, 5))
        assert name == 'netcheck' or .299 * red + .587 * green + .114 * blue > 140 or name in {'ericsson', 'vodafone'}, name
    # Heatmap is light everywhere: its headers and bands too.
    assert all(min(int(color[index:index + 2], 16) for index in (1, 3, 5)) > 190 for color in THEMES['heatmap']['bands'].values())
    assert set(STYLES) == set(THEMES)


def _slides_titled(deck, prefix):
    return [slide for slide in deck.slides if slide.shapes.title is not None
            and slide.shapes.title.text_frame.text.split('\n')[0].startswith(prefix)]


def test_powerpoint_tables_have_the_scorecard_design_of_the_theme():
    job, result = synthetic_result(levels=('Operator',))
    for theme in ('netcheck', 'ericsson', 'midnight'):
        deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, theme=theme,
                                                              scorings=['best_network'], environments=['Combined'])))
        breakdown = _slides_titled(deck, 'Scoring Tables — Breakdown')[0]
        table = next(shape.table for shape in breakdown.shapes if shape.has_table and len(shape.table.columns) > 3)
        headers = [cell.text for cell in table.rows[0].cells]
        assert headers[0] == 'KPI' and not {'CATEGORY', 'Type of KPI', 'Score weight\n(%)', 'Max score'} & set(headers)
        labels = [row.cells[0].text for row in list(table.rows)[1:]]
        # A section per category (its subtotal first), then its KPIs with their badge, weight and maximum.
        assert labels[0].startswith('CLASSIC CALLS') and ' pts' in labels[0]
        assert labels[1].startswith(('R  ', 'D  ')) and labels[1].endswith(' pts')
        assert labels[-1].startswith('Weighted score')
        # The scores fill their cells as far as their share of the maximum.
        fills = [row.cells[1]._tc.tcPr.find(qn('a:gradFill')) for row in list(table.rows)[2:-1]]
        assert any(fill is not None for fill in fills), theme
        gap_slide = _slides_titled(deck, 'GAP Analysis — All vs')[0]
        gap_table = next(shape.table for shape in gap_slide.shapes if shape.has_table)
        assert gap_table.rows[0].cells[0].text == 'NetCheck KPIs'
        assert 'Type of KPI' not in [cell.text for cell in gap_table.rows[0].cells]
        assert gap_table.rows[1].cells[0].text.endswith('KPIs')
        # The cells are set apart by the lines of the theme, never a dark grid.
        borders = {line.get('val') for line in table._tbl.iter(qn('a:srgbClr')) if line.getparent().getparent().tag == qn('a:lnL')}
        assert '65717A' not in borders


def test_drop_columns_keeps_a_label_merged_across_them():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(2, 4, 0, 0, 4000000, 800000).table
    table.cell(1, 0).merge(table.cell(1, 2))
    table.cell(1, 0).text = 'Total KPI GAP'
    drop_columns(table, {0, 2})
    assert len(table.columns) == 2
    assert table.cell(1, 0).text == 'Total KPI GAP' and not table.cell(1, 0).is_merge_origin


def test_exports_keep_the_chosen_scorings_and_environments_and_name_their_sections():
    job, result = synthetic_result(levels=('Operator',))
    full = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE)))
    chosen = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, scorings=['best_network'],
                                                            environments=['Combined'])))
    assert len(chosen.slides) < len(full.slides)
    titles = [slide.shapes.title.text_frame.text for slide in chosen.slides if slide.shapes.title is not None]
    assert not any('Most Reliable' in title for title in titles)
    sections = [str(slide._element.cSld.get('name')) for slide in chosen.slides]
    assert all(name.startswith(SECTION_TAG_PREFIX) for name in sections)
    names = list(dict.fromkeys(name[len(SECTION_TAG_PREFIX):] for name in sections))
    assert names == ['cover', 'summary', 'charts', 'tables', 'gap', 'points']
    with pytest.raises(ValueError, match='None of the selected environments'):
        export_scoring_powerpoint(job, result, TEMPLATE, environments=['Nowhere'])


def test_scoring_exports_are_composed_as_the_bundle_report_template(monkeypatch):
    from src.modules import bundle_report_templates, pptx_compose

    job, result = synthetic_result(levels=('Operator',))
    content = export_scoring_powerpoint(job, result, TEMPLATE, scorings=['best_network'], environments=['Combined'])
    # The PowerPoint keeps the covers of the module and composes the rest as the Scoring section of a Bundle Report Template.
    original = Presentation(BytesIO(content))
    composed = Presentation(BytesIO(bundle_report_templates.module_document('scoring', content, 'powerpoint', 'Scoring & GAP Analysis')))
    titles = [pptx_compose.title_parts(slide)[0] for slide in composed.slides]
    assert titles[0] == 'Scoring & GAP Analysis' and any('Best Network Scoring' in title for title in titles)
    # It ends with the closing slide of the template, as a Bundle Report Template.
    assert 1 < len(composed.slides) <= len(original.slides) + 1 and composed.slides[-1].slide_layout.name == 'Black logo end slide'

    # A slide image keeps its blocks and loses its empty bands.
    image = Image.new('RGB', (400, 300), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 20, 380, 60), fill='black')
    draw.rectangle((20, 250, 380, 270), fill='black')
    raw = BytesIO()
    image.save(raw, format='PNG')
    with Image.open(BytesIO(pptx_to_docx_module.compact_slide_image(raw.getvalue()))) as compact:
        assert compact.width == 400 and compact.height < 120

    # Word: the cover of the module fills a page of its own without margins, then portrait pages hold as many
    # components as fit, as in the Word version of a Bundle Report Template.
    monkeypatch.setattr(pptx_to_docx_module, 'portrait_covers', lambda _content, positions: [raw.getvalue()] * len(positions))
    monkeypatch.setattr(pptx_to_docx_module, 'slide_images',
                        lambda deck, dpi=160: [raw.getvalue()] * len(Presentation(BytesIO(deck)).slides))
    document = Document(BytesIO(bundle_report_templates.module_document('scoring', content, 'word', 'Scoring & GAP Analysis')))
    cover_page, *pages = document.sections
    assert cover_page.left_margin == cover_page.top_margin == 0
    assert pages and pages[0].left_margin > 0 and pages[0].page_height > pages[0].page_width


def test_compared_operators_share_slides_without_leaving_one_alone():
    from src.modules.scoring_summary_slides import operator_frames, operator_groups

    sizes = {count: [len(group) for group in operator_groups(list(range(count)))] for count in range(1, 11)}
    # Up to six on one slide, more on slides of four at most shared evenly: never one alone.
    assert sizes == {1: [1], 2: [2], 3: [3], 4: [4], 5: [5], 6: [6], 7: [4, 3], 8: [4, 4], 9: [3, 3, 3], 10: [4, 3, 3]}
    # One centred, two or three side by side, four in two rows of two, five or six in two rows of three.
    rows = {count: len({round(top, 2) for _left, top, _width, _height in operator_frames(count)}) for count in range(1, 7)}
    columns = {count: len({round(left, 2) for left, _top, _width, _height in operator_frames(count)}) for count in range(1, 7)}
    assert rows == {1: 1, 2: 1, 3: 1, 4: 2, 5: 2, 6: 2} and columns == {1: 1, 2: 2, 3: 3, 4: 2, 5: 3, 6: 3}
    assert all(left + width <= 12.91 and top + height <= 7.05 for count in range(1, 7)
               for left, top, width, height in operator_frames(count))


def test_slides_of_one_campaign_keep_the_latest_campaign():
    from src.modules.scoring_summary_slides import latest_campaign_items

    items = [{'operator': 'VF', 'context': {'campaign': '2026-Q2'}}, {'operator': 'VF', 'context': {'campaign': '2026-Q3'}},
             {'operator': 'VF SA', 'context': {'campaign': '2026-Q2'}}, {'operator': 'Lebara', 'context': {'campaign': '2026-Q1'}},
             {'operator': 'EE', 'context': {}}]
    # The operators not measured in the latest campaign are left out of those slides.
    assert latest_campaign_items(items) == [items[1], items[4]]
    assert latest_campaign_items(items[:1]) == items[:1]


def test_operator_rules_reach_every_campaign_and_trend_labels_stand_apart():
    from src.modules.scoring_insight_slides import add_trend_slides

    job, result = synthetic_result(levels=('Operator', 'Campaign'))
    deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, theme='ericsson', scorings=['best_network'],
                                                          environments=['Combined'])))
    summary = next(slide for slide in deck.slides if slide.shapes.title is not None
                   and slide.shapes.title.text_frame.text.startswith('Scoring Tables'))
    table = next(shape.table for shape in summary.shapes if shape.has_table)
    # PowerPoint draws the bottom line of a merged cell under its first column only: the cells under an operator have its
    # colour on top, under every campaign of the operator.
    ruled = 0
    for row_index in range(len(table.rows) - 1):
        for column, cell in enumerate(table.rows[row_index].cells):
            if cell.is_merge_origin and cell.span_width > 1:
                line = cell._tc.tcPr.find(qn('a:lnB'))
                # The rule of an operator (2.25 pt), not the white lines between cards.
                if line is None or line.get('w') != str(int(2.25 * 12700)):
                    continue
                colour = line.find(qn('a:solidFill'))[0].get('val')
                tops = [table.cell(row_index + 1, spanned)._tc.tcPr.find(qn('a:lnT')) for spanned in range(column, column + cell.span_width)]
                assert all(top is not None and top.find(qn('a:solidFill'))[0].get('val') == colour for top in tops)
                ruled += 1
    assert ruled

    # Two lines that end close together have their names moved apart, one up and one down.
    from pptx import Presentation as NewPresentation
    from src.modules.cdr_reporting import _remove_all_slides

    presentation = NewPresentation(TEMPLATE)
    _remove_all_slides(presentation)

    def new_slide(target, _title, _subtitle):
        return target.slides.add_slide(target.slide_layouts[0])

    trend = {'title': '', 'campaigns': ['2026-Q2', '2026-Q3'], 'maximum': 1000, 'series': [
        {'label': 'VF', 'color': '#E60000', 'points': [846.3, 863.1]},
        {'label': 'Three', 'color': '#1A1A1A', 'points': [838.0, 861.4]},
        {'label': 'BT', 'color': '#00A6A6', 'points': [940.6, 938.5], 'is_reference': True}]}
    add_trend_slides(presentation, [trend], scoring_label='Best Network', subtitle='', new_slide=new_slide)
    chart = next(shape.chart for shape in presentation.slides[0].shapes if shape.has_chart)
    offsets = {}
    for series in chart.plots[0].series:
        label = series.points[1].data_label._dLbl
        layout = None if label is None else label.find(qn('c:layout'))
        offsets[series.name] = 0.0 if layout is None else float(layout.find(qn('c:manualLayout')).find(qn('c:y')).get('val'))
    assert offsets['VF'] < 0 < offsets['Three'] and offsets['BT'] == 0.0


def test_bar_fills_keep_their_edges_at_whole_percents():
    from pptx import Presentation as NewPresentation
    from pptx.util import Inches

    from src.modules.scoring_pptx_scorecard import diverging_bar, share_bar

    slide = NewPresentation().slides.add_slide(NewPresentation().slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(.4)).table
    # PowerPoint draws fills with their stops at whole percents on screen: a loss bar ends at the middle of the cell
    # and the next stop is one percent further, so the bar never runs on to the right.
    diverging_bar(table.cell(0, 0), '#FFB3B3', .223, '#F5F5F5', gain=False)
    share_bar(table.cell(0, 1), '#4F46E5', .4321, '#FFFFFF')
    for cell in table.rows[0].cells:
        positions = [int(stop.get('pos')) for stop in cell._tc.tcPr.iter(qn('a:gs'))]
        assert all(position % 1000 == 0 for position in positions)
        assert all(after - before >= 1000 for before, after in zip(positions, positions[1:]))
    loss = [(int(stop.get('pos')), stop[0].get('val')) for stop in table.cell(0, 0)._tc.tcPr.iter(qn('a:gs'))]
    assert loss == [(0, 'F5F5F5'), (39000, 'F5F5F5'), (40000, 'FFB3B3'), (50000, 'FFB3B3'), (51000, 'F5F5F5'), (100000, 'F5F5F5')]


def test_exports_show_the_operators_each_section_of_the_page_has_enabled():
    from src.modules.pptx_compose import title_parts

    job, result = synthetic_result(levels=('Operator', 'Campaign'))

    def operators_of(deck, title):
        """The operators named over the panels of the first slide with ``title``."""
        slide = next(slide for slide in deck.slides if title_parts(slide)[0].startswith(title))
        return {shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text in labels}

    deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, scorings=['best_network'], environments=['Combined'])))
    labels = {'EE', 'O2 UK', 'Three UK', 'Vodafone UK'}
    # By default the Summary shows every operator, the reference (EE) included, as the page does.
    assert operators_of(deck, 'Position — per environment') == labels
    # The operators the page has enabled in a section are those of its slides.
    job['_operator_choices'] = {'summary-operators': ['O2 UK', 'Vodafone UK']}
    chosen = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, scorings=['best_network'], environments=['Combined'])))
    assert operators_of(chosen, 'Position — per environment') == {'O2 UK', 'Vodafone UK'}


def test_reliable_and_diff_kpis_carry_badges_of_distinct_colours():
    from src.modules.pptx_compose import title_parts

    # Every theme gives the R and D badges two different colours, neither the colour of the KPI names.
    for name, current in STYLES.items():
        assert current['reliable'].upper() != current['diff'].upper(), name
        assert current['ink'].upper() not in {current['reliable'].upper(), current['diff'].upper()}, name
    job, result = synthetic_result(levels=('Operator',))
    deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, scorings=['best_network'],
                                                          environments=['Combined'], theme='ericsson')))
    for title in ('KPI GAP Profile', 'Points lost per KPI'):
        slide = next(slide for slide in deck.slides if title_parts(slide)[0].startswith(title))
        table = next(shape for shape in slide.shapes if shape.name.startswith('Scoring KPI GAP')).table
        badges = {}
        for row in list(table.rows)[2:]:
            runs = [run for paragraph in row.cells[0].text_frame.paragraphs for run in paragraph.runs]
            assert not any(run.font.underline for run in runs)
            if runs and runs[0].text.strip() in {'R', 'D'}:
                badges[runs[0].text.strip()] = str(runs[0].font.color.rgb)
        # Its KPIs start with the R or D of their type, as in the GAP tables, instead of being underlined.
        assert badges == {'R': STYLES['ericsson']['reliable'].lstrip('#'), 'D': STYLES['ericsson']['diff'].lstrip('#')}
        note = next(shape for shape in slide.shapes if shape.name == 'Scoring KPI GAP Profile Note')
        assert 'Underlined' not in note.text_frame.text


def test_gap_of_one_operator_follows_its_latest_campaign_and_fills_the_slide():
    import re as regex

    from src.modules.pptx_compose import title_parts

    job, result = synthetic_result(levels=('Operator', 'Campaign'))
    deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, scorings=['best_network'],
                                                          environments=['Combined'])))
    gap_slides = [slide for slide in deck.slides if title_parts(slide)[0].startswith('GAP Analysis')]
    individual = next(slide for slide in gap_slides if not title_parts(slide)[0].startswith('GAP Analysis — All vs'))
    table_shape = next(shape for shape in individual.shapes if shape.has_table)
    table = table_shape.table
    campaigns = [cell.text_frame.text for cell in list(table.rows[1].cells)[1:]]
    assert campaigns == sorted(campaigns) and len(campaigns) > 1

    def number(text):
        found = regex.search(r'[-−]?\d+\.\d+', text)
        return float(found.group().replace('−', '-')) if found else None

    # The KPIs are ordered by their GAP in the latest campaign (the last column), and the slide says so.
    values = [number(row.cells[len(campaigns)].text_frame.text) for row in list(table.rows)[2:-1]]
    known = [value for value in values if value is not None]
    assert known == sorted(known) and values[:len(known)] == known
    assert any(shape.has_text_frame and shape.text_frame.text.startswith('KPIs ordered by GAP in the latest campaign')
               for shape in individual.shapes)
    for slide in gap_slides:
        table_shape = next(shape for shape in slide.shapes if shape.has_table)
        heading = next(shape for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text == 'GAP color scale')
        segment = next(shape for shape in slide.shapes if shape.name == 'GAP Color Scale Segment 1')
        # The color scale is one line under the table, which takes the height of the slide down to it.
        assert abs(segment.top - heading.top) < 914400 * .05
        bottom = (table_shape.top + sum(row.height for row in table_shape.table.rows)) / 914400
        assert 6.4 < bottom <= 6.71 and bottom < heading.top / 914400
        sizes = [run.font.size.pt for row in list(table_shape.table.rows)[2:] for cell in row.cells
                 for paragraph in cell.text_frame.paragraphs for run in paragraph.runs if run.font.size]
        assert max(sizes) >= 6


def test_the_groups_below_the_operators_are_set_apart_down_the_header_and_the_body():
    result = _run_node(r"""
const table = make('table', '', {className: 'scoring-comparison-table scoring-scorecard'});
const thead = make('thead');
const operator = name => make('th', name, {colSpan: 4, className: 'scoring-hierarchy-header scoring-operator-header'});
const region = name => make('th', name, {colSpan: 2, className: 'scoring-hierarchy-header'});
thead.append(row(make('th', 'Category', {rowSpan: 3, className: 'scoring-scorecard-corner'}), operator('EE'), operator('O2')),
  row(region('North'), region('South'), region('North'), region('South')),
  row(...Array.from({length: 8}, (_value, index) => make('th', index % 2 ? '2026-Q3' : '2026-Q2', {className: 'scoring-hierarchy-header'}))));
const tbody = make('tbody');
tbody.append(row(make('td', 'Voice total'), ...Array.from({length: 8}, () => make('td', '10.00', {dataset: {column: 'score'}}))));
table.append(thead, tbody);
markScorecardGroupEdges(table);
console.log(JSON.stringify({
  head: table.tHead.rows.map(tr => tr.cells.map(cell => `${cell.dataset.groupEdge || ''}${cell.dataset.groupHeader === undefined ? '' : '*'}`)),
  body: table.tBodies[0].rows[0].cells.map(cell => cell.dataset.groupEdge || ''),
}));
""")
    # Each operator (major) and each region below it (minor) opens a group in every header row below it and in the body;
    # the regions are the labels ruled across their columns (*).
    assert result['head'] == [['', 'major', 'major'], ['major*', 'minor*', 'major*', 'minor*'],
                              ['major', '', 'minor', '', 'major', '', 'minor', '']]
    assert result['body'] == ['', 'major', '', 'minor', '', 'major', '', 'minor', '']
    css = SCORING_TEMPLATE.read_text(encoding='utf-8')
    assert ':is(thead, tbody, tfoot) tr > [data-group-edge="minor"] { border-left: 8px solid transparent !important;' in css
    assert 'thead tr > th[data-group-header] { box-shadow: inset 0 -1.5px 0' in css


def test_the_powerpoint_scorecard_sets_apart_the_groups_below_the_operators():
    from tests.test_scoring_exports import _mapping_groups, _result, scoring_configuration
    base = _result()
    rows = [{**row, 'region': region, 'campaign': campaign} for region in ('North', 'South')
            for campaign in ('UK_Q2_2026', 'UK_Q3_2026') for row in base['scoring']]
    levels = ['Operator', 'Region', 'Campaign']
    result = {'aggregation_contract_version': 2, 'aggregation_levels': levels, 'scoring': rows, 'totals': [],
              'configuration': scoring_configuration(), 'warnings': []}
    job = {'aggregation_levels': levels, 'aggregation_contract_version': 2, 'nr_mode': 'NSA', 'baseline_operator': 'EE',
           'campaigns': ['UK_Q2_2026', 'UK_Q3_2026'], 'configuration': scoring_configuration()}
    presentation = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, operator_mapping_groups=_mapping_groups(),
                                                                  environment='DriveCity', theme='netcheck')))
    table = next(shape.table for slide in presentation.slides for shape in slide.shapes
                 if shape.has_table and any(cell.text_frame.text == 'South' for row in shape.table.rows for cell in row.cells))

    def width(cell, edge):
        line = cell._tc.get_or_add_tcPr().find(qn(f'a:{edge}'))
        return int(line.get('w')) if line is not None and line.find(qn('a:solidFill')) is not None else 0
    regions = next(index for index, row in enumerate(table.rows) if any(cell.text_frame.text == 'South' for cell in row.cells))
    named = {}
    for column, cell in enumerate(table.rows[regions].cells):
        named.setdefault(cell.text_frame.text, []).append(column)
    total = len(table.rows) - 1
    # A line before each region (thin) and each operator (stronger) down to the total, and each region ruled under its name.
    assert width(table.cell(total, named['South'][0]), 'lnL') == 9525
    assert width(table.cell(total, named['North'][1]), 'lnL') == 19050
    assert width(table.cell(regions, named['North'][1]), 'lnB') == 12700


def test_ericsson_heads_its_scorecards_in_white_and_highlights_stay_in_their_box():
    # The Ericsson scorecards (page and exports) have a white header: the operators in dark ink over the rule of their
    # colour, the levels below them in grey.
    assert (STYLES['ericsson']['header'], STYLES['ericsson']['header_label'], STYLES['ericsson']['header_ink']) == ('#FFFFFF', '#242424', '#767676')
    css = SCORING_TEMPLATE.read_text(encoding='utf-8')
    assert '[data-scoring-theme="ericsson"] .scoring-comparison-table.scoring-scorecard thead tr > th { background: #0068bf' not in css
    assert '[data-scoring-theme="ericsson"] .scoring-scorecard thead .scoring-operator-chip { color: #242424; }' in css
    # In every theme the bar, card or highlight of a cell stays inside its box, never in the gap that opens a group, and
    # a highlighted total keeps the ink of the cells.
    assert ('> td[data-column][data-group-edge]:is([data-group-edge="major"], [data-group-edge="minor"]) { background-clip: padding-box !important;'
            in css)
    assert 'tfoot tr > td[data-column="score"][data-operator-rank] { color: var(--sc-ink) !important; }' in css


def test_ericsson_and_vodafone_score_cells_take_the_colours_of_the_legend():
    from src.modules.scoring_pptx_scorecard import _band_fill
    from src.modules.scoring_themes import use_theme
    assert STYLES['ericsson'].get('legend_fill') and STYLES['vodafone'].get('legend_fill') and not STYLES['netcheck'].get('legend_fill')
    with use_theme('ericsson'):
        assert _band_fill({'threshold_band': 'Medium'}) == THEMES['ericsson']['bands']['Medium']
    css = SCORING_TEMPLATE.read_text(encoding='utf-8')
    # On the page, the cell keeps the colour of its band in the legend (--cell-tint) in both themes.
    assert 'background: linear-gradient(90deg, var(--cell-tint, color-mix(in srgb, var(--band-dot, #0068bf) 34%, #fff))' in css
    assert 'background: var(--cell-tint, color-mix(in srgb, var(--band-dot, #bdbdbd) 24%, #fff)) !important;' in css


def test_the_legend_of_the_bands_looks_as_their_cells_in_every_theme():
    from pptx.util import Inches

    from src.modules.scoring_pptx_scorecard import style_legend_band
    from src.modules.scoring_themes import use_theme

    def legend_cell(theme):
        table = Presentation().slides.add_slide(Presentation().slide_layouts[6]).shapes.add_table(1, 1, 0, 0, Inches(1), Inches(.2)).table
        cell = table.cell(0, 0)
        cell.text = 'Medium'
        with use_theme(theme):
            style_legend_band(cell, 'Medium', 8)
        return cell
    # NetCheck: the data bar of its cells; Paper: the band in the coloured, underlined text of its cells.
    assert legend_cell('netcheck')._tc.find('.//' + qn('a:gradFill')) is not None
    run = legend_cell('paper').text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == THEMES['paper']['band_strong']['Medium'].lstrip('#') and run.font.underline
    # Ericsson keeps the colour of the legend, which its cells take.
    assert legend_cell('ericsson')._tc.find('.//' + qn('a:gradFill')) is None
    script = SCORING_SCRIPT.read_text(encoding='utf-8')
    css = SCORING_TEMPLATE.read_text(encoding='utf-8')
    assert 'swatch.dataset.band = label;' in script
    assert '[data-scoring-workspace] .scoring-legend-swatch[data-band="Medium"] { --band-dot: var(--st-band-medium); }' in css
    assert 'background: linear-gradient(90deg, var(--band-dot), color-mix(in srgb, var(--band-dot) 14%, #fff)) !important; }' in css
