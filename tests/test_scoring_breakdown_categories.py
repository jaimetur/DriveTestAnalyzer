import pytest
from pptx import Presentation

from src.modules.scoring_exports import _hierarchy_score_tables, _score_tables, _table_for_mode
from src.modules.scoring_views import build_scoring_views
from tests.test_scoring_table_exports import _job, _result


@pytest.mark.parametrize('hierarchy', [False, True])
def test_breakdown_opens_each_category_with_its_subtotal(hierarchy):
    job = _job()
    result = _result()
    if hierarchy:
        job['aggregation_levels'] = ['Operator', 'Region', 'Campaign']
        result.update({'aggregation_contract_version': 2, 'aggregation_levels': job['aggregation_levels']})
    views = build_scoring_views(job, result)
    matrices = views['hierarchy_score_tables' if hierarchy else 'score_tables']
    matrix = _table_for_mode(next(item for item in matrices
                                  if item['context']['environment'] == 'DriveCity'), 'expanded')
    presentation = Presentation()
    renderer = _hierarchy_score_tables if hierarchy else _score_tables
    renderer(presentation, [matrix], views.get('threshold_legend', []), show_gap_values=False)
    table = next(shape.table for shape in presentation.slides[0].shapes if shape.has_table)
    # No Category column: the scorecard starts with the KPI and each category opens a section with its subtotal.
    assert table.cell(0, 0).text == 'KPI'
    header_rows = next(index for index in range(len(table.rows)) if table.cell(index, 0).text != 'KPI'
                       and not table.cell(index, 0).is_spanned)
    labels = [table.cell(index, 0).text for index in range(header_rows, len(table.rows) - 1)]
    categories = list(dict.fromkeys(row['category'] for row in matrix['rows']))
    # A section opens with its category in bold; a KPI starts with the badge of its type.
    starts = [index for index in range(len(labels))
              if table.cell(header_rows + index, 0).text_frame.paragraphs[0].runs[0].font.bold
              and table.cell(header_rows + index, 0).text_frame.paragraphs[0].runs[0].text.strip() not in {'R', 'D'}]
    assert [labels[index].split('   ')[0] for index in starts] == categories
    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(labels)
        kpis = [row['kpi'] for row in matrix['rows'] if row['category'] == categories[position] and row.get('row_type') != 'category']
        assert [label.split('  ', 1)[-1].split('   ')[0] for label in labels[start + 1:end]] == kpis
        assert labels[start].endswith(' pts')


def test_breakdown_slides_keep_up_to_fifty_bars_together():
    from src.modules.scoring_exports import BREAKDOWN_BARS_PER_SLIDE, _breakdown_pages

    categories = [f'C{index}' for index in range(1, 11)]
    assert BREAKDOWN_BARS_PER_SLIDE == 50
    # 10 categories × 5 operators = 50 bars: one slide.
    assert _breakdown_pages(categories, 5) == [categories]
    # 12 × 5 = 60 bars: two balanced slides of 30 bars.
    twelve = [f'C{index}' for index in range(1, 13)]
    assert _breakdown_pages(twelve, 5) == [twelve[:6], twelve[6:]]
    # Not splitting keeps every category together; a category wider than the limit stays alone.
    assert _breakdown_pages(twelve, 5, split=False) == [twelve]
    assert _breakdown_pages(['A', 'B'], 60) == [['A'], ['B']]
