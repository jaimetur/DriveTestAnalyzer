import json
import re
import shutil
import subprocess
import sys
import threading
import time
import zipfile
from io import BytesIO, StringIO
from pathlib import Path
from threading import Event, Thread

import pandas as pd
import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE_TYPE

import src.DriveTestAnalyzer as core
from src.modules.ppt_designer import DesignDefinition, filter_frame


def definition(**changes):
    return DesignDefinition(name='Comparison', template='Design test', datasets={'data': [1]}, **changes)


def test_design_uses_combined_tables_without_projection_or_warmup_queue():
    source = (Path(__file__).parents[1] / 'src/modules/ppt_designer.py').read_text(encoding='utf-8')
    assert 'drivetest-analyzers.sqlite3' not in source
    assert 'def ensure_projection' not in source
    assert 'def enqueue_prefetch' not in source
    assert 'def prefetch_workspace_designs' not in source
    assert 'def reporting_source(snapshot, kind, task_repository):' in source


def test_design_filter_panel_state_is_scoped_to_the_authenticated_session():
    first_user = core.SessionUser(username='super', role='super-admin')
    second_user = core.SessionUser(username='super', role='super-admin')
    assert first_user.session_marker
    assert first_user.session_marker != second_user.session_marker


def test_design_ppt_chart_filters_use_complete_job_geography_and_cross_filter():
    node_binary = shutil.which('node')
    if node_binary is None:
        pytest.skip('Node.js is required for this browser filter test')
    script_path = Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js'
    harness = r"""
const assert = require('node:assert/strict');
const fs = require('node:fs');
const source = fs.readFileSync(process.argv[1], 'utf8');
const fields = ['nr_mode', 'design', 'template', 'scope', 'region', 'city'];
const controls = Object.fromEntries(fields.map(field => [field, {
  dataset: {designPptChartFilter: field}, options: [], value: '',
  replaceChildren(...choices) { this.options = choices; },
}]));
const document = {querySelectorAll: () => Object.values(controls)};
const option = (value, textContent) => ({value, textContent});
const start = source.indexOf('  const designPptChartFilterControls =');
const end = source.indexOf('  const designPptJobTimestampLabel =', start);
assert.ok(start >= 0 && end > start);
eval(source.slice(start, end) + `
  const jobs = [
    {nr_mode: 'SA', design_name: 'Main Cities', template: 'A', scope: 'Operator Comparison',
      cover: {regions: 'All Regions', cities: 'Cities: Belfast, Bristol, Cardiff, Edinburgh, Leeds, London, Sheffield'}},
    {nr_mode: 'SA', design_name: 'Full Report', template: 'B', scope: 'Operator Comparison',
      cover: {regions: 'All Regions', cities: 'All Cities'}},
  ];
  const labels = field => controls[field].options.map(choice => choice.textContent);
  refreshDesignPptChartFilterChoices(jobs);
  assert.deepEqual(labels('city'), ['All', 'All Cities',
    'Cities: Belfast, Bristol, Cardiff, Edinburgh, Leeds, London, Sheffield']);
  assert.deepEqual(labels('region'), ['All', 'All Regions']);
  assert.deepEqual(labels('nr_mode'), ['All', 'SA']);
  designPptChartFilterState.design = 'main cities';
  refreshDesignPptChartFilterChoices(jobs);
  assert.deepEqual(labels('city'), ['All',
    'Cities: Belfast, Bristol, Cardiff, Edinburgh, Leeds, London, Sheffield']);
  assert.deepEqual(labels('template'), ['All', 'A']);
  designPptChartFilterState.design = '';
  designPptChartFilterState.city = 'all cities';
  refreshDesignPptChartFilterChoices(jobs);
  assert.deepEqual(labels('design'), ['All', 'Full Report']);
  assert.deepEqual(jobs.filter(matchesDesignPptChartFilters), [jobs[1]]);
`);
"""
    completed = subprocess.run(
        [node_binary, '-e', harness, str(script_path)], capture_output=True, text=True, check=False,
    )
    assert completed.returncode == 0, completed.stderr


def test_design_library_can_change_nr_mode_and_template_with_confirmation_controls():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')

    assert "technologySelect.className = 'ds-design-definition-select ds-design-nr-mode-select'" in script
    assert "templateSelect.className = 'ds-design-definition-select ds-design-template-select'" in script
    assert 'will invalidate every report and chart created inside this Design' in script
    assert 'Its Scope, dates and filters are kept (a new NR Mode uses its newest campaigns).' in script
    assert "api(`/${encodeURIComponent(id)}/template`, 'PATCH'" in script
    assert '.ds-design-definition-select{' in stylesheet

    app_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')
    assert "activeCell.textContent = operation ? `${operation}(${selected[0]})` : selected[0];" in app_script
    assert "if (editedCell?.dataset.catalogueField === 'Source Dataset' && !assistanceFocused) normaliseCatalogueRows();" in app_script
    assert "const sessionMarker = document.body.dataset.authenticatedSession || 'anonymous';" in app_script
    assert "${sessionScoped ? `:${sessionMarker}` : ''}" in app_script

    design_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')
    assert "const authenticatedSession = document.body.dataset.authenticatedSession || 'anonymous';" in design_script
    assert ':open:${authenticatedSession}`' in design_script
    assert ':filters-open:${authenticatedSession}`' in design_script
    assert "sessionStorage.setItem(filtersOpenStorageKey, isOpen ? 'open' : 'closed')" in design_script
    assert 'await openDesign(last);' in design_script


def test_template_visual_controls_are_available_for_every_chart_type():
    root = Path(__file__).parents[1]
    app_script = (root / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')
    chart_script = (root / 'src/web_interface/static/js/canvas_charts.js').read_text(encoding='utf-8')

    assert 'const syncConditionalVisualCells = (row, {clear = true} = {}) =>' in app_script
    assert "'Axis X Range': true" in app_script
    assert "'Axis Y Range': true" in app_script
    assert "'Label Position': true" in app_script
    assert "'Label Format': true" in app_script
    assert "'Legend Format': true" in app_script
    assert "cell.contentEditable = enabled ? 'true' : 'false';" in app_script
    assert "if (!enabled && clear && cell.textContent.trim()) cell.textContent = '';" in app_script
    assert 'const syncConditionalVisualControls = () =>' in app_script
    assert "axis_x_range: true" in app_script
    assert "axis_y_range: true" in app_script
    assert "label_position: true" in app_script
    assert "if (control.name === 'chart_type') syncConditionalVisualControls();" in app_script
    assert "['axis_x_range', 'Axis X Range']" in app_script
    assert "['label_position', 'Label Position']" in app_script
    assert "['label_format', 'Label Format']" in app_script
    assert "['legend_format', 'Legend Format']" in app_script
    assert "['exclude_null_empty', 'Exclude Null/Empty']" in app_script
    assert "['exclude_zero', 'Exclude Zero']" in app_script
    assert 'function drawConfiguredPointLabel(' in chart_script
    assert "drawConfiguredPointLabel(context, series.name" in chart_script


def test_saving_the_embedded_template_rebuilds_the_design_immediately():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')

    assert 'const rebuildDesignAfterTemplateSave = async () =>' in script
    assert "const expandedChartIndex = !$('ds-chart-expanded-overlay').hidden" in script
    assert "await openExpandedChart(refreshedChart, null, 'design', {preserveFocus: true});" in script
    assert 'function expandedChartOverlay(show, {preserveFocus = false} = {})' in script
    assert "if (event.data?.type === 'drivetest-analyzer:template-saved')" in script
    assert 'await rebuildDesignAfterTemplateSave();' in script
    assert 'templateEditorSaved = false;' in script


def test_compact_landscape_design_comments_are_docked_to_the_bottom():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')

    assert '@media (orientation:landscape) and (max-height:600px) {' in stylesheet
    assert '#ds-viewer.ds-overlay {' in stylesheet
    assert 'grid-template-columns:minmax(10rem,1fr) auto;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-content {' in stylesheet
    assert 'flex-direction:column;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-main {' in stylesheet
    assert 'grid-template-rows:auto minmax(0,1fr);' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-positioned .ds-chart {' in stylesheet
    assert 'left:var(--ds-chart-left)!important;' in stylesheet
    assert 'height:var(--ds-chart-height)!important;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-panel {' in stylesheet
    assert 'padding:.5rem .65rem 3.75rem!important;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-content>.ds-slide-comments {' in stylesheet
    assert 'position:fixed;' in stylesheet
    assert 'bottom:max(.35rem,env(safe-area-inset-bottom,0px));' in stylesheet
    assert 'left:max(.35rem,env(safe-area-inset-left,0px));' in stylesheet
    assert 'right:max(.35rem,env(safe-area-inset-right,0px));' in stylesheet
    assert 'max-height:min(70dvh,16rem);' in stylesheet


def test_design_job_panels_use_the_stack_spacing_without_an_empty_filter_row():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')

    assert '#ds-filter-home:not(:has(> #ds-filter-panel:not([hidden]))) {' in stylesheet
    assert '.ppt-designer :is(.ds-ppt-jobs-panel,.ds-ppt-charts-panel) {' in stylesheet
    assert 'margin-top:0;' in stylesheet


def test_design_discard_restores_saved_scope_dates_and_filters_before_navigation():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')

    assert "const discardPart = () => {" in script
    assert "copyDefinitionFields(definition, savedDesignDefinition(), filterDefinitionFields);" in script
    assert "else if (choice === 'secondary') discardPart('filters');" in script
    # There is no separate Dataset Universe: Scope, vendor comparison and dates are saved with the filters.
    assert "const filterDefinitionFields = ['filters', 'custom_fields', 'hidden_filters', 'scope', 'vendor_comparison', 'date_from', 'date_to'];" in script
    assert "discardPart('universe')" not in script


def test_design_prepare_retries_one_transient_proxy_501_response():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')

    assert "path.startsWith('/prepare?') && response.status === 501" in script
    assert "await new Promise(resolve => window.setTimeout(resolve, 250));" in script


def test_auto_calculated_field_editor_is_one_wide_dialog_shared_by_every_entry_point():
    root = Path(__file__).parents[1] / 'src/web_interface'
    stylesheet = (root / 'static/css/app.css').read_text(encoding='utf-8')
    script = (root / 'static/js/app.js').read_text(encoding='utf-8')
    editor = (root / 'static/js/auto_calculated_fields.js').read_text(encoding='utf-8')
    base = (root / 'templates/base.html').read_text(encoding='utf-8')

    # One editor (auto_calculated_fields.js) for the Workspace panel, PPT Designer and the PPT Templates editor, which
    # opens it in the page holding it when embedded.
    assert "js/auto_calculated_fields.js" in base
    assert 'window.AutoCalculatedFields = {open};' in editor
    assert 'window.AutoCalculatedFields?.open({' in script
    assert 'window.top.AutoCalculatedFields' in script
    assert 'calculated-dimensions-overlay' not in script and 'Save & Materialize' not in script
    # Wide, as tall as the window, below the confirm and info dialogs it opens; the fields beside the chosen one.
    dialog_rule = stylesheet.split('.acf-dialog {', 1)[1].split('}', 1)[0]
    assert 'width: min(1320px, calc(100vw - 2rem));' in dialog_rule
    assert 'height: min(900px, calc(100dvh - 2rem));' in dialog_rule
    assert '.acf-overlay { z-index: 9990; }' in stylesheet
    assert '.acf-body { display: grid; grid-template-columns: minmax(15rem, 19rem) minmax(0, 1fr);' in stylesheet
    # One Save per field, which checks, tries and saves through the field endpoints.
    for endpoint in ("request('/check'", "request('/preview'", "request('/save'", "request('/delete'", "request('/reorder'"):
        assert endpoint in editor
    assert ".auto-calculated-field-job:not([data-materialization-job-key])" in script
    assert "if (!['processing', 'failed', 'stopped'].includes(status)) return 100;" in script


def test_designs_prepare_ahead_only_in_the_idle_warm_up_and_compact_panel_headers_stay_aligned():
    design_module = (Path(__file__).parents[1] / 'src/modules/ppt_designer.py').read_text(encoding='utf-8')
    app_stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/app.css').read_text(encoding='utf-8')
    design_stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')

    # A Design is prepared when it is opened, or ahead by the warm-up while the server is idle (test_design_warmup):
    # no change queues its preparation in the background.
    assert 'core.register_idle_task(warm_up_when_idle, DESIGN_WARMUP_IDLE_SECONDS)' in design_module
    assert 'core.submit_background_task(delayed_retry)' not in design_module
    assert 'core.submit_background_task(run)' not in design_module
    assert '.collapsible-summary .collapse-chip {' in app_stylesheet
    assert 'position: absolute; top: .7rem; right: .75rem;' in app_stylesheet
    assert '.report-charts-panel .report-charts-chart-count .pill {' in app_stylesheet
    assert '#ds-viewer.ds-presentation-active .ds-viewer-primary-controls>.ds-viewer-tool-actions {' in design_stylesheet
    assert 'visibility:visible!important;' in design_stylesheet


def test_compact_landscape_presentation_settings_are_vertically_scrollable():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')

    assert '#ds-presentation-overlay.ds-overlay {' in stylesheet
    assert '#ds-presentation-overlay .ds-presentation-dialog {' in stylesheet
    assert 'max-height:100%;' in stylesheet
    assert 'overflow-y:auto;' in stylesheet
    assert 'overscroll-behavior:contain;' in stylesheet
    assert '-webkit-overflow-scrolling:touch;' in stylesheet
    assert '#ds-presentation-overlay .ds-presentation-dialog>.section-head {' in stylesheet
    assert 'position:sticky;' in stylesheet
    assert '#ds-presentation-overlay .ds-presentation-options {' in stylesheet
    assert 'grid-template-columns:repeat(2,minmax(0,1fr));' in stylesheet


def test_compact_portrait_design_fills_the_available_viewport():
    stylesheet = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')

    assert '@media (orientation:portrait) and (max-width:640px) {' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-content {' in stylesheet
    assert 'flex:1 1 auto;' in stylesheet
    assert 'overflow:hidden;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-main {' in stylesheet
    assert 'grid-template-columns:minmax(0,3fr) minmax(0,4fr);' in stylesheet
    assert 'grid-template-rows:auto minmax(0,1fr);' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-chart-stage.ds-positioned,' in stylesheet
    assert 'height:100%;' in stylesheet
    assert 'aspect-ratio:auto;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-positioned .ds-chart {' in stylesheet
    assert 'left:var(--ds-chart-left)!important;' in stylesheet
    assert 'width:var(--ds-chart-width)!important;' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-primary-controls {' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-navigation-footer {' in stylesheet
    assert '#ds-viewer:not(.ds-presentation-active) .ds-viewer-right-actions {' in stylesheet
    assert 'grid-template-columns:repeat(3,minmax(0,1fr));' in stylesheet


def test_design_end_navigation_controls_use_line_svgs():
    template = (Path(__file__).parents[1] / 'src/web_interface/templates/ppt_designer.html').read_text(encoding='utf-8')

    assert 'id="ds-first" class="ds-slide-nav-button" title="First slide" aria-label="First slide"><svg class="ds-slide-arrow-icon"' in template
    assert 'id="ds-last" class="ds-slide-nav-button" title="Last slide" aria-label="Last slide"><svg class="ds-slide-arrow-icon"' in template
    assert 'id="ds-chart-expanded-first" title="First chart" aria-label="First chart"><svg class="ds-chart-expanded-arrow-icon"' in template
    assert 'id="ds-chart-expanded-last" title="Last chart" aria-label="Last chart"><svg class="ds-chart-expanded-arrow-icon"' in template
    assert 'M5 5v14M18 6l-6 6 6 6M12 6l-6 6 6 6' in template
    assert 'M19 5v14M6 6l6 6-6 6M12 6l6 6-6 6' in template
    assert '⏮' not in template
    assert '⏭' not in template


def without_starter_designs():
    """Start without the Designs every workspace is given (assets/ppt-designs)."""
    with core.repository.connection() as connection:
        connection.execute("DELETE FROM workspace_state WHERE key = 'ppt_designs_v2'")


def setup_design(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('sample.csv', BytesIO(b'Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\nA,London,10,HTTP DL,2026-09-01\nB,Leeds,20,HTTP DL,2026-09-02\nA,London,30,HTTP DL,2026-09-03\n'), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('Design test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Comparison,,Title and 2 columns + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
        '1,Comparison,,Title and 2 columns + Comments,Rate again,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
        '2,Next,,Title and 1 column + Comments,Rate next,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    return definition().model_dump(mode='json')


def test_filter_empty_missing_and_inclusive_dates():
    frame = pd.DataFrame({'City': ['London', 'Leeds'], 'Test_Start_Time': ['2026-09-01 23:59', '2026-09-02']})
    assert filter_frame(frame, definition(filters={'City': []})).empty
    assert filter_frame(frame, definition(filters={'Missing': ['Yes']})).empty
    assert len(filter_frame(frame, definition(date_from='2026-09-01', date_to='2026-09-01'))) == 1
    assert len(filter_frame(frame, definition(filters={'city': ['London']}))) == 1
    radio = pd.DataFrame({'RAT_A': ['LTE', 'NR'], 'technology_primary': ['4G', '5G']})
    assert len(filter_frame(radio, definition(filters={'RAT': ['NR']}))) == 1
    assert filter_frame(radio, definition(filters={'Technology': ['5G']})).empty
    geography = pd.DataFrame({
        'city': ['', None], 'G_Level_4': ['London', 'Leeds'],
        'region': ['', None], 'G_Level_2': ['England', 'England'],
    })
    assert len(filter_frame(geography, definition(filters={'City': ['London'], 'Region': ['England']}))) == 1


def test_filter_frame_ignores_date_range_but_keeps_timestamp_field_filters(monkeypatch):
    monkeypatch.setenv('IGNORE_EVENT_TIME_FILTERING', 'true')
    frame = pd.DataFrame({
        'Event_Start_Time': ['2026-09-01 10:00:00', None],
        'Event_End_Time': ['2026-09-01 11:00:00', None],
        'Value': [1, 2],
    })

    filtered = filter_frame(frame, definition(
        date_from='2026-09-01', date_to='2026-09-01',
        filters={'Event_Start_Time': ['2026-09-01 10:00:00']},
    ))

    assert len(filtered) == 1


def test_design_page_exposes_disabled_event_time_filtering(client, monkeypatch):
    monkeypatch.setenv('IGNORE_EVENT_TIME_FILTERING', 'true')
    setup_design(client)

    response = client.get('/ppt-designer')

    assert response.status_code == 200
    assert '"ignore_event_time_filtering": true' in response.text
    assert 'data-ignore-event-time-filtering="true"' in response.text


def test_design_export_uses_the_admin_import_archive_format(client):
    payload = setup_design(client)
    design_id = 'portable-design'
    saved = client.put(f'/api/ppt-designer/{design_id}', json=payload)
    assert saved.status_code == 200

    exported = client.get(f'/api/ppt-designer/{design_id}/export')

    assert exported.status_code == 200
    with zipfile.ZipFile(BytesIO(exported.content)) as archive:
        manifest = json.loads(archive.read('manifest.json'))
        assert manifest['format'] == 'drivetest-analyzer-export'
        assert manifest['kind'] == 'designs'
        assert manifest['workspace_components'] == ['designs']
        document = json.loads(archive.read(manifest['archive_path']))
    assert document['format'] == 'drivetest-analyzer-designs'
    # The Design travels as saved: its CDRs follow its NR Mode and Campaign, they are not part of it. Who created
    # it is kept apart from the definition (the save answers with it, as the list does).
    stored = {key: value for key, value in saved.json()['definition'].items() if key != 'created_by'}
    assert document['designs'] == {design_id: stored}
    assert 'datasets' not in document['designs'][design_id]
    assert document['designs'][design_id]['auto_cdrs'] is True


def test_design_library_order_is_saved_listed_and_carried_by_packages(client):
    from src.modules.ppt_designer import read_design_order

    payload = setup_design(client)
    for design_id, name in (('design-b', 'Bravo'), ('design-a', 'Alpha'), ('design-c', 'Charlie')):
        assert client.put(f'/api/ppt-designer/{design_id}', json={**payload, 'name': name}).status_code == 200
    places = lambda: {key: value['position'] for key, value in client.get('/api/ppt-designer').json().items()
                      if key in {'design-a', 'design-b', 'design-c'}}
    # Until its editors order it, no Design has a place (the page sorts them by NR Mode and name).
    assert places() == {'design-a': None, 'design-b': None, 'design-c': None}

    # Ids it does not know are left out; the Designs it does not name keep following them.
    response = client.post('/api/ppt-designer/order', json={'ids': ['design-c', 'missing', 'design-a']})
    assert response.status_code == 200 and response.json() == {'order': ['design-c', 'design-a']}
    assert places() == {'design-c': 1, 'design-a': 2, 'design-b': None}
    [event] = [row for row in core.repository.list_logs(limit=20) if row['action'] == 'reorder_designs']
    assert json.loads(event['details']) == {'previous_order': [], 'order': ['Charlie', 'Alpha']}
    # A new order changes no Design, and saving a Design keeps its place.
    assert client.put('/api/ppt-designer/design-a', json={**payload, 'name': 'Alpha 2'}).status_code == 200
    assert places() == {'design-c': 1, 'design-a': 2, 'design-b': None}
    assert 'position' not in json.loads(core.repository.get_workspace_state('ppt_designs_v2'))['design-a']

    # Packages carry the order: their Designs keep it, before those only the workspace has.
    workspace = core.active_workspace
    package = core._design_archive_payload(workspace)
    assert json.loads(package)['order'] == ['design-c', 'design-a']
    client.post('/api/ppt-designer/order', json={'ids': ['design-b']})
    assert read_design_order(core.repository) == ['design-b', 'design-c', 'design-a']
    core._import_workspace_designs(workspace, package, 'super')
    assert read_design_order(core.repository) == ['design-c', 'design-a', 'design-b']

    # A deleted Design leaves its place.
    assert client.delete('/api/ppt-designer/design-a').status_code == 200
    assert places() == {'design-c': 1, 'design-b': 2}


def test_design_saves_scope_and_dates_and_chooses_its_cdrs(client):
    payload = setup_design(client)
    payload['date_from'], payload['date_to'] = 'Oldest', 'Newest'

    saved = client.put('/api/ppt-designer/symbolic-dates', json=payload)

    assert saved.status_code == 200, saved.text
    assert {key: saved.json()['definition'][key] for key in ('scope', 'date_from', 'date_to')} == {
        key: payload[key] for key in ('scope', 'date_from', 'date_to')
    }
    stored = client.get('/api/ppt-designer').json()['symbolic-dates']
    # The CDRs are never saved: they follow the NR Mode and the Campaign filter (a former Design without
    # campaigns in its CDRs follows the newest two campaigns).
    assert 'datasets' not in stored and stored['auto_cdrs'] is True
    assert stored['filters']['Campaign'] == ['@latest-2-campaigns']
    prepared = client.post('/api/ppt-designer/prepare', json=stored)
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['date_bounds'] == {'min': '2026-09-01', 'max': '2026-09-03'}
    assert prepared.json()['rows']['data'] == 3
    assert [row['id'] for row in prepared.json()['cdrs']['data']] == [1]


def test_design_all_values_expands_when_saved_universe_gains_a_new_value(client):
    payload = setup_design(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    payload['custom_fields'] = ['Test_Name']
    payload['filters'] = {}
    design_id = 'expand-all-values'
    saved = client.put(f'/api/ppt-designer/{design_id}', json=payload)
    assert saved.status_code == 200, saved.text
    assert 'Test_Name' not in saved.json()['definition']['filters']

    initial = client.post('/api/ppt-designer/prepare', json=payload)
    assert initial.status_code == 200, initial.text
    assert initial.json()['options']['Test_Name'] == ['HTTP DL']
    assert initial.json()['rows']['data'] == 3

    uploaded = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('new-test.csv', BytesIO(
            b'Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\n'
            b'C,Bristol,40,HTTP UL,2026-09-04\n'
        ), 'text/csv'),
    })
    assert uploaded.status_code == 200, uploaded.text
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(2)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(2)['status'] == 'ready'

    # The new CDR of the NR Mode is chosen without editing the Design.
    expanded = client.get('/api/ppt-designer').json()[design_id]
    assert 'Test_Name' not in expanded['filters']
    prepared = client.post('/api/ppt-designer/prepare', json=expanded)
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['options']['Test_Name'] == ['HTTP DL', 'HTTP UL']
    assert prepared.json()['rows']['data'] == 4


def test_design_template_change_preserves_saved_scope_dates_and_filters(client):
    payload = setup_design(client)
    core.repository.add_report_template('Design SA test', core.repository.report_template_content('Design test',
    ), is_default=False)
    payload.update({
        'scope': 'multivendor',
        'auto_cdrs': True,
        'filters': {'City': ['London'], 'Campaign': ['UK_Q2_2026']},
        'custom_fields': ['Mean_Data_Rate'],
        'hidden_filters': ['Vendor'],
        'date_from': '2026-09-01',
        'date_to': '2026-09-03',
        'slide_comments': {'1': ['Keep this note']},
    })
    assert client.put('/api/ppt-designer/change-template', json=payload).status_code == 200
    saved_before_change = client.get('/api/ppt-designer').json()['change-template']

    changed = client.patch('/api/ppt-designer/change-template/template', json={
        'nr_mode': 'SA', 'template': 'Design SA test',
    })

    assert changed.status_code == 200, changed.text
    definition_payload = changed.json()['definition']
    assert changed.json()['invalidated'] is True
    assert definition_payload['nr_mode'] == 'SA'
    assert 'technology' not in definition_payload and 'template_technology' not in definition_payload
    assert definition_payload['template'] == 'Design SA test'
    for field in ('scope', 'custom_fields', 'hidden_filters', 'date_from', 'date_to', 'slide_comments'):
        assert definition_payload[field] == saved_before_change[field]
    # The campaigns of NSA do not apply to SA: the newest one of SA does (Multivendor Operator_Vendor).
    assert definition_payload['filters'] == {'City': ['London'], 'Campaign': ['@latest-campaign']}
    assert 'datasets' not in definition_payload

    core.repository.add_report_template('Design SA copy', core.repository.report_template_content('Design test',
    ), is_default=False)
    client.put('/api/ppt-designer/change-template', json={**definition_payload, 'filters': {'Campaign': ['UK_Q3_SA_2026']}})
    same_mode = client.patch('/api/ppt-designer/change-template/template', json={
        'technology': 'sa', 'template': 'Design SA copy',
    }).json()['definition']
    # Changing only the template keeps the Campaign.
    assert same_mode['filters'] == {'Campaign': ['UK_Q3_SA_2026']}


def test_design_vendor_comparison_defaults_and_persists_vendor_only(client):
    payload = setup_design(client)
    assert payload['vendor_comparison'] == 'operator_vendor'
    payload.update(scope='multivendor', vendor_comparison='vendor_only')

    response = client.put('/api/ppt-designer/vendor-pooling', json=payload)

    assert response.status_code == 200, response.text
    assert response.json()['definition']['vendor_comparison'] == 'vendor_only'
    saved = json.loads(core.repository.get_workspace_state('ppt_designs_v2'))
    assert saved['vendor-pooling']['vendor_comparison'] == 'vendor_only'


def test_vendor_comparison_prompt_uses_confirm_and_secondary_choices():
    root = Path(__file__).parents[1]
    design_script = (root / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')
    app_script = (root / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')

    assert "confirmLabel: 'Vendor (All Operators Combined)', secondaryLabel: 'Operator_Vendor'" in design_script
    assert "if (choice !== 'confirm' && choice !== 'secondary') return;" in design_script
    assert "exportDefinition.vendor_comparison = choice === 'secondary' ? 'operator_vendor' : 'vendor_only';" in design_script
    assert "const handleAccept = () => close(hasAlternatives ? 'confirm' : true);" in app_script
    assert "const handleSecondary = () => close('secondary');" in app_script


def test_design_ppt_job_vendor_comparison_survives_retry_without_manifest(client, monkeypatch):
    payload = setup_design(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    payload.update(scope='multivendor', vendor_comparison='vendor_only')
    design_id = 'vendor-only-ppt-retry'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    monkeypatch.setattr(core, 'submit_background_task', lambda *_args, **_kwargs: None)

    queued = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={'definition': payload})
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    with core.repository.connection() as connection:
        row = connection.execute(
            'SELECT vendor_comparison, output_path FROM design_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()
        connection.execute('UPDATE design_ppt_jobs SET status = ? WHERE id = ?', ('failed', job_id))
    assert row['vendor_comparison'] == 'vendor_only'
    assert not (Path(row['output_path']).parent / 'design-charts' / 'manifest.json').exists()

    retried = client.post(f'/api/ppt-designer/ppt-jobs/{job_id}/retry')
    assert retried.status_code == 202, retried.text
    with core.repository.connection() as connection:
        restored = connection.execute(
            'SELECT vendor_comparison, status FROM design_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()
    assert restored['vendor_comparison'] == 'vendor_only'
    assert restored['status'] == 'queued'


def test_design_library_ppt_scope_builds_its_automatic_dataset_universe(client):
    payload = setup_design(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'], 'City': ['London', 'Leeds', 'London'],
        'Mean_Data_Rate': [10, 20, 30], 'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
        'Event_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    design_id = 'library-ppt-scope'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200

    scope_only = {**payload, 'scope': 'single', 'datasets': {}, 'date_from': None, 'date_to': None}
    prepared = client.post(
        f'/api/ppt-designer/prepare?design_id={design_id}&use_scope_universe=1', json=scope_only,
    )

    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['rows']['data'] == 3
    assert prepared.json()['date_bounds'] == {'min': '2026-09-01', 'max': '2026-09-03'}


def test_design_library_geography_options_are_loaded_in_one_request(client):
    payload = setup_design(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'

    response = client.post('/api/ppt-designer/geography-options', json=payload)

    assert response.status_code == 200, response.text
    assert response.json() == {
        'campaigns': [], 'campaigns_pending': False,
        'operators': ['A', 'B'], 'vendors': ['A - All', 'B - All'], 'regions': [], 'cities': ['Leeds', 'London'],
    }
    # The export dialog and the Design filters list every campaign of the NR Mode and the CDRs they choose,
    # with their stage and campaigns, without preparing the Design.
    automatic = client.post('/api/ppt-designer/auto-cdrs', json=payload)
    assert automatic.status_code == 200, automatic.text
    assert [(row['name'], row['stage']) for row in automatic.json()['cdrs']['data']] == [('sample.csv', 'final')]
    assert automatic.json()['cdrs']['voice'] == [] and automatic.json()['options']['Campaign'] == []
    assert automatic.json()['options']['Operator'] == ['A', 'B']


def test_design_filter_options_batch_returns_every_requested_field(client):
    payload = setup_design(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    # Each filter lists the values the other filters leave (Leeds leaves Operator B); its own choice does not
    # narrow its list.
    payload['filters'] = {'City': ['Leeds']}

    response = client.post('/api/ppt-designer/filter-options/batch', json={
        'definition': payload, 'fields': ['Operator', 'City'],
    })

    assert response.status_code == 200, response.text
    assert response.json()['options'] == {'Operator': ['B'], 'City': ['Leeds', 'London']}


def test_design_campaigns_keep_their_nr_mode_suffix(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('campaigns.csv', BytesIO(
            b'Campaign,Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\n'
            b'NetCheck_UK_2026_Q2_SA,A,London,10,HTTP DL,2026-09-01\n'
            b'NetCheck UK 2026 Q1,B,Leeds,20,HTTP DL,2026-09-02\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('Design test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Comparison,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    payload = definition().model_dump(mode='json')

    options = client.post('/api/ppt-designer/geography-options', json=payload)

    assert options.status_code == 200, options.text
    # Compact labels keep an SA/NSA suffix; a Campaign filter limits them.
    assert options.json()['campaigns'] == ['26-Q1', '26-Q2-SA']
    payload['filters'] = {'Campaign': ['NetCheck_UK_2026_Q2_SA']}
    filtered = client.post('/api/ppt-designer/geography-options', json=payload)
    assert filtered.json()['campaigns'] == ['26-Q2-SA']

    # The PPT cover lists the Campaigns above the Scope in a larger font.
    core.repository.add_report_template('Campaign cover', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Quarterly review,Template subtitle,Title Page,,,,Title Slide,,,,,Top\n'
    ).encode(), is_default=False)
    payload = {**definition().model_dump(mode='json'), 'template': 'Campaign cover'}
    assert client.put('/api/ppt-designer/campaign-cover', json=payload).status_code == 200
    queued = client.post('/api/ppt-designer/campaign-cover/export-ppt', json={'definition': payload})
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.05)
    assert job['status'] == 'ready', job
    assert job['cover']['campaigns'] == 'Campaigns: 26-Q1, 26-Q2-SA'
    # The job Filters list the Campaigns right above the Operators.
    assert job['filters'].index('Campaigns: 26-Q1, 26-Q2-SA') + 1 == job['filters'].index('Operator: All Operators')
    with core.repository.connection() as connection:
        row = connection.execute('SELECT output_path FROM design_ppt_jobs WHERE id = ?', (job_id,)).fetchone()
    # Exports comparing several Campaigns name them at the end of the file name.
    assert row['output_path'].endswith(' - Operator Comparison - 26-Q1_vs_26-Q2-SA.pptx')
    slide = Presentation(row['output_path']).slides[0]
    details = {shape.name: shape for shape in slide.shapes if shape.name.startswith('design-ppt-')}
    campaigns, scope = details['design-ppt-campaigns'], details['design-ppt-scope']
    assert campaigns.text == 'Campaigns: 26-Q1, 26-Q2-SA'
    # Campaigns sit between the subtitle and the template's decorative line,
    # while the Scope and geography stay below that line.
    subtitle = next(shape for shape in slide.placeholders if shape.placeholder_format.type == 4)
    separator = next(shape for shape in slide.slide_layout.shapes if shape.shape_type == MSO_SHAPE_TYPE.LINE)
    assert subtitle.top + subtitle.height <= campaigns.top
    assert campaigns.top + campaigns.height <= separator.top < scope.top
    assert campaigns.text_frame.paragraphs[0].runs[0].font.size > scope.text_frame.paragraphs[0].runs[0].font.size
    # White on the red of the cover, its label bold.
    assert all(run.font.color.rgb == RGBColor(255, 255, 255) for run in campaigns.text_frame.paragraphs[0].runs)
    assert campaigns.text_frame.paragraphs[0].runs[0].font.bold


def test_design_filter_catalogue_includes_values_beyond_legacy_profile_limit(client, monkeypatch):
    setup_design(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    cities = [f'City {index:03d}' for index in range(225)]
    rows = '\n'.join(
        f'A,{city},{index + 1},HTTP DL,2026-09-01'
        for index, city in enumerate(cities)
    )
    csv_content = ('Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\n' + rows + '\n').encode()
    uploaded = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('many-cities.csv', BytesIO(csv_content), 'text/csv'),
    })
    assert uploaded.status_code == 200
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(2)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(2)['status'] == 'ready'
    processed_options = json.loads(core.repository.get_dataset(2)['filter_options_json'])
    assert len(processed_options['city']) == 225
    assert core.repository.list_distinct_dataset_row_values(2, 'City', limit=None) == cities
    core.repository.replace_reporting_rows(2, 'data', pd.read_csv(BytesIO(csv_content)))
    with core.repository.connection() as connection:
        connection.execute(
            'UPDATE dataset_profiles SET filter_options_json = ? WHERE dataset_id = 2',
            (json.dumps({'city': cities[:50]}),),
        )
    monkeypatch.setattr('src.modules.ppt_designer.DESIGN_PROFILE_SELECTION_THRESHOLD', 1)
    payload = definition().model_dump(mode='json')
    payload['datasets'] = {'data': [2]}

    prepared = client.post('/api/ppt-designer/prepare', json=payload)

    assert prepared.status_code == 200, prepared.text
    # Every CDR of the NR Mode offers its Cities, beyond the former profile limit.
    assert set(cities) <= set(prepared.json()['options']['City'])


def test_design_ppt_filename_summarizes_complete_geography_selections(client, monkeypatch):
    payload = setup_design(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        if core.repository.get_dataset(1)['status'] == 'ready':
            break
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    design_id = 'geography-ppt-name'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=['North', 'South'], cities=['Leeds', 'London'],
    )
    # Labels and file names are resolved by the background job; run it inline.
    monkeypatch.setattr(core, 'submit_background_task', lambda callback, *args: callback(*args))

    def queued_name(regions, cities, operators=None):
        filters = {'Region': regions, 'City': cities}
        if operators is not None:
            filters['Operator'] = operators
        export_definition = {**payload, 'filters': filters}
        response = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
            'definition': export_definition,
            'selected_regions': regions,
            'selected_cities': cities,
        })
        assert response.status_code == 202, response.text
        with core.repository.connection() as connection:
            row = connection.execute(
                'SELECT output_file FROM design_ppt_jobs WHERE id = ?',
                (response.json()['job_id'],),
            ).fetchone()
        return row['output_file']

    assert re.fullmatch(
        r'\d{8}_\d{6} - PPT Designer - NSA - Comparison - Operator Comparison - All Regions\.pptx',
        queued_name(['South', 'North'], ['London', 'Leeds']),
    )
    first_north_export = queued_name(['North'], ['London'])
    assert re.fullmatch(
        r'\d{8}_\d{6} - PPT Designer - NSA - Comparison - Operator Comparison - North\.pptx',
        first_north_export,
    )
    second_north_export = queued_name(['North'], ['London'], ['A'])
    assert re.fullmatch(
        r'\d{8}_\d{6} - PPT Designer - NSA - Comparison - Operator Comparison - North\.pptx',
        second_north_export,
    )
    assert second_north_export != first_north_export


def test_design_ppt_filename_summarizes_partial_selections_and_fits_filesystem(client, monkeypatch):
    payload = setup_design(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    payload['name'] = 'Validation Design ' + 'Long Name ' * 9
    design_id = 'compact-ppt-name'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    core.repository.replace_cdr_catalogue(
        1,
        vendors=['Vendor A', 'Vendor B', 'Vendor C'],
        regions=['North', 'South', 'East'],
        cities=['Belfast', 'Bristol', 'Cardiff', 'Edinburgh', 'Leeds', 'London', 'Sheffield', 'York'],
    )
    monkeypatch.setattr(core, 'submit_background_task', lambda callback, *args: callback(*args))
    selections = {
        'selected_operators': ['A'],
        'selected_vendors': ['Vendor A', 'Vendor B'],
        'selected_regions': ['North', 'South'],
        'selected_cities': ['Belfast', 'Bristol', 'Cardiff', 'Edinburgh', 'Leeds', 'London', 'Sheffield'],
    }
    queued = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
        'definition': payload, **selections,
    })
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    with core.repository.connection() as connection:
        row = connection.execute(
            'SELECT output_file, output_path, filters_json FROM design_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()
    output_path = Path(row['output_path'])
    assert row['output_file'].endswith(
        ' - Operator Comparison - North + South.pptx'
    )
    assert len(output_path.parent.name.encode('utf-8')) <= 240
    assert len(output_path.name.encode('utf-8')) <= 255
    output_path.parent.mkdir(parents=True, exist_ok=True)
    assert 'Vendor: Vendor A, Vendor B' in json.loads(row['filters_json'])
    assert 'Region: North, South' in json.loads(row['filters_json'])
    assert 'City: Belfast, Bristol, Cardiff, Edinburgh, Leeds, London, Sheffield' in json.loads(row['filters_json'])

    with core.repository.connection() as connection:
        connection.execute('UPDATE design_ppt_jobs SET status = ? WHERE id = ?', ('failed', job_id))
    relaunched = client.post(f'/api/ppt-designer/ppt-jobs/{job_id}/retry')
    assert relaunched.status_code == 202, relaunched.text
    with core.repository.connection() as connection:
        retried = connection.execute(
            'SELECT output_file, filters_json FROM design_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()
    assert retried['output_file'].endswith(
        ' - Operator Comparison - North + South.pptx'
    )
    assert 'City: Belfast, Bristol, Cardiff, Edinburgh, Leeds, London, Sheffield' in json.loads(retried['filters_json'])

    unicode_definition = {**payload, 'name': 'É' * 120}
    unicode_job = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
        'definition': unicode_definition, **selections,
    })
    assert unicode_job.status_code == 202, unicode_job.text
    with core.repository.connection() as connection:
        unicode_row = connection.execute(
            'SELECT output_path FROM design_ppt_jobs WHERE id = ?', (unicode_job.json()['job_id'],),
        ).fetchone()
    unicode_path = Path(unicode_row['output_path'])
    assert len(unicode_path.parent.name.encode('utf-8')) <= 240
    unicode_path.parent.mkdir(parents=True, exist_ok=True)


def test_design_ppt_dialog_selections_override_saved_design_filters(client, monkeypatch):
    payload = setup_design(client)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    design_id = 'dialog-ppt-filters'
    payload['filters'] = {
        'Operator': ['A'], 'Vendor': ['Vendor A'],
        'Region': ['North'], 'City': ['London'],
    }
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=['North', 'South'], cities=['Leeds', 'London'],
    )
    submitted = []

    def run_inline(callback, *args):
        submitted.append(args)
        callback(*args)

    monkeypatch.setattr(core, 'submit_background_task', run_inline)

    queued = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
        'definition': payload,
        'selected_operators': ['B'],
        'selected_vendors': ['Vendor B'],
        'selected_regions': ['South'],
        'selected_cities': ['Leeds'],
    })

    assert queued.status_code == 202, queued.text
    job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == queued.json()['job_id'])
    assert job['filters'][-4:] == [
        'Operator: B', 'Vendor: Vendor B', 'Region: South', 'City: Leeds',
    ]
    assert job['cover']['regions'] == 'Region: South'
    assert job['cover']['cities'] == 'City: Leeds'
    with core.repository.connection() as connection:
        output_file = connection.execute(
            'SELECT output_file FROM design_ppt_jobs WHERE id = ?', (queued.json()['job_id'],),
        ).fetchone()['output_file']
    assert output_file.endswith(' - Operator Comparison - South.pptx')
    # The dialog selections travel with the job and override the saved filters.
    assert submitted[0][6] == {
        'Operator': ['B'], 'Vendor': ['Vendor B'], 'Region': ['South'], 'City': ['Leeds'],
    }
    saved_filters = client.get('/api/ppt-designer').json()[design_id]['filters']
    assert {field: values for field, values in saved_filters.items() if field != 'Campaign'} == payload['filters']


def test_design_ppt_all_labels_only_consider_the_selected_cdrs(client):
    payload = setup_design(client)
    # A CDR of the other NR Mode carries an extra Operator and City that the
    # Design's CDRs do not contain; its saved filters still list them.
    other = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('other_SA.csv', BytesIO(
            b'Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\nC,York,5,HTTP DL,2026-09-01\n'
        ), 'text/csv'),
    })
    assert other.status_code == 200
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and any(core.repository.get_dataset(item)['status'] != 'ready' for item in (1, 2)):
        time.sleep(0.05)
    payload['filters'] = {'Operator': ['A', 'B', 'C'], 'City': ['Leeds', 'London', 'York']}
    design_id = 'selected-cdr-labels'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200

    options = client.post('/api/ppt-designer/geography-options', json=payload).json()
    assert options['operators'] == ['A', 'B']
    # Selecting every value the selected CDRs offer reads as "All", whether
    # the export uses the saved filters or the dialog's explicit selection.
    for selections in ({}, {'selected_operators': ['A', 'B'], 'selected_cities': ['Leeds', 'London']}):
        queued = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={'definition': payload, **selections})
        assert queued.status_code == 202, queued.text
        job_id = queued.json()['job_id']
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline:
            job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
            if job['status'] in {'ready', 'failed'}:
                break
            time.sleep(0.05)
        assert 'Operator: All Operators' in job['filters'], job['filters']
        assert 'City: All Cities' in job['filters'], job['filters']


def test_design_ppt_cover_uses_scope_and_catalogue_geography(client):
    payload = setup_design(client)
    core.repository.add_report_template('Structural design', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Quarterly review,Template subtitle,Title Page,,,,Title Slide,,,,,Top\n'
        '2,,,Black logo end slide,,,,,,,,,\n'
    ).encode(), is_default=False)
    payload['template'] = 'Structural design'
    design_id = 'structural-ppt-cover'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=['North', 'South'], cities=['Leeds', 'London'],
    )

    queued = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={'definition': payload})
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    while time.monotonic() < deadline:
        job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.05)
    assert job['status'] == 'ready', job
    assert 'Operator: All Operators' in job['filters']
    assert 'Vendor: All Vendors' in job['filters']
    assert 'Region: All Regions' in job['filters']
    assert 'City: All Cities' in job['filters']
    # Jobs that stored no Operator or Vendor line were not restricted by them.
    stored = [line for line in job['filters'] if not line.startswith(('Operator:', 'Vendor:'))]
    with core.repository.connection() as connection:
        connection.execute(
            'UPDATE design_ppt_jobs SET filters_json = ? WHERE id = ?', (json.dumps(stored), job_id),
        )
    legacy = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
    assert legacy['filters'].index('Operator: All Operators') < legacy['filters'].index('Vendor: All Vendors')
    assert legacy['filters'].index('Vendor: All Vendors') < legacy['filters'].index('Region: All Regions')
    with core.repository.connection() as connection:
        row = connection.execute('SELECT output_path FROM design_ppt_jobs WHERE id = ?', (job_id,)).fetchone()
    slide, end = Presentation(row['output_path']).slides
    # The End Slide is the closing slide of the template as it is, without a counter.
    assert end.slide_layout.name == 'Black logo end slide'
    assert not any(shape.has_text_frame and shape.text_frame.text.strip() for shape in end.shapes)
    placeholders = {shape.placeholder_format.type: shape for shape in slide.placeholders}
    title = placeholders.get(1) or placeholders[3]
    assert title.text == 'Quarterly review'
    assert placeholders[4].text == 'Template subtitle'
    details = {shape.name: shape for shape in slide.shapes if shape.name.startswith('design-ppt-')}
    # The Scope is listed above the Regions and Cities, white on the red of the cover.
    assert details['design-ppt-scope'].text == 'Operator Comparison'
    assert details['design-ppt-region'].text == 'All Regions'
    assert details['design-ppt-city'].text == 'All Cities'
    assert details['design-ppt-scope'].top < details['design-ppt-region'].top < details['design-ppt-city'].top
    assert details['design-ppt-scope'].left == details['design-ppt-region'].left == title.left
    assert details['design-ppt-city'].left == title.left
    assert details['design-ppt-scope'].text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(255, 255, 255)
    # The job's viewer snapshot shows exactly the same cover lines.
    # This CDR has no Campaign column, so no Campaigns line is drawn.
    assert 'design-ppt-campaigns' not in details
    expected_cover = {'campaigns': '', 'scope': 'Operator Comparison', 'regions': 'All Regions', 'cities': 'All Cities'}
    assert client.get(f'/api/ppt-designer/ppt-jobs/{job_id}/charts.json').json()['cover'] == expected_cover
    listed_job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
    assert listed_job['cover'] == expected_cover
    manifest_path = Path(row['output_path']).parent / 'design-charts' / 'manifest.json'
    legacy = json.loads(manifest_path.read_text(encoding='utf-8'))
    legacy.pop('cover')
    manifest_path.write_text(json.dumps(legacy), encoding='utf-8')
    # Jobs exported before the cover was stored rebuild it from their Filters.
    assert client.get(f'/api/ppt-designer/ppt-jobs/{job_id}/charts.json').json()['cover'] == expected_cover
    assert next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)['cover'] == expected_cover
    # Like the viewer, the details follow the subtitle (below the template line).
    assert details['design-ppt-scope'].top > placeholders[4].top + placeholders[4].height
    assert placeholders[4].text_frame.paragraphs[0].runs[0].font.color.theme_color == MSO_THEME_COLOR.ACCENT_4
    # The cover subtitle does not inherit the layout's condensed letter spacing.
    assert placeholders[4].text_frame.paragraphs[0].runs[0].font._rPr.get('spc') == '0'
    assert details['design-ppt-region'].text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(255, 255, 255)


def test_expanded_design_chart_apply_builds_a_new_temporary_model(client):
    payload = setup_design(client)
    prepared = client.post('/api/ppt-designer/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    token = prepared.json()['token']

    original = client.get(f'/api/ppt-designer/chart/{token}/0')
    assert original.status_code == 200, original.text
    assert original.headers['cache-control'] == 'no-store'
    assert len(original.json()['series']) == 2
    context = client.get(f'/api/ppt-designer/chart/{token}/0/filter-context')
    assert context.status_code == 200, context.text
    assert context.json()['kpi'] == 'Mean_Data_Rate'
    assert context.json()['exclude_null_empty'] is False
    assert context.json()['exclude_zero'] is False

    preview = client.post(f'/api/ppt-designer/chart/{token}/0/filter-preview', json={
        'chart_title': 'Filtered operator preview',
        'filters': 'Operator = A',
        'rows_split': 'Test_Name × City × Operator',
        'columns_split': '',
        'legend': 'Operator',
        'legend_position': 'Right',
    })
    assert preview.status_code == 200, preview.text
    assert preview.json()['title'] == 'Filtered operator preview'
    assert len(preview.json()['series']) == 1
    assert preview.json()['series'][0]['name'] == 'HTTP DL · London · A'
    assert preview.json()['legend']['position'] == 'right'

    retained_kpi = client.post(f'/api/ppt-designer/chart/{token}/0/filter-preview', json={
        'kpi': '', 'rows_split': 'Operator', 'columns_split': '',
    })
    assert retained_kpi.status_code == 200, retained_kpi.text
    assert retained_kpi.json()['metric'] == 'Mean Data Rate'

    updated = client.post(f'/api/ppt-designer/chart/{token}/0/update-template', json={
        'chart_title': 'Updated template chart', 'rows_split': 'Test_Name × Operator',
        'columns_split': 'City', 'legend_position': 'Right',
        'exclude_null_empty': 'Yes', 'exclude_zero': 'Yes',
    })
    assert updated.status_code == 200, updated.text
    entry = core.load_template_catalogue(core.repository.report_template_content('Design test'))[0]
    assert (entry.chart_title, entry.rows_split, entry.columns_split, entry.legend_position) == (
        'Updated template chart', 'Test_Name × Operator', 'City', 'right',
    )
    assert entry.exclude_null_empty is True
    assert entry.exclude_zero is True
    refreshed_context = client.get(f'/api/ppt-designer/chart/{token}/0/filter-context')
    assert refreshed_context.status_code == 200, refreshed_context.text
    assert refreshed_context.json()['chart_title'] == 'Updated template chart'
    assert refreshed_context.json()['rows_split'] == 'Test_Name × Operator'
    assert refreshed_context.json()['exclude_null_empty'] is True
    assert refreshed_context.json()['exclude_zero'] is True
    assert updated.json()['updated_at']


def test_expanded_design_chart_apply_can_render_outside_the_proxy_request(client):
    payload = setup_design(client)
    prepared = client.post('/api/ppt-designer/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    token = prepared.json()['token']

    queued = client.post(
        f'/api/ppt-designer/chart/{token}/0/filter-preview?background=true',
        json={'filters': 'Operator = A', 'rows_split': 'Operator'},
    )

    assert queued.status_code == 202, queued.text
    assert queued.json()['state'] == 'processing'
    completed = client.get(
        f"/api/ppt-designer/chart-preview-jobs/{queued.json()['job_id']}"
    )
    assert completed.status_code == 200, completed.text
    assert completed.json()['series'][0]['name'] == 'A'


def test_update_template_synchronizes_chart_definition_across_live_snapshots(client):
    payload = setup_design(client)
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'], 'City': ['London', 'Leeds', 'London'],
        'Mean_Data_Rate': [10, 20, 30], 'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
        'Event_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    first = client.post('/api/ppt-designer/prepare', json=payload)
    second = client.post('/api/ppt-designer/prepare', json=payload)
    assert first.status_code == second.status_code == 200
    first_token = first.json()['token']
    second_token = second.json()['token']

    filtered = client.post(
        f'/api/ppt-designer/chart/{first_token}/0/update-template',
        json={'filters': 'Operator = A'},
    )
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()['synced_snapshots'] == 2
    second_context = client.get(
        f'/api/ppt-designer/chart/{second_token}/0/filter-context'
    )
    assert second_context.status_code == 200, second_context.text
    assert second_context.json()['filters'] == 'Operator = A'

    cleared = client.post(
        f'/api/ppt-designer/chart/{first_token}/0/update-template',
        json={'filters': ''},
    )
    assert cleared.status_code == 200, cleared.text
    reopened_context = client.get(
        f'/api/ppt-designer/chart/{second_token}/0/filter-context'
    )
    assert reopened_context.status_code == 200, reopened_context.text
    assert reopened_context.json()['filters'] == ''


def test_design_refresh_rebuilds_all_models_and_chart_refresh_rebuilds_only_one(client, monkeypatch):
    payload = setup_design(client)
    preview = client.post('/api/ppt-designer/prepare', json=payload).json()
    token = preview['token']
    indexes = [
        chart['index']
        for slide in preview['slides']
        for chart in slide['charts']
        if chart['available']
    ]
    for index in indexes:
        assert client.get(f'/api/ppt-designer/chart/{token}/{index}').status_code == 200

    import src.modules.ppt_designer as designs_module
    original = designs_module.catalog_chart_payload
    calls = []

    def tracked(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(designs_module, 'catalog_chart_payload', tracked)

    current = indexes[1]
    core.repository.replace_operator_mapping_group(None, 'Alpha', ['A'], '#123456')
    core.repository.replace_operator_mapping_group(None, 'Beta', ['B'], '#654321')
    core.repository.move_chart_mapping_group('operator', 'Beta', 'up')
    refreshed_chart = client.post(f'/api/ppt-designer/chart/{token}/{current}/refresh')
    assert refreshed_chart.status_code == 200, refreshed_chart.text
    assert len(calls) == 1
    assert [series['name'] for series in refreshed_chart.json()['series']] == ['Beta', 'Alpha']
    assert {
        series['name']: series['colour'] for series in refreshed_chart.json()['series']
    } == {'Alpha': '#123456', 'Beta': '#654321'}
    assert client.get(f'/api/ppt-designer/chart/{token}/{indexes[0]}').status_code == 200
    assert len(calls) == 1

    calls.clear()
    refreshed_design = client.post(f'/api/ppt-designer/charts/{token}/refresh')
    assert refreshed_design.status_code == 200, refreshed_design.text
    assert refreshed_design.json() == {'refreshed': len(indexes)}
    assert len(calls) == len(indexes)


def test_design_editing_controls_match_workspace_editor_roles(client):
    controls = (
        'id="ds-edit"',
        'data-workspace-manage-calculated-dimensions>Auto-Calculated Fields',
        'id="ds-chart-expanded-edit"',
        'id="ds-chart-expanded-auto-fields"',
    )
    for role in ('user-viewer', 'user-editor', 'admin', 'super-admin'):
        session_id = f'design-controls-{role}'
        if role != 'super-admin':
            core.repository.create_user(session_id, 'test-password', role)
            account = next(row for row in core.repository.list_users() if row['username'] == session_id)
            core.repository.set_user_workspace_access(int(account['id']), [core.active_workspace.id])
        core.SESSIONS[session_id] = core.SessionUser(username=session_id, role=role)
        client.cookies.set(core.SESSION_COOKIE, session_id)
        page = client.get('/ppt-designer')
        assert page.status_code == 200
        for control in controls:
            assert (control in page.text) == (role != 'user-viewer')
        assert ('id="ds-ppt-jobs-delete-all"' in page.text) == (role != 'user-viewer')
        assert ('"can_manage": true' in page.text) == (role != 'user-viewer')
        expected_status = 403 if role == 'user-viewer' else 404
        assert client.post('/api/ppt-designer/ppt-jobs/999999/delete').status_code == expected_status
        expected_chart_status = 403 if role == 'user-viewer' else 410
        assert client.post('/api/ppt-designer/chart/missing/0/update-template', json={}).status_code == expected_chart_status
        if role != 'user-viewer':
            assert client.post('/api/ppt-designer/ppt-jobs/delete-all').status_code == 200


def test_designs_lifecycle_and_layout(client):
    payload = setup_design(client)
    without_starter_designs()
    page = client.get('/ppt-designer')
    assert page.status_code == 200
    assert re.search(r'data-authenticated-session="[^"]+"', page.text)
    assert 'id="page-panel-navigator"' in page.text
    assert 'Modules &amp; Sections' in page.text and 'class="module-tree-sections"' in page.text
    assert page.text.index('>CDR Analysis<') < page.text.index('>PPT Designer<') < page.text.index('>Network Insights<')
    assert '>PPT Reporting (old)<' not in page.text
    assert 'id="ds-nr-mode"' in page.text
    assert 'id="ds-designs-body"' in page.text
    assert '>PPT Generation Jobs<' in page.text
    assert 'id="ds-ppt-jobs-body"' in page.text
    assert '>Charts Panel<' in page.text
    assert 'id="ds-ppt-chart-job"' in page.text
    assert 'id="ds-ppt-charts-body"' in page.text
    assert 'id="ds-ppt-chart-viewer"' not in page.text
    assert page.text.count('id="ds-chart-expanded-overlay"') == 1
    assert page.text.index('<th>Status</th>') < page.text.index('<th>Actions</th>')
    assert 'colspan="5" class="form-note">Loading Designs' in page.text
    assert 'id="ds-library"' not in page.text
    assert 'id="ds-save"' in page.text
    assert '>Save Filters<' in page.text
    assert page.text.index('id="ds-unapplied-filters-badge"') < page.text.index('id="ds-unsaved-filters-badge"')
    assert 'id="confirm-secondary"' in page.text
    assert 'id="confirm-tertiary"' in page.text
    assert '"filter_aliases"' in page.text
    assert '"Region": ["Region", "G_Level_2", "G Level 2"]' in page.text
    assert '>Apply Filters<' in page.text
    # There is no Dataset Universe: the Campaign filter chooses the CDRs, listed read-only.
    assert 'id="ds-apply-universe"' not in page.text and 'id="ds-save-universe"' not in page.text
    assert 'id="ds-auto-cdrs"' in page.text and '>CDRs Used<' in page.text
    assert page.text.index('>Apply Filters<') < page.text.index('>Save Filters<')
    assert 'id="ds-apply-filters" title="Apply the current PPT filters without saving them" disabled' in page.text
    assert 'id="ds-refresh"' not in page.text
    assert '>Import PPT Design<' in page.text
    assert 'Total Designs: 0' in page.text
    assert '>PPT Scope & Filters<' in page.text
    assert 'id="ds-active-design-heading">PPT Filters<' in page.text
    filter_panel = page.text.split('id="ds-filter-panel"', 1)[1].split('>', 1)[0]
    assert ' open' in filter_panel
    assert ' hidden' in filter_panel
    assert 'data-panel-state-storage="session"' in filter_panel
    assert 'id="ds-ppt-dataset-overlay"' in page.text
    assert 'id="ds-ppt-dataset-choices"' not in page.text
    assert 'id="ds-ppt-dataset-scope"' in page.text
    assert '>Select PPT Scope and Filters<' in page.text
    # The PPT dialog groups the Scope above the PPT Filters: the Campaign (which chooses the CDRs) first, then the
    # dates, which are filters, then the other filters.
    universe_group = page.text.index('class="ds-ppt-group ds-ppt-universe-group"')
    filters_group = page.text.index('class="ds-ppt-group ds-ppt-filters-group"')
    assert universe_group < filters_group < page.text.index('id="ds-ppt-campaign-section"') < page.text.index('id="ds-ppt-dates-section"')
    assert page.text.index('id="ds-ppt-dates-section"') < page.text.index('id="ds-ppt-operator-section"') < page.text.index('id="ds-ppt-city-section"')
    # Main PPT Filters can reveal every other filter, including Additional Filters.
    assert '>Main PPT Filters</h3><button type="button" id="ds-ppt-show-all-filters"' in page.text
    # The heading names the Design's NR Mode and the Design.
    assert '<span class="ds-ppt-eyebrow-part" id="ds-ppt-dataset-nr-mode"></span><span class="ds-ppt-eyebrow-part ds-ppt-eyebrow-design" id="ds-ppt-dataset-design"></span><span class="ds-ppt-eyebrow-part">PowerPoint export</span>' in page.text
    assert page.text.index('id="ds-ppt-city-section"') < page.text.index('id="ds-ppt-extra-main-filters"') < page.text.index('id="ds-ppt-additional-group"')
    # Save Filters sits between Cancel and Generate PPT.
    assert 'id="ds-ppt-save-universe"' not in page.text
    actions = [page.text.index(f'id="{name}"') for name in ('ds-ppt-dataset-cancel', 'ds-ppt-save-filters', 'ds-ppt-dataset-confirm')]
    assert actions == sorted(actions)
    dialog_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')
    assert "return [key, saved?.length ? {values: saved, all: false} : {values: [], all: true}];" in dialog_script
    assert "options = (await api('/filter-options/batch', 'POST', {definition: exportUniverse(), fields: [...extraFilters.keys()]})).options || {};" in dialog_script
    assert "for (const [field, values] of Object.entries(universeChoice.extra_filters || {})) withPptSelection(exportDefinition, field, values);" in dialog_script
    assert "saveFiltersButton.disabled = busy || (!universeChanged() && !Object.keys(changedFilterFields()).length);" in dialog_script
    # An opened selector menu is fixed below its field and may overflow the dialog.
    assert "menu.style.top = `${shellBounds.bottom + gap - origin.top}px`;" in dialog_script
    assert 'id="ds-ppt-date-from"' in page.text
    assert 'id="ds-ppt-date-to"' in page.text
    assert '>Dataset Universe<' not in page.text
    assert '>Comparison &amp; CDRs<' in page.text
    assert '>Select PPT Filters<' in page.text
    assert page.text.index('>Comparison &amp; CDRs<') < page.text.index('>PPT Scope<')
    assert page.text.index('>Select PPT Filters<') < page.text.index('>Default Filters<')
    assert '>PPT Scope<' in page.text
    assert '>Select Comparison Scope<' in page.text
    # The Date filters open the Default Filters.
    assert page.text.index('class="ds-scope-control"') < page.text.index('>Default Filters<') < page.text.index('id="ds-sources"') < page.text.index('id="ds-default-facets"')
    assert '>Default Filters<' in page.text
    assert '>Additional Filters<' in page.text
    assert '>Clear Filters<' in page.text
    assert '>Reload Saved Filters<' in page.text
    assert 'id="ds-design-name"' not in page.text
    assert 'title="Save and apply the current PPT filters"' in page.text
    assert 'title="Remove all filter restrictions"' in page.text
    assert 'title="Restore filters and additional fields from the last saved Design"' in page.text
    assert 'id="ds-preparing-rows"' in page.text
    assert 'class="form-note ds-filter-help"' in page.text
    assert 'id="ds-default-facets"' in page.text
    assert 'id="ds-additional-facets"' in page.text
    assert 'Select field to add new filter' in page.text
    assert 'id="ds-custom-field" multiple size="1" data-multiselect-single="true"' in page.text
    assert 'hidden_filters' in DesignDefinition.model_fields
    assert 'slide_comments' in DesignDefinition.model_fields
    assert 'id="ds-view" class="ds-view-design-action" title="Open the PPT viewer" disabled' in page.text
    assert 'id="ds-generate-ppt" class="ds-generate-ppt-action" title="Generate a PowerPoint presentation for this Design" disabled' in page.text
    assert page.text.index('id="ds-generate-ppt"') < page.text.index('id="ds-view"')
    assert 'id="ds-preparing"' in page.text
    assert 'id="ds-preparing-title"' in page.text
    assert '>Preparing PPT dataset<' in page.text
    assert 'id="ds-preparing-progress"' in page.text
    assert 'id="ds-preparing-progress-bar"' in page.text
    assert 'id="ds-preparing-progress-label">0%</span>' in page.text
    assert 'id="ds-preparing-refresh"' in page.text
    assert 'id="ds-preparing-out-of-sync"' in page.text
    assert 'id="ds-viewer-preparing"' in page.text
    assert 'id="ds-viewer-preparing-progress"' in page.text
    assert 'id="ds-viewer-preparing-detail"' in page.text
    assert '>Auto-Calculated Fields<' in page.text
    assert 'class="ds-viewer-icon-action ds-viewer-refresh-action"' in page.text
    assert 'class="ds-viewer-tool-actions" role="group" aria-label="PPT actions"' in page.text
    assert 'id="ds-viewer-close" class="report-chart-viewer-close" title="Close"' in page.text
    assert page.text.index('id="ds-floating-filters"') < page.text.index('id="ds-viewer-export-ppt"')
    assert page.text.index('id="ds-viewer-export-ppt"') < page.text.index('id="ds-viewer-refresh"')
    assert page.text.index('id="ds-viewer-refresh"') < page.text.index('id="ds-presentation"')
    assert 'id="ds-viewer-export-ppt" class="ds-viewer-top-action ds-viewer-top-ppt"' in page.text
    assert '>Generate PPT</button>' in page.text
    assert 'id="ds-viewer-refresh" class="ds-viewer-top-action ds-viewer-top-refresh"' in page.text
    assert '>Refresh Design</button>' in page.text
    assert 'id="ds-floating-filters" class="ds-viewer-top-action ds-viewer-top-filters"' in page.text
    assert 'class="ds-viewer-top-action ds-viewer-top-auto-fields" data-workspace-manage-calculated-dimensions' in page.text
    assert 'id="ds-edit" class="ds-viewer-top-action ds-viewer-top-edit"' in page.text
    assert page.text.index('id="ds-presentation"') < page.text.index('id="ds-first"')
    assert '>PPT Filters</button>' in page.text
    assert 'id="ds-prev" class="ds-slide-nav-button" title="Previous slide" aria-label="Previous slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-next" class="ds-slide-nav-button" title="Next slide" aria-label="Next slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-first" class="ds-slide-nav-button" title="First slide" aria-label="First slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-last" class="ds-slide-nav-button" title="Last slide" aria-label="Last slide"><svg class="ds-slide-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-overlay"' in page.text
    assert 'class="ds-chart-expanded-meta" id="ds-chart-expanded-meta"' in page.text
    assert 'id="ds-chart-expanded-close" class="report-chart-viewer-close" title="Close"' in page.text
    assert 'id="ds-chart-expanded-canvas"' in page.text
    assert 'id="ds-chart-expanded-data"' in page.text
    assert 'id="ds-chart-expanded-zoom"' in page.text
    assert 'id="ds-chart-expanded-filters"' in page.text
    assert 'id="ds-chart-expanded-edit"' in page.text
    assert 'id="ds-chart-expanded-first"' in page.text
    assert 'id="ds-chart-expanded-last"' in page.text
    assert 'id="ds-chart-expanded-pan-left"' in page.text
    assert 'id="ds-chart-expanded-pan-right"' in page.text
    assert 'id="ds-chart-expanded-first" title="First chart" aria-label="First chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-prev" title="Previous chart" aria-label="Previous chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-next" title="Next chart" aria-label="Next chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert 'id="ds-chart-expanded-last" title="Last chart" aria-label="Last chart"><svg class="ds-chart-expanded-arrow-icon"' in page.text
    assert '⏮' not in page.text
    assert '⏭' not in page.text
    assert 'id="ds-chart-expanded-canvas-shell"' in page.text
    assert 'id="ds-chart-expanded-position"' in page.text
    assert 'id="ds-data-table"' in page.text
    assert '<span>Export CSV</span>' in page.text
    assert 'class="ds-chart-expanded-footer"' in page.text
    assert 'id="ds-chart-expanded-controls"' in page.text
    assert 'id="ds-chart-expanded-controls"><button type="button" id="ds-chart-expanded-data"' in page.text
    assert 'data-design-ppt-chart-filter="nr_mode"' in page.text
    assert 'data-design-ppt-chart-filter="design"' in page.text
    assert 'data-design-ppt-chart-filter="template"' in page.text
    assert 'data-design-ppt-chart-filter="scope"' in page.text
    assert 'data-design-ppt-chart-filter="region"' in page.text
    assert 'data-design-ppt-chart-filter="city"' in page.text
    assert 'id="ds-ppt-charts-filters"' in page.text
    assert '>View Filters</button>' in page.text
    assert 'id="ds-ppt-chart-job-picker"' in page.text
    assert 'id="ds-multivendor-overlay"' not in page.text
    assert '>Keep Current Selection</button>' not in page.text
    assert '>Use Selected Datasets</button>' not in page.text
    assert '>Use Latest Datasets</button>' not in page.text
    assert 'ds-viewer-refresh-action' in page.text
    design_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')
    assert "controls.append(data, expand, zoom)" in design_script
    assert "savedDefinition = definitionFingerprint(savedRuntimeDesignDefinition(designs[id])); dirty = false;" in design_script
    assert "await warmDesignModels();" not in design_script
    assert "Rendering PPT Charts" in design_script
    assert "overlay('ds-viewer', true);\n    if (needsPreparation) await prepare();" in design_script
    assert "window.addEventListener('beforeunload', rememberScroll);" in design_script
    assert "document.addEventListener('visibilitychange'" in design_script
    assert "document.documentElement.scrollHeight - window.innerHeight" in design_script
    assert "const restorePageState = navigationEntry?.type === 'reload';" in design_script
    assert "if (!restorePageState) resetScroll();" in design_script
    assert "if (restorePageState) restoreScroll(); else resetScroll();" in design_script
    assert "sessionStorage.removeItem(scrollStorageKey);" in design_script
    assert "last = sessionStorage.getItem(openStorageKey) || '';" in design_script
    assert 'const restoreOpenDesign = restorePageState;' not in design_script
    assert "if (designs[last]) await openDesign(last);" in design_script
    assert "preview_snapshot = replace(" in (Path(__file__).parents[1] / 'src/modules/ppt_designer.py').read_text(encoding='utf-8')
    assert "The template owns these required chart attributes." in (Path(__file__).parents[1] / 'src/modules/ppt_designer.py').read_text(encoding='utf-8')
    app_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')
    assert 'function setupPagePanelNavigator()' in app_script
    # Modules & Sections: the drawer of every module with its sections (rendered by the page) in the colour of the module
    # of the page; choosing a module or a section closes it.
    assert "'module-tab-ppt-designer': 'designs'" in app_script
    assert "navigator.querySelectorAll('a[href]').forEach((link) => link.addEventListener('click', () => setOpen(false)));" in app_script
    assert 'setupModuleNavigator' not in app_script
    assert 'setupPagePanelNavigator();' in app_script
    design_template = (Path(__file__).parents[1] / 'src/web_interface/templates/ppt_designer.html').read_text(encoding='utf-8')
    assert design_template.index('id="ds-apply-filters"') < design_template.index('id="ds-generate-ppt"')
    assert 'id="ds-ppt-charts-panel" open data-panel-state-key="ppt-designer:ppt-charts" hidden' in design_template
    assert "$('ds-ppt-charts-panel').hidden = false;" in design_script
    assert "$('ds-ppt-charts-panel').hidden = true;" in design_script
    assert "designFiltersOpen = false;" in design_script
    assert "$('ds-filter-panel').hidden = true;" in design_script
    chart_builder_template = (Path(__file__).parents[1] / 'src/web_interface/templates/chart_builder.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">Chart Builder</p><h2>Ad-Hoc Charts from your CDRs</h2>' in chart_builder_template
    assert 'data-page-panel-label="Interactive Preview"' in chart_builder_template
    assert '<p class="eyebrow">Chart Builder</p>\n        <h2>Interactive Preview</h2>' in chart_builder_template
    assert '<p class="eyebrow">Chart Definition</p>' in chart_builder_template
    workspace_template = (Path(__file__).parents[1] / 'src/web_interface/templates/workspace.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">Workspaces Management</p>\n        <h2>Select Workspace</h2>' in workspace_template
    datasets_analysis_template = (Path(__file__).parents[1] / 'src/web_interface/templates/datasets_analysis.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">CDR Analysis</p>\n    <h2>KPI Analysis of a Single CDR</h2>' in datasets_analysis_template
    assert '<p class="eyebrow">Dataset</p>\n            <h2>Select CDR to Analyze</h2>' in datasets_analysis_template
    app_logs_template = (Path(__file__).parents[1] / 'src/web_interface/templates/app_logs.html').read_text(encoding='utf-8')
    assert '<p class="eyebrow">Application activity</p>\n        <h2>App Logs</h2>' in app_logs_template
    documentation_template = (Path(__file__).parents[1] / 'src/web_interface/templates/doc_view.html').read_text(encoding='utf-8')
    assert 'class="panel doc-panel" data-page-panel-label="{{ doc_name }}"' in documentation_template
    assert "const selectAllOrNone = () => {\n      cancelAutoClose();" in app_script
    assert "dispatchNativeChange();\n      menu.hidden = true;\n      syncTrigger();\n      trigger.focus();" not in app_script
    assert 'window.createUnifiedDatasetViewer' in app_script
    assert 'window.createUnifiedDatasetViewer' in design_script
    assert "exportControl: $('ds-data-download')" in design_script
    assert "window.showLoadingOverlay('Loading Filtered Dataset'" in design_script
    assert "await renderData();\n      overlay('ds-data-overlay', true);" in design_script
    app_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/app.css').read_text(encoding='utf-8')
    assert '.page-panel-navigator{position:fixed;' in app_styles
    assert '.page-panel-navigator-tab{' in app_styles
    assert '.page-panel-navigator[data-theme="datasets"]' in app_styles
    assert '.page-panel-navigator[data-theme="designs"]' in app_styles
    assert '.page-panel-navigator[data-theme="reporting"]' not in app_styles
    assert '.page-panel-navigator .page-panel-main-menu{' in app_styles
    assert '.preview-column-filter-menu { position: fixed; z-index: 10010;' in app_styles
    assert "\\s*×\\s*|\\s+\\bx\\b\\s+" in app_script
    assert "const configuredValues = multiFields.has(key)" in app_script
    # The CDR columns of a chart definition are offered by Chart Builder; the viewer has no Chart Definition any more.
    assert "columns_by_source" not in design_script
    assert "control.dataset.previewDisplay = parsed.value.trim();" in app_script
    assert "control.dispatchEvent(new Event('preview-selection-change'));" in app_script
    assert "would erase text the user has just entered but has not completed." in app_script
    reporting_source = (Path(__file__).parents[1] / 'src/modules/cdr_reporting.py').read_text(encoding='utf-8')
    assert 'This chart type has no interactive renderer' not in reporting_source
    app_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/app.css').read_text(encoding='utf-8')
    design_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')
    # The charts are changed in the Designer: the expanded viewer has no Chart Definition tab.
    assert 'ds-chart-filter-panel' not in design_styles
    assert '#ds-data-overlay .ds-data-dialog { position: absolute; inset: 5%;' in design_styles
    assert '.ds-chart-expanded-overlay.ds-overlay { position: fixed; z-index: 9000;' in design_styles
    assert '#ds-viewer.ds-overlay{z-index:8800}' in design_styles
    assert '#ds-data-overlay.ds-overlay, #ds-filter-overlay.ds-overlay, #ds-presentation-overlay.ds-overlay { z-index: 9100 !important; }' in design_styles
    assert '#ds-editor-overlay.ds-overlay { z-index: 9200 !important; }' in design_styles
    assert '.ppt-designer .ds-chart-expanded-navigation { gap: 0.65rem; }' in design_styles
    assert '.ppt-designer .ds-chart-expanded-arrow-icon {' in design_styles
    assert '.ppt-designer .ds-slide-navigation { gap: 0.65rem; }' in design_styles
    assert '.ppt-designer .ds-slide-arrow-icon {' in design_styles
    assert '.ds-viewer-tool-actions { display: flex;' in design_styles
    assert '.ppt-designer .ds-viewer-tool-actions .ds-viewer-icon-action {' in design_styles
    assert ':is(#ds-viewer.ppt-designer, #ds-chart-expanded-overlay) .ds-viewer-top-action::before {' in design_styles
    assert ':is(#ds-viewer.ppt-designer, #ds-chart-expanded-overlay) .ds-viewer-top-filters::before {' in design_styles
    assert ':is(#ds-viewer.ppt-designer, #ds-chart-expanded-overlay) .ds-viewer-top-auto-fields::before {' in design_styles
    assert ':is(#ds-viewer.ppt-designer, #ds-chart-expanded-overlay) .ds-viewer-top-edit::before {' in design_styles
    assert 'class="ds-presentation-navigation-separator"' in page.text
    assert '#ds-viewer.ppt-designer:not(.ds-presentation-active) .ds-presentation-navigation-separator {' in design_styles
    assert 'border-right:1px solid #b8a8ca;' in design_styles
    assert '.ppt-designer .ds-chart-expanded-pan-up {' in design_styles
    assert '.ppt-designer .ds-chart-expanded-pan-down {' in design_styles
    assert '.ppt-designer .ds-chart-pan-button {' in design_styles
    assert '.ds-chart-expanded-overlay .ds-chart-expanded-dialog { position: relative;' in design_styles
    assert '.ds-chart-expanded-overlay .ds-chart-expanded-canvas { border:' in design_styles
    assert '#ds-chart-expanded-close.report-chart-viewer-close, .ppt-designer .ds-viewer-header #ds-viewer-close.report-chart-viewer-close {' in design_styles
    assert '#ds-subtitle{min-height:1.4em}.ds-viewer-panel>.ds-actions{display:grid;' in design_styles
    assert '#ds-subtitle:empty{display:none}' not in design_styles
    assert 'text-transform: uppercase;' in design_styles
    assert 'height: 14rem;' in design_styles
    assert '.multiselect-option[hidden], .multiselect-group-label[hidden], .multiselect-subgroup-label[hidden] { display: none !important; }' in app_styles
    task_panel_start = app_script.index("const root = document.getElementById('background-task-panels');")
    task_listener = app_script.index("window.addEventListener('drivetest-analyzer:background-task'")
    assert task_listener > task_panel_start
    auto_fields_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/auto_calculated_fields.js').read_text(encoding='utf-8')
    # Closing the Auto-calculated Fields editor asks only while a field has unsaved changes.
    assert "return confirmDialog(`Discard the unsaved changes of ${name}?`, {" in auto_fields_script
    assert "const confirmTertiary = document.getElementById('confirm-tertiary');" in app_script
    assert "const handleTertiary = () => close('tertiary');" in app_script
    assert "overlay.addEventListener('pointerdown', (event) => { if (event.target === overlay) void requestClose(); });" in auto_fields_script
    assert "window.parent.postMessage({type: 'drivetest-analyzer:template-saved'}, window.location.origin);" in app_script
    assert "openTemplateEditor(expandedChart?.focus_row, sourceDefinition)" in design_script
    assert 'const hasAppliedUnsavedFilterChanges = () => Boolean(' in design_script
    assert "confirmLabel: 'Use Current Filters'" in design_script
    assert "secondaryLabel: 'Use Saved Filters'" in design_script
    assert 'wideActions: true' in design_script
    assert 'preparation_token: preparationToken' in design_script
    assert "const preview = node('div', undefined, 'ds-ppt-chart-thumbnail');" in design_script
    assert 'globalThis.renderCanvasChart(canvas, model);' in design_script
    assert 'preview.replaceChildren(cachedImage(chart))' in design_script
    # A chart is changed in the Designer only: the expanded viewer no longer applies a temporary definition.
    assert 'rememberRenderedChartPayload' not in design_script
    assert 'ds-chart-definition-toggle' not in design_script
    assert "card.ondblclick = safe(async event =>" in design_script
    assert "const syncExpandedChartNavigation" in design_script
    assert 'const setExpandedChartHeader = (chart, title = \'\') =>' in design_script
    assert "$('ds-chart-expanded-meta').textContent" in design_script
    assert "navigateExpandedChart(expandedCharts().length - 1)" in design_script
    assert "const chartPanDirections = ['left', 'right', 'up', 'down'];" in design_script
    assert "expandedPanUp.id = 'ds-chart-expanded-pan-up';" in design_script
    assert "expandedPanDown.id = 'ds-chart-expanded-pan-down';" in design_script
    assert 'bindChartPanControls(expandedCanvas, {' in design_script
    assert 'bindChartPanControls(canvas, panButtons);' in design_script
    assert 'globalThis.panCanvasChart?.(canvas, direction);' in design_script
    assert "visible === 'ds-chart-expanded-overlay' && !editing && event.key === 'ArrowLeft'" in design_script
    assert "visible === 'ds-chart-expanded-overlay' && !editing && event.key === 'ArrowRight'" in design_script
    assert "ds-chart-filter-panel" not in page.text
    assert "/filter-preview" not in design_script
    assert "operation && value ? `${operation}(${value})` : value" in app_script
    assert "operation.dataset.previewKpiAggregation = ''" in app_script
    assert "['COUNTD', 'COUNTD']" in app_script
    assert "/update-template" not in design_script
    assert 'renderExpandedChartDefinition' not in design_script
    assert "expandedCanvasShell.classList.add('ds-hover')" in design_script
    assert "addEventListener('focusout', event =>" not in design_script
    assert "document.addEventListener('pointerdown', event => {" in design_script
    # The image of the chart of the larger view, from the buttons above it.
    assert "['ds-chart-expanded-copy', true], ['ds-chart-expanded-download', false]" in design_script
    assert "`Chart ${index + 1} / ${charts.length}`" in design_script
    assert "if (event.target === event.currentTarget) expandedChartOverlay(false);" in design_script
    assert 'const closeOnOutsidePointer = (id, close) =>' in design_script
    assert "closeOnOutsidePointer('ds-filter-overlay', closeFilters);" in design_script
    assert "closeOnOutsidePointer('ds-editor-overlay', closeTemplateEditor);" in design_script
    assert "$('ds-unapplied-filters-badge').hidden = !hasUnappliedFilterChanges();" in design_script
    assert "$('ds-unsaved-filters-badge').hidden = !hasUnsavedFilterChanges();" in design_script
    # The Scope and dates are applied and saved with the filters: there are no Universe badges.
    assert "ds-unapplied-universe-badge" not in design_script
    assert "if (!$('ds-filter-overlay').hidden) await closeFilters();" in design_script
    assert "This Design has unsaved changes. Close Adaptative Filters without saving them?" not in design_script
    assert "This Report Template has unsaved changes. Close the editor without saving them?" in design_script
    assert 'const templateChanged = templateEditorSaved;' in design_script
    assert "if (templateChanged && expandedChartMode !== 'ppt') {" in design_script
    assert design_script.count('forgetPrepared();') >= 2
    assert "'Reload the current Report Template, rebuild its slides and render every chart in this Design again?'" in design_script
    assert "event.data?.type === 'drivetest-analyzer:template-saved'" in design_script
    chart_script = (Path(__file__).parents[1] / 'src/web_interface/static/js/canvas_charts.js').read_text(encoding='utf-8')
    assert "const percent = (value, digits = 1) => `${(Number(value) * 100).toFixed(digits)}%`;" in chart_script
    # Clicking a legend entry hides or shows its series for every chart type;
    # a hidden value also hides the rows and columns where it is a level.
    assert 'function legendHiddenPayload(payload, state) {' in chart_script
    assert 'const visible = legendHiddenPayload(payload, hiddenLegendItems.get(canvas));' in chart_script
    assert 'drawPayload(context, visible, state, {x: 0, y: 0, scale: 1});' in chart_script
    assert "return !shown.has(text) && (valueHidden(text) || text.split(' · ').some(valueHidden));" in chart_script
    assert 'function axisItemAt(canvas, event)' in chart_script
    assert "pushAxisHit('column', level, value, left, bandTop - 2, right - left, height);" in chart_script
    assert 'attachLegendToggle(canvas);' in chart_script
    # Legend and axis menus act on the whole Design; a hidden value's Show
    # submenu opens only while the pointer moves over it.
    assert "action(menu, 'Keep only', `Show only ${target} in every slide of the Design`" in design_script
    assert "entry.addEventListener('pointermove', event => { if (event.pointerType !== 'touch') open(); });" in design_script
    assert "parent.addEventListener('focus', open);" not in design_script
    assert 'Apply to all design' not in design_script
    # Hidden dynamic values remove their grid cards and reflow the slide.
    assert 'function slideCards(slide) {' in design_script
    assert 'const positions = prepared?.grid_layouts?.[result.layout];' in design_script
    assert "if (slide && !$('ds-viewer').hidden && slideCardsSignature(slide) !== renderedSlideCards) renderSlide();" in design_script
    assert 'if (legendItemAt(canvas, event)) return false;' in chart_script
    assert 'if (bar.hidden) return;' in chart_script
    assert 'function drawInsideHorizontalBarLabel(' in chart_script
    assert "const tooltipPercent = value => `${(Number(value) * 100).toFixed(2)}%`;" in chart_script
    assert 'value: tooltipPercent(ratio)' in chart_script
    assert 'lines.push(tooltipPercent(point.cumulative))' in chart_script
    assert "const dynamic = Boolean(payload.dynamic)" in chart_script
    assert "const repeatedHierarchyValue = dynamic" in chart_script
    assert "const boundaryStyle = (level, levels) =>" in chart_script
    assert "left + changedLevel * columnWidth" in chart_script
    assert "const columnKeys = Array.isArray(payload.column_keys)" in chart_script
    assert 'function sameImmediateHierarchyParent(left, right)' in chart_script
    assert 'cell.row >= drag.parent.start && cell.row < drag.parent.end' in chart_script
    assert 'sameImmediateHierarchyParent(drag.sourceKey, targetKey)' in chart_script
    assert 'function chartPanState(canvas)' in chart_script
    assert 'function panChart(canvas, direction)' in chart_script
    assert "['left', 'right', 'up', 'down'].includes(direction)" in chart_script
    assert "if (direction === 'up') camera.panY += LOGICAL_HEIGHT * 0.16;" in chart_script
    assert "if (direction === 'down') camera.panY -= LOGICAL_HEIGHT * 0.16;" in chart_script
    assert 'globalThis.panCanvasChart = panChart;' in chart_script
    assert 'function legendLayout(legend, fontSize = LEGEND_TEXT_SIZE)' in chart_script
    assert "position = items.length ? String(legend?.position || 'none').toLowerCase() : 'none';" in chart_script
    assert "left: position === 'left' ? sideLegendWidth + 20 : 70" in chart_script
    assert "right: position === 'right' ? LOGICAL_WIDTH - sideLegendWidth - 20 : LOGICAL_WIDTH - 60" in chart_script
    assert "top: position === 'top' ? 80 + rows * rowHeight + 18 : 82" in chart_script
    assert "bottom: position === 'bottom' ? LOGICAL_HEIGHT - rows * rowHeight - 22 : LOGICAL_HEIGHT - 40" in chart_script
    # Axis labels wrap, shrink or turn vertical with reserved room; never truncated.
    assert "function axisLevelPlan(context, entries, maximumHeight, format = {}, level = 0)" in chart_script
    assert "function hierarchicalAxisPlans(context, keys, width, format = {})" in chart_script
    assert "function hierarchyRowLabelSize()" in chart_script
    assert "function drawFullHierarchyLabel(context, value, x, y, width, size" in chart_script
    assert "const aggregationFont = (context, format = {}, level = 0) =>" in chart_script
    assert "bar.colour, 24, payload.label_position" in chart_script
    assert "series.colour, 21, payload.label_position" in chart_script
    assert "bucket.colour, 17, payload.label_position" in chart_script
    assert "drawFullHierarchyLabel(context, value, x + 4" in chart_script
    assert "const rowLabelGap = rowLevels > 1 ? 20 : 0;" in chart_script
    assert "const rowLabelArea = rowLevels ? Math.min(440, Math.max(230" in chart_script
    assert "drawFullHierarchyLabel(context, value, labelLeft + 3" in chart_script
    assert "const usableBottom = layout.position === 'bottom' ? layout.bottom : LOGICAL_HEIGHT - 16;" in chart_script
    assert 'Math.min(250, columnWidth * .72)' in chart_script
    assert "const colour = payload.cell_colours?.[rowIndex]?.[columnIndex] || '#4E79A7';" in chart_script
    assert "drawAxisLabel(context, plan, value, (left + right) / 2, bandTop + Math.max(0, (height - 4 - plan.height) / 2), right - left, '#405765', payload.legend_format);" in chart_script
    assert "drawAxisLabel(context, leafPlan, columnKey.at(-1), chartLeft + (columnIndex + .5) * columnWidth, chartTop + chartHeight + 4, columnWidth, '#4E6271', payload.legend_format);" in chart_script
    assert "line(context, left, bandTop + height - 3, right, bandTop + height - 3, '#C8D2D9');" in chart_script
    assert "const lineTop = layout.top + headerOffset(Math.min(changed, upperLevels));" in chart_script
    assert "else dashedVertical(context, cellLeft, lineTop, chartTop + chartHeight + 22);" in chart_script
    assert 'function drawOutsideBarLabel(context, value, x, y, colour, size = 12, format = {})' in chart_script
    assert "context.fillStyle = 'rgba(255, 255, 255, 0.94)'" in chart_script
    assert "context.fillText(`${(cumulative * 100).toFixed(0)}%`, plot.left - 14, y - 10)" in chart_script
    assert 'function drawConfiguredBarLabel(' in chart_script
    assert 'labelWidth + 8 > width && size + 8 <= width && labelWidth + 8 <= height' in chart_script
    assert "verticalLabel(context, value, x + width / 2, centreY, format.color || '#FFFFFF', size, false, format);" in chart_script
    assert 'function drawAdjacentStackLabel(' in chart_script
    assert 'sideSpace, colour, occupied, size = 10' in chart_script
    assert "if (ratio > 0) drawAdjacentStackLabel(" in chart_script
    assert 'series.colour, sideLabels' in chart_script
    assert 'Math.floor((layout.right - left) / Math.max(headers.length, 1))' in chart_script
    assert 'function selectionStartAllowed(canvas, event)' in chart_script
    assert 'return logicalY >= 90;' in chart_script
    assert "canvas.classList.toggle('ds-chart-selection-blocked', !selectionStartAllowed(canvas, event));" in chart_script
    assert 'function applyDateBounds(bounds)' in design_script
    assert "button.disabled = Boolean((minimum && value < minimum) || (maximum && value > maximum));" in design_script
    assert 'applyDateBounds(payload.date_bounds)' in design_script
    assert 'function datePicker(key, label)' in design_script
    assert "previous.addEventListener('click', () => { month.setMonth(month.getMonth() - 1); render(); });" in design_script
    assert "button.addEventListener('click', () => { input.value = iso; definition[key] = iso; automatic.setAttribute('aria-pressed', 'false'); updateFilterControlState(wrapper, dateState(key), true); menu.hidden = true; filterChanged(); });" in design_script
    assert 'const current = (hasStoredSelection ? selected : available).filter(value => value && value !== MAIN_CITIES);' in design_script
    assert "const filterControlState = (current, applied, saved, equal = sameFilterValues) =>" in design_script
    assert "if (!equal(current, applied)) return 'unapplied';" in design_script
    assert "return equal(applied, saved) ? '' : 'applied-unsaved';" in design_script
    assert "const scopeState = () => filterControlState(" in design_script
    assert "const dateState = key => filterControlState(" in design_script
    assert "const hasUnappliedFilterChanges = () =>" in design_script
    assert "filterStateFingerprint(definition) !== filterStateFingerprint(appliedDefinition())" in design_script
    # Scope, vendor comparison and dates are part of the Design filters.
    assert "scope: selection.scope || 'single', vendor_comparison: selection.vendor_comparison || 'operator_vendor'," in design_script
    assert 'const resolveUnappliedFilterChanges = async () =>' in design_script
    assert "title: 'Unapplied PPT filters'" in design_script
    assert "confirmLabel: 'Apply Filters and Continue'" in design_script
    assert "secondaryLabel: 'Discard and Continue'" in design_script
    assert "tertiaryLabel: 'Save and Continue'" in design_script
    assert "await prepare();\n      return 'applied';" in design_script
    assert "if (hasUnsavedFilterChanges()) await save();" in design_script
    assert "return 'saved';" in design_script
    assert 'restoreAppliedFilterSelection();' in design_script
    assert 'const openActiveDesignViewer = async () =>' in design_script
    assert "if (!$('ds-filter-overlay').hidden) await closeFilters();" in design_script
    assert 'const needsPreparation = !prepared || preparationStateFingerprint(definition) !== appliedFilterState;' in design_script
    assert "setPreparationState('preparing', needsDataPreparation ? 'data' : 'rendering');" in design_script
    assert "bind('ds-view', openActiveDesignViewer);" in design_script
    assert "async function queueDesignPptExport(id, item, {chooseScope = false} = {})" in design_script
    assert "const universeChoice = await chooseDesignPptUniverse(item, id);" in design_script
    assert "title: 'Choose PowerPoint Scope'" not in design_script
    assert "updateFilterControlState(facet, filterState(field));" in design_script
    assert "updateFilterControlState(wrapper, dateState(key), true);" in design_script
    # Saving makes what the server saved the saved state, and applies it.
    assert "savedDefinition = definitionFingerprint(result.definition);" in design_script
    assert "if (hasUnappliedFilterChanges()) await preparePart();" in design_script
    assert 'await prepare();' in design_script
    assert 'if (next.length === current.length && next.every(value => current.includes(value))) return;' in design_script
    assert 'function filterChanged() {' in design_script
    assert "? 'Scope, date and filter changes are ready to apply.'" in design_script
    assert ": 'The PPT Scope, dates and filters are applied.'" in design_script
    assert 'if (preparing && cached && currentFilterState !== preparingFilterState)' not in design_script
    assert "status('Restored the previously prepared filters.');" in design_script
    assert "const payload = await api('/filter-options', 'POST', {definition, field});" in design_script
    assert "facetOptionRequests.has(field) ? 'Loading values…'" in design_script
    assert "const filterAliases = config.filter_aliases || {};" in design_script
    assert "facet.dataset.aliasTooltip = aliasTooltip;" in design_script
    design_styles = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')
    assert '.ds-facet[data-alias-tooltip]::after' in design_styles
    assert 'const preparedPayloadKey = (id, fingerprint) =>' in design_script
    assert 'const preparedStateFingerprint = value =>' in design_script
    assert 'const canonicalSetValues = values =>' in design_script
    assert 'const canonicalSelectionDefinition = value =>' in design_script
    assert 'selection.filters = Object.fromEntries' in design_script
    assert "if (!definitionValue.date_from) definitionValue.date_from = 'Oldest';" in design_script
    assert "if (!definitionValue.date_to) definitionValue.date_to = 'Newest';" in design_script
    assert "const automaticLabel = key === 'date_from' ? 'Use oldest' : 'Use newest';" in design_script
    assert "const automaticValue = key === 'date_from' ? 'Oldest' : 'Newest';" in design_script
    assert 'const eventTimeFilteringDisabled = Boolean(config.ignore_event_time_filtering);' in design_script
    assert "wrapper.classList.add('is-event-time-filtering-disabled');" in design_script
    assert "eventTimeFilteringDisabled && ['eventstarttime', 'eventendtime']" not in design_script
    assert "if (value === 'Oldest') return dateBounds?.min ? `Oldest (${dateBounds.min})` : 'Oldest';" in design_script
    assert "if (value === 'Newest') return dateBounds?.max ? `Newest (${dateBounds.max})` : 'Newest';" in design_script
    assert "input.value = dateInputDisplayValue(key, definition[key]);" in design_script
    assert 'const persistedDesignDefinition = value => canonicalDesignDefinition(value);' in design_script
    assert "delete persisted.scope;" not in design_script
    assert "delete persisted.datasets;" not in design_script
    assert 'delete preparedDefinition.name;' in design_script
    assert 'delete preparedDefinition.slide_comments;' in design_script
    assert 'entries.push({fingerprint, token: payload.token});' in design_script
    assert 'const selectionStateFingerprint = value =>' in design_script
    assert 'appliedSelectionState = selectionStateFingerprint(definition);' in design_script
    assert 'const needsDataPreparation = selectionStateFingerprint(definition) !== appliedSelectionState;' in design_script
    assert "params.set('rendering_only', '1');" in design_script
    assert "params.set('preparation_id', preparationToken);" in design_script
    assert "design_name: definition?.name || 'Design'," in design_script
    assert "let preparationToken = '';" in design_script
    assert "window.dispatchEvent(new Event('drivetest-analyzer:refresh-background-tasks'));" in design_script
    assert 'let designPptJobsLoaded = false, designPptJobsRefreshing = false;' in design_script
    assert "previous.status !== 'ready'" in design_script
    assert "renderDesignPptJobs(jobs, newlyReady[0]?.id || '');" in design_script
    assert "const previous = String(preferredJobId || select.value);" in design_script
    design_module = (Path(__file__).parents[1] / 'src/modules/ppt_designer.py').read_text(encoding='utf-8')
    assert 'for saved_id, saved_definition in read_designs(task_repository).items()' not in design_module
    assert "'design_name': task['name']," in design_module
    assert "'Waiting to prepare PPT dataset'" in design_module
    assert "'progress': task.get('progress', 0)," in design_module
    assert "@app.get('/api/ppt-designer/preparation-progress/{preparation_id}')" in design_module
    assert "'Counting reduced and filtered Universe rows in the combined CDR tables'" in design_module
    assert "f'Counting Filtered Universe CDR-{kind.title()} rows in the combined table'" in design_module
    assert "progress(78, 'Restoring cached row counts and filter options')" in design_module
    assert "update_preparation_progress(94, 'Saving the reusable PPT preparation cache')" in design_module
    assert 'with lock, task_repository.connection() as connection:' not in design_module
    assert "job['progress'] = round(82 + job['completed'] * 18 / max(job['total'], 1))" not in design_module
    assert 'direct_preparation_tasks: dict[str, dict] = {}' in design_module
    assert "preparation_id: str | None = None," in design_module
    assert "'label': 'Rendering PPT Charts' if job['total'] else 'Preparing PPT dataset'," not in design_module
    assert "phase === 'rendering' ? 'Rendering PPT Charts' : 'Preparing PPT dataset'" in design_script
    assert "if (preparingFilterState === requestedFilterState) return preparing;" in design_script
    assert "if (backgroundPreparationToken === preparationToken) dismissPreparationStatus();" in design_script
    assert "bind('ds-apply-filters', async () => {" in design_script
    assert "if (!hasUnappliedFilterChanges()) return;" in design_script
    assert "bind('ds-apply-universe'" not in design_script
    assert "bind('ds-reload-universe'" not in design_script
    assert "await api(`/prepare?${params}`, 'POST', definition, controller.signal);" in design_script
    assert "window.dispatchEvent(new Event('drivetest-analyzer:refresh-background-tasks'));" in design_script
    assert "title: 'Unsaved PPT filters'" in design_script
    assert "confirmLabel: 'Save Filters'" in design_script
    assert "title: 'Unsaved Dataset Universe'" not in design_script
    assert "if (link.classList.contains('topnav-link-logout'))" in design_script
    assert 'for (const key of [openStorageKey, filtersOpenStorageKey, scrollStorageKey, preparedStorageKey, universeStorageKey])' in design_script
    assert "secondaryLabel: 'Discard'" in design_script
    assert "window.location.assign(target.href);" in design_script
    assert "$('ds-apply-filters').disabled = !hasUnappliedFilterChanges() || filterActionBusy;" in design_script
    assert "const designStatuses = new Map();" in design_script
    assert 'const designPreparationTokens = new Map();' in design_script
    assert "if (!designPreparationTokens.has(id)) designStatuses.set(id, value);" in design_script
    assert "bind('ds-generate-ppt', async () => {" in design_script
    assert "heading.replaceChildren(document.createTextNode(name ? 'PPT Filters: ' : 'PPT Filters'));" in design_script
    assert "heading.append(node('span', name.toUpperCase(), 'ds-active-design-name'))" in design_script
    assert 'setActiveDesignHeading(definition.name);' in design_script
    assert "setActiveDesignHeading('');" in design_script
    assert "if (activePrepared) designStatuses.set(id, ['ready', 'rendering', 'charts-queued'].includes(value?.state) ? value : {state: 'charts-queued', label: 'Data Cached'});" in design_script
    assert "setDesignStatus(activeId, 'charts-queued', 'Data Cached');" in design_script
    assert "view.dataset.designViewId = id;" in design_script
    assert 'button.disabled = false;' in design_script
    assert "const setViewEnabled = enabled => { $('ds-view').disabled = !definition; syncDesignViewActions(); syncDesignPptActions(); };" in design_script
    assert "const payload = await api(statusPath, 'POST', statusDefinitions);" in design_script
    assert "window.setInterval(refreshDesignStatuses, 2000);" in design_script
    assert "const designName = String(task.design_name || '');" in app_script
    assert "? `PPT Design “${designName}”: ${String(task.label || 'Background task')}`" in app_script
    assert 'completedByWorkspace.forEach((entries) =>' in app_script
    assert 'entry.group.tasks.splice(Math.min(entry.position + offset, entry.group.tasks.length), 0, entry.task);' in app_script
    selection_key_source = design_module[design_module.index('def persistent_selection_key'):design_module.index('def selected_date_bounds')]
    assert "'scope': definition.scope," not in selection_key_source
    assert "'schema': DESIGN_SELECTION_CACHE_VERSION," in selection_key_source
    assert 'DESIGN_SELECTION_CACHE_VERSION = 18' in design_module
    assert "kind: sorted([" in selection_key_source
    assert "kind: sorted(set(dataset_ids))" in selection_key_source
    assert "field: sorted(set(values))" in selection_key_source
    assert '{entry_key}:{operator_mapping_key}' in design_module
    assert "'rendering_only': rendering_only," in design_module
    assert 'def materialize_selection(' in design_module
    assert 'use_profile_options=False, known_full_row_counts=None, progress=None,' in design_module
    assert 'use_profile_options=use_profile_options,' in design_module
    assert 'DESIGN_CHART_RENDER_WORKERS = max(1, min(2, (os.cpu_count() or 2) - 1))' in design_module
    assert 'DESIGN_PREVIEW_MANIFEST_VERSION = 10' in design_module
    assert "thread_name_prefix='ppt-designer-chart'," not in design_module
    assert 'def schedule_next_prefetch() -> None:' not in design_module
    assert 'def enqueue_prefetch(' not in design_module
    assert "void api(`/prefetched/${encodeURIComponent(activeId)}/priority`" not in design_script
    assert 'const proximityIndexes = (length, origin) =>' in design_script
    assert 'if (origin + distance < length) indexes.push(origin + distance);' in design_script
    assert 'if (origin - distance >= 0) indexes.push(origin - distance);' in design_script
    # Neighbouring slides preload one chart at a time, as deferrable requests,
    # and a visible chart is requested again with high priority.
    assert "await loadChartPayload(chart, 'low').catch(() => undefined);" in design_script
    assert "if (!request || (priority === 'high' && request.chartPriority === 'low')) {" in design_script
    assert "priority=low" in design_script
    assert "await loadChartPayload(charts[target], 'low').catch(() => undefined);" in design_script
    assert 'scheduleNearbySlidePreload(preloadToken, preloadOrigin, visibleLoads);' in design_script
    assert 'scheduleNearbyExpandedChartPreload(contextKey, chart);' in design_script
    assert 'template_where, template_parameters, template_filters_applied = chart_filter_sql(' in design_module
    assert 'template_filters_applied=template_filters_applied,' in design_module
    assert 'database_path, table_name, source_columns = reporting_source(' in design_module
    assert 'ATTACH DATABASE' not in design_module
    assert 'workspace_read_connection(database_path)' in design_module
    assert 'aggregation_columns = chart_aggregation_columns(' in design_module
    assert "thread_name_prefix='ppt-designer-data'," not in design_module
    design_css = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')
    assert '.ppt-designer :is(.ds-unsaved-filters-badge,.ds-unsaved-universe-badge)' in design_css
    assert '.ppt-designer :is(.ds-unapplied-filters-badge,.ds-unapplied-universe-badge)' in design_css
    assert '.ds-scope-control.ds-filter-applied-unsaved select' in design_css
    assert '.ds-date-picker.ds-date-picker-applied-unsaved .ds-date-picker-control>input' in design_css
    assert '.ppt-designer .ds-design-close::after' in design_css
    assert '.ppt-designer .ds-active-design-name{color:#12664f;font-family:inherit;font-size:1.12em;font-weight:900}' in design_css
    assert '.ppt-designer .ds-design-close{background:linear-gradient(135deg,#a9273a,#dc5361)' in design_css
    assert '.ppt-designer .ds-design-open{background:linear-gradient(145deg,#e87912,#ffad2f)' in design_css
    assert '.ppt-designer .ds-design-view{background:linear-gradient(145deg,#167957,#29ae7d)' in design_css
    assert '.ds-design-status-loading-data{border-color:#d3aa45;background:#fff1c9;color:#77570a}' in design_css
    assert '.ds-design-status-data-queued,.ds-design-status-charts-queued{border-color:#aaa3b2;background:#f0edf2;color:#655e6c}' in design_css
    assert '.ds-filter-groups{display:grid;grid-template-rows:max-content max-content minmax(0,1fr);gap:14px}' in design_css
    assert '.ds-default-filter-panel{border-color:#afd1e8;background:#f1f8fd}' in design_css
    assert "d='M12 2v10'" in design_css
    assert "title: 'Refresh Design?'" in design_script
    assert "api(`/charts/${encodeURIComponent(token)}/refresh`, 'POST')" in design_script
    assert "api(`/chart/${encodeURIComponent(token)}/${chart.index}/refresh`, 'POST')" in design_script
    assert "definition[key] = automaticValue;" in design_script
    assert "input.value = dateInputDisplayValue(key, automaticValue);" in design_script
    # The Campaign of a new Design follows its newest two campaigns (the newest one for Multivendor Operator_Vendor),
    # as its newest CDRs did; a Scope change updates a Campaign that still is the default of the former Scope.
    assert "const defaultCampaignFilter = (scope = 'single', vendorComparison = 'operator_vendor') => [" in design_script
    assert "scope === 'multivendor' && vendorComparison === 'operator_vendor' ? LATEST_CAMPAIGN : LATEST_TWO_CAMPAIGNS," in design_script
    assert "if (selectedScope !== definition.scope)" in design_script
    assert "followDefaultCampaign(previousScope, previousComparison);" in design_script
    assert "latestDatasetsForScope" not in design_script
    assert "chooseScopeDatasets" not in design_script
    assert 'async function restorePrepared(id) {' in design_script
    assert 'const preparedPayloads = new Map();' in design_script
    assert 'const restoreRememberedPrepared = async id =>' in design_script
    assert 'const inMemory = preparedPayloads.get(preparedPayloadKey(id, fingerprint));' in design_script
    assert 'const restoredNow = await restorePrepared(activeId);' in design_script
    assert "const payload = await api(`/prepared/${encodeURIComponent(inMemory.payload.token)}`);" in design_script
    assert 'forgetPreparedToken(inMemory.payload.token);' in design_script
    assert 'forgetPrepared();' in design_script
    assert "await prepare();\n      token = prepared?.token || '';" in design_script
    assert "api(`/prefetched/${encodeURIComponent(id)}`, 'POST', definition)" in design_script
    assert 'if (error.status === 409) return false;' in design_script
    assert 'No CDR ${chart.source[0].toUpperCase()}${chart.source.slice(1)} dataset has been selected for this chart.' in design_script
    assert "api(`/prepared/${encodeURIComponent(cachedEntry.token)}`)" in design_script
    assert 'const monitorPreparationProgress = token =>' in design_script
    assert 'const designNeedsRefresh = () => Boolean(' in design_script
    assert "$('ds-preparing-out-of-sync').hidden = !stale;" in design_script
    assert "bind('ds-preparing-refresh', async () => {" in design_script
    assert 'label.textContent = known ? `${detail} · ${percent}%` : detail;' in design_script
    assert "const known = progress !== null && progress !== undefined && progress !== '' && Number.isFinite(numeric);" in design_script
    assert "$('ds-preparing-detail').textContent = detail;" in design_script
    assert "emitPreparationStatus(payload.status || 'processing', payload.detail, token, queued ? null : payload.progress);" in design_script
    assert 'api(`/preparation-progress/${encodeURIComponent(token)}`)' in design_script
    assert 'monitorPreparationProgress(preparationToken);' in design_script
    # The preparation is registered (202) before its progress is polled.
    assert design_script.index("await api(`/prepare?${params}`, 'POST', definition, controller.signal);") < design_script.index('monitorPreparationProgress(preparationToken);')
    assert 'const payload = await awaitBackgroundPreparation(preparationToken, controller.signal);' in design_script
    assert "design_work_gate.acquire(timeout=0.2)" in design_module
    assert "status(''); await prepare();" in design_script
    assert "setPreparationState('preparing', needsDataPreparation ? 'data' : 'rendering');" in design_script
    assert 'void refreshDesignStatuses();' in design_script
    assert 'void refreshDesignPptJobs();' in design_script
    assert 'await Promise.all([refreshDesignStatuses(), refreshDesignPptJobs()]);' not in design_script
    assert 'savedDefinition = definitionFingerprint(definition); updateDirtyState();' not in design_script
    assert 'the controls so their state agrees with the Apply and Save buttons.\n    sources(); facets();' in design_script
    assert "setPreparationState('preparing');" in design_script
    assert "bind('ds-refresh',prepare);" not in design_script
    assert 'const setPreparationRows = payload =>' in design_script
    assert 'setPreparationRows(payload);' in design_script
    assert 'applyPreparedPayload(cached.payload);' not in design_script
    assert "['CDRs used', universe, 'ds-preparing-universe-label']" in design_script
    assert "['Filtered rows', filtered, 'ds-preparing-filtered-label']" in design_script
    assert 'rows.hidden = !rows.textContent;' in design_script
    assert 'const zoomResetIcon = () =>' in design_script
    assert "if (request.filter_column) parameters.set('filter_column', request.filter_column);" in design_script
    assert "Clear ${activeFilters} Filter${activeFilters === 1 ? '' : 's'}" in app_script
    assert "window.createUnifiedDatasetViewer" in design_script
    assert "const reset = node('button', undefined, 'ds-chart-zoom-button ds-chart-zoom-reset');" in design_script
    assert "bind('ds-clear-filters', () => { definition.filters = {Campaign: defaultCampaignFilter(definition.scope, definition.vendor_comparison)}; facets(); filterChanged(); });" in design_script
    assert "bind('ds-last-saved-filters', () => {" in design_script
    assert 'id="ds-reload-universe"' not in page.text
    assert "definition.datasets = structuredClone(saved.datasets || {});" not in design_script
    # Reload Saved Filters restores the saved Scope, dates and filters.
    assert "copyDefinitionFields(definition, savedDesignDefinition(), filterDefinitionFields);" in design_script
    design_css = (Path(__file__).parents[1] / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')
    assert '.ds-design-status-ready{' in design_css
    assert '.ds-design-status-rendering{' in design_css
    assert '.ds-chart-controls button:not(:disabled){cursor:pointer!important}' in design_css
    assert '.ds-preparing .ds-preparing-universe-label{color:#60408d}' in design_css
    assert '.ds-preparing .ds-preparing-filtered-label{color:#d7a921}' in design_css
    assert '.ppt-designer .ds-chart-zoom-reset svg' in design_css
    assert '.ppt-designer .ds-chart-data{width:2rem;min-width:2rem;height:2rem;min-height:2rem}' in design_css
    assert '.ds-source-filter.ds-filter-unsaved select,.ds-source-filter.ds-filter-unsaved .multiselect-trigger' in design_css
    assert '.ds-date-picker.ds-date-picker-unsaved .ds-date-picker-control>input,.ds-facet.ds-filter-unsaved .multiselect-trigger' in design_css
    assert '.ppt-designer .ds-date-auto-action{position:absolute;right:.32rem;top:50%' in design_css
    assert '.ppt-designer .ds-generate-ppt-action{' in design_css
    assert ':is(.ds-view-design-action,.ds-generate-ppt-action)::before{width:1.35rem' in design_css
    assert '#ds-apply-filters::before' in design_css
    assert '#ds-save::before' in design_css
    assert '#ds-clear-filters::before' in design_css
    assert '#ds-last-saved-filters::before' in design_css
    assert '.ds-slide-content.ds-comments-right .ds-viewer-navigation-footer {' in design_css
    assert 'grid-column:1 / -1;' in design_css
    assert '.ds-slide-content.ds-comments-right>.ds-slide-comments:not([open]) #ds-comments-title {' in design_css
    assert 'transform:translate(-50%,-50%) rotate(90deg);' in design_css
    assert 'clip-path:polygon(100% 0,32% 50%,100% 100%,76% 100%,8% 50%,76% 0);' in design_css
    assert '.ds-slide-comments,\n.ds-slide-comments>summary,\n.ds-slide-comments-body {\n  background:#c4b0dc;' in design_css
    assert "panel?.classList.toggle('ds-has-comments', comments.length > 0);" in design_script
    assert 'if (panel && presentation.active) panel.open = presentation.showComments && comments.length > 0;' in design_script
    assert 'commentsPanel.open = presentationCommentsWasOpen;' in design_script
    assert '#ds-viewer.ds-presentation-active.ds-presentation-comments-enabled .ds-slide-comments.ds-has-comments {' in design_css
    assert '#ds-viewer.ds-presentation-active .ds-slide-comments { display:none!important; }' in design_css
    assert 'id="ds-presentation-comments"' in page.text
    assert 'id="ds-presentation-delay" type="number" min="1" max="300"' in page.text
    assert 'id="ds-presentation-transition-duration" type="number" min="0.1" max="10" step="0.1" value="1"' in page.text
    assert '<option value="slide-left">Slide from left</option>' in page.text
    assert '<option value="zoom">Zoom</option>' in page.text
    assert '<option value="rise">Rise</option>' in page.text
    assert '<option value="blur">Blur</option>' in page.text
    assert '<option value="rotate">Soft rotate</option>' in page.text
    assert '<option value="flip" selected>Flip</option>' in page.text
    assert '<option value="bounce">Bounce</option>' in page.text
    assert '<option value="wipe">Wipe</option>' in page.text
    assert '<option value="mosaic">Mosaic</option>' in page.text
    assert '<option value="curtain">Curtain</option>' in page.text
    assert '<option value="blinds">Blinds</option>' in page.text
    assert '<option value="random">Random</option>' in page.text
    assert '<option value="none">None</option><option value="random">Random</option><optgroup label="Effects">' in page.text
    assert '<option value="flip" selected>Flip</option>' in page.text
    assert 'id="ds-presentation-stop-viewer"' in page.text
    assert 'id="ds-presentation-toggle-viewer"' in page.text
    assert "presentation.showComments = $('ds-presentation-comments').value === 'yes';" in design_script
    assert "$('ds-viewer').style.setProperty('--ds-presentation-transition-duration', `${presentation.transitionDuration}s`);" in design_script
    assert 'animation:ds-slide-fade var(--ds-presentation-transition-duration,1s) ease both' in design_css
    assert "const randomPresentationEffects = ['fade', 'slide', 'slide-left', 'zoom', 'rise', 'blur', 'rotate', 'flip', 'bounce', 'wipe', 'mosaic', 'curtain', 'blinds'];" in design_script
    assert "if (presentation.effect !== 'random') return presentation.effect;" in design_script
    assert "lastRandomPresentationEffect = choices[Math.floor(Math.random() * choices.length)] || 'fade';" in design_script
    assert "localStorage.getItem(presentationEffectStorageKey)" in design_script
    assert "localStorage.setItem(presentationEffectStorageKey, event.target.value)" in design_script
    assert "lastRandomPresentationEffect = '';\n    slideIndex = 0;" in design_script
    assert "bind('ds-presentation-stop-viewer', stopPresentation);" in design_script
    assert "bind('ds-presentation-toggle-viewer', togglePresentation);" in design_script
    assert "if (presentation.running) pausePresentation(); else resumePresentation();" in design_script
    assert "event.code === 'Space' || event.key === ' '" in design_script
    assert "if (visible === 'ds-viewer' && !editing && event.key === 'F8')" in design_script
    assert 'if (!presentation.active) startPresentation();' in design_script
    assert "presentation.active && event.key === 'F7'" in design_script
    assert "presentation.active && event.key === 'F9'" in design_script
    assert 'if (presentation.active) navigatePresentationSlide(-1);' in design_script
    assert 'if (presentation.active) navigatePresentationSlide(1);' in design_script
    assert '[data-presentation-effect="slide-left"]' in design_css
    assert '[data-presentation-effect="zoom"]' in design_css
    assert '[data-presentation-effect="rise"]' in design_css
    assert '[data-presentation-effect="blur"]' in design_css
    assert ':is(#ds-presentation,#ds-presentation-toggle-viewer,#ds-presentation-stop-viewer)' in design_css
    assert '.ds-presentation-stop-action::after' in design_css
    assert ':is(#ds-presentation-toggle-viewer,#ds-presentation-stop-viewer)[hidden] {' in design_css
    assert "function designViewerBrand(className = 'ds-structural-brand')" in design_script
    assert "const brand = designViewerBrand('ds-chart-brand'); card.append(brand);" in design_script
    assert '.ds-chart-brand{' in design_css
    assert 'position:absolute;z-index:2;top:2.5rem;right:12px;' in design_css
    assert '.ds-chart>img,.ds-chart>canvas{z-index:1}' in design_css
    assert '.ds-chart>.ds-chart-controls{top:calc(2.9rem + 9px)}' in design_css
    assert '.ds-chart>.ds-chart-brand>img{' in design_css
    assert 'position:static!important;inset:auto!important;width:2.65rem!important;height:2.65rem!important;object-fit:contain!important' in design_css
    assert "canvas.addEventListener('designchartlayout', event => {" in design_script
    assert "brand.style.top = `${top + height / 2}px`;" in design_script
    assert 'filter:none' in design_css
    assert 'id="ds-filter-close-action"' in page.text
    assert '#ds-filter-close-action[hidden]{display:none!important}' in design_css
    assert "bind('ds-filter-close-action', closeFilters);" in design_script
    assert "$('ds-generate-ppt').hidden = !fromLibrary; $('ds-view').hidden = !fromLibrary; $('ds-filter-close-action').hidden = fromLibrary; overlay('ds-filter-overlay', true);" in design_script
    assert "panel.hidden = !designFiltersOpen;" in design_script
    assert '#ds-add-filter::before' in design_css
    assert '#ds-ppt-jobs-delete-all::before' in design_css
    assert '.ppt-designer .ds-ppt-charts-filters::before{' in design_css
    assert '#ds-filter-float .ds-filter-help{margin:26px 0 5px}' in design_css
    assert '#ds-filter-float .ds-data-panel>summary{pointer-events:none;cursor:default}' in design_css
    assert "panel.querySelector('summary').tabIndex = -1;" in design_script
    assert '.ds-chart-expanded-canvas.ds-hover .ds-chart-controls,.ds-chart-expanded-canvas:focus-within .ds-chart-controls{opacity:1;visibility:visible;transform:translateY(0);transition-delay:0s;pointer-events:auto}' in design_css
    saved = client.put('/api/ppt-designer/test', json=payload)
    assert saved.status_code == 200
    assert saved.json()['definition']['hidden_filters'] == []
    assert saved.json()['definition']['slide_comments'] == {}
    listed = client.get('/api/ppt-designer')
    assert listed.json()['test']['name'] == 'Comparison'
    # The saved Design comes back as the list gives it, so the page does not take it for an unsaved change (its place in
    # the library comes with the list only, and the page keeps it apart from the definition).
    assert saved.json()['definition'] == {key: value for key, value in listed.json()['test'].items() if key != 'position'}
    assert "const appliedDefinition = () => applyingDesignDefinition || appliedDesignDefinition || savedDesignDefinition();" in design_script
    # A running preparation does not hold Apply and Save Filters: a new preparation replaces it.
    preparation = design_script[design_script.index('  async function prepareDefinition() {'):design_script.index('  async function openDesign(')]
    assert 'filterActionBusy' not in preparation and 'preparationRunning = true; updateFilterActionState();' in preparation
    assert listed.headers['cache-control'] == 'no-store, max-age=0, must-revalidate'
    renamed = client.patch('/api/ppt-designer/test/name', json={'name': 'Renamed comparison'})
    assert renamed.status_code == 200
    assert renamed.json()['name'] == 'Renamed comparison'
    assert client.get('/api/ppt-designer').json()['test']['name'] == 'Renamed comparison'
    comments = client.patch('/api/ppt-designer/test/comments', json={'slide_comments': {'1': ['Review city outliers']}})
    assert comments.status_code == 200
    assert comments.json()['slide_comments'] == {'1': ['Review city outliers']}
    assert client.get('/api/ppt-designer').json()['test']['slide_comments'] == {'1': ['Review city outliers']}
    result = client.post('/api/ppt-designer/prepare', json=payload)
    assert result.status_code == 200, result.text
    preview = result.json()
    assert preview['filter_fields'] == ['Operator', 'Operator_Vendor', 'Vendor_Operator', 'Vendor', 'Market', 'Region', 'Cluster', 'City', 'Campaign', 'RAT', 'Session Type', 'Call Status']
    assert 'Technology' not in preview['options']
    assert preview['options']['Operator'] == ['A', 'B']
    assert preview['options']['City'] == ['Leeds', 'London']
    assert preview['date_bounds'] == {'min': '2026-09-01', 'max': '2026-09-03'}
    assert len(preview['slides']) == 2
    assert len(preview['slides'][0]['charts']) == 2
    assert preview['slides'][0]['charts'][0]['focus_row'] == 0
    assert preview['slides'][0]['charts'][0]['cdr_source']
    assert preview['slides'][0]['charts'][0]['chart_type']
    assert preview['slides'][0]['charts'][1]['focus_row'] == 1
    assert preview['slides'][0]['charts'][0]['position'][0] < preview['slides'][0]['charts'][1]['position'][0]
    token = preview['token']
    restored = client.get(f'/api/ppt-designer/prepared/{token}')
    assert restored.status_code == 200
    assert restored.json() == preview
    image = client.get(f'/api/ppt-designer/preview/{token}/0.png')
    assert image.status_code == 200, image.text if image.status_code != 200 else ''
    assert image.content.startswith(b'\x89PNG')
    interactive = client.get(f'/api/ppt-designer/chart/{token}/0')
    assert interactive.status_code == 200, interactive.text
    assert interactive.json()['type'] == 'cdf'
    assert interactive.json()['renderer'] == 'catalog-v2'
    assert interactive.json()['series']
    # Charts and their images are kept in memory only.
    assert not (Path(core.repository.db_path).parent / '.design-data-cache').exists()
    entry = core.load_template_catalogue(next(row['content'] for row in core.repository.list_report_templates() if row['name'] == 'Design test'))[0]
    from src.modules.cdr_reporting import is_empty_catalog_chart
    assert not is_empty_catalog_chart(image.content, entry)
    data = client.get(f'/api/ppt-designer/data/{token}/0').json()
    assert data['total'] == 3
    assert data['page_size'] == 100
    assert data['unfiltered_total'] == 3
    assert set(data['column_metadata']) == set(data['columns'])
    assert all({'label', 'kind', 'rule', 'pinned', 'class_name'} <= set(item) for item in data['column_metadata'].values())
    # The NR Mode of the CDR of each row is a pinned field of the rows of every chart.
    assert data['column_metadata']['NR_Mode']['pinned'] and {row[data['columns'].index('NR_Mode')] for row in data['rows']} == {'NSA'}
    city_values = client.get(f'/api/ppt-designer/data/{token}/0', params={
        'column_filters': json.dumps({'Operator': ['A']}), 'filter_column': 'City',
    }).json()
    assert city_values['filter_values'] == ['London']
    assert client.get(f'/api/ppt-designer/data/{token}/0?download=true').headers['content-type'].startswith('text/csv')
    assert client.delete('/api/ppt-designer/test').status_code == 200
    assert client.get('/api/ppt-designer').json() == {}


def test_ppt_reporting_old_is_no_longer_a_module(client):
    # Not even a super-admin, who sees every module, has it: its former addresses open PPT Designer.
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    workspace_page = client.get('/workspace')
    assert workspace_page.status_code == 200
    assert 'href="/ppt-reporting-old"' not in workspace_page.text
    assert 'PPT Reporting (old)' not in workspace_page.text
    for former in ('/ppt-reporting-old', '/e2e-reporting'):
        response = client.get(former, follow_redirects=False)
        assert response.status_code == 307
        assert response.headers['location'] == '/ppt-designer'
    assert client.get('/api/ppt-reporting-old/jobs').status_code == 404
    assert 'ppt-reporting-old' not in core.FEATURE_KEYS
    # A Modules Profile that still names it (saved before, or from a package of before) keeps only the other modules.
    profile_id = core.repository.save_modules_profile(None, 'Legacy reporting', '', ['workspace', 'ppt-reporting-old'])
    profile = next(item for item in core.repository.list_modules_profiles() if item['id'] == profile_id)
    assert profile['modules'] == ['workspace']
    assert 'value="ppt-reporting-old"' not in client.get('/super-admin').text


def test_a_workspace_of_before_drops_the_jobs_of_ppt_reporting_old_and_keeps_its_documents(tmp_path):
    import sqlite3

    from src.modules.repository import Repository

    database = tmp_path / 'workspace.db'
    document = tmp_path / 'output' / 'reports' / 'reports-old' / 'report' / 'report.pptx'
    document.parent.mkdir(parents=True)
    document.write_bytes(b'report')
    with sqlite3.connect(database) as connection:
        connection.execute('CREATE TABLE generated_jobs (id INTEGER PRIMARY KEY AUTOINCREMENT, job_type TEXT, output_path TEXT)')
        connection.execute("INSERT INTO generated_jobs (job_type, output_path) VALUES ('report', ?)", (str(document),))
        # Workspaces older still kept them in two tables.
        connection.execute('CREATE TABLE report_chart_jobs (id INTEGER PRIMARY KEY, generation TEXT)')
    repository = Repository(database)
    repository.initialize()
    tables = set(repository.list_database_tables())
    assert not {'generated_jobs', 'report_runs', 'report_chart_jobs'} & tables
    assert 'datasets' in tables and document.read_bytes() == b'report'
    with sqlite3.connect(database) as connection:
        events = connection.execute("SELECT username, details FROM audit_logs WHERE action = 'retired_report_jobs_dropped'").fetchall()
    assert [(username, json.loads(details)['tables']) for username, details in events] == [
        ('system', {'generated_jobs': 1, 'report_chart_jobs': 0})]
    # Opening it again leaves it as it is.
    repository.initialize()
    assert 'generated_jobs' not in repository.list_database_tables()
    with sqlite3.connect(database) as connection:
        assert connection.execute("SELECT COUNT(*) FROM audit_logs WHERE action = 'retired_report_jobs_dropped'").fetchone()[0] == 1


def test_ready_design_exports_ppt_and_persistent_chart_files(client, monkeypatch):
    payload = setup_design(client)
    payload['slide_comments'] = {'1': ['Review city outliers', 'Validate campaign coverage']}
    payload['filters'] = {'City': ['London']}
    design_id = 'ppt-designer'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200

    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    initial_status = client.get('/api/ppt-designer/statuses').json()[design_id]
    # Loading its data, data cached with its charts waiting to be rendered, or ready, as the background work goes.
    assert initial_status['state'] in {'loading-data', 'charts-queued', 'ready'}
    assert initial_status['label'] == 'Ready' or initial_status['label'].startswith(('Pre-Caching Universe ', 'Data Cached · Charts'))

    with core.repository.connection() as connection:
        connection.execute('UPDATE dataset_profiles SET vendor_mapping_applied = 1')
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=[], cities=['Leeds', 'London'],
    )
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'],
        'Vendor': ['Vendor A', 'Vendor B', 'Vendor A'],
        'City': ['London', 'Leeds', 'London'],
        'Mean_Data_Rate': [10, 20, 30],
        'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
        # Ingestion always normalizes this field; a reused snapshot filters
        # its resolved calendar dates against it.
        'event_start_time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    export_payload = json.loads(json.dumps(payload))
    export_payload['scope'] = 'multivendor'
    prepared = client.post(
        f'/api/ppt-designer/prepare?design_id={design_id}', json=export_payload,
    )
    assert prepared.status_code == 200, prepared.text
    prepared_payload = prepared.json()
    for slide in prepared_payload['slides']:
        for chart in slide['charts']:
            if chart['available']:
                rendered = client.get(f"/api/ppt-designer/chart/{prepared_payload['token']}/{chart['index']}")
                assert rendered.status_code == 200, rendered.text

    queued = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
        'definition': export_payload,
        'preparation_token': prepared_payload['token'],
    })
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']
    job = None
    while time.monotonic() < deadline:
        jobs = client.get('/api/ppt-designer/ppt-jobs').json()['jobs']
        job = next(item for item in jobs if item['id'] == job_id)
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.05)
    assert job is not None and job['status'] == 'ready', job
    assert job['slides'] == 2
    assert job['charts'] == 3
    assert job['nr_mode'] == 'NSA'
    assert job['scope'] == 'Multivendor Comparison'
    assert job['template'] == 'Design test'
    assert re.fullmatch(r'\d{8}_\d{6}', job['timestamp'])
    assert job['filters'] == [
        'CDR Data: sample.csv', 'Date: Oldest to Newest',
        'Operator: All Operators', 'Vendor: All Vendors', 'City: London',
        'Vendor comparison: Operator_Vendor',
    ]
    assert job['duration_seconds'] is not None
    background_groups = client.get('/api/background-tasks').json()['groups']
    background_task = next(
        task for group in background_groups for task in group['tasks']
        if task['id'] == f'design-ppt:{core.active_workspace.id}:{job_id}'
    )
    assert background_task['design_name'] == 'Comparison'
    assert background_task['label'] == 'Generating PPT Design'
    assert background_task['detail'] == 'Completed'
    assert background_task['duration_seconds'] is not None

    with core.repository.connection() as connection:
        row = connection.execute('SELECT output_path FROM design_ppt_jobs WHERE id = ?', (job_id,)).fetchone()
    output_path = Path(row['output_path'])
    charts_dir = output_path.parent / 'design-charts'
    assert output_path.is_file()
    assert re.fullmatch(
        r'\d{8}_\d{6} - PPT Designer - NSA - Comparison - Multivendor Comparison\.pptx',
        output_path.name,
    )
    assert output_path.parent.name == output_path.stem
    assert output_path.parent.parent == Path(core.repository.db_path).parent / 'output' / 'reports' / 'ppt-designs'
    assert len(list(charts_dir.glob('*.png'))) == 3
    assert len(list(charts_dir.glob('*.hover.json'))) == 3
    assert len(list(charts_dir.glob('*.model.json'))) == 3
    manifest = json.loads((charts_dir / 'manifest.json').read_text(encoding='utf-8'))
    assert manifest['generate_tooltips'] is True
    assert manifest['design_id'] == design_id
    assert re.fullmatch(r'[0-9a-f]{64}', manifest['preview_fingerprint'])
    assert manifest['definition']['filters'] == {'City': ['London']}
    assert manifest['definition']['scope'] == 'multivendor'
    assert len(manifest['charts']) == 3
    assert manifest['charts'][0]['entry_index'] == 0
    assert manifest['charts'][0]['focus_row'] == 0
    presentation = Presentation(output_path)
    assert not [
        shape for slide in presentation.slides for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        and any((shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom))
    ]
    for shape in (
        shape for slide in presentation.slides for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ):
        with Image.open(BytesIO(shape.image.blob)) as chart_image:
            assert abs(chart_image.width / chart_image.height - shape.width / shape.height) < 0.01
    commentary = next(
        shape for shape in presentation.slides[0].shapes
        if shape.is_placeholder and shape.placeholder_format.idx == 10
    )
    assert [paragraph.text for paragraph in commentary.text_frame.paragraphs] == [
        'Review city outliers', 'Validate campaign coverage',
    ]

    ppt = client.get(job['download_url'])
    assert ppt.status_code == 200
    assert ppt.content.startswith(b'PK')
    assert client.get(job['charts_url']).status_code == 200
    charts_manifest = client.get(job['charts_api_url'])
    assert charts_manifest.status_code == 200
    charts_payload = charts_manifest.json()
    assert charts_payload['job']['id'] == job_id
    assert len(charts_payload['charts']) == 3
    assert client.get(charts_payload['charts'][0]['image_url']).status_code == 200
    assert charts_payload['charts'][0]['payload_url']
    assert charts_payload['charts'][0]['data_url']
    assert charts_payload['charts'][0]['focus_row'] == 0
    chart_model = client.get(charts_payload['charts'][0]['payload_url'])
    assert chart_model.status_code == 200
    # Charts read only their own columns; a dataset page reads every field.
    chart_data = client.get(f"/api/ppt-designer{charts_payload['charts'][0]['data_url']}")
    assert chart_data.status_code == 200
    assert chart_data.json()['total'] == 2
    filtered_data = client.get(
        f"/api/ppt-designer{charts_payload['charts'][0]['data_url']}",
        params={
            'column_filters': json.dumps({'Operator': ['A']}),
            'include_filter_values': 'true',
        },
    )
    assert filtered_data.status_code == 200
    assert filtered_data.json()['chart_total'] == 2
    assert filtered_data.json()['total'] == 2
    assert filtered_data.json()['filter_values']['Operator'] == ['A']
    empty_data = client.get(
        f"/api/ppt-designer{charts_payload['charts'][0]['data_url']}",
        params={'column_filters': json.dumps({'Operator': ['B']})},
    )
    assert empty_data.status_code == 200
    assert empty_data.json()['chart_total'] == 2
    assert empty_data.json()['total'] == 0
    assert empty_data.json()['rows'] == []
    assert chart_model.json()['title'] == charts_payload['charts'][0]['title']
    archive = client.get(job['charts_download_url'])
    assert archive.status_code == 200
    with zipfile.ZipFile(BytesIO(archive.content)) as bundle:
        assert len([name for name in bundle.namelist() if name.endswith('.png')]) == 3

    deleted = client.post('/api/ppt-designer/ppt-jobs/delete-all')
    assert deleted.status_code == 200
    assert deleted.json()['deleted'] == 1
    assert client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] == []
    assert not output_path.parent.exists()


def test_relaunching_design_ppt_uses_the_modified_report_template(client):
    payload = setup_design(client)
    design_id = 'updated-template-ppt'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200

    deadline = time.monotonic() + 15
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    assert core.repository.get_dataset(1)['status'] == 'ready'
    with core.repository.connection() as connection:
        connection.execute('UPDATE dataset_profiles SET vendor_mapping_applied = 1')
    core.repository.replace_cdr_catalogue(
        1, vendors=['Vendor A', 'Vendor B'], regions=[], cities=['Leeds', 'London'],
    )
    core.repository.replace_reporting_rows(1, 'data', pd.DataFrame({
        'Operator': ['A', 'B', 'A'], 'Vendor': ['Vendor A', 'Vendor B', 'Vendor A'],
        'City': ['London', 'Leeds', 'London'], 'Mean_Data_Rate': [10, 20, 30],
        'Test_Name': ['HTTP DL'] * 3,
        'Test_Start_Time': ['2026-09-01', '2026-09-02', '2026-09-03'],
    }))
    export_payload = {**payload, 'scope': 'multivendor'}
    prepared = client.post(
        f'/api/ppt-designer/prepare?design_id={design_id}', json=export_payload,
    )
    assert prepared.status_code == 200, prepared.text
    prepared_payload = prepared.json()
    for slide in prepared_payload['slides']:
        for chart in slide['charts']:
            if chart['available']:
                assert client.get(
                    f"/api/ppt-designer/chart/{prepared_payload['token']}/{chart['index']}"
                ).status_code == 200
    queued = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
        'definition': export_payload,
        'preparation_token': prepared_payload['token'],
    })
    assert queued.status_code == 202, queued.text
    job_id = queued.json()['job_id']

    def wait_for_job():
        while time.monotonic() < deadline:
            job = next(
                item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs']
                if item['id'] == job_id
            )
            if job['status'] in {'ready', 'failed'}:
                return job
            time.sleep(0.05)
        return job

    original_job = wait_for_job()
    assert original_job['status'] == 'ready'
    assert original_job['scope'] == 'Multivendor Comparison'
    original_charts = client.get(
        f'/api/ppt-designer/ppt-jobs/{job_id}/charts.json'
    ).json()['charts']
    assert original_charts[0]['title'] == 'Rate'

    template = core.repository.report_template_content('Design test')
    core.repository.set_report_template_content('Design test', template.replace(b',Rate,CDR-Data,', b',Updated rate,CDR-Data,', 1),
    )

    relaunched = client.post(f'/api/ppt-designer/ppt-jobs/{job_id}/retry')
    assert relaunched.status_code == 202, relaunched.text
    deadline = time.monotonic() + 15
    job = wait_for_job()
    assert job['status'] == 'ready', job
    assert job['scope'] == 'Multivendor Comparison'
    refreshed_charts = client.get(job['charts_api_url']).json()['charts']
    assert refreshed_charts[0]['title'] == 'Updated rate'
    refreshed_model = client.get(refreshed_charts[0]['payload_url'])
    assert refreshed_model.status_code == 200, refreshed_model.text
    assert refreshed_model.json()['title'] == 'Updated rate'
    with core.repository.connection() as connection:
        output_path = Path(connection.execute(
            'SELECT output_path FROM design_ppt_jobs WHERE id = ?', (job_id,),
        ).fetchone()['output_path'])
    manifest = json.loads((output_path.parent / 'design-charts' / 'manifest.json').read_text(encoding='utf-8'))
    assert manifest['definition']['scope'] == 'multivendor'
    assert client.get(job['download_url']).content.startswith(b'PK')


def test_design_is_prepared_only_when_opened_and_then_reuses_manifest(client):
    payload = setup_design(client)
    design_id = 'on-demand-design'
    saved = client.put(f'/api/ppt-designer/{design_id}', json=payload)
    assert saved.status_code == 200

    response = client.get(f'/api/ppt-designer/prefetched/{design_id}')
    assert response.status_code == 409
    assert not any(
        task.get('design_name') == payload['name']
        for group in client.get('/api/background-tasks').json()['groups']
        for task in group['tasks']
    )

    prepared = client.post(f'/api/ppt-designer/prepare?design_id={design_id}', json=payload)
    assert prepared.status_code == 200, prepared.text
    response = client.get(f'/api/ppt-designer/prefetched/{design_id}')
    assert response.status_code == 200, response.text
    assert response.json()['slides']
    # Its preparation is kept in memory only.
    assert not (Path(core.repository.db_path).parent / '.design-data-cache').exists()
    token = prepared.json()['token']
    for _ in range(130):
        assert client.get('/api/ppt-designer/statuses').status_code == 200
    assert client.get(f'/api/ppt-designer/prepared/{token}').status_code == 200


def test_closed_design_status_uses_its_remembered_session_universe(client):
    payload = setup_design(client)
    design_id = 'session-universe-design'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200

    session_definition = json.loads(json.dumps(payload))
    session_definition['date_from'] = '2026-09-02'
    session_definition['date_to'] = '2026-09-03'
    prepared = client.post(
        f'/api/ppt-designer/prepare?design_id={design_id}', json=session_definition,
    )
    assert prepared.status_code == 200, prepared.text

    background_status = client.get('/api/ppt-designer/statuses').json()[design_id]
    assert background_status['state'] in {'loading-data', 'charts-queued', 'ready'}
    session_status = client.post('/api/ppt-designer/statuses', json={
        design_id: session_definition,
    }).json()[design_id]
    assert session_status['state'] in {'loading-data', 'charts-queued', 'ready'}


def test_saving_design_prepares_nothing_and_renders_no_chart(client, monkeypatch):
    payload = setup_design(client)
    import src.modules.ppt_designer as designs_module
    monkeypatch.setattr(
        designs_module,
        'catalog_chart_payload',
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError('Saving must not render charts.')),
    )
    design_id = 'saved-design'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    assert client.get(f'/api/ppt-designer/prefetched/{design_id}').status_code == 409
    # It has CDRs to show: it is prepared when opened, nothing in the background.
    assert client.get('/api/ppt-designer/statuses').json()[design_id] == {'state': 'ready', 'label': 'Ready'}


def test_direct_design_preparation_separates_queue_and_execution_timestamps(client, monkeypatch):
    payload = setup_design(client)
    entered = Event()
    release = Event()
    original = core.Repository.list_dataset_row_columns

    def delayed(repository, *args, **kwargs):
        # Only the preparations requested here wait. TestClient runs them in
        # AnyIO worker threads; background pre-caching and warm-up threads
        # started by the setup (whatever their name) must not stand in for them.
        if not threading.current_thread().name.startswith('AnyIO'):
            return original(repository, *args, **kwargs)
        entered.set()
        assert release.wait(5)
        return original(repository, *args, **kwargs)

    monkeypatch.setattr(core.Repository, 'list_dataset_row_columns', delayed)
    responses = {}
    running = Thread(target=lambda: responses.setdefault(
        'running', client.post(
            '/api/ppt-designer/prepare?preparation_id=direct-running', json=payload,
        ),
    ))
    queued = Thread(target=lambda: responses.setdefault(
        'queued', client.post(
            '/api/ppt-designer/prepare?preparation_id=direct-queued', json=payload,
        ),
    ))
    try:
        running.start()
        assert entered.wait(5)
        running_task = next(
            task for group in client.get('/api/background-tasks').json()['groups']
            for task in group['tasks'] if task['id'] == 'direct-running'
        )
        assert running_task['status'] == 'processing'
        assert running_task['started_at'] >= running_task['queued_at']

        queued.start()
        deadline = time.monotonic() + 5
        queued_task = None
        while time.monotonic() < deadline:
            queued_task = next((
                task for group in client.get('/api/background-tasks').json()['groups']
                for task in group['tasks'] if task['id'] == 'direct-queued'
            ), None)
            if queued_task is not None:
                break
            time.sleep(0.02)
        assert queued_task is not None
        assert queued_task['status'] == 'queued'
        assert queued_task['queued_at'] is not None
        assert queued_task['started_at'] is None
    finally:
        release.set()
        running.join(5)
        if queued.ident is not None:
            queued.join(5)

    assert not running.is_alive()
    assert not queued.is_alive()


def test_applying_filters_prepares_only_data_and_renders_charts_on_demand(client, monkeypatch):
    payload = setup_design(client)
    import src.modules.ppt_designer as designs_module

    core.repository.set_workspace_state('ppt_designs_v2', json.dumps({'filtered-design': payload}))
    uncached = client.get('/api/ppt-designer/statuses')
    assert uncached.status_code == 200
    assert uncached.json()['filtered-design'] == {'state': 'ready', 'label': 'Ready'}

    calls = []
    original = designs_module.catalog_chart_payload

    def tracked(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(designs_module, 'catalog_chart_payload', tracked)
    response = client.post('/api/ppt-designer/prepare?design_id=filtered-design', json=payload)
    assert response.status_code == 200, response.text

    assert calls == []

    token = response.json()['token']
    for index in range(3):
        rendered = client.get(f'/api/ppt-designer/chart/{token}/{index}')
        assert rendered.status_code == 200, rendered.text
    assert len(calls) == 3
    # The chart models are kept in memory only.
    assert not (Path(core.repository.db_path).parent / '.design-data-cache').exists()

    payload['filters'] = {'City': ['London']}
    filtered = client.post('/api/ppt-designer/prepare?design_id=filtered-design', json=payload)
    assert filtered.status_code == 200, filtered.text
    assert len(calls) == 3

    payload['filters'] = {}
    calls_before_restore = len(calls)
    restored = client.post('/api/ppt-designer/prefetched/filtered-design', json=payload)
    assert restored.status_code == 200, restored.text
    assert len(calls) == calls_before_restore

    core.repository.update_dataset_profile(1, progress=100)
    assert client.get('/api/ppt-designer/statuses').json()['filtered-design']['state'] == 'ready'


def test_adding_a_dataset_builds_a_new_chart_model_with_every_campaign(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    for file_name, campaign, date_value in (
        ('campaign-q1.csv', '2026-Q1', '2026-02-01'),
        ('campaign-q2.csv', '2026-Q2', '2026-05-01'),
    ):
        response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
            'dataset_files': (
                file_name,
                BytesIO(
                    f'Operator,Campaign,Mean_Data_Rate,Test_Start_Time\nA,{campaign},10,{date_value}\n'.encode()
                ),
                'text/csv',
            ),
        })
        assert response.status_code == 200
    core.repository.add_report_template('Campaign cache test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Campaigns,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,Average Vertical Bars,,Operator,Campaign,,Right\n'
    ).encode(), is_default=False)
    datasets = sorted(core.repository.list_datasets(), key=lambda row: int(row['id']))
    first_id, second_id = (int(row['id']) for row in datasets[-2:])
    payload = DesignDefinition(
        name='Campaign cache test', template='Campaign cache test', datasets={'data': [first_id]},
    ).model_dump(mode='json')

    first = client.post('/api/ppt-designer/prepare', json=payload)
    assert first.status_code == 200, first.text
    first_model = client.get(f"/api/ppt-designer/chart/{first.json()['token']}/0")
    assert first_model.status_code == 200, first_model.text
    assert {bar['key'][-1] for bar in first_model.json()['bars']} == {'26-Q1'}

    payload['datasets']['data'] = [first_id, second_id]
    payload['date_from'], payload['date_to'] = 'Oldest', 'Newest'
    combined = client.post('/api/ppt-designer/prepare', json=payload)
    assert combined.status_code == 200, combined.text
    combined_model = client.get(f"/api/ppt-designer/chart/{combined.json()['token']}/0")
    assert combined_model.status_code == 200, combined_model.text
    assert {bar['key'][-1] for bar in combined_model.json()['bars']} == {'26-Q1', '26-Q2'}

    assert not (Path(core.repository.db_path).parent / '.design-data-cache').exists()


def test_design_api_session_expires_on_application_process_restart(client):
    setup_design(client)
    import src.DriveTestAnalyzer as app_module

    app_module.SESSIONS.clear()

    response = client.get('/api/ppt-designer', follow_redirects=False)
    assert response.status_code == 401
    assert response.json()['detail'] == 'Your session has expired. Please sign in again.'


def test_design_preview_identifies_title_and_transition_slides(client):
    payload = setup_design(client)
    core.repository.add_report_template('Structural design', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Quarterly review,Network results,Title Page,,,,Title Slide,,,,,Top\n'
        '2,Voice performance,,Title Only,,,,Transition Slide,,,,,Top\n'
        '3,,,Black logo end slide,,,,,,,,,\n'
    ).encode(), is_default=False)
    payload['template'] = 'Structural design'

    preview = client.post('/api/ppt-designer/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    slides = preview.json()['slides']
    assert [(slide['title'], slide['structural_type'], slide['charts']) for slide in slides] == [
        ('Quarterly review', 'title slide', []),
        ('Voice performance', 'transition slide', []),
        ('', 'end slide', []),
    ]


def test_design_state_migrates_from_legacy_storage(client):
    payload = setup_design(client)
    without_starter_designs()
    core.repository.set_workspace_state('ppt_design_sets_v1', '{"legacy": ' + json.dumps(payload) + '}')
    result = client.get('/api/ppt-designer')
    assert result.status_code == 200
    assert result.json()['legacy']['name'] == 'Comparison'
    # Listed with who created it (nobody known here: only admins and super-admins delete it) and its place in the
    # library (none until its editors order it).
    listed = {key: {**definition, 'created_by': '', 'position': None}
              for key, definition in json.loads(core.repository.get_workspace_state('ppt_designs_v2')).items()}
    assert listed == result.json()


def test_design_migrates_all_saved_definitions_to_current_default_filters(client):
    payload = setup_design(client)
    payload['filters'] = {'Technology': ['5G'], 'Zone': ['North'], 'Operator': ['A']}
    payload['custom_fields'] = ['Technology', 'Zone', 'Region', 'Mean_Data_Rate']
    payload['hidden_filters'] = ['Technology', 'Region', 'Zone', 'City', 'Campaign', 'RAT', 'Call Status', 'Former custom filter']
    with core.repository.connection() as connection:
        connection.execute("DELETE FROM workspace_state WHERE key = 'ppt_design_default_filters_v6'")
    core.repository.set_workspace_state('ppt_designs_v2', json.dumps({'legacy': payload}))

    result = client.get('/api/ppt-designer')

    assert result.status_code == 200
    migrated = result.json()['legacy']
    # Its CDRs, which hold no campaign, become the newest two campaigns.
    assert migrated['filters'] == {'Operator': ['A'], 'Region': ['North'], 'Campaign': ['@latest-2-campaigns']}
    assert migrated['custom_fields'] == ['Mean_Data_Rate']
    assert migrated['hidden_filters'] == ['Former custom filter']
    assert {key: migrated[key] for key in ('scope', 'date_from', 'date_to')} == {
        key: payload[key] for key in ('scope', 'date_from', 'date_to')
    }
    assert 'datasets' not in migrated and migrated['auto_cdrs'] is True
    assert core.repository.get_workspace_state('ppt_design_default_filters_v6') == '1'
    assert {**json.loads(core.repository.get_workspace_state('ppt_designs_v2'))['legacy'], 'created_by': '', 'position': None} == migrated


def test_design_custom_fields_and_snapshot_filters(client):
    payload = setup_design(client)
    core.write_workspace_calculated_dimensions([{'name': '7-cities', 'sources': ['cdr-data'], 'rules': [{'when': 'City IN (London)', 'value': 'Yes'}], 'default': 'No'}])
    payload['custom_fields'] = ['7-cities']
    preview = client.post('/api/ppt-designer/prepare', json=payload)
    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['7-cities'] == ['No', 'Yes']
    original = preview.json()['token']
    payload['filters'] = {'7-cities': ['Yes']}
    preview = client.post('/api/ppt-designer/prepare', json=payload).json()
    assert preview['rows']['data'] == 2
    assert preview['universe_rows']['data'] == 3
    assert preview['options']['7-cities'] == ['No', 'Yes']
    for index in (0, 1, 2):
        assert client.get(f"/api/ppt-designer/data/{preview['token']}/{index}").json()['total'] == 2
    assert client.get(f'/api/ppt-designer/data/{original}/0').json()['total'] == 3
    payload['filters'] = {'7-cities': []}
    assert client.post('/api/ppt-designer/prepare', json=payload).json()['rows']['data'] == 0
    payload['filters'] = {'Test_Name': ['HTTP DL']}
    payload['custom_fields'] = ['Test_Name']
    preview = client.post('/api/ppt-designer/prepare', json=payload).json()
    assert 'Test_Name' in preview['available_fields']
    assert preview['options']['Test_Name'] == ['HTTP DL']
    payload['hidden_filters'] = ['Market']
    preview = client.post('/api/ppt-designer/prepare', json=payload).json()
    assert 'Market' not in preview['options']


def test_design_source_specific_custom_fields_do_not_repeat_combined_preparation(client, monkeypatch):
    payload = setup_design(client)
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'voice'}, files={
        'dataset_files': ('voice.csv', BytesIO(
            b'Operator,City,Mean_Call_Setup_Time,Session_Type\nA,London,1.2,VoLTE\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.write_workspace_calculated_dimensions([
        {
            'name': 'Test Family', 'sources': ['cdr-data'],
            'rules': [{'when': 'Test_Name CONTAINS HTTP', 'value': 'HTTP'}], 'default': '',
        },
        {
            'name': 'Call Family', 'sources': ['cdr-voice'],
            'rules': [{'when': 'Session_Type CONTAINS VoLTE', 'value': 'VoLTE'}], 'default': 'Call',
        },
    ])
    payload['datasets']['voice'] = [2]
    payload['custom_fields'] = ['Test Family', 'Call Family']
    first = client.post('/api/ppt-designer/prepare', json=payload)
    assert first.status_code == 200, first.text

    from src.modules.repository import Repository

    def unexpected_combined_rebuild(*_args, **_kwargs):
        raise AssertionError('Fields from another CDR type must not prepare source columns.')

    monkeypatch.setattr(Repository, 'copy_dataset_rows_to_reporting', unexpected_combined_rebuild)
    payload['filters'] = {'City': ['London']}
    repeated = client.post('/api/ppt-designer/prepare', json=payload)
    assert repeated.status_code == 200, repeated.text


def test_design_loads_new_filter_values_without_preparing_a_snapshot(client):
    payload = setup_design(client)
    payload['custom_fields'] = ['mean data rate']
    payload['filters'] = {'CITY': ['lOnDoN']}

    response = client.post('/api/ppt-designer/filter-options', json={
        'definition': payload,
        'field': 'MEAN-DATA-RATE',
    })

    assert response.status_code == 200, response.text
    assert response.json() == {'field': 'MEAN-DATA-RATE', 'values': ['10', '30']}


def test_design_rat_filter_uses_dataset_preview_column_precedence(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('radio.csv', BytesIO(
            b'Operator,City,Campaign,RAT_A,RAT,Sample_RAT_A,Call_Status,Mean_Data_Rate\n'
            b'A,London,Spring,ENDC,LTE,NR,Completed,10\n'
            b'B,Leeds,Summer,NR,LTE,ENDC,Failed,20\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('Radio Design', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Radio,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    payload = definition().model_dump(mode='json')
    payload['template'] = 'Radio Design'

    preview = client.post('/api/ppt-designer/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['Campaign'] == ['Spring', 'Summer']
    assert preview.json()['options']['RAT'] == ['ENDC', 'NR']
    assert preview.json()['options']['Call Status'] == ['Completed', 'Failed']


def test_design_falls_back_to_combined_rows_for_values_missing_from_profiles(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('profile-gap.csv', BytesIO(
            b'Operator,G_Level_2,G_Level_4,Campaign,RAT,Call_Status,Mean_Data_Rate,Test_Start_Time\n'
            b'A,England,London,Spring,ENDC,Completed,10,2026-09-01\n'
            b'B,England,Leeds,Summer,NR,Failed,20,2026-09-02\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    with core.repository.connection() as connection:
        connection.execute("UPDATE dataset_profiles SET filter_options_json = '{}' WHERE dataset_id = 1")
    core.repository.add_report_template('Profile fallback Design', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Profile fallback,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    payload = definition().model_dump(mode='json')
    payload['template'] = 'Profile fallback Design'

    preview = client.post('/api/ppt-designer/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['Region'] == ['England']
    assert preview.json()['options']['City'] == ['Leeds', 'London']
    assert preview.json()['options']['Campaign'] == ['Spring', 'Summer']
    assert preview.json()['options']['RAT'] == ['ENDC', 'NR']
    assert preview.json()['options']['Call Status'] == ['Completed', 'Failed']


def test_design_validates_template_dates_and_sources(client):
    payload = setup_design(client)
    payload['date_from'], payload['date_to'] = '2026-09-03', '2026-09-01'
    assert client.post('/api/ppt-designer/prepare', json=payload).status_code == 400
    payload['date_from'] = payload['date_to'] = None
    # A Campaign that no CDR of the NR Mode holds chooses no CDR.
    payload['filters'] = {'Campaign': ['UK_Q1_2030']}
    assert client.post('/api/ppt-designer/prepare', json=payload).status_code == 400
    payload['template'] = 'Unknown'
    assert client.put('/api/ppt-designer/test', json=payload).status_code == 400
    assert client.get('/api/ppt-designer/data/missing/0').status_code == 410


def test_design_sql_selection_leaves_nr_mode_to_explicit_user_filters(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'voice'}, files={
        'dataset_files': ('voice.csv', BytesIO(
            b'Operator,Mean_Call_Setup_Time,RAT_A,Session_Type,L1_Call_Mode_A\n'
            b'A,1.0,EN-DC,WhatsApp Voice,\n'
            b'A,2.0,NR,WhatsApp Voice,\n'
            b'A,3.0,LTE,Native Voice,VoLTE\n'
            b'A,4.0,NR,Native Voice,VoNR\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('Voice Design', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Voice,,Title and 1 column + Comments,Calls,CDR-Voice,Mean_Call_Setup_Time,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    payload = DesignDefinition(
        name='Voice', template='Voice Design', datasets={'voice': [1]}, nr_mode='NSA',
    ).model_dump(mode='json')
    nsa = client.post('/api/ppt-designer/prepare', json=payload).json()
    assert nsa['rows']['voice'] == 4
    assert nsa['universe_rows']['voice'] == 4
    payload['nr_mode'] = 'SA'
    # An SA Design only uses SA CDRs, but never filters rows by RAT.
    assert client.post('/api/ppt-designer/prepare', json=payload).status_code == 400
    core.repository.update_dataset_profile(1, nr_mode='SA')
    sa = client.post('/api/ppt-designer/prepare', json=payload).json()
    assert sa['rows']['voice'] == 4
    assert sa['universe_rows']['voice'] == 4
    design_module = (Path(__file__).parents[1] / 'src/modules/ppt_designer.py').read_text(encoding='utf-8')
    assert 'def nr_mode_sql(' not in design_module


def test_design_snapshot_access_and_legacy_redirect(client):
    payload = setup_design(client)
    response = client.get('/dashboard?dataset_id=1', follow_redirects=False)
    assert response.status_code == 307
    assert response.headers['location'] == '/datasets-analysis?dataset_id=1'
    assert client.get(response.headers['location']).status_code == 200
    preview = client.post('/api/ppt-designer/prepare', json=payload).json()
    client.get('/logout')
    client.post('/login', data={'username': 'demo', 'password': 'demo123'})
    assert client.get(f"/api/ppt-designer/data/{preview['token']}/0").status_code == 410
    assert client.get(f"/api/ppt-designer/preview/{preview['token']}/0.png").status_code == 410
    assert client.get(f"/api/ppt-designer/chart/{preview['token']}/0").status_code == 410
    assert client.get(f"/api/ppt-designer/prepared/{preview['token']}").status_code == 410


def test_design_reuses_persistent_sql_selection_and_invalidates_dataset_versions(client, monkeypatch):
    payload = setup_design(client)
    first = client.post('/api/ppt-designer/prepare', json=payload)
    assert first.status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 1
    repeated = client.post('/api/ppt-designer/prepare', json=payload)
    assert repeated.status_code == 200
    assert repeated.json()['token'] != first.json()['token']
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 1
    payload['name'] = 'Renamed comparison'
    assert client.post('/api/ppt-designer/prepare', json=payload).status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 1

    from src.modules.repository import Repository

    def unexpected_combined_rebuild(*_args, **_kwargs):
        raise AssertionError('Default Filter value changes must not prepare source columns.')

    monkeypatch.setattr(Repository, 'copy_dataset_rows_to_reporting', unexpected_combined_rebuild)
    payload['filters'] = {'City': ['London']}
    assert client.post('/api/ppt-designer/prepare', json=payload).json()['rows']['data'] == 2
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 2
    payload['filters'] = {'City': ['Leeds']}
    assert client.post('/api/ppt-designer/prepare', json=payload).json()['rows']['data'] == 1
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 3
    payload['filters'] = {'City': ['London']}
    assert client.post('/api/ppt-designer/prepare', json=payload).json()['rows']['data'] == 2
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 3
    core.repository.set_workspace_state('design_cache_test', 'updated')
    assert client.post('/api/ppt-designer/prepare', json=payload).status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 3
    core.repository.update_dataset_profile(1, progress=100)
    assert client.post('/api/ppt-designer/prepare', json=payload).status_code == 200
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) AS count FROM design_filter_selections').fetchone()['count'] == 4


def test_design_selection_cache_ignores_filter_value_order(client):
    payload = setup_design(client)
    payload['filters'] = {'City': ['London', 'Leeds']}

    first = client.post('/api/ppt-designer/prepare', json=payload)
    assert first.status_code == 200
    payload['filters']['City'].reverse()
    repeated = client.post('/api/ppt-designer/prepare', json=payload)
    assert repeated.status_code == 200

    with core.repository.connection() as connection:
        assert connection.execute(
            'SELECT COUNT(*) AS count FROM design_filter_selections'
        ).fetchone()['count'] == 1


def test_design_selection_uses_direct_sql_predicate(client, monkeypatch):
    payload = setup_design(client)
    import src.modules.ppt_designer as designs_module

    monkeypatch.setattr(designs_module, 'DESIGN_PROFILE_SELECTION_THRESHOLD', 1)
    payload['filters'] = {'City': ['London']}
    preview = client.post('/api/ppt-designer/prepare', json=payload)
    assert preview.status_code == 200
    assert preview.json()['rows']['data'] == 2
    assert preview.json()['rows_exact'] is True
    with core.repository.connection() as connection:
        selection = connection.execute(
            'SELECT id FROM design_filter_selections ORDER BY id DESC LIMIT 1'
        ).fetchone()
        assert selection is not None
        assert connection.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = 'design_filter_selection_rows'"
        ).fetchone() is None
        assert 'materialized' not in {
            row['name'] for row in connection.execute('PRAGMA table_info(design_filter_selections)')
        }
    data = client.get(f"/api/ppt-designer/data/{preview.json()['token']}/0")
    assert data.status_code == 200
    assert data.json()['total'] == 2


def test_design_profile_facets_show_values_outside_the_saved_filter(client, monkeypatch):
    payload = setup_design(client)
    import src.modules.ppt_designer as designs_module

    monkeypatch.setattr(designs_module, 'DESIGN_PROFILE_SELECTION_THRESHOLD', 1)
    with core.repository.connection() as connection:
        connection.execute("UPDATE dataset_profiles SET filter_options_json = '{}' WHERE dataset_id = 1")
    payload['filters'] = {'Operator': ['A']}

    preview = client.post('/api/ppt-designer/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    assert preview.json()['rows']['data'] == 2
    assert preview.json()['options']['Operator'] == ['A', 'B']


def combined_scans(monkeypatch, caller):
    """The dataset ids and fields of the combined-table reads ``caller`` (a function of ppt_designer) makes."""
    import src.modules.cdr_store as cdr_store_module

    scans = []
    original = cdr_store_module.combined_scope

    def recording_scope(dataset_ids):
        frame = sys._getframe(1)
        if frame.f_code.co_name == caller:
            scans.append((sorted(int(item) for item in dataset_ids), list(frame.f_locals.get('field_order') or [])))
        return original(dataset_ids)

    monkeypatch.setattr(cdr_store_module, 'combined_scope', recording_scope)
    return scans


def test_design_profile_facets_read_values_missing_from_profiles_from_each_cdr(client, monkeypatch):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('call-status.csv', BytesIO(
            b'Operator,City,Call_Status,Mean_Data_Rate,Test_Start_Time\n'
            b'A,London,Completed,10,2026-09-01\n'
            b'B,Leeds,Failed,20,2026-09-02\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200
    core.repository.add_report_template('Design test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Comparison,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline and core.repository.get_dataset(1)['status'] != 'ready':
        time.sleep(0.05)
    import src.modules.ppt_designer as designs_module

    monkeypatch.setattr(designs_module, 'DESIGN_PROFILE_SELECTION_THRESHOLD', 1)
    # The upload profile lists no Call Status, and the Cluster catalogue of the CDR was read and is empty.
    with core.repository.connection() as connection:
        connection.execute(
            'UPDATE dataset_profiles SET filter_options_json = ? WHERE dataset_id = 1',
            (json.dumps({'operator': ['A', 'B'], 'city': ['Leeds', 'London']}),),
        )
    core.repository.set_cdr_catalogue_clusters(1, [])
    scans = combined_scans(monkeypatch, 'combined_filter_options')

    preview = client.post('/api/ppt-designer/prepare', json=definition().model_dump(mode='json'))

    assert preview.status_code == 200, preview.text
    assert preview.json()['options']['Call Status'] == ['Completed', 'Failed']
    assert preview.json()['options']['Cluster'] == []
    # Call Status is read from the CDR and Cluster is answered by its catalogue: neither is looked for in the
    # combined tables.
    scanned = {field for _ids, fields in scans for field in fields}
    assert not scanned & {'Call Status', 'Cluster'}, scans


def test_design_date_limits_of_each_cdr_are_kept_for_other_selections(client, monkeypatch):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    for index, (campaign, first_day, last_day) in enumerate((
            ('Spring', '2026-09-01', '2026-09-03'), ('Summer', '2026-08-15', '2026-08-20')), start=1):
        response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
            'dataset_files': (f'{campaign}.csv', BytesIO((
                'Campaign,Operator,City,Mean_Data_Rate,Test_Name,Test_Start_Time\n'
                f'{campaign},A,London,10,HTTP DL,{first_day}\n{campaign},B,Leeds,20,HTTP DL,{last_day}\n'
            ).encode()), 'text/csv'),
        })
        assert response.status_code == 200
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline and core.repository.get_dataset(index)['status'] != 'ready':
            time.sleep(0.05)
    core.repository.add_report_template('Design test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Comparison,,Title and 1 column + Comments,Rate,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,,Top\n'
    ).encode(), is_default=False)
    scans = combined_scans(monkeypatch, 'selected_date_bounds')
    payload = definition().model_dump(mode='json')

    both = client.post('/api/ppt-designer/prepare', json={**payload, 'filters': {'Campaign': ['Spring', 'Summer']}})

    assert both.status_code == 200, both.text
    assert both.json()['date_bounds'] == {'min': '2026-08-15', 'max': '2026-09-03'}
    assert [ids for ids, _fields in scans] == [[1, 2]]
    # Another selection of the same CDRs takes the limits kept for each of them, without a read.
    summer = client.post('/api/ppt-designer/prepare', json={**payload, 'filters': {'Campaign': ['Summer']}})

    assert summer.status_code == 200, summer.text
    assert summer.json()['date_bounds'] == {'min': '2026-08-15', 'max': '2026-08-20'}
    assert [ids for ids, _fields in scans] == [[1, 2]]


def test_design_operator_facets_show_and_filter_the_mapped_operators(client):
    payload = setup_design(client)
    core.repository.replace_operator_mapping_group(None, 'Alpha', ['A'])

    preview = client.post('/api/ppt-designer/prepare', json=payload)

    assert preview.status_code == 200, preview.text
    # The combined table stores A; the Operator Maps show it as Alpha everywhere.
    assert preview.json()['options']['Operator'] == ['Alpha', 'B']
    options = client.post('/api/ppt-designer/filter-options', json={
        'definition': payload, 'field': 'Operator',
    })
    assert options.status_code == 200, options.text
    assert options.json()['values'] == ['Alpha', 'B']
    # Choosing the mapped label selects the rows stored with its source spelling.
    filtered = client.post('/api/ppt-designer/prepare', json={**payload, 'filters': {'Operator': ['Alpha']}})
    assert filtered.status_code == 200, filtered.text
    assert sum(filtered.json()['rows'].values()) > 0


def test_design_operator_patterns_take_the_rows_of_the_spellings_they_match(client):
    payload = setup_design(client)
    core.repository.replace_operator_mapping_group(None, 'Alpha', ['A'])

    # Two rows of A (Alpha in the Operator Maps) and one of B. An exclusion of the label also leaves out its spelling.
    for patterns, rows in ((['-Alpha'], 1), (['-A'], 1), (['NOT Alp*'], 1), (['Al*'], 2), (['B*'], 1), (['Z*'], 0)):
        prepared = client.post('/api/ppt-designer/prepare', json={**payload, 'filters': {'Operator': patterns}})
        assert prepared.status_code == 200, prepared.text
        assert prepared.json()['rows']['data'] == rows, patterns
        assert client.get(f"/api/ppt-designer/data/{prepared.json()['token']}/0").json()['total'] == rows, patterns


def test_design_reuses_normalized_snapshot_for_every_chart(client, monkeypatch):
    payload = setup_design(client)
    import src.modules.ppt_designer as designs_module

    calls = []
    original = designs_module.normalise_operator_aliases

    def tracked(frame):
        calls.append(len(frame))
        return original(frame)

    monkeypatch.setattr(designs_module, 'normalise_operator_aliases', tracked)
    preview = client.post('/api/ppt-designer/prepare', json=payload).json()
    for index in (0, 1, 2):
        assert client.get(f"/api/ppt-designer/preview/{preview['token']}/{index}.png").status_code == 200
    assert calls == [3]
    assert client.get(f"/api/ppt-designer/preview/{preview['token']}/0.png").status_code == 200
    assert calls == [3]


def test_design_is_available_to_workspace_users(client):
    setup_design(client)
    client.get('/logout')
    client.post('/login', data={'username': 'admin', 'password': 'admin123'})
    assert client.get('/ppt-designer').status_code == 200
    assert client.get('/api/ppt-designer').status_code == 200
    assert client.put('/api/ppt-designer/test', json=definition().model_dump(mode='json')).status_code == 200
    assert 'href="/ppt-designer"' in client.get('/datasets-analysis').text
    client.get('/logout')
    client.post('/login', data={'username': 'demo', 'password': 'demo123'})
    assert client.get('/ppt-designer').status_code == 200
    assert 'href="/ppt-designer"' in client.get('/datasets-analysis').text


def test_design_preparation_can_run_in_the_background(client):
    payload = setup_design(client)

    queued = client.post(
        '/api/ppt-designer/prepare?background=1&preparation_id=background-test', json=payload,
    )

    assert queued.status_code == 202
    assert queued.json() == {'preparation_id': 'background-test', 'status': 'queued'}
    deadline = time.monotonic() + 15
    progress = {}
    while time.monotonic() < deadline:
        response = client.get('/api/ppt-designer/preparation-progress/background-test')
        if response.status_code == 200:
            progress = response.json()
            if progress.get('status') in {'completed', 'failed'}:
                break
        time.sleep(0.05)
    assert progress['status'] == 'completed', progress
    prepared = client.get(f"/api/ppt-designer/prepared/{progress['token']}")
    assert prepared.status_code == 200
    assert prepared.json()['rows']['data'] == 3


def test_adding_combined_columns_keeps_prepared_design_caches_valid(client):
    setup_design(client)
    before = core.repository.get_workspace_state('combined_reporting_updated_data')

    core.mark_combined_reporting_updated(core.repository, 'data', rows_changed=False)

    assert core.repository.get_workspace_state('combined_reporting_updated_data') == before
    assert core.repository.get_workspace_state('combined_reporting_columns_updated_data')
    core.mark_combined_reporting_updated(core.repository, 'data', rows_changed=True)
    assert core.repository.get_workspace_state('combined_reporting_updated_data') != before


def test_design_viewer_cover_matches_the_exported_ppt_cover():
    root = Path(__file__).parents[1] / 'src/web_interface/static'
    script = (root / 'js/ppt_designer.js').read_text(encoding='utf-8')
    stylesheet = (root / 'css/ppt_designer.css').read_text(encoding='utf-8')

    assert "coverGeographyLabel('Region', 'Region', 'Regions', 'All Regions')" in script
    assert "coverGeographyLabel('City', 'City', 'Cities', 'All Cities')" in script
    assert "const coverLabels = isCover ? viewerGeographyLabels() : [];" in script
    # Campaigns sit right above the decorative line, the other details below it, all in white (labels bold).
    assert "if (campaigns) details.append(labelledLine(campaigns, 'ds-structural-campaigns'));" in script
    assert script.index("'ds-structural-campaigns'") < script.index("node('span', undefined, 'ds-structural-rule')")
    # Regions and Cities use the same per-CDR catalogue as the PPT cover.
    assert "api('/geography-options', 'POST', canonicalDesignDefinition(appliedDesignDefinition || definition))" in script
    # The Vodafone-Three red with the Ericsson mark and the logos, as the PPT covers and transitions.
    assert 'background:linear-gradient(128deg,#6f0f16 0%,#9c1b22 45%,#d42b21 100%)' in stylesheet
    assert "svg.setAttribute('class', 'ds-structural-watermark')" in script and 'config.operators_logo' in script
    assert '.ds-structural-slide p.ds-structural-campaigns { margin:0; overflow:hidden; color:#fff;' in stylesheet
    assert '.ds-structural-slide .ds-structural-geography p { margin:0; overflow:hidden; color:#fff;' in stylesheet
    assert 'color:#8bf0a6' not in stylesheet and 'color:#52ddf0' not in stylesheet and 'color:#ff8c42' not in stylesheet
    # The subtitle in the template's yellow; the cover 54/24 pt and each transition 48/40/20 pt of a 13.33 in slide.
    assert '.ds-structural-slide p.ds-structural-subtitle { margin:0; color:#ffc700; font-size:min(2.08cqw,3.7cqh);' in stylesheet
    assert '.ds-structural-cover h3.ds-structural-title { font-size:min(5.63cqw,10cqh); }' in stylesheet
    assert '.ds-structural-slide p.ds-structural-number { margin:0 0 min(2.64cqw,4.7cqh); color:#fff; font-size:min(5cqw,8.9cqh);' in stylesheet
    # Each transition after the cover has its counter: 01, 02…
    assert "String(Math.max(1, index + 1)).padStart(2, '0'), 'ds-structural-number'" in script
    assert "['scope', scope]," in script


def test_design_viewer_scope_selector_confirms_before_rerendering():
    root = Path(__file__).parents[1] / 'src/web_interface'
    template = (root / 'templates/ppt_designer.html').read_text(encoding='utf-8')
    script = (root / 'static/js/ppt_designer.js').read_text(encoding='utf-8')

    assert template.index('id="ds-viewer-scope"') < template.index('id="ds-floating-filters"')
    assert "title: 'Change PPT Scope', confirmLabel: 'Change Scope', tone: 'warning'" in script
    assert "if (!accepted) { control.value = current; return; }" in script
    assert "await preparePart();" in script
    assert "for (const id of ['ds-scope', 'ds-ppt-dataset-scope', 'ds-viewer-scope'])" in script


def test_transition_slide_subtitle_uses_the_template_yellow_accent():
    from pptx import Presentation as PptxPresentation
    from src.modules.cdr_reporting import _named_slide_layout, _set_structural_slide_text

    deck = PptxPresentation(Path(__file__).parents[1] / 'assets/powerpoint-templates/Template_VodafoneThree.pptx')
    slide = deck.slides.add_slide(_named_slide_layout(deck, 'Title Only'))
    _set_structural_slide_text(slide, 'Executive Summary', 'Key findings')

    title = next(shape for shape in slide.placeholders if shape.placeholder_format.type in {1, 3})
    subtitle = title.text_frame.paragraphs[-1]
    assert subtitle.text == 'Key findings'
    assert all(run.font.color.theme_color == MSO_THEME_COLOR.ACCENT_4 for run in subtitle.runs)


def test_report_template_editor_offers_find_and_replace():
    root = Path(__file__).parents[1] / 'src/web_interface'
    template = (root / 'templates/admin.html').read_text(encoding='utf-8')
    script = (root / 'static/js/app.js').read_text(encoding='utf-8')

    assert 'data-catalogue-find-toggle' in template
    assert 'data-catalogue-replace-all' in template
    assert "if ((event.ctrlKey || event.metaKey) && !event.altKey && event.key.toLowerCase() === 'f') {" in script
    assert "title: 'Replace All', confirmLabel: 'Replace All', tone: 'warning'" in script
    assert 'const setEditorCellValue = (cell, value) => {' in script


def test_design_filters_open_in_the_viewer_of_the_design_never_in_the_table():
    root = Path(__file__).parents[1] / 'src/web_interface'
    template = (root / 'templates/ppt_designer.html').read_text(encoding='utf-8')
    script = (root / 'static/js/ppt_designer.js').read_text(encoding='utf-8')

    assert '<div id="ds-filter-parking" hidden><div id="ds-filter-home">' in template
    assert "const parkFilterPanel = () => { $('ds-filter-parking').append($('ds-filter-home')); };" in script
    # No row of the table holds the filters; a Design opens without them, and one just created, duplicated or imported
    # opens in its viewer with its PPT Filters.
    assert 'filterRowCell' not in script and "'Open Filters'" not in script and "'Close Filters'" not in script
    assert 'async function openDesign(id, {showFilters = false} = {}) {' in script
    assert "if (!$('ds-viewer').hidden) openFloatingFilters();" in script
    assert script.count('await openDesignFilters(id)') == 3
    assert '>PPT Filters</button>' in template


def test_cover_geography_falls_back_to_design_filter_values(client):
    payload = setup_design(client)
    # A stale or alias-less catalogue has no Cities; the Design's own City
    # filter values (resolved through aliases) complete the cover geography.
    core.repository.replace_cdr_catalogue(1, vendors=[], regions=[], cities=[])

    options = client.post('/api/ppt-designer/geography-options', json=payload)

    assert options.status_code == 200, options.text
    assert options.json()['cities'] == ['Leeds', 'London']
    assert options.json()['regions'] == []

    core.repository.add_report_template('Cover only', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Quarterly review,Template subtitle,Title Page,,,,Title Slide,,,,,Top\n'
    ).encode(), is_default=False)
    payload['template'] = 'Cover only'
    assert client.put('/api/ppt-designer/cover-only', json=payload).status_code == 200
    job_id = client.post('/api/ppt-designer/cover-only/export-ppt', json={'definition': payload}).json()['job_id']
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.05)
    assert job['status'] == 'ready', job
    with core.repository.connection() as connection:
        row = connection.execute('SELECT output_path FROM design_ppt_jobs WHERE id = ?', (job_id,)).fetchone()
    names = {shape.name: shape.text for shape in Presentation(row['output_path']).slides[0].shapes if shape.name.startswith('design-ppt-')}
    assert names == {'design-ppt-scope': 'Operator Comparison', 'design-ppt-city': 'All Cities'}


def test_ppt_cover_treats_every_filter_value_as_all_cities(client):
    payload = setup_design(client)
    # The catalogue lists one more City than the Design's filter options.
    core.repository.replace_cdr_catalogue(1, vendors=[], regions=[], cities=['Leeds', 'London', 'Manchester'])
    core.repository.add_report_template('Cover only', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Quarterly review,Template subtitle,Title Page,,,,Title Slide,,,,,Top\n'
    ).encode(), is_default=False)
    payload.update(template='Cover only', filters={'City': ['leeds ', 'London']})
    assert client.put('/api/ppt-designer/all-cities', json=payload).status_code == 200

    job_id = client.post('/api/ppt-designer/all-cities/export-ppt', json={'definition': payload}).json()['job_id']
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.05)
    assert job['status'] == 'ready', job
    assert 'City: All Cities' in job['filters']
    assert client.get(f'/api/ppt-designer/ppt-jobs/{job_id}/charts.json').json()['cover']['cities'] == 'All Cities'


def test_design_names_are_unique_whatever_their_nr_mode(client):
    payload = setup_design(client)
    assert client.put('/api/ppt-designer/nsa-one', json=payload).status_code == 200
    duplicate = client.put('/api/ppt-designer/nsa-two', json=payload)
    assert duplicate.status_code == 409
    assert 'A Design with this name already exists' in duplicate.json()['detail']
    # A Design that reads the other NR Mode (with the same template) needs a name of its own.
    sa_payload = {**payload, 'nr_mode': 'SA'}
    assert client.put('/api/ppt-designer/sa-one', json=sa_payload).status_code == 409
    assert client.put('/api/ppt-designer/sa-one', json={**sa_payload, 'name': 'Design SA'}).status_code == 200
    assert client.get('/api/ppt-designer').json()['sa-one']['nr_mode'] == 'SA'


def test_report_template_changes_record_the_last_editor(client):
    import csv
    from src.modules.cdr_reporting import CATALOG_HEADERS

    client.post('/login', data={'username': 'super', 'password': 'super123'})
    row = {header: '' for header in CATALOG_HEADERS}
    row.update({'Slide': '1', 'Slide Tittle': 'T', 'Layout': 'Title Page', 'Chart type': 'Title Slide'})
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=CATALOG_HEADERS)
    writer.writeheader()
    writer.writerow(row)
    content = buffer.getvalue().encode()
    response = client.post(
        '/ppt-designer/templates', data={'catalogue_name': 'Editor audit'},
        files={'catalogue_file': ('x.csv', BytesIO(content), 'text/csv')}, follow_redirects=False,
    )
    assert response.status_code == 303
    row = next(row for row in core.repository.list_report_templates() if row['name'] == 'Editor audit')
    assert row['updated_by'] == 'super'
    page = client.get('/ppt-designer')
    # Templates have no NR Mode: the library (PPT Templates of PPT Designer) has no column for it.
    assert '<th>Template name</th><th>Created</th><th>Last Updated</th><th>Last Updated by</th>' in page.text


def test_combined_rf_design_keeps_samples_and_selected_geography(client):
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    sources = [
        ('data', 'LTE_PCell_RSRP_Avg', [-100, -90, -80]),
        ('voice', '4G_RSRP_Avg_A', [-60]),
        ('speech', 'Playing_RSRP_Avg', [-70]),
    ]
    datasets = {}
    for kind, field, values in sources:
        content = 'Operator,City,Campaign,Test_Start_Time,' + field + '\n'
        content += ''.join(f'EE,London,UK_Q1_2026,2026-01-01,{value}\n' for value in values)
        content += f'EE,Leeds,UK_Q1_2026,2026-01-01,-120\n'
        response = client.post('/datasets-analysis/upload', data={'dataset_kinds': kind}, files={
            'dataset_files': (f'RF_{kind}_NSA.csv', BytesIO(content.encode()), 'text/csv'),
        })
        assert response.status_code == 200, response.text
        datasets[kind] = [max(int(row['id']) for row in core.repository.list_datasets())]
    core.repository.add_report_template('Pooled RF test', (
        'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '1,Combined RF,,Title and 1 column + Comments,RSRP,CDR-All,LTE_RSRP,Histogram Line,Bin Size = 5,Operator,Campaign,Operator,Right\n'
    ).encode(), is_default=False)
    payload = DesignDefinition(name='Pooled RF', template='Pooled RF test', datasets=datasets,
                                  filters={'City': ['London']}).model_dump(mode='json')
    prepared = client.post('/api/ppt-designer/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    assert prepared.json()['slides'][0]['charts'][0]['available'] is True
    token = prepared.json()['token']
    rendered = client.get(f'/api/ppt-designer/chart/{token}/0')
    assert rendered.status_code == 200, rendered.text
    assert 'series' in rendered.json(), rendered.json()
    series = rendered.json()['series'][0]
    assert series['samples'] == 5
    assert sum(series['counts']) == 5
    context = client.get(f'/api/ppt-designer/chart/{token}/0/filter-context')
    assert context.status_code == 200, context.text
    assert 'LTE_RSRP' in context.json()['columns_by_source']['cdr-all']
    assert len(context.json()['dataset_ids']) == 3
    preview = client.post(f'/api/ppt-designer/chart/{token}/0/filter-preview', json={
        'filters': 'Bin Size = 5; LTE_RSRP >= -85',
    })
    assert preview.status_code == 200, preview.text


def test_dynamic_histogram_design_filters_operators_and_updates_source_template_row(client):
    from src.modules.cdr_reporting import CatalogEntry, catalogue_csv
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    content = (
        'Operator,City,Campaign,Test_Start_Time,LTE_PCell_RSRP_Avg,NR_PCell_RSRP_Avg\n'
        'A,London,UK_Q1_2026,2026-01-01,-100,-95\n'
        'A,London,UK_Q2_2026,2026-04-01,-90,-85\n'
        'B,London,UK_Q1_2026,2026-01-01,-105,\n'
        'B,London,UK_Q2_2026,2026-04-01,-95,\n'
        'C,Leeds,UK_Q1_2026,2026-01-01,-110,-100\n'
    )
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('Dynamic_RF_NSA.csv', BytesIO(content.encode()), 'text/csv'),
    })
    assert response.status_code == 200, response.text
    dataset_id = max(int(row['id']) for row in core.repository.list_datasets())
    entries = [CatalogEntry(
        slide=1, slide_title='Dynamic histograms', slide_subtitle='',
        layout='2 rows + dynamic columns, comments down', chart_title=f'{radio} RSRP',
        cdr_source='CDR-Data', kpi=field, chart_type='Histogram Bars',
        filters='Bin Size = 5', rows_split='Operator', columns_split='Campaign',
        legend='Campaign', dynamic_field='Operator',
    ) for radio, field in [('LTE', 'LTE_PCell_RSRP_Avg'), ('NR', 'NR_PCell_RSRP_Avg')]]
    core.repository.add_report_template('Dynamic RF test', catalogue_csv(entries), is_default=False)
    payload = DesignDefinition(name='Dynamic RF', template='Dynamic RF test',
                                  datasets={'data': [dataset_id]}, filters={'City': ['London']}).model_dump(mode='json')
    prepared = client.post('/api/ppt-designer/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    slide = prepared.json()['slides'][0]
    assert slide['layout'] == 'Title + 2 rows + 2 columns + comments down'
    assert [chart['title'] for chart in slide['charts']] == ['LTE RSRP – A', 'LTE RSRP – B', 'NR RSRP – A', 'NR RSRP – B']
    assert [chart['focus_row'] for chart in slide['charts']] == [0, 0, 1, 1]
    token = prepared.json()['token']
    model = client.get(f'/api/ppt-designer/chart/{token}/1').json()
    assert all(item['key'][0] == 'B' for item in model['series'])
    missing = client.get(f'/api/ppt-designer/chart/{token}/3').json()
    assert missing['type'] == 'empty'
    saved = client.post(f'/api/ppt-designer/chart/{token}/1/update-template', json={
        'chart_title': 'LTE RSRP – B', 'filters': 'Bin Size = 10',
    })
    assert saved.status_code == 200, saved.text
    restored = core.load_template_catalogue(core.repository.report_template_content('Dynamic RF test'))
    assert len(restored) == 2
    assert restored[0].filters == 'Bin Size = 10'
    assert restored[0].chart_title == 'LTE RSRP'
    assert restored[1].filters == 'Bin Size = 5'
    for index in (0, 1):
        context = client.get(f'/api/ppt-designer/chart/{token}/{index}/filter-context')
        assert context.status_code == 200, context.text
        assert context.json()['filters'] == 'Bin Size = 10'


def test_dynamic_vendor_only_editor_keeps_operator_rows_after_bin_edit(client):
    from src.modules.cdr_reporting import CatalogEntry, catalogue_csv
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('Dynamic_RF_NSA.csv', BytesIO(
            b'Operator,City,Campaign,Test_Start_Time,LTE_PCell_RSRP_Avg,NR_PCell_RSRP_Avg\n'
            b'A,London,UK_Q1_2026,2026-01-01,-100,-95\n'
            b'A,London,UK_Q2_2026,2026-04-01,-90,-85\n'
            b'B,London,UK_Q1_2026,2026-01-01,-105,\n'
            b'B,London,UK_Q2_2026,2026-04-01,-95,\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200, response.text
    dataset_id = max(int(row['id']) for row in core.repository.list_datasets())
    with core.repository.connection() as connection:
        connection.execute('UPDATE dataset_profiles SET vendor_mapping_applied = 1 WHERE dataset_id = ?', (dataset_id,))
    core.repository.replace_cdr_catalogue(dataset_id, vendors=['A_Ericsson', 'B_Huawei'], regions=[], cities=['London'])
    core.repository.replace_reporting_rows(dataset_id, 'data', pd.DataFrame({
        'Operator': ['A', 'A', 'B', 'B'], 'Operator_Vendor': ['A_Ericsson', 'A_Ericsson', 'B_Huawei', 'B_Huawei'],
        'Vendor': ['Ericsson', 'Ericsson', 'Huawei', 'Huawei'], 'City': ['London'] * 4,
        'Campaign': ['UK_Q1_2026', 'UK_Q2_2026', 'UK_Q1_2026', 'UK_Q2_2026'],
        'Test_Start_Time': ['2026-01-01', '2026-04-01', '2026-01-01', '2026-04-01'],
        'LTE_PCell_RSRP_Avg': [-100, -90, -105, -95], 'NR_PCell_RSRP_Avg': [-95, -85, None, None],
    }))
    entries = [CatalogEntry(
        slide=1, slide_title='Dynamic histograms', slide_subtitle='',
        layout='2 rows + dynamic columns, comments down', chart_title=f'{radio} RSRP',
        cdr_source='CDR-Data', kpi=field, chart_type='Histogram Bars', filters='Bin Size = 5',
        rows_split='Operator', columns_split='Campaign', legend='Campaign', dynamic_field='Operator',
    ) for radio, field in [('LTE', 'LTE_PCell_RSRP_Avg'), ('NR', 'NR_PCell_RSRP_Avg')]]
    core.repository.add_report_template('Dynamic RF vendor test', catalogue_csv(entries), is_default=False)
    payload = DesignDefinition(name='Dynamic RF vendor', template='Dynamic RF vendor test',
                                  datasets={'data': [dataset_id]}, filters={'City': ['London']}).model_dump(mode='json')
    ordinary = client.post('/api/ppt-designer/prepare', json=payload)
    assert ordinary.status_code == 200, ordinary.text
    vendor_payload = {**payload, 'scope': 'multivendor', 'vendor_comparison': 'vendor_only'}
    vendor = client.post('/api/ppt-designer/prepare', json=vendor_payload)
    assert vendor.status_code == 200, vendor.text
    slide = vendor.json()['slides'][0]
    assert [chart['title'] for chart in slide['charts']] == [
        'LTE RSRP – Ericsson', 'LTE RSRP – Huawei', 'NR RSRP – Ericsson', 'NR RSRP – Huawei',
    ]
    token = vendor.json()['token']
    saved = client.post(f'/api/ppt-designer/chart/{token}/1/update-template', json={
        'chart_title': 'LTE RSRP – Huawei', 'filters': 'Bin Size = 10',
    })
    assert saved.status_code == 200, saved.text
    restored = core.load_template_catalogue(core.repository.report_template_content('Dynamic RF vendor test'))
    assert [entry.dynamic_columns_field for entry in restored] == ['Operator', 'Operator']
    assert [entry.rows_split for entry in restored] == ['Operator', 'Operator']
    assert [entry.filters for entry in restored] == ['Bin Size = 10', 'Bin Size = 5']
    for index in range(4):
        context = client.get(f'/api/ppt-designer/chart/{token}/{index}/filter-context')
        assert context.status_code == 200, context.text
        assert context.json()['dynamic_rows_field'] == ''
        assert context.json()['dynamic_columns_field'] == 'Vendor'
        expected_filter = 'Bin Size = 10' if index < 2 else 'Bin Size = 5'
        assert context.json()['filters'].startswith(expected_filter)


def _compact_read_design(client):
    from src.modules.cdr_reporting import CatalogEntry, catalogue_csv
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    common = dict(slide_subtitle='', cdr_source='CDR-Data', rows_split='Operator', columns_split='Campaign', legend='Campaign')
    entries = [
        CatalogEntry(slide=1, slide_title='RSRP', layout='Title + 1 rows + 2 columns', chart_title='CDF',
                     kpi='LTE_PCell_RSRP_Avg', chart_type='CDF Line', filters='LTE_PCell_RSRP_Avg >= -160', **common),
        CatalogEntry(slide=1, slide_title='RSRP', layout='Title + 1 rows + 2 columns', chart_title='Histogram',
                     kpi='LTE_PCell_RSRP_Avg', chart_type='Histogram Line', filters='Bin Size = 5', **common),
        CatalogEntry(slide=2, slide_title='RSRP', layout='Title + 1 rows + 3 columns', chart_title='Average',
                     kpi='LTE_PCell_RSRP_Avg', chart_type='Average Vertical Bars', filters='', **common),
        CatalogEntry(slide=2, slide_title='RSRP', layout='Title + 1 rows + 3 columns', chart_title='Classes',
                     kpi='LTE_PCell_RSRP_Avg', chart_type='Distribution Stacked Vertical Bars',
                     filters='Buckets = -110,-100,-90,-80', **{**common, 'columns_split': 'Campaign × Rate Bucket', 'legend': 'Rate Bucket'}),
        CatalogEntry(slide=2, slide_title='RSRP', layout='Title + 1 rows + 3 columns', chart_title='Threshold',
                     kpi='LTE_PCell_RSRP_Avg', chart_type='Threshold Stacked Vertical Bars', filters='Threshold = -100', **common),
    ]
    # Registered before the upload, so the combined CDR table includes its fields.
    core.repository.add_report_template('Compact reads test', catalogue_csv(entries), is_default=False)
    rows = ['Operator,City,Campaign,Test_Start_Time,LTE_PCell_RSRP_Avg,Call_Status']
    for index in range(240):
        operator = 'ABC'[index % 3]
        campaign = 'UK_Q1_2026' if index % 2 else 'UK_Q2_2026'
        # Many repeated measurements, so grouped reads collapse rows.
        rows.append(f'{operator},London,{campaign},2026-0{1 + index % 6}-01,{-80 - (index % 17) * 2.5},{"Completed" if index % 5 else "Failed"}')
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('Compact_reads_NSA.csv', BytesIO('\n'.join(rows).encode()), 'text/csv'),
    })
    assert response.status_code == 200, response.text
    # The combined CDR table is rebuilt in the background after an upload.
    deadline = time.monotonic() + 30
    while not core.BACKGROUND_TASK_SCHEDULER.is_idle and time.monotonic() < deadline:
        time.sleep(0.05)
    dataset_id = max(int(row['id']) for row in core.repository.list_datasets())
    return DesignDefinition(name='Compact reads', template='Compact reads test',
                               datasets={'data': [dataset_id]}).model_dump(mode='json')


def test_cached_and_grouped_reads_produce_identical_chart_models(client, monkeypatch):
    import src.modules.ppt_designer as designs_module

    payload = _compact_read_design(client)
    prepared = client.post('/api/ppt-designer/prepare', json=payload)
    assert prepared.status_code == 200, prepared.text
    token = prepared.json()['token']
    indexes = [chart['index'] for slide in prepared.json()['slides'] for chart in slide['charts']]
    assert len(indexes) == 5
    cached = [client.post(f'/api/ppt-designer/chart/{token}/{index}/refresh').json() for index in indexes]
    # The columns of the CDRs are kept in memory only.
    assert not list(designs_module.column_cache_dir(core.repository.db_path).rglob('*.pkl'))
    monkeypatch.setattr(designs_module, 'DESIGN_COLUMN_CACHE', False)
    compact = [client.post(f'/api/ppt-designer/chart/{token}/{index}/refresh').json() for index in indexes]
    monkeypatch.setattr(designs_module, 'DESIGN_COMPACT_READS', False)
    complete = [client.post(f'/api/ppt-designer/chart/{token}/{index}/refresh').json() for index in indexes]
    assert all(model['type'] != 'empty' for model in complete)
    assert cached == compact == complete


@pytest.mark.parametrize('changes', [
    {'filters': {'Operator': ['A', 'C']}},
    {'filters': {'Campaign': ['UK_Q1_2026']}},
    {'date_from': '2026-02-01', 'date_to': '2026-04-30'},
    {'filters': {'Operator': ['B'], 'City': ['London']}, 'date_from': '2026-03-01'},
])
def test_column_cache_serves_changed_universes_without_sqlite(client, monkeypatch, changes):
    import src.modules.ppt_designer as designs_module

    payload = _compact_read_design(client)
    baseline = client.post('/api/ppt-designer/prepare', json=payload)
    assert baseline.status_code == 200, baseline.text
    for chart in (chart for slide in baseline.json()['slides'] for chart in slide['charts']):
        assert client.get(f"/api/ppt-designer/chart/{baseline.json()['token']}/{chart['index']}").status_code == 200
    changed = {**payload, **changes}
    prepared = client.post('/api/ppt-designer/prepare', json=changed)
    assert prepared.status_code == 200, prepared.text
    token = prepared.json()['token']
    indexes = [chart['index'] for slide in prepared.json()['slides'] for chart in slide['charts']]
    # Every column and filter of the changed universe is already cached: no
    # CDR row is read from SQLite again.
    from src.modules import design_column_cache
    original_build, original_read = design_column_cache.ColumnCache._build, designs_module.pd.read_sql_query
    reads = []
    monkeypatch.setattr(design_column_cache.ColumnCache, '_build', lambda *args: reads.append('build') or original_build(*args))
    monkeypatch.setattr(designs_module.pd, 'read_sql_query', lambda query, *args, **kwargs: (
        reads.append(query) if 'reporting_rows_' in query else None) or original_read(query, *args, **kwargs))
    cached = [client.get(f'/api/ppt-designer/chart/{token}/{index}').json() for index in indexes]
    assert reads == []
    monkeypatch.setattr(design_column_cache.ColumnCache, '_build', original_build)
    monkeypatch.setattr(designs_module.pd, 'read_sql_query', original_read)
    monkeypatch.setattr(designs_module, 'DESIGN_COLUMN_CACHE', False)
    monkeypatch.setattr(designs_module, 'DESIGN_COMPACT_READS', False)
    sqlite = [client.post(f'/api/ppt-designer/chart/{token}/{index}/refresh').json() for index in indexes]
    assert cached == sqlite


def _wait_for(condition, timeout=60.0):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if condition():
            return True
        time.sleep(0.2)
    return condition()


def test_preloaded_charts_wait_for_the_visible_slide(client, monkeypatch):
    import src.modules.ppt_designer as designs_module

    payload = setup_design(client)
    preview = client.post('/api/ppt-designer/prepare', json=payload).json()
    token = preview['token']
    indexes = [chart['index'] for slide in preview['slides'] for chart in slide['charts'] if chart['available']]
    assert len(indexes) >= 2
    visible_started, release = threading.Event(), threading.Event()
    original = designs_module.catalog_chart_payload
    order = []

    def gated(frame, entry, *args, **kwargs):
        order.append(entry.chart_title)
        if len(order) == 1:
            visible_started.set()
            assert release.wait(10)
        return original(frame, entry, *args, **kwargs)

    monkeypatch.setattr(designs_module, 'catalog_chart_payload', gated)
    results = {}
    visible = threading.Thread(target=lambda: results.setdefault('visible', client.get(f'/api/ppt-designer/chart/{token}/{indexes[0]}')))
    visible.start()
    assert visible_started.wait(10)
    preload = threading.Thread(target=lambda: results.setdefault('preload', client.get(f'/api/ppt-designer/chart/{token}/{indexes[1]}?priority=low')))
    preload.start()
    time.sleep(0.8)
    # The preload is held while the visible chart is being calculated.
    assert 'preload' not in results and len(order) == 1
    release.set()
    visible.join(10)
    preload.join(10)
    assert results['visible'].status_code == 200
    assert results['preload'].status_code == 200
    assert len(order) == 2
    # A preload of a chart that is already calculated is served at once.
    assert client.get(f'/api/ppt-designer/chart/{token}/{indexes[0]}?priority=low').status_code == 200
    assert len(order) == 2


@pytest.mark.parametrize('changes', [
    {},
    {'filters': {'Operator': ['A', 'C']}},
    {'filters': {'Campaign': ['UK_Q1_2026'], 'City': ['London']}, 'date_from': '2026-02-01', 'date_to': '2026-05-31'},
])
def test_column_cache_preparation_counts_match_sqlite(client, monkeypatch, changes):
    import src.modules.ppt_designer as designs_module

    payload = {**_compact_read_design(client), **changes}
    # The first preparation caches every CDR's filter columns.
    assert client.post('/api/ppt-designer/prepare', json=payload).status_code == 200

    def prepared_universe():
        with core.repository.connection() as connection:
            connection.execute('DELETE FROM design_filter_selections')
        response = client.post('/api/ppt-designer/prepare', json=payload)
        assert response.status_code == 200, response.text
        body = response.json()
        return {key: value for key, value in body.items() if key not in {'token', 'selection_id'}}

    counted = []
    original_count = designs_module.count_rows
    monkeypatch.setattr(designs_module, 'count_rows', lambda *args, **kwargs: counted.append(1) or original_count(*args, **kwargs))
    cached = prepared_universe()
    # A complete, unfiltered universe already knows its counts.
    assert bool(counted) == bool(changes)
    monkeypatch.setattr(designs_module, 'DESIGN_COLUMN_CACHE', False)
    assert cached == prepared_universe()


def test_ppt_export_keeps_legend_entries_hidden_in_the_viewer(client):
    payload = setup_design(client)
    design_id = 'legend-ppt'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    prepared = client.post(f'/api/ppt-designer/prepare?design_id={design_id}', json=payload).json()
    slide = next(slide for slide in prepared['slides'] if slide['charts'])
    first = slide['charts'][0]
    first_key = f"{slide['number']}|{first['title']}|0"
    # Hidden in every chart, except the first chart, where it is shown again.
    legend = {'design': {'A': True}, 'charts': {first_key: {'A': False}}}

    def exported_models(response):
        assert response.status_code == 202, response.text
        job_id = response.json()['job_id']
        deadline = time.monotonic() + 30
        job = None
        while time.monotonic() < deadline:
            job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
            if job['status'] in {'ready', 'failed'}:
                break
            time.sleep(0.05)
        assert job is not None and job['status'] == 'ready', job
        with core.repository.connection() as connection:
            row = connection.execute('SELECT output_path FROM design_ppt_jobs WHERE id = ?', (job_id,)).fetchone()
        charts_dir = Path(row['output_path']).parent / 'design-charts'
        return job_id, [json.loads(path.read_text(encoding='utf-8')) for path in sorted(charts_dir.glob('*.model.json'))]

    job_id, models = exported_models(client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
        'definition': payload, 'preparation_token': prepared['token'], 'legend_visibility': legend,
    }))
    with_a = [model for model in models if any(item.get('label') == 'A' and item.get('colour') for item in model['legend']['items'])]
    assert len(with_a) >= 2
    assert with_a[0].get('hidden_legend_items') is None
    assert all(model.get('hidden_legend_items') == ['A'] for model in with_a[1:])
    # A relaunched job keeps the legend state it was created with.
    _job_id, relaunched = exported_models(client.post(f'/api/ppt-designer/ppt-jobs/{job_id}/retry'))
    assert [model.get('hidden_legend_items') for model in relaunched] == [model.get('hidden_legend_items') for model in models]


def test_dynamic_grid_reflow_helpers():
    from types import SimpleNamespace
    from src.modules.ppt_designer import dynamic_card_hidden, dynamic_chart_values, legend_value_hidden, reflowed_dynamic_grid

    # Hidden values close whole grid columns or rows; the others keep their order.
    assert reflowed_dynamic_grid('Title + 2 rows + 4 columns', [False, True, False, False] * 2) == (
        'Title + 2 rows + 3 columns', [0, None, 1, 2, 3, None, 4, 5],
    )
    assert reflowed_dynamic_grid('Title + 3 rows + 2 columns + comments right', [False, False, True, True, False, False]) == (
        'Title + 2 rows + 2 columns + comments right', [0, 1, None, None, 2, 3],
    )
    assert reflowed_dynamic_grid('Title + 1 rows + 2 columns', [True, True]) == ('', [None, None])
    assert reflowed_dynamic_grid('Title + 1 rows + 2 columns', [False, False]) is None
    assert reflowed_dynamic_grid('Title + 1 rows + 2 columns', [True]) is None
    # Same matching as the Canvas renderer: any " · " level, unless shown again.
    assert legend_value_hidden(['VF_SA'], [], 'London · VF_SA')
    assert not legend_value_hidden(['VF_SA'], ['VF_SA'], 'VF_SA')
    vendor_only = SimpleNamespace(dynamic_row_value=None, dynamic_column_value='VF', dynamic_value='VF', chart_title='LTE – VF - All')
    assert dynamic_chart_values(vendor_only) == [['VF', 'VF - All']]
    assert dynamic_card_hidden(dynamic_chart_values(vendor_only), ['VF - All'], [])
    both = SimpleNamespace(dynamic_row_value='London', dynamic_column_value='EE', dynamic_value='EE', chart_title='LTE – London / EE')
    assert dynamic_chart_values(both) == [['London'], ['EE']]


def test_hidden_dynamic_values_close_their_grid_cards_in_the_ppt(client):
    from src.modules.cdr_reporting import CatalogEntry, catalogue_csv
    client.post('/login', data={'username': 'super', 'password': 'super123'})
    response = client.post('/datasets-analysis/upload', data={'dataset_kinds': 'data'}, files={
        'dataset_files': ('Dynamic_Grid_NSA.csv', BytesIO(
            b'Operator,City,Campaign,Test_Start_Time,LTE_PCell_RSRP_Avg,NR_PCell_RSRP_Avg\n'
            b'A,London,UK_Q1_2026,2026-01-01,-100,-95\n'
            b'B,London,UK_Q1_2026,2026-01-01,-105,-97\n'
            b'C,London,UK_Q1_2026,2026-01-01,-110,-100\n'
        ), 'text/csv'),
    })
    assert response.status_code == 200, response.text
    dataset_id = max(int(row['id']) for row in core.repository.list_datasets())
    entries = [CatalogEntry(
        slide=1, slide_title='Dynamic histograms', slide_subtitle='',
        layout='2 rows + dynamic columns, comments down', chart_title=f'{radio} RSRP',
        cdr_source='CDR-Data', kpi=field, chart_type='Histogram Bars', filters='Bin Size = 5',
        rows_split='Operator', columns_split='Campaign', legend='Campaign', dynamic_field='Operator',
    ) for radio, field in [('LTE', 'LTE_PCell_RSRP_Avg'), ('NR', 'NR_PCell_RSRP_Avg')]]
    core.repository.add_report_template('Dynamic grid legend test', catalogue_csv(entries), is_default=False)
    payload = DesignDefinition(name='Dynamic grid', template='Dynamic grid legend test',
                                  datasets={'data': [dataset_id]}).model_dump(mode='json')
    design_id = 'dynamic-grid-legend'
    assert client.put(f'/api/ppt-designer/{design_id}', json=payload).status_code == 200
    prepared = client.post(f'/api/ppt-designer/prepare?design_id={design_id}', json=payload).json()
    slide = prepared['slides'][0]
    assert slide['layout'] == 'Title + 2 rows + 3 columns + comments down'
    assert [chart['dynamic_values'] for chart in slide['charts']] == [[['A']], [['B']], [['C']]] * 2
    smaller = 'Title + 2 rows + 2 columns + comments down'
    assert len(prepared['grid_layouts'][smaller]) == 4

    def exported(legend):
        response = client.post(f'/api/ppt-designer/{design_id}/export-ppt', json={
            'definition': payload, 'preparation_token': prepared['token'], 'legend_visibility': legend,
        })
        assert response.status_code == 202, response.text
        job_id = response.json()['job_id']
        deadline = time.monotonic() + 30
        job = None
        while time.monotonic() < deadline:
            job = next(item for item in client.get('/api/ppt-designer/ppt-jobs').json()['jobs'] if item['id'] == job_id)
            if job['status'] in {'ready', 'failed'}:
                break
            time.sleep(0.05)
        assert job is not None and job['status'] == 'ready', job
        with core.repository.connection() as connection:
            output = Path(connection.execute('SELECT output_path FROM design_ppt_jobs WHERE id = ?', (job_id,)).fetchone()['output_path'])
        return output, json.loads((output.parent / 'design-charts' / 'manifest.json').read_text(encoding='utf-8'))

    # Hiding Operator B removes its grid column from the PPT and its viewer.
    output, manifest = exported({'design': {'B': True}})
    assert [chart['title'] for chart in manifest['charts']] == ['LTE RSRP – A', 'LTE RSRP – C', 'NR RSRP – A', 'NR RSRP – C']
    viewer_slide = manifest['slides'][0]
    assert viewer_slide['layout'] == smaller
    assert [chart['position'] for chart in viewer_slide['charts']] == prepared['grid_layouts'][smaller]
    grid_slides = [item for item in Presentation(output).slides if item.slide_layout.name == smaller]
    assert len(grid_slides) == 1
    assert sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in grid_slides[0].shapes) == 4
    # A slide whose cards are all hidden is left out.
    _output, manifest = exported({'design': {'A': True, 'B': True, 'C': True}})
    assert manifest['charts'] == [] and manifest['slides'] == []


def test_design_closed_during_its_cache_lookup_is_not_prepared():
    script = (Path(__file__).parents[1] / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')

    # A Design closed (or another one opened) while its preparation checks the persistent cache is not prepared,
    # and the Scope and filter panels draw nothing once no Design is open.
    prepare = script[script.index('  async function prepare() {'):script.index('    const pending = (async () => {')]
    assert prepare.startswith('  async function prepare() {\n    if (!definition) return null;')
    assert "if (!definition || activeId !== lookupDesignId) {" in prepare
    assert "throw new DOMException('The Design is no longer open.', 'AbortError');" in prepare
    restore = script[script.index('  async function restorePrepared(id) {'):script.index('  function filterChanged() {')]
    # The answer of a Design closed or changed meanwhile is never shown (it returns before applying it).
    assert restore.index('if (activeId !== id || preparedStateFingerprint(definition) !== fingerprint) return false;\n'
                         '      // (Recorded before the payload is shown') < restore.index('applyPreparedPayload(payload);')
    assert "const host = $('ds-sources'); host.replaceChildren();\n    if (!definition) { renderAutoCdrs(); return; }" in script
    assert 'defaultHost.replaceChildren(); additionalHost.replaceChildren();\n' in script
    facets = script[script.index('  function facets() {'):script.index('    definition.hidden_filters ||= [];')]
    assert 'if (!definition) return;' in facets


def test_view_design_lists_its_slides_at_the_left_with_the_navigation_below_the_comments():
    root = Path(__file__).parents[1]
    template = (root / 'src/web_interface/templates/ppt_designer.html').read_text(encoding='utf-8')
    script = (root / 'src/web_interface/static/js/ppt_designer.js').read_text(encoding='utf-8')
    styles = (root / 'src/web_interface/static/css/ppt_designer.css').read_text(encoding='utf-8')
    # The slides at the left of the slide, and Presentation with First / Previous / Next / Last below its comments.
    assert '<div class="ds-viewer-body"><nav class="ds-slide-thumbnails" id="ds-slide-thumbnails" aria-label="Slides of the Design"></nav>' in template
    viewer = template[template.index('<div class="ds-viewer-body">'):template.index('<div class="ds-chart-expanded-overlay"')]
    assert viewer.index('id="ds-charts"') < viewer.index('</details>') < viewer.index('<footer class="ds-viewer-navigation-footer">')
    assert viewer.index('<footer class="ds-viewer-navigation-footer">') < viewer.index('id="ds-presentation"') < viewer.index('id="ds-last"')
    # Each slide in small: the covers and transitions as in the viewer, the charts once calculated, only those in sight.
    assert "if (!slide.charts.length) { structuralDesign(frame, slide); return; }" in script
    assert "loadChartPayload(chart, 'low').then(payload => {" in script and "thumbnailObserver.observe(item)" in script
    assert "item.onclick = () => { stopPresentation(); slideIndex = index; renderSlide(); };" in script
    assert "    renderComments();\n    renderSlideThumbnails();\n" in script
    assert '#ds-viewer:not(.ds-presentation-active) .ds-slide-thumbnails { display:flex;' in styles
    assert '.ds-viewer-body { display:contents; }\n.ds-slide-thumbnails { display:none; }' in styles
    # The keys of the presentations of the Help.
    assert "const backKey = ['ArrowLeft', 'ArrowUp', 'PageUp'].includes(event.key);" in script
    assert "const forwardKey = ['ArrowRight', 'ArrowDown', 'PageDown'].includes(event.key)" in script
    assert "slideIndex = event.key === 'Home' ? 0 : prepared.slides.length - 1;" in script
    # The wheel over the slide changes slide, Ctrl/Cmd+wheel zooms the chart under the pointer, and the focus follows the slide.
    assert "$('ds-charts').addEventListener('wheel', event => {" in script
    assert "globalThis.zoomCanvasChartAt(canvas, (globalThis.getCanvasChartZoom?.(canvas) || 1) * Math.exp(-event.deltaY * 0.0015), event.clientX, event.clientY);" in script
    assert "nav.querySelector('.is-current')?.focus({preventScroll: true});" in script
    charts = (root / 'src/web_interface/static/js/canvas_charts.js').read_text(encoding='utf-8')
    assert 'function zoomChartAt(canvas, requestedZoom, clientX, clientY) {' in charts and 'globalThis.zoomCanvasChartAt = zoomChartAt;' in charts


@pytest.mark.parametrize('changes', [
    {},
    {'filters': {'Operator': ['A', 'C']}},
    {'date_from': '2026-02-01', 'date_to': '2026-04-30'},
])
def test_designs_give_the_same_results_once_the_copies_are_the_only_store(client, monkeypatch, changes):
    from src.modules import cdr_store

    payload = {**_compact_read_design(client), **changes}
    # The keys of the caches (they follow the revisions of the data) and the times.
    unstable = {'token', 'selection_key', 'prepared_at', 'generated_at', 'elapsed', 'elapsed_ms', 'cached', 'from_cache',
                'restored', 'timing'}

    def stable(value):
        if isinstance(value, dict):
            return {key: stable(item) for key, item in value.items() if key not in unstable}
        if isinstance(value, list):
            return [stable(item) for item in value]
        return value

    def results():
        prepared = client.post('/api/ppt-designer/prepare', json=payload)
        assert prepared.status_code == 200, prepared.text
        token = prepared.json()['token']
        indexes = [chart['index'] for slide in prepared.json()['slides'] for chart in slide['charts']]
        charts = [client.post(f'/api/ppt-designer/chart/{token}/{index}/refresh').json() for index in indexes]
        pages = [client.get(f'/api/ppt-designer/data/{token}/{index}', params={'include_filter_values': True}).json()
                 for index in indexes]
        options = client.post('/api/ppt-designer/filter-options/batch', json={
            'definition': payload, 'fields': ['Operator', 'Campaign', 'City']}).json()
        return stable(prepared.json()), charts, stable(pages), options

    with_tables = results()
    # The copies become the only store: the CDR and combined tables are emptied, and the Designs read again.
    core.repository.set_workspace_state(cdr_store.ONLY_COPIES_STATE_KEY, '1')
    repository = core.Repository(core.active_workspace.database_path, global_db_path=core.repository.global_db_path)
    cdr_store.start_store_sync(repository, lambda run: run(), lambda: cdr_store.included_cdrs(repository))
    assert cdr_store.empty_combined_table(core.repository, 'data')
    with core.repository.connection() as connection:
        assert connection.execute('SELECT COUNT(*) FROM main.reporting_rows_data').fetchone()[0] == 0
        assert cdr_store.shadows(connection)
    core.mark_combined_reporting_updated(core.repository, 'data', rows_changed=True)

    prepared, charts, pages, options = results()
    assert {key: (prepared.get(key), with_tables[0].get(key)) for key in {*prepared, *with_tables[0]}
            if prepared.get(key) != with_tables[0].get(key)} == {}
    assert charts == with_tables[1]
    assert pages == with_tables[2]
    assert options == with_tables[3]
