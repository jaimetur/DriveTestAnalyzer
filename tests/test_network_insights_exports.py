from pathlib import Path
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from pptx import Presentation

import io

from PIL import Image

from src.modules.network_insights_export import export_network_insights_powerpoint, export_network_insights_word, rf_quality_table

TEMPLATE = Path(__file__).resolve().parents[1] / 'assets/powerpoint-templates/Template_01.pptx'


def _png(*_args) -> bytes:
    output = io.BytesIO()
    Image.new('RGB', (160, 90), (230, 236, 245)).save(output, 'PNG')
    return output.getvalue()


def _analysis(comparison=False):
    rows = [{
        'operator': f'3 · Edinburgh · LTE · 2026-Q{1 + index % 2}', 'samples': 5751,
        'rsrp_mean': -89.1, 'rsrp_median': -90.4, 'rsrp_p10': -104.8,
        'low_coverage_share': 2.4, 'sinr_mean': 7.3, 'sinr_median': 6.9,
        'sinr_p10': -0.8, 'high_interference_share': 13.8, 'observed_enodebs': 140, 'observed_cells': 512,
        'deltas': {'rsrp_median': 1.2, 'low_coverage_share': -0.3, 'sinr_median': 0.2, 'high_interference_share': -1.1},
    } for index in range(25)]
    analysis = {'rf_rows': rows, 'maps': {}, 'spectrum': {}, 'charts': {'rsrp_cdf': {'type': 'cdf'}, 'sinr_cdf': {'type': 'cdf'}}}
    if comparison:
        analysis['comparison'] = {'latest': '2026-Q2', 'previous': '2026-Q1'}
    return analysis


def test_network_ppt_preserves_dashboard_template_and_fits_group_cells(tmp_path):
    output = tmp_path / 'network.pptx'
    source = _analysis()
    export_network_insights_powerpoint(output, source, {}, {'group_label': 'Operator → City → Technology → Campaign'}, _png, template=TEMPLATE)
    deck, reference = Presentation(output), Presentation(TEMPLATE)
    assert (deck.slide_width, deck.slide_height) == (reference.slide_width, reference.slide_height)
    assert deck.slides[0].slide_layout.name == 'Title Page'
    assert deck.slides[-1].slide_layout.name == 'Black logo end slide'
    assert len(deck.slide_layouts) == len(reference.slide_layouts)
    columns, values, _bars = rf_quality_table(source)
    # The RF Quality table sits under the CDFs and continues on further slides; no separate overview table is exported.
    tables = [shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table and shape.table.cell(0, 1).text == 'Samples']
    assert len(tables) > 1
    actual = []
    for table in tables:
        assert [cell.text for cell in table.rows[0].cells] == columns
        assert table.columns[0].width > table.columns[1].width
        for row in table.rows:
            assert all(cell.text_frame.word_wrap is False for cell in row.cells)
            assert all(len(cell.text_frame.paragraphs) == 1 for cell in row.cells)
            assert all(run.font.size.pt >= 8 for cell in row.cells for run in cell.text_frame.paragraphs[0].runs)
        actual.extend([[cell.text for cell in row.cells] for row in list(table.rows)[1:]])
    assert actual == values
    # The same XML content; python-pptx may write a part again with another XML declaration.
    canonical = lambda xml: etree.tostring(etree.fromstring(xml), method='c14n')
    with zipfile.ZipFile(output) as generated, zipfile.ZipFile(TEMPLATE) as original:
        for part in ('ppt/theme/theme1.xml', 'ppt/slideMasters/slideMaster1.xml'):
            assert canonical(generated.read(part)) == canonical(original.read(part))


def test_network_word_uses_landscape_fixed_single_line_tables(tmp_path):
    output = tmp_path / 'network.docx'
    source = _analysis()
    export_network_insights_word(output, source, {}, {}, _png)
    document = Document(output)
    section = document.sections[0]
    assert section.page_width > section.page_height
    columns, values, _bars = rf_quality_table(source)
    tables = [table for table in document.tables if [cell.text for cell in table.rows[0].cells] == columns]
    assert len(tables) == 1
    table = tables[0]
    assert table.autofit is False
    assert table.columns[0].width > table.columns[1].width
    assert sum(column.width for column in table.columns) <= section.page_width - section.left_margin - section.right_margin + 9144
    assert [[cell.text for cell in row.cells] for row in list(table.rows)[1:]] == values
    assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    for row in table.rows:
        assert row._tr.trPr.find(qn('w:cantSplit')) is not None
        for cell in row.cells:
            assert cell._tc.tcPr.find(qn('w:noWrap')) is not None
            assert len(cell.paragraphs) == 1 and '\n' not in cell.text
            assert all(run.font.size.pt >= 8 for run in cell.paragraphs[0].runs)


def test_network_summaries_include_all_grouped_deployment_views_and_exclude_full_inventory(client, tmp_path, monkeypatch):
    import pandas as pd
    import src.DriveTestAnalyzer as app_module
    from src.modules import network_insights as ni, network_insights_export as export
    from src.modules.repository import local_now_iso

    client.post('/login', data={'username': 'admin', 'password': 'admin123'}, follow_redirects=False)
    repo = app_module.repository

    def ready(name, kind, frame):
        source = tmp_path / name
        source.write_text('source', encoding='utf-8')
        identifier, _ = repo.add_dataset(name, str(source), 'admin')
        repo.replace_dataset_rows(identifier, frame)
        repo.update_dataset_profile(identifier, status='ready', progress=100, dataset_kind=kind,
                                    nr_mode='NSA' if kind == 'data' else None, row_count=len(frame),
                                    column_count=len(frame.columns), processed_at=local_now_iso())
        return identifier

    data_id = ready('CDR.xlsx', 'data', pd.DataFrame({'Operator': ['VF'] * 3, 'Campaign': ['2026-Q1'] * 3,
                                                     'LTE_PCell_RSRP_Avg': [-90, -95, -100], 'LTE_PCell_SINR_Avg': [10, 12, 8]}))
    ready('mapping.xlsx', 'mapping_vodafone', pd.DataFrame({'Site_ID': ['S1'], 'Cell ID': ['1'], 'eMOCNScenario': ['NNS']}))
    ready('clusters.geojson', 'clusters', pd.DataFrame({'Cluster_Field': ['Cluster']}))
    (tmp_path / 'clusters.geojson').write_text('{"type":"FeatureCollection","features":[{"type":"Feature","properties":{"Name":"Cluster A"},"geometry":{"type":"Polygon","coordinates":[[[0,0],[1,0],[1,1],[0,1],[0,0]]]}}]}', encoding='utf-8')
    captured = []

    def writer(destination, _analysis, deployment, _selection, _render):
        captured.append(deployment)
        destination.write_bytes(b'export')
        return destination

    monkeypatch.setattr(export, 'export_network_insights_word', writer)
    monkeypatch.setattr(export, 'export_network_insights_powerpoint', writer)
    for kind in ['word', 'powerpoint']:
        app_module.write_network_insights_summary({'datasets': {'data': [data_id]}}, kind, tmp_path / f'{kind}.output')
    for deployment in captured:
        # The grouped views come from the inventory snapshot, as on the page.
        assert [group['group'] for group in deployment['groupings']] == list(ni.SUMMARY_DEPLOYMENT_GROUPS)
        tables = export.deployment_tables(deployment)
        assert len(tables) == len(ni.SUMMARY_DEPLOYMENT_GROUPS)
        assert {columns[0] for _title, columns, _rows in tables} == {ni.ii.DEPLOYMENT_DIMENSIONS[group] for group in ni.SUMMARY_DEPLOYMENT_GROUPS}
        assert deployment['snapshot']['selection'][0]['cells'] == 1
        assert all('Full Site' not in title for title, _columns, _rows in tables)
        assert deployment['cluster_inventories']['rows'] == [{'operator': 'Vodafone', 'sites': 1, 'cells': 1}]


def test_cluster_sites_sources_are_both_present_in_ppt_and_word(tmp_path):
    analysis = _analysis()
    deployment = {'cluster_inventories': {'rows': [{'operator': 'Vodafone', 'sites': 37201, 'cells': 516898},
                                                 {'operator': 'Three', 'sites': 16107, 'cells': 140127}]}}
    ppt = tmp_path / 'density.pptx'
    word = tmp_path / 'density.docx'
    export_network_insights_powerpoint(ppt, analysis, deployment, {}, lambda *_args: b'', template=TEMPLATE)
    export_network_insights_word(word, analysis, deployment, {}, lambda *_args: b'')
    ppt_tables = [shape.table for slide in Presentation(ppt).slides for shape in slide.shapes if shape.has_table]
    word_tables = Document(word).tables
    for tables in [ppt_tables, word_tables]:
        observed = [table for table in tables if table.cell(0, 1).text == 'Observed eNodeBs']
        inventory = [table for table in tables if table.cell(0, 1).text == 'Inventory sites']
        assert sum(len(table.rows) - 1 for table in observed) == 25
        assert len(inventory) == 1
        assert [[cell.text for cell in row.cells] for row in list(inventory[0].rows)[1:]] == [
            ['Vodafone', '37,201', '516,898'], ['Three', '16,107', '140,127']]


def test_serving_cells_5g_availability_and_snapshot_tables_are_exported(tmp_path):
    analysis = _analysis()
    analysis['serving'] = {
        'overshoot_km': 5, 'max_distance_km': 30,
        'summary': [{'operator': 'VF', 'samples': 10, 'with_cell': 10, 'matched': 9, 'matched_share': 90.0, 'median_km': 0.8,
                     'p90_km': 2.1, 'overshoot_share': 10.0, 'hosted_elsewhere_share': 20.0}],
        'by_band': [{'operator': 'VF', 'value': 'B20', 'samples': 9, 'median_km': 0.8, 'overshoot_share': 10.0}],
        'by_vendor': [], 'by_host': [],
        'overshooting_cells': [{'operator': 'VF', 'cell': '6692112', 'site': '57744', 'band': 'B3', 'vendor': 'Ericsson', 'network': 'VF',
                                'host': 'O2', 'samples_beyond': 3, 'samples': 4, 'median_km': 7.5, 'max_km': 9.1}],
        'nr_availability': {'rows': [{'operator': 'VF', 'anchored': 9, 'on_nr_sites': 6, 'without_nr_on_nr_sites': 2, 'gap_share': 33.3,
                                      'without_nr_elsewhere_share': 80.0}], 'sites': []},
        'distance_cdf': {'type': 'empty', 'message': 'No chart'},
    }
    deployment = {'snapshot': {
        'selection': [{'operator': 'VF', 'snapshot_date': '2026-09-17', 'file_name': 'cells.zip', 'sites': 2, 'cells': 6,
                       'technologies': {'LTE': 5, 'NR': 1}}],
        'host_matrix': {'owners': ['VF'], 'hosts': ['O2', 'VF'], 'cells': {'VF|VF': 5, 'VF|O2': 1}, 'sites': {}},
        'quality': [{'operator': 'VF', 'cells': 6, 'not_in_service': 0, 'locked_or_disabled': 0, 'without_position': 0, 'without_azimuth': 1,
                     'without_band': 0, 'without_vendor': 0, 'duplicate_identities': 0}],
        'spectrum': [{'operator': 'VF', 'technology': 'LTE', 'band': 'B20', 'band_class': 'Low', 'sites': 2, 'cells': 5, 'carriers': 1,
                      'typical_bandwidth_mhz': 10, 'deployed_mhz': 10}],
        'density': {'region': [{'area': 'North', 'operator': 'VF', 'sites': 2, 'nr_sites': 1, 'area_km2': None, 'sites_per_km2': None,
                                'median_isd_km': 1.3, 'p90_isd_km': 1.3}]},
    }}
    ppt, word = tmp_path / 'serving.pptx', tmp_path / 'serving.docx'
    export_network_insights_powerpoint(ppt, analysis, deployment, {}, lambda *_args: b'', template=TEMPLATE)
    export_network_insights_word(word, analysis, deployment, {}, lambda *_args: b'')
    ppt_headers = {shape.table.cell(0, 1).text for slide in Presentation(ppt).slides for shape in slide.shapes if shape.has_table}
    word_headers = {table.cell(0, 1).text for table in Document(word).tables}
    for headers in (ppt_headers, word_headers):
        assert {'Samples', 'Cell', 'Samples on LTE cells', 'Snapshot date', 'Technology', 'Network'} <= headers


def test_nr_cell_of_nsa_samples_mimo_and_cell_status_tables_are_exported(tmp_path):
    analysis = _analysis()
    analysis['maps'] = {'operator': 'VF', 'cell_status': {'radius_km': 3, 'known_share': 34.2, 'rows': [
        {'site': 'VF2', 'network': 'VF', 'host': 'VF', 'vendor': 'Ericsson', 'summary': 'Locked: 1 cell (B20)', 'on_air': True,
         'area_kind': 'coverage', 'area': 'Leeds', 'distance_km': 0.66, 'states': [{'state': 'Locked', 'cells': 1}]}]}}
    analysis['serving'] = {
        'overshoot_km': 5, 'max_distance_km': 30, 'summary': [], 'by_band': [], 'by_vendor': [], 'by_host': [], 'overshooting_cells': [],
        'nr_leg': {'rows': [{'operator': 'VF', 'with_nr_pci': 12, 'found': 12, 'found_share': 100.0, 'anchor_site_share': 50.0,
                             'only_candidate_share': 25.0, 'nearest_share': 25.0, 'median_km': 0.4, 'p90_km': 2.1,
                             'vendors': [{'vendor': 'Ericsson', 'share': 100.0}]}],
                   'by_vendor': [{'operator': 'VF', 'value': 'Ericsson', 'samples': 12, 'median_km': 0.4, 'overshoot_share': 0.0}]},
        'mimo': [{'operator': 'VF', 'technology': 'LTE', 'mimo': '4T', 'tx_antennas': 4, 'modes': ['4T4R'], 'samples': 3, 'mean_rank': 1.5,
                  'rank2_share': 50.0, 'mean_layers': None, 'rank_use': 37.5, 'rank_classes': []}],
        'nr_availability': {'rows': [{'operator': 'VF', 'anchored': 9, 'on_nr_sites': 6, 'without_nr_on_nr_sites': 2, 'gap_share': 33.3,
                                      'without_nr_elsewhere_share': 80.0, 'nr_found_share': 100.0, 'nr_on_anchor_share': 50.0,
                                      'nr_median_km': 1.2}], 'sites': [],
                            'nr_sites': [{'operator': 'VF', 'site': 'VF0', 'vendor': 'Ericsson', 'samples': 6, 'from_anchor_share': 50.0,
                                          'median_km': 1.9}]},
        'distance_cdf': {'type': 'empty', 'message': 'No chart'},
    }
    ppt, word = tmp_path / 'phase3.pptx', tmp_path / 'phase3.docx'
    export_network_insights_powerpoint(ppt, analysis, {}, {}, lambda *_args: b'', template=TEMPLATE)
    export_network_insights_word(word, analysis, {}, {}, lambda *_args: b'')
    ppt_tables = [shape.table for slide in Presentation(ppt).slides for shape in slide.shapes if shape.has_table]
    for tables in (ppt_tables, Document(word).tables):
        headers = {table.cell(0, 1).text for table in tables}
        assert {'Samples with NR PCI', 'NR vendor', 'Technology', 'Network → Host', 'NR site'} <= headers
        five_g = next(table for table in tables if table.cell(0, 1).text == 'Samples on LTE cells')
        assert [cell.text for cell in five_g.rows[0].cells][-3:] == ['NR cell found', 'NR on the anchor site', 'Median km to the NR site']
        status = next(table for table in tables if table.cell(0, 1).text == 'Network → Host')
        assert [cell.text for cell in status.rows[1].cells][3:6] == ['Locked: 1 cell (B20)', 'Yes', 'Coverage · Leeds']


def _bandwidth_spectrum():
    classes = [{'label': label, 'lower': lower, 'upper': upper, 'colour': colour, 'width': width} for label, lower, upper, colour, width in (
        ('200+ MHz', 200, None, '#1E6B3A', 225), ('140–199 MHz', 140, 200, '#3E9455', 169.5), ('100–139 MHz', 100, 140, '#8CC68A', 119.5),
        ('60–99 MHz', 60, 100, '#F2C66D', 79.5), ('40–59 MHz', 40, 60, '#EE9E4E', 49.5), ('25–39 MHz', 25, 40, '#E2685A', 32),
        ('Under 25 MHz', 0, 25, '#B5132B', 12.5))]
    cell = lambda column, environment, shares, narrow: {  # noqa: E731
        'column': column, 'environment': environment, 'tests': 20, 'known': 20, 'counts': [int(share / 5) for share in shares],
        'shares': shares, 'narrow_share': narrow, 'median_mhz': 120.0}
    scope = {'tests': 80, 'bands': [
        {'operator': '3 · 2026-Q3', 'technology': 'NR', 'band': 'n78', 'duplex': 'TDD', 'frequency_mhz': 3500, 'samples': 40,
         'unknown': 2, 'bandwidths': {'40': 8, '100': 30}, 'typical_bandwidth_mhz': 100.0}],
        'ladders': {'columns': ['3', 'EE'], 'environments': ['Cities', 'Towns'], 'unknown': 3, 'without_environment': 0, 'tests': 80, 'cells': [
            cell('3', 'Cities', [10, 40, 30, 10, 0, 0, 10], 10.0), cell('3', 'Towns', [0, 20, 30, 20, 0, 0, 30], 30.0),
            cell('EE', 'Cities', [0, 0, 50, 20, 20, 0, 10], 10.0), cell('EE', 'Towns', [0, 0, 20, 30, 20, 10, 20], 20.0)]}}
    return {'bandwidth': {'classes': classes, 'narrow_below_mhz': 25, 'scope_labels': {'download': 'Download tests', 'all': 'All Data tests'},
                          'scopes': {'download': scope, 'all': scope}, 'column_label': 'Operator',
                          'column_colours': {'3': '#000000', 'EE': '#00B2A9'}}}


def test_bandwidth_views_and_nr_sites_by_duplex_mode_are_exported(tmp_path):
    from src.modules.network_insights_export import bandwidth_ladder_image, bandwidth_tables, nr_duplex_table

    analysis = {**_analysis(), 'group_label': 'Operator → Campaign', 'spectrum': _bandwidth_spectrum()}
    tables = {title: (columns, rows) for title, columns, rows in bandwidth_tables(analysis)}
    columns, rows = tables['Spectrum · Bandwidth per band (Download tests)']
    assert columns == ['Operator → Campaign', 'Technology', 'Band', 'Duplex', 'Samples', '40 MHz', '100 MHz', 'Unknown', 'Typical']
    assert rows[0][5:] == ['8 (20%)', '30 (75%)', '2', '100 MHz']
    columns, rows = tables['Spectrum · Total bandwidth, LTE and NR (Download tests)']
    assert columns[:3] == ['Operator', 'Environment', 'Tests'] and columns[-2:] == ['Below 25 MHz', 'Median']
    assert rows[1][:2] == ['3', 'Towns'] and rows[1][-2] == '30.0%'
    image = bandwidth_ladder_image(analysis)
    with Image.open(image) as raster:
        assert raster.width > raster.height > 400
    deployment = {'nr_duplex': {'radius_m': 50, 'rows': [
        {'operator': name, 'sites': 10, 'tdd_sites': 8, 'fdd_sites': 6, 'both_sites': 4, 'tdd_only_sites': 4, 'fdd_only_sites': 2,
         'physical_sites': 9, 'physical_tdd_sites': 7, 'physical_fdd_sites': 6} for name in ('VF', 'All networks')]}}
    assert nr_duplex_table(deployment)[0][1][-3:] == ['Physical sites (within 50 m)', 'Physical NR TDD', 'Physical NR FDD']

    output = tmp_path / 'network.pptx'
    export_network_insights_powerpoint(output, analysis, deployment, {}, _png, template=TEMPLATE)
    titles = [shape.text_frame.text for slide in Presentation(output).slides for shape in slide.shapes
              if shape.has_text_frame and shape.text_frame.text.startswith('Spectrum ·')]
    assert {'Spectrum · NR sites by duplex mode', 'Spectrum · Bandwidth per band (Download tests)',
            'Spectrum · Total bandwidth, LTE and NR (Download tests)'} <= set(titles)
    pictures = [shape for slide in Presentation(output).slides for shape in slide.shapes if shape.shape_type == 13]
    assert pictures
    output = tmp_path / 'network.docx'
    export_network_insights_word(output, analysis, deployment, {}, _png)
    document = Document(output)
    headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name.startswith('Heading')]
    assert 'Spectrum · Total bandwidth, LTE and NR (Download tests)' in headings and 'Spectrum · NR sites by duplex mode' in headings
    assert document.inline_shapes
