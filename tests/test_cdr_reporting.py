from __future__ import annotations

import json
from contextlib import contextmanager
from dataclasses import replace
import pandas as pd
import csv
import pytest
from PIL import Image, ImageDraw, ImageFont
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import MagicMock, call, patch
from pptx import Presentation

from src.modules.cdr_reporting import CATALOG_HEADERS, TRAILING_DYNAMIC_CATALOG_HEADERS, CatalogEntry, _apply_catalog_filters, _apply_catalog_grouping, _cdf_campaign_line_widths, _cdf_plot_geometry, _cdf_terminal_x_maximum, _cdf_visible_points, _draw_adjacent_stacked_bar_label, _draw_chart_legend, _draw_configured_bar_label, _draw_inside_bar_label, _draw_top_column_group_separators, _hierarchical_complete_keys, _hierarchical_unique_keys, _hierarchy_caption_spans, _hierarchy_group_colours, _hierarchy_spans, _horizontal_legend_columns, _layout_chart_frames, _legend_dimensions, _legend_labels, _named_slide_layout, _render_cdf_line, _render_failure_count, _render_failure_count_hierarchy, _render_map, _render_mean_column, _render_stacked_distribution, _render_status_100, _render_table, _resolved_legend_items, _series_colours, _series_line_dashes, _status_chart_categories, assign_cdr_vendors, catalog_chart_hover_targets, catalog_chart_payload, catalogue_csv, convert_catalog_csv, ensure_vendor_group, enrich_multivendor, load_catalog_csv, normalise_operator_aliases, parse_axis_range, parse_calculated_dimensions, parse_catalog_csv, parse_catalog_filters, parse_catalog_grouping, parse_kpi_expression, parse_label_format, parse_label_position, parse_legend_position, parse_template_boolean, prepare_catalog_chart_preview_frame, prepare_multivendor_catalog_entry, render_catalog_chart_preview, vendor_from_cells
from src.modules.cdr_reporting import _draw_fitted_aggregation_header


CHART_MAPPING_ATTRS = {
    'operator_mapping_groups': [
        {'canonical': 'VF', 'aliases': ['Vodafone', 'Vodafone UK', 'VFUK'], 'position': 0, 'color': '#E15759'},
        {'canonical': '3', 'aliases': ['Three', 'Three UK', '3 UK'], 'position': 1, 'color': '#F28E2B'},
        {'canonical': 'EE', 'aliases': ['EE UK', 'Everything Everywhere'], 'position': 2, 'color': '#76B7B2'},
        {'canonical': 'O2', 'aliases': ['Telefonica', 'Telefonica O2'], 'position': 3, 'color': '#4E79A7'},
    ],
    'vendor_mapping_groups': [
        {'canonical': 'Ericsson', 'aliases': [], 'position': 0, 'color': '#2E8B57'},
        {'canonical': 'Huawei', 'aliases': [], 'position': 1, 'color': '#E15759'},
        {'canonical': 'Samsung', 'aliases': [], 'position': 2, 'color': '#7B3FB5'},
        {'canonical': 'NSN', 'aliases': [], 'position': 3, 'color': '#4E79A7'},
        {'canonical': 'Mixed Vendor', 'aliases': ['Mixed'], 'position': 4, 'color': '#D9A514'},
        {'canonical': 'Other Vendor', 'aliases': ['Other'], 'position': 5, 'color': '#D9A514'},
        {'canonical': '(blank)', 'aliases': ['Blank'], 'position': 6, 'color': '#7A8791'},
    ],
}


def chart_frame(data: dict[str, object]) -> pd.DataFrame:
    frame = pd.DataFrame(data)
    frame.attrs.update(CHART_MAPPING_ATTRS)
    return frame


def test_kpi_expression_supports_explicit_aggregation_aliases() -> None:
    assert parse_kpi_expression('COUNTD(Test_ID)') == ('Test_ID', 'countd')
    assert parse_kpi_expression('MEAN(`Mean Data Rate`)') == ('Mean Data Rate', 'mean')
    assert parse_kpi_expression('AVERAGE(Mean_Data_Rate)') == ('Mean_Data_Rate', 'mean')
    assert parse_kpi_expression('Test_ID') == ('Test_ID', None)


def test_table_payload_counts_test_ids_across_the_declared_row_hierarchy() -> None:
    frame = chart_frame({
        'Benchmark': ['UK_Q1_2026'] * 4,
        'Subscriber': ['EE'] * 4,
        'G Level 4': ['Belfast'] * 4,
        'Test Result': ['Completed'] * 4,
        'Test_ID': ['A', 'A', 'B', None],
    })
    base = dict(
        slide=3, slide_title='Validation', slide_subtitle='', layout='', chart_title='Test count',
        cdr_source='CDR-Data', chart_type='Table', legend='', filters='',
        rows_split='Benchmark × Subscriber × G Level 4 × Test Result',
        # A KPI accidentally repeated as a column dimension must not split its
        # own aggregation. Cell Assistance now makes the intended KPI syntax explicit.
        columns_split='Test_ID', legend_position='Top',
    )

    count = catalog_chart_payload(frame, CatalogEntry(kpi='COUNT(Test_ID)', **base), prefiltered=True)
    distinct = catalog_chart_payload(frame, CatalogEntry(kpi='COUNTD(Test_ID)', **base), prefiltered=True)

    assert count['headers'] == ['Benchmark', 'Subscriber', 'G Level 4', 'Test Result', 'COUNT(Test_ID)']
    assert count['rows'] == [['UK_Q1_2026', 'EE', 'Belfast', 'Completed', '3']]
    assert distinct['rows'] == [['UK_Q1_2026', 'EE', 'Belfast', 'Completed', '2']]


def test_dynamic_table_pivots_columns_and_exposes_hierarchy_metadata() -> None:
    frame = chart_frame({
        'Benchmark': ['UK_Q1_2026'] * 6,
        'Subscriber': ['EE'] * 6,
        'Test_Result': ['Completed', 'Cutoff', 'Failed', 'Completed', 'Failed', 'Completed'],
        'G_Level_4': ['Bristol', 'Bristol', 'Bristol', 'Belfast', 'Belfast', 'Belfast'],
        'Test_ID': ['A', 'B', 'C', 'D', 'E', 'F'],
    })
    entry = CatalogEntry(
        slide=3, slide_title='Validation', slide_subtitle='', layout='', chart_title='Test count',
        cdr_source='CDR-Data', kpi='COUNT(Test_ID)', chart_type='Dynamic Table', legend='Subscriber', filters='',
        rows_split='Benchmark × Subscriber × G Level 4 × Test_Result', columns_split='G Level 4',
        legend_position='',
    )

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['type'] == 'table'
    assert model['dynamic'] is True
    assert model['row_dimension_count'] == 3
    assert model['column_heading'] == 'G Level 4'
    assert model['legend']['position'] == 'top'
    assert model['headers'] == ['Benchmark', 'Subscriber', 'Test_Result', 'Belfast', 'Bristol']
    assert model['column_keys'] == [['Belfast'], ['Bristol']]
    assert model['rows'] == [
        ['UK_Q1_2026', 'EE', 'Completed', '2', '1'],
        ['UK_Q1_2026', 'EE', 'Cutoff', '', '1'],
        ['UK_Q1_2026', 'EE', 'Failed', '1', '1'],
    ]


def test_dynamic_table_uses_operator_mapping_order_by_default() -> None:
    frame = chart_frame({
        'Subscriber': ['VF_SA', 'VF_UK', 'O2', 'EE', '3'],
        'Test_Result': ['Completed'] * 5,
        'G_Level_4': ['London'] * 5,
        'Test_ID': ['A', 'B', 'C', 'D', 'E'],
    })
    frame.attrs['operator_mapping_groups'] = [
        {'canonical': '3', 'aliases': [], 'position': 0, 'color': '#F28E2B'},
        {'canonical': 'EE', 'aliases': [], 'position': 1, 'color': '#76B7B2'},
        {'canonical': 'O2', 'aliases': [], 'position': 2, 'color': '#4E79A7'},
        {'canonical': 'VF_UK', 'aliases': [], 'position': 3, 'color': '#E15759'},
        {'canonical': 'VF_SA', 'aliases': [], 'position': 4, 'color': '#8000FF'},
    ]
    entry = CatalogEntry(
        slide=3, slide_title='Validation', slide_subtitle='', layout='', chart_title='Test count',
        cdr_source='CDR-Data', kpi='COUNT(Test_ID)', chart_type='Dynamic Table', legend='', filters='',
        rows_split='Subscriber × Test_Result', columns_split='G Level 4',
        legend_position='',
    )

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert [row[0] for row in model['rows']] == ['3', 'EE', 'O2', 'VF_UK', 'VF_SA']
    column_model = catalog_chart_payload(
        frame,
        replace(entry, rows_split='Test_Result', columns_split='Subscriber'),
        prefiltered=True,
    )
    assert column_model['headers'][1:] == ['3', 'EE', 'O2', 'VF_UK', 'VF_SA']
    assert column_model['column_keys'] == [['3'], ['EE'], ['O2'], ['VF_UK'], ['VF_SA']]


def test_failure_count_hierarchy_uses_operator_mapping_order() -> None:
    frame = chart_frame({
        'Operator': ['VF_SA', 'VF_UK', 'O2', 'EE', '3'],
        'Campaign': ['2026-Q2'] * 5,
        'Session_Type': ['VoLTE'] * 5,
        'City': ['London'] * 5,
        'Call_Status': ['Failed'] * 5,
        'Test_ID': ['A', 'B', 'C', 'D', 'E'],
    })
    frame.attrs['operator_mapping_groups'] = [
        {'canonical': '3', 'aliases': [], 'position': 0, 'color': '#F28E2B'},
        {'canonical': 'EE', 'aliases': [], 'position': 1, 'color': '#76B7B2'},
        {'canonical': 'O2', 'aliases': [], 'position': 2, 'color': '#4E79A7'},
        {'canonical': 'VF_UK', 'aliases': [], 'position': 3, 'color': '#E15759'},
        {'canonical': 'VF_SA', 'aliases': [], 'position': 4, 'color': '#8000FF'},
    ]
    entry = CatalogEntry(
        slide=3, slide_title='Validation', slide_subtitle='', layout='', chart_title='Failures',
        cdr_source='CDR-Voice', kpi='COUNT(Test_ID)', chart_type='Count Stacked Horizontal Bars',
        legend='Call Status', filters='', rows_split='Session Type × City',
        columns_split='Operator × Campaign', legend_position='Right',
    )

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['column_keys'] == [
        ['3', '26-Q2'], ['EE', '26-Q2'], ['O2', '26-Q2'],
        ['VF_UK', '26-Q2'], ['VF_SA', '26-Q2'],
    ]


def test_failure_count_hierarchy_keeps_operator_mapping_order_when_city_rows_are_sparse() -> None:
    """Column order must not depend on the first operator present in a city."""
    frame = chart_frame({
        'Operator': ['VF_SA', '3', 'O2', 'EE', 'VF_UK'],
        'Campaign': ['2026-Q2'] * 5,
        'Session_Type': ['MultiRAB', 'MultiRAB', 'VoLTE', 'VoLTE', 'VoLTE'],
        'City': ['Belfast', 'Belfast', 'Bristol', 'Bristol', 'London'],
        'Call_Status': ['Failed'] * 5,
        'Test_ID': ['A', 'B', 'C', 'D', 'E'],
    })
    frame.attrs['operator_mapping_groups'] = [
        {'canonical': '3', 'aliases': [], 'position': 0, 'color': '#F28E2B'},
        {'canonical': 'EE', 'aliases': [], 'position': 1, 'color': '#76B7B2'},
        {'canonical': 'O2', 'aliases': [], 'position': 2, 'color': '#4E79A7'},
        {'canonical': 'VF_UK', 'aliases': [], 'position': 3, 'color': '#E15759'},
        {'canonical': 'VF_SA', 'aliases': [], 'position': 4, 'color': '#8000FF'},
    ]
    entry = CatalogEntry(
        slide=3, slide_title='Validation', slide_subtitle='', layout='', chart_title='Failures',
        cdr_source='CDR-Voice', kpi='COUNT(Test_ID)', chart_type='Count Stacked Horizontal Bars',
        legend='Call Status', filters='', rows_split='Session Type × City',
        columns_split='Operator × Campaign', legend_position='Right',
    )

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['column_keys'] == [
        ['3', '26-Q2'], ['EE', '26-Q2'], ['O2', '26-Q2'],
        ['VF_UK', '26-Q2'], ['VF_SA', '26-Q2'],
    ]


def test_multivendor_dynamic_table_places_operator_only_identities_after_vendors() -> None:
    frame = chart_frame({
        'Operator_Vendor': ['EE', 'O2', 'VF_SA', '3_Huawei', '3_Ericsson', 'VF_Ericsson'],
        'Test_Result': ['Completed'] * 6,
        'G_Level_4': ['London'] * 6,
        'Test_ID': ['A', 'B', 'C', 'D', 'E', 'F'],
    })
    frame.attrs['operator_mapping_groups'].append({
        'canonical': 'VF_SA', 'aliases': [], 'position': 4, 'color': '#8000FF',
    })
    frame.attrs['operator_mappings'] = {'vf': 'VF', 'vf_sa': 'VF_SA'}
    entry = CatalogEntry(
        slide=3, slide_title='Validation', slide_subtitle='', layout='', chart_title='Test count',
        cdr_source='CDR-Data', kpi='COUNT(Test_ID)', chart_type='Dynamic Table', legend='', filters='',
        rows_split='Operator_Vendor × Operator × Test_Result', columns_split='G Level 4',
        legend_position='',
    )

    model = catalog_chart_payload(frame, entry, multivendor=True, prefiltered=True)

    assert [row[:2] for row in model['rows']] == [
        ['Ericsson', 'VF'], ['Ericsson', '3'], ['Huawei', '3'],
        ['EE', 'EE'], ['O2', 'O2'], ['VF_SA', 'VF_SA'],
    ]


def test_dynamic_table_rejects_an_unreadable_number_of_pivot_columns() -> None:
    frame = chart_frame({
        'Benchmark': ['UK_Q1_2026'] * 21,
        'G_Level_4': [f'City {index}' for index in range(21)],
        'Test_ID': [f'Test {index}' for index in range(21)],
    })
    entry = CatalogEntry(
        slide=3, slide_title='Validation', slide_subtitle='', layout='', chart_title='Test count',
        cdr_source='CDR-Data', kpi='COUNT(Test_ID)', chart_type='Dynamic Table', legend='', filters='',
        rows_split='Benchmark', columns_split='G Level 4', legend_position='',
    )

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['type'] == 'empty'
    assert model['message'] == (
        'Dynamic Table has 21 distinct column values for G Level 4. '
        'Add a filter to reduce the table to 20 columns or fewer.'
    )
from src.modules.repository import Repository


def with_default_calculated_dimensions(entry: CatalogEntry) -> CatalogEntry:
    payload = json.loads((Path(__file__).parents[1] / 'assets' / 'autocalculated-fields' / 'default-autocalculated-fields.json').read_text(encoding='utf-8'))
    return replace(entry, calculated_dimensions=parse_calculated_dimensions(payload))


def test_vendor_formula_is_shared_by_vodafone_and_three() -> None:
    lookup = {'eri': 'Ericsson', 'nok': 'Nokia', 'hua': 'Huawei'}
    for operator, prefix in (('Vodafone UK', 'Vodafone'), ('3', '3')):
        assert vendor_from_cells(operator, 'eri -> eri', lookup) == f'{prefix}_Ericsson'
        assert vendor_from_cells(operator, 'nok -> nok', lookup) == f'{prefix}_Nokia'
        assert vendor_from_cells(operator, 'eri -> unknown', lookup) == f'{prefix}_Ericsson (Mixed)'
        assert vendor_from_cells(operator, 'nok -> eri', lookup) == f'{prefix}_Ericsson (Mixed)'
        assert vendor_from_cells(operator, 'nok -> hua', lookup) == f'{prefix}_Mixed (non-Ericsson)'
        assert vendor_from_cells(operator, 'unknown -> nok', lookup) == f'{prefix}_Mixed (non-Ericsson)'
        assert vendor_from_cells(operator, '', lookup) == f'{prefix}_Mixed (non-Ericsson)'
    assert vendor_from_cells('O2', 'eri -> unknown', lookup) == 'O2'


def test_report_operator_aliases_use_only_workspace_configuration() -> None:
    frame = chart_frame({
        "Operator": ["Vodafone", "Vodafone UK", "o2 - de", "O2(UK)", "Telefónica", "Three", "Three UK", "3 UK", "EE UK", "Everything Everywhere"],
        "Campaign": ["UK_Q4_2025", "UK_Q2_2026", "UK_Q4_2025", "UK_Q2_2026", "UK_Q2_2026", "UK_Q4_2025", "UK_Q2_2026", "UK_Q4_2025", "UK_Q2_2026", "UK_Q4_2025"],
        "Call_Status": ["Completed"] * 10,
    })
    entry = CatalogEntry(
        1, "", "", "", "", "CDR-Voice", "Call_Status", "100% Stacked Vertical Bars", "",
        "Operator IN (VF, O2, 3, EE)", "Operator", "Campaign",
    )

    normalised = normalise_operator_aliases(frame, {
        'vodafone': 'VF', 'vodafone uk': 'VF', 'o2(uk)': 'O2',
        'o2 - de': 'O2', 'telefónica': 'O2', 'three': '3', 'three uk': '3', '3 uk': '3',
        'ee uk': 'EE', 'everything everywhere': 'EE',
    })
    filtered = _apply_catalog_filters(normalised, entry, False, "Call_Status")
    grouped, primary, series = _apply_catalog_grouping(filtered, entry, False, "Call_Status")

    assert filtered["Operator"].tolist() == ["VF", "VF", "O2", "O2", "O2", "3", "3", "3", "EE", "EE"]
    assert set(grouped[primary]) == {"VF", "O2", "3", "EE"}
    assert set(grouped[series]) == {"25-Q4", "26-Q2"}


def test_catalog_filters_accept_case_separators_and_subscriber_spelling_alias() -> None:
    frame = chart_frame({
        'Suscriber': ['Alpha User', 'Beta User'],
        'Test_Result': ['Completed', 'Failed'],
        'Campaign': ['UK_Q3_2026', 'UK_Q4_2026'],
    })
    entry = CatalogEntry(
        slide=1, slide_title='', slide_subtitle='', layout='', chart_title='',
        cdr_source='CDR-Data', kpi='test result', chart_type='Table', legend='',
        filters='SUBSCRIBER = alpha user; test result = completed; campaign = uk_q3_2026',
        rows_split='', columns_split='',
    )

    filtered = _apply_catalog_filters(frame, entry, False, 'test result')

    assert filtered['Suscriber'].tolist() == ['Alpha User']


def test_catalog_timestamp_conditions_remain_active_when_date_range_filtering_is_disabled(monkeypatch) -> None:
    monkeypatch.setenv('IGNORE_EVENT_TIME_FILTERING', 'true')
    frame = chart_frame({
        'Event_Start_Time': ['2026-09-01 10:00:00', '2026-09-02 10:00:00'],
        'Metric': [1, 2],
    })
    entry = CatalogEntry(
        slide=1, slide_title='', slide_subtitle='', layout='', chart_title='',
        cdr_source='CDR-Data', kpi='Metric', chart_type='Table', legend='',
        filters='Event_Start_Time = 2026-09-01 10:00:00', rows_split='', columns_split='',
    )

    filtered = _apply_catalog_filters(frame, entry, False, 'Metric')

    assert filtered['Metric'].tolist() == [1]


def test_h3g_is_not_normalised_as_operator_three() -> None:
    frame = chart_frame({"Operator": ["H3G", "H3G UK", "Three UK"]})

    assert normalise_operator_aliases(frame)["Operator"].tolist() == ["H3G", "H3G UK", "Three UK"]


def test_vendor_aliases_normalise_only_with_workspace_configuration() -> None:
    frame = chart_frame({"vendor": ["Vodafone_Ericsson", "Three UK_Nokia", "O2 (UK)_Huawei", "EE_Ericsson", "H3G_Huawei"]})

    assert normalise_operator_aliases(frame)["vendor"].tolist() == frame["vendor"].tolist()
    assert normalise_operator_aliases(frame, {
        'vodafone': 'VF', 'three uk': '3', 'o2 (uk)': 'O2', 'ee': 'EE',
    })["vendor"].tolist() == [
        "VF_Ericsson", "3_Nokia", "O2_Huawei", "EE_Ericsson", "H3G_Huawei",
    ]


def test_combined_vendor_identity_preserves_an_operator_with_an_underscore() -> None:
    frame = chart_frame({'Vendor': ['VF_SA_Ericsson', 'VF_NSA_Huawei']})
    frame.attrs['operator_mappings'] = {'vf_sa': 'VF_SA', 'vf_nsa': 'VF_NSA'}
    frame.attrs['vendor_mappings'] = {'ericsson': 'Ericsson', 'huawei': 'Huawei'}

    normalized = normalise_operator_aliases(frame)

    assert normalized['Vendor'].tolist() == ['VF_SA_Ericsson', 'VF_NSA_Huawei']


def test_vendor_legend_colours_match_the_vendor_bar_colours() -> None:
    entry = CatalogEntry(
        1, "", "", "", "", "CDR-Speech", "LQ", "Average Vertical Bars",
        "Vendor", "", "Vendor", "", "Right",
    )
    frame = chart_frame({
        "__catalog_row_0": ["VF_Ericsson", "VF_Samsung", "VF_Huawei", "3_Huawei"],
        "LQ": [4.64, 4.65, 4.64, 4.64],
    })
    frame.attrs["catalogue_dimension_labels"] = {"__catalog_row_0": "Vendor"}
    keys = [(value,) for value in frame["__catalog_row_0"]]

    legend = _resolved_legend_items(entry, frame, "LQ")

    assert {caption: colour for caption, colour, _width in legend} == {
        key[0]: colour for key, colour in _series_colours(keys, ["__catalog_row_0"], frame).items()
    }


def test_reporting_cache_resolves_separator_variants_and_refreshes_derived_dimensions(tmp_path) -> None:
    repository = Repository(tmp_path / 'workspace.db')
    repository.replace_dataset_rows(1, pd.DataFrame({
        'RAT_A': ['EN-DC'], 'G_Level_4': ['London'], 'Call Family': [None],
    }))
    repository.copy_dataset_rows_to_reporting(1, 'voice', ['RAT_A', 'G Level 4', 'Call Family'])

    # Simulate a CDR that received its derived field after its first reporting
    # copy. The cache must not retain the original null value indefinitely.
    repository.replace_dataset_rows(1, pd.DataFrame({
        'RAT_A': ['EN-DC'], 'G_Level_4': ['London'], 'Call Family': ['VoLTE'],
    }))
    repository.copy_dataset_rows_to_reporting(1, 'voice', ['RAT_A', 'G Level 4', 'Call Family'])

    loaded = repository.load_reporting_rows('voice', [1], ['G Level 4', 'Call Family'])
    assert loaded.to_dict(orient='records') == [{'G Level 4': 'London', 'Call Family': 'VoLTE'}]


def test_dataset_storage_coalesces_case_only_columns_without_sqlite_suffixes(tmp_path) -> None:
    repository = Repository(tmp_path / 'workspace.db')
    repository.replace_dataset_rows(1, pd.DataFrame({
        'Campaign': ['UK_Q3_2026'], 'campaign': ['normalized-copy'],
        'Vendor': ['source-vendor'], 'vendor': ['VF_Ericsson'],
    }))

    assert repository.list_dataset_row_columns(1) == ['Campaign', 'Vendor']
    stored = repository.load_dataset_rows(1, ['Campaign', 'Vendor'], {})
    # Vendor (the vendor alone) wins over a lower-case helper sharing its SQLite name.
    assert stored.to_dict(orient='records') == [{'Campaign': 'UK_Q3_2026', 'Vendor': 'source-vendor'}]


def test_vendor_only_removes_equivalent_operator_alias_prefixes() -> None:
    from src.modules.column_names import vendor_only_value

    assert vendor_only_value('Vodafone_Ericsson', 'Vodafone UK') == 'Ericsson'
    assert vendor_only_value('VF_Samsung', 'Vodafone UK') == 'Samsung'
    assert vendor_only_value('3_Nokia', 'Three UK') == 'Nokia'
    assert vendor_only_value('Huawei', 'Vodafone UK') == 'Huawei'


@pytest.mark.parametrize('chart_type', ['CDF Line', 'Threshold Stacked Vertical Bars', 'Histogram Bars'])
def test_vendor_only_grouping_pools_two_operators_for_radio_quality_charts(chart_type) -> None:
    entry = CatalogEntry(
        slide=1, slide_title='', slide_subtitle='',
        layout='2 rows + dynamic columns, comments down', chart_title='RSRP',
        cdr_source='CDR-All', kpi='LTE_RSRP', chart_type=chart_type,
        legend='Operator', filters='', rows_split='Operator',
        columns_split='Campaign', legend_position='Right', dynamic_field='Operator',
    )
    prepared = prepare_multivendor_catalog_entry(entry, 'vendor_only')
    frame = chart_frame({
        'Operator': ['EE', 'EE', 'O2', 'O2', 'VF'],
        'Operator_Vendor': ['EE_Ericsson', 'EE_Ericsson', 'O2_Ericsson', 'O2_Ericsson', 'VF_Huawei'],
        'Vendor': ['Ericsson', 'Ericsson', 'Ericsson', 'Ericsson', 'Huawei'],
        'Campaign': ['2026-Q1'] * 5,
        'LTE_RSRP': [-110, -100, -90, -80, -70],
    })

    grouped, primary, _series = _apply_catalog_grouping(frame, prepared, True, 'LTE_RSRP')

    assert prepared.dynamic_field == 'Vendor'
    assert grouped[primary].drop_duplicates().tolist() == ['Ericsson', 'Huawei']


def test_reporting_cache_repairs_an_empty_requested_source_column(tmp_path) -> None:
    repository = Repository(tmp_path / 'workspace.db')
    repository.replace_dataset_rows(1, pd.DataFrame({'RAT': ['EN-DC'], 'G_Level_4': ['London']}))
    repository.copy_dataset_rows_to_reporting(1, 'data', ['RAT', 'G Level 4'])
    with repository.connection() as connection:
        connection.execute('UPDATE reporting_rows_data SET G_Level_4 = NULL WHERE dataset_id = 1')

    repository.copy_dataset_rows_to_reporting(1, 'data', ['RAT', 'G Level 4'])

    loaded = repository.load_reporting_rows('data', [1], ['G Level 4'])
    assert loaded.to_dict(orient='records') == [{'G Level 4': 'London'}]


def test_reporting_cache_checks_requested_columns_in_bounded_table_scans(tmp_path) -> None:
    class TracedRepository(Repository):
        def __init__(self, db_path: Path) -> None:
            super().__init__(db_path)
            self.statements: list[str] = []

        @contextmanager
        def connection(self):
            with super().connection() as connection:
                connection.set_trace_callback(self.statements.append)
                yield connection

    repository = TracedRepository(tmp_path / 'workspace.db')
    repository.replace_dataset_rows(1, pd.DataFrame({
        'first': ['A', 'B'], 'second': [None, None], 'third': ['C', None],
    }))
    repository.copy_dataset_rows_to_reporting(1, 'data', ['first', 'second', 'third'])
    repository.statements.clear()

    assert repository.copy_dataset_rows_to_reporting(1, 'data', ['first', 'second', 'third']) is False

    presence_queries = [
        statement for statement in repository.statements
        if statement.startswith('SELECT EXISTS(')
    ]
    assert len(presence_queries) == 2
    assert '"first" IS NOT NULL' in presence_queries[0]
    assert '"second" IS NOT NULL' in presence_queries[0]
    assert '"third" IS NOT NULL' in presence_queries[0]
    assert 'FROM "dataset_rows_1"' in presence_queries[1]


def test_multivendor_enrichment_by_operator() -> None:
    cdr = pd.DataFrame({
        'Operator': ['Vodafone UK', '3'],
        'RAT_A': ['LTE EN-DC', 'NR'],
        'Cell_ID_A': ['100 -> 100', '200 -> 201'],
    })
    vodafone_mapping = pd.DataFrame({
        'source_sheet': ['4G'],
        'eNodeB ID': [0],
        'Local Cell ID': [100],
        'OP/ Vendor': ['Ericsson'],
    })
    three_mapping = pd.DataFrame({
        'Cid__ECI': [200, 201],
        'Vendor': ['Nokia', 'Ericsson'],
    })

    # The rows of each NR Mode are those of its CDRs (no call-mode classification).
    nsa, sa = cdr.iloc[[0]], cdr.iloc[[1]]

    assert enrich_multivendor(nsa, vodafone_mapping, three_mapping)['Operator_Vendor'].tolist() == ['Vodafone_Ericsson']
    assert enrich_multivendor(sa, vodafone_mapping, three_mapping)['Operator_Vendor'].tolist() == ['3_Ericsson (Mixed)']


def test_combined_reporting_frame_keeps_every_data_attempt_before_template_filters(tmp_path) -> None:
    import src.DriveTestAnalyzer as app_module

    repository = Repository(tmp_path / 'workspace.db')
    dataset_id = 1
    rows = pd.DataFrame({
        'source_sheet': ['EE', 'EE', 'EE', 'TMP_CLIPBOARD'],
        'Campaign': ['UK_Q2_2026'] * 4,
        'Operator': ['EE'] * 4,
        'Test_Name': ['FDFS HTTPS UL ST'] * 4,
        'Test_Result': ['Cutoff'] * 4,
        'G_Level_4': ['London'] * 4,
        'RAT': ['EN-DC', 'LTE', 'NR SA', 'EN-DC'],
    })
    repository.replace_dataset_rows(dataset_id, rows)
    entry = CatalogEntry(
        slide=5, slide_title='Data failures', slide_subtitle='', layout='', chart_title='',
        cdr_source='CDR-Data', kpi='Test_Result', chart_type='100% Stacked Vertical Bars',
        legend='Test_Result', filters=(
            'Test Name CONTAINS FDFS; Test_Result IN (Completed, Cutoff, Failed); '
            'Operator IN (Vodafone UK, 3, EE); G Level 4 IN (London)'
        ), rows_split='Test_Name',
        columns_split='Operator × Campaign', legend_position='Right',
    )

    frame = app_module._combined_reporting_frame(
        [{'id': dataset_id, 'dataset_kind': 'data'}], [entry], False, repository,
    )

    assert frame['RAT'].tolist() == ['EN-DC', 'LTE', 'NR SA']
    assert frame['source_sheet'].tolist() == ['EE', 'EE', 'EE']
    preview, summary = app_module.preview_catalog_chart_data(frame, entry, limit=100)
    assert summary['matched_rows'] == 3
    assert preview['Filter · Test_Result'].tolist() == ['Cutoff', 'Cutoff', 'Cutoff']
    assert preview['Resolved Columns Split'].tolist() == ['EE · 26-Q2'] * 3


def test_workspace_vendor_assignment_writes_the_normalized_vendor_field() -> None:
    cdr = pd.DataFrame({
        'Operator': ['Vodafone UK', '3', 'O2 (UK)'],
        'Cell_ID_A': ['100 -> 100', '200 -> 201', '300'],
    })
    vodafone_mapping = pd.DataFrame({
        'source_sheet': ['4G'], 'eNodeB ID': [0], 'Local Cell ID': [100], 'OP/ Vendor': ['Ericsson'],
    })
    three_mapping = pd.DataFrame({'Cid__ECI': [200, 201], 'Vendor': ['Nokia', 'Ericsson']})

    mapped = assign_cdr_vendors(cdr, vodafone_mapping, three_mapping)

    assert mapped['Operator_Vendor'].tolist() == ['Vodafone_Ericsson', '3_Ericsson (Mixed)', 'O2 (UK) - All']
    assert mapped['Vendor'].tolist() == ['Ericsson', 'Ericsson (Mixed)', 'O2 (UK) - All']
    assert mapped.columns[:3].tolist() == ['Operator_Vendor', 'Vendor', 'Operator']


def test_vendor_mapping_files_apply_to_every_spelling_of_their_operator() -> None:
    # Tests configured in SA or VoNR mode name the same networks differently; they share their cells.
    cdr = pd.DataFrame({
        'Operator': ['Vodafone VoNR', 'VF_SA', 'Vodafone SA', 'Three SA', 'O2 (UK)'],
        'Cell_ID_A': ['100 -> 100', '100', '100 -> 999', '200 -> 200', '300'],
    })
    vodafone_mapping = pd.DataFrame({
        'source_sheet': ['4G'], 'eNodeB ID': [0], 'Local Cell ID': [100], 'OP/ Vendor': ['Ericsson'],
    })
    three_mapping = pd.DataFrame({'Cid__ECI': [200], 'Vendor': ['Nokia']})

    mapped = assign_cdr_vendors(cdr, vodafone_mapping, three_mapping)

    assert mapped['Operator_Vendor'].tolist() == [
        'Vodafone VoNR_Ericsson', 'VF_SA_Ericsson', 'Vodafone SA_Ericsson (Mixed)', 'Three SA_Nokia', 'O2 (UK) - All',
    ]
    assert mapped['Vendor'].tolist() == ['Ericsson', 'Ericsson', 'Ericsson (Mixed)', 'Nokia', 'O2 (UK) - All']


def test_workspace_vendor_assignment_replaces_source_vendor_collisions() -> None:
    cdr = pd.DataFrame({
        'source_sheet': ['Vodafone'],
        'Vendor': ['legacy source value'],
        'vendor__2': ['normalised source value'],
        'Operator': ['3'],
        'Cell_ID_A': ['200 -> 200'],
    })
    three_mapping = pd.DataFrame({'Cid__ECI': [200], 'Vendor': ['Nokia']})

    mapped = assign_cdr_vendors(cdr, None, three_mapping)

    assert mapped.columns[:3].tolist() == ['source_sheet', 'Operator_Vendor', 'Vendor']
    assert 'vendor__2' not in mapped.columns
    assert mapped.loc[0, 'Operator_Vendor'] == '3_Nokia'
    assert mapped.loc[0, 'Vendor'] == 'Nokia'


def test_workspace_vendor_assignment_supports_a_single_selected_mapping() -> None:
    cdr = pd.DataFrame({
        'Operator': ['Vodafone UK', '3'],
        'Cell_ID_A': ['100 -> 100', '200 -> 200'],
    })
    vodafone_mapping = pd.DataFrame({
        'source_sheet': ['4G'], 'eNodeB ID': [0], 'Local Cell ID': [100], 'OP/ Vendor': ['Ericsson'],
    })

    mapped = assign_cdr_vendors(cdr, vodafone_mapping, None)

    assert mapped.loc[0, 'Operator_Vendor'] == 'Vodafone_Ericsson'
    assert mapped['Operator_Vendor'].tolist() == ['Vodafone_Ericsson', '3 - All']


def test_workspace_vendor_assignment_accepts_equivalent_global_cell_id_columns() -> None:
    cdr = pd.DataFrame({
        'Operator': ['3'],
        'Global CI': ['200 -> 200'],
    })
    three_mapping = pd.DataFrame({'Cid__ECI': [200], 'Vendor': ['Nokia']})

    mapped = assign_cdr_vendors(cdr, None, three_mapping)

    assert mapped['Operator_Vendor'].tolist() == ['3_Nokia']


def test_catalogue_converter_migrates_legacy_headers_and_grouping() -> None:
    legacy = (
        'Slide,Slide title,Slide subtitle,Layout,CDR Source,KPI,Chart Type,Filters,Grouping\n'
        '8,Quality,Voice,Title and 1 column + Comments,CDR-Voice,Call_Status,100% Stacked Vertical Bars,,Operator × Campaign\n'
    )

    converted = convert_catalog_csv(legacy)
    entries = parse_catalog_csv(converted)

    assert converted.decode('utf-8').splitlines()[0] == ','.join(CATALOG_HEADERS)
    assert entries[0].slide_title == 'Quality'
    assert entries[0].chart_title == ''
    assert entries[0].legend == ''
    assert entries[0].rows_split == 'Operator'
    assert entries[0].columns_split == 'Campaign'


def test_catalogue_converter_assigns_layouts_for_missing_legacy_layouts() -> None:
    legacy = (
        'Slide,Slide title,Slide subtitle,Layout,CDR Source,KPI,Chart Type,Filters,Grouping\n'
        '9,Quality,,,CDR-Voice,Call_Status,100% Stacked Vertical Bars,,Operator × Campaign\n'
        '9,Quality,,,CDR-Voice,Call_Setup_Time,Average Vertical Bars,,Operator × Campaign\n'
    )

    entries = parse_catalog_csv(convert_catalog_csv(legacy))

    assert {entry.layout for entry in entries} == {'Title + 1 rows + 2 columns + comments down'}


def test_vendor_group_fills_unmapped_operators_in_the_official_vendor_field() -> None:
    frame = chart_frame({'Operator': ['Vodafone UK', 'O2 (UK)'], 'Operator_Vendor': ['Vodafone_Ericsson', pd.NA]})

    grouped = ensure_vendor_group(frame)

    assert grouped['Operator_Vendor'].tolist() == ['Vodafone_Ericsson', 'O2 (UK) - All']
    assert grouped['Vendor'].tolist() == ['Ericsson', 'O2 (UK) - All']


def test_vodafone_mapping_derives_gcid_from_4g_enodeb_and_local_cell() -> None:
    cdr = pd.DataFrame({'Operator': ['Vodafone UK'], 'Cell_IDs_A': ['3330049']})
    vodafone_mapping = pd.DataFrame({
        'source_sheet': ['4G'],
        'eNodeB ID': [13008],
        'Local Cell ID': [1],
        'OP/ Vendor': ['Samsung'],
    })
    three_mapping = pd.DataFrame({'Cid__ECI': [1], 'Vendor': ['Nokia']})

    assert enrich_multivendor(cdr, vodafone_mapping, three_mapping)['Operator_Vendor'].tolist() == ['Vodafone_Samsung']


def test_vodafone_mapping_keys_huawei_cells_by_their_cell_id_and_nr_cells_by_nci() -> None:
    # The CDRs report eNodeB × 256 + Cell ID; Huawei's Local Cell ID (7) is internal.
    cdr = pd.DataFrame({'Operator': ['Vodafone UK', 'Vodafone SA'],
                        'Cell_IDs_A': [str(1491 * 256 + 110), str(16628 * 16384 + 515)]})
    vodafone_mapping = pd.DataFrame({
        'source_sheet': ['4G', '5G'],
        'eNodeB ID': [1491, None],
        'gNodeB ID': [None, 16628],
        'Cell ID': [110, 515],
        'Local Cell ID': [7, 10],
        'OP/ Vendor': ['Huawei', 'Ericsson'],
    })
    three_mapping = pd.DataFrame({'Cid__ECI': [1], 'Vendor': ['Nokia']})

    assert enrich_multivendor(cdr, vodafone_mapping, three_mapping)['Operator_Vendor'].tolist() == ['Vodafone_Huawei', 'Vodafone SA_Ericsson']


def test_catalogue_csv_requires_the_report_chart_contract_columns() -> None:
    catalogue = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Completed Call Ratio,Voice quality,Title and 1 column + Comments,Completed call ratio,CDR-Voice,Call_Status,100% Stacked Vertical Bars,Call Family = VoLTE,Operator,Campaign,Completed/Dropped/Failed,\n'
    ).encode('utf-8')

    entries = parse_catalog_csv(catalogue)

    assert entries[0].slide == 8
    assert entries[0].slide_title == 'Completed Call Ratio'
    assert entries[0].slide_subtitle == 'Voice quality'
    assert entries[0].chart_title == 'Completed call ratio'
    assert entries[0].legend == 'Completed/Dropped/Failed'
    assert entries[0].source_kind == 'voice'
    assert entries[0].chart_type == '100% Stacked Vertical Bars'


def test_catalogue_parses_legend_position_and_accepts_prior_schema() -> None:
    current = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Quality,,Title and 1 column + Comments,,CDR-Voice,Call_Status,100% Stacked Vertical Bars,,Operator,Campaign,Completed/Dropped/Failed,Left\n'
    )
    entry = parse_catalog_csv(current)[0]

    assert entry.legend_position == 'left'
    assert parse_legend_position('Bottom') == 'bottom'
    with pytest.raises(ValueError, match='Legend Position'):
        parse_legend_position('Centre')

    previous = 'Slide,Slide tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Legend,Filters,Grouping_Rows,Grouping_Columns' + '\n8,Quality,,Title and 1 column + Comments,,CDR-Voice,Call_Status,100% Stacked Vertical Bars,,,Operator,Campaign\n'
    assert parse_catalog_csv(previous)[0].legend_position == ''
    two_columns = ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n8,Quality,,Title and 2 columns + Comments,,CDR-Voice,Call_Status,100% Stacked Vertical Bars,,Operator,Campaign,,\n'
    assert parse_catalog_csv(two_columns)[0].legend_position == ''


def test_catalogue_cdf_axis_ranges_are_optional_and_backward_compatible() -> None:
    rangeless = (
        'Slide,Slide Tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position\n'
        '8,Quality,,Title and 1 column + Comments,Quality,CDR-Speech,LQ,CDF Line,,Operator,,,Top\n'
    )
    legacy_entry = parse_catalog_csv(rangeless)[0]
    assert legacy_entry.axis_x_range == ''
    assert legacy_entry.axis_y_range == ''

    current = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Quality,,Title and 1 column + Comments,Quality,CDR-Speech,LQ,CDF Line,,Operator,,,Top,,,,"[0.01,]","[75,100]"\n'
    )
    entry = parse_catalog_csv(current)[0]
    assert parse_axis_range(entry.axis_x_range, 'x') == (0.01, None)
    assert parse_axis_range(entry.axis_y_range, 'y') == (75.0, 100.0)
    assert b'Legend Position,Legend Format,Label Position,Label Format,Axis X Range,Axis Y Range' in catalogue_csv([entry])


def test_catalogue_bar_label_position_is_validated_and_serialised() -> None:
    assert parse_label_position('Middle') == 'middle'
    with pytest.raises(ValueError, match='Label'):
        parse_label_position('Centre')
    content = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Failures,,Title and 1 column + Comments,Failures,CDR-Voice,Call_Status,Count Stacked Horizontal Bars,,Operator,,,Top,,Down,,,\n'
    )
    entry = parse_catalog_csv(content)[0]
    assert entry.label_position == 'down'
    assert b',Down,,' in catalogue_csv([entry])


def test_catalogue_label_format_is_validated_and_serialised() -> None:
    assert parse_label_format('["#1a2b3c", "Verdana", "Large", "Bold"]') == '["#1A2B3C","Verdana","Large","Bold"]'
    with pytest.raises(ValueError, match='Label Format'):
        parse_label_format('["blue"]')
    row = {
        'Slide': '8', 'Slide Tittle': 'Failures', 'Slide Subtittle': '', 'Layout': 'Title and 1 column + Comments',
        'Chart Tittle': 'Failures', 'Source Dataset': 'CDR-Voice', 'KPI': 'Call_Status',
        'Chart type': 'Count Stacked Horizontal Bars', 'Filters': '', 'Rows Split': 'Operator',
        'Columns Split': '', 'Legend': '', 'Legend Position': 'Top', 'Label Position': 'Down',
        'Label Format': '["#1a2b3c", "Verdana", "Large", "Bold"]', 'Axis X Range': '', 'Axis Y Range': '',
        'Exclude Null/Empty': '', 'Exclude Zero': '',
    }
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=CATALOG_HEADERS)
    writer.writeheader()
    writer.writerow(row)
    entry = parse_catalog_csv(buffer.getvalue())[0]
    assert entry.label_format == '["#1A2B3C","Verdana","Large","Bold"]'
    assert b'#1A2B3C' in catalogue_csv([entry])


def test_catalogue_legend_format_is_validated_and_serialised() -> None:
    row = {
        'Slide': '8', 'Slide Tittle': 'Failures', 'Slide Subtittle': '', 'Layout': 'Title and 1 column + Comments',
        'Chart Tittle': 'Failures', 'Source Dataset': 'CDR-Voice', 'KPI': 'Call_Status',
        'Chart type': 'Count Stacked Horizontal Bars', 'Filters': '', 'Rows Split': 'Operator',
        'Columns Split': '', 'Legend': '', 'Legend Position': 'Top',
        'Legend Format': '["#1a2b3c", "Verdana", "Large", "Bold"]', 'Label Position': '', 'Label Format': '',
        'Axis X Range': '', 'Axis Y Range': '', 'Exclude Null/Empty': '', 'Exclude Zero': '',
    }
    buffer = StringIO()
    writer = csv.DictWriter(buffer, fieldnames=CATALOG_HEADERS)
    writer.writeheader()
    writer.writerow(row)
    entry = parse_catalog_csv(buffer.getvalue())[0]
    assert entry.legend_format == '["#1A2B3C","Verdana","Large","Bold"]'
    assert b'Legend Format' in catalogue_csv([entry])


def test_catalogue_null_and_zero_exclusions_are_independent_and_backward_compatible() -> None:
    assert parse_template_boolean('', 'Exclude Zero') is False
    assert parse_template_boolean('Yes', 'Exclude Zero') is True
    with pytest.raises(ValueError, match='Exclude Zero'):
        parse_template_boolean('Sometimes', 'Exclude Zero')

    previous_visual_schema = (
        'Slide,Slide Tittle,Slide Subtittle,Layout,Chart Tittle,CDR source,KPI,Chart type,Filters,Rows Split,Columns Split,Legend,Legend Position,Label,Axis X Range,Axis Y Range\n'
        '8,Quality,,Layout,Quality,CDR-Data,Metric,CDF Line,,Operator,,,Top,,,\n'
    )
    previous = parse_catalog_csv(previous_visual_schema)[0]
    assert previous.exclude_null_empty is False
    assert previous.exclude_zero is False

    content = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Quality,,Layout,Quality,CDR-Data,Metric,CDF Line,,Operator,,,Top,,,,,,Yes,Yes\n'
    )
    entry = parse_catalog_csv(content)[0]
    assert entry.exclude_null_empty is True
    assert entry.exclude_zero is True
    assert b'Exclude Null/Empty,Exclude Zero' in catalogue_csv([entry])


def test_null_and_zero_exclusions_filter_plotted_values_independently() -> None:
    frame = chart_frame({'Operator': ['A'] * 4, 'Metric': [None, 0.0, 1.0, 2.0]})
    base = dict(
        slide=1, slide_title='Chart', slide_subtitle='', layout='Layout', chart_title='Chart',
        cdr_source='CDR-Data', kpi='Metric', chart_type='CDF Line', legend='', filters='',
        rows_split='Operator', columns_split='',
    )

    without_nulls, _ = prepare_catalog_chart_preview_frame(
        frame, CatalogEntry(exclude_null_empty=True, **base),
    )
    without_zeroes, _ = prepare_catalog_chart_preview_frame(
        frame, CatalogEntry(exclude_zero=True, **base),
    )
    without_either, _ = prepare_catalog_chart_preview_frame(
        frame, CatalogEntry(exclude_null_empty=True, exclude_zero=True, **base),
    )

    assert without_nulls['Metric'].tolist() == [0.0, 1.0, 2.0]
    assert without_zeroes['Metric'].isna().sum() == 1
    assert without_zeroes['Metric'].dropna().tolist() == [1.0, 2.0]
    assert without_either['Metric'].tolist() == [1.0, 2.0]

    payload = catalog_chart_payload(without_either, CatalogEntry(exclude_null_empty=True, exclude_zero=True, **base), prefiltered=True)
    assert payload['series'][0]['samples'] == 2
    assert payload['domain']['x'][0] == 1.0

    scatter_frame = chart_frame({
        'Operator': ['A'] * 5,
        'Y': [10.0, 0.0, 20.0, None, 30.0],
        'X': [1.0, 2.0, 0.0, 3.0, None],
    })
    scatter = catalog_chart_payload(
        scatter_frame,
        CatalogEntry(
            slide=1, slide_title='Scatter', slide_subtitle='', layout='Layout', chart_title='Scatter',
            cdr_source='CDR-Data', kpi='Y vs X', chart_type='Scatter', legend='', filters='',
            rows_split='Operator', columns_split='', exclude_null_empty=True, exclude_zero=True,
        ),
    )
    assert scatter['series'][0]['points'] == [[1.0, 10.0]]


def test_catalogue_accepts_axis_ranges_and_labels_for_every_chart_type() -> None:
    non_cdf_range = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Quality,,Title and 1 column + Comments,Quality,CDR-Data,Metric,Average Vertical Bars,,Operator,,,Top,,,,"[0.01,]",\n'
    )
    bar_entry = parse_catalog_csv(non_cdf_range)[0]
    assert bar_entry.axis_x_range == '[0.01,]'

    non_bar_label = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Quality,,Title and 1 column + Comments,Quality,CDR-Data,Metric,CDF Line,,Operator,,,Top,,Down,,,\n'
    )
    cdf_entry = parse_catalog_csv(non_bar_label)[0]
    assert cdf_entry.label_position == 'down'


def test_chart_payload_applies_cdf_ranges_and_bar_label_override() -> None:
    frame = chart_frame({'Operator': ['A'] * 4, 'Metric': [0.0, 1.0, 2.0, 3.0]})
    base = dict(
        slide=1, slide_title='Chart', slide_subtitle='', layout='Layout', chart_title='Chart',
        cdr_source='CDR-Data', kpi='Metric', legend='', filters='',
        rows_split='Operator', columns_split='',
    )
    cdf = catalog_chart_payload(
        frame,
        CatalogEntry(chart_type='CDF Line', axis_x_range='[0.01,]', axis_y_range='[50,100]', **base),
        prefiltered=True,
    )
    bars = catalog_chart_payload(
        frame,
        CatalogEntry(
            chart_type='Average Vertical Bars', axis_y_range='[1,5]',
            label_position='down', **base,
        ),
        prefiltered=True,
    )

    assert cdf['domain'] == {'x': [0.01, 3.0], 'y': [0.5, 1.0]}
    assert cdf['series'][0]['x'][0] == 0.01
    assert bars['label_position'] == 'down'
    assert bars['domain']['y'] == [1.0, 5.0]
    assert bars['axis_ranges']['y'] == [1.0, 5.0]

    scatter_frame = chart_frame({
        'Operator': ['A', 'B'], 'Y': [10.0, 20.0], 'X': [1.0, 2.0],
    })
    scatter = catalog_chart_payload(
        scatter_frame,
        CatalogEntry(
            chart_type='Scatter', kpi='Y vs X', axis_x_range='[0,3]',
            axis_y_range='[5,25]', label_position='top',
            **{key: value for key, value in base.items() if key != 'kpi'},
        ),
    )
    assert scatter['domain'] == {'x': [0.0, 3.0], 'y': [5.0, 25.0]}
    assert scatter['label_position'] == 'top'


def test_cdf_visible_points_preserve_the_vertical_step_at_the_automatic_minimum() -> None:
    values = [0.0, 0.0, 0.0, 1.0]

    assert _cdf_visible_points(values, 0.0, 1.0) == [
        (0.0, 0.25), (0.0, 0.5), (0.0, 0.75), (1.0, 1.0),
    ]
    assert _cdf_visible_points(values, 0.01, 1.0) == [
        (0.01, 0.75), (1.0, 1.0),
    ]

    frame = chart_frame({'Operator': ['A'] * 4, 'Metric': values})
    payload = catalog_chart_payload(
        frame,
        CatalogEntry(
            slide=1, slide_title='CDF', slide_subtitle='', layout='Layout', chart_title='CDF',
            cdr_source='CDR-Data', kpi='Metric', chart_type='CDF Line', legend='', filters='',
            rows_split='Operator', columns_split='',
        ),
        prefiltered=True,
    )
    assert payload['domain']['x'] == [0.0, 1.0]
    assert payload['series'][0]['x'] == [0.0, 0.0, 0.0, 1.0]
    assert payload['series'][0]['y'] == [0.25, 0.5, 0.75, 1.0]


@pytest.mark.parametrize('value', ['0,1', '[,]', '[2,1]', '[nan,3]'])
def test_catalogue_rejects_invalid_axis_ranges(value: str) -> None:
    with pytest.raises(ValueError, match='Axis X Range'):
        parse_axis_range(value, 'x')


def test_legend_parser_supports_manual_captions_and_dimension_selection() -> None:
    assert _legend_labels('Completed/Dropped/Failed') == ('Completed', 'Dropped', 'Failed')
    assert _legend_dimensions('Completed/Dropped/Failed') == ()
    assert _legend_dimensions('Operator, Campaign') == ('Operator', 'Campaign')
    assert _legend_labels('Operator, Campaign') == ()


def test_resolved_legend_uses_selected_chart_field_values_and_empty_disables_it() -> None:
    frame = chart_frame({'Test_Name': ['FDFS', 'FDTT', 'FDFS'], 'Test_Result': ['Completed', 'Failed', 'Completed']})
    base = dict(
        slide=5, slide_title='', slide_subtitle='', layout='', chart_title='', cdr_source='CDR-Data',
        kpi='Test_Result', chart_type='100% Stacked Vertical Bars', filters='',
        rows_split='Test_Name', columns_split='', legend_position='top',
    )
    assert _resolved_legend_items(CatalogEntry(legend='', **base), frame, 'Test_Result') == []
    assert [item[0] for item in _resolved_legend_items(CatalogEntry(legend='Test_Name', **base), frame, 'Test_Result')] == ['FDFS', 'FDTT']


def test_resolved_filter_legend_is_text_only() -> None:
    entry = CatalogEntry(
        slide=5, slide_title='', slide_subtitle='', layout='', chart_title='', cdr_source='CDR-Data',
        kpi='Test_Result', chart_type='100% Stacked Vertical Bars', legend='Operator',
        filters='Operator IN (Vodafone UK, EE);', rows_split='Test_Name', columns_split='',
        legend_position='top',
    )
    items = _resolved_legend_items(entry, pd.DataFrame({'Test_Name': ['FDFS']}), 'Test_Result')
    assert items == [('Operator IN (Vodafone UK, EE)', '', 0)]


def test_cdf_resolved_legend_reproduces_historical_and_latest_line_widths() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['EE', 'EE'],
        '__catalog_column_0': ['2026-Q1', '2026-Q2'],
        'Campaign': ['2026-Q1', '2026-Q2'],
        'KPI': [1.0, 2.0],
    })
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': ('Operator',),
        '__catalog_column_0': ('Campaign',),
    }
    entry = CatalogEntry(
        slide=1, slide_title='', slide_subtitle='', layout='', chart_title='', cdr_source='CDR-Data',
        kpi='KPI', chart_type='CDF Lines', legend='Operator, Campaign', filters='',
        rows_split='Operator', columns_split='Campaign', legend_position='top',
    )
    items = _resolved_legend_items(entry, frame, 'KPI')
    assert [(caption, width) for caption, _colour_value, width in items] == [
        ('EE · 2026-Q1', 1), ('EE · 2026-Q2', 4),
    ]


def test_cdf_resolved_legend_decreases_four_campaign_widths_per_operator() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['EE'] * 4 + ['O2'] * 3,
        '__catalog_column_0': [
            '2026-Q1', '2026-Q2', '2026-Q3', '2026-Q4',
            '2025-Q4', '2026-Q1', '2026-Q2',
        ],
        'Campaign': [
            '2026-Q1', '2026-Q2', '2026-Q3', '2026-Q4',
            '2025-Q4', '2026-Q1', '2026-Q2',
        ],
        'KPI': [1.0, 2.0, 3.0, 4.0, 1.5, 2.5, 3.5],
    })
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': ('Operator',),
        '__catalog_column_0': ('Campaign',),
    }
    entry = CatalogEntry(
        slide=1, slide_title='', slide_subtitle='', layout='', chart_title='', cdr_source='CDR-Data',
        kpi='KPI', chart_type='CDF Lines', legend='Operator, Campaign', filters='',
        rows_split='Operator', columns_split='Campaign', legend_position='top',
    )

    items = _resolved_legend_items(entry, frame, 'KPI')

    assert {caption: width for caption, _colour_value, width in items} == {
        'EE · 2026-Q1': 1,
        'EE · 2026-Q2': 2,
        'EE · 2026-Q3': 3,
        'EE · 2026-Q4': 4,
        'O2 · 2025-Q4': 2,
        'O2 · 2026-Q1': 3,
        'O2 · 2026-Q2': 4,
    }


def test_interactive_cdf_model_uses_reporting_hierarchy_palette_and_legend() -> None:
    entry = CatalogEntry(
        1, 'Speech', '', '', 'POLQA CDF', 'CDR-Speech', 'LQ', 'CDF Line',
        'Vendor, Campaign', '', 'Vendor', 'Campaign', 'Right',
    )
    frame = chart_frame({
        'Vendor': ['VF_Ericsson'] * 4 + ['VF_Huawei'] * 4,
        'Campaign': ['UK_Q2_2026', 'UK_Q2_2026', 'UK_Q1_2026', 'UK_Q1_2026'] * 2,
        'LQ': [1.2, 4.2, 1.0, 4.0, 1.4, 4.4, 1.1, 4.1],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert (model['renderer'], model['type'], model['width'], model['height'], model['title']) == (
        'catalog-v2', 'cdf', 1600, 900, 'POLQA CDF',
    )
    assert [(series['key'], series['colour'], series['width'], series['dash']) for series in model['series']] == [
        (['VF_Ericsson', '26-Q1'], '#2E8B57', 1, []),
        (['VF_Ericsson', '26-Q2'], '#2E8B57', 4, []),
        (['VF_Huawei', '26-Q1'], '#E15759', 1, []),
        (['VF_Huawei', '26-Q2'], '#E15759', 4, []),
    ]
    assert {key: value for key, value in model['legend'].items() if key != 'format'} == {
        'position': 'right',
        'line_markers': True,
        'items': [
            {'label': 'VF_Ericsson · 26-Q1', 'colour': '#2E8B57', 'width': 1, 'dash': []},
            {'label': 'VF_Ericsson · 26-Q2', 'colour': '#2E8B57', 'width': 4, 'dash': []},
            {'label': 'VF_Huawei · 26-Q1', 'colour': '#E15759', 'width': 1, 'dash': []},
            {'label': 'VF_Huawei · 26-Q2', 'colour': '#E15759', 'width': 4, 'dash': []},
        ],
    }
    assert all(series['x'] and series['y'] and series['samples'] == 2 for series in model['series'])


@pytest.mark.parametrize('chart_type', ['CDF Line', 'Histogram Line'])
def test_interactive_cdf_model_uses_progressive_campaign_widths(chart_type) -> None:
    entry = CatalogEntry(
        1, 'Data', '', '', 'Rate CDF', 'CDR-Data', 'Rate', chart_type,
        'Operator, Campaign', '', 'Operator', 'Campaign', 'Top',
    )
    campaigns = ['2026-Q1', '2026-Q2', '2026-Q3', '2026-Q4']
    frame = chart_frame({
        'Operator': ['EE'] * 8,
        'Campaign': [campaign for campaign in campaigns for _ in range(2)],
        'Rate': [1.0, 2.0, 1.1, 2.1, 1.2, 2.2, 1.3, 2.3],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert [series['width'] for series in model['series']] == [1, 2, 3, 4]
    assert [item['width'] for item in model['legend']['items']] == [1, 2, 3, 4]


def test_multi_cdf_payload_bounds_high_cardinality_identifier_groups() -> None:
    entries = 250
    entry = CatalogEntry(
        1, 'Interactivity score', '', '', 'Interactivity score', 'CDR-Data',
        'Twamp Interactivity Score | Packet Delay Variation Score | Packet Loss Score',
        'Multi KPI CDF Lines', 'Operator', '',
        'Test Info × Test_Name × Latency Score × Operator', '', 'Right',
    )
    frame = chart_frame({
        'Test_Info': [f'test-{index}' for index in range(entries)],
        'Test_Name': ['Interactivity'] * entries,
        'Latency_Score': [str(index % 4) for index in range(entries)],
        'Operator': ['EE' if index % 2 else 'VF' for index in range(entries)],
        'Twamp_Interactivity_Score': list(range(entries)),
        'Packet_Delay_Variation_Score': list(range(entries)),
        'Packet_Loss_Score': list(range(entries)),
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['type'] == 'multi_cdf'
    assert model['legend']['items'] == []
    assert all(len(panel['series']) <= 120 for panel in model['panels'])
    assert all({series['key'][-1] for series in panel['series']} == {'EE', 'VF'} for panel in model['panels'])


@pytest.mark.parametrize('chart_type', ['CDF Line', 'Histogram Line'])
def test_multivendor_cdf_legend_keeps_operator_and_vendor_for_each_curve(chart_type) -> None:
    entry = CatalogEntry(
        1, 'Speech', '', '', 'Interactivity', 'CDR-Speech', 'LQ', chart_type,
        'Operator', '', 'Operator', 'Campaign', 'Bottom',
    )
    frame = chart_frame({
        'Operator_Vendor': ['VF_Ericsson'] * 3 + ['VF_Huawei'] * 3 + ['3_Ericsson'] * 3,
        'Campaign': ['2026 Q1'] * 9,
        'LQ': [1.0, 2.0, 3.0] * 3,
    })

    model = catalog_chart_payload(frame, entry, multivendor=True, prefiltered=True)

    labels = [item['label'] for item in model['legend']['items']]
    assert labels == ['Ericsson · VF · 26-Q1', 'Ericsson · 3 · 26-Q1', 'Huawei · VF · 26-Q1']
    assert [series['name'] for series in model['series']] == [
        'Ericsson · VF · 26-Q1', 'Ericsson · 3 · 26-Q1', 'Huawei · VF · 26-Q1',
    ]
    assert [item['colour'] for item in model['legend']['items']] == ['#2E8B57', '#2E8B57', '#E15759']
    assert [item['dash'] for item in model['legend']['items']] == [[], [18, 8], []]


def test_interactive_status_model_preserves_reporting_row_and_column_aggregation() -> None:
    entry = CatalogEntry(
        1, 'Voice', '', '', 'Completed Ratio', 'CDR-Voice', 'Test_Result',
        '100% Stacked Vertical Bars', 'Test_Result', '', 'Call Family',
        'Operator × Campaign', 'Right',
    )
    entry = replace(entry, axis_y_range='[25,75]', label_position='middle')
    frame = chart_frame({
        'Call Family': ['VoLTE'] * 8,
        'Operator': ['VF'] * 4 + ['3'] * 4,
        'Campaign': ['UK_Q2_2026', 'UK_Q2_2026', 'UK_Q1_2026', 'UK_Q1_2026'] * 2,
        'Test_Result': [
            'Completed', 'Failed', 'Completed', 'Completed',
            'Failed', 'Failed', 'Completed', 'Failed',
        ],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert (model['type'], model['mode'], model['row_keys']) == ('status_100', 'hierarchy', [['VoLTE']])
    assert model['domain']['y'] == [.25, .75]
    assert model['label_position'] == 'middle'
    assert model['column_keys'] == [
        ['VF', '26-Q1'], ['VF', '26-Q2'],
        ['3', '26-Q1'], ['3', '26-Q2'],
    ]
    assert model['cells'][0] == [
        pytest.approx([1.0, 0.0]), pytest.approx([0.5, 0.5]),
        pytest.approx([0.5, 0.5]), pytest.approx([0.0, 1.0]),
    ]
    assert model['states'] == [
        {'name': 'Completed', 'colour': '#2C9A62'},
        {'name': 'Failed', 'colour': '#E15759'},
    ]
    assert model['legend']['items'] == [
        {'label': 'Completed', 'colour': '#2C9A62', 'width': 2},
        {'label': 'Failed', 'colour': '#E15759', 'width': 2},
    ]


def test_interactive_status_model_preserves_sql_aggregated_row_weights() -> None:
    entry = CatalogEntry(
        1, 'Data', '', '', 'Completed Ratio', 'CDR-Data', 'Test_Result',
        '100% Stacked Vertical Bars', 'Test_Result', '', 'Type_of_Test',
        'Operator × Campaign', 'Right',
    )
    frame = chart_frame({
        'Type_of_Test': ['Browsing'] * 6,
        'Operator': ['VF'] * 3 + ['3'] * 3,
        'Campaign': ['2026 Q1'] * 2 + ['2026 Q2'] + ['2026 Q1'] + ['2026 Q2'] * 2,
        'Test_Result': ['Completed', 'Failed', 'Completed', 'Failed', 'Completed', 'Completed'],
    })
    weighted = (
        frame.groupby(
            ['Type_of_Test', 'Operator', 'Campaign', 'Test_Result'],
            sort=False,
            dropna=False,
        )
        .size()
        .rename('__catalog_weight')
        .reset_index()
    )
    weighted.attrs.update(CHART_MAPPING_ATTRS)

    expanded_model = catalog_chart_payload(frame, entry, prefiltered=True)
    weighted_model = catalog_chart_payload(weighted, entry, prefiltered=True)

    assert weighted_model == expanded_model


def test_interactive_status_legend_uses_the_colours_of_its_plotted_states() -> None:
    entry = CatalogEntry(
        1, 'Voice', '', '', 'Failed ratio', 'CDR-Voice', 'Test_Result',
        '100% Stacked Vertical Bars', 'Operator', '', 'Test Name', 'Operator', 'Right',
    )
    frame = chart_frame({
        'Test Name': ['httpBrowser', 'httpBrowser', 'VideoStreaming', 'VideoStreaming'],
        'Operator': ['O2', 'O2', 'VF', 'VF'],
        'Test_Result': ['Completed', 'Failed', 'Completed', 'Failed'],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['legend']['items'] == [
        {'label': state['name'], 'colour': state['colour'], 'width': 2}
        for state in model['states']
    ]


def test_interactive_distribution_model_uses_reporting_buckets_and_nested_keys() -> None:
    entry = CatalogEntry(
        1, 'Data', '', '', 'Rate distribution', 'CDR-Data', 'Mean_Data_Rate',
        'Distribution Stacked Vertical Bars', 'Buckets', 'Buckets = 1,5,20;',
        'Operator', 'Campaign × Rate Bucket', 'Bottom',
    )
    entry = replace(entry, axis_y_range='[10,90]', label_position='down')
    frame = chart_frame({
        'Operator': ['VF'] * 4 + ['3'] * 4,
        'Campaign': ['UK_Q2_2026', 'UK_Q2_2026', 'UK_Q1_2026', 'UK_Q1_2026'] * 2,
        'Mean_Data_Rate': [2, 25, .5, 7, 2, 2, .2, 25],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['type'] == 'distribution'
    assert model['domain']['y'] == [.1, .9]
    assert model['label_position'] == 'down'
    assert model['keys'] == [
        ['VF', '26-Q1'], ['VF', '26-Q2'],
        ['3', '26-Q1'], ['3', '26-Q2'],
    ]
    assert [bucket['name'] for bucket in model['buckets']] == ['<1', '1-5', '5-20', '20+']
    assert model['cells'] == [
        pytest.approx([.5, 0, .5, 0]), pytest.approx([0, .5, 0, .5]),
        pytest.approx([.5, 0, 0, .5]), pytest.approx([0, 1, 0, 0]),
    ]
    assert model['legend']['position'] == 'bottom'
    assert [item['label'] for item in model['legend']['items']] == ['<1', '1-5', '5-20', '20+']
    assert [item['colour'] for item in model['legend']['items']] == [
        bucket['colour'] for bucket in model['buckets']
    ]


def test_distribution_upper_bound_buckets_match_tableau_formula_and_palette() -> None:
    entry = CatalogEntry(
        1, 'FDTT DL', '', '', 'FDTT http DL MT', 'CDR-Data', 'Mean_Data_Rate',
        'Distribution Stacked Vertical Bars', 'Buckets', 'Buckets < 2,5,20,100;',
        'Operator', 'Campaign × Rate Bucket', 'Right',
    )
    frame = chart_frame({
        'Operator': ['EE'] * 6,
        'Campaign': ['UK_Q2_2026'] * 6,
        'Mean_Data_Rate': [.5, 2, 5, 20, 100, 101],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert [bucket['name'] for bucket in model['buckets']] == [
        'Above', 'below100', 'below20', 'below5', 'below2',
    ]
    assert [bucket['colour'] for bucket in model['buckets']] == [
        '#4E79A7', '#76B7B2', '#E15759', '#F28E2B', '#FF9DA7',
    ]
    assert model['cells'] == [pytest.approx([.2, .2, .2, .2, .2])]
    assert model['legend']['items'] == [
        {'label': 'below2', 'colour': '#FF9DA7', 'width': 2},
        {'label': 'below5', 'colour': '#F28E2B', 'width': 2},
        {'label': 'below20', 'colour': '#E15759', 'width': 2},
        {'label': 'below100', 'colour': '#76B7B2', 'width': 2},
        {'label': 'Above', 'colour': '#4E79A7', 'width': 2},
    ]


def test_distribution_inclusive_upper_bound_buckets_are_supported() -> None:
    entry = CatalogEntry(
        1, 'Custom', '', '', 'Custom distribution', 'CDR-Data', 'Rate',
        'Distribution Stacked Vertical Bars', 'Buckets', 'Buckets <= 1,3,10,20;',
        'Operator', 'Campaign × Rate Bucket', 'Right',
    )
    frame = chart_frame({
        'Operator': ['EE'] * 5,
        'Campaign': ['2026-Q2'] * 5,
        'Rate': [1, 3, 10, 20, 21],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert [bucket['name'] for bucket in model['buckets']] == [
        'Above', 'below20', 'below10', 'below3', 'below1',
    ]
    assert model['cells'] == [pytest.approx([.2, .2, .2, .2, .2])]


def test_static_distribution_draws_horizontal_white_percentage_labels_inside_segments() -> None:
    frame = chart_frame({
        'Operator': ['EE'] * 10,
        'Campaign': ['2026-Q2'] * 10,
        '__catalog_stack': ['Above'] * 8 + ['below100'] * 2,
    })
    frame.attrs['catalogue_distribution_buckets'] = [
        'below2', 'below5', 'below20', 'below100', 'Above',
    ]

    with patch('src.modules.cdr_reporting._draw_inside_horizontal_bar_label', return_value=True) as draw_label:
        _render_stacked_distribution(
            'Distribution', frame, 'Operator', 'Campaign', '__catalog_stack',
        )

    visible = [item for item in draw_label.call_args_list if item.kwargs['height'] > 0]
    assert [item.args[1] for item in visible] == ['80.00%', '20.00%']
    assert {item.kwargs['fill'] for item in visible} == {'#FFFFFF'}
    assert {item.kwargs['font'].size for item in visible} == {14}


def test_interactive_mean_model_omits_groups_without_valid_samples() -> None:
    entry = CatalogEntry(
        1, 'Browsing', '', '', 'Browsing time', 'CDR-Data', 'Duration',
        'Average Vertical Bars', 'Vendor', '', 'Vendor', 'Campaign', 'Right',
    )
    frame = chart_frame({
        'Vendor': ['VF_Ericsson'] * 2 + ['VF_Huawei'] * 2,
        'Campaign': ['UK_Q1_2026', 'UK_Q2_2026'] * 2,
        'Duration': [1.5, 2.5, 3.0, None],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    # The JSON response rejects NaN: a group without samples has no bar.
    json.dumps(model, allow_nan=False)
    assert [(bar['key'], bar['value']) for bar in model['bars']] == [
        (['VF_Ericsson', '26-Q1'], pytest.approx(1.5)),
        (['VF_Ericsson', '26-Q2'], pytest.approx(2.5)),
        (['VF_Huawei', '26-Q1'], pytest.approx(3.0)),
    ]


def test_interactive_mean_model_uses_reporting_aggregation_and_vendor_palette() -> None:
    entry = CatalogEntry(
        1, 'Speech', '', '', 'Average POLQA', 'CDR-Speech', 'LQ',
        'Average Vertical Bars', 'Vendor', '', 'Vendor', 'Campaign', 'Right',
    )
    frame = chart_frame({
        'Vendor': ['VF_Ericsson'] * 4 + ['VF_Huawei'] * 4,
        'Campaign': ['UK_Q2_2026', 'UK_Q2_2026', 'UK_Q1_2026', 'UK_Q1_2026'] * 2,
        'LQ': [4.0, 4.4, 3.8, 4.0, 3.0, 3.4, 2.8, 3.0],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['type'] == 'mean_bar'
    assert [(bar['key'], bar['value'], bar['colour']) for bar in model['bars']] == [
        (['VF_Ericsson', '26-Q1'], pytest.approx(3.9), '#2E8B57'),
        (['VF_Ericsson', '26-Q2'], pytest.approx(4.2), '#2E8B57'),
        (['VF_Huawei', '26-Q1'], pytest.approx(2.9), '#E15759'),
        (['VF_Huawei', '26-Q2'], pytest.approx(3.2), '#E15759'),
    ]
    assert [(item['label'], item['colour']) for item in model['legend']['items']] == [
        ('VF_Ericsson', '#2E8B57'), ('VF_Huawei', '#E15759'),
    ]


def test_interactive_mean_model_nests_column_only_hierarchies_by_parent() -> None:
    entry = CatalogEntry(
        1, 'Speech', '', '', 'Average POLQA', 'CDR-Speech', 'LQ',
        'Average Vertical Bars', 'Vendor', '', '', 'Operator × Campaign', 'Right',
    )
    frame = chart_frame({
        # This is the natural order after appending one CDR per campaign.
        'Operator': ['VF', '3', 'EE', 'VF', '3', 'EE'],
        'Campaign': ['UK_Q1_2026'] * 3 + ['UK_Q2_2026'] * 3,
        'LQ': [4.0, 3.8, 4.2, 4.1, 3.9, 4.3],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['mode'] == 'hierarchy'
    assert model['row_keys'] == [[]]
    assert model['column_keys'] == [
        ['VF', '26-Q1'], ['VF', '26-Q2'],
        ['3', '26-Q1'], ['3', '26-Q2'],
        ['EE', '26-Q1'], ['EE', '26-Q2'],
    ]
    assert model['cell_colours'] == [[
        '#E15759', '#E15759',
        '#F28E2B', '#F28E2B',
        '#76B7B2', '#76B7B2',
    ]]
    filtered_entry = replace(entry, filters='Operator IN (VF, 3, EE)')
    filtered = _apply_catalog_filters(frame, filtered_entry, False, 'LQ')
    filtered_model = catalog_chart_payload(filtered, filtered_entry, prefiltered=True)
    assert filtered_model['column_keys'] == model['column_keys']
    assert filtered_model['cell_colours'] == model['cell_colours']


def test_bar_value_label_rotates_when_it_only_fits_vertically() -> None:
    image = Image.new('RGB', (200, 200), 'white')
    draw = ImageDraw.Draw(image)

    with patch('src.modules.cdr_reporting._draw_vertical_label') as draw_vertical:
        drawn = _draw_inside_bar_label(
            image, draw, '104.62', x=20, y=20, width=30, height=150,
            fill='white', font=ImageFont.load_default(),
        )

    assert drawn is True
    draw_vertical.assert_called_once()


def test_configured_vertical_bar_label_rotates_when_many_bars_make_it_too_wide() -> None:
    image = Image.new('RGB', (200, 200), 'white')
    draw = ImageDraw.Draw(image)

    with patch('src.modules.cdr_reporting._draw_vertical_label') as draw_vertical:
        _draw_configured_bar_label(
            image, draw, '104.62', x=20, y=20, width=30, height=150,
            colour='#4E79A7', font=ImageFont.load_default(), position='middle',
            horizontal=False, automatic=lambda: None,
        )

    draw_vertical.assert_called_once()


def test_tiny_stacked_label_uses_clear_space_beside_its_bar() -> None:
    draw = MagicMock()
    draw.textbbox.return_value = (0, 0, 28, 12)

    drawn = _draw_adjacent_stacked_bar_label(
        draw, '0.6%', x=100, y=20, width=80, height=2,
        bar_top=20, bar_bottom=180, side_space=40, fill='#C83E4D', font=ImageFont.load_default(),
    )

    assert drawn is True
    draw.rectangle.assert_not_called()
    assert draw.text.call_args.args[0][0] == pytest.approx(184)
    assert draw.text.call_args.kwargs['fill'] == '#C83E4D'


def test_tiny_stacked_label_is_hidden_without_lateral_space() -> None:
    draw = MagicMock()
    draw.textbbox.return_value = (0, 0, 28, 12)

    drawn = _draw_adjacent_stacked_bar_label(
        draw, '0.6%', x=100, y=20, width=80, height=2,
        bar_top=20, bar_bottom=180, side_space=5, fill='#C83E4D',
        font=ImageFont.load_default(),
    )

    assert drawn is False
    draw.text.assert_not_called()


def test_top_column_group_separator_is_solid_from_the_header_to_the_plot() -> None:
    draw = MagicMock()

    _draw_top_column_group_separators(
        draw,
        [('VF', '2026-Q1'), ('VF', '2026-Q2'), ('3', '2026-Q1')],
        ['__catalog_column_0', '__catalog_column_1'],
        left=155, width=1165, top=280, bottom=680,
    )

    assert call((155 + 2 * 1165 / 3, 212, 155 + 2 * 1165 / 3, 680), fill='#AEBBC4', width=2) in draw.line.call_args_list
    assert call((155 + 1165 / 3, 242, 155 + 1165 / 3, 250), fill='#AEBBC4', width=1) in draw.line.call_args_list


def test_nested_column_group_separator_is_dashed_below_its_parent_header() -> None:
    draw = MagicMock()

    _draw_top_column_group_separators(
        draw,
        [('VF', 'Ericsson', '2026-Q1'), ('VF', 'Ericsson', '2026-Q2')],
        ['__catalog_row_0', '__catalog_row_1', '__catalog_column_0'],
        left=155, width=1165, top=280, bottom=680,
    )

    assert draw.line.call_args_list[0] == call.line((155 + 1165 / 2, 242, 155 + 1165 / 2, 250), fill='#AEBBC4', width=1)
    assert all(item.kwargs.get('width') == 1 for item in draw.line.call_args_list)


def test_interactive_failure_model_matches_reporting_legend_plot_geometry() -> None:
    entry = CatalogEntry(
        1, 'Failures', '', '', 'Voice failures', 'CDR-Voice', 'Call_Status',
        'Count Stacked Horizontal Bars', 'Call_Status', '', 'Call Family',
        'Operator × Campaign', 'Right',
    )
    entry = replace(entry, axis_x_range='[0,4]', label_position='top')
    frame = chart_frame({
        'Call Family': ['VoLTE', 'VoLTE', 'MultiRAB'],
        'Operator': ['VF', 'VF', '3'],
        'Campaign': ['2026 Q1', '2026 Q2', '2026 Q2'],
        'Call_Status': ['Failed', 'Dropped', 'Completed'],
    })

    field_legend = catalog_chart_payload(frame, entry, prefiltered=True)
    manual_legend = catalog_chart_payload(
        frame, replace(entry, legend='Failed/Dropped'), prefiltered=True,
    )

    assert field_legend['type'] == manual_legend['type'] == 'failure_count'
    assert field_legend['domain']['x'] == [0.0, 4.0]
    assert field_legend['label_position'] == 'top'
    assert field_legend['plot_legend_position'] == 'none'
    assert manual_legend['plot_legend_position'] == 'right'
    assert [state['name'] for state in manual_legend['states']] == ['Failed', 'Dropped']


def test_interactive_map_model_preserves_operator_colours_and_osm_geometry() -> None:
    entry = CatalogEntry(
        1, 'Map', '', '', 'Coverage', 'CDR-Data', 'Latitude vs Longitude',
        'Map', 'Operator', '', 'Operator', 'Campaign', 'Right',
    )
    entry = replace(entry, axis_x_range='[-2,0]', axis_y_range='[50,55]', label_position='up')
    frame = chart_frame({
        'Latitude': [51.5, 51.51, 53.8],
        'Longitude': [-.12, -.11, -1.55],
        'Operator': ['VF', 'VF', '3'],
        'Campaign': ['2026 Q2'] * 3,
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert model['type'] == 'map'
    assert model['domain'] == {'x': [-2.0, 0.0], 'y': [50.0, 55.0]}
    assert model['label_position'] == 'up'
    assert [(series['name'], series['colour'], len(series['points'])) for series in model['series']] == [
        ('VF', '#E15759', 2), ('3', '#F28E2B', 1),
    ]
    assert model['basemap']['provider'] == 'OpenStreetMap'
    tile_left, tile_top, tile_right, tile_bottom = model['basemap']['tile_range']
    assert (tile_right - tile_left + 1) * (tile_bottom - tile_top + 1) <= 48
    assert len(model['basemap']['world_bounds']) == 4


def test_interactive_map_keeps_manual_legend_labels_only_when_they_cover_every_series() -> None:
    entry = CatalogEntry(
        1, 'Map', '', '', 'Coverage', 'CDR-Data', 'Latitude vs Longitude',
        'Map', 'Success/Failure', '', 'Operator', '', 'Right',
    )
    frame = chart_frame({
        'Latitude': [51.5, 51.51, 53.8],
        'Longitude': [-.12, -.11, -1.55],
        'Operator': ['VF', '3', 'EE'],
    })

    model = catalog_chart_payload(frame, entry, prefiltered=True)

    assert [(item['label'], item['colour']) for item in model['legend']['items']] == [
        (series['name'], series['colour']) for series in model['series']
    ]


def test_cdf_plot_reserves_space_for_side_legends() -> None:
    left, _top, width, _height = _cdf_plot_geometry('right')
    assert left + width < 1320
    left, _top, width, _height = _cdf_plot_geometry('left')
    assert left >= 400
    assert _cdf_plot_geometry('top') == (100, 135, 1320, 590)


def test_threshold_legend_uses_the_resolved_value_and_chart_segment_colours() -> None:
    entry = CatalogEntry(
        slide=8, slide_title='', slide_subtitle='', layout='', chart_title='POLQA <1.6',
        cdr_source='CDR-Speech', kpi='POLQA < 1.6 vs >= 1.6',
        chart_type='Threshold Stacked Vertical Bars', legend='Threshold', filters='',
        rows_split='Operator', columns_split='Campaign', legend_position='bottom',
    )
    assert _resolved_legend_items(entry, pd.DataFrame({'POLQA': [1.2, 2.1]}), 'POLQA') == [
        ('< 1.6', '#E15759', 2),
        ('≥ 1.6', '#59A14F', 2),
    ]


def test_distribution_bucket_legend_uses_resolved_bucket_colours_not_filter_text() -> None:
    frame = chart_frame({
        '__catalog_stack': ['20-100', '100+', '5-20', '1-5', '<1', '20-100'],
    })
    entry = CatalogEntry(
        slide=10, slide_title='', slide_subtitle='', layout='', chart_title='FDTT DL (7s)',
        cdr_source='CDR-Data', kpi='FDTT_Sustainable_MDR',
        chart_type='Distribution Stacked Vertical Bars', legend='Buckets',
        filters='Buckets = 1,5,20,100;', rows_split='Operator',
        columns_split='Campaign x Rate Bucket', legend_position='bottom',
    )
    items = _resolved_legend_items(entry, frame, 'FDTT_Sustainable_MDR')
    assert [caption for caption, _colour_value, _width in items] == ['<1', '1-5', '5-20', '20-100', '100+']
    assert all(colour for _caption, colour, _width in items)


def test_mean_chart_renders_selected_dimension_legend_at_requested_position() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['VF', '3'], '__catalog_column_0': ['2026-Q2', '2026-Q2'],
        'Metric': [1.0, 2.0], '__catalog_primary': ['VF', '3'], '__catalog_series': ['2026-Q2', '2026-Q2'],
    })
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': 'Operator', '__catalog_column_0': 'Campaign',
    }

    with patch('src.modules.cdr_reporting._draw_chart_legend') as draw_legend:
        _render_mean_column(
            'Average', frame, '__catalog_primary', '__catalog_series', 'Metric',
            legend_dimensions=('Operator',), legend_position='right',
        )

    assert {item[0] for item in draw_legend.call_args.args[1]} == {'VF', '3'}
    assert draw_legend.call_args.args[2] == 'right'


def test_status_chart_draws_legend_at_the_catalogue_position() -> None:
    frame = chart_frame({
        'Operator': ['Vodafone'],
        'Campaign': ['2026 Q1'],
        'Call_Status': ['Completed'],
    })

    with patch('src.modules.cdr_reporting._draw_chart_legend') as draw_legend:
        _render_status_100('Status', frame, 'Operator', 'Campaign', legend_position='left')

    assert draw_legend.call_args.args[2] == 'left'


def test_chart_legend_normalises_title_case_position_from_interactive_preview() -> None:
    draw = MagicMock()

    _draw_chart_legend(draw, [('VF', '#E60000', 1)], 'Left')

    assert draw.rectangle.call_args.args[0][0] == 26


def test_catalogue_accepts_structural_slides_and_rejects_chart_configuration_on_them() -> None:
    title = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n1,Quarterly report,NSA analysis,Title Page,,,,Title Slide,,,,,\n',
    )[0]
    transition = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n2,Voice analysis,Seven cities,Title Only,,,,Transition Slide,,,,,\n',
    )[0]

    assert title.structural_type == 'title slide'
    assert transition.structural_type == 'transition slide'
    with pytest.raises(ValueError, match='cannot define chart, CDR, KPI'):
        parse_catalog_csv(
            ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n1,Quarterly report,,Title Page,,CDR-Data,LQ,Title Slide,,,,,\n',
        )


def test_catalogue_reads_structural_layouts_without_chart_type_as_their_structural_slides() -> None:
    cover, transition, divider, implicit, explicit, moved = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n1,Charts Imported from Tableau,Automatically Generated by GenAI,Title Page,,,,,,,,,'
        + '\n2,Voice,,Transition,,,,,,,,,\n3,Data,,Title Only,,,,,,,,,'
        + '\n18,,,Black logo end slide,,,,,,,,,\n19,,,Black logo end slide,,,,End Slide,,,,,'
        # A Transition row moved to the closing layout, still with its former Chart type, is an End Slide.
        + '\n20,,,Black logo end slide,,,,Transition Slide,,,,,\n',
    )

    assert (cover.structural_type, cover.chart_type) == ('title slide', 'Title Slide')
    assert [transition.chart_type, divider.chart_type] == ['Transition Slide', 'Transition Slide']
    assert (implicit.structural_type, implicit.chart_type) == ('end slide', 'End Slide')
    assert explicit.structural_type == 'end slide'
    assert (moved.structural_type, moved.chart_type) == ('end slide', 'End Slide')
    saved = catalogue_csv([cover, implicit]).decode('utf-8')
    assert 'Title Slide' in saved and 'End Slide' in saved
    with pytest.raises(ValueError, match=r'^Slide: 2 -> Slide Tittle is required for a Transition Slide\.$'):
        parse_catalog_csv(','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n2,,,Title Only,,,,Transition Slide,,,,,\n')
    with pytest.raises(ValueError, match='cannot define chart, CDR, KPI'):
        parse_catalog_csv(
            ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n1,Quarterly report,,Title Page,,CDR-Data,LQ,Title Slide,,,,,\n',
        )


def test_catalogue_validation_names_the_slide_and_chart_of_the_row() -> None:
    header = ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
    chart = '3,Rate,,Title + 1 rows + 2 columns,A,CDR-Data,Mean_Data_Rate,CDF Line,,Operator,,'

    with pytest.raises(ValueError, match=r'^Slide: 3 - Chart: 2 -> KPI and Chart type are required'):
        parse_catalog_csv(f'{header}\n{chart}\n3,Rate,,Title + 1 rows + 2 columns,B,CDR-Data,,CDF Line,,Operator,,\n')
    with pytest.raises(ValueError, match=r"^Slide: 3 - Chart: 1 -> Chart type 'Radar' is not supported\."):
        parse_catalog_csv(f"{header}\n{chart.replace('CDF Line', 'Radar')}\n")
    with pytest.raises(ValueError, match=r'^Slide: x \(after Slide: 3\) -> The Slide value must be a positive whole number\.'):
        parse_catalog_csv(f'{header}\n{chart}\nx,Rate,,Title Page,,,,Title Slide,,,,,\n')
    with pytest.raises(ValueError, match=r'^Slide: 4 - Chart: 1 -> The row must define a supported CDR chart'):
        parse_catalog_csv(f'{header}\n4,Notes,,Title + 1 rows + 1 columns,,,,,,,,,\n')


def test_catalogue_filter_and_grouping_contract_is_parsed_and_applied() -> None:
    conditions = parse_catalog_filters('Session_Type IN (VoLTE, MultiRAB); LQ >= 1.6')
    grouping = parse_catalog_grouping('City × Operator × Campaign')
    assert [(item.column, item.operator, item.values) for item in conditions] == [
        ('Session_Type', 'IN', ('VoLTE', 'MultiRAB')), ('LQ', '>=', ('1.6',)),
    ]
    assert grouping.dimensions == ('City', 'Operator', 'Campaign')
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n8,Quality,,Title and 1 column + Comments,Quality by city,CDR-Speech,LQ,Average Vertical Bars,Session_Type IN (VoLTE); LQ >= 1.6,City,Operator × Campaign,,\n',
    )[0]
    frame = chart_frame({
        'Session_Type': ['VoLTE', 'WhatsApp'], 'LQ': [3.2, 4.0], 'City': ['London', 'Leeds'],
        'Operator': ['EE', 'O2'], 'Campaign': ['Q1', 'Q1'],
    })
    filtered = _apply_catalog_filters(frame, entry, False, 'LQ')
    grouped, primary, series = _apply_catalog_grouping(filtered, entry, False, 'LQ')
    assert len(grouped) == 1
    assert grouped[primary].tolist() == ['London']
    assert grouped[series].tolist() == ['EE · Q1']
    assert '__catalog_stack' not in grouped.columns


def test_catalogue_filters_accept_lines_but_reject_a_missing_condition_separator() -> None:
    conditions = parse_catalog_filters('Test Name CONTAINS FDFS\nTest_Result IN (Completed, Dropped, Failed)')
    assert [condition.column for condition in conditions] == ['Test Name', 'Test_Result']

    malformed = (
        'Test Name CONTAINS FDFS;'
        'Test_Result IN (Completed, Dropped, Failed)Operator IN (Vodafone UK, 3, EE);'
        'G Level 4 IN (Belfast, Bristol, Cardiff, Edinburgh, London, Leeds, Sheffield);'
    )
    with pytest.raises(ValueError, match='semicolon is required'):
        parse_catalog_filters(malformed)

    row = ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + (
        '\n5,Failures,,Title and 1 column + Comments,Failures,CDR-Data,Test_Result,'
        '100% Stacked Vertical Bars,"Test_Result IN (Completed, Dropped, Failed)Operator IN (EE, 3)",'
        'Operator,Campaign,,Top\n'
    )
    with pytest.raises(ValueError, match=r'Slide: 5 - Chart: 1 ->.*semicolon is required'):
        parse_catalog_csv(row)
    editable = parse_catalog_csv(row, validate_filters=False)
    assert editable[0].filters.endswith('Operator IN (EE, 3)')


def test_multivendor_rendering_rewrites_display_and_grouping_and_excludes_unresolved_vendors() -> None:
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Operator comparison,Operator subtitle,Title and 1 column + Comments,Operator chart,CDR-Speech,LQ,Average Vertical Bars,Operator = Vodafone UK,Operator,Operator × Campaign,Operator,\n',
    )[0]
    rendered = prepare_multivendor_catalog_entry(entry)

    assert rendered.slide_title == 'Vendor comparison'
    assert rendered.slide_subtitle == 'Vendor subtitle'
    assert rendered.chart_title == 'Vendor chart'
    assert rendered.legend == 'Operator_Vendor, Operator'
    assert rendered.rows_split == 'Operator_Vendor × Operator'
    assert rendered.columns_split == 'Operator_Vendor × Operator × Campaign'
    assert rendered.filters == 'Operator = Vodafone UK; Vendor NOT CONTAINS (Mixed, Other)'

    frame = chart_frame({
        'Operator': ['Vodafone UK', 'Vodafone UK', '3'],
        'Operator_Vendor': ['Vodafone_Ericsson', 'Vodafone_Mixed Vendor', '3_Nokia'],
        'Vendor': ['Ericsson', 'Mixed Vendor', 'Nokia'],
        'Campaign': ['UK_Q2_SA_2026', 'UK_Q2_SA_2026', 'UK_Q2_SA_2026'],
        'LQ': [3.8, 3.6, 3.5],
    })
    frame = normalise_operator_aliases(frame, {'vodafone uk': 'VF', 'vodafone': 'VF'})
    filtered = _apply_catalog_filters(frame, rendered, True, 'LQ')
    grouped, primary, series = _apply_catalog_grouping(filtered, rendered, True, 'LQ')
    assert grouped[primary].tolist() == ['Ericsson · VF']
    assert grouped[series].tolist() == ['Ericsson · VF · 26-Q2-SA']

    already_filtered = replace(entry, filters='vendor NOT CONTAINS (Mixed, Other)')
    assert prepare_multivendor_catalog_entry(already_filtered).filters == already_filtered.filters


def test_vendor_filters_accept_full_or_operator_independent_vendor_values() -> None:
    frame = chart_frame({
        'Operator_Vendor': [
            'VF_Ericsson', 'VF_Huawei', 'VF_Mixed Vendor', '3_Ericsson',
            '3_Huawei', '3_Samsung', '3_Mixed Vendor', 'O2_NSN',
        ],
        'Vendor': [
            'Ericsson', 'Huawei', 'Mixed Vendor', 'Ericsson', 'Huawei',
            'Samsung', 'Mixed Vendor', 'NSN',
        ],
        'LQ': [4.0] * 8,
    })
    bare_entry = CatalogEntry(
        1, '', '', '', '', 'CDR-Speech', 'LQ', 'Average Vertical Bars',
        'Vendor IN (Ericsson, Huawei, Samsung, NSN, Mixed Vendor)', '', 'Operator', '',
    )
    full_entry = replace(bare_entry, filters='Operator_Vendor IN (3_Ericsson, 3_Huawei, 3_Mixed Vendor, VF_Ericsson, VF_Huawei, VF_Mixed Vendor, O2_NSN)')

    bare = _apply_catalog_filters(frame, bare_entry, True, 'LQ')
    full = _apply_catalog_filters(frame, full_entry, True, 'LQ')

    assert bare['Operator_Vendor'].tolist() == frame['Operator_Vendor'].tolist()
    assert full['Operator_Vendor'].tolist() == [
        'VF_Ericsson', 'VF_Huawei', 'VF_Mixed Vendor', '3_Ericsson',
        '3_Huawei', '3_Mixed Vendor', 'O2_NSN',
    ]


def test_real_vendor_filter_does_not_include_unmapped_operator_only_rows() -> None:
    frame = chart_frame({
        'Operator': ['O2', 'EE'],
        'Operator_Vendor': ['O2', 'EE_Ericsson'],
        'Vendor': ['O2 - All', 'Ericsson'],
        'LQ': [3.0, 4.0],
    })
    entry = CatalogEntry(
        1, '', '', '', '', 'CDR-Speech', 'LQ', 'Average Vertical Bars',
        '', 'Vendor = Ericsson', 'Operator', '',
    )

    filtered = _apply_catalog_filters(frame, entry, False, 'LQ')

    assert filtered['Operator'].tolist() == ['EE']


def test_vendor_filter_alias_round_trips_through_catalogue_csv() -> None:
    entry = CatalogEntry(
        1, 'Vendor filter', '', 'Title + 1 rows + 1 columns', '', 'CDR-Speech', 'LQ', 'Average Vertical Bars',
        '', 'Vendor IN (VF_Ericsson, Ericsson)', 'Operator', '',
    )

    parsed = parse_catalog_csv(catalogue_csv([entry]))[0]

    assert parsed.filters == 'Vendor IN (VF_Ericsson, Ericsson)'
    assert parse_catalog_csv(catalogue_csv([parsed]))[0].filters == parsed.filters


def test_multivendor_operator_filters_match_vendor_prefixes_and_keep_full_grouping_values() -> None:
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Operator comparison,,Title and 1 column + Comments,Operator chart,CDR-Speech,LQ,Average Vertical Bars,"Operator IN (Vodafone UK, 3, O2)",Operator,Operator × Campaign,Operator,Top\n',
    )[0]
    rendered = prepare_multivendor_catalog_entry(entry)
    frame = chart_frame({
        'Operator': ['Vodafone UK', 'Vodafone UK', '3', 'O2', 'EE'],
        'Operator_Vendor': ['Vodafone_Ericsson', 'Vodafone_Huawei', '3_Nokia', 'O2_Ericsson', 'EE_Nokia'],
        'Vendor': ['Ericsson', 'Huawei', 'Nokia', 'Ericsson', 'Nokia'],
        'Campaign': ['2025 Q4', '2026 Q1', '2025 Q4', '2026 Q1', '2026 Q1'],
        'LQ': [3.8, 3.7, 3.6, 3.5, 3.4],
    })

    frame = normalise_operator_aliases(frame, {'vodafone uk': 'VF', 'vodafone': 'VF'})
    filtered = _apply_catalog_filters(frame, rendered, True, 'LQ')
    grouped, primary, series = _apply_catalog_grouping(filtered, rendered, True, 'LQ')

    assert filtered['Operator_Vendor'].tolist() == ['VF_Ericsson', 'VF_Huawei', '3_Nokia', 'O2_Ericsson']
    assert grouped[primary].tolist() == ['Ericsson · VF', 'Ericsson · O2', 'Huawei · VF', 'Nokia · 3']
    assert grouped[series].tolist() == [
        'Ericsson · VF · 25-Q4', 'Ericsson · O2 · 26-Q1',
        'Huawei · VF · 26-Q1', 'Nokia · 3 · 25-Q4',
    ]
    assert [caption for caption, _colour, _width in _resolved_legend_items(rendered, grouped, 'LQ')] == [
        'Ericsson · VF', 'Ericsson · O2', 'Huawei · VF', 'Nokia · 3',
    ]


def test_multivendor_grouping_orders_vendors_then_operators() -> None:
    entry = CatalogEntry(
        1, "", "", "", "", "CDR-Speech", "LQ", "Average Vertical Bars",
        "", "", "Operator", "Campaign", "Top",
    )
    frame = chart_frame({
        "Operator_Vendor": ["O2_NSN", "Lebara_NSN", "EE_NSN", "3_Ericsson", "VF_Huawei"],
        "Campaign": ["2026 Q2"] * 5,
        "LQ": [4.0] * 5,
    })

    grouped, primary, _series = _apply_catalog_grouping(
        frame, prepare_multivendor_catalog_entry(entry), True, "LQ",
    )

    assert grouped["__catalog_row_0"].tolist() == ["Ericsson", "Huawei", "NSN", "NSN", "NSN"]
    assert grouped["__catalog_row_1"].tolist() == ["3", "VF", "EE", "O2", "Lebara"]


def test_chart_grouping_uses_workspace_order_for_subscribers_and_combined_vendors() -> None:
    subscriber_entry = CatalogEntry(
        1, '', '', '', '', 'CDR-Data', 'Mean_Data_Rate', 'Average Vertical Bars',
        '', '', 'Subscriber', 'Campaign', 'Top',
    )
    subscriber_frame = chart_frame({
        'Subscriber': ['Alpha', 'Beta'], 'Campaign': ['2026 Q1', '2026 Q1'],
        'Mean_Data_Rate': [1.0, 2.0],
    })
    subscriber_frame.attrs['operator_mapping_groups'] = [
        {'canonical': 'Beta', 'aliases': [], 'position': 0, 'color': '#112233'},
        {'canonical': 'Alpha', 'aliases': [], 'position': 1, 'color': '#445566'},
    ]

    grouped, primary, _series = _apply_catalog_grouping(
        subscriber_frame, subscriber_entry, False, 'Mean_Data_Rate',
    )
    assert grouped[primary].tolist() == ['Beta', 'Alpha']

    vendor_entry = replace(subscriber_entry, rows_split='Operator_Vendor')
    vendor_frame = chart_frame({
        'Operator_Vendor': ['VF_Ericsson', 'VF_Huawei'], 'Campaign': ['2026 Q1', '2026 Q1'],
        'Mean_Data_Rate': [1.0, 2.0],
    })
    vendor_frame.attrs['vendor_mapping_groups'] = [
        {'canonical': 'Huawei', 'aliases': [], 'position': 0, 'color': '#ABCDEF'},
        {'canonical': 'Ericsson', 'aliases': [], 'position': 1, 'color': '#123456'},
    ]

    grouped, primary, _series = _apply_catalog_grouping(
        vendor_frame, vendor_entry, False, 'Mean_Data_Rate',
    )
    assert grouped[primary].tolist() == ['VF_Huawei', 'VF_Ericsson']


def test_cdf_campaigns_share_the_exact_workspace_theme_colour() -> None:
    frame = chart_frame({'operator': [], 'campaign': []})
    frame.attrs['operator_mapping_groups'] = [
        {'canonical': 'Carrier', 'aliases': [], 'position': 0, 'color': '#204060'},
    ]
    frame.attrs['catalogue_dimension_labels'] = {
        'operator': ('Operator',), 'campaign': ('Campaign',),
    }

    colours = _series_colours(
        [('Carrier', 'Q1'), ('Carrier', 'Q2')], ['operator', 'campaign'], frame, line_chart=True,
    )

    assert colours[('Carrier', 'Q1')] == '#204060'
    assert colours[('Carrier', 'Q2')] == '#204060'


def test_vf_uk_cdf_quarters_share_colour_and_use_width_for_recency() -> None:
    frame = chart_frame({'operator': [], 'campaign': []})
    frame.attrs['operator_mapping_groups'] = [
        {'canonical': 'VF_UK', 'aliases': [], 'position': 0, 'color': '#C83E4D'},
    ]
    frame.attrs['catalogue_dimension_labels'] = {
        'operator': ('Operator',), 'campaign': ('Campaign',),
    }
    keys = [('VF_UK', '2026-Q1'), ('VF_UK', '2026-Q2')]

    colours = _series_colours(keys, ['operator', 'campaign'], frame, line_chart=True)
    widths = _cdf_campaign_line_widths(
        [(key, [key[1]]) for key in keys], ['operator', 'campaign'], frame,
    )

    assert set(colours.values()) == {'#C83E4D'}
    assert widths[('VF_UK', '2026-Q2')] > widths[('VF_UK', '2026-Q1')]


def test_multivendor_grouping_uses_the_same_vendor_order_for_each_operator() -> None:
    entry = CatalogEntry(
        1, "", "", "", "", "CDR-Speech", "LQ", "Average Vertical Bars",
        "Operator_Vendor", "", "Operator_Vendor", "Campaign", "Top",
    )
    frame = chart_frame({
        "Operator_Vendor": [
            "VF_Samsung", "VF_NSN", "VF_Huawei", "VF_Ericsson",
            "3_Huawei", "3_Samsung", "3_Ericsson",
        ],
        "Campaign": ["2026 Q2"] * 7,
        "LQ": [4.0] * 7,
    })

    grouped, primary, _series = _apply_catalog_grouping(frame, prepare_multivendor_catalog_entry(entry), True, "LQ")

    assert grouped[primary].drop_duplicates().tolist() == [
        "Ericsson · VF", "Ericsson · 3", "Huawei · VF", "Huawei · 3",
        "Samsung · VF", "Samsung · 3", "NSN · VF",
    ]


def test_multivendor_grouping_does_not_split_an_exact_underscore_operator() -> None:
    entry = CatalogEntry(
        1, "", "", "", "", "CDR-Speech", "LQ", "Average Vertical Bars",
        "Operator_Vendor", "", "Operator_Vendor", "Campaign", "Top",
    )
    frame = chart_frame({
        "Operator": ["VF_SA"],
        "Operator_Vendor": ["VF_SA"],
        "Campaign": ["2026 Q2"],
        "LQ": [4.0],
    })
    frame.attrs['operator_mappings'] = {'vf': 'VF', 'vf_sa': 'VF_SA'}
    frame.attrs['vendor_mappings'] = {'sa': 'SA'}
    frame.attrs['operator_mapping_groups'].append({
        'canonical': 'VF_SA', 'aliases': [], 'position': 4, 'color': '#8000FF',
    })
    frame.attrs['vendor_mapping_groups'].append({
        'canonical': 'SA', 'aliases': [], 'position': 7, 'color': '#123456',
    })

    grouped, primary, _series = _apply_catalog_grouping(
        frame, prepare_multivendor_catalog_entry(entry), True, "LQ",
    )

    assert grouped[primary].tolist() == ['VF_SA · VF_SA']


def test_vendor_grouping_keeps_each_operator_together_outside_multivendor_mode() -> None:
    entry = CatalogEntry(
        1, "", "", "", "", "CDR-Speech", "LQ", "Average Vertical Bars",
        "Operator_Vendor", "", "Operator_Vendor", "Campaign", "Top",
    )
    frame = chart_frame({
        "Operator_Vendor": [
            "VF_Ericsson", "VF_Huawei", "3_Ericsson", "3_Huawei",
            "3_Samsung", "VF_Samsung", "VF_NSN",
        ],
        "Campaign": ["2026 Q2"] * 7,
        "LQ": [4.0] * 7,
    })

    grouped, primary, _series = _apply_catalog_grouping(frame, entry, False, "LQ")

    assert grouped[primary].drop_duplicates().tolist() == [
        "VF_Ericsson", "VF_Huawei", "VF_Samsung", "VF_NSN",
        "3_Ericsson", "3_Huawei", "3_Samsung",
    ]


def test_rows_only_grouping_uses_one_all_series_without_repeating_the_category() -> None:
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n8,Quality,,Title and 1 column + Comments,Quality,CDR-Speech,LQ,Average Vertical Bars,,Operator,,,\n',
    )[0]
    frame = chart_frame({'Operator': ['EE', 'O2'], 'LQ': [3.2, 3.8]})

    grouped, primary, series = _apply_catalog_grouping(frame, entry, False, 'LQ')

    assert grouped[primary].tolist() == ['EE', 'O2']
    assert grouped[series].tolist() == ['(all)', '(all)']


def test_campaign_grouping_displays_only_year_and_quarter() -> None:
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n8,Quality,,Title and 1 column + Comments,Quality,CDR-Speech,LQ,Average Vertical Bars,,Operator,Campaign,,\n',
    )[0]
    frame = chart_frame({
        'Operator': ['EE', '3', 'Vodafone UK'],
        'Campaign': ['UK_Q2_SA_2026', 'UK_Q4_2025', '2024 Q3 NSA'],
        'LQ': [3.2, 3.8, 4.0],
    })

    grouped, _primary, series = _apply_catalog_grouping(frame, entry, False, 'LQ')

    assert grouped['Campaign'].tolist() == ['2024 Q3 NSA', 'UK_Q4_2025', 'UK_Q2_SA_2026']
    assert grouped['__catalog_column_0'].tolist() == ['24-Q3-NSA', '25-Q4', '26-Q2-SA']
    assert grouped[series].tolist() == ['24-Q3-NSA', '25-Q4', '26-Q2-SA']


def test_cdf_renders_a_curve_for_each_complete_rows_and_columns_combination() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['Vodafone', 'Vodafone', 'Vodafone', 'Vodafone', 'O2', 'O2', 'O2', 'O2'],
        '__catalog_column_0': ['2025', '2025', '2026', '2026', '2025', '2025', '2026', '2026'],
        '__catalog_primary': ['unused'] * 8,
        '__catalog_series': ['unused'] * 8,
        'Campaign': ['2025', '2025', '2026', '2026', '2025', '2025', '2026', '2026'],
        'Metric': [1.0, 2.0, 1.5, 2.5, 1.2, 2.2, 1.7, 2.7],
    })

    with patch('src.modules.cdr_reporting._draw_chart_legend') as draw_legend:
        _render_cdf_line('CDF', frame, '__catalog_primary', '__catalog_series', 'Metric')

    legend_items = draw_legend.call_args.args[1]
    assert [item[0] for item in legend_items] == [
        'Vodafone · 2025', 'Vodafone · 2026', 'O2 · 2025', 'O2 · 2026',
    ]


def test_cdf_uses_emphasised_lines_when_only_one_campaign_is_rendered() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['Vodafone', 'Vodafone', 'O2', 'O2'],
        '__catalog_primary': ['unused'] * 4, '__catalog_series': ['unused'] * 4,
        'Campaign': ['2026 Q2'] * 4, 'Metric': [1.0, 2.0, 1.2, 2.2],
    })

    with patch('src.modules.cdr_reporting._draw_chart_legend') as draw_legend:
        _render_cdf_line('CDF', frame, '__catalog_primary', '__catalog_series', 'Metric')

    assert {item[2] for item in draw_legend.call_args.args[1]} == {4}


def test_cdf_renderer_uses_progressive_widths_for_four_campaigns() -> None:
    campaigns = ['2026-Q1', '2026-Q2', '2026-Q3', '2026-Q4']
    frame = chart_frame({
        '__catalog_row_0': ['Vodafone'] * 8,
        '__catalog_column_0': [campaign for campaign in campaigns for _ in range(2)],
        '__catalog_primary': ['unused'] * 8,
        '__catalog_series': ['unused'] * 8,
        'Campaign': [campaign for campaign in campaigns for _ in range(2)],
        'Metric': [1.0, 2.0, 1.1, 2.1, 1.2, 2.2, 1.3, 2.3],
    })
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': ('Operator',),
        '__catalog_column_0': ('Campaign',),
    }

    with patch('src.modules.cdr_reporting._draw_chart_legend') as draw_legend:
        _render_cdf_line('CDF', frame, '__catalog_primary', '__catalog_series', 'Metric')

    assert [item[2] for item in draw_legend.call_args.args[1]] == [1, 2, 3, 4]


def test_cdf_trims_only_after_80_percent_of_curves_exceed_99_percent() -> None:
    # With four curves, ceil(4 * 80%) requires all four curves.  Three
    # converged curves must retain the fourth curve's meaningful tail.
    assert _cdf_terminal_x_maximum(
        [[1.0] * 100 + [10.0], [2.0] * 100 + [10.0], [3.0] * 100 + [10.0], [4.0] * 70 + [30.0] * 31],
        1.0, 30.0, minimum_separation=0.08,
    ) == 30.0
    # With five curves, ceil(5 * 80%) requires four curves and permits the
    # converged tail to be removed once four are above 99%.
    assert _cdf_terminal_x_maximum(
        [[1.0] * 100 + [10.0], [2.0] * 100 + [10.0], [3.0] * 100 + [10.0], [4.0] * 100 + [10.0], [5.0] * 70 + [30.0] * 31],
        1.0, 30.0, minimum_separation=0.08,
    ) == 4.0
    # Exactly 99% does not satisfy the strictly-above-99% threshold.
    assert _cdf_terminal_x_maximum(
        [[1.0] * 99 + [10.0], [2.0] * 99 + [10.0], [3.0] * 99 + [10.0]],
        1.0, 10.0, minimum_separation=0.08,
    ) == 10.0


def test_outcome_series_colours_use_green_for_success_and_red_for_failure() -> None:
    frame = pd.DataFrame()
    keys = [('Success',), ('Failure',), ('Completed',), ('Dropped',)]

    colours = _series_colours(keys, ['__catalog_column_0'], frame)

    assert colours == {
        ('Success',): '#2C9A62',
        ('Failure',): '#E15759',
        ('Completed',): '#197A4A',
        ('Dropped',): '#F28E2B',
    }


def test_status_categories_colour_failure_and_success_outcomes_semantically() -> None:
    data = pd.DataFrame({'Test_Result': ['Failure', 'Success', 'Dropped']})

    _result, states, colours = _status_chart_categories(
        data, 'Test_Result', quality=False, threshold=1.6,
    )

    palette = dict(zip(states, colours, strict=True))
    assert palette['Success'] == '#2C9A62'
    assert palette['Dropped'] == '#F28E2B'
    assert palette['Failed'] == '#E15759'


def test_operator_vendor_column_groups_keep_campaign_bars_in_their_operator_palette() -> None:
    frame = chart_frame({})
    colours = _hierarchy_group_colours([
        ('Vodafone_Ericsson', '2026 Q1'), ('Vodafone_Ericsson', '2026 Q2'),
        ('Vodafone_Huawei', '2026 Q1'), ('Vodafone_Huawei', '2026 Q2'),
        ('3_Ericsson', '2026 Q1'), ('3_Ericsson', '2026 Q2'),
    ], frame=frame)

    assert colours['Vodafone_Ericsson'] == '#E15759'
    assert colours['Vodafone_Huawei'] == '#8C3637'
    assert colours['3_Ericsson'] == '#F28E2B'


def test_exact_operator_mapping_wins_over_underscore_prefix_fallback() -> None:
    frame = chart_frame({})
    frame.attrs['operator_mapping_groups'].append({
        'canonical': 'VF_SA', 'aliases': ['VF SA UK'], 'position': 4, 'color': '#8000FF',
    })
    frame.attrs['catalogue_dimension_labels'] = {'__catalog_row_0': ('Operator',)}

    colours = _series_colours(
        [('VF',), ('VF_SA',)], ['__catalog_row_0'], frame, line_chart=True,
    )

    assert colours == {('VF',): '#E15759', ('VF_SA',): '#8000FF'}


def test_vendor_is_the_complete_suffix_after_the_longest_configured_operator() -> None:
    frame = chart_frame({})
    frame.attrs['operator_mapping_groups'].extend([
        {'canonical': 'VF_SA', 'aliases': [], 'position': 4, 'color': '#8000FF'},
        {'canonical': 'VF_NSA', 'aliases': [], 'position': 5, 'color': '#0080FF'},
    ])
    frame.attrs['vendor_mapping_groups'].append({
        'canonical': 'Open_RAN', 'aliases': [], 'position': 7, 'color': '#12AB34',
    })
    frame.attrs['catalogue_dimension_labels'] = {'__catalog_row_0': ('Vendor',)}

    colours = _series_colours(
        [('VF_SA_Open_RAN',), ('VF_NSA_Open_RAN',)], ['__catalog_row_0'], frame,
    )

    assert colours == {
        ('VF_SA_Open_RAN',): '#12AB34',
        ('VF_NSA_Open_RAN',): '#12AB34',
    }


def test_neutral_operator_colour_matches_between_campaign_bars_and_legend() -> None:
    bars = _hierarchy_group_colours([
        ('O2', '2026 Q1'), ('O2', '2026 Q2'),
        ('Vodafone', '2026 Q1'), ('Vodafone', '2026 Q2'),
        ('3', '2026 Q1'), ('3', '2026 Q2'),
        ('EE', '2026 Q1'), ('EE', '2026 Q2'),
        ('Lebara', '2026 Q1'), ('Lebara', '2026 Q2'),
    ])
    legend = _hierarchy_group_colours([
        ('O2',), ('Vodafone',), ('3',), ('EE',), ('Lebara',),
    ])

    assert bars['Lebara'] == legend['Lebara']


def test_table_renders_percentages_for_a_categorical_metric() -> None:
    frame = chart_frame({
        'Test': ['Browse', 'Browse', 'Transfer'],
        'Result': ['Completed', 'Failed', 'Completed'],
    })

    rendered = _render_table('Result ratio', frame, 'Test', None, 'Result')

    assert rendered.getbuffer().nbytes > 0


def test_chart_colours_use_vendor_families_for_multi_operator_dimensions() -> None:
    keys = [('Vodafone', 'Ericsson'), ('Vodafone', 'Huawei'), ('3', 'Ericsson'), ('3', 'Huawei')]
    frame = chart_frame({'__catalog_row_0': [], '__catalog_row_1': []})
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': ('Operator',), '__catalog_row_1': ('Vendor',),
    }

    colours = _series_colours(keys, ['__catalog_row_0', '__catalog_row_1'], frame)

    assert colours[('Vodafone', 'Ericsson')] == '#2E8B57'
    assert colours[('3', 'Ericsson')] == '#2E8B57'
    assert colours[('Vodafone', 'Huawei')] == '#E15759'
    assert colours[('3', 'Huawei')] == '#E15759'

    line_colours = _series_colours(keys, ['__catalog_row_0', '__catalog_row_1'], frame, line_chart=True)
    assert line_colours == colours

    line_dashes = _series_line_dashes(keys, ['__catalog_row_0', '__catalog_row_1'], frame)
    assert line_dashes[('Vodafone', 'Ericsson')] == ()
    assert line_dashes[('Vodafone', 'Huawei')] == ()
    assert line_dashes[('3', 'Ericsson')] == (18, 8)
    assert line_dashes[('3', 'Huawei')] == (18, 8)


def test_chart_colours_use_vendor_families_for_one_operator() -> None:
    keys = [('Vodafone', 'Ericsson'), ('Vodafone', 'Huawei'), ('Vodafone', 'Samsung'), ('Vodafone', 'NSN')]
    frame = chart_frame({'__catalog_row_0': [], '__catalog_row_1': []})
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': ('Operator',), '__catalog_row_1': ('Vendor',),
    }

    colours = _series_colours(keys, ['__catalog_row_0', '__catalog_row_1'], frame)

    assert colours == {
        ('Vodafone', 'Ericsson'): '#2E8B57', ('Vodafone', 'Huawei'): '#E15759',
        ('Vodafone', 'Samsung'): '#7B3FB5', ('Vodafone', 'NSN'): '#4E79A7',
    }


def test_multivendor_chart_uses_operator_colour_when_vendor_is_missing() -> None:
    keys = [('Vodafone', '(blank)'), ('3', '(blank)'), ('EE', 'EE')]
    frame = chart_frame({'__catalog_row_0': [], '__catalog_row_1': []})
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': ('Operator',), '__catalog_row_1': ('Vendor',),
    }

    colours = _series_colours(keys, ['__catalog_row_0', '__catalog_row_1'], frame)

    assert colours == {
        ('Vodafone', '(blank)'): '#E15759',
        ('3', '(blank)'): '#F28E2B',
        ('EE', 'EE'): '#76B7B2',
    }


def test_cdf_uses_solid_first_operator_and_distinct_later_operators_per_vendor() -> None:
    keys = [
        ('Ericsson', '3'), ('Huawei', '3'), ('Samsung', '3'),
        ('Ericsson', 'VF'), ('Huawei', 'VF'), ('Samsung', 'VF'),
        ('Ericsson', 'EE'), ('Huawei', 'EE'),
        ('NSN', 'VF'), ('EE', 'EE'), ('O2', 'O2'), ('VF_SA', 'VF_SA'),
    ]
    frame = chart_frame({'__catalog_row_0': [], '__catalog_row_1': []})
    frame.attrs['operator_mapping_groups'].append({
        'canonical': 'VF_SA', 'aliases': [], 'position': 4, 'color': '#8000FF',
    })
    frame.attrs['operator_mapping_groups'] = [
        {**group, 'position': index}
        for index, group in enumerate(sorted(
            frame.attrs['operator_mapping_groups'],
            key=lambda group: {'3': 0, 'VF': 1, 'EE': 2, 'O2': 3, 'VF_SA': 4}.get(str(group['canonical']), 9),
        ))
    ]
    frame.attrs['catalogue_dimension_labels'] = {
        '__catalog_row_0': ('Vendor',), '__catalog_row_1': ('Operator',),
    }

    dashes = _series_line_dashes(keys, ['__catalog_row_0', '__catalog_row_1'], frame)

    assert dashes[('Ericsson', '3')] == ()
    assert dashes[('Huawei', '3')] == ()
    assert dashes[('Samsung', '3')] == ()
    assert dashes[('Ericsson', 'VF')] == (18, 8)
    assert dashes[('Huawei', 'VF')] == (18, 8)
    assert dashes[('Samsung', 'VF')] == (18, 8)
    assert dashes[('Ericsson', 'EE')] == (2, 10)
    assert dashes[('Huawei', 'EE')] == (2, 10)
    assert dashes[('NSN', 'VF')] == ()
    assert dashes[('EE', 'EE')] == ()
    assert dashes[('O2', 'O2')] == ()
    assert dashes[('VF_SA', 'VF_SA')] == ()


def test_chart_colours_detect_a_single_operator_from_composite_vendor_values() -> None:
    keys = [('3_Ericsson',), ('3_Huawei',), ('3_Nokia',)]
    frame = chart_frame({'__catalog_column_0': []})
    frame.attrs['catalogue_dimension_labels'] = {'__catalog_column_0': ('Vendor',)}

    colours = _series_colours(keys, ['__catalog_column_0'], frame)

    assert colours[('3_Ericsson',)] == '#2E8B57'
    assert len(set(colours.values())) == 3


def test_primary_vendor_dimension_assigns_semantic_colours_to_special_vendors() -> None:
    keys = [
        ('(blank)',), ('Vodafone_Ericsson',), ('Vodafone_Mixed Vendor',), ('Vodafone_Huawei',),
        ('Vodafone_Other Vendor',), ('Vodafone_NSN',), ('3_Ericsson',), ('3_Mixed Vendor',),
        ('3_Huawei',), ('3_Samsung',),
    ]
    frame = chart_frame({'__catalog_column_0': []})
    frame.attrs['catalogue_dimension_labels'] = {'__catalog_column_0': ('Vendor',)}

    colours = _series_colours(keys, ['__catalog_column_0'], frame)

    assert colours == {
        ('(blank)',): '#7A8791', ('Vodafone_Ericsson',): '#2E8B57',
        ('Vodafone_Mixed Vendor',): '#D9A514', ('Vodafone_Huawei',): '#E15759',
        ('Vodafone_Other Vendor',): '#D9A514', ('Vodafone_NSN',): '#4E79A7',
        ('3_Ericsson',): '#2E8B57', ('3_Mixed Vendor',): '#D9A514',
        ('3_Huawei',): '#E15759', ('3_Samsung',): '#7B3FB5',
    }


def test_horizontal_line_legend_reduces_columns_for_long_cdf_labels() -> None:
    captions = [
        'VF_UK · Ericsson · 2026-Q2', 'VF_UK · Huawei · 2026-Q2',
        'VF_UK · Samsung · 2026-Q2', 'VF_SA · Ericsson · 2026-Q2',
        'VF_SA · Huawei · 2026-Q2', 'VF_SA · Samsung · 2026-Q2',
    ]

    assert _horizontal_legend_columns(captions, 11, line_markers=True) < 6


def test_reporting_query_columns_splits_map_coordinates() -> None:
    import src.DriveTestAnalyzer as app_module

    entry = CatalogEntry(
        9, 'Map', '', '', 'DL Thput vs RF', 'CDR-Data',
        'Test_Start_Latitude vs Test_Start_Longitude', 'Map', '', 'Test Name IN (FDTT http DL MT)', 'Test_Result', '', 'Right',
    )

    columns = app_module.reporting_query_columns('data', [entry], False)

    assert 'Test_Start_Latitude' in columns
    assert 'Test_Start_Longitude' in columns
    assert 'Test_Start_Latitude vs Test_Start_Longitude' not in columns


def test_map_renderer_keeps_a_colour_key_for_every_filtered_point() -> None:
    frame = chart_frame({
        'Latitude': [51.5, 51.51, 51.52],
        'Longitude': [-0.12, -0.11, -0.10],
        'Operator': ['VF', 'EE', 'VF'],
        'Campaign': ['2026 Q2', '2026 Q2', '2026 Q1'],
    })

    image = _render_map('Coverage', frame, 'Operator', 'Campaign', 'Latitude', 'Longitude')

    assert image.getvalue().startswith(b'\x89PNG\r\n\x1a\n')


def test_reporting_query_columns_splits_multi_kpi_cdf_metrics() -> None:
    import src.DriveTestAnalyzer as app_module

    entry = CatalogEntry(
        43, 'FDTT DL SINR', '', '', 'FDTT DL SINR', 'CDR-Data',
        'NR_PCell_SINR_Avg | LTE_PCell_SINR_Avg', 'Multi KPI CDF Lines',
        'Operator', 'Test_Result IN (Completed)', 'Operator', '', 'Right',
    )

    columns = app_module.reporting_query_columns('data', [entry], False)

    assert {'NR_PCell_SINR_Avg', 'LTE_PCell_SINR_Avg'} <= set(columns)
    assert 'NR_PCell_SINR_Avg | LTE_PCell_SINR_Avg' not in columns


def test_reporting_query_columns_include_calculated_dimension_dependencies() -> None:
    import src.DriveTestAnalyzer as app_module

    entry = CatalogEntry(
        22, 'LTE PCell ARFCN', '', '', 'LTE PCell ARFCN', 'CDR-Data',
        'First LTE PCC ARFCN', '100% Stacked Vertical Bars',
        'First LTE PCC ARFCN', 'Tput Above IN (above100, above20)', 'Operator', '', 'Right',
    )

    columns = app_module.reporting_query_columns('data', [entry], False)

    assert {'LTE_PCC_EARFCN', 'Mean_Data_Rate', 'Test_Name'} <= set(columns)


def test_catalogue_filter_contract_supports_not_in_and_not_contains() -> None:
    conditions = parse_catalog_filters('Session_Type NOT IN (WhatsApp, SMS); Vendor NOT CONTAINS (Mixed, Other); Campaign NOT CONTAINS legacy')
    assert [(item.column, item.operator, item.values) for item in conditions] == [
        ('Session_Type', 'NOT IN', ('WhatsApp', 'SMS')),
        ('Vendor', 'NOT CONTAINS', ('Mixed', 'Other')),
        ('Campaign', 'NOT CONTAINS', ('legacy',)),
    ]
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS) + '\n8,Quality,,Title and 1 column + Comments,Quality,CDR-Speech,LQ,Average Vertical Bars,Session_Type NOT IN (WhatsApp); Campaign NOT CONTAINS legacy,Operator,Campaign,,\n',
    )[0]
    frame = chart_frame({
        'Session_Type': ['VoLTE', 'WhatsApp', 'VoLTE'], 'Campaign': ['Q1', 'Q1', 'legacy-Q2'],
        'LQ': [3.2, 4.0, 3.8], 'Operator': ['EE', 'EE', 'O2'],
    })
    assert _apply_catalog_filters(frame, entry, False, 'LQ')['Operator'].tolist() == ['EE']


def test_campaign_filters_accept_compact_permutations_without_merging_sa_and_nsa() -> None:
    frame = chart_frame({
        'Campaign': ['UK_Q2_2026', 'UK_Q2_2026_SA', '2026_Q2_NSA_UK'],
        'Metric': [1, 2, 3],
    })
    plain = replace(CatalogEntry(1, '', '', '', '', 'CDR-Data', 'Metric', 'Average Vertical Bars', '', '', '', '', '', ''), filters='Campaign = 2026-Q2')
    sa = replace(plain, filters='Campaign = 2026-Q2_SA')
    modes = replace(plain, filters='Campaign IN (2026-Q2_SA, UK_Q2_2026_NSA)')

    assert _apply_catalog_filters(frame, plain, False, 'Metric')['Metric'].tolist() == [1]
    assert _apply_catalog_filters(frame, sa, False, 'Metric')['Metric'].tolist() == [2]
    assert _apply_catalog_filters(frame, modes, False, 'Metric')['Metric'].tolist() == [2, 3]


def test_tableau_result_group_filter_uses_the_workbook_bins() -> None:
    dimensions = parse_calculated_dimensions([{
        'name': 'Result Group', 'sources': ['cdr-data'],
        'rules': [
            {'when': 'Test_Result IN (Completed, Visible Completed)', 'value': 'Success'},
            {'when': 'Test_Result IN (Cutoff, Failed)', 'value': 'Failure'},
        ],
    }])
    entry = CatalogEntry(
        1, 'Failures', '', '', 'Failures', 'CDR-Data', 'Result Group',
        '100% Stacked Vertical Bars', 'Result Group', 'Result Group NOT IN (Success)',
        'Operator', '', 'Right', calculated_dimensions=dimensions,
    )
    frame = chart_frame({
        'Operator': ['EE'] * 5,
        'Test_Result': ['Completed', 'Visible Completed', 'Cutoff', 'Failed', 'Unknown'],
    })

    filtered = _apply_catalog_filters(frame, entry, False, None)

    assert filtered['Test_Result'].tolist() == ['Cutoff', 'Failed', 'Unknown']


def test_multi_kpi_cdf_lines_render_each_tableau_measure() -> None:
    entry = CatalogEntry(
        1, 'Radio quality', '', '', 'Radio quality', 'CDR-Data',
        'NR SINR | LTE SINR', 'Multi KPI CDF Lines', 'Operator', '',
        'Operator', 'Campaign', 'Right',
    )
    frame = chart_frame({
        'Operator': ['EE', 'EE', 'VF', 'VF'],
        'Campaign': ['2026 Q1', '2026 Q2', '2026 Q1', '2026 Q2'],
        'NR SINR': [4.0, 5.0, 6.0, 7.0],
        'LTE SINR': [8.0, 9.0, 10.0, 11.0],
    })

    image = render_catalog_chart_preview(frame, entry)
    targets = catalog_chart_hover_targets(frame, entry)

    assert image.startswith(b'\x89PNG\r\n\x1a\n')
    assert {target['label'] for target in targets} == {'NR SINR', 'LTE SINR'}
    assert any(target['x'] < 750 for target in targets)
    assert any(target['x'] > 750 for target in targets)


def test_design_canvas_report_renderer_uses_design_payload() -> None:
    entry = CatalogEntry(
        1, 'Radio quality', '', '', 'Radio quality', 'CDR-Data',
        'NR SINR', 'CDF Line', 'Operator', '', 'Operator', 'Campaign', 'Right',
    )
    frame = chart_frame({
        'Operator': ['EE', 'VF'],
        'Campaign': ['2026 Q1', '2026 Q1'],
        'NR SINR': [4.0, 8.0],
    })

    with patch('src.modules.cdr_reporting._render_canvas_payload_png', return_value=b'canvas-png') as render:
        image = render_catalog_chart_preview(frame, entry, renderer='canvas')

    assert image == b'canvas-png'
    payload = render.call_args.args[0]
    assert payload['renderer'] == 'catalog-v2'
    assert payload['type'] == 'cdf'
    assert payload['title'] == 'Radio quality'


def test_canvas_chart_renderer_uses_configured_node_path(monkeypatch, tmp_path: Path) -> None:
    import src.modules.cdr_reporting as reporting

    node = tmp_path / 'node'
    node.write_text('', encoding='utf-8')
    node.chmod(0o755)
    monkeypatch.setenv('DRIVETEST_ANALYZER_NODE_PATH', str(node))

    assert reporting._node_executable() == str(node)


def test_report_renderer_rejects_unknown_engine() -> None:
    entry = CatalogEntry(
        1, 'Radio quality', '', '', 'Radio quality', 'CDR-Data',
        'NR SINR', 'CDF Line', '', '', 'Operator', 'Campaign', 'Right',
    )
    frame = chart_frame({'Operator': ['EE'], 'Campaign': ['2026 Q1'], 'NR SINR': [4.0]})

    with pytest.raises(ValueError, match="Expected 'pil' or 'canvas'"):
        render_catalog_chart_preview(frame, entry, renderer='unknown')


def test_report_renderer_defaults_to_canvas_and_keeps_pil_override(monkeypatch) -> None:
    import src.modules.cdr_reporting as reporting

    monkeypatch.delenv(reporting.REPORT_CHART_RENDERER_ENV, raising=False)
    assert reporting.report_chart_renderer_name() == 'canvas'
    monkeypatch.setenv(reporting.REPORT_CHART_RENDERER_ENV, 'pil')
    assert reporting.report_chart_renderer_name() == 'pil'


def test_canvas_hits_are_reused_by_static_chart_tooltips() -> None:
    import src.modules.cdr_reporting as reporting

    entry = CatalogEntry(
        1, 'Throughput', '', '', 'Throughput', 'CDR-Data',
        'Mean_Data_Rate', 'Average Vertical Bars', '', '', 'Operator', 'Campaign', 'Right',
    )
    frame = chart_frame({'Operator': ['EE'], 'Campaign': ['2026 Q1'], 'Mean_Data_Rate': [42.0]})
    hits = [{
        'kind': 'rectangle', 'x': 10, 'y': 20, 'width': 30, 'height': 40,
        'label': 'EE · 2026 Q1', 'series': 'EE', 'value': '42.00',
    }]

    with patch('src.modules.cdr_reporting._render_canvas_payload', return_value=(b'canvas-png', hits)):
        image, targets = reporting.render_catalog_chart_preview_with_hover(frame, entry)

    assert image == b'canvas-png'
    assert targets == [{
        'kind': 'bar', 'x': 10.0, 'y': 20.0, 'width': 30.0, 'height': 40.0,
        'label': 'EE · 2026 Q1', 'legend': 'EE', 'value': '42.00',
    }]


def test_not_contains_filter_excludes_each_comma_separated_term() -> None:
    entry = CatalogEntry(
        1, 'Quality', '', 'Title and 1 column', '', 'CDR-Speech', 'LQ', 'CDF Line',
        '', 'Vendor NOT CONTAINS (Mixed, Other)', 'Vendor', 'Campaign', 'Top',
    )
    frame = chart_frame({
        'Vendor': ['Vodafone_Ericsson', 'Vodafone_Mixed Vendor', '3_Other Vendor'],
        'Vendor_Only': ['Ericsson', 'Mixed Vendor', 'Other Vendor'],
        'Campaign': ['2026 Q1'] * 3,
        'LQ': [3.5, 3.6, 3.7],
    })

    assert _apply_catalog_filters(frame, entry, False, 'LQ')['Vendor'].tolist() == ['Vodafone_Ericsson']


def test_catalogue_call_family_uses_documented_netcheck_session_values() -> None:
    entry = with_default_calculated_dimensions(parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n8,Completed Call Ratio,,Title and 1 column + Comments,,CDR-Voice,Call_Status,100% Stacked Vertical Bars,"Call Family IN (VoLTE, MultiRAB, WhatsApp)",Call Family,Operator × Campaign,,\n',
    )[0])
    frame = chart_frame({
        'Session_Type': ['CALL', 'MultiRAB CALL', 'WhatsApp CALL'],
        'L1_Call_Mode_A': ['VoLTE', '', ''],
        'Operator': ['EE', 'EE', 'EE'],
        'Campaign': ['Q1', 'Q1', 'Q1'],
        'Call_Status': ['Completed', 'Completed', 'Completed'],
    })

    filtered = _apply_catalog_filters(frame, entry, False, 'Call_Status')
    grouped, primary, series = _apply_catalog_grouping(filtered, entry, False, 'Call_Status')

    assert grouped[primary].tolist() == ['MultiRAB', 'VoLTE', 'WhatsApp']
    assert grouped['__catalog_row_0'].tolist() == ['MultiRAB', 'VoLTE', 'WhatsApp']
    assert grouped['__catalog_column_0'].tolist() == ['EE', 'EE', 'EE']
    assert grouped['__catalog_column_1'].tolist() == ['Q1', 'Q1', 'Q1']

    with patch('src.modules.cdr_reporting._render_status_100_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'nested-chart')
        chart = _render_status_100('Completed Call Ratio', grouped, primary, series)

    assert chart.getvalue() == b'nested-chart'
    assert hierarchy_renderer.call_args.args[2] == ['__catalog_row_0']
    assert hierarchy_renderer.call_args.args[3] == ['__catalog_column_0', '__catalog_column_1']


def test_status_chart_uses_nested_columns_without_a_row_grouping() -> None:
    entry = CatalogEntry(
        8, 'Completed Call Ratio', '', 'Title and 1 column + Comments', '', 'CDR-Voice',
        'Call_Status', '100% Stacked Vertical Bars', '', '', '', 'Operator × Campaign',
    )
    frame = chart_frame({
        'Operator': ['Vodafone', 'Vodafone', 'O2', 'O2'],
        'Campaign': ['2025 Q4', '2026 Q1', '2025 Q4', '2026 Q1'],
        'Call_Status': ['Completed', 'Failed', 'Completed', 'Dropped'],
    })
    grouped, primary, series = _apply_catalog_grouping(frame, entry, False, 'Call_Status')

    with patch('src.modules.cdr_reporting._render_status_100_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'nested-columns')
        chart = _render_status_100('Completed Call Ratio', grouped, primary, series)

    assert chart.getvalue() == b'nested-columns'
    assert hierarchy_renderer.call_args.args[2] == []
    assert hierarchy_renderer.call_args.args[3] == ['__catalog_column_0', '__catalog_column_1']


def test_status_chart_keeps_row_only_hierarchy_on_the_left() -> None:
    entry = CatalogEntry(
        8, 'Status', '', '', '', 'CDR-Data', 'Test_Result',
        '100% Stacked Vertical Bars', '', '', 'City × G Level 1', '',
    )
    frame = chart_frame({
        'City': ['London', 'London'],
        'G_Level_1': ['Drive', 'Connecting Roads'],
        'Test_Result': ['Completed', 'Failed'],
    })
    grouped, primary, series = _apply_catalog_grouping(frame, entry, False, 'Test_Result')

    with patch('src.modules.cdr_reporting._render_status_100_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'row-hierarchy')
        chart = _render_status_100('Status', grouped, primary, series)

    assert chart.getvalue() == b'row-hierarchy'
    assert hierarchy_renderer.call_args.args[2] == ['__catalog_row_0', '__catalog_row_1']
    assert hierarchy_renderer.call_args.args[3] == ['__catalog_single_column']


def test_status_chart_leaves_status_row_inclusion_to_the_template_filter() -> None:
    frame = chart_frame({
        'Operator': ['EE'] * 8,
        'Campaign': ['2026-Q2'] * 8,
        'Test_Result': ['Completed', 'Failed', None, float('nan'), '', ' NaN ', 'Not executed', 'Unknown'],
        '__catalog_column_0': ['EE'] * 8,
    })

    with patch('src.modules.cdr_reporting._render_status_100_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'filtered-status-chart')
        chart = _render_status_100('Data failures', frame, 'Operator', 'Campaign')

    rendered_data = hierarchy_renderer.call_args.args[1]
    assert chart.getvalue() == b'filtered-status-chart'
    assert len(rendered_data.index) == 8
    assert rendered_data['state'].iloc[:2].tolist() == ['Completed', 'Failed']
    assert rendered_data['state'].iloc[2:].isna().all()


def test_status_chart_honours_selected_kpi_when_other_status_columns_exist() -> None:
    frame = chart_frame({
        'Operator': ['EE', 'EE'],
        'Campaign': ['2026-Q2', '2026-Q2'],
        'Call_Status': ['Completed', 'Completed'],
        'Test_Result': ['Completed', 'Failed'],
        '__catalog_column_0': ['EE', 'EE'],
    })

    with patch('src.modules.cdr_reporting._render_status_100_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'selected-status-kpi')
        chart = _render_status_100('Data failures', frame, 'Operator', 'Campaign', metric='Test_Result')

    rendered_data = hierarchy_renderer.call_args.args[1]
    assert chart.getvalue() == b'selected-status-kpi'
    assert rendered_data['state'].tolist() == ['Completed', 'Failed']


def test_status_chart_maps_cutoff_to_a_visible_failure_segment_without_filtering_rows() -> None:
    frame = chart_frame({
        'Operator': ['EE'] * 8,
        'Campaign': ['2026-Q2'] * 8,
        'Test_Result': ['Incomplete', 'Unsuccessful', 'Aborted', 'Cancelled', 'Timeout', 'Cutoff', 'Failed', 'Completed'],
        '__catalog_column_0': ['EE'] * 8,
    })

    with patch('src.modules.cdr_reporting._render_status_100_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'negative-statuses')
        _render_status_100('Data failures', frame, 'Operator', 'Campaign', metric='Test_Result')

    states = hierarchy_renderer.call_args.args[1]['state']
    assert states.tolist() == ['Incomplete', 'Unsuccessful', 'Aborted', 'Cancelled', 'Timeout', 'Cutoff', 'Failed', 'Completed']
    assert hierarchy_renderer.call_args.args[4] == (
        'Completed', 'Cutoff', 'Failed', 'Aborted', 'Cancelled', 'Incomplete', 'Timeout', 'Unsuccessful',
    )
    assert hierarchy_renderer.call_args.args[5] == (
        '#2C9A62', '#C83E4D', '#E15759', '#D8555F', '#E26A70', '#AE2F42', '#F08A8F', '#8F2035',
    )


def test_data_cutoffs_are_excluded_from_tableau_status_denominator() -> None:
    frame = chart_frame({
        'Operator': ['EE'] * 1000,
        'Campaign': ['2026-Q2'] * 1000,
        'Test_Result': ['Completed'] * 692 + ['Cutoff'] * 278 + ['Failed'] * 30,
        '__catalog_column_0': ['EE'] * 1000,
    })

    with patch('src.modules.cdr_reporting._render_status_100_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'tableau-status-ratio')
        _render_status_100('Data failures', frame, 'Operator', 'Campaign', metric='Test_Result')

    states = hierarchy_renderer.call_args.args[1]['state']
    assert states.eq('Completed').sum() == 692
    assert states.eq('Failed').sum() == 30


def test_hierarchical_grouping_keeps_campaign_bars_together_per_operator() -> None:
    frame = chart_frame({
        '__catalog_column_0': ['Vodafone', 'O2', 'Vodafone', 'O2'],
        '__catalog_column_1': ['2025 Q4', '2025 Q4', '2026 Q1', '2026 Q1'],
    })

    keys = _hierarchical_unique_keys(frame, ['__catalog_column_0', '__catalog_column_1'])

    assert keys == [
        ('Vodafone', '2025 Q4'),
        ('Vodafone', '2026 Q1'),
        ('O2', '2025 Q4'),
        ('O2', '2026 Q1'),
    ]


def test_hierarchy_spans_nest_each_level_without_compound_captions() -> None:
    keys = [
        ('City', 'Drive', 'EE'), ('City', 'Drive', '3'),
        ('City', 'Connecting Roads', 'EE'), ('Rural', 'Drive', 'EE'),
    ]

    assert _hierarchy_caption_spans(keys, 0) == [(0, 3, 'City'), (3, 4, 'Rural')]
    assert _hierarchy_caption_spans(keys, 1) == [
        (0, 2, 'Drive'), (2, 3, 'Connecting Roads'), (3, 4, 'Drive'),
    ]
    assert _hierarchy_caption_spans(keys, 2) == [
        (0, 1, 'EE'), (1, 2, '3'), (2, 3, 'EE'), (3, 4, 'EE'),
    ]


def test_campaign_aggregation_is_ordered_oldest_to_newest_for_every_chart_renderer() -> None:
    entry = CatalogEntry(
        5, 'Data failures', '', '', '', 'CDR-Data', 'Test_Result',
        '100% Stacked Vertical Bars', '', '', 'Test_Name', 'Operator × Campaign',
    )
    frame = chart_frame({
        'Test_Name': ['FDFS', 'FDFS', 'FDFS', 'FDFS'],
        'Operator': ['EE', 'VF', 'EE', 'VF'],
        'Campaign': ['UK_Q2_2026', 'UK_Q2_2026', 'UK_Q1_2026', 'UK_Q1_2026'],
        'Test_Result': ['Completed'] * 4,
    })

    grouped, _, _ = _apply_catalog_grouping(frame, entry, False, 'Test_Result')
    keys = _hierarchical_unique_keys(grouped, ['__catalog_column_0', '__catalog_column_1'])

    assert keys == [
        ('VF', '26-Q1'), ('VF', '26-Q2'),
        ('EE', '26-Q1'), ('EE', '26-Q2'),
    ]


def test_nsa_speech_catalogue_filters_produce_samples_and_use_latest_campaign() -> None:
    entries = [with_default_calculated_dimensions(entry) for entry in load_catalog_csv(Path(__file__).parent / 'fixtures' / 'NSA Slide Template.csv')]
    speech = pd.DataFrame({
        'sample': ['volte', 'multirab', 'whatsapp-old', 'whatsapp-latest', 'whatsapp-sa', 'o2-latest'],
        'Session_Type': ['CALL', 'MultiRAB CALL', 'WhatsApp CALL', 'WhatsApp CALL', 'WhatsApp CALL', 'WhatsApp CALL'],
        'L1_Call_Mode_A': ['VoLTE', 'VoLTE', 'VoIP', 'VoIP', 'VoIP', 'VoIP'],
        'Sample_RAT_A': [None, None, 'EN-DC', 'EN-DC', 'NR SA', 'EN-DC'],
        'Call_Status': ['Completed'] * 6,
        'Operator': ['Vodafone UK', '3', 'EE', 'EE', 'EE', 'O2 (UK)'],
        'Campaign': ['UK_Q3_2025', 'UK_Q3_2025', 'UK_Q3_2025', 'UK_Q4_2025', 'UK_Q4_2025', 'UK_Q4_2025'],
        'LQ': [3.8, 3.7, 3.9, 4.0, 4.1, 4.2],
    })
    # The NSA CDR: its rows without the one of an SA CDR.
    nsa = speech.loc[speech['sample'] != 'whatsapp-sa']
    filtered_by_entry = {
        (entry.slide, entry.chart_type, index): _apply_catalog_filters(nsa, entry, False, 'LQ')
        for index, entry in enumerate(entries)
        if entry.slide in {7, 8, 9}
    }

    assert all(not frame.empty for frame in filtered_by_entry.values())
    latest_whatsapp = [frame for (slide, _chart, _index), frame in filtered_by_entry.items() if slide == 8][2]
    assert latest_whatsapp['sample'].tolist() == ['whatsapp-latest']


def test_layout_chart_frames_are_always_ordered_by_visual_rows_then_columns() -> None:
    presentation = Presentation('assets/powerpoint-templates/Template_VodafoneThree.pptx')
    layout = _named_slide_layout(presentation, 'Title and 2 columns and 2 rows + Comments right')

    frames = _layout_chart_frames(layout)

    assert len(frames) == 4
    assert frames[0][0] < frames[1][0]
    assert frames[0][1] < frames[2][1]
    assert frames[2][0] < frames[3][0]


def test_failure_count_uses_row_and_column_hierarchies_without_flattening() -> None:
    entry = with_default_calculated_dimensions(parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n9,Voice failures per Q/city,,Title and 1 column + Comments,Failures,CDR-Voice,Call_Status,Count Stacked Horizontal Bars,,Call Family × G Level 4,Operator × Campaign,Failed/Dropped,\n',
    )[0])
    frame = chart_frame({
        'Session_Type': ['VoLTE', 'VoLTE', 'MultiRAB CALL'],
        'G_Level_4': ['London', 'London', 'Belfast'],
        'Operator': ['EE', 'EE', '3'],
        'Campaign': ['Q2', 'Q2', 'Q3'],
        'Call_Status': ['Failed', 'Dropped', 'Failed'],
    })
    grouped, primary, series = _apply_catalog_grouping(frame, entry, False, 'Call_Status')

    with patch('src.modules.cdr_reporting._render_failure_count_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'nested-failure-chart')
        chart = _render_failure_count('Voice failures per Q/city', grouped, primary, series)

    assert chart.getvalue() == b'nested-failure-chart'
    assert hierarchy_renderer.call_args.args[2] == ['__catalog_row_0', '__catalog_row_1']
    assert hierarchy_renderer.call_args.args[3] == ['__catalog_column_0', '__catalog_column_1']


def test_failure_count_hover_targets_cover_rendered_horizontal_segments() -> None:
    entry = with_default_calculated_dimensions(parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n9,Voice failures per Q/city,,Title and 1 column + Comments,Failures,CDR-Voice,Call_Status,Count Stacked Horizontal Bars,,Call Family × G Level 4,Operator × Campaign,Failed/Dropped,Right\n',
    )[0])
    frame = chart_frame({
        'Session_Type': ['VoLTE', 'VoLTE', 'VoLTE'], 'G_Level_4': ['London', 'London', 'Belfast'],
        'Operator': ['EE', 'EE', '3'], 'Campaign': ['Q2', 'Q2', 'Q3'],
        'Call_Status': ['Failed', 'Dropped', 'Completed'],
    })

    targets = catalog_chart_hover_targets(frame, entry)

    assert {(target['legend'], target['value']) for target in targets} == {('Failed', '1'), ('Dropped', '1')}
    assert all(target['width'] > 0 and target['height'] > 0 for target in targets)


def test_failure_count_hover_targets_use_the_renderer_width_for_field_legends() -> None:
    entry = with_default_calculated_dimensions(parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n9,Voice failures,,Title and 1 column + Comments,Failures,CDR-Voice,Call_Status,Count Stacked Horizontal Bars,,Call Family,Operator × Campaign,Call_Status,Right\n',
    )[0])
    frame = chart_frame({
        'Session_Type': ['VoLTE', 'VoLTE'], 'Operator': ['EE', '3'],
        'Campaign': ['Q2', 'Q2'], 'Call_Status': ['Failed', 'Failed'],
    })

    targets = catalog_chart_hover_targets(frame, entry)

    # A field legend is resolved after rendering, so the horizontal plot keeps
    # the full 1250-pixel renderer width instead of reserving a right lane.
    assert sorted(target['x'] for target in targets) == [289.0, 914.0]


def test_status_100_hover_targets_support_multi_level_row_grouping() -> None:
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n12,Success Ratio per Type of Test (Vendor Split),,Title and 1 column + Comments,Success Ratio per Type of Test (Vendor Split),CDR-Data,Result Group,100% Stacked Vertical Bars,,Type_of_Test × Vendor,,Result Group,Right\n',
    )[0]
    frame = chart_frame({
        'Type_of_Test': ['FTP', 'HTTP', 'FTP', 'HTTP'],
        'Vendor': ['EE', 'EE', '3', '3'],
        'Result Group': ['Completed', 'Failed', 'Completed', 'Cutoff'],
    })

    targets = catalog_chart_hover_targets(frame, entry)

    assert len(targets) == 12
    assert {target['label'] for target in targets} == {'FTP · EE', 'HTTP · EE', 'FTP · 3', 'HTTP · 3'}
    assert any(target['height'] > 0 for target in targets)


def test_status_100_hover_targets_support_tableau_design_geography_hierarchy() -> None:
    entry = parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n1,SR Design,,Title and 3 columns + Comments,Success Ratio per glevel,CDR-Data,Result Group,100% Stacked Vertical Bars,,G_Level_2 × G_Level_1 × Operator,,Result Group,Right\n',
    )[0]
    frame = chart_frame({
        'G_Level_2': ['England', 'England', 'Scotland', 'Scotland'],
        'G_Level_1': ['North', 'North', 'Central', 'Central'],
        'Operator': ['EE', '3', 'EE', '3'],
        'Result Group': ['Completed', 'Failed', 'Completed', 'Cutoff'],
    })

    targets = catalog_chart_hover_targets(frame, entry)

    assert len(targets) == 12
    assert {target['label'] for target in targets} == {
        'England · North · EE', 'England · North · 3',
        'Scotland · Central · EE', 'Scotland · Central · 3',
    }
    assert any(target['height'] > 0 for target in targets)


def test_cdf_hover_targets_use_the_same_clipped_domain_as_the_renderer() -> None:
    entry = CatalogEntry(
        13, 'POLQA CDF', '', '', '', 'CDR-Speech', 'LQ', 'CDF Line', 'Operator × Campaign', '', 'Operator', 'Campaign', 'Top',
    )
    # 100 / 101 samples exceeds 99%, so the common terminal value remains
    # inside the clipped renderer domain under the stricter CDF policy.
    values = list(range(1, 101)) + [1000]
    frame = chart_frame({
        'Session_Type': ['WhatsApp CALL'] * 303, 'Operator': ['VF'] * 101 + ['3'] * 101 + ['EE'] * 101,
        'Campaign': ['UK_Q2_2026'] * 303, 'LQ': values * 3,
    })

    targets = catalog_chart_hover_targets(frame, entry)
    low = 1.0
    high = _cdf_terminal_x_maximum([values, values, values], low, 1000.0)
    left, top, width, height = _cdf_plot_geometry('top')
    vf_last = [target for target in targets if target['legend'].startswith('VF')][-1]

    assert vf_last['x'] == pytest.approx(left + (100.0 - low) / (high - low) * width)
    assert vf_last['y'] == pytest.approx(top + height - 100 / 101 * height)
    assert all(target['x'] <= left + width for target in targets)


def test_cdf_hover_targets_are_bounded_per_series() -> None:
    entry = CatalogEntry(
        13, 'POLQA CDF', '', '', '', 'CDR-Speech', 'LQ', 'CDF Line', 'Operator', '', 'Operator', '', 'Top',
    )
    frame = chart_frame({
        'Session_Type': ['WhatsApp CALL'] * 2_000,
        'Operator': ['VF'] * 1_000 + ['EE'] * 1_000,
        'LQ': list(range(1_000)) * 2,
    })

    targets = catalog_chart_hover_targets(frame, entry)

    assert len([target for target in targets if target['legend'].startswith('VF')]) == 120
    assert len([target for target in targets if target['legend'].startswith('EE')]) == 120


def test_failure_count_keeps_zero_count_hierarchy_categories_from_all_filtered_rows() -> None:
    entry = with_default_calculated_dimensions(parse_catalog_csv(
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n9,Voice failures,,Title and 1 column + Comments,Failures,CDR-Voice,Call_Status,Count Stacked Horizontal Bars,,Call Family,Operator × Campaign,Failed/Dropped,\n',
    )[0])
    frame = chart_frame({
        'Session_Type': ['VoLTE', 'VoLTE'], 'Operator': ['Vodafone', '3'],
        'Campaign': ['Q2', 'Q2'], 'Call_Status': ['Completed', 'Completed'],
    })
    grouped, primary, series = _apply_catalog_grouping(frame, entry, False, 'Call_Status')

    with patch('src.modules.cdr_reporting._render_failure_count_hierarchy') as hierarchy_renderer:
        hierarchy_renderer.return_value = BytesIO(b'zero-count-grid')
        chart = _render_failure_count('Voice failures', grouped, primary, series)

    assert chart.getvalue() == b'zero-count-grid'
    assert hierarchy_renderer.call_args.kwargs['comparison_frame'] is grouped
    assert hierarchy_renderer.call_args.args[1].empty


def test_failure_count_keeps_explicit_in_filter_categories_with_zero_matching_rows() -> None:
    entry = CatalogEntry(
        4, 'Failures', '', '', '', 'CDR-Voice', 'Call_Status', 'Count Stacked Horizontal Bars', '',
        'Call_Status IN (Failed, Dropped); G Level 4 IN (Belfast, Bristol, Cardiff)',
        'Call Family × G Level 4', 'Operator × Campaign',
    )
    frame = chart_frame({
        'Call Family': ['VoLTE'], 'G_Level_4': ['Belfast'], 'Operator': ['EE'],
        'Campaign': ['2026 Q2'], 'Call_Status': ['Failed'],
    })

    filtered = _apply_catalog_filters(frame, entry, False, 'Call_Status')
    grouped, _primary, _series = _apply_catalog_grouping(filtered, entry, False, 'Call_Status')

    assert grouped.attrs['catalogue_dimension_values']['__catalog_row_1'] == ['Belfast', 'Bristol', 'Cardiff']
    assert _hierarchical_complete_keys(grouped, ['__catalog_row_0', '__catalog_row_1']) == [
        ('VoLTE', 'Belfast'), ('VoLTE', 'Bristol'), ('VoLTE', 'Cardiff'),
    ]


def test_failure_hierarchy_reserves_a_right_legend_lane() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['VoLTE'],
        '__catalog_column_0': ['Vodafone'],
        '__catalog_column_1': ['2026 Q2'],
        '__catalog_failure_state': ['Failed'],
    })

    with patch('src.modules.cdr_reporting._draw_chart_legend') as draw_legend:
        _render_failure_count_hierarchy(
            'Voice failures per Q/city', frame,
            ['__catalog_row_0'], ['__catalog_column_0', '__catalog_column_1'],
            legend_position='right',
        )

    # A short right-side legend receives only the lane it needs, allowing the
    # failure matrix to use the otherwise blank portion of the canvas.
    assert draw_legend.call_args.kwargs['side_x'] > 1289


def test_failure_hierarchy_keeps_complete_column_headers() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['VoLTE'],
        '__catalog_column_0': ['Very Long Operator Name'],
        '__catalog_column_1': ['2026-09-24 Campaign Full Name'],
        '__catalog_failure_state': ['Failed'],
    })

    with patch('src.modules.cdr_reporting._draw_fitted_aggregation_header', return_value=20) as draw_header:
        _render_failure_count_hierarchy(
            'Voice failures', frame,
            ['__catalog_row_0'], ['__catalog_column_0', '__catalog_column_1'],
        )

    assert [call.args[2] for call in draw_header.call_args_list] == [
        'Very Long Operator Name', '2026-09-24 Campaign Full Name',
    ]
    assert [call.kwargs['level'] for call in draw_header.call_args_list] == [0, 1]


def test_failure_hierarchy_header_fonts_decrease_at_each_column_level() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['VoLTE'],
        '__catalog_column_0': ['Ericsson'],
        '__catalog_column_1': ['VF_UK'],
        '__catalog_column_2': ['2026-Q2'],
        '__catalog_failure_state': ['Failed'],
    })
    sizes = []

    def record_header(*args, **kwargs):
        size = _draw_fitted_aggregation_header(*args, **kwargs)
        sizes.append((kwargs['level'], size))
        return size

    with patch('src.modules.cdr_reporting._draw_fitted_aggregation_header', side_effect=record_header):
        _render_failure_count_hierarchy(
            'Voice failures', frame,
            ['__catalog_row_0'],
            ['__catalog_column_0', '__catalog_column_1', '__catalog_column_2'],
        )

    assert [level for level, _size in sizes] == [0, 1, 2]
    assert sizes[0][1] > sizes[1][1] > sizes[2][1]


def test_failure_hierarchy_reserves_space_below_column_headers() -> None:
    frame = chart_frame({
        '__catalog_row_0': [f'Row {index:02d}' for index in range(21)],
        '__catalog_column_0': ['VF'] * 21,
        '__catalog_column_1': ['2026-Q2'] * 21,
        '__catalog_failure_state': ['Failed'] * 21,
    })

    with patch('src.modules.cdr_reporting._draw_fitted_aggregation_header', return_value=20) as draw_header:
        with patch('src.modules.cdr_reporting._draw_configured_bar_label') as draw_bar_label:
            _render_failure_count_hierarchy(
                'Voice failures', frame,
                ['__catalog_row_0'], ['__catalog_column_0', '__catalog_column_1'],
            )

    leaf_header = next(call for call in draw_header.call_args_list if call.args[2] == '2026-Q2')
    first_bar_y = min(call.kwargs['y'] for call in draw_bar_label.call_args_list)
    assert leaf_header.kwargs['y'] + 26 + 8 < first_bar_y


def test_failure_hierarchy_uses_dashed_child_boundaries_within_one_operator() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['VoLTE', 'VoLTE'],
        '__catalog_column_0': ['Vodafone', 'Vodafone'],
        '__catalog_column_1': ['2026 Q2', '2026 Q1'],
        '__catalog_failure_state': ['Failed', 'Dropped'],
    })

    with patch('src.modules.cdr_reporting._draw_dashed_vertical_line') as draw_dashed:
        _render_failure_count_hierarchy(
            'Voice failures per campaign', frame,
            ['__catalog_row_0'], ['__catalog_column_0', '__catalog_column_1'],
        )

    assert draw_dashed.call_count == 1


def test_failure_hierarchy_uses_dashed_child_boundaries_within_one_row_group() -> None:
    frame = chart_frame({
        '__catalog_row_0': ['VoLTE', 'VoLTE', 'MultiRAB'],
        '__catalog_row_1': ['Belfast', 'Bristol', 'Belfast'],
        '__catalog_column_0': ['VF', 'VF', 'VF'],
        '__catalog_failure_state': ['Failed', 'Dropped', 'Failed'],
    })

    with patch('src.modules.cdr_reporting._draw_dashed_horizontal_line') as draw_dashed:
        _render_failure_count_hierarchy(
            'Voice failures per city', frame,
            ['__catalog_row_0', '__catalog_row_1'], ['__catalog_column_0'],
        )

    assert draw_dashed.call_count == 1
    assert draw_dashed.call_args.args[2] == 20 + (285 - 28) / 2


def test_nsa_catalogue_splits_template_screenshots_into_individual_charts() -> None:
    entries = [with_default_calculated_dimensions(entry) for entry in load_catalog_csv(Path(__file__).parent / 'fixtures' / 'NSA Slide Template.csv')]
    slide_ten = [entry for entry in entries if entry.slide == 10]
    slide_thirteen = [entry for entry in entries if entry.slide == 13]

    assert len(slide_ten) == 2
    assert {entry.layout for entry in slide_ten} == {'Title + 1 rows + 2 columns + comments down'}
    assert len(slide_thirteen) == 3
    assert {entry.layout for entry in slide_thirteen} == {'Title + 1 rows + 3 columns + comments down'}
    assert {slide: sum(entry.slide == slide for entry in entries) for slide in range(12, 17)} == {
        12: 2, 13: 3, 14: 3, 15: 4, 16: 2,
    }


def test_catalogue_uses_explicit_title_and_transition_slides() -> None:
    entries = [with_default_calculated_dimensions(entry) for entry in load_catalog_csv(Path(__file__).parent / 'fixtures' / 'NSA Slide Template.csv')]
    structural = [entry.chart_type for entry in entries if not entry.source_kind]
    title = next(entry for entry in entries if entry.slide == 1)
    conclusions = next(entry for entry in entries if entry.slide == 17)

    assert set(structural) == {'Title Slide', 'Transition Slide'}
    assert (title.chart_type, title.layout) == ('Title Slide', 'Title Page')
    assert (conclusions.chart_type, conclusions.layout) == ('Transition Slide', 'Title Only')


def test_grid_layout_without_comments_keeps_placeholder_ten_as_a_chart() -> None:
    from src.modules.cdr_reporting import (
        _clear_commentary,
        _named_slide_layout,
        _remove_template_chart_placeholders,
        _set_commentary,
    )

    presentation = Presentation('assets/powerpoint-templates/Template_VodafoneThree.pptx')
    layout = _named_slide_layout(presentation, 'Title + 1 rows + 1 columns')
    assert layout is not None
    slide = presentation.slides.add_slide(layout)
    chart_placeholder = next(shape for shape in slide.placeholders if shape.placeholder_format.idx == 10)
    original_xml = chart_placeholder._element.xml

    _clear_commentary(slide)
    _set_commentary(slide, ['Keep this chart placeholder'])
    assert next(shape for shape in slide.placeholders if shape.placeholder_format.idx == 10)._element.xml == original_xml

    _remove_template_chart_placeholders(slide)

    assert not any(shape.placeholder_format.idx == 10 for shape in slide.placeholders)


def test_template_chart_image_preview_uses_combined_reporting_rows(client, monkeypatch) -> None:
    import src.DriveTestAnalyzer as app_module

    client.post('/login', data={'username': 'super', 'password': 'super123'}, follow_redirects=False)
    catalogue_content = (
        ','.join(TRAILING_DYNAMIC_CATALOG_HEADERS)
        + '\n1,Preview slide,,Title and 1 column + Comments,Preview chart,CDR-Data,Mean_Data_Rate,Average Vertical Bars,,Operator,Campaign,,Top\n'
    )
    created = client.post(
        '/ppt-designer/templates', data={'catalogue_name': 'Combined Preview'},
        files={'catalogue_file': ('combined-preview.csv', BytesIO(catalogue_content.encode()), 'text/csv')},
        follow_redirects=False,
    )
    assert created.status_code == 303

    observed: dict[str, object] = {}
    monkeypatch.setattr(app_module.repository, 'list_datasets', lambda: [
        {'id': 7, 'status': 'ready', 'dataset_kind': 'data', 'updated_at': '', 'processed_at': '', 'normalization_version': ''},
    ])
    monkeypatch.setattr(app_module.repository, 'load_dataset_rows', lambda *_args: pytest.fail('Individual CDR tables must not be loaded.'))
    monkeypatch.setattr(
        app_module,
        '_combined_reporting_frame',
        lambda datasets, entries, multivendor: observed.update({
            'dataset_ids': [dataset['id'] for dataset in datasets],
            'entries': entries,
            'multivendor': multivendor,
        }) or pd.DataFrame(),
    )
    monkeypatch.setattr(app_module, '_cached_filtered_chart_frame', lambda _key, frame, *_args: frame)
    monkeypatch.setattr(app_module, 'render_catalog_chart_preview', lambda *_args, **_kwargs: b'PNG')
    monkeypatch.setattr(
        app_module,
        'preview_catalog_chart_data',
        lambda *_args, **_kwargs: (pd.DataFrame([{'Operator': 'EE'}]), {'filter_values': {}}),
    )

    data_preview = client.post(
        '/ppt-designer/templates/Combined%20Preview/chart-preview',
        json={'catalogue_content': catalogue_content, 'row_index': 0},
    )
    data_export = client.post(
        '/ppt-designer/templates/Combined%20Preview/chart-preview',
        json={'catalogue_content': catalogue_content, 'row_index': 0, 'download': True},
    )

    preview = client.post(
        '/ppt-designer/templates/Combined%20Preview/chart-image-preview',
        json={'catalogue_content': catalogue_content, 'row_index': 0},
    )

    assert data_preview.status_code == 200
    assert data_preview.json()['rows'] == [{'Operator': 'EE'}]
    assert set(data_preview.json()['column_metadata']) == {'Operator'}
    assert data_export.headers['content-type'].startswith('text/csv')
    assert data_export.text == 'Operator\nEE\n'
    assert preview.status_code == 200
    assert preview.content == b'PNG'
    assert observed['dataset_ids'] == [7]
    assert observed['multivendor'] is False
    # The previews read the CDRs of the Preview NR Mode of the editor (NSA by default): this CDR is NSA.
    sa_preview = client.post(
        '/ppt-designer/templates/Combined%20Preview/chart-image-preview',
        json={'catalogue_content': catalogue_content, 'row_index': 0, 'nr_mode': 'SA'},
    )
    assert sa_preview.status_code == 400 and 'No processed SA CDR-Data datasets' in sa_preview.json()['detail']


@pytest.mark.parametrize('rows,columns,cells', [
    ('Operator', '', [[15.0], [35.0]]),
    ('', 'Operator', [[15.0, 35.0]]),
    ('Operator', 'Campaign', [[10.0, 20.0], [30.0, 40.0]]),
])
def test_average_vertical_bars_have_a_value_in_each_bar_with_rows_columns_or_both(rows, columns, cells) -> None:
    from src.modules.cdr_reporting import CatalogEntry, catalog_chart_payload

    frame = pd.DataFrame({'Operator': ['EE', 'EE', 'O2', 'O2'], 'Campaign': ['C1', 'C2', 'C1', 'C2'], 'Mean_Data_Rate': [10.0, 20.0, 30.0, 40.0]})
    entry = CatalogEntry(slide=1, slide_title='Rates', slide_subtitle='', layout='', chart_title='Rate', cdr_source='CDR-Data',
                         kpi='Mean_Data_Rate', chart_type='Average Vertical Bars', legend='', filters='', rows_split=rows, columns_split=columns)
    assert catalog_chart_payload(frame, entry)['cells'] == cells
