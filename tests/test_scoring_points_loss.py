"""Where the scoring points are lost: points-lost maps per City, ITL3 area and custom Clusters."""
from __future__ import annotations

import json
import random
from io import BytesIO

import pandas as pd
import pytest
from pptx import Presentation

from src.modules import scoring_points_loss as points_loss
from src.modules.scoring import calculate_scoring
from src.modules.scoring_exports import export_scoring_powerpoint, export_scoring_report
from src.modules.scoring_reports import default_scenario
from src.modules.scoring_views import build_scoring_views
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_api import scoring_api  # noqa: F401  (fixture)
from tests.test_scoring_exports import TEMPLATE

PLACES = {'Leeds': (53.80, -1.55), 'Manchester': (53.48, -2.24), 'London': (51.51, -0.13)}


def _result(levels=('Operator',), roads=True):
    """Three operators; Vodafone and Three fail calls mostly in Leeds and on a road towards it (without the road, no
    Connecting Roads results: All Environments scales the City up to the full maximum)."""
    random.seed(7)
    rows = []
    for operator in ('EE', 'Vodafone UK', 'Three UK'):
        for city, (latitude, longitude) in PLACES.items():
            for index in range(40):
                failed = operator != 'EE' and ((city == 'Leeds' and index < 8) or (city == 'Manchester' and index < 2))
                rows.append({'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': 'City',
                             'City': city, 'Session_Type': 'CALL', 'Call_Status': 'Failed' if failed else 'Completed',
                             'Call_Start_Latitude_A': latitude + random.uniform(-.02, .02),
                             'Call_Start_Longitude_A': longitude + random.uniform(-.02, .02)})
        for index in range(30 if roads else 0):
            rows.append({'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': 'Connectionroad',
                         'City': 'Sheffield to Leeds', 'Session_Type': 'CALL',
                         'Call_Status': 'Failed' if operator != 'EE' and index % 6 == 0 else 'Completed',
                         'Call_Start_Latitude_A': 53.38 + index * .014, 'Call_Start_Longitude_A': -1.47 - index * .003})
    result = calculate_scoring({'voice': pd.DataFrame(rows)}, list(levels), configuration=scoring_configuration())
    job = {'id': 1, 'levels': list(levels), 'baseline_operator': 'EE', 'configuration': result['configuration'],
           'campaigns': ['2026-Q2']}
    return job, result


def test_points_lost_are_spread_over_the_cities_routes_and_itl3_areas():
    job, result = _result()
    maps = build_scoring_views(job, result)['insights']['points_loss_maps']
    vodafone = {item['field']: item for item in maps if item['operator'] == 'Vodafone UK' and item['environment'] == 'Combined'}
    assert {'City', 'Area'} <= set(vodafone) and vodafone['Area']['area_label'] == 'ITL3 area'
    cities = {area['name']: area for area in vodafone['City']['areas']}
    # Leeds loses most of the city points; the road is a route between cities.
    assert next(iter(cities)) in {'Leeds', 'Sheffield to Leeds'}
    assert cities['Leeds']['points'] > cities['Manchester']['points'] and 'London' not in cities
    assert cities['Sheffield to Leeds']['kind'] == 'route' and cities['Sheffield to Leeds']['route']
    # The map areas (ITL3 in the UK) share the same total: the tests are placed in the polygon that contains them.
    areas = {area['name']: area['points'] for area in vodafone['Area']['areas']}
    assert vodafone['Area']['total'] == pytest.approx(vodafone['City']['total'])
    # Leeds also takes the end of the road that enters it.
    assert cities['Leeds']['points'] < areas['Leeds'] < cities['Leeds']['points'] + cities['Sheffield to Leeds']['points']
    # EE completes every call: it loses no points and gets no map.
    assert not any(item['operator'] == 'EE' for item in maps)
    # Each environment has its own maps; All Environments names, beside each area, where it loses points.
    environments = {item['environment'] for item in maps} - {'Combined'}
    assert environments and all(cities[name]['environments'] for name in cities)
    for environment in environments:
        city = next(item for item in maps if item['operator'] == 'Vodafone UK' and item['environment'] == environment
                    and item['field'] == 'City')
        assert all(not area['environments'] for area in city['areas'])
    # The map areas of each City's tests (of every operator), to show it alone on the map.
    assert cities['Leeds']['area_tests'] == {'Leeds': 120}
    assert set(cities['Sheffield to Leeds']['area_tests']) >= {'Sheffield', 'Leeds'}
    # The points each City loses in each area add up to its points, and a route loses them along its way.
    assert cities['Leeds']['area_losses'] == {'Leeds': pytest.approx(cities['Leeds']['points'])}
    road = cities['Sheffield to Leeds']['area_losses']
    assert len(road) > 1 and sum(road.values()) == pytest.approx(cities['Sheffield to Leeds']['points'])
    assert areas['Leeds'] == pytest.approx(cities['Leeds']['points'] + road.get('Leeds', 0))


def test_map_areas_take_the_points_of_the_city_or_route_that_loses_most_in_them():
    areas = [
        {'name': 'London', 'points': 12.5, 'kind': 'place',
         'area_tests': {'Westminster and City of London': 60, 'Camden': 39, 'West Surrey': 1},
         'area_losses': {'Westminster and City of London': 8.0, 'Camden': 4.4, 'West Surrey': .1}},
        {'name': 'Coventry', 'points': 5.5, 'kind': 'place', 'area_tests': {'Coventry': 50}, 'area_losses': {'Coventry': 5.5}},
        {'name': 'Coventry to Corby', 'points': 1.7, 'kind': 'route', 'area_tests': {'Coventry': 10, 'Leicestershire': 30},
         'area_losses': {'Coventry': 1.2, 'Leicestershire': .5}},
    ]
    # The points of the city or route that loses most among those measured in each area (2% of its tests or more).
    assert points_loss.row_area_values(areas) == {
        'Westminster and City of London': 12.5, 'Camden': 12.5, 'Coventry': 5.5, 'Leicestershire': 1.7}



def test_the_itl3_areas_shipped_for_the_uk_cover_it():
    from src.modules import map_areas

    assert map_areas.bundled_layers()['GBR']['level_label'] == 'ITL3 area'
    boundaries = map_areas.bundled_rings('GBR')
    assert len(boundaries) == 182 and 'Leeds' in boundaries
    names = map_areas.area_index(None).names_of(pd.Series([53.80, 51.51, 0.0]), pd.Series([-1.55, -0.13, 0.0]))
    assert names.iloc[0] == 'Leeds' and names.iloc[1] == 'Westminster and City of London' and pd.isna(names.iloc[2])


def test_custom_cluster_polygons_are_read_from_a_geojson(tmp_path):
    path = tmp_path / 'clusters.geojson'
    path.write_text(json.dumps({'type': 'FeatureCollection', 'crs': {'type': 'name', 'properties': {'name': 'EPSG:4326'}},
                                'features': [{'type': 'Feature', 'properties': {'Cluster': 'North'},
                                              'geometry': {'type': 'Polygon', 'coordinates': [[[-2.5, 53], [-1, 53], [-1, 54.5], [-2.5, 54.5], [-2.5, 53]]]}}]}))
    boundaries = points_loss.mapping_boundaries(path, 'Cluster')
    assert list(boundaries) == ['North'] and boundaries['North'][0][0] == [-2.5, 53.0]
    assert points_loss.map_boundaries({'points_loss': {'boundaries': {'Cluster': boundaries}}}, 'Cluster') == boundaries


def _titles(content):
    return [slide.shapes.title.text_frame.text.split('\n')[0] for slide in Presentation(BytesIO(content)).slides
            if slide.shapes.title is not None]


def test_powerpoint_maps_the_points_lost_of_the_compared_operators():
    job, result = _result()
    content = export_scoring_powerpoint(job, result, TEMPLATE)
    titles = _titles(content)
    assert 'Points Lost per City — Vodafone UK' in titles and 'Points Lost per City — Three UK' in titles
    assert not any(title.startswith('Points Lost per City — EE') for title in titles)
    deck = Presentation(BytesIO(content))
    slide = next(slide for slide in deck.slides if slide.shapes.title is not None
                 and slide.shapes.title.text_frame.text.startswith('Points Lost per City — Vodafone UK'))
    names = {shape.name for shape in slide.shapes}
    assert {'Scoring Points Lost Map', 'Scoring Points Lost Bars', 'Scoring Points Lost Scale'} <= names
    # The operator, large, above the map and the ranking.
    heading = next(shape for shape in slide.shapes if shape.name == 'Scoring Points Lost Operator')
    assert heading.text_frame.text == 'Vodafone UK' and heading.text_frame.paragraphs[0].runs[0].font.size.pt == 24
    # Its figures are National points: the part of the National loss in each place, not the score of the place on its own.
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert any('by the National points Vodafone UK loses' in text and 'of its National score lost in total' in text for text in texts)
    national = next(shape for shape in slide.shapes if shape.name == 'Scoring Points Lost National Note')
    assert national.text_frame.text.startswith("National points: the part of Vodafone UK's National loss in the tests of each place")
    scenario = default_scenario()
    for options in scenario['scorings'].values():
        options['points']['map'] = False
    without = export_scoring_report([{'job': job, 'result': result, 'scenario': scenario}], TEMPLATE)
    assert not any(title.startswith('Points Lost per') for title in _titles(without))


def test_job_api_serves_the_map_layers_without_the_raw_shares(scoring_api, monkeypatch):
    from types import ModuleType

    from src.modules import saved_filters

    client = scoring_api['client']
    repository = scoring_api['repository']
    repository.replace_scoring_configuration(scoring_configuration())
    _job, computed = _result()
    engine = ModuleType('src.modules.scoring')
    engine.METHOD_VERSION = 'points-loss-api-test'
    engine.required_input_columns = lambda kind, levels: [*levels, 'score']
    engine.calculate_scoring = lambda frames, levels, *, baseline_operator, configuration=None: computed
    monkeypatch.setattr(saved_filters, '_scoring_engine', lambda: engine)
    created = client.post('/api/scoring/saved-filters', json={
        'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator'],
        'nr_mode': 'NSA', 'baseline_operator': 'EE',
    })
    assert created.status_code == 200, created.text
    job_id = created.json()['saved_filter']['id']
    assert saved_filters.calculate_saved_filter(repository, job_id)['status'] == 'completed'

    payload = client.get(f'/api/scoring/saved-filters/{job_id}').json()
    assert 'points_loss' not in payload
    assert any(item['field'] == 'Area' for item in payload['views']['insights']['points_loss_maps'])
    # UK tests use the ITL3 areas shipped with the application, over the outline of the country.
    layer = client.get(f'/api/scoring/saved-filters/{job_id}/boundaries/Area').json()
    assert layer['field'] == 'Area' and len(layer['boundaries']) == 182 and layer['background']
    assert client.get(f'/api/scoring/saved-filters/{job_id}/boundaries/ITL3').status_code == 404
    # Without a Clusters dataset applied to the CDRs the Cluster layer has no polygons.
    assert client.get(f'/api/scoring/saved-filters/{job_id}/boundaries/Cluster').json()['boundaries'] == {}
    assert client.get(f'/api/scoring/saved-filters/{job_id}/boundaries/Unknown').status_code == 404


def test_the_web_map_links_the_ranking_and_the_areas():
    from pathlib import Path

    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    # Each column is headed by its operator; rows and areas choose each other, and chosen rows centre the map.
    assert "title.className = 'scoring-loss-map-operator';" in script
    # Its figures are National points (the part of the National loss in each place), as those of the performance owners.
    assert "bold('National points'), ': the part of the National loss in the tests of each place, not its own score" in script
    assert "'National points lost',\n      ...(rated ? ['Per 1,000 tests'] : []), 'Share']" in script
    assert "document.createTextNode(' National points lost')" in script
    assert 'function lossMapSelection(card, svg, rows, initialLinks)' in script and 'selection.zoom.fit({x, y, width, height});' in script
    # Every area keeps its colour: choosing rows or areas only dims the others.
    assert "marks.forEach(mark => mark.classList.toggle('is-dimmed', !lit.has(mark.dataset.area)));" in script
    # One choice above the maps, the city or route that loses most by default; changing it keeps the chosen rows.
    assert "for (const [value, text] of [['top', 'Points of the city or route that loses most in each area']," in script
    assert "cards.forEach(card => card.setColourMode?.(modeSelect.value));" in script
    assert "mode = real && rank === 'points' && insightSelections.get('points-loss-colour') === 'losses' ? 'losses' : 'top';" in script
    assert 'if (event.shiftKey && onSelect) onSelect(area, event); else apply(area);' in script
    # All Environments names the environments of each area.
    assert "...(withEnvironments ? ['Type'] : [])" in script


def test_a_level_with_a_single_value_keeps_the_maps():
    # One campaign: the views leave the Campaign level out, in the score rows and in the points-lost shares alike.
    job, result = _result(levels=('Operator', 'Campaign'))
    assert result['points_loss']['shares'][0]['campaign'] == '2026-Q2'
    maps = build_scoring_views(job, result)['insights']['points_loss_maps']
    assert {item['operator'] for item in maps if item['field'] == 'City'} == {'Vodafone UK', 'Three UK'}
    assert all(not item['context'] for item in maps)
    # The saved result keeps its campaign.
    assert result['points_loss']['shares'][0]['campaign'] == '2026-Q2'


def test_operators_of_some_tables_only_keep_their_colour():
    from src.modules.scoring_insights import operator_styles

    # VF SA is measured in the 2026-Q2 campaign only, not in the first score table.
    tables = [
        {'operator_styles': {'VF': {'label': 'VF', 'color': '#E60000'}}},
        {'operator_styles': {'VF': {'label': 'VF', 'color': '#123456'}, 'VF SA': {'label': 'VF SA', 'color': '#8B1A1A'}}},
        {},
    ]
    styles = operator_styles(tables)
    assert styles['VF SA']['color'] == '#8B1A1A'
    assert styles['VF']['color'] == '#E60000'


def test_points_lost_views_count_in_the_points_of_a_scaled_combined():
    job, result = _result(roads=False)
    assert result['environment_scaling']['scaled_environments']
    insights = build_scoring_views(job, result)['insights']
    stake = {operator['operator']: operator['lost'] for item in insights['points_at_stake'] if item['environment'] == 'Combined'
             for operator in item['operators']}
    assert stake['Vodafone UK'] > 0
    # The maps, Where to act and the points lost per KPI of All Environments lose the points of Where the points go.
    maps = [item for item in insights['points_loss_maps'] if item['environment'] == 'Combined']
    plans = [item for item in insights['act_plans'] if item['environment'] == 'Combined']
    profiles = [item for item in insights['kpi_gap_profiles'] if item['environment'] == 'Combined']
    assert maps and plans and profiles
    for item in maps:
        assert item['total'] == pytest.approx(stake[item['operator']])
    for item in plans:
        assert item['lost'] == pytest.approx(stake[item['operator']])
    for item in profiles:
        assert sum(row['total_gap_to_maximum'] for row in item['to_maximum']) == pytest.approx(stake[item['operator']])
    # Each environment keeps its own points.
    city = next(item for item in insights['points_loss_maps'] if item['environment'] != 'Combined'
                and item['operator'] == 'Vodafone UK' and item['field'] == 'City')
    assert city['total'] < stake['Vodafone UK']


def test_the_reference_operator_has_its_map_and_its_gap_to_maximum():
    from pathlib import Path

    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    # The reference operator keeps its Points Lost Map and its points lost per KPI; the KPI GAP Profile leaves it out,
    # as its GAP to itself would be empty.
    assert "const maps = insightItems(payload, 'points_loss_maps', environment);" in script
    assert "const profiles = insightItems(payload, 'kpi_gap_profiles', environment);" in script
    assert "const shown = inScope.filter(item => chosen.includes(item.operator) && item.operator !== item.reference);" in script
    assert "for (const profile of shown) grid.append(profileTable(profile, 'maximum'));" in script
    # A choice of operators saved before the reference was offered leaves the reference enabled.
    assert "const newlyOffered = item => offered === '*' ? false : Array.isArray(offered) ? !offered.includes(item.operator) : isReference(item);" in script
    job, result = _result()
    profiles = build_scoring_views(job, result)['insights']['kpi_gap_profiles']
    reference = [item for item in profiles if item['operator'] == 'EE']
    assert reference and all(item['to_maximum'] and not item['to_reference'] for item in reference)


def test_exports_include_the_reference_operator_only_when_it_is_a_focus_operator():
    job, result = _result()

    def export(focus):
        # The Focus operators of the Operator Maps.
        groups = [{'canonical': name, 'aliases': [], 'position': index, 'color': '#365F91', 'role': 'focus'}
                  for index, name in enumerate(focus)]
        return _titles(export_scoring_report([{'job': job, 'result': result, 'scenario': default_scenario()}],
                                             TEMPLATE, groups))

    titles = export(['Vodafone UK', 'EE'])
    # The reference has its points lost per KPI, and no GAP to itself.
    assert any(title.startswith('Points lost per KPI — ') and 'EE' in title.split(' | ')[0] for title in titles)
    assert any(title.startswith('KPI GAP Profile — Vodafone UK vs EE') for title in titles)
    titles = export(['Vodafone UK'])
    assert not any(title.startswith('Points lost per KPI — ') and 'EE' in title.split(' | ')[0] for title in titles)


def test_places_rank_by_their_points_lost_per_1000_tests(monkeypatch):
    from src.modules import scoring_insights

    monkeypatch.setattr(scoring_insights, 'MIN_RATE_TESTS', 35)
    job, result = _result()
    maps = build_scoring_views(job, result)['insights']['points_loss_maps']
    of = {item['field']: item for item in maps if item['operator'] == 'Vodafone UK' and item['environment'] == 'Combined'}
    assert of['City']['rate_tests'] == 1000 and of['City']['rate_minimum'] == 35
    cities = {area['name']: area for area in of['City']['areas']}
    # 40 calls in each city: its points per 1,000 of them; the 30 calls of the road are under the minimum.
    assert cities['Leeds']['tests'] == 40
    assert cities['Leeds']['rate'] == pytest.approx(cities['Leeds']['points'] / 40 * 1000)
    assert cities['Sheffield to Leeds']['tests'] == 30 and cities['Sheffield to Leeds']['rate'] is None
    # The map areas have no tests of their own.
    assert 'rate_tests' not in of['Area'] and 'rate' not in of['Area']['areas'][0]
    # The PowerPoint follows each map by its per 1,000 tests slide, which leaves the road out.
    content = export_scoring_powerpoint(job, result, TEMPLATE)
    titles = _titles(content)
    position = titles.index('Points Lost per City — Vodafone UK')
    assert titles[position + 1] == 'Points Lost per City, per 1,000 tests — Vodafone UK'
    deck = Presentation(BytesIO(content))
    slide = [slide for slide in deck.slides if slide.shapes.title is not None][position + 1]
    texts = ' '.join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
    assert 'loses per 1,000 tests' in texts and '(35 tests or more) hold' in texts
    bars = next(shape for shape in slide.shapes if shape.name == 'Scoring Points Lost Bars').chart
    assert 'Sheffield to Leeds' not in ' '.join(bars.plots[0].categories)


def test_results_without_the_tests_of_each_place_keep_one_ranking():
    job, result = _result()
    for entry in result['points_loss']['serving_tests']:
        for field in ('City', 'Region', 'Cluster'):
            entry['counts'].pop(field, None)
    maps = build_scoring_views(job, result)['insights']['points_loss_maps']
    city = next(item for item in maps if item['operator'] == 'Vodafone UK' and item['field'] == 'City')
    assert 'rate_tests' not in city and 'rate' not in city['areas'][0]
    titles = _titles(export_scoring_powerpoint(job, result, TEMPLATE))
    assert 'Points Lost per City — Vodafone UK' in titles and not any('per 1,000 tests' in title for title in titles)


def test_the_map_of_cdr_types_put_together_from_their_contributions_is_the_one_added_at_once():
    import numpy as np

    rng = np.random.default_rng(5)
    city, region, site = (points_loss.AREA_COLUMN.format(field) for field in ('City', 'Region', points_loss.SITE))

    def frame(rows: int, cities: list[str]) -> pd.DataFrame:
        return pd.DataFrame({
            points_loss.LATITUDE: rng.uniform(50, 56, rows), points_loss.LONGITUDE: rng.uniform(-5, 1, rows),
            city: rng.choice(cities, rows), region: rng.choice(['North', 'South', None], rows),
            points_loss.PLACE: rng.choice([True, False], rows),
            site: rng.choice(['EE\x1eS1', 'EE\x1eS2', None], rows)})

    # London fills part of its 60 points with the first CDR type, the rest from the others; Leeds only shows later.
    frames = [frame(40, ['London']), frame(3000, ['London', 'Leeds']), frame(7, ['Leeds', 'York'])]
    added = points_loss.AreaGeometry()
    for each in frames:
        added.add(each)
    applied = points_loss.AreaGeometry()
    for each in frames:
        # As a process sends it: only the points some room takes.
        applied.apply(points_loss.compact_contribution(points_loss.contributions(each)))
    assert applied.document() == added.document()
    assert applied.used_areas == added.used_areas
    assert len(applied.areas['City']['London']['points']) == points_loss.MAX_AREA_POINTS


def test_the_points_of_an_area_sent_to_another_process_are_those_of_all_its_tests():
    import numpy as np

    for count in (1, 7, 59, 60, 61, 119, 120, 121, 1000, 12345):
        latitudes, longitudes = np.linspace(50, 51, count), np.linspace(-1, 0, count)
        whole = points_loss._AreaPoints(latitudes, longitudes)
        compact = points_loss._AreaPoints(latitudes.copy(), longitudes.copy())
        compact.compact()
        assert len(compact.latitudes) <= count
        for room in range(1, points_loss.MAX_AREA_POINTS + 1):
            assert compact.take(room) == whole.take(room)
