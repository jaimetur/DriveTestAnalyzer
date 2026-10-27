"""National & Geographical Areas Trends: the National scoring next to chosen cities and each Region, Cluster or City."""
from __future__ import annotations

import json
import shutil
import subprocess
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

from src.modules import scoring as scoring_engine
from src.modules.scoring import most_reliable_result
from src.modules.scoring_area_summary import (
    area_summary_levels, build_area_summary, calculate_area_summary, calculate_geographical, geographical_insights,
    geographical_levels, geographical_passes, has_area_filters, normalize_area_summary_request, whole_country_filters,
)
from src.modules.scoring_exports import export_scoring_report
from src.modules.scoring_reports import default_scenario, has_geographic_filters, normalize_report_configuration
from tests.scoring_synthetic import synthetic_result
from tests.test_scoring_exports import TEMPLATE
from tests.test_scoring_api import scoring_api  # noqa: F401  (fixture)
from tests.test_scoring_most_reliable import _reliable_configuration

SUMMARY_TITLE = 'Best Network Trends — National & Geographical Areas'


def _calls(city, region, statuses):
    return [{'Campaign': '2026-Q2', 'Operator': operator, 'G_Level_1': 'Drive', 'G_Level_2': 'City',
             'Session_Type': 'CALL', 'Call_Status': status, 'City': city, 'Region': region}
            for operator, values in statuses.items() for status in values]


ROWS = [
    *_calls('London', 'Vendor A', {'EE': ['Completed'] * 4, 'Three UK': ['Completed', 'Completed', 'Completed', 'Failed']}),
    *_calls('Leeds', 'VMO2 North', {'EE': ['Completed', 'Failed'], 'Three UK': ['Completed', 'Completed']}),
]


def _calculation():
    configuration = _reliable_configuration()
    kwargs = {'baseline_operator': 'EE', 'configuration': configuration}
    frames = {'voice': pd.DataFrame(ROWS)}
    result = scoring_engine.calculate_scoring(frames, ['Operator'], **kwargs)
    job = {'levels': ['Operator'], 'aggregation_levels': ['Operator'], 'aggregation_contract_version': 2,
           'baseline_operator': 'EE', 'configuration': configuration}
    return job, result, frames, kwargs


def _kpi_values(result):
    return {(row.get('city') or row.get('region'), row['operator'], row['environment'], row['kpi_code']): row['value']
            for row in result['scoring'] if row['value'] is not None}


def test_area_summary_request_is_scored_on_the_whole_country():
    # Main Cities (the flag or its token among the cities) are kept as such, and resolved when it is calculated.
    assert normalize_area_summary_request({'breakdown': 'Region', 'cities': [], 'main_cities': True}, {}) == {
        'breakdown': 'Region', 'cities': [], 'time_split': 'Campaign', 'main_cities': True}
    assert normalize_area_summary_request({'breakdown': 'None', 'cities': ['@main-cities', 'Leeds', 'leeds']}, {}) == {
        'breakdown': None, 'cities': ['Leeds'], 'time_split': 'Campaign', 'main_cities': True}
    from src.modules.scoring_area_summary import resolve_main_cities
    assert resolve_main_cities({'breakdown': 'Region', 'cities': ['Leeds'], 'main_cities': True}, ['London', 'leeds']) == {
        'breakdown': 'Region', 'cities': ['Leeds', 'London']}
    # Always included: without a breakdown or separate cities (one saved when they could be left out), each Region and London.
    assert normalize_area_summary_request(None, {}) == normalize_area_summary_request({'breakdown': 'None', 'cities': []}, {}) == {
        'breakdown': 'Region', 'cities': ['London'], 'time_split': 'Campaign'}
    assert normalize_area_summary_request({'breakdown': 'region', 'cities': ['London', 'london', ' ']}, {}) == {
        'breakdown': 'Region', 'cities': ['London'], 'time_split': 'Campaign'}
    # Area filters are allowed: the summary is scored without them.
    assert normalize_area_summary_request({'breakdown': 'Region'}, {'City': ['London']}) == {
        'breakdown': 'Region', 'cities': [], 'time_split': 'Campaign'}
    assert whole_country_filters({'City': ['London'], 'Region': ['North'], 'Operator': ['EE']}) == {'Operator': ['EE']}
    assert has_area_filters({'Cluster': ['C1']}) and not has_area_filters({'Operator': ['EE'], 'City': []})
    # Operator or Campaign filters keep the whole country.
    assert normalize_area_summary_request({'cities': ['London']}, {'Operator': ['EE']}) == {
        'breakdown': None, 'cities': ['London'], 'time_split': 'Campaign'}
    assert area_summary_levels(['Operator', 'Campaign'], {'breakdown': 'Region', 'cities': ['London']}) == ['Region', 'City']
    assert area_summary_levels(['Operator', 'Region'], {'breakdown': 'Region', 'cities': []}) == []


def test_each_area_is_scored_with_its_own_measurements():
    job, result, frames, kwargs = _calculation()
    summary = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                     {'breakdown': 'Region', 'cities': ['london', 'Bristol'], 'time_split': 'All'}, kwargs)
    assert [item['kind'] for item in summary['passes']] == ['National', 'City', 'Region']
    national, cities, regions = summary['passes']
    # Every selected CDR in one value, as the job: the National calculation is the job's own, not calculated again.
    assert national['reuses_job'] and 'result' not in national
    assert {row['city'] for row in cities['result']['scoring']} == {'London'}
    assert {row['region'] for row in regions['result']['scoring']} == {'Vendor A', 'VMO2 North'}
    assert 'points_loss' not in cities['result'] and 'configuration' not in cities['result']
    # London scores with its own calls only, the same as a calculation filtered to London.
    london = scoring_engine.calculate_scoring({'voice': pd.DataFrame([row for row in ROWS if row['City'] == 'London'])},
                                              ['Operator', 'City'], **kwargs)
    assert _kpi_values(cities['result']) == _kpi_values(london)
    # A chosen city without measurements is left out with a warning.
    assert summary['warnings'] == []
    missing = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                     {'breakdown': None, 'cities': ['Bristol']}, kwargs)
    assert [item['kind'] for item in missing['passes']] == ['National']
    assert missing['warnings'] == ['National & Geographical Areas Trends: the selected CDRs have no Bristol measurements.']


def test_a_job_with_area_filters_never_reuses_its_rows_for_the_summary():
    job, result, frames, kwargs = _calculation()
    # The rows of the whole country (read apart): the National calculation is not the filtered job's own.
    summary = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                     {'breakdown': None, 'cities': [], 'time_split': 'All'}, kwargs, job_filtered=True)
    national = summary['passes'][0]
    assert not national['reuses_job']
    assert _kpi_values(national['result']) == _kpi_values(result)


def test_a_job_split_by_no_area_has_its_geographical_areas():
    from src.modules.scoring_views import build_scoring_views
    job, result, frames, kwargs = _calculation()
    assert [item['kind'] for item in geographical_passes(['Operator', 'Campaign'])] == ['National', 'Region', 'City']
    # Whatever the Split by: its levels without the areas (the scoring Nationwide), split by the Breakdown, and the
    # Separate cities alone.
    assert [item['levels'] for item in geographical_passes(['Operator', 'City', 'Campaign'])] == [
        ['Operator', 'Campaign'], ['Operator', 'Campaign', 'Region'], ['Operator', 'Campaign', 'City']]
    assert [item['kind'] for item in geographical_passes(['Operator'], {'breakdown': 'City', 'cities': ['Leeds']})] == ['National', 'City']
    assert geographical_passes(['Operator'], {'breakdown': 'None', 'cities': ['Leeds']}) == [
        {'kind': 'National', 'field': None, 'levels': ['Operator']},
        {'kind': 'City', 'field': 'city', 'levels': ['Operator', 'City'], 'cities': ['Leeds']}]
    assert geographical_levels(['Operator']) == ['Region', 'City'] and geographical_levels(['Operator', 'Region']) == ['City']
    assert geographical_levels(['Operator'], {'breakdown': 'Cluster', 'cities': []}) == ['Cluster']
    geographical = calculate_geographical(scoring_engine.calculate_scoring, frames, job['levels'], kwargs)
    national, regions, london = geographical['passes']
    assert _kpi_values(national['result']) == _kpi_values(result)
    assert {row['region'] for row in regions['result']['scoring']} == {'Vendor A', 'VMO2 North'}
    # London is scored with its own calls only.
    assert {row['city'] for row in london['result']['scoring']} == {'London'} and london['cities'] == ['London']
    result['geographical'] = geographical
    # The Geographical Areas views: the card of the whole country first, then a card per Region and one of London.
    cards = geographical_insights(job, result)['location_cards']
    assert {(item['level'], tuple(card['label'] for card in item['cards'])) for item in cards} >= {
        ('National', ('National',)), ('Region', ('Vendor A', 'VMO2 North')), ('City', ('London',))}
    assert cards[0]['level'] == 'National'
    nationwide = {item['operator']: item['total'] for item in cards[0]['cards'][0]['operators']}
    assert set(nationwide) == {row['operator'] for row in result['scoring']}
    views = build_scoring_views(job, result)
    assert {item['level'] for item in views['insights']['location_cards']} >= {'National', 'Region', 'City'}
    # Results calculated before the scoring Nationwide: the job's own calculation of the whole country gives it.
    earlier = {'passes': [item for item in geographical['passes'] if item['kind'] != 'National']}
    earlier_cards = geographical_insights(job, {**result, 'geographical': earlier})['location_cards']
    assert earlier_cards[0]['level'] == 'National'
    assert {item['operator']: item['total'] for item in earlier_cards[0]['cards'][0]['operators']} == nationwide
    # A job split by an area shows them too, instead of its own rows; one calculated before them keeps its own.
    split = {**job, 'levels': ['Operator', 'Region'], 'aggregation_levels': ['Operator', 'Region']}
    by_region = scoring_engine.calculate_scoring(frames, ['Operator', 'Region'], **kwargs)
    split_views = build_scoring_views(split, {**by_region, 'geographical': geographical})
    assert {item['level'] for item in split_views['insights']['location_cards']} == {'National', 'Region', 'City'}
    assert {item['level'] for item in build_scoring_views(split, by_region)['insights']['location_cards']} == {'Region'}
    # Without area filters, the calculation equal to the job's own is read from it.
    own = calculate_geographical(scoring_engine.calculate_scoring, frames, ['Operator', 'Region'], kwargs, reuse_job=True)
    assert [(item['kind'], bool(item.get('reuses_job'))) for item in own['passes']] == [
        ('National', False), ('Region', True), ('City', False)]
    reused_cards = build_scoring_views(split, {**by_region, 'geographical': own})['insights']['location_cards']
    assert {item['level'] for item in reused_cards} == {'National', 'Region', 'City'}
    # With an area summary of the same rows, its calculations per Region and of London are read from it.
    summary = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                     {'breakdown': 'Region', 'cities': ['London'], 'time_split': 'All'}, kwargs)
    calls = []
    counting = lambda source, levels, **options: calls.append(levels) or scoring_engine.calculate_scoring(source, levels, **options)
    reused = calculate_geographical(counting, frames, job['levels'], kwargs, summary=summary)
    assert calls == [] and all(item['reuses_summary'] and 'result' not in item for item in reused['passes'][1:])
    # (Its National calculation is the job's own: so is the scoring Nationwide.)
    assert reused['passes'][0]['kind'] == 'National' and reused['passes'][0]['reuses_job']
    reused_views = build_scoring_views(job, {**result, 'geographical': reused, 'area_summary': summary})
    assert {item['level'] for item in reused_views['insights']['location_cards']} >= {'National', 'Region', 'City'}
    # Not one split by period, nor of other cities.
    assert len(calculate_geographical(counting, frames, job['levels'], kwargs, summary={**summary, 'time_split': 'Weekly'})['passes']) == 3
    assert len(calls) == 3
    # The separate cities of the National & Geographical Areas Trends are scored apart instead of London.
    assert geographical_passes(['Operator', 'Campaign'], {'breakdown': 'Region', 'cities': ['Leeds', ' ']})[2]['cities'] == ['Leeds']
    chosen = calculate_geographical(scoring_engine.calculate_scoring, frames, job['levels'], kwargs,
                                    areas={'breakdown': 'Region', 'cities': ['Leeds', 'London']})
    assert {row['city'] for row in chosen['passes'][2]['result']['scoring']} == {'Leeds', 'London'}
    chosen_cards = geographical_insights(job, {**result, 'geographical': chosen})['location_cards']
    assert any(item['level'] == 'City' and {card['label'] for card in item['cards']} == {'Leeds', 'London'} for item in chosen_cards)
    # Without measurements of London (or with filters leaving none), only the Regions.
    leeds = {'voice': pd.DataFrame([row for row in ROWS if row['City'] == 'Leeds'])}
    assert [item['kind'] for item in calculate_geographical(scoring_engine.calculate_scoring, leeds, job['levels'], kwargs)['passes']] == ['National', 'Region']


def test_the_summary_lists_national_the_chosen_cities_and_the_breakdown():
    job, result, frames, kwargs = _calculation()
    result['area_summary'] = calculate_area_summary(scoring_engine.calculate_scoring, frames, job['levels'],
                                                    {'breakdown': 'Region', 'cities': ['London']}, kwargs)
    summaries = build_area_summary(job, result)
    combined = next(item for item in summaries if item['environment'] == 'Combined')
    # National, the breakdown, then the cities.
    assert [area['label'] for area in combined['areas']] == ['National', 'Vendor A', 'VMO2 North', 'London']
    national = {row['operator']: row for row in combined['areas'][0]['rows']}
    # The reference first; the GAP is each operator's points minus the reference's.
    assert combined['areas'][0]['rows'][0]['operator'] == 'EE' and national['EE']['gap'] is None
    assert national['Three UK']['gap'] == pytest.approx(national['Three UK']['points'] - national['EE']['points'])
    # One campaign: no change from a previous campaign.
    assert combined['previous_campaign'] is None and national['EE']['delta'] is None
    # The Most Reliable scoring rates the area calculations with its own points.
    reliable = most_reliable_result(result, 'EE')
    reliable_passes = [item for item in reliable['area_summary']['passes'] if 'result' in item]
    assert reliable_passes and all(item['result']['scoring_kind'] == 'most_reliable' for item in reliable_passes)


def _synthetic_report(campaigns=('2026-Q1', '2026-Q2'), with_summary=True):
    levels = ('Operator', 'Campaign')
    job, result = synthetic_result(levels=levels, campaigns=campaigns)
    _job, london = synthetic_result(levels=('Operator', 'City', 'Campaign'), cities=('London',), campaigns=campaigns)
    _job, areas = synthetic_result(levels=('Operator', 'City', 'Campaign'), cities=('Vendor A', 'VMO2 North'),
                                   campaigns=campaigns, seed=5)
    rename = lambda row: {('region' if key == 'city' else key): value for key, value in row.items()}
    regions = {**areas, 'scoring': [rename(row) for row in areas['scoring']],
               'totals': [rename(row) for row in areas['totals']]}
    summary = {'breakdown': 'Region', 'cities': ['London'], 'passes': [
        {'kind': 'National', 'field': None, 'levels': list(levels), 'reuses_job': True},
        {'kind': 'City', 'field': 'city', 'levels': [*levels, 'City'], 'cities': ['London'], 'reuses_job': False,
         'result': {key: value for key, value in london.items() if key != 'configuration'}},
        {'kind': 'Region', 'field': 'region', 'levels': [*levels, 'Region'], 'reuses_job': False,
         'result': {key: value for key, value in regions.items() if key != 'configuration'}},
    ]}
    # The job of a scenario without the National & area summary has none.
    if with_summary:
        result['area_summary'] = summary
    scenario = default_scenario(name='National')
    scenario['scorings']['most_reliable']['enabled'] = False
    scenario = normalize_report_configuration({'scenarios': [scenario]})['scenarios'][0]
    return Presentation(BytesIO(export_scoring_report([{'scenario': scenario, 'job': job, 'result': result}], TEMPLATE)))


def test_an_area_not_measured_in_the_latest_campaign_keeps_its_trend_without_a_score():
    levels = ('Operator', 'Campaign')
    job, result = synthetic_result(levels=levels, campaigns=('2026-Q1', '2026-Q2'))
    # London was measured in the first campaign only (the latest one is still being measured).
    _job, london = synthetic_result(levels=('Operator', 'City', 'Campaign'), cities=('London',), campaigns=('2026-Q1',))
    result['area_summary'] = {'breakdown': None, 'cities': ['London'], 'passes': [
        {'kind': 'National', 'field': None, 'levels': list(levels), 'reuses_job': True},
        {'kind': 'City', 'field': 'city', 'levels': [*levels, 'City'], 'cities': ['London'], 'reuses_job': False,
         'result': {key: value for key, value in london.items() if key != 'configuration'}},
    ]}
    combined = next(item for item in build_area_summary(job, result) if item['environment'] == 'Combined')
    national, city = combined['areas']
    assert (national['measured'], city['label'], city['measured']) == (True, 'London', False)
    assert {row['operator'] for row in city['rows']} == {row['operator'] for row in national['rows']}
    for row in city['rows']:
        assert (row['points'], row['delta'], row['gap']) == (None, None, None)
        assert row['trend'][0] is not None and row['trend'][1] is None


def test_an_area_without_an_environment_is_scaled_to_the_full_maximum():
    levels = ('Operator', 'Campaign')
    job, result = synthetic_result(levels=levels, campaigns=('2026-Q1',))
    # London has no Connecting Roads measurements at all: its score is scaled to the 1,000 points of the methodology.
    _job, london = synthetic_result(levels=('Operator', 'City', 'Campaign'), cities=('London',), campaigns=('2026-Q1',))
    london['scoring'] = [row for row in london['scoring'] if row['environment'] != 'DriveConnectionroad']
    london['totals'] = [row for row in london['totals'] if row['environment'] != 'DriveConnectionroad']
    result['area_summary'] = {'breakdown': None, 'cities': ['London'], 'passes': [
        {'kind': 'National', 'field': None, 'levels': list(levels), 'reuses_job': True},
        {'kind': 'City', 'field': 'city', 'levels': [*levels, 'City'], 'cities': ['London'], 'reuses_job': False,
         'result': {key: value for key, value in london.items() if key != 'configuration'}},
    ]}
    combined = next(item for item in build_area_summary(job, result) if item['environment'] == 'Combined')
    city = combined['areas'][1]
    ee = next(row for row in city['rows'] if row['operator'] == 'EE')
    city_points = sum(row['weighted_points'] for row in london['scoring']
                      if row['operator'] == 'EE' and row['environment'] == 'DriveCity')
    assert combined['maximum'] == pytest.approx(1000) and combined['scaled'] is True
    assert ee['points'] == pytest.approx(city_points * 1000 / 650)


@pytest.mark.skipif(not shutil.which('node'), reason='Node.js is required')
def test_the_areas_are_national_then_the_breakdown_then_the_cities():
    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    sources = []
    for name in ('areaCategory', 'areaChartOrder'):
        start = script.index(f'  function {name}(')
        sources.append(script[start:script.index('\n  }\n', start) + 4])
    program = ''.join(sources) + r"""
const area = (label, kind) => ({label, kind});
const one = [area('National', 'National'), area('London', 'City'), area('North', 'Region'), area('South', 'Region')];
const two = [area('National', 'National'), area('London', 'City'), area('Leeds', 'City'), area('North', 'Cluster')];
const cities = [area('National', 'National'), area('London', 'City'), area('Leeds', 'City')];
process.stdout.write(JSON.stringify([one, two, cities].map(areas => areaChartOrder(areas).map(item => item.label))));
"""
    orders = json.loads(subprocess.run(['node', '-e', program], capture_output=True, text=True, check=True).stdout)
    assert orders == [['National', 'North', 'South', 'London'], ['National', 'North', 'London', 'Leeds'],
                      ['National', 'London', 'Leeds']]
    # The table: each category in its colour, a thick rule where one starts.
    assert "className: `scoring-area-name is-${category}${measured ? '' : ' is-unmeasured'}`}," in script
    assert "tr.is-category-start td { border-top: 4px solid #5b6b78 !important; }" in (
        Path(__file__).resolve().parents[1] / 'src/web_interface/templates/scoring.html').read_text(encoding='utf-8')


def test_report_adds_the_national_and_areas_slide():
    presentation = _synthetic_report()
    slides = [slide for slide in presentation.slides
              if slide.shapes.title is not None and slide.shapes.title.text_frame.text.startswith(SUMMARY_TITLE)]
    assert len(slides) == 1
    table = next(shape.table for shape in slides[0].shapes if shape.has_table)
    assert [table.cell(0, column).text for column in range(len(table.columns))] == [
        'Area', 'Operator', 'Score\n26-Q2', 'Δ vs\n26-Q1', 'Δ vs\nEE']
    areas = [table.cell(row, 0).text for row in range(1, len(table.rows)) if table.cell(row, 0).text]
    assert areas == ['National', 'Vendor A', 'VMO2 North', 'London']
    # Every operator: the Focus operators choose none of the trends.
    operators = {table.cell(row, 1).text for row in range(1, len(table.rows))}
    assert operators == {'EE', 'O2 UK', 'Vodafone UK', 'Three UK'}
    charts = [shape.name for shape in slides[0].shapes if shape.has_chart]
    assert charts == [f'Scoring Area Summary Chart {area}' for area in areas]


def test_jobs_without_area_summary_have_no_slide():
    presentation = _synthetic_report(with_summary=False)
    assert not any(slide.shapes.title is not None and slide.shapes.title.text_frame.text.startswith(SUMMARY_TITLE)
                   for slide in presentation.slides)


def test_scenarios_of_the_whole_country_calculate_their_area_summary(monkeypatch):
    import src.DriveTestAnalyzer as core

    class Repository:
        def list_main_cities(self):
            return ['Leeds', 'Bristol']

    scenario = normalize_report_configuration({'scenarios': [default_scenario(name='National')]})['scenarios'][0]
    # By default a scenario of the whole country has it, with the four regions and London.
    assert scenario['area_summary'] == {'enabled': True, 'time_split': 'Campaign', 'breakdown': 'Region',
                                        'cities': ['London'], 'main_cities': False, 'campaigns': ['@all-campaigns']}
    # The Scoring Trends take All available campaigns by default.
    assert core._scenario_area_summary(Repository(), scenario) == {'breakdown': 'Region', 'cities': ['London'],
                                                                   'time_split': 'Campaign', 'campaigns': ['@all-campaigns']}
    # Main Cities stay as such in the Saved Filter of the scenario: the workspace's when it is calculated.
    scenario['area_summary'] = {'enabled': True, 'breakdown': 'None', 'cities': ['London'], 'main_cities': True}
    assert core._scenario_area_summary(Repository(), scenario) == {
        'breakdown': None, 'cities': ['London'], 'time_split': 'Campaign', 'main_cities': True}
    # The campaigns of the Scoring Trends join them; alone, they give the scoring trend of those campaigns.
    scenario['area_summary'] = {**scenario['area_summary'], 'campaigns': ['@latest-2-campaigns']}
    assert core._scenario_area_summary(Repository(), scenario)['campaigns'] == ['@latest-2-campaigns']
    # Always included: one saved without areas takes them from the normalization of the Saved Filter (Region and London).
    trends_only = {**scenario, 'area_summary': {'enabled': False, 'campaigns': ['2026-Q1', '2026-Q2']}}
    assert core._scenario_area_summary(Repository(), trends_only) == core._scenario_area_summary(
        Repository(), {**trends_only, 'context_filters': {'Region': ['Vendor A']}}) == {
        'breakdown': None, 'cities': [], 'time_split': 'Campaign', 'campaigns': ['2026-Q1', '2026-Q2']}
    scenario['area_summary'] = {'enabled': True, 'breakdown': 'None', 'cities': ['London'], 'main_cities': True}
    # A scenario with area filters has them too: they are scored on the whole country.
    for filtered in ({'context_filters': {'Region': ['Vendor A']}}, {'main_cities': True}):
        assert has_geographic_filters({**scenario, **filtered})
        assert core._scenario_area_summary(Repository(), {**scenario, **filtered}) == {
            'breakdown': None, 'cities': ['London'], 'time_split': 'Campaign', 'main_cities': True}
    london = normalize_report_configuration({'scenarios': [default_scenario(name='London', context_filters={
        'City': ['London']})]})['scenarios'][0]
    assert london['area_summary']['enabled'] is True
    # Also when it was saved unchecked.
    scenario['area_summary']['enabled'] = False
    assert core._scenario_area_summary(Repository(), scenario)['main_cities'] is True
    assert normalize_report_configuration({'scenarios': [{**scenario, 'name': 'Saved'}]})['scenarios'][0]['area_summary']['enabled'] is True

    calculated = []

    def create(_repository, _selected, levels, *_args, area_summary=None, **_kwargs):
        calculated.append(area_summary)
        return {'id': 1, 'status': 'completed', 'levels': list(levels)}, True

    monkeypatch.setattr(core, 'validate_complete_scoring_cdr_selection', lambda *_args, **_kwargs: [1])
    monkeypatch.setattr(core, 'save_saved_filter', create)
    monkeypatch.setattr(core, 'get_saved_filter', lambda _repository, job_id, include_result=False: {
        'id': job_id, 'status': 'completed', 'result': {}})
    core._scenario_saved_filter(Repository(), {'dataset_ids': [1], 'nr_mode': 'NSA'},
                               normalize_report_configuration({'scenarios': [default_scenario()]})['scenarios'][0], 'super')
    # Also when the scenario is calculated again without a Campaign level of a single value.
    assert calculated and all(item == {'breakdown': 'Region', 'cities': ['London'], 'time_split': 'Campaign',
                                       'campaigns': ['@all-campaigns']} for item in calculated)


def _timed_calls():
    rows = []
    for day, statuses in (('2026-02-10 09:00:00', {'EE': ['Completed', 'Completed'], 'Three UK': ['Failed', 'Completed']}),
                          ('2026-02-17 09:00:00', {'EE': ['Completed', 'Failed'], 'Three UK': ['Completed', 'Completed']})):
        for row in _calls('London', 'Vendor A', statuses):
            rows.append({**row, 'Campaign': 'UK_Q1_2026', 'Event_Start_Time': day})
    return rows


def test_the_time_split_divides_the_campaign_level_by_period():
    from src.modules.scoring import normalize_time_split

    configuration = _reliable_configuration()
    rows = _timed_calls()
    kwargs = {'baseline_operator': 'EE', 'configuration': configuration}
    weekly = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator', 'Campaign'],
                                              time_split='Weekly', **kwargs)
    assert {row['campaign'] for row in weekly['scoring']} == {'2026-W07', '2026-W08'}
    # The campaigns of the CDRs stay listed as they are.
    assert weekly['campaigns'] == ['UK_Q1_2026']
    # Each week is scored with its own tests.
    first_week = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows[:4])}, ['Operator'], **kwargs)
    week_values = {(row['operator'], row['kpi_code']): row['value'] for row in weekly['scoring']
                   if row['campaign'] == '2026-W07' and row['environment'] == 'DriveCity' and row['value'] is not None}
    assert week_values == {(row['operator'], row['kpi_code']): row['value'] for row in first_week['scoring']
                           if row['environment'] == 'DriveCity' and row['value'] is not None}
    monthly = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator', 'Campaign'],
                                               time_split='monthly', **kwargs)
    assert {row['campaign'] for row in monthly['scoring']} == {'2026-02'}
    # Without the Campaign level every selected CDR is aggregated into one value.
    pooled = scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, ['Operator'], time_split='Weekly', **kwargs)
    assert {row['campaign'] for row in pooled['scoring']} == {None}
    # Rows without a start time cannot be placed in a period.
    untimed = scoring_engine.calculate_scoring({'voice': pd.DataFrame(ROWS)}, ['Operator', 'Campaign'],
                                               time_split='Weekly', **kwargs)
    assert 'Voice: no test start time; its rows cannot be split by weekly period.' in untimed['warnings']
    assert normalize_time_split('Campaign') is None and normalize_time_split('YEARLY') == 'Yearly'


def test_the_area_summary_has_its_own_time_split(monkeypatch):
    import src.DriveTestAnalyzer as core

    configuration = _reliable_configuration()
    rows = _timed_calls()
    kwargs = {'baseline_operator': 'EE', 'configuration': configuration}
    frames = {'voice': pd.DataFrame(rows)}
    # The job itself splits by campaign; only the area summary is split by week.
    request = normalize_area_summary_request({'cities': ['London'], 'time_split': 'weekly'}, {})
    assert request == {'breakdown': None, 'cities': ['London'], 'time_split': 'Weekly'}
    assert set(area_summary_levels(['Operator', 'Campaign'], request)) == {'City', 'Event_Start_Time', 'Call_Start_Time',
                                                                            'Test_Start_Time'}
    summary = calculate_area_summary(scoring_engine.calculate_scoring, frames, ['Operator', 'Campaign'], request, kwargs)
    national, london = summary['passes']
    assert not national['reuses_job'] and summary['time_split'] == 'Weekly'
    assert {row['campaign'] for row in national['result']['scoring']} == {'2026-W07', '2026-W08'}
    assert {row['campaign'] for row in london['result']['scoring']} == {'2026-W07', '2026-W08'}
    # All selected CDRs: one value; the job split by campaign keeps its own result for itself.
    pooled = calculate_area_summary(scoring_engine.calculate_scoring, frames, ['Operator'],
                                    {'breakdown': None, 'cities': ['London'], 'time_split': 'All'}, kwargs)
    assert pooled['passes'][0]['reuses_job'] and pooled['passes'][0]['levels'] == ['Operator']
    assert normalize_area_summary_request({'cities': ['London']}, {})['time_split'] == 'Campaign'

    scenario = normalize_report_configuration({'scenarios': [default_scenario()]})['scenarios'][0]
    assert 'time_split' not in scenario and scenario['area_summary']['time_split'] == 'Campaign'
    scenario['area_summary']['time_split'] = 'Monthly'
    assert core._scenario_area_summary(None, scenario) == {'breakdown': 'Region', 'cities': ['London'],
                                                           'time_split': 'Monthly', 'campaigns': ['@all-campaigns']}


def test_the_scoring_page_calculates_and_shows_the_national_and_areas(scoring_api, monkeypatch):
    from types import ModuleType

    from src.modules import saved_filters

    client = scoring_api['client']
    repository = scoring_api['repository']
    repository.replace_scoring_configuration(_reliable_configuration())
    engine = ModuleType('src.modules.scoring')
    engine.METHOD_VERSION = 'scoring-area-summary-api-test'
    engine.required_input_columns = lambda kind, levels: [*levels, 'score']
    engine.calculate_scoring = lambda frames, levels, *, baseline_operator, configuration=None, **kwargs: (
        scoring_engine.calculate_scoring({'voice': pd.DataFrame(ROWS)}, levels, baseline_operator=baseline_operator,
                                         configuration=configuration, **kwargs))
    monkeypatch.setattr(saved_filters, '_scoring_engine', lambda: engine)
    request = {'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator'], 'nr_mode': 'NSA',
               'baseline_operator': 'EE', 'area_summary': {'time_split': 'All', 'breakdown': 'Region', 'cities': []}}
    # Area filters are allowed: the summary is scored on the whole country.
    filtered = client.post('/api/scoring/saved-filters', json={**request, 'context_filters': {'City': ['London']}})
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()['saved_filter']['area_summary'] == {'breakdown': 'Region', 'cities': [], 'time_split': 'All'}
    client.delete(f"/api/scoring/saved-filters/{filtered.json()['saved_filter']['id']}")
    created = client.post('/api/scoring/saved-filters', json=request)
    assert created.status_code == 200, created.text
    job_id = created.json()['saved_filter']['id']
    assert created.json()['saved_filter']['area_summary'] == {'breakdown': 'Region', 'cities': [], 'time_split': 'All'}
    saved_filters.calculate_saved_filter(repository, job_id)
    assert client.post('/api/scoring/saved-filters/match', json=request).json()['saved_filter']['id'] == job_id
    # A calculation without the summary is another job.
    assert client.post('/api/scoring/saved-filters/match', json={**request, 'area_summary': None}).json()['saved_filter'] is None
    result = client.get(f'/api/scoring/saved-filters/{job_id}').json()
    assert 'area_summary' not in result
    summaries = result['views']['insights']['area_summaries']
    combined = next(item for item in summaries if item['environment'] == 'Combined')
    assert [area['label'] for area in combined['areas']] == ['National', 'Vendor A', 'VMO2 North']
    reliable = result['scorings']['most_reliable']['views']['insights']['area_summaries']
    assert reliable and reliable[0]['areas'][0]['label'] == 'National'


def test_the_calculation_selection_keeps_the_national_and_area_summary(scoring_api):
    client = scoring_api['client']
    # Always included: each Region and London by default.
    assert client.get('/api/scoring/selection').json()['selection']['area_summary'] == {
        'enabled': True, 'time_split': 'Campaign', 'breakdown': 'Region', 'cities': ['London'], 'main_cities': False,
        'campaigns': ['@all-campaigns']}
    selection = client.get('/api/scoring/selection').json()['selection']
    summary = {'enabled': True, 'time_split': 'weekly', 'breakdown': 'cluster', 'cities': ['London', 'Leeds'],
               'campaigns': ['@latest-2-campaigns', '2026-Q1']}
    saved = client.put('/api/scoring/selection', json={'selection': {**selection, 'area_summary': summary}}).json()
    # The campaigns of the Scoring Trends are kept with them (an automatic choice alone).
    assert saved['selection']['area_summary'] == {'enabled': True, 'time_split': 'Weekly', 'breakdown': 'Cluster',
                                                  'cities': ['London', 'Leeds'], 'main_cities': False,
                                                  'campaigns': ['@latest-2-campaigns']}
    assert client.get('/api/scoring/selection').json()['selection']['area_summary']['time_split'] == 'Weekly'
    page = client.get('/scoring').text
    # The Geographical Areas & Scoring Trends: the areas of both tabs, then the campaigns and time split of the trends.
    for marker in ('data-scoring-trends', 'Geographical Areas &amp; Scoring Trends', 'data-scoring-area-breakdown',
                   'data-scoring-area-cities', 'data-scoring-trend-campaigns', 'data-scoring-area-time-split',
                   'data-result-tab="trends"', 'data-result-pane="trends"'):
        assert marker in page
    assert 'data-scoring-area-enabled' not in page and 'Include National' not in page
    assert page.index('data-scoring-area-cities') < page.index('data-scoring-trend-campaigns') < page.index('data-scoring-area-time-split')


def test_national_and_areas_trends_score_campaigns_of_their_own(scoring_api, monkeypatch):
    from types import ModuleType

    from src.modules import saved_filters

    client = scoring_api['client']
    repository = scoring_api['repository']
    repository.replace_scoring_configuration(_reliable_configuration())
    earlier = [scoring_api['add_ready_cdr'](name=f'UK_Q1_2026_NSA_{kind.title()}.csv', kind=kind, campaign='2026-Q1')
               for kind in ('data', 'voice', 'speech')]
    engine = ModuleType('src.modules.scoring')
    engine.METHOD_VERSION = 'scoring-area-trends-api-test'
    engine.required_input_columns = lambda kind, levels: [*levels, 'score']

    def calculate_scoring(frames, levels, *, baseline_operator, configuration=None, **kwargs):
        # The calls of every campaign of the CDRs read.
        campaigns = sorted({value for frame in frames.values() if 'Campaign' in frame for value in frame['Campaign']})
        rows = [{**row, 'Campaign': campaign} for campaign in campaigns for row in ROWS]
        return scoring_engine.calculate_scoring({'voice': pd.DataFrame(rows)}, levels, baseline_operator=baseline_operator,
                                                configuration=configuration, **kwargs)
    engine.calculate_scoring = calculate_scoring
    monkeypatch.setattr(saved_filters, '_scoring_engine', lambda: engine)
    request = {'dataset_ids': scoring_api['complete_dataset_ids'], 'aggregation_levels': ['Operator', 'Campaign'],
               'nr_mode': 'NSA', 'baseline_operator': 'EE', 'context_filters': {'Campaign': ['2026-Q2']},
               'area_summary': {'time_split': 'Campaign', 'breakdown': 'Region', 'cities': [], 'campaigns': ['@all-campaigns']}}
    created = client.post('/api/scoring/saved-filters', json=request)
    assert created.status_code == 200, created.text
    job_id = created.json()['saved_filter']['id']
    assert created.json()['saved_filter']['area_summary']['campaigns'] == ['@all-campaigns']
    assert sorted(created.json()['saved_filter']['trend_campaigns']) == ['2026-Q1', '2026-Q2']
    saved_filters.calculate_saved_filter(repository, job_id)
    stored = saved_filters.get_saved_filter(repository, job_id, include_internal_snapshot=True)
    assert stored['status'] == 'completed', stored['error']
    assert sorted(stored['area_trends']['dataset_ids']) == sorted([*earlier, *scoring_api['complete_dataset_ids']])
    result = client.get(f'/api/scoring/saved-filters/{job_id}').json()
    # The job keeps its own campaign; the trends score every campaign.
    assert {row['campaign'] for row in result['scoring']} == {'2026-Q2'}
    insights = result['views']['insights']
    combined = next(item for item in insights['area_summaries'] if item['environment'] == 'Combined')
    assert combined['campaigns'] == ['2026-Q1', '2026-Q2']
    assert [area['label'] for area in combined['areas']] == ['National', 'Vendor A', 'VMO2 North']
    # The scoring trend follows the campaigns of the trends, from two campaigns on.
    trend = next(item for item in insights['campaign_trends'] if item['environment'] == 'Combined')
    assert len(trend['campaigns']) == 2 and len(trend['series']) == 2 and all(None not in series['points'] for series in trend['series'])
    reliable = result['scorings']['most_reliable']['views']['insights']
    assert reliable['campaign_trends'] and reliable['area_summaries'][0]['campaigns'] == ['2026-Q1', '2026-Q2']
    # Without campaigns of its own a summary keeps the form (and the Saved Filter) of the summaries saved before.
    assert normalize_area_summary_request({'breakdown': 'Region', 'campaigns': []}, {}) == {
        'breakdown': 'Region', 'cities': [], 'time_split': 'Campaign'}
    # Campaigns saved alone (when the areas could be left out) take each Region and London; an automatic choice stands alone.
    assert normalize_area_summary_request({'breakdown': 'None', 'campaigns': ['2026-Q1', '@latest-2-campaigns']},
                                          {'City': ['London']}) == {
        'breakdown': 'Region', 'cities': ['London'], 'time_split': 'Campaign', 'campaigns': ['@latest-2-campaigns']}


def test_the_scoring_nationwide_comes_first_on_the_page_and_in_the_slides():
    from io import BytesIO
    from pathlib import Path

    from pptx import Presentation

    from src.modules.scoring_exports import export_scoring_powerpoint
    from tests.test_scoring_exports import TEMPLATE

    job, result, frames, kwargs = _calculation()
    result['geographical'] = calculate_geographical(scoring_engine.calculate_scoring, frames, job['levels'], kwargs)
    deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE)))
    titles = [slide.shapes.title.text_frame.paragraphs[0].text if slide.shapes.title is not None else '' for slide in deck.slides]
    assert titles.index('Best Network Scoring Nationwide') < titles.index('Best Network Scoring per Region')
    slide = deck.slides[titles.index('Best Network Scoring Nationwide')]
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    assert any(text.startswith('The Best Network scoring of the whole country') for text in texts)
    # On the page: first, across the panel and drawn at its width; the notes above the cards keep their room.
    root = Path(__file__).resolve().parents[1] / 'src/web_interface'
    script = (root / 'static/js/scoring.js').read_text(encoding='utf-8')
    template = (root / 'templates/scoring.html').read_text(encoding='utf-8')
    assert "group.level === 'National' ? `${scoringLabel()} Scoring Nationwide`" in script
    assert 'makeLocationCardChart(card, group.maximum, {wide: true, width: available})' in script
    assert '.scoring-location-cards.is-national { grid-template-columns: 1fr; }' in template
    assert '.scoring-insight-note.scoring-location-note { margin: 0 0 .75rem;' in template
