"""Summary, Geographical Areas, GAP by part and Points Lost insights of Scoring & GAP Analysis, and their slides.

They are calculated from synthetic NetCheck CDRs with NetCheck 2026: four operators in eight cities and on four
Connecting Roads routes over three campaigns, with the serving cell of every test and a few sites that fail often.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import zlib

import pytest
from pptx import Presentation

from src.modules.scoring_exports import export_scoring_powerpoint
from src.modules.scoring_insights import MAX_LEVERS, MAX_ROUTE_STRETCHES, MAX_SITES
from src.modules.scoring_views import build_scoring_views
from tests.scoring_cdr_synthetic import ROUTES, synthetic_cdr_result
from tests.test_scoring_exports import TEMPLATE


@pytest.fixture(scope='module')
def national():
    job, result = synthetic_cdr_result(levels=('Operator', 'Campaign'))
    return job, result, build_scoring_views(job, result)['insights']


@pytest.fixture(scope='module')
def cities():
    job, result = synthetic_cdr_result(levels=('Operator', 'City', 'Campaign'), campaigns=('2026-Q2', '2026-Q3'))
    return job, result, build_scoring_views(job, result)['insights']


def _combined(items, operator=None):
    return [item for item in items if item['environment'] == 'Combined' and (operator is None or item['operator'] == operator)]


def test_standings_rank_every_operator_and_measure_its_distances(national):
    _job, _result, insights = national
    standing = _combined(insights['standings'])[0]
    assert (standing['campaign'], standing['previous_campaign']) == ('2026-Q3', '2026-Q2')
    rows = standing['operators']
    assert [row['rank'] for row in rows] == [1, 2, 3, 4]
    assert [row['total'] for row in rows] == sorted((row['total'] for row in rows), reverse=True)
    reference = next(row for row in rows if row['operator'] == 'EE')
    assert reference['gap_to_reference'] is None
    for index, row in enumerate(rows):
        assert row['total'] == pytest.approx(row['data'] + row['voice'])
        assert row['delta'] == pytest.approx(row['total'] - row['previous'])
        assert row['takeaway'].startswith(row['label'])
        if index:
            assert row['gap_to_above'] == pytest.approx(row['total'] - rows[index - 1]['total'])
        if row['operator'] != 'EE':
            assert row['gap_to_reference'] == pytest.approx(row['total'] - reference['total'])
            # The GAP of every part of the scoring adds up to the GAP to the reference.
            assert sum(row['parts'].values()) == pytest.approx(row['gap_to_reference'], abs=.05)


def test_position_matrix_ranks_each_service_of_each_environment(national):
    _job, _result, insights = national
    matrix = _combined(insights['position_matrices'])[0]
    environments = [row['environment'] for row in matrix['rows']]
    assert environments[-1] == 'Combined' and len(environments) == 3
    for row in matrix['rows']:
        for service, cells in row['cells'].items():
            assert sorted(cell['rank'] for cell in cells.values()) == list(range(1, len(cells) + 1))
            for cell in cells.values():
                assert cell['percent'] == pytest.approx(cell['points'] / cell['maximum'] * 100)
                assert cell['missed'] == pytest.approx(cell['maximum'] - cell['points'])


def test_levers_are_the_kpis_behind_the_reference_that_win_most(national):
    _job, _result, insights = national
    levers = _combined(insights['levers'], 'Three UK')[0]
    gaps = [row['gap'] for row in levers['levers']]
    assert 0 < len(gaps) <= MAX_LEVERS and all(gap < 0 for gap in gaps) and gaps == sorted(gaps)
    assert levers['total_behind'] <= sum(gaps) + 1e-9


def test_campaign_scope_counts_the_tests_of_the_latest_campaign(national):
    _job, _result, insights = national
    scope = next(item for item in insights['campaign_scope'] if item['environment'] == 'Combined')
    assert scope['campaigns'] == ['2026-Q1', '2026-Q2', '2026-Q3'] and scope['latest_campaign'] == '2026-Q3'
    assert scope['window']['first'] == '2026-02-09' and scope['window']['days'] > 3
    # 8 cities × 6 positions and 4 routes × 9 positions, twice, with 7 Data tests and 4 calls each.
    assert scope['tests']['Vodafone UK']['data'] == {'Drive - City': 672, 'Drive - Connecting Roads': 504}
    assert scope['tests']['Vodafone UK']['voice'] == {'Drive - City': 384, 'Drive - Connecting Roads': 288}
    assert scope['places'] == {'cities': 8, 'routes': 4}


def test_gap_parts_add_up_to_the_total_gap_of_each_campaign(national):
    _job, _result, insights = national
    item = _combined(insights['gap_parts'], 'O2 UK')[0]
    assert item['by_reliability'] is True and item['campaigns'] == ['2026-Q1', '2026-Q2', '2026-Q3']
    assert [part['label'] for part in item['parts']] == ['Data reliability', 'Data quality', 'Voice reliability', 'Voice quality']
    for index, total in enumerate(item['totals']):
        assert sum(part['values'][index] for part in item['parts']) == pytest.approx(total)


def test_points_at_stake_share_every_point_lost(national):
    _job, _result, insights = national
    stake = _combined(insights['points_at_stake'])[0]
    assert [kind['key'] for kind in stake['kinds']] == ['failed', 'below_par', 'quality']
    for row in stake['operators']:
        assert sum(row['parts'].values()) == pytest.approx(row['lost'])
        assert row['lost'] + row['points'] == pytest.approx(stake['maximum'])


def test_sites_vendors_and_stretches_find_where_the_points_are_lost(national):
    _job, _result, insights = national
    sites = next(item for item in _combined(insights['sites_to_review'], 'Vodafone UK') if item['context'].get('campaign') == '2026-Q3')
    assert 0 < len(sites['rows']) <= MAX_SITES
    points = [row['points'] for row in sites['rows']]
    assert points == sorted(points, reverse=True)
    # The site of Leeds that fails many tests in every campaign is the one to review first.
    leeds = f"{zlib.crc32(b'Leeds') % 900 + 100}03"
    first = sites['rows'][0]
    assert (first['site'], first['city'], first['priority']) == (leeds, 'Leeds', 'High')
    assert [entry['campaign'] for entry in first['history']] == ['2026-Q1', '2026-Q2']
    vendors = _combined(insights['vendor_losses'], 'Vodafone UK')[-1]
    assert {row['name'] for row in vendors['rows']} <= {'Ericsson', 'Nokia'}
    assert sum(row['share'] for row in vendors['rows']) == pytest.approx(100, abs=.2)
    kinds = {row['name'] for item in _combined(insights['site_type_losses'], 'Three UK') for row in item['rows']}
    assert {'H3G', 'NNS · Vodafone'} <= kinds
    stretches = _combined(insights['route_stretches'], 'Vodafone UK')[-1]
    assert {route['route'] for route in stretches['routes']} <= set(ROUTES)
    assert stretches['routes'][0]['route'] == 'Manchester to Leeds'
    for route in stretches['routes']:
        shown = [row['points'] for row in route['stretches']]
        assert len(shown) <= MAX_ROUTE_STRETCHES and shown == sorted(shown, reverse=True)
        # The measured cities around the route name its map, and each numbered stretch its nearest city.
        assert route['landmarks'] and all({'name', 'latitude', 'longitude'} <= set(place) for place in route['landmarks'])
        assert all(row['label'] for row in route['stretches'])
    assert {'Manchester', 'Leeds'} <= {place['name'] for place in stretches['routes'][0]['landmarks']}


def test_the_route_stretch_map_colours_each_stretch_by_its_points_lost(national) -> None:
    from PIL import Image

    from src.modules.scoring_summary_slides import render_stretch_map, stretch_colour
    from src.modules.scoring_themes import THEMES
    scale = THEMES['ericsson']['scale']
    # From the third colour of the quality scale (little lost) through the second to the first (the most lost).
    assert (stretch_colour(0, scale), stretch_colour(.5, scale), stretch_colour(1, scale)) == (scale[2], scale[1], scale[0])
    _job, _result, insights = national
    route = _combined(insights['route_stretches'], 'Vodafone UK')[-1]['routes'][0]
    square = [[-3.0, 52.5], [-0.5, 52.5], [-0.5, 54.5], [-3.0, 54.5], [-3.0, 52.5]]
    image = Image.open(BytesIO(render_stretch_map(route, '#E60000', [], boundaries={'Area': [square]})))
    assert image.size == (1100, 900)
    # The map areas are the land: white over the grey sea.
    assert image.getpixel((5, 5)) == (233, 238, 243) and (255, 255, 255) in {colour for _count, colour in image.getcolors(1 << 20)}


def test_locations_are_ranked_and_compared_with_the_reference(cities):
    _job, _result, insights = cities
    rankings = _combined(insights['location_rankings'], 'Three UK')
    ranking = next(item for item in rankings if item['level'] == 'City')
    percents = [row['percent'] for row in ranking['rows']]
    assert len(percents) == 12 and percents == sorted(percents, reverse=True)
    for row in ranking['rows']:
        assert row['gap'] == pytest.approx(row['percent'] - row['reference_percent'])
        assert row['delta'] is not None and row['vendors']
    assert {row['kind'] for row in ranking['rows']} == {'place', 'route'}


def test_powerpoint_follows_the_order_of_the_results_tabs(national):
    job, result, _insights = national
    presentation = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE)))
    titles = [slide.shapes.title.text_frame.text.replace('\n', ' ') for slide in presentation.slides
              if slide.shapes.title is not None]

    def first(prefix):
        return next(index for index, title in enumerate(titles) if title.startswith(prefix))

    order = ['How the measurements ran', 'Where to act — ', 'Best Network — Standings', 'Position — per environment',
             'Biggest levers — vs EE', 'Best Network Scoring per Service', 'Scoring Tables — Summary',
             # The comparison of the campaigns closes the Scoring Tables, as on the page.
             'Best Network Campaign Comparison', 'GAP Analysis — ', 'GAP by part — vs EE',
             'KPI GAP Profile — ', 'Points Lost — where the points go', 'Points lost per KPI — ', 'Points Lost per City — ',
             'Points Lost — per Vendor', 'Points Lost — per kind of site', 'Sites to review — ', 'Route stretches — ']
    positions = [first(prefix) for prefix in order]
    assert positions == sorted(positions)


def test_the_route_of_where_to_act_is_drawn_as_the_routes_of_points_lost(national, monkeypatch) -> None:
    import src.modules.scoring_summary_slides as summary_slides
    from src.modules.scoring_act_slides import render_owner_route_map
    _job, _result, insights = national
    route = next(item for item in insights['route_focus'] if item['routes'])['routes'][0]
    calls = []
    original = summary_slides.render_stretch_map
    monkeypatch.setattr(summary_slides, 'render_stretch_map', lambda *args, **kwargs: calls.append(kwargs) or original(*args, **kwargs))
    from PIL import Image
    image = Image.open(BytesIO(render_owner_route_map(route, lambda _owner: '#4472C4', None)))
    # One map for both: its top stretches numbered, each owner as the ring of its circles, on a square for the slide.
    assert image.size == (1100, 1100)
    assert calls and calls[0]['ranked'] == route['top'] and calls[0]['measure'] == 'Points lost' and calls[0]['ring_color']
    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    # On the page: the same map, opened larger on a double-click, its stretches chosen in the table beside it, and the
    # operators of every view of the results in the order of the Operator Maps.
    assert "return makeStretchMap(route, '#64748b', [], {ranked: route.top, ringColor: ownerColor" in script
    assert "const routeMap = event.target.closest('svg.scoring-stretch-map');" in script
    assert 'function linkStretchTable(map, table, stretches) {' in script and 'body.append(linkStretchTable(map, stretchTable, route.top), side);' in script
    assert '.sort((left, right) => position(left?.operator) - position(right?.operator));' in script
    template = (Path(__file__).resolve().parents[1] / 'src/web_interface/templates/scoring.html').read_text(encoding='utf-8')
    assert '[data-scoring-theme="umlaut"] .scoring-insight-table th { background: #fff !important; color: #3f3f4c !important;' in template


def test_owners_take_their_colours_in_one_order_and_cards_without_a_breakdown_ask_to_recalculate() -> None:
    from src.modules.scoring_act_slides import owner_palette
    from src.modules.scoring_themes import THEMES, use_theme
    from src.modules.scoring_where_to_act import owner_cards, owner_order
    # The owners of the UK measurements first, Vendor B before Vendor A (their colours swapped), the others by name.
    assert owner_order(['Vendor A', 'Region X', 'Vendor B', 'VMO2 South', 'VMO2 North', '']) == [
        'VMO2 North', 'VMO2 South', 'Vendor B', 'Vendor A', 'Region X']
    with use_theme('ericsson'):
        color = owner_palette(['Vendor A', 'Vendor B', 'VMO2 North', 'VMO2 South'])
        categories = THEMES['ericsson']['categories']
        assert (color('Vendor B'), color('Vendor A'), color('Nobody')) == (categories[2], categories[3], '#9AA3AD')
    # Results calculated before the points of each owner were kept per City or route: no zeros, None.
    region = {'field': 'Region', 'operator': 'VF', 'label': 'VF', 'color': '#E60000', 'is_reference': False, 'environment': 'Combined',
              'context': {}, 'areas': [{'name': 'Vendor A', 'points': 5.0, 'share': 1.0}], 'total': 5.0}
    cards = owner_cards([region], {}, [], lambda _context: 1.0)[0]['cards']
    assert cards[0]['owner'] == 'Vendor A' and cards[0]['settings'] is None
    cards = owner_cards([{**region, 'city_losses': {'Vendor A': {'Leeds': 5.0}}}], {'Leeds': {'kind': 'place', 'setting': 'Cities'}},
                        [], lambda _context: 1.0)[0]['cards']
    assert cards[0]['settings'] == {'Cities': 5.0, 'Towns': 0.0, 'Roads and railways': 0.0}
