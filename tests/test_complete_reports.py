import json
import time
import zipfile
from collections import Counter
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

import src.DriveTestAnalyzer as core
from src.modules import complete_reports, pptx_compose
from src.modules import network_insights_export as ni
from src.modules import non_qualified_calls_export as nq_export
from src.modules import report_tasks
from src.modules.cdr_reporting import _remove_all_slides
from src.modules.pptx_to_docx import PackItem, pack_pages, packed_component_document
from src.modules.scoring_reports import normalize_report_configuration

TEMPLATE = Path(__file__).resolve().parents[1] / 'assets' / 'powerpoint-templates' / 'Template_01.pptx'
PNG = (b'\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00'
       b'\x0cIDATx\x9cc\xf8\xcf\xc0\x00\x00\x03\x01\x01\x00\xc9\xfe\x92\xef\x00\x00\x00\x00IEND\xaeB`\x82')


def login(client, username='super', password='super123'):
    client.cookies.clear()
    response = client.post('/login', data={'username': username, 'password': password}, follow_redirects=False)
    assert response.status_code == 303


def _deck(build) -> bytes:
    presentation = Presentation(TEMPLATE)
    _remove_all_slides(presentation)
    build(presentation)
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _table_pages(presentation, title='Serving Cells · By serving band', rows=30, panel='serving'):
    start = len(presentation.slides)
    ni._table_slides(presentation, title, ['Group', 'Band', 'Samples'], [[f'Op{index}', f'B{index}', '1,000'] for index in range(rows)], 'Scope')
    ni._tag_report_slides(presentation, start, panel)


def _small_table(presentation, title, panel):
    start = len(presentation.slides)
    slide = ni._slide(presentation, title, 'Scope')
    left, top, _width, _height = ni._content_frame(slide)
    ni._add_table(slide, ['Network', 'Sites'], [['EE', '10'], ['VF', '12']], left, top, 5)
    ni._tag_report_slides(presentation, start, panel)


def _unique_parts(content: bytes) -> bool:
    names = zipfile.ZipFile(BytesIO(content)).namelist()
    return not [name for name, count in Counter(names).items() if count > 1]


def test_definitions_keep_the_order_and_options_of_every_section():
    defaults = complete_reports.normalize_definition({})
    assert [section['key'] for section in defaults['sections']] == list(complete_reports.SECTION_KEYS)
    assert defaults['filters'] == {'nr_mode': 'NSA', 'operators': [], 'operator_vendors': [], 'vendor_operators': [], 'vendors': [],
                                   'regions': [], 'clusters': [], 'cities': [], 'main_cities': False}
    # Each section has its own Campaigns.
    assert all(section['options']['campaigns'] == [] and 'nr_mode' not in section['options'] for section in defaults['sections'])
    scoring = next(section for section in defaults['sections'] if section['key'] == 'scoring')
    assert [item['key'] for item in scoring['items']] == ['summary', 'charts', 'tables', 'locations', 'trends', 'gap', 'points']
    assert all(all(item['components'].values()) for item in scoring['items'])
    assert scoring['options']['scorings'] == ['best_network'] and scoring['options']['environments'] == ['all']
    # Saved order first, then what the report did not have yet; unknown keys are dropped. A report saved when the
    # Campaigns were a common filter gives them to the sections without their own.
    definition = complete_reports.normalize_definition({
        'title': '  Monthly  ', 'filters': {'nr_mode': 'sa', 'campaigns': ['@latest-campaign', ''], 'operators': ['EE', 'EE']},
        'sections': [
            {'key': 'network_insights', 'enabled': True, 'items': [
                {'key': 'spectrum', 'enabled': True, 'components': {'licensed': False}}, {'key': 'unknown'}],
             'options': {'technology': 'nr', 'group': ['vendor', 'nonsense'], 'overshoot_km': 500, 'campaigns': ['2026-Q2']}},
            {'key': 'dashboards', 'enabled': False, 'items': [{'key': 'd1', 'label': 'Weekly'}, {'key': 'd1'}]},
            {'key': 'bogus'},
        ]})
    assert definition['title'] == 'Monthly'
    assert definition['filters']['nr_mode'] == 'SA' and definition['filters']['operators'] == ['EE']
    assert 'campaigns' not in definition['filters']
    campaigns = {section['key']: section['options']['campaigns'] for section in definition['sections']}
    assert campaigns == {'network_insights': ['2026-Q2'], 'dashboards': ['@latest-campaign'], 'scoring': ['@latest-campaign'],
                         'non_qualified_calls': ['@latest-campaign']}
    # A report saved when the NR Mode was an option of each section takes that of the first one.
    assert complete_reports.normalize_definition({'sections': [{'key': 'scoring', 'options': {'nr_mode': 'SA'}}]})['filters']['nr_mode'] == 'SA'
    assert [section['key'] for section in definition['sections']] == ['network_insights', 'dashboards', 'scoring', 'non_qualified_calls']
    network = definition['sections'][0]
    assert network['items'][0]['key'] == 'spectrum' and network['items'][0]['components']['licensed'] is False
    assert network['items'][0]['components']['observed'] is True
    assert len(network['items']) == len(ni.REPORT_PANELS)
    assert network['options']['technology'] == 'nr' and network['options']['group'] == ['vendor']
    assert network['options']['overshoot_km'] == 100
    assert definition['sections'][1]['items'] == [{'key': 'd1', 'label': 'Weekly', 'enabled': True, 'nr_mode': ''}]
    chosen = [section['key'] for section in complete_reports.enabled_sections(definition)]
    # Dashboards is unchecked; Scoring, Network Insights and Non-Qualified Calls have checked sub-modules.
    assert chosen == ['network_insights', 'scoring', 'non_qualified_calls']


def test_reports_saved_before_the_scoring_trends_keep_their_trend_after_the_locations():
    definition = complete_reports.normalize_definition({'sections': [{'key': 'scoring', 'items': [
        {'key': 'locations', 'components': {}}, {'key': 'charts', 'components': {'trend': False}}],
        'options': {'area_summary': {'enabled': True, 'campaigns': ['@all-campaigns', '2026-Q1']}}}]})
    scoring = next(section for section in definition['sections'] if section['key'] == 'scoring')
    keys = [item['key'] for item in scoring['items']]
    assert keys[:3] == ['locations', 'trends', 'charts']
    trends = next(item for item in scoring['items'] if item['key'] == 'trends')
    assert trends['components'] == {'trend': False, 'areas': True}
    # The campaigns of the Scoring Trends: an automatic choice stands alone.
    assert scoring['options']['area_summary']['campaigns'] == ['@all-campaigns']
    lines = complete_reports._option_lines(scoring, None, None)
    assert lines[-1] == 'Scoring Trends: All available campaigns · National & Areas Trends: each Region, time split Campaign'


def test_section_lines_and_filters_describe_the_report():
    definition = complete_reports.normalize_definition({'subtitle': ' Q3 benchmark ', 'filters': {'regions': ['North'], 'main_cities': True}})
    assert definition['subtitle'] == 'Q3 benchmark'
    assert complete_reports.filter_lines(definition) == ['NR Mode: NSA', 'Regions: North · Main Cities']
    # The covers name every common filter, All when it has no value.
    assert complete_reports.common_filter_items(definition) == [
        'NR Mode: NSA', 'Operators: All', 'Operator_Vendor: All', 'Vendor_Operator: All', 'Vendors: All', 'Regions: North',
        'Clusters: All', 'Cities: Main Cities']
    scoring = next(section for section in definition['sections'] if section['key'] == 'scoring')
    scoring['options'].update(scorings=['best_network', 'most_reliable'], environments=['all', 'Cities'], campaigns=['@latest-campaign'])
    lines = complete_reports.section_lines(scoring)
    assert lines[0] == 'Campaigns: Latest campaign'
    assert lines[1] == 'Best Network, Most Reliable Network · Environments: All Environments, Cities'


def test_scoring_scenarios_keep_their_chosen_environments():
    configuration = normalize_report_configuration({'scenarios': [{'name': 'A', 'scorings': {
        'best_network': {'enabled': True, 'environments': 'split', 'environment_names': ['All Environments', 'Cities']}}}]})
    options = configuration['scenarios'][0]['scorings']['best_network']
    assert options['environment_names'] == ['All Environments', 'Cities']
    # Configurations saved before have none.
    assert 'environment_names' not in configuration['scenarios'][0]['scorings']['most_reliable']


def test_network_insights_slides_are_named_after_their_panel_and_component():
    assert ni.report_component('Serving Cells · By serving band') == 'serving.breakdowns'
    assert ni.report_component('Serving Cells · NR cell of the NSA samples') == 'serving.nr_leg'
    assert ni.report_component('5G Availability · Sites with the most samples without NR') == 'nr_availability.sites_without_nr'
    assert ni.report_component('RF Quality · LTE') == 'rf_quality.cdfs'
    assert ni.report_component('Bandwidth of each test: LTE + NR') == 'spectrum.bandwidth'
    assert ni.report_component('Under 25 MHz by location (Download tests)') == 'spectrum.bandwidth_splits'
    keys = {f'{panel}.{component}' for panel, _title, components in ni.REPORT_PANELS for component, _label in components}
    assert {key for _start, key in ni._COMPONENT_TITLES if not key.startswith('selection.')} <= keys


def test_network_insights_export_keeps_only_the_chosen_components(tmp_path):
    analysis = {'serving': {'overshoot_km': 5, 'summary': [{'operator': 'EE', 'samples': 10, 'matched_share': 90, 'median_km': 1,
                                                           'p90_km': 2, 'overshoot_share': 1, 'hosted_elsewhere_share': 0}],
                            'overshooting_cells': [{'operator': 'EE', 'cell': 'C1', 'site': 'S1', 'band': 'B1', 'vendor': 'V',
                                                    'network': 'EE', 'host': 'EE', 'samples_beyond': 1, 'samples': 2,
                                                    'median_km': 1, 'max_km': 6}]},
                'sections': [{}], 'group_label': 'Operator'}
    full, chosen = tmp_path / 'full.pptx', tmp_path / 'chosen.pptx'
    ni.export_network_insights_powerpoint(full, analysis, {}, {}, lambda *_args: b'', template=TEMPLATE)
    ni.export_network_insights_powerpoint(chosen, analysis, {}, {}, lambda *_args: b'', template=TEMPLATE,
                                          components=['serving.overshooting'])
    titles = [pptx_compose.title_parts(slide)[0] for slide in Presentation(chosen).slides]
    assert 'Serving Cells · Overshooting cells' in titles
    assert 'Serving Cells · Distance and overshooting' not in titles
    assert len(Presentation(full).slides) > len(Presentation(chosen).slides)
    tags = {pptx_compose.title_parts(slide)[0]: pptx_compose.slide_tag(slide) for slide in Presentation(chosen).slides}
    assert tags['Serving Cells · Overshooting cells'] == 'serving' and tags['Analysis Selection'] == 'selection'


def test_non_qualified_calls_slides_keep_the_chosen_components(tmp_path, monkeypatch):
    image = BytesIO(PNG)
    panels = [('summary', 'executive', 'Executive Summary', image), ('root', 'netcheck', 'Root Cause Analysis', BytesIO(PNG)),
              ('root', 'analysts', 'Root Cause Analysis', BytesIO(PNG))]
    monkeypatch.setattr(nq_export, 'keyed_report_panels', lambda *_args: panels)
    destination = tmp_path / 'nq.pptx'
    nq_export.export_powerpoint(destination, ['Every call'], {}, [], {'granularity': 'week'}, {}, template=TEMPLATE,
                                components=['summary.executive', 'root.analysts'])
    slides = list(Presentation(destination).slides)
    assert [pptx_compose.slide_tag(slide) for slide in slides] == ['cover', 'summary', 'root', '']
    assert nq_export._failure_component('Voice · Technology and classification') == 'technology'
    keys = {key for _panel, _title, components in nq_export.REPORT_PANELS for key, _label in components}
    assert {'executive', 'netcheck', 'script', 'analysts', 'pipeline', 'details', 'activity'} <= keys


def test_slides_are_copied_with_their_charts_and_pictures_into_one_deck():
    def build(presentation):
        _table_pages(presentation)
        slide = ni._slide(presentation, 'Chart slide')
        data = CategoryChartData()
        data.categories = ['a', 'b']
        data.add_series('S', (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data)
        slide.shapes.add_picture(BytesIO(PNG), Inches(8), Inches(2), Inches(1), Inches(1))
        slide.notes_slide.notes_text_frame.text = 'Speaker notes'

    source = pptx_compose.deck(_deck(build))
    composer = pptx_compose.DeckComposer(TEMPLATE)
    composer.add_layout_slide('Title Page')
    copies = [composer.add_slide(slide) for slide in source.slides]
    # Slides added by python-pptx between the copies keep their own part names.
    composer.add_layout_slide('Title Page')
    composer.add_slide(source.slides[0])
    composer.remove_slide(copies[1])
    composer.add_layout_slide('Black logo end slide')
    content = composer.save()
    assert _unique_parts(content)
    reopened = Presentation(BytesIO(content))
    chart_slide = next(slide for slide in reopened.slides if pptx_compose.title_parts(slide)[0] == 'Chart slide')
    assert any(shape.has_chart for shape in chart_slide.shapes)
    assert chart_slide.shapes[-1].shape_type == 13  # The picture.
    assert chart_slide.notes_slide.notes_text_frame.text == 'Speaker notes'
    assert [slide.slide_layout.name for slide in reopened.slides][0] == 'Title Page'


def test_a_table_spread_over_several_slides_becomes_one_table_on_one_slide():
    source = pptx_compose.deck(_deck(_table_pages))
    assert len(source.slides) == 3 and all(pptx_compose.is_continued(slide) for slide in source.slides)
    composer = pptx_compose.DeckComposer(TEMPLATE)
    pages = [composer.add_slide(slide) for slide in source.slides]
    assert pptx_compose.join_table_pages(composer, pages)
    assert len(composer.presentation.slides) == 1
    table = next(shape for shape in pptx_compose.content_shapes(pages[0]) if shape.has_table).table
    # One header and every row.
    assert len(table.rows) == 31 and table.cell(0, 0).text == 'Group' and table.cell(30, 0).text == 'Op29'
    scale = pptx_compose.merge_onto_slide(composer, pages[0], [pptx_compose.Block(pages[0], pptx_compose.content_shapes(pages[0]))],
                                          captions=False)
    assert 0.3 < scale < 1
    shape = pptx_compose.content_shapes(pages[0])[0]
    assert shape.top + shape.height <= composer.height


def test_wide_tables_join_their_parts_and_long_tables_continue_side_by_side():
    columns = ['Operator → Campaign', *[f'Location number {index}' for index in range(14)]]
    source = pptx_compose.deck(_deck(lambda presentation: ni._table_slides(
        presentation, 'Under 25 MHz', columns, [[f'VF · Q{row}', *[f'{row + index}%' for index in range(14)]] for row in range(12)], 'Scope')))
    assert len(source.slides) == 2
    composer = pptx_compose.DeckComposer(TEMPLATE)
    pages = [composer.add_slide(slide) for slide in source.slides]
    assert pptx_compose.join_table_pages(composer, pages)
    table = pptx_compose.content_shapes(pages[0])[0].table
    # The second part adds its columns (without repeating the first one) to the same rows.
    assert [table.cell(0, index).text for index in range(len(table.columns))] == columns
    assert len(table.rows) == 13 and table.cell(12, 14).text == '24%'
    frame = pptx_compose.content_frame(pages[0], composer.width, composer.height, [pptx_compose.bounds(pptx_compose.content_shapes(pages[0]))])
    assert pptx_compose.fit_table(pages[0], frame) > 0.9
    shape = pptx_compose.content_shapes(pages[0])[0]
    assert shape.left + shape.width <= frame[0] + frame[2] + 10 and shape.top + shape.height <= frame[1] + frame[3]

    source = pptx_compose.deck(_deck(lambda presentation: _table_pages(presentation, rows=70)))
    pages = [composer.add_slide(slide) for slide in source.slides]
    assert pptx_compose.join_table_pages(composer, pages)
    frame = pptx_compose.content_frame(pages[0], composer.width, composer.height, [pptx_compose.bounds(pptx_compose.content_shapes(pages[0]))])
    scale = pptx_compose.fit_table(pages[0], frame)
    tables = pptx_compose.content_shapes(pages[0])
    # The rows continue in tables beside the first one, each under the header, at a larger size than in one table.
    assert len(tables) > 1 and scale > 0.5
    assert all(shape.table.cell(0, 0).text == 'Group' for shape in tables)
    assert [shape.table.cell(row, 0).text for shape in tables for row in range(1, len(shape.table.rows))] == [f'Op{index}' for index in range(70)]
    assert all(shape.top + shape.height <= frame[1] + frame[3] and shape.left + shape.width <= frame[0] + frame[2] + 10 for shape in tables)
    assert len({shape.shape_id for shape in tables}) == len(tables)


def test_a_table_too_long_for_a_legible_font_continues_on_more_slides():
    def build(presentation):
        _table_pages(presentation, rows=400)

    section = complete_reports.SectionDeck('network_insights', 'Network Insights', [], decks=[(None, _deck(build))], order=['serving'])
    component = complete_reports._components(complete_reports._section_entries(section))[0]
    composer = pptx_compose.DeckComposer(TEMPLATE)
    pages = complete_reports._place_component(composer, component)
    assert len(pages) > 1 and len(composer.presentation.slides) == len(pages)
    titles = [pptx_compose.title_parts(page)[0] for page in pages]
    assert titles == [f'Serving Cells · By serving band · {number}/{len(pages)}' for number in range(1, len(pages) + 1)]
    # Every row once, in its order, each slide under the header and with values of a legible size.
    tables = [shape.table for page in pages for shape in pptx_compose.content_shapes(page)]
    assert all(table.cell(0, 0).text == 'Group' for table in tables)
    assert [table.cell(row, 0).text for table in tables for row in range(1, len(table.rows))] == [f'Op{index}' for index in range(400)]
    sizes = [run.font.size.pt for table in tables for row in list(table.rows)[1:] for cell in row.cells
             for paragraph in cell.text_frame.paragraphs for run in paragraph.runs if run.font.size]
    assert min(sizes) >= pptx_compose.LEGIBLE_FONT / 100
    # The slides of a table that continues are not compacted with other components.
    content, _outline, _covers = complete_reports.compose_powerpoint(TEMPLATE, 'Report', '', [], [section], {})
    assert len(Presentation(BytesIO(content)).slides) == 2 + len(pages) + 1


def test_pages_that_would_be_too_small_on_one_slide_continue_on_more_slides():
    def pages(width, *heights):
        def build(presentation):
            for number, height in enumerate(heights, start=1):
                slide = ni._slide(presentation, f'Sites · {number}/{len(heights)}', 'Scope')
                left, top, _width, _height = ni._content_frame(slide)
                slide.shapes.add_picture(BytesIO(PNG), Inches(left), Inches(top), Inches(width), Inches(height))
        return [composer.add_slide(slide) for slide in pptx_compose.deck(_deck(build)).slides]

    composer = pptx_compose.DeckComposer(TEMPLATE)
    minimum = complete_reports.PAGE_MINIMUM_SCALE
    # Small pages share one slide; pages as large as the slide each keep a slide of their own, and a short last page
    # joins the one before it; pages of the same size share the slides evenly.
    assert pptx_compose.page_groups(composer, pages(5, 2, 2, 2, 2), minimum) == [[0, 1, 2, 3]]
    assert pptx_compose.page_groups(composer, pages(12, 4.5, 4.5, 4.5, 4.5), minimum) == [[0], [1], [2], [3]]
    assert pptx_compose.page_groups(composer, pages(12, 4.5, 4.5, 4.5, 0.4), minimum) == [[0], [1], [2, 3]]
    assert pptx_compose.page_groups(composer, pages(12, 2.2, 2.2, 2.2, 2.2), minimum) == [[0, 1], [2, 3]]


def test_tables_side_by_side_continue_together_on_as_few_slides_as_keep_them_legible():
    counts = {'VF': 110, 'Three': 60, 'VMO2': 14}

    def build(presentation):
        start = len(presentation.slides)
        ni._paired_table_slides(presentation, 'Network Deployment · Cluster', [
            (f'{operator} (2026-09-17)', ['Cluster', 'Sites', 'Cells'], [[f'{operator}_Cluster{index:03d}', '10', '1,000'] for index in range(rows)])
            for operator, rows in counts.items()], 'Scope')
        ni._tag_report_slides(presentation, start, 'deployment')

    section = complete_reports.SectionDeck('network_insights', 'Network Insights', [], decks=[(None, _deck(build))], order=['deployment'])
    component = complete_reports._components(complete_reports._section_entries(section))[0]
    assert len(component.slides) == 10
    composer = pptx_compose.DeckComposer(TEMPLATE)
    pages = complete_reports._place_component(composer, component)
    assert 1 < len(pages) < 10
    rows = {operator: [] for operator in counts}
    for page in pages:
        captions = {shape.text_frame.text for shape in pptx_compose.content_shapes(page) if not shape.has_table and shape.has_text_frame}
        for shape in (shape for shape in pptx_compose.content_shapes(page) if shape.has_table):
            names = [shape.table.cell(row, 0).text for row in range(1, len(shape.table.rows))]
            operator = names[0].split('_')[0]
            rows[operator].extend(names)
            assert shape.table.cell(0, 0).text == 'Cluster' and f'{operator} (2026-09-17)' in captions
        # A table that has run out of rows leaves its place empty, with its caption.
        assert len(captions) == len([shape for shape in pptx_compose.content_shapes(page) if shape.has_table])
    assert rows == {operator: [f'{operator}_Cluster{index:03d}' for index in range(count)] for operator, count in counts.items()}


def test_components_join_the_pages_of_a_chart_or_table_but_not_slides_of_the_same_title():
    def build(presentation):
        _small_table(presentation, 'RF Quality', 'rf_quality')
        _small_table(presentation, 'RF Quality · 2/2', 'rf_quality')
        _small_table(presentation, 'Position', 'rf_quality')
        _small_table(presentation, 'Position', 'rf_quality')

    section = complete_reports.SectionDeck('network_insights', 'Network Insights', [], decks=[(None, _deck(build))], order=['rf_quality'])
    components = complete_reports._components(complete_reports._section_entries(section))
    assert [len(component.slides) for component in components] == [2, 1, 1]
    assert components[0].title == ('RF Quality', 'Scope')


def test_sections_follow_the_order_of_their_sub_modules_inside_each_block():
    def build(presentation):
        cover = presentation.slides.add_slide(next(layout for layout in presentation.slide_layouts if layout.name == 'Title Page'))
        pptx_compose.tag(cover, 'cover')
        _small_table(presentation, 'Spectrum table', 'spectrum')
        _small_table(presentation, 'Overview cards', 'overview')
        _small_table(presentation, 'Snapshot table', 'snapshot')

    section = complete_reports.SectionDeck('network_insights', 'Network Insights', [], decks=[(None, _deck(build))],
                                           order=['snapshot', 'spectrum', 'overview'])
    entries = complete_reports._section_entries(section)
    # The report cover of the module is left out.
    assert [entry.item for entry in entries] == ['snapshot', 'spectrum', 'overview']


def test_the_powerpoint_has_a_cover_a_cover_per_section_and_compacts_small_components():
    def build(presentation):
        _table_pages(presentation)
        _small_table(presentation, 'Network Snapshot', 'snapshot')
        _small_table(presentation, 'Network Snapshot · Inventory quality', 'snapshot')

    section = complete_reports.SectionDeck('network_insights', 'Network Insights', ['Technology: LTE'], decks=[(None, _deck(build))],
                                           order=['serving', 'snapshot'], item_titles={'snapshot': 'Network Snapshot'})
    content, outline, covers = complete_reports.compose_powerpoint(
        TEMPLATE, 'Monthly', 'Network benchmark', ['NR Mode: NSA', 'Operators: All'], [section], {}, 'Generated today')
    assert _unique_parts(content)
    titles = [item['title'] for item in outline]
    # Cover, section cover, the joined table, both small tables on one slide and the closing slide.
    assert titles[:4] == ['Monthly', 'Network Insights', 'Serving Cells · By serving band', 'Network Snapshot']
    assert len(titles) == 5 and covers == [0, 1]
    deck = Presentation(BytesIO(content))
    cover, section_cover = deck.slides[0], deck.slides[1]
    # The report cover: a large title, the subtitle and every common filter; the section cover adds its options.
    cover_title = next(shape for shape in cover.placeholders if shape.placeholder_format.type in {1, 3})
    assert cover_title.text_frame.paragraphs[0].runs[0].font.size.pt == complete_reports.COVER_TITLE_PT
    cover_texts = [shape.text_frame.text for shape in cover.shapes if shape.has_text_frame]
    assert 'Network benchmark' in cover_texts
    assert 'Common filters\nNR Mode: NSA · Operators: All' in cover_texts and 'Generated today' in cover_texts
    section_texts = [shape.text_frame.text for shape in section_cover.shapes if shape.has_text_frame]
    assert 'Technology: LTE' in section_texts and 'Common filters\nNR Mode: NSA · Operators: All' in section_texts
    snapshot = deck.slides[3]
    texts = [shape.text_frame.text for shape in snapshot.shapes if shape.has_text_frame]
    assert any(text.startswith('Network Snapshot · Inventory quality') for text in texts)
    # A caption does not repeat the title of the slide: the subtitle of the first table is its caption then.
    assert 'Scope' in texts and not any(text.startswith('Network Snapshot\n') for text in texts)


def _item(width, height):
    return PackItem('image', image=PNG, width=width, height=height)


def test_word_pages_take_consecutive_components_of_one_section_without_splitting_them():
    items = [PackItem('heading1', 'Scoring'), _item(3.4, 4), _item(7, 2.5), _item(3.4, 4), _item(7, 5),
             PackItem('heading1', 'Network Insights'), _item(7, 14), _item(7, 2), PackItem('heading2', 'Best Network'), _item(7, 9)]
    pages = pack_pages(items, 7.37, 10.4)
    flattened = [item for rows, _scale in pages for row in rows for item in row]
    assert sorted(map(id, flattened)) == sorted(map(id, items))
    order = {id(item): position for position, item in enumerate(items)}
    page_ranges = [sorted(order[id(item)] for row in rows for item in row) for rows, _scale in pages]
    # Every page holds a run of consecutive items (reordered inside the page at most).
    assert all(numbers == list(range(numbers[0], numbers[-1] + 1)) for numbers in page_ranges)
    assert all(earlier[-1] < later[0] for earlier, later in zip(page_ranges, page_ranges[1:]))
    # A section starts a new page; a heading never ends a page.
    assert [rows[0][0].kind for rows, _scale in pages if rows[0][0].kind == 'heading1'] == ['heading1', 'heading1']
    assert all(rows[-1][0].kind == 'image' for rows, _scale in pages)
    # The two narrow components of the first page share a row (1, 3, 2): reordered inside the page, they fit at full size.
    first_rows, first_scale = pages[0]
    assert [len(row) for row in first_rows] == [1, 2, 1] and first_scale == 1
    # One taller than the page fills a page on its own, smaller.
    assert any(len(row) == 2 for rows, _scale in pages for row in rows)
    tall = next((rows, scale) for rows, scale in pages if any(item.height == 14 for row in rows for item in row))
    assert tall[1] < 1
    document = packed_component_document(items)
    assert document[:2] == b'PK'


def test_favourites_are_saved_listed_and_travel_as_their_own_element(client):
    login(client)
    options = client.get('/api/reporting/complete-report/options')
    assert options.status_code == 200, options.text
    payload = options.json()
    assert payload['favourites'] == [] and payload['can_edit'] is True
    assert [section['key'] for section in payload['catalogue']['sections']] == list(complete_reports.SECTION_KEYS)
    definition = payload['defaults']
    definition['sections'][0]['items'] = [{'key': 'd1', 'label': 'Weekly Dashboard', 'enabled': True}]
    created = client.post('/api/reporting/complete-report/favourites', json={'name': 'Monthly', 'definition': definition})
    assert created.status_code == 200, created.text
    favourite = created.json()['favourite']
    assert client.post('/api/reporting/complete-report/favourites', json={'name': 'monthly', 'definition': definition}).status_code == 400
    renamed = client.put(f"/api/reporting/complete-report/favourites/{favourite['id']}", json={'name': 'Monthly report', 'definition': definition})
    assert renamed.status_code == 200 and renamed.json()['favourite']['name'] == 'Monthly report'
    actions = [row['action'] for row in core.repository.list_logs(limit=20)]
    assert 'complete_report_favourite_saved' in actions
    # The favourites travel as the Complete Report Favourites element (not with the Reporting Jobs), and are added or
    # replaced by name.
    assert 'complete_reports' not in json.loads(report_tasks.export_tasks_document(core.repository))
    document = json.loads(complete_reports.export_favourites_document(core.repository))
    assert document['format'] == complete_reports.FAVOURITES_FORMAT
    assert [item['name'] for item in document['favourites']] == ['Monthly report']
    assert client.delete(f"/api/reporting/complete-report/favourites/{favourite['id']}").status_code == 200
    assert complete_reports.list_favourites(core.repository) == []
    assert core._restore_workspace_complete_report_favourites(core.active_workspace, json.dumps(document).encode(), 'super') == 1
    restored = complete_reports.list_favourites(core.repository)
    assert [item['name'] for item in restored] == ['Monthly report']
    assert restored[0]['definition']['sections'][0]['items'][0]['label'] == 'Weekly Dashboard'
    assert 'complete_report_favourite_deleted' in [row['action'] for row in core.repository.list_logs(limit=20)]
    # Reporting Jobs documents that carried them still bring them.
    complete_reports.delete_favourite(core.repository, restored[0]['id'])
    jobs = {'format': 'drivetest-analyzer-reporting-jobs', 'version': 1, 'reporting_jobs': [],
            'complete_reports': document['favourites']}
    core._restore_workspace_reporting_jobs(core.active_workspace, json.dumps(jobs).encode(), 'super')
    assert [item['name'] for item in complete_reports.list_favourites(core.repository)] == ['Monthly report']
    # The element, the workspace packages and the backups carry them.
    assert 'complete-report-favourites' in core.WORKSPACE_ELEMENT_EXPORT_TARGETS
    assert core.archive_workspace_components_for_target('complete-report-favourites') == ['complete_report_favourites']
    assert 'complete_report_favourites' in core.full_workspace_archive_components()
    output = BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        core._archive_workspace_complete_report_favourites(archive, core.active_workspace, 'workspaces/Default')
    assert zipfile.ZipFile(BytesIO(output.getvalue())).namelist() == [f'workspaces/Default/{core.CR_FAVOURITES_ARCHIVE_NAME}']
    page = client.get('/admin').text
    assert 'Complete Report Favourites (from active workspace)' in page and 'Scoring &amp; GAP Favourites Filters (from active workspace)' in page


def test_word_pages_do_not_repeat_the_title_of_components_that_share_it():
    def component(group, height=3.0):
        return PackItem('image', 'Biggest levers', image=PNG, width=7, height=height, group=group, untitled=PNG,
                        untitled_width=7, untitled_height=height - 0.5)

    items = [PackItem('heading1', 'Scoring'), component('levers'), component('levers'), component('other'), component('levers')]
    shown = [item for rows, _scale in pack_pages(items, 7.37, 10.4) for row in rows for item in row if item.kind == 'image']
    # On the same page, a component after one with the same title and context goes without its title.
    assert [item.height for item in shown] == [3.0, 2.5, 3.0, 3.0]
    # One that starts a page shows its title.
    pages = pack_pages([component('levers', 7), component('levers', 7)], 7.37, 10.4)
    assert [[item.height for row in rows for item in row] for rows, _scale in pages] == [[7], [7]]


def test_previews_on_linux_use_installed_fonts_instead_of_missing_ones(monkeypatch):
    from pptx import Presentation as NewPresentation
    from src.modules import pptx_to_docx

    presentation = NewPresentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(0, 0, 914400, 914400).text_frame.paragraphs[0]
    for text, font in (('GAP Analysis', 'Ericsson Hilda Light'), (' vs EE', 'Arial')):
        run = paragraph.add_run()
        run.text, run.font.name = text, font
    output = BytesIO()
    presentation.save(output)
    content = output.getvalue()
    # Where LibreOffice finds stand-ins itself (macOS, Windows), the document is converted as it is.
    monkeypatch.setattr(pptx_to_docx, 'installed_fonts', lambda: None)
    assert pptx_to_docx.with_installed_fonts(content) == content
    # On Linux a missing font would come out with spaced letters: an installed one stands in for it.
    monkeypatch.setattr(pptx_to_docx, 'installed_fonts', lambda: frozenset({'liberation sans', 'arial'}))
    converted = NewPresentation(BytesIO(pptx_to_docx.with_installed_fonts(content)))
    runs = converted.slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert [run.font.name for run in runs] == ['Liberation Sans', 'Arial']
    # The fonts of the theme too; the references to them are kept.
    converted_zip = zipfile.ZipFile(BytesIO(pptx_to_docx.with_installed_fonts(content)))
    theme = next(converted_zip.read(name) for name in converted_zip.namelist() if name.startswith('ppt/theme/'))
    assert b'typeface="Liberation Sans"' in theme and b'typeface="Calibri"' not in theme
    # The Docker image installs LibreOffice Writer too, which renders the previews of the Word documents.
    dockerfile = (core.PROJECT_ROOT / 'docker/dockerfile').read_text(encoding='utf-8')
    assert 'libreoffice-impress-nogui' in dockerfile and 'libreoffice-writer-nogui' in dockerfile


def test_long_decks_are_converted_by_several_libreoffice_runs_at_once(monkeypatch):
    import pymupdf
    from pptx import Presentation as NewPresentation
    from src.modules import pptx_to_docx

    presentation = NewPresentation()
    for number in range(20):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 914400, 914400).text_frame.text = f'Slide {number + 1}'
    output = BytesIO()
    presentation.save(output)
    runs = []

    def convert(soffice, content, suffix, slides):
        # Each run gets consecutive slides and a time limit that grows with them.
        deck = NewPresentation(BytesIO(content))
        texts = [slide.shapes[0].text_frame.text for slide in deck.slides]
        runs.append((texts, slides))
        with pymupdf.open() as document:
            for text in texts:
                document.new_page().insert_text((72, 72), text)
            return document.tobytes()

    monkeypatch.setattr(pptx_to_docx, '_soffice', lambda: '/usr/bin/soffice')
    monkeypatch.setattr(pptx_to_docx, '_convert_to_pdf', convert)
    monkeypatch.setattr(pptx_to_docx, 'libreoffice_workers', lambda: 3)
    pdf = pptx_to_docx.document_pdf(output.getvalue(), 'pptx')
    assert len(runs) == 3 and all(slides == len(texts) for texts, slides in runs)
    with pymupdf.open(stream=pdf, filetype='pdf') as document:
        # The PDFs of the runs are joined in the order of the deck.
        assert [document[index].get_text().strip() for index in range(document.page_count)] == [f'Slide {n + 1}' for n in range(20)]
    # More processors than runs of six slides: the runs keep six slides at least.
    runs.clear()
    monkeypatch.setattr(pptx_to_docx, 'libreoffice_workers', lambda: 8)
    pptx_to_docx.document_pdf(output.getvalue(), 'pptx')
    assert len(runs) == 3  # Twenty slides give three runs of at least six.
    # The number of runs at once can be set for the machine.
    monkeypatch.undo()
    monkeypatch.setenv('DRIVETEST_ANALYZER_LIBREOFFICE_WORKERS', '2')
    assert pptx_to_docx.libreoffice_workers() == 2


def test_a_report_being_generated_can_be_cancelled_and_tells_its_processing_time(client, monkeypatch):
    login(client)
    started = __import__('threading').Event()
    release = __import__('threading').Event()

    def slow_build(*args, progress=lambda *_: None, **kwargs):
        started.set()
        release.wait(5)
        progress(50, 'Composing the PowerPoint')
        raise AssertionError('A cancelled report goes no further.')

    monkeypatch.setattr(complete_reports, 'build_report', slow_build)
    definition = complete_reports.normalize_definition({})
    job = client.post('/api/reporting/complete-report/jobs', json={'definition': definition}).json()['job']
    assert job['elapsed_seconds'] >= 0
    assert started.wait(5)
    cancelled = client.post(f"/api/reporting/complete-report/jobs/{job['id']}/cancel").json()['job']
    assert cancelled['status'] == 'cancelled'
    release.set()
    time.sleep(0.3)
    # The generation stops at its next step and does not change the job any more.
    assert client.get(f"/api/reporting/complete-report/jobs/{job['id']}").json()['job']['status'] == 'cancelled'
    assert 'complete_report_cancelled' in [row['action'] for row in core.repository.list_logs(limit=20)]
    page = (core.PROJECT_ROOT / 'src/web_interface/templates/report_jobs.html').read_text(encoding='utf-8')
    # While it runs, the dialog offers Cancel and Hide (it keeps being generated) and shows its processing time.
    assert all(f'id="{name}"' in page for name in ('cr-dialog-cancel', 'cr-dialog-hide', 'cr-dialog-done', 'cr-dialog-elapsed', 'cr-open-last'))
    script = (core.PROJECT_ROOT / 'src/web_interface/static/js/complete_report.js').read_text(encoding='utf-8')
    assert "title: 'Complete Report Ready'" in script and 'Processed in ${duration(job.elapsed_seconds)}' in script
    # Every page tells when a report hidden in Reporting is ready.
    app_script = (core.PROJECT_ROOT / 'src/web_interface/static/js/app.js').read_text(encoding='utf-8')
    assert "(function watchCompleteReport()" in app_script and '/reporting?complete_report_job=' in app_script


def test_one_action_generates_both_formats_and_the_preview_chooses_one():
    page = (core.PROJECT_ROOT / 'src/web_interface/templates/report_jobs.html').read_text(encoding='utf-8')
    assert page.count('data-cr-export') == 1 and 'data-cr-export="powerpoint"' not in page
    assert 'data-cr-format="powerpoint"' in page and 'data-cr-format="word"' in page
    script = (core.PROJECT_ROOT / 'src/web_interface/static/js/complete_report.js').read_text(encoding='utf-8')
    # Switching the format shows what is already generated: its pages, its sections and its download.
    assert "download.href = `/api/reporting/complete-report/jobs/${job.id}/download/${previewFormat}`" in script
    assert "dialog.classList.toggle('is-word', previewFormat === 'word')" in script
    assert 'function switchFormat(format)' in script


def test_a_report_without_sections_is_refused_and_the_table_is_listed(client):
    login(client)
    definition = complete_reports.normalize_definition({})
    for section in definition['sections']:
        section['enabled'] = False
    response = client.post('/api/reporting/complete-report/jobs', json={'definition': definition})
    assert response.status_code == 400
    page = client.get('/admin')
    assert page.status_code == 200
    assert 'Complete Reports' in page.text


def test_a_report_job_generates_the_chosen_sections_and_previews_them(client, monkeypatch):
    login(client)

    def network_summary(selection, export_kind, destination, components=None, compose=True):
        # The Complete Report composes the deck of the module itself.
        assert export_kind == 'powerpoint' and components == ['snapshot.snapshot'] and not compose
        # The NR Mode of the report and the Campaigns of the section.
        assert selection['nr_mode'] == 'SA' and selection['campaigns'] == ['2026-Q3']
        Path(destination).write_bytes(_deck(lambda presentation: _small_table(presentation, 'Network Snapshot', 'snapshot')))
        return {}

    monkeypatch.setattr(core, 'write_network_insights_summary', network_summary)
    monkeypatch.setattr('src.modules.pptx_to_docx.document_pdf', lambda content, suffix: b'%PDF-1.4')
    rendered = []
    monkeypatch.setattr('src.modules.pptx_to_docx.pdf_page_images',
                        lambda pdf, dpi, pages=None: rendered.append((dpi, pages)) or [PNG] * (len(pages) if pages is not None else 2))
    # The Word version is packed with the PowerPoint, before the preview opens.
    packed = packed_component_document([PackItem('heading1', 'Network Insights')])
    monkeypatch.setattr(complete_reports, 'compose_word', lambda *args, **kwargs: (
        packed, [{'title': 'Network Insights', 'pages': [{'index': 0, 'title': 'Network Snapshot'}, {'index': 1, 'title': 'Network Snapshot'}]}]))
    definition = complete_reports.normalize_definition({'filters': {'nr_mode': 'SA'}})
    for section in definition['sections']:
        section['enabled'] = section['key'] == 'network_insights'
        section['options'].update(campaigns=['2026-Q3'])
        for item in section['items']:
            item['enabled'] = item['key'] == 'snapshot'
            if 'components' in item:
                item['components'] = {key: key == 'snapshot' for key in item['components']}
    response = client.post('/api/reporting/complete-report/jobs', json={'definition': definition, 'name': 'Snapshot'})
    assert response.status_code == 200, response.text
    job_id = response.json()['job']['id']
    import time
    for _attempt in range(100):
        job = client.get(f'/api/reporting/complete-report/jobs/{job_id}').json()['job']
        if job['status'] in {'ready', 'failed'}:
            break
        time.sleep(0.1)
    assert job['status'] == 'ready', job
    assert [item['title'] for item in job['outline']][:3] == ['Snapshot', 'Network Insights', 'Network Snapshot']
    # Both formats are ready with their previews: the slides of the PowerPoint and the pages of the Word document.
    assert {name: item['previews'] for name, item in job['formats'].items()} == {'powerpoint': 2, 'word': 2}
    assert job['formats']['word']['navigation'][0]['title'] == 'Network Insights'
    assert job['formats']['powerpoint']['file_name'].endswith('.pptx') and job['formats']['word']['file_name'].endswith('.docx')
    for export_format in ('powerpoint', 'word'):
        assert client.get(f'/api/reporting/complete-report/jobs/{job_id}/preview/{export_format}/1').content == PNG
    # Opened on the whole screen, a page is rendered large from the preview PDF, once (the first one is ready).
    for _attempt in range(2):
        assert client.get(f'/api/reporting/complete-report/jobs/{job_id}/preview/word/1?large=true').content == PNG
    assert rendered == [(complete_reports.PREVIEW_DPI['pptx'], None), (complete_reports.LARGE_PREVIEW_DPI['pptx'], [0]),
                        (complete_reports.PREVIEW_DPI['docx'], None), (complete_reports.LARGE_PREVIEW_DPI['docx'], [0]),
                        (complete_reports.LARGE_PREVIEW_DPI['docx'], [1])]
    assert client.get(f'/api/reporting/complete-report/jobs/{job_id}/preview/pdf/1').status_code == 404
    for export_format in ('powerpoint', 'word'):
        download = client.get(f'/api/reporting/complete-report/jobs/{job_id}/download/{export_format}')
        assert download.status_code == 200 and download.content[:2] == b'PK'
    assert download.headers['content-disposition'].endswith('.docx')
    actions = [row['action'] for row in core.repository.list_logs(limit=20)]
    assert {'complete_report_requested', 'complete_report_generated', 'complete_report_downloaded'} <= set(actions)


def test_favourites_are_an_artifact_of_reporting_jobs(client, monkeypatch, tmp_path):
    login(client)

    def network_summary(selection, export_kind, destination, components=None, compose=True):
        Path(destination).write_bytes(_deck(lambda presentation: _small_table(presentation, 'Network Snapshot', 'snapshot')))
        return {}

    monkeypatch.setattr(core, 'write_network_insights_summary', network_summary)
    definition = complete_reports.normalize_definition({})
    for section in definition['sections']:
        section['enabled'] = section['key'] == 'network_insights'
    favourite = complete_reports.save_favourite(core.repository, None, 'Weekly', definition, 'super')
    provider = report_tasks.ARTIFACT_PROVIDERS['complete_report']
    assert provider['values'](core.SessionUser(username='super', role='super-admin')) == {'favourites': [[str(favourite['id']), 'Weekly']]}
    artifacts = provider['generate']({'formats': ['powerpoint'], 'options': {'filters': {'favourites': [str(favourite['id'])]}}},
                                     tmp_path, '20261017_120000', core.SessionUser(username='super', role='super-admin'))
    assert [item['status'] for item in artifacts] == ['ready'], artifacts
    assert artifacts[0]['title'] == 'Complete Report · Weekly (PPT)'
    assert (tmp_path / artifacts[0]['file_name']).read_bytes()[:2] == b'PK'
    assert not list(tmp_path.glob('.complete-report-*'))


def test_imported_favourites_find_their_dashboards_by_name_only_when_their_id_is_unknown(client):
    login(client)
    definition = complete_reports.normalize_definition({'sections': [{'key': 'dashboards', 'items': [
        {'key': 'elsewhere', 'label': 'Weekly'}, {'key': 'sa-weekly', 'label': 'Weekly'}]}]})
    imported = complete_reports.import_entries(core.repository, [{'name': 'Imported', 'definition': definition}], 'super',
                                               {'nsa-weekly': ('Weekly', 'NSA'), 'sa-weekly': ('Other', 'SA')})
    assert imported == 1
    favourite = complete_reports.list_favourites(core.repository)[0]
    dashboards = next(section for section in favourite['definition']['sections'] if section['key'] == 'dashboards')
    assert [item['key'] for item in dashboards['items']] == ['nsa-weekly', 'sa-weekly']


def test_favourites_keep_their_dashboards_checked_and_ordered_in_another_workspace(client):
    from src.modules.ppt_dashboards import STATE_KEY

    login(client)

    def dashboards(ids):
        names = ['NetCheck CDR Report', 'NetCheck CDR Report (Main Cities)', 'RF Quality (RSRP & SINR)']
        return {ids[index * 2 + offset]: {'name': name, 'technology': mode}
                for index, name in enumerate(names) for offset, mode in enumerate(('nsa', 'sa'))}

    # The source workspace: every Dashboard in NSA and in SA with the same name; the NSA report checks two of them,
    # RF Quality first.
    source = [f'source-{index}' for index in range(6)]
    core.repository.set_workspace_state(STATE_KEY, json.dumps(dashboards(source)))
    definition = complete_reports.normalize_definition({'sections': [{'key': 'dashboards', 'items': [
        {'key': source[4], 'enabled': True}, {'key': source[0], 'enabled': True}, {'key': source[2], 'enabled': False},
        {'key': source[1], 'enabled': False}, {'key': source[3], 'enabled': False}, {'key': source[5], 'enabled': False}]}]})
    complete_reports.save_favourite(core.repository, None, 'Final CDR', definition, 'super')
    document = complete_reports.export_favourites_document(core.repository)
    exported = json.loads(document)['favourites'][0]['definition']['sections'][0]['items']
    # Each Dashboard travels with its name and NR Mode.
    assert [(item['label'], item['nr_mode']) for item in exported[:2]] == [('RF Quality (RSRP & SINR)', 'NSA'), ('NetCheck CDR Report', 'NSA')]
    # Another workspace has the same Dashboards with other IDs, in another order.
    target = [f'target-{index}' for index in range(6)]
    core.repository.set_workspace_state(STATE_KEY, json.dumps(dict(reversed(list(dashboards(target).items())))))
    complete_reports.delete_favourite(core.repository, complete_reports.list_favourites(core.repository)[0]['id'])
    assert complete_reports.import_favourites_document(core.repository, document, 'super') == 1
    items = complete_reports.list_favourites(core.repository)[0]['definition']['sections'][0]['items']
    # They find the Dashboards of their NR Mode, in the same order and checked as they were.
    assert [(item['key'], item['enabled']) for item in items] == [
        (target[4], True), (target[0], True), (target[2], False), (target[1], False), (target[3], False), (target[5], False)]
    # Documents of earlier versions do not name the NR Mode: the checked Dashboards take those of the report's NR Mode.
    legacy = json.loads(document)
    for item in legacy['favourites'][0]['definition']['sections'][0]['items']:
        item.pop('nr_mode')
    complete_reports.delete_favourite(core.repository, complete_reports.list_favourites(core.repository)[0]['id'])
    complete_reports.import_favourites_document(core.repository, json.dumps(legacy), 'super')
    items = complete_reports.list_favourites(core.repository)[0]['definition']['sections'][0]['items']
    assert [item['key'] for item in items if item['enabled']] == [target[4], target[0]]
    assert len({item['key'] for item in items}) == 6
    # Saving with the name of a favourite overwrites it once the user confirms it, on the page.
    script = (core.PROJECT_ROOT / 'src/web_interface/static/js/complete_report.js').read_text(encoding='utf-8')
    assert "title: 'Overwrite Favourite'" in script
    assert "existing ? `/api/reporting/complete-report/favourites/${existing.id}`" in script
    assert "nr_mode: catalogueItem.nr_mode || ''" in script


def test_a_stored_scheduled_reporting_title_becomes_reporting_once(client):
    from src.modules.module_labels import (
        MODULE_LABELS_STATE_KEY, REPORTING_TITLE_MIGRATION_KEY, load_module_labels, migrate_reporting_title,
    )

    core.repository.set_application_state(REPORTING_TITLE_MIGRATION_KEY, '')
    profile = next(item for item in core.repository.list_modules_profiles() if item['name'] == 'Non-Qualified Calls')
    core.repository.save_modules_profile(profile['id'], profile['name'], 'Workspace, Non-Qualified Calls, Scheduled Reporting, Builders & Bugs Report.',
                                         profile['modules'])
    core.repository.set_application_state(MODULE_LABELS_STATE_KEY, json.dumps({
        'reporting': {'title': 'Scheduled Reporting', 'short_title': 'Scheduled Reporting', 'tab_color': '#123456'}}))
    assert migrate_reporting_title(core.repository) is True
    labels = load_module_labels(core.repository)
    assert labels['reporting']['title'] == 'Reporting' and labels['reporting']['short_title'] == 'Reporting'
    assert labels['reporting']['tab_color'] == '#123456'
    profile = next(item for item in core.repository.list_modules_profiles() if item['name'] == 'Non-Qualified Calls')
    assert profile['description'] == 'Workspace, Non-Qualified Calls, Reporting, Builders & Bugs Report.'
    # A title chosen afterwards stays.
    core.repository.set_application_state(MODULE_LABELS_STATE_KEY, json.dumps({'reporting': {'title': 'Scheduled Reporting'}}))
    assert migrate_reporting_title(core.repository) is False
    assert load_module_labels(core.repository)['reporting']['title'] == 'Scheduled Reporting'



def test_covers_fill_word_pages_of_their_own():
    items = [PackItem('cover', image=PNG), PackItem('cover', image=PNG), _item(7, 3), _item(7, 3), PackItem('cover', image=PNG),
             _item(7, 2)]
    pages = pack_pages(items, 7.37, 10.4)
    assert [[item.kind for row in rows for item in row] for rows, _scale in pages] == [
        ['cover'], ['cover'], ['image', 'image'], ['cover'], ['image']]
    from docx import Document
    document = Document(BytesIO(packed_component_document(items[1:], PNG)))
    # Covers without margins, each with a section of its own; the pages of components with margins.
    margins = [section.left_margin for section in document.sections]
    assert margins[:3] == [0, 0, Inches(0.45)] and len(document.sections) == 5
    # No empty paragraph closes a section (it would add a blank page after a full-page cover).
    assert all(paragraph.runs for paragraph in document.paragraphs)


def test_long_titles_fit_one_line_and_short_ones_keep_their_size():
    from src.modules.slide_titles import fit_slide_title

    presentation = Presentation(TEMPLATE)
    layout = next(item for item in presentation.slide_layouts if item.name == 'Title + 1 rows + 1 columns')
    short = presentation.slides.add_slide(layout)
    short.shapes.title.text = 'Overview'
    assert fit_slide_title(short) is None
    long = presentation.slides.add_slide(layout)
    long.shapes.title.text = 'Spectrum · Total bandwidth, LTE and NR (Download tests) per Operator and Environment'
    size = fit_slide_title(long)
    assert size is not None and 18 <= size < 40
    assert {run.font.size.pt for run in long.shapes.title.text_frame.paragraphs[0].runs} == {size}


def test_dark_layouts_draw_the_white_ericsson_logo_and_light_slides_the_dark_one():
    from src.modules.scoring_exports import _slide
    from src.modules.template_logos import WHITE_LOGO_SHAPES

    presentation = Presentation(TEMPLATE)
    layouts = {layout.name: layout for layout in presentation.slide_masters[0].slide_layouts}
    for name in ('Title Page', 'Transition', 'Black Transition', 'Title Only'):
        assert layouts[name]._element.get('showMasterSp') == '0'
        assert {shape.name for shape in layouts[name].shapes} >= WHITE_LOGO_SHAPES
    # Content layouts keep the dark logos of the master.
    assert layouts['Title + 1 rows + 1 columns']._element.get('showMasterSp') is None
    _remove_all_slides(presentation)
    slide = _slide(presentation, 'Scoring Tables — Summary', 'Best Network')
    assert slide._element.get('showMasterSp') == '0'
    assert {'Logo', 'Graphic 6'} <= {shape.name for shape in slide.shapes}
    output = BytesIO()
    presentation.save(output)
    assert _unique_parts(output.getvalue())


def test_slides_with_pictures_pillow_cannot_read_are_copied():
    from src.modules.scoring_exports import _slide

    # The Scoring slides draw the WMF logo of the master themselves.
    source = pptx_compose.deck(_deck(lambda presentation: _slide(presentation, 'Scoring Tables — Summary', 'Best Network')))
    composer = pptx_compose.DeckComposer(TEMPLATE)
    copied = composer.add_slide(source.slides[0])
    assert {'Logo', 'Graphic 6'} <= {shape.name for shape in copied.shapes}
    content = composer.save()
    assert _unique_parts(content)
    assert len(Presentation(BytesIO(content)).slides) == 1


def test_a_second_export_reuses_the_sections_whose_data_has_not_changed(client, monkeypatch):
    import time

    login(client)
    calls = []

    def network_summary(selection, export_kind, destination, components=None, compose=True):
        calls.append(selection['campaigns'])
        Path(destination).write_bytes(_deck(lambda presentation: _small_table(presentation, 'Network Snapshot', 'snapshot')))
        return {}

    monkeypatch.setattr(core, 'write_network_insights_summary', network_summary)
    monkeypatch.setattr(core, 'network_insights_summary_key', lambda selection: 'analysis-' + ','.join(selection['campaigns']))
    monkeypatch.setattr('src.modules.pptx_to_docx.document_pdf', lambda content, suffix: None)
    definition = complete_reports.normalize_definition({})
    for section in definition['sections']:
        section['enabled'] = section['key'] == 'network_insights'
        section['options']['campaigns'] = ['2026-Q3']

    def export(**extra):
        job_id = client.post('/api/reporting/complete-report/jobs', json={'definition': definition, **extra}).json()['job']['id']
        for _attempt in range(100):
            job = client.get(f'/api/reporting/complete-report/jobs/{job_id}').json()['job']
            if job['status'] in {'ready', 'failed'}:
                return job
            time.sleep(0.1)
        return job

    composed = []
    original_compose = complete_reports.compose_powerpoint
    monkeypatch.setattr(complete_reports, 'compose_powerpoint', lambda *args, **kwargs: composed.append(1) or original_compose(*args, **kwargs))
    first, second = export(), export()
    assert first['status'] == second['status'] == 'ready'
    assert first['reused'] == [] and second['reused'] == ['Network Insights'] and calls == [['2026-Q3']]
    # The same sections, titles and filters reuse the document composed (and, for Word, packed) the first time.
    assert first['document_reused'] is False and second['document_reused'] is True and len(composed) == 1
    # Another identity (other Campaigns), or Generate again, generates it.
    assert export(refresh=True)['reused'] == [] and len(calls) == 2
    definition['sections'][2]['options']['campaigns'] = ['2026-Q2']
    assert export()['reused'] == [] and calls[-1] == ['2026-Q2']


def test_a_saved_filter_of_the_same_cdrs_and_filters_is_used_whatever_its_split_by(monkeypatch):
    def item(**changes):
        return {'id': 1, 'status': 'completed', 'dataset_ids': [3, 1, 2], 'nr_mode': 'NSA', 'baseline_operator': 'EE',
                'scoring_profile_id': 'by-campaign', 'context_filters': {'Campaign': ['@latest-campaign'], 'Operator': ['VF']},
                'levels': ['Operator', 'City', 'Campaign'], 'area_summary': None, **changes}

    starred = item(id=7)
    monkeypatch.setattr(core, 'list_saved_filters', lambda _repository: [
        item(id=5, status='failed'), item(id=6, dataset_ids=[1, 2]), item(id=8, scoring_profile_id='netcheck'), starred, item(id=9)])
    base = {'nr_mode': 'nsa', 'baseline_operator': 'ee', 'scoring_profile_id': 'by-campaign'}
    # The Campaign filter is not compared: the CDRs are the campaigns.
    assert core._matching_saved_filter(None, base, [1, 2, 3], {'Operator': ['vf']}) is starred
    assert core._matching_saved_filter(None, base, [1, 2, 3], {'Operator': ['EE']}) is None
    assert core._matching_saved_filter(None, {**base, 'scoring_profile_id': None}, [1, 2, 3], {'Operator': ['VF']}) is None


def test_favourite_elements_are_exported_and_imported_on_their_own(client):
    import time

    login(client)
    definition = complete_reports.normalize_definition({})
    complete_reports.save_favourite(core.repository, None, 'Quarterly', definition, 'super')
    from src.modules.scoring_reports import load_report_state

    configurations = [item['name'] for item in load_report_state(core.repository)['configurations']]
    for target, label in (('complete-report-favourites', 'Complete Report Favourites'),
                          ('scoring-favourites-filters', 'Scoring & GAP Favourites Filters'),
                          ('scoring-report-templates', 'Scoring & GAP Report Templates')):
        exported = client.get(f'/admin/import-export/export?export_target={target}')
        assert exported.status_code == 200, exported.text
        with zipfile.ZipFile(BytesIO(exported.content)) as archive:
            manifest = json.loads(archive.read('manifest.json'))
            document = json.loads(archive.read(manifest['archive_path']))
        assert manifest['kind'] == target and manifest['source_workspace']['name'] == 'Default'
        inspection = client.post('/admin/import-export/inspect',
                                 files={'package': (f'{target}.zip', BytesIO(exported.content), 'application/zip')})
        assert inspection.status_code == 200, inspection.text
        imported = client.post('/admin/import-export/import/jobs', data={
            'upload_id': inspection.headers['X-Import-Upload-Id'], 'confirmed_import': 'true', 'workspace_ids': 'default'})
        assert imported.status_code == 200, imported.text
        for _attempt in range(100):
            status = client.get(imported.json()['status_url']).json()
            if status['status'] in {'ready', 'failed'}:
                break
            time.sleep(0.01)
        assert status['status'] == 'ready', status
        assert label in json.dumps(status)
        if target == 'complete-report-favourites':
            assert [item['name'] for item in document['favourites']] == ['Quarterly']
            assert [item['name'] for item in complete_reports.list_favourites(core.repository)] == ['Quarterly']
        elif target == 'scoring-favourites-filters':
            assert document['saved_filters'] == []
        else:
            # The saved configurations of the Scoring report editor travel on their own too.
            assert document['format'] == 'drivetest-analyzer-scoring-report-configurations'
            assert [item['name'] for item in document['configurations']] == configurations and configurations
            assert [item['name'] for item in load_report_state(core.repository)['configurations']] == configurations
