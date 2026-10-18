"""Midnight, the dark theme: its content slides are night blue with the white logos of the template, and their tables,
cards and text turn dark when the deck is finished (theme_dark)."""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from src.modules import report_structure
from src.modules.cdr_reporting import _remove_all_slides
from src.modules.pptx_compose import content_shapes
from src.modules.scoring_themes import THEMES, use_theme
from src.modules.template_logos import WHITE_LOGO_SHAPES, use_dark_logos
from src.modules.theme_dark import dark_theme_of, darken_presentation
from src.modules.theme_titles import style_title
from tests.test_scoring_exports import TEMPLATE


def _deck(theme: str, layout_name: str = 'Title Only'):
    deck = Presentation(TEMPLATE)
    _remove_all_slides(deck)
    layout = next(layout for layout in deck.slide_layouts if layout.name == layout_name)
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = 'Best Network — Standings'
    if layout_name == 'Title Only':
        # As the Scoring slides: a white background with the dark logos of the master.
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        use_dark_logos(slide)
    card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1))
    card.name = 'Card'
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    card.text_frame.text = 'EE 906.6'
    card.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x1F, 0x2A, 0x37)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(2), Inches(3), Inches(1))
    band.name = 'Band'
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0xE6, 0x00, 0x00)
    band.text_frame.text = 'VF'
    band.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    with use_theme(theme):
        style_title(slide)
    return deck, slide


def _fill(shape) -> str:
    return str(shape.fill.fore_color.rgb)


def _text(shape) -> str:
    return str(shape.text_frame.paragraphs[0].runs[0].font.color.rgb)


def test_a_dark_theme_paints_its_content_slides_night_blue_with_the_white_logos():
    midnight = THEMES['midnight']
    for layout_name in ('Title Only', 'Title + 1 rows + 1 columns'):
        deck, slide = _deck('midnight', layout_name)
        assert str(slide.background.fill.fore_color.rgb) == midnight['slide'].lstrip('#')
        names = {shape.name for shape in slide.shapes}
        # The white logos of the dark layouts, never the dark ones of the master nor the mark of the red background.
        assert WHITE_LOGO_SHAPES <= names and not names & {'Logo', 'Graphic 6'}
        assert not any(name.startswith('Ericsson watermark') for name in names)
        assert slide._element.get('showMasterSp') == '0'
        assert dark_theme_of(slide) is midnight
    # A light theme leaves the slide white.
    _deck_light, light = _deck('ericsson')
    assert dark_theme_of(light) is None and not WHITE_LOGO_SHAPES & {shape.name for shape in light.shapes}


def test_finishing_a_dark_deck_turns_its_content_dark_and_keeps_the_strong_colours():
    midnight = THEMES['midnight']
    deck, slide = _deck('midnight')
    report_structure.finish(deck)
    shapes = {shape.name: shape for shape in slide.shapes}
    # A white card takes the cards of the theme and its dark text the light ink.
    assert _fill(shapes['Card']) == midnight['surface'].lstrip('#') and _text(shapes['Card']) == midnight['ink'].lstrip('#')
    # An operator colour and white text on it stay as they are.
    assert _fill(shapes['Band']) == 'E60000' and _text(shapes['Band']) == 'FFFFFF'
    before = slide._element.xml
    darken_presentation(deck)
    assert slide._element.xml == before


def test_finishing_a_light_deck_leaves_it_as_it_is():
    deck, slide = _deck('netcheck')
    report_structure.finish(deck)
    shapes = {shape.name: shape for shape in slide.shapes}
    assert _fill(shapes['Card']) == 'FFFFFF' and _text(shapes['Card']) == '1F2A37'


def test_the_logos_of_a_dark_slide_are_not_content_to_compose():
    _deck_dark, slide = _deck('midnight')
    names = {shape.name for shape in content_shapes(slide)}
    assert names == {'Card', 'Band'}
