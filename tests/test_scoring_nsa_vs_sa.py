"""Scoring & GAP Analysis NSA vs SA: the NSA and SA CDRs of the chosen campaigns scored together, each campaign with the
NR Mode of its CDRs, so the NSA and SA tests of an operator are compared side by side and never added together."""
from __future__ import annotations

from pathlib import Path

import pytest

from src.modules import saved_filters, scoring_history
from src.modules.column_names import use_campaign_map
from src.modules.nr_mode import (
    NSA_VS_SA, campaign_with_nr_mode, normalize_scoring_nr_mode, nr_mode_label, nr_modes_of,
)
from src.modules.repository import Repository, local_now_iso
from src.modules.scoring_selection import expand_scoring_campaigns, load_scoring_selection, normalize_scoring_selection, scoring_auto_cdrs
from tests.scoring_fixtures import load_initial_scoring_configuration
from tests.test_scoring_history import _rows


def _add_cdr(repository: Repository, folder: Path, kind: str, nr_mode: str, campaign: str, seed: int) -> int:
    frame = _rows(kind, seed=seed)
    frame['Campaign'] = campaign
    name = f'NetCheck_UK_CDR_{kind.title()}_{campaign}_{nr_mode}.xlsx'
    source = folder / name
    source.write_text(f'{kind} {nr_mode} {campaign}', encoding='utf-8')
    dataset_id, _created = repository.add_dataset(name, str(source), 'super')
    repository.replace_dataset_rows(dataset_id, frame)
    repository.update_dataset_profile(dataset_id, status='ready', progress=100, dataset_kind=kind, nr_mode=nr_mode,
                                      row_count=len(frame), column_count=len(frame.columns), processed_at=local_now_iso())
    repository.replace_cdr_catalogue(dataset_id, vendors=[], regions=[], cities=[], campaigns=[campaign])
    return dataset_id


def _workspace(tmp_path: Path, sa_campaign: str = 'UK_Q3_SA_2026') -> tuple[Repository, dict[str, list[int]]]:
    repository = Repository(tmp_path / 'workspace.db', global_db_path=tmp_path / 'global.db')
    repository.initialize()
    repository.replace_scoring_configuration(load_initial_scoring_configuration())
    ids = {
        'NSA': [_add_cdr(repository, tmp_path, kind, 'NSA', 'UK_Q3_NSA_2026', 3) for kind in ('data', 'voice', 'speech')],
        'SA': [_add_cdr(repository, tmp_path, kind, 'SA', sa_campaign, 7) for kind in ('data', 'voice', 'speech')],
    }
    return repository, ids


def _totals(result: dict, campaign: str) -> dict[tuple, float]:
    return {(row['operator'], row['environment']): round(row['weighted_points'] or 0, 6)
            for row in result['totals'] if row['campaign'] == campaign}


def test_nr_mode_helpers_name_the_nr_mode_of_each_campaign():
    assert normalize_scoring_nr_mode('NSA+SA') == normalize_scoring_nr_mode('nsa vs sa') == NSA_VS_SA
    assert normalize_scoring_nr_mode('sa') == 'SA' and normalize_scoring_nr_mode('LTE') is None
    assert nr_modes_of(NSA_VS_SA) == ('NSA', 'SA') and nr_modes_of('NSA') == ('NSA',)
    assert nr_mode_label(NSA_VS_SA) == 'NSA vs SA' and nr_mode_label('SA') == 'SA'
    # Campaigns that name their NR Mode, and NSA campaigns without one, stay as they are.
    assert campaign_with_nr_mode('UK_Q3_SA_2026', 'SA') == 'UK_Q3_SA_2026'
    assert campaign_with_nr_mode('UK_Q3_NSA_2026', 'NSA') == 'UK_Q3_NSA_2026'
    assert campaign_with_nr_mode('UK_Q2_2026', 'NSA') == 'UK_Q2_2026'
    # An SA campaign without its mode gets it, and one naming the other mode takes the mode of its CDR.
    assert campaign_with_nr_mode('UK_Q2_2026', 'SA') == 'UK_Q2_2026 SA'
    assert campaign_with_nr_mode('UK_Q2_SA_2026', 'NSA') == 'UK_Q2_NSA_2026'
    assert campaign_with_nr_mode('2026-W07', 'SA') == '2026-W07 SA'
    assert campaign_with_nr_mode(None, 'SA') is None


def test_nsa_vs_sa_is_offered_with_both_nr_modes_and_chooses_the_cdrs_of_both(tmp_path):
    repository, ids = _workspace(tmp_path)
    chosen = scoring_auto_cdrs(repository, NSA_VS_SA, {'Campaign': ['@latest-campaign']})
    assert chosen['nr_mode'] == NSA_VS_SA and chosen['nr_modes'] == ['NSA', 'SA', NSA_VS_SA]
    assert sorted(chosen['dataset_ids']) == sorted([*ids['NSA'], *ids['SA']])
    # Latest campaign is the newest quarter, with its NSA and SA campaigns.
    assert sorted(expand_scoring_campaigns(repository, NSA_VS_SA, ['@latest-campaign'])) == ['UK_Q3_NSA_2026', 'UK_Q3_SA_2026']
    assert expand_scoring_campaigns(repository, 'SA', ['@latest-campaign']) == ['UK_Q3_SA_2026']
    # A selection in NSA vs SA is always split by Campaign.
    selection, _warnings = normalize_scoring_selection(repository, {'nr_mode': NSA_VS_SA, 'aggregation_levels': ['Operator']})
    assert selection['nr_mode'] == NSA_VS_SA and 'Campaign' in selection['aggregation_levels']


def test_nsa_vs_sa_is_not_offered_with_one_nr_mode(tmp_path):
    repository = Repository(tmp_path / 'workspace.db', global_db_path=tmp_path / 'global.db')
    repository.initialize()
    repository.replace_scoring_configuration(load_initial_scoring_configuration())
    for kind in ('data', 'voice', 'speech'):
        _add_cdr(repository, tmp_path, kind, 'NSA', 'UK_Q3_NSA_2026', 3)
    chosen = scoring_auto_cdrs(repository, NSA_VS_SA, {})
    assert chosen['nr_mode'] == 'NSA' and chosen['nr_modes'] == ['NSA']
    selection, warnings = normalize_scoring_selection(repository, {'nr_mode': NSA_VS_SA})
    assert selection['nr_mode'] == 'NSA' and warnings


@pytest.mark.parametrize('sa_campaign, sa_label', [('UK_Q3_SA_2026', 'UK_Q3_SA_2026'),
                                                    # An SA campaign that does not name its mode gets it.
                                                    ('UK_Q3_NSA_2026', 'UK_Q3_SA_2026')])
def test_nsa_vs_sa_scores_each_nr_mode_apart_as_its_own_calculation(tmp_path, sa_campaign, sa_label):
    repository, ids = _workspace(tmp_path, sa_campaign)
    both, _ = saved_filters.save_saved_filter(repository, [*ids['NSA'], *ids['SA']], ['Operator'], NSA_VS_SA,
                                              baseline_operator='EE')
    assert both['nr_mode'] == NSA_VS_SA and 'Campaign' in both['levels']
    result = saved_filters.calculate_saved_filter(repository, both['id'])
    assert result['status'] == 'completed', result['error']
    combined = result['result']
    assert {row['campaign'] for row in combined['totals']} == {'UK_Q3_NSA_2026', sa_label}

    # The NSA and SA scores are those of each NR Mode on its own: the SA operators are no peers of the NSA ones.
    nsa, _ = saved_filters.save_saved_filter(repository, ids['NSA'], ['Operator', 'Campaign'], 'NSA', baseline_operator='EE')
    sa, _ = saved_filters.save_saved_filter(repository, ids['SA'], ['Operator', 'Campaign'], 'SA', baseline_operator='EE')
    nsa_result = saved_filters.calculate_saved_filter(repository, nsa['id'])['result']
    sa_result = saved_filters.calculate_saved_filter(repository, sa['id'])['result']
    assert _totals(combined, 'UK_Q3_NSA_2026') == _totals(nsa_result, 'UK_Q3_NSA_2026')
    assert _totals(combined, sa_label) == _totals(sa_result, sa_campaign)
    # The GAP compares each operator with the reference of its own NR Mode.
    gaps = {(row['operator'], row['campaign'], row['environment']): row['gap_points'] for row in combined['gap_totals']
            if row.get('gap_points') is not None}
    reference = {(row['operator'], row['environment']): row['gap_points'] for row in sa_result['gap_totals']
                 if row.get('gap_points') is not None}
    assert reference and all(gaps[(operator, sa_label, environment)] == pytest.approx(value)
                             for (operator, environment), value in reference.items())


def test_nsa_vs_sa_from_the_scoring_history_keeps_the_nr_modes_apart(tmp_path):
    repository, ids = _workspace(tmp_path, 'UK_Q3_NSA_2026')
    scoring_history.sync_history(repository)
    configuration = repository.get_scoring_profile()['configuration']
    dataset_ids = [*ids['NSA'], *ids['SA']]
    from_history = scoring_history.calculate_scoring_from_history(
        repository, dataset_ids, ['Operator', 'Campaign'], configuration=configuration, nr_mode_campaigns=True)
    assert {row['campaign'] for row in from_history['totals']} == {'UK_Q3_NSA_2026', 'UK_Q3_SA_2026'}
    weekly = scoring_history.calculate_scoring_from_history(
        repository, dataset_ids, ['Operator', 'Campaign'], configuration=configuration, nr_mode_campaigns=True,
        time_split='Weekly')
    assert {row['campaign'] for row in weekly['totals']} == {'2026-W28', '2026-W29', '2026-W28 SA', '2026-W29 SA'}
    sources, _campaigns, _fingerprint = saved_filters._source_snapshot(repository, dataset_ids, expected_nr_mode=NSA_VS_SA)
    for source in sources:
        source['levels'] = ['Operator', 'Campaign']
    from src.modules import scoring
    frames = saved_filters._load_source_frames(repository, sources, lambda *args: None, scoring.load_input_columns, {},
                                              nr_mode_campaigns=True)
    by_rows = scoring.calculate_scoring(frames, ['Operator', 'Campaign'], configuration=configuration)
    assert _totals(by_rows, 'UK_Q3_SA_2026') == pytest.approx(_totals(from_history, 'UK_Q3_SA_2026'), rel=1e-2)


def test_an_operator_without_a_network_in_an_nr_mode_is_left_out_of_its_cdrs(tmp_path):
    # Three has no SA (the NR Modes of its label in the Operator Maps): its tests of the SA CDRs are out of scope in the rows
    # read and in the Scoring History, its NSA tests stay, and the results calculated before are out of date.
    import src.DriveTestAnalyzer as core
    from src.modules import scoring
    from src.modules.nr_mode import NR_MODE_COLUMN

    repository, ids = _workspace(tmp_path, 'UK_Q3_NSA_2026')
    scoring_history.sync_history(repository)
    configuration = repository.get_scoring_profile()['configuration']
    dataset_ids = [*ids['NSA'], *ids['SA']]
    _sources, _campaigns, before = saved_filters._source_snapshot(repository, dataset_ids, expected_nr_mode=NSA_VS_SA)
    repository.replace_operator_mapping_group(None, 'Three', ['Three UK'], '#000000', None, ['NSA'])
    sources, _campaigns, after = saved_filters._source_snapshot(repository, dataset_ids, expected_nr_mode=NSA_VS_SA)
    assert after != before
    for source in sources:
        source['levels'] = ['Operator', 'Campaign']
    frames = saved_filters._load_source_frames(repository, sources, lambda *args: None, scoring.load_input_columns, {},
                                              nr_mode_campaigns=True)
    assert all(set(frame.loc[frame[NR_MODE_COLUMN] == 'SA', 'Operator']) == {'EE', 'Vodafone UK'}
               and 'Three UK' in set(frame.loc[frame[NR_MODE_COLUMN] == 'NSA', 'Operator']) for frame in frames.values())
    by_rows = scoring.calculate_scoring(frames, ['Operator', 'Campaign'], configuration=configuration)
    from_history = scoring_history.calculate_scoring_from_history(
        repository, dataset_ids, ['Operator', 'Campaign'], configuration=configuration, nr_mode_campaigns=True)
    for result in (by_rows, from_history):
        operators = {(row['campaign'], row['operator']) for row in result['totals']}
        assert ('UK_Q3_NSA_2026', 'Three UK') in operators and ('UK_Q3_SA_2026', 'Three UK') not in operators
    assert _totals(by_rows, 'UK_Q3_SA_2026') == pytest.approx(_totals(from_history, 'UK_Q3_SA_2026'), rel=1e-2)
    # The headings of the panels of the page name the pairs out of scope of the NR Modes of the result.
    saved, _ = saved_filters.save_saved_filter(repository, dataset_ids, ['Operator'], NSA_VS_SA)
    groups = repository.list_operator_mapping_groups()
    payload = core._saved_filter_payload(saved_filters.calculate_saved_filter(repository, saved['id']), groups, [], None)
    nsa, sa = payload['nr_mode_parts']
    assert payload['out_of_scope'] == sa['out_of_scope'] == ['SA · Three'] and 'out_of_scope' not in nsa
    single, _ = saved_filters.save_saved_filter(repository, ids['NSA'], ['Operator'], 'NSA')
    assert 'out_of_scope' not in core._saved_filter_payload(
        saved_filters.calculate_saved_filter(repository, single['id']), groups, [], None)


def test_nsa_vs_sa_needs_campaign_labels_that_tell_the_nr_modes_apart(tmp_path):
    repository, ids = _workspace(tmp_path)
    dataset_ids = [*ids['NSA'], *ids['SA']]
    with use_campaign_map({'format': '{year}-Q{quarter}', 'modes_in_label': []}):
        with pytest.raises(ValueError, match='Campaign Map shows the NSA and SA campaigns alike'):
            saved_filters.save_saved_filter(repository, dataset_ids, ['Operator'], NSA_VS_SA)
    with pytest.raises(ValueError, match='time split other than All'):
        saved_filters.save_saved_filter(repository, dataset_ids, ['Operator'], NSA_VS_SA,
                                        area_summary={'time_split': 'All'})
    # A single NR Mode still needs CDRs of that mode only.
    with pytest.raises(ValueError, match='does not match the selected NSA NR Mode'):
        saved_filters.save_saved_filter(repository, dataset_ids, ['Operator'], 'NSA')


def test_nsa_vs_sa_saved_filters_travel_with_their_nr_mode(tmp_path):
    source_folder = tmp_path / 'source'
    source_folder.mkdir()
    source, ids = _workspace(source_folder)
    saved, _ = saved_filters.save_saved_filter(source, [*ids['NSA'], *ids['SA']], ['Operator'], NSA_VS_SA)
    document = saved_filters.saved_filters_document(source)
    [entry] = document['saved_filters']
    assert entry['nr_mode'] == NSA_VS_SA
    target_folder = tmp_path / 'target'
    target_folder.mkdir()
    target = Repository(target_folder / 'workspace.db', global_db_path=target_folder / 'global.db')
    target.initialize()
    target.replace_scoring_configuration(load_initial_scoring_configuration())
    assert saved_filters.import_saved_filters(target, document) == 1
    [imported] = saved_filters.list_saved_filters(target)
    assert imported['nr_mode'] == NSA_VS_SA


def test_the_saved_selection_keeps_nsa_vs_sa(tmp_path):
    repository, _ids = _workspace(tmp_path)
    from src.modules.scoring_selection import save_scoring_selection
    selection, _warnings = normalize_scoring_selection(repository, {'nr_mode': NSA_VS_SA, 'context_filters': {
        'Campaign': ['@latest-campaign']}})
    assert save_scoring_selection(repository, selection)
    loaded, warnings = load_scoring_selection(repository)
    assert loaded['nr_mode'] == NSA_VS_SA and not warnings and len(loaded['dataset_ids']) == 6


def test_reporting_jobs_and_bundle_report_templates_keep_nsa_vs_sa():
    from src.modules import bundle_report_templates, report_tasks
    from src.modules.scoring_reports import default_report_configuration
    definition = report_tasks.normalize_definition({'scoring': [
        {'nr_mode': 'NSA+SA', 'report': default_report_configuration()},
        {'nr_mode': 'LTE', 'report': default_report_configuration()}]})
    assert [entry['nr_mode'] for entry in definition['scoring']] == [NSA_VS_SA, 'NSA']
    assert report_tasks.scoring_entry_name({'nr_mode': NSA_VS_SA}) == 'NSA vs SA'
    template = bundle_report_templates.normalize_definition({'filters': {'nr_mode': 'SA'}, 'sections': [
        {'key': 'scoring', 'options': {'nsa_vs_sa': True}}]})
    [scoring] = [section for section in template['sections'] if section['key'] == 'scoring']
    # The section compares NSA vs SA; the report (and its other sections) keeps its NR Mode.
    assert scoring['options']['nsa_vs_sa'] and template['filters']['nr_mode'] == 'SA'
    assert 'NSA vs SA' in bundle_report_templates.section_lines(scoring)[1]
    assert not bundle_report_templates.normalize_definition({})['sections'][1]['options'].get('nsa_vs_sa')


def test_documents_name_nsa_vs_sa():
    from datetime import datetime
    from src.modules.cdr_report_filenames import build_scoring_report_filename
    from src.modules.scoring_exports import _scoring_filter_subtitle
    assert ' - NSA-vs-SA - ' in build_scoring_report_filename(datetime(2026, 10, 20), NSA_VS_SA, {})
    assert ' - NSA - ' in build_scoring_report_filename(datetime(2026, 10, 20), 'NSA', {})
    assert _scoring_filter_subtitle({'nr_mode': NSA_VS_SA}).startswith('Non-Standalone vs Standalone')


def _same_operator_result(tmp_path):
    from src.modules.nr_mode import SAME_OPERATOR_NSA
    repository, ids = _workspace(tmp_path)
    saved, _ = saved_filters.save_saved_filter(repository, [*ids['NSA'], *ids['SA']], ['Operator'], NSA_VS_SA,
                                               baseline_operator=SAME_OPERATOR_NSA)
    calculated = saved_filters.calculate_saved_filter(repository, saved['id'])
    assert calculated['status'] == 'completed', calculated['error']
    return repository, ids, calculated


def test_same_operator_in_nsa_gives_the_gap_of_each_operator_in_sa_to_itself_in_nsa(tmp_path):
    from src.modules.nr_mode import SAME_OPERATOR_NSA, SAME_OPERATOR_NSA_LABEL
    from src.modules.scoring_views import build_scoring_views
    repository, ids, calculated = _same_operator_result(tmp_path)
    result = calculated['result']
    assert not [warning for warning in result['warnings'] if 'unavailable' in warning]
    totals = {(row['operator'], row['campaign'], row['environment'], row['category']): row['weighted_points']
              for row in result['totals']}
    gaps = [row for row in result['gap_totals'] if row.get('gap_points') is not None]
    # Only the SA campaign has GAPs: each operator minus itself in the NSA campaign of the quarter.
    assert gaps and {row['campaign'] for row in gaps} == {'UK_Q3_SA_2026'}
    for row in gaps:
        nsa = totals[(row['operator'], 'UK_Q3_NSA_2026', row['environment'], row['category'])]
        assert row['gap_points'] == pytest.approx(row['operator_points'] - nsa)
        assert row['baseline_actual_operator'] == row['operator']

    views = build_scoring_views(calculated, result, [])
    sa_tables = [table for table in views['score_tables'] if table['context'].get('campaign') == 'UK_Q3_SA_2026']
    nsa_tables = {table['context']['environment']: table for table in views['score_tables']
                  if table['context'].get('campaign') == 'UK_Q3_NSA_2026'}
    assert sa_tables and all(table['baseline_operator'] == SAME_OPERATOR_NSA_LABEL for table in views['score_tables'])
    for table in sa_tables:
        twin = nsa_tables[table['context']['environment']]
        for operator in table['operators']:
            sa_points, nsa_points = table['total']['values'][operator]['points'], twin['total']['values'][operator]['points']
            if table['total']['gaps'][operator] is not None:
                assert table['total']['gaps'][operator] == pytest.approx(sa_points - nsa_points, abs=1e-6)
    # The GAP tables are those of the SA campaign (the NSA one is the reference), one per operator.
    assert {table['context'].get('campaign') for table in views['gap_summary_tables']} == {'UK_Q3_SA_2026'}
    assert {table['title'].split(' — ')[0] for table in views['gap_tables']} >= {f'GAP Analysis: EE vs {SAME_OPERATOR_NSA_LABEL}'}
    # In the matrices split by Operator and Campaign, each SA column is compared with its operator's NSA column.
    [matrix] = [table for table in views['hierarchy_gap_tables'] if table['context']['environment'] == 'Combined']
    columns = {column['id']: column for column in matrix['hierarchy_columns']}
    by_path = {(column['operator'], column['context'].get('campaign')): column['id'] for column in columns.values()}
    [scores] = [table for table in views['hierarchy_score_tables'] if table['context']['environment'] == 'Combined']
    for (operator, campaign), column_id in by_path.items():
        assert columns[column_id]['is_reference'] == (campaign == 'UK_Q3_NSA_2026')
        if campaign == 'UK_Q3_SA_2026' and matrix['total']['gaps'][column_id] is not None:
            reference_id = by_path[(operator, 'UK_Q3_NSA_2026')]
            assert matrix['total']['gaps'][column_id] == pytest.approx(
                scores['total']['values'][column_id]['points'] - scores['total']['values'][reference_id]['points'], abs=1e-6)
    assert views['insights'] is not None


def test_same_operator_in_nsa_needs_nsa_vs_sa_and_exports(tmp_path):
    from io import BytesIO
    from pptx import Presentation
    from src.modules.nr_mode import SAME_OPERATOR_NSA, SAME_OPERATOR_NSA_LABEL
    from src.modules.scoring_exports import export_scoring_powerpoint
    repository, ids, calculated = _same_operator_result(tmp_path)
    with pytest.raises(ValueError, match='choose NSA vs SA'):
        saved_filters.save_saved_filter(repository, ids['SA'], ['Operator'], 'SA', baseline_operator=SAME_OPERATOR_NSA)
    selection, warnings = normalize_scoring_selection(repository, {'nr_mode': NSA_VS_SA, 'baseline_operator': SAME_OPERATOR_NSA})
    assert selection['baseline_operator'] == SAME_OPERATOR_NSA and not warnings
    selection, warnings = normalize_scoring_selection(repository, {'nr_mode': 'NSA', 'baseline_operator': SAME_OPERATOR_NSA})
    assert selection['baseline_operator'] != SAME_OPERATOR_NSA and warnings
    template = Path(__file__).resolve().parents[1] / 'assets/powerpoint-templates/Template_VodafoneThree.pptx'
    content = export_scoring_powerpoint(calculated, calculated['result'], template, [], theme='netcheck')
    text = ' '.join(shape.text_frame.text for slide in Presentation(BytesIO(content)).slides for shape in slide.shapes
                    if shape.has_text_frame)
    assert SAME_OPERATOR_NSA_LABEL in text


def test_the_page_gets_the_results_of_each_nr_mode_on_their_own(tmp_path):
    import src.DriveTestAnalyzer as core
    repository, ids = _workspace(tmp_path)
    saved, _ = saved_filters.save_saved_filter(repository, [*ids['NSA'], *ids['SA']], ['Operator'], NSA_VS_SA,
                                               area_summary={'time_split': 'Campaign', 'breakdown': 'Region', 'cities': []})
    calculated = saved_filters.calculate_saved_filter(repository, saved['id'])
    payload = core._saved_filter_payload(calculated, [], [], None)
    nsa, sa = payload['nr_mode_parts']
    assert (nsa['nr_mode'], sa['nr_mode']) == ('NSA', 'SA')
    # Each part is the result of its NR Mode: its campaign only, in every view.
    assert {row['campaign'] for row in nsa['totals']} == {'UK_Q3_NSA_2026'}
    assert {row['campaign'] for row in sa['totals']} == {'UK_Q3_SA_2026'}
    assert {table['context'].get('campaign') for table in sa['views']['score_tables']} <= {'UK_Q3_SA_2026', None}
    assert sa['campaigns'] == ['UK_Q3_SA_2026'] and sa['saved_filter']['nr_mode'] == 'SA'
    assert sa['views']['insights']['area_summaries'] is not None
    # The whole result stays for the rest of the page.
    assert {row['campaign'] for row in payload['totals']} == {'UK_Q3_NSA_2026', 'UK_Q3_SA_2026'}
    single, _ = saved_filters.save_saved_filter(repository, ids['NSA'], ['Operator'], 'NSA')
    assert 'nr_mode_parts' not in core._saved_filter_payload(
        saved_filters.calculate_saved_filter(repository, single['id']), [], [], None)


def test_each_panel_of_nsa_goes_before_the_same_panel_of_sa():
    import json
    import shutil
    import subprocess
    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    start = script.index('  function interleaveNrModeBlocks(')
    source = script[start:script.index('\n  }', start) + len('\n  }')]
    node = shutil.which('node')
    if not node:
        pytest.skip('Node.js is required to check the NSA vs SA layout.')
    block = lambda name, cls: {'tagName': 'SECTION', 'className': cls, 'name': name}  # noqa: E731
    nsa = [block('nsa table 1', 'table'), block('nsa table 2', 'table'), block('nsa chart', 'chart')]
    sa = [block('sa table 1', 'table'), block('sa comparison', 'comparison'), block('sa table 2', 'table'),
          block('sa chart', 'chart')]
    program = source + '\nconst {nsa, sa} = JSON.parse(require("fs").readFileSync(0, "utf8"));\n' \
        'process.stdout.write(JSON.stringify(interleaveNrModeBlocks(nsa, sa).map(item => item.name)));'
    completed = subprocess.run([node, '-e', program], input=json.dumps({'nsa': nsa, 'sa': sa}), capture_output=True,
                               text=True, check=False)
    assert completed.returncode == 0, completed.stderr
    # Each NSA panel with the SA panel of its kind and place; an SA panel without one of NSA after the SA panel before it.
    assert json.loads(completed.stdout) == ['nsa table 1', 'sa table 1', 'sa comparison', 'nsa table 2', 'sa table 2',
                                            'nsa chart', 'sa chart']


def test_nsa_and_sa_trends_share_one_axis_of_quarters_in_the_same_chart():
    import json
    import shutil
    import subprocess
    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')

    def function(name):
        start = script.index(f'  function {name}(')
        return script[start:script.index('\n  }', start) + len('\n  }')]
    node = shutil.which('node')
    if not node:
        pytest.skip('Node.js is required to check the NSA vs SA trends.')
    program = '\n'.join(function(name) for name in ('nrModeQuarter', 'nrModeQuarters', 'nrModeSeries')) + '''
const quarters = nrModeQuarters(['26-Q2', '26-Q3'], ['26-Q2-SA', '26-Q3-SA']);
const series = nrModeSeries([{label: 'EE', points: [930, 915]}, {label: 'VF', points: [860, 830]}],
  [{label: 'VF', points: [820]}, {label: 'EE', points: [910, 905]}, {label: 'VF SM', points: [null, 800]}],
  {NSA: ['26-Q2', '26-Q3'], SA: ['26-Q3-SA', '26-Q2-SA']}, quarters, 'points');
process.stdout.write(JSON.stringify({quarters, series: series.map(item => [item.label, item.points, item.dashed])}));'''
    completed = subprocess.run([node, '-e', program], capture_output=True, text=True, check=False)
    assert completed.returncode == 0, completed.stderr
    result = json.loads(completed.stdout)
    # 26-Q2 and 26-Q2-SA are one point of the axis; each operator has its NSA line and then its dashed SA line.
    assert result['quarters'] == ['26-Q2', '26-Q3']
    assert result['series'] == [['EE NSA', [930, 915], False], ['EE SA', [905, 910], True], ['VF NSA', [860, 830], False],
                                ['VF SA', [None, 820], True], ['VF SM SA', [800, None], True]]


def _nr_mode_slides(content: bytes) -> list[tuple[str, str]]:
    from io import BytesIO
    from pptx import Presentation
    from src.modules.scoring_exports import SECTION_TAG_PREFIX
    return [(str(slide._element.cSld.get('name') or '')[len(SECTION_TAG_PREFIX):],
             slide.shapes.title.text_frame.paragraphs[0].text if slide.shapes.title is not None else '')
            for slide in Presentation(BytesIO(content)).slides]


def test_documents_show_each_nr_mode_as_the_results_of_the_page(tmp_path):
    from src.modules.scoring_exports import export_scoring_report
    from src.modules.scoring_reports import normalize_report_configuration
    repository, ids = _workspace(tmp_path)
    saved, _ = saved_filters.save_saved_filter(repository, [*ids['NSA'], *ids['SA']], ['Operator'], NSA_VS_SA,
                                               area_summary={'time_split': 'Campaign', 'breakdown': 'Region', 'cities': []})
    calculated = saved_filters.calculate_saved_filter(repository, saved['id'])
    assert calculated['status'] == 'completed', calculated['error']
    report = normalize_report_configuration({'scenarios': [{'name': 'NSA vs SA'}], 'nr_layout': 'interleaved'})
    # NSA & SA in the same chart unless chosen otherwise; without nr_layout, the whole result as it is.
    assert (report['nr_layout'], report['nr_charts']) == ('interleaved', 'same')
    assert 'nr_layout' not in normalize_report_configuration({'scenarios': [{'name': 'NSA vs SA'}], 'nr_layout': 'other'})
    template = Path(__file__).resolve().parents[1] / 'assets/powerpoint-templates/Template_VodafoneThree.pptx'
    entries = [{'scenario': report['scenarios'][0], 'job': calculated, 'result': calculated['result']}]
    whole = _nr_mode_slides(export_scoring_report(entries, template, []))
    assert not [title for _tag, title in whole if title.startswith(('NSA · ', 'SA · '))]

    def titles(slides, section):
        return [title for tag, title in slides if tag == section]
    # Each panel of NSA followed by the same panel of SA, each tagged; the covers once and, in the same chart, the Scoring
    # Charts of the whole result.
    interleaved = _nr_mode_slides(export_scoring_report(entries, template, [], nr_layout='interleaved', nr_charts='same'))
    assert titles(interleaved, 'cover') == titles(whole, 'cover')
    assert titles(interleaved, 'tables') == ['NSA · Scoring Tables — Summary', 'SA · Scoring Tables — Summary',
                                             'NSA · Scoring Tables — Breakdown', 'SA · Scoring Tables — Breakdown']
    assert titles(interleaved, 'charts') == titles(whole, 'charts')
    summary = titles(interleaved, 'summary')
    assert summary[:2] == ['NSA · How the measurements ran', 'SA · How the measurements ran']
    # Every panel of NSA, then every one of SA; in separate charts, the Scoring Charts of each NR Mode.
    grouped = _nr_mode_slides(export_scoring_report(entries, template, [], nr_layout='grouped', nr_charts='separate'))
    charts = titles(grouped, 'charts')
    half = len(charts) // 2
    assert charts and [title.split(' · ', 1)[0] for title in charts] == ['NSA'] * half + ['SA'] * half
    assert [title.split(' · ', 1)[1] for title in charts[:half]] == [title.split(' · ', 1)[1] for title in charts[half:]]
    assert titles(grouped, 'tables') == ['NSA · Scoring Tables — Summary', 'NSA · Scoring Tables — Breakdown',
                                         'SA · Scoring Tables — Summary', 'SA · Scoring Tables — Breakdown']
    # Export Selected Scoring Report follows the NSA vs SA and NSA & SA charts of the page too.
    from src.modules.scoring_exports import export_scoring_powerpoint
    single = _nr_mode_slides(export_scoring_powerpoint(calculated, calculated['result'], template, [],
                                                       nr_layout='interleaved', nr_charts='separate'))
    assert titles(single, 'tables')[:2] == ['NSA · Scoring Tables — Summary', 'SA · Scoring Tables — Summary']
    assert {title.split(' · ', 1)[0] for title in titles(single, 'charts')} == {'NSA', 'SA'}
    assert not [title for _tag, title in _nr_mode_slides(export_scoring_powerpoint(calculated, calculated['result'], template, []))
                if title.startswith(('NSA · ', 'SA · '))]


def test_the_charts_name_the_operators_out_of_scope_at_their_foot(tmp_path):
    # Three has no SA: the charts of the operators keep no place for its SA bar and name it at their foot, on the page (the
    # pairs of the result) and in the PowerPoint, while the charts of an NSA scoring say nothing.
    from io import BytesIO

    from pptx import Presentation

    import src.DriveTestAnalyzer as core
    from src.modules.scoring_exports import export_scoring_powerpoint

    repository, ids = _workspace(tmp_path, 'UK_Q3_NSA_2026')
    repository.replace_operator_mapping_group(None, 'Three', ['Three UK'], '#000000', None, ['NSA'])
    saved, _ = saved_filters.save_saved_filter(repository, [*ids['NSA'], *ids['SA']], ['Operator'], NSA_VS_SA)
    calculated = saved_filters.calculate_saved_filter(repository, saved['id'])
    groups = repository.list_operator_mapping_groups()
    payload = core._saved_filter_payload(calculated, groups, [], None)
    assert payload['out_of_scope'] == ['SA · Three']
    assert not any(column.get('out_of_scope') for matrix in payload['views']['hierarchy_score_tables']
                   for column in matrix['hierarchy_columns'])
    template = Path(__file__).resolve().parents[1] / 'assets/powerpoint-templates/Template_VodafoneThree.pptx'

    def chart_notes(job, result):
        # (All Environments carries the Scoring per Service charts of every environment.)
        content = export_scoring_powerpoint(job, result, template, groups, nr_layout='interleaved', nr_charts='same',
                                            environments=['Combined'])
        return [shape.text_frame.text for slide in Presentation(BytesIO(content)).slides
                if slide.shapes.title is not None and 'Scoring per Service' in slide.shapes.title.text_frame.text
                for shape in slide.shapes if shape.has_text_frame and 'Out of scope' in shape.text_frame.text]
    notes = chart_notes(calculated, calculated['result'])
    assert notes and all(note.startswith('Out of scope: SA · Three: no network in that NR Mode') for note in notes)
    single, _ = saved_filters.save_saved_filter(repository, ids['NSA'], ['Operator'], 'NSA')
    nsa = saved_filters.calculate_saved_filter(repository, single['id'])
    assert chart_notes(nsa, nsa['result']) == []


def test_documents_draw_the_nsa_and_sa_trends_on_one_axis_of_quarters():
    from io import BytesIO
    from pptx import Presentation
    from src.modules.scoring_exports import _nr_mode_area_summaries, _nr_mode_trends, _remove_all_slides, _slide
    from src.modules.scoring_insight_slides import add_trend_slides
    nsa = [{'environment': 'Combined', 'title': '', 'campaigns': ['2026-Q2', '2026-Q3'], 'maximum': 1000,
            'series': [{'label': 'EE', 'operator': 'EE', 'points': [500.0, 520.0]}]}]
    sa = [{'environment': 'Combined', 'title': '', 'campaigns': ['2026-Q3-SA'], 'maximum': 1000,
           'series': [{'label': 'EE', 'operator': 'EE', 'points': [610.0]}, {'label': 'VF', 'operator': 'VF', 'points': [400.0]}]}]
    [trend] = _nr_mode_trends(nsa, sa)
    # 2026-Q3 and 2026-Q3-SA are one point; each operator has its NSA line and its dashed SA line.
    assert trend['campaigns'] == ['2026-Q2', '2026-Q3']
    assert [(series['label'], series['points'], series['dashed']) for series in trend['series']] == [
        ('EE NSA', [500.0, 520.0], False), ('EE SA', [None, 610.0], True), ('VF SA', [None, 400.0], True)]
    areas = [{'environment': 'Combined', 'campaigns': ['2026-Q3'], 'previous_campaign': None, 'scaled': False,
              'areas': [{'label': 'National', 'rows': [{'label': 'EE', 'trend': [500.0]}]}]}]
    sa_areas = [{'environment': 'Combined', 'campaigns': ['2026-Q3-SA'], 'previous_campaign': None, 'scaled': False,
                 'areas': [{'label': 'National', 'rows': [{'label': 'EE', 'trend': [610.0]}]},
                           {'label': 'London', 'rows': [{'label': 'EE', 'trend': [590.0]}]}]}]
    [summary] = _nr_mode_area_summaries(areas, sa_areas)
    assert summary['campaigns'] == ['2026-Q3'] and summary['latest_campaign'] == 'latest campaign'
    assert [(area['label'], [row['label'] for row in area['rows']]) for area in summary['areas']] == [
        ('National', ['EE NSA', 'EE SA']), ('London', ['EE SA'])]
    template = Path(__file__).resolve().parents[1] / 'assets/powerpoint-templates/Template_VodafoneThree.pptx'
    presentation = Presentation(template)
    _remove_all_slides(presentation)
    add_trend_slides(presentation, [trend], scoring_label='Best Network', subtitle='', new_slide=_slide)
    output = BytesIO()
    presentation.save(output)
    [chart] = [shape.chart for shape in Presentation(BytesIO(output.getvalue())).slides[0].shapes if shape.has_chart]
    assert [bool(series.format.line.dash_style) for series in chart.series] == [False, True, True]


def test_bundle_report_templates_show_each_nr_mode_as_chosen(monkeypatch, tmp_path):
    from types import SimpleNamespace
    from src.modules import bundle_report_templates, scoring_selection
    definition = bundle_report_templates.normalize_definition({'sections': [
        {'key': 'scoring', 'options': {'nsa_vs_sa': True, 'nr_layout': 'grouped', 'nr_charts': 'unknown'}}]})
    [scoring] = [section for section in definition['sections'] if section['key'] == 'scoring']
    # Each panel: NSA then SA and NSA & SA in the same chart by default, as the page.
    assert (scoring['options']['nr_layout'], scoring['options']['nr_charts']) == ('grouped', 'same')
    assert 'NSA vs SA (All NSA, then all SA, NSA & SA in the same chart)' in bundle_report_templates.section_lines(scoring)[1]
    defaults = bundle_report_templates.default_options('scoring')
    assert (defaults['nr_layout'], defaults['nr_charts']) == ('interleaved', 'same')
    captured = []

    def build(_repository, base, report, _username):
        captured.append((base, report))
        return b'deck', 'scoring.pptx', {}
    monkeypatch.setattr(scoring_selection, 'expand_scoring_campaigns', lambda *_args: ['UK_Q3_NSA_2026', 'UK_Q3_SA_2026'])
    monkeypatch.setattr(bundle_report_templates, 'complete_campaign_cdrs', lambda *_args: ([1, 2], []))
    core = SimpleNamespace(build_scoring_report_document=build)
    generator = bundle_report_templates.SectionGenerator(core, SimpleNamespace(), SimpleNamespace(username='super'),
                                                         definition, tmp_path)
    generator.scoring(scoring)
    scoring['options']['nsa_vs_sa'] = False
    monkeypatch.setattr(bundle_report_templates, 'resolve_campaign_filter', lambda *_args: [])
    generator.scoring(scoring)
    (nsa_vs_sa_base, nsa_vs_sa_report), (_base, report) = captured
    assert nsa_vs_sa_base['nr_mode'] == NSA_VS_SA
    assert (nsa_vs_sa_report['nr_layout'], nsa_vs_sa_report['nr_charts']) == ('grouped', 'same')
    # Without NSA vs SA, the report has a single NR Mode.
    assert 'nr_layout' not in report
    # The editor offers both choices below Compare NSA vs SA while it is checked.
    script = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/bundle_report_template.js').read_text(encoding='utf-8')
    assert "selectField('NSA vs SA', brtOptions.nr_layouts || [], saved.nr_layout)" in script
    assert "selectField('NSA & SA charts', brtOptions.nr_charts || [], saved.nr_charts)" in script
    assert 'nr_layout: nrLayout.input.value, nr_charts: nrCharts.input.value' in script
    # The Scoring page sends the choices of its results with both exports.
    page = (Path(__file__).resolve().parents[1] / 'src/web_interface/static/js/scoring.js').read_text(encoding='utf-8')
    assert "const nrViews = selection.nr_mode === 'NSA+SA' && typeof selectedNrLayout !== 'undefined'" in page
    assert "if (typeof nrLayoutControl !== 'undefined' && nrLayoutControl && !nrLayoutControl.hidden) {" in page
    assert "url.searchParams.set('nr_layout', selectedNrLayout);" in page
