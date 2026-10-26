"""NQ Prioritization: the scoring impact, the repeated issues, the strategies (Dual Queue Cities first and Roads first,
Points Lost Based and Vodafone Proposal) and the Prioritization Rules of the Non-Qualified Calls."""
import io
import json
from pathlib import Path

import pandas as pd
import pytest

import src.DriveTestAnalyzer as core
from src.modules import non_qualified_calls as nq
from src.modules import nq_catalog, nq_prioritization, scoring, scoring_exclusions, scoring_history
from src.modules.repository import local_now_iso
from tests.scoring_cdr_synthetic import synthetic_cdr_frames
from tests.scoring_fixtures import load_initial_scoring_configuration
from tests.test_non_qualified_calls import enable_module, login, workspace_user


@pytest.fixture(autouse=True)
def background_runs_at_once(monkeypatch):
    """The background work runs at once, so each test reads what it has just calculated."""
    monkeypatch.setattr(core, 'submit_background_task', lambda callback, *args: callback(*args))


# -- the parts without a workspace ---------------------------------------------------------------
def test_impact_levels_are_a_pareto_of_the_points_with_ties_together():
    points = {'a': 5.0, 'b': 3.0, 'c': 1.0, 'd': 1.0, 'e': 0.0}
    # 10 points: a holds the first 50% (High), b reaches 80% (Medium), c and d share their level (Low), e loses none.
    assert nq_prioritization.impact_levels(points, 50, 80) == {
        'a': 'High', 'b': 'Medium', 'c': 'Low', 'd': 'Low', 'e': 'No impact'}
    # Calls with the same points share their level, also across the High share.
    assert set(nq_prioritization.impact_levels({'a': 1.0, 'b': 1.0, 'c': 1.0}, 50, 80).values()) == {'High'}
    assert nq_prioritization.impact_levels({'a': 0.0}, 50, 80) == {'a': 'No impact'}


def test_bad_tests_follow_the_direction_of_the_kpi():
    success = {'calculation': {'formula': '100 * SUM(Call_Status == "Completed") / COUNT(Call_Status)'}, 'direction': 'higher_is_better'}
    polqa = {'calculation': {'formula': '100 * SUM(LQ <= 1.6) / COUNT(Session_Type)'}, 'direction': 'lower_is_better'}
    assert nq_prioritization.bad_tests(success, 95.0, 200) == 10
    assert nq_prioritization.bad_tests(polqa, 2.5, 200) == 5


def test_the_rules_of_the_first_version_become_the_steps_of_points_lost_based():
    # Rules of their own saved before the strategies: Points Lost Based takes them as the steps of its only list, one per
    # rule and Call Type, and a last tracked step with the Priority the rules gave their tracked calls.
    custom = nq_prioritization.CUSTOM
    settings = nq_prioritization.normalize_settings({'rules': [
        {'name': 'London Voice', 'services': ['voice'], 'nr_modes': ['NSA'], 'cities': ['London'], 'criterion': 'all', 'priority': 'P1'},
        {'name': 'VF Data', 'services': ['data'], 'operators': ['VF', '-VF SM'], 'criterion': 'repeated_or_high',
         'priority': 'P2', 'track_priority': 'P5'},
    ]})
    assert settings['strategy'] == custom and settings['final_lists'][custom] is False
    steps = nq_prioritization.strategy_steps(custom, settings['steps'][custom])
    assert [(step['priority'], step['nr_mode'], step['call_type'], step['conditions'], step['joins'], step['queue'], step['operators'])
            for step in steps if step['list'] == 'campaign'] == [
        ('P1', 'NSA', 'voice', ['all'], [], 'Voice', ['@focused']), ('P1', 'NSA', 'whatsapp', ['all'], [], 'Voice', ['@focused']),
        ('P2', 'any', 'data', ['repeated', 'high'], ['or'], 'Data', ['VF']), ('P5', 'any', 'any', ['all'], [], '', ['@focused'])]
    takes = lambda operators, operator: '@focused' in operators or operator in operators
    data = strategy_call(kind='data')
    assert nq_prioritization.strategy_decision(custom, {**data, 'repeated': True}, steps, takes, False)[:4] == (
        'Prioritized', 'Campaign 3 · Data · Cluster OR High impact', 'P2', 'Data')
    assert nq_prioritization.strategy_decision(custom, data, steps, takes, False)[:3] == ('Track only', 'Campaign 4 · Tracked', 'P5')
    # The default rules of the first version give way to the default strategy.
    assert nq_prioritization.normalize_settings({'rules': nq_prioritization.default_rules()})['strategy'] == nq_prioritization.DEFAULT_STRATEGY


def test_settings_are_validated():
    defaults = nq_prioritization.normalize_settings(None)
    assert nq_prioritization.STRATEGIES == {'cities_first': 'Dual Queue (Cities first)', 'roads_first': 'Dual Queue (Roads first)',
                                            'custom': 'Dual Queue (Scoring Based)', 'vodafone_proposal': 'Vodafone Proposal'}
    assert defaults['strategy'] == 'cities_first'
    assert tuple(defaults[key] for key in ('high_share', 'medium_share', 'min_calls', 'previous_quarters', 'spot_meters',
                                           'consecutive_minutes', 'hotspot_join', 'capacity')) == (50, 80, 2, 1, 250, 15, 'or', 20)
    assert nq_prioritization.normalize_settings({'hotspot_join': ' AND '})['hotspot_join'] == 'and'
    # The Hotspot time and distance of the first version give way to the Consecutive time and the Cluster distance.
    assert 'hotspot_minutes' not in nq_prioritization.normalize_settings({'hotspot_minutes': 30, 'hotspot_meters': 500})
    # Dual Queue (Scoring Based) and Vodafone Proposal use their list 1 alone; the other Dual Queue strategies take list 2 once the
    # Final CDR is scored.
    assert defaults['final_lists'] == {'cities_first': True, 'roads_first': True, 'custom': False, 'vodafone_proposal': False}
    assert defaults['excluded_regions'] == ['Vendor B']
    for settings, message in (({'high_share': 60, 'medium_share': 50}, 'above the High share'), ({'min_calls': 1}, 'goes from 2'),
                              ({'consecutive_minutes': 0}, 'Consecutive time'), ({'spot_meters': 10}, 'Cluster distance'),
                              ({'hotspot_join': 'xor'}, 'Hotspot join'), ({'consecutive_minutes': 241}, 'Consecutive time'),
                              ({'excluded_modes': [{'nr_mode': '5G', 'operator': 'VF'}]}, 'pair out of scope')):
        with pytest.raises(ValueError, match=message):
            nq_prioritization.normalize_settings(settings, strict=True)
    assert nq_prioritization.normalize_settings({'excluded_modes': [{'nr_mode': 'sa', 'operator': ' Three '}, {'nr_mode': 'SA', 'operator': 'three'}]})[
        'excluded_modes'] == [{'nr_mode': 'SA', 'operator': 'Three'}]
    assert nq_prioritization.normalize_settings({'high_share': 40})['strategy'] == 'cities_first'
    with pytest.raises(ValueError, match='Choose the strategy'):
        nq_prioritization.normalize_settings({'strategy': 'fifo'}, strict=True)
    # Every field of every step of the strategies: by default its own Priority and the Focused Operators.
    steps = nq_prioritization.normalize_settings({'strategy': 'roads_first', 'steps': {'roads_first': {
        'campaign-6': {'priority': 'p2', 'operators': [' VF ', 'Three', 'three']}}}})['steps']
    assert steps['roads_first']['campaign-6'] == {'number': 6, 'priority': 'P2', 'nr_mode': 'NSA', 'call_type': 'data', 'area': 'any',
                                                  'conditions': ['cluster'], 'joins': [], 'queue': 'Data',
                                                  'operators': ['VF', 'Three']}
    # NSA Voice (P5) goes before NSA Data (P6) in list 1, as in list 2.
    assert [(steps['cities_first'][key]['priority'], steps['cities_first'][key]['call_type']) for key in ('campaign-5', 'campaign-6')] == [
        ('P5', 'voice'), ('P6', 'data')]
    assert steps['roads_first']['final-tracked'] == {'number': 11, 'priority': 'P8', 'nr_mode': 'any', 'call_type': 'any', 'area': 'any',
                                                     'conditions': ['all'], 'joins': [], 'queue': '', 'operators': ['@focused']}
    assert (steps['roads_first']['campaign-1']['call_type'], steps['roads_first']['campaign-1']['area']) == ('voice', 'roads')
    # The Data steps of list 1 of both Dual Queue strategies take the failures of a hotspot.
    for strategy in ('cities_first', 'roads_first'):
        assert (steps[strategy]['campaign-2']['conditions'], steps[strategy]['campaign-2']['joins']) == (['cluster'], [])
    # List 2 of every strategy, by the impact levels, is the only list of Dual Queue (Scoring Based); Vodafone Proposal takes every
    # SA, then NSA, Voice and WhatsApp call and the Data calls of a cluster, in its Unique queue.
    ordered = lambda strategy, prefix: [steps[strategy][key] for key in sorted(steps[strategy], key=lambda key: steps[strategy][key]['number'])
                                        if key.startswith(prefix)]
    list_2 = [('P1', 'SA', 'voice', ['high'], 'Voice'), ('P2', 'SA', 'data', ['high'], 'Data'), ('P3', 'SA', 'whatsapp', ['high'], 'Voice'),
              ('P4', 'SA', 'voice', ['medium'], 'Voice'), ('P4', 'SA', 'whatsapp', ['medium'], 'Voice'), ('P4', 'SA', 'voice', ['all'], 'Voice'),
              ('P4', 'SA', 'whatsapp', ['all'], 'Voice'), ('P5', 'NSA', 'voice', ['all'], 'Voice'), ('P6', 'NSA', 'data', ['high'], 'Data'),
              ('P7', 'NSA', 'whatsapp', ['all'], 'Voice'), ('P8', 'any', 'any', ['all'], '')]
    for strategy, prefix in (('cities_first', 'final-'), ('roads_first', 'final-'), ('vodafone_proposal', 'final-'), ('custom', 'final-'),
                             ('custom', 'campaign-')):
        assert [(step['priority'], step['nr_mode'], step['call_type'], step['conditions'], step['queue'])
                for step in ordered(strategy, prefix)] == list_2
    assert [(step['priority'], step['nr_mode'], step['call_type'], step['conditions'], step['queue']) for step in ordered('vodafone_proposal', 'campaign-')] == [
        ('P1', 'SA', 'voice', ['all'], 'Unique'), ('P1', 'SA', 'whatsapp', ['all'], 'Unique'), ('P2', 'SA', 'data', ['repeated'], 'Unique'),
        ('P3', 'NSA', 'voice', ['all'], 'Unique'), ('P3', 'NSA', 'whatsapp', ['all'], 'Unique'), ('P4', 'NSA', 'data', ['repeated'], 'Unique'),
        ('P5', 'any', 'any', ['all'], 'Unique'), ('P8', 'any', 'any', ['all'], '')]
    # A Call Type of the first version of the steps (Voice · Cities) gives its type and area apart.
    legacy = nq_prioritization.normalize_steps({'cities_first': {'campaign-3': {'call_type': 'voice_cities'}}}, strict=True)
    assert (legacy['cities_first']['campaign-3']['call_type'], legacy['cities_first']['campaign-3']['area']) == ('voice', 'city')
    # Biggest impact, a condition of the first steps, is the High impact now.
    assert nq_prioritization.normalize_steps({'cities_first': {'final-1': {'conditions': ['biggest']}}})['cities_first']['final-1'][
        'conditions'] == ['high']
    # The conditions keep the order chosen with the operator before each one (OR when not given, as the steps saved before
    # the operators), and Every call stands alone.
    chosen = nq_prioritization.normalize_steps({'cities_first': {
        'final-2': {'conditions': ['repeated', 'high', 'consecutive'], 'joins': ['AND']}, 'final-1': {'conditions': ['high', 'all'], 'joins': ['and']},
        'final-8': {'conditions': ['high', 'repeated']}}}, strict=True)['cities_first']
    assert (chosen['final-2']['conditions'], chosen['final-2']['joins']) == (['repeated', 'high', 'consecutive'], ['and', 'or'])
    assert (chosen['final-1']['conditions'], chosen['final-1']['joins']) == (['all'], [])
    assert chosen['final-8']['joins'] == ['or']
    for step, message in (({'nr_mode': '5G'}, 'NR Mode'), ({'call_type': 'video'}, 'Call Type'), ({'area': 'rural'}, 'Area'),
                          ({'queue': 'Ping'}, 'queue'), ({'conditions': ['cluster', 'consecutive'], 'joins': ['xor']}, 'AND or OR'),
                          ({'conditions': ['fast']}, 'conditions'), ({'number': 21}, 'from 1 to 20')):
        with pytest.raises(ValueError, match=message):
            nq_prioritization.normalize_steps({'cities_first': {'campaign-2': step}}, strict=True)
    # A list saved before a step joined the defaults keeps its own steps: the new default step is not added when the list
    # already has the same step under another key (it would show twice).
    every = ['all']
    own = {f'campaign-{number}': {'number': number, 'priority': priority, 'nr_mode': mode, 'call_type': kind, 'area': 'any',
                                  'conditions': conditions, 'joins': [], 'queue': 'Unique'}
           for number, priority, mode, kind, conditions in ((1, 'P1', 'SA', 'voice', every), (2, 'P1', 'SA', 'whatsapp', every),
                                                           (3, 'P2', 'SA', 'data', ['repeated']), (4, 'P3', 'NSA', 'voice', every),
                                                           (5, 'P3', 'NSA', 'whatsapp', every), (6, 'P4', 'NSA', 'data', ['repeated']))}
    own['campaign-tracked'] = {'number': 7, 'priority': 'P5', 'nr_mode': 'any', 'call_type': 'any', 'area': 'any', 'conditions': every,
                               'joins': [], 'queue': 'Unique'}
    own['campaign-s1'] = {'number': 8, 'priority': 'P8', 'nr_mode': 'any', 'call_type': 'any', 'area': 'any', 'conditions': every,
                          'joins': [], 'queue': ''}
    vodafone = [step for step in nq_prioritization.strategy_steps('vodafone_proposal', own) if step['list'] == 'campaign']
    assert [(step['priority'], step['queue']) for step in vodafone].count(('P5', 'Unique')) == 1 and len(vodafone) == 8
    # The same once loaded and saved again: the new default step stays deleted, it does not come back with every Save.
    loaded = nq_prioritization.normalize_steps({'vodafone_proposal': own})
    for steps in (loaded, nq_prioritization.normalize_steps(loaded)):
        assert steps['vodafone_proposal']['campaign-7'] == {'deleted': True}
        vodafone = [step for step in nq_prioritization.strategy_steps('vodafone_proposal', steps['vodafone_proposal'])
                    if step['list'] == 'campaign']
        assert [(step['priority'], step['queue']) for step in vodafone].count(('P5', 'Unique')) == 1 and len(vodafone) == 8
    # Focused Operators stands alone.
    assert nq_prioritization.normalize_steps({'cities_first': {'campaign-1': {'operators': ['VF', '@focused']}}})[
        'cities_first']['campaign-1']['operators'] == ['@focused']
    # A list saved exactly as former default steps takes the current ones (every Save keeps every step): list 2 and the list 1 of
    # Dual Queue (Scoring Based) as they were, and the list 1 of Vodafone Proposal with the Voice and Data queues.
    def saved(rows, prefix):
        return {f'{prefix}-{"tracked" if index == len(rows) else index}': {
            'priority': row[0], 'nr_mode': row[1], 'call_type': row[2], 'area': row[3], 'conditions': list(row[4]),
            'joins': nq_prioritization.joins_of(list(row[4]), list(row[6]) if len(row) > 6 else None), 'queue': row[5]}
            for index, row in enumerate(rows, start=1)}
    for strategy, list_name in (('cities_first', 'final'), ('cities_first', 'campaign'), ('roads_first', 'campaign'), ('custom', 'campaign'),
                                ('vodafone_proposal', 'campaign')):
        former = saved(nq_prioritization._former_rows(strategy, list_name)[0], list_name)
        migrated = nq_prioritization.normalize_steps({strategy: former})[strategy]
        assert migrated == nq_prioritization.normalize_steps({})[strategy], (strategy, list_name)


def test_the_steps_take_the_focus_operators_by_name_the_sa_steps_those_with_sa():
    # An operator uses SA when at least 10 % of its calls of the SA campaigns ran on NR SA (Three drives them on LTE).
    rows = [('VF', 'NR SA', 40), ('Vodafone UK', 'LTE', 60), ('Three', 'LTE', 95), ('Three', 'NR SA', 3), ('Three UK', 'EN-DC', 2),
            ('O2', 'NE-DC', 30), ('O2', 'LTE', 70)]
    aliases = {'vf': {'vf', 'vodafone uk'}, 'vodafone uk': {'vf', 'vodafone uk'}, 'three': {'three', 'three uk'}, 'three uk': {'three', 'three uk'}}
    assert nq_prioritization.operators_with_sa(rows, aliases) == {'vf', 'vodafone uk', 'o2'}
    assert nq_prioritization.operators_with_sa([('VF', ''), ('Three', None)]) is None

    class Repository:
        def __init__(self, groups, calls):
            self.groups, self.calls = groups, calls

        def list_operator_mapping_groups(self):
            return self.groups

        def connection(self):
            import contextlib
            import sqlite3
            connection = sqlite3.connect(':memory:')
            connection.execute('CREATE TABLE nq_call_latest (operator TEXT, technology TEXT, nr_mode TEXT)')
            connection.executemany('INSERT INTO nq_call_latest VALUES (?, ?, ?)', self.calls)
            return contextlib.closing(connection)
    groups = [{'canonical': 'VF', 'aliases': ['Vodafone UK'], 'role': 'focus'}, {'canonical': 'VF SM', 'aliases': [], 'role': 'focus'},
              {'canonical': 'Three', 'aliases': ['Three UK'], 'role': 'focus'}, {'canonical': 'EE', 'aliases': [], 'role': ''}]
    calls = [('VF', 'NR SA', 'SA'), ('Vodafone UK', 'LTE', 'SA'), ('Three', 'LTE', 'SA'), ('Three UK', 'EN-DC', 'SA'), ('VF', 'LTE', 'NSA'),
             ('Three', 'LTE', 'NSA'), ('EE', 'NR SA', 'SA')]
    # VF SM has no call: the NSA and tracked steps take VF and Three, the SA steps VF.
    operators = nq_prioritization.default_step_operators(Repository(groups, calls))
    assert operators == {'SA': ['VF'], 'NSA': ['VF', 'Three'], 'any': ['VF', 'Three']}
    steps = nq_prioritization.normalize_settings({}, operators=operators)['steps']
    for strategy in ('cities_first', 'roads_first'):
        assert {key: step['operators'] for key, step in steps[strategy].items()
                if key in {'campaign-1', 'campaign-5', 'campaign-tracked', 'final-4', 'final-7', 'final-8', 'final-tracked'}} == {
            'campaign-1': ['VF'], 'campaign-5': ['VF', 'Three'], 'campaign-tracked': ['VF', 'Three'], 'final-4': ['VF'], 'final-7': ['VF'],
            'final-8': ['VF', 'Three'], 'final-tracked': ['VF', 'Three']}
    # A Three call of an SA campaign takes its NSA step, as the SA steps take VF only.
    takes = lambda chosen, operator: operator in chosen
    decided = nq_prioritization.strategy_decision('cities_first', strategy_call(operator='Three'),
                                                  nq_prioritization.strategy_steps('cities_first', steps['cities_first']), takes)
    assert decided[1:5] == ('Campaign 5 · NSA Voice', 'P5', 'Voice', 5)
    # Without Focus operators, Focused Operators (every operator).
    assert nq_prioritization.default_step_operators(Repository([], calls))['SA'] == ['@focused']


def test_every_call_type_takes_every_condition():
    # Clusters, hotspots, consecutive failures, the impact levels, drops and setup failures are conditions of every call type
    # and both lists.
    assert set(nq_prioritization.CAMPAIGN_CONDITIONS) == set(nq_prioritization.FINAL_CONDITIONS) == {
        'all', 'consecutive', 'repeated', 'cluster', 'high', 'medium', 'low', 'drop', 'setup'}
    assert (nq_prioritization.CONDITION_LABELS['repeated'], nq_prioritization.CONDITION_LABELS['cluster']) == ('Cluster', 'Hotspots')
    holds = nq_prioritization.condition_holds
    voice = {'kind': 'voice', 'cluster_size': 2, 'consecutive': 'Yes', 'result': 'Dropped', 'repeated': True, 'level': 'Medium'}
    assert holds(['cluster'], voice) and holds(['consecutive'], voice) and holds(['drop'], voice) and not holds(['setup'], voice)
    assert holds(['repeated'], voice) and holds(['medium'], voice) and not holds(['high'], voice) and not holds(['low'], voice)
    assert holds(['drop'], {'result': 'Cutoff'}) and holds(['setup'], {'result': 'Failed'}) and holds(['low'], {'level': ''})
    # The impact levels the conditions read: a Pareto of the Expected Gain of each campaign, operator and service, the calls in
    # scope compared only with each other (a call out of scope never moves their levels), the Excluded Calls without level.
    def row(key, points, region='Vendor A', state='scored', service='voice'):
        return {'call_key': key, 'campaign_key': 'SA|2026-q3', 'operator': 'VF', 'service': service, 'region': region, 'nr_mode': 'SA',
                'impact_state': state, 'expected_points': points}
    rows = [row('a', 6), row('b', 3), row('c', 1), row('far', 10, region='Vendor B'), row('x', 9, state='excluded'), row('d', 2, service='data')]
    assert nq_prioritization.expected_levels(rows, 50, 80) == {'far': 'High', 'a': 'Medium', 'b': 'Low', 'c': 'Low', 'd': 'High'}
    out = lambda item: item['region'] == 'Vendor B'
    assert nq_prioritization.expected_levels(rows, 50, 80, out) == {'a': 'High', 'b': 'Medium', 'c': 'Low', 'far': 'High', 'd': 'High'}
    # The Scoring Impact of the calls follows the same groups.
    scored = [{**item, 'impact_points': item['expected_points']} for item in rows]
    assert nq_prioritization.scored_levels(scored, 50, 80, out) == {'a': 'High', 'b': 'Medium', 'c': 'Low', 'far': 'High', 'd': 'High'}


def test_the_calls_out_of_scope_stay_out_of_the_queues(strategized, monkeypatch):
    # The calls of the regions out of the scope of the team, those of the NR Mode · operator pairs chosen (SA · Three) and,
    # always, those of the operators that are not Focus operators of the Operator Maps.
    settings = {**nq_prioritization.default_settings(), 'excluded_regions': ['Vendor B'], 'excluded_modes': [{'nr_mode': 'SA', 'operator': 'Three'}]}
    scope = nq_prioritization.scope_of(core.repository, settings)
    assert scope('Vendor B', 'NSA', 'VF') == ('Out of scope · Vendor B', 'region')
    assert scope('Vendor A', 'SA', 'Three') == ('Out of scope · SA · Three', 'mode')
    assert scope('Vendor A', 'NSA', 'Three') == scope('Vendor A', 'SA', 'VF') == ('', '')
    with monkeypatch.context() as patch:
        patch.setattr(nq_prioritization, 'focus_names', lambda repository: {'vf', 'three'})
        assert nq_prioritization.scope_of(core.repository, settings)('Vendor A', 'NSA', 'EE') == ('Out of scope · NSA · EE', 'operator')
    assert nq_prioritization.strategy_decision('cities_first', {**strategy_call(), 'scope': 'Out of scope · Vendor B'}) == (
        'No rule', 'Out of scope · Vendor B', '', '', None, '')
    client = strategized['client']
    # A third of the calls in a region, first in scope, then out of it.
    region = 'Vendor B'
    with core.repository.connection() as connection:
        connection.execute(f"UPDATE {nq.NQ_CALL_LATEST_TABLE} SET region = CASE WHEN rowid % 3 = 0 THEN ? ELSE 'Vendor A' END", (region,))
    assert nq_prioritization.region_choices(core.repository) == ['Vendor A', 'Vendor B']
    settings = nq_prioritization.save_settings(core.repository, {**nq_prioritization.load_settings(core.repository), 'excluded_regions': []}, 'super')
    response = client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': {**settings, 'excluded_regions': [region]}})
    assert response.status_code == 200, response.text
    with core.repository.connection() as connection:
        rows = connection.execute(f'SELECT p.prioritization, p.priority_rule, p.queue, l.region FROM {nq_prioritization.PRIORITY_TABLE} p '
                                  f'JOIN {nq.NQ_CALL_LATEST_TABLE} l ON l.call_key = p.call_key').fetchall()
    out = [row for row in rows if row['region'] == region]
    assert out and {(row['prioritization'], row['priority_rule'], row['queue']) for row in out} == {('No rule', f'Out of scope · {region}', '')}
    assert all(not str(row['priority_rule']).startswith('Out of scope') for row in rows if row['region'] != region)
    panel = nq_prioritization.panel(core.repository, {}, 'super')
    assert next(entry for entry in panel['rules'] if entry['name'] == 'Out of scope')['calls'] == len(out)
    log = next(row for row in core.repository.list_logs(limit=50) if row['action'] == 'nq_prioritization_settings')
    details = json.loads(log['details'])
    assert 'excluded_regions' in details['changed'] and details['excluded_regions'] == [region]
    # The #Calls of the steps leave them out too.
    counts = nq_prioritization.step_counts(core.repository, nq_prioritization.load_settings(core.repository))
    assert counts['out_of_scope'] and counts['scope']['regions'][0][0] == region


def test_nq_calls_in_scope_only_leaves_the_calls_out_of_scope_out_of_the_module(strategized, monkeypatch):
    client = strategized['client']
    region = 'Vendor B'
    with core.repository.connection() as connection:
        connection.execute(f"UPDATE {nq.NQ_CALL_LATEST_TABLE} SET region = CASE WHEN rowid % 3 = 0 THEN ? ELSE 'Vendor A' END", (region,))
        latest = [dict(row) for row in connection.execute(f'SELECT call_key, region, nr_mode, operator FROM {nq.NQ_CALL_LATEST_TABLE}')]
    settings = {**nq_prioritization.load_settings(core.repository), 'excluded_regions': [region],
                'excluded_modes': [{'nr_mode': 'NSA', 'operator': 'Three'}]}
    nq_prioritization.save_settings(core.repository, settings, 'super')
    # The SQL condition takes the calls scope_of leaves in scope, also with Focus operators.
    for focus in (None, {'vf'}):
        with monkeypatch.context() as patch:
            patch.setattr(nq_prioritization, 'focus_names', lambda repository, focus=focus: focus)
            reason = nq_prioritization.scope_of(core.repository, nq_prioritization.load_settings(core.repository))
            clause, params = nq_prioritization.scope_sql(core.repository)
            with core.repository.connection() as connection:
                kept = {row[0] for row in connection.execute(f'SELECT call_key FROM {nq.NQ_CALL_LATEST_TABLE} WHERE {clause}', params)}
            assert kept == {row['call_key'] for row in latest if not reason(row['region'], row['nr_mode'], row['operator'])[0]}
    reason = nq_prioritization.scope_of(core.repository, nq_prioritization.load_settings(core.repository))
    in_scope = [row for row in latest if not reason(row['region'], row['nr_mode'], row['operator'])[0]]
    assert 0 < len(in_scope) < len(latest)
    # The table, the panels and the values of the filters take only the calls in scope.
    every = client.post('/api/non-qualified-calls/calls', json={'filters': {}}).json()['total']
    scoped = client.post('/api/non-qualified-calls/calls', json={'filters': {'in_scope': True}}).json()['total']
    assert every == len(latest) and scoped == len(in_scope)
    panel = nq_prioritization.panel(core.repository, {'in_scope': True}, 'super')
    assert panel['totals']['calls'] == len(in_scope)
    assert next(entry for entry in panel['rules'] if entry['name'] == 'Out of scope')['calls'] == 0
    values = client.post('/api/non-qualified-calls/filters/values', json={'filters': {'in_scope': True}}).json()['values']
    assert region not in values['region']
    # The NQ rates count the calls in scope only once the Region of every call is indexed (the qualified calls, which are
    # not in the latest calls, keep the Region they were indexed with).
    with core.repository.connection() as connection:
        connection.execute(f"UPDATE {nq.NQ_CALL_POPULATION_TABLE} SET region = COALESCE((SELECT l.region FROM {nq.NQ_CALL_LATEST_TABLE} l "
                           f"WHERE l.call_key = {nq.NQ_CALL_POPULATION_TABLE}.call_key), region)")
    nq._rate_counts_cache.clear()
    every_rate = sum(item['total']['nq'] for item in nq.nq_rates(core.repository, {})['matrices'] if item['service'] != 'all')
    scoped_rate = sum(item['total']['nq'] for item in nq.nq_rates(core.repository, {'in_scope': True})['matrices'] if item['service'] != 'all')
    assert scoped_rate < every_rate
    # The switch is kept with the filters (Reset Filters leaves it), and App Events records each change.
    response = client.put('/api/non-qualified-calls/filters', json={'filters': {'in_scope': True}})
    assert response.status_code == 200 and response.json()['filters'] == {'in_scope': True}
    log = next(row for row in core.repository.list_logs(limit=20) if row['action'] == 'nq_calls_scope_only')
    assert json.loads(log['details']) == {'in_scope': True}


def test_the_operators_of_a_step_take_the_names_of_their_labels_and_the_focus_operators(client):
    takes = nq_prioritization.operator_takes(core.repository)
    assert takes(['VF'], 'VF') and takes(['vf'], 'VF') and not takes(['VF'], 'Three') and not takes([], 'VF')
    # Focused Operators: every operator while the workspace has no Focus operator.
    assert takes(['@focused'], 'EE')
    assert nq_prioritization.operators_text(['VF', 'Three']) == 'VF & Three' and nq_prioritization.operators_text([]) == 'No operator'
    with pytest.raises(ValueError, match='from P1 to P8'):
        nq_prioritization.normalize_settings({'steps': {'roads_first': {'campaign-1': {'priority': 'P9'}}}}, strict=True)


def strategy_call(**values):
    call = {'operator': 'VF', 'nr_mode': 'SA', 'kind': 'voice', 'area': 'city', 'excluded': False, 'final': False, 'level': 'Low',
            'repeated': False, 'cluster_size': None, 'consecutive': ''}
    return {**call, **values}


def test_strategies_give_each_call_its_step_queue_and_priority():
    decide = nq_prioritization.strategy_decision
    cities, roads = nq_prioritization.CITIES_FIRST, nq_prioritization.ROADS_FIRST
    # List 1: the step is the Priority, P1 to P7, cities or roads first.
    assert decide(cities, strategy_call()) == ('Prioritized', 'Campaign 1 · SA Voice · Cities', 'P1', 'Voice', 1, 'campaign')
    assert decide(cities, strategy_call(area='roads'))[1:5] == ('Campaign 3 · SA Voice · Connecting Roads', 'P3', 'Voice', 3)
    assert decide(roads, strategy_call(area='roads'))[1:5] == ('Campaign 1 · SA Voice · Connecting Roads', 'P1', 'Voice', 1)
    assert decide(roads, strategy_call())[2] == 'P3'
    assert decide(cities, strategy_call(kind='whatsapp', area='roads'))[2:5] == ('P4', 'Voice', 4)
    assert decide(cities, strategy_call(nr_mode='NSA'))[2:5] == ('P5', 'Voice', 5)
    assert decide(cities, strategy_call(nr_mode='NSA', kind='whatsapp'))[2:5] == ('P7', 'Voice', 7)
    # Data, in both strategies: the failures of a hotspot go to the Data queue, the other failures are tracked.
    cluster = strategy_call(kind='data', cluster_size=2)
    for strategy in (cities, roads):
        assert decide(strategy, cluster)[1:5] == ('Campaign 2 · SA Data · Hotspots', 'P2', 'Data', 2)
        assert decide(strategy, {**cluster, 'cluster_size': 1, 'repeated': True, 'consecutive': 'Yes'})[:4] == (
            'Track only', 'Campaign 8 · Tracked', 'P8', '')
        assert decide(strategy, {**cluster, 'nr_mode': 'NSA'})[2:5] == ('P6', 'Data', 6)
    single = strategy_call(kind='data', cluster_size=1)
    for call in (single, strategy_call(kind='ping', cluster_size=4), strategy_call(excluded=True)):
        assert decide(roads, call)[:3] == ('Track only', 'Campaign 8 · Tracked', 'P8')
    # The operators of the steps: Three has no SA, so its SA calls take the NSA steps; an operator of no step is out of the
    # queues, or tracked when the Tracked step takes it.
    configured = nq_prioritization.normalize_steps({'cities_first': {
        **{f'campaign-{number}': {'operators': ['VF']} for number in range(1, 5)},
        **{f'campaign-{number}': {'operators': ['VF', 'Three']} for number in range(5, 8)},
        'campaign-tracked': {'operators': ['VF', 'Three'], 'priority': 'P8'}, 'campaign-5': {'operators': ['VF', 'Three'], 'priority': 'P2'}}})
    steps = nq_prioritization.strategy_steps(cities, configured['cities_first'])
    takes = lambda operators, operator: operator in operators
    assert decide(cities, strategy_call(operator='Three'), steps, takes)[1:5] == ('Campaign 5 · NSA Voice', 'P2', 'Voice', 5)
    assert decide(cities, strategy_call(operator='Three', kind='ping'), steps, takes)[:3] == ('Track only', 'Campaign 8 · Tracked', 'P8')
    assert decide(cities, strategy_call(operator='EE'), steps, takes) == ('No rule', '', '', '', None, '')
    assert decide(cities, strategy_call(), steps, takes)[2] == 'P1'
    # Every field of a step changes what it takes: every data failure, only those of a cluster once the Final CDR is scored,
    # another NR Mode, Call Type or queue, and its step orders it inside its Priority.
    chosen = nq_prioritization.normalize_steps({'cities_first': {
        'campaign-2': {'conditions': ['all']}, 'final-2': {'conditions': ['repeated']},
        'campaign-4': {'queue': ''}, 'campaign-7': {'priority': 'P4', 'number': 3, 'queue': ''},
        'campaign-5': {'nr_mode': 'any', 'call_type': 'data', 'conditions': ['consecutive'], 'queue': 'Voice'},
        'final-5': {'area': 'roads'}}})['cities_first']
    steps = nq_prioritization.strategy_steps(cities, chosen)
    assert decide(cities, single, steps)[1:3] == ('Campaign 2 · SA Data', 'P2')
    assert decide(cities, {**single, 'final': True, 'level': 'High'}, steps)[:3] == ('Track only', 'Final 11 · Tracked', 'P8')
    assert decide(cities, {**single, 'final': True, 'repeated': True}, steps)[1] == 'Final 2 · SA Data · Cluster'
    assert decide(cities, strategy_call(kind='whatsapp'), steps)[:5] == ('Track only', 'Campaign 3 · Tracked · NSA WhatsApp', 'P4', '', 3)
    assert decide(cities, {**single, 'nr_mode': 'NSA', 'consecutive': 'Yes'}, steps)[1:5] == (
        'Campaign 5 · Data · Consecutive failures', 'P5', 'Voice', 5)
    # The Area of a step takes the calls of any type that failed there: a medium impact WhatsApp call in a city is no longer
    # taken by step 5 of list 2, which takes those on Connecting Roads, but by step 7, every other SA WhatsApp call.
    whatsapp = strategy_call(kind='whatsapp', final=True, level='Medium')
    assert decide(cities, {**whatsapp, 'area': 'roads'}, steps)[1:5] == ('Final 5 · SA WhatsApp · Connecting Roads · Medium impact', 'P4', 'Voice', 5)
    assert decide(cities, whatsapp, steps)[1:5] == ('Final 7 · SA WhatsApp', 'P4', 'Voice', 7)
    # AND and OR mix, AND first: the consecutive failures of a hotspot during the campaign; once the Final CDR is scored, the
    # data failures with a High impact and of a cluster, or consecutive.
    chosen = nq_prioritization.normalize_steps({'cities_first': {
        'campaign-2': {'conditions': ['consecutive', 'cluster'], 'joins': ['and']},
        'final-2': {'conditions': ['high', 'repeated', 'consecutive'], 'joins': ['and', 'or']}}})['cities_first']
    steps = nq_prioritization.strategy_steps(cities, chosen)
    hotspot = strategy_call(kind='data', cluster_size=3, consecutive='No')
    assert decide(cities, hotspot, steps)[1:3] == ('Campaign 8 · Tracked', 'P8')
    assert decide(cities, {**hotspot, 'consecutive': 'Yes'}, steps)[1:3] == ('Campaign 2 · SA Data · Consecutive failures AND Hotspots', 'P2')
    assert decide(cities, {**single, 'consecutive': 'Yes'}, steps)[2] == 'P8'
    high_data = {**single, 'final': True, 'level': 'High'}
    assert decide(cities, high_data, steps)[2] == 'P8'
    assert decide(cities, {**high_data, 'repeated': True}, steps)[1:3] == ('Final 2 · SA Data · High impact AND Cluster OR Consecutive failures', 'P2')
    assert decide(cities, {**single, 'final': True, 'consecutive': 'Yes'}, steps)[2] == 'P2'
    # AND ( and OR ( start a group: one level of parentheses, AND first inside each group and between the groups.
    chosen = nq_prioritization.normalize_steps({'cities_first': {
        'final-1': {'conditions': ['high', 'medium', 'repeated'], 'joins': ['or', 'and_group']},
        'final-2': {'conditions': ['consecutive', 'high', 'repeated'], 'joins': ['and_group', 'or']}}}, strict=True)['cities_first']
    steps = nq_prioritization.strategy_steps(cities, chosen)
    voice = strategy_call(final=True)
    assert decide(cities, {**voice, 'level': 'Medium', 'repeated': True}, steps)[1:3] == (
        'Final 1 · SA Voice · (High impact OR Medium impact) AND Cluster', 'P1')
    assert decide(cities, {**voice, 'level': 'High'}, steps)[2] != 'P1'
    data = {**single, 'final': True, 'consecutive': 'Yes'}
    assert decide(cities, {**data, 'repeated': True}, steps)[1:3] == ('Final 2 · SA Data · Consecutive failures AND (High impact OR Cluster)', 'P2')
    assert decide(cities, data, steps)[2] != 'P2'
    assert nq_prioritization.step_conditions_text(['consecutive', 'cluster', 'high', 'repeated'], ['or', 'and_group', 'or']) == (
        '(Consecutive failures OR Hotspots) AND (High impact OR Cluster)')
    # List 2, once the Final CDR of the campaign is scored: the same in both strategies, by the impact levels.
    final = strategy_call(final=True)
    for strategy in (cities, roads):
        assert decide(strategy, {**final, 'level': 'High'})[1:6] == ('Final 1 · SA Voice · High impact', 'P1', 'Voice', 1, 'final')
        assert decide(strategy, {**final, 'kind': 'data', 'level': 'High'})[2:5] == ('P2', 'Data', 2)
        assert decide(strategy, {**final, 'level': 'Medium'})[2:5] == ('P4', 'Voice', 4)
        assert decide(strategy, final)[1:5] == ('Final 6 · SA Voice', 'P4', 'Voice', 6)
        assert [decide(strategy, {**final, 'kind': 'whatsapp', 'level': level})[2:5] for level in ('High', 'Medium', 'Low')] == [
            ('P3', 'Voice', 3), ('P4', 'Voice', 5), ('P4', 'Voice', 7)]
        assert decide(strategy, {**final, 'nr_mode': 'NSA'})[2:5] == ('P5', 'Voice', 8)
        assert decide(strategy, {**final, 'nr_mode': 'NSA', 'kind': 'data', 'level': 'High'})[2:5] == ('P6', 'Data', 9)
        assert decide(strategy, {**final, 'nr_mode': 'NSA', 'kind': 'whatsapp'})[2:5] == ('P7', 'Voice', 10)
        assert decide(strategy, {**final, 'kind': 'data', 'repeated': True})[:3] == ('Track only', 'Final 11 · Tracked', 'P8')
    # Dual Queue (Scoring Based) uses list 2 as its only list, also once the Final CDR is scored; Vodafone Proposal takes every SA,
    # then NSA, Voice and WhatsApp call, the Data calls of a cluster and every other call with P5 in its Unique queue, and tracks
    # the Ping tests and Excluded Calls with P8.
    custom, vodafone = nq_prioritization.CUSTOM, nq_prioritization.VODAFONE_PROPOSAL
    assert decide(custom, {**final, 'level': 'High'}, final_list=False)[1:6] == ('Campaign 1 · SA Voice · High impact', 'P1', 'Voice', 1, 'campaign')
    assert decide(custom, {**final, 'nr_mode': 'NSA', 'kind': 'data', 'level': 'High'}, final_list=False)[2:5] == ('P6', 'Data', 9)
    assert decide(vodafone, final, final_list=False)[1:6] == ('Campaign 1 · SA Voice', 'P1', 'Unique', 1, 'campaign')
    assert decide(vodafone, strategy_call(kind='whatsapp'))[2:5] == ('P1', 'Unique', 2)
    assert decide(vodafone, {**single, 'repeated': True})[1:5] == ('Campaign 3 · SA Data · Cluster', 'P2', 'Unique', 3)
    assert decide(vodafone, strategy_call(nr_mode='NSA', kind='whatsapp'))[2:5] == ('P3', 'Unique', 5)
    assert decide(vodafone, {**single, 'nr_mode': 'NSA', 'repeated': True})[2:5] == ('P4', 'Unique', 6)
    assert decide(vodafone, single)[:5] == ('Prioritized', 'Campaign 7 · Every call', 'P5', 'Unique', 7)
    for call in (strategy_call(kind='ping'), strategy_call(excluded=True)):
        assert decide(vodafone, call)[:4] == ('Track only', 'Campaign 8 · Tracked', 'P8', '')


def failure(key, minute, latitude, site='', operator='VF', day=1, **values):
    return {'call_key': key, 'operator': operator, 'nr_mode': 'SA', 'campaign': '2026-Q3', 'place_key': f'site:|{site}' if site else '',
            'start_time': f'2026-09-0{day} 10:{minute:02d}:00', 'latitude': latitude, 'longitude': -0.12, **values}


def test_hotspots_link_the_failures_close_in_time_or_distance_whatever_their_call_type():
    calls = [
        failure('a', 0, 51.5000), failure('b', 10, 51.5010), failure('c', 24, 51.5020),  # each within 15 min and 250 m of the one before
        failure('d', 50, 51.5030),  # 26 minutes later
        failure('e', 52, 51.6000, site='s1'), failure('f', 55, 51.7000, site='s1'),  # far, but ended at the same site
        failure('g', 0, 51.5000, operator='EE'), failure('h', 1, 51.5000, day=2),  # another operator, another day
        failure('k', 30, 51.9000), failure('m', 40, 51.9900),  # 10 minutes apart, 10 km apart
        failure('i', 0, 51.5000, start_time='')]  # no date
    found = nq_prioritization.hotspots(calls, join='and')
    assert found['a'][0] == found['b'][0] == found['c'][0] and found['a'][1] == 3
    assert found['d'][1] == 1 and found['e'][0] == found['f'][0] and found['e'][1] == 2
    assert found['g'][1] == found['h'][1] == found['i'][1] == found['k'][1] == found['m'][1] == 1
    # A failure linked to two others joins them, whichever was first: a chain of links is one hotspot.
    chain = nq_prioritization.hotspots([failure('x', 0, 51.5000), failure('y', 14, 51.5040), failure('z', 7, 51.5020)], join='and')
    assert len({key for key, _size in chain.values()}) == 1 and chain['x'] == ('cluster:x', 3)
    # OR (the default): close in time wherever they are, or close in distance at any time of the campaign.
    either = nq_prioritization.hotspots(calls)
    assert either['k'][0] == either['m'][0] == either['d'][0] == either['a'][0]
    assert either['i'][0] == either['h'][0] == either['a'][0] and either['g'][1] == 1


def test_consecutive_failures_have_another_failure_minutes_apart_wherever_it_was():
    calls = [failure('a', 0, 51.5000), failure('b', 9, 52.0000),  # 9 minutes apart, 55 km apart
             failure('c', 40, 51.5000), failure('d', 41, 53.0000), failure('e', 41, 51.5000, operator='EE'),
             failure('f', 0, 51.5000, day=2), failure('g', 0, 51.5, day=3, start_time='')]
    found = nq_prioritization.consecutive_failures(calls, 10)
    assert (found['a'], found['b'], found['c'], found['d'], found['e'], found['f'], found['g']) == ('Yes', 'Yes', 'Yes', 'Yes', 'No', 'No', '')
    assert nq_prioritization.consecutive_failures(calls, 5)['a'] == 'No'


def test_clusters_take_the_quarter_of_the_campaign_and_leave_the_ping_tests_out():
    # A call of the campaign of Q2 driven in July is in Q2; the date counts only when the campaign does not name its quarter.
    assert nq_prioritization.quarter_of('2026-07-02 10:00:00', 'UK_Q2_2026') == nq_prioritization.quarter_of('', '2026-Q2-SA') == 2026 * 4 + 1
    assert nq_prioritization.quarter_of('2026-07-02 10:00:00', 'Drive test') == 2026 * 4 + 2
    def call(key, test_name='HTTP Browsing'):
        return {**failure(key, 0, 51.5000), 'service': 'data', 'test_name': test_name, 'session_type': '', 'quarter': 8106,
                'campaign_label': 'Q3', 'place_key': 'spot:250:1:1'}
    settings = nq_prioritization.default_settings()
    linked = nq_prioritization.linked_values([call('a'), call('p', 'Ping')], settings)
    assert linked['repetitions']['a'] == (1, 1) and not any('p' in linked[name] for name in ('repetitions', 'hotspots', 'runs', 'clusters'))
    linked = nq_prioritization.linked_values([call('a'), call('b')], settings)
    assert linked['repetitions']['a'] == (2, 1) and linked['clusters']['a'] == linked['clusters']['b'] == ('group:a', 2)
    assert linked['clusters_all']['a'] == ('group-all:a', 2)
    # A call out of scope is in no group and makes none.
    linked = nq_prioritization.linked_values([call('a'), call('b')], settings, out_of_scope=lambda row: row['call_key'] == 'b')
    assert linked['repetitions']['a'] == (1, 1) and linked['hotspots']['a'][1] == linked['runs']['a'][1] == 1 and 'a' not in linked['clusters']
    assert not any('b' in linked[name] for name in ('repetitions', 'hotspots', 'runs', 'clusters'))


def test_clusters_hold_the_calls_linked_by_distance_or_site_and_those_they_count():
    def call(key, latitude, place='', quarter=8106, operator='EE'):
        return {'call_key': key, 'operator': operator, 'nr_mode': 'SA', 'place_key': place or 'spot:250:1:1', 'quarter': quarter,
                'campaign_label': 'Q', 'latitude': latitude, 'longitude': -0.12, 'start_time': f'2026-09-01 10:{key[-1]}0'}
    # a and b 220 m apart, b and c 220 m apart: one cluster through b; d at the site of e far away; f alone.
    calls = [call('a1', 51.5000), call('b2', 51.5020), call('c3', 51.5040), call('d4', 52.0, 'site:|s9'), call('e5', 53.0, 'site:|s9'),
             call('f6', 54.0), call('g7', 51.5000, operator='VF')]
    counts = nq_prioritization.repetitions(calls, 1, 250)
    found = nq_prioritization.cluster_groups(calls, counts, 2, 1, 250)
    assert found['a1'] == found['b2'] == found['c3'] == ('group:a1', 3) and found['d4'] == found['e5'] == ('group:d4', 2)
    assert 'f6' not in found and 'g7' not in found
    # A cluster of the next campaign counts a call of the one before, which belongs to it although it is alone in its own.
    calls = [call('x1', 51.5000, quarter=8105), call('y2', 51.5010, quarter=8106)]
    counts = nq_prioritization.repetitions(calls, 1, 250)
    assert counts['x1'][0] == 1 and nq_prioritization.cluster_groups(calls, counts, 2, 1, 250)['x1'] == ('group:x1', 2)
    assert nq_prioritization.cluster_groups(calls, nq_prioritization.repetitions(calls, 0, 250), 2, 0, 250) == {}


def test_queues_take_priority_then_the_newest_campaign_step_gain_and_oldest():
    def row(key, priority, campaign, step, gain, start):
        return {'call_key': key, 'priority': priority, 'campaign': campaign, 'queue_step': step, 'expected_points': gain,
                'start_time': start}
    rows = [row('old-q2', 'P1', '2026-Q2', 1, 2.0, '2026-06-01'), row('p3', 'P3', '2026-Q3', 3, 9.0, '2026-09-01'),
            row('low', 'P1', '2026-Q3', 1, 0.1, '2026-09-01'), row('high', 'P1', '2026-Q3', 1, 0.5, '2026-09-03'),
            row('none-new', 'P1', '2026-Q3', 1, None, '2026-09-05'), row('none-old', 'P1', '2026-Q3', 1, None, '2026-09-02')]
    order = [item['call_key'] for item in nq_prioritization.order_queue(rows, [f'P{number}' for number in range(1, 9)])]
    assert order == ['high', 'low', 'none-old', 'none-new', 'old-q2', 'p3']


def test_final_campaigns_and_the_latest_one_before_a_campaign():
    finals = {('SA', '2026-q2-sa'): {'voice', 'data'}, ('SA', '2025-q4-sa'): {'voice'}, ('NSA', '2026-q2'): {'voice'},
              ('', '2026-q1'): {'speech'}}
    assert nq_prioritization.has_final(finals, 'SA', '2026-q2-sa', 'data') and not nq_prioritization.has_final(finals, 'SA', '2026-q2-sa', 'speech')
    assert nq_prioritization.has_final(finals, 'NSA', '2026-q1') and not nq_prioritization.has_final(finals, 'SA', '2026-q3-sa')
    assert nq_prioritization.previous_final_campaign(finals, {}, 'SA', 'UK_Q3_SA_2026') == '2026-q2-sa'
    assert nq_prioritization.previous_final_campaign(finals, {('SA', '2025-q4-sa'): 'UK_Q4_2025_SA'}, 'SA', 'UK_Q1_SA_2026') == 'UK_Q4_2025_SA'
    assert nq_prioritization.previous_final_campaign(finals, {}, 'NSA', 'UK_Q3_NSA_2026') == '2026-q2'
    assert nq_prioritization.previous_final_campaign(finals, {}, 'SA', 'UK_Q4_2025_SA') == ''


def test_repetitions_count_the_calls_of_the_operator_near_the_call_in_its_campaign_and_those_before():
    calls = [
        {'call_key': 'a', 'operator': 'EE', 'place_key': 'site:|s1', 'quarter': 8100, 'campaign_label': 'Q1'},
        {'call_key': 'b', 'operator': 'EE', 'place_key': 'site:|s1', 'quarter': 8101, 'campaign_label': 'Q2'},
        {'call_key': 'c', 'operator': 'EE', 'place_key': 'site:|s1', 'quarter': 8103, 'campaign_label': 'Q4'},
        {'call_key': 'd', 'operator': 'VF', 'place_key': 'site:|s1', 'quarter': 8101, 'campaign_label': 'Q2'},
        {'call_key': 'e', 'operator': 'EE', 'place_key': '', 'quarter': 8101, 'campaign_label': 'Q2'},
    ]
    # The Cluster period counts campaigns: the one before Q4 (8103) is Q2 (8101), as Q3 has none.
    found = nq_prioritization.repetitions(calls, 1)
    assert found['a'] == (1, 1) and found['b'] == (2, 2) and found['c'] == (2, 2) and found['d'] == (1, 1)
    assert 'e' not in found
    assert nq_prioritization.repetitions(calls, 0)['c'] == (1, 1) and nq_prioritization.repetitions(calls, 2)['c'] == (3, 3)
    # Within the Cluster distance of the call, at another site or at a spot, or at the same site wherever it is.
    def call(key, latitude, place, quarter=8101):
        return {'call_key': key, 'operator': 'EE', 'place_key': place, 'quarter': quarter, 'campaign_label': 'Q2',
                'latitude': latitude, 'longitude': -0.12}
    near = [call('p', 51.5000, 'site:|s2'), call('q', 51.5020, 'spot:250:1:1'), call('r', 51.5040, 'site:|s3'),
            call('s', 52.0000, 'site:|s2'), call('t', 51.5010, 'site:|s4', quarter=8104)]
    found = nq_prioritization.repetitions(near, 1, 250)
    # t, of the next campaign, counts those of the campaign before it near it; they never count the calls of later campaigns.
    assert (found['p'], found['q'], found['r'], found['s'], found['t']) == ((3, 1), (3, 1), (2, 1), (2, 1), (3, 1))
    assert nq_prioritization.repetitions(near, 1, 500)['p'] == (4, 1)


def test_former_descriptions_of_the_repetitions_take_the_campaigns(client):
    enable_module()
    nq.ensure_nq_tables(core.repository)
    former = dict(nq.V34_FORMER_DESCRIPTIONS)
    fields = nq.list_fields(core.repository)
    for field in fields:
        field['description'] = former.get(field['key'], field['description'])
    nq._store_fields(core.repository, nq_catalog.normalize_fields(fields, {field['key']: field for field in nq.list_fields(core.repository)},
                                                                  nq.CATALOG_CALL_FIELDS))
    # A workspace of the version before: the descriptions still the former defaults take the current ones, once.
    core.repository.set_workspace_state(nq.ANALYSIS_CENTER_V34_STATE_KEY, '')
    core.repository.set_workspace_state(nq.SCHEMA_STATE_KEY, '2026-10-26.1')
    nq.ensure_nq_tables(core.repository)
    defaults = {field['key']: field['description'] for field in nq_catalog.DEFAULT_FIELDS}
    descriptions = {field['key']: field['description'] for field in nq.list_fields(core.repository)}
    assert all(descriptions[key] == defaults[key] and 'campaign' in defaults[key] for key in former)


def test_steps_saved_with_consecutive_failures_and_cluster_take_hotspots_once(client):
    enable_module()
    settings = nq_prioritization.default_settings()
    settings['steps']['roads_first']['campaign-2'] = {**settings['steps']['roads_first']['campaign-2'], 'conditions': ['consecutive', 'repeated'],
                                                      'joins': ['and'], 'priority': 'P3'}
    settings['steps']['custom']['campaign-3'] = {**settings['steps']['custom']['campaign-3'], 'conditions': ['repeated', 'consecutive'], 'joins': ['or']}
    settings['steps']['custom']['campaign-4'] = {**settings['steps']['custom']['campaign-4'], 'conditions': ['consecutive', 'repeated', 'high'],
                                                 'joins': ['and', 'or']}
    core.repository.set_workspace_state(nq_prioritization.SETTINGS_STATE_KEY, json.dumps(settings))
    loaded = nq_prioritization.load_settings(core.repository)
    assert (loaded['steps']['roads_first']['campaign-2']['conditions'], loaded['steps']['roads_first']['campaign-2']['priority']) == (['cluster'], 'P3')
    assert loaded['steps']['custom']['campaign-3']['conditions'] == ['cluster']
    # Only the steps with those two conditions alone.
    assert loaded['steps']['custom']['campaign-4']['conditions'] == ['consecutive', 'repeated', 'high']
    log = next(row for row in core.repository.list_logs(limit=20) if row['action'] == 'nq_prioritization_settings')
    assert 'roads_first · campaign-2 · Consecutive failures and Cluster → Hotspots' in json.loads(log['details'])['steps_changed']
    # Once: a step saved with them afterwards keeps them.
    stored = json.loads(core.repository.get_workspace_state(nq_prioritization.SETTINGS_STATE_KEY))
    stored['steps']['roads_first']['campaign-2']['conditions'] = ['consecutive', 'repeated']
    core.repository.set_workspace_state(nq_prioritization.SETTINGS_STATE_KEY, json.dumps(stored))
    assert nq_prioritization.load_settings(core.repository)['steps']['roads_first']['campaign-2']['conditions'] == ['consecutive', 'repeated']


def test_workspaces_of_the_previous_schema_get_the_runs_of_consecutive_failures(client):
    enable_module()
    nq.ensure_nq_tables(core.repository)
    with core.repository.connection() as connection:
        for column in ('consecutive_key', 'consecutive_size', 'repeat_key', 'repeat_all_key'):
            connection.execute(f'ALTER TABLE {nq_prioritization.PRIORITY_TABLE} DROP COLUMN {column}')
    core.repository.set_workspace_state(nq.SCHEMA_STATE_KEY, '2026-10-26.2')
    nq.ensure_nq_tables(core.repository)
    with core.repository.connection() as connection:
        columns = {row[1] for row in connection.execute(f'PRAGMA table_info({nq_prioritization.PRIORITY_TABLE})')}
    assert {'consecutive_key', 'consecutive_size', 'repeat_key', 'repeat_all_key'} <= columns and 'drive' not in columns


def test_place_of_a_call_falls_back_from_the_site_of_its_last_cell_to_a_spot():
    def lookup(call, node):
        return {'site': '47258', 'host': 'Vodafone'} if node == ('eNB', 1000) and call['operator'] == 'VF' else None
    assert nq_prioritization.place_of({'cell_id': str(1000 * 256 + 1), 'operator': 'VF', 'site_id': 'S9'}, lookup, 100) == (
        'site:vodafone|47258', 'Site 47258 (Vodafone)')
    # The site that served its start is the same place as that of the cell where another call ended.
    assert nq_prioritization.place_of({'cell_id': str(2000 * 256 + 1), 'operator': 'VF', 'site_id': '47258',
                                       'host_network': 'Vodafone'}, lookup, 100) == ('site:vodafone|47258', 'Site 47258 (Vodafone)')
    assert nq_prioritization.place_of({'cell_id': str(2000 * 256 + 1), 'operator': 'VF', 'site_id': 'S9', 'host_network': 'H3G'},
                                      lookup, 100) == ('site:h3g|s9', 'Site S9 (H3G)')
    assert nq_prioritization.place_of({'cell_id': str(2000 * 256 + 1), 'operator': 'VF'}, lookup, 100) == ('node:eNB 2000', 'eNB 2000')
    near = nq_prioritization.place_of({'latitude': 51.50001, 'longitude': -0.12001}, lookup, 100)
    assert near[0].startswith('spot:') and near == nq_prioritization.place_of({'latitude': 51.50002, 'longitude': -0.12002}, lookup, 100)
    assert nq_prioritization.place_of({}, lookup, 100) == ('', '')


def test_the_site_of_a_node_is_that_of_a_network_serving_the_operator_near_the_call(client, tmp_path):
    """Ticket #9: node numbers are only unique within a network. The eNB 4643 of a Vodafone site in Cardiff and of an H3G
    site near Hull (named ENB0, the placeholder of the H3G rows of the VFUK mapping) never put the calls of both at one
    site: each call takes the node of a network that serves its operator (its own first) within the Serving Networks
    distance of it, and a node takes the sites of the newest snapshot that lists it."""
    from src.modules import network_inventory

    cardiff, hull, london = (51.4857, -3.2041), (53.8031, -0.5652), (51.5072, -0.1276)

    def add_inventory(name, kind, rows, snapshot_date):
        dataset_id = add_cdr(tmp_path, name, kind, pd.DataFrame(rows), '')
        core.repository.update_dataset_profile(dataset_id, snapshot_date=snapshot_date)

    add_inventory('Multivendor_Mapping_VFUK.csv', 'mapping_vodafone', [
        {'Network': 'Vodafone', 'Host Network': 'Vodafone', 'Site ID': 47258, 'eNodeB': 'V0001320H047258', 'eNodeB ID': 4643,
         'Cell ID': 1, 'Latitude': cardiff[0], 'Longitude': cardiff[1], 'Vendor': 'Ericsson'},
        {'Network': 'H3G', 'Host Network': 'H3G', 'Site ID': 43880, 'eNodeB': 'ENB0', 'OssName': 'ENB0', 'eNodeB ID': 4643,
         'Cell ID': 1, 'Latitude': hull[0], 'Longitude': hull[1], 'Vendor': 'Samsung'},
        {'Network': 'Telefonica', 'Host Network': 'Vodafone', 'Site ID': 18890, 'eNodeB ID': 5306, 'Cell ID': 1,
         'Latitude': hull[0], 'Longitude': hull[1], 'Vendor': 'Ericsson'},
    ], '2026-08-20')
    network_inventory.ensure_inventory_cells(core.repository)
    lookup = nq.node_site_lookup(core.repository)

    def site(operator, node, position=None):
        found = lookup({'operator': operator, 'latitude': position and position[0], 'longitude': position and position[1]}, ('eNB', node))
        return found and (found['site'], found['host'])

    assert site('VF', 4643, (51.49, -3.21)) == ('47258', 'Vodafone')
    # Its own Vodafone node is too far: the H3G node near it serves Vodafone through eMOCN.
    assert site('VF', 4643, (53.84, -0.47)) == ('43880', 'H3G')
    assert site('3', 4643, (53.84, -0.47)) == ('43880', 'H3G') and site('3', 4643, (51.49, -3.21)) == ('47258', 'Vodafone')
    # Without a position its own network first; no node of its networks near it, no site.
    assert site('VF', 4643) == ('47258', 'Vodafone') and site('VF', 4643, london) is None
    # The Telefonica nodes serve O2, never Vodafone.
    assert site('O2', 5306, hull) == ('18890', 'Vodafone') and site('VF', 5306, hull) is None
    # Each call of the eNB is at its own site: the start site of a call at the same site is the same place.
    places = {nq_prioritization.place_of({'operator': 'VF', 'cell_id': str(4643 * 256 + 1), 'latitude': lat, 'longitude': lon},
                                         lookup, 250)[1] for lat, lon in ((51.49, -3.21), (53.84, -0.47))}
    assert places == {'Site 47258 (Vodafone)', 'Site 43880 (H3G)'}
    assert nq_prioritization.place_of({'operator': 'VF', 'cell_id': '1', 'site_id': '47258', 'host_network': 'Vodafone'},
                                      lookup, 250)[1] == 'Site 47258 (Vodafone)'
    # A newer snapshot that lists the node gives its sites.
    add_inventory('260917_Cell_Config_LTE_NR_KMZ.csv', 'network_inventory', [
        {'Network': 'H3G', 'Host Network': 'H3G', 'Site ID': 'H43880', 'eNodeB ID': 4643, 'Cell ID': 1,
         'Latitude': hull[0], 'Longitude': hull[1], 'Vendor': 'Samsung'}], '2026-09-17')
    network_inventory.ensure_inventory_cells(core.repository)
    lookup = nq.node_site_lookup(core.repository)
    assert site('VF', 4643, (53.84, -0.47)) == ('H43880', 'H3G') and site('VF', 4643, (51.49, -3.21)) == ('47258', 'Vodafone')


# -- in a workspace -----------------------------------------------------------------------------------
def add_cdr(tmp_path: Path, name: str, kind: str, rows: pd.DataFrame, nr_mode: str) -> int:
    source = tmp_path / name
    source.write_text('test source', encoding='utf-8')
    dataset_id, _created = core.repository.add_dataset(name, str(source), 'super')
    core.repository.replace_dataset_rows(dataset_id, rows)
    core.repository.update_dataset_profile(dataset_id, status='ready', progress=100, dataset_kind=kind, nr_mode=nr_mode,
                                           row_count=len(rows), column_count=len(rows.columns), processed_at=local_now_iso())
    return dataset_id


def nq_frames(prefix: str, campaigns=('2026-Q3',), seed=11) -> dict[str, pd.DataFrame]:
    """Synthetic CDR rows with their result and a JOIN_ID per call (the samples of a Speech call share it)."""
    frames = synthetic_cdr_frames(seed=seed, campaigns=campaigns, repeats=1)
    for kind, frame in frames.items():
        if kind == 'voice':
            frame['status'] = frame['Call_Status']
            frame['JOIN_ID'] = [f'{prefix}V{index}' for index in range(len(frame))]
        elif kind == 'data':
            frame['status'] = frame['Test_Result']
            frame['JOIN_ID'] = [f'{prefix}D{index}' for index in range(len(frame))]
        else:
            frame['status'] = ['Completed' if value > 1.6 else 'Failed' for value in frame['LQ']]
            frame['JOIN_ID'] = [f'{prefix}S{index // 2}' for index in range(len(frame))]
    return frames


def prepare(client, tmp_path, strategy: str) -> dict:
    """A workspace with NSA and SA CDRs of a campaign, its Scoring History and the calls indexed and prioritized with a
    strategy."""
    enable_module()
    core.repository.replace_scoring_configuration(load_initial_scoring_configuration())
    core.repository.set_workspace_state(nq_prioritization.SETTINGS_STATE_KEY,
                                        json.dumps({**nq_prioritization.default_settings(), 'strategy': strategy}))
    frames = {'NSA': nq_frames('N'), 'SA': nq_frames('A', seed=12)}
    for mode, kinds in frames.items():
        for kind, rows in kinds.items():
            add_cdr(tmp_path, f'NetCheck_UK_CDR_{kind.title()}_{mode}_2026_Q3.xlsx', kind, rows, mode)
    scoring_history.sync_history(core.repository)
    nq.sync_nq_calls(core.repository)
    login(client)
    summary = nq_prioritization.calculate(core.repository)
    return {'client': client, 'frames': frames, 'summary': summary}


@pytest.fixture
def prioritized(client, tmp_path):
    """Prioritized with the Custom rules."""
    return prepare(client, tmp_path, nq_prioritization.CUSTOM)


@pytest.fixture
def strategized(client, tmp_path):
    """Prioritized with the default strategy, Dual Queue (Cities first)."""
    return prepare(client, tmp_path, nq_prioritization.DEFAULT_STRATEGY)


def priority_rows() -> dict[str, dict]:
    with core.repository.connection() as connection:
        return {str(row['call_key']): dict(row) for row in connection.execute(
            f'SELECT p.*, l.service, l.nr_mode, l.operator, l.join_id, l.latitude, l.longitude, l.campaign, l.start_time, l.test_name, '
            f'l.session_type, l.region FROM {nq_prioritization.PRIORITY_TABLE} p '
            f'JOIN {nq.NQ_CALL_LATEST_TABLE} l ON l.call_key = p.call_key')}


def test_every_call_holds_its_share_of_the_points_its_kpis_lose(prioritized):
    rows = priority_rows()
    assert rows and prioritized['summary']['scored'] == len(rows)
    assert {row['impact_state'] for row in rows.values()} == {'scored'}
    campaigns = nq_prioritization.status(core.repository)['campaigns']
    assert sorted((item['nr_mode'], item['campaign'], item['state']) for item in campaigns) == [
        ('NSA', '2026-Q3', 'scored'), ('SA', '2026-Q3', 'scored')]
    # The points of the call success ratio the operators lose in the scoring of the NSA CDRs (calculated from their rows),
    # which every Non-Qualified Voice call shares.
    from src.modules.scoring_config import validate_scoring_configuration
    from src.modules.scoring_methodologies import methodology_for
    configuration = validate_scoring_configuration(methodology_for(core.repository.get_scoring_profiles(), '2026-Q3')['configuration'])
    result = scoring.calculate_scoring(prioritized['frames']['NSA'], ['Campaign', 'Operator'], configuration=configuration,
                                       operator_mappings=core.repository.list_operator_mappings())
    lost = {(row['operator'], row['environment']): row['max_points'] - row['weighted_points'] for row in result['scoring']
            if row['kpi_code'] == 'K1' and row['weighted_points'] is not None}
    held: dict[tuple, float] = {}
    for row in rows.values():
        if row['nr_mode'] != 'NSA':
            continue
        for item in json.loads(row['impact_json']):
            if item['code'] == 'K1':
                held[(row['operator'], item['environment'])] = held.get((row['operator'], item['environment']), 0) + item['points']
    assert held and set(held) <= set(lost)
    for key, points in held.items():
        assert points == pytest.approx(lost[key], abs=0.01)
    # The level of the calls of each campaign, operator and service is a Pareto of their points.
    groups: dict[tuple, dict[str, list[float]]] = {}
    for row in rows.values():
        group = groups.setdefault((row['campaign_key'], row['operator'], row['service']), {})
        group.setdefault(row['impact_level'], []).append(row['impact_points'])
    assert any('High' in group for group in groups.values())
    for group in groups.values():
        if 'High' in group:
            assert min(group['High']) >= max(group.get('Medium', [0]) + group.get('Low', [0]))
        assert all(points == 0 for points in group.get('No impact', []))


def test_steps_prioritize_the_calls_and_fill_their_priority_until_it_is_chosen_by_hand(prioritized):
    # Dual Queue (Scoring Based): its only list orders every call, also those of a campaign with its Final CDR, by its impact level.
    rows = priority_rows()
    steps = {step['rule']: step for step in nq_prioritization.strategy_steps(nq_prioritization.CUSTOM) if step['list'] == 'campaign'}
    prioritized_rows = [row for row in rows.values() if row['prioritization'] == 'Prioritized']
    assert prioritized_rows and {row['queue_list'] for row in prioritized_rows} == {'campaign'}
    for row in prioritized_rows:
        step = steps[row['priority_rule']]
        assert (row['suggested_priority'], row['queue'], row['queue_step']) == (step['priority'], step['queue'], step['number'])
    nsa_data = [row for row in rows.values() if row['nr_mode'] == 'NSA' and row['service'] == 'data']
    assert nsa_data and {row['priority_rule'] for row in nsa_data} <= {'Campaign 9 · NSA Data · High impact', 'Campaign 11 · Tracked'}
    sa_voice = [row for row in rows.values() if row['priority_rule'] == 'Campaign 1 · SA Voice · High impact']
    assert sa_voice and {row['suggested_priority'] for row in sa_voice} == {'P1'}
    # The Priority of every open call is the one of its step, recorded in its history.
    key = sa_voice[0]['call_key']
    with core.repository.connection() as connection:
        values = dict(connection.execute("SELECT call_key, value FROM nq_analysis_field_values WHERE field_key = 'priority'").fetchall())
        history = connection.execute("SELECT new_value, changed_by FROM nq_call_history WHERE call_key = ? AND field = 'field:priority'",
                                     (key,)).fetchall()
    assert values[key] == 'P1' and len(values) == sum(1 for row in rows.values() if row['suggested_priority'])
    assert [tuple(row) for row in history] == [('P1', 'system')]
    # Chosen by hand: the steps keep it, also when they change.
    client = prioritized['client']
    call = client.get(f'/api/non-qualified-calls/calls/{key}').json()
    response = client.patch(f'/api/non-qualified-calls/calls/{key}', json={'changes': {'fields': {'priority': 'P3'}},
                                                                           'version': call['call']['version']})
    assert response.status_code == 200, response.text
    settings = nq_prioritization.load_settings(core.repository)
    settings['steps']['custom']['campaign-1']['priority'] = 'P2'
    response = client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': settings})
    assert response.status_code == 200, response.text
    with core.repository.connection() as connection:
        values = dict(connection.execute("SELECT call_key, value FROM nq_analysis_field_values WHERE field_key = 'priority'").fetchall())
    assert values[key] == 'P3'
    assert {values[row['call_key']] for row in sa_voice if row['call_key'] != key} <= {'P2'}
    assert priority_rows()[key]['priority_filled'] == nq_prioritization.BY_HAND
    # A step takes a Priority from P1 to P8.
    settings['steps']['custom']['campaign-1']['priority'] = 'P9'
    response = client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': settings})
    assert response.status_code == 400 and 'from P1 to P8' in response.json()['detail']


def test_repeated_issues_count_every_service_and_nr_mode_of_the_operator(prioritized):
    # An NR Mode · operator pair out of scope: its calls leave the groups at once, as the settings are saved.
    operators = [row['operator'] for row in priority_rows().values() if row['nr_mode'] == 'SA']
    settings = nq_prioritization.load_settings(core.repository)
    response = prioritized['client'].put('/api/non-qualified-calls/prioritization/settings', json={'settings': {
        **settings, 'excluded_modes': [*settings['excluded_modes'], {'nr_mode': 'SA', 'operator': max(set(operators), key=operators.count)}]}})
    assert response.status_code == 200, response.text
    rows = priority_rows()
    repeated = [row for row in rows.values() if row['repeated'] == 'Yes']
    assert repeated and all(row['repetitions'] >= 2 for row in repeated)
    places: dict[tuple, set] = {}
    for row in rows.values():
        if row['place_key']:
            places.setdefault((row['operator'], row['place_key']), set()).add((row['service'], row['nr_mode']))
    # Some places repeat across services and across NR Modes.
    assert any(len({service for service, _mode in found}) > 1 for found in places.values())
    assert any(len({mode for _service, mode in found}) > 1 for found in places.values())
    # The calls of the operator and NR Mode within the Cluster distance of the call or at its site (one quarter here), the Ping
    # tests apart.
    meters = nq_prioritization.load_settings(core.repository)['spot_meters']
    kinds = {key: nq_prioritization.call_kind(row['service'], row['test_name'], row['session_type']) for key, row in rows.items()}
    # The Ping tests and the calls out of scope are in no cluster and make none.
    out = nq_prioritization.scope_test(core.repository, nq_prioritization.load_settings(core.repository))
    counted = {key: row for key, row in rows.items() if kinds[key] != 'ping' and not out(row) and row['place_key']}
    assert any(out(row) for row in rows.values())
    assert all(row['repetitions'] is None and row['repeated'] == 'No' and not row['cluster_key'] and not row['consecutive']
               for key, row in rows.items() if kinds[key] == 'ping' or out(row))
    for row in counted.values():
        site = '' if row['place_key'].startswith('spot:') else row['place_key']
        assert row['repetitions'] == sum(1 for other in counted.values() if other['operator'] == row['operator'] and other['nr_mode'] == row['nr_mode']
                                         and ((site and other['place_key'] == site) or nq_prioritization._meters(row, other) <= meters))


def test_hotspots_and_consecutive_failures_mix_the_call_types(prioritized):
    rows = priority_rows()
    kinds = {key: nq_prioritization.call_kind(row['service'], row['test_name'], row['session_type']) for key, row in rows.items()}
    # The Ping tests and the calls out of scope are neither consecutive nor in a hotspot.
    out = nq_prioritization.scope_test(core.repository, nq_prioritization.load_settings(core.repository))
    assert {(row['consecutive'], row['cluster_key']) for key, row in rows.items() if kinds[key] == 'ping' or out(row)} <= {('', '')}
    linked = {key: row for key, row in rows.items() if kinds[key] != 'ping' and not out(row)}
    assert {row['consecutive'] for row in linked.values()} <= {'Yes', 'No'} and any(row['consecutive'] == 'Yes' for row in linked.values())
    assert all(row['cluster_key'] and row['cluster_size'] >= 1 for row in linked.values())
    # A failure is consecutive in a run of 2 or more; with the OR join every run falls within one hotspot.
    assert all((row['consecutive'] == 'Yes') == ((row['consecutive_size'] or 0) >= 2) for row in linked.values())
    runs: dict[str, set] = {}
    for row in linked.values():
        runs.setdefault(row['consecutive_key'], set()).add(row['cluster_key'])
    assert all(len(found) == 1 for found in runs.values())
    # A hotspot mixes the call types of its operator, NR Mode and campaign.
    spots: dict[str, set] = {}
    for key, row in linked.items():
        spots.setdefault(row['cluster_key'], set()).add((row['operator'], row['nr_mode'], kinds[key]))
    assert all(len({(operator, mode) for operator, mode, _kind in found}) == 1 for found in spots.values())
    client = prioritized['client']
    settings = nq_prioritization.load_settings(core.repository)
    # With the clusters of the campaign alone (a Cluster period of 0), a failure is in a hotspot of OR exactly when it is
    # consecutive or in a cluster, as the rules take them (each of its NR Mode).
    response = client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': {**settings, 'previous_quarters': 0}})
    assert response.status_code == 200, response.text
    current = {key: row for key, row in priority_rows().items() if key in linked}
    assert all((row['consecutive'] == 'Yes' or row['repeated'] == 'Yes') == (row['cluster_size'] >= 2) for row in current.values())
    # The settings apply at once, from the stored values: a shorter Consecutive time leaves fewer consecutive failures and
    # the AND join smaller hotspots than OR (the default).
    assert settings['hotspot_join'] == 'or'
    before = sum(row['consecutive'] == 'Yes' for row in linked.values())
    response = client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': {
        **settings, 'consecutive_minutes': 1, 'hotspot_join': 'and'}})
    assert response.status_code == 200, response.text
    after = {key: row for key, row in priority_rows().items() if key in linked}
    assert sum(row['consecutive'] == 'Yes' for row in after.values()) <= before
    assert max(row['cluster_size'] for row in after.values()) <= max(row['cluster_size'] for row in linked.values())


def test_panel_filters_and_exports_show_the_prioritization(prioritized):
    from pptx import Presentation

    from src.modules import non_qualified_calls_export

    client = prioritized['client']
    panel = client.post('/api/non-qualified-calls/prioritization', json={'filters': {}}).json()
    rows = priority_rows()
    assert panel['totals']['calls'] == len(rows)
    assert panel['totals']['prioritized'] == sum(row['prioritization'] == 'Prioritized' for row in rows.values())
    # Dual Queue (Scoring Based) has one list: its steps, then the calls out of scope, each with its fields and its calls.
    steps = [step for step in nq_prioritization.strategy_steps(nq_prioritization.CUSTOM) if step['list'] == 'campaign']
    assert [rule['name'] for rule in panel['rules'][:len(steps)]] == [step['rule'] for step in steps]
    first = panel['rules'][0]
    assert (first['kind'], first['number'], first['list'], first['priority'], first['nr_mode'], first['call_type'], first['condition'],
            first['queue']) == ('step', 1, 'campaign', 'P1', 'SA', 'Voice', 'High impact', 'Voice')
    assert all(rule['waiting'] <= rule['open'] <= rule['calls'] for rule in panel['rules'])
    assert panel['rules'][len(steps)]['kind'] == 'out_of_scope'
    assert [item['name'] for item in panel['queues']] == ['Voice', 'Data'] and not panel['queue']
    voice = panel['queues'][0]
    assert voice['rows'] and len(voice['rows']) <= 20 and voice['rows'][0]['rule'] in {step['rule'] for step in steps}
    assert panel['issues'] and panel['kpis']
    # Its values are fields of the calls: columns, filters and breakdowns of the Summary.
    fields = {field['source_ref']: field for field in nq.list_fields(core.repository) if field['source'] == 'priority'}
    assert set(fields) == set(nq_catalog.PRIORITY_REFS)
    assert [option['name'] for option in fields['priority_rule']['options']] == nq_prioritization.rule_names(core.repository)
    filtered = client.post('/api/non-qualified-calls/calls', json={
        'filters': {'fields': {fields['prioritization']['key']: ['Track only']}}, 'page_size': 100}).json()
    assert filtered['total'] == panel['totals']['tracked']
    assert {call['values'][fields['prioritization']['key']] for call in filtered['calls']} == {'Track only'}
    labels = [breakdown['label'] for breakdown in filtered['breakdowns']]
    assert {'By Prioritization', 'By Scoring Impact', 'By Repeated Issue'} <= set(labels)
    ordered = client.post('/api/non-qualified-calls/calls', json={
        'filters': {}, 'sort': f"field:{fields['impact_points']['key']}", 'direction': 'desc', 'page_size': 10}).json()
    points = [call['values'][fields['impact_points']['key']] for call in ordered['calls']]
    assert points == sorted(points, reverse=True)
    # Every cluster (of every NR Mode together, as no NR Mode is filtered): its key filters its calls.
    issue = panel['issues'][0]
    assert issue['key'].startswith('group-all:') and issue['calls'] >= 2
    cluster = client.post('/api/non-qualified-calls/calls', json={'filters': {'clustered_issue': [issue['key']]}, 'page_size': 100}).json()
    assert cluster['total'] == issue['calls'] == issue['size']
    # Every hotspot, the most calls first, with its calls of each call type; its key filters its calls.
    spots = panel['hotspots']
    assert spots and [item['calls'] for item in spots] == sorted((item['calls'] for item in spots), reverse=True)
    spot = spots[0]
    assert spot['calls'] >= 2 and sum(spot['kinds'].values()) == spot['calls'] and set(spot['kinds']) == {'Voice', 'WhatsApp', 'Data'}
    assert spot['first'] <= spot['last'] and spot['place']
    calls = client.post('/api/non-qualified-calls/calls', json={'filters': {'hotspot': [spot['key']]}, 'page_size': 100}).json()
    assert calls['total'] == spot['calls'] == spot['size']
    # Every run of consecutive failures, with its calls of each call type; its key filters its calls.
    run = panel['consecutive'][0]
    assert run['calls'] >= 2 and sum(run['kinds'].values()) == run['calls'] and run['key'].startswith('run:')
    calls = client.post('/api/non-qualified-calls/calls', json={'filters': {'consecutive_issue': [run['key']]}, 'page_size': 100}).json()
    assert calls['total'] == run['calls'] == run['size']
    # The exports follow the order of the page: the queues, the steps (a table per list, then the calls outside the lists),
    # the KPIs, then the hotspots, the consecutive issues and the clustered issues (the hotspots first with the OR join).
    tables = non_qualified_calls_export.prioritization_tables(panel)
    order = [component for component, *_rest in tables]
    assert order.index('queue') < order.index('rules') < order.index('kpis') < order.index('hotspots') < min(order.index('consecutive'), order.index('issues'))
    # Between the consecutive and the clustered issues, the one with more calls first.
    total = {'consecutive': sum(item['calls'] for item in panel['consecutive']), 'issues': sum(item['calls'] for item in panel['issues'])}
    first, second = sorted(('consecutive', 'issues'), key=order.index)
    assert total[first] >= total[second]
    assert next(subtitle for component, subtitle, *_rest in tables if component == 'issues').startswith('Clustered Issues')
    assert next(subtitle for component, subtitle, *_rest in tables if component == 'hotspots').startswith(
        'Hotspots (Consecutives OR Clustered Issues)')
    # The clusters of the panel follow the NR Mode of the filters: one NR Mode filtered counts its calls alone.
    assert panel['cluster_mode'] == ''
    sa = client.post('/api/non-qualified-calls/prioritization', json={'filters': {'nr_mode': ['SA']}}).json()
    assert (sa['cluster_mode'], sa['cluster_mode_source']) == ('SA', 'NR Mode filter')
    # One NR Mode counts the clusters as the rules do (its own calls); several count every NR Mode together.
    assert panel['totals']['repeated'] == sum(bool(row['repeat_all_key']) for row in rows.values()) >= sum(bool(row['repeat_key']) for row in rows.values())
    assert sa['totals']['repeated'] == sum(bool(row['repeat_key']) for row in rows.values() if row['nr_mode'] == 'SA')
    # The other filters never choose it, whatever NR Modes they leave; a campaign of both NR Modes neither.
    single = next(item for item in sa['issues'] if set(item['nr_modes']) == {'SA'})
    by_site = client.post('/api/non-qualified-calls/prioritization', json={'filters': {'place': [f"{single['operator']}||{single['place']}"]}}).json()
    assert by_site['cluster_mode'] == ''
    campaign = next(row['campaign'] for row in rows.values() if row['nr_mode'] == 'SA')
    assert client.post('/api/non-qualified-calls/prioritization', json={'filters': {'campaign': [campaign]}}).json()['cluster_mode'] == ''
    # With the AND join, the consecutive and the clustered issues (the more general tables) go first.
    joined = non_qualified_calls_export.prioritization_tables({**panel, 'settings': {**panel['settings'], 'hotspot_join': 'and'}})
    order_and = [component for component, *_rest in joined]
    assert max(order_and.index('consecutive'), order_and.index('issues')) < order_and.index('hotspots')
    assert next(subtitle for component, subtitle, *_rest in joined if component == 'hotspots').startswith('Hotspots (Consecutives AND')
    # The three tables share their columns.
    columns = {component: columns for component, _subtitle, columns, *_rest in tables}
    assert [columns[name][0] for name in ('hotspots', 'consecutive', 'issues')] == ['Hotspot ID', 'Consecutive ID', 'Cluster ID']
    assert columns['hotspots'][1:] == columns['consecutive'][1:] == columns['issues'][1:] == [
        'Operator', 'NR Mode', 'Campaign', 'Site(s)', 'From', 'To', 'Calls', 'Voice', 'WhatsApp', 'Data', 'Points lost', 'Open', 'Main Root Category']
    # The IDs of the groups: short, of their letter, one for each group of a table.
    ids = [item['id'] for item in panel['issues']]
    assert len(set(ids)) == len(ids) and all(value.startswith('C-') for value in ids) and spot['id'].startswith('H-') and run['id'].startswith('R-')
    assert sum(issue['kinds'].values()) == issue['calls'] == sum(issue['campaign_names'].values()) and issue['first'] <= issue['last']
    # From and To: the day, the month in letters and the time, without the year.
    assert non_qualified_calls_export.group_time('2026-11-20 16:18:44.5') == '20 Nov 16:18'
    assert non_qualified_calls_export.group_time('') == '—'
    subtitles = [subtitle for component, subtitle, *_rest in tables if component == 'rules']
    assert subtitles == ['Dual Queue (Scoring Based) · List 1 · Weekly & Final CDRs', 'Dual Queue (Scoring Based) · Outside the lists']
    columns = next(columns for component, _subtitle, columns, *_rest in tables if component == 'rules')
    assert columns == ['Priority', 'Step', 'NR Mode', 'Operators', 'Call Type', 'Area', 'Condition', 'Calls', 'Open', 'Waiting', 'Points lost', 'Queue']
    # Campaigns of one NR Mode choose it.
    with core.repository.connection() as connection:
        connection.execute(f"UPDATE {nq.NQ_CALL_LATEST_TABLE} SET campaign = campaign || '-SA' WHERE nr_mode = 'SA'")
    by_campaign = client.post('/api/non-qualified-calls/prioritization', json={'filters': {'campaign': [f'{campaign}-SA']}}).json()
    assert (by_campaign['cluster_mode'], by_campaign['cluster_mode_source']) == ('SA', 'campaign filter')
    with core.repository.connection() as connection:
        connection.execute(f"UPDATE {nq.NQ_CALL_LATEST_TABLE} SET campaign = substr(campaign, 1, length(campaign) - 3) WHERE nr_mode = 'SA'")
    deck = client.post('/api/non-qualified-calls/export/powerpoint', json={'filters': {}})
    titles = [slide.shapes.title.text_frame.text for slide in Presentation(io.BytesIO(deck.content)).slides if slide.shapes.title is not None]
    assert any(title.startswith('NQ Prioritization & Queues') for title in titles)


def test_excluded_calls_lose_no_points(prioritized):
    rows = priority_rows()
    target = next(row for row in rows.values() if row['service'] == 'voice' and row['impact_points'])
    scoring_exclusions.create_list(core.repository, name='Faulty phone', calls=[{'join_id': target['join_id']}], username='super')
    scoring_exclusions.match_pending(core.repository)
    scoring_history.sync_history(core.repository)
    assert nq_prioritization.needs_calculation(core.repository)
    nq_prioritization.calculate(core.repository)
    row = priority_rows()[target['call_key']]
    assert (row['impact_state'], row['impact_level'], row['impact_points']) == ('excluded', 'Excluded from Scoring', 0.0)
    assert not nq_prioritization.needs_calculation(core.repository)


def test_campaigns_wait_for_the_scoring_history(client, tmp_path):
    enable_module()
    core.repository.replace_scoring_configuration(load_initial_scoring_configuration())
    for kind, rows in nq_frames('W').items():
        add_cdr(tmp_path, f'NetCheck_UK_CDR_{kind.title()}_2026_Q3.xlsx', kind, rows, 'NSA')
    nq.sync_nq_calls(core.repository)
    nq_prioritization.calculate(core.repository)
    campaign = nq_prioritization.status(core.repository)['campaigns'][0]
    assert campaign['state'] == 'pending' and 'Scoring History' in campaign['detail']
    assert {row['impact_level'] for row in priority_rows().values()} == {'Not scored'}
    # The repetitions do not wait.
    assert any(row['repeated'] == 'Yes' for row in priority_rows().values())


def test_only_admins_change_the_rules_and_editors_recalculate(prioritized):
    client = prioritized['client']
    workspace_user('editor', 'user-editor')
    login(client, 'editor', 'editor123')
    settings = nq_prioritization.load_settings(core.repository)
    assert client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': settings}).status_code == 403
    response = client.post('/api/non-qualified-calls/prioritization/recalculate')
    assert response.status_code == 200, response.text
    assert client.get('/api/non-qualified-calls/prioritization/settings').json()['settings']['strategy'] == nq_prioritization.CUSTOM
    assert 'nq_prioritization_recalculate' in [row['action'] for row in core.repository.list_logs(limit=50)]
    # For now EDANREC changes the Prioritization Rules too, without being an admin.
    workspace_user('EDANREC', 'user-editor')
    login(client, 'EDANREC', 'EDANREC123')
    assert client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': settings}).status_code == 200


def test_the_panel_chooses_the_strategy_of_the_module(strategized):
    client = strategized['client']
    panel = client.post('/api/non-qualified-calls/prioritization', json={'filters': {}}).json()
    assert panel['strategy']['key'] == 'cities_first' and panel['strategies']['vodafone_proposal'] == 'Vodafone Proposal'
    # The other settings stay as saved: only the strategy changes, and every call takes its step at once.
    nq_prioritization.save_settings(core.repository, {**nq_prioritization.load_settings(core.repository), 'high_share': 40}, 'super')
    response = client.put('/api/non-qualified-calls/prioritization/strategy', json={'strategy': 'vodafone_proposal'})
    assert response.status_code == 200, response.text
    settings = nq_prioritization.load_settings(core.repository)
    assert (settings['strategy'], settings['high_share']) == ('vodafone_proposal', 40)
    assert {row['queue'] for row in priority_rows().values() if row['queue']} == {'Unique'}
    log = next(row for row in core.repository.list_logs(limit=50) if row['action'] == 'nq_prioritization_strategy')
    assert json.loads(log['details'])['to'] == 'Vodafone Proposal'
    assert client.put('/api/non-qualified-calls/prioritization/strategy', json={'strategy': 'fifo'}).status_code == 400
    # The same users as the Prioritization Rules: admins and super-admins, and for now EDANREC.
    workspace_user('editor', 'user-editor')
    login(client, 'editor', 'editor123')
    assert client.put('/api/non-qualified-calls/prioritization/strategy', json={'strategy': 'roads_first'}).status_code == 403
    page = client.get('/non-qualified-calls').text
    assert 'id="nq-priority-panel-strategy" aria-label="Strategy of the NQ Prioritization" disabled' in page
    workspace_user('EDANREC', 'user-editor')
    login(client, 'EDANREC', 'EDANREC123')
    assert client.put('/api/non-qualified-calls/prioritization/strategy', json={'strategy': 'roads_first'}).status_code == 200
    assert nq_prioritization.load_settings(core.repository)['strategy'] == 'roads_first'


def test_strategies_travel_with_the_configuration(prioritized):
    settings = nq_prioritization.load_settings(core.repository)
    settings['steps']['custom']['campaign-1']['priority'] = 'P2'
    settings.update(high_share=40, strategy=nq_prioritization.VODAFONE_PROPOSAL, strategy_names={'custom': 'Impact first'},
                    excluded_modes=[{'nr_mode': 'SA', 'operator': 'Three'}])
    nq_prioritization.save_settings(core.repository, settings, 'super')
    document = nq.export_configuration_document(core.repository)
    assert json.loads(document)['prioritization']['high_share'] == 40
    core.repository.set_workspace_state(nq_prioritization.SETTINGS_STATE_KEY, json.dumps(nq_prioritization.default_settings()))
    nq.import_configuration_document(core.repository, document)
    restored = nq_prioritization.load_settings(core.repository)
    assert (restored['high_share'], restored['strategy'], restored['strategy_names'], restored['excluded_modes']) == (
        40, nq_prioritization.VODAFONE_PROPOSAL, {'custom': 'Impact first'}, [{'nr_mode': 'SA', 'operator': 'Three'}])
    assert restored['steps']['custom']['campaign-1']['priority'] == 'P2'


def test_workspaces_of_earlier_versions_get_the_tables_and_fields(client, tmp_path):
    enable_module()
    nq.ensure_nq_tables(core.repository)
    keys = nq.V23_FIELD_KEYS
    fields = [field for field in nq.list_fields(core.repository) if field['key'] not in keys]
    nq._store_fields(core.repository, fields)
    with core.repository.connection() as connection:
        connection.execute(f'DROP TABLE {nq_prioritization.PRIORITY_TABLE}')
        connection.execute("DELETE FROM workspace_state WHERE key IN (?, ?)", (nq.ANALYSIS_CENTER_V23_STATE_KEY, nq.SCHEMA_STATE_KEY))
    nq.ensure_nq_tables(core.repository)
    with core.repository.connection() as connection:
        assert connection.execute("SELECT 1 FROM sqlite_master WHERE name = ?", (nq_prioritization.PRIORITY_TABLE,)).fetchone()
    order = [field['key'] for field in nq.list_fields(core.repository)]
    assert keys <= set(order) and order.index('prioritization') == order.index('priority') + 1
    shown = {field['key'] for field in nq.list_fields(core.repository) if field['in_table']}
    assert {'prioritization', 'scoring_points_lost', 'repeated_issue'} <= shown and 'scoring_impact' not in shown


def test_workspaces_of_earlier_versions_get_the_strategy_fields_and_priorities(client, tmp_path):
    enable_module()
    nq.ensure_nq_tables(core.repository)
    fields = [field for field in nq.list_fields(core.repository) if field['key'] not in nq.V24_FIELD_KEYS]
    priority = next(field for field in fields if field['key'] == 'priority')
    # P1 has a colour of its own, the others none.
    priority['options'] = [{'name': f'P{number}', 'color': '#123456' if number == 1 else ''} for number in range(1, 6)]
    nq._store_fields(core.repository, fields)
    with core.repository.connection() as connection:
        # The table of the previous version, without the columns of the strategies.
        connection.execute(f'DROP TABLE {nq_prioritization.PRIORITY_TABLE}')
        connection.execute(f'CREATE TABLE {nq_prioritization.PRIORITY_TABLE} (call_key TEXT PRIMARY KEY, campaign_key TEXT NOT NULL DEFAULT \'\')')
        connection.execute("DELETE FROM workspace_state WHERE key IN (?, ?, ?)",
                           (nq.ANALYSIS_CENTER_V24_STATE_KEY, nq.ANALYSIS_CENTER_V25_STATE_KEY, nq.SCHEMA_STATE_KEY))
    nq.ensure_nq_tables(core.repository)
    with core.repository.connection() as connection:
        columns = {str(row[1]) for row in connection.execute(f'PRAGMA table_info({nq_prioritization.PRIORITY_TABLE})')}
    assert {column for column, _definition in nq_prioritization.STRATEGY_COLUMNS} <= columns
    order = [field['key'] for field in nq.list_fields(core.repository)]
    assert order.index('prioritization_queue') == order.index('prioritization_rule') + 1
    assert order.index('expected_gain') == order.index('prioritization_queue') + 1
    priority = next(field for field in nq.list_fields(core.repository) if field['key'] == 'priority')
    assert [option['name'] for option in priority['options']] == [f'P{number}' for number in range(1, 9)]
    # The priorities without a colour take those of the queues and steps; a colour chosen stays.
    assert [option['color'] for option in priority['options']] == ['#123456', *nq_catalog.PRIORITY_COLORS[1:]]


def test_the_default_strategy_fills_the_queues_and_priorities_by_step(strategized):
    rows = priority_rows()
    steps = {step['rule']: step for step in nq_prioritization.strategy_steps(nq_prioritization.CITIES_FIRST)}
    sa_voice = [row for row in rows.values() if row['nr_mode'] == 'SA' and row['service'] in {'voice', 'speech'}]
    assert sa_voice and {row['priority_rule'] for row in sa_voice} <= set(steps)
    for row in rows.values():
        if row['prioritization'] == 'Prioritized':
            step = steps[row['priority_rule']]
            assert (row['suggested_priority'], row['queue'], row['queue_step'], row['queue_list']) == (
                step['priority'], step['queue'], step['number'], step['list'])
        elif row['prioritization'] == 'Track only':
            assert (row['suggested_priority'], row['queue']) == ('P8', '')
    # The CDRs of the workspace are Final CDRs, so list 2 orders the scored calls, and their Expected Gain is their points.
    assert {row['queue_list'] for row in rows.values() if row['queue_list']} == {'final'}
    assert all(row['expected_points'] == pytest.approx(row['impact_points']) for row in rows.values() if row['final_list'])
    data = [row for row in rows.values() if row['service'] == 'data']
    assert any((row['cluster_size'] or 0) >= 2 for row in data)
    # The panel shows the steps and a queue per team, in the order the analysts take its calls.
    client = strategized['client']
    panel = client.post('/api/non-qualified-calls/prioritization', json={'filters': {}}).json()
    assert panel['strategy'] == {'key': 'cities_first', 'label': 'Dual Queue (Cities first)'}
    assert [rule['name'] for rule in panel['rules']][:len(steps)] == list(steps)
    assert [item['name'] for item in panel['queues']] == ['Voice', 'Data'] and not panel['queue']
    voice = panel['queues'][0]
    assert voice['rows'] and [row['position'] for row in voice['rows']] == list(range(1, len(voice['rows']) + 1))
    assert {row['rule'] for row in voice['rows']} <= {rule for rule, step in steps.items() if step['queue'] == 'Voice'}
    ranks = [int(row['priority'][1:]) for row in voice['rows']]
    assert ranks == sorted(ranks)
    fields = {field['source_ref']: field for field in nq.list_fields(core.repository) if field['source'] == 'priority'}
    assert [option['name'] for option in fields['queue']['options']] == ['Voice', 'Data', 'Unique']
    assert [option['name'] for option in fields['priority_rule']['options']] == list(steps)


def test_admins_switch_the_strategy_at_once_and_calls_in_analysis_keep_their_priority(strategized):
    client = strategized['client']
    rows = priority_rows()
    medium = [row for row in rows.values() if row['priority_rule'] == 'Final 4 · SA Voice · Medium impact']
    started = next(row for row in rows.values() if row['prioritization'] == 'Prioritized' and row['suggested_priority'] != 'P1')
    call = client.get(f"/api/non-qualified-calls/calls/{started['call_key']}").json()
    response = client.patch(f"/api/non-qualified-calls/calls/{started['call_key']}",
                            json={'changes': {'status': 'Under Analysis'}, 'version': call['call']['version']})
    assert response.status_code == 200, response.text
    before = started['suggested_priority']
    # Vodafone Proposal gives every SA Voice call P1 at once, in its Unique queue; the call in analysis keeps the Priority it had.
    settings = {**nq_prioritization.load_settings(core.repository), 'strategy': nq_prioritization.VODAFONE_PROPOSAL}
    response = client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': settings})
    assert response.status_code == 200, response.text
    with core.repository.connection() as connection:
        values = dict(connection.execute("SELECT call_key, value FROM nq_analysis_field_values WHERE field_key = 'priority'").fetchall())
    assert values[started['call_key']] == before
    assert medium and all(values[row['call_key']] == 'P1' for row in medium if row['call_key'] != started['call_key'])
    assert {row['queue'] for row in priority_rows().values() if row['queue']} == {'Unique'}
    panel = nq_prioritization.panel(core.repository, {}, 'super')
    assert [item['name'] for item in panel['queues']] == ['Unique']
    log = next(row for row in core.repository.list_logs(limit=50) if row['action'] == 'nq_prioritization_strategy')
    assert json.loads(log['details'])['from'] == 'Dual Queue (Cities first)' and json.loads(log['details'])['to'] == 'Vodafone Proposal'
    # Only admins and super-admins choose it.
    workspace_user('editor', 'user-editor')
    login(client, 'editor', 'editor123')
    settings['strategy'] = 'roads_first'
    assert client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': settings}).status_code == 403


def test_list_1_orders_a_campaign_without_its_final_cdr_by_the_gain_of_the_latest_final_one(client, tmp_path):
    enable_module()
    core.repository.replace_scoring_configuration(load_initial_scoring_configuration())
    # The Final CDRs of Q2 and the Weekly CDRs of Q3, in SA.
    for campaign, stage, seed in (('2026-Q2', 'Final', 21), ('2026-Q3', 'Weekly', 22)):
        for kind, rows in nq_frames(campaign[-2:], campaigns=(campaign,), seed=seed).items():
            add_cdr(tmp_path, f'NetCheck_UK_{stage}_CDR_{kind.title()}_SA_{campaign}.xlsx', kind, rows, 'SA')
    scoring_history.sync_history(core.repository)
    nq.sync_nq_calls(core.repository)
    nq_prioritization.calculate(core.repository)
    rows = priority_rows()
    q2 = [row for row in rows.values() if '2026-q2' in row['campaign_key']]
    q3 = [row for row in rows.values() if '2026-q3' in row['campaign_key']]
    assert q2 and q3
    assert {row['final_list'] for row in q2} == {1} and {row['final_list'] for row in q3} == {0}
    assert {row['queue_list'] for row in q3 if row['queue_list']} == {'campaign'}
    estimated = [row for row in q3 if row['expected_points'] is not None]
    assert estimated and {row['expected_campaign'] for row in estimated} == {'2026-Q2'}
    # A Q3 Voice call is expected to lose what the same failure lost in Q2.
    assert any(row['expected_points'] > 0 for row in estimated)


def test_admins_change_the_priority_and_operators_of_each_step(strategized):
    client = strategized['client']
    settings = nq_prioritization.load_settings(core.repository)
    settings['steps']['cities_first']['final-1'] = {'priority': 'P3', 'operators': ['@focused']}
    # NSA Voice takes no operator: its calls are tracked, as the Tracked step takes the Focused Operators.
    settings['steps']['cities_first']['final-8'] = {'priority': 'P5', 'operators': []}
    settings['steps']['cities_first']['final-2']['conditions'] = ['repeated']
    settings['steps']['cities_first']['campaign-2']['conditions'] = ['consecutive', 'repeated']
    settings['steps']['cities_first']['campaign-2']['joins'] = ['or']
    response = client.put('/api/non-qualified-calls/prioritization/settings', json={'settings': settings})
    assert response.status_code == 200, response.text
    rows = priority_rows()
    high = [row for row in rows.values() if row['priority_rule'] == 'Final 1 · SA Voice · High impact']
    assert high and {row['suggested_priority'] for row in high} == {'P3'}
    assert not any(row['priority_rule'] == 'Final 8 · NSA Voice' for row in rows.values())
    nsa_voice = [row for row in rows.values() if row['nr_mode'] == 'NSA' and row['service'] == 'voice' and row['final_list']]
    # WhatsApp calls are Voice calls too, and keep their step.
    assert nsa_voice and {(row['prioritization'], row['priority_rule']) for row in nsa_voice} <= {
        ('Track only', 'Final 11 · Tracked'), ('Prioritized', 'Final 10 · NSA WhatsApp')}
    assert any(row['prioritization'] == 'Track only' for row in nsa_voice)
    log = next(row for row in core.repository.list_logs(limit=50) if row['action'] == 'nq_prioritization_settings')
    changed = json.loads(log['details'])['steps_changed']
    assert 'Dual Queue (Cities first) · Final 1 · SA Voice · High impact: Priority P1 → P3' in changed
    assert 'Dual Queue (Cities first) · Final 8 · NSA Voice: Operators Focused Operators → No operator' in changed
    assert 'Dual Queue (Cities first) · Final 2 · SA Data · High impact: Condition High impact → Cluster' in changed
    assert any('Campaign 2 · SA Data' in line and 'Condition Hotspots → Consecutive failures OR Cluster' in line for line in changed)
    # A step saved with Consecutive failures and Cluster after the Hotspots took them keeps them.
    assert nq_prioritization.load_settings(core.repository)['steps']['cities_first']['campaign-2']['conditions'] == ['consecutive', 'repeated']
