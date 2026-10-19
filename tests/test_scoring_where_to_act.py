"""Where to act in the Summary of Scoring & GAP Analysis: the plan, the performance owners, the places ranked and the
route that loses most, on the page payload and in the PowerPoint exports."""
from __future__ import annotations

from io import BytesIO

import pytest
from pptx import Presentation

from src.modules.scoring import calculate_scoring
from src.modules.scoring_exports import export_scoring_powerpoint
from src.modules.scoring_points_loss import setting_of
from src.modules.scoring_views import build_scoring_views
from tests.scoring_cdr_synthetic import synthetic_cdr_frames
from tests.scoring_fixtures import scoring_configuration
from tests.test_scoring_exports import TEMPLATE

OWNERS = {'London': 'Vendor A', 'Bristol': 'Vendor A', 'Cardiff': 'Vendor A', 'London to Birmingham': 'Vendor A',
          'Bristol to Cardiff': 'VMO2 South', 'Birmingham': 'Vendor B', 'Manchester': 'Vendor B', 'Leeds': 'Vendor B',
          'Manchester to Leeds': 'Vendor B', 'Glasgow': 'VMO2 North', 'Edinburgh': 'VMO2 North', 'Glasgow to Edinburgh': 'VMO2 North'}
BIG = {'London', 'Birmingham', 'Manchester', 'Glasgow'}


@pytest.fixture(scope='module')
def owned():
    """The synthetic CDRs with their performance owners (Region) and settings, a third of London walked, split by City."""
    frames = synthetic_cdr_frames(campaigns=('2026-Q2', '2026-Q3'), repeats=1)
    for kind, frame in frames.items():
        routes = frame['G_Level_2'] != 'City'
        frame.loc[routes, 'G_Level_2'] = 'Connectionroad'
        frame.loc[(frame['City'] == 'London') & (frame.index % 3 == 0), 'G_Level_1'] = 'Walk'
        frame['Region'] = frame['City'].map(OWNERS)
        frame.loc[routes & (frame.index % 4 == 0), 'Region'] = 'VMO2 South'
        frame['G_Level_3'] = [('Big Cities' if city in BIG else 'Medium Cities') if level == 'City' else 'Roads'
                              for city, level in zip(frame['City'], frame['G_Level_2'])]
    levels = ['Operator', 'City', 'Campaign']
    result = calculate_scoring(frames, levels, configuration=scoring_configuration())
    job = {'id': 1, 'levels': levels, 'baseline_operator': 'EE', 'configuration': result['configuration'],
           'campaigns': ['2026-Q2', '2026-Q3']}
    return job, result, build_scoring_views(job, result, [])['insights']


def test_the_setting_of_a_test():
    assert setting_of('Big Cities', 'City') == 'Cities' and setting_of('Small Cities', 'City') == 'Towns'
    assert setting_of('Highway', 'Connecting Roads') == 'Roads' and setting_of('Train', 'Train') == 'Railways'
    assert setting_of('', 'Connecting Roads') == 'Roads' and setting_of('', '') == ''


def test_the_places_keep_their_owners_settings_and_modes(owned):
    _job, result, _insights = owned
    places = result['points_loss']['areas']['City']
    assert places['London']['setting'] == 'Cities' and places['London']['modes'] == ['Drive', 'Walk']
    assert places['Leeds']['setting'] == 'Towns' and places['Leeds']['region_tests'] == {'Vendor B': places['Leeds']['tests']}
    assert places['Manchester to Leeds']['setting'] == 'Roads' and set(places['Manchester to Leeds']['region_tests']) == {'Vendor B', 'VMO2 South'}


def test_where_to_act_adds_up_the_places_as_national_points(owned):
    _job, _result, insights = owned
    plans = [item for item in insights['act_plans'] if item['environment'] == 'Combined']
    cards = [item for item in insights['owner_cards'] if item['environment'] == 'Combined']
    # One plan and one set of cards per operator, whatever the number of cities.
    assert sorted(item['operator'] for item in plans) == sorted(item['operator'] for item in cards) == sorted(
        {'EE', 'O2 UK', 'Three UK', 'Vodafone UK'})
    plan = next(item for item in plans if item['operator'] == 'Vodafone UK')
    themes = [row['theme'] for row in plan['rows']]
    assert themes[-2:] == ['Sites to review', 'Route stretches'] and 'Voice calls' in themes
    assert plan['headline'].endswith('; then the sites to review, High priority first')
    owners = next(item for item in cards if item['operator'] == 'Vodafone UK')
    # The plan and the owners count the same points (the places by their share of the tests).
    planned = sum(row['points'] for row in plan['rows'] if row['points'])
    assert abs(planned - owners['total']) / planned < .01
    assert {card['owner'] for card in owners['cards']} == {'Vendor A', 'Vendor B', 'VMO2 North', 'VMO2 South'}
    vendor_b = next(card for card in owners['cards'] if card['owner'] == 'Vendor B')
    assert set(vendor_b['settings']) == {'Cities', 'Towns', 'Roads and railways'} and vendor_b['settings']['Towns'] > 0
    assert vendor_b['main'].startswith(('cities', 'towns', 'roads and railways'))


def test_the_places_are_ranked_per_setting_with_their_owners(owned):
    _job, _result, insights = owned
    ranking = next(item for item in insights['place_rankings']
                   if item['environment'] == 'Combined' and item['operator'] == 'Vodafone UK')
    groups = {group['setting']: group for group in ranking['groups']}
    assert list(groups) == ['Cities', 'Towns', 'Roads']
    assert {row['location'] for row in groups['Cities']['rows']} == BIG
    london = next(row for row in groups['Cities']['rows'] if row['location'] == 'London')
    assert london['note'] == '(drive + walk)' and london['owners'] == [{'name': 'Vendor A', 'share': 100.0}]
    assert [row['percent'] for row in groups['Towns']['rows']] == sorted((row['percent'] for row in groups['Towns']['rows']), reverse=True)
    assert groups['Roads']['headline'].endswith('loses the most points') or 'none of the' in groups['Roads']['headline']


def test_the_route_that_loses_most_with_its_owners(owned):
    _job, _result, insights = owned
    focus = next(item for item in insights['route_focus'] if item['environment'] == 'Combined' and item['operator'] == 'Vodafone UK')
    route = focus['routes'][0]
    assert route['points'] == max(item['points'] for item in focus['routes'])
    assert route['top'] and all(row.get('region') for row in route['top'])
    assert abs(sum(owner['points'] for owner in route['owners']) - route['points']) < .05
    assert route['headline'].startswith('Points spread along the route:') and route['review']
    assert all(len(entry) == 4 for entry in route['all_stretches'])
    assert route['landmarks'] and all(place['name'] in OWNERS for place in route['landmarks'])


def test_the_summary_slides_of_where_to_act(owned):
    job, result, _insights = owned
    deck = Presentation(BytesIO(export_scoring_powerpoint(job, result, TEMPLATE, theme='umlaut')))
    titles = [slide.shapes.title.text_frame.paragraphs[0].text for slide in deck.slides if slide.shapes.title is not None]
    for operator in ('Vodafone UK', 'Three UK'):
        assert f'Where to act — {operator}' in titles and f'Where to act, by performance owner — {operator}' in titles
        assert f'Cities ranked — {operator}' in titles and f'Towns ranked — {operator}' in titles
        assert any(title.endswith(f': where the points are — {operator}') for title in titles)
    # The reference is not one of the operators compared by default.
    assert 'Where to act — EE' not in titles
    route = next(slide for slide in deck.slides if slide.shapes.title is not None
                 and slide.shapes.title.text_frame.paragraphs[0].text.endswith(': where the points are — Vodafone UK'))
    names = {shape.name for shape in route.shapes}
    assert {'Scoring Owner Route Map', 'Scoring Route Owners Table', 'Scoring Route Stretches Table', 'Scoring Act Review'} <= names
    without = Presentation(BytesIO(export_scoring_powerpoint(
        {**job, '_operator_choices': {'act-operator': ['EE']}}, result, TEMPLATE)))
    titles = [slide.shapes.title.text_frame.paragraphs[0].text for slide in without.slides if slide.shapes.title is not None]
    # The operator chosen on the page, the reference included.
    assert 'Where to act — EE' in titles and 'Where to act — Vodafone UK' not in titles
    # The route that loses most follows the operators chosen in its own section, whatever that of Where to act.
    own = Presentation(BytesIO(export_scoring_powerpoint(
        {**job, '_operator_choices': {'act-operator': ['EE'], 'route-operators': ['Three UK']}}, result, TEMPLATE)))
    titles = [slide.shapes.title.text_frame.paragraphs[0].text for slide in own.slides if slide.shapes.title is not None]
    assert 'Where to act — EE' in titles
    assert [title for title in titles if ': where the points are — ' in title and title.endswith('Three UK')]
    assert not [title for title in titles if ': where the points are — ' in title and not title.endswith('Three UK')]


def test_where_to_act_of_a_report_follows_the_focus_operators_of_the_operator_maps(owned):
    from src.modules.scoring_exports import export_scoring_report
    from src.modules.scoring_reports import default_scenario, normalize_report_configuration

    job, result, _insights = owned

    def act_titles(roles):
        groups = [{'canonical': name, 'aliases': aliases, 'position': index, 'color': '#365F91', **({'role': role} if role else {})}
                  for index, (name, (aliases, role)) in enumerate(roles.items())]
        entry = {'scenario': normalize_report_configuration({'scenarios': [default_scenario()]})['scenarios'][0], 'job': job, 'result': result}
        deck = Presentation(BytesIO(export_scoring_report([entry], TEMPLATE, groups)))
        titles = [slide.shapes.title.text_frame.paragraphs[0].text for slide in deck.slides if slide.shapes.title is not None]
        return sorted({title.split(' — ', 1)[1] for title in titles if title.startswith('Where to act — ')})

    # The labels with the Focus role, every one of them (also by a name of the CDRs they stand for, shown with its label).
    assert act_titles({'Three UK': ([], 'focus'), 'O2 UK': ([], '')}) == ['Three UK']
    assert act_titles({'Three': (['Three UK'], 'focus'), 'Vodafone UK': ([], 'focus'), 'O2 UK': ([], 'exclude')}) == ['Three', 'Vodafone UK']
    # Without them, or when none of them has results, every operator but the reference.
    assert act_titles({}) == act_titles({'Another operator': ([], 'focus')}) == ['O2 UK', 'Three UK', 'Vodafone UK']


def test_focus_operators_are_names_or_patterns_and_every_operator_without_a_match():
    from src.modules.scoring_exports import focus_matcher

    operators = ['EE', 'VF', 'VF SA', 'VF VoNR', 'Three', '3 UK', 'VMO2']
    focus = focus_matcher(['VF*', 'three*', '3*'], operators)
    assert [name for name in operators if focus(name)] == ['VF', 'VF SA', 'VF VoNR', 'Three', '3 UK']
    # A label (the name of the operator in the Operator Maps) matches too; other characters are taken as they are.
    assert focus('Vodafone', 'VF UK') and not focus('E.E') and focus_matcher(['E.E'], ['E.E'])('E.E')
    # Without Focus operators, or when none of the operators of the results is one of them, every operator.
    assert all(focus_matcher([], operators)(name) for name in operators)
    assert all(focus_matcher(['Orange*'], operators)(name) for name in operators)
